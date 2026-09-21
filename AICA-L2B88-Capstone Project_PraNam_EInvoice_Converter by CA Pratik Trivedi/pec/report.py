"""Validation report: Field | Source | Mapped Value | Target Field | Status | Remarks, plus
items, issues and reconciliation. Written with openpyxl as a NEW workbook (never the template)."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import APP_NAME, APP_VERSION

FILL = {"GREEN": PatternFill("solid", fgColor="C6EFCE"), "AMBER": PatternFill("solid", fgColor="FFEB9C"),
        "RED": PatternFill("solid", fgColor="FFC7CE")}
HEAD = PatternFill("solid", fgColor="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F = "Arial"


def _table(ws, start_row, headers, rows, widths, status_col=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(start_row, c, h)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(name=F, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BOX
            if status_col and c == status_col and v in FILL:
                cell.fill = FILL[v]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 1)


def write_report(path: Path, res, supplier, engine, info: dict, problems: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    o, t = res.options, res.totals
    nic_name, _ = engine.output_names(res)
    ws["A1"] = f"{APP_NAME} v{APP_VERSION} - Validation Report"
    ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = "NIC/GePP input file prepared. This is NOT an e-Invoice: the IRN is generated only by the official NIC/IRIS system."
    ws["A2"].font = Font(name=F, italic=True, size=10, color="C00000")
    rows = [
        ("Generated on", dt.datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Source file", Path(res.source.path).name),
        ("Source sheet", res.source.sheet),
        ("NIC template", Path(engine.template).name),
        ("NIC template SHA-256 (unchanged)", info.get("template_sha256", "")),
        ("Output file", nic_name),
        ("Invoice No.", res.header.get("colDocno", "")),
        ("Invoice Date", res.header.get("colDocdate", "")),
        ("Transaction / Supply Type", f"{o.transaction_type} / {o.supply_type}"),
        ("Supplier GSTIN", supplier.gstin),
        ("Buyer", res.header.get("colBLegalname", "")),
        ("Country code", res.header.get("colCntryCode", "")),
        ("Currency", res.header.get("colForCur", "")),
        ("Exchange rate (INR per 1 FC)", str(o.exchange_rate or "")),
        ("Invoice amount (FC)", str(t.get("amount_fc", ""))),
        ("Total taxable value (INR)", res.header.get("colTotTaxval", "")),
        ("Total IGST (INR)", res.header.get("colTigstval", "")),
        ("Total invoice value (INR)", res.header.get("colTinvoiceval", "")),
        ("Line items", str(len(res.items))),
        ("GST quantity basis / UQC", f"{o.quantity_basis} / {o.uqc}"),
        ("Items sheet: rows written from", f"row {info['Items']['first_row']} ({info['Items']['rows_written']} rows); "
                                           f"{info['Items']['pre_existing_rows']} pre-existing rows in template copy were "
                                           f"{'replaced' if o.write_mode == 'replace' else 'kept'}"),
        ("RED / AMBER issues", f"{res.issues.n('RED')} / {res.issues.n('AMBER')}"),
        ("Post-write verification", "PASSED - macros, ActiveX controls and all other parts byte-identical; values read back"
                                    if not problems else "PROBLEMS: " + " | ".join(problems[:10])),
    ]
    if o.test_note:
        rows.insert(0, ("TEST RUN NOTE", o.test_note))
    for i, (k, v) in enumerate(rows, 4):
        ws.cell(i, 1, k).font = Font(name=F, bold=True, size=10)
        c = ws.cell(i, 2, v)
        c.font = Font(name=F, size=10, color="C00000" if k in ("TEST RUN NOTE",) or (k.startswith("Post") and problems) else "000000",
                      bold=k == "TEST RUN NOTE")
        c.alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110

    fm = wb.create_sheet("Field Mapping")
    _table(fm, 1, ["Field", "Source", "Mapped Value", "Target Field", "Status", "Remarks"],
           [(m.field, m.source, m.value, m.target, m.status, m.remarks) for m in res.field_log],
           [24, 40, 34, 46, 10, 70], status_col=5)

    it = wb.create_sheet("Items")
    by_item = {}
    for i in res.issues:
        if i.item_no:
            lvl = by_item.get(i.item_no, "GREEN")
            by_item[i.item_no] = "RED" if "RED" in (lvl, i.level) else ("AMBER" if "AMBER" in (lvl, i.level) else "GREEN")
    src_rows = {s.sl_no: s.row for s in res.source.items}
    _table(it, 1, ["Sl", "Source row", "Description", "HSN", "Qty", "Unit", "Unit Price (INR)", "Gross (INR)",
                   "Taxable (INR)", "GST %", "IGST (INR)", "Item Total (INR)", "Status"],
           [(int(r["colProdSlno"]), src_rows.get(int(r["colProdSlno"])), r.get("colProddesc"), r.get("colHsn"),
             r.get("colQuantity"), r.get("colUnit"), r.get("colUnitPrice"), r.get("colTotal"), r.get("colAssValue"),
             r.get("colGstrate"), r.get("colIgst"), r.get("colTolitemval"), by_item.get(int(r["colProdSlno"]), "GREEN"))
            for r in res.items],
           [5, 8, 50, 11, 10, 11, 13, 13, 13, 7, 11, 13, 9], status_col=13)

    isx = wb.create_sheet("Issues")
    order = {"RED": 0, "AMBER": 1, "GREEN": 2}
    _table(isx, 1, ["Status", "Item", "Field", "Message"],
           [(i.level, i.item_no or "", i.field, i.message) for i in sorted(res.issues, key=lambda x: order[x.level])],
           [9, 6, 22, 130], status_col=1)

    rc = wb.create_sheet("Reconciliation")
    _table(rc, 1, ["Check", "Value A", "Value B", "Difference", "Status"], res.recon, [70, 18, 18, 14, 10], status_col=5)
    wb.save(path)
