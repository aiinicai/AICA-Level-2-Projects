"""Reads the commercial-invoice workbook WITHOUT assuming fixed rows or columns.

* Header fields are found by their printed labels (profile['fields']).
* The item table is found by scoring rows for known column headings (profile['items']).
* Formula cells: the cached result stored by Excel is used. If Excel never calculated the
  cell (no cached value) the value is flagged - it is never treated as zero.
* Merged cells resolve to their top-left anchor. Hidden rows are read but flagged.
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .logsetup import get_logger
from .models import Issues, Party, SourceInvoice, SourceItem, SourceValue
from .util import clean_text, split_block_lines, to_decimal

log = get_logger()
NUMERIC_ROLES = {"cartons", "pack_size", "packs_per_carton", "rate_per_piece", "rate_per_carton", "rate",
                 "qty", "amount", "discount", "taxable_value", "gst_rate", "igst", "cgst", "sgst", "cess",
                 "net_weight", "gross_weight"}


class SourceReadError(Exception):
    """Raised with a user-readable message when the workbook cannot be processed at all."""


class SheetView:
    def __init__(self, ws_v, ws_f):
        self.ws_v, self.ws_f = ws_v, ws_f
        self.title = ws_v.title
        self._anchor: dict[tuple[int, int], tuple[int, int, int, int]] = {}
        for rng in ws_f.merged_cells.ranges:
            b = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self._anchor[(r, c)] = b
        self.max_row = ws_f.max_row
        self.max_col = ws_f.max_column

    def bounds(self, r, c):
        return self._anchor.get((r, c), (r, c, r, c))

    def ref(self, r, c) -> str:
        return f"{self.title}!{get_column_letter(c)}{r}"

    def get(self, r, c) -> SourceValue:
        ar, ac, _, _ = self.bounds(r, c)
        raw_f = self.ws_f.cell(ar, ac).value
        val = self.ws_v.cell(ar, ac).value
        is_formula = isinstance(raw_f, str) and raw_f.startswith("=")
        sv = SourceValue(val, self.ref(ar, ac))
        if is_formula:
            if val is None:
                sv.formula_uncached = True
                sv.note = "formula without a calculated value - open and save the file in Excel"
            else:
                sv.note = "formula result (cached value)"
        return sv

    def text(self, r, c) -> str:
        return clean_text(self.get(r, c).value)

    def hidden(self, r) -> bool:
        d = self.ws_f.row_dimensions.get(r)
        return bool(d is not None and d.hidden)

    def text_cells(self):
        seen = set()
        for row in self.ws_f.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                b = self.bounds(cell.row, cell.column)
                if (b[0], b[1]) in seen:
                    continue
                seen.add((b[0], b[1]))
                t = clean_text(self.ws_v.cell(b[0], b[1]).value if not isinstance(cell.value, str)
                               or not cell.value.startswith("=") else self.ws_v.cell(b[0], b[1]).value)
                if t:
                    yield b[0], b[1], t


def _match(patterns, text) -> bool:
    return any(re.search(p, text, re.I) for p in patterns)


def r_is_header_fragment(sv, r, roles) -> bool:
    """A PDF/Word column heading often prints over two or three lines ('Kind of' / 'Packages' /
    'GMS'); those stray lines sit below the heading row and are not items."""
    from openpyxl.utils import column_index_from_string as ci
    filled = [(c, sv.text(r, ci(col))) for col, c in [(col, role) for role, col in roles.items()]
              if sv.text(r, ci(col))]
    if not filled or len(filled) > 3:
        return False
    for role, t in filled:
        col = ci(roles[role])
        if t.upper() not in sv.header_text(r - 2, col, span=3).upper():
            return False
    return True


class SourceReader:
    def __init__(self, path: str | Path, profile: dict, currency_codes: set[str] | None = None, ocr_exe: str = ""):
        self.path = Path(path)
        self.profile = profile
        self.currency_codes = currency_codes or set()
        self.ocr_exe = ocr_exe

    # ------------------------------------------------------------------ open
    def _open(self):
        if not self.path.exists():
            raise SourceReadError(f"File not found: {self.path.name}")
        if self.path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise SourceReadError("Only .xlsx and .xlsm invoice files are supported.")
        try:
            wb_f = load_workbook(self.path, data_only=False)
            wb_v = load_workbook(self.path, data_only=True)
        except PermissionError:
            raise SourceReadError("The invoice file is locked. Close it in Excel and try again.")
        except Exception as e:
            log.exception("open failed")
            raise SourceReadError(f"The invoice file could not be read as an Excel workbook ({type(e).__name__}).")
        return wb_f, wb_v

    def _pick_sheet(self, names, include, exclude):
        for pat in include:
            for n in names:
                if re.search(pat, n.strip(), re.I) and not _match(exclude, n):
                    return n
        return None

    # ------------------------------------------------------------------ main
    def read(self) -> SourceInvoice:
        return self.read_all()[0]

    def read_all(self) -> list[SourceInvoice]:
        """A PDF may hold many invoices (one per page); each is returned separately."""
        if not self.path.exists():
            raise SourceReadError(f"File not found: {self.path.name}")
        if self.path.suffix.lower() in {".xlsx", ".xlsm"}:
            return [self._read_excel()]
        from .doc_grids import load_documents
        try:
            grids = load_documents(self.path, self.ocr_exe)
        except RuntimeError as e:
            raise SourceReadError(str(e))
        except ImportError as e:
            name = getattr(e, "name", "") or str(e)
            raise SourceReadError(
                f"A required package is not installed on this computer: '{name}'. "
                "Double-click install_requirements.bat in the tool's folder (or run "
                "'pip install -r requirements.txt'), then try again. Version 2.0 needs opencv-python-headless, "
                "numpy, pypdfium2, pytesseract, pillow, pdfplumber and python-docx.")
        except Exception as e:
            log.exception("document read")
            import traceback
            tb = traceback.extract_tb(e.__traceback__)
            where = f" at {tb[-1].filename.split('/')[-1].split(chr(92))[-1]}:{tb[-1].lineno} in {tb[-1].name}()" if tb else ""
            raise SourceReadError(f"The file could not be read ({type(e).__name__}: {e}){where}. "
                                  "The full details are in logs\\app.log - send that file if the problem repeats.")
        out, failures = [], []
        for i, grid in enumerate(grids, 1):
            inv = SourceInvoice(path=str(self.path), sheet=grid.title)
            try:
                self._fill(grid, inv)
                out.append(inv)
            except SourceReadError as e:
                failures.append(f"{grid.title}: {e}")
        if not out:
            raise SourceReadError(failures[0] if failures else "No invoice could be read from this file.")
        for inv in out:
            for fmsg in failures:
                inv.read_issues.amber("Multi-invoice file", f"A page could not be read - {fmsg}")
        return out

    def _read_excel(self) -> SourceInvoice:
        wb_f, wb_v = self._open()
        names = wb_f.sheetnames
        sheet = self._pick_sheet(names, self.profile["invoice_sheet"], self.profile.get("exclude_sheet", []))
        if not sheet:
            sheet = self._best_sheet(wb_f, wb_v, names)
        sv = SheetView(wb_v[sheet], wb_f[sheet])
        inv = SourceInvoice(path=str(self.path), sheet=sheet)
        self._fill(sv, inv)
        pk = self._pick_sheet(names, self.profile.get("packing_sheet", []), [])
        if pk and pk != sheet:
            inv.packing_sheet = pk
            self._read_packing(SheetView(wb_v[pk], wb_f[pk]), inv)
        inv.signature["sheets"] = names
        wb_f.close(); wb_v.close()
        return inv

    def _best_sheet(self, wb_f, wb_v, names):
        """No sheet called 'Invoice': use the one that actually holds an item table."""
        for n in names:
            try:
                sv = SheetView(wb_v[n], wb_f[n])
                if self._detect_table(sv, self.profile["items"]["columns"], (("description",), ("amount", "taxable_value"))):
                    return n
            except Exception:
                continue
        raise SourceReadError(f"No invoice sheet with a line-item table was found. Sheets: {', '.join(names)}")

    def _fill(self, sv, inv: SourceInvoice):
        fg = getattr(sv, "field_grid", None) or sv
        cells = list(sv.text_cells())
        inv.all_text = [(sv.ref(r, c), t) for r, c, t in cells]
        self._read_fields(sv, cells, inv)
        if fg is not sv:                      # second pass on the other reading of the page
            self._read_fields(fg, list(fg.text_cells()), inv, only_missing=True)
        self._read_items(sv, inv)
        inv.signature.update({
            "invoice_sheet": inv.sheet, "header_row": inv.header_row, "columns": dict(inv.columns),
            "header_labels": {k: v.upper() for k, v in inv.header_labels.items()},
            "field_cells": {k: v.cell for k, v in inv.fields.items() if v.cell},
            "port_of_loading_text": inv.fields.get("port_of_loading", SourceValue()).text,
        })
        for k, v in inv.fields.items():
            if v.formula_uncached:
                inv.read_issues.red(k, f"'{k}' at {v.cell} is a formula with no calculated value. Open the invoice in Excel, "
                                       "let it recalculate, save, and analyse again.")

    # ------------------------------------------------------------------ fields
    def _read_fields(self, sv, cells, inv: SourceInvoice, only_missing: bool = False):
        specs = self.profile["fields"]
        all_labels = [s["label"] for s in specs.values()] + self.profile["items"]["stop_labels"]
        for name, spec in specs.items():
            if only_missing and (name in inv.fields or getattr(inv, name, None) and getattr(inv, name).name):
                continue
            hits = [(r, c, t) for r, c, t in cells if re.search(spec["label"], t, re.I)]
            near = spec.get("near")
            if near and near in inv.fields and inv.fields[near].cell:
                m = re.search(r"([A-Z]+)(\d+)$", inv.fields[near].cell)
                nr = int(m.group(2))
                hits.sort(key=lambda h: abs(h[0] - nr))
            mode = spec["mode"]
            if mode == "block_below":
                party = self._block(sv, hits, all_labels)
                if hasattr(inv, name):
                    setattr(inv, name, party)
                if party.name:
                    inv.fields[name] = SourceValue(" | ".join([party.name] + party.lines), party.cell)
                continue
            val = self._value(sv, hits, spec)
            if val is None and spec.get("cell_of") in inv.fields:
                ref = inv.fields[spec["cell_of"]].cell
                for r, c, t in cells:
                    if sv.ref(r, c) == ref:
                        mm = re.search(spec["cell_regex"], t)
                        if mm:
                            val = SourceValue(mm.group(1), ref, "found in the same cell as the invoice number")
                        break
            if val is not None:
                inv.fields[name] = val

    def _value(self, sv: SheetView, hits, spec) -> SourceValue | None:
        rx, mode = spec.get("regex"), spec["mode"]
        for r, c, text in hits:
            b = sv.bounds(r, c)
            if rx and mode in {"inline", "inline_right", "inline_below"}:
                m = re.search(rx, text, re.I)
                if m and clean_text(m.group(1)):
                    src = sv.get(r, c)
                    return SourceValue(clean_text(m.group(1)), src.cell, "extracted from label text",
                                       src.formula_uncached)
            if mode in {"inline_right", "inline", "inline_below"}:
                cc, taken, acc, first = b[3] + 1, 0, "", None
                while cc <= sv.max_col and taken < 4:
                    cand = sv.get(r, cc)
                    ct = clean_text(cand.value)
                    if not ct:
                        cc += 1
                        continue
                    if self._looks_like_label(ct) and not acc:
                        break
                    if self._looks_like_label(ct):
                        break
                    taken += 1
                    first = first or cand
                    acc = (acc + " " + ct).strip()
                    vp = spec.get("value_pattern")
                    if not vp:
                        return cand
                    if re.fullmatch(vp, acc, re.I):            # e.g. '10 October' + '2022'
                        return SourceValue(acc, first.cell, "read from the cells next to the label")
                    cc = sv.bounds(r, cc)[3] + 1
            if spec.get("gather_column"):
                # the value may start on a row above/below the label (multi-line addresses)
                for cc2 in range(b[3] + 1, min(b[3] + 6, sv.max_col) + 1):
                    for rr in (r - 1, r + 1, r - 2, r + 2):
                        if not 1 <= rr <= sv.max_row:
                            continue
                        t = sv.text(rr, cc2)
                        if t and not self._looks_like_label(t) and not any(sv.text(rr, c) for c in range(1, b[1] + 1)):
                            return self._gather_column(sv, rr, cc2, sv.get(rr, cc2))
            if mode in {"below", "inline_below", "below_or_above"}:
                cand = sv.get(b[2] + 1, b[1])
                ct = clean_text(cand.value)
                if ct and self._looks_like_label(ct):
                    ct = ""
                if spec.get("below_regex") and ct:
                    mm = re.search(spec["below_regex"], ct, re.I)
                    if mm:
                        return SourceValue(clean_text(mm.group(1)), cand.cell, "extracted from cell below label")
                elif ct or cand.formula_uncached:
                    return cand
            if mode == "below_or_above" and b[0] > 1:
                ab = sv.bounds(b[0] - 1, b[1])
                cand = sv.get(ab[0], ab[1])
                ct = clean_text(cand.value)
                if ct and not re.search(spec["label"], ct, re.I) and len(ct) < 80:
                    cand.note = "value printed above its label"
                    return cand
        return None

    def _gather_column(self, sv, row, col, cand) -> SourceValue:
        """Multi-line addresses print on the rows above and below the row holding the label."""
        lines = [clean_text(cand.value)]
        for step in (-1, 1):
            r = row + step
            while 1 <= r <= sv.max_row and abs(r - row) <= 3:
                t = sv.text(r, col)
                if not t or self._looks_like_label(t) or any(sv.text(r, c) for c in range(1, col)):
                    break
                if step == -1:
                    lines.insert(0, t)
                else:
                    lines.append(t)
                r += step
        return SourceValue(", ".join(x for x in lines if x), cand.cell, "address read from the rows around the label")

    def _looks_like_label(self, text: str) -> bool:
        specs = self.profile["fields"]
        return any(re.search(sp["label"], text, re.I) for sp in specs.values())

    def _block(self, sv: SheetView, hits, all_labels) -> Party:
        for r, c, _ in hits:
            b = sv.bounds(r, c)
            rr, lines, first_cell, skipped = b[2] + 1, [], "", 0
            while rr <= sv.max_row and len(lines) < 8:
                cand = sv.get(rr, b[1])
                t = clean_text(cand.value)
                if not t:
                    if not lines and skipped == 0:
                        skipped = 1; rr += 1; continue
                    break
                if _match(all_labels, t):
                    break
                first_cell = first_cell or cand.cell
                lines.extend(split_block_lines(cand.value))
                rr = sv.bounds(rr, b[1])[2] + 1
            email = phone = ""
            kept = []
            for ln in lines:
                em = re.search(r"[A-Za-z0-9+_.\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}", ln)
                if em and re.search(r"E-?MAIL|@", ln, re.I):
                    email = email or em.group(0); continue
                ph = re.match(r"^\s*(TEL|TELE|PHONE|PH|MOB|MOBILE|CONTACT|CELL|FAX)\b[^0-9+]*([+0-9][0-9 ()\-]{5,})", ln, re.I)
                if ph:
                    phone = phone or re.sub(r"\D", "", ph.group(2)); continue
                kept.append(ln)
            lines = kept
            if lines:
                name_is_first = skipped == 0
                if name_is_first:
                    return Party(lines[0], lines[1:], first_cell, email, phone)
                return Party("", lines, first_cell, email, phone)
            if False:
                if name_is_first:
                    return Party(lines[0], lines[1:], first_cell)
                return Party("", lines, first_cell)       # a blank first row means the name is missing
        return Party()

    # ------------------------------------------------------------------ item table
    def _detect_table(self, sv: SheetView, columns: dict, need: tuple):
        best = None
        for r in range(1, sv.max_row + 1):
            roles, labels, taken = {}, {}, set()
            for c in range(1, sv.max_col + 1):
                b = sv.bounds(r, c)
                if (b[0], b[1]) != (r, c):
                    continue
                own = sv.text(r, c)
                texts = [own]
                if getattr(sv, "is_text_grid", False):
                    texts.append(sv.header_text(r, c))
                texts = [t for t in texts if t and len(t) <= 120]
                if not texts:
                    continue
                for role, pats in columns.items():
                    if role in taken:
                        continue
                    hit = next((t for t in texts if _match(pats, t)), None)
                    if hit:
                        roles[role] = get_column_letter(c); labels[role] = own or hit; taken.add(role)
                        break
            enough = len(roles) >= 3 or (("amount" in roles or "taxable_value" in roles) and
                                        re.search(r"DESCRIPTION|PARTICULARS", labels.get("description", ""), re.I))
            if not enough:
                continue
            if all(n in roles for n in need[0]) and any(n in roles for n in need[1]):
                if best is None or len(roles) > len(best[1]):
                    best = (r, roles, labels)
        return best

    def _read_rows(self, sv: SheetView, header_row, roles: dict, issues: Issues, collect_hsn=True):
        from openpyxl.utils import column_index_from_string as ci
        spec = self.profile["items"]
        desc_col = ci(roles["description"])
        other_cols = sorted(ci(c) for r_, c in roles.items() if r_ != "description")
        lo = max([c for c in other_cols if c < desc_col], default=0)
        hi = min([c for c in other_cols if c > desc_col], default=sv.max_col + 1)
        text_grid = getattr(sv, "is_text_grid", False)

        def desc_of(r):
            """On a PDF/Word grid the description often spills into neighbouring columns."""
            if not text_grid:
                return sv.get(r, desc_col)
            parts, cell = [], None
            for c in range(lo + 1, hi):
                t = sv.text(r, c)
                if t:
                    parts.append(t)
                    cell = cell or sv.get(r, c)
            if not parts:
                return sv.get(r, desc_col)
            return SourceValue(" ".join(parts), cell.cell if cell else "", cell.note if cell else "")

        start = sv.bounds(header_row, desc_col)[2] + 1
        rows, blank, stop_row = [], 0, None
        if getattr(sv, "is_text_grid", False):
            while r_is_header_fragment(sv, start, roles):
                start += 1
        r = start
        while r <= sv.max_row:
            dsv = desc_of(r)
            d = clean_text(dsv.value)
            line = " ".join(sv.text(r, c) for c in range(1, sv.max_col + 1) if sv.text(r, c))
            if _match(spec["stop_labels"], d or "") or (rows and _match(spec["stop_labels"], line)):
                stop_row = r
                break
            if sv.bounds(r, desc_col)[0] != r:          # continuation of a vertically merged cell
                r += 1; continue
            vals = {role: sv.get(r, ci(col)) for role, col in roles.items() if role != "description"}
            nonblank = {k: v for k, v in vals.items() if clean_text(v.value) or v.formula_uncached}
            if not d and not nonblank:
                blank += 1
                if blank > spec.get("max_blank_rows", 15):
                    break
                r += 1; continue
            blank = 0
            numeric_present = any(k in NUMERIC_ROLES for k in nonblank)
            money = any(k in {"amount", "taxable_value", "rate", "qty", "cartons"} for k in nonblank)
            if not d and not nonblank:
                blank += 1
                if blank > spec.get("max_blank_rows", 15):
                    break
                r += 1
                continue
            rows.append((r, dsv, vals))
            r += 1
        if getattr(sv, "is_text_grid", False):
            rows = self._stitch(rows, issues, "sl_no" in roles)
        kept = []
        for rr, dsv, vals in rows:
            filled = {k for k, v in vals.items() if clean_text(v.value)}
            money = {"amount", "taxable_value", "rate", "qty", "cartons"} & filled
            if not clean_text(dsv.value):
                if money:                      # a totals / summary line closes the table
                    stop_row = stop_row or rr
                    break
                continue                        # blank or decorative row
            if not money and "hsn" not in filled:
                issues.amber("Line items", f"Row {rr} ('{clean_text(dsv.value)[:40]}') carries no quantity, rate or "
                                           "amount - it was not treated as an item.")
                continue
            kept.append((rr, dsv, vals))
        return kept, stop_row

    @staticmethod
    def _stitch(rows, issues: Issues, by_serial: bool = False):
        """PDF/Word item lines wrap: an amount may print on the next line and a long description
        may run over two or three lines. Fragments that carry no column in common with the row
        above belong to that row."""
        out: list = []
        stitched: list[bool] = []
        if by_serial:
            # a new item starts only where the serial-number column has a value
            for r, dsv, vals in rows:
                has_sl = clean_text(vals.get("sl_no").value) if vals.get("sl_no") else ""
                if out and not has_sl:
                    pr, pdsv, pvals = out[-1]
                    if clean_text(dsv.value):
                        pdsv = SourceValue(f"{clean_text(pdsv.value)} {clean_text(dsv.value)}".strip(),
                                           pdsv.cell or dsv.cell, "text continued on the next line")
                    for k, v in vals.items():
                        if clean_text(v.value) and not clean_text(pvals.get(k).value if pvals.get(k) else ""):
                            pvals[k] = v
                    out[-1] = (pr, pdsv, pvals)
                    continue
                out.append((r, dsv, vals))
            return out
        for r, dsv, vals in rows:
            filled = {k for k, v in vals.items() if clean_text(v.value)}
            if clean_text(dsv.value):
                filled.add("description")
            if not filled:
                continue
            if out:
                pr, pdsv, pvals = out[-1]
                prev = {k for k, v in pvals.items() if clean_text(v.value)}
                if clean_text(pdsv.value):
                    prev.add("description")
                incomplete = "amount" not in prev or stitched[-1]
                wrapped_desc = filled == {"description"} and "description" in prev and incomplete
                if wrapped_desc or (filled.isdisjoint(prev) and "description" not in filled):
                    if clean_text(dsv.value):
                        pdsv = SourceValue(f"{clean_text(pdsv.value)} {clean_text(dsv.value)}".strip(),
                                           pdsv.cell or dsv.cell, "text continued on the next line")
                    for k, v in vals.items():
                        if clean_text(v.value) and k not in prev:
                            pvals[k] = v
                    out[-1] = (pr, pdsv, pvals)
                    stitched[-1] = True
                    continue
            out.append((r, dsv, vals))
            stitched.append(False)
        return out

    def _read_items(self, sv: SheetView, inv: SourceInvoice):
        cols = self.profile["items"]["columns"]
        found = self._detect_table(sv, cols, (("description",), ("amount", "taxable_value")))
        if not found and getattr(sv, "table_rows", None):
            # A scan whose column headings OCR could not read: the rows are handed over as they
            # are so they can be completed on the check screen instead of being lost.
            for n, r in enumerate(sv.table_rows, 1):
                line = " ".join(sv.text(r, c) for c in range(1, sv.max_col + 1) if sv.text(r, c)).strip()
                if line:
                    inv.items.append(SourceItem(n, r, SourceValue(line, sv.ref(r, 1), "row read from the scan"),
                                                SourceValue()))
            inv.columns = {"description": "A"}
            inv.header_row = sv.table_rows[0] if sv.table_rows else 0
            inv.read_issues.red("Item table", f"The column headings on this scan could not be read, so the "
                                              f"{len(inv.items)} table rows are shown as text. Split them into "
                                              "description, HSN, quantity, rate and value on the check screen.")
            return
        if not found:
            raise SourceReadError("The line-item table could not be located (no header row with a "
                                  "Description column and an Amount/Value column). Review the mapping profile.")
        inv.header_row, inv.columns, inv.header_labels = found
        rows, stop_row = self._read_rows(sv, inv.header_row, inv.columns, inv.read_issues)
        after = stop_row or ((rows[-1][0] + 1) if rows else None)
        if after:
            self._read_totals(sv, inv, after)
        for n, (r, dsv, vals) in enumerate(rows, 1):
            hsn_sv = vals.pop("hsn", SourceValue())
            was_num = isinstance(hsn_sv.value, (int, float)) and not isinstance(hsn_sv.value, bool)
            if was_num:
                hsn_sv = SourceValue(str(int(hsn_sv.value)) if float(hsn_sv.value).is_integer()
                                     else str(hsn_sv.value), hsn_sv.cell, "HSN stored as a number in Excel")
            item = SourceItem(n, r, dsv, hsn_sv, was_num, vals, sv.hidden(r))
            if hasattr(sv, "confidence"):
                from openpyxl.utils import column_index_from_string as cix
                weak = []
                for role, col in inv.columns.items():
                    cf = sv.confidence(r, cix(col))
                    if cf < 70:
                        weak.append(f"{role} ({cf:.0f}%)")
                if weak:
                    inv.read_issues.amber("OCR confidence", f"Item {n}: the scan was read with low confidence for "
                                          + ", ".join(weak) + ". Check these values in the item table before generating.", n)
            inv.items.append(item)

        from openpyxl.utils import column_index_from_string as ci
        # currency from the amount / rate column headings, only if it is a real master code
        for role in ("amount", "rate_per_carton", "rate_per_piece", "rate"):
            label = inv.header_labels.get(role, "")
            for tok in re.findall(r"\b[A-Z]{3}\b", label.upper()):
                if tok in self.currency_codes:
                    inv.currency = SourceValue(tok, sv.ref(inv.header_row, ci(inv.columns[role])),
                                               f"read from column heading '{label}'")
                    break
            if inv.currency.value:
                break

        if not inv.currency.value and inv.items:
            sym = {"$": "USD", "US$": "USD", "\u20ac": "EUR", "EURO": "EUR", "EUR": "EUR", "\u00a3": "GBP",
                   "GBP": "GBP", "\u20b9": "INR", "RS": "INR", "INR": "INR", "AED": "AED", "AUD": "AUD",
                   "CAD": "CAD", "SGD": "SGD", "JPY": "JPY", "CHF": "CHF"}
            for role in ("amount", "taxable_value"):
                for itx in inv.items:
                    sv_amt = itx.values.get(role)
                    t = clean_text(sv_amt.value) if sv_amt else ""
                    if not t:
                        continue
                    tok = re.match(r"\s*([^\s\d.,-]+|[A-Za-z]{2,4})", t)
                    code = sym.get(tok.group(1).upper().strip(".")) if tok else None
                    if code and (not self.currency_codes or code in self.currency_codes):
                        inv.currency = SourceValue(code, sv_amt.cell,
                                                   f"currency symbol printed with the amount ('{t[:12]}')")
                        break
                if inv.currency.value:
                    break
            if not inv.currency.value:
                for itx in inv.items:
                    raw = str((itx.values.get("amount") or itx.values.get("taxable_value") or SourceValue()).value or "")
                    if "(cid:" in raw or "\u20b9" in raw:
                        inv.currency = SourceValue("INR", "", "rupee symbol printed with the amount")
                        break

        # printed invoice total (usually a SUM on or just below the TOTAL row)
        if stop_row and "amount" in inv.columns:
            ac = ci(inv.columns["amount"])
            found = []
            for rr in range(stop_row, min(stop_row + 5, sv.max_row) + 1):
                cand = sv.get(rr, ac)
                if cand.formula_uncached or to_decimal(cand.value) is not None:
                    found.append(cand)
            if found:
                inv.source_total = found[0]
                if len(found) > 1:
                    inv.source_grand_total = found[-1]

    def _read_totals(self, sv, inv: SourceInvoice, stop_row: int):
        """Reads the CGST / SGST / IGST / GRAND TOTAL / 'value in INR @ rate' lines that many
        invoices print under the item table instead of as columns."""
        from openpyxl.utils import column_index_from_string as ci
        pats = {"tax_cgst": r"\bCGST\b", "tax_sgst": r"\bSGST\b", "tax_igst": r"\bIGST\b",
                "tax_cess": r"\bCESS\b", "grand_total": r"GRAND\s*TOTAL|^\s*TOTAL\s*$|INVOICE\s*TOTAL",
                "total_taxable": r"TOTAL\s*TAXABLE"}
        for r in range(stop_row, min(stop_row + 14, sv.max_row) + 1):
            line = " ".join(sv.text(r, c) for c in range(1, sv.max_col + 1) if sv.text(r, c)).strip()
            if not line:
                continue
            m = re.search(r"IN\s*INR\s*@\s*([0-9]+(?:\.[0-9]+)?)", line, re.I)
            if m and "exchange_rate" not in inv.fields:
                inv.fields["exchange_rate"] = SourceValue(m.group(1), sv.ref(r, 1), "rate printed with the INR value")
                inr = re.findall(r"(?:RS\.?|INR)\s*\.?\s*(-?[0-9][0-9,]*(?:\.[0-9]+)?)", line, re.I)
                if inr:
                    inv.fields["inr_total"] = SourceValue(inr[-1], sv.ref(r, 1), "INR value printed on the invoice")
            nums = [to_decimal(sv.text(r, c)) for c in range(1, sv.max_col + 1) if sv.text(r, c)]
            nums = [n for n in nums if n is not None]
            for key, pat in pats.items():
                if key not in inv.fields and re.search(pat, line, re.I) and nums:
                    rate = re.search(pat + r"\s*@?\s*([0-9.]+)\s*%", line, re.I)
                    inv.fields[key] = SourceValue(nums[-1], sv.ref(r, 1), line[:60])
                    if rate and key.startswith("tax_"):
                        inv.fields[key + "_rate"] = SourceValue(rate.group(1), sv.ref(r, 1), "rate printed on the tax line")

    def _read_packing(self, sv: SheetView, inv: SourceInvoice):
        cols = {k: v for k, v in self.profile["items"]["columns"].items()
                if k in {"description", "cartons", "pack_size", "packs_per_carton", "net_weight", "gross_weight"}}
        found = self._detect_table(sv, cols, (("description", "net_weight"), ("net_weight",)))
        if not found:
            inv.read_issues.amber("Packing list", f"Sheet '{sv.title}' found but its item table could not be read.")
            return
        hr, roles, _ = found
        rows, _ = self._read_rows(sv, hr, roles, Issues())
        if len(rows) != len(inv.items):
            inv.read_issues.amber("Packing list", f"Packing list has {len(rows)} rows but the invoice has "
                                                  f"{len(inv.items)} items - net weights not linked.")
            return
        for it, (r, dsv, vals) in zip(inv.items, rows):
            it.packing_desc = clean_text(dsv.value)
            nw = vals.get("net_weight")
            it.packing_net_kg = to_decimal(nw.value) if nw else None
