import re
import math
from io import BytesIO
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import streamlit as st

st.set_page_config(page_title="SURL Checker + OCR", page_icon="🧾", layout="wide")
st.title("Search for Unrecorded Liabilities (SURL) Checker")
st.caption("Capstone prototype with multi-format ingestion and optional OCR | Use fictional or approved redacted data only")
st.warning("This tool screens for potential exceptions; it does not determine by itself that a liability is unrecorded. Validate extracted data against the source document, confirm population completeness, inspect source evidence, and follow your firm's audit methodology.")

SUPPORTED = ["csv", "xlsx", "xls", "pdf", "png", "jpg", "jpeg", "tif", "tiff", "bmp", "docx", "txt"]


def _configure_tesseract():
    # Windows installers often do not add Tesseract to PATH; fall back to the default install locations.
    import shutil
    import os
    try:
        import pytesseract
    except ImportError:
        return
    if shutil.which("tesseract"):
        return
    for candidate in [r"C:\Program Files\Tesseract-OCR\tesseract.exe", r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                      os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe")]:
        if os.path.exists(candidate):
            pytesseract.pytesseract.tesseract_cmd = candidate
            return


_configure_tesseract()

with st.sidebar:
    st.header("Engagement settings")
    period_end = st.date_input("Year-end date", value=date(2025, 12, 31))
    search_start = st.date_input("Subsequent search start", value=date(2026, 1, 1))
    search_end = st.date_input("Subsequent search end", value=date(2026, 3, 31))
    overall_mat = st.number_input("Overall materiality (engagement currency)", min_value=0.0, value=250000.0, step=10000.0)
    perf_mat = st.number_input("Performance materiality", min_value=0.0, value=175000.0, step=10000.0)
    tolerable = st.number_input("SURL tolerable misstatement", min_value=0.0, value=100000.0, step=5000.0)
    sig_threshold = st.number_input("Individually significant threshold", min_value=0.0, value=50000.0, step=5000.0)
    coverage_rate = st.slider("Illustrative non-statistical planning coverage rate", 0.05, 1.0, 0.25, 0.05)
    use_ocr = st.checkbox("Enable OCR fallback for scanned PDFs/images", value=True)
    st.header("Workpaper settings")
    entity_name = st.text_input("Entity name (workpaper header)", value="Example Company")
    wp_scope = st.number_input("Subsequent payments scope: amounts greater than", min_value=0.0, value=20000.0, step=1000.0)
    aging_scope = st.number_input("AP aging / unprocessed scope: amounts greater than", min_value=0.0, value=0.0, step=1000.0)
    st.caption("Planning coverage is not an audit standard or statistical confidence level.")

st.subheader("1. Upload your source populations")
st.markdown("Each slot accepts **CSV, Excel, text-based PDF, scanned PDF, image, DOCX or TXT**. Select the correct population for each file. If a PDF contains a table, the app attempts table extraction; if it is scanned or text extraction is poor, it attempts OCR when enabled.")
col1, col2, col3 = st.columns(3)
with col1:
    ap_file = st.file_uploader("Year-end AP aging", type=SUPPORTED, key="ap")
    disb_file = st.file_uploader("Subsequent disbursements / bank register", type=SUPPORTED, key="disb")
with col2:
    subsequent_file = st.file_uploader("Subsequent AP aging (optional)", type=SUPPORTED, key="subap")
    voucher_file = st.file_uploader("Unprocessed invoices / vouchers (optional)", type=SUPPORTED, key="vouchers")
with col3:
    tb_file = st.file_uploader("Trial balance (optional)", type=SUPPORTED, key="tb")
    extra_docs = st.file_uploader("Other invoice / supporting documents (optional; multiple)", type=SUPPORTED, accept_multiple_files=True, key="extra")

st.info("OCR requires the Tesseract OCR engine to be installed on your computer in addition to the Python packages. See README.txt for setup. OCR output may contain errors and must be reviewed.")


def _bytes(file):
    file.seek(0)
    return file.read()


def _clean_headers(df):
    df = df.copy()
    df.columns = [str(c).strip() if str(c).strip() else f"Column_{i+1}" for i, c in enumerate(df.columns)]
    return df


def _ocr_image(image_bytes):
    try:
        from PIL import Image
        import pytesseract
        import io
        image = Image.open(io.BytesIO(image_bytes))
        return pytesseract.image_to_string(image), "OCR"
    except Exception as exc:
        return "", f"OCR unavailable/error: {exc}. Install Tesseract OCR and verify it is on PATH."


def _text_to_invoice_row(text, source_name, page_no=None):
    # Conservative heuristic extraction: keep full text for human review; fields may be blank or wrong.
    lines = [re.sub(r"\s+", " ", x).strip(" |\t") for x in text.splitlines() if x.strip()]
    joined = "\n".join(lines)
    def first(patterns):
        for pattern in patterns:
            m = re.search(pattern, joined, flags=re.I | re.M)
            if m:
                return m.group(1).strip(" :#-\t")
        return ""
    invoice = first([
        r"(?:invoice|inv\.?|bill)\s*(?:no\.?|number|#|ref(?:erence)?)\s*[:#\-]?\s*([A-Z0-9][A-Z0-9\-/_.]+)",
        r"(?:document|voucher)\s*(?:no\.?|number|#)\s*[:#\-]?\s*([A-Z0-9][A-Z0-9\-/_.]+)",
    ])
    date_text = first([
        r"(?:invoice date|bill date|document date|date of invoice)\s*[:#\-]?\s*([0-9]{1,2}[/-][0-9]{1,2}[/-][0-9]{2,4}|[0-9]{1,2}[-/][A-Za-z]{3,9}[-/][0-9]{4}|[0-9]{4}-[0-9]{2}-[0-9]{2})",
        r"(?:invoice date|bill date|document date|date of invoice)\s*[:#\-]?\s*([A-Za-z]{3,9}\s+[0-9]{1,2},?\s+[0-9]{4})",
    ])
    amount = ""
    amount_patterns = [
        r"(?:invoice amount|grand total|invoice total|total due|amount due|net payable|total amount|balance due)\s*[:$€£₹\s]*(?:USD|INR|EUR|GBP|CAD|AUD)?\s*([-+]?\(?\d[\d,]*(?:\.\d{2})?\)?)",
        r"(?:total)\s*[:$€£₹\s]+([-+]?\(?\d[\d,]*(?:\.\d{2})?\)?)",
    ]
    amount_text = first(amount_patterns)
    if amount_text:
        amount = re.sub(r"[^0-9.()-]", "", amount_text)
        if amount.startswith("(") and amount.endswith(")"):
            amount = "-" + amount[1:-1]
    vendor = ""
    # If labels are available, use the line following/containing a vendor label.
    # Prefer labels at the start of a line so disclaimers such as "All supplier names..." are not captured.
    vendor = first([r"^(?:vendor|supplier|billed by|from)(?:\s+name)?\s*[:\-]?\s+([^\n]{2,100})",
                    r"\b(?:vendor|supplier|billed by|from)(?:\s+name)?\s*[:\-]\s*([^\n]{2,100})"])
    if vendor:
        vendor = re.split(r"\s+(?:vendor id|invoice|bill to|ship to|invoice date|date)\b", vendor, maxsplit=1, flags=re.I)[0].strip(" :-")
    if not vendor:
        # Conservative fallback: first non-generic line, retained as a low-confidence guess.
        generic = re.compile(r"invoice|tax invoice|bill to|ship to|date|gst|vat|pan|phone|email|address|total|amount|page|terms|due date", re.I)
        for line in lines[:8]:
            if len(line) >= 3 and not generic.search(line) and not re.fullmatch(r"[\W\d_]+", line):
                vendor = line[:100]
                break
    return {
        "Source File": source_name,
        "Page": page_no if page_no is not None else "",
        "Vendor Name": vendor,
        "Invoice No": invoice,
        "Invoice Date": date_text,
        "Invoice Amount": amount,
        "Extracted Text": joined,
        "Extraction Method": "OCR/text heuristic - verify all fields",
    }


_FORM_LABEL = re.compile(r"^(supplier|vendor|billed by|invoice|bill|amount|currency|date|due|service|receipt|description|total|status|po\b|purchase order|terms)", re.I)


def _is_key_value_table(table):
    # Invoice forms are often laid out as label/value grids; those are not population tables.
    first_cells = [str(row[0] or "").strip() for row in table if row]
    return bool(first_cells) and sum(bool(_FORM_LABEL.match(c)) for c in first_cells) >= max(2, len(first_cells) / 2)


def read_document(file, enable_ocr=True):
    """Return (DataFrame, extraction notes). Never silently treat OCR guesses as verified."""
    if file is None:
        return None, []
    name = file.name
    ext = Path(name).suffix.lower().lstrip(".")
    raw = _bytes(file)
    notes = []
    if ext == "csv":
        try:
            df = pd.read_csv(BytesIO(raw))
        except UnicodeDecodeError:
            df = pd.read_csv(BytesIO(raw), encoding="latin-1")
        return _clean_headers(df), ["CSV read as a table."]
    if ext in ("xlsx", "xls"):
        try:
            df = pd.read_excel(BytesIO(raw), sheet_name=0)
            return _clean_headers(df), ["Excel: first worksheet read. If the relevant table is on another sheet, save/upload that sheet separately."]
        except Exception as exc:
            return pd.DataFrame({"Extraction Error": [str(exc)]}), [f"Excel read failed: {exc}"]
    if ext == "txt":
        text = raw.decode("utf-8", errors="replace")
        return pd.DataFrame([_text_to_invoice_row(text, name)]), ["TXT text converted to one heuristic invoice row; verify fields."]
    if ext == "docx":
        try:
            from docx import Document
            doc = Document(BytesIO(raw))
            rows = []
            for table in doc.tables:
                for row in table.rows:
                    rows.append([cell.text.strip() for cell in row.cells])
            if rows and len(rows[0]) > 1:
                width = max(len(r) for r in rows)
                padded = [r + [""] * (width-len(r)) for r in rows]
                df = pd.DataFrame(padded[1:], columns=padded[0])
                return _clean_headers(df), ["DOCX table extracted. Check the header row and values."]
            text = "\n".join(p.text for p in doc.paragraphs)
            return pd.DataFrame([_text_to_invoice_row(text, name)]), ["DOCX text converted to one heuristic invoice row; verify fields."]
        except Exception as exc:
            return pd.DataFrame({"Extraction Error": [str(exc)]}), [f"DOCX read failed: {exc}"]
    if ext == "pdf":
        frames = []
        extracted_pages = []
        try:
            import pdfplumber
            with pdfplumber.open(BytesIO(raw)) as pdf:
                for i, page in enumerate(pdf.pages, start=1):
                    text = page.extract_text() or ""
                    tables = page.extract_tables() or []
                    if tables:
                        for table in tables:
                            if table and len(table) > 1 and not _is_key_value_table(table):
                                header = [str(x or f"Column_{j+1}").strip() for j, x in enumerate(table[0])]
                                body = [[str(x or "").strip() for x in row] for row in table[1:] if row]
                                width = len(header)
                                body = [r[:width] + [""] * max(0, width-len(r)) for r in body]
                                if body:
                                    frames.append(pd.DataFrame(body, columns=header))
                    if text.strip():
                        extracted_pages.append((i, text))
        except Exception as exc:
            notes.append(f"PDF text/table extraction issue: {exc}")
        if frames:
            # Combine extracted tables. PDF table headers can differ across pages; concat safely.
            table_df = pd.concat(frames, ignore_index=True, sort=False)
            notes.append("PDF table(s) extracted. Check repeated header rows, split columns, and amounts.")
            if extracted_pages:
                notes.append("Text layer was also detected; OCR was not needed for pages with readable text.")
            return _clean_headers(table_df), notes
        if extracted_pages:
            if enable_ocr and sum(len(t.strip()) for _, t in extracted_pages) < 40:
                extracted_pages = []
            else:
                rows = [_text_to_invoice_row(t, name, page_no=i) for i, t in extracted_pages]
                notes.append("PDF text extracted; invoice fields are heuristic guesses. Review the extracted text.")
                return pd.DataFrame(rows), notes
        if enable_ocr:
            try:
                import pypdfium2 as pdfium
                pdf = pdfium.PdfDocument(BytesIO(raw))
                rows = []
                for i in range(len(pdf)):
                    page = pdf[i]
                    bitmap = page.render(scale=2.5)
                    pil_image = bitmap.to_pil()
                    try:
                        import pytesseract
                        text = pytesseract.image_to_string(pil_image)
                        if text.strip():
                            rows.append(_text_to_invoice_row(text, name, page_no=i+1))
                    finally:
                        pil_image.close()
                if rows:
                    notes.append("Scanned PDF processed with OCR. Verify every extracted field against the page image; OCR is not guaranteed accurate.")
                    return pd.DataFrame(rows), notes
                notes.append("OCR completed but no readable text was detected.")
            except Exception as exc:
                notes.append(f"Scanned PDF OCR failed: {exc}. Confirm Tesseract is installed and available on PATH.")
        return pd.DataFrame({"Source File": [name], "Extracted Text": [""]}), notes + ["No table or readable text was extracted from this PDF."]
    if ext in ("png", "jpg", "jpeg", "tif", "tiff", "bmp"):
        if not enable_ocr:
            return pd.DataFrame({"Source File": [name], "Extracted Text": [""]}), ["Image received, but OCR is disabled."]
        text, method = _ocr_image(raw)
        if text.strip():
            return pd.DataFrame([_text_to_invoice_row(text, name)]), ["Image OCR performed. Verify every field against the original image."]
        return pd.DataFrame({"Source File": [name], "Extracted Text": [text]}), [method]
    return None, [f"Unsupported format: {ext}"]


def norm(s):
    return "".join(ch.lower() for ch in str(s) if ch.isalnum())


def find_col(df, candidates):
    lookup = {norm(c): c for c in df.columns}
    for cand in candidates:
        if norm(cand) in lookup:
            return lookup[norm(cand)]
    for cand in candidates:
        for k, v in lookup.items():
            if norm(cand) in k or k in norm(cand):
                return v
    return None


def standardize(df, kind):
    if df is None:
        return None
    df = df.copy()
    mappings = {
        "ap": {
            "vendor_id": ["Vendor ID", "Vendor Code", "Supplier ID"],
            "vendor": ["Vendor Name", "Vendor", "Supplier Name", "Supplier", "Billed By"],
            "invoice": ["Invoice No.", "Invoice No", "Invoice Number", "Invoice", "Invoice Ref"],
            "invoice_date": ["Invoice Date", "Bill Date", "Document Date"],
            "service_date": ["Receipt / Service Date", "Service Date", "Goods Receipt Date", "Receipt Date"],
            "invoice_amount": ["Invoice Amount", "Gross Invoice Amount", "Amount", "Total Due", "Invoice Total"],
            "open_balance": ["Open Balance YE", "Open Balance", "Outstanding Amount", "Balance", "Amount Due"],
        },
        "disb": {
            "payment_id": ["Payment ID", "Disbursement ID", "Reference ID", "Transaction ID"],
            "payment_date": ["Payment Date", "Disbursement Date", "Date", "Transaction Date", "Value Date"],
            "vendor_id": ["Vendor ID", "Vendor Code", "Supplier ID"],
            "vendor": ["Vendor Name", "Vendor", "Supplier Name", "Supplier", "Payee", "Beneficiary"],
            "invoice": ["Invoice Ref", "Invoice No", "Invoice Number", "Reference", "Narration", "Description"],
            "amount": ["Payment Amount", "Amount", "Disbursement Amount", "Debit", "Paid Amount", "Transaction Amount"],
            "description": ["Description", "Narration", "Notes", "Source Description", "Particulars"],
        },
        "subap": {
            "vendor_id": ["Vendor ID", "Vendor Code", "Supplier ID"],
            "vendor": ["Vendor Name", "Vendor", "Supplier Name", "Supplier"],
            "invoice": ["Invoice No.", "Invoice No", "Invoice Number", "Invoice"],
            "invoice_date": ["Invoice Date", "Bill Date", "Document Date"],
            "service_date": ["Receipt / Service Date", "Service Date", "Goods Receipt Date", "Receipt Date"],
            "invoice_amount": ["Invoice Amount", "Amount", "Total Due"],
            "open_balance": ["Open Balance 31-Mar", "Open Balance", "Outstanding Amount", "Balance"],
        },
        "vouchers": {
            "voucher_id": ["Voucher ID", "Voucher Number", "Voucher No"],
            "vendor_id": ["Vendor ID", "Vendor Code", "Supplier ID"],
            "vendor": ["Vendor Name", "Vendor", "Supplier Name", "Supplier", "Billed By"],
            "invoice": ["Invoice No.", "Invoice No", "Invoice Number", "Invoice", "Invoice Ref"],
            "date_received": ["Date Received", "Received Date", "Invoice Received Date"],
            "amount": ["Amount", "Invoice Amount", "Gross Amount", "Total Due", "Invoice Total"],
            "service_date": ["Receipt / Service Date", "Service Date", "Goods Receipt Date", "Receipt Date"],
            "description": ["Description", "Notes", "Narration", "Particulars"],
        },
    }
    if kind not in mappings:
        return df
    # Merge headings that differ only in punctuation (e.g. "Invoice No." from a file and "Invoice No" from OCR rows).
    for c in list(df.columns):
        for other in list(df.columns):
            if other != c and c in df.columns and other in df.columns and norm(other) == norm(c):
                df[c] = df[c].combine_first(df[other])
                df = df.drop(columns=[other])
    result = pd.DataFrame(index=df.index)
    required = {"ap": ["vendor", "invoice", "open_balance"], "disb": ["payment_date", "vendor", "amount"], "subap": ["vendor", "invoice", "open_balance"], "vouchers": ["vendor", "amount"]}[kind]
    missing = []
    for target, candidates in mappings[kind].items():
        col = find_col(df, candidates)
        result[target] = df[col] if col is not None else np.nan
        if col is None and target in required:
            missing.append(target)
    if missing:
        st.error(f"{kind.upper()} file is missing required fields: {', '.join(missing)}. The app needs a table with recognizable column headings. For scanned PDFs, review OCR output and use a structured CSV/Excel file where possible.")
        return None
    for col in result.columns:
        if "date" in col and not pd.api.types.is_datetime64_any_dtype(result[col]):
            # ISO dates (e.g. CSV exports) must not be read day-first, or 2026-01-08 becomes 1 August.
            iso = pd.to_datetime(result[col], errors="coerce", format="ISO8601")
            result[col] = iso.fillna(pd.to_datetime(result[col], errors="coerce", dayfirst=True, format="mixed"))
    for col in ["amount", "invoice_amount", "open_balance"]:
        if col in result.columns:
            result[col] = pd.to_numeric(result[col].astype(str).str.replace(r"[^0-9.()\-]", "", regex=True).str.replace(r"^\((.*)\)$", r"-\1", regex=True), errors="coerce")
            result[col] = result[col].fillna(0)
    for col in ["vendor", "invoice", "vendor_id", "description"]:
        if col in result.columns:
            result[col] = result[col].fillna("").astype(str).str.strip()
    vendor_key = result["vendor"].astype(str) if "vendor" in result else pd.Series("", index=result.index)
    invoice_key = result["invoice"].astype(str) if "invoice" in result else pd.Series("", index=result.index)
    result["match_key"] = vendor_key.str.upper().str.replace(r"[^A-Z0-9]", "", regex=True) + "|" + invoice_key.str.upper().str.replace(r"[^A-Z0-9]", "", regex=True)
    return result


def display_extract(label, file):
    if file is None:
        return None
    with st.expander(f"Extraction preview — {label}: {file.name}", expanded=False):
        df, notes = read_document(file, enable_ocr=use_ocr)
        for note in notes:
            st.caption(note)
        if df is not None:
            st.dataframe(df.head(100), use_container_width=True)
            st.caption("Preview is limited to the first 100 rows. Validate the full source and all extracted values before relying on the result.")
        return df

def write_workpaper(book, window, invoice_info, ap_keys, subap, vouchers):
    """Add a 'SURL Workpaper' sheet laid out like the firm's AP - SURL template (Error/Conclusion left for the auditor)."""
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    ws = book.create_sheet("SURL Workpaper", 0)
    thin = Side(style="thin")
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    for col, width in zip("ABCDEFGHIJKL", [4, 11, 14, 12, 28, 14, 12, 15, 16, 18, 18, 55]):
        ws.column_dimensions[col].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    def put(r, c, value, bold=False, fmt=None):
        if isinstance(value, float) and np.isnan(value) or value is pd.NaT:
            value = None
        if isinstance(value, pd.Timestamp):
            value = value.to_pydatetime()
        cell = ws.cell(r, c, value)
        if bold:
            cell.font = Font(bold=True)
        if fmt:
            cell.number_format = fmt
        return cell

    def section(r, heading, procedure, scope):
        put(r, 1, heading, bold=True).font = Font(bold=True, size=12)
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=12)
        put(r, 2, procedure).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 48
        r += 2
        for c, value in scope:
            put(r, c, value, bold=isinstance(value, str),
                fmt="#,##0.00" if isinstance(value, float) else "mm/dd/yyyy" if isinstance(value, date) else None)
        return r + 2

    def table(r, headers, groups):
        for c, h in enumerate(headers, start=2):
            cell = put(r, c, h, bold=True)
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.row_dimensions[r].height = 32
        r += 1
        for label, rows in groups:
            if label:
                put(r, 2, label, bold=True)
                r += 1
            for i, values in enumerate(rows + [[]] * max(0, 5 - len(rows)), start=1):
                put(r, 2, i)
                for c, value in enumerate(values, start=3):
                    fmt = "mm/dd/yyyy" if isinstance(value, (pd.Timestamp, date)) else "#,##0.00" if isinstance(value, (int, float, np.number)) and not isinstance(value, bool) else None
                    put(r, c, value, fmt=fmt).alignment = Alignment(wrap_text=c == 12, vertical="top")
                r += 1
            r += 1
        return r + 1

    def period_note(key, svc):
        if key in ap_keys:
            return "Included in YE AP aging — agree to AP ledger and period-end recording"
        if pd.isna(svc):
            return "Not in YE AP aging; service/receipt date unknown — obtain delivery or service evidence"
        if svc.date() <= period_end:
            return "Not in YE AP aging; service/receipt on or before YE — potential unrecorded liability, investigate"
        return "Not in YE AP aging; service/receipt after YE — likely post-period, verify"

    r = 1
    put(r, 1, f"AUDIT Template - {entity_name}").font = Font(bold=True, size=14)
    put(r + 1, 1, "Accounts Payable - SURL", bold=True)
    put(r + 2, 1, period_end, bold=True, fmt="mm/dd/yyyy").alignment = Alignment(horizontal="left")
    put(r + 3, 1, "Prepared from SURL Checker screening output. Error and Conclusion columns are for the auditor; the tool note is a screening indicator, not an audit conclusion.").font = Font(italic=True, size=9)
    r += 5

    # Section 1 - checks / wires issued subsequent to year-end
    r = section(r, "Testing of Checks / Wires Issued Subsequent to YE",
                "Obtain and examine supporting detail for material disbursements after the balance-sheet date and determine whether the goods or services on the paid invoices were received on or before the balance-sheet date. If so, determine whether the liability is properly recorded at year end.",
                [(8, "Scope:  Amounts greater than"), (10, float(wp_scope)), (11, "from dates between"), (12, f"{search_start:%m/%d/%Y} and {search_end:%m/%d/%Y}")])

    def pay_row(p):
        info = invoice_info.get(p.match_key, {})
        return [p.payment_id, p.payment_date, p.vendor, p.invoice, info.get("invoice_date"), info.get("invoice_amount", p.amount),
                p.Service_receipt_date, "", "", f"{p.Risk_flag}: {p.Screening_note}"]
    ordered = window.sort_values("amount", ascending=False)
    isi = [pay_row(p) for p in ordered[ordered.amount > wp_scope].itertuples()]
    extra = [pay_row(p) for p in ordered[(ordered.amount <= wp_scope) & (ordered.Risk_flag != "Routine")].itertuples()]
    r = table(r, ["Selection #", "Check/Wire Number", "Date Paid", "Vendor", "Invoice Number", "Invoice Date", "Invoice Amount", "Delivery or Service Date", "Error", "Conclusion", "Tool screening note (not a conclusion)"],
              [("ISI Items", isi), ("Additional selection below scope to incorporate a level of unpredictability (screening-flagged items shown; auditor to add others)", extra)])

    # Section 2 - AP aging as of the SURL date
    r = section(r, "Testing of Accounts Payable Aging as of SURL Date",
                "As of the final date the SURL is extended through, obtain that day's trade accounts payable aging. If the goods or services were received on or before the balance-sheet date, determine whether the liability is included in the accounts payable listing (or the listing of accrued liabilities). Document the source and selection criteria for items tested below.",
                [(8, "Scope:  Amounts greater than"), (10, float(aging_scope)), (11, "as of"), (12, f"{search_end:%m/%d/%Y} (last day through which SURL is updated)")])
    aging_rows = []
    if subap is not None:
        for s in subap[subap.open_balance > aging_scope].sort_values("open_balance", ascending=False).itertuples():
            aging_rows.append(["n/a", "n/a", s.vendor, s.invoice, s.invoice_date, s.invoice_amount, "", "",
                               f"Open balance {s.open_balance:,.2f}; service/receipt {s.service_date:%d-%b-%Y}. " if pd.notna(s.service_date) else f"Open balance {s.open_balance:,.2f}. "])
            aging_rows[-1][-1] += period_note(s.match_key, s.service_date)
    else:
        aging_rows.append(["", "", "Subsequent AP aging not uploaded", "", "", "", "", "", ""])
    r = table(r, ["Selection #", "Check/Wire Number", "Date Paid", "Vendor", "Invoice Number", "Invoice Date", "Invoice Amount", "Error", "Conclusion", "", "Tool screening note (not a conclusion)"],
              [("", [row[:8] + [None, row[8]] for row in aging_rows])])

    # Section 3 - unprocessed / un-vouchered invoices as of the SURL date
    r = section(r, "Testing of Unprocessed / Un-vouchered Invoices as of SURL Date",
                "As of the final date the SURL is extended through, obtain the listing of unprocessed / un-vouchered invoices. If the goods or services were received on or before the balance-sheet date, determine whether the liability is included in the accounts payable listing (or the listing of accrued liabilities). Document the source and selection criteria for items tested below.",
                [(8, "Scope:  Amounts greater than"), (10, float(aging_scope)), (11, "as of"), (12, f"{search_end:%m/%d/%Y}")])
    voucher_rows = []
    if vouchers is not None:
        for v in vouchers[vouchers.amount > aging_scope].sort_values("amount", ascending=False).itertuples():
            voucher_rows.append([v.voucher_id, v.date_received, v.vendor, v.invoice, invoice_info.get(v.match_key, {}).get("invoice_date"),
                                 v.amount, v.service_date, "", "", period_note(v.match_key, v.service_date)])
    else:
        voucher_rows.append(["", "", "Unprocessed invoice listing not uploaded", "", "", "", "", "", "", ""])
    r = table(r, ["Selection #", "Voucher Number", "Date Received", "Vendor", "Invoice Number", "Invoice Date", "Invoice Amount", "Delivery or Service Date", "Error", "Conclusion", "Tool screening note (not a conclusion)"],
              [("", voucher_rows)])

    for label in ["Prepared by / date:", "Reviewed by / date:"]:
        put(r, 2, label, bold=True)
        r += 2


# Extract all uploaded files once; show a preview to make OCR/extraction visible to the user.
files = {"ap": ap_file, "disb": disb_file, "subap": subsequent_file, "vouchers": voucher_file, "tb": tb_file}
raw_data = {}
for key, label in [("ap", "Year-end AP aging"), ("disb", "Subsequent disbursements"), ("subap", "Subsequent AP aging"), ("vouchers", "Unprocessed invoices / vouchers"), ("tb", "Trial balance")]:
    raw_data[key] = display_extract(label, files[key]) if files[key] else None
extra_data = []
if extra_docs:
    st.subheader("Other supporting documents")
    for f in extra_docs:
        df = display_extract("Supporting document", f)
        if df is not None:
            extra_data.append((f.name, df))

# Optional bridge: reviewer can correct OCR suggestions before adding them to the voucher population.
reviewed_ocr_vouchers = None
if extra_data:
    candidate_frames = []
    for filename, frame in extra_data:
        if all(c in frame.columns for c in ["Vendor Name", "Invoice No", "Invoice Amount"]):
            tmp = pd.DataFrame({
                "Source File": filename,
                "Vendor Name": frame["Vendor Name"],
                "Invoice No": frame["Invoice No"],
                "Invoice Date": frame["Invoice Date"] if "Invoice Date" in frame.columns else "",
                "Invoice Amount": frame["Invoice Amount"],
                "Extracted Text": frame["Extracted Text"] if "Extracted Text" in frame.columns else "",
            })
            candidate_frames.append(tmp)
    if candidate_frames:
        st.subheader("Review OCR-extracted invoice candidates")
        st.caption("Correct the vendor, invoice number and amount against each original document before selecting the option below. OCR results are unverified until you perform that review.")
        candidates = pd.concat(candidate_frames, ignore_index=True, sort=False)
        reviewed = st.data_editor(candidates, use_container_width=True, num_rows="dynamic", key="reviewed_ocr_candidates")
        include_ocr = st.checkbox("Include these reviewed OCR invoice candidates in the unprocessed-invoice screening population", value=False, key="include_ocr_vouchers")
        if include_ocr:
            reviewed_ocr_vouchers = pd.DataFrame({
                "Vendor Name": reviewed.get("Vendor Name", ""),
                "Invoice No": reviewed.get("Invoice No", ""),
                "Amount": reviewed.get("Invoice Amount", ""),
                "Description": "Reviewed OCR invoice candidate; verify service/receipt period and ledger",
                "Source File": reviewed.get("Source File", ""),
            })
            st.warning("OCR-derived candidates are now included as unprocessed-invoice candidates. Confirm the original document, service period, receipt/acceptance evidence and ledger before reaching any conclusion.")

if ap_file and disb_file:
    ap = standardize(raw_data["ap"], "ap")
    disb = standardize(raw_data["disb"], "disb")
    subap = standardize(raw_data["subap"], "subap") if raw_data["subap"] is not None else None
    voucher_raw = raw_data["vouchers"].copy() if raw_data["vouchers"] is not None else None
    if reviewed_ocr_vouchers is not None:
        voucher_raw = pd.concat([voucher_raw, reviewed_ocr_vouchers], ignore_index=True, sort=False) if voucher_raw is not None else reviewed_ocr_vouchers
    vouchers = standardize(voucher_raw, "vouchers") if voucher_raw is not None else None
    tb = raw_data["tb"]
    if ap is not None and disb is not None:
        st.subheader("2. Population overview")
        if "payment_date" in disb:
            window = disb[(disb.payment_date.dt.date >= search_start) & (disb.payment_date.dt.date <= search_end)].copy()
        else:
            window = disb.copy()
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("YE AP rows", f"{len(ap):,}")
        m2.metric("YE AP open balance", f"{ap.open_balance.sum():,.2f}")
        m3.metric("Payments in search window", f"{len(window):,}")
        m4.metric("Disbursement value in window", f"{window.amount.sum():,.2f}")
        ap_keys = set(ap.match_key.dropna())
        v_keys = set(vouchers.match_key.dropna()) if vouchers is not None else set()
        ap_vendor_inv = ap[["match_key", "open_balance"]].groupby("match_key", as_index=False).open_balance.sum()
        ap_map = dict(zip(ap_vendor_inv.match_key, ap_vendor_inv.open_balance))
        # Service/receipt dates from complementary populations support the period (cut-off) assessment.
        service_map = {}
        for src in [subap, vouchers, ap]:
            if src is not None and "service_date" in src:
                service_map.update(src.dropna(subset=["service_date"]).set_index("match_key").service_date.to_dict())
        if "payment_date" in window:
            window["days_after_year_end"] = (window.payment_date - pd.Timestamp(period_end)).dt.days
        else:
            window["days_after_year_end"] = np.nan
        has_keys = (window.vendor.str.strip() != "") & (window.invoice.str.strip() != "")
        window["YE_AP_match"] = np.where(~has_keys, "Insufficient info", np.where(window.match_key.isin(ap_keys), "Yes", "No"))
        window["Unprocessed_voucher_match"] = window.match_key.isin(v_keys).map({True: "Yes", False: "No"})
        window["YE_AP_open_balance_match"] = pd.to_numeric(window.match_key.map(ap_map), errors="coerce").fillna(0)
        window["Amount_vs_AP_balance_difference"] = window.amount - window.YE_AP_open_balance_match
        window["Service_receipt_date"] = pd.to_datetime(window.match_key.map(service_map), errors="coerce")
        window["Service_period"] = np.where(window.Service_receipt_date.isna(), "Unknown",
                                            np.where(window.Service_receipt_date.dt.date <= period_end, "On/before YE", "After YE"))
        window["Date_bucket"] = pd.cut(window.days_after_year_end, bins=[-999, 30, 60, 9999], labels=["0-30 days", "31-60 days", "61+ days"])
        post_period = (window.YE_AP_match == "No") & (window.Service_period == "After YE")
        window["Risk_flag"] = np.where(post_period & (window.amount < sig_threshold), "Post-period (verify)",
                                       np.where((window.YE_AP_match != "Yes") | (window.Unprocessed_voucher_match == "Yes") | (window.amount >= sig_threshold), "Review", "Routine"))
        window["Suggested_selection"] = np.where(window.Risk_flag == "Review", "SELECT", "")
        window["Screening_note"] = np.select(
            [window.YE_AP_match == "Insufficient info", post_period, (window.YE_AP_match == "No") & (window.Unprocessed_voucher_match == "Yes"), (window.YE_AP_match == "No"), (window.amount >= sig_threshold), (window.YE_AP_match == "Yes")],
            ["Insufficient vendor/invoice information to match — obtain invoice and supporting documents", "Not in YE AP aging; service/receipt date is after year-end — likely post-period, verify service period and contrary evidence", "Not in YE AP aging; also in unprocessed voucher listing — inspect service period and ledger", "Not matched to YE AP aging — inspect invoice, receipt/service evidence and ledger", "Above significant-item threshold — consider specific-item testing", "Matched to YE AP aging — trace payment and confirm period-end recording"],
            default="Review source evidence and audit rationale")
        st.subheader("3. Suggested selections")
        st.caption("Risk-based screening only. Add auditor-directed or random items if required. No claim of statistical representativeness is made.")
        st.dataframe(window, use_container_width=True)
        selected = window[window.Suggested_selection == "SELECT"].copy()
        population_n = len(window)
        planning_count = min(population_n, max(4, math.ceil(population_n * coverage_rate))) if population_n else 0
        st.subheader("4. Sampling planning helper")
        a, b, c = st.columns(3)
        a.metric("Population count", f"{population_n}")
        b.metric("Illustrative planning count", f"{planning_count}")
        c.metric("Risk-flagged items", f"{len(selected)}")
        st.info("This count is a user-adjustable, non-statistical planning aid only. There is no universal formula that derives a defensible audit sample size from materiality and population count alone. Engagement team judgment is required.")
        st.subheader("5. Complementary population review")
        if subap is not None:
            st.write(f"Subsequent AP aging rows: **{len(subap)}** | open balance: **{subap.open_balance.sum():,.2f}**")
            st.dataframe(subap.head(100), use_container_width=True)
        else:
            st.warning("Subsequent AP aging not uploaded. Consider including it as complementary evidence.")
        if vouchers is not None:
            st.write(f"Unprocessed invoice/voucher rows: **{len(vouchers)}** | amount total: **{vouchers.amount.sum():,.2f}**")
            st.dataframe(vouchers.head(100), use_container_width=True)
        else:
            st.warning("Unprocessed invoice/voucher population not uploaded. A payments-only search may miss unpaid/unprocessed liabilities.")
        if tb is not None:
            st.write(f"Trial balance rows: **{len(tb)}**")
            st.dataframe(tb.head(100), use_container_width=True)
        else:
            st.warning("Trial balance not uploaded. Tie AP control and accrual accounts to the ledger and investigate relevant expense accounts.")
        if extra_data:
            st.subheader("6. Other supporting-document extracts")
            st.write(f"Extracted {len(extra_data)} additional document(s). These are shown for manual review and are not automatically included in the core match unless mapped into a structured population.")
        def make_report():
            buff = BytesIO()
            with pd.ExcelWriter(buff, engine="openpyxl") as writer:
                window.to_excel(writer, index=False, sheet_name="Payment Screening")
                selected.to_excel(writer, index=False, sheet_name="Suggested Sample")
                raw_data["ap"].to_excel(writer, index=False, sheet_name="YE AP Aging Source")
                raw_data["disb"].to_excel(writer, index=False, sheet_name="Disbursements Source")
                if raw_data["subap"] is not None: raw_data["subap"].to_excel(writer, index=False, sheet_name="Subsequent AP Source")
                if voucher_raw is not None: voucher_raw.to_excel(writer, index=False, sheet_name="Unprocessed Source")
                if tb is not None: tb.to_excel(writer, index=False, sheet_name="Trial Balance Source")
                for i, (filename, frame) in enumerate(extra_data[:10], start=1):
                    safe_name = re.sub(r"[^A-Za-z0-9]", "", Path(filename).stem)[:20] or f"Extra{i}"
                    frame.to_excel(writer, index=False, sheet_name=f"Extra{i}_{safe_name}"[:31])
                summary = pd.DataFrame([
                    ["Year-end date", str(period_end)], ["Search window start", str(search_start)], ["Search window end", str(search_end)],
                    ["Overall materiality", overall_mat], ["Performance materiality", perf_mat], ["SURL tolerable misstatement", tolerable],
                    ["Significant-item threshold", sig_threshold], ["YE AP population count", len(ap)], ["YE AP open balance", ap.open_balance.sum()],
                    ["Disbursement population in window", len(window)], ["Disbursement value in window", window.amount.sum()],
                    ["Suggested selected count", len(selected)], ["Suggested selected value", selected.amount.sum()],
                    ["OCR caution", "OCR and heuristic field extraction can be inaccurate. Compare extracted values to original documents."],
                    ["Tool limitation", "Risk-based screening only; auditor validation required."],
                ], columns=["Field", "Value"])
                summary.to_excel(writer, index=False, sheet_name="Summary")
                invoice_info = {}
                for src in [vouchers, subap, ap]:  # later sources (YE AP) take precedence
                    if src is None:
                        continue
                    for row in src.itertuples():
                        info = invoice_info.setdefault(row.match_key, {})
                        if pd.notna(getattr(row, "invoice_date", pd.NaT)):
                            info["invoice_date"] = row.invoice_date
                        amount = getattr(row, "invoice_amount", getattr(row, "amount", 0))
                        if amount:
                            info["invoice_amount"] = amount
                write_workpaper(writer.book, window, invoice_info, ap_keys, subap, vouchers)
            buff.seek(0)
            return buff
        st.download_button("Download SURL screening report (.xlsx)", data=make_report(), file_name="SURL_Screening_Report.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
else:
    st.info("Upload at least year-end AP aging and subsequent disbursement register to run screening. You can upload structured files or documents; however, the core populations need usable columns. The ZIP includes sample Excel populations and four fictional invoice test cases.")

st.divider()
st.subheader("Supported formats and important limits")
st.markdown("- **Structured data:** CSV, XLSX, XLS. Excel reads the first worksheet only.\n- **Documents:** text PDFs, scanned PDFs, PNG/JPG/TIFF/BMP images, DOCX and TXT. PDF tables are attempted first; scanned files use OCR if configured.\n- **OCR fields:** vendor, invoice number, date and amount are heuristic extraction only. The app also preserves extracted text for review. Layouts vary, and multi-invoice/complex documents may not parse correctly.\n- **Not yet supported:** password-protected files, handwritten invoices, image-only tables requiring complex layout reconstruction, and every possible proprietary file type. Convert these to a supported format or enter a reviewed structured extract.\n- **Privacy:** use fictional or approved redacted data only; check your firm's data handling policy before uploading client documents to any tool.")
