"""Draft Rule 11(g) reporting text from the findings, for the auditor to review, edit and paste into the audit report.

Inputs: the batch-check result (checks.Result, from an uploaded or pulled extract), and/or the monitor's Rule 11(g)
register (rule11g.build_register), plus a few judgements only the auditor can supply (Inputs).

It never decides for the auditor. It proposes a conclusion with reasons, and gives the alternative wordings so the
auditor can pick. Wording follows the structure of the ICAI Implementation Guide illustrations (revised 2024); clause
numbering in the report differs by firm format, so cross-references are left in [square brackets].
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

import pandas as pd

import checks

DISCLAIMER = ("DRAFT for the auditor's professional judgement. Not to be pasted into the final report unreviewed. "
              "Align the wording and clause references with the ICAI Implementation Guide on Reporting under Rule 11(g) "
              "and your firm's report format.")


@dataclass
class Inputs:
    company: str = "[Company name]"
    fy_end: str = "March 31, 20XX"
    software: str = "[accounting software, e.g. TallyPrime]"
    has_feature: str = "Yes"            # Yes / No
    db_level: str = "Yes"               # Yes / No / Not applicable
    sooc_report: str = "Not applicable"  # Obtained / Not obtained / Not applicable (third-party software or hosting)
    mgmt_explanation: str = "Awaited"   # Satisfactory / Not satisfactory / Awaited
    preserved: str = "Yes"              # Yes / No / Evidence awaited
    report_preservation: bool = True    # preservation sentence applies from the year the rule requires it
    manual_notes: str = ""


@dataclass
class Draft:
    conclusion: str = ""            # unmodified / modified / adverse / pending
    why: list = field(default_factory=list)
    caution: list = field(default_factory=list)
    paras: dict = field(default_factory=dict)       # name -> text
    open_items: list = field(default_factory=list)


def _facts(res, reg, inp: Inputs) -> dict:
    f = {"a1_n": 0, "a1_amt": 0.0, "gaps": [], "high": 0, "generic": 0, "b3": 0, "reg_attention": [], "reg_missing": []}
    if res is not None:
        ex = res.exceptions
        a1 = ex[ex.Test == "A1"]
        f["a1_n"], f["a1_amt"] = len(a1), float(a1.Amount.fillna(0).sum())
        f["gaps"] = [g for g in res.gaps if g["vouchers_without_trail"]] or list(res.gaps)
        s = res.summary.set_index("Test")
        f["high"] = int(s.loc[["B1", "B3", "C1", "C2", "D2"], "High"].sum())
        f["generic"] = int(s.loc["D2", "Exceptions"])
        f["b3"] = int(s.loc["B3", "Exceptions"])
        f["indicator"] = checks.indicator(res)[0]
    if reg is not None and len(reg):
        f["reg_attention"] = list(reg[reg.Status == "Attention"].Ref)
        f["reg_missing"] = list(reg[reg.Status == "Not evidenced"].Ref)
    return f


def build(res=None, reg=None, inp: Inputs | None = None) -> Draft:
    inp = inp or Inputs()
    f = _facts(res, reg, inp)
    d = Draft()
    problems = []        # facts that justify an exception ("except that ...")
    tamper = []

    # ---- classify
    if inp.has_feature == "No":
        d.conclusion = "adverse"
        d.why.append("You recorded that the accounting software has no audit-trail (edit log) feature.")
    else:
        if f["a1_n"]:
            problems.append(f"the audit trail did not record the creation of {f['a1_n']} voucher(s) aggregating "
                            f"{checks.inr(f['a1_amt'])} that appear in the voucher register")
        for g in f["gaps"]:
            problems.append(f"no audit-trail activity was recorded from {g['start']:%d %B %Y} to {g['end']:%d %B %Y} "
                            f"({g['working_days']} working days)")
        if "R3" in f["reg_attention"]:
            problems.append("the monitoring record shows the audit trail could not be evidenced as operating for part of the year")
        if "R5" in f["reg_attention"]:
            problems.append("changes detected in the books were not all found in the software's own Edit Log")
        if inp.db_level == "No":
            problems.append("the audit trail feature was not enabled at the database level to log direct data changes")
        if inp.sooc_report == "Not obtained":
            problems.append("an independent service-auditor's report covering the audit trail controls of the third-party "
                            "software/hosting provider was not available")
        if f["high"] or f["generic"]:
            tamper.append(f"{f['high']} high-risk change(s), of which {f['generic']} were made under generic/shared user IDs")
        if inp.mgmt_explanation == "Not satisfactory" and (f["high"] or f["b3"]):
            problems.append("instances were noted where records were altered or deleted after the books were finalised "
                            "without a satisfactory explanation, indicating the audit trail may have been tampered with")

        d.conclusion = "modified" if problems else "unmodified"
        if problems:
            d.why.append("Exceptions found: " + "; ".join(problems) + ".")
        else:
            d.why.append("The tests found no completeness gap and no exception that justifies a modification on the data available.")
        if tamper and inp.mgmt_explanation in ("Awaited", "") and not problems:
            d.conclusion = "pending"
            d.why.append(f"{tamper[0]}. Management's explanation is awaited, so the 'not tampered' sentence cannot be concluded yet.")
        elif tamper and inp.mgmt_explanation in ("Awaited", ""):
            d.why.append(f"{tamper[0]}. Management's explanation is awaited, so the 'not tampered' sentence stays open even though the exceptions above already justify an exception.")
        elif tamper and inp.mgmt_explanation == "Satisfactory":
            d.why.append(f"{tamper[0]}; management's explanations were recorded as satisfactory, so they do not by themselves modify the wording. "
                         "Keep the explanations and your corroboration on file.")
    if f["reg_missing"]:
        d.open_items.append("Not evidenced in the monitor register: " + ", ".join(f["reg_missing"]))
    if inp.preserved == "Evidence awaited":
        d.open_items.append("Preservation of the audit trail (retention) is not yet evidenced.")
        if d.conclusion == "unmodified":
            d.conclusion = "pending"
            d.why.append("Preservation evidence is awaited.")
    if inp.preserved == "No":
        problems.append("the audit trail has not been preserved by the Company as per the statutory requirements for record retention")
        d.conclusion = "modified" if d.conclusion != "adverse" else "adverse"
    if f.get("indicator") == "Green":
        d.caution.append("A Green indicator only means these data tests found nothing. It does not replace your ITGC, configuration, "
                         "access and preservation evidence.")
    if not inp.report_preservation:
        d.caution.append("The preservation sentence is left out because you indicated it is not yet applicable for this year.")
    d.caution.append("Detection of a change is not proof of tampering. Conclude only after management's explanation and your corroboration.")

    # ---- wording
    base = (f"Based on our examination which included test checks, the Company has used {inp.software} for maintaining its books "
            "of account for the financial year ended " + inp.fy_end + " which has a feature of recording audit trail (edit log) "
            "facility and the same has operated throughout the year for all relevant transactions recorded in the software")
    tamp = "Further, during the course of our audit we did not come across any instance of audit trail feature being tampered with."
    pres = (" Additionally, the audit trail has been preserved by the Company as per the statutory requirements for record retention."
            if inp.report_preservation else "")
    unmod = base + ". " + tamp + pres

    parts = [p for p in problems if "tampered" not in p]
    tampered = [p for p in problems if "tampered" in p]
    mod = base
    if parts:
        mod += ", except that " + "; and ".join(parts)
    mod += ". "
    mod += (("Further, " + "; ".join(tampered) + ".") if tampered else tamp)
    if not inp.report_preservation:
        pass
    elif inp.preserved == "No":
        mod += " However, " + [p for p in problems if "preserved" in p][0] + "."
    else:
        mod += pres
    adv = (f"Based on our examination, the Company has used {inp.software} for maintaining its books of account for the financial "
           f"year ended {inp.fy_end} which does not have a feature of recording audit trail (edit log) facility. Accordingly, we are "
           "unable to comment on whether an audit trail operated throughout the year, was not tampered with, or has been preserved.")

    d.paras["A. Rule 11(g) paragraph: unmodified"] = unmod
    d.paras["B. Rule 11(g) paragraph: with exception (modified)"] = mod
    d.paras["C. Rule 11(g) paragraph: software has no audit-trail feature"] = adv
    d.paras["_suggested"] = {"unmodified": "A. Rule 11(g) paragraph: unmodified",
                             "modified": "B. Rule 11(g) paragraph: with exception (modified)",
                             "adverse": "C. Rule 11(g) paragraph: software has no audit-trail feature",
                             "pending": "B. Rule 11(g) paragraph: with exception (modified)"}[d.conclusion]

    # ---- basis for qualification (only meaningful when an exception exists)
    if problems:
        basis = ("Basis for Qualified Opinion\n"
                 f"As described in [Note/paragraph reference] and in our report on Other Legal and Regulatory Requirements, "
                 f"{'; '.join(problems)}. "
                 "These matters affect [the reliability of the books of account / our ability to rely on the accounting software] and "
                 "our opinion on [clause (b) of paragraph 2(A), proper books of account], is modified accordingly. "
                 "We conducted our audit in accordance with the Standards on Auditing specified under section 143(10) of the Act. "
                 "[Add management's explanation and the possible effect on the financial statements, if quantified.]\n\n"
                 "Consider whether the matter is pervasive; if so, an adverse opinion on the other-requirements clause may be warranted. "
                 "Whether the Rule 11(g) exception also modifies your opinion on the financial statements depends on the effect on "
                 "your audit evidence: that is your judgement under SA 705.")
    else:
        basis = "Not required: no exception identified on the information entered. Re-check once the open items are cleared."
    d.paras["D. Basis for qualification (use only if you modify)"] = basis

    # ---- emphasis / other matter
    eom = []
    if inp.sooc_report == "Obtained":
        eom.append("Other Matter: We have relied on the independent service auditor's report (SOC 2 / SAE 3402) of "
                   "[service provider] for the operation of audit-trail controls of [software/hosting] for the period [dates]. "
                   "Our opinion is not modified in respect of this matter.")
    if tamper and inp.mgmt_explanation == "Satisfactory":
        eom.append("Emphasis of Matter (if you judge it fundamental to users' understanding): We draw attention to [Note X], which "
                   "describes the changes made to the books of account after they were finalised, the management's explanation and "
                   "the approvals obtained. Our opinion is not modified in respect of this matter.")
    if not eom:
        eom.append("No Emphasis of Matter or Other Matter paragraph is indicated by the information entered. An Emphasis of Matter "
                   "paragraph cannot replace the Rule 11(g) statement or a required modification.")
    d.paras["E. Emphasis of Matter / Other Matter (optional)"] = "\n\n".join(eom)

    # ---- management letter
    ml = ["To those charged with governance: observations on the accounting software and audit trail (ITGC)", ""]
    n = 1
    def item(t):
        nonlocal n
        ml.append(f"{n}. {t}")
        n += 1
    if f["a1_n"]:
        item(f"{f['a1_n']} voucher(s) aggregating {checks.inr(f['a1_amt'])} have no corresponding 'created' entry in the audit trail. "
             "Recommendation: confirm that the audit trail cannot be switched off or bypassed, and explain how these were posted.")
    for g in f["gaps"]:
        item(f"No audit-trail activity from {g['start']:%d %b %Y} to {g['end']:%d %b %Y}. Recommendation: confirm the software "
             "was running with the audit trail enabled, and keep configuration evidence.")
    if f["generic"]:
        item(f"{f['generic']} alteration/deletion exception(s) under generic or shared user IDs. Recommendation: individual user IDs, "
             "remove shared and default accounts, review privileged access at least annually.")
    if f["high"]:
        item("High-risk changes after finalisation, back-dating, amount and date changes were noted for explanation. Recommendation: "
             "lock the books after finalisation (period/voucher lock), require approval for post-closure changes, review the Edit Log monthly.")
    if inp.db_level == "No":
        item("The audit trail does not cover direct database-level changes. Recommendation: enable database-level logging or restrict "
             "direct database access.")
    item("Back-up and retention: daily back-up of the books and audit trail on a server physically located in India (Rule 3(5)) and "
         "retention for eight financial years (s.128(5)). Recommendation: document the back-up location and test restores.")
    item("Management should evidence its own review of the audit trail, and obtain a service-auditor's report where software or hosting "
         "is provided by a third party.")
    d.paras["F. Management letter / ITGC observations"] = "\n".join(ml)
    if inp.manual_notes.strip():
        d.paras["G. Your notes"] = inp.manual_notes.strip()
    return d


def to_docx(d: Draft, company: str) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading(f"Draft Rule 11(g) reporting text: {company}", 1)
    doc.add_paragraph(DISCLAIMER).italic = True
    doc.add_heading("Suggested conclusion: " + d.conclusion.upper(), 2)
    for w in d.why:
        doc.add_paragraph(w, style="List Bullet")
    for c in d.caution:
        doc.add_paragraph("Caution: " + c, style="List Bullet")
    for k, v in d.paras.items():
        if k.startswith("_"):
            continue
        doc.add_heading(k + ("  [suggested]" if k == d.paras["_suggested"] else ""), 3)
        for para in str(v).split("\n"):
            doc.add_paragraph(para)
    if d.open_items:
        doc.add_heading("Open items", 3)
        for o in d.open_items:
            doc.add_paragraph(o, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
