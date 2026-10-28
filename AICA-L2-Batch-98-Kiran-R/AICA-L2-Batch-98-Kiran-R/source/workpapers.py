"""Suggested answers and working papers for the guided Rule 11(g) engagement.

suggest()  -> for every control / question: a suggested answer, the reason, and whether it comes from the data or needs an inquiry.
paper()    -> a working paper for one control: objective, procedure, evidence tables taken from the loaded data, result, conclusion.
Nothing here changes the engagement; the auditor confirms each suggestion in the app.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime

import pandas as pd

import guided as G

GENERIC = ("admin", "support", "tally", "guest", "test", "default")
CLOUD = re.compile(r"cloud|saas|hosted|aws|azure|gcp|vendor|third|remote server|tally on cloud|tallyprime server", re.I)


@dataclass
class Ctx:
    res: object = None            # checks.Result
    log: pd.DataFrame = None      # prepared voucher log
    reg: pd.DataFrame = None      # prepared register
    masters: pd.DataFrame = None
    raw: pd.DataFrame = None      # Tally's own Edit Log lines (with Version)
    fy_start: date = date(2025, 4, 1)
    fy_end: date = date(2026, 3, 31)
    company: str = ""
    source: str = ""              # "Pull from Tally" / "Upload files" / "Sample data (fictional)"


def _n(df) -> int:
    return 0 if df is None else len(df)


def _sample(src: str) -> bool:
    return str(src).startswith("Sample")


# --------------------------------------------------------------------------- evidence tables
def month_table(c: Ctx) -> pd.DataFrame:
    rows = []
    reg_d = pd.to_datetime(c.reg["VoucherDate"], errors="coerce") if _n(c.reg) else pd.Series(dtype="datetime64[ns]")
    log_t = pd.to_datetime(c.log["LogTime"], errors="coerce") if _n(c.log) else pd.Series(dtype="datetime64[ns]")
    for i in range(12):
        m0 = pd.Timestamp(c.fy_start) + pd.DateOffset(months=i)
        m1 = m0 + pd.DateOffset(months=1)
        v = int(((reg_d >= m0) & (reg_d < m1)).sum())
        e = int(((log_t >= m0) & (log_t < m1)).sum())
        rows.append({"Month": m0.strftime("%b-%Y"), "Vouchers dated in the month": v, "Edit Log entries made in the month": e,
                     "Observation": "No audit-trail entries made in this month" if e == 0 else ""})
    return pd.DataFrame(rows)


def completeness_table(c: Ctx) -> pd.DataFrame:
    rows = []
    src = c.raw if _n(c.raw) else c.log
    for col, what in [("User", "who"), ("LogTime", "date and time"), ("Action", "created / altered / deleted"),
                      ("Version", "version number"), ("Field", "field changed"), ("OldValue", "old value"), ("NewValue", "new value")]:
        if src is None or col not in src.columns:
            rows.append({"Attribute": what, "Column": col, "Rows": _n(src), "Filled": 0, "Filled %": 0.0, "Note": "column not in the loaded Edit Log"})
            continue
        s = src[col]
        filled = int(s.notna().sum() - (s.astype(str).str.strip() == "").sum()) if len(s) else 0
        filled = max(filled, 0)
        rows.append({"Attribute": what, "Column": col, "Rows": len(s), "Filled": filled,
                     "Filled %": round(100 * filled / len(s), 1) if len(s) else 0.0, "Note": ""})
    return pd.DataFrame(rows)


def a1_table(c: Ctx) -> pd.DataFrame:
    if c.res is None or c.res.exceptions is None:
        return pd.DataFrame()
    x = c.res.exceptions
    return x[x.Test == "A1"][[k for k in ["VoucherType", "VoucherNo", "VoucherDate", "Detail"] if k in x.columns]].reset_index(drop=True)


def version_table(c: Ctx) -> pd.DataFrame:
    if not _n(c.raw) or "Version" not in c.raw.columns:
        return pd.DataFrame()
    v = pd.to_numeric(c.raw["Version"], errors="coerce")
    g = c.raw.assign(_v=v).groupby("MasterID")["_v"].max()
    t = g.value_counts().sort_index().rename_axis("Highest version of the voucher").reset_index(name="Vouchers")
    return t


def coverage_by_type(c: Ctx) -> pd.DataFrame:
    if not _n(c.reg):
        return pd.DataFrame()
    r = c.reg.groupby("VoucherType").size().rename("In register")
    if _n(c.log):
        cr = c.log[c.log.Action == "Created"].groupby("VoucherType").size().rename("With a Created entry")
        r = pd.concat([r, cr], axis=1).fillna(0).astype(int)
    r["Difference"] = r.get("With a Created entry", 0) - r["In register"]
    return r.reset_index()


def company_activity(c: Ctx) -> pd.DataFrame:
    m = c.masters
    if not _n(m):
        return pd.DataFrame()
    return m[m.Object == "Company"] if "Object" in m else pd.DataFrame()


def config_changes(c: Ctx) -> pd.DataFrame:
    ca = company_activity(c)
    if not _n(ca):
        return pd.DataFrame()
    pat = re.compile(r"edit.?log|audit|security|feature|password|user|backup|restore|split|repair", re.I)
    txt = ca.astype(str).agg(" ".join, axis=1)
    return ca[txt.map(lambda s: bool(pat.search(s)))]


def users_table(c: Ctx) -> pd.DataFrame:
    if c.res is None or c.res.users is None:
        return pd.DataFrame()
    u = c.res.users.copy()
    u["Looks generic / shared"] = u["User"].astype(str).str.lower().map(lambda s: any(g in s for g in GENERIC))
    return u


def tally_sample(c: Ctx, n=25) -> pd.DataFrame:
    if not _n(c.log):
        return pd.DataFrame()
    cols = [k for k in ["LogTime", "User", "Action", "VoucherType", "VoucherNo", "VoucherDate", "Amount"] if k in c.log.columns]
    s = c.log.sample(min(n, len(c.log)), random_state=7)[cols].sort_values("LogTime")
    s["Agreed to Tally screen (✓)"] = ""
    return s.reset_index(drop=True)


def blank_checklist(items: list[str]) -> pd.DataFrame:
    return pd.DataFrame({"Evidence to obtain": items, "Obtained (Y/N)": "", "Reference / file": "", "Remarks": ""})


# --------------------------------------------------------------------------- suggestions
def suggest(e: G.Engagement, c: Ctx) -> dict:
    """{'controls': {id: (result, note, basis)}, 'app': {key: (value, why)}, 'prep': {key: (value, why)}}.
    basis = 'data' (computed from the loaded Edit Log) or 'inquiry' (needs management evidence; result is only a starting point)."""
    out = {"controls": {}, "app": {}, "prep": {}}
    real = not _sample(c.source)
    has_log = _n(c.log) > 0
    S = out["controls"]

    # applicability
    nm = (c.company or e.company or "").lower()
    if re.search(r"\bllp\b|partnership|trust\b|society|proprietor", nm):
        out["app"]["entity"] = (G.APPLICABILITY_QS[0][2][2], "the name looks like an LLP, firm, trust or society")
    elif re.search(r"pvt|private|ltd|limited|company|corporation", nm):
        out["app"]["entity"] = (G.APPLICABILITY_QS[0][2][0], "the name ends like a company (Pvt/Ltd/Limited)")
    if c.source and real:
        out["app"]["mode"] = (G.APPLICABILITY_QS[1][2][0], "the vouchers were read from accounting software (TallyPrime / uploaded export)")
    out["app"]["fs"] = ("Standalone", "default; change to Consolidated only for a parent reporting on group accounts")

    # C1 feature present
    if has_log:
        t = pd.to_datetime(c.log["LogTime"], errors="coerce")
        S["C1"] = ("Effective", f"Edit Log entries exist: {len(c.log):,} entries from {t.min():%d-%b-%Y} to {t.max():%d-%b-%Y}, each with user and date-time.", "data")
    else:
        S["C1"] = ("Not tested", "No Edit Log loaded yet (Input tab).", "data")

    # C2 cannot be disabled / C7 configuration changes
    ca = company_activity(c)
    cc = config_changes(c)
    if _n(ca):
        if _n(cc):
            S["C2"] = ("Exception", f"Company Activity shows {len(cc)} configuration-type event(s); read them in the working paper and confirm none switched the Edit Log off.", "data")
        else:
            S["C2"] = ("Effective", f"Company Activity ({len(ca)} events) shows no change to the Edit Log or security settings.", "data")
        S["C7"] = ("Not tested", f"{len(cc)} configuration-type event(s) in Company Activity to be matched with management's approvals.", "inquiry")
    else:
        S["C2"] = ("Not tested", "Company Activity is not loaded (optional TDL pull). Check which Tally users can change company features or security.", "inquiry")
        S["C7"] = ("Not tested", "Needs management's change-approval records.", "inquiry")

    # C3 operated throughout the year
    if has_log and c.res is not None:
        gaps = [g for g in c.res.gaps if g.get("vouchers_without_trail")]
        mt = month_table(c)
        empty = mt[(mt["Edit Log entries made in the month"] == 0) & (mt["Vouchers dated in the month"] > 0)]
        if gaps:
            S["C3"] = ("Exception", f"{len(gaps)} run(s) of working days with no audit-trail activity while vouchers exist (test A2).", "data")
        elif len(empty):
            S["C3"] = ("Not tested", f"No logging gap by test A2, but no entries were made in: {', '.join(empty['Month'])}. Ask when the trail was enabled.", "data")
        else:
            S["C3"] = ("Effective", "Test A2 found no run of working days without trail activity; entries exist in every month.", "data")
    else:
        S["C3"] = ("Not tested", "Edit Log not loaded.", "data")

    # C4 all transactions covered
    if has_log and _n(c.reg):
        a1 = len(a1_table(c))
        multi = len(version_table(c)) > 0 and (version_table(c).iloc[:, 0] > 1).any()
        if a1:
            S["C4"] = ("Exception", f"{a1} voucher(s) in the register have no 'Created' entry in the Edit Log (test A1).", "data")
        else:
            S["C4"] = ("Effective", "Every voucher in the register has a Created entry" + ("; altered vouchers show every version." if multi else "."), "data")
    else:
        S["C4"] = ("Not tested", "Needs both the voucher register and the Edit Log.", "data")

    # C5 who / when / what
    if has_log:
        ct = completeness_table(c).set_index("Column")
        who, when = ct.loc["User", "Filled %"], ct.loc["LogTime", "Filled %"]
        what = ct.loc["OldValue", "Filled %"] if "OldValue" in ct.index else 0
        if min(who, when) < 100:
            S["C5"] = ("Exception", f"User filled {who}%, date-time {when}%: some entries lack who or when.", "data")
        elif what >= 99:
            S["C5"] = ("Effective", "User, date-time and old/new values are complete.", "data")
        else:
            S["C5"] = ("Not tested", f"User and date-time are 100% complete, but the export has no old/new values ({what}%). "
                       "Take a sample from Tally's 'Differences Between Edit Log Versions' screen (see the working paper).", "data")
    else:
        S["C5"] = ("Not tested", "Edit Log not loaded.", "data")

    # C6 unique user IDs
    if has_log and c.res is not None:
        u = users_table(c)
        gen = u[u["Looks generic / shared"]]
        if _n(gen):
            S["C6"] = ("Exception", f"Generic / shared user ID(s) in the Edit Log: {', '.join(gen['User'].astype(str))}.", "data")
        else:
            S["C6"] = ("Effective", f"{len(u)} user ID(s) in the Edit Log, none looks generic. Match them to named staff.", "data")
    else:
        S["C6"] = ("Not tested", "Edit Log not loaded.", "data")

    for cid, why in [("C8", "Folder and back-up permissions need inspection on the client's PC or server."),
                     ("C9", "Needs management's confirmation that no tool writes to Tally data directly."),
                     ("C10", "Needs evidence of management's periodic review of the Edit Log."),
                     ("C11", "Needs the Tally release history and any migration note.")]:
        S[cid] = ("Not tested", why, "inquiry")

    inv = (e.inventory or [{}])[0]
    host = " ".join(str(v) for v in inv.values())
    if inv.get("Software") and not CLOUD.search(host) and "in-house" in str(inv.get("In-house / outsourced", "")).lower():
        S["C12"] = ("Not applicable", "The inventory shows TallyPrime in-house, on the company's own premises (confirm in step 2).", "inquiry")
    else:
        S["C12"] = ("Not tested", "Check whether the software is hosted or run by a third party; if so obtain the SOC report.", "inquiry")

    # preservation: never evidenced from the data alone
    for k, q, _ in G.PRESERVATION_QS:
        out["prep"][k] = ("Not yet evidenced", "needs the back-up schedule, server location and a retrieval test; nothing in the Edit Log shows it")
    return out


# --------------------------------------------------------------------------- working papers
def _hdr(e: G.Engagement, c: Ctx, ref, area, test, icai, how):
    return [("Company", c.company or e.company or "-"), ("Financial year", f"{c.fy_start:%d-%b-%Y} to {c.fy_end:%d-%b-%Y}"),
            ("Reference", f"{ref} · {area}"), ("Test", test), ("ICAI reference", icai), ("How to evidence in TallyPrime", how),
            ("Data source", c.source or "-"), ("Prepared on", f"{date.today():%d-%b-%Y}")]


PROCEDURES = {
    "C1": ["Open TallyPrime → F3 Company → Alter (or Gateway → F11 Features) and note that Edit Log is enabled; take a dated screenshot.",
           "Open Gateway → Display More Reports → Edit Log Summary and confirm it lists vouchers with versions, users and date-times.",
           "Agree the extract below with that screen for a few entries."],
    "C2": ["List the Tally users and their security levels (Gateway → Alt+K Company → Users and Passwords).",
           "Identify who can alter company features or security; confirm those roles are restricted.",
           "Read the Company Activity events below for any change to the Edit Log or security settings."],
    "C3": ["Review the month-by-month table: every month the books were written should have Edit Log entries.",
           "For each run of working days without activity (test A2), obtain management's explanation.",
           "Ask management for the date the Edit Log was first enabled and agree it with the first entry below."],
    "C4": ["Agree the count of vouchers by type in the register with those having a Created entry in the Edit Log.",
           "Investigate every voucher in the exceptions list (no Created entry).",
           "Confirm that altered vouchers show every version, not only the latest (version table)."],
    "C5": ["Check that user, date-time and action are filled for all entries (completeness table).",
           "Tick-and-tie the sample below to Tally's Edit Log screens.",
           "For the 'data changed', open Tally's 'Differences Between Edit Log Versions' on a sample of altered vouchers and record old and new values."],
    "C6": ["Agree each user ID below to a named person and a role.",
           "Investigate generic or shared IDs that created, altered or deleted vouchers.",
           "Inspect Tally security: are passwords enforced and shared logins avoided?"],
    "C7": ["Read the configuration-type events from Company Activity below.",
           "Obtain management's approval for each, or confirm that there were none."],
    "C8": ["Inspect operating-system permissions on the Tally data folder and on the back-up location.",
           "Confirm who can delete or copy Tally data files and whether access is logged."],
    "C9": ["Ask management whether any ODBC write-back, third-party utility or file replacement is used to change Tally data.",
           "Inspect the Tally data folder for unexpected tools; record the written answer."],
    "C10": ["Obtain evidence that management reviews the Edit Log periodically (monthly sign-off, exception follow-up).",
            "Read the test results below; they are what a reviewer would be expected to follow up."],
    "C11": ["Obtain the Tally release / licence history and any migration note for the year.",
            "Assess whether an upgrade or migration affected the audit trail."],
    "C12": ["Decide whether the software is hosted or run by a third party.",
            "If it is, obtain the SOC 1/SOC 2 report covering audit trail and the whole year; check scope and exceptions.",
            "For on-premise TallyPrime, record 'Not applicable' with the inventory as evidence."],
}

EVIDENCE_TO_OBTAIN = {
    "C8": ["Folder permission screenshot (data folder)", "Back-up location and who can access it", "Evidence that access is logged"],
    "C9": ["Management's written confirmation on direct data changes", "List of third-party tools installed on the Tally PC"],
    "C10": ["Review sign-offs for sampled months", "Follow-up of exceptions noted by management"],
    "C11": ["Tally release history", "Migration or upgrade note, if any", "Assessment of the impact on audit trail"],
    "C12": ["SOC 1 / SOC 2 / SAE 3402 report (if hosted)", "Hosting agreement or declaration that the software is on-premise"],
    "C7": ["Management's approval for each configuration change"],
}


def paper(cid: str, e: G.Engagement, c: Ctx, suggestion: tuple | None = None) -> dict:
    ctrl = next(x for x in G.CONTROLS if x.id == cid)
    tables: list[tuple[str, pd.DataFrame]] = []
    has_log = _n(c.log) > 0
    if cid == "C1" and has_log:
        t = pd.to_datetime(c.log["LogTime"], errors="coerce")
        tables.append(("Edit Log at a glance", pd.DataFrame({"Item": ["Entries", "First entry", "Last entry", "Users", "Actions"],
                       "Value": [len(c.log), f"{t.min():%d-%b-%Y %H:%M}", f"{t.max():%d-%b-%Y %H:%M}", c.log["User"].nunique(),
                                 ", ".join(f"{k}: {v}" for k, v in c.log["Action"].value_counts().items())]})))
        tables.append(("Extract to agree with Tally's screen", tally_sample(c, 15)))
    elif cid == "C2":
        tables.append(("Configuration-type events in Company Activity", config_changes(c)))
        tables.append(("All Company Activity", company_activity(c).head(200)))
        tables.append(("Users", users_table(c)))
    elif cid == "C3":
        tables.append(("Month by month", month_table(c)))
        if c.res is not None and c.res.gaps:
            tables.append(("Runs of working days with no trail activity (test A2)", pd.DataFrame(c.res.gaps)))
    elif cid == "C4":
        tables.append(("Register against Edit Log, by voucher type", coverage_by_type(c)))
        tables.append(("Vouchers with no Created entry (test A1)", a1_table(c)))
        tables.append(("Versions per voucher", version_table(c)))
    elif cid == "C5":
        tables.append(("Completeness of the trail", completeness_table(c)))
        tables.append(("Random sample of 25 entries (tick and tie)", tally_sample(c, 25)))
        tables.append(("Differences-screen sample (fill in)", pd.DataFrame(
            {"Voucher": ["", "", "", "", ""], "Version": "", "Field": "", "Old value": "", "New value": "", "Agrees to the Edit Log?": ""})))
    elif cid == "C6":
        tables.append(("Users in the Edit Log", users_table(c)))
        if c.res is not None and c.res.exceptions is not None:
            d2 = c.res.exceptions[c.res.exceptions.Test == "D2"]
            tables.append(("Changes by generic / shared IDs (test D2)", d2))
    elif cid == "C7":
        tables.append(("Configuration-type events", config_changes(c)))
        if _n(c.masters):
            tables.append(("Master-data changes by object and user", c.masters.groupby(["Object", "User"]).size().rename("Changes").reset_index()))
    elif cid == "C10" and c.res is not None and c.res.summary is not None:
        tables.append(("Test results management should have followed up", c.res.summary))
    elif cid == "C12":
        tables.append(("Software inventory (step 2)", pd.DataFrame(e.inventory)))
    if cid in EVIDENCE_TO_OBTAIN:
        tables.append(("Evidence to obtain", blank_checklist(EVIDENCE_TO_OBTAIN[cid])))
    tables = [(n, d) for n, d in tables if d is not None and (len(d) or n in ("Evidence to obtain", "Differences-screen sample (fill in)"))]
    cur = e.controls.get(cid, {})
    return {"id": cid, "title": f"{cid} · {ctrl.area}", "header": _hdr(e, c, cid, ctrl.area, ctrl.test, ctrl.ref, ctrl.how),
            "procedure": PROCEDURES.get(cid, []), "tables": tables,
            "suggested": suggestion, "result": cur.get("result", "Not tested"), "confirmed": cur.get("confirmed", True),
            "note": cur.get("note", ""), "conclusion": (e.wp_notes or {}).get(cid, "")}


# --------------------------------------------------------------------------- Excel
def _write(xw, p: dict, sheet: str):
    import openpyxl.styles as S
    ws = xw.book.create_sheet(sheet[:31])
    xw.sheets[ws.title] = ws
    bold = S.Font(bold=True)
    r = 1
    ws.cell(r, 1, p["title"]).font = S.Font(bold=True, size=14); r += 2
    for k, v in p["header"]:
        ws.cell(r, 1, k).font = bold; ws.cell(r, 2, str(v)); r += 1
    r += 1
    ws.cell(r, 1, "Procedure").font = bold; r += 1
    for i, s in enumerate(p["procedure"], 1):
        ws.cell(r, 1, i); ws.cell(r, 2, s); r += 1
    r += 1
    ws.cell(r, 1, "Result").font = bold
    ws.cell(r, 2, f"{p['result']}" + ("" if p["confirmed"] else "  (suggested, not yet confirmed by the auditor)")); r += 1
    if p["suggested"]:
        ws.cell(r, 1, "Suggested").font = bold; ws.cell(r, 2, f"{p['suggested'][0]} — {p['suggested'][1]}"); r += 1
    ws.cell(r, 1, "Note").font = bold; ws.cell(r, 2, p["note"]); r += 1
    ws.cell(r, 1, "Auditor's conclusion").font = bold; ws.cell(r, 2, p["conclusion"]); r += 2
    for name, df in p["tables"]:
        ws.cell(r, 1, name).font = bold; r += 1
        d = df.copy()
        for col in d.columns:
            if pd.api.types.is_datetime64_any_dtype(d[col]):
                d[col] = d[col].dt.strftime("%d-%b-%Y %H:%M")
        for j, col in enumerate(d.columns, 1):
            ws.cell(r, j, str(col)).font = bold
        r += 1
        for row in d.itertuples(index=False):
            for j, v in enumerate(row, 1):
                ws.cell(r, j, v if isinstance(v, (int, float, str)) or v is None else str(v))
            r += 1
        r += 1
    ws.cell(r, 1, "Prepared by / date").font = bold; ws.cell(r + 1, 1, "Reviewed by / date").font = bold
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 22


def papers_xlsx(papers: list[dict]) -> bytes:
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as xw:
        pd.DataFrame({"Rule 11(g) working papers": [p["title"] for p in papers],
                      "Result": [p["result"] + ("" if p["confirmed"] else " (suggested)") for p in papers]}).to_excel(xw, sheet_name="Index", index=False)
        for p in papers:
            _write(xw, p, p["id"])
    return bio.getvalue()
