"""
exporter.py
-----------
Builds the final downloadable summary in Word (.docx) or PDF format,
reflecting whatever verdicts/remarks the user has set in the GUI.
"""

from __future__ import annotations

import datetime
from typing import List

from models import (
    AnalysisResult,
    ClausePoint,
    SECTION_MAIN_SUMMARY,
    SECTION_GOODS,
    SECTION_DOCUMENTS,
    SECTION_CHARGES,
    SUBSECTION_ORDER,
)

TABLE_HEADERS = ["Sl No.", "Point No.", "Point Header", "Summary", "Analysis / Remarks", "Final Verdict", "User Remarks"]


def _rows_for(clauses: List[ClausePoint]) -> List[List[str]]:
    rows = []
    for c in clauses:
        rows.append([
            str(c.sl_no),
            c.point_no,
            c.header,
            c.summary,
            c.analysis_note,
            c.effective_verdict(),
            c.user_remarks or "",
        ])
    return rows


# Column widths for the Word table, in inches. Landscape A4 with 1.27cm
# margins gives ~10.7" of usable width - these sum to a bit under that so
# the table never gets squeezed back down to portrait-width columns (which
# is what causes text to wrap letter-by-letter and become unreadable).
_WORD_COL_WIDTHS_IN = [0.45, 0.75, 1.3, 2.5, 2.5, 1.3, 1.35]


def _set_landscape(document) -> None:
    from docx.shared import Cm
    from docx.enum.section import WD_ORIENT

    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(1.27)
    section.right_margin = Cm(1.27)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)


def _make_repeating_header(row) -> None:
    """Marks a table row to repeat as a header on every page it spans
    across (python-docx has no first-class API for this - it's the
    <w:tblHeader/> row property)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    trPr = row._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    trPr.append(header_flag)


def _set_col_widths(table) -> None:
    """python-docx only reliably renders column widths when they're set on
    every individual cell (not just on table.columns), so set both."""
    from docx.shared import Inches

    table.autofit = False
    widths = [Inches(w) for w in _WORD_COL_WIDTHS_IN]
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    for col, width in zip(table.columns, widths):
        col.width = width


def _style_header_row(row) -> None:
    from docx.shared import Pt

    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)


def _style_body_row(row) -> None:
    from docx.shared import Pt
    from docx.enum.table import WD_ALIGN_VERTICAL

    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.size = Pt(9)


def _add_table(doc, clauses: List[ClausePoint]):
    table = doc.add_table(rows=1, cols=len(TABLE_HEADERS))
    table.style = "Light Grid Accent 1"
    hdr_row = table.rows[0]
    for i, h in enumerate(TABLE_HEADERS):
        hdr_row.cells[i].text = h
    _style_header_row(hdr_row)
    _make_repeating_header(hdr_row)
    for row_values in _rows_for(clauses):
        cells = table.add_row().cells
        for i, val in enumerate(row_values):
            cells[i].text = val
        _style_body_row(table.rows[-1])
    _set_col_widths(table)
    return table


def export_to_word(result: AnalysisResult, out_path: str) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    _set_landscape(doc)

    # Slightly smaller base font so a landscape page still comfortably fits
    # a wide, readable table.
    normal_style = doc.styles["Normal"]
    normal_style.font.size = Pt(10)

    doc.add_heading("LC Analysis Summary", level=0)

    meta = doc.add_paragraph()
    meta.add_run("Document type: ").bold = True
    meta.add_run(result.lc_type + "\n")
    meta.add_run("Documentary Credit No.: ").bold = True
    meta.add_run((result.dc_number or "N/A") + "\n")
    meta.add_run("Source file: ").bold = True
    meta.add_run(result.source_filename + "\n")
    meta.add_run("Generated on: ").bold = True
    meta.add_run(datetime.datetime.now().strftime("%d-%b-%Y %H:%M"))

    def add_section(title_text: str, clauses: List[ClausePoint]):
        doc.add_heading(title_text, level=1)
        if not clauses:
            doc.add_paragraph("No points identified under this section.")
            return
        _add_table(doc, clauses)

    add_section(SECTION_MAIN_SUMMARY, result.by_section(SECTION_MAIN_SUMMARY))
    add_section(SECTION_GOODS, result.by_section(SECTION_GOODS))

    doc.add_heading(SECTION_DOCUMENTS, level=1)
    doc_clauses = result.by_section(SECTION_DOCUMENTS)
    any_sub = False
    for sub in SUBSECTION_ORDER:
        sub_clauses = [c for c in doc_clauses if c.subsection == sub]
        if not sub_clauses:
            continue
        any_sub = True
        doc.add_heading(sub, level=2)
        _add_table(doc, sub_clauses)
    if not any_sub:
        doc.add_paragraph("No document requirement points identified.")

    add_section(SECTION_CHARGES, result.by_section(SECTION_CHARGES))

    doc.save(out_path)


def _group_by_verdict(clauses: List[ClausePoint]):
    """Splits clauses into the 3 buckets the finalized summary is grouped
    into: confirmed-correct, needs-amendment, and everything else (Needs
    Clarification / Informational / any other verdict)."""
    correct, amendment, other = [], [], []
    for c in clauses:
        v = c.effective_verdict()
        if v == "Correct / In Order":
            correct.append(c)
        elif v == "Requires Amendment":
            amendment.append(c)
        else:
            other.append(c)
    return correct, amendment, other


def export_final_summary_word(result: AnalysisResult, out_path: str) -> None:
    """Builds the 'finalized' single Word file the user reviews and signs
    off on: all points across the whole document (not split by section),
    grouped into three tables by their final verdict - correct/in order,
    requires amendment, and other remarks - reflecting exactly what the
    user set with the verdict radio buttons and remarks boxes in the GUI."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    _set_landscape(doc)

    normal_style = doc.styles["Normal"]
    normal_style.font.size = Pt(10)

    doc.add_heading("LC Analysis - Finalized Summary", level=0)

    meta = doc.add_paragraph()
    meta.add_run("Document type: ").bold = True
    meta.add_run(result.lc_type + "\n")
    meta.add_run("Documentary Credit No.: ").bold = True
    meta.add_run((result.dc_number or "N/A") + "\n")
    meta.add_run("Source file: ").bold = True
    meta.add_run(result.source_filename + "\n")
    meta.add_run("Generated on: ").bold = True
    meta.add_run(datetime.datetime.now().strftime("%d-%b-%Y %H:%M"))

    correct, amendment, other = _group_by_verdict(result.clauses)

    def add_group(title_text: str, clauses: List[ClausePoint], empty_msg: str):
        doc.add_heading(title_text, level=1)
        if not clauses:
            doc.add_paragraph(empty_msg)
            return
        _add_table(doc, clauses)

    add_group(
        f"1. Points Confirmed Correct / In Order ({len(correct)})", correct,
        "No points have been finalised as Correct / In Order.",
    )
    add_group(
        f"2. Points Requiring Amendment ({len(amendment)})", amendment,
        "No points are currently marked as requiring amendment.",
    )
    add_group(
        f"3. Other Remarks (Needs Clarification / Informational) ({len(other)})", other,
        "No points fall under this category.",
    )

    doc.save(out_path)


def _find_field_text(result: AnalysisResult, tag: str) -> str:
    """Looks up the raw text captured for a given SWIFT field tag (e.g. '50'
    for Applicant, '59' for Beneficiary) from the already-parsed clauses, so
    the amendment letter can auto-fill the addressee/signatory blocks."""
    for c in result.clauses:
        if c.point_no == tag:
            return (c.raw_text or c.summary or "").strip()
    return ""


def generate_amendment_letter(
    result: AnalysisResult, out_path: str, amendment_clauses: List[ClausePoint]
) -> None:
    """Drafts a formal letter - addressed to the Applicant / Opener of the
    Letter of Credit, per standard trade practice that only the applicant
    can instruct their bank to amend a credit - listing every point the user
    has marked 'Requires Amendment' together with the remarks they typed as
    the specific amendment being requested."""
    from docx import Document
    from docx.shared import Pt, Cm

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    normal_style = doc.styles["Normal"]
    normal_style.font.size = Pt(11)

    applicant_text = _find_field_text(result, "50") or (
        "[Name & Address of the Applicant / Opener of the Letter of Credit - "
        "not found automatically, please fill in]"
    )
    beneficiary_text = _find_field_text(result, "59") or "[Your Company Name - Beneficiary]"
    dc_number = result.dc_number or "[LC Number]"

    doc.add_paragraph(datetime.datetime.now().strftime("%d-%b-%Y"))
    doc.add_paragraph("")

    doc.add_paragraph("To,")
    for line in applicant_text.splitlines():
        line = line.strip()
        if line:
            doc.add_paragraph(line)
    doc.add_paragraph("")

    sub_par = doc.add_paragraph()
    sub_par.add_run(f"Sub: Request for Amendment(s) to Letter of Credit No. {dc_number}").bold = True
    doc.add_paragraph("")

    doc.add_paragraph("Dear Sir/Madam,")
    doc.add_paragraph("")
    doc.add_paragraph(
        f"We refer to the captioned Letter of Credit No. {dc_number} opened by you in our "
        "favour. On scrutiny of its terms and conditions, we request you to kindly arrange "
        "for the following amendment(s) through your bankers at the earliest, to enable us "
        "to comply with the credit terms and effect timely shipment / presentation of "
        "documents without any discrepancy:"
    )
    doc.add_paragraph("")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    table.autofit = False
    headers = ["Sl No.", "LC Point / Field", "Present Clause (As per LC)", "Amendment Requested"]
    widths = [Cm(1.5), Cm(3.0), Cm(5.5), Cm(6.0)]
    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        hdr_row.cells[i].text = h
    _style_header_row(hdr_row)
    _make_repeating_header(hdr_row)
    for i, c in enumerate(amendment_clauses, start=1):
        cells = table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = f"{c.point_no} - {c.header}"
        cells[2].text = c.summary
        remark = (c.user_remarks or "").strip()
        cells[3].text = remark if remark else "(Please specify the exact amendment required.)"
        _style_body_row(table.rows[-1])
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    for col, w in zip(table.columns, widths):
        col.width = w

    doc.add_paragraph("")
    doc.add_paragraph(
        "We shall be grateful if the above amendment(s) is/are arranged at the earliest to "
        "avoid any delay in shipment or discrepancy in documents at the time of negotiation / "
        "presentation. Please treat this as urgent."
    )
    doc.add_paragraph("")
    doc.add_paragraph("Thanking you,")
    doc.add_paragraph("Yours faithfully,")
    doc.add_paragraph("")
    doc.add_paragraph("")
    for line in beneficiary_text.splitlines():
        line = line.strip()
        if line:
            doc.add_paragraph(line)
    doc.add_paragraph("(Authorised Signatory)")

    doc.save(out_path)


def export_to_pdf(result: AnalysisResult, out_path: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak,
    )

    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=8, leading=10)
    header_style = ParagraphStyle("cellHeader", parent=styles["BodyText"], fontSize=8, leading=10, textColor=colors.white)

    doc = SimpleDocTemplate(
        out_path, pagesize=landscape(A4),
        leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm,
    )
    elements = []

    elements.append(Paragraph("LC Analysis Summary", styles["Title"]))
    elements.append(Paragraph(f"Document type: {result.lc_type}", styles["Normal"]))
    elements.append(Paragraph(f"Documentary Credit No.: {result.dc_number or 'N/A'}", styles["Normal"]))
    elements.append(Paragraph(f"Source file: {result.source_filename}", styles["Normal"]))
    elements.append(Paragraph(
        f"Generated on: {datetime.datetime.now().strftime('%d-%b-%Y %H:%M')}", styles["Normal"]
    ))
    elements.append(Spacer(1, 0.5 * cm))

    col_widths = [1.5 * cm, 2.2 * cm, 4 * cm, 6.5 * cm, 6.5 * cm, 3 * cm, 4 * cm]

    def build_table(clauses: List[ClausePoint]):
        data = [[Paragraph(h, header_style) for h in TABLE_HEADERS]]
        for row in _rows_for(clauses):
            data.append([Paragraph(str(v).replace("\n", "<br/>"), cell_style) for v in row])
        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E4A62")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F8")]),
        ]))
        return table

    def add_section(title_text: str, clauses: List[ClausePoint]):
        elements.append(Paragraph(title_text, styles["Heading1"]))
        if not clauses:
            elements.append(Paragraph("No points identified under this section.", styles["Normal"]))
        else:
            elements.append(build_table(clauses))
        elements.append(Spacer(1, 0.5 * cm))

    add_section(SECTION_MAIN_SUMMARY, result.by_section(SECTION_MAIN_SUMMARY))
    add_section(SECTION_GOODS, result.by_section(SECTION_GOODS))

    elements.append(Paragraph(SECTION_DOCUMENTS, styles["Heading1"]))
    doc_clauses = result.by_section(SECTION_DOCUMENTS)
    any_sub = False
    for sub in SUBSECTION_ORDER:
        sub_clauses = [c for c in doc_clauses if c.subsection == sub]
        if not sub_clauses:
            continue
        any_sub = True
        elements.append(Paragraph(sub, styles["Heading2"]))
        elements.append(build_table(sub_clauses))
        elements.append(Spacer(1, 0.3 * cm))
    if not any_sub:
        elements.append(Paragraph("No document requirement points identified.", styles["Normal"]))

    add_section(SECTION_CHARGES, result.by_section(SECTION_CHARGES))

    doc.build(elements)
