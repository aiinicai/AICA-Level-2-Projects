"""Trial Balance Review & Audit Exception Analyzer — single-file application."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

REQUIRED = ["Ledger Name", "Group", "Opening Balance", "Debit", "Credit", "Closing Balance"]
REPORT_COLUMNS = ["Ledger Name", "Group", "Closing Balance", "Exception Type", "Remarks"]
DISCLAIMER = "Exceptions are indicators requiring professional review, not conclusions of accounting error or fraud."
DEFAULT_RULES = {
    "cash": "Debit", "bank": "Debit", "cash/bank": "Debit",
    "expense": "Debit", "expenses": "Debit", "income": "Credit",
    "sales": "Credit", "liability": "Credit", "liabilities": "Credit",
    "asset": "Debit", "assets": "Debit", "equity": "Credit",
    "receivables": "Debit", "payables": "Credit",
}


def normalized(value):
    return " ".join(str(value).split()).casefold()


def paise(value):
    """Strict two-decimal input; integer paise avoids float comparison errors."""
    if isinstance(value, bool):
        raise ValueError("Boolean values are not amounts.")
    text = str(value).strip()
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text):
        raise ValueError("Use a numeric amount without currency symbols or commas.")
    try:
        number = Decimal(text)
        if not number.is_finite() or abs(number) > Decimal("1000000000000"):
            raise ValueError("Amount must be finite and within +/- 1,000,000,000,000.")
        cents = number * 100
        if cents != cents.to_integral_value():
            raise ValueError("Use no more than two decimal places.")
        return int(cents)
    except InvalidOperation as exc:
        raise ValueError("Invalid numeric amount.") from exc


def read_trial_balance(source):
    """Accept uploaded file/bytes. Read first sheet and return validated data."""
    raw = source if isinstance(source, bytes) else source.getvalue()
    if not raw:
        raise ValueError("The uploaded file is empty. Upload a populated .xlsx workbook.")
    if len(raw) > 10 * 1024 * 1024:
        raise ValueError("Please upload a workbook smaller than 10 MB.")
    book = None
    try:
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        sheet = book.worksheets[0]
        rows = sheet.iter_rows()
        first = next(rows, ())
        headers = [str(c.value).strip() if c.value is not None else "" for c in first]
        if not any(headers):
            raise ValueError("The first worksheet is empty. Put the required headers in row 1.")
        missing = [c for c in REQUIRED if c not in headers]
        if missing:
            raise ValueError("Missing required columns: " + ", ".join(missing))
        if any(headers.count(c) != 1 for c in REQUIRED):
            raise ValueError("Required column headers must not be duplicated.")
        positions = [headers.index(c) for c in REQUIRED]
        for row in rows:
            for pos in positions:
                if pos < len(row) and row[pos].data_type == "f":
                    raise ValueError("Formula cells are not supported. Paste values before uploading.")
        data = pd.read_excel(BytesIO(raw), engine="openpyxl", dtype=object, keep_default_na=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Cannot read this Excel file. Upload a valid, unencrypted .xlsx workbook.") from exc
    finally:
        if book is not None:
            book.close()
    data.columns = [str(c).strip() for c in data.columns]
    data = data[REQUIRED].copy()
    data = data.map(lambda v: pd.NA if isinstance(v, str) and not v.strip() else v)
    data = data.dropna(how="all")
    if data.empty:
        raise ValueError("No ledger rows found. Add ledger data below the headers.")
    for column in REQUIRED:
        for index, value in data[column].items():
            location = f"Excel row {index + 2}, {column}"
            if pd.isna(value):
                raise ValueError(f"{location}: blank values are not allowed; enter 0 for zero amounts.")
            if column in REQUIRED[:2]:
                data.at[index, column] = str(value).strip()
            else:
                try:
                    cents = paise(value)
                except ValueError as exc:
                    raise ValueError(f"{location}: {exc}") from exc
                if column in ("Debit", "Credit") and cents < 0:
                    raise ValueError(f"{location}: period Debit/Credit must be non-negative.")
                data.at[index, column] = cents / 100
    for column in REQUIRED[2:]:
        data[column] = data[column].astype(float)
    return data.reset_index(drop=True)


def analyze(data, materiality=500000, round_multiple=10000, rules=None):
    """One report row per exception per source row; mismatch is one TB-level row."""
    threshold, multiple = paise(materiality), paise(round_multiple)
    if threshold < 0 or multiple <= 0:
        raise ValueError("Materiality must be non-negative and round multiple must be positive.")
    rules = DEFAULT_RULES if rules is None else rules
    rules = {normalized(group): nature for group, nature in rules.items()}
    if any(nature not in ("Debit", "Credit", "Ignore") for nature in rules.values()):
        raise ValueError("Expected nature must be Debit, Credit or Ignore.")
    debit = sum(paise(v) for v in data["Debit"])
    credit = sum(paise(v) for v in data["Credit"])
    difference = debit - credit
    duplicate = data["Ledger Name"].map(normalized).duplicated(keep=False)
    output = []
    counts = {"High-Value Items": 0, "Negative Balances": 0, "Unusual Balances": 0}
    if difference:
        output.append({"Ledger Name": "[Overall Trial Balance]", "Group": "[Overall TB]",
                       "Closing Balance": None, "Exception Type": "Trial Balance Difference Identified",
                       "Remarks": f"Debit minus Credit = INR {difference / 100:,.2f}. Period totals only."})
    for index, row in data.iterrows():
        close, dr, cr = (paise(row[c]) for c in ("Closing Balance", "Debit", "Credit"))
        def flag(kind, remarks):
            output.append({"Ledger Name": row["Ledger Name"], "Group": row["Group"],
                           "Closing Balance": close / 100, "Exception Type": kind,
                           "Remarks": f"Ledger row {index + 1}: {remarks}"})
        if close < 0:
            counts["Negative Balances"] += 1
            flag("Negative Balance", "Credit-signed closing balance; may be normal for this group.")
        if duplicate.iloc[index]:
            flag("Duplicate Ledger Name", "Repeated name after ignoring case and repeated whitespace; verify distinct accounts.")
        if abs(close) > threshold:
            counts["High-Value Items"] += 1
            flag("High Value", f"Absolute closing balance exceeds INR {threshold / 100:,.2f}; not a full audit materiality assessment.")
        sides = [name for name, amount in (("Debit", dr), ("Credit", cr)) if amount != 0 and amount % multiple == 0]
        if sides:
            flag("Round Figure Transaction – Review Required",
                 f"{', '.join(sides)} period total is a non-zero exact multiple of INR {multiple / 100:,.2f}. Aggregate screening only, not voucher testing.")
        if close == 0 and (dr != 0 or cr != 0):
            flag("Zero Closing Balance", "Period movement exists with zero closing balance; review if relevant.")
        nature = rules.get(normalized(row["Group"]), "Ignore")
        if (nature == "Debit" and close < 0) or (nature == "Credit" and close > 0):
            counts["Unusual Balances"] += 1
            flag("Unusual Balance – Review Required", f"Configured expected nature is {nature}; closing sign is contrary. Review classification/context.")
    report = pd.DataFrame(output, columns=REPORT_COLUMNS)
    largest = max(data["Closing Balance"], key=lambda v: abs(paise(v)), default=0)
    summary = {"Total Ledgers": len(data), "Total Debit": debit / 100, "Total Credit": credit / 100,
               "Difference": difference / 100, "Highest Closing Balance": float(largest),
               "Total Exceptions": len(report), **counts}
    return report, summary


def filter_report(report, types=None, groups=None):
    result = report
    if types:
        result = result[result["Exception Type"].isin(types)]
    if groups:
        result = result[result["Group"].isin(groups)]
    return result.reset_index(drop=True)


def excel_bytes(sheets):
    """Export literal text (including names starting '=') rather than Excel formulas."""
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    elif isinstance(cell.value, (float, int)):
                        cell.number_format = '#,##0.00;[Red]-#,##0.00'
            for column in sheet.columns:
                letter = column[0].column_letter
                width = max(len(str(c.value or "")) for c in column) + 2
                sheet.column_dimensions[letter].width = min(75, max(16, width))
    return buffer.getvalue()


def sample_data():
    # Name, group, opening, period debit, period credit. Closing = opening + Dr - Cr.
    rows = [
        ("Cash in Hand", "Cash", 10000, 40000, 60000),
        ("Main Bank", "Bank", 150000, 800000, 200000),
        ("Petty Cash", "Cash", 2000, 5123, 3123),
        ("Trade Receivables", "Receivables", 125000, 225500, 100500),
        ("Inventory", "Assets", 200000, 310250, 110250),
        ("Plant and Machinery", "Assets", 500000, 250000, 0),
        ("Office Equipment", "Assets", 60000, 12345, 0),
        ("Rent Expense", "Expenses", 0, 120000, 0),
        ("Salaries Expense", "Expenses", 0, 360000, 0),
        ("Electricity Expense", "Expenses", 0, 45678, 0),
        ("Travel Expense", "Expenses", 0, 12345, 23456),
        ("Sales Revenue", "Sales", 0, 0, 900000),
        ("Service Income", "Income", 0, 45678, 12345),
        ("Interest Income", "Income", 0, 0, 12345),
        ("Trade Payables", "Payables", -125000, 50500, 100500),
        ("Bank Loan", "Liabilities", -500000, 100000, 0),
        ("Customer Advances", "Liabilities", -10000, 25000, 5000),
        ("Owner Capital", "Equity", -612000, 0, 0),
        ("Clearing Account", "Assets", 0, 20000, 20000),
        ("Dormant Account", "Assets", 0, 0, 0),
        ("Office Supplies", "Expenses", 0, 6789, 0),
        ("Office Supplies", "Expenses", 0, 4321, 0),
        ("Insurance Prepaid", "Assets", 15000, 15555, 5555),
        ("Security Deposit", "Assets", 25000, 0, 0),
    ]
    return pd.DataFrame([[*row, row[2] + row[3] - row[4]] for row in rows], columns=REQUIRED)




def report_bytes(report, summary, rules, materiality, multiple, types=None, groups=None, source=""):
    """One export path shared by the desktop UI and tests."""
    settings = pd.DataFrame([
        ("Source filename", source), ("Materiality (INR)", materiality), ("Round multiple (INR)", multiple),
        ("Exception Type filter", "; ".join(types or []) or "All"),
        ("Ledger Group filter", "; ".join(groups or []) or "All"),
        ("Sign convention", "Positive = debit; negative = credit"),
        ("Scope", "Exceptions is filtered; All Exceptions and Summary are unfiltered"),
        ("Disclaimer", DISCLAIMER),
    ], columns=["Setting", "Value"])
    return excel_bytes({
        "Exceptions": filter_report(report, types, groups), "All Exceptions": report,
        "Summary": pd.DataFrame(summary.items(), columns=["Metric", "Value"]), "Settings": settings,
        "Expected Nature": pd.DataFrame(rules.items(), columns=["Group", "Expected Nature"]),
    })


class AnalyzerWindow:
    """Small desktop UI. Tkinter is imported only when the GUI is started."""

    def __init__(self, root):
        import tkinter as tk
        from tkinter import ttk, filedialog, messagebox
        self.tk, self.ttk = tk, ttk
        self.dialog, self.message = filedialog, messagebox
        self.root = root
        self.data = None
        self.report = pd.DataFrame(columns=REPORT_COLUMNS)
        self.summary, self.rules = {}, {}
        self.source = ""
        self.dirty = True
        root.title("Trial Balance Review & Audit Exception Analyzer")
        root.geometry("1180x760")
        root.minsize(960, 680)
        style = ttk.Style(root)
        if "clam" in style.theme_names():
            style.theme_use("clam")
        style.configure("Title.TLabel", font=("Segoe UI", 17, "bold"), foreground="#17365D")
        style.configure("Metric.TLabel", font=("Segoe UI", 14, "bold"), foreground="#17365D")
        style.configure("Treeview", rowheight=25)
        style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
        style.configure("TButton", padding=5)
        shell = ttk.Frame(root, padding=12)
        shell.pack(fill="both", expand=True)
        ttk.Label(shell, text="Trial Balance Review & Audit Exception Analyzer", style="Title.TLabel").pack(anchor="w")
        ttk.Label(shell, text="ICAI capstone demonstration | Desktop edition | INR | No browser required").pack(anchor="w", pady=(2, 5))
        ttk.Label(shell, text=DISCLAIMER, foreground="#9A4D00").pack(anchor="w")
        ttk.Label(shell, text="Positive opening/closing = Debit; negative = Credit. Period Debit/Credit must be non-negative.\n"
                  "Round-figure checks use ledger totals, not individual vouchers. Paste values; exclude totals rows.").pack(anchor="w", pady=(4, 8))
        toolbar = ttk.Frame(shell)
        toolbar.pack(fill="x")
        ttk.Button(toolbar, text="Open Trial Balance (.xlsx)", command=self.open_file).pack(side="left", padx=(0, 6))
        ttk.Button(toolbar, text="Save Sample Excel Format", command=self.save_sample).pack(side="left", padx=6)
        self.export_button = ttk.Button(toolbar, text="Download Audit Exception Report", command=self.export, state="disabled")
        self.export_button.pack(side="left", padx=6)
        self.file_label = tk.StringVar(value="No workbook opened")
        ttk.Label(shell, textvariable=self.file_label).pack(anchor="w", pady=6)
        settings = ttk.LabelFrame(shell, text="Review settings — click Analyze / Refresh after changing settings", padding=8)
        settings.pack(fill="x")
        self.materiality = tk.StringVar(value="500000")
        self.multiple = tk.StringVar(value="10000")
        ttk.Label(settings, text="High-value threshold (INR)").grid(row=0, column=0, sticky="w")
        ttk.Entry(settings, textvariable=self.materiality, width=16).grid(row=0, column=1, padx=8)
        ttk.Label(settings, text="Round-figure multiple (INR)").grid(row=0, column=2, padx=(10, 0))
        ttk.Entry(settings, textvariable=self.multiple, width=14).grid(row=0, column=3, padx=8)
        ttk.Button(settings, text="Analyze / Refresh", command=self.refresh).grid(row=0, column=4, padx=6)
        self.rule_button = ttk.Button(settings, text="Expected Nature Rules", command=self.edit_rules, state="disabled")
        self.rule_button.grid(row=0, column=5, padx=6)
        self.materiality.trace_add("write", self.mark_dirty)
        self.multiple.trace_add("write", self.mark_dirty)
        dashboard = ttk.Frame(shell)
        dashboard.pack(fill="x", pady=10)
        self.metrics = {}
        for index, name in enumerate(["Total Ledgers", "Total Exceptions", "High-Value Items", "Negative Balances", "Unusual Balances"]):
            dashboard.columnconfigure(index, weight=1)
            card = ttk.LabelFrame(dashboard, text=name, padding=8)
            card.grid(row=0, column=index, sticky="ew", padx=3)
            self.metrics[name] = tk.StringVar(value="—")
            ttk.Label(card, textvariable=self.metrics[name], style="Metric.TLabel").pack(anchor="w")
        self.summary_label = tk.StringVar(value="Open an Excel workbook to begin.")
        ttk.Label(shell, textvariable=self.summary_label).pack(anchor="w")
        ttk.Label(shell, text="Dashboard is unfiltered; exceptions count flags, not unique ledgers. Highest balance = largest magnitude, shown signed.").pack(anchor="w", pady=(2, 6))
        self.notebook = ttk.Notebook(shell)
        self.notebook.pack(fill="both", expand=True)
        exceptions_tab = ttk.Frame(self.notebook, padding=6)
        data_tab = ttk.Frame(self.notebook, padding=6)
        self.notebook.add(exceptions_tab, text="Audit Exceptions")
        self.notebook.add(data_tab, text="Uploaded Trial Balance")
        filters = ttk.Frame(exceptions_tab)
        filters.pack(fill="x", pady=(0, 6))
        ttk.Label(filters, text="Exception Type").pack(side="left")
        # First item means all; indices avoid collisions with literal group names.
        self.type_box = ttk.Combobox(filters, values=["All types"], state="readonly", width=43)
        self.type_box.current(0)
        self.type_box.pack(side="left", padx=6)
        ttk.Label(filters, text="Ledger Group").pack(side="left", padx=(8, 0))
        self.group_box = ttk.Combobox(filters, values=["All groups"], state="readonly", width=24)
        self.group_box.current(0)
        self.group_box.pack(side="left", padx=6)
        self.type_box.bind("<<ComboboxSelected>>", self.apply_filters)
        self.group_box.bind("<<ComboboxSelected>>", self.apply_filters)
        ttk.Button(filters, text="Clear Filters", command=self.clear_filters).pack(side="left", padx=6)
        self.count_label = tk.StringVar(value="No report yet")
        ttk.Label(exceptions_tab, textvariable=self.count_label).pack(anchor="w", pady=(0, 5))
        self.exception_tree = self.make_table(exceptions_tab, REPORT_COLUMNS)
        self.data_tree = self.make_table(data_tab, REQUIRED)
        self.exception_tree.bind("<Double-1>", self.show_remark)
        ttk.Label(exceptions_tab, text="Double-click an exception for full remarks. Scroll horizontally/vertically to view all columns and rows.").pack(anchor="w", pady=(4, 0))
        self.status = tk.StringVar(value="Ready. Use the sample workbook for your first review.")
        ttk.Label(shell, textvariable=self.status, foreground="#17365D").pack(anchor="w", pady=(8, 0))

    def make_table(self, parent, columns):
        frame = self.ttk.Frame(parent)
        frame.pack(fill="both", expand=True)
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        tree = self.ttk.Treeview(frame, columns=columns, show="headings", selectmode="browse")
        for name in columns:
            width = 420 if name == "Remarks" else 290 if name == "Exception Type" else 180
            tree.heading(name, text=name)
            tree.column(name, width=width, minwidth=100, stretch=False,
                        anchor="e" if "Balance" in name or name in ("Debit", "Credit") else "w")
        vertical = self.ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        horizontal = self.ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        tree.grid(row=0, column=0, sticky="nsew")
        vertical.grid(row=0, column=1, sticky="ns")
        horizontal.grid(row=1, column=0, sticky="ew")
        return tree

    @staticmethod
    def fill_table(tree, frame):
        items = tree.get_children()
        if items:
            tree.delete(*items)
        for row in frame.itertuples(index=False, name=None):
            values = []
            for column, value in zip(frame.columns, row):
                if pd.isna(value):
                    values.append("")
                elif column in REQUIRED[2:]:
                    values.append(f"{value:,.2f}")
                else:
                    values.append(str(value))
            tree.insert("", "end", values=values)

    def mark_dirty(self, *_):
        self.dirty = True
        self.export_button.configure(state="disabled")
        if self.data is not None:
            self.status.set("Settings changed — displayed results are previous results. Click Analyze / Refresh before exporting.")

    def open_file(self):
        path = self.dialog.askopenfilename(parent=self.root, title="Open Trial Balance",
                                          filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        try:
            selected = Path(path)
            if selected.suffix.lower() != ".xlsx":
                raise ValueError("Select an .xlsx workbook.")
            if selected.stat().st_size > 10 * 1024 * 1024:
                raise ValueError("Please select a workbook smaller than 10 MB.")
            data = read_trial_balance(selected.read_bytes())
        except (ValueError, OSError) as exc:
            suffix = "\n\nPreviously opened data remains unchanged." if self.data is not None else ""
            self.message.showerror("Cannot open Trial Balance", str(exc) + suffix, parent=self.root)
            return
        self.data, self.source = data, selected.name
        self.file_label.set(f"Opened: {self.source} | {len(data)} ledger rows")
        self.rules = {group: DEFAULT_RULES.get(normalized(group), "Ignore") for group in sorted(data["Group"].unique())}
        self.report = pd.DataFrame(columns=REPORT_COLUMNS)
        self.summary = {}
        for value in self.metrics.values():
            value.set("—")
        self.summary_label.set("Workbook loaded. Run analysis to calculate totals.")
        self.fill_table(self.data_tree, data)
        self.rule_button.configure(state="normal")
        self.clear_filters()
        self.mark_dirty()
        self.refresh()

    def refresh(self):
        if self.data is None:
            self.message.showinfo("Open a workbook", "Open a Trial Balance Excel file first.", parent=self.root)
            return False
        try:
            report, summary = analyze(self.data, self.materiality.get(), self.multiple.get(), self.rules)
        except ValueError as exc:
            self.message.showerror("Review settings", str(exc), parent=self.root)
            return False
        self.report, self.summary = report, summary
        for name, value in self.metrics.items():
            value.set(str(summary[name]))
        self.summary_label.set(
            f"Total Debit: INR {summary['Total Debit']:,.2f}    |    Total Credit: INR {summary['Total Credit']:,.2f}\n"
            f"Debit − Credit: INR {summary['Difference']:,.2f}    |    Highest Closing Balance: INR {summary['Highest Closing Balance']:,.2f}")
        self.type_box.configure(values=["All types"] + sorted(report["Exception Type"].unique()))
        self.group_box.configure(values=["All groups"] + sorted(report["Group"].unique()))
        self.clear_filters()
        self.dirty = False
        self.export_button.configure(state="normal")
        if summary["Difference"]:
            self.status.set(f"Trial Balance Difference Identified: INR {summary['Difference']:,.2f}. Review required.")
        else:
            self.status.set("Period debit and credit totals agree. This alone does not establish correctness.")
        return True

    def selected_filters(self):
        types = [self.type_box.get()] if self.type_box.current() > 0 else []
        groups = [self.group_box.get()] if self.group_box.current() > 0 else []
        return types, groups

    def apply_filters(self, *_):
        types, groups = self.selected_filters()
        filtered = filter_report(self.report, types, groups)
        self.fill_table(self.exception_tree, filtered)
        self.count_label.set(f"Showing {len(filtered)} of {len(self.report)} exception rows. Overall mismatch uses group [Overall TB].")

    def clear_filters(self):
        self.type_box.current(0)
        self.group_box.current(0)
        self.apply_filters()

    def show_remark(self, _event=None):
        selected = self.exception_tree.selection()
        if selected:
            row = self.exception_tree.item(selected[0], "values")
            self.message.showinfo(str(row[3]), f"Ledger: {row[0]}\nGroup: {row[1]}\nClosing: {row[2]}\n\n{row[4]}", parent=self.root)

    def edit_rules(self):
        if self.data is None:
            return
        popup = self.tk.Toplevel(self.root)
        popup.title("Expected Nature Rules")
        popup.geometry("610x430")
        popup.transient(self.root)
        popup.grab_set()
        self.ttk.Label(popup, text="Select a group, choose its expected nature, then click Set Nature.\n"
                       "Unmapped groups default to Ignore. Consider overdrafts and contra accounts.", padding=10).pack(anchor="w")
        local_rules = self.rules.copy()
        table = self.make_table(popup, ["Group", "Expected Nature"])
        def populate():
            self.fill_table(table, pd.DataFrame(local_rules.items(), columns=["Group", "Expected Nature"]))
        populate()
        controls = self.ttk.Frame(popup, padding=10)
        controls.pack(fill="x")
        nature = self.ttk.Combobox(controls, values=["Debit", "Credit", "Ignore"], state="readonly", width=12)
        nature.current(0)
        nature.pack(side="left")
        def set_nature():
            selected = table.selection()
            if not selected:
                self.message.showinfo("Select group", "Select a group row first.", parent=popup)
                return
            group = table.item(selected[0], "values")[0]
            for key in local_rules:
                if normalized(key) == normalized(group):
                    local_rules[key] = nature.get()
            populate()
        def commit():
            self.rules = local_rules
            self.mark_dirty()
            popup.destroy()
            self.refresh()
        self.ttk.Button(controls, text="Set Nature", command=set_nature).pack(side="left", padx=8)
        self.ttk.Button(controls, text="Apply Rules", command=commit).pack(side="right", padx=4)
        self.ttk.Button(controls, text="Cancel", command=popup.destroy).pack(side="right", padx=4)

    def save_sample(self):
        self.save_workbook("sample_trial_balance.xlsx", lambda: excel_bytes({"Trial Balance": sample_data()}))

    def save_workbook(self, filename, producer):
        path = self.dialog.asksaveasfilename(parent=self.root, title="Save Excel workbook",
                                            initialfile=filename, defaultextension=".xlsx",
                                            filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        try:
            if Path(path).suffix.lower() != ".xlsx":
                raise ValueError("Use the .xlsx extension for the saved workbook.")
            Path(path).write_bytes(producer())
        except Exception as exc:
            self.message.showerror("Cannot save workbook", f"{exc}\n\nClose the file in Excel and check folder permissions.", parent=self.root)
            return
        self.message.showinfo("Workbook saved", f"Saved successfully:\n{path}", parent=self.root)

    def export(self):
        if self.data is None or self.dirty:
            self.message.showinfo("Refresh required", "Open a workbook and click Analyze / Refresh before exporting.", parent=self.root)
            return
        types, groups = self.selected_filters()
        self.save_workbook("audit_exception_report.xlsx", lambda: report_bytes(
            self.report, self.summary, self.rules, self.materiality.get(), self.multiple.get(),
            types, groups, self.source))


def main():
    try:
        import tkinter as tk
        root = tk.Tk()
    except (ImportError, RuntimeError) as exc:
        raise SystemExit(f"Tkinter is unavailable: {exc}\nInstall Python with Tcl/Tk support and try: python -m tkinter")
    except Exception as exc:
        raise SystemExit(f"Cannot open the desktop window: {exc}\nRun this application in a desktop session, not a headless server.")
    AnalyzerWindow(root)
    root.mainloop()


if __name__ == "__main__":
    main()
