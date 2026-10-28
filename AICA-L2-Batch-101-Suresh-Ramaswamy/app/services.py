"""Workflow layer used by the UI (and by tests): ingest files, merge, classify, dashboard stats."""
import logging
import shutil
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from typing import List, Optional

from . import gstr1
from . import hsn as hsn_mod
from .autofill import autofill_lines
from .classifier import classify_all, is_implemented
from .db import Database
from .exporters import make_project_folders
from .extractors import ReadResult
from .extractors.excel_reader import read_workbook
from .extractors.pdf_reader import read_pdf
from .models import Line, STATUS_APPROVED, STATUS_DUP, STATUS_PENDING, group_invoices, invoice_value
from .utils import sha256_file, ZERO

log = logging.getLogger(__name__)
EXCEL_EXT = {".xlsx", ".xlsm", ".xls", ".csv"}
PDF_EXT = {".pdf"}


def file_type(path: Path) -> str:
    e = Path(path).suffix.lower()
    return "Excel" if e in EXCEL_EXT else "PDF" if e in PDF_EXT else "Unsupported"


def register_file(db: Database, client: dict, project: dict, src: Path):
    """Hash, de-duplicate and copy a source file into the project folder.
    Returns (file_id or None, message)."""
    src = Path(src)
    ftype = file_type(src)
    if ftype == "Unsupported":
        return None, f"{src.name}: unsupported file type"
    sha = sha256_file(src)
    dup = db.file_by_hash(project["id"], sha)
    if dup:
        return None, f"{src.name}: identical file already uploaded as '{dup['filename']}' - skipped"
    root = make_project_folders(client, project)
    dest = root / "source_files" / src.name
    n = 1
    while dest.exists():
        dest = root / "source_files" / f"{src.stem}_{n}{src.suffix}"
        n += 1
    shutil.copy2(src, dest)
    fid = db.add_file(project["id"], src.name, dest, sha, ftype)
    db.audit(project["id"], "File uploaded", f"{src.name} sha256={sha[:12]}")
    return fid, "queued"


def extract_file(db: Database, client: dict, project: dict, fid: int, settings: dict,
                 override: Optional[dict] = None) -> ReadResult:
    rec = db.q("SELECT * FROM files WHERE id=?", (fid,))[0]
    path = Path(rec["stored_path"])
    report = None
    if rec["ftype"] == "Excel" and not override:
        from .extractors import reports
        report = reports.detect_report(path)
    if report:
        from .extractors import reports
        from .schema import load_schema
        res = ReadResult(kind=report)
        schema = load_schema()
        if report == reports.HSN_REPORT:
            res.tables, msgs = reports.read_hsn_report(path, schema)
        else:
            res.tables, msgs, res.docs = reports.read_output_register(path, client, schema)
        res.messages += msgs
        for d in db.datasets(project["id"]):
            if d["filename"] == rec["filename"]:
                db.delete_dataset(d["id"])
        if res.tables:
            db.add_dataset(project["id"], rec["filename"], rec["sha256"] if "sha256" in rec.keys() else "", res.tables)
            db.audit(project["id"], "Summary report read", f"{rec['filename']}: " +
                     ", ".join(f"{k}={len(v)}" for k, v in res.tables.items()))
    elif rec["ftype"] == "Excel":
        res = read_workbook(path, client, settings, override)
    else:
        res = read_pdf(path, client, settings)
    for l in res.lines:
        l.source_file = rec["filename"]
    hsn_mod.apply_to_lines(res.lines, db.hsn_master(client["id"]))          # description / unit / rate from memory
    filled = autofill_lines(res.lines, client["gstin"][:2], settings)
    if filled:
        res.messages.append(f"{filled} line(s): place of supply / tax completed automatically")
    if res.needs_mapping and not res.lines:
        status = "Needs column mapping"
    elif not res.lines and not res.tables:
        status = "Failed"
    else:
        status = "Extracted"
    db.update_file(fid, status=status, records=len(res.lines) or sum(len(v) for v in res.tables.values()),
                   message="; ".join(res.messages)[:500], processed_at=__import__("datetime").datetime.now()
                   .strftime("%Y-%m-%d %H:%M:%S"))
    log.info("Extracted file id=%s type=%s records=%d status=%s", fid, rec["ftype"], len(res.lines), status)
    return res


def merge_lines(existing: List[Line], new: List[Line]) -> int:
    """Append new lines. Lines identical to an existing line from ANOTHER file are flagged for review
    instead of being merged or dropped. Returns number flagged."""
    nid = max((l.id for l in existing), default=0) + 1
    sigs = defaultdict(list)
    for l in existing:
        sigs[(l.key(), l.hsn, l.taxable, l.rate, l.igst, l.cgst, l.sgst)].append(l)
    flagged = 0
    for l in new:
        l.id = nid
        nid += 1
        sig = (l.key(), l.hsn, l.taxable, l.rate, l.igst, l.cgst, l.sgst)
        other = [x for x in sigs.get(sig, []) if x.source_file != l.source_file]
        if other:
            l.review_status = STATUS_DUP
            l.note = (l.note + "; " if l.note else "") + f"Same as line in {other[0].source_file}"
            flagged += 1
        existing.append(l)
        sigs[sig].append(l)
    return flagged


def review_count(lines: List[Line]) -> int:
    return sum(1 for l in lines if l.review_status in (STATUS_PENDING, STATUS_DUP) or l.confidence < 0.6)


def dashboard(lines: List[Line], issues, files, project: dict) -> dict:
    live = [l for l in lines if l.review_status != "Rejected"]
    invs = group_invoices(live)
    appr_inv = group_invoices([l for l in live if l.review_status == STATUS_APPROVED])
    b2b = group_invoices([l for l in live if l.section == "B2B"])
    sections = defaultdict(lambda: [0, ZERO])
    for k, g in invs.items():
        for s in {l.section for l in g}:
            sections[s][0] += 1
            sections[s][1] += sum((l.taxable for l in g if l.section == s), ZERO)
    return {
        "files": len(files), "invoices": len(invs), "approved": len(appr_inv),
        "to_review": len(invs) - len(appr_inv),
        "taxable": sum((l.taxable for l in live), ZERO), "igst": sum((l.igst for l in live), ZERO),
        "cgst": sum((l.cgst for l in live), ZERO), "sgst": sum((l.sgst for l in live), ZERO),
        "cess": sum((l.cess for l in live), ZERO),
        "b2b_count": len(b2b), "b2b_value": sum((invoice_value(g) for g in b2b.values()), ZERO),
        "errors": sum(1 for i in issues if i.severity == "Error"),
        "warnings": sum(1 for i in issues if i.severity == "Warning"),
        "csv_status": project.get("csv_status", "Not generated") if project else "-",
        "sections": dict(sections),
    }


# ---------------------------------------------------------------- master template
def import_master_template(db: Database, client: dict, project: dict, src: Path, settings: dict, schema: dict,
                           skip_errors: bool = False):
    """Read and validate a filled Master Template. Nothing is stored when there is any error.

    Returns (result, dataset_id | None, message). `result.docs` holds the Documents-sheet series - the caller
    decides whether they replace or are added to the existing register."""
    src = Path(src)
    sha = sha256_file(src)
    dup = db.dataset_by_hash(project["id"], sha)
    if dup:
        res = gstr1.TemplateResult()
        return res, None, f"{src.name}: identical file already imported as '{dup['filename']}' - skipped"
    ctx = gstr1.make_ctx(client, project, settings, schema)
    res = gstr1.read_master_template(src, schema, ctx, src.name, skip_errors=skip_errors)
    if res.errors:
        return res, None, f"{len(res.errors)} error(s) - nothing was imported"
    if res.total_rows == 0:
        return res, None, "The template contains no data rows"
    root = make_project_folders(client, project)
    dest = root / "source_files" / src.name
    n = 1
    while dest.exists():
        dest = root / "source_files" / f"{src.stem}_{n}{src.suffix}"
        n += 1
    shutil.copy2(src, dest)
    ds_id = db.add_dataset(project["id"], src.name, sha, res.tables)
    db.audit(project["id"], "Master template imported",
             f"{src.name}: " + ", ".join(f"{k}={v}" for k, v in sorted(res.counts.items()))
             + (f"; {res.skipped} row(s) with errors skipped" if res.skipped else ""))
    return res, ds_id, "imported"


def merge_docs(existing, new, mode: str):
    """mode: 'replace' | 'add'. Suggested ('derived') series are always dropped."""
    keep = [d for d in existing if d.origin != "derived"]
    return list(new) if mode == "replace" else keep + list(new)


def is_master_template(path: Path) -> bool:
    """True if an Excel file looks like the GSTR-1 Master Template (so it is not read as an invoice register)."""
    if Path(path).suffix.lower() not in (".xlsx", ".xlsm"):
        return False
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        names = {n.lower().replace(" ", "") for n in wb.sheetnames}
        wb.close()
        return "returninfo" in names and {"b2b", "b2cs"} <= names
    except Exception:  # noqa
        return False


def remove_template_rows(db: Database, sets: list, targets) -> int:
    """Delete stored Master Template rows. targets = {(file name, 'B2B row 7'), ...}; for invoices / notes the whole
    invoice goes. Returns the number of rows removed; the database is updated."""
    targets = set(targets)
    removed = 0
    for ds in sets:
        changed = False
        for key, rows in list(ds["tables"].items()):
            meta = gstr1.TABLE_META.get(key, {})
            hit = [r for r in rows if (r.get("__file"), r.get("__ref")) in targets]
            if not hit:
                continue
            ids = {r.get(meta["id"]) for r in hit} if meta.get("family") in ("invoice", "note") else set()
            keep = [r for r in rows if r not in hit and not (ids and r.get(meta["id"]) in ids)]
            removed += len(rows) - len(keep)
            ds["tables"][key] = keep
            changed = True
        if changed:
            db.update_dataset(ds["id"], ds["tables"])
    return removed


def learn_hsn(db: Database, client_id: int, lines) -> int:
    """Remember description / unit / rate of every approved HSN of this client for the next files."""
    n = 0
    for l in lines:
        if l.review_status != STATUS_APPROVED or not l.hsn:
            continue
        db.remember_hsn(client_id, l.hsn, l.description, l.uqc, l.rate if l.taxable else "")
        n += 1
    return n
