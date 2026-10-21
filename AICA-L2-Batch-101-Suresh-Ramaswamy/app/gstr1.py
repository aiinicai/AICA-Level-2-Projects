"""GSTR-1 tables engine.

* Master Template  - ONE Excel workbook (one sheet per GSTR-1 table) that the client fills in; written and read here.
* Normalisation    - every cell is converted to a canonical value and checked (dates, GSTIN, place of supply, ...).
* Validation       - row rules per table + cross-table / cross-source checks.
* Tables           - rows from the template and from extracted invoices are merged and aggregated into the
                     tables written as CSV for the Offline Tool.
* Summary / recon  - tax summary per table and reconciliation checks.
* Header check     - compares the schema headings with the official sample CSV files.

Canonical row = dict keyed by the schema 'source' names; dates are ISO (YYYY-MM-DD), place of supply is the
2-digit state code, money is a 2-decimal string, rate/qty are plain decimal strings.
"""
import calendar
import csv
import re
from collections import OrderedDict, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Dict, List, Optional

from . import autofill
from . import hsn as hsn_mod
from .config import STATE_CODES
from .gstin import validate_gstin, checksum_char
from .models import DocSeries, STATUS_APPROVED, group_invoices, invoice_value
from .utils import (D, q2, parse_date, fmt_date, pos_label, state_from_text, parse_rates, norm_invoice_no,
                    MONTHS, ZERO)
from .validation import Issue, ERROR, WARNING, INFO

IGST_TYPES = {"SEZ supplies with payment", "SEZ supplies without payment", "Deemed Exp",
              "Intra-State supplies attracting IGST"}
INV_SUMS = ("taxable", "cess")
LINE_COLS = {"rate", "taxable", "cess", "ecom_gstin", "applicable_pct"}   # may differ between rows of one invoice

# family: invoice | note | summary | exemp | hsn | docs
TABLE_META = {
    "b2b": dict(family="invoice", id="invoice_no", date="invoice_date", value="invoice_value", sums=INV_SUMS,
                amt=("taxable",), tax="pos_type", type_col="invoice_type", what="Invoice"),
    "b2cl": dict(family="invoice", id="invoice_no", date="invoice_date", value="invoice_value", sums=INV_SUMS,
                 amt=("taxable",), tax="igst", what="Invoice"),
    "exp": dict(family="invoice", id="invoice_no", date="invoice_date", value="invoice_value", sums=INV_SUMS,
                amt=("taxable",), tax="igst", what="Invoice"),
    "cdnr": dict(family="note", id="note_no", date="note_date", value="note_value", sums=INV_SUMS,
                 amt=("taxable",), tax="pos_type", type_col="note_supply_type", what="Note"),
    "cdnur": dict(family="note", id="note_no", date="note_date", value="note_value", sums=INV_SUMS,
                  amt=("taxable",), tax="igst", what="Note"),
    "b2ba": dict(family="invoice", id="invoice_no", date="invoice_date", value="invoice_value", sums=INV_SUMS,
                 amt=("taxable",), tax="pos_type", type_col="invoice_type", what="Invoice", amend=True),
    "b2cla": dict(family="invoice", id="invoice_no", date="invoice_date", value="invoice_value", sums=INV_SUMS,
                  amt=("taxable",), tax="igst", what="Invoice", amend=True),
    "expa": dict(family="invoice", id="invoice_no", date="invoice_date", value="invoice_value", sums=INV_SUMS,
                 amt=("taxable",), tax="igst", what="Invoice", amend=True),
    "cdnra": dict(family="note", id="note_no", date="note_date", value="note_value", sums=INV_SUMS,
                  amt=("taxable",), tax="pos_type", type_col="note_supply_type", what="Note", amend=True),
    "cdnura": dict(family="note", id="note_no", date="note_date", value="note_value", sums=INV_SUMS,
                   amt=("taxable",), tax="igst", what="Note", amend=True),
    "b2cs": dict(family="summary", sums=INV_SUMS, amt=("taxable",), tax="pos"),
    "b2csa": dict(family="summary", sums=INV_SUMS, amt=("taxable",), tax="pos", amend=True),
    "at": dict(family="summary", sums=("gross_advance", "cess"), amt=("gross_advance",), tax="pos", gross=True),
    "atadj": dict(family="summary", sums=("gross_adjusted", "cess"), amt=("gross_adjusted",), tax="pos", gross=True),
    "ata": dict(family="summary", sums=("gross_advance", "cess"), amt=("gross_advance",), tax="pos", gross=True,
                amend=True),
    "atadja": dict(family="summary", sums=("gross_adjusted", "cess"), amt=("gross_adjusted",), tax="pos",
                   gross=True, amend=True),
    "exemp": dict(family="exemp", sums=("nil_rated", "exempted", "non_gst"), amt=("nil_rated", "exempted", "non_gst")),
    "hsn": dict(family="hsn", sums=("qty", "total_value", "taxable", "igst", "cgst", "sgst", "cess"),
                nokey=("description",), amt=("taxable",)),
    "docs": dict(family="docs"),
}
DATA_TABLES = [k for k in TABLE_META if k != "docs"]
IMPL_SECTIONS = {"B2B", "B2CL", "B2CS", "CDNR", "CDNUR", "EXP", "NIL", "AT", "ATADJ"}
_INVNO = re.compile(r"^[A-Za-z0-9/\-]{1,16}$")
_NORM = lambda s: re.sub(r"[^a-z0-9]", "", str(s).lower())  # noqa: E731


def no_tax_row(row: dict) -> bool:
    """True when no tax is charged on the row: reverse charge, exports / SEZ supplies without payment of tax."""
    return (row.get("reverse_charge") == "Y" or row.get("export_type") == "WOPAY"
            or row.get("ur_type") == "EXPWOP"
            or row.get("invoice_type") == "SEZ supplies without payment"
            or row.get("note_supply_type") == "SEZ supplies without payment")


# ====================================================================== parsing helpers
def _blank(v) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and v != v:
        return True
    return isinstance(v, str) and v.strip() == ""


def _txt(v) -> str:
    if _blank(v):
        return ""
    if isinstance(v, bool):
        return "Y" if v else "N"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "strftime"):
        return v.strftime("%d-%m-%Y")
    return str(v).strip()


def parse_decimal(raw):
    """-> (Decimal|None, 'blank'|'ok'|'bad')"""
    if _blank(raw):
        return None, "blank"
    if isinstance(raw, bool):
        return None, "bad"
    if isinstance(raw, (int, Decimal)):
        return Decimal(raw), "ok"
    if isinstance(raw, float):
        return Decimal(str(raw)), "ok"
    d = D(str(raw).strip(), None)
    if d is None or d.is_nan() or d.is_infinite():
        return None, "bad"
    return d, "ok"


def plain(d: Decimal) -> str:
    if d == 0:
        return "0"
    s = format(d.normalize(), "f")
    return s


def parse_month(raw) -> str:
    """-> 'YYYY-MM' or ''"""
    if _blank(raw):
        return ""
    if hasattr(raw, "year") and hasattr(raw, "month"):
        return f"{raw.year:04d}-{raw.month:02d}"
    s = str(raw).strip()
    m = re.match(r"^([A-Za-z]{3,9})[\s\-/,.]*(\d{2}|\d{4})$", s)
    if m:
        mon = m.group(1)[:3].lower()
        idx = [x.lower() for x in MONTHS].index(mon) + 1 if mon in [x.lower() for x in MONTHS] else 0
        if idx:
            y = int(m.group(2))
            return f"{(y + 2000 if y < 100 else y):04d}-{idx:02d}"
    m = re.match(r"^(\d{1,2})[\-/.](\d{4}|\d{2})$", s)
    if m and 1 <= int(m.group(1)) <= 12:
        y = int(m.group(2))
        return f"{(y + 2000 if y < 100 else y):04d}-{int(m.group(1)):02d}"
    m = re.match(r"^(\d{4})[\-/.](\d{1,2})$", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    return ""


def parse_fy(raw) -> str:
    s = _txt(raw)
    m = re.match(r"^(\d{2}|\d{4})\s*[-/]\s*(\d{2}|\d{4})$", s)
    if not m:
        return ""
    a, b = int(m.group(1)), int(m.group(2))
    a = a + 2000 if a < 100 else a
    b = b % 100
    if b != (a + 1) % 100:
        return ""
    return f"{a}-{b:02d}"


def fmt_month(ym: str, schema: dict) -> str:
    if not ym:
        return ""
    try:
        y, m = ym.split("-")
        mon = MONTHS[int(m) - 1]
        style = schema.get("month_format", "Mon-YY")
        return style.replace("YYYY", y).replace("YY", y[2:]).replace("Mon", mon)
    except (ValueError, IndexError):
        return ym


def _choices(schema, key, source):
    for c in schema["tables"][key]["columns"]:
        if c["source"] == source:
            return c.get("choices", [])
    return []


def normalize_cell(col: dict, raw, schema: dict):
    """-> (canonical value, error message or None)."""
    kind = col.get("kind", "text")
    if kind == "date":
        if _blank(raw):
            return "", None
        d = parse_date(raw)
        return (d, None) if d else ("", f"'{_txt(raw)}' is not a valid date (use DD-MM-YYYY)")
    if kind in ("money", "rate", "qty", "int"):
        d, st = parse_decimal(raw)
        if st == "blank":
            return "", None
        if st == "bad":
            return "", f"'{_txt(raw)}' is not a number"
        if kind == "money":
            return str(q2(d)), None
        if kind == "int":
            if d != d.to_integral() or d < 0:
                return "", f"'{_txt(raw)}' must be a whole number (0 or more)"
            return str(int(d)), None
        return plain(d), None
    if kind == "month":
        if _blank(raw):
            return "", None
        v = parse_month(raw)
        return (v, None) if v else ("", f"'{_txt(raw)}' is not a valid month (use e.g. Sep-2026)")
    s = _txt(raw)
    if s == "":
        return "", None
    if kind == "gstin":
        return s.upper().replace(" ", ""), None
    if kind == "pos":
        t = ("0" + s) if (s.isdigit() and len(s) == 1) else s
        code = state_from_text(t)
        return (code, None) if code else ("", f"'{s}' is not a recognised state / place of supply")
    if kind == "choice":
        choices = col.get("choices", [])
        low = {c.lower(): c for c in choices}
        al = {k.lower(): v for k, v in col.get("aliases", {}).items()}
        if s.lower() in low:
            return low[s.lower()], None
        if s.lower() in al:
            return al[s.lower()], None
        comp = {_NORM(c): c for c in choices}
        if _NORM(s) in comp:
            return comp[_NORM(s)], None
        shown = ", ".join(choices[:6]) + (" ..." if len(choices) > 6 else "")
        return "", f"'{s}' is not an allowed value. Allowed: {shown}"
    if kind == "uqc":
        codes = {u.split("-")[0]: u for u in schema.get("uqc", [])}
        code = s.upper().split("-")[0].strip()
        if code in codes:
            return codes[code], None
        for u in schema.get("uqc", []):
            if _NORM(u) == _NORM(s):
                return u, None
        return "", f"'{s}' is not a valid unit of measure (UQC). Use a code such as NOS, KGS, PCS"
    if kind == "fy":
        v = parse_fy(s)
        return (v, None) if v else ("", f"'{s}' is not a financial year (use e.g. 2026-27)")
    if kind == "apct":
        m = re.match(r"^(\d{1,3})\s*%?$", s)
        return (f"{m.group(1)}%", None) if m else ("", f"'{s}' is not valid (leave blank or enter 65 / 65%)")
    if kind == "hsn":
        return re.sub(r"\s+", "", s), None
    return s, None


# ====================================================================== context
@dataclass
class Ctx:
    gstin: str
    state: str
    start: date
    end: date
    tol: Decimal
    rates: list
    b2cl_threshold: Decimal
    hsn_min: int
    schema: dict
    month: int = 0
    year: int = 0
    auto_pos: bool = True


def period_bounds(month: int, year: int, frequency: str = ""):
    if frequency and frequency.lower().startswith("quarter"):
        qe = ((month - 1) // 3 + 1) * 3
        qs = qe - 2
        return date(year, qs, 1), date(year, qe, calendar.monthrange(year, qe)[1])
    return date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])


def make_ctx(client: dict, project: dict, settings: dict, schema: dict) -> Ctx:
    m, y = (project["month"], project["year"]) if project else (date.today().month, date.today().year)
    start, end = period_bounds(m, y, (project or {}).get("frequency", ""))
    return Ctx(gstin=client["gstin"], state=client["gstin"][:2], start=start, end=end,
               tol=D(settings.get("rounding_tolerance", "1.00")), rates=parse_rates(settings.get("valid_rates", "")),
               b2cl_threshold=D(settings.get("b2cl_threshold", "100000")),
               hsn_min=int(D(settings.get("hsn_min_digits", "4"))), schema=schema, month=m, year=y,
               auto_pos=autofill.flag(settings, "auto_pos"))


# ====================================================================== sample data
def _mk_gstin(prefix14: str) -> str:
    return prefix14 + checksum_char(prefix14)


def sample_data(own_gstin: str, month: int, year: int) -> Dict[str, List[dict]]:
    """Demo rows (input form) for every table; they pass validation and reconcile with each other."""
    own = own_gstin[:2] if own_gstin else "27"
    oth = "29" if own != "29" else "27"
    c_inter = _mk_gstin(oth + "AABCT1332L1Z")
    c_intra = _mk_gstin(own + "AAACR5055K1Z")
    ecom = _mk_gstin("07AAACE1234E1Z")
    pm, py = (month - 1, year) if month > 1 else (12, year - 1)
    fy_start = year if month >= 4 else year - 1
    fy = f"{fy_start}-{str(fy_start + 1)[-2:]}"
    d = lambda day: date(year, month, day)          # noqa: E731
    pd_ = date(py, pm, 20)
    pmon = f"{MONTHS[pm - 1]}-{py}"
    return {
        "b2b": [
            dict(recipient_gstin=c_inter, recipient_name="Karnataka Traders", invoice_no="INV001", invoice_date=d(5),
                 invoice_value=11800, pos=oth, reverse_charge="N", invoice_type="Regular", rate=18, taxable=10000),
            dict(recipient_gstin=c_intra, recipient_name="Local Retail Co", invoice_no="INV002", invoice_date=d(8),
                 invoice_value=6425, pos=own, reverse_charge="N", invoice_type="Regular", rate=18, taxable=5000),
            dict(recipient_gstin=c_intra, recipient_name="Local Retail Co", invoice_no="INV002", invoice_date=d(8),
                 invoice_value=6425, pos=own, reverse_charge="N", invoice_type="Regular", rate=5, taxable=500),
            dict(recipient_gstin=_mk_gstin("24AABCS1429B1Z"), recipient_name="SEZ Unit Gujarat", invoice_no="INV003",
                 invoice_date=d(9), invoice_value=11800, pos="24", reverse_charge="N",
                 invoice_type="SEZ supplies with payment", rate=18, taxable=10000)],
        "b2cl": [dict(invoice_no="INV101", invoice_date=d(10), invoice_value=177000, pos=oth, rate=18, taxable=150000)],
        "b2cs": [
            dict(supply_type="OE", pos=own, rate=18, taxable=40000), dict(supply_type="OE", pos=own, rate=5, taxable=10000),
            dict(supply_type="OE", pos=oth, rate=18, taxable=20000),
            dict(supply_type="E", pos=own, rate=18, taxable=5000, ecom_gstin=ecom)],
        "exp": [
            dict(export_type="WPAY", invoice_no="INV201", invoice_date=d(15), invoice_value=236000, port_code="INMAA1",
                 ship_bill_no="1234567", ship_bill_date=d(14), rate=18, taxable=200000),
            dict(export_type="WOPAY", invoice_no="INV202", invoice_date=d(16), invoice_value=100000, port_code="INMAA1",
                 ship_bill_no="1234568", ship_bill_date=d(15), rate=18, taxable=100000)],
        "cdnr": [
            dict(recipient_gstin=c_inter, recipient_name="Karnataka Traders", note_no="CN001", note_date=d(20),
                 note_type="C", pos=oth, reverse_charge="N", note_supply_type="Regular", note_value=1180, rate=18,
                 taxable=1000),
            dict(recipient_gstin=c_intra, recipient_name="Local Retail Co", note_no="DN001", note_date=d(21),
                 note_type="D", pos=own, reverse_charge="N", note_supply_type="Regular", note_value=590, rate=18,
                 taxable=500)],
        "cdnur": [dict(ur_type="B2CL", note_no="CNU01", note_date=d(22), note_type="C", pos=oth, note_value=11800,
                       rate=18, taxable=10000)],
        "exemp": [
            dict(description="Inter-State supplies to registered persons", nil_rated=2000, exempted=0, non_gst=0),
            dict(description="Intra-State supplies to registered persons", nil_rated=5000, exempted=3000, non_gst=0),
            dict(description="Inter-State supplies to unregistered persons", nil_rated=0, exempted=1000, non_gst=0),
            dict(description="Intra-State supplies to unregistered persons", nil_rated=0, exempted=0, non_gst=500)],
        "at": [dict(pos=own, rate=18, gross_advance=11800)],
        "atadj": [dict(pos=own, rate=18, gross_adjusted=5900)],
        "hsn": [
            dict(hsn="8471", description="Computers", uqc="NOS", qty=10, total_value=472000, rate=18, taxable=400000,
                 igst=60000, cgst=6000, sgst=6000),
            dict(hsn="8473", description="Computer parts", uqc="NOS", qty=25, total_value=152810, rate=18,
                 taxable=129500, igst=20000, cgst=1655, sgst=1655),
            dict(hsn="2106", description="Food preparations", uqc="KGS", qty=100, total_value=11025, rate=5,
                 taxable=10500, cgst=262.5, sgst=262.5),
            dict(hsn="0401", description="Milk (nil rated)", uqc="LTR", qty=500, total_value=11500, rate=0,
                 taxable=11500)],
        "docs": [
            dict(nature="Invoices for outward supply", from_no="INV001", to_no="INV202", total=202, cancelled=5),
            dict(nature="Credit Note", from_no="CN001", to_no="CN002", total=2, cancelled=0),
            dict(nature="Debit Note", from_no="DN001", to_no="DN001", total=1, cancelled=0)],
        "b2ba": [dict(recipient_gstin=c_inter, recipient_name="Karnataka Traders", orig_invoice_no="INV009",
                      orig_invoice_date=pd_, invoice_no="INV009", invoice_date=pd_, invoice_value=12980, pos=oth,
                      reverse_charge="N", invoice_type="Regular", rate=18, taxable=11000)],
        "b2cla": [dict(orig_invoice_no="INV090", orig_invoice_date=pd_, pos=oth, invoice_no="INV090",
                       invoice_date=pd_, invoice_value=59000, rate=18, taxable=50000)],
        "b2csa": [dict(fy=fy, orig_month=pmon, pos=own, supply_type="OE", rate=18, taxable=3000)],
        "expa": [dict(export_type="WPAY", orig_invoice_no="INV190", orig_invoice_date=pd_, invoice_no="INV190",
                      invoice_date=pd_, invoice_value=118000, port_code="INMAA1", ship_bill_no="1234000",
                      ship_bill_date=pd_, rate=18, taxable=100000)],
        "cdnra": [dict(recipient_gstin=c_inter, recipient_name="Karnataka Traders", orig_note_no="CN090",
                       orig_note_date=pd_, note_no="CN090", note_date=pd_, note_type="C", pos=oth,
                       reverse_charge="N", note_supply_type="Regular", note_value=590, rate=18, taxable=500)],
        "cdnura": [dict(ur_type="B2CL", orig_note_no="CNU90", orig_note_date=pd_, note_no="CNU90", note_date=pd_,
                        note_type="C", pos=oth, note_value=5900, rate=18, taxable=5000)],
        "ata": [dict(fy=fy, orig_month=pmon, pos=own, rate=18, gross_advance=2360)],
        "atadja": [dict(fy=fy, orig_month=pmon, pos=own, rate=18, gross_adjusted=1180)],
    }


# ====================================================================== master template (write)
HELP = {
    "gstin": "15-character GSTIN/UIN", "invnum": "Max 16 characters: letters, digits, / and - only",
    "date": "Date, e.g. 05-09-2026 (a real Excel date also works)", "money": "Amount in rupees (2 decimals)",
    "rate": "GST rate %, e.g. 18, 5, 0.25", "qty": "Quantity", "int": "Whole number", "pos": "Pick the state from the list",
    "month": "Month, e.g. Sep-2026", "fy": "Financial year, e.g. 2026-27", "apct": "Leave blank (or 65 for 65%)",
    "hsn": "HSN/SAC code - 4, 6 or 8 digits (enter as TEXT)", "uqc": "Unit - pick from the list",
}

INSTRUCTIONS = [
    "GSTR-1 MASTER TEMPLATE  -  how to fill",
    "",
    "1. Fill ONLY the sheets that apply to the client. Empty sheets are ignored. Do not rename sheets, do not delete, "
    "reorder or rename the column headings in row 1.",
    "2. Orange headings = mandatory columns. Grey headings = optional (a default is used when blank).",
    "3. Hover over a heading to see what is expected. Columns with drop-down lists should be picked from the list.",
    "4. One row per invoice PER TAX RATE (B2B, B2CL, EXP, CDNR, CDNUR): an invoice with two rates takes two rows with "
    "the same invoice number, date, value and recipient. Repeat the invoice value on every row of the invoice.",
    "5. B2CS (Table 7) is a SUMMARY: one row per Type + Place of supply + Rate (+ E-commerce GSTIN). Credit notes to "
    "unregistered customers that are not B2CL are entered here as negative taxable values.",
    "6. Exempt-Nil-NonGST (Table 8): one row per description with the value of supplies (taxable value) in the "
    "matching column. Zero-rated exports/SEZ supplies go to EXP/B2B, not here.",
    "7. AT / ATADJ: enter GROSS advance (tax inclusive) received / adjusted, state-wise and rate-wise.",
    "8. HSN (Table 12): HSN/SAC-wise summary with quantity, value, taxable value and tax amounts. Enter HSN codes "
    "as text so that leading zeros are kept. The total of HSN taxable value should agree with the outward supplies "
    "(the application reconciles this).",
    "9. Documents (Table 13): one row per document series (invoices, credit notes, debit notes, ...) with first and "
    "last serial number, total issued and cancelled. Take these from the client's invoice register.",
    "10. Amendment sheets (B2BA, B2CLA, B2CSA, EXPA, CDNRA, CDNURA, ATA, ATADJA): enter the ORIGINAL details and the "
    "REVISED details exactly as they appear / should appear on the portal.",
    "11. Taxes are calculated from Taxable Value x Rate; you do not type IGST/CGST/SGST except in the HSN sheet "
    "(and even there they may be left blank - they are then calculated).",
    "12. MINIMUM DATA: for B2B / CDN type the recipient GSTIN, invoice (note) number, date, rate and taxable value. "
    "Place of Supply is taken from the first two digits of the recipient GSTIN, and Invoice Value is calculated from "
    "taxable value + tax + cess, when those cells are left blank. B2CS / advances with a blank Place of Supply are "
    "taken as the supplier's own state (the application points this out). B2CL needs the Place of Supply (inter-state).",
    "13. Save the file as .xlsx and upload it in the application (tab 'Master Template'). The whole file is "
    "checked; if any row has an error, nothing is imported and every problem is listed with its sheet and row number "
    "- or tick 'skip the rows with errors' to import the good invoices and get the bad ones listed.",
    "",
    "The sample version of this file contains demo data. Never upload the sample for a real client.",
]


def write_master_template(path, schema: dict, client: Optional[dict] = None, project: Optional[dict] = None,
                          sample: bool = False) -> Path:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.properties import PageSetupProperties

    p = Path(path)
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, 1):
        c = ws0.cell(i, 1, line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=14)
    ws0.column_dimensions["A"].width = 130
    ws0.cell(len(INSTRUCTIONS) + 2, 1, "Sheets in this workbook").font = Font(bold=True)
    r0 = len(INSTRUCTIONS) + 3
    for k, key in enumerate(schema["tables"]):
        ws0.cell(r0 + k, 1, f"{schema['tables'][key]['sheet']}  -  {schema['tables'][key]['title']}")

    wi = wb.create_sheet("Return Info")
    wi["A1"] = "Return information (checked when the file is uploaded)"
    wi["A1"].font = Font(bold=True, size=12)
    gst = client["gstin"] if client else ""
    info_rows = [("GSTIN of supplier", gst), ("Legal / trade name", (client or {}).get("legal_name") or
                                              (client or {}).get("name", "")),
                 ("Return month (1-12)", project["month"] if project else ""),
                 ("Return year", project["year"] if project else ""),
                 ("Filing frequency", (project or {}).get("frequency", "")),
                 ("Sample file (YES = demo data - do not upload)", "YES" if sample else "NO")]
    for i, (a, b) in enumerate(info_rows, 3):
        wi.cell(i, 1, a).font = Font(bold=True)
        wi.cell(i, 2, b)
    wi.column_dimensions["A"].width = 48
    wi.column_dimensions["B"].width = 34

    lists = OrderedDict()           # name -> values

    def list_ref(name, values):
        if name not in lists:
            lists[name] = list(values)
        col = get_column_letter(list(lists).index(name) + 1)
        return f"Lists!${col}$2:${col}${len(lists[name]) + 1}"

    fill_req = PatternFill("solid", fgColor="F8CBAD")
    fill_opt = PatternFill("solid", fgColor="D9D9D9")
    month = project["month"] if project else (date.today().month - 1 or 12)
    year = project["year"] if project else date.today().year
    sdata = sample_data(gst or "", month, year) if sample else {}
    MAXR = 2000

    for key, spec in schema["tables"].items():
        ws = wb.create_sheet(spec["sheet"])
        ws.freeze_panes = "A2"
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        for ci, col in enumerate(spec["columns"], 1):
            h = ws.cell(1, ci, col["header"])
            h.font = Font(bold=True)
            h.fill = fill_req if col.get("required") else fill_opt
            h.alignment = Alignment(wrap_text=True, vertical="center")
            tip = HELP.get(col.get("kind"), "")
            if col.get("choices") and col["kind"] == "choice":
                tip = "Allowed: " + ", ".join(col["choices"])
            if col.get("default") not in (None, ""):
                tip += f"  (blank = {col['default']})"
            if tip:
                h.comment = Comment(tip, "GSTR-1 template", width=260, height=90)
            L = get_column_letter(ci)
            kind = col.get("kind")
            w = {"gstin": 19, "invnum": 17, "date": 13, "money": 15, "rate": 8, "qty": 11, "int": 11, "pos": 28,
                 "month": 12, "fy": 11, "apct": 11, "hsn": 11, "uqc": 24}.get(kind, 26)
            if kind == "choice":
                w = min(40, max(14, max(len(c) for c in col["choices"]) + 3))
            ws.column_dimensions[L].width = max(w, min(34, len(col["header"]) * 0.9 + 4))
            cd = ws.column_dimensions[L]
            if kind in ("gstin", "invnum", "text", "hsn", "apct", "month", "fy", "choice", "uqc", "pos"):
                cd.number_format = "@"
            elif kind == "date":
                cd.number_format = "DD-MMM-YY"
            elif kind == "money":
                cd.number_format = "#,##0.00"
            rng = f"{L}2:{L}{MAXR}"
            dv = None
            if kind == "choice" and col.get("choices"):
                dv = DataValidation(type="list", formula1=list_ref(f"{key}.{col['source']}", col["choices"]),
                                    allow_blank=True)
            elif kind == "pos":
                dv = DataValidation(type="list", formula1=list_ref("POS", [pos_label(c) for c in STATE_CODES]),
                                    allow_blank=True)
            elif kind == "uqc":
                dv = DataValidation(type="list", formula1=list_ref("UQC", schema.get("uqc", [])), allow_blank=True)
            if dv:
                dv.showErrorMessage = True
                dv.errorTitle = col["header"]
                dv.error = "Please pick a value from the list."
                ws.add_data_validation(dv)
                dv.add(rng)
        ws.row_dimensions[1].height = 34
        for ri, row in enumerate(sdata.get(key, []), 2):
            for ci, col in enumerate(spec["columns"], 1):
                v = row.get(col["source"])
                if v is None:
                    continue
                if col["kind"] == "pos":
                    v = pos_label(v)
                c = ws.cell(ri, ci, v)
                if hasattr(v, "year"):
                    c.number_format = "DD-MMM-YY"

    wl = wb.create_sheet("Lists")
    for ci, (name, vals) in enumerate(lists.items(), 1):
        wl.cell(1, ci, name).font = Font(bold=True)
        wl.column_dimensions[get_column_letter(ci)].width = 30
        for ri, v in enumerate(vals, 2):
            wl.cell(ri, ci, v)
    wb.save(p)
    return p


# ====================================================================== master template (read)
@dataclass
class TemplateResult:
    tables: Dict[str, List[dict]] = field(default_factory=dict)
    docs: List[DocSeries] = field(default_factory=list)
    issues: List[Issue] = field(default_factory=list)
    counts: Dict[str, int] = field(default_factory=dict)
    skipped: int = 0                     # rows left out because of errors (skip_errors mode)

    @property
    def errors(self):
        return [i for i in self.issues if i.severity == ERROR]

    @property
    def warnings(self):
        return [i for i in self.issues if i.severity == WARNING]

    @property
    def total_rows(self):
        return sum(len(v) for v in self.tables.values()) + len(self.docs)


def _find_header(matrix, want_norms):
    for idx, (rn, vals) in enumerate(matrix[:20]):
        got = {_NORM(v) for v in vals if not _blank(v)}
        if len(got & want_norms) >= max(2, len(want_norms) // 2):
            return idx
    return None


def _sheet_matrix(ws):
    return [(i, list(r)) for i, r in enumerate(ws.iter_rows(values_only=True), 1)]


def read_master_template(path, schema: dict, ctx: Ctx, filename: str = "", skip_errors: bool = False) -> TemplateResult:
    from openpyxl import load_workbook
    p = Path(path)
    fname = filename or p.name
    res = TemplateResult()

    def err(sev, code, msg, ref="", action=""):
        res.issues.append(Issue(sev, code, msg, fname, "", action, 0, ref))

    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        err(ERROR, "FILE_TYPE", "The master template must be an .xlsx file", "", "Open it in Excel and Save As .xlsx")
        return res
    try:
        wb = load_workbook(p, data_only=True)
    except Exception as e:  # noqa
        err(ERROR, "FILE_OPEN", f"The file could not be opened: {e}", "", "Check that the file is not corrupted or open elsewhere")
        return res

    sheets = {_NORM(ws.title): ws for ws in wb.worksheets}
    # ---- return info guard
    info = {}
    if "returninfo" in sheets:
        for _, vals in _sheet_matrix(sheets["returninfo"]):
            if len(vals) >= 2 and not _blank(vals[0]):
                info[_NORM(vals[0])] = vals[1]
    sample = _txt(next((v for k, v in info.items() if k.startswith("samplefile")), "")).upper()
    if sample.startswith("YES"):
        err(ERROR, "SAMPLE_FILE", "This is the SAMPLE template with demo data", "Return Info",
            "Download the blank template and fill that instead")
    g = _txt(next((v for k, v in info.items() if k.startswith("gstin")), "")).upper()
    if g and g != ctx.gstin:
        err(ERROR, "TPL_GSTIN", f"The template is for GSTIN {g} but the open return is for {ctx.gstin}",
            "Return Info", "Use the right client's return, or correct the GSTIN on the Return Info sheet")
    tm = next((v for k, v in info.items() if k.startswith("returnmonth")), None)
    ty = next((v for k, v in info.items() if k.startswith("returnyear")), None)
    dm, dy = parse_decimal(tm)[0], parse_decimal(ty)[0]
    if dm is not None and dy is not None and (int(dm), int(dy)) != (ctx.month, ctx.year):
        err(ERROR, "TPL_PERIOD", f"The template is for {int(dm):02d}/{int(dy)} but the open return is "
                                 f"{ctx.month:02d}/{ctx.year}", "Return Info",
            "Open the correct return, or correct the month/year on the Return Info sheet")
    if not g:
        err(INFO, "TPL_NO_INFO", "The Return Info sheet is blank - GSTIN and period could not be cross-checked",
            "Return Info", "Download the template from this application after opening the return")

    found = 0
    deferred = []
    for key in schema["tables"]:
        spec = schema["tables"][key]
        ws = sheets.get(_NORM(spec["sheet"])) or sheets.get(_NORM(key))
        if ws is None:
            continue
        found += 1
        if key == "docs":
            _read_docs_sheet(ws, spec, schema, fname, res)
            continue
        matrix = _sheet_matrix(ws)
        cols = spec["columns"]
        want = {_NORM(c["header"]): c for c in cols}
        hi = _find_header(matrix, set(want))
        if hi is None:
            if any(any(not _blank(v) for v in vals) for _, vals in matrix):
                err(ERROR, "HEADERS", f"Column headings of sheet '{spec['sheet']}' not found", spec["sheet"],
                    "Do not change the headings in row 1 - copy your data into a fresh template")
            continue
        idx = {}
        for ci, h in enumerate(matrix[hi][1]):
            c = want.get(_NORM(h)) if not _blank(h) else None
            if c and c["source"] not in idx:
                idx[c["source"]] = ci
        missing = [c["header"] for c in cols if c["source"] not in idx]
        if missing:
            err(ERROR, "HEADERS", f"Sheet '{spec['sheet']}': missing column heading(s): {', '.join(missing)}",
                spec["sheet"], "Copy your data into a fresh template downloaded from this application")
            continue
        rows = []
        for rn, vals in matrix[hi + 1:]:
            if all(_blank(vals[idx[c["source"]]]) if idx[c["source"]] < len(vals) else True for c in cols):
                continue
            row = OrderedDict()
            ok = True
            for c in cols:
                ci = idx[c["source"]]
                raw = vals[ci] if ci < len(vals) else None
                v, e = normalize_cell(c, raw, schema)
                if e:
                    err(ERROR, "BAD_VALUE", f"{c['header']}: {e}", f"{spec['sheet']} row {rn}",
                        "Correct the value in the template")
                    ok = False
                if v == "" and c.get("default") not in (None, ""):
                    v, _ = normalize_cell(c, c["default"], schema)
                if v == "" and c.get("required") and not e:
                    if c["source"] in DERIVABLE:
                        deferred.append((row, c, spec))        # may still be worked out from other cells
                    else:
                        err(ERROR, "REQUIRED", f"{c['header']} is required", f"{spec['sheet']} row {rn}",
                            "Fill the mandatory (orange) column")
                    ok = False
                row[c["source"]] = v
            row["__file"] = fname
            row["__ref"] = f"{spec['sheet']} row {rn}"
            rows.append(row)
        res.tables[key] = rows
        res.counts[key] = len(rows)
    if not found:
        err(ERROR, "NO_SHEETS", "No GSTR-1 sheets were recognised in this workbook", "",
            "Use the template downloaded from this application")
    # ---- complete what can be worked out (place of supply, invoice value, HSN tax) then check the rest
    derive_template(res, ctx, fname)
    for row, c, spec in deferred:
        if row.get(c["source"], "") == "":
            err(ERROR, "REQUIRED", f"{c['header']} is required", row.get("__ref", spec["sheet"]),
                "Fill the mandatory (orange) column")
    # ---- rule validation of the rows read
    if res.errors and skip_errors:
        _skip_bad_rows(res, ctx, fname)
    elif not res.errors:
        for key, rows in res.tables.items():
            res.issues.extend(validate_table(key, rows, ctx))
        res.issues.extend(_validate_docs(res.docs, ctx, fname))
    return res


DERIVABLE = {"pos", "invoice_value", "note_value"}
_POS_SUPPLIER_TABLES = {"b2cs", "b2csa", "at", "atadj", "ata", "atadja"}
_INTER_TABLES = {"b2cl", "b2cla", "exp", "expa", "cdnur", "cdnura"}


def _inter_share(tables: Dict[str, List[dict]], ctx: Ctx) -> Dict[str, Decimal]:
    """Share of inter-state taxable value per rate in the template (used to split HSN tax into IGST / CGST+SGST)."""
    inter, total = defaultdict(lambda: ZERO), defaultdict(lambda: ZERO)
    for key in ("b2b", "b2ba", "cdnr", "cdnra", "b2cl", "b2cla", "exp", "expa", "cdnur", "cdnura", "b2cs", "b2csa"):
        for r in tables.get(key, []):
            if no_tax_row(r) or not r.get("taxable"):
                continue
            t = D(r.get("taxable"))
            rate = plain(D(r.get("rate")))
            is_inter = (key in _INTER_TABLES or r.get("invoice_type") in IGST_TYPES
                        or r.get("note_supply_type") in IGST_TYPES or (r.get("pos") and r.get("pos") != ctx.state))
            total[rate] += t
            total["*"] += t
            if is_inter:
                inter[rate] += t
                inter["*"] += t
    out = {k: (inter[k] / total[k] if total[k] else ZERO) for k in total}
    return out


def derive_template(res: "TemplateResult", ctx: Ctx, fname: str) -> None:
    """Fill blank cells that follow from other cells, so the client needs to supply the minimum:
    * Place of supply        <- state code of the recipient GSTIN (unregistered B2CS / advances: supplier's state)
    * Invoice / note value   <- sum of taxable value + tax + cess of the invoice's rows
    * HSN IGST/CGST/SGST     <- taxable x rate, split by the inter/intra-state share of that rate in this file"""
    def info(code, msg, sheet, action):
        res.issues.append(Issue(INFO, code, msg, fname, "", action, 0, sheet))

    share = _inter_share(res.tables, ctx)
    for key, rows in res.tables.items():
        meta, spec = TABLE_META.get(key), ctx.schema["tables"].get(key)
        if not meta or not spec:
            continue
        srcs = {c["source"] for c in spec["columns"]}
        sheet = spec["sheet"]
        n_gstin = n_sup = 0
        if "pos" in srcs and ctx.auto_pos:
            for r in rows:
                if r.get("pos"):
                    continue
                p = autofill.pos_from_gstin(r.get("recipient_gstin", ""))
                if p:
                    r["pos"], n_gstin = p, n_gstin + 1
                elif key in _POS_SUPPLIER_TABLES:
                    r["pos"], n_sup = ctx.state, n_sup + 1
        if n_gstin:
            info("POS_FROM_GSTIN", f"{sheet}: place of supply filled from the recipient GSTIN in {n_gstin} row(s)",
                 sheet, "Change the cell if the goods were delivered to a different state")
        if n_sup:
            info("POS_ASSUMED", f"{sheet}: place of supply was blank in {n_sup} row(s) and was set to the supplier's "
                 f"state ({ctx.state}) - confirm", sheet, "Enter the real place of supply where it differs")
        # ---- invoice / note value
        if meta.get("family") in ("invoice", "note"):
            idc, vc = meta["id"], meta["value"]
            groups = defaultdict(list)
            for r in rows:
                groups[r.get(idc, "")].append(r)
            n_val = 0
            for idv, grp in groups.items():
                have = next((x.get(vc) for x in grp if x.get(vc) not in (None, "")), "")
                if have:
                    for x in grp:
                        if x.get(vc) in (None, ""):
                            x[vc] = have
                    continue
                comp = ZERO
                for x in grp:
                    t = D(x.get("taxable"))
                    comp += t + D(x.get("cess"))
                    if not no_tax_row(x):
                        comp += q2(t * D(x.get("rate")) / 100)
                for x in grp:
                    x[vc] = str(q2(comp))
                n_val += 1
            if n_val:
                info("VALUE_CALC", f"{sheet}: {meta['what'].lower()} value calculated (taxable + tax + cess) for "
                     f"{n_val} {meta['what'].lower()}(s)", sheet, "Check against the actual invoice value if it includes round-off")
        # ---- HSN tax heads
        if key == "hsn":
            n_hsn = 0
            for r in rows:
                if not r.get("description") and r.get("hsn"):
                    r["description"] = hsn_mod.short(hsn_mod.describe(r["hsn"]))
                t, rt = D(r.get("taxable")), D(r.get("rate"))
                if not t or not rt or D(r.get("igst")) or D(r.get("cgst")) or D(r.get("sgst")):
                    continue
                tot = q2(t * rt / 100)
                sh = share.get(plain(rt), share.get("*", ZERO))
                ig = q2(tot * sh)
                half = q2((tot - ig) / 2)
                r["igst"], r["cgst"], r["sgst"] = str(ig), str(half), str(half)
                if not r.get("total_value") or D(r.get("total_value")) == 0:
                    r["total_value"] = str(q2(t + tot + D(r.get("cess"))))
                n_hsn += 1
            if n_hsn:
                info("HSN_TAX_CALC", f"{sheet}: IGST/CGST/SGST calculated as taxable x rate for {n_hsn} row(s) "
                     "(split by the inter-state share of that rate in this file)", sheet,
                     "Compare the HSN tax total with the GSTR-1 Summary")


_ROWREF = re.compile(r" row \d+$")


def _drop_rows(res: "TemplateResult", refs: set) -> int:
    """Remove rows whose __ref is in refs; for invoices / notes the whole invoice (all its rows) goes."""
    n = 0
    for key, rows in list(res.tables.items()):
        meta = TABLE_META.get(key, {})
        hit = [r for r in rows if r.get("__ref") in refs]
        if not hit:
            continue
        ids = {r.get(meta["id"]) for r in hit} if meta.get("family") in ("invoice", "note") else set()
        keep = [r for r in rows if r not in hit and not (ids and r.get(meta["id"]) in ids)]
        n += len(rows) - len(keep)
        res.tables[key] = keep
        res.counts[key] = len(keep)
    return n


def _skip_bad_rows(res: "TemplateResult", ctx: Ctx, fname: str) -> None:
    """skip_errors mode: leave out every row (whole invoice) that has an error and import the rest.
    File-level problems (headings, GSTIN / period of the file, Documents sheet) can never be skipped."""
    cell_issues = list(res.issues)
    skipped: List[Issue] = []
    for _ in range(6):
        errs = [i for i in cell_issues if i.severity == ERROR]
        if any(not _ROWREF.search(i.ref or "") for i in errs):
            res.issues = cell_issues + skipped
            return
        if errs:
            res.skipped += _drop_rows(res, {i.ref for i in errs})
            skipped += [Issue(WARNING, "ROW_SKIPPED", f"NOT IMPORTED - {i.message}", i.source_file, i.invoice_no,
                              "Correct this row in the client's file and upload it again", 0, i.ref) for i in errs]
            cell_issues = [i for i in cell_issues if i.severity != ERROR]
        rule: List[Issue] = []
        for key, rows in res.tables.items():
            rule.extend(validate_table(key, rows, ctx))
        rule.extend(_validate_docs(res.docs, ctx, fname))
        rerrs = [i for i in rule if i.severity == ERROR]
        if not rerrs:
            res.issues = cell_issues + skipped + rule
            return
        cell_issues = cell_issues + rerrs
    res.issues = cell_issues + skipped


def _read_docs_sheet(ws, spec, schema, fname, res):
    from .exporters import parse_docs_matrix
    matrix = _sheet_matrix(ws)
    try:
        series, problems = parse_docs_matrix(matrix, schema, fname)
    except ValueError as e:
        if not any(any(not _blank(v) for v in vals) for _, vals in matrix[1:]):
            return                                   # heading-only / empty Documents sheet
        res.issues.append(Issue(ERROR, "HEADERS", str(e).splitlines()[0], fname, "",
                                "Do not change the headings of the Documents sheet", 0, "Documents"))
        return
    for rn_msg in problems:
        res.issues.append(Issue(ERROR, "DOCS_ROW", rn_msg, fname, "", "Correct the Documents sheet", 0,
                                "Documents"))
    res.docs = series
    res.counts["docs"] = len(series)


def _validate_docs(docs, ctx, fname):
    out = []
    for d in docs:
        if d.cancelled > d.total:
            out.append(Issue(ERROR, "DOCS_CANCEL", "Cancelled documents exceed total issued", fname, d.from_no,
                             "Correct the Documents sheet", 0, "Documents"))
    return out


# ====================================================================== validation
def _ident(r, meta):
    return r.get(meta.get("id", ""), "") if meta.get("id") else ""


def validate_table(key: str, rows: List[dict], ctx: Ctx) -> List[Issue]:
    spec = ctx.schema["tables"][key]
    meta = TABLE_META[key]
    hdr = {c["source"]: c["header"] for c in spec["columns"]}
    out: List[Issue] = []

    def add(r, sev, code, msg, action=""):
        out.append(Issue(sev, code, msg, r.get("__file", ""), _ident(r, meta), action, 0, r.get("__ref", "")))

    fam = meta["family"]
    for r in rows:
        g = r.get
        if g("recipient_gstin"):
            lvl, msg = validate_gstin(g("recipient_gstin"))
            if lvl == "error":
                add(r, ERROR, "GSTIN_INVALID", f"Recipient GSTIN {g('recipient_gstin')}: {msg}",
                    "Correct the GSTIN (unregistered customers belong in B2CL/B2CS)")
            elif lvl == "warn":
                add(r, WARNING, "GSTIN_CHECKDIGIT", f"Recipient GSTIN {g('recipient_gstin')}: {msg}",
                    "Verify the GSTIN on the GST portal")
            if g("recipient_gstin") == ctx.gstin:
                add(r, ERROR, "GSTIN_SELF", "Recipient GSTIN equals the supplier GSTIN", "Correct the recipient")
        if g("ecom_gstin"):
            lvl, msg = validate_gstin(g("ecom_gstin"))
            if lvl == "error":
                add(r, ERROR, "ECOM_GSTIN", f"E-Commerce GSTIN {g('ecom_gstin')}: {msg}", "Correct the GSTIN")
        for idc in ("invoice_no", "note_no", "orig_invoice_no", "orig_note_no"):
            if g(idc) and not _INVNO.match(g(idc)):
                add(r, ERROR, "INV_NO_FORMAT", f"{hdr.get(idc, idc)} '{g(idc)}' must be max 16 characters using "
                    "letters, digits, '/' and '-' only", "Correct the number exactly as issued")
        for dc in ("invoice_date", "note_date"):
            if g(dc):
                try:
                    dd = date.fromisoformat(g(dc))
                except ValueError:
                    continue
                if dd > ctx.end:
                    add(r, ERROR, "DATE_AFTER_PERIOD", f"{hdr.get(dc, dc)} {fmt_date(g(dc), 'DD-MM-YYYY')} is after the "
                        f"end of the return period ({ctx.end:%d-%m-%Y})", "Report it in the correct period")
                elif dd < ctx.start and not meta.get("amend"):
                    add(r, WARNING, "DATE_BEFORE_PERIOD", f"{hdr.get(dc, dc)} {fmt_date(g(dc), 'DD-MM-YYYY')} is before "
                        f"the return period", "Confirm it was not reported earlier (late reporting is allowed)")
        if g("rate") not in (None, "") and ctx.rates and D(g("rate")) not in ctx.rates:
            add(r, WARNING, "RATE_UNUSUAL", f"Rate {g('rate')}% is not one of the configured GST slabs",
                "Confirm the rate (Settings > valid rates)")
        amt_src = meta.get("amt", ("taxable",))[0] if fam not in ("exemp", "docs") else ""
        if amt_src and g(amt_src) not in (None, ""):
            a = D(g(amt_src))
            if a < 0 and key not in ("b2cs", "b2csa", "hsn"):
                add(r, WARNING, "NEGATIVE", f"Negative {hdr.get(amt_src, amt_src)}", "Credit notes go to CDNR/CDNUR")
            if a == 0 and fam in ("invoice", "note"):
                add(r, WARNING, "ZERO_VALUE", f"{hdr.get(amt_src, amt_src)} is zero", "Remove the row or fill the value")
        # ---------------- table specific
        if key in ("b2cl", "b2cla"):
            if g("pos") and g("pos") == ctx.state:
                add(r, ERROR, "B2CL_INTRA", "B2CL supplies must be inter-state but the place of supply equals the "
                    "supplier's state", "Report intra-state B2C sales in B2CS")
            if key == "b2cl" and g("invoice_value") and D(g("invoice_value")) <= ctx.b2cl_threshold:
                add(r, WARNING, "B2CL_LOW", f"Invoice value is not above the B2CL threshold "
                    f"(Rs {ctx.b2cl_threshold:,.0f}) - normally belongs in B2CS", "Check the classification")
        if key in ("b2cs", "b2csa"):
            if g("supply_type") == "E" and not g("ecom_gstin"):
                add(r, ERROR, "B2CS_ECOM", "Type E (e-commerce) needs the E-Commerce GSTIN", "Enter the operator's GSTIN")
            if g("supply_type") == "OE" and g("ecom_gstin"):
                add(r, ERROR, "B2CS_ECOM", "Type OE must not have an E-Commerce GSTIN",
                    "Use Type E, or clear the E-Commerce GSTIN")
            if g("taxable") and D(g("taxable")) < 0:
                add(r, INFO, "B2CS_NEG", "Negative taxable value (net of credit notes)", "Confirm it is intended")
        if key in ("exp", "expa"):
            if bool(g("ship_bill_no")) != bool(g("ship_bill_date")):
                add(r, WARNING, "SB_PARTIAL", "Shipping bill number and date should be given together",
                    "Fill both or leave both blank")
            if g("port_code") and not re.match(r"^[A-Za-z0-9]{6}$", g("port_code")):
                add(r, WARNING, "PORT_CODE", f"Port code '{g('port_code')}' should be 6 characters", "Check the port code")
            if g("ship_bill_no") and not re.match(r"^\d{1,7}$", g("ship_bill_no")):
                add(r, WARNING, "SB_NO", f"Shipping bill number '{g('ship_bill_no')}' should be up to 7 digits",
                    "Check the shipping bill number")
            if g("export_type") == "WPAY" and g("rate") not in (None, "") and D(g("rate")) == 0:
                add(r, WARNING, "EXP_WPAY_ZERO", "Export with payment of tax but rate is 0", "Check type / rate")
        if key in ("cdnur", "cdnura") and g("ur_type") == "B2CL" and g("pos") == ctx.state:
            add(r, ERROR, "CDNUR_INTRA", "UR type B2CL needs an inter-state place of supply",
                "Intra-state notes are adjusted in B2CS")
        if key in ("at", "atadj", "ata", "atadja"):
            if g(meta["amt"][0]) not in (None, "") and D(g(meta["amt"][0])) <= 0:
                add(r, WARNING, "ADV_ZERO", "Advance amount is zero or negative", "Check the amount")
        if key == "exemp":
            for s in meta["sums"]:
                if g(s) not in (None, "") and D(g(s)) < 0:
                    add(r, ERROR, "EXEMP_NEG", f"{hdr[s]} cannot be negative", "Correct the amount")
        if key == "hsn":
            h = g("hsn", "")
            if not h.isdigit() or len(h) not in (4, 6, 8):
                add(r, ERROR, "HSN_INVALID", f"HSN/SAC '{h}' must be 4, 6 or 8 digits "
                    "(enter it as text so leading zeros are kept)", "Correct the HSN/SAC")
            elif h[:2] == "00":
                add(r, WARNING, "HSN_PLACEHOLDER", f"HSN '{h}' is not a real HSN/SAC code (chapter 00 does not exist) - "
                    f"it looks like a placeholder in the client's billing software"
                    + (f". {r.get('__ref')}" if r.get("__ref") else ""),
                    "Get the correct HSN for these products from the client, fix it in the client's product master, "
                    "and re-export the report (or correct the HSN here before filing)")
            elif len(h) < ctx.hsn_min:
                add(r, ERROR, "HSN_DIGITS", f"HSN/SAC '{h}' has {len(h)} digits; at least {ctx.hsn_min} digits are "
                    "required for this taxpayer", "Give a longer HSN/SAC (Settings > minimum HSN digits)")
            if h and not h.startswith("99") and g("uqc") in ("", None, "NA"):
                add(r, WARNING, "NO_UQC", f"UQC missing/NA for goods HSN {h}; OTH-OTHERS will be used", "Pick the unit")
            if h and not h.startswith("99") and g("qty") not in (None, "") and D(g("qty")) <= 0:
                add(r, WARNING, "NO_QTY", f"Quantity missing for HSN {h}", "Enter the total quantity")
            t, rt = D(g("taxable")), D(g("rate"))
            tx = D(g("igst")) + D(g("cgst")) + D(g("sgst"))
            exp_tax = q2(t * rt / 100)
            if abs(exp_tax - tx) > max(ctx.tol, abs(exp_tax) * Decimal("0.001")):
                add(r, WARNING, "HSN_TAX", f"HSN {h}: tax {tx} differs from taxable x rate ({rt}%) = {exp_tax}",
                    "Check taxable value / rate / tax amounts")
            if g("cgst") and g("sgst") and D(g("cgst")) != D(g("sgst")):
                add(r, WARNING, "HSN_CGST_SGST", f"HSN {h}: CGST and SGST differ", "CGST and SGST should be equal")
            if g("total_value") not in (None, "") and D(g("total_value")) + ctx.tol < t + tx + D(g("cess")):
                add(r, WARNING, "HSN_TOTAL", f"HSN {h}: total value is less than taxable value + taxes",
                    "Check the total value")
        if key in ("b2ba", "b2cla", "expa") and not g("orig_invoice_no"):
            add(r, ERROR, "ORIG_MISSING", "Original invoice details are required for an amendment", "Fill them")

    # ---------------- group level (invoices / notes)
    if fam in ("invoice", "note"):
        idc, vc = meta["id"], meta["value"]
        groups = defaultdict(list)
        for r in rows:
            groups[r.get(idc, "")].append(r)
        cons = [c["source"] for c in spec["columns"] if c["source"] not in LINE_COLS]
        for idv, grp in groups.items():
            first = grp[0]
            for s in cons:
                vals = {x.get(s, "") for x in grp}
                if len(vals) > 1:
                    add(first, ERROR, "MIXED_VALUES", f"{meta['what']} {idv}: different '{hdr[s]}' on different rows "
                        f"({', '.join(sorted(map(str, vals))[:4])})", "Repeat identical invoice details on all rows")
            seen = defaultdict(int)
            for x in grp:
                seen[(x.get("rate"), x.get("ecom_gstin"), x.get("applicable_pct"))] += 1
            if any(n > 1 for n in seen.values()):
                add(first, WARNING, "DUP_ROW", f"{meta['what']} {idv}: more than one row with the same rate "
                    "(they will be added together)", "Merge the rows if this is a mistake")
            # value vs computed
            if grp[0].get(vc) not in (None, ""):
                comp = ZERO
                for x in grp:
                    t = D(x.get("taxable"))
                    comp += t + D(x.get("cess"))
                    if not no_tax_row(x):
                        comp += q2(t * D(x.get("rate")) / 100)
                val = D(grp[0].get(vc))
                if abs(val - comp) > ctx.tol:
                    add(first, WARNING, "VALUE_DIFF", f"{meta['what']} {idv}: value {val} differs from taxable + tax "
                        f"+ cess = {q2(comp)}", "Check that all rate-wise rows of the invoice are entered")
    return out


def cross_checks(sets: List[dict], lines, ctx: Ctx) -> List[Issue]:
    occ = defaultdict(set)
    ns_of = {"b2b": "inv", "b2cl": "inv", "exp": "inv", "cdnr": "note", "cdnur": "note"}
    for s in sets:
        for key, ns in ns_of.items():
            idc = TABLE_META[key]["id"]
            for r in s["tables"].get(key, []):
                if r.get(idc):
                    occ[(ns, norm_invoice_no(r[idc]))].add((f"{s['filename']} [{spec_sheet(ctx, key)}]", r[idc]))
    for l in lines or []:
        if l.review_status != STATUS_APPROVED or not l.invoice_no:
            continue
        if l.note_type == "" and l.section in ("B2B", "B2CL", "EXP"):
            occ[("inv", norm_invoice_no(l.invoice_no))].add(("approved invoice lines", l.invoice_no))
        elif l.note_type and l.section in ("CDNR", "CDNUR"):
            occ[("note", norm_invoice_no(l.invoice_no))].add(("approved invoice lines", l.invoice_no))
    out = []
    for (ns, idn), where in occ.items():
        if len(where) > 1:
            disp = sorted(where)[0][1]
            out.append(Issue(ERROR, "DUP_ACROSS_SOURCES",
                             f"{'Invoice' if ns == 'inv' else 'Note'} {disp} is reported in more than one place: "
                             + "; ".join(sorted(w for w, _ in where)), "", disp,
                             "Keep it in one place only (remove it from the other file / reject the line)", 0, ""))
    return out


def spec_sheet(ctx, key):
    return ctx.schema["tables"][key]["sheet"]


def validate_sets(sets: List[dict], lines, ctx: Ctx) -> List[Issue]:
    """Rule validation of all stored template rows + cross checks (re-run on every refresh)."""
    out: List[Issue] = []
    for s in sets:
        for key, rows in s["tables"].items():
            if key in TABLE_META and key in ctx.schema["tables"]:
                out.extend(validate_table(key, rows, ctx))
    out.extend(cross_checks(sets, lines, ctx))
    return out


# ====================================================================== tables from invoice lines
def _uqc_full(schema, text, hsn):
    if hsn.startswith("99"):
        return "NA"
    codes = {u.split("-")[0]: u for u in schema.get("uqc", [])}
    code = (text or "").upper().split("-")[0].strip()
    return codes.get(code, schema.get("default_uqc", "OTH-OTHERS"))


def lines_to_tables(lines, client: dict, settings: dict, schema: dict) -> Dict[str, List[dict]]:
    own = client["gstin"][:2]
    T: Dict[str, List[dict]] = {k: [] for k in schema["tables"] if k != "docs"}
    inv_types = {c.lower(): c for c in _choices(schema, "b2b", "invoice_type")}
    inv_types.setdefault("regular", "Regular B2B")
    appr = [l for l in lines if l.review_status == STATUS_APPROVED and l.section in IMPL_SECTIONS]

    def tag(l):
        return {"__file": l.source_file, "__ref": l.source_ref}

    def buckets(ls, with_ecom=True):
        b: "OrderedDict[tuple, list]" = OrderedDict()
        for x in ls:
            b.setdefault((x.rate, x.ecom_gstin if with_ecom else ""), []).append(x)
        return b

    for _, grp in group_invoices(appr).items():
        val = invoice_value(grp)
        by_sec: "OrderedDict[str, list]" = OrderedDict()
        for l in grp:
            by_sec.setdefault(l.section, []).append(l)
        for sec, ls in by_sec.items():
            f = ls[0]
            itype = inv_types.get((f.invoice_type or "").lower(), "Regular B2B")
            if sec == "B2B":
                for (rate, ecom), xs in buckets(ls).items():
                    T["b2b"].append(dict(
                        recipient_gstin=f.recipient_gstin, recipient_name=f.recipient_name,
                        invoice_no=f.invoice_no.strip(), invoice_date=f.invoice_date, invoice_value=str(q2(val)),
                        pos=f.pos, reverse_charge=f.reverse_charge or "N", applicable_pct="", invoice_type=itype,
                        ecom_gstin=ecom, rate=plain(rate), taxable=str(q2(sum((x.taxable for x in xs), ZERO))),
                        cess=str(q2(sum((x.cess for x in xs), ZERO))), **tag(f)))
            elif sec == "B2CL":
                for (rate, ecom), xs in buckets(ls).items():
                    T["b2cl"].append(dict(
                        invoice_no=f.invoice_no.strip(), invoice_date=f.invoice_date, invoice_value=str(q2(val)),
                        pos=f.pos, applicable_pct="", rate=plain(rate),
                        taxable=str(q2(sum((x.taxable for x in xs), ZERO))),
                        cess=str(q2(sum((x.cess for x in xs), ZERO))), ecom_gstin=ecom, **tag(f)))
            elif sec == "EXP":
                for (rate, _e), xs in buckets(ls, False).items():
                    et = f.export_type or ("WPAY" if any(x.tax_total for x in ls) else "WOPAY")
                    T["exp"].append(dict(
                        export_type=et, invoice_no=f.invoice_no.strip(), invoice_date=f.invoice_date,
                        invoice_value=str(q2(val)), port_code=f.port_code, ship_bill_no=f.ship_bill_no,
                        ship_bill_date=f.ship_bill_date, applicable_pct="", rate=plain(rate),
                        taxable=str(q2(sum((x.taxable for x in xs), ZERO))),
                        cess=str(q2(sum((x.cess for x in xs), ZERO))), **tag(f)))
            elif sec == "CDNR":
                for (rate, _e), xs in buckets(ls, False).items():
                    T["cdnr"].append(dict(
                        recipient_gstin=f.recipient_gstin, recipient_name=f.recipient_name,
                        note_no=f.invoice_no.strip(), note_date=f.invoice_date, note_type=f.note_type or "C",
                        pos=f.pos, reverse_charge=f.reverse_charge or "N", note_supply_type=itype,
                        note_value=str(q2(val)), applicable_pct="", rate=plain(rate),
                        taxable=str(q2(sum((x.taxable for x in xs), ZERO))),
                        cess=str(q2(sum((x.cess for x in xs), ZERO))), **tag(f)))
            elif sec == "CDNUR":
                urt = ("EXPWOP" if f.export_type == "WOPAY" else "EXPWP") if (f.export_type or f.pos == "96") else "B2CL"
                for (rate, _e), xs in buckets(ls, False).items():
                    T["cdnur"].append(dict(
                        ur_type=urt, note_no=f.invoice_no.strip(), note_date=f.invoice_date,
                        note_type=f.note_type or "C", pos=f.pos, note_value=str(q2(val)), applicable_pct="",
                        rate=plain(rate), taxable=str(q2(sum((x.taxable for x in xs), ZERO))),
                        cess=str(q2(sum((x.cess for x in xs), ZERO))), **tag(f)))
            elif sec == "B2CS":
                sign = Decimal(-1) if f.note_type == "C" else Decimal(1)
                for (rate, ecom), xs in buckets(ls).items():
                    T["b2cs"].append(dict(
                        supply_type="E" if ecom else "OE", pos=f.pos, applicable_pct="", rate=plain(rate),
                        taxable=str(q2(sign * sum((x.taxable for x in xs), ZERO))),
                        cess=str(q2(sign * sum((x.cess for x in xs), ZERO))), ecom_gstin=ecom, **tag(f)))
            elif sec == "NIL":
                sign = Decimal(-1) if f.note_type == "C" else Decimal(1)
                reg = bool(f.recipient_gstin) and validate_gstin(f.recipient_gstin)[0] != "error"
                inter = bool(f.pos) and f.pos != own
                desc = (f"{'Inter' if inter else 'Intra'}-State supplies to {'registered' if reg else 'unregistered'} "
                        "persons")
                for x in ls:
                    row = dict(description=desc, nil_rated="0.00", exempted="0.00", non_gst="0.00", **tag(x))
                    col = {"Exempted": "exempted", "Non-GST": "non_gst"}.get(x.nil_kind, "nil_rated")
                    row[col] = str(q2(sign * x.taxable))
                    T["exemp"].append(row)
            elif sec in ("AT", "ATADJ"):
                key, amt = ("at", "gross_advance") if sec == "AT" else ("atadj", "gross_adjusted")
                for (rate, _e), xs in buckets(ls, False).items():
                    T[key].append({"pos": f.pos, "applicable_pct": "", "rate": plain(rate),
                                   amt: str(q2(sum((x.line_total for x in xs), ZERO))),
                                   "cess": str(q2(sum((x.cess for x in xs), ZERO))), **tag(f)})
        # ---- HSN (all outward lines of the invoice except advances)
        for l in grp:
            if not l.hsn or l.section in ("AT", "ATADJ"):
                continue
            sign = Decimal(-1) if l.note_type == "C" else Decimal(1)
            T["hsn"].append(dict(
                hsn=l.hsn, description=(l.description or hsn_mod.describe(l.hsn))[:30], uqc=_uqc_full(schema, l.uqc, l.hsn),
                qty=plain(sign * l.qty), total_value=str(q2(sign * l.line_total)), rate=plain(l.rate),
                taxable=str(q2(sign * l.taxable)), igst=str(q2(sign * l.igst)), cgst=str(q2(sign * l.cgst)),
                sgst=str(q2(sign * l.sgst)), cess=str(q2(sign * l.cess)), **tag(l)))
    return T


# ====================================================================== aggregation
def _fmt_sum(col, v: Decimal) -> str:
    return plain(v) if col.get("kind") == "qty" else str(q2(v))


def aggregate(key: str, rows: List[dict], schema: dict) -> List[dict]:
    meta = TABLE_META[key]
    cols = schema["tables"][key]["columns"]
    cmap = {c["source"]: c for c in cols}
    srcs = [c["source"] for c in cols]
    sums = [s for s in meta.get("sums", ()) if s in cmap]
    nokey = meta.get("nokey", ())
    default_uqc = schema.get("default_uqc", "OTH-OTHERS")
    out: "OrderedDict[tuple, dict]" = OrderedDict()
    for r in rows:
        r = dict(r)
        if key == "hsn":
            if not r.get("uqc"):
                r["uqc"] = "NA" if r.get("hsn", "").startswith("99") else default_uqc
            if r.get("total_value") in ("", None):
                r["total_value"] = str(D(r.get("taxable")) + D(r.get("igst")) + D(r.get("cgst")) + D(r.get("sgst"))
                                       + D(r.get("cess")))
        k = tuple(str(r.get(s, "")) for s in srcs if s not in sums and s not in nokey)
        if k not in out:
            out[k] = {s: r.get(s, "") for s in srcs}
            out[k]["__file"] = r.get("__file", "")
            out[k]["__ref"] = r.get("__ref", "")
            for s in sums:
                out[k][s] = ZERO
        for s in sums:
            out[k][s] += D(r.get(s))
    res = []
    for row in out.values():
        for s in sums:
            row[s] = _fmt_sum(cmap[s], row[s])
        res.append(row)
    if key == "exemp":
        order = _choices(schema, "exemp", "description")
        have = {r["description"]: r for r in res}
        if any(D(r[s]) != 0 for r in res for s in sums):
            res = [have.get(d_) or {"description": d_, **{s: "0.00" for s in sums}, "__file": "", "__ref": ""}
                   for d_ in order]
        else:
            res = []
    return res


def build_tables(lines, sets: List[dict], client: dict, settings: dict, schema: dict):
    """-> (pre, final): pre = concatenated rows, final = aggregated rows ready for CSV."""
    pre: Dict[str, List[dict]] = {k: [] for k in schema["tables"] if k != "docs"}
    for s in sets:
        for key, rows in s["tables"].items():
            if key in pre:
                pre[key].extend(rows)
    for key, rows in lines_to_tables(lines, client, settings, schema).items():
        pre[key].extend(rows)
    final = {k: aggregate(k, rows, schema) for k, rows in pre.items()}
    return pre, final


# ====================================================================== tax summary
def row_tax(key: str, row: dict, own_state: str) -> dict:
    meta = TABLE_META[key]
    rate, cess = D(row.get("rate")), D(row.get("cess"))
    if meta.get("gross"):
        gross = D(row.get(meta["amt"][0]))
        tax = q2(gross * rate / (100 + rate)) if rate else ZERO
        taxable = gross - tax
    else:
        taxable = D(row.get("taxable"))
        tax = q2(taxable * rate / 100)
    mode = meta.get("tax")
    if mode == "igst":
        inter = True
    elif mode == "pos_type":
        inter = (row.get("pos") != own_state) or (row.get(meta.get("type_col", ""), "") in IGST_TYPES)
    else:
        inter = bool(row.get("pos")) and row.get("pos") != own_state
    if no_tax_row(row):
        tax = ZERO
    if inter:
        igst, cgst, sgst = tax, ZERO, ZERO
    else:
        cgst = q2(tax / 2)
        sgst = tax - cgst
        igst = ZERO
    sign = Decimal(-1) if row.get("note_type") == "C" else Decimal(1)
    return {"taxable": sign * taxable, "igst": sign * igst, "cgst": sign * cgst, "sgst": sign * sgst,
            "cess": sign * cess}


def summarize(final: Dict[str, List[dict]], own_state: str, schema: dict) -> List[dict]:
    out = []
    for key, rows in final.items():
        if not rows:
            continue
        meta = TABLE_META[key]
        t = {k: ZERO for k in ("taxable", "igst", "cgst", "sgst", "cess")}
        for r in rows:
            if meta["family"] == "hsn":
                x = {"taxable": D(r.get("taxable")), "igst": D(r.get("igst")), "cgst": D(r.get("cgst")),
                     "sgst": D(r.get("sgst")), "cess": D(r.get("cess"))}
            elif meta["family"] == "exemp":
                x = {"taxable": sum((D(r.get(s)) for s in meta["sums"]), ZERO), "igst": ZERO, "cgst": ZERO,
                     "sgst": ZERO, "cess": ZERO}
            else:
                x = row_tax(key, r, own_state)
            for k in t:
                t[k] += x[k]
        docs = len({r.get(meta["id"]) for r in rows}) if meta.get("id") else len(rows)
        out.append({"key": key, "title": schema["tables"][key]["title"], "rows": len(rows), "docs": docs,
                    **{k: q2(v) for k, v in t.items()},
                    "total": q2(sum(t.values(), ZERO)), "amend": bool(meta.get("amend"))})
    return out


# ====================================================================== reconciliation
def reconcile(lines, pre, final, docs, summary, settings) -> List[dict]:
    rows = []

    def row(check, a, b, note="", tol=Decimal("1"), info=False):
        diff = q2(a - b)
        st = "INFO" if info else ("OK" if abs(diff) <= tol else "DIFFERENCE")
        rows.append({"check": check, "source": str(q2(a)), "processed": str(q2(b)), "difference": str(diff),
                     "status": st, "note": note})

    # 1. aggregation integrity per table
    for key, frows in final.items():
        meta = TABLE_META[key]
        if not frows and not pre.get(key):
            continue
        for s in meta.get("amt", ()):
            row(f"{key.upper()}: {s} - input rows vs generated CSV",
                sum((D(r.get(s)) for r in pre.get(key, [])), ZERO), sum((D(r.get(s)) for r in frows), ZERO),
                "Aggregation check - must be zero")
    # 2. extracted lines
    live = [l for l in lines if l.review_status != "Rejected"]
    appr = [l for l in lines if l.review_status == STATUS_APPROVED]
    per_file = defaultdict(lambda: ZERO)
    for l in live:
        per_file[l.source_file] += l.taxable
    for f, v in per_file.items():
        row(f"Taxable value - extracted vs approved ({f})", v,
            sum((l.taxable for l in appr if l.source_file == f), ZERO), "Difference = lines not yet approved")
    no_hsn = sum((l.taxable for l in appr if not l.hsn and l.section not in ("AT", "ATADJ")), ZERO)
    if appr:
        row("Approved invoice-line taxable value without HSN (not in HSN table)", ZERO, no_hsn,
            "Should be zero - fill the HSN/SAC on these lines")
    manual = sum((l.taxable for l in appr if l.section not in IMPL_SECTIONS), ZERO)
    if manual:
        row("Approved lines in sections not generated (manual reporting)", ZERO, manual,
            "Report these in the Offline Tool manually")
    # 3. HSN vs outward supplies
    if final.get("hsn"):
        out_taxable = ZERO
        for s in summary:
            if s["key"] in ("b2b", "b2cl", "b2cs", "exp", "exemp", "cdnr", "cdnur"):
                if s["key"] in ("cdnr", "cdnur"):
                    for r in final[s["key"]]:
                        sign = Decimal(-1) if r.get("note_type") == "C" else Decimal(1)
                        out_taxable += sign * D(r.get("taxable"))
                else:
                    out_taxable += s["taxable"]
        hsn_taxable = sum((D(r.get("taxable")) for r in final["hsn"]), ZERO)
        row("HSN taxable value vs outward supplies (B2B+B2CL+B2CS+Exports+Nil/Exempt/Non-GST +/- notes)",
            out_taxable, hsn_taxable,
            "A difference needs explanation (missing HSN rows / mismatched tables / advances)")
    elif any(final.get(k) for k in ("b2b", "b2cl", "b2cs", "exp")):
        rows.append({"check": "HSN table", "source": "-", "processed": "-", "difference": "-", "status": "DIFFERENCE",
                     "note": "Outward supplies are present but there is no HSN summary - Table 12 is mandatory"})
    # 4. documents vs invoices
    inv_docs = [d for d in docs if "invoice" in d.nature.lower() and "inward" not in d.nature.lower()]
    listed = set()
    for key in ("b2b", "b2cl", "exp"):
        listed |= {norm_invoice_no(r.get("invoice_no")) for r in final.get(key, [])}
    if inv_docs:
        reg_net = sum(d.net for d in inv_docs)
        tot_listed = len(listed)
        # invoices that are reported only as totals (Table 8 nil / exempt / non-GST and B2CS) - not listed one by one
        only_totals = {(norm_invoice_no(l.invoice_no)) for l in lines
                       if l.review_status == STATUS_APPROVED and l.invoice_no and l.section in ("NIL", "B2CS")} - listed
        diff = reg_net - tot_listed - len(only_totals)
        rows.append({"check": "Document register vs invoices (listed B2B/B2CL/EXP + reported as totals NIL/B2CS)",
                     "source": f"{tot_listed} + {len(only_totals)}", "processed": str(reg_net), "difference": str(diff),
                     "status": "OK" if diff == 0 else "INFO",
                     "note": ("Every document in the register is accounted for" if diff == 0 else
                              "Difference = numbers in the register that are not in the approved invoices (cancelled / "
                              "not yet approved / missing)")})
    elif not docs:
        rows.append({"check": "Documents register", "source": "-", "processed": "0", "difference": "-",
                     "status": "DIFFERENCE", "note": "No document series entered (Table 13 is mandatory)"})
    return rows


# ====================================================================== CSV output
def format_cell(col: dict, value, schema: dict) -> str:
    kind = col.get("kind", "text")
    v = "" if value is None else value
    if v == "":
        return ""
    if kind == "date":
        return fmt_date(v, schema.get("date_format", "DD-Mon-YY"))
    if kind == "pos":
        return pos_label(v)
    if kind == "month":
        return fmt_month(v, schema)
    if kind == "money":
        return str(q2(D(v)))
    if kind in ("rate", "qty"):
        return plain(D(v))
    if kind == "int":
        return str(int(D(v)))
    om = col.get("output_map")
    if om:
        return om.get(str(v), str(v))
    return str(v)


def write_csv(path, key: str, rows: List[dict], schema: dict) -> None:
    cols = schema["tables"][key]["columns"]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([c["header"] for c in cols])
        for r in rows:
            w.writerow([format_cell(c, r.get(c["source"], ""), schema) for c in cols])


# ====================================================================== official header check
def _read_header(path: Path):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            with open(path, encoding=enc, newline="") as f:
                return [h.strip() for h in next(csv.reader(f), [])]
        except UnicodeDecodeError:
            continue
    return []


def find_official_files(folder, schema: dict) -> Dict[str, Optional[Path]]:
    folder = Path(folder)
    files = [p for p in folder.rglob("*.csv")] if folder.exists() else []
    by_stem = {}
    for p in files:
        by_stem.setdefault(_NORM(p.stem), p)
    out = {}
    for key, spec in schema["tables"].items():
        cands = [_NORM(Path(spec["filename"]).stem), _NORM(key)] + [_NORM(a) for a in spec.get("aliases", [])]
        out[key] = next((by_stem[c] for c in cands if c in by_stem), None)
    return out


def compare_headers(folder, schema: dict) -> List[dict]:
    res = []
    found = find_official_files(folder, schema)
    for key, spec in schema["tables"].items():
        mine = [c["header"] for c in spec["columns"]]
        p = found.get(key)
        if not p:
            res.append({"key": key, "file": "", "status": "not found", "official": [], "mine": mine,
                        "missing": [], "extra": [], "detail": "No official sample file for this table in the folder"})
            continue
        off = _read_header(p)
        if off == mine:
            res.append({"key": key, "file": p.name, "status": "match", "official": off, "mine": mine,
                        "missing": [], "extra": [], "detail": "Identical"})
            continue
        n_off, n_mine = [_NORM(x) for x in off], [_NORM(x) for x in mine]
        if sorted(n_off) == sorted(n_mine):
            res.append({"key": key, "file": p.name, "status": "adoptable", "official": off, "mine": mine,
                        "missing": [], "extra": [],
                        "detail": "Same columns; spelling/case/order differ - can be adopted automatically"})
        else:
            res.append({"key": key, "file": p.name, "status": "different", "official": off, "mine": mine,
                        "missing": [h for h in off if _NORM(h) not in n_mine],
                        "extra": [h for h in mine if _NORM(h) not in n_off],
                        "detail": "Columns differ - needs manual change of the schema (see lists)"})
    return res


def adopt_headers(schema: dict, comparison: List[dict]) -> int:
    """Rewrite headings/order for 'match' and 'adoptable' tables. Returns the number of verified tables."""
    n = 0
    for c in comparison:
        spec = schema["tables"][c["key"]]
        if c["status"] == "adoptable":
            bynorm = {_NORM(col["header"]): col for col in spec["columns"]}
            spec["columns"] = [dict(bynorm[_NORM(h)], header=h) for h in c["official"]]
        if c["status"] in ("match", "adoptable"):
            spec["header_verified"] = True
            n += 1
        else:
            spec["header_verified"] = False
    if n:
        schema["headers_verified_on"] = datetime.now().strftime("%Y-%m-%d")
    return n
