"""Upload a whole year in one go: a folder (or a .zip) with the client's Excel / PDF files.

scan()   - lists every file, works out what it is and which month(s) it belongs to (nothing is imported yet)
import_entries() - creates the missing month returns and imports each file into its month:
                   Master Templates -> the template import, invoice files / reports -> upload + extract,
                   filed GSTR-1 / 3B / 9 PDFs -> the filed-returns store used by the reconciliation.
A register that covers several months is imported into each of them and only that month's invoices are kept.
"""
import logging
import re
import shutil
import tempfile
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from . import returns as R
from .utils import parse_date, fin_year

log = logging.getLogger(__name__)

SUPPORTED = {".xlsx", ".xlsm", ".xls", ".csv", ".pdf"}
MON = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
KIND_LABEL = {"master": "Master Template", "excel": "Excel invoices / register", "report": "Billing summary report",
              "pdf": "PDF invoice", "GSTR1": "Filed GSTR-1", "GSTR3B": "Filed GSTR-3B", "GSTR9": "Filed GSTR-9",
              "skip": "Not used"}


def month_label(m, y):
    return f"{MON[m - 1].capitalize()} {y}"


def parse_month_label(text):
    """'Jul 2026' / 'July 2026' / '07-2026' / '2026-07' -> (7, 2026) or None."""
    t = str(text).strip().lower()
    m = re.match(r"^([a-z]{3})[a-z]*[\s\-_/]*(20\d\d)$", t)
    if m and m.group(1) in MON:
        return MON.index(m.group(1)) + 1, int(m.group(2))
    m = re.match(r"^(\d{1,2})[\s\-_/.](20\d\d)$", t) or None
    if m and 1 <= int(m.group(1)) <= 12:
        return int(m.group(1)), int(m.group(2))
    m = re.match(r"^(20\d\d)[\s\-_/.](\d{1,2})$", t)
    if m and 1 <= int(m.group(2)) <= 12:
        return int(m.group(2)), int(m.group(1))
    return None


def month_from_name(name: str):
    """Month in a file name: Sep2026, September_2026, 09-2026, 2026-09, 092026 (portal style)."""
    s = Path(name).stem.lower()
    m = re.search(r"(?<![a-z])(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[\s\-_.]*'?(20\d\d|\d\d)(?!\d)", s)
    if m:
        y = int(m.group(2))
        return MON.index(m.group(1)) + 1, (y if y > 99 else 2000 + y)
    m = re.search(r"(?<!\d)(20\d\d)[\-_.](0[1-9]|1[0-2])(?!\d)", s)
    if m:
        return int(m.group(2)), int(m.group(1))
    m = re.search(r"(?<!\d)(0[1-9]|1[0-2])[\-_.]?(20\d\d)(?!\d)", s)
    if m:
        return int(m.group(1)), int(m.group(2))
    return None


def safe_extract_zip(zpath, dest) -> list:
    """Unzip without letting a name escape the folder (../ or absolute paths are skipped)."""
    out = []
    dest = Path(dest).resolve()
    with zipfile.ZipFile(zpath) as z:
        for info in z.infolist():
            if info.is_dir() or "__MACOSX" in info.filename or Path(info.filename).name.startswith((".", "~$")):
                continue
            target = (dest / info.filename).resolve()
            if dest not in target.parents:
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with z.open(info) as src, open(target, "wb") as f:
                shutil.copyfileobj(src, f)
            out.append(target)
    return out


def list_files(source) -> tuple:
    """-> (files, temp_dir or None). source: a folder or a .zip file (zips inside the folder are opened too)."""
    tmp = None
    files = []
    if isinstance(source, (list, tuple)):
        for p in map(Path, source):
            if p.suffix.lower() == ".zip":
                tmp = tmp or Path(tempfile.mkdtemp(prefix="cagst_year_"))
                files += safe_extract_zip(p, tmp / p.stem)
            elif p.is_file():
                files.append(p)
        return sorted(files, key=lambda p: p.name.lower()), tmp
    source = Path(source)
    if source.is_file() and source.suffix.lower() == ".zip":
        tmp = Path(tempfile.mkdtemp(prefix="cagst_year_"))
        files = safe_extract_zip(source, tmp)
    elif source.is_dir():
        for p in sorted(source.rglob("*")):
            if p.is_file() and not p.name.startswith((".", "~$")):
                if p.suffix.lower() == ".zip":
                    tmp = tmp or Path(tempfile.mkdtemp(prefix="cagst_year_"))
                    files += safe_extract_zip(p, tmp / p.stem)
                else:
                    files.append(p)
    return sorted(files, key=lambda p: p.name.lower()), tmp


def _excel_months(path: Path) -> Counter:
    """Months of the dates found in an Excel / CSV file (invoice dates)."""
    c = Counter()
    try:
        if path.suffix.lower() == ".csv":
            import csv
            with open(path, newline="", encoding="utf-8-sig", errors="ignore") as f:
                for i, row in enumerate(csv.reader(f)):
                    for v in row:
                        d = parse_date(v) if re.search(r"\d", v or "") and len(v) <= 20 else ""
                        if d:
                            c[(int(d[5:7]), int(d[:4]))] += 1
                    if i > 20000:
                        break
            return c
        if path.suffix.lower() == ".xls":
            import pandas as pd
            for _, df in pd.read_excel(path, sheet_name=None, header=None).items():
                for v in df.values.ravel():
                    if hasattr(v, "strftime"):
                        d = parse_date(v)
                        if d:
                            c[(int(d[5:7]), int(d[:4]))] += 1
            return c
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                for v in row:
                    if hasattr(v, "strftime"):
                        d = parse_date(v)
                    elif isinstance(v, str) and re.match(r"^\s*\d{1,4}[/\-.]\w{1,9}[/\-.]\d{2,4}\s*$", v):
                        d = parse_date(v)
                    else:
                        continue
                    if d:
                        c[(int(d[5:7]), int(d[:4]))] += 1
                if i > 20000:
                    break
        wb.close()
    except Exception as e:  # noqa
        log.warning("Date scan failed for %s: %s", path.name, e)
    return c


def _pdf_months(path: Path) -> Counter:
    c = Counter()
    try:
        texts, _, _ = R.pdf_pages_text(path, {"ocr_enabled": "0"})
    except Exception:  # noqa
        return c
    for t in texts:
        for m in re.finditer(r"\b(\d{1,2}[/\-.]\d{1,2}[/\-.]20\d\d|20\d\d-\d{2}-\d{2}|\d{1,2}[\-\s][A-Za-z]{3}[\-\s]20\d\d)\b", t):
            d = parse_date(m.group(1))
            if d:
                c[(int(d[5:7]), int(d[:4]))] += 1
    return c


def _template_month(path: Path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = next((w for w in wb.worksheets if w.title.lower().replace(" ", "") == "returninfo"), None)
        mo = yr = None
        if ws is not None:
            for row in ws.iter_rows(values_only=True):
                k = str(row[0] or "").lower()
                v = row[1] if len(row) > 1 else None
                if k.startswith("return month") and v not in (None, ""):
                    mo = int(float(v))
                if k.startswith("return year") and v not in (None, ""):
                    yr = int(float(v))
        wb.close()
        if mo and yr and 1 <= mo <= 12:
            return mo, yr
    except Exception:  # noqa
        pass
    return None


def folder_month(p: Path):
    """Month from the name of the folder the file is in (or its parent): 'Aug 2026', '2026-08', '08-2026'."""
    for d in (p.parent, p.parent.parent):
        if d.name.startswith("cagst_year_"):
            break
        mm = parse_month_label(d.name) or month_from_name(d.name + ".x")
        if mm:
            return mm
    return None


def fy_over(fy, today=None):
    """True once the financial year has ended (GSTR-9 is filed only after 31 March)."""
    from datetime import date
    today = today or date.today()
    return today > date(int(fy[:4]) + 1, 3, 31)


def in_fy(m, y, fy):
    s = int(fy[:4])
    return (y == s and m >= 4) or (y == s + 1 and m <= 3)


def scan(source, fy: str, settings=None, progress=None) -> tuple:
    """-> (entries, temp_dir). Each entry: dict(path, name, kind, months [(m, y)], counts, include, note, doc)."""
    from .services import is_master_template
    from .extractors import reports
    from .utils import sha256_file
    files, tmp = list_files(source)
    entries = []
    seen = {}
    for i, p in enumerate(files):
        if progress:
            progress(i, len(files), p.name)
        e = {"path": str(p), "name": p.name, "kind": "skip", "months": [], "counts": {}, "include": False, "note": "",
             "doc": None}
        ext = p.suffix.lower()
        try:
            sha = sha256_file(p)
        except Exception:  # noqa
            sha = None
        if sha and sha in seen:
            e["note"] = f"Same file as {seen[sha]} - left out"
            entries.append(e)
            continue
        if sha:
            seen[sha] = p.name
        if ext not in SUPPORTED:
            e["note"] = "File type not used"
            entries.append(e)
            continue
        try:
            if ext == ".pdf":
                doc = None
                if R.looks_like_return(p):
                    doc = R.read_return(p, settings)
                else:
                    pages, _, _ = R.pdf_pages_text(p, {"ocr_enabled": "0"})
                    if R.detect_form("\n".join(pages[:1]), p.name):
                        doc = R.read_return(p, settings)
                if doc and doc["form"]:
                    e.update(kind=doc["form"], doc=doc, include=True, note="; ".join(doc["messages"]))
                    if doc["form"] == "GSTR9" and doc.get("fy") and not fy_over(doc["fy"]):
                        e["include"] = False
                        e["note"] = (f"FY {doc['fy']} is not over yet - a GSTR-9 cannot be filed for it. Left out "
                                     "(a draft can still be included by selecting it and clicking Include).")
                    if doc["form"] == "GSTR9":
                        e["months"] = [(3, int(doc["fy"][:4]) + 1)] if doc.get("fy") else []
                    elif doc.get("month") and doc.get("year"):
                        e["months"] = [(doc["month"], doc["year"])]
                else:
                    e["kind"] = "pdf"
                    mm = month_from_name(p.name)
                    c = Counter({mm: 1}) if mm else _pdf_months(p)
                    e["counts"] = dict(c)
                    e["months"] = [c.most_common(1)[0][0]] if c else []
                    e["include"] = True
            elif is_master_template(p):
                e["kind"] = "master"
                mm = _template_month(p) or month_from_name(p.name)
                e["months"] = [mm] if mm else []
                e["include"] = True
            else:
                rep = reports.detect_report(p) if ext in (".xlsx", ".xlsm", ".xls") else None
                c = _excel_months(p)
                mm = month_from_name(p.name)
                if rep:
                    e["kind"] = "report"
                    e["months"] = [mm] if mm else ([c.most_common(1)[0][0]] if c else [])
                    if not mm and len(c) > 1:
                        e["note"] = "Summary report: the month is taken from the most frequent date - check it"
                else:
                    e["kind"] = "excel"
                    months = [k for k, n in c.items() if n >= 1]
                    if mm and (not months or mm in c):
                        months = months or [mm]
                    e["months"] = sorted(months, key=lambda k: (k[1], k[0]))
                    e["counts"] = {k: c[k] for k in e["months"]}
                    if len(e["months"]) > 1:
                        e["note"] = "Covers several months - each month gets its own invoices"
                e["include"] = True
        except Exception as ex:  # noqa
            log.exception("Scan failed for %s", p)
            e["note"] = f"Could not read: {ex}"
            e["include"] = False
        fm = folder_month(p) if (e["kind"] in ("excel", "report", "pdf") or
                                 (e["kind"] == "master" and not e["months"])) else None
        if fm and e["include"] is not None and e["kind"] != "skip":
            if e["months"] and set(e["months"]) != {fm}:
                e["note"] = (e["note"] + "; " if e["note"] else "") + \
                    f"Dates in the file are {', '.join(month_label(*m) for m in e['months'])} - month taken from the folder"
            elif not e["months"]:
                e["note"] = (e["note"] + "; " if e["note"] else "") + "Month taken from the folder name"
            e["months"] = [fm]
            e["include"] = not str(e["note"]).startswith("Could not read")
        if e["include"]:
            if not e["months"]:
                e["note"] = (e["note"] + "; " if e["note"] else "") + "Month not found - double-click to choose it"
                e["include"] = False
            else:
                out = [mm for mm in e["months"] if not in_fy(mm[0], mm[1], fy)]
                if out and len(out) == len(e["months"]):
                    e["note"] = (e["note"] + "; " if e["note"] else "") + f"Not in FY {fy} - skipped"
                    e["include"] = False
                elif out:
                    e["months"] = [mm for mm in e["months"] if mm not in out]
                    e["note"] = (e["note"] + "; " if e["note"] else "") + \
                        f"{len(out)} month(s) outside FY {fy} ignored"
        entries.append(e)
    return entries, tmp


def target_month(m, y, client):
    """Quarterly (QRMP) clients: the return is the quarter-end month."""
    if str(client.get("return_type", "")).lower().startswith("quarterly"):
        q = {1: 3, 2: 3, 3: 3, 4: 6, 5: 6, 6: 6, 7: 9, 8: 9, 9: 9, 10: 12, 11: 12, 12: 12}[m]
        return q, y
    return m, y


def import_entries(db, client, entries, settings, schema, progress=None, approve_clean=False) -> list:
    """Import the included entries. Returns a log: list of (file, month label, result text, ok).
    approve_clean: approve the imported lines that have no errors (same as 'Approve all without errors')."""
    from .classifier import classify_all
    from .exporters import register_table13
    from .services import import_master_template, merge_docs, merge_lines, register_file, extract_file, review_count
    out = []
    projects = {(p["month"], p["year"]): p for p in db.projects(client["id"])}
    freq = client.get("return_type") or "Monthly"

    def project_for(m, y):
        m, y = target_month(m, y, client)
        if (m, y) not in projects:
            pid = db.create_project(client["id"], fin_year(m, y), m, y, freq, schema.get("schema_version", ""))
            projects[(m, y)] = db.get_project(pid)
        return projects[(m, y)]

    todo = [e for e in entries if e["include"] and e["months"]]
    touched = {}
    for i, e in enumerate(todo):
        if progress:
            progress(i, len(todo), e["name"])
        path = Path(e["path"])
        try:
            if e["kind"] in ("GSTR1", "GSTR3B", "GSTR9"):
                doc = e["doc"] or R.read_return(path, settings)
                m, y = e["months"][0]
                doc["month"], doc["year"] = (3, y) if e["kind"] == "GSTR9" else (m, y)
                doc["fy"] = fin_year(m, y)
                gst = doc.get("gstin")
                warn = f" - GSTIN in the file is {gst}, not {client['gstin']}" if gst and gst != client["gstin"] else ""
                R.store_return(db, client, doc)
                out.append((e["name"], "FY " + doc["fy"] if e["kind"] == "GSTR9" else month_label(m, y),
                            f"{KIND_LABEL[e['kind']]} stored for the reconciliation{warn}", not warn))
                continue
            for (m, y) in e["months"]:
                p = project_for(m, y)
                lab = month_label(p["month"], p["year"])
                if e["kind"] == "master":
                    res, ds_id, msg = import_master_template(db, client, p, path, settings, schema, skip_errors=False)
                    if ds_id is None:
                        out.append((e["name"], lab, f"NOT imported: {msg} - open {lab}, tab 4, to see the problems",
                                    False))
                        continue
                    docs = db.load_docs(p["id"])
                    if res.docs and not any(d.origin not in ("derived",) for d in docs):
                        db.save_docs(p["id"], merge_docs(docs, res.docs, "replace"))
                    out.append((e["name"], lab, f"Master Template imported: {res.total_rows} rows"
                                + (f", {len(res.warnings)} warning(s)" if res.warnings else ""), True))
                    continue
                touched[p["id"]] = p
                fid, msg = register_file(db, client, p, path)
                if not fid:
                    out.append((e["name"], lab, msg, "already uploaded" in msg))
                    continue
                res = extract_file(db, client, p, fid, settings)
                if res.needs_mapping and not res.lines:
                    out.append((e["name"], lab, f"Needs column mapping - open {lab}, tab 5, select the file and click "
                                                "'Map columns'", False))
                    continue
                new = res.lines
                if len(e["months"]) > 1:
                    keep = [l for l in new if not l.invoice_date or
                            (int(l.invoice_date[5:7]), int(l.invoice_date[:4])) == (m, y) or
                            target_month(int(l.invoice_date[5:7]), int(l.invoice_date[:4]), client) == (p["month"], p["year"])]
                    new = keep
                lines = db.load_lines(p["id"])
                name = db.q("SELECT filename FROM files WHERE id=?", (fid,))[0]["filename"]
                for l in new:
                    l.source_file = name
                lines = [l for l in lines if l.source_file != name]
                flagged = merge_lines(lines, new)
                classify_all(lines, settings)
                db.save_lines(p["id"], lines)
                docs = db.load_docs(p["id"])
                if res.docs:
                    docs = merge_docs([d for d in docs if d.source_file != name], res.docs, "add")
                t13 = register_table13(path, schema) if path.suffix.lower() in (".xlsx", ".xlsm") else None
                if t13 and t13[0] and not t13[1] and len(e["months"]) == 1:
                    for d in t13[0]:
                        d.source_file = name
                    docs = merge_docs([d for d in docs if not (d.origin == "register" and d.source_file == name)],
                                      t13[0], "add")
                db.save_docs(p["id"], docs)
                db.update_file(fid, records=len(new), review_count=review_count(new),
                               message=("; ".join(res.messages) + (f"; {flagged} possible duplicate(s)" if flagged else "")
                                        + ("; year upload: this month's invoices only" if len(e["months"]) > 1 else ""))[:500])
                db.audit(p["id"], "Year upload", f"{name}: {len(new)} records")
                what = f"{len(new)} invoice line(s)" if new else (
                    "summary tables read" if res.tables else "no records found")
                out.append((e["name"], lab, what + (f", {flagged} possible duplicate(s)" if flagged else ""),
                            bool(new or res.tables)))
        except Exception as ex:  # noqa
            log.exception("Year import failed for %s", e["name"])
            out.append((e["name"], ", ".join(month_label(*mm) for mm in e["months"]), f"Error: {ex}", False))
    if approve_clean:
        from .models import STATUS_APPROVED, STATUS_PENDING
        from .services import learn_hsn
        from .validation import validate_all
        for pid, p in touched.items():
            lines = db.load_lines(pid)
            issues = validate_all(lines, client, db.load_docs(pid), settings, (p["month"], p["year"]))
            bad = {i.line_id for i in issues if i.severity == "Error" and i.line_id}
            ok = [l for l in lines if l.review_status == STATUS_PENDING and l.id not in bad]
            for l in ok:
                l.review_status = STATUS_APPROVED
            if ok:
                learn_hsn(db, client["id"], ok)
                db.save_lines(pid, lines)
                db.audit(pid, "Year upload", f"{len(ok)} line(s) without errors approved")
                out.append(("(approval)", month_label(p["month"], p["year"]),
                            f"{len(ok)} line(s) approved, {len(bad)} with errors left pending", not bad))
    return out


def write_log(path, client, fy, log_rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Year upload"
    ws.append([f"{client['name']} | {client['gstin']} | FY {fy} | year upload {datetime.now():%d-%m-%Y %H:%M}"])
    ws.append([])
    ws.append(["File", "Month", "Result", "OK"])
    for r in log_rows:
        ws.append([r[0], r[1], r[2], "Yes" if r[3] else "CHECK"])
    for col, w in zip("ABCD", (45, 14, 90, 8)):
        ws.column_dimensions[col].width = w
    wb.save(path)
