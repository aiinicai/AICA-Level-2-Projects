"""v2.4.6 additions on top of 2.4.5: help lines, compact tabs, GSTR-9C check.  python -m tests.test_v246"""
import os, sys, tempfile
from decimal import Decimal
from pathlib import Path

tmp = Path(tempfile.mkdtemp())
os.environ["CAGST_DATA"] = str(tmp / "data")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from tests import tk_stub  # noqa: E402
tk_stub.install()
from app import annual, recon, returns, ui  # noqa: E402
from app.gstin import checksum_char  # noqa: E402

own = "33DEMOA1234B1Z"
own += checksum_char(own)
app = ui.App()
names = [app.nb.tab(t, "text") for t in app.nb.tabs()]
assert len(names) == 15 and "15 YTD & Annual" in names, names
assert ui.App.TAB_HELP["6 Invoice Review"].startswith("WORK")
for k, v in (("name", "Demo Co"), ("gstin", own), ("legal_name", "Demo Co"), ("reg_type", "Regular"), ("return_type", "Monthly")):
    app.cv[k].set(v)
app.save_client()
app.proj_client.current(0)
app.pm.set("April"); app.py.set("2025")
app.new_return()
fy = "2025-26"
app.an_fy.set(fy)
app.refresh_annual()
assert any(n.startswith("6 GSTR-9 vs GSTR-9C") for n, *_ in recon.RECON_TABLES)
assert len(app.an_rc_trees) == len(recon.RECON_TABLES)

# no GSTR-9, no 9C -> nothing for table 6
y = annual.build_year(app.db, app.client, fy, app.settings, app.schema)
assert recon.ninec_rows(app.db, app.client, fy, None) == []
# GSTR-9 uploaded but 9C not typed -> MISSING row
g9 = {"data": {"5N": {"value": "1000000"}}, "fy": fy}
r = recon.ninec_rows(app.db, app.client, fy, g9)
assert r[0]["status"] == recon.MISSING
# type 9C figures (through the screen's storage) -> compare
e = annual.load_entries(app.db, app.client["id"], fy)
e[annual.ANNUAL_KEY] = {"fs_turnover": "1050000", "c9_adj": "-50000"}
annual.save_entries(app.db, app.client["id"], fy, e)
r = recon.ninec_rows(app.db, app.client, fy, g9)
assert r[-1]["status"] == recon.OK and r[-1]["d_value"] == 0, r[-1]
e[annual.ANNUAL_KEY]["c9_adj"] = "-40000"
annual.save_entries(app.db, app.client["id"], fy, e)
r = recon.ninec_rows(app.db, app.client, fy, g9)
assert r[-1]["status"] == recon.DIFF and r[-1]["d_value"] == Decimal("10000"), r[-1]
assert "5R" in recon.how_to_fix(r[-1])
# stored GSTR-9 via the real store -> table builds through recon.build and the screen
returns.save_filed(app.db, app.client["id"], fy, {"GSTR9": {"FY": g9}})
rows = recon.build(app.db, app.client, fy, y)
assert any(x["area"] == "GSTR-9 vs GSTR-9C" and x["status"] == recon.DIFF for x in rows)
app.refresh_annual()
app.ninec_entry()
print("v2.4.6 TESTS PASSED")
fd = ui.filedialog
out = tmp / "ytd.xlsx"
fd.asksaveasfilename = lambda **k: str(out)
app.annual_export()
from openpyxl import load_workbook  # noqa: E402
sn = load_workbook(out).sheetnames
assert "6 GSTR-9 vs GSTR-9C" in sn and "Filed returns" in sn, sn
print("export sheets:", sn)
