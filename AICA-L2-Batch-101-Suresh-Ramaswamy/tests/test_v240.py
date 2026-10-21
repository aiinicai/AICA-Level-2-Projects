"""v2.4.0: whole-year upload, filed GSTR-1 / 3B / 9 readers and the reconciliation:  python -m tests.test_v240
Uses the made-up DEMO returns of tools/make_demo_returns.py (no real client data)."""
import os, shutil, sys, tempfile, zipfile
from decimal import Decimal as Dc
from pathlib import Path
tmp = Path(tempfile.mkdtemp()); os.environ["CAGST_DATA"] = str(tmp / "data")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / "tools"))
import make_demo_returns as demo  # noqa: E402
from app import returns as R, recon, yearimport as Y, annual  # noqa: E402
from app.db import Database  # noqa: E402
from app.schema import load_schema  # noqa: E402

# ---------------------------------------------------------------- number / text helpers
assert R.amt("1,25,000.50") == Dc("125000.50") and R.amt("-12,34,567.80") == Dc("-1234567.80")
assert R.amt("680899.01]") == Dc("680899.01") and R.amt("500") is None and R.amt("Invoice") is None
assert R.vec([Dc("10"), Dc("1.8"), Dc("0")]) == {"value": "10", "igst": "1.8", "cgst": "0", "sgst": "0", "cess": "0"}
assert R.period_from_text("Financial year 2025-26\nTax period March", "x.pdf") == ("2025-26", 3, 2026)
assert R.period_from_text("", "GSTR3B_33AAAAA0000A1Z5_082026.pdf") == ("2026-27", 8, 2026)
assert R.detect_form("Form GSTR-3B [See rule 61(5)]") == "GSTR3B" and R.detect_form("FORM GSTR-1") == "GSTR1"
assert R.detect_form("Form GSTR-9 [See rule 80]") == "GSTR9" and R.detect_form("", "GSTR9C_x.pdf") == ""
# watermark letters (rotated / very large) are dropped before the rows are rebuilt
spans = [(30, 100, 60, 108, "Total", 7.5, False), (300, 100, 360, 108, "1,000.00", 7.5, False),
         (200, 60, 330, 300, "N", 250, False), (400, 98, 420, 110, "F", 9, True)]
assert R.spans_to_text(spans) == "Total  1,000.00", R.spans_to_text(spans)

# ---------------------------------------------------------------- the three readers on the demo returns
def doc(name):
    rows, _ = demo.FILES[name]
    d = R.read_return_text([demo.to_text(rows())], name)
    d["ocr"] = False
    return d

G = demo.GSTIN
j1, j3 = doc(f"GSTR1_{G}_072026.pdf"), doc(f"GSTR3B_{G}_072026.pdf")
a1, a3 = doc(f"GSTR1_{G}_082026.pdf"), doc(f"GSTR3B_{G}_082026.pdf")
g9 = R.read_return_text([demo.to_text(demo.g9_rows(demo.FY, demo.G9))], f"GSTR9_{G}_032027.pdf"); g9["ocr"] = False
assert (j1["form"], j1["fy"], j1["month"], j1["year"], j1["gstin"], j1["arn_date"]) == ("GSTR1", "2026-27", 7, 2026, G, "11/08/2026")
assert (j3["form"], j3["month"], a3["month"], g9["form"], g9["fy"]) == ("GSTR3B", 7, 8, "GSTR9", "2026-27")
f = R.g1_figures(j1["data"])
assert f["b2b"]["value"] == Dc("15500.00") and f["b2cl"]["igst"] == Dc("27000.00") and f["expwop"]["value"] == Dc("100000.00")
assert f["sezwp"]["igst"] == Dc("1800.00") and f["n_reg"]["value"] == Dc("-500.00") and f["nu_b2cl"]["igst"] == Dc("-1800.00")
assert (f["nil"]["value"], f["exempt"]["value"], f["nongst"]["value"], f["docs"]) == (Dc("7000.00"), Dc("4000.00"), Dc("500.00"), 200)
e = R.g1_expected_3b(f)
assert e["a"]["value"] == Dc("235000.00") and R.tax_of(e["a"]) == Dc("40935.00"), e["a"]
assert e["b"]["value"] == Dc("310000.00") and e["b"]["igst"] == Dc("37800.00") and e["c"]["value"] == Dc("11000.00")
t = R.g3b_figures(j3["data"])
assert t["3.1"]["a"]["igst"] == Dc("40935.00") and t["pay"]["A_igst_cash"] == Dc("58735.00") and t["itc"]["A5"]["igst"] == Dc("20000.00")
assert g9["data"]["4A"]["value"] == "225330.00" and g9["data"]["4A"]["igst"] == "39200.00" and g9["data"]["9B"]["payable"] == "10.20"

# a portal-style GSTR-1 (made-up figures): reverse-charge supplier (GTA) - 3.1(a) value with no tax, amendment + notes included
REAL = """FORM GSTR-1
Financial year 2025-26
Tax period March
4B - Taxable outward supplies made to registered persons attracting tax on reverse charge - B2B Reverse charge
Total  50  Invoice  10,00,000.00  30,000.00  10,000.00  10,000.00  0.00
8 - Nil rated, exempted and non GST outward supplies
Total  1,00,000.00
-  Exempted  1,00,000.00
9A - Amendment to taxable outward supplies made to registered person in returns of earlier tax periods in table 4 - B2B Reverse charge
Amended amount - Total  3  Invoice  2,00,000.00  6,000.00  2,000.00  2,000.00  0.00
Net differential amount (Amended - Original)  -50,000.00  -1,500.00  -500.00  -500.00  0.00
9B - Credit/Debit Notes (Registered) - CDNR
Total - Net off debit/credit notes (Debit notes - Credit notes)  2  Note  -20,000.00  -600.00  -200.00  -200.00  0.00
Credit / Debit notes issued to registered person for taxable outward supplies in table 4 other than table 6 - B2B Reverse charge
Net Total (Debit notes - Credit notes)  2  Note  -20,000.00  -600.00  -200.00  -200.00  0.00
Total Liability (Outward supplies other than Reverse charge)  10,30,000.00  0.00  0.00  0.00  0.00"""
r = R.read_return_text([REAL], "GSTR1_x_032026.pdf")
e = R.g1_expected_3b(R.g1_figures(r["data"]))
assert e["a"]["value"] == Dc("930000.00") and R.tax_of(e["a"]) == 0, e["a"]        # what GSTR-3B 3.1(a) must show
assert e["a"]["value"] + e["c"]["value"] == Dc("1030000.00")                         # = the return's own total

# ---------------------------------------------------------------- reconciliation rows
rows = recon.month_rows("Jul 2026", None, j1, j3)
assert all(x["status"] in ("OK", "MISSING") for x in rows), [(x["check"], x["d_value"], x["d_tax"]) for x in rows if x["status"] != "OK"]
rows = recon.month_rows("Aug 2026", None, a1, a3)
bad = [x for x in rows if x["status"] == "DIFF"]
assert len(bad) == 1 and bad[0]["check"].startswith("3.1(a)") and bad[0]["d_value"] == Dc("30.00") and bad[0]["d_tax"] == Dc("2.40"), bad
rows = recon.annual_rows("2026-27", {"2026-07": j1, "2026-08": a1}, {"2026-07": j3, "2026-08": a3}, g9)
by = {x["check"]: x for x in rows}
for k in ("4A Supplies to unregistered persons (B2C)", "4B Supplies to registered persons (B2B)", "4C Exports with payment of tax",
          "4D SEZ supplies with payment of tax", "4F Advances (net)", "4J - 4I Debit notes less credit notes (net)",
          "5A Exports without payment of tax", "5D Exempted", "5E Nil rated", "5F Non-GST supply", "6A ITC availed (GSTR-3B 4A)",
          "9 Tax paid"):
    assert by[k]["status"] == "OK", (k, by[k])
assert by["9 Tax payable"]["d_tax"] == Dc("2.40") and by["9 Payable vs paid in GSTR-9"]["status"] == "DIFF"

# ---------------------------------------------------------------- month detection and zip safety
assert Y.month_from_name("Sales_Register_Sep2026.xlsx") == (9, 2026) and Y.month_from_name("GSTR1_X_072026.pdf") == (7, 2026)
assert Y.month_from_name("register 2026-08.xlsx") == (8, 2026) and Y.month_from_name("Invoice_PDF-001.pdf") is None
assert Y.parse_month_label("Aug 2026") == (8, 2026) and Y.parse_month_label("2026-07") == (7, 2026)
assert Y.folder_month(Path("/x/2026-08 August/a.xlsx")) == (8, 2026) and Y.folder_month(Path("/x/y/a.xlsx")) is None
assert Y.target_month(8, 2026, {"return_type": "Quarterly (QRMP)"}) == (9, 2026) and Y.target_month(8, 2026, {}) == (8, 2026)
zp = tmp / "evil.zip"
with zipfile.ZipFile(zp, "w") as z:
    z.writestr("../escape.txt", "x")
    z.writestr("ok/inside.txt", "y")
got = Y.safe_extract_zip(zp, tmp / "unz")
assert [p.name for p in got] == ["inside.txt"] and not (tmp / "escape.txt").exists()

# ---------------------------------------------------------------- whole-year upload of the sample pack (Excel files)
S = ROOT / "sample_data" / "Demo_Samples_Pack"
yr = tmp / "year"
for sub, names in (("2026-07 July", ["03_Demo_Master_Template_FILLED_Jul2026.xlsx"]),
                   ("2026-08 August", ["02_Demo_HSN_Report_WITH_WRONG_CODES.xlsx", "07_Demo_GST_Output_Register.xlsx"]),
                   ("2026-09 September", ["01_Demo_Sales_Register_Sep2026.xlsx"])):
    (yr / sub).mkdir(parents=True)
    for n in names:
        shutil.copy(S / n, yr / sub / n)
(yr / "notes.txt").write_text("not used")
zpath = tmp / "year.zip"
with zipfile.ZipFile(zpath, "w") as z:
    for p in yr.rglob("*"):
        if p.is_file():
            z.write(p, p.relative_to(yr))
db = Database(); settings = db.get_settings(); schema = load_schema()
cid = db.save_client(dict(name="Demo Traders", gstin=G, legal_name="", trade_name="", state_code="33",
                          reg_type="Regular", return_type="Monthly"))
client = [c for c in db.clients() if c["id"] == cid][0]
entries, t_ = Y.scan(zpath, "2026-27", settings)
months = {e["name"]: [Y.month_label(*m) for m in e["months"]] for e in entries if e["include"]}
assert months == {"01_Demo_Sales_Register_Sep2026.xlsx": ["Sep 2026"], "02_Demo_HSN_Report_WITH_WRONG_CODES.xlsx": ["Aug 2026"],
                  "03_Demo_Master_Template_FILLED_Jul2026.xlsx": ["Jul 2026"], "07_Demo_GST_Output_Register.xlsx": ["Aug 2026"]}, months
assert [e["kind"] for e in entries if e["name"] == "notes.txt"] == ["skip"]
log = Y.import_entries(db, client, entries, settings, schema, approve_clean=True)
assert sum(1 for x in log if x[3]) >= 4, log
shutil.rmtree(t_, ignore_errors=True)
assert sorted((p["month"], p["year"]) for p in db.projects(cid)) == [(7, 2026), (8, 2026), (9, 2026)]
# a register covering two months goes to both months, each keeping its own invoices
from openpyxl import load_workbook
wb = load_workbook(S / "01_Demo_Sales_Register_Sep2026.xlsx")
ws = wb.active
for row in ws.iter_rows(min_row=2):
    for c in row:
        if hasattr(c.value, "strftime") and c.value.day <= 10:
            c.value = c.value.replace(month=10)
two = tmp / "two" / "Register_Sep_Oct.xlsx"; two.parent.mkdir()
wb.save(two)
e2, _ = Y.scan(two.parent, "2026-27", settings)
assert [Y.month_label(*m) for m in e2[0]["months"]] == ["Sep 2026", "Oct 2026"], e2[0]
Y.import_entries(db, client, e2, settings, schema)
octp = [p for p in db.projects(cid) if p["month"] == 10][0]
assert db.load_lines(octp["id"]) and all(l.invoice_date[5:7] == "10" for l in db.load_lines(octp["id"]))
# filed returns stored -> GSTR-3B fills the 3B figures of the month; reconciliation with the books
for d in (j1, j3, a1, a3, g9):
    R.store_return(db, client, d)
ent = annual.load_entries(db, cid, "2026-27")
assert ent["2026-07"]["tb_taxable"] == "556500.00" and ent["2026-07"]["_src"] == "GSTR-3B PDF", ent["2026-07"]
y = annual.build_year(db, client, "2026-27", settings, schema)
rows = recon.build(db, client, "2026-27", y)
jul_books = [x for x in rows if x["period"] == "Jul 2026" and x["area"] == "Books vs GSTR-1"]
assert jul_books and all(x["status"] == "OK" for x in jul_books), [(x["check"], x["d_value"], x["d_tax"]) for x in jul_books]
assert any(x["period"] == "Sep 2026" and x["status"] == "MISSING" for x in rows)
lst = recon.filed_list(db, cid, "2026-27")
assert [x["form"] for x in lst] == ["GSTR1", "GSTR1", "GSTR3B", "GSTR3B", "GSTR9"]
assert not Y.fy_over("2026-27", __import__("datetime").date(2026, 10, 20)) and Y.fy_over("2025-26", __import__("datetime").date(2026, 10, 20))
out = tmp / "ytd.xlsx"
annual.export_year(out, client, y)
recon.export_sheets(out, client, "2026-27", rows, lst)
assert load_workbook(out).sheetnames[-2:] == ["Missing - information", "Filed returns"]
R.remove_return(db, cid, "2026-27", "GSTR9", "FY")
assert "FY" not in R.load_filed(db, cid, "2026-27").get("GSTR9", {})
print("ALL v2.4.0 TESTS PASSED")
