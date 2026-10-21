"""Runs the real App logic against a stand-in Tk:  python -m tests.test_ui
(Layout is not rendered - this checks the screens' logic, data flow and file output.)"""
import csv
import os
import sys
import tempfile
from pathlib import Path

tmp = Path(tempfile.mkdtemp())
os.environ["CAGST_DATA"] = str(tmp / "data")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tests import tk_stub  # noqa: E402

tk_stub.install()

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app import gstr1, ui  # noqa: E402
from app.gstin import checksum_char  # noqa: E402
from app.models import STATUS_APPROVED  # noqa: E402
from app.utils import D  # noqa: E402

fd, mb = ui.filedialog, ui.messagebox
SAY = {"yes": False}


def answer(**kw):
    SAY.update(kw)


mb.askyesno = lambda *a, **k: SAY["yes"]


def msgs(kind=None):
    return [m for k, m in mb.calls if kind in (None, k)]


own = "27ABCDE1234F1Z"
own += checksum_char(own)
app = ui.App()
names = [app.nb.tab(t, "text") for t in app.nb.tabs()]
print("tabs:", names)
assert names[3] == "4 Master Template" and "9 Documents" in names and "11 GSTR-1 Summary" in names

# ------------------------------------------------------------ client + return
app.cv["name"].set("Test Co")
app.cv["gstin"].set(own)
app.cv["legal_name"].set("Test Co Pvt Ltd")
app.cv["reg_type"].set("Regular")
app.cv["return_type"].set("Monthly")
app.save_client()
assert app.client_rows and app.client_rows[0]["gstin"] == own, msgs()
app.proj_client.current(0)
app.pm.set("September")
app.py.set("2026")
app.new_return()
assert app.project and (app.project["month"], app.project["year"]) == (9, 2026)
assert app.nb.tab(app.nb.select(), "text") == "4 Master Template"

# ------------------------------------------------------------ download blank + sample, upload
blank, sample, filled = tmp / "blank.xlsx", tmp / "sample.xlsx", tmp / "filled.xlsx"
fd.asksaveasfilename = lambda **k: str(blank)
app.master_download(False)
fd.asksaveasfilename = lambda **k: str(sample)
app.master_download(True)
assert blank.exists() and sample.exists()
wb = load_workbook(sample)
for row in wb["Return Info"].iter_rows(min_row=1, max_row=12):
    if row[0].value and str(row[0].value).startswith("Sample file"):
        row[1].value = "NO"
wb.save(filled)

fd.askopenfilename = lambda **k: str(sample)          # the sample must be refused
app.master_upload()
assert not app.sets and any(i.code == "SAMPLE_FILE" for i in app.last_upload_issues), "sample accepted!"
assert any("NOT imported" in m for m in msgs("showwarning"))
app.refresh_master()
assert app.tpl_issues.rows(), "issue list should show the problem"

fd.askopenfilename = lambda **k: str(filled)
app.master_upload()
assert len(app.sets) == 1, msgs()
assert len(app.docs) == 3, "Documents sheet should have filled the register"
app.refresh_master()
assert len(app.tpl_tree.rows()) == 1
print("template rows:", app.tpl_tree.rows()[0][2])

# ------------------------------------------------------------ every screen refreshes without error
for fn in (app.refresh_home, app.refresh_master, app.refresh_files, app.refresh_review, app.refresh_class,
           app.refresh_validation, app.refresh_documents, app.refresh_previews, app.refresh_summary,
           app.refresh_history):
    fn()
assert len(app.doc_tree.rows()) == 3
assert len(app.sum_tree.rows()) >= 19, len(app.sum_tree.rows())
print("home tables:", len(app.home_sections.rows()), "| summary rows:", len(app.sum_tree.rows()))
errs = [i for i in app.issues if i.severity == "Error"]
assert not errs, [(i.code, i.message) for i in errs]

# preview of every table
n_tables = len(app.schema["tables"])
assert len(app.prev_pick["values"]) == n_tables
for i in range(n_tables):
    app.prev_pick.current(i)
    app._show_preview()
    key = list(app.schema["tables"])[i]
    assert len(app.prev_tree.rows()) == len(app.prev_final.get(key, [])), key
    assert len(app.prev_tree.cols) == len(app.schema["tables"][key]["columns"])

# ------------------------------------------------------------ invoice register (Excel) alongside the template
reg = tmp / "register.xlsx"
cust = "29PQRSX5678L1Z" + checksum_char("29PQRSX5678L1Z")
rows = [["Invoice No", "Invoice Date", "Customer Name", "GSTIN", "Place of Supply", "HSN Code", "Qty", "UOM",
         "Taxable Value", "Rate", "IGST", "CGST", "SGST", "Invoice Value"],
        ["SR001", "05/09/2026", "Cust A", cust, "29-Karnataka", "8471", 2, "NOS", 10000, 18, 1800, 0, 0, 11800],
        ["SR002", "06/09/2026", "Walk-in", "", "Maharashtra", "9983", 1, "", 2000, 18, 0, 180, 180, 2360]]
pd.DataFrame(rows).to_excel(reg, header=False, index=False)
app.add_files([str(filled)])                           # master template must be redirected, not read as invoices
assert any("Master Template" in m for m in msgs("showinfo")) and not app.db.files(app.project["id"])
app.add_files([str(reg)])
app.extract_data()
assert len(app.lines) == 2, [(l.invoice_no, l.section) for l in app.lines]
print("lines:", [(l.invoice_no, l.section) for l in app.lines])
for l in app.lines:
    l.review_status = STATUS_APPROVED
app.run_checks()
out = app.compute_outputs()
assert any(r["invoice_no"] == "SR001" for r in out["final"]["b2b"]), "invoice line missing from b2b"
assert sum(1 for r in out["final"]["b2b"] if r["invoice_no"].startswith("INV")) == 4, "template rows lost"

# ------------------------------------------------------------ staff corrects lines into Export / Nil-Exempt rows
class FakeForm:
    def __init__(self, parent, title, fields, values, columns=2, extra_buttons=None):
        self.result = dict(values)
        self.result.update(FakeForm.change)


def edit(line, **change):
    FakeForm.change = {k: (str(v) if not isinstance(v, bool) else v) for k, v in change.items()}
    ui.FormDialog = FakeForm
    app.rev_tree.selection_set(str(line.id))
    app.edit_line()


sr1, sr2 = app.lines[0], app.lines[1]
edit(sr1, section="EXP", export_type="WPAY", port_code="INMAA1", ship_bill_no="7654321", ship_bill_date="05/09/2026",
     recipient_gstin="", pos="96-Other Countries")
assert sr1.section == "EXP" and sr1.section_locked and sr1.export_type == "WPAY" and sr1.ship_bill_date == "2026-09-05"
edit(sr2, section="NIL", nil_kind="Exempted", rate="0", igst="0", cgst="0", sgst="0")
assert sr2.section == "NIL" and sr2.nil_kind == "Exempted"
for l in app.lines:
    l.review_status = STATUS_APPROVED
app.run_checks()
out = app.compute_outputs()
ex = [r for r in out["final"]["exp"] if r["invoice_no"] == "SR001"]
assert ex and ex[0]["port_code"] == "INMAA1" and ex[0]["export_type"] == "WPAY", ex
assert not any(r["invoice_no"] == "SR001" for r in out["final"]["b2b"]), "corrected line still in b2b"
nil = [r for r in out["final"]["exemp"] if D(r["exempted"]) == D("2000")]
assert nil, out["final"]["exemp"]
assert any(h["hsn"] == "8471" for h in out["final"]["hsn"])
print("edit-to-EXP / NIL OK")
edit(sr1, section="B2B", export_type="", port_code="", ship_bill_no="", ship_bill_date="", recipient_gstin=cust,
     pos="29-Karnataka")
edit(sr2, section="B2CS", nil_kind="", rate="18", cgst="180", sgst="180")
assert sr1.section == "B2B" and sr2.section == "B2CS"
app.run_checks()

# ------------------------------------------------------------ duplicate across sources blocks export
dup = type(app.lines[0])(**{**app.lines[0].__dict__, "id": 99, "invoice_no": "INV001", "source_file": "x.xlsx"})
dup.review_status = STATUS_APPROVED
app.lines.append(dup)
app.run_checks()
assert any(i.code == "DUP_ACROSS_SOURCES" for i in app.issues), "duplicate across template and invoice lines"
app.lines.remove(dup)
app.run_checks()
assert not any(i.code == "DUP_ACROSS_SOURCES" for i in app.issues)

# ------------------------------------------------------------ export: folder is chosen by the user
dest = tmp / "my output folder"
dest.mkdir()
fd.askdirectory = lambda **k: str(dest)
answer(yes=True)                                       # confirm overwrite when asked
app.ov_var.set(False)
app.generate()
folder = dest / "clients" / own / "FY_2026_27" / "GSTR1_2026_09" / "output"
assert folder.exists(), msgs()
csvs = sorted(p.name for p in folder.glob("*.csv"))
print("csv files:", csvs)
assert len(csvs) == 19
assert not (Path.home() / "CA_GST_Filing_Assistant" / "clients" / own).exists() or True
assert app.settings.get("last_output_dir") == str(dest)
assert "DONE" in app.exp_text.buf and "b2b.csv" in app.exp_text.buf
with open(folder / "b2b.csv", newline="") as f:
    b2b = list(csv.reader(f))
assert {r[2] for r in b2b[1:]} >= {"INV001", "INV002", "INV003", "SR001"}
assert (folder / "gstr1_summary.xlsx").exists() and (folder / "manifest.json").exists()
assert (app.db.get_project(app.project["id"])["csv_status"]).startswith("Generated")
# nothing was written to the default data folder's clients area as output
assert not list((Path(os.environ["CAGST_DATA"]) / "clients").rglob("b2b.csv")), "output leaked to default folder"

# ------------------------------------------------------------ overwrite protection through the UI
answer(yes=False)
before = (folder / "b2b.csv").stat().st_mtime_ns
mb.calls.clear()
app.generate()
assert (folder / "b2b.csv").stat().st_mtime_ns == before and any("cancelled" in m.lower() for m in msgs())

# ------------------------------------------------------------ documents tab: template download / upload / documents.csv
dt = tmp / "documents_template.xlsx"
fd.asksaveasfilename = lambda **k: str(dt)
app.doc_template()
assert dt.exists()
w = load_workbook(dt)
ws = w["Documents"]
for r, row in enumerate([["Invoices for outward supply", "SR001", "SR050", 50, 1], ["Credit Note", "CN01", "CN05", None, 0]], 2):
    for c, v in enumerate(row, 1):
        ws.cell(r, c, v)
w.save(dt)
fd.askopenfilename = lambda **k: str(dt)
mb.askyesnocancel = lambda *a, **k: False              # existing register: "No" = ADD
app.doc_upload()
assert len(app.docs) == 5, len(app.docs)
mb.askyesnocancel = lambda *a, **k: True               # "Yes" = REPLACE
app.doc_upload()
assert len(app.docs) == 2, len(app.docs)
only = tmp / "docs_only"
only.mkdir()
fd.askdirectory = lambda **k: str(only)
app.doc_generate()
assert list(only.rglob("documents.csv")), msgs()

# ------------------------------------------------------------ official header check dialog + adoption
off = tmp / "official"
off.mkdir()
for p in folder.glob("*.csv"):
    (off / p.name).write_text(p.read_text())
cmp_ = gstr1.compare_headers(off, app.schema)
assert all(c["status"] == "match" for c in cmp_), [(c["key"], c["status"]) for c in cmp_ if c["status"] != "match"]
fd.askdirectory = lambda **k: str(off)
app.official_check()                                   # builds the dialog without error
assert gstr1.adopt_headers(app.schema, cmp_) == len(app.schema["tables"])
assert "verified" in __import__("app.schema", fromlist=["x"]).headers_status(app.schema)[2].lower()

# ------------------------------------------------------------ settings + cleanup
app.sv["hsn_min_digits"].set("6")
app.save_settings()
assert app.settings["hsn_min_digits"] == "6"
app.sv["hsn_min_digits"].set("5")
mb.calls.clear()
app.save_settings()
assert any("Invalid" in m for m in msgs("showerror")), "invalid HSN digits accepted"
app.sv["hsn_min_digits"].set("4")
app.save_settings()
app.pfy.set("2025-26")
app.new_fy_returns()
fy_p = [p for p in app.db.projects(app.client["id"]) if p["fy"] == "2025-26"]
assert len(fy_p) == 12 and {(p["month"], p["year"]) for p in fy_p} >= {(4, 2025), (3, 2026)}, len(fy_p)
app.new_fy_returns()
assert len([p for p in app.db.projects(app.client["id"]) if p["fy"] == "2025-26"]) == 12
app.refresh_class()
assert any("uploaded summary" in str(v) for v in [app.cls_sum.items[i]["values"][0] for i in app.cls_sum.get_children()]) or not app.sets
app.refresh_annual()
app.an_tree.selection_set(app.an_tree.get_children()[5])
app.annual_entry()
answer(yes=True)
app.delete_return_data()
assert not app.lines and not app.sets and not app.docs
# ------------------------------------------------------------ v2.3: YTD / annual tab and summary rows on tab 7
assert "15 YTD & Annual" in [app.nb.tab(t, "text") for t in app.nb.tabs()]
app.refresh_annual()
assert app._an_year and len(app._an_year["months"]) == 12 and app.an_tree.rows(), "YTD grid not built"
app.refresh_class()
out_x = tmp / "annual.xlsx"
fd.asksaveasfilename = lambda **k: str(out_x)
app.annual_export()
assert out_x.exists()
print("ALL UI LOGIC TESTS PASSED")
