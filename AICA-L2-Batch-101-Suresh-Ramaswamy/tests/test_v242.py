"""v2.4.2: complete demo year FY 2025-26 (12 months of books, GSTR-1, GSTR-3B and the GSTR-9) - the reconciliation
must find exactly the planted differences:  python -m tests.test_v242"""
import os, sys, tempfile
from datetime import date
from pathlib import Path
tmp = Path(tempfile.mkdtemp()); os.environ["CAGST_DATA"] = str(tmp / "data")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / "tools"))
import make_demo_fy2025 as fy  # noqa: E402
fy.MNAME_SHORT = {m[:3] for m in fy.MNAME}
from make_demo_returns import to_text, g1_rows, g3b_rows, g9_rows, GSTIN  # noqa: E402
from app import returns as R, recon, yearimport as Y, annual  # noqa: E402
from app.db import Database  # noqa: E402
from app.schema import load_schema  # noqa: E402

year = tmp / "FY 2025-26"
for i, (m, y) in enumerate(fy.MONTHS):
    f = year / f"{y}-{m:02d} {fy.MNAME[m - 1]}"
    f.mkdir(parents=True)
    fy.books_template(i, f / f"Master_Template_{fy.MNAME[m - 1][:3]}{y}.xlsx")
db = Database(); settings = db.get_settings(); schema = load_schema()
cid = db.save_client(dict(name="Demo Traders", gstin=GSTIN, legal_name="", trade_name="", state_code="33",
                          reg_type="Regular", return_type="Monthly"))
client = [c for c in db.clients() if c["id"] == cid][0]
entries, _ = Y.scan(year, "2025-26", settings)
assert len([e for e in entries if e["include"] and e["kind"] == "master"]) == 12
log = Y.import_entries(db, client, entries, settings, schema)
assert all(x[3] for x in log), [x for x in log if not x[3]]
for i, (m, y) in enumerate(fy.MONTHS):
    g1, g3 = fy.returns_of(i)
    for rows, name in ((g1_rows(fy.MNAME[m - 1], fy.FY, "11/01/2026", g1), f"GSTR1_{GSTIN}_{m:02d}{y}.pdf"),
                       (g3b_rows(fy.MNAME[m - 1], fy.FY, "20/01/2026", g3), f"GSTR3B_{GSTIN}_{m:02d}{y}.pdf")):
        d = R.read_return_text([to_text(rows)], name)
        assert (d["month"], d["year"]) == (m, y), (name, d["month"], d["year"])
        R.store_return(db, client, d)
g9 = R.read_return_text([to_text(g9_rows(fy.FY, fy.gstr9()))], f"GSTR9_{GSTIN}_032026.pdf")
assert Y.fy_over(g9["fy"], date(2026, 10, 20))
R.store_return(db, client, g9)
yv = annual.build_year(db, client, "2025-26", settings, schema)
assert sum(1 for r in yv["months"] if r["has_data"]) == 12
rows = recon.build(db, client, "2025-26", yv)
bad = sorted((r["period"], r["check"], str(r["d_value"]), str(r["d_tax"])) for r in rows if r["status"] != "OK")
want = sorted([("Aug 2025", "3.1(a) Outward taxable supplies", "10000.00", "1800.00"),
               ("Sep 2025", "3.1(a) Outward taxable supplies", "-10000.00", "-1800.00"),
               ("Oct 2025", "B2B (Table 4 + 6B + 6C)", "25000.00", "4500.00"),
               ("Oct 2025", "Documents issued, net (Table 13)", "1", "None"),
               ("FY 2025-26", "4B B2B", "25000.00", "4500.00"),
               ("FY 2025-26", "Output tax for the year (books vs GSTR-9 Table 9)", "None", "4500.00")])
assert bad == want, bad
assert all(r["status"] == "OK" for r in rows if r["area"] == "GSTR-9 vs books"), \
    [(r["check"], r["d_value"]) for r in rows if r["area"] == "GSTR-9 vs books" and r["status"] != "OK"]
assert "made good" in recon.how_to_fix([r for r in rows if r["period"] == "Sep 2025" and r["status"] == "DIFF"][0])
# export of 'GSTR-1 vs 3B vs Books': taxable (net) is not reduced by the exempt value a second time
from openpyxl import load_workbook
out = tmp / "ytd.xlsx"
annual.export_year(out, client, yv)
ws = load_workbook(out)["Books vs GSTR-3B"]
apr = [r for r in ws.iter_rows(values_only=True) if r[0] == "Apr 2025"][0]
assert apr[1] + apr[2] == apr[4] == 195000 and apr[7] == 0, apr
tot = [r for r in ws.iter_rows(values_only=True) if r[0] == "TOTAL"][0]
assert tot[1] == sum(r[1] for r in ws.iter_rows(values_only=True) if r[0] and r[0][:3] in fy.MNAME_SHORT), tot
print("ALL v2.4.2 TESTS PASSED")
