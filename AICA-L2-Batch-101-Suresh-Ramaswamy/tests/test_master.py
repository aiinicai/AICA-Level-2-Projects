"""Master Template round trip for every GSTR-1 table.   Run:  python -m tests.test_master"""
import csv
import os
import sys
import tempfile
from pathlib import Path

tmp = Path(tempfile.mkdtemp())
os.environ["CAGST_DATA"] = str(tmp / "data")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from app import gstr1  # noqa: E402
from app.db import Database  # noqa: E402
from app.exporters import export_project  # noqa: E402
from app.gstin import checksum_char  # noqa: E402
from app.schema import load_schema  # noqa: E402
from app.services import import_master_template, merge_docs  # noqa: E402
from app.validation import validate_all  # noqa: E402

own = "27ABCDE1234F1Z"
own += checksum_char(own)
client = {"name": "Test Co", "gstin": own, "legal_name": "Test Co Pvt Ltd", "trade_name": "", "state_code": "27",
          "reg_type": "Regular", "return_type": "Monthly"}
db = Database()
client["id"] = db.save_client(client)
schema = load_schema()
pid = db.create_project(client["id"], "2026-27", 9, 2026, "Monthly", schema["schema_version"])
project = db.get_project(pid)
settings = db.get_settings()
ctx = gstr1.make_ctx(client, project, settings, schema)

# ---------------------------------------------------------------- 1. sample file is refused
sample = tmp / "sample.xlsx"
gstr1.write_master_template(sample, schema, client, project, sample=True)
wb = load_workbook(sample)
assert {"Instructions", "Return Info", "Lists"} <= set(wb.sheetnames)
assert {spec["sheet"] for spec in schema["tables"].values()} <= set(wb.sheetnames), "a table sheet is missing"
print("sheets:", wb.sheetnames)
res = gstr1.read_master_template(sample, schema, ctx, "sample.xlsx")
assert any(i.code == "SAMPLE_FILE" for i in res.errors), "sample template must be refused"
res_db = import_master_template(db, client, project, sample, settings, schema)
assert res_db[1] is None, "sample must not be imported"

# ---------------------------------------------------------------- 2. blank template: headings only, nothing to import
blank = tmp / "blank.xlsx"
gstr1.write_master_template(blank, schema, client, project, sample=False)
r0 = gstr1.read_master_template(blank, schema, ctx, "blank.xlsx")
assert not r0.errors and r0.total_rows == 0, (r0.errors, r0.total_rows)

# ---------------------------------------------------------------- 3. the "client-filled" file = sample data with the flag cleared
filled = tmp / "filled.xlsx"
wb = load_workbook(sample)
ws = wb["Return Info"]
for row in ws.iter_rows(min_row=1, max_row=12):
    if row[0].value and str(row[0].value).startswith("Sample file"):
        row[1].value = "NO"
wb.save(filled)
res = gstr1.read_master_template(filled, schema, ctx, "filled.xlsx")
for i in res.issues:
    print(i.severity, i.code, i.ref, "-", i.message)
assert not res.errors, [i.message for i in res.errors]
print("rows read:", res.counts)
for k in ("b2b", "b2cl", "b2cs", "exp", "cdnr", "cdnur", "exemp", "at", "atadj", "hsn", "docs",
          "b2ba", "b2cla", "b2csa", "expa", "cdnra", "cdnura", "ata", "atadja"):
    assert res.counts.get(k, 0) > 0, f"table {k} was not read"

msg = import_master_template(db, client, project, filled, settings, schema)
assert msg[1], msg[2]
assert import_master_template(db, client, project, filled, settings, schema)[1] is None   # same file twice -> skipped
sets = db.datasets(pid)
docs = merge_docs([], res.docs, "replace")

# ---------------------------------------------------------------- 4. generate every CSV
issues = gstr1.validate_sets(sets, [], ctx) + validate_all([], client, docs, settings, (9, 2026))
out_base = tmp / "chosen_folder"
out_base.mkdir()
r = export_project(client, project, [], sets, docs, issues, settings, out_base, confirm_overwrite=lambda n: True)
out = r["folder"]
print(sorted(p.name for p in out.iterdir()))
expect = {spec["filename"] for k, spec in schema["tables"].items()}
got = {p.name for p in out.glob("*.csv")}
assert expect <= got, f"missing CSV files: {expect - got}"
assert {"gstr1_summary.xlsx", "validation_report.xlsx", "reconciliation_report.xlsx", "manifest.json"} <= \
    {p.name for p in out.iterdir()}


def read(name):
    with open(out / name, newline="", encoding="utf-8") as f:
        return list(csv.reader(f))


for key, spec in schema["tables"].items():
    rows = read(spec["filename"])
    assert rows[0] == [c["header"] for c in spec["columns"]], f"{key}: header differs"
    assert len(rows) > 1, f"{key}: no data rows"
    assert all(len(x) == len(rows[0]) for x in rows), f"{key}: ragged rows"

b2b = read("b2b.csv")
h = b2b[0]
inv2 = [x for x in b2b[1:] if x[h.index("Invoice Number")] == "INV002"]
assert len(inv2) == 2, "invoice with two rates must have two rows"
assert {x[h.index("Rate")] for x in inv2} == {"18", "5"}
assert inv2[0][h.index("Invoice Value")] == inv2[1][h.index("Invoice Value")] == "6425.00"
assert any(x[h.index("Invoice Type")] == "SEZ supplies with payment" for x in b2b[1:])

b2cs = read("b2cs.csv")
assert len(b2cs) - 1 == 4, "b2cs summary rows"
ex = read("exemp.csv")
assert len(ex) - 1 == 4 and ex[0][1].startswith("Nil Rated")
print("summary:")
for s in r["summary"]:
    print(f"  {s['key']:8} rows={s['rows']:3} taxable={s['taxable']:>12} igst={s['igst']:>10} cgst={s['cgst']:>9} "
          f"sgst={s['sgst']:>9}")
bad = [x for x in r["recon"] if x["status"] == "DIFFERENCE"]
print("recon differences:", [(x["check"], x["difference"]) for x in bad])
assert not [x for x in bad if "input rows vs generated" in x["check"]], "aggregation integrity failed"

# ---------------------------------------------------------------- 5. bad file: row-level errors, nothing imported
bad_path = tmp / "bad.xlsx"
wb = load_workbook(filled)
b = wb[schema["tables"]["b2b"]["sheet"]]
hdr = [c.value for c in b[1]]
b.cell(2, hdr.index("GSTIN/UIN of Recipient") + 1).value = "NOTAGSTIN"
b.cell(3, hdr.index("Invoice date") + 1).value = "31-02-2026"
b.cell(4, hdr.index("Taxable Value") + 1).value = "abc"
b.cell(5, hdr.index("Place Of Supply") + 1).value = "Atlantis"
c = wb[schema["tables"]["cdnr"]["sheet"]]
ch = [x.value for x in c[1]]
c.cell(2, ch.index("Note Type") + 1).value = "X"
wb["Return Info"]["B3"].value = "29" + own[2:]      # wrong GSTIN
wb.save(bad_path)
ctx2 = gstr1.make_ctx(client, project, settings, schema)
rb = gstr1.read_master_template(bad_path, schema, ctx2, "bad.xlsx")
for i in rb.errors:
    print("  ERR", i.ref, "|", i.message)
assert len(rb.errors) >= 5
assert all(i.ref for i in rb.errors), "every error must carry sheet + row"
m2 = import_master_template(db, client, project, bad_path, settings, schema)
assert m2[1] is None and len(db.datasets(pid)) == 1, "bad file must not be imported"

# ---------------------------------------------------------------- 6. overwrite protection
try:
    export_project(client, project, [], sets, docs, issues, settings, out_base)
    raise AssertionError("overwrote without confirmation")
except ValueError:
    pass
print("ALL MASTER TEMPLATE TESTS PASSED")
