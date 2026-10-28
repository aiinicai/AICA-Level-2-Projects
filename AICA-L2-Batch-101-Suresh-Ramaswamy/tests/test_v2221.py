"""v2.2.21: one printed tax invoice per sheet; document numbers like DM/001/2026-2027."""
import sys, tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from openpyxl import Workbook
from app.extractors.excel_reader import read_workbook
from app.series import build_runs, split_no

assert split_no("DM/001/2026-2027") == ("DM/", 1, 3, "/2026-2027")
assert split_no("INV/0007") == ("INV/", 7, 4, "")
runs = build_runs([(f"DM/{n:03d}/2026-2027", "X", "Invoices for outward supply") for n in (1, 2, 3, 5)], "prefix")
assert [(r["from_no"], r["to_no"], r["total"]) for r in runs] == [("DM/001/2026-2027", "DM/003/2026-2027", 3),
                                                                  ("DM/005/2026-2027", "DM/005/2026-2027", 1)], runs

wb = Workbook(); wb.remove(wb.active)
for no, d, vals in (("001", "07.04.2026", (1000, 2000)), ("002", "10.04.2026", (500,))):
    ws = wb.create_sheet(f"{no}-VB")
    ws["C5"], ws["I5"] = "Regd. Off\nDEMO PLASTICS", "TAX INVOICE"
    ws["I6"] = f" INVOICE NO :  DM/{no}/2026-2027        DATE :{d}\n PAN : ELSPP6300K"
    ws["C8"] = "GSTIN : 33DEMOA1234B1ZD"
    ws["C10"] = "CONSIGNEE :\n To,\n     M/s.DEMO AUTO PARTS\n     GSTIN:33AAAAA0000A1Z5\n     State : Tamilnadu"
    for c, h in zip("CDHIJMO", ("S.No", "Description / Part Name", "HSN/SAC Code", "Quantity", "Unit ", "Basic Rate/Unit", "Assessable Value (Rs.)")):
        ws[f"{c}14"] = h
    for i, v in enumerate(vals):
        r = 15 + 2 * i
        ws[f"C{r}"], ws[f"D{r}"], ws[f"H{r}"], ws[f"I{r}"], ws[f"J{r}"], ws[f"M{r}"], ws[f"O{r}"] = i + 1, f"ITEM {i}", 84501900, 10, "Nos", v / 10, v
    t = sum(vals)
    ws["I33"], ws["O33"] = "Total assessable  value", t
    ws["I34"], ws["O34"] = "SGST @ 9%", t * .09
    ws["I35"], ws["O35"] = "CGST @ 9%", t * .09
    ws["I37"], ws["O37"] = "TOTAL INVOICE VALUE", t * 1.18
p = Path(tempfile.mkdtemp()) / "inv.xlsx"; wb.save(p)
res = read_workbook(p, {"gstin": "33DEMOA1234B1ZD", "name": "DM"}, {})
assert res.kind == "invoice-per-sheet" and len(res.lines) == 3, (res.kind, res.messages)
l = res.lines[0]
assert (l.invoice_no, str(l.invoice_date), l.recipient_gstin, l.recipient_name, l.pos, l.rate) == \
       ("DM/001/2026-2027", "2026-04-07", "33AAAAA0000A1Z5", "DEMO AUTO PARTS", "33", 18), l.__dict__
assert float(l.cgst) == 90.0 and float(l.sgst) == 90.0 and float(res.lines[1].taxable) == 2000
print("ALL v2.2.21 TESTS PASSED")
