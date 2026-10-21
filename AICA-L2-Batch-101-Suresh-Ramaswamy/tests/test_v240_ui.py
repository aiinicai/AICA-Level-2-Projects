"""v2.4.0 screens (stand-in Tk): whole-year upload and the reconciliation tables on tab 15:
python -m tests.test_v240_ui"""
import os, shutil, sys, tempfile
from pathlib import Path
tmp = Path(tempfile.mkdtemp()); os.environ["CAGST_DATA"] = str(tmp / "data")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / "tools"))
from tests import tk_stub  # noqa: E402
tk_stub.install()
from app import ui, returns as R  # noqa: E402
import make_demo_returns as demo  # noqa: E402

fd, mb = ui.filedialog, ui.messagebox
app = ui.App()
assert app.nb.tab(app.nb.tabs()[-1], "text") == "15 YTD & Annual"
app.cv["name"].set("Demo Traders"); app.cv["gstin"].set(demo.GSTIN); app.cv["legal_name"].set("Demo Traders")
app.cv["reg_type"].set("Regular"); app.cv["return_type"].set("Monthly")
app.save_client()
app.proj_client.current(0); app.pm.set("July"); app.py.set("2026"); app.new_return()
app.refresh_annual()
app.an_fy.set("2026-27")

# year folder: Excel files in month folders + the demo filed returns (read through the text reader here)
S = ROOT / "sample_data" / "Demo_Samples_Pack"
yr = tmp / "FY 2026-27"
(yr / "2026-07 July").mkdir(parents=True); (yr / "2026-09 September").mkdir(); (yr / "Filed").mkdir()
shutil.copy(S / "03_Demo_Master_Template_FILLED_Jul2026.xlsx", yr / "2026-07 July")
shutil.copy(S / "01_Demo_Sales_Register_Sep2026.xlsx", yr / "2026-09 September")
for name in demo.FILES:
    (yr / "Filed" / name).write_bytes(b"%PDF-1.4 demo " + name.encode())
texts = {name: demo.to_text(rows()) for name, (rows, _) in demo.FILES.items()}
g9name = f"GSTR9_{demo.GSTIN}_032027.pdf"                      # FY 2026-27 is not over: must be left out
(yr / "Filed" / g9name).write_bytes(b"%PDF-1.4 demo g9")
texts[g9name] = demo.to_text(demo.g9_rows(demo.FY, demo.G9))
R.pdf_pages_text = lambda p, s=None: ([texts.get(Path(p).name, "")], False, [])

fd.askdirectory = lambda **k: str(yr)
seen = {}


class FakeDialog:
    def __init__(self, parent, fy, entries, client):
        seen["entries"] = entries
        self.result = [e for e in entries if e["include"]]
        self.approve_clean = True


ui.YearImportDialog = FakeDialog
mb.askyesno = lambda *a, **k: False
app.year_upload("folder")
kinds = sorted(e["kind"] for e in seen["entries"] if e["include"])
assert kinds == ["GSTR1", "GSTR1", "GSTR3B", "GSTR3B", "excel", "master"], kinds
g9e = [e for e in seen["entries"] if e["kind"] == "GSTR9"][0]
assert not g9e["include"] and "not over" in g9e["note"], g9e["note"]
assert len(app.an_filed.rows()) == 4, app.an_filed.rows()
assert any(r[0] == "MISSING" and r[2] == "GSTR-9 not uploaded" for r in app.an_rc_trees[-1][3].rows())
rc = [r for _, _, _, t in app.an_rc_trees for r in t.rows()]
assert any(r[0] == "DIFF" and r[1] == "Aug 2026" and r[2].startswith("3.1(a)") for r in rc), rc
assert any(r[0] == "MISSING" and r[1] == "Sep 2026" for r in rc)
assert "reconciliation" in app.an_info.get()
assert app.fill_recon_table() == len(app._an_recon)          # every row is in exactly one table
t13 = {name: tree for name, _, _, tree in app.an_rc_trees}
assert any(r[0] == "DIFF" and r[1] == "Aug 2026" for r in t13["2 GSTR-1 vs GSTR-3B"].rows())
assert all(r[0] == "OK" for r in t13["1 Books vs GSTR-1"].rows())
assert any(r[0] == "MISSING" for r in t13["Missing - information"].rows())
# export with the two new sheets
out = tmp / "ytd.xlsx"
fd.asksaveasfilename = lambda **k: str(out)
app.annual_export()
from openpyxl import load_workbook
assert load_workbook(out).sheetnames[-7:] == ["1 Books vs GSTR-1", "2 GSTR-1 vs GSTR-3B", "3 GSTR-3B tax paid",
                                              "4 GSTR-9 vs GSTR-1+3B", "5 Books vs GSTR-9", "Missing - information",
                                              "Filed returns"], load_workbook(out).sheetnames
# remove a filed return
app.an_filed.selection_set(["fr3"]) if hasattr(app.an_filed, "selection_set") else None
mb.askyesno = lambda *a, **k: True
app.remove_filed_return()
print("ALL v2.4.0 UI TESTS PASSED")
