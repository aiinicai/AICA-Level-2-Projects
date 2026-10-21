import sys, tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from app.series import build_runs, gap_report
from app.models import DocSeries
from app.exporters import register_table13
from app.schema import load_schema

entries = [(f"INV/2627/{n:04d}", "X", "Invoices for outward supply") for n in (894, 896, 897, 898)]
runs = build_runs(entries, "prefix")
assert len(gap_report(runs)) == 1                       # 895 missing
typed = DocSeries(nature="Invoices for outward supply", from_no="INV/2627/0894", to_no="INV/2627/0898",
                  total=5, cancelled=1, origin="manual")
assert gap_report(runs, [typed]) == []                  # accounted for as cancelled
typed0 = DocSeries(nature="Invoices for outward supply", from_no="INV/2627/0894", to_no="INV/2627/0898",
                   total=5, cancelled=0, origin="manual")
assert len(gap_report(runs, [typed0])) == 1             # no cancelled declared -> still reported
der = DocSeries(nature="Invoices for outward supply", from_no="INV/2627/0894", to_no="INV/2627/0898",
                total=5, cancelled=1, origin="derived")
assert len(gap_report(runs, [der])) == 1                # unverified series do not count

# a Smart Billing style register with a Table 13 sheet
from openpyxl import Workbook
tmp = Path(tempfile.mkdtemp())
wb = Workbook()
wb.active.title = "Sales Register"
wb.active.append(["Invoice No", "Invoice Date"])
ws = wb.create_sheet("Documents Issued (Table 13)")
ws.append(["Nature of document", "Sr. No. From", "Sr. No. To", "Total Number", "Cancelled", "Net Issued"])
ws.append(["Invoices for outward supply", "INV/2627/0894", "INV/2627/0898", 5, 1, 4])
p = tmp / "reg.xlsx"
wb.save(p)
res = register_table13(p, load_schema())
assert res and not res[1] and res[0][0].total == 5 and res[0][0].cancelled == 1, res
wb2 = Workbook()
wb2.active.title = "Sheet1"
p2 = tmp / "plain.xlsx"
wb2.save(p2)
assert register_table13(p2, load_schema()) is None
print("ALL v2.2.20 TESTS PASSED")
