"""Bank audit report exports: distinct-account exposure, never summed finding exposures."""
import csv
import json
from decimal import Decimal
from pathlib import Path
from html import escape
from audit_engine import dashboard,stamp

def safe(value):
    if isinstance(value,Decimal):return float(value)
    if isinstance(value,(dict,list)):value=json.dumps(value,default=str)
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):return "'"+value
    return value

def account_rows(result):
    out=[]
    for a in result['account_summary']:
        fs=[f for f in result['findings'] if f['account_id']==a['account_id']]
        coverage=result['coverage_reviews'].get(a['account_id'],{})
        out.append(dict(**a,total_findings=len(fs),open_findings=sum(f['review_status']!='Reviewed' for f in fs),
          coverage_reviewer=coverage.get('reviewer',''),coverage_assessment=coverage.get('assessment',''),
          coverage_reviewed_utc=coverage.get('reviewed_utc',''),
          scope='No automated classification or clean opinion; supplied data only.'))
    return out

def sections(result):
    ds=dashboard(result)
    output=[('Audit overview',[dict(metric=k,value=v) for k,v in ds.items()]),('Account coverage',account_rows(result)),
      ('Findings',result['findings']),('Drawing power',result['drawing_power']),('Conduct',result['conduct']),('Term dues',result['term_dues']),
      ('Policy assumptions',[dict(setting=k,value=v) for k,v in result['policy'].items()]),('Review events',result['events']),('AI drafts',result['ai_notes'])]
    output.extend((k+' source',v) for k,v in result['raw'].items())
    output.append(('Run metadata',[dict(item='Audit date',value=result['as_of']),dict(item='Created/exported UTC',value=stamp()),
       dict(item='Scope',value='Illustrative bank-policy screens; no NPA classification, provision calculation, document authentication or final audit opinion.'),
       *[dict(item=k+' SHA256',value=v) for k,v in result.get('input_hashes',{}).items()]]))
    return output

def export_excel(result,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook();wb.remove(wb.active)
    for name,rows in sections(result):
        rows=rows or [{'message':'No supplied records / no findings for this schedule'}]
        headers=list(dict.fromkeys(k for r in rows for k in r))
        ws=wb.create_sheet(name[:31]);ws.append([safe(h) for h in headers])
        for r in rows:ws.append([safe(r.get(k,'')) for k in headers])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='173B50');c.font=Font(bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True)
        ws.row_dimensions[1].height=32
        for i,k in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(i)].width=min(60,max(18,len(k)+3,max(len(str(r.get(k,''))) for r in rows)+2))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(wrap_text=True,vertical='top')
                if isinstance(cell.value,(float,Decimal)):cell.number_format='#,##0.00;[Red](#,##0.00);0.00'
                if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='EFF4F7')
    wb.save(path)

def export_csv(result,folder):
    p=Path(folder);p.mkdir(parents=True,exist_ok=True)
    for name,rs in sections(result):
        rows=rs or [{'message':'No supplied records / no findings for this schedule'}]
        heads=list(dict.fromkeys(k for r in rows for k in r))
        with (p/(name.replace(' ','_')+'.csv')).open('w',newline='',encoding='utf-8-sig') as f:
            w=csv.writer(f);w.writerow([safe(k) for k in heads]);w.writerows([safe(r.get(k,'')) for k in heads] for r in rows)

def export_html(result,path):
    e=lambda x:escape(str(x))
    html='''<!doctype html><html><head><meta charset="utf-8"><title>Bank advances audit review</title><style>
    body{font:15px Arial;margin:32px;color:#173b50}p{line-height:1.5}h1{font-size:28px}.stats{display:flex;flex-wrap:wrap;gap:12px}.stats div{background:#edf4f6;padding:12px;min-width:140px}table{border-collapse:collapse;width:100%;font-size:12px}td,th{border:1px solid #ccd9df;padding:8px;text-align:left;vertical-align:top}th{background:#173b50;color:white}.high{color:#9c2626}@media print{body{margin:8px}thead{display:table-header-group}}
    </style></head><body><h1>Bank Advances Audit and Exception Review</h1>'''
    html+='<p>Audit date '+e(result['as_of'])+' | '+e(result['policy']['profile_name'])+' | UTC '+e(stamp())+'</p>'
    html+='<div class="stats">'+''.join('<div><b>'+e(k.replace('_',' '))+'</b><p>'+e(v)+'</p></div>' for k,v in dashboard(result).items())+'</div>'
    html+='<p>Exposure totals count each supplied account once. Multiple findings do not multiply loan exposure. Criteria are supplied audit screens and sanction assumptions, not an automatic RBI classification/provisioning decision. Missing history is not a clean result.</p>'
    for title,rows,columns in [('Account coverage',account_rows(result),['account_id','borrower_name','outstanding','total_findings','open_findings','coverage_reviewer','coverage_assessment']),
       ('Audit findings',result['findings'],['finding_id','priority','title','condition','review_status','disposition','reviewer','auditor_conclusion','recommended_action']),
       ('Proposed reporting points',[f for f in result['findings'] if f['review_status']=='Reviewed' and f['disposition']=='Potential reporting point'],['finding_id','borrower_name','proposed_reporting_point','management_response'])]:
        html+='<h2>'+e(title)+'</h2><table><thead><tr>'+''.join('<th>'+e(c.replace('_',' '))+'</th>' for c in columns)+'</tr></thead><tbody>'
        html+=''.join('<tr>'+''.join('<td>'+e(r.get(c,''))+'</td>' for c in columns)+'</tr>' for r in rows)
        if not rows:html+='<tr><td colspan="'+str(len(columns))+'">No items in this section.</td></tr>'
        html+='</tbody></table>'
    html+='<h2>Scope and assumptions</h2><p>Only supplied facilities/records are covered. Source references are not authenticated documents. Proposed observations and bank responses require independent auditor evaluation. Local sessions/event logs are editable and not tamper-proof.</p></body></html>'
    Path(path).write_text(html,encoding='utf-8')
