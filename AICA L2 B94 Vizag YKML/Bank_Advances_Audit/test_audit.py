import copy
import csv
import io
import json
import tempfile
import unittest
from datetime import date,timedelta
from decimal import Decimal
from pathlib import Path
from audit_engine import *
from reports import export_excel,export_html,export_csv
from ai_review import request_draft
ROOT=Path(__file__).resolve().parent

class AuditTests(unittest.TestCase):
    def setUp(self):
        self.paths={k:ROOT/'sample_data'/f'{k}_demo.csv' for k in SCHEMAS}
        self.r=load_audit(self.paths,'2026-09-30')
    def account(self,aid):return next(a for a in self.r['raw']['accounts'] if a['account_id']==aid)
    def has(self,aid,code,result=None):return any(f['finding_id']==aid+':'+code for f in (result or self.r)['findings'])
    def rerun(self,data=None,p=None):return run_audit(data or self.r['raw'],'2026-09-30',p)
    def valid_review(self):return dict(review_status='Reviewed',disposition='Potential reporting point',reviewer='Demo auditor',
      evidence_assessment='Fictional sanction and DP register inspected; supporting certificate still requested.',management_response='Branch says approval is pending.',
      auditor_conclusion='The supplied source shows a DP difference; investigate the sanction basis.',recommended_action='Obtain current statement and document applicable deduction method.',
      proposed_reporting_point='DP computation requires reconciliation to sanctioned method; branch evidence pending.')
    def test_demo_controls(self):
        d=dashboard(self.r)
        self.assertEqual((d['accounts'],d['borrowers'],d['findings']),(8,7,18))
        self.assertEqual(d['outstanding'],Decimal(4650000))
        self.assertEqual(d['high_priority_exposure'],Decimal(3550000))
        self.assertEqual(d['high_priority_accounts'],6)
        self.assertEqual(d['accounts_pending_coverage'],8)
    def test_dp_before_after_creditors(self):
        a=copy.deepcopy(self.account('CC001'));s=self.r['raw']['stocks'][0]
        self.assertEqual(drawing_power(a,s)['capped_audit_dp'],Decimal(705000))
        a['creditor_method']='After margin'
        self.assertEqual(drawing_power(a,s)['capped_audit_dp'],Decimal(655000))
    def test_dp_capped_by_sanction(self):
        dp=next(x for x in self.r['drawing_power'] if x['account_id']=='CC003')
        self.assertEqual(dp['uncapped_dp'],Decimal(1500000));self.assertEqual(dp['capped_audit_dp'],Decimal(1000000))
    def test_creditors_cannot_make_stock_component_negative(self):
        a=self.account('CC001');s=copy.deepcopy(self.r['raw']['stocks'][0]);s['trade_creditors']=Decimal(2000000)
        self.assertEqual(drawing_power(a,s)['stock_dp'],Decimal(0))
        a=copy.deepcopy(a);a['creditor_method']='After margin'
        self.assertEqual(drawing_power(a,s)['stock_dp'],Decimal(0))
    def test_full_continuous_excess(self):
        self.assertTrue(self.has('CC003','CONDUCT_CONTINUOUS'))
        c=next(x for x in self.r['conduct'] if x['account_id']=='CC003')
        self.assertEqual(c['trailing_continuous_excess_days'],90)
    def test_equality_breaks_continuous_excess(self):
        data=copy.deepcopy(self.r['raw'])
        row=next(d for d in data['daily'] if d['account_id']=='CC003' and d['date']==date(2026,8,1))
        row['closing_balance']=min(row['sanctioned_limit'],row['drawing_power'])
        r=self.rerun(data)
        self.assertFalse(self.has('CC003','CONDUCT_CONTINUOUS',r))
        self.assertEqual(next(x for x in r['conduct'] if x['account_id']=='CC003')['excess_days_observed'],89)
    def test_missing_date_prevents_full_window_conclusion(self):
        data=copy.deepcopy(self.r['raw']);data['daily']=[d for d in data['daily'] if not(d['account_id']=='CC003' and d['date']==date(2026,8,1))]
        r=self.rerun(data)
        self.assertTrue(self.has('CC003','CONDUCT_COVERAGE',r));self.assertFalse(self.has('CC003','CONDUCT_CONTINUOUS',r))
    def test_partial_history_does_not_run_no_credit_screen(self):
        data=copy.deepcopy(self.r['raw'])
        for d in data['daily']:
            if d['account_id']=='CC004':d['credits']=Decimal(0)
        r=self.rerun(data)
        self.assertTrue(self.has('CC004','CONDUCT_COVERAGE',r));self.assertFalse(self.has('CC004','CONDUCT_NO_CREDITS',r))
    def test_historical_limits_not_current_dp_used(self):
        data=copy.deepcopy(self.r['raw'])
        for d in data['daily']:
            if d['account_id']=='CC003':d['sanctioned_limit']=Decimal(2000000);d['drawing_power']=Decimal(2000000)
        self.assertFalse(self.has('CC003','CONDUCT_CONTINUOUS',self.rerun(data)))
    def test_no_credit_and_interest_shortfall(self):
        self.assertTrue(self.has('OD005','CONDUCT_NO_CREDITS'));self.assertTrue(self.has('OD005','CONDUCT_CREDIT_SHORTFALL'))
    def test_repayment_age_and_partial_amounts(self):
        d=next(x for x in self.r['term_dues'] if x['due_id']=='DUE06')
        self.assertEqual((d['days_since_due'],d['unpaid_amount']),(92,Decimal(110000)))
        d=next(x for x in self.r['term_dues'] if x['due_id']=='DUE07');self.assertEqual(d['unpaid_amount'],Decimal(22000))
    def test_exact_term_screen_boundary_not_over(self):
        data=copy.deepcopy(self.r['raw']);data['dues'][0]['due_date']=date(2026,7,2)
        r=self.rerun(data)
        self.assertFalse(self.has('TL006','REPAYMENT_AGED',r));self.assertTrue(self.has('TL006','REPAYMENT_UNPAID',r))
    def test_future_and_paid_dues_not_flagged(self):
        self.assertFalse(any(f['account_id']=='TL008' for f in self.r['findings']))
        self.assertNotIn('DUE08B',{x['due_id'] for x in self.r['term_dues']})
    def test_borrower_cross_facility_review_no_classification_change(self):
        self.assertTrue(self.has('TL007','BORROWER_REVIEW'))
        self.assertTrue(self.has('CC003','BORROWER_REVIEW'))
        self.assertTrue(all(a['bank_classification']=='Standard' for a in self.r['raw']['accounts']))
    def test_insurance_expiry_boundary(self):
        data=copy.deepcopy(self.r['raw']);doc=next(x for x in data['documents'] if x['account_id']=='CC002' and x['document_type']=='Insurance')
        doc['valid_until']=date(2026,9,30);r=self.rerun(data)
        self.assertFalse(any(f['account_id']=='CC002' and f['title'].startswith('Insurance') for f in r['findings']))
    def test_na_document_not_automatic_clearance(self):
        data=copy.deepcopy(self.r['raw']);doc=next(x for x in data['documents'] if x['account_id']=='CC001' and x['document_type']=='Insurance')
        doc['status']='Not applicable';doc['remarks']='Branch assertion only.';r=self.rerun(data)
        self.assertTrue(any(f['account_id']=='CC001' and f['title'].startswith('Insurance') for f in r['findings']))
    def test_unknown_account_and_duplicate_day_rejected(self):
        data=copy.deepcopy(self.r['raw']);data['daily'][0]['account_id']='UNKNOWN'
        with self.assertRaises(ValueError):self.rerun(data)
        data=copy.deepcopy(self.r['raw']);data['daily'].append(copy.deepcopy(data['daily'][0]))
        with self.assertRaises(ValueError):self.rerun(data)
    def test_policy_change_affects_tests(self):
        p=copy.deepcopy(DEFAULT_POLICY);p['term_overdue_review_days']=100;p['stock_stale_days']=150
        r=self.rerun(p=p)
        self.assertFalse(self.has('TL006','REPAYMENT_AGED',r));self.assertFalse(self.has('CC002','DP_STALE',r))
    def test_invalid_policy_rejected(self):
        for k,v in [('conduct_window_days',0),('stock_stale_days','90.5'),('tolerance_inr','-1')]:
            p=copy.deepcopy(DEFAULT_POLICY);p[k]=v
            with self.assertRaises(ValueError):validate_policy(p)
    def test_no_finding_is_not_a_completed_account_review(self):
        self.assertFalse(any(f['account_id']=='CC001' for f in self.r['findings']))
        review_coverage(self.r,'CC001','Demo auditor','Fictional source population and policy reviewed; no exceptions under supplied screens.')
        self.assertEqual(dashboard(self.r)['accounts_pending_coverage'],7)
        self.assertEqual(dashboard(self.r)['closing_state'],'Pending audit review')
    def test_review_required_fields_and_proposed_reporting_wording(self):
        fid='CC002:DP_DIFFERENCE';v=self.valid_review()
        for change in [dict(reviewer=''),dict(evidence_assessment=''),dict(auditor_conclusion=''),dict(recommended_action=''),dict(proposed_reporting_point=''),dict(disposition='Unresolved')]:
            with self.assertRaises(ValueError):update_review(self.r,fid,{**v,**change})
        update_review(self.r,fid,v)
        self.assertEqual(dashboard(self.r)['potential_reporting_points'],1)
        self.assertEqual(dashboard(self.r)['unresolved_findings'],17)
    def test_session_hash_date_policy_and_reviews(self):
        update_review(self.r,'CC002:DP_DIFFERENCE',self.valid_review());review_coverage(self.r,'CC001','Auditor','Source completeness checked.')
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'session.json';save_session(self.r,p)
            restored=restore_session(load_audit(self.paths,'2026-09-30'),p)
            self.assertEqual(dashboard(restored)['potential_reporting_points'],1)
            self.assertEqual(dashboard(restored)['accounts_pending_coverage'],7)
            changed=load_audit(self.paths,'2026-09-30');changed['policy']['stock_stale_days']=91
            with self.assertRaises(ValueError):restore_session(changed,p)
            changed=load_audit(self.paths,'2026-09-30');changed['input_hashes']['daily']='changed'
            with self.assertRaises(ValueError):restore_session(changed,p)
    def test_prompt_selected_account_and_no_npa_decision(self):
        text=draft_prompt(self.r,'CC002:DP_DIFFERENCE')
        self.assertIn('CC002',text);self.assertNotIn('Example Healthy Traders',text)
        self.assertIn('Do not call a screen an NPA decision',text)
    def test_excel_html_csv_exports_and_injection(self):
        self.r['findings'][0]['management_response']='=HYPERLINK("unsafe")'
        self.r['findings'][0]['auditor_conclusion']='<script>alert(1)</script>'
        with tempfile.TemporaryDirectory() as d:
            x=Path(d)/'review.xlsx';export_excel(self.r,x)
            from openpyxl import load_workbook
            wb=load_workbook(x,read_only=True)
            try:
                self.assertEqual(wb.sheetnames[0],'Audit overview');self.assertIn('Policy assumptions',wb.sheetnames)
                for sheet in wb:
                    for row in sheet:self.assertFalse(any(c.data_type=='f' for c in row))
            finally:wb.close()
            h=Path(d)/'review.html';export_html(self.r,h);self.assertNotIn('<script>',h.read_text());self.assertIn('&lt;script&gt;',h.read_text())
            export_csv(self.r,Path(d)/'csv');self.assertTrue((Path(d)/'csv'/'Findings.csv').exists())
    def test_api_mock_not_live(self):
        def opener(req,timeout):
            body=json.loads(req.data);self.assertFalse(body['store'])
            return io.StringIO(json.dumps(dict(status='completed',id='mock',model='mock',output=[dict(type='message',content=[dict(type='output_text',text='UNAPPROVED draft')])],usage={})))
        self.assertIn('UNAPPROVED',request_draft('fictional prompt','fake-key','mock',opener=opener)['text'])

if __name__=='__main__':unittest.main()
