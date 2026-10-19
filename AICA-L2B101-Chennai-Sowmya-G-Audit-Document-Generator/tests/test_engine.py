"""Automated tests for audit_engine.py.

Run from the project folder:   python -m pytest -v tests
All test data is fictitious (see tests/make_fixtures.py).
"""
import io
import itertools
import os
import sys

import pytest
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
import audit_engine as eng                                  # noqa: E402

FIX = os.path.join(HERE, "fixtures")


# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------

def details(**changes):
    """A complete, valid set of details for a dummy company."""
    d = {
        "company": "MEENAKSHI TEXTILE TRADERS PRIVATE LIMITED",
        "cin": "U74999TN2019PTC123456",
        "address": "NO 12, SOUTH MASI STREET, MADURAI, Tamil Nadu, India, 625001",
        "constitution": "Private", "private": True, "fy": "2025-2026", "result": "profit",
        "cash_flow": True, "opinion": "Unmodified", "basis_text": "", "eom": "",
        "caro": True, "ifc": False, "litigation": False,
        "audit_trail": eng.AUDIT_TRAIL_OPTIONS[0], "audit_trail_from": "",
        "firm": "Sowmya G & Associates", "auditor": "CA. Sowmya G", "designation": "Proprietor",
        "frn": "012345S", "membership": "234567", "udin": "26234567BKABCD1234",
        "place": "Madurai", "firm_address1": "Office address line 1,",
        "firm_address2": "Madurai - 6250XX.", "phone": "+91 00000 00000",
        "email": "audit@example.com", "website": "www.example.com",
        "logo": eng.default_logo(),
        "engagement_date": "11/08/2026", "mrl_date": "14/08/2026", "report_date": "05/09/2026",
        "directors": [{"name": "KARTHIK SUNDARAM", "din": "01234567", "designation": "Director"},
                      {"name": "VALLIAMMAI RAMANATHAN", "din": "07654321",
                       "designation": "Managing Director"}],
    }
    for key, _, default in eng.CARO_ITEMS:
        d[key] = default
    d.update(changes)
    d.update(eng.fy_details(d["fy"]))
    return d


def text_of(maker, d):
    """Generate a document in memory and return all its text (body, tables, headers)."""
    buffer = io.BytesIO()
    maker(d, buffer)
    doc = Document(io.BytesIO(buffer.getvalue()))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        for table in section.header.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts), doc


# ------------------------------------------------------------------
# 1. Reading the MCA master data PDF
# ------------------------------------------------------------------

MCA_CASES = {
    "mca_v3_one_line_address.pdf": "NO 5, WEST VELI STREET, MADURAI, Tamil Nadu, India, 625001",
    "mca_v3_two_line_address.pdf": "NO 12, SOUTH MASI STREET, NEAR MEENAKSHI AMMAN TEMPLE, "
                                   "MADURAI, Madurai, Tamil Nadu, India, 625001",
    "mca_v3_three_line_address.pdf": "DOOR NO 4/122, PLOT NO 7, SRI RAM NAGAR 2ND STREET, "
                                     "K.PUDUR, OPP. TO GOVT HIGHER SECONDARY SCHOOL, MADURAI, "
                                     "Madurai, Tamil Nadu, India, 625007",
    "bordered_table.pdf": "PLOT 14, SIPCOT INDUSTRIAL ESTATE, KAPPALUR, MADURAI, Tamil Nadu, "
                          "India, 625008",
    "value_below_label.pdf": "12 ANNA SALAI, CHENNAI 600002",
}


@pytest.mark.parametrize("name,address", MCA_CASES.items())
def test_registered_office_read_in_full(name, address):
    found, _ = eng.extract_master_data(os.path.join(FIX, name))
    assert found["address"] == address


@pytest.mark.parametrize("name", MCA_CASES)
def test_books_of_account_address_never_taken(name):
    found, _ = eng.extract_master_data(os.path.join(FIX, name))
    assert "GODOWN" not in found["address"]
    assert "books" not in found["address"].lower()


def test_mca_v3_directors_name_and_din_only():
    found, _ = eng.extract_master_data(os.path.join(FIX, "mca_v3_two_line_address.pdf"))
    assert found["directors"] == [{"din": "01234567", "name": "KARTHIK SUNDARAM"},
                                  {"din": "07654321", "name": "VALLIAMMAI RAMANATHAN"}]


def test_bordered_table_directors_without_designation_or_category():
    found, _ = eng.extract_master_data(os.path.join(FIX, "bordered_table.pdf"))
    names = [x["name"] for x in found["directors"]]
    assert names == ["PRIYA NATARAJAN", "ARUL MURUGAN"]
    for word in ("Managing", "Director", "Promoter", "Professional"):
        assert all(word not in n for n in names)


def test_cin_company_and_constitution():
    found, _ = eng.extract_master_data(os.path.join(FIX, "mca_v3_one_line_address.pdf"))
    assert found["cin"] == "L17111TN2015PLC098765"
    assert found["company"] == "PANDIAN SPINNING MILLS LIMITED"
    assert found["constitution"] == "Public"


def test_constitution_from_cin_when_not_printed():
    found, _ = eng.extract_master_data(os.path.join(FIX, "value_below_label.pdf"))
    assert found["constitution"] == "Section 8 company"


def test_scanned_pdf_has_no_text():
    found, lines = eng.extract_master_data(os.path.join(FIX, "scanned_no_text.pdf"))
    assert lines == 0 and found["company"] == ""


# ------------------------------------------------------------------
# 2. Small helpers
# ------------------------------------------------------------------

def test_financial_year_details():
    assert eng.fy_details("2025-2026") == {"year_end": "31st March 2026",
                                           "year_end_short": "31.03.2026",
                                           "next_fy": "2026-2027"}


@pytest.mark.parametrize("text,ok", [("05/09/2026", True), ("31/02/2026", False),
                                     ("2026-09-05", False), ("", False)])
def test_date_check(text, ok):
    assert eng.check_date(text) is ok


def test_spacing_tidied_like_master_data():
    assert eng.tidy("NO 5,WEST  VELI STREET ,MADURAI") == "NO 5, WEST VELI STREET, MADURAI"


def test_safe_folder_name():
    assert eng.safe_name("A & B (India) Pvt. Ltd.") == "A_B_India_Pvt_Ltd"


# ------------------------------------------------------------------
# 3. Letter of Engagement
# ------------------------------------------------------------------

def test_engagement_letter():
    text, doc = text_of(eng.make_engagement_letter, details())
    assert "SOWMYA G & ASSOCIATES" in text                      # auditor's letterhead
    assert "financial year 2025-2026" in text
    assert "head office" not in text and "elsewhere" not in text
    assert "Artificial Intelligence" in text and "redacted" in text
    assert "Digital Personal Data Protection Act, 2023" in text
    assert "KARTHIK SUNDARAM" in text and "VALLIAMMAI RAMANATHAN" in text


# ------------------------------------------------------------------
# 4. DSC authorisation letter
# ------------------------------------------------------------------

@pytest.mark.parametrize("count", [1, 2, 4])
def test_one_dsc_page_per_director(count):
    people = [{"name": "DIRECTOR %d" % i, "din": "1000000%d" % i, "designation": "Director"}
              for i in range(count)]
    text, _ = text_of(eng.make_dsc_letters, details(directors=people))
    assert text.count("Authorization to obtain and use Digital Signature Certificate") == count
    assert "Financial Year 2025-2026 & Financial Year 2026-2027" in text
    assert "CIN: U74999TN2019PTC123456" in text                  # company letterhead


# ------------------------------------------------------------------
# 5. Letter of representation
# ------------------------------------------------------------------

def test_representation_letter_blanks_for_client_to_fill():
    text, doc = text_of(eng.make_mrl, details())
    assert "Depreciation amounting to Rs. ________/-" in text
    assert "Physical balance of cash as on 31.03.2026 was Rs. ________/-" in text
    related = [t for t in doc.tables
               if t.rows[0].cells[0].text == "Name of Party"]
    assert len(related) == 1 and len(related[0].rows) == 6        # header + 5 blank rows
    assert all(c.text == "" for row in related[0].rows[1:] for c in row.cells)
    bank = [t for t in doc.tables if "Bank & A/c No." in [c.text for c in t.rows[0].cells]]
    assert len(bank) == 2                                          # FDs and bank accounts


def test_representation_letter_follows_company_facts():
    text, _ = text_of(eng.make_mrl, details(caro_inventory=False, caro_borrow=False))
    assert "The company does not hold any inventory." in text
    assert "There are no borrowings from banks or financial institutions." in text
    text, _ = text_of(eng.make_mrl, details(caro_inventory=True, caro_borrow=True))
    assert "Inventories were physically verified" in text
    assert "The term loans have been used only for the purpose" in text


@pytest.mark.parametrize("option,expected", [
    (0, "has operated throughout the year"),
    (1, "was enabled with effect from 15/06/2025"),
    (2, "does not have a feature of recording audit trail"),
])
def test_representation_letter_audit_trail_wording(option, expected):
    d = details(audit_trail=eng.AUDIT_TRAIL_OPTIONS[option], audit_trail_from="15/06/2025")
    text, _ = text_of(eng.make_mrl, d)
    assert expected in text
    if option == 1:
        assert "01/04/2025 to 14/06/2025" in text


# ------------------------------------------------------------------
# 6. Auditor's report - every combination of the main choices
# ------------------------------------------------------------------

COMBINATIONS = list(itertools.product(
    eng.OPINIONS,                       # 4 opinions
    [0, 1, 2],                          # audit trail answers
    [True, False],                      # private / public
    [(True, True), (True, False), (False, True), (False, False)],   # CARO, IFC
    [True, False],                      # cash flow statement
))


@pytest.mark.parametrize("opinion,trail,private,annexures,cash_flow", COMBINATIONS)
def test_audit_report_combinations(opinion, trail, private, annexures, cash_flow):
    caro, ifc = annexures
    d = details(opinion=opinion, private=private, caro=caro, ifc=ifc, cash_flow=cash_flow,
                audit_trail=eng.AUDIT_TRAIL_OPTIONS[trail], audit_trail_from="15/06/2025",
                basis_text="" if opinion == "Unmodified" else "The matter is described here.")
    text, _ = text_of(eng.make_audit_report, d)

    assert "INDEPENDENT AUDITOR'S REPORT" in text
    assert "SOWMYA G & ASSOCIATES" in text                     # auditor's letterhead
    assert "Ind AS" not in text                                # Accounting Standards only
    assert "books of account are to be maintained" not in text

    heading = {"Unmodified": "Opinion", "Qualified": "Qualified Opinion",
               "Adverse": "Adverse Opinion", "Disclaimer": "Disclaimer of Opinion"}[opinion]
    assert "\n" + heading + "\n" in "\n" + text + "\n"
    if opinion != "Unmodified":
        assert "The matter is described here." in text

    # Cash flow statement mentioned only when prepared
    assert ("Statement of Cash Flows" in text) is cash_flow

    # Annexure lettering: CARO first, then IFC
    if caro:
        assert 'ANNEXURE "A"' in text and "Companies (Auditor's Report) Order, 2020" in text
        assert "(xxi)" in text
    else:
        assert "is not applicable to the Company." in text
    if ifc:
        assert ('ANNEXURE "B"' if caro else 'ANNEXURE "A"') in text
        assert "Report on the Internal Financial Controls" in text
    elif private:
        assert "G.S.R. 583(E)" in text

    # Section 197 and 177
    assert ("provisions of section 197 are not applicable" in text) is private
    if caro and not private:
        assert "sections 177 and 188" in text

    # Audit trail: clause (b) exception and Rule 11(g) wording
    exception = "except for the matter stated in paragraph 2(h)(vi)" in text
    assert exception is (trail != 0)
    if trail == 1:
        assert "01/04/2025 to 14/06/2025" in text
    if trail == 2:
        assert "unable to comment" in text


def test_caro_clauses_follow_company_facts():
    clauses = dict((label, txt) for label, txt, _ in eng.caro_clauses(details(
        caro_ppe=False, caro_inventory=False, caro_loans=False, caro_csr=False)))
    assert "does not hold any Property, Plant and Equipment" in clauses["(i)"]
    assert "not made investments" in clauses["(iii)"]
    assert "section 135 of the Act are not applicable" in clauses["(xx)"]


def test_word_files_open_and_are_a4():
    for maker in (eng.make_engagement_letter, eng.make_dsc_letters, eng.make_mrl,
                  eng.make_audit_report):
        _, doc = text_of(maker, details(ifc=True))
        section = doc.sections[0]
        assert round(section.page_width.inches, 2) == 8.27
        assert round(section.page_height.inches, 2) == 11.69
