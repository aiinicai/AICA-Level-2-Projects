"""
Automated tests for the greytHR Payroll Reconciliation Tool.

Run from the application folder:   python -m pytest -v tests
The real-file tests need the two salary statements in tests/data/ (names below).
The formula tests need LibreOffice (soffice) to recalculate the workbook; they are
skipped automatically if LibreOffice is not installed.
"""
import io
import os
import shutil
import subprocess
import sys
import tempfile
from decimal import Decimal

import openpyxl
import pandas as pd
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
import app  # noqa: E402

DATA = os.path.join(HERE, "data")
AUG = os.path.join(DATA, "Salary Statement-August 2026.xls")
SEP = os.path.join(DATA, "Salary Statement-September 26.xls")
real_files = pytest.mark.skipif(not (os.path.exists(AUG) and os.path.exists(SEP)),
                                reason="real salary statements not in tests/data")
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


def read(path):
    with open(path, "rb") as handle:
        return handle.read()


def run_real():
    return app.run_reconciliation(read(AUG), os.path.basename(AUG), "August 2026",
                                  read(SEP), os.path.basename(SEP), "September 2026")


def recalc(workbook_bytes):
    """Recalculate every formula with LibreOffice and return a values-only workbook."""
    folder = tempfile.mkdtemp()
    src = os.path.join(folder, "in.xlsx")
    out_dir = os.path.join(folder, "out")
    with open(src, "wb") as handle:
        handle.write(workbook_bytes)
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out_dir, src],
                   check=True, capture_output=True, timeout=300)
    return openpyxl.load_workbook(os.path.join(out_dir, "in.xlsx"), data_only=True)


def independent_totals(path):
    """Totals computed straight from the file with pandas - no app code involved."""
    df = pd.read_excel(path, dtype={"Employee No": str})
    df = df[df["Employee No"].notna()]
    return df.set_index("Employee No")


def i(value):
    """numpy number -> plain int so it can be compared with Decimal."""
    return int(value)


# ---------------------------------------------------------------------------
# Synthetic file builder
# ---------------------------------------------------------------------------
HEAD = ["Employee No", "Join Date", "Left?", "DAYS IN MONTH", "EMP EFFECTIVE WORKDAYS",
        "BASIC", "HRA", "BONUS", "GROSS", "PF", "INCOME TAX", "TOTAL DEDUCTIONS", "NET PAY"]


def row(eid, basic, hra, bonus=0, pf=1800, tax=0, days=30, paid=30, gross=None, td=None, net=None, join="01 Jan 2020"):
    g = basic + hra + bonus if gross is None else gross
    t = -(pf + tax) if td is None else td
    n = g + t if net is None else net
    return [eid, join, False, days, paid, basic, hra, bonus, g, pf, tax, t, n]


def csv_bytes(rows, head=HEAD, title_rows=0, total=True):
    lines = [["Salary Statement"] + [""] * (len(head) - 1)] * title_rows + [head] + rows
    if total:
        tot = ["", "", "", "Grand Total"] + [""] * (len(head) - 4)
        for i in range(5, len(head)):
            tot[i] = sum(Decimal(str(r[i])) for r in rows if str(r[i]).strip() not in ("", "abc"))
        lines.append(tot)
    buf = io.StringIO()
    pd.DataFrame(lines).to_csv(buf, header=False, index=False)
    return buf.getvalue().encode("utf-8")


def recon(prev_rows, curr_rows, **kw):
    return app.run_reconciliation(csv_bytes(prev_rows, **kw), "prev.csv", "Prev",
                                  csv_bytes(curr_rows, **kw), "curr.csv", "Curr")


BASE_PREV = [row("101", 50000, 20000), row("102", 40000, 16000, tax=2000), row("103", 30000, 12000)]
BASE_CURR = [row("101", 50000, 20000), row("102", 44000, 17600, tax=2500), row("104", 25000, 10000)]


# ===========================================================================
# Real August / September files
# ===========================================================================
@real_files
def test_columns_detected_and_classified():
    res = run_real()
    prev, curr = res["prev"], res["curr"]
    assert prev["header_row"] == 1 and curr["header_row"] == 1
    assert len(prev["headings"]) == 43 and len(curr["headings"]) == 40
    assert len(prev["earn_cols"]) == 25 and len(prev["ded_cols"]) == 8
    assert len(curr["earn_cols"]) == 24 and len(curr["ded_cols"]) == 7
    # Regression: the old substring rule treated these as attendance columns ("deveLOPment")
    for month in (prev, curr):
        assert "PROFESSIONAL DEVELOPMENT" in month["earn_cols"]
        assert "PROFESSIONAL DEVELOPMENT Arrears" in month["earn_cols"]
        roles = {c["heading"]: c["role"] for c in month["columns"]}
        for info in ("Join Date", "Left?", "Status", "DAYS IN MONTH", "EMP EFFECTIVE WORKDAYS"):
            assert roles[info] == app.ROLE_INFO
    assert {c["heading"]: c["role"] for c in prev["columns"]}["LOP"] == app.ROLE_INFO
    assert prev["convention"] == "NEGATIVE" and curr["convention"] == "NEGATIVE"
    assert prev["grand_total_row"] == 189 and curr["grand_total_row"] == 196


@real_files
def test_old_5000_gap_traced_to_professional_development():
    res = run_real()
    r1198 = res["prev"]["by_id"]["1198"][0]
    assert r1198["gross"] == 469343
    assert sum(r1198["earn"].values()) == 469343
    old_rule = ["lop", "days", "status", "left?", "join date", "department", "designation", "location",
                "date of joining", "date of leaving", "paid days", "arrear days", "days in month",
                "emp effective workdays", "workdays"]
    dropped = [k for k in r1198["earn"] if any(w in k.lower() for w in old_rule)]
    assert dropped == ["PROFESSIONAL DEVELOPMENT", "PROFESSIONAL DEVELOPMENT Arrears"]
    assert sum(v for k, v in r1198["earn"].items() if k not in dropped) == 464343  # the old app's figure


@real_files
def test_employee_counts_and_matching():
    res = run_real()
    assert res["headcount"] == {"prev": 187, "curr": 194, "joiners": 9, "leavers": 2, "dup_prev": 0,
                                "dup_curr": 0, "calc": 194, "diff": 0}
    assert res["joiners"] == ["8873", "8951", "8978", "8980", "8990", "8991", "8992", "9000", "9011"]
    assert res["leavers"] == ["1548", "2739"]
    assert len(res["continuing"]) == 185
    assert not res["dup_ids"]


@real_files
def test_source_checks_all_zero():
    res = run_real()
    assert len(res["checks"]) == 187 + 194
    assert all(c["earn_var"] == 0 and c["ded_var"] == 0 and c["net_var"] == 0 for c in res["checks"])
    assert res["fail_count"] == 0


@real_files
def test_bridges_against_independent_pandas_figures():
    aug, sep = independent_totals(AUG), independent_totals(SEP)
    cont = aug.index.intersection(sep.index)
    joiners, leavers = sep.index.difference(aug.index), aug.index.difference(sep.index)
    res = run_real()
    g, n = res["gross"], res["net"]
    assert g["prev"] == i(aug["GROSS"].sum()) == 50942601
    assert g["curr"] == i(sep["GROSS"].sum()) == 48168536
    assert g["joiners"] == i(sep.loc[joiners, "GROSS"].sum()) == 1137925
    assert g["leavers"] == -i(aug.loc[leavers, "GROSS"].sum()) == -2804386
    assert g["cont"] == i((sep.loc[cont, "GROSS"] - aug.loc[cont, "GROSS"]).sum()) == -1107604
    assert g["diff"] == 0 and g["unexplained"] == 0
    assert n["prev"] == i(aug["NET PAY"].sum()) and n["curr"] == i(sep["NET PAY"].sum())
    assert n["joiners"] == i(sep.loc[joiners, "NET PAY"].sum()) == 1017499
    assert n["leavers"] == -i(aug.loc[leavers, "NET PAY"].sum()) == -2272182
    assert n["cont"] == -640000 and n["diff"] == 0 and n["unexplained"] == 0
    # deductions as positive amounts
    assert n["ded_prev"] == -i(aug["TOTAL DEDUCTIONS"].sum()) == 11236493
    assert n["ded_curr"] == -i(sep["TOTAL DEDUCTIONS"].sum()) == 10357111
    assert n["cont_ded_total"] == -i((sep.loc[cont, "TOTAL DEDUCTIONS"] - aug.loc[cont, "TOTAL DEDUCTIONS"]).sum())


@real_files
def test_continuing_movements_and_components():
    res = run_real()
    rows = {r["id"]: r for r in res["cont_rows"]}
    r = rows["1293"]
    assert r["gross_change"] == -1445427 and r["onetime"] == -1445427
    r = rows["5821"]  # 18 of 30 paid days in September plus leave encashment
    assert r["gross_change"] == 103758 and r["onetime"] == 165888
    assert r["structure"] + r["paid_impact"] == 103758 - 165888
    for r in res["cont_rows"]:
        parts = r["structure"] + r["paid_impact"] + r["arrears"] + r["onetime"] + r["variable"] + r["unclassified"]
        assert parts + r["gross_unexplained"] == r["gross_change"]
        assert r["net_change"] == r["gross_change"] - r["ded_change"]
    assert all(t["diff"] == 0 for t in res["earn_table"] + res["ded_table"])
    only = {t["component"]: (t["in_prev"], t["in_curr"]) for t in res["earn_table"] + res["ded_table"]}
    for gone in ("GRATUITY", "Excellence Award", "SDM SDE Bonus", "OTHER DEDUCTION"):
        assert only[gone] == (True, False)
    for new in ("REFERRAL BONUS", "Recognition Bonus"):
        assert only[new] == (False, True)
    assert len(res["differences"]) == 133
    assert sum(1 for r in res["cont_rows"] if not r["has_difference"]) == 52


@real_files
def test_review_items_found():
    res = run_real()
    review = {(i["check"], i["emp"]) for i in res["issues"] if i["severity"] == "REVIEW"}
    assert ("Left? flag", "7945") in review and ("Join Date changed", "7945") in review
    for eid in ("5821", "8621", "8706"):
        assert ("Left? flag", eid) in review


@real_files
def test_xlsx_and_csv_give_identical_results():
    raw = pd.read_excel(AUG, header=None, dtype=object)
    title = pd.DataFrame([["ACME Pvt Ltd"] + [None] * (raw.shape[1] - 1),
                          ["Salary Statement August 2026"] + [None] * (raw.shape[1] - 1),
                          [None] * raw.shape[1]])
    shifted = pd.concat([title, raw], ignore_index=True)
    xbuf = io.BytesIO()
    shifted.to_excel(xbuf, header=False, index=False)
    csv = raw.copy()
    for col in range(7, raw.shape[1]):  # amounts as text with thousands separators
        csv[col] = [v if i == 0 or pd.isna(v) else f"{int(v):,}" for i, v in enumerate(raw[col])]
    cbuf = io.StringIO()
    csv.to_csv(cbuf, header=False, index=False)
    for data, name in ((xbuf.getvalue(), "aug.xlsx"), (cbuf.getvalue().encode(), "aug.csv")):
        month = app.parse_month(data, name, "Aug")
        ref = app.parse_month(read(AUG), "aug.xls", "Aug")
        assert len(month["records"]) == 187
        assert month["header_row"] == (4 if name.endswith("xlsx") else 1)
        assert sum(r["gross"] for r in month["records"]) == sum(r["gross"] for r in ref["records"])
        assert sum(r["net"] for r in month["records"]) == sum(r["net"] for r in ref["records"])
        assert month["earn_cols"] == ref["earn_cols"] and month["ded_cols"] == ref["ded_cols"]
        assert [r["id"] for r in month["records"]] == [r["id"] for r in ref["records"]]


@real_files
@pytest.mark.skipif(not SOFFICE, reason="LibreOffice not installed")
def test_workbook_formulas_recalculate_and_agree():
    res = run_real()
    wb = recalc(res["workbook"])
    assert wb.sheetnames == ["Dashboard", "Summary", "Gross Reconciliation", "Net Reconciliation", "Earning Components",
                             "Deduction Components", "All Differences", "All Employees", "Joiners", "Leavers",
                             "Data Checks", "Column Mapping", "Source Data"]
    errors, agrees = [], []
    for ws in wb:
        for line in ws.iter_rows():
            for cell in line:
                if isinstance(cell.value, str) and (cell.value.startswith("#") or cell.value.startswith("Err:")):
                    errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    assert not errors, errors[:10]
    for name in ("Summary", "Gross Reconciliation", "Net Reconciliation"):
        ws = wb[name]
        for line in ws.iter_rows(values_only=True):
            if len(line) >= 4 and line[3] in ("Yes", "NO"):
                agrees.append((name, line[0], line[3]))
    assert agrees and all(a[2] == "Yes" for a in agrees), [a for a in agrees if a[2] != "Yes"]
    summary = {line[0]: line[1] for line in wb["Summary"].iter_rows(values_only=True) if line[0]}
    assert summary["OVERALL"] == "FULLY RECONCILED"
    assert summary["Calculated September 2026 gross"] == 48168536
    assert summary["Actual September 2026 net pay"] == 37811425
    assert wb["Data Checks"]["B3"].value == "PASS"
    # every employee-level formula difference is zero
    ae = wb["All Employees"]
    assert ae.max_row == 186
    for line in ae.iter_rows(min_row=2, values_only=True):
        assert abs(line[20]) < 0.005 and abs(line[28]) < 0.005  # U and AC unexplained columns


@real_files
def test_source_files_unchanged():
    before = (os.path.getsize(AUG), read(AUG)[:4096], os.path.getsize(SEP), read(SEP)[:4096])
    run_real()
    assert before == (os.path.getsize(AUG), read(AUG)[:4096], os.path.getsize(SEP), read(SEP)[:4096])


# ===========================================================================
# Synthetic edge cases
# ===========================================================================
def test_zero_variance_is_a_match():
    res = recon(BASE_PREV, BASE_CURR)
    assert res["status"]["overall"] is True
    assert res["fail_count"] == 0


def test_one_rupee_component_variance_is_reported():
    curr = [row("101", 50000, 20000, gross=70001)] + BASE_CURR[1:]  # GROSS 1 more than components
    curr[0][-1] = 70001 - 1800
    res = recon(BASE_PREV, curr)
    fails = [i for i in res["issues"] if i["check"] == "Earnings components vs GROSS"]
    assert len(fails) == 1 and fails[0]["emp"] == "101" and fails[0]["amount"] == Decimal("-1")
    assert res["status"]["data_checks"] is False and res["status"]["overall"] is False


def test_one_paisa_net_variance_is_reported():
    curr = list(BASE_CURR)
    curr[0] = row("101", 50000, 20000, net=Decimal("68199.99"))
    res = recon(BASE_PREV, curr)
    fails = [i for i in res["issues"] if i["check"] == "GROSS - deductions vs NET PAY"]
    assert len(fails) == 1 and fails[0]["amount"] == Decimal("0.01")
    assert res["status"]["overall"] is False


def test_one_rupee_movement_is_an_exception_not_suppressed():
    curr = [row("101", 50001, 20000)] + BASE_CURR[1:]
    res = recon(BASE_PREV, curr)
    ids = [r["id"] for r in res["differences"]]
    assert "101" in ids
    r101 = next(r for r in res["cont_rows"] if r["id"] == "101")
    assert r101["gross_change"] == 1 and "BASIC +1" in r101["what_changed"]


def test_unchanged_employee_is_not_an_exception():
    res = recon(BASE_PREV, BASE_CURR)
    assert "101" not in [r["id"] for r in res["differences"]]


def test_signed_deductions_and_net():
    res = recon(BASE_PREV, BASE_CURR)
    assert res["prev"]["convention"] == "NEGATIVE"
    r102 = res["prev"]["by_id"]["102"][0]
    assert r102["ded_total"] == 3800 and r102["net"] == 56000 - 3800
    check = next(c for c in res["checks"] if c["id"] == "102" and c["month"] == "Prev")
    assert check["ded_var"] == 0 and check["net_var"] == 0


def test_positive_deduction_convention_supported():
    prev = [row("101", 50000, 20000, td=1800)]
    curr = [row("101", 50000, 20000, td=1800)]
    for r in prev + curr:
        r[-1] = 70000 - 1800
    res = recon(prev, curr)
    assert res["prev"]["convention"] == "POSITIVE" and res["status"]["overall"] is True


def test_joiners_leavers_and_bridge():
    res = recon(BASE_PREV, BASE_CURR)
    assert res["joiners"] == ["104"] and res["leavers"] == ["103"]
    g = res["gross"]
    assert g["joiners"] == 35000 and g["leavers"] == -42000 and g["cont"] == 5600
    assert g["prev"] + g["joiners"] + g["leavers"] + g["cont"] == g["curr"] and g["diff"] == 0
    n = res["net"]
    assert n["cont"] == 5600 - 500 and n["diff"] == 0


def test_new_and_missing_components():
    head_new = HEAD[:8] + ["REFERRAL BONUS"] + HEAD[8:]
    prev = [row("101", 50000, 20000, bonus=5000)]
    curr_row = row("101", 50000, 20000)
    curr_row = curr_row[:8] + [3000] + curr_row[8:]
    curr_row[9] = 73000
    curr_row[-1] = 73000 - 1800
    curr_head = [h for h in head_new if h != "BONUS"]
    del curr_row[7]
    res = app.run_reconciliation(csv_bytes(prev), "p.csv", "Prev", csv_bytes([curr_row], head=curr_head), "c.csv",
                                 "Curr")
    table = {t["component"]: t for t in res["earn_table"]}
    assert table["BONUS"]["in_curr"] is False and table["BONUS"]["cont"] == -5000
    assert table["REFERRAL BONUS"]["in_prev"] is False and table["REFERRAL BONUS"]["cont"] == 3000
    assert res["status"]["overall"] is True
    assert {i["column"] for i in res["issues"] if i["check"] == "Component in one month only"} == \
        {"BONUS", "REFERRAL BONUS"}


def test_duplicate_ids_flagged_never_overwritten():
    curr = BASE_CURR + [row("102", 1000, 0, pf=0)]
    res = recon(BASE_PREV, curr)
    assert "102" in res["dup_ids"]
    assert len(res["curr"]["by_id"]["102"]) == 2
    assert any(i["check"] == "Duplicate Employee No" and i["severity"] == "FAIL" for i in res["issues"])
    g = res["gross"]
    assert g["dup_prev"] == -56000 and g["dup_curr"] == 61600 + 1000
    assert g["diff"] == 0 and res["headcount"]["diff"] == 0
    assert res["status"]["overall"] is False  # duplicates are a data-check failure


def test_invalid_amount_reported_not_silent():
    curr = [list(r) for r in BASE_CURR]
    curr[0][6] = "abc"
    res = recon(BASE_PREV, curr, total=False)
    bad = [i for i in res["issues"] if i["check"] == "Invalid amount"]
    assert len(bad) == 1 and bad[0]["emp"] == "101" and bad[0]["column"] == "HRA"
    assert res["status"]["overall"] is False


def test_blank_amount_reported():
    curr = [list(r) for r in BASE_CURR]
    curr[2][7] = ""  # BONUS blank
    res = recon(BASE_PREV, curr)
    assert any(i["check"] == "Blank amount treated as 0" and i["column"] == "BONUS" for i in res["issues"])


def test_leading_zeros_and_numeric_ids():
    prev = [row("00101", 50000, 20000), row(102.0, 40000, 16000)]
    curr = [row("00101", 50000, 20000), row("102", 40000, 16000)]
    res = recon(prev, curr)
    assert res["continuing"] == ["102", "00101"] or set(res["continuing"]) == {"00101", "102"}
    assert not res["joiners"] and not res["leavers"]


def test_header_below_title_rows_and_grand_total_check():
    res = recon(BASE_PREV, BASE_CURR, title_rows=3)
    assert res["prev"]["header_row"] == 4 and res["prev"]["grand_total_row"] == 8
    assert res["status"]["overall"] is True


def test_grand_total_mismatch_detected():
    data = csv_bytes(BASE_PREV).decode().splitlines()
    data[-1] = data[-1].replace("120000", "120005")  # tamper with the grand-total BASIC
    month = app.parse_month("\n".join(data).encode(), "p.csv", "Prev")
    curr = app.parse_month(csv_bytes(BASE_CURR), "c.csv", "Curr")
    res = app.reconcile(month, curr)
    assert any(i["check"] == "Grand-total row" and i["column"] == "BASIC" for i in res["issues"])


def test_missing_required_column_message():
    head = [h if h != "NET PAY" else "SOMETHING" for h in HEAD]
    with pytest.raises(app.ReconError) as err:
        app.parse_month(csv_bytes(BASE_PREV, head=head), "p.csv", "Prev")
    assert "NET PAY" in str(err.value) and "Headings found" in str(err.value)


def test_names_used_when_present_and_never_invented():
    res = recon(BASE_PREV, BASE_CURR)
    assert all(r["name"] == "" for r in res["cont_rows"])
    head = ["Employee No", "Employee Name"] + HEAD[1:]
    prev = [[r[0], f"Person {r[0]}"] + r[1:] for r in BASE_PREV]
    curr = [[r[0], f"Person {r[0]}"] + r[1:] for r in BASE_CURR]
    res = app.run_reconciliation(csv_bytes(prev, head=head), "p.csv", "P", csv_bytes(curr, head=head), "c.csv", "C")
    assert next(r for r in res["cont_rows"] if r["id"] == "101")["name"] == "Person 101"


def test_lop_split_has_no_double_counting():
    prev = [row("101", 30000, 12000)]
    curr = [row("101", 20000, 8000, paid=20)]  # 20 of 30 days paid
    res = recon(prev, curr)
    r = res["cont_rows"][0]
    assert r["paid_impact"] == -14000 and r["structure"] == 0
    assert r["gross_unexplained"] == 0


def test_same_labels_rejected():
    with pytest.raises(app.ReconError):
        app.run_reconciliation(csv_bytes(BASE_PREV), "p.csv", "Aug", csv_bytes(BASE_CURR), "c.csv", "aug")


@pytest.mark.skipif(not SOFFICE, reason="LibreOffice not installed")
def test_synthetic_failure_shows_in_workbook():
    curr = [row("101", 50000, 20000, gross=70001)] + BASE_CURR[1:]
    curr[0][-1] = 70001 - 1800
    wb = recalc(recon(BASE_PREV, curr)["workbook"])
    summary = {line[0]: line[1] for line in wb["Summary"].iter_rows(values_only=True) if line[0]}
    assert summary["OVERALL"].startswith("NOT RECONCILED")
    assert wb["Data Checks"]["B3"].value == "FAIL"


# ===========================================================================
# Version 2.1: Indian number format, management dashboard, register library
# ===========================================================================
@real_files
@pytest.mark.skipif(not SOFFICE, reason="LibreOffice not installed")
def test_dashboard_values_and_assertions():
    res = app.run_reconciliation(read(AUG), "aug.xls", "August 2026", read(SEP), "sep.xls", "September 2026",
                                 client="Test Client")
    wb = recalc(res["workbook"])
    rows = {line[0]: line[1:] for line in wb["Dashboard"].iter_rows(values_only=True, max_col=5) if line[0]}
    assert rows["OVERALL RESULT"][0] == "FULLY RECONCILED"
    assertions = [k for k, v in rows.items() if v[0] in ("✔ Yes", "✘ No", "Not checked")]
    assert len(assertions) == 7 and all(rows[k][0] == "✔ Yes" for k in assertions)
    assert rows["Employees paid"][:3] == (187, 194, 7)
    assert rows["Gross payroll"][:2] == (50942601, 48168536)
    assert rows["Total deductions"][:2] == (11236493, 10357111)
    assert rows["Net pay (take-home)"][:2] == (39706108, 37811425)
    assert rows["Income tax (TDS) deducted"][:2] == (8767075, 7866042)
    assert rows["Provident fund deducted (PF, VPF, PF arrears)"][:2] == (1061512 + 5291 + 58000, 1110032 + 3735 + 72998)
    assert rows["Bonus and one-time payments (incl. gratuity, leave encashment)"][:2] == \
        (902624 + 1372160 + 55369 + 1445427, 215603 + 25000 + 380952)
    assert abs(rows["Deductions as % of gross"][1] - 10357111 / 48168536) < 1e-12
    assert abs(rows["Attrition (leavers ÷ opening employees)"][1] - 2 / 187) < 1e-12
    assert abs(rows["Continuing employees whose pay changed"][1] - 133 / 185) < 1e-12
    assert rows["Leavers"][0] == -2804386 and rows["Unexplained"][0] == 0
    assert rows["Closing gross (September 2026)"][0] == 48168536
    assert "Test Client" in wb["Dashboard"]["A2"].value


def test_indian_number_format_applied():
    res = recon(BASE_PREV, BASE_CURR)
    wb = openpyxl.load_workbook(io.BytesIO(res["workbook"]))
    gr = wb["Gross Reconciliation"]
    assert gr["B6"].number_format == app.AMT
    assert "##\\,##\\,##0.00" in app.AMT and app.AMT.startswith("[>=10000000]")
    codes = {rule.dxf.numFmt.formatCode for cf in gr.conditional_formatting for rule in cf.rules
             if rule.dxf is not None and rule.dxf.numFmt is not None}
    assert '##\\,##\\,##0.00;[Red]-##\\,##\\,##0.00' in codes


def test_fmt_inr_and_inr_text():
    assert app.fmt_inr(Decimal("14454270")) == "+1,44,54,270"
    assert app.inr_text(Decimal("-2804386")) == "-28,04,386"
    assert app.inr_text(Decimal("123456789.5"), True) == "12,34,56,789.50"
    assert app.inr_text(Decimal("999")) == "999"


def test_period_key():
    assert app.period_key("August 2026") == "2026-08"
    assert app.period_key("Sep-26") == "2026-09"
    assert app.period_key("Previous month") == ""


def test_library_save_dedupe_read_delete(tmp_path):
    lib = str(tmp_path / "lib")
    data = csv_bytes(BASE_PREV)
    entry, new = app.library_save(data, "prev.csv", "ACME Ltd", "August 2026", lib_dir=lib)
    assert new and entry["employees"] == 3 and entry["gross"] == "168000" and entry["period"] == "2026-08"
    again, new2 = app.library_save(data, "copy of prev.csv", "acme ltd", "Aug 2026", lib_dir=lib)
    assert not new2 and again["id"] == entry["id"]
    other, new3 = app.library_save(data, "prev.csv", "Other Co", "August 2026", lib_dir=lib)
    assert new3 and other["id"] != entry["id"]
    curr_entry, _ = app.library_save(csv_bytes(BASE_CURR), "curr.csv", "ACME Ltd", "September 2026", lib_dir=lib)
    index = app.library_index(lib)
    assert len(index) == 3
    assert app.library_read(entry, lib) == data
    res = app.run_reconciliation(app.library_read(entry, lib), entry["filename"], entry["label"],
                                 app.library_read(curr_entry, lib), curr_entry["filename"], curr_entry["label"])
    assert res["status"]["overall"] is True
    assert app.library_delete(other["id"], lib) is True
    assert len(app.library_index(lib)) == 2
    assert not os.path.exists(os.path.join(lib, other["stored_as"]))
    assert app.library_delete("nope", lib) is False


def test_library_detects_changed_file(tmp_path):
    lib = str(tmp_path / "lib")
    entry, _ = app.library_save(csv_bytes(BASE_PREV), "prev.csv", "ACME", "August 2026", lib_dir=lib)
    with open(os.path.join(lib, entry["stored_as"]), "ab") as handle:
        handle.write(b"tampered")
    with pytest.raises(app.ReconError):
        app.library_read(entry, lib)
