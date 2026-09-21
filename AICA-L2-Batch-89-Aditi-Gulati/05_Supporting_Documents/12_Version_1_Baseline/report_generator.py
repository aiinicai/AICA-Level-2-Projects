"""SurakshaScan v1 report generator (reconstructed from bytecode)."""
from docx import Document
from docx.shared import Pt
from datetime import date


def generate_report(institution_name, scan_data, ai_data, output_path):
    doc = Document()
    doc.add_heading("DPDP Readiness Report", level=0)
    p = doc.add_paragraph(institution_name)
    p.runs[0].font.size = Pt(16)
    p.runs[0].font.bold = True
    doc.add_paragraph("Report generated on: " + date.today().strftime("%d %B %Y"))
    doc.add_paragraph(
        "This report reflects a point-in-time readiness assessment. It does not "
        "certify compliance with the Digital Personal Data Protection Act, 2023 "
        "or the DPDP Rules, 2025."
    )

    doc.add_heading("1. Website Scan Findings", level=1)
    for page in scan_data["pages"]:
        doc.add_heading(page["url"], level=2)
        doc.add_paragraph("Status: " + str(page["status_code"]))
        doc.add_paragraph("Privacy/policy links found: "
                          + (", ".join(page["policy_links"]) or "None"))
        doc.add_paragraph("Cookie/consent wording detected: "
                          + ("Yes" if page["cookie_wording_found"] else "No"))
        doc.add_paragraph("Forms found: " + str(len(page["forms"])))
        for i, form in enumerate(page["forms"], start=1):
            names = ", ".join(f["name"] for f in form["fields"])
            doc.add_paragraph("   Form " + str(i) + " fields: " + names,
                              style="List Bullet")
        doc.add_paragraph("Trackers found: "
                          + (", ".join(page["trackers"]) or "None detected"))

    if scan_data["policy_paths_found"]:
        doc.add_paragraph("Common policy URL patterns that exist: "
                          + ", ".join(scan_data["policy_paths_found"]))
    else:
        doc.add_paragraph("No common privacy-policy URL patterns were found to exist.")

    doc.add_heading("2. AI Policy Content Analysis", level=1)
    if ai_data.get("source") == "simulated":
        doc.add_paragraph("Note: AI analysis is not yet connected. "
                          "The items below are placeholders.")
    for item, status in ai_data["checks"].items():
        doc.add_paragraph(item + ": " + status)

    doc.add_heading("3. Scope and Limitations", level=1)
    doc.add_paragraph(
        "This assessment covers publicly accessible web pages only. It does not "
        "access logged-in areas, internal systems, or documents not published on "
        "the website. Absence of a finding is not evidence of absence. This report "
        "should be read together with the manual questionnaire covering internal "
        "practices."
    )
    doc.save(output_path)
    print("Report saved to: " + output_path)
