"""Local Word and PDF exports, with no Google/Office/cloud dependency."""
from io import BytesIO
from pathlib import Path
import datetime as dt
import html

FIRM='ABC Chartered Accountants'
COLUMNS=['Name of Client','Types of Work','Time From','Time To','Hours worked (HH:MM)','Billable or Non-billable']
NOTE='Hours are HH:MM, not decimal hours. Earlier duration-only entries retain their recorded hours. Blank From/To fields and Not recorded billing indicate information not captured originally. No start/end times have been inferred.'

def hours(minutes): return f'{minutes//60}:{minutes%60:02d}'
def summary(r): return [
 ['Total hours worked',hours(r['minutes'])],['Billable hours',hours(r['billable'])],
 ['Non-billable hours',hours(r['nonBillable'])],['Billing not recorded - earlier entries',hours(r['unclassified'])],
 ['Idle time - included in total hours',hours(r['idle'])]]

def export_report(r,format): return pdf_report(r) if format=='pdf' else word_report(r)

def word_report(r):
 from docx import Document
 from docx.shared import Inches, Pt, RGBColor
 from docx.enum.section import WD_ORIENT
 from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 doc=Document(); doc.core_properties.author=FIRM; doc.core_properties.title=r['heading'];
 for border in list(doc.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
 section=doc.sections[0]; section.orientation=WD_ORIENT.LANDSCAPE
 section.page_width=Inches(11.69); section.page_height=Inches(8.27)
 section.left_margin=section.right_margin=Inches(.45); section.top_margin=section.bottom_margin=Inches(.45)
 normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(9)
 normal.paragraph_format.space_after=Pt(7)
 for key in ['Title','Heading 1','Heading 2']:
  style=doc.styles[key]; style.font.name='Arial'; style.font.color.rgb=RGBColor.from_string('172C49')
  style.paragraph_format.space_before=Pt(8); style.paragraph_format.space_after=Pt(7)
 doc.styles['Title'].font.size=Pt(17); doc.styles['Heading 2'].font.size=Pt(11)
 p=doc.add_paragraph(); run=p.add_run(FIRM.upper()); run.bold=True; run.font.size=Pt(12); run.font.color.rgb=RGBColor.from_string('145B73')
 doc.add_paragraph(r['heading'],style='Title')
 def table(rows,widths,header=False,total=False):
  t=doc.add_table(rows=0,cols=len(widths)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
  for column,width in zip(t.columns,widths): column.width=Inches(width)
  for i,values in enumerate(rows):
   row=t.add_row(); is_header=header and i==0; is_total=total and i==len(rows)-1
   if is_header:
    repeat=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
   cant=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
   for j,(cell,value) in enumerate(zip(row.cells,values)):
    cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    props=cell._tc.get_or_add_tcPr(); shading=OxmlElement('w:shd')
    shading.set(qn('w:fill'),'172C49' if is_header else 'E7F1F7' if is_total else 'F5F8FC' if i%2 else 'FFFFFF'); props.append(shading)
    borders=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    props.append(borders)
    margins=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     el=OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margins.append(el)
    props.append(margins)
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
    if header and j in (2,3,4):
     from docx.enum.text import WD_ALIGN_PARAGRAPH
     p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run(str(value)); run.bold=bool(is_header or is_total); run.font.size=Pt(8.5)
    run.font.color.rgb=RGBColor.from_string('FFFFFF' if is_header else '243449')
  doc.add_paragraph().paragraph_format.space_after=Pt(2)
  return t
 table([['Name of '+r['type'],r['name']],['Time Period',r['from']+' to '+r['to']]],[2.0,8.79])
 if r.get('clientFilter'): doc.add_paragraph('Client filter: '+r['clientFilter'])
 if r.get('assignmentFilter'): doc.add_paragraph('Assignment filter: '+r['assignmentFilter'])
 doc.add_paragraph('Generated: '+dt.datetime.now(dt.timezone(dt.timedelta(hours=5,minutes=30))).strftime('%d %b %Y, %H:%M'))
 for g in r['groups']:
  title='Date: '+g['date']+('   |   Staff: '+g['staff'] if r['type']=='Client' else '')
  doc.add_paragraph(title,style='Heading 2')
  table([COLUMNS]+g['rows']+[['TOTAL','','','',hours(g['minutes']),'']],[2.3,3.05,.85,.85,1.3,2.44],True,True)
 if not r['groups']: doc.add_paragraph('No timesheet entries found for the selected staff/client and period.')
 doc.add_paragraph('PERIOD TOTAL',style='Heading 2'); totals=table(summary(r),[5.4,5.39]);
 for row in totals.rows[:-1]:
  for cell in row.cells:
   for paragraph in cell.paragraphs: paragraph.paragraph_format.keep_with_next=True
 doc.add_paragraph(NOTE)
 footer=section.footer.paragraphs[0]
 footer.add_run(FIRM+' | Time Sheet Report | '+r['from']+' to '+r['to']+' | Page ').font.size=Pt(8)
 field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
 output=BytesIO(); doc.save(output); return output.getvalue()

def pdf_report(r):
 from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,KeepTogether
 from reportlab.lib import colors
 from reportlab.lib.styles import ParagraphStyle
 from reportlab.lib.pagesizes import A4,landscape
 from reportlab.pdfbase import pdfmetrics
 from reportlab.pdfbase.ttfonts import TTFont
 import reportlab
 fonts=Path(reportlab.__file__).parent/'fonts'
 if 'ABCRegular' not in pdfmetrics.getRegisteredFontNames():
  pdfmetrics.registerFont(TTFont('ABCRegular',str(fonts/'Vera.ttf')))
  pdfmetrics.registerFont(TTFont('ABCBold',str(fonts/'VeraBd.ttf')))
  pdfmetrics.registerFontFamily('ABCRegular',normal='ABCRegular',bold='ABCBold',italic='ABCRegular',boldItalic='ABCBold')
 regular=ParagraphStyle('regular',fontName='ABCRegular',fontSize=8.5,leading=12,textColor=colors.HexColor('#243449'),spaceAfter=6)
 title=ParagraphStyle('title',parent=regular,fontName='ABCBold',fontSize=15,leading=20,textColor=colors.HexColor('#172C49'),spaceAfter=12)
 section=ParagraphStyle('section',parent=regular,fontName='ABCBold',fontSize=10,leading=14,textColor=colors.HexColor('#145B73'),spaceBefore=10,spaceAfter=8,keepWithNext=True)
 header_style=ParagraphStyle('header',parent=regular,fontName='ABCBold',fontSize=8,leading=11,textColor=colors.white,spaceAfter=0)
 cell_style=ParagraphStyle('cell',parent=regular,fontSize=8,leading=11,spaceAfter=0)
 small=ParagraphStyle('small',parent=regular,fontSize=7,leading=10,textColor=colors.HexColor('#596B80'))
 output=BytesIO(); page_width,page_height=landscape(A4); width=page_width-64
 doc=SimpleDocTemplate(output,pagesize=landscape(A4),leftMargin=32,rightMargin=32,topMargin=30,bottomMargin=38,title=r['heading'],author=FIRM)
 def p(text,style=regular): return Paragraph(html.escape(str(text)).replace('\n','<br/>'),style)
 def table(rows,widths,header=False,total=False):
  matrix=[[p(value,header_style if header and i==0 else cell_style) for value in row] for i,row in enumerate(rows)]
  t=Table(matrix,colWidths=[width*x/sum(widths) for x in widths],repeatRows=1 if header else 0,hAlign='LEFT',splitByRow=1,splitInRow=1)
  commands=[('GRID',(0,0),(-1,-1),.5,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]
  for i in range(len(rows)):
   background='#172C49' if header and i==0 else '#E7F1F7' if total and i==len(rows)-1 else '#F5F8FC' if i%2 else '#FFFFFF'
   commands.append(('BACKGROUND',(0,i),(-1,i),colors.HexColor(background)))
  t.setStyle(TableStyle(commands)); return t
 story=[p(FIRM.upper(),section),p(r['heading'],title),table([['Name of '+r['type'],r['name']],['Time Period',r['from']+' to '+r['to']]],[150,628]),Spacer(1,8)]
 if r.get('clientFilter'): story.append(p('Client filter: '+r['clientFilter']))
 if r.get('assignmentFilter'): story.append(p('Assignment filter: '+r['assignmentFilter']))
 story.append(p('Generated: '+dt.datetime.now(dt.timezone(dt.timedelta(hours=5,minutes=30))).strftime('%d %b %Y, %H:%M'),small))
 for g in r['groups']:
  text='Date: '+g['date']+('   |   Staff: '+g['staff'] if r['type']=='Client' else '')
  story.extend([p(text,section),table([COLUMNS]+g['rows']+[['TOTAL','','','',hours(g['minutes']),'']],[180,228,65,65,100,140],True,True),Spacer(1,8)])
 if not r['groups']: story.append(p('No timesheet entries found for the selected staff/client and period.'))
 story.extend([p('PERIOD TOTAL',section),table(summary(r),[390,388]),Spacer(1,10),p(NOTE,small)])
 def footer(canvas,document):
  canvas.setFont('ABCRegular',7); canvas.setFillColor(colors.HexColor('#596B80'))
  canvas.drawString(32,20,FIRM+' | Time Sheet Report | '+r['from']+' to '+r['to'])
  canvas.drawRightString(page_width-32,20,'Page '+str(document.page))
 doc.build(story,onFirstPage=footer,onLaterPages=footer); return output.getvalue()
