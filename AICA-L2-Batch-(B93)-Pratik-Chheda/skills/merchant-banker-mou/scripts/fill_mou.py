#!/usr/bin/env python3
"""Fill the approved Sample Advisory MOU with BRLM template.

The template's wording, numbering, tables and formatting are never rewritten:
each placeholder ([●], [1956/2013], the date blanks, [ADVISOR NAME], the
invoice window) is replaced in place inside its own paragraph, even when Word
has split the brackets across several runs.

Usage:
    python fill_mou.py content.json OUTPUT.docx [--template path.docx]

The script exits with code 2 and lists what is left if any placeholder is
still unfilled, so a half filled MOU is never produced silently.
"""
import argparse
import datetime as dt
import json
import os
import sys

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(HERE, "..", "assets", "MOU_with_BRLM_template.docx")
BLANK = "[●]"  # [●]

ADVISOR = {
    "name": "Sample Advisory Private Limited",
    "companies_act": "2013",
    "cin": "U00000MH2020PTC000000",
    "registered_office": "Registered office address, City 400001",
    "signatory_name": "Mr. A. Partner",
    "signatory_designation": "Director",
}

# ---------------------------------------------------------------- words
ONES = ["", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten",
        "eleven", "twelve", "thirteen", "fourteen", "fifteen", "sixteen", "seventeen", "eighteen", "nineteen"]
TENS = ["", "", "twenty", "thirty", "forty", "fifty", "sixty", "seventy", "eighty", "ninety"]


def _below_hundred(n):
    return ONES[n] if n < 20 else (TENS[n // 10] + (" " + ONES[n % 10] if n % 10 else ""))


def _below_thousand(n):
    h, r = divmod(n, 100)
    parts = []
    if h:
        parts.append(ONES[h] + " hundred")
    if r:
        parts.append(_below_hundred(r))
    return " ".join(parts)


def indian_words(n):
    """Integer to words in the Indian system (lakh, crore), lower case."""
    n = int(round(n))
    if n == 0:
        return "zero"
    parts = []
    for size, label in ((10 ** 7, "crore"), (10 ** 5, "lakh"), (1000, "thousand")):
        q, n = divmod(n, size)
        if q:
            parts.append((indian_words(q) if q >= 100 else _below_hundred(q)) + " " + label)
    if n:
        parts.append(_below_thousand(n))
    return " ".join(parts)


def rupees_words(n):
    return indian_words(n).title().replace(" And ", " and ")


def inr(n):
    """1500000 -> 15,00,000"""
    s = str(int(round(n)))
    if len(s) <= 3:
        return s
    head, tail = s[:-3], s[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return ",".join(groups) + "," + tail


def period(n, unit):
    n = int(n)
    unit = unit if n == 1 else unit + "s"
    return f"{n} ({indian_words(n)}) {unit}"


def pct(v):
    v = float(v)
    return (str(int(v)) if v.is_integer() else f"{v:g}") + "%"


def ordinal(d):
    return str(d) + ("th" if 11 <= d % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(d % 10, "th"))


# ---------------------------------------------------------------- run safe replacement
def replace_once(paragraph, old, new):
    runs = paragraph.runs
    texts = [r.text for r in runs]
    full = "".join(texts)
    i = full.find(old)
    if i < 0:
        return False
    j, pos, first = i + len(old), 0, True
    for run, t in zip(runs, texts):
        s, e = pos, pos + len(t)
        pos = e
        if e <= i or s >= j:
            continue
        a, b = max(i, s) - s, min(j, e) - s
        if first:
            run.text = t[:a] + new + t[b:]
            first = False
        else:
            run.text = t[:a] + t[b:]
    return True


def fill(paragraph, token, values):
    for v in values:
        if not replace_once(paragraph, token, v):
            raise ValueError(f"Expected another '{token}' in: {paragraph.text[:90]}")


def text_of(p):
    return "".join(r.text for r in p.runs)


# ---------------------------------------------------------------- main fill
def build(content, template):
    doc = Document(template)
    adv = dict(ADVISOR, **(content.get("advisor") or {}))
    brlm = content["brlm"]
    issuer = content["issuer"]
    when = dt.date.fromisoformat(content.get("effective_date") or dt.date.today().isoformat())
    month = when.strftime("%B")
    retainer = content["retainer_inr"]
    success_pct = content["success_fee_pct"]
    success_inr = content.get("success_fee_inr")
    miles = content.get("milestones") or []
    inv_days = int(content.get("invoice_days", 15))

    for p in doc.paragraphs:
        t = text_of(p)
        if "is made and executed on this" in t:
            fill(p, "___ day of __________, 2026", [f"{ordinal(when.day)} day of {month}, {when.year}"])
        elif "hereinafter referred to as the “Advisor”" in t:
            fill(p, BLANK, [adv["name"]])
            fill(p, "[1956/2013]", [str(adv["companies_act"])])
            fill(p, BLANK, [adv["cin"], adv["registered_office"]])
        elif "Book Running Lead Manager" in t and "SEBI Registration No." in t:
            fill(p, BLANK, [brlm["name"], brlm["sebi_registration"]])
            fill(p, "[1956/2013]", [str(brlm.get("companies_act", "2013"))])
            fill(p, BLANK, [brlm["cin"], brlm["registered_office"]])
        elif "proposed Initial Public Offering" in t:
            fill(p, BLANK, [issuer])
        elif "Confidentiality obligations shall survive" in t:
            fill(p, BLANK, [period(content.get("confidentiality_years", 2), "year")])
        elif "solicit, entice, recruit" in t:
            fill(p, BLANK, [period(content.get("non_poaching_months", 12), "month")])
        elif "shall automatically terminate upon the earliest of" in t:
            fill(p, BLANK, [period(content.get("term_months", 24), "month")])
        elif "terminate this MOU without cause by giving" in t:
            fill(p, BLANK, [period(content.get("notice_days", 30), "day") + "’",
                            period(content.get("cure_days", 15), "day")])
        elif "seat and venue of arbitration" in t:
            city = content.get("arbitration_city", "Mumbai")
            fill(p, BLANK, [city, content.get("courts_city", city)])
        elif "forms an integral and binding part" in t:
            fill(p, "___ __________, 2026", [f"{ordinal(when.day)} {month}, {when.year}"])
        elif t.strip().startswith("Fixed Retainership Fee"):
            fill(p, BLANK, [inr(retainer), rupees_words(retainer)])
        elif t.strip().startswith("Success / Milestone Fee"):
            if success_inr:
                fill(p, BLANK, [pct(success_pct).rstrip("%"), inr(success_inr)])
            else:
                fill(p, f"{BLANK}% (or INR {BLANK})", [pct(success_pct)])
        elif "Invoices shall be payable within" in t:
            fill(p, "[15 (fifteen) business days]", [period(inv_days, "business day")])

    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            first = cells[0].text.strip()
            words = first.split()
            if len(words) > 1 and words[0] == "Milestone" and words[1].rstrip(":").isdigit():
                k = int(words[1].rstrip(":")) - 1
                m = miles[k] if k < len(miles) else {}
                parts = []
                if m.get("pct") not in (None, ""):
                    parts.append(pct(m["pct"]))
                if m.get("inr") not in (None, ""):
                    parts.append("INR " + inr(m["inr"]))
                cell_fill(cells[1], f"{BLANK}% / INR {BLANK}", [" / ".join(parts) or "Nil"])
                continue
            for col, cell in enumerate(cells[:2]):
                party = adv if col == 0 else brlm
                ct = cell.text
                if "[ADVISOR NAME]" in ct:
                    cell_fill(cell, "[ADVISOR NAME]", [adv["name"]])
                elif ct.strip() == "For and on behalf of":
                    p = cell.paragraphs[0]
                    run = p.runs[-1] if p.runs else p.add_run("")
                    run.text = run.text.rstrip() + " " + party["name"]
                elif ct.startswith("For and on behalf of") and BLANK in ct:
                    cell_fill(cell, BLANK, [party["name"]])
                elif "Authorized Signatory" in ct and BLANK in ct:
                    cell_fill(cell, BLANK, [party["signatory_name"], party["signatory_designation"]])
    return doc


def cell_fill(cell, token, values):
    values = list(values)
    for p in cell.paragraphs:
        while values and token in text_of(p):
            replace_once(p, token, values.pop(0))
    if values:
        raise ValueError(f"Expected another '{token}' in table cell: {cell.text[:90]}")


def leftovers(doc):
    marks = [BLANK, "[1956/2013]", "[ADVISOR NAME]", "__________, 20", "[15 (fifteen)"]
    found = []
    def scan(p, where):
        t = text_of(p)
        for m in marks:
            if m in t:
                found.append(f"{where}: {t.strip()[:100]}")
    for p in doc.paragraphs:
        scan(p, "text")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    scan(p, "table")
    return found


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("content")
    ap.add_argument("output")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    a = ap.parse_args()
    with open(a.content, encoding="utf-8") as f:
        content = json.load(f)
    miles = content.get("milestones") or []
    shares = [float(m["pct"]) for m in miles if m.get("pct") not in (None, "")]
    if shares and len(shares) == len(miles) and round(sum(shares), 2) != 100:
        sys.exit(f"Milestone percentages add up to {sum(shares):g}%, not 100%")
    doc = build(content, a.template)
    left = leftovers(doc)
    if left:
        print("Unfilled placeholders remain:\n  " + "\n  ".join(left))
        sys.exit(2)
    doc.save(a.output)
    print("Saved " + a.output)


if __name__ == "__main__":
    main()
