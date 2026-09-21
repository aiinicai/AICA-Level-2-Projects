from io import BytesIO
from types import SimpleNamespace as NS
from unittest.mock import Mock
import pytest
from docx import Document
from openpyxl import Workbook
from services.case_models import Authority, Research, EvidenceReview, DocumentFinding, ReviewResult
from services.models import Draft
from services.case_service import CaseService
from services.research_service import ResearchService, official_url
from services.evidence_service import extract_supporting_document
from services.demo_service import sample_analysis
from services.demo_case import demo_strategy
from utils.validators import AppError


def authority(url="https://www.indiacode.nic.in/example", kind="Act", position="Mixed / neutral"):
    return Authority(title="Test authority", kind=kind, reference="Test reference", court_or_issuer="Test issuer", date_or_effective_period="Unverified", issue="Test issue", position=position, proposition="Test proposition", relevance_and_distinctions="Verify factual fit", pinpoint="Not established", source_url=url, status_checks="Not checked")


def provider(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test-only")
    monkeypatch.setenv("OPENAI_MODEL", "test-model")
    client = Mock()
    monkeypatch.setattr("services.llm_service.OpenAI", Mock(return_value=client))
    return client


def test_official_sources_are_strict():
    assert official_url("https://www.sci.gov.in/judgment.pdf")
    assert not official_url("https://sci.gov.in.evil.example/judgment")
    assert not official_url("javascript:alert(1)")
    assert not official_url("https://user:password@sci.gov.in/a")


def test_live_search_required_and_unreturned_sources_removed(monkeypatch):
    client = provider(monkeypatch)
    trusted = "https://www.indiacode.nic.in/example"
    client.responses.parse.return_value = NS(
        output=[NS(type="web_search_call", status="completed", action=NS(sources=[{"url": trusted}]))],
        output_parsed=Research(authorities=[authority(trusted), authority("https://www.sci.gov.in/not-retrieved")], gaps=[], search_summary="Test report"),
    )
    result, timestamp = ResearchService().research("Abstract public tax law question")
    assert len(result.authorities) == 1
    assert len(result.gaps) == 3  # excluded URL, missing favorable and adverse cases
    args = client.responses.parse.call_args.kwargs
    assert args["tool_choice"] == "required" and args["store"] is False
    assert args["input"] == "Abstract public tax law question"
    assert "case" not in args and "notice" not in args
    assert timestamp.endswith("+00:00")


def test_no_tool_call_cannot_masquerade_as_research(monkeypatch):
    client = provider(monkeypatch)
    client.responses.parse.return_value = NS(output=[], output_parsed=Research(authorities=[], gaps=[], search_summary="Unfounded"))
    with pytest.raises(AppError, match="No completed web search"):
        ResearchService().research("Research question")


def test_evidence_findings_need_source_quotes(monkeypatch):
    provider(monkeypatch)
    monkeypatch.setattr(CaseService, "request_case", lambda *_: EvidenceReview(findings=[DocumentFinding(observation="False observation", source_quote="Not in document", relevance="Test", effect="Supports")], missing_information=[]))
    with pytest.raises(AppError, match="traced"):
        CaseService().review_document("Actual bank statement content", sample_analysis())


def test_final_reply_uses_reviewed_version(monkeypatch):
    provider(monkeypatch)
    replies = [Draft(draft_reply="Initial incomplete draft"), ReviewResult(revised_reply="Corrected draft with placeholders", review_notes=["Confirm evidence before filing"])]
    mock = Mock(side_effect=replies)
    monkeypatch.setattr(CaseService, "request_case", mock)
    draft, notes = CaseService().comprehensive_draft({"notice": "Fictional"}, demo_strategy(sample_analysis(), {}), "Detailed")
    assert draft == "Corrected draft with placeholders"
    assert notes == ["Confirm evidence before filing"]
    assert mock.call_count == 2


def test_strategy_cannot_omit_notice_issue(monkeypatch):
    provider(monkeypatch)
    analysis = sample_analysis()
    incomplete = demo_strategy(analysis, {})
    incomplete.issues = incomplete.issues[:1]
    monkeypatch.setattr(CaseService, "request_case", lambda *_: incomplete)
    with pytest.raises(AppError, match="every notice issue"):
        CaseService().prepare_strategy({"notice": analysis.model_dump(mode="json")})


def test_strategy_cannot_invent_research_authority(monkeypatch):
    provider(monkeypatch)
    analysis = sample_analysis()
    strategy = demo_strategy(analysis, {})
    strategy.issues[0].authorities_to_use = ["Invented judgment"]
    monkeypatch.setattr(CaseService, "request_case", lambda *_: strategy)
    with pytest.raises(AppError, match="outside the research"):
        CaseService().prepare_strategy({"notice": analysis.model_dump(mode="json"), "research": None})


def test_final_review_rejects_new_unresearched_link(monkeypatch):
    provider(monkeypatch)
    mock = Mock(side_effect=[Draft(draft_reply="First draft"), ReviewResult(revised_reply="See https://www.sci.gov.in/invented.pdf", review_notes=[])])
    monkeypatch.setattr(CaseService, "request_case", mock)
    with pytest.raises(AppError, match="outside the research record"):
        CaseService().comprehensive_draft({}, demo_strategy(sample_analysis(), {}), "Detailed")


def test_supporting_docx_includes_tables():
    doc = Document()
    doc.add_paragraph("Interest certificate")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Interest: 100"
    buffer = BytesIO()
    doc.save(buffer)
    text, warnings = extract_supporting_document(buffer.getvalue(), "certificate.docx")
    assert "Interest certificate" in text and "Interest: 100" in text
    assert warnings


def test_supporting_excel_reads_all_sheets_and_does_not_invent_formula_values():
    workbook = Workbook()
    workbook.active.append(["Interest", 100, "=B1*2"])
    workbook.create_sheet("Other").append(["Difference", 50])
    buffer = BytesIO()
    workbook.save(buffer)
    text, warnings = extract_supporting_document(buffer.getvalue(), "ledger.xlsx")
    assert "=B1*2" in text and "Sheet: Other" in text and "50" in text
    assert "not calculated" in warnings[0]


def test_supporting_csv_and_invalid_type():
    text, _ = extract_supporting_document(b"Item,Amount\nInterest,100", "ledger.csv")
    assert "Row 2: Interest | 100" in text
    with pytest.raises(AppError):
        extract_supporting_document(b"any", "script.exe")
