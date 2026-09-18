import io
import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from typing import List, Dict, Any, Tuple
import database as db
import analytics

EXCEL_COLUMNS_SPEC = [
    ("Borrower Name", "borrower_name", True),
    ("Bank Name", "bank_name", True),
    ("Allotment Letter Date", "allotment_date", False),
    ("Acceptance Letter Date", "acceptance_date", False),
    ("Data Requirement Sent Date (Bank)", "data_req_bank_date", False),
    ("Data Requirement Sent Date (Borrower)", "data_req_borrower_date", False),
    ("Pending Details / Documents", "pending_details", False),
    ("Team Person Name", "team_person_name", False),
    ("Team Person Number", "team_person_number", False),
    ("Target Visit Date", "target_visit_date", False),
    ("Actual Visit Date", "actual_visit_date", False),
    ("Audit Status", "audit_status", False),
    ("Review with GC / Partner Date", "review_partner_date", False),
    ("Report Target Date", "report_target_date", False),
    ("Report Submission Date", "report_submission_date", False),
    ("Remarks / Delay Reason", "remarks", False)
]

def format_date_for_db(val: Any) -> str:
    if pd.isna(val) or val is None or str(val).strip() == "" or str(val).lower() == "nat":
        return ""
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.strftime("%Y-%m-%d")
    val_str = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%b-%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.datetime.strptime(val_str[:10] if len(val_str)>=10 else val_str, fmt).strftime("%Y-%m-%d")
        except Exception:
            continue
    return val_str[:10]

def parse_and_import_excel(file_bytes: bytes, user_name: str) -> Tuple[int, int, List[str]]:
    errors = []
    imported = 0
    skipped = 0
    
    try:
        df = pd.read_excel(io.BytesIO(file_bytes))
    except Exception as e:
        return 0, 0, [f"Failed to read Excel file: {str(e)}"]
        
    if df.empty:
        return 0, 0, ["Excel sheet is empty."]
        
    col_map = {}
    for col in df.columns:
        norm_col = str(col).strip().lower().replace(" ", "").replace("_", "").replace("-", "").replace("/", "")
        for title, key, _ in EXCEL_COLUMNS_SPEC:
            norm_target = title.strip().lower().replace(" ", "").replace("_", "").replace("-", "").replace("/", "")
            if norm_col == norm_target or norm_target in norm_col or norm_col in norm_target:
                col_map[key] = col
                break
                
    if "borrower_name" not in col_map:
        return 0, 0, ["Required column 'Borrower Name' not found in uploaded file."]
    if "bank_name" not in col_map:
        return 0, 0, ["Required column 'Bank Name' not found in uploaded file."]
        
    for idx, row in df.iterrows():
        borrower = str(row.get(col_map.get("borrower_name"), "")).strip()
        bank = str(row.get(col_map.get("bank_name"), "")).strip()
        
        if not borrower or borrower.lower() == "nan" or not bank or bank.lower() == "nan":
            skipped += 1
            continue
            
        record = {
            "borrower_name": borrower,
            "bank_name": bank,
            "allotment_date": format_date_for_db(row.get(col_map.get("allotment_date")) if "allotment_date" in col_map else ""),
            "acceptance_date": format_date_for_db(row.get(col_map.get("acceptance_date")) if "acceptance_date" in col_map else ""),
            "data_req_bank_date": format_date_for_db(row.get(col_map.get("data_req_bank_date")) if "data_req_bank_date" in col_map else ""),
            "data_req_borrower_date": format_date_for_db(row.get(col_map.get("data_req_borrower_date")) if "data_req_borrower_date" in col_map else ""),
            "pending_details": "" if pd.isna(row.get(col_map.get("pending_details"))) else str(row.get(col_map.get("pending_details"))),
            "team_person_name": "" if pd.isna(row.get(col_map.get("team_person_name"))) else str(row.get(col_map.get("team_person_name"))).strip(),
            "team_person_number": "" if pd.isna(row.get(col_map.get("team_person_number"))) else str(row.get(col_map.get("team_person_number"))).strip(),
            "target_visit_date": format_date_for_db(row.get(col_map.get("target_visit_date")) if "target_visit_date" in col_map else ""),
            "actual_visit_date": format_date_for_db(row.get(col_map.get("actual_visit_date")) if "actual_visit_date" in col_map else ""),
            "audit_status": str(row.get(col_map.get("audit_status"))).strip() if (("audit_status" in col_map) and not pd.isna(row.get(col_map.get("audit_status")))) else "Not Started",
            "review_partner_date": format_date_for_db(row.get(col_map.get("review_partner_date")) if "review_partner_date" in col_map else ""),
            "report_target_date": format_date_for_db(row.get(col_map.get("report_target_date")) if "report_target_date" in col_map else ""),
            "report_submission_date": format_date_for_db(row.get(col_map.get("report_submission_date")) if "report_submission_date" in col_map else ""),
            "remarks": "" if pd.isna(row.get(col_map.get("remarks"))) else str(row.get(col_map.get("remarks"))).strip()
        }
        
        if record["audit_status"] not in db.STATUS_CHOICES:
            record["audit_status"] = "Not Started"
            
        try:
            db.create_assignment(record, user_name=f"{user_name} (Excel Import)")
            imported += 1
        except Exception as e:
            errors.append(f"Row {idx+2} ({borrower}): {str(e)}")
            skipped += 1
            
    return imported, skipped, errors

def generate_sample_excel_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock Audit Master"
    
    headers = [col[0] for col in EXCEL_COLUMNS_SPEC]
    ws.append(headers)
    
    sample_rows = [
        [
            "Acme Steels Pvt Ltd", "State Bank of India", "2026-08-01", "2026-08-02",
            "2026-08-03", "2026-08-03", "Stock statements for July 2026 pending",
            "Priya Verma", "+91 98111 22334", "2026-08-10", "2026-08-10",
            "Stock Verification Completed", "2026-08-14", "2026-08-18", "", "Partner review scheduled with CA Gaurav Chaudhary"
        ],
        [
            "Bharat Textiles Ltd", "Punjab National Bank", "2026-08-05", "2026-08-06",
            "2026-08-07", "2026-08-07", "Audited balance sheet FY 25-26 awaited",
            "Amit Gupta", "+91 98222 33445", "2026-08-12", "",
            "Data Awaited from Borrower", "", "2026-08-20", "", "Borrower requested 3 days extension"
        ],
        [
            "Delta Pharmaceuticals", "HDFC Bank", "2026-07-20", "2026-07-21",
            "2026-07-22", "2026-07-22", "None",
            "Sneha Patel", "+91 98333 44556", "2026-07-28", "2026-07-28",
            "Report Submitted", "2026-08-04", "2026-08-08", "2026-08-06", "Submitted to Circle Office on time"
        ]
    ]
    
    for r in sample_rows:
        ws.append(r)
        
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    for row in ws.iter_rows(min_row=1, max_row=len(sample_rows)+1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    ws.row_dimensions[1].height = 30
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

def export_audits_to_excel(assignments: List[Dict[str, Any]], export_type: str = "Complete") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{export_type} Audits"
    
    title_cell = ws.cell(row=1, column=1, value=f"GS STOCK AUDIT TRACKER - {export_type.upper()} REPORT")
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    
    gen_cell = ws.cell(row=2, column=1, value=f"Partner: CA Gaurav Chaudhary | Generated: {datetime.datetime.now().strftime('%d-%b-%Y %I:%M %p')} | Total: {len(assignments)}")
    gen_cell.font = Font(name="Calibri", size=10, italic=True, color="64748B")
    
    display_cols = [
        ("ID", "id"),
        ("Borrower Name", "borrower_name"),
        ("Bank Name", "bank_name"),
        ("Allotment Date", "allotment_date"),
        ("Acceptance Date", "acceptance_date"),
        ("Data Req (Bank)", "data_req_bank_date"),
        ("Data Req (Borrower)", "data_req_borrower_date"),
        ("Pending Details", "pending_details"),
        ("Team Person", "team_person_name"),
        ("Team Mobile", "team_person_number"),
        ("Target Visit Date", "target_visit_date"),
        ("Actual Visit Date", "actual_visit_date"),
        ("Audit Status", "audit_status"),
        ("Partner Review Date", "review_partner_date"),
        ("Report Target Date", "report_target_date"),
        ("Submission Date", "report_submission_date"),
        ("Remarks / Delay", "remarks")
    ]
    
    start_row = 4
    for col_idx, (col_title, _) in enumerate(display_cols, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=col_title)
        cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    
    for row_idx, item in enumerate(assignments, start=start_row + 1):
        for col_idx, (_, key) in enumerate(display_cols, start=1):
            val = item.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=str(val if val is not None else ""))
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            
            if key == "audit_status":
                if val in ["Report Submitted", "Closed"]:
                    cell.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
                elif val == "Delayed":
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                elif val in ["Visit Planned", "Stock Verification Completed", "Report Preparation"]:
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                    
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, 40), 12)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

def generate_dashboard_pdf(assignments: List[Dict[str, Any]], partner_name: str = "CA Gaurav Chaudhary") -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        name="RepTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        textColor=colors.HexColor("#1E3A8A"),
        spaceAfter=6
    )
    
    sub_style = ParagraphStyle(
        name="RepSub",
        parent=styles["Normal"],
        fontSize=10,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=14
    )
    
    cell_style = ParagraphStyle(
        name="CellText",
        parent=styles["Normal"],
        fontSize=8,
        leading=10
    )
    
    elements = []
    
    elements.append(Paragraph("GS STOCK AUDIT TRACKER - EXECUTIVE MIS DASHBOARD", title_style))
    elements.append(Paragraph(f"Lead Partner: {partner_name} | Date: {datetime.datetime.now().strftime('%d-%b-%Y %I:%M %p')} | Total Assignments: {len(assignments)}", sub_style))
    
    total = len(assignments)
    completed = sum(1 for a in assignments if a.get("audit_status") in ["Report Submitted", "Closed"])
    delayed = sum(1 for a in assignments if a.get("audit_status") == "Delayed")
    pending = total - completed
    reports_pending = sum(1 for a in assignments if a.get("audit_status") in ["Report Preparation", "Review Pending", "Partner Review Completed"])
    visits_pending = sum(1 for a in assignments if a.get("audit_status") in ["Not Started", "Data Awaited from Bank", "Data Awaited from Borrower", "Documents Under Review", "Visit Planned"])
    
    summary_data = [
        ["Total Assignments", "Completed Audits", "Pending Audits", "Delayed Audits", "Reports Pending", "Visits Pending"],
        [str(total), str(completed), str(pending), str(delayed), str(reports_pending), str(visits_pending)]
    ]
    
    sum_table = Table(summary_data, colWidths=[125, 125, 125, 125, 125, 125])
    sum_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor("#F8FAFC")),
        ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, 1), 12),
        ('TEXTCOLOR', (0, 1), (-1, 1), colors.HexColor("#0F172A")),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
    ]))
    
    elements.append(sum_table)
    elements.append(Spacer(1, 16))
    
    table_headers = ["ID", "Borrower Name", "Bank", "Team Person", "Target Visit", "Actual Visit", "Report Target", "Status"]
    table_rows = [table_headers]
    
    for a in assignments[:25]:
        table_rows.append([
            str(a.get("id")),
            Paragraph(str(a.get("borrower_name", "")), cell_style),
            Paragraph(str(a.get("bank_name", "")), cell_style),
            str(a.get("team_person_name", "")),
            str(a.get("target_visit_date", "") or "-"),
            str(a.get("actual_visit_date", "") or "-"),
            str(a.get("report_target_date", "") or "-"),
            Paragraph(str(a.get("audit_status", "")), cell_style)
        ])
        
    tbl = Table(table_rows, colWidths=[35, 170, 110, 110, 75, 75, 75, 100])
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#334155")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    
    elements.append(tbl)
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()
