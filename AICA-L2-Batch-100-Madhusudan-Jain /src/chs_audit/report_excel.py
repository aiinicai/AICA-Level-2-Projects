"""Write every schedule, the findings and the AI drafts into one audit workbook."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SEV_FILL = {"High": "F8D7DA", "Medium": "FFF3CD", "Low": "D1E7DD", "Info": "E2E3E5"}


def _autofit(ws, max_width=60):
    for col in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, length + 2), max_width)


def _style_header(ws):
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame], texts: dict[str, str], findings_df: pd.DataFrame):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        findings_df.to_excel(xw, sheet_name="Audit_Observations", index=False)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
        for name, text in texts.items():
            pd.DataFrame({name: text.split("\n")}).to_excel(xw, sheet_name=name[:31], index=False)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for ws in wb.worksheets:
        _style_header(ws)
        _autofit(ws)
        if ws.title == "Audit_Observations":
            sev_col = [c.value for c in ws[1]].index("severity") + 1
            for row in ws.iter_rows(min_row=2):
                sev = row[sev_col - 1].value
                if sev in SEV_FILL:
                    row[sev_col - 1].fill = PatternFill("solid", fgColor=SEV_FILL[sev])
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions["D"].width = 90
            ws.column_dimensions["E"].width = 45
            ws.column_dimensions["F"].width = 60
        if ws.title.startswith("Draft_"):
            ws.column_dimensions["A"].width = 140
            for row in ws.iter_rows(min_row=2):
                row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
