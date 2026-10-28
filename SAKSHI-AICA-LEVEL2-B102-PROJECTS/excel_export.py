"""
Export verification results to a formatted Excel workbook (openpyxl).

Udyam numbers and every portal value are written as text exactly as the
portal returned them.  Only a date of registration in the portal's
DD/MM/YYYY form is turned into a real Excel date (shown as DD-MM-YYYY);
anything else stays as text so nothing is guessed.
"""

import datetime
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from udyam_core import (CAT_ERROR, CAT_NOT_VERIFIED, CAT_PENDING, CAT_VERIFIED,
                        COLUMN_TITLES, OUTPUT_COLUMNS, status_category)

CATEGORY_FILL = {
    CAT_VERIFIED: "E6F4EA",
    CAT_NOT_VERIFIED: "FFF4E5",
    CAT_ERROR: "FDECEE",
    CAT_PENDING: "E8EEF8",
}
WIDTHS = {"udyam_number": 24, "enterprise_name": 38, "enterprise_type": 14,
          "organisation_type": 20,
          "major_activity": 16, "social_category": 15,
          "date_of_registration": 16, "enterprise_address": 60,
          "status": 18, "error": 48}
DATE_RE = re.compile(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$")


def _as_date(text):
    m = DATE_RE.match((text or "").strip())
    if not m:
        return None
    d, mth, y = (int(g) for g in m.groups())
    try:
        return datetime.date(y, mth, d)
    except ValueError:
        return None


def export_results(rows, path):
    """Write rows (dicts keyed by OUTPUT_COLUMNS) to an .xlsx file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Verification Results"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1E4DB7")
    thin = Side(style="thin", color="D5DBE5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    headers = ["Category"] + [COLUMN_TITLES[c] for c in OUTPUT_COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
    ws.row_dimensions[1].height = 32

    for r, row in enumerate(rows, start=2):
        cat = status_category(row.get("status"))
        ws.cell(row=r, column=1, value=cat)
        for c, col in enumerate(OUTPUT_COLUMNS, start=2):
            value = row.get(col, "") or ""
            cell = ws.cell(row=r, column=c)
            date = _as_date(value) if col == "date_of_registration" else None
            if date:
                cell.value = date
                cell.number_format = "DD-MM-YYYY"
            else:
                cell.value = value
                cell.number_format = "@"          # keep as text
            cell.alignment = wrap
            cell.border = border
        fill = PatternFill("solid", fgColor=CATEGORY_FILL[cat])
        for c in (1, 1 + OUTPUT_COLUMNS.index("status") + 1):
            ws.cell(row=r, column=c).fill = fill
        ws.cell(row=r, column=1).border = border
        ws.cell(row=r, column=1).alignment = wrap

    ws.column_dimensions["A"].width = 20
    for c, col in enumerate(OUTPUT_COLUMNS, start=2):
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS.get(col, 18)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"

    # Summary sheet
    sm = wb.create_sheet("Summary")
    sm.append(["MSME Verification System - results summary"])
    sm["A1"].font = Font(bold=True, size=14, color="1E4DB7")
    sm.append(["Exported on", datetime.datetime.now().strftime("%d-%m-%Y %I:%M %p")])
    sm.append([])
    sm.append(["Category", "Count"])
    for cell in sm[4]:
        cell.font = head_font
        cell.fill = head_fill
    counts = {k: 0 for k in CATEGORY_FILL}
    for row in rows:
        counts[status_category(row.get("status"))] += 1
    for cat, n in counts.items():
        sm.append([cat, n])
        sm.cell(row=sm.max_row, column=1).fill = PatternFill(
            "solid", fgColor=CATEGORY_FILL[cat])
    sm.append(["Total", len(rows)])
    sm.cell(row=sm.max_row, column=1).font = Font(bold=True)
    sm.cell(row=sm.max_row, column=2).font = Font(bold=True)

    # Verified enterprises by type (Micro / Small / Medium)
    sm.append([])
    sm.append(["Verified - Type of Enterprise", "Count"])
    for cell in sm[sm.max_row]:
        cell.font = head_font
        cell.fill = head_fill
    by_type = {"Micro": 0, "Small": 0, "Medium": 0, "Not shown": 0}
    for row in rows:
        if status_category(row.get("status")) == CAT_VERIFIED:
            kind = row.get("enterprise_type") or "Not shown"
            by_type[kind if kind in by_type else "Not shown"] += 1
    for kind, n in by_type.items():
        sm.append([kind, n])
    sm.append([])
    sm.append(["'Verified' rows are only those for which the official Udyam "
               "portal returned registration details (or that you entered "
               "manually from the portal)."])
    sm.column_dimensions["A"].width = 26
    sm.column_dimensions["B"].width = 22

    wb.save(path)
    return counts
