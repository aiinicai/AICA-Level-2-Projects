"""Convert the model's Markdown answer to HTML (for the screen) and to a
Word document (for editing and filing)."""
from __future__ import annotations

import re
from pathlib import Path

import markdown as md
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def to_html(text: str) -> str:
    html = md.markdown(text, extensions=["tables", "sane_lists", "nl2br"])
    # highlight [[placeholders]] the CA must fill
    return re.sub(r"\[\[(.+?)\]\]", r'<mark class="fill">[[\1]]</mark>', html)


INLINE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|__.+?__|\*[^*\s][^*]*?\*|_[^_\s][^_]*?_|`[^`]+`|\[\[.+?\]\])")


def _runs(par, text: str, bold=False):
    for piece in INLINE.split(text):
        if not piece:
            continue
        b, i, hl = bold, False, False
        if piece.startswith("***") and piece.endswith("***"):
            piece, b, i = piece[3:-3], True, True
        elif (piece.startswith("**") and piece.endswith("**")) or (piece.startswith("__") and piece.endswith("__")):
            piece, b = piece[2:-2], True
        elif piece.startswith("[[") and piece.endswith("]]"):
            hl = True
        elif (piece.startswith("*") and piece.endswith("*")) or (piece.startswith("_") and piece.endswith("_")):
            piece, i = piece[1:-1], True
        elif piece.startswith("`") and piece.endswith("`"):
            piece = piece[1:-1]
        run = par.add_run(piece)
        run.bold = b or None
        run.italic = i or None
        if hl:
            rpr = run._r.get_or_add_rPr()
            h = OxmlElement("w:highlight")
            h.set(qn("w:val"), "yellow")
            rpr.append(h)


def _table(doc, rows: list[list[str]]):
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(cols):
            val = row[ci] if ci < len(row) else ""
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _runs(p, val, bold=(ri == 0))
            if re.fullmatch(r"\s*[₹Rs.\-\d,()\s]+\s*", val) and re.search(r"\d", val):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()


def to_docx(text: str, out: Path, title: str | None = None, footer: str | None = None) -> Path:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Cambria"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria")
    st.font.size = Pt(11.5)
    st.paragraph_format.space_after = Pt(6)
    for s in doc.sections:
        s.left_margin = s.right_margin = Pt(72)
        s.top_margin = s.bottom_margin = Pt(64)
        if footer:
            fp = s.footer.paragraphs[0]
            fp.text = footer
            fp.runs[0].font.size = Pt(8)
            fp.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x70)
    for lvl in (1, 2, 3):
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.name = "Cambria"
        hs.font.color.rgb = RGBColor(0x1C, 0x2A, 0x3A)
        hs.font.size = Pt({1: 15, 2: 13, 3: 12}[lvl])
    if title:
        doc.add_heading(title, level=1)

    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    para_buf: list[str] = []

    def flush():
        if para_buf:
            p = doc.add_paragraph()
            for k, ln in enumerate(para_buf):
                if k:
                    p.add_run().add_break()
                _runs(p, ln.strip())
            para_buf.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            flush()
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|?\s*:?-{2,}", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            _table(doc, rows)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            flush()
            lvl = min(len(m.group(1)), 3)
            h = doc.add_heading(level=lvl)
            _runs(h, m.group(2).strip("# "))
            i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            flush()
            i += 1
            continue
        m = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)", ln)
        if m:
            flush()
            indent = len(m.group(1).replace("\t", "    ")) // 2
            numbered = m.group(2)[0].isdigit()
            style = "List Number" if numbered else "List Bullet"
            if indent:
                style += " 2"
            p = doc.add_paragraph(style=style)
            _runs(p, m.group(3))
            i += 1
            continue
        if s.startswith(">"):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            _runs(p, s.lstrip("> "))
            i += 1
            continue
        para_buf.append(ln)
        i += 1
    flush()
    doc.save(out)
    return out
