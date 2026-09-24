# Copyright © 2026 Arjun Seth. All rights reserved. Viewable for ICAI evaluation only.
# No permission to copy, modify, redistribute or use commercially without written permission.
"""The disclosure report as a formatted Excel workbook (openpyxl).

One sheet per part of the report (Summary, Lease Register, Maturity, Liability Movement, ROU Assets, Lease Costs,
Consolidated Memo, Checks, Notes), drawn generically from ``core.reports`` sections: navy headers, zebra rows, bold
totals, Arial, amounts as real numbers in the user's number format and the lease currency. Landscape, fit to one page
wide, with page numbers. Values only (no formulas), so it looks the same in Protected View.
"""
import math
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border

from core.dashboard_report import COLORS
from core.dashboard_xlsx import (
    _BOX,
    _NAVY_LINE,
    _THIN,
    _fill,
    _font,
    _header,
    _print_setup,
    _section,
    _title,
    number_format,
)
from core.reports import report_file_stem

SHEETS = [
    ("Summary", ["details", "summary", "averages"]),
    ("Lease Register", ["register"]),
    ("Maturity", ["maturity"]),
    ("Liability Movement", ["liability"]),
    ("ROU Assets", ["rou"]),
    ("Lease Costs", ["costs"]),
    ("Consolidated Memo", ["memo"]),
    ("Checks", ["checks"]),
    ("Notes", ["notes"]),
]
_KIND_FORMATS = {"count": "#,##0", "percent": "0.0%", "rate": "0.00%", "years": "0.0", "date": "dd-mmm-yyyy"}


def _format_for(kind: str, report: dict):
    return number_format(report) if kind == "money" else _KIND_FORMATS.get(kind)


def _row(ws, row: int, first_col: int, values: list, kinds: list, report: dict, zebra: bool, total: bool = False) -> None:
    for offset, (value, kind) in enumerate(zip(values, kinds)):
        cell = ws.cell(row=row, column=first_col + offset, value="-" if value is None else value)
        cell.font = _font(bold=total)
        cell.border = Border(left=_THIN, right=_THIN, bottom=_THIN, top=_NAVY_LINE if total else _THIN)
        if total:
            cell.fill = _fill(COLORS["band"])
        elif zebra:
            cell.fill = _fill(COLORS["zebra"])
        numeric = kind != "text" and value is not None
        cell.alignment = Alignment(horizontal="right" if numeric or (kind != "text" and value is None) else "left", vertical="center", wrap_text=kind == "text")
        fmt = _format_for(kind, report)
        if fmt and value is not None:
            cell.number_format = fmt


def _measure(sections: list) -> dict:
    """Column widths for a sheet: the longest text in each column across its tables (capped)."""
    widths = {}
    for section in sections:
        for block in section["blocks"]:
            if block["type"] != "table":
                continue
            for index, (header, kind) in enumerate(block["columns"]):
                samples = [len(header) * 0.9 + 2] + [len(str(row[index])) + 2 for row in block["rows"] + ([block["total"]] if block.get("total") else []) if row[index] is not None and kind == "text"]
                if kind != "text":
                    samples.append(20)
                widths[index] = max(widths.get(index, 0), min(max(samples), 46))
    return widths


def _write_sheet(ws, sections: list, report: dict, title: str) -> None:
    widths = _measure(sections)
    ncols = max([len(widths), 6])
    for index in range(ncols):
        ws.column_dimensions[chr(ord("A") + index)].width = max(widths.get(index, 16), 16 if index else 26)
    total_width = sum(ws.column_dimensions[chr(ord("A") + i)].width for i in range(ncols))
    _title(ws, title, "{}  |  as at {}  |  {}".format(report["framework_title"], report["as_at_label"], report["currency_name"]), ncols)
    row = 4
    for section in sections:
        _section(ws, row, section["title"], 1, ncols)
        row += 2
        for block in section["blocks"]:
            kind = block["type"]
            if kind in ("text", "heading", "bullets"):
                lines = [block["text"]] if kind != "bullets" else ["\u2022  " + item for item in block["items"]]
                for line in lines:
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
                    cell = ws.cell(row=row, column=1, value=line)
                    cell.font = _font(bold=kind == "heading", color=COLORS["navy"] if kind == "heading" else "1A2233", size=10 if kind != "heading" else 11)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    ws.row_dimensions[row].height = 15 * max(1, math.ceil(len(line) / max(60, total_width * 1.05)))
                    row += 1
                row += 1
            elif kind == "table":
                if block.get("title"):
                    ws.cell(row=row, column=1, value=block["title"]).font = _font(size=10.5, bold=True, color=COLORS["navy"])
                    row += 1
                columns, kinds = block["columns"], [k for _, k in block["columns"]]
                _header(ws, row, 1, [h for h, _ in columns], right_from=1)
                for offset, (_, k) in enumerate(columns):
                    ws.cell(row=row, column=1 + offset).alignment = Alignment(horizontal="right" if k != "text" else "left", vertical="center", wrap_text=True)
                for index, values in enumerate(block["rows"]):
                    _row(ws, row + 1 + index, 1, values, kinds, report, zebra=index % 2 == 1)
                row += len(block["rows"]) + 1
                if block.get("total"):
                    _row(ws, row, 1, block["total"], kinds, report, zebra=False, total=True)
                    row += 1
                for note in block.get("notes", []):
                    ws.cell(row=row, column=1, value=note).font = _font(italic=True, color=COLORS["muted"], size=9)
                    row += 1
                row += 2
    _print_setup(ws, report)


def build_report_workbook(report: dict) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    by_id = {section["id"]: section for section in report["sections"]}
    for name, ids in SHEETS:
        sections = [by_id[i] for i in ids if i in by_id]
        if not sections:
            continue
        ws = wb.create_sheet(name)
        _write_sheet(ws, sections, report, report["title"])
        if name == "Summary":  # a short sheet: print it on one page instead of spilling a block onto a second
            ws.page_setup.fitToHeight = 1
        if name == "Lease Register":  # keep the header row in view and repeat it on every printed page
            header_row = next(r for r in range(1, 20) if ws.cell(row=r, column=1).value == "Lease ID")
            ws.freeze_panes = "A{}".format(header_row + 1)
            ws.print_title_rows = "{0}:{0}".format(header_row)
    wb.properties.title, wb.properties.creator = report["title"], "LeaseIQ Pro"
    wb.properties.subject = "{} - as at {}".format(report["currency_name"], report["as_at_label"])
    wb.properties.created = report["generated_at"]
    return wb


def build_report_excel(report: dict) -> bytes:
    buffer = BytesIO()
    build_report_workbook(report).save(buffer)
    return buffer.getvalue()


def report_excel_name(report: dict) -> str:
    return report_file_stem(report) + ".xlsx"
