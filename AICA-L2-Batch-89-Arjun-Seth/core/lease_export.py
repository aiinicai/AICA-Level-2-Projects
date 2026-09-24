# Copyright © 2026 Arjun Seth. All rights reserved. Viewable for ICAI evaluation only.
# No permission to copy, modify, redistribute or use commercially without written permission.
"""Export ONE lease's schedules and journal entries as an Excel workbook or as CSV files.

Pure Python. Built from the lease's SAVED results (``core.workflow.get_lease_export_data``), so the files always match the
screens.

* Excel: Summary (details, key results, classification and its five tests), a schedule per framework (every stored column,
  flow totals), and the journal entries per framework (Day 1 legs, then every month) with debit / credit totals.
  Real numbers in the user's number format and the lease currency; Arial; landscape; header rows frozen and repeated.
* CSV (four files: schedule and journal for each of Ind AS 116 and ASC 842, also as one zip): machine-friendly - one header
  row, ISO dates, plain decimals in the currency's decimals with no thousands separators, UTF-8 with a byte-order mark so
  Excel opens it correctly. Each row starts with the lease ID and currency, so files from many leases can be stacked.
"""
import csv
import io
import re
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border

from core.dashboard_report import COLORS
from core.dashboard_xlsx import _NAVY_LINE, _THIN, _fill, _font, _header, _print_setup, _section, _title, number_format
from core.formatting import CURRENCIES, NUMBER_FORMATS, currency_decimals, parse_currency
from core.results import FRAMEWORK_LABELS, FRAMEWORKS, format_tests

FRAMEWORK_STEMS = {"IND_AS_116": "IndAS116", "ASC_842": "ASC842"}

_GROSS_COLUMNS = [
    ("period", "Period", "count"), ("period_date", "Date", "date"), ("escalated_contractual_rent", "Contractual rent", "money"),
    ("net_cash_payment", "Cash payment", "money"), ("opening_liability", "Opening liability", "money"),
    ("interest_expense", "Interest", "money"), ("principal_repayment", "Principal repayment", "money"),
    ("closing_liability", "Closing liability", "money"), ("rou_gross_cost", "ROU gross cost", "money"),
    ("accum_amortization_opening", "Accumulated amortization - opening", "money"), ("amortization_expense", "Amortization", "money"),
    ("accum_amortization_closing", "Accumulated amortization - closing", "money"), ("rou_net_carrying_value", "ROU net carrying value", "money"),
    ("security_deposit_opening", "Security deposit - opening", "money"), ("security_deposit_closing", "Security deposit - closing", "money"),
]
_NET_COLUMNS = [
    ("period", "Period", "count"), ("period_date", "Date", "date"), ("escalated_contractual_rent", "Contractual rent", "money"),
    ("net_cash_payment", "Cash payment", "money"), ("opening_liability", "Opening liability", "money"),
    ("interest_expense", "Interest", "money"), ("principal_repayment", "Principal repayment", "money"),
    ("closing_liability", "Closing liability", "money"), ("single_lease_cost", "Single lease cost", "money"),
    ("rou_reduction_plug", "ROU reduction", "money"), ("rou_net_carrying_value", "ROU net carrying value", "money"),
    ("security_deposit_opening", "Security deposit - opening", "money"), ("security_deposit_closing", "Security deposit - closing", "money"),
]
SCHEDULE_COLUMNS = {"gross_accum": _GROSS_COLUMNS, "net_direct": _NET_COLUMNS}
# flows that make sense to total over the life of the lease (balances do not)
_FLOW_FIELDS = {"escalated_contractual_rent", "net_cash_payment", "interest_expense", "principal_repayment", "amortization_expense", "single_lease_cost", "rou_reduction_plug"}
JOURNAL_COLUMNS = [
    ("lease", "Lease ID", "text"), ("currency", "Currency", "text"), ("framework", "Framework", "text"), ("entry", "Entry", "text"),
    ("date", "Entry date", "date"), ("line", "Line", "count"), ("account", "Account", "text"), ("debit", "Debit", "money"),
    ("credit", "Credit", "money"), ("narration", "Narration", "text"), ("balanced", "Balanced", "text"),
]


def safe_name(text: str) -> str:
    """A file-name-safe version of a lease ID ('LS-0005' stays, '#12' becomes '12')."""
    return re.sub(r"[^A-Za-z0-9_-]", "", str(text or "")) or "lease"


# --------------------------------------------------------------------------- #
# shaping the saved rows
# --------------------------------------------------------------------------- #
def build_lease_export(data: dict, number_style: str = "indian", prepared_by: str = "", generated_at: datetime = None) -> dict:
    """Shape a lease's saved results (``get_lease_export_data``) into the tables the files are written from.

    Raises ValueError if ``data`` is None (a lease with no calculation results has nothing to export).
    """
    if not data:
        raise ValueError("This lease has no calculation results yet, so there is nothing to export.")
    stored, case, journal = data["stored"], data["case_data"], data["journal_rows"]
    generated_at = generated_at or datetime.now()
    ref = stored["case"]["lease_ref"] or "#{}".format(stored["case"]["case_id"])
    code = parse_currency(stored["case"]["currency"]) or str(stored["case"]["currency"] or "INR").upper()
    decimals = currency_decimals(code)
    calc, classification = stored["calculation"], stored["classification"]

    meta = {
        "title": "Lease {} - schedules and journal entries".format(ref),
        "lease_ref": ref, "safe_ref": safe_name(ref), "currency": code, "symbol": CURRENCIES[code][1] if code in CURRENCIES else code,
        "currency_name": "{} - {}".format(code, CURRENCIES[code][0]) if code in CURRENCIES else code, "decimals": decimals,
        "number_style": number_style, "number_format_label": NUMBER_FORMATS.get(number_style, number_style),
        "prepared_by": prepared_by, "generated_at": generated_at, "as_at_label": generated_at.strftime("%d-%b-%Y"),
    }

    def yes_no(flag):
        return "Yes" if flag else "No"

    details = [
        ("Lease ID", ref, "text"), ("Status", case.get("status") or stored["case"]["status"], "text"), ("Lessor", case.get("lessor") or "-", "text"),
        ("Lessee", case.get("lessee") or "-", "text"), ("Asset type", case.get("asset_type") or "-", "text"), ("Currency", meta["currency_name"], "text"),
        ("Commencement date", case.get("commencement_date"), "date"), ("End date", case.get("end_date"), "date"), ("Term (months)", case.get("term_months"), "count"),
        ("Monthly rent, first lease year", case.get("base_rent"), "money"), ("Annual rent escalation", case.get("escalation"), "rate"),
        ("Advance rent paid", case.get("prepaid_rent"), "money"), ("Months covered by advance rent", case.get("prepaid_rent_months"), "count"),
        ("Refundable security deposit", case.get("deposit"), "money"), ("Initial direct costs", case.get("idc"), "money"),
        ("Lease incentives received", case.get("incentives"), "money"), ("Restoration obligation", case.get("restoration_cost"), "money"),
        ("Discount rate (incremental borrowing rate, per year)", case.get("ibr"), "rate"), ("Asset fair value", case.get("asset_fair_value"), "money"),
        ("Asset economic life (months)", case.get("asset_economic_life_months"), "count"),
        ("Reporting framework", {"BOTH": "Ind AS 116 and ASC 842", "IND_AS_116": "Ind AS 116 only", "ASC_842": "ASC 842 only"}.get(case.get("reporting_framework"), "-"), "text"),
        ("ASC 842 classification set manually", yes_no(classification.get("is_override")), "text"),
    ]
    results = [
        ("Lease liability at commencement", calc["lease_liability_initial"], "money"), ("Right-of-use asset at commencement", calc["rou_asset_gross"], "money"),
        ("Security deposit - present value", calc["security_deposit_pv"], "money"), ("Security deposit - discount", calc["security_deposit_discount"], "money"),
        ("Monthly discount rate", calc["monthly_rate"], "rate"), ("Single straight-line lease cost per month", calc["single_lease_cost_per_month"], "money"),
        ("Total contractual rent", calc["total_contractual_rent"], "money"), ("Total undiscounted lease payments", calc["total_undiscounted_payments"], "money"),
    ]

    schedules = {}
    for framework in FRAMEWORKS:
        block = stored["schedules"][framework]
        method = block["rou_method"] or "net_direct"
        columns = SCHEDULE_COLUMNS[method]
        schedules[framework] = {
            "label": FRAMEWORK_LABELS[framework], "rou_method": method, "columns": columns,
            "rows": [{field: row.get(field) for field, _, _ in columns} for row in block["rows"]],
        }

    journals = {}
    for framework in FRAMEWORKS:
        day1 = sorted((r for r in journal if r["entry_type"] == "DAY1"), key=lambda r: (r["leg"] or "", r["line_no"]))
        periodic = sorted((r for r in journal if r["entry_type"] == "PERIODIC" and r["framework"] == framework), key=lambda r: (r["period_number"], r["line_no"]))
        rows = []
        for r in day1 + periodic:
            rows.append(
                {
                    "lease": ref, "currency": code, "framework": FRAMEWORK_LABELS[framework],
                    "entry": "Day 1 - Leg {}".format(r["leg"]) if r["entry_type"] == "DAY1" else "Period {}".format(r["period_number"]),
                    "date": r["entry_date"], "line": r["line_no"], "account": r["account"], "debit": float(r["debit"]), "credit": float(r["credit"]),
                    "narration": r["narration"] or "", "balanced": yes_no(r["is_balanced"]),
                }
            )
        journals[framework] = {"label": FRAMEWORK_LABELS[framework], "rows": rows}

    return {
        "meta": meta, "details": details, "results": results,
        "classification": {
            "asc842": classification["asc842_classification"], "ind_as116": classification["ind_as116_exemption"],
            "rationale": classification.get("ind_as116_rationale") or "", "is_override": bool(classification.get("is_override")),
            "tests": format_tests(classification["tests"]),
        },
        "schedules": schedules, "journals": journals,
    }


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #
def _csv_value(value, kind: str, decimals: int) -> str:
    if value is None:
        return ""
    if kind == "money":
        return "{:.{d}f}".format(float(value), d=decimals)
    if kind == "date":
        return value.isoformat() if hasattr(value, "isoformat") else str(value)
    if kind == "count":
        return str(int(value))
    return str(value)


def _csv_bytes(header: list, rows: list) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")  # the byte-order mark makes Excel read it as UTF-8


def schedule_csv(export: dict, framework: str) -> bytes:
    """One framework's schedule: Lease ID, Currency, then every stored column."""
    meta, block = export["meta"], export["schedules"][framework]
    header = ["Lease ID", "Currency"] + [label for _, label, _ in block["columns"]]
    rows = [[meta["lease_ref"], meta["currency"]] + [_csv_value(row[field], kind, meta["decimals"]) for field, _, kind in block["columns"]] for row in block["rows"]]
    return _csv_bytes(header, rows)


def journal_csv(export: dict, framework: str) -> bytes:
    """One framework's journal entries: the Day 1 legs, then every monthly entry (ready for a ledger import)."""
    meta = export["meta"]
    header = [label for _, label, _ in JOURNAL_COLUMNS]
    rows = [[_csv_value(row[field], kind, meta["decimals"]) for field, _, kind in JOURNAL_COLUMNS] for row in export["journals"][framework]["rows"]]
    return _csv_bytes(header, rows)


def csv_files(export: dict) -> dict:
    """{file name: bytes} for the four CSV files."""
    stem = "LeaseIQ_{}".format(export["meta"]["safe_ref"])
    files = {}
    for framework in FRAMEWORKS:
        tag = FRAMEWORK_STEMS[framework]
        files["{}_Schedule_{}.csv".format(stem, tag)] = schedule_csv(export, framework)
        files["{}_Journal_{}.csv".format(stem, tag)] = journal_csv(export, framework)
    return files


def csv_zip(export: dict) -> bytes:
    """The four CSV files in one zip."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in csv_files(export).items():
            archive.writestr(name, content)
    return buffer.getvalue()


def zip_file_name(export: dict) -> str:
    return "LeaseIQ_{}_CSV.zip".format(export["meta"]["safe_ref"])


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
_KIND_FORMATS = {"count": "#,##0", "date": "dd-mmm-yyyy", "rate": "0.00%"}


def _fmt(kind: str, meta: dict):
    return number_format(meta) if kind == "money" else _KIND_FORMATS.get(kind)


def _write_table(ws, top: int, headers: list, kinds: list, rows: list, meta: dict, total: list = None) -> int:
    """Header + rows (+ totals). Returns the last row used. Numbers are right-aligned, text left."""
    _header(ws, top, 1, headers, right_from=len(headers) + 1)
    for offset, kind in enumerate(kinds):
        ws.cell(row=top, column=1 + offset).alignment = Alignment(horizontal="right" if kind in ("money", "count", "rate") else "left", vertical="center", wrap_text=True)
    for index, values in enumerate(rows):
        row = top + 1 + index
        for offset, (value, kind) in enumerate(zip(values, kinds)):
            cell = ws.cell(row=row, column=1 + offset, value=value)
            cell.font = _font()
            cell.border = Border(left=_THIN, right=_THIN, bottom=_THIN, top=_THIN)
            if index % 2 == 1:
                cell.fill = _fill(COLORS["zebra"])
            cell.alignment = Alignment(horizontal="right" if kind in ("money", "count", "rate") else "left", vertical="center")
            if _fmt(kind, meta) and value is not None:
                cell.number_format = _fmt(kind, meta)
    last = top + len(rows)
    if total is not None:
        last += 1
        for offset, (value, kind) in enumerate(zip(total, kinds)):
            cell = ws.cell(row=last, column=1 + offset, value=value)
            cell.font, cell.fill = _font(bold=True), _fill(COLORS["band"])
            cell.border = Border(left=_THIN, right=_THIN, bottom=_THIN, top=_NAVY_LINE)
            cell.alignment = Alignment(horizontal="right" if kind in ("money", "count", "rate") else "left", vertical="center")
            if _fmt(kind, meta) and value is not None:
                cell.number_format = _fmt(kind, meta)
    return last


def _summary_sheet(wb, export: dict) -> None:
    meta = export["meta"]
    ws = wb.active
    ws.title = "Summary"
    for letter, width in zip("ABCD", (46, 40, 16, 16)):
        ws.column_dimensions[letter].width = width
    _title(ws, meta["title"], "{}  |  generated {} by {}".format(meta["currency_name"], meta["generated_at"].strftime("%d-%b-%Y %H:%M"), meta["prepared_by"] or "-"), 4)

    def key_values(row: int, heading: str, items: list) -> int:
        _section(ws, row, heading, 1, 2)
        row += 1
        for index, (label, value, kind) in enumerate(items):
            ws.cell(row=row, column=1, value=label).font = _font(bold=True, color=COLORS["muted"])
            cell = ws.cell(row=row, column=2, value="-" if value is None else value)
            cell.font = _font()
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if value is not None and _fmt(kind, meta):
                cell.number_format = _fmt(kind, meta)
            row += 1
        return row + 1

    row = key_values(4, "Lease details", export["details"])
    row = key_values(row, "Key results at commencement", export["results"])
    classification = export["classification"]
    row = key_values(row, "Classification", [
        ("ASC 842", classification["asc842"] + (" (set manually)" if classification["is_override"] else ""), "text"),
        ("Ind AS 116", classification["ind_as116"], "text"), ("Ind AS 116 rationale", classification["rationale"] or "-", "text"),
    ])
    _section(ws, row, "ASC 842 classification tests (a finance lease if ANY test is met)", 1, 4)
    tests = [[t["Test"], t["Computed"], t["Threshold"], t["Met"]] for t in classification["tests"]]
    last = _write_table(ws, row + 1, ["Test", "Computed", "Threshold", "Met"], ["text", "text", "text", "text"], tests, meta)
    ws.cell(row=last + 2, column=1, value="Day 1 journal entries are the same under both frameworks. Figures are values saved when the lease was calculated.").font = _font(italic=True, color=COLORS["muted"], size=9)
    _print_setup(ws, meta)


def _schedule_sheet(wb, export: dict, framework: str) -> None:
    meta, block = export["meta"], export["schedules"][framework]
    ws = wb.create_sheet("Schedule - {}".format(block["label"]))
    columns = block["columns"]
    _title(ws, "{} - {} schedule".format(meta["lease_ref"], block["label"]), "{}  |  ROU asset method: {}".format(
        meta["currency_name"], "gross cost and accumulated amortization" if block["rou_method"] == "gross_accum" else "single lease cost, asset reduced directly"), len(columns))
    for index, (_, label, kind) in enumerate(columns):
        ws.column_dimensions[chr(ord("A") + index)].width = 9 if kind == "count" else 13 if kind == "date" else 19
    kinds = [kind for _, _, kind in columns]
    rows = [[row[field] for field, _, _ in columns] for row in block["rows"]]
    total = ["Total"] + [None] * (len(columns) - 1)
    for index, (field, _, _) in enumerate(columns):
        if field in _FLOW_FIELDS:
            total[index] = sum(row[field] or 0.0 for row in block["rows"])
    last = _write_table(ws, 4, [label for _, label, _ in columns], kinds, rows, meta, total)
    ws.freeze_panes = "C5"  # the header row, the period and the date stay in view
    ws.auto_filter.ref = "A4:{}{}".format(chr(ord("A") + len(columns) - 1), last - 1)
    _print_setup(ws, meta, title_rows="4:4")


def _journal_sheet(wb, export: dict, framework: str) -> None:
    meta, block = export["meta"], export["journals"][framework]
    ws = wb.create_sheet("Journal - {}".format(block["label"]))
    columns = [c for c in JOURNAL_COLUMNS if c[0] not in ("lease", "currency", "framework")]
    for letter, width in zip("ABCDEFGH", (18, 13, 7, 40, 19, 19, 72, 10)):
        ws.column_dimensions[letter].width = width
    _title(ws, "{} - {} journal entries".format(meta["lease_ref"], block["label"]), "{}  |  Day 1 legs, then every monthly entry".format(meta["currency_name"]), len(columns))
    rows = [[row[field] for field, _, _ in columns] for row in block["rows"]]
    debit_total, credit_total = sum(r["debit"] for r in block["rows"]), sum(r["credit"] for r in block["rows"])
    total = ["Total", None, None, None, debit_total, credit_total, "Debits equal credits" if abs(debit_total - credit_total) < 0.01 * max(1, len(block["rows"])) else "DEBITS DO NOT EQUAL CREDITS", None]
    last = _write_table(ws, 4, [label for _, label, _ in columns], [kind for _, _, kind in columns], rows, meta, total)
    for row in range(5, last):  # long narrations wrap instead of being clipped
        ws.cell(row=row, column=7).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:H{}".format(last - 1)
    _print_setup(ws, meta, title_rows="4:4")


def build_lease_workbook(export: dict) -> Workbook:
    wb = Workbook()
    _summary_sheet(wb, export)
    for framework in FRAMEWORKS:
        _schedule_sheet(wb, export, framework)
    for framework in FRAMEWORKS:
        _journal_sheet(wb, export, framework)
    meta = export["meta"]
    wb.properties.title, wb.properties.creator = meta["title"], "LeaseIQ Pro"
    wb.properties.subject, wb.properties.created = meta["currency_name"], meta["generated_at"]
    return wb


def build_lease_excel(export: dict) -> bytes:
    buffer = io.BytesIO()
    build_lease_workbook(export).save(buffer)
    return buffer.getvalue()


def excel_file_name(export: dict) -> str:
    return "LeaseIQ_{}_Export.xlsx".format(export["meta"]["safe_ref"])
