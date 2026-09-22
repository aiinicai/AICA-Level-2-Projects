# Copyright © 2026 Arjun Seth. All rights reserved. Viewable for ICAI evaluation only.
# No permission to copy, modify, redistribute or use commercially without written permission.
"""The dashboard report as a formatted Excel workbook (openpyxl).

Four sheets: Dashboard (details, key figures, status, classification, maturity - each with a native Excel chart),
Interest vs Principal (60 months + line chart), Future Payments (12 / 24 / 60 months + chart) and Notes.
Numbers are written as VALUES (not formulas) in the user's number format and the lease currency, so the
workbook looks the same in Excel's Protected View and in any viewer. Arial throughout; print-ready.
"""
import math
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

from core.branding import FOOTER_TEXT
from core.dashboard_report import CLASSIFICATION_HEX, COLORS, STATUS_HEX, file_stem

FONT = "Arial"
_THIN = Side(style="thin", color=COLORS["line"])
_NAVY_LINE = Side(style="medium", color=COLORS["navy"])
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CHART_ROWS = 15  # rows a chart (about 7 cm high) occupies


def _font(**kwargs):
    return Font(name=FONT, size=kwargs.pop("size", 10), **kwargs)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


# --------------------------------------------------------------------------- #
# number formats
# --------------------------------------------------------------------------- #
def number_format(report: dict, with_symbol: bool = True, decimals: int = None) -> str:
    """The Excel format for an amount: the currency's symbol, decimals, and Indian (12,34,567) or international grouping.
    ``decimals`` overrides the currency's own (charts use whole numbers)."""
    places = report["decimals"] if decimals is None else decimals
    fraction = "." + "0" * places if places else ""
    prefix = '"{} "'.format(report["symbol"].replace('"', "")) if with_symbol else ""
    if report["number_style"] == "indian":
        return "[>=10000000]{p}##\\,##\\,##\\,##0{f};[>=100000]{p}##\\,##\\,##0{f};{p}##,##0{f}".format(p=prefix, f=fraction)
    return "{p}#,##0{f}".format(p=prefix, f=fraction)


PERCENT_FORMAT = "0.0%"
CHANGE_FORMAT = "+0.0%;-0.0%;0.0%"
MONTH_FORMAT = "mmm yyyy"
COUNT_FORMAT = "#,##0"


# --------------------------------------------------------------------------- #
# building blocks
# --------------------------------------------------------------------------- #
def _title(ws, text: str, subtitle: str, last_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(row=1, column=1, value=text)
    cell.font, cell.fill = _font(size=16, bold=True, color="FFFFFF"), _fill(COLORS["navy"])
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32
    for col in range(2, last_col + 1):
        ws.cell(row=1, column=col).fill = _fill(COLORS["navy"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    sub = ws.cell(row=2, column=1, value=subtitle)
    sub.font, sub.alignment = _font(italic=True, color=COLORS["muted"]), Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 20


def _section(ws, row: int, text: str, first_col: int, last_col: int) -> None:
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = Border(bottom=_NAVY_LINE)
    cell = ws.cell(row=row, column=first_col, value=text)
    cell.font = _font(size=12, bold=True, color=COLORS["navy"])
    ws.row_dimensions[row].height = 22


def _header(ws, row: int, first_col: int, headers: list, right_from: int = 1) -> None:
    for offset, text in enumerate(headers):
        cell = ws.cell(row=row, column=first_col + offset, value=text)
        cell.font, cell.fill, cell.border = _font(bold=True, color="FFFFFF"), _fill(COLORS["navy"]), _BOX
        cell.alignment = Alignment(horizontal="right" if offset >= right_from else "left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _body_row(ws, row: int, first_col: int, values: list, formats: list, zebra: bool, bold: bool = False, total: bool = False) -> None:
    for offset, (value, fmt) in enumerate(zip(values, formats)):
        cell = ws.cell(row=row, column=first_col + offset, value=value)
        cell.font = _font(bold=bold or total)
        cell.border = Border(left=_THIN, right=_THIN, bottom=_THIN, top=_NAVY_LINE if total else _THIN)
        if total:
            cell.fill = _fill(COLORS["band"])
        elif zebra:
            cell.fill = _fill(COLORS["zebra"])
        cell.alignment = Alignment(horizontal="right" if offset > 0 else "left", vertical="center")
        if fmt:
            cell.number_format = fmt


def _table(ws, row: int, first_col: int, headers: list, rows: list, formats: list, total: list = None) -> int:
    """Header + body (+ optional total row). Returns the last row used."""
    _header(ws, row, first_col, headers)
    for index, values in enumerate(rows):
        _body_row(ws, row + 1 + index, first_col, values, formats, zebra=index % 2 == 1)
    last = row + len(rows)
    if total is not None:
        last += 1
        _body_row(ws, last, first_col, total, formats, zebra=False, total=True)
    return last


def _widths(ws, widths: dict) -> None:
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def _print_setup(ws, report: dict, title_rows: str = None) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top, ws.page_margins.bottom = 0.6, 0.7
    ws.oddFooter.left.text, ws.oddFooter.left.size = FOOTER_TEXT, 8
    ws.oddFooter.center.text, ws.oddFooter.center.size = "{} - {}".format(report["title"], report.get("month_label") or report["as_at_label"]), 8
    ws.oddFooter.right.text, ws.oddFooter.right.size = "Page &P of &N", 8
    ws.print_options.horizontalCentered = True
    if title_rows:
        ws.print_title_rows = title_rows
    ws.sheet_view.showGridLines = False


def _style_axes(chart) -> None:
    chart.x_axis.delete = False  # openpyxl 3.1 hides axes in Excel unless told otherwise
    chart.y_axis.delete = False


def _doughnut(ws, title: str, header_row: int, first: int, last: int, colors: list, anchor: str) -> None:
    chart = DoughnutChart()
    chart.title, chart.holeSize, chart.height, chart.width = title, 55, 7.0, 12.0
    chart.add_data(Reference(ws, min_col=2, min_row=header_row, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    series = chart.series[0]
    for index, color in enumerate(colors):
        point = DataPoint(idx=index)
        point.graphicalProperties.solidFill = color
        series.dPt.append(point)
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent, chart.dataLabels.showVal = True, False
    chart.dataLabels.showCatName = chart.dataLabels.showSerName = chart.dataLabels.showLeaderLines = False
    chart.legend.position = "r"
    ws.add_chart(chart, anchor)


# --------------------------------------------------------------------------- #
# the sheets
# --------------------------------------------------------------------------- #
def _dashboard_sheet(wb, report: dict) -> None:
    ws = wb.active
    ws.title = "Dashboard"
    _widths(ws, {"A": 30, "B": 20, "C": 20, "D": 22, "E": 3, **{letter: 11 for letter in "FGHIJKL"}})
    _title(ws, report["title"], "Portfolio dashboard for {}  |  {}".format(report["month_label"], report["currency_label"]), 12)
    money = number_format(report)

    row = 4
    _section(ws, row, "Report details", 1, 4)
    details = [
        ("Prepared by", report["prepared_by"] or "-"),
        ("Generated on", report["generated_at"].strftime("%d-%b-%Y %H:%M")),
        ("Currency", report["currency_label"]),
        ("Number format", report["number_format_label"]),
        ("Leases included in the amounts", report["include_label"]),
        ("Reporting view", report["framework_label"]),
        ("Leases counted in the amounts", "{} of {} in {}".format(report["leases_counted"], report["total_leases"], report["currency"])),
    ]
    for offset, (label, value) in enumerate(details, start=1):
        ws.cell(row=row + offset, column=1, value=label).font = _font(bold=True, color=COLORS["muted"])
        ws.merge_cells(start_row=row + offset, start_column=2, end_row=row + offset, end_column=4)
        cell = ws.cell(row=row + offset, column=2, value=value)
        cell.font, cell.alignment = _font(), Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row + offset].height = 18
    row += len(details) + 2

    _section(ws, row, "Key figures - {}".format(report["month_label"]), 1, 4)
    kpi_rows = []
    for kpi in report["kpis"]:
        if kpi["kind"] == "count":
            kpi_rows.append([kpi["label"], kpi["value"], "-", kpi["note"] or "-"])
        else:
            change = kpi["change_percent"] / 100.0 if kpi["change_percent"] is not None else "n/a"
            kpi_rows.append([kpi["label"], kpi["value"], kpi["previous"], change])
    header_row = row + 1
    _header(ws, header_row, 1, ["Metric", "Current month", "Previous month", "Change vs previous month"])
    for index, values in enumerate(kpi_rows):
        line = header_row + 1 + index
        kind = report["kpis"][index]["kind"]
        _body_row(ws, line, 1, values, [None, COUNT_FORMAT if kind == "count" else money, money if kind == "money" else None, CHANGE_FORMAT], zebra=index % 2 == 1)
    row = header_row + len(kpi_rows) + 2

    def block(title, headers, rows, formats, total, chart):
        nonlocal row
        _section(ws, row, title, 1, 4)
        start = row
        if not rows:
            ws.cell(row=row + 1, column=1, value="No leases in {} yet.".format(report["currency"])).font = _font(italic=True, color=COLORS["muted"])
            row += 3
            return
        header_at = row + 1
        last = _table(ws, header_at, 1, headers, rows, formats, total)
        if chart:
            chart(header_at, header_at + 1, header_at + len(rows), "F{}".format(start))
        row = max(last, start + _CHART_ROWS) + 2

    block(
        "Leases by Status",
        ["Status", "Leases", "Share of leases"],
        [[item["label"], item["count"], item["percent"] / 100.0] for item in report["status"]],
        [None, COUNT_FORMAT, PERCENT_FORMAT],
        ["Total", report["status_total"], 1.0 if report["status_total"] else 0.0],
        lambda h, f, l, anchor: _doughnut(ws, "Leases by Status", h, f, l, [STATUS_HEX.get(i["label"], "9AA5B1") for i in report["status"]], anchor),
    )
    block(
        "Leases by Classification (ASC 842)",
        ["Classification", "Leases", "Share of leases"],
        [[item["label"], item["count"], item["percent"] / 100.0] for item in report["classification"]],
        [None, COUNT_FORMAT, PERCENT_FORMAT],
        ["Total", report["classification_total"], 1.0 if report["classification_total"] else 0.0],
        lambda h, f, l, anchor: _doughnut(ws, "Leases by Classification (ASC 842)", h, f, l, [CLASSIFICATION_HEX.get(i["label"], "9AA5B1") for i in report["classification"]], anchor),
    )

    def maturity_chart(header_at, first, last, anchor):
        chart = BarChart()
        chart.type, chart.title, chart.height, chart.width, chart.legend = "col", "Lease Liability Maturity (Next 5 Years)", 7.0, 14.5, None
        chart.add_data(Reference(ws, min_col=3, min_row=header_at, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        chart.series[0].graphicalProperties.solidFill = COLORS["blue"]
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.dataLabels.showSerName = chart.dataLabels.showCatName = chart.dataLabels.showLegendKey = False
        chart.dataLabels.numFmt = number_format(report, with_symbol=False, decimals=0)
        small = CharacterProperties(sz=700)  # 7 pt: the labels stay clear of each other
        chart.dataLabels.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=small), endParaRPr=small)])
        chart.dataLabels.position = "outEnd"
        chart.y_axis.number_format, chart.y_axis.scaling.min = number_format(report, with_symbol=False, decimals=0), 0
        chart.gapWidth = 60
        _style_axes(chart)
        ws.add_chart(chart, anchor)

    block(
        "Lease Liability Maturity (Next 5 Years) - undiscounted payments",
        ["Year", "Period", "Payments"],
        [[b["label"], b["range"], b["value"]] for b in report["maturity"]],
        [None, None, money],
        ["Total", "", report["maturity_total"]],
        maturity_chart,
    )
    _print_setup(ws, report)


def _monthly_sheet(wb, report: dict) -> None:
    ws = wb.create_sheet("Interest vs Principal")
    _widths(ws, {"A": 14, "B": 22, "C": 22, "D": 22, "E": 3, **{letter: 11 for letter in "FGHIJKLM"}})
    _title(ws, "Interest vs. Principal - Next 60 Months (Portfolio)", "{}  |  amounts in {}".format(report["month_label"], report["currency_label"]), 13)
    ws.merge_cells("A3:D3")
    ws["A3"] = report["crossover_text"]
    ws["A3"].font, ws["A3"].alignment = _font(italic=True, color=COLORS["muted"]), Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 30
    money = number_format(report)
    monthly = report["monthly"]
    rows = [[m, float(i), float(p), float(c)] for m, i, p, c in zip(monthly["month"], monthly["interest"], monthly["principal"], monthly["cash"])]
    totals = report["monthly_totals"]
    last = _table(
        ws, 5, 1, ["Month", "Interest expense", "Principal repayment", "Total payment"], rows,
        [MONTH_FORMAT, money, money, money], ["Total (60 months)", totals["interest"], totals["principal"], totals["cash"]],
    )
    for line in range(6, 6 + len(rows)):
        ws.cell(row=line, column=1).alignment = Alignment(horizontal="left")
    ws.freeze_panes = "A6"
    chart = LineChart()
    chart.title, chart.height, chart.width = "Interest vs. Principal - Next 60 Months", 9.5, 19.0
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=5, max_row=5 + len(rows)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=6, max_row=5 + len(rows)))
    for series, color in zip(chart.series, (COLORS["interest"], COLORS["principal"])):
        series.graphicalProperties.line.solidFill, series.graphicalProperties.line.width, series.smooth = color, 28575, False
    chart.x_axis.number_format, chart.x_axis.tickLblSkip, chart.x_axis.tickMarkSkip = MONTH_FORMAT, 6, 6
    chart.y_axis.number_format, chart.y_axis.scaling.min = number_format(report, with_symbol=False, decimals=0), 0
    chart.legend.position = "t"
    _style_axes(chart)
    ws.add_chart(chart, "F5")
    _print_setup(ws, report, title_rows="5:5")


def _payments_sheet(wb, report: dict) -> None:
    ws = wb.create_sheet("Future Payments")
    _widths(ws, {"A": 20, "B": 22, "C": 22, "D": 22, "E": 16, "F": 16})
    _title(ws, "Future Payments Overview", "Payments over the next 12, 24 and 60 months  |  amounts in {}".format(report["currency_label"]), 6)
    money = number_format(report)
    rows = [[p["label"], p["total"], p["interest"], p["principal"], p["interest_share"], p["principal_share"]] for p in report["payments"]]
    last = _table(
        ws, 4, 1, ["Period", "Total payments", "Interest component", "Principal component", "Interest share", "Principal share"],
        rows, [None, money, money, money, PERCENT_FORMAT, PERCENT_FORMAT],
    )
    ws.cell(row=last + 1, column=1, value="Total payments = interest component + principal component in every period.").font = _font(italic=True, color=COLORS["muted"])
    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
    chart.title, chart.height, chart.width = "Interest and principal by period", 8.5, 16.0
    chart.add_data(Reference(ws, min_col=3, max_col=4, min_row=4, max_row=4 + len(rows)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(rows)))
    for series, color in zip(chart.series, (COLORS["interest"], COLORS["principal"])):
        series.graphicalProperties.solidFill = color
    chart.y_axis.number_format, chart.y_axis.scaling.min = number_format(report, with_symbol=False, decimals=0), 0
    chart.legend.position = "b"
    _style_axes(chart)
    ws.add_chart(chart, "A{}".format(last + 3))
    _print_setup(ws, report)


def _notes_sheet(wb, report: dict) -> None:
    ws = wb.create_sheet("Notes")
    _widths(ws, {"A": 6, "B": 120})
    _title(ws, "Notes and definitions", report["title"], 2)
    for number, text in enumerate(report["notes"], start=1):
        ws.cell(row=3 + number, column=1, value=number).font = _font(bold=True, color=COLORS["navy"])
        cell = ws.cell(row=3 + number, column=2, value=text)
        cell.font, cell.alignment = _font(), Alignment(wrap_text=True, vertical="top")
        ws.cell(row=3 + number, column=1).alignment = Alignment(horizontal="center", vertical="top")
        ws.row_dimensions[3 + number].height = 16 * max(1, math.ceil(len(text) / 115))
    _print_setup(ws, report)


def build_workbook(report: dict) -> Workbook:
    wb = Workbook()
    _dashboard_sheet(wb, report)
    _monthly_sheet(wb, report)
    _payments_sheet(wb, report)
    _notes_sheet(wb, report)
    wb.properties.title, wb.properties.creator = report["title"], "LeaseIQ Pro"
    wb.properties.subject = "{} - {}".format(report["currency_label"], report["month_label"])
    wb.properties.created = report["generated_at"] if isinstance(report["generated_at"], datetime) else datetime.now()
    return wb


def build_excel(report: dict) -> bytes:
    """The formatted workbook as bytes, ready to be downloaded."""
    buffer = BytesIO()
    build_workbook(report).save(buffer)
    return buffer.getvalue()


def excel_file_name(report: dict) -> str:
    return file_stem(report) + ".xlsx"
