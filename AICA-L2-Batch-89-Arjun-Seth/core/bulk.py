# Copyright © 2026 Arjun Seth. All rights reserved. Viewable for ICAI evaluation only.
# No permission to copy, modify, redistribute or use commercially without written permission.
"""Bulk import of leases from an Excel / CSV file.

Pure Python + pandas + openpyxl: no UI framework, no database, no AI. The rows are already
structured data, so they skip the AI extraction and go through the SAME rules as the
validation screen (``core/validation.py``) - a lease imported in bulk is held to exactly
the same standard as one validated by hand.

* ``build_template_bytes()``  the downloadable .xlsx template
* ``analyse_bulk_file()``     read + validate every row, BEFORE anything is saved
* ``valid_frame()`` / ``error_frame()``  the two preview tables
"""
import io
import os
import re
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core import validation as rules
from core.formatting import CURRENCIES, CURRENCY_CODES, DEFAULT_CURRENCY, parse_currency
from core.extraction import _normalise_value  # the same cell-cleaning rules the AI path uses

MAX_BULK_ROWS = 200  # keeps one import to a sensible time
FRAMEWORK_CHOICES = ("IND_AS_116", "ASC_842", "BOTH")
OVERRIDE_CHOICES = ("FINANCE LEASE", "OPERATING LEASE")
TEMPLATE_COLUMNS = [key for _, fields in rules.REVIEW_SECTIONS for key, _, _ in fields] + [
    "framework",
    "classification_override",
]
# A lease with no parties is not a usable record, so lessor and lessee are required here too.
REQUIRED_COLUMNS = ("lessor", "lessee") + rules.REQUIRED_FIELDS
DATA_SHEET = "Leases"


class BulkFileError(Exception):
    """The whole file cannot be used (unreadable, wrong columns, too big). Message is safe to show."""


# --------------------------------------------------------------------------- #
# The template
# --------------------------------------------------------------------------- #
_KIND_FORMAT = {
    "date": "YYYY-MM-DD (or a real Excel date)",
    "int": "whole number",
    "amount": "number, e.g. 100000 (no symbols needed)",
    "percent": "percent, e.g. 5% (or 0.05)",
    "yn": "Y or N",
    "text": "text",
    "longtext": "text",
    "currency": "3-letter code such as INR, USD, GBP (blank = your default currency)",
}
_EXAMPLE_ROW = {
    "lessor": "Sunrise Realty Private Limited",
    "lessee": "Bluewave Retail Private Limited",
    "asset_type": "Office space",
    "currency": "INR",
    "commencement_date": date(2026, 5, 1),
    "lease_end_date": date(2031, 4, 30),
    "lease_term_months": 60,
    "base_rent": 100000,
    "escalation": 0.05,
    "prepaid_rent": 200000,
    "prepaid_rent_months": 2,
    "deposit": 300000,
    "idc": 150000,
    "incentives": 0,
    "restoration_cost": 50000,
    "ibr": 0.09,
    "asset_fair_value": 10000000,
    "asset_economic_life_months": 480,
    "ownership_transfers": "N",
    "bargain_purchase_option": "N",
    "specialized_asset": "N",
    "low_value_election": "N",
    "options": "Renewal for 3 years by mutual agreement",
    "cpi_details": None,
    "framework": "BOTH",
    "classification_override": None,
}


def _column_notes() -> dict:
    """{column: (meaning, format, required?)} for the Instructions sheet."""
    notes = {}
    for key in TEMPLATE_COLUMNS:
        if key == "framework":
            notes[key] = ("Which framework(s) to report under", "IND_AS_116, ASC_842 or BOTH (blank = BOTH)", "No")
        elif key == "classification_override":
            notes[key] = (
                "Force the ASC 842 classification instead of using the five tests",
                "FINANCE LEASE or OPERATING LEASE (blank = let the tests decide)",
                "No",
            )
        else:
            kind = rules.FIELD_KIND[key]
            if key in REQUIRED_COLUMNS:
                required = "YES"
            elif key in rules.DEFAULT_ZERO_FIELDS:
                required = "No (blank = 0)"
            else:
                required = "No"
            notes[key] = (rules.FIELD_LABEL[key], _KIND_FORMAT[kind], required)
    return notes


def build_template_bytes() -> bytes:
    """The .xlsx template: a 'Leases' sheet to fill in, an 'Instructions' sheet and an 'Example' sheet.

    The importer reads ONLY the 'Leases' sheet, so the example can never be imported by accident.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DATA_SHEET

    required_fill = PatternFill("solid", fgColor="1F3A93")
    optional_fill = PatternFill("solid", fgColor="7F8C9A")
    header_font = Font(bold=True, color="FFFFFF")
    last_row = MAX_BULK_ROWS + 1
    number_formats = {"date": "yyyy-mm-dd", "percent": "0.00%", "amount": "#,##0.00", "int": "0"}

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=column)
        cell.font = header_font
        cell.fill = required_fill if column in REQUIRED_COLUMNS else optional_fill
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[get_column_letter(index)].width = max(16, len(column) + 4)
        number_format = number_formats.get(rules.FIELD_KIND.get(column, ""))
        if number_format:
            for row in range(2, last_row + 1):
                sheet.cell(row=row, column=index).number_format = number_format
    sheet.freeze_panes = "A2"

    def add_list(column: str, choices: str) -> None:
        letter = get_column_letter(TEMPLATE_COLUMNS.index(column) + 1)
        validation = DataValidation(type="list", formula1='"{}"'.format(choices), allow_blank=True)
        validation.add("{0}2:{0}{1}".format(letter, last_row))
        sheet.add_data_validation(validation)

    for column in (k for k, kind in rules.FIELD_KIND.items() if kind == "yn"):
        add_list(column, "Y,N")
    add_list("currency", ",".join(CURRENCY_CODES))
    add_list("framework", ",".join(FRAMEWORK_CHOICES))
    add_list("classification_override", ",".join(OVERRIDE_CHOICES))

    guide = workbook.create_sheet("Instructions")
    guide.append(["Column", "Meaning", "Format", "Required?"])
    for cell in guide[1]:
        cell.font = header_font
        cell.fill = required_fill
    for column, (meaning, format_text, required) in _column_notes().items():
        guide.append([column, meaning, format_text, required])
    guide.append([])
    for line in (
        "Fill in one lease per row on the 'Leases' sheet. Do not rename or delete the column headings.",
        "Dark blue headings are required. Blank optional amounts are treated as 0.",
        "Dates: use YYYY-MM-DD (e.g. 2026-05-01) or real Excel dates. Slash dates are read day-first (01/05/2026 = 1 May 2026).",
        "Rates (escalation, ibr): type 5% or 0.05. A plain number of 1 or more is read as a percentage.",
        "The 'Example' sheet shows one complete row. It is never imported.",
        "Maximum {} leases per file.".format(MAX_BULK_ROWS),
        "Currency: pick a code from the dropdown; a blank currency uses your default currency (Settings page). "
        "Amounts are never converted between currencies.",
    ):
        guide.append([line])
    guide.append([])
    guide.append(["Supported currencies"])
    guide.cell(row=guide.max_row, column=1).font = Font(bold=True)
    for code, (name, symbol, _) in CURRENCIES.items():
        guide.append([code, "{} ({})".format(name, symbol) if symbol != code else name])
    for letter, width in (("A", 28), ("B", 62), ("C", 48), ("D", 18)):
        guide.column_dimensions[letter].width = width

    example = workbook.create_sheet("Example")
    example.append(TEMPLATE_COLUMNS)
    for cell in example[1]:
        cell.font = header_font
        cell.fill = optional_fill
    example.append([_EXAMPLE_ROW.get(column) for column in TEMPLATE_COLUMNS])
    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        example.column_dimensions[get_column_letter(index)].width = max(16, len(column) + 4)
        number_format = number_formats.get(rules.FIELD_KIND.get(column, ""))
        if number_format:
            example.cell(row=2, column=index).number_format = number_format

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Reading the file
# --------------------------------------------------------------------------- #
def _is_blank(value) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(value, str) and not value.strip()


def _normalise_header(name) -> str:
    return re.sub(r"\s+", "_", str(name).strip().rstrip("*").strip().lower())


def read_bulk_file(filename: str, data: bytes) -> pd.DataFrame:
    """Read an .xlsx (sheet 'Leases', else the first sheet) or .csv into a table of raw cells."""
    if not data:
        raise BulkFileError("The file is empty.")
    extension = os.path.splitext(filename or "")[1].lower()
    read_options = {"dtype": object, "keep_default_na": False, "na_values": [""]}
    try:
        if extension == ".csv":
            try:
                frame = pd.read_csv(io.BytesIO(data), encoding="utf-8-sig", **read_options)
            except UnicodeDecodeError:
                frame = pd.read_csv(io.BytesIO(data), encoding="latin-1", **read_options)
        elif extension == ".xlsx":
            excel = pd.ExcelFile(io.BytesIO(data))
            sheet = DATA_SHEET if DATA_SHEET in excel.sheet_names else excel.sheet_names[0]
            frame = pd.read_excel(excel, sheet_name=sheet, **read_options)
        else:
            raise BulkFileError("Please upload an .xlsx or .csv file (got '{}').".format(extension or filename))
    except BulkFileError:
        raise
    except pd.errors.EmptyDataError:
        raise BulkFileError("The file has no data.") from None
    except Exception as exc:
        raise BulkFileError("Could not read the file ({}). Please use the template.".format(type(exc).__name__)) from None

    raw_headings = [str(column) for column in frame.columns]
    normalised = [_normalise_header(column) for column in raw_headings]
    duplicated = {name for name in normalised if normalised.count(name) > 1}
    for heading in raw_headings:  # pandas silently renames a repeated heading to 'lessor.1'
        mangled = re.fullmatch(r"(.+)\.(\d+)", heading)
        if mangled and _normalise_header(mangled.group(1)) in normalised:
            duplicated.add(_normalise_header(mangled.group(1)))
    frame.columns = normalised
    duplicated = sorted(duplicated)
    if duplicated:
        raise BulkFileError("These column headings appear more than once: {}.".format(", ".join(duplicated)))
    return frame


# --------------------------------------------------------------------------- #
# Checking one row
# --------------------------------------------------------------------------- #
def normalise_cell(key: str, raw):
    """Return ``(value, problem, note)`` for one cell. ``problem`` is set when it cannot be used."""
    if _is_blank(raw):
        return None, None, None
    kind = rules.FIELD_KIND[key]
    if kind == "currency":
        code = parse_currency(str(raw))
        if code is None:
            return None, "'{}' is not a supported currency (use a code such as INR, USD, GBP, EUR)".format(str(raw).strip()), None
        return code, None, None
    if kind == "date":
        if isinstance(raw, datetime):  # also pandas Timestamp
            return raw.date(), None, None
        if isinstance(raw, date):
            return raw, None, None
        value, note = _normalise_value("date", raw)
        return (date.fromisoformat(value), None, None) if value else (None, note, None)
    if isinstance(raw, (datetime, date)):
        return None, "expected a {} but found a date".format(_KIND_FORMAT[kind].split(",")[0].split(" (")[0]), None
    value, note = _normalise_value({"percent": "fraction", "longtext": "text"}.get(kind, kind), raw)
    if value is None:
        return None, note or "could not be read", None
    return value, None, note  # a note (e.g. "Read '5' as a percentage") is informational only


def _parse_choice(raw, choices: dict, column: str):
    """Match a free-typed choice ('Ind AS 116', 'finance') against ``choices``; returns (value, problem)."""
    if _is_blank(raw):
        return None, None
    squashed = re.sub(r"[^a-z0-9]", "", str(raw).lower())
    if squashed in choices:
        return choices[squashed], None
    allowed = sorted(set(choices.values()))
    return None, "{}: '{}' is not one of {}".format(column, str(raw).strip(), ", ".join(allowed))


_FRAMEWORK_CHOICES = {
    "indas116": "IND_AS_116", "indas": "IND_AS_116", "asc842": "ASC_842", "asc": "ASC_842", "both": "BOTH",
}
_OVERRIDE_CHOICES = {
    "financelease": "FINANCE LEASE", "finance": "FINANCE LEASE",
    "operatinglease": "OPERATING LEASE", "operating": "OPERATING LEASE",
}


def _check_row(record: dict, row_number: int, default_currency: str = DEFAULT_CURRENCY):
    """Return ('valid', dict) or ('error', dict) for one spreadsheet row."""
    values, problems, warnings = {}, [], []
    for key in rules.FIELD_KIND:
        value, problem, note = normalise_cell(key, record.get(key))
        values[key] = value
        if problem:
            problems.append("{}: {}".format(key, problem))
        elif note:
            warnings.append("{}: {}".format(key, note))
    if values.get("currency") is None:
        values["currency"] = default_currency  # a blank currency means the user's default currency
    framework, problem = _parse_choice(record.get("framework"), _FRAMEWORK_CHOICES, "framework")
    if problem:
        problems.append(problem)
    override, problem = _parse_choice(record.get("classification_override"), _OVERRIDE_CHOICES, "classification_override")
    if problem:
        problems.append(problem)

    already_flagged = {p.split(":", 1)[0] for p in problems}
    for key in REQUIRED_COLUMNS:
        if values.get(key) is None and key not in already_flagged:
            problems.append("{}: required but blank".format(key))
    if not problems:  # only judge the logic once every cell is readable, to avoid noisy follow-on errors
        errors, logic_warnings = rules.validate_values(values)
        problems.extend(errors)
        warnings.extend(logic_warnings)

    if problems:
        return "error", {
            "row": row_number,
            "lessor": "" if _is_blank(record.get("lessor")) else str(record.get("lessor")).strip(),
            "lessee": "" if _is_blank(record.get("lessee")) else str(record.get("lessee")).strip(),
            "problems": problems,
        }
    return "valid", {
        "row": row_number,
        "values": values,
        "framework": framework or "BOTH",
        "override": override,
        "warnings": warnings,
    }


def analyse_bulk_file(filename: str, data: bytes, default_currency: str = DEFAULT_CURRENCY) -> dict:
    """Read the file and check EVERY row without saving anything.

    Returns ``{"total_rows", "valid": [...], "errors": [...], "ignored_columns": [...]}``.
    A blank currency becomes ``default_currency``. Raises BulkFileError if the file as a whole cannot be used. Spreadsheet row numbers
    count the heading as row 1, so the first lease is row 2.
    """
    default_code = parse_currency(default_currency)
    if default_code is None:
        raise ValueError("default_currency must be a supported currency code, got {!r}".format(default_currency))
    frame = read_bulk_file(filename, data)
    columns = list(frame.columns)
    missing = [column for column in REQUIRED_COLUMNS if column not in columns]
    if missing:
        raise BulkFileError(
            "The file is missing required column(s): {}. Please start from the template.".format(", ".join(missing))
        )
    ignored = [column for column in columns if column not in TEMPLATE_COLUMNS]

    records = [
        (position + 2, record)
        for position, record in enumerate(frame.to_dict("records"))
        if not all(_is_blank(value) for value in record.values())
    ]
    if len(records) > MAX_BULK_ROWS:
        raise BulkFileError(
            "The file has {} leases; the limit is {} per file. Please split it.".format(len(records), MAX_BULK_ROWS)
        )
    if not records:
        raise BulkFileError("The file has no lease rows (only the headings).")

    valid, errors = [], []
    for row_number, record in records:
        outcome, item = _check_row(record, row_number, default_code)
        (valid if outcome == "valid" else errors).append(item)
    return {"total_rows": len(records), "valid": valid, "errors": errors, "ignored_columns": ignored}


# --------------------------------------------------------------------------- #
# The two preview tables
# --------------------------------------------------------------------------- #
def valid_frame(report: dict) -> pd.DataFrame:
    rows = []
    for item in report["valid"]:
        values = item["values"]
        rows.append(
            {
                "Row": item["row"],
                "Lessor": values["lessor"],
                "Lessee": values["lessee"],
                "Currency": values["currency"],
                "Commencement": values["commencement_date"].isoformat(),
                "Term (months)": values["lease_term_months"],
                "Monthly rent": values["base_rent"],
                "Framework": item["framework"],
                "Override": item["override"] or "",
                "Warnings": "; ".join(item["warnings"]),
            }
        )
    return pd.DataFrame(
        rows,
        columns=[
            "Row", "Lessor", "Lessee", "Currency", "Commencement", "Term (months)", "Monthly rent", "Framework", "Override", "Warnings",
        ],
    )


def error_frame(report: dict) -> pd.DataFrame:
    rows = [
        {"Row": item["row"], "Lessor": item["lessor"], "Lessee": item["lessee"], "Problems": "; ".join(item["problems"])}
        for item in report["errors"]
    ]
    return pd.DataFrame(rows, columns=["Row", "Lessor", "Lessee", "Problems"])
