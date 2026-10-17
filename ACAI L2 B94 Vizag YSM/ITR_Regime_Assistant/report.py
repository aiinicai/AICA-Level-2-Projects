"""Excel working paper; deterministic values exported as numbers, text as literals."""
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from engine import calculate,scenarios,narrative

def export_report(p,path,reviewer='',decision='',notes='',draft=''):
 r=calculate(p)
 if r['blocked']: raise ValueError('Resolve scope blockers before exporting a tax working paper.')
 w=Workbook(); ws=w.active; ws.title='Comparison'
 ws.append(['ITR preparation working paper','AY 2026-27 / FY 2025-26'])
 ws.append(['Case',p['case']]); ws.append(['Numeric preference',r['winner']]); ws.append(['Difference',float(r['saving'])]); ws.append(['Status',r['status']])
 ws.append(['Particular','Old','New'])
 for k,v in r['Old'].items():
  if k not in ('bands','claims'): ws.append([k,float(v),float(r['New'][k])])
 ws=w.create_sheet('Slabs')
 ws.append(['Regime','From','To (blank = above)','Rate','Amount in band','Tax'])
 for regime in ('Old','New'):
  for b in r[regime]['bands']: ws.append([regime,float(b['lower']),b['upper'],b['rate'],float(b['taxable']),float(b['tax'])])
 ws=w.create_sheet('Claims'); ws.append(['Old regime bucket','Capped eligible amount before aggregate GTI cap'])
 for k,v in r['Old']['claims'].items(): ws.append([k,float(v)])
 ws=w.create_sheet('Scenarios'); rows=scenarios(p); ws.append(list(rows[0]))
 for row in rows: ws.append([float(v) if hasattr(v,'as_tuple') else v for v in row.values()])
 ws=w.create_sheet('Inputs'); ws.append(['Field','Value'])
 for k,v in p.items(): ws.append([k,v])
 ws=w.create_sheet('Review'); ws.append(['Reviewer',reviewer]); ws.append(['Chosen regime',decision]); ws.append(['Reasons / evidence references',notes]); ws.append(['Form candidate',r['form']])
 ws.append(['Not a portal return','Interest/fees, filing eligibility and e-verification require separate review'])
 for item in r['readiness']+r['warnings']: ws.append(['Pending / warning',item])
 ws=w.create_sheet('AI draft'); ws.append(['UNAPPROVED AI / external draft',draft or 'Not generated'])
 ws=w.create_sheet('Notes')
 for line in narrative(p,r).splitlines(): ws.append([line])
 for ws in w:
  ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='17324D')
  for row in ws:
   for c in row:
    if isinstance(c.value,str): c.data_type='s'
  for col in ws.columns:
   ws.column_dimensions[col[0].column_letter].width=min(95,max(18,max(len(str(c.value or '')) for c in col)+2))
 w.save(path)
