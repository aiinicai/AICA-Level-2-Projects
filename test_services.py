from datetime import date
from io import BytesIO
import base64
import re
from unittest.mock import Mock
import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from services.demo_service import SAMPLE_TEXT, sample_analysis, sample_draft
from services.document_service import generate_docx
from services.llm_service import LLMService, consolidate, split_notice, verify_evidence
from services.models import Draft, Evidence
from services.pdf_service import extract_text_from_pdf
from utils.helpers import clipboard_html, deadline_label, safe_filename
from utils.validators import AppError, MAX_UPLOAD_BYTES, validate_pdf


def pdf_bytes(pages: list[str]) -> bytes:
    output = BytesIO()
    pdf = canvas.Canvas(output)
    for content in pages:
        text = pdf.beginText(40, 790)
        for line in content.splitlines():
            text.textLine(line)
        pdf.drawText(text)
        pdf.showPage()
    pdf.save()
    return output.getvalue()


def test_searchable_pdf_does_not_ocr(monkeypatch):
    mock = Mock(side_effect=AssertionError("OCR must not run"))
    monkeypatch.setattr("services.pdf_service.ocr_page", mock)
    result = extract_text_from_pdf(pdf_bytes([SAMPLE_TEXT, SAMPLE_TEXT]))
    assert result.page_count == 2 and result.status == "text"
    assert result.text.count("--- Page") == 2
    assert result.text.index("--- Page 1") < result.text.index("--- Page 2")
    mock.assert_not_called()


def test_mixed_pdf_selectively_reads_scanned_page(monkeypatch):
    mock = Mock(return_value="Scanned page contains an additional request. " * 5)
    monkeypatch.setattr("services.pdf_service.ocr_page", mock)
    result = extract_text_from_pdf(pdf_bytes([SAMPLE_TEXT, "", SAMPLE_TEXT]))
    assert result.ocr_pages == (2,)
    assert result.page_count == 3 and "additional request" in result.text
    assert mock.call_args.args[1] == 1
    assert mock.call_count == 1


def test_unreadable_page_is_flagged(monkeypatch):
    monkeypatch.setattr("services.pdf_service.ocr_page", lambda *_: "")
    result = extract_text_from_pdf(pdf_bytes([SAMPLE_TEXT, ""]))
    assert "Page 2" in result.warnings[0]
    with pytest.raises(AppError, match="reliably read"):
        extract_text_from_pdf(pdf_bytes([""]))


@pytest.mark.parametrize("password", ["secret", ""])
def test_encrypted_pdf_rejected(password):
    writer = PdfWriter()
    writer.append(PdfReader(BytesIO(pdf_bytes([SAMPLE_TEXT]))))
    writer.encrypt(password)
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(AppError, match="encrypted|password"):
        extract_text_from_pdf(stream.getvalue())


@pytest.mark.parametrize("data,name,mime", [(b"", "a.pdf", None), (b"%PDF-no", "a.txt", None), (b"not a PDF", "a.pdf", None), (b"%PDF-no", "a.pdf", "image/png"), (b"%PDF-" + b"x" * MAX_UPLOAD_BYTES, "a.pdf", None)], ids=["empty", "extension", "signature", "mime", "oversized"])
def test_invalid_uploads(data, name, mime):
    with pytest.raises(AppError):
        validate_pdf(data, name, mime)


def test_corrupt_pdf_is_friendly():
    with pytest.raises(AppError, match="Unable to read PDF"):
        extract_text_from_pdf(b"%PDF-1.7\nbroken")


def test_verified_demo_evidence_and_fabricated_section():
    assert verify_evidence(sample_analysis(), SAMPLE_TEXT).response_deadline.exact_date == date(2026, 9, 27)
    analysis = sample_analysis()
    analysis.sections_mentioned.append(Evidence(text="Section 999", source_quote="Section 999"))
    with pytest.raises(AppError, match="section"):
        verify_evidence(analysis, SAMPLE_TEXT)


def test_receipt_deadline_never_infers_date():
    analysis = sample_analysis()
    quote = "Submit within 15 days of receipt. Notice dated 27 September 2026."
    analysis.response_deadline.source_quote = quote
    analysis.response_deadline.deadline_text = quote
    result = verify_evidence(analysis, SAMPLE_TEXT + quote)
    assert result.response_deadline.exact_date is None


def test_conflicting_deadlines_preserve_both():
    first, second = sample_analysis(), sample_analysis()
    second.response_deadline.exact_date = date(2026, 10, 1)
    second.response_deadline.deadline_text = "Respond by 1 October 2026."
    merged = consolidate([first, second])
    assert merged.response_deadline.exact_date is None
    assert "27 September" in merged.response_deadline.deadline_text
    assert "1 October" in merged.response_deadline.deadline_text
    assert len(merged.issues) == 2


def test_notice_issue_date_is_not_a_reply_deadline():
    analysis = sample_analysis()
    analysis.response_deadline.exact_date = date(2026, 9, 12)
    analysis.response_deadline.source_quote = "Notice dated 12 September 2026"
    analysis.response_deadline.deadline_text = "Notice dated 12 September 2026"
    assert verify_evidence(analysis, SAMPLE_TEXT).response_deadline.exact_date is None


def test_chunks_cover_all_content_including_last_page():
    text = "\n".join(f"line {i:06}: more notice content and deadlines" for i in range(3000))
    chunks = split_notice(text)
    for line in text.splitlines():
        assert any(line in chunk for chunk in chunks)
    assert chunks[-1].endswith(text.splitlines()[-1])


def test_edited_reply_is_the_word_content():
    reply = sample_draft() + "\nUSER EDIT: Verify ₹1,23,456."
    document = Document(BytesIO(generate_docx(reply)))
    assert "\n".join(p.text for p in document.paragraphs) == reply
    assert next(p for p in document.paragraphs if p.text.startswith("Subject:")).runs[0].bold
    assert document.sections[0].left_margin.inches == pytest.approx(0.9)


def test_empty_word_rejected():
    with pytest.raises(AppError):
        generate_docx(" ")


def test_clipboard_does_not_embed_executable_document_content():
    malicious = '</script><img src=x onerror="alert(1)"> ₹'
    result = clipboard_html(malicious)
    assert malicious not in result
    payload = re.search(r'atob\("([^"]+)"\)', result).group(1)
    assert base64.b64decode(payload).decode() == malicious


def test_days_remaining_and_filenames():
    assert deadline_label(date(2026, 9, 27), date(2026, 9, 20))[0] == "7 days remaining"
    assert deadline_label(date(2026, 9, 20), date(2026, 9, 20))[0] == "Due today"
    assert deadline_label(date(2026, 9, 19), date(2026, 9, 20)) == ("1 day overdue", "error")
    assert "/" not in safe_filename("../../2025-26")


def test_structured_api_contract(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test-placeholder")
    monkeypatch.setenv("OPENAI_MODEL", "test-model")
    provider = Mock()
    provider.chat.completions.parse.return_value.choices = [Mock(message=Mock(parsed=sample_analysis(), refusal=None))]
    monkeypatch.setattr("services.llm_service.OpenAI", Mock(return_value=provider))
    service = LLMService()
    analysis = service.analyze(SAMPLE_TEXT)
    assert analysis.is_tax_notice
    args = provider.chat.completions.parse.call_args.kwargs
    assert args["store"] is False
    assert args["model"] == "test-model"
    provider.chat.completions.parse.return_value.choices = [Mock(message=Mock(parsed=Draft(draft_reply="Reviewed reply"), refusal=None))]
    assert service.draft(analysis) == "Reviewed reply"


def test_provider_refusal_and_failure_are_friendly(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test-placeholder")
    monkeypatch.setenv("OPENAI_MODEL", "test-model")
    provider = Mock()
    provider.chat.completions.parse.return_value.choices = [Mock(message=Mock(parsed=None, refusal="private provider detail"))]
    monkeypatch.setattr("services.llm_service.OpenAI", Mock(return_value=provider))
    with pytest.raises(AppError, match="complete analysis"):
        LLMService().analyze(SAMPLE_TEXT)
    provider.chat.completions.parse.side_effect = ValueError("private raw response")
    with pytest.raises(AppError) as exc:
        LLMService().analyze(SAMPLE_TEXT)
    assert "private" not in str(exc.value)


def test_non_tax_document_rejected(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test-placeholder")
    monkeypatch.setenv("OPENAI_MODEL", "test-model")
    monkeypatch.setattr("services.llm_service.OpenAI", Mock())
    analysis = sample_analysis()
    analysis.is_tax_notice = False
    monkeypatch.setattr(LLMService, "_request", lambda *_: analysis)
    with pytest.raises(AppError, match="does not appear"):
        LLMService().analyze(SAMPLE_TEXT)
