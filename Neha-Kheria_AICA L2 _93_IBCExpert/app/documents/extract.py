"""Local text extraction and OCR. No cloud or network code exists here."""
from __future__ import annotations

import csv
import html
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

from app.core.errors import DocumentError


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    method: str
    ocr_status: str
    confidence: float | None = None


def _decode_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentError("Text encoding could not be recognized.")


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs = []
    for paragraph in root.findall(".//w:p", namespaces):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespaces))
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs)


def _xlsx_text(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentError("XLSX extraction requires the bundled openpyxl component.") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    lines.append("\t".join(values))
    finally:
        workbook.close()
    return "\n".join(lines)


def _pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentError("PDF extraction requires the bundled pypdf component.") from exc
    try:
        reader = PdfReader(str(path), strict=False)
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:
        raise DocumentError("PDF is corrupt, encrypted without a password, or unsupported.") from exc


def run_image_ocr(path: Path, command: str = "tesseract", language: str = "eng") -> ExtractionResult:
    executable = shutil.which(command)
    if not executable:
        return ExtractionResult("", "ocr-unavailable", "UNAVAILABLE")
    try:
        process = subprocess.run(
            [executable, str(path), "stdout", "-l", language],
            capture_output=True, text=True, timeout=300, check=False,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise DocumentError("Local OCR engine failed or timed out.") from exc
    if process.returncode != 0:
        message = (process.stderr or "unknown OCR error").strip()[-500:]
        raise DocumentError(f"Local OCR engine could not process the document: {message}")
    return ExtractionResult(process.stdout.strip(), "tesseract-ocr", "COMPLETE")


def _ocr_pdf(path: Path, command: str, language: str) -> ExtractionResult:
    try:
        import pymupdf as fitz
    except ImportError:
        return ExtractionResult("", "pdf-ocr-unavailable", "UNAVAILABLE")
    pages: list[str] = []
    with tempfile.TemporaryDirectory(prefix="ibc-ocr-") as temp:
        try:
            document = fitz.open(path)
            for number, page in enumerate(document):
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Path(temp) / f"page-{number + 1}.png"
                pixmap.save(image)
                result = run_image_ocr(image, command, language)
                if result.ocr_status != "COMPLETE":
                    return result
                pages.append(result.text)
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError("Scanned PDF could not be rendered for local OCR.") from exc
    return ExtractionResult("\n\n".join(pages), "pdf-tesseract-ocr", "COMPLETE")


def extract_text(path: Path, media_type: str, *, ocr_command: str = "tesseract", ocr_language: str = "eng", force_ocr: bool = False) -> ExtractionResult:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".csv"}:
        return ExtractionResult(_decode_text(path), "text-decoder", "NOT_REQUIRED")
    if suffix in {".html", ".htm"}:
        raw = _decode_text(path)
        text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
        text = re.sub(r"(?s)<[^>]+>", " ", text)
        return ExtractionResult(html.unescape(re.sub(r"\s+", " ", text)).strip(), "html-text", "NOT_REQUIRED")
    if suffix == ".docx":
        return ExtractionResult(_docx_text(path), "docx-xml", "NOT_REQUIRED")
    if suffix == ".xlsx":
        return ExtractionResult(_xlsx_text(path), "openpyxl", "NOT_REQUIRED")
    if suffix == ".pdf":
        text = "" if force_ocr else _pdf_text(path)
        if force_ocr or len(re.sub(r"\s", "", text)) < 30:
            ocr = _ocr_pdf(path, ocr_command, ocr_language)
            if ocr.ocr_status == "COMPLETE":
                return ocr
            return ExtractionResult(text, "pypdf-low-text", ocr.ocr_status)
        return ExtractionResult(text, "pypdf", "NOT_REQUIRED")
    if media_type.startswith("image/"):
        return run_image_ocr(path, ocr_command, ocr_language)
    return ExtractionResult("", "unsupported-for-text", "NOT_REQUIRED")
