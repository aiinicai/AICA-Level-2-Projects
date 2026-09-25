"""Structured local-file import staging and manual review.

This module deliberately has no network client, URL loader, downloader, remote scheduler,
or automatic database updater. Every source is a user-selected local file. Structured rows
are staged for human review before they can create application/legal records.
"""
from __future__ import annotations

import csv
import json
import re
import sqlite3
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from app.core.errors import ValidationError
from app.core.time import SystemClock, to_utc_iso
from app.db.connection import transaction
from app.documents.vault import DocumentVault
from app.legal.service import LegalService
from app.plugins.manager import EventBus
from app.services.clients import ClientService

SUPPORTED_STRUCTURED = {".csv": "CSV", ".xlsx": "XLSX", ".docx": "DOCX", ".json": "JSON"}
TARGET_FIELDS: dict[str, tuple[str, ...]] = {
    "CLIENT": (
        "name", "corporate_debtor", "cin", "pan", "gst", "registered_office",
        "industry", "primary_email", "primary_phone", "notes",
    ),
    "LEGAL_STATUTE": (
        "title", "short_title", "statute_type", "jurisdiction", "enactment_date",
        "commencement_date", "identifier",
    ),
    "LEGAL_JUDGMENT": (
        "title", "neutral_citation", "reported_citation", "court", "bench",
        "judgment_date", "case_number", "parties", "text_content", "holding_summary",
    ),
}
REQUIRED_FIELDS = {
    "CLIENT": {"name"},
    "LEGAL_STATUTE": {"title", "statute_type"},
    "LEGAL_JUDGMENT": {"title", "court", "text_content"},
}

ALIASES = {
    "client name": "name", "client": "name", "name": "name",
    "corporate debtor": "corporate_debtor", "corporate_debtor": "corporate_debtor",
    "cin": "cin", "pan": "pan", "gst": "gst", "gstin": "gst",
    "registered office": "registered_office", "address": "registered_office",
    "industry": "industry", "email": "primary_email", "primary email": "primary_email",
    "phone": "primary_phone", "mobile": "primary_phone", "notes": "notes",
    "title": "title", "short title": "short_title", "statute type": "statute_type",
    "jurisdiction": "jurisdiction", "enactment date": "enactment_date",
    "commencement date": "commencement_date", "identifier": "identifier",
    "neutral citation": "neutral_citation", "reported citation": "reported_citation",
    "court": "court", "tribunal": "court", "bench": "bench", "judgment date": "judgment_date",
    "case number": "case_number", "parties": "parties", "text": "text_content",
    "text content": "text_content", "judgment text": "text_content", "holding summary": "holding_summary",
}


def _norm_header(value: Any) -> str:
    text = re.sub(r"[_\-]+", " ", str(value or "").strip().lower())
    return re.sub(r"\s+", " ", text)


def _clean_cell(value: Any) -> str | None:
    if value is None:
        return None
    if hasattr(value, "isoformat") and not isinstance(value, (str, bytes)):
        try:
            return value.isoformat()
        except Exception:
            pass
    text = str(value).strip()
    return text if text else None


def _records_csv(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    reader = csv.DictReader(raw.splitlines())
    if not reader.fieldnames:
        raise ValidationError("CSV file does not contain a header row.")
    headers = [str(h or "").strip() for h in reader.fieldnames]
    rows = [{str(k or "").strip(): _clean_cell(v) for k, v in row.items()} for row in reader]
    return headers, rows


def _records_xlsx(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValidationError("XLSX import requires the bundled openpyxl dependency.") from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        first = next(iterator, None)
        if not first:
            raise ValidationError("Workbook is empty.")
        headers = [str(v or "").strip() for v in first]
        if not any(headers):
            raise ValidationError("Workbook first row must contain column headings.")
        rows: list[dict[str, Any]] = []
        for values in iterator:
            row = {headers[i]: _clean_cell(values[i] if i < len(values) else None) for i in range(len(headers)) if headers[i]}
            if any(v is not None for v in row.values()):
                rows.append(row)
        return [h for h in headers if h], rows
    finally:
        try:
            workbook.close()
        except Exception:
            pass


def _records_docx(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ValidationError("DOCX import requires the bundled python-docx dependency.") from exc
    doc = Document(str(path))
    if doc.tables:
        table = doc.tables[0]
        if not table.rows:
            raise ValidationError("DOCX table is empty.")
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if not any(headers):
            raise ValidationError("DOCX table first row must contain column headings.")
        rows = []
        for row in table.rows[1:]:
            values = [cell.text.strip() or None for cell in row.cells]
            item = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers)) if headers[i]}
            if any(v is not None for v in item.values()):
                rows.append(item)
        return [h for h in headers if h], rows
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if not paragraphs:
        raise ValidationError("DOCX file contains no table or readable text.")
    return ["text_content"], [{"text_content": "\n\n".join(paragraphs)}]


def _records_json(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValidationError("JSON file is unreadable or invalid.") from exc
    if isinstance(data, dict) and isinstance(data.get("records"), list):
        data = data["records"]
    elif isinstance(data, dict):
        data = [data]
    if not isinstance(data, list) or not all(isinstance(item, dict) for item in data):
        raise ValidationError("JSON structured import must contain an object, a list of objects, or a records list.")
    headers: list[str] = []
    seen = set()
    rows = []
    for item in data:
        row = {}
        for key, value in item.items():
            key = str(key)
            if key not in seen:
                headers.append(key); seen.add(key)
            if isinstance(value, (dict, list)):
                row[key] = json.dumps(value, ensure_ascii=False, sort_keys=True)
            else:
                row[key] = _clean_cell(value)
        rows.append(row)
    if not headers:
        raise ValidationError("JSON file contains no fields.")
    return headers, rows


def read_structured_records(path: Path | str, *, max_rows: int = 5000) -> tuple[str, list[str], list[dict[str, Any]]]:
    source = Path(path)
    source_format = SUPPORTED_STRUCTURED.get(source.suffix.lower())
    if not source_format:
        raise ValidationError("Structured mapping supports CSV, XLSX, DOCX and JSON files.")
    loaders = {"CSV": _records_csv, "XLSX": _records_xlsx, "DOCX": _records_docx, "JSON": _records_json}
    headers, rows = loaders[source_format](source)
    if len(rows) > max_rows:
        raise ValidationError(f"Structured import is limited to {max_rows} rows per file for safe local review.")
    if not rows:
        raise ValidationError("Structured import contains no data rows.")
    return source_format, headers, rows


@dataclass
class StructuredImportService:
    connection: sqlite3.Connection
    vault: DocumentVault
    events: EventBus
    clients: ClientService
    legal: LegalService
    clock: SystemClock = SystemClock()

    def target_fields(self, target_type: str) -> tuple[str, ...]:
        if target_type not in TARGET_FIELDS:
            raise ValidationError("Structured import target is unsupported.")
        return TARGET_FIELDS[target_type]

    def suggest_mapping(self, headers: Iterable[str], target_type: str) -> dict[str, str]:
        allowed = set(self.target_fields(target_type))
        result = {}
        for header in headers:
            norm = _norm_header(header)
            candidate = ALIASES.get(norm)
            if candidate in allowed:
                result[str(header)] = candidate
            elif norm.replace(" ", "_") in allowed:
                result[str(header)] = norm.replace(" ", "_")
        return result

    def profiles(self, target_type: str | None = None) -> list[dict[str, Any]]:
        sql = "SELECT * FROM import_mapping_profiles"
        params: list[Any] = []
        if target_type:
            self.target_fields(target_type)
            sql += " WHERE target_type=?"; params.append(target_type)
        sql += " ORDER BY name COLLATE NOCASE"
        result = []
        for row in self.connection.execute(sql, params).fetchall():
            item = dict(row); item["mapping"] = json.loads(item.pop("mapping_json")); result.append(item)
        return result

    def save_profile(self, name: str, target_type: str, source_format: str, mapping: dict[str, str]) -> int:
        name = str(name or "").strip()
        if not name or len(name) > 200:
            raise ValidationError("Mapping profile name is required and must be under 200 characters.")
        self._validate_mapping(target_type, mapping)
        if source_format not in set(SUPPORTED_STRUCTURED.values()):
            raise ValidationError("Mapping profile source format is unsupported.")
        now = to_utc_iso(self.clock.now())
        encoded = json.dumps(mapping, ensure_ascii=False, sort_keys=True)
        with transaction(self.connection):
            self.connection.execute(
                """INSERT INTO import_mapping_profiles(profile_uuid,name,target_type,source_format,mapping_json,created_at,updated_at)
                   VALUES(?,?,?,?,?,?,?) ON CONFLICT(name,target_type,source_format) DO UPDATE SET
                   mapping_json=excluded.mapping_json,updated_at=excluded.updated_at""",
                (str(uuid.uuid4()), name, target_type, source_format, encoded, now, now),
            )
            row = self.connection.execute(
                "SELECT id FROM import_mapping_profiles WHERE name=? AND target_type=? AND source_format=?",
                (name, target_type, source_format),
            ).fetchone()
            return int(row["id"])

    def _validate_mapping(self, target_type: str, mapping: dict[str, str]) -> None:
        allowed = set(self.target_fields(target_type))
        if not isinstance(mapping, dict):
            raise ValidationError("Column mapping must be a JSON object.")
        invalid = {str(v) for v in mapping.values() if v not in allowed}
        if invalid:
            raise ValidationError("Mapping contains unsupported target field(s): " + ", ".join(sorted(invalid)))
        if len(set(mapping.values())) != len(mapping.values()):
            raise ValidationError("Each target field may be mapped from only one source column.")

    def _start_import(self, source: Path, import_type: str, total: int) -> int:
        import hashlib
        now = to_utc_iso(self.clock.now())
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        cur = self.connection.execute(
            """INSERT INTO imports(import_uuid,import_type,source_name,source_hash,status,total_items,imported_items,rejected_items,started_at)
               VALUES(?,?,?,?,?,?,?,?,?)""",
            (str(uuid.uuid4()), import_type, source.name, digest, "RUNNING", total, 0, 0, now),
        )
        return int(cur.lastrowid)

    def stage_structured(self, source: Path | str, *, target_type: str, mapping: dict[str, str] | None = None,
                         profile_name: str | None = None, user_id: int | None = None) -> dict[str, Any]:
        path = Path(source)
        if not path.is_file():
            raise ValidationError("Selected structured import source is not a local file.")
        source_format, headers, rows = read_structured_records(path)
        chosen = dict(mapping or self.suggest_mapping(headers, target_type))
        self._validate_mapping(target_type, chosen)
        required = REQUIRED_FIELDS[target_type]
        mapped_targets = set(chosen.values())
        now = to_utc_iso(self.clock.now())
        # Save the reusable mapping before staging rows. save_profile owns its own
        # transaction, so it must not be called inside another transaction.
        if profile_name:
            self.save_profile(profile_name, target_type, source_format, chosen)
        with transaction(self.connection):
            import_id = self._start_import(path, f"STRUCTURED_{source_format}_{target_type}", len(rows))
            for index, raw in enumerate(rows, start=2):
                payload = {target: _clean_cell(raw.get(source_header)) for source_header, target in chosen.items()}
                missing = sorted(field for field in required if not payload.get(field))
                reason = "Missing required mapped value(s): " + ", ".join(missing) if missing else "Review mapped values before creating the record."
                title = payload.get("name") or payload.get("title") or f"{target_type.replace('_', ' ').title()} row {index}"
                excerpt = json.dumps(raw, ensure_ascii=False)[:2000]
                self.connection.execute(
                    """INSERT INTO manual_review_queue(review_uuid,import_id,item_type,target_type,title,payload_json,source_excerpt,reason,status,created_at,updated_at)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (str(uuid.uuid4()), import_id, "STRUCTURED_ROW", target_type, str(title)[:500],
                     json.dumps(payload, ensure_ascii=False, sort_keys=True), excerpt, reason, "PENDING", now, now),
                )
            self.connection.execute(
                "UPDATE imports SET status='REVIEW_REQUIRED',completed_at=? WHERE id=?", (now, import_id)
            )
        self.events.publish("structured_import.staged", {"import_id": import_id, "source_name": path.name, "target_type": target_type, "rows": len(rows)})
        return {"import_id": import_id, "source_format": source_format, "headers": headers, "mapping": chosen, "review_items": len(rows)}

    def stage_legal_source_document(self, source: Path | str, *, client_id: int | None = None, matter_id: int | None = None,
                                    category: str | None = "Legal source for review", tags: list[str] | None = None,
                                    user_id: int | None = None) -> dict[str, Any]:
        path = Path(source)
        if not path.is_file():
            raise ValidationError("Selected source is not a local file.")
        now = to_utc_iso(self.clock.now())
        # Keep the import-history row and vault ingestion as separate transactions: the
        # vault deliberately owns its transaction and encrypted-file rollback behavior.
        with transaction(self.connection):
            import_id = self._start_import(path, "LEGAL_SOURCE_DOCUMENT_REVIEW", 1)
        try:
            duplicate_id = self.vault.find_duplicate(path, client_id=client_id, matter_id=matter_id)
            document_id = self.vault.ingest(path, client_id=client_id, matter_id=matter_id, category=category,
                                            tags=tags or [], source_label="User-selected legal source for manual review", user_id=user_id)
            text_rows = self.connection.execute(
                "SELECT text_content FROM document_text WHERE document_id=? ORDER BY page_number LIMIT 4", (document_id,)
            ).fetchall()
            excerpt = "\n".join(str(r["text_content"]) for r in text_rows)[:4000]
            with transaction(self.connection):
                self.connection.execute(
                    """INSERT INTO manual_review_queue(review_uuid,import_id,document_id,item_type,target_type,title,payload_json,source_excerpt,reason,status,created_at,updated_at)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (str(uuid.uuid4()), import_id, document_id, "LEGAL_SOURCE_DOCUMENT", "LEGAL_SOURCE_DOCUMENT", path.name,
                     json.dumps({"document_id": document_id, "filename": path.name}, sort_keys=True), excerpt,
                     "Review the locally extracted source and manually classify/add legal metadata. No legal proposition is auto-verified.",
                     "PENDING", now, now),
                )
                self.connection.execute("UPDATE imports SET status='REVIEW_REQUIRED',completed_at=? WHERE id=?", (now, import_id))
        except Exception as exc:
            with transaction(self.connection):
                self.connection.execute(
                    "UPDATE imports SET status='FAILED',rejected_items=1,error_log=?,completed_at=? WHERE id=?",
                    (str(exc)[:4000], to_utc_iso(self.clock.now()), import_id),
                )
            raise
        self.events.publish("legal_source.staged", {"import_id": import_id, "document_id": document_id, "source_name": path.name})
        return {"import_id": import_id, "document_id": document_id, "duplicate": duplicate_id is not None, "review_items": 1}

    def list_review(self, status: str = "PENDING", limit: int = 250) -> list[dict[str, Any]]:
        if status not in {"PENDING", "RESOLVED", "REJECTED", "ALL"}:
            raise ValidationError("Review status filter is invalid.")
        if not 1 <= limit <= 1000:
            raise ValidationError("Review result limit is invalid.")
        where = "" if status == "ALL" else "WHERE q.status=?"
        params: list[Any] = [] if status == "ALL" else [status]
        params.append(limit)
        rows = self.connection.execute(
            f"""SELECT q.*,i.source_name,i.import_type,d.original_filename AS document_filename
                FROM manual_review_queue q
                LEFT JOIN imports i ON i.id=q.import_id
                LEFT JOIN documents d ON d.id=q.document_id
                {where} ORDER BY q.created_at DESC,q.id DESC LIMIT ?""", params
        ).fetchall()
        return [dict(row) for row in rows]

    def get_review(self, review_id: int) -> dict[str, Any]:
        row = self.connection.execute(
            """SELECT q.*,i.source_name,i.import_type,d.original_filename AS document_filename
               FROM manual_review_queue q LEFT JOIN imports i ON i.id=q.import_id
               LEFT JOIN documents d ON d.id=q.document_id WHERE q.id=?""", (review_id,)
        ).fetchone()
        if row is None:
            raise ValidationError("Review item was not found.")
        item = dict(row)
        item["payload"] = json.loads(item["payload_json"])
        return item

    def update_payload(self, review_id: int, payload: dict[str, Any], *, user_id: int | None = None) -> None:
        item = self.get_review(review_id)
        if item["status"] != "PENDING":
            raise ValidationError("Only pending review items can be edited.")
        if item["target_type"] in TARGET_FIELDS:
            allowed = set(TARGET_FIELDS[item["target_type"]])
            unknown = set(payload) - allowed
            if unknown:
                raise ValidationError("Review payload contains unsupported field(s): " + ", ".join(sorted(unknown)))
        encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True)
        if len(encoded) > 2_000_000:
            raise ValidationError("Review payload is too large.")
        self.connection.execute("UPDATE manual_review_queue SET payload_json=?,updated_at=? WHERE id=?", (encoded, to_utc_iso(self.clock.now()), review_id))
        self.connection.commit()

    def resolve(self, review_id: int, *, user_id: int | None = None, note: str | None = None) -> dict[str, Any]:
        item = self.get_review(review_id)
        if item["status"] != "PENDING":
            raise ValidationError("Review item has already been resolved or rejected.")
        payload = item["payload"]
        created_type: str | None = None
        created_id: int | None = None
        if item["target_type"] == "CLIENT":
            created_id = self.clients.create(payload, user_id=user_id); created_type = "client"
        elif item["target_type"] == "LEGAL_STATUTE":
            data = dict(payload); data["verification_status"] = "REVIEW_REQUIRED"
            created_id = self.legal.add_statute(data, user_id=user_id); created_type = "statute"
        elif item["target_type"] == "LEGAL_JUDGMENT":
            data = dict(payload); data["verification_status"] = "REVIEW_REQUIRED"
            created_id = self.legal.add_judgment(data, user_id=user_id); created_type = "judgment"
        elif item["target_type"] == "LEGAL_SOURCE_DOCUMENT":
            created_id = item.get("document_id"); created_type = "document"
        else:
            raise ValidationError("Review target cannot be resolved automatically.")
        now = to_utc_iso(self.clock.now())
        with transaction(self.connection):
            self.connection.execute(
                """UPDATE manual_review_queue SET status='RESOLVED',resolved_at=?,resolved_by_user_id=?,resolution_note=?,
                   created_entity_type=?,created_entity_id=?,updated_at=? WHERE id=?""",
                (now, user_id, (note or "")[:2000] or None, created_type, created_id, now, review_id),
            )
            if item.get("import_id"):
                counts = self.connection.execute(
                    """SELECT COUNT(*) total,
                       SUM(CASE WHEN status='RESOLVED' THEN 1 ELSE 0 END) resolved,
                       SUM(CASE WHEN status='REJECTED' THEN 1 ELSE 0 END) rejected,
                       SUM(CASE WHEN status='PENDING' THEN 1 ELSE 0 END) pending
                       FROM manual_review_queue WHERE import_id=?""", (item["import_id"],)
                ).fetchone()
                final_status = "REVIEW_REQUIRED" if int(counts["pending"] or 0) else ("PARTIAL" if int(counts["rejected"] or 0) else "COMPLETE")
                self.connection.execute(
                    "UPDATE imports SET status=?,imported_items=?,rejected_items=? WHERE id=?",
                    (final_status, int(counts["resolved"] or 0), int(counts["rejected"] or 0), item["import_id"]),
                )
        self.events.publish("manual_review.resolved", {"review_id": review_id, "created_entity_type": created_type, "created_entity_id": created_id})
        return {"review_id": review_id, "created_entity_type": created_type, "created_entity_id": created_id}

    def reject(self, review_id: int, *, user_id: int | None = None, note: str | None = None) -> None:
        item = self.get_review(review_id)
        if item["status"] != "PENDING":
            raise ValidationError("Review item has already been resolved or rejected.")
        now = to_utc_iso(self.clock.now())
        with transaction(self.connection):
            self.connection.execute(
                "UPDATE manual_review_queue SET status='REJECTED',resolved_at=?,resolved_by_user_id=?,resolution_note=?,updated_at=? WHERE id=?",
                (now, user_id, (note or "")[:2000] or None, now, review_id),
            )
            if item.get("import_id"):
                counts = self.connection.execute(
                    """SELECT SUM(CASE WHEN status='RESOLVED' THEN 1 ELSE 0 END) resolved,
                       SUM(CASE WHEN status='REJECTED' THEN 1 ELSE 0 END) rejected,
                       SUM(CASE WHEN status='PENDING' THEN 1 ELSE 0 END) pending
                       FROM manual_review_queue WHERE import_id=?""", (item["import_id"],)
                ).fetchone()
                final_status = "REVIEW_REQUIRED" if int(counts["pending"] or 0) else "PARTIAL"
                self.connection.execute(
                    "UPDATE imports SET status=?,imported_items=?,rejected_items=? WHERE id=?",
                    (final_status, int(counts["resolved"] or 0), int(counts["rejected"] or 0), item["import_id"]),
                )
        self.events.publish("manual_review.rejected", {"review_id": review_id})
