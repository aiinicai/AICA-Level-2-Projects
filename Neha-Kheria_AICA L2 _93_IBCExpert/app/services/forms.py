"""Persistent form/template merge, versioning and real document exports."""
from __future__ import annotations

import hashlib
import json
import re
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.core.errors import StorageError, ValidationError
from app.core.time import Clock, SystemClock, to_utc_iso
from app.db.connection import transaction
from app.security.audit import AuditService

TOKEN = re.compile(r"{{\s*([A-Za-z0-9_.]+)\s*}}")


def _resolve(data: dict[str, Any], dotted: str) -> str:
    value: Any = data
    for part in dotted.split("."):
        if not isinstance(value, dict) or part not in value:
            return ""
        value = value[part]
    return "" if value is None else str(value)


def merge_template(template: str, data: dict[str, Any]) -> str:
    return TOKEN.sub(lambda match: _resolve(data, match.group(1)), template)


@dataclass
class FormService:
    connection: sqlite3.Connection
    audit: AuditService
    clock: Clock = SystemClock()

    def add_template(self, *, name: str, template: str, template_type: str = "TEXT", classification: str = "PRACTITIONER_TEMPLATE", verification_status: str = "REVIEW_REQUIRED", form_code: str | None = None, citation_id: int | None = None) -> int:
        if template_type not in {"TEXT", "HTML", "DOCX", "XLSX"}:
            raise ValidationError("Form template type is invalid.")
        if classification == "VERIFIED_OFFICIAL" and (verification_status != "VERIFIED" or citation_id is None):
            raise ValidationError("An official form requires VERIFIED status and an associated citation.")
        now = to_utc_iso(self.clock.now())
        cursor = self.connection.execute(
            """INSERT INTO forms(name,form_code,classification,template_type,template_content,
               verification_status,citation_id,active,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?,1,?,?)""",
            (name.strip(), form_code, classification, template_type, template.encode("utf-8"), verification_status, citation_id, now, now),
        )
        form_id = int(cursor.lastrowid)
        self.audit.append("FORM_TEMPLATE_ADDED", f"Added form template '{name}'.", entity_type="form", entity_id=form_id, details={"classification": classification})
        return form_id

    def matter_merge_data(self, matter_id: int) -> dict[str, Any]:
        row = self.connection.execute(
            """SELECT m.*,c.name AS client_name,c.corporate_debtor,c.cin,c.gst,c.registered_office,
                      c.industry,c.primary_email,c.primary_phone
               FROM matters m JOIN clients c ON c.id=m.client_id WHERE m.id=?""", (matter_id,)
        ).fetchone()
        if row is None:
            raise ValidationError("Matter was not found.")
        matter = {key: row[key] for key in row.keys() if key not in {"client_name", "corporate_debtor", "cin", "gst", "registered_office", "industry", "primary_email", "primary_phone"}}
        client = {"name": row["client_name"], "corporate_debtor": row["corporate_debtor"], "cin": row["cin"], "gst": row["gst"], "registered_office": row["registered_office"], "industry": row["industry"], "primary_email": row["primary_email"], "primary_phone": row["primary_phone"]}
        return {"client": client, "matter": matter}

    def generate(self, form_id: int, matter_id: int, *, edited_content: str | None = None) -> int:
        form = self.connection.execute("SELECT * FROM forms WHERE id=? AND active=1", (form_id,)).fetchone()
        if form is None:
            raise ValidationError("Active form template was not found.")
        data = self.matter_merge_data(matter_id)
        template = bytes(form["template_content"] or b"").decode("utf-8")
        content = edited_content if edited_content is not None else merge_template(template, data)
        if len(content) > 5_000_000:
            raise ValidationError("Generated form is too large.")
        version = self.connection.execute("SELECT COALESCE(MAX(version_number),0)+1 FROM form_versions WHERE form_id=? AND matter_id=?", (form_id, matter_id)).fetchone()[0]
        now = to_utc_iso(self.clock.now())
        digest = hashlib.sha256(content.encode("utf-8")).hexdigest()
        cursor = self.connection.execute(
            """INSERT INTO form_versions(form_id,matter_id,version_number,merged_data_json,
               editable_content,output_type,content_hash,created_at) VALUES(?,?,?,?,?,?,?,?)""",
            (form_id, matter_id, version, json.dumps(data, ensure_ascii=False, default=str), content.encode("utf-8"), form["template_type"], digest, now),
        )
        version_id = int(cursor.lastrowid)
        self.audit.append("FORM_GENERATED", f"Generated form version {version}.", entity_type="form_version", entity_id=version_id, details={"form_id": form_id, "matter_id": matter_id})
        return version_id

    def content(self, version_id: int) -> str:
        row = self.connection.execute("SELECT editable_content FROM form_versions WHERE id=?", (version_id,)).fetchone()
        if row is None:
            raise ValidationError("Generated form version was not found.")
        return bytes(row[0]).decode("utf-8")

    def export(self, version_id: int, destination: Path, output_type: str) -> Path:
        row = self.connection.execute(
            """SELECT fv.*,f.name,f.form_code FROM form_versions fv JOIN forms f ON f.id=fv.form_id WHERE fv.id=?""", (version_id,)
        ).fetchone()
        if row is None:
            raise ValidationError("Generated form version was not found.")
        output_type = output_type.upper()
        if output_type not in {"TXT", "DOCX", "PDF", "XLSX"}:
            raise ValidationError("Requested form export type is not supported.")
        destination.mkdir(parents=True, exist_ok=True)
        safe = re.sub(r"[^A-Za-z0-9._ -]", "_", row["form_code"] or row["name"]).strip() or "form"
        target = destination / f"{safe}-v{row['version_number']}.{output_type.lower()}"
        content = bytes(row["editable_content"]).decode("utf-8")
        try:
            if output_type == "TXT":
                target.write_text(content, encoding="utf-8")
            elif output_type == "DOCX":
                from docx import Document
                document = Document()
                for paragraph in content.splitlines() or [""]:
                    document.add_paragraph(paragraph)
                document.save(target)
            elif output_type == "PDF":
                from reportlab.lib.pagesizes import A4
                from reportlab.pdfbase.pdfmetrics import stringWidth
                from reportlab.pdfgen.canvas import Canvas
                canvas = Canvas(str(target), pagesize=A4)
                width, height = A4
                y = height - 54
                for original_line in content.splitlines() or [""]:
                    words = original_line.split() or [""]
                    line = ""
                    for word in words:
                        candidate = (line + " " + word).strip()
                        if stringWidth(candidate, "Helvetica", 10) > width - 108 and line:
                            canvas.drawString(54, y, line); y -= 14; line = word
                            if y < 54: canvas.showPage(); y = height - 54
                        else:
                            line = candidate
                    canvas.drawString(54, y, line); y -= 14
                    if y < 54: canvas.showPage(); y = height - 54
                canvas.save()
            else:
                from openpyxl import Workbook
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Generated Form"
                for index, line in enumerate(content.splitlines(), 1):
                    sheet.cell(index, 1, line)
                workbook.save(target)
        except (ImportError, OSError) as exc:
            raise StorageError(f"{output_type} export component is unavailable or the file cannot be written.") from exc
        digest = hashlib.sha256(target.read_bytes()).hexdigest()
        if target.stat().st_size == 0:
            raise StorageError("Generated export file is empty.")
        self.connection.execute("UPDATE form_versions SET output_type=?,output_path=? WHERE id=?", (output_type, str(target), version_id))
        self.audit.append("FORM_EXPORTED", f"Exported generated form as {output_type}.", entity_type="form_version", entity_id=version_id, details={"sha256": digest})
        return target
