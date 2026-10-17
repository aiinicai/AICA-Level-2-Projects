"""Module A part 2 - the remaining FS-pack checks. Class A = rules; class B/C = ASSISTED (software extracts evidence, reviewer confirms)."""
from __future__ import annotations
import re
from datetime import date

from .loaders import find_row, row_label_and_numbers
from .models import Finding
from .rules_fs import NotRun, Ctx, MONTHS, _sofp_table, _pair, _cy_first, _soce_table, _cf_table


def _f(cid, title, sev, detail, loc, exp="", found=""):
    return Finding(cid, title, sev, detail, loc, str(exp), str(found))


def _rv(cid, title, detail, loc="FS pack", exp="", found=""):
    return Finding(cid, title, "Info", detail, loc, str(exp), str(found))


def _assist(fn):
    fn.assist = True
    return fn


def _cy(c, row):
    a, b, loc = _pair(row)
    return (a if _cy_first(c, row) else b), loc


def _pl(c, rx):
    r = find_row(c.doc.tables, rx, lambda t: any(re.search(r"^revenue|^profit for the year|^gross profit", x[0] or "", re.I) for x in t.rows if x))
    return r


def _text(c):
    return re.sub(r"[ \t]+", " ", c.doc.all_text)


def _dates(txt):
    out = []
    for m in re.finditer(r"(\d{1,2})\s+(January|February|March|April|May|June|July|August|September|October|November|December)\s+(20\d{2})", txt, re.I):
        out.append(date(int(m.group(3)), MONTHS[m.group(2).lower()], int(m.group(1))))
    return out


# ------------------------------------------------------------------ class A
def sa12_dates(c: Ctx):
    t = _text(c)
    aud = c.audit_region()
    ad = _dates(aud[-1200:]) if aud else []
    appr = [d for m in re.finditer(r"(?:approved|authori[sz]ed)[^.]{0,80}?(?:for issue|issue)[^.]{0,30}", t, re.I) for d in _dates(t[m.start(): m.end() + 60])]
    drep = [d for m in re.finditer(r"directors.? report", t, re.I) for d in _dates(t[m.start(): m.start() + 3000])][:3]
    if not ad:
        raise NotRun("Auditor's report date not found")
    out, adt = [], max(ad)
    for lab, ds in (("approval for issue", appr), ("Directors' report", drep)):
        if ds and max(ds) > adt:
            out.append(_f("SA-12", f"{lab} is dated after the auditor's report", "High", f"{lab} {max(ds):%d %B %Y} vs auditor's report {adt:%d %B %Y}", "Dates", f"<= {adt:%d %B %Y}", f"{max(ds):%d %B %Y}"))
    if appr and not re.search(r"subsequent events?|events after the reporting", t, re.I):
        out.append(_f("SA-12", "Approval date present but no subsequent-events note (ISA 560 / IAS 10)", "Medium", "No 'subsequent events' wording found", "Notes", "subsequent events note", "absent"))
    return out


def _docx_tables(c):
    if c.doc.kind != "docx":
        raise NotRun("Needs the Word version of the pack (bold/alignment cannot be read from a PDF)")
    import docx
    return docx.Document(c.doc.path)


def sa17_format(c: Ctx):
    d = _docx_tables(c)
    out = []
    for ti, t in enumerate(d.tables[:4], 1):
        hdr = t.rows[0].cells
        cy = next((i for i, cell in enumerate(hdr) if re.fullmatch(r"\s*%d\s*" % c.cy, cell.text)), None)
        py = next((i for i, cell in enumerate(hdr) if re.fullmatch(r"\s*%d\s*" % c.py, cell.text)), None)
        if cy is None or py is None:
            continue
        bold = lambda cell: any(r.bold for p in cell.paragraphs for r in p.runs)
        if not bold(hdr[cy]):
            out.append(_f("SA-17", "Current-year column heading is not bold", "Low", f"Table {ti}: '{c.cy}' heading", f"Table {ti}", "bold", "regular"))
        if bold(hdr[py]):
            out.append(_f("SA-17", "Prior-year column heading is bold", "Low", f"Table {ti}: '{c.py}' heading", f"Table {ti}", "regular", "bold"))
        for ri, row in enumerate(t.rows[1:], 2):
            for ci in (cy, py):
                cell = row.cells[ci]
                tx = cell.text.strip()
                if re.fullmatch(r"\(?-?\d{4,}\)?", tx):
                    out.append(_f("SA-17", "Number without thousands separator", "Low", f"Table {ti}, row {ri}: '{tx}'", f"Table {ti}, row {ri}", "1,234", tx))
                if tx and re.search(r"\d", tx) and cell.paragraphs and cell.paragraphs[0].alignment is not None and int(cell.paragraphs[0].alignment) == 0:
                    out.append(_f("SA-17", "Figure is left-aligned", "Low", f"Table {ti}, row {ri}: '{tx}'", f"Table {ti}, row {ri}", "right", "left"))
    if not any(True for _ in d.tables):
        raise NotRun("No tables")
    return out[:8]


def sa19_case(c: Ctx):
    out = []
    for tb in c.doc.tables:
        for ri, r in enumerate(tb.rows):
            lab = (r[0] if r else "").strip()
            words = lab.split()
            if len(words) >= 3 and sum(1 for w in words[1:] if w[:1].isupper() and w.lower() not in ("and", "of", "the", "ifrs", "vat", "uae", "ppe")) >= 2 and not lab.isupper():
                out.append(_f("SA-19", "Line item not in sentence case", "Low", f"'{lab}'", tb.loc + f", row {ri + 1}", "Sentence case", lab))
            if "  " in lab or re.search(r"\b(\w{3,})\s+\1\b", lab, re.I):
                out.append(_f("SA-19", "Double space or repeated word in a line item", "Low", f"'{lab}'", tb.loc + f", row {ri + 1}", "", lab))
    for m in re.finditer(r"\b(\w{3,})\s+\1\b", _text(c)):
        if m.group(1).lower() not in ("that", "had"):
            out.append(_f("SA-19", "Repeated word in text", "Low", f"'{m.group(0)}'", "Text", "", m.group(0)))
    return out[:8]


def sa23_profit(c: Ctx):
    vals = {}
    for lab, rx, flt in (("P&L", r"^profit (for the year|after tax)|^net profit|^profit/\(loss\) for the year", None),):
        r = find_row(c.doc.tables, rx, lambda t: not _soce_table(t) and not _cf_table(t))
        if r:
            vals["P&L"] = (_cy(c, r)[0], r[2])
    for t in c.doc.tables:
        if _soce_table(t):
            for ri, row in enumerate(t.rows):
                lab, nums = row_label_and_numbers(row)
                if re.match(r"^(profit|net profit|total comprehensive income)", lab, re.I) and nums:
                    vals["SOCE"] = (nums[-1] if len(nums) < 3 else nums[-2] if False else nums[1] if len(nums) >= 2 else nums[0], f"{t.loc}, row {ri + 1}")
                    break
    m = re.search(r"(?:net )?(?:profit|loss)[^.]{0,40}?(?:of|was|amounted to)\s+(?:AED\s*)?(\d[\d,]{3,})", _text(c), re.I)
    if m:
        vals["Directors' report"] = (float(m.group(1).replace(",", "")), "Directors' report")
    if len(vals) < 2:
        raise NotRun("Net profit not found in at least two places (P&L, SOCE, Directors' report)")
    base = vals.get("P&L") or next(iter(vals.values()))
    out = []
    for k, (v, loc) in vals.items():
        if v is not None and base[0] is not None and abs(abs(v) - abs(base[0])) > 1:
            out.append(_f("SA-23", f"Net profit per {k} differs from the P&L", "High", f"{k} {v:,.0f} vs P&L {base[0]:,.0f}", loc, f"{base[0]:,.0f}", f"{v:,.0f}"))
    return out


def sa25_notes(c: Ctx):
    refs = {}
    for t in c.doc.tables:
        if not (_sofp_table(t) or table_has_pl(t)):
            continue
        for ri, r in enumerate(t.rows):
            if len(r) >= 3 and re.fullmatch(r"\d{1,2}", (r[1] or "").strip()) and (r[0] or "").strip():
                refs.setdefault(int(r[1]), f"{t.loc}, row {ri + 1}: '{r[0].strip()}'")
    if not refs:
        raise NotRun("No 'Note' column found in the statements")
    nt = c.notes_text()
    have = {int(m) for m in re.findall(r"(?:^|\s)(\d{1,2})\.\s+[A-Z]", nt)}
    have |= {int(m) for m in re.findall(r"\bNote\s+(\d{1,2})\b", nt)}
    out = []
    if not have:
        return [_f("SA-25", "Statements refer to notes but no numbered notes were found", "High", f"Note references {sorted(refs)}", "Notes", "numbered notes", "none")]
    for n, loc in sorted(refs.items()):
        if n not in have:
            out.append(_f("SA-25", f"Note {n} is referred to in the statements but not found in the notes", "High", loc, loc.split(":")[0], f"Note {n}", "absent"))
    return out[:8]


def table_has_pl(t):
    return any(re.match(r"^revenue|^profit for the year", (r[0] or ""), re.I) for r in t.rows if r)


def sa26_dirrep(c: Ctx):
    t = _text(c)
    m = re.search(r"revenue[^.]{0,40}?(?:of|was|amounted to)\s+(?:AED\s*)?(\d[\d,]{4,})", t, re.I)
    if not m:
        raise NotRun("No revenue figure stated in the Directors' report narrative")
    r = find_row(c.doc.tables, r"^revenue", table_has_pl)
    if not r:
        raise NotRun("Revenue row of the P&L not found")
    v, loc = _cy(c, r)
    d = float(m.group(1).replace(",", ""))
    return [_f("SA-26", "Revenue in the Directors' report differs from the P&L", "High", f"'{m.group(0)[:70]}' vs P&L {v:,.0f}", "Directors' report", f"{v:,.0f}", f"{d:,.0f}")] if abs(d - v) > 1 else []


def _note_table(c, title_rx):
    nt = c.doc.text
    for tb in c.doc.tables:
        first = " ".join(" ".join(r) for r in tb.rows[:2])
        if re.search(title_rx, first, re.I):
            return tb
    return None


def sa27_ga(c: Ctx):
    tb = None
    for t in c.doc.tables:
        labs = [(r[0] or "").lower() for r in t.rows if r]
        if any(re.search(r"general and administrative|administrative expenses", " ".join(r), re.I) for r in t.rows[:2]) or (sum(1 for l in labs if re.search(r"salar|rent|audit fee|professional|utilities|insurance", l)) >= 3):
            tb = t
            break
    if not tb:
        raise NotRun("General and administrative expenses breakdown not found")
    rows = []
    for ri, r in enumerate(tb.rows):
        lab, nums = row_label_and_numbers(r)
        if lab and nums and not re.match(r"^total", lab, re.I):
            rows.append((lab, nums[0] if _cy_first(c, (lab, nums, "")) else nums[-1], ri + 1))
    out = []
    vals = [abs(v) for _, v, _ in rows]
    body = [r for r in rows if not re.match(r"^(other|miscellaneous|sundry)", r[0], re.I)]
    if [abs(v) for _, v, _ in body] != sorted([abs(v) for _, v, _ in body], reverse=True):
        out.append(_f("SA-27", "G&A expenses are not in descending order", "Low", f"Order: {[l for l, _, _ in rows][:6]}", tb.loc, "descending", "mixed"))
    gp = find_row(c.doc.tables, r"^gross profit", table_has_pl)
    oth = [r for r in rows if re.match(r"^(other|miscellaneous|sundry)", r[0], re.I)]
    if gp and oth:
        g, _ = _cy(c, gp)
        if g and abs(oth[0][1]) > 0.05 * abs(g):
            out.append(_f("SA-27", "'Other' G&A expenses exceed 5% of gross profit", "Medium", f"'{oth[0][0]}' {abs(oth[0][1]):,.0f} vs 5% of gross profit {0.05 * abs(g):,.0f}", f"{tb.loc}, row {oth[0][2]}", "<= 5%", f"{abs(oth[0][1]) / abs(g):.1%}"))
    if not any(re.search(r"audit", l, re.I) for l, _, _ in rows):
        out.append(_f("SA-27", "Audit fee not recorded in G&A", "Medium", "No audit fee line in the G&A breakdown", tb.loc, "audit fee line", "absent"))
    return out


def sa28_forex(c: Ctx):
    out = []
    fi = _note_table(c, r"finance income|other income")
    fc = _note_table(c, r"finance costs?|finance expense")
    if not (fi or fc):
        raise NotRun("Finance income / finance cost note tables not found")
    if fi and any(re.search(r"exchange (loss|difference)|forex loss", " ".join(r), re.I) for r in fi.rows):
        out.append(_f("SA-28", "Exchange loss shown in finance income", "Medium", "'exchange loss' appears in the finance income table", fi.loc, "loss in finance cost", "in finance income"))
    if fc and any(re.search(r"exchange gain|forex gain", " ".join(r), re.I) for r in fc.rows):
        out.append(_f("SA-28", "Exchange gain shown in finance cost", "Medium", "'exchange gain' appears in the finance cost table", fc.loc, "gain in finance income", "in finance cost"))
    return out


def sa32_cash(c: Ctx):
    t = c.notes_text()
    if not re.search(r"cash and (bank|cash equivalents)", t, re.I):
        raise NotRun("Cash and bank note not found")
    if not re.search(r"outside (the )?UAE|inside (the )?UAE|within (the )?UAE|in the UAE|overseas", t, re.I):
        return [_f("SA-32", "Cash and bank note does not split balances inside / outside the UAE", "Low", "No UAE / outside-UAE wording in the cash note", "Notes - cash and bank", "UAE / outside UAE split", "absent")]
    return []


def sa37_ecl(c: Ctx):
    t = c.notes_text()
    if not find_row(c.doc.tables, r"trade (and other )?receivables|accounts receivable", _sofp_table):
        raise NotRun("No trade receivables on the SOFP")
    out = []
    if not re.search(r"expected credit loss|ECL|impairment (allowance|loss)|provision for (doubtful|impairment)", t, re.I):
        out.append(_f("SA-37", "Receivables without an ECL / impairment disclosure (IFRS 9)", "High", "No expected credit loss wording in the notes", "Notes - receivables", "ECL note", "absent"))
    ag = None
    for tb in c.doc.tables:
        if any(re.search(r"(0\s*[-–]\s*30|1\s*[-–]\s*30|30 days|90 days|over 365|> ?365)", " ".join(r)) for r in tb.rows[:2]):
            ag = tb
            break
    if not ag:
        out.append(_f("SA-37", "No ageing analysis of receivables", "Medium", "No ageing table (0-30, 31-60 ... days) found", "Notes - receivables", "ageing table", "absent"))
    return out


def sa40_rp(c: Ctx):
    v = c.eng.get("tp_related_party_total")
    if v is None:
        raise NotRun("Enter 'tp_related_party_total' in the fact sheet (from the TP report) to compare related-party balances")
    r = find_row(c.doc.tables, r"related part(y|ies)|due (from|to) related|amounts? due (from|to) (shareholder|director)", None)
    if not r:
        raise NotRun("No related-party row in the statements/notes")
    a = abs(r[1][-2] if len(r[1]) >= 2 and _cy_first(c, r) else r[1][-1])
    return [_f("SA-40", "Related-party amount in the FS differs from the TP report", "High", f"FS {a:,.0f} vs TP report {float(v):,.0f}", r[2], f"{float(v):,.0f}", f"{a:,.0f}")] if abs(a - float(v)) > 1 else []


def sa43_ct(c: Ctx):
    v = c.eng.get("ct_payable_working")
    if v is None:
        raise NotRun("Enter 'ct_payable_working' (from the CT working sheet) in the fact sheet to compare")
    r = find_row(c.doc.tables, r"(corporate|income) tax (payable|expense)|current tax|provision for (corporate )?tax", None)
    if not r:
        raise NotRun("No corporate tax row in the statements/notes")
    a = abs(_cy(c, r)[0] or 0)
    out = []
    if abs(a - float(v)) > 1:
        out.append(_f("SA-43", "CT payable per FS differs from the CT working", "High", f"FS {a:,.0f} vs CT working {float(v):,.0f}", r[2], f"{float(v):,.0f}", f"{a:,.0f}"))
    rv = c.eng.get("ct_return_payable")
    if rv is not None and abs(float(rv) - float(v)) > 1:
        out.append(_f("SA-43", "CT working differs from the return", "High", f"Working {float(v):,.0f} vs return {float(rv):,.0f}", "CT return", f"{float(v):,.0f}", f"{float(rv):,.0f}"))
    return out


def sa45_vat(c: Ctx):
    v = c.eng.get("vat_turnover")
    if v is None:
        raise NotRun("Enter 'vat_turnover' (total of box-1 to box-7 turnover for the year) in the fact sheet")
    r = find_row(c.doc.tables, r"^revenue|^sales", table_has_pl)
    if not r:
        raise NotRun("Revenue row not found")
    a = _cy(c, r)[0]
    d = abs(a - float(v))
    return [_f("SA-45", "Revenue per books does not reconcile to VAT returns", "Medium", f"Revenue {a:,.0f} vs VAT turnover {float(v):,.0f}; difference {d:,.0f} ({d / max(abs(a), 1):.1%}). Explain: zero-rated/exempt, timing, cross-year, out-of-scope.", r[2], f"{float(v):,.0f}", f"{a:,.0f}")] if d > max(1, 0.005 * abs(a)) else []


def sa03_owners(c: Ctx):
    sh = c.client.get("shareholders") or []
    if not sh:
        raise NotRun("Shareholders not in the fact sheet")
    t = _text(c)
    pct = [float(x) for x in re.findall(r"(\d{1,3}(?:\.\d+)?)\s*%", t[t.lower().find("share capital"):][:3000])] if "share capital" in t.lower() else []
    out = []
    if pct and 0 < len(pct) <= 8 and abs(sum(pct) - 100) > 0.01 and sum(pct) not in (200.0,):
        out.append(_f("SA-03", "Shareholding percentages do not add up to 100%", "High", f"Percentages {pct} = {sum(pct):g}%", "Share capital note", "100%", f"{sum(pct):g}%"))
    return out


# ------------------------------------------------------------------ class B / C assisted
def _find(c, rx, n=3, w=140):
    t = _text(c)
    return [re.sub(r"\s+", " ", t[max(0, m.start() - 20): m.end() + w]).strip() for m in re.finditer(rx, t, re.I)][:n]


@_assist
def sa11_opinion(c: Ctx):
    a = c.audit_region()
    if not a:
        raise NotRun("Auditor's report not found")
    op = "modified" if re.search(r"qualified opinion|adverse opinion|disclaimer of opinion|except for", a, re.I) else "unmodified"
    items = {"Basis for opinion": r"basis for (qualified |adverse )?opinion", "Key audit matters": r"key audit matters", "Emphasis of matter": r"emphasis of matter", "Material uncertainty - going concern": r"material uncertainty|going concern",
             "Other information": r"other information", "Responsibilities of management": r"responsibilit\w+ of (management|those charged)", "Auditor's responsibilities": r"auditor.?s responsibilit"}
    have = {k: bool(re.search(v, a, re.I)) for k, v in items.items()}
    return [_rv("SA-11", "Auditor's report wording", f"Opinion looks {op}. Sections present: {[k for k, v in have.items() if v]}. Missing: {[k for k, v in have.items() if not v]}. ISA 700/705/706/570: confirm wording, going concern, related party and inventory paragraphs.", "Independent auditor's report")]


@_assist
def sa31_law(c: Ctx):
    auth = str(c.client.get("authority", ""))
    hits = _find(c, r"(federal (decree[- ])?law|commercial companies law|law no\.?\s*\(?\d+\)?|DMCC (law|company regulations)|regulations? \d{4})", 3)
    return [_rv("SA-31", "Statement of compliance - law cited", f"Authority {auth}. Wording found: {hits or 'none'}. Confirm it names the correct law for this authority (e.g. Federal Decree-Law No. 32 of 2021 for mainland / the free-zone company law).", "Note 1")]


@_assist
def sa33_policies(c: Ctx):
    pol = {"IFRS 9 / ECL": r"IFRS 9|expected credit loss", "IFRS 15 revenue": r"IFRS 15|revenue from contracts", "IFRS 16 leases": r"IFRS 16|right-of-use|lease", "IAS 24 related parties": r"IAS 24|related part", "IAS 12 / corporate tax": r"IAS 12|corporate tax|deferred tax", "IAS 1 going concern": r"going concern", "IAS 19 / EOSB": r"IAS 19|end of service|gratuity", "IAS 16 PPE": r"IAS 16|property, plant"}
    t = c.notes_text()
    have = {k: bool(re.search(v, t, re.I)) for k, v in pol.items()}
    miss = [k for k, v in have.items() if not v]
    return [_rv("SA-33", "Accounting policy coverage", f"Policies mentioned: {[k for k, v in have.items() if v]}. Not found in the notes: {miss or 'none'}. Confirm applicability and IFRS-compliant wording.", "Note 2/3")]


@_assist
def sa36_legal(c: Ctx):
    h = _find(c, r"litigation|legal (case|claim|proceeding|fees)|contingent liabilit|lawsuit", 4)
    return [_rv("SA-36", "Legal cases and contingencies", f"References found: {h or 'none - confirm none exist (legal confirmation)'}. Check unusual legal fees in the G&A note.", "Notes - contingencies")]


@_assist
def sa38_loans(c: Ctx):
    h = _find(c, r"(bank (loan|borrowing|facilit)|term loan|overdraft|borrowings)", 3)
    cl = bool(re.search(r"current portion|non-current|repayable (within|after)", _text(c), re.I))
    return [_rv("SA-38", "Bank loans", f"Mentions: {h or 'none'}. Current/non-current split wording present: {cl}. Obtain bank confirmation; check covenants and classification (IAS 1).", "Notes - borrowings")]


@_assist
def sa39_group(c: Ctx):
    h = _find(c, r"(subsidiar|associate|joint venture|branch|goodwill|liquidat)", 5)
    return [_rv("SA-39", "Associates, subsidiaries, branches, goodwill", f"References: {h or 'none'}. Check the licence/MOA for branches and group companies, and whether any entity was liquidated or deregistered in the year.", "Notes")]


@_assist
def sa41_ppe(c: Ctx):
    r = find_row(c.doc.tables, r"property, plant|fixed assets", _sofp_table)
    h = _find(c, r"(depreciation|useful li(fe|ves)|right-of-use|IFRS 16|amortis)", 3, 100)
    return [_rv("SA-41", "PPE, depreciation and leases", f"PPE on SOFP: {_cy(c, r)[0]:,.0f} ({r[2]}). " if r else "No PPE line found. " + f"Policy wording: {h or 'none'}. Agree the PPE schedule roll-forward (opening + additions - disposals - depreciation = closing) and IFRS 16 lease schedule.")]


@_assist
def sa42_dues(c: Ctx):
    r = find_row(c.doc.tables, r"end of service|employees.? benefits|gratuity", None)
    h = _find(c, r"(WPS|wages protection|gratuity|end of service)", 3, 100)
    return [_rv("SA-42", "Statutory dues, WPS and gratuity", (f"EOSB {abs(_cy(c, r)[0] or 0):,.0f} at {r[2]}. " if r else "No EOSB balance found. ") + f"Wording: {h or 'none'}. Recompute gratuity (UAE Labour Law: 21 days/yr for first 5 years, 30 days after) and check WPS status.")]


@_assist
def sa44_taxes(c: Ctx):
    vat = find_row(c.doc.tables, r"vat|duties and taxes", None)
    ct = find_row(c.doc.tables, r"corporate tax|income tax", None)
    return [_rv("SA-44", "VAT and CT payable to tick to the FTA portal", f"VAT/duties: {vat[1][:2] if vat else 'not found'} ({vat[2] if vat else ''}); corporate tax: {ct[1][:2] if ct else 'not found'} ({ct[2] if ct else ''}). Agree each to the EmaraTax statement of account.")]


@_assist
def sa35_conf(c: Ctx):
    rows = []
    for rx in (r"cash and (bank|cash equivalents)", r"trade (and other )?receivables", r"trade (and other )?payables", r"(bank (loan|borrowing)|borrowings)"):
        r = find_row(c.doc.tables, rx, _sofp_table)
        if r:
            rows.append(f"{r[0]} {abs(_cy(c, r)[0] or 0):,.0f} ({r[2]})")
    return [_rv("SA-35", "Balances needing external confirmation", "; ".join(rows) or "no balances found", "SOFP", "confirmations on file", f"{len(rows)} balances to confirm")]


@_assist
def sa34_wp(c: Ctx):
    rows = []
    for t in c.doc.tables:
        if _sofp_table(t) or table_has_pl(t):
            for ri, r in enumerate(t.rows):
                lab, nums = row_label_and_numbers(r)
                if lab and nums and abs(nums[0]) >= 100000 and not re.match(r"^total", lab, re.I):
                    rows.append(f"{lab} ({abs(nums[0]):,.0f})")
    return [_rv("SA-34", "Material lines - confirm a completed working paper exists for each", "; ".join(rows[:14]) or "no lines found", "SOFP / P&L", "working paper saved", f"{len(rows)} material lines")]


@_assist
def sa48_collection(c: Ctx):
    return [_rv("SA-48", "Collection status", "Confirm with admin that fees are collected before release; record name and date in the review log.", "Engagement file")]


@_assist
def sa51_scepticism(c: Ctx):
    h = _find(c, r"(estimate|judgement|useful li|expected credit loss|provision|impairment|fair value|related part)", 6, 90)
    return [_rv("SA-51", "Estimates and judgement areas disclosed", f"{len(h)} statements found: " + " | ".join(h[:4]) + ". Consider management override, unusual transactions and estimation uncertainty.", "Notes")]


@_assist
def sa18_index(c: Ctx):
    cont = len(re.findall(r"\(continued\)", _text(c), re.I))
    idx = bool(re.search(r"^\s*(table of )?contents?\s*$|^\s*index\s*$", c.doc.text, re.I | re.M))
    extra = ""
    if c.doc.kind != "docx":
        from .refs import printed_map
        pm = printed_map(c.doc)
        seq = [pm[k] for k in sorted(pm)]
        bad = [(a, b) for a, b in zip(sorted(pm), sorted(pm)[1:]) if pm[b] != pm[a] + (b - a)]
        extra = f" Printed page numbers read on {len(pm)} pages; sequence breaks at PDF pages {bad[:3] or 'none'}."
    return [_rv("SA-18", "Index, page numbers and continuation labels", f"Index/contents page present: {idx}. '(continued)' labels: {cont}." + extra + " Word: update the index field before PDF conversion.")]


A2_REGISTRY = {"SA-03": sa03_owners, "SA-12": sa12_dates, "SA-17": sa17_format, "SA-19": sa19_case, "SA-23": sa23_profit, "SA-25": sa25_notes,
               "SA-26": sa26_dirrep, "SA-27": sa27_ga, "SA-28": sa28_forex, "SA-32": sa32_cash, "SA-37": sa37_ecl, "SA-40": sa40_rp,
               "SA-43": sa43_ct, "SA-45": sa45_vat}
ASSIST2_REGISTRY = {"SA-11": sa11_opinion, "SA-18": sa18_index, "SA-31": sa31_law, "SA-33": sa33_policies, "SA-34": sa34_wp, "SA-35": sa35_conf, "SA-36": sa36_legal,
                    "SA-38": sa38_loans, "SA-39": sa39_group, "SA-41": sa41_ppe, "SA-42": sa42_dues, "SA-44": sa44_taxes, "SA-48": sa48_collection, "SA-51": sa51_scepticism}
