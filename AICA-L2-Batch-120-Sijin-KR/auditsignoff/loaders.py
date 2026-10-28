"""Document loaders. Everything is read locally; nothing leaves the machine."""
from __future__ import annotations
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

NUM_RE = re.compile(r"^\(?-?[\d,]+(\.\d+)?\)?$")


@dataclass
class Run:
    text: str
    font: str | None
    size: float | None
    bold: bool | None


@dataclass
class Para:
    idx: int
    text: str
    style: str
    justified: bool
    runs: list
    where: str            # "body" | "header" | "footer"
    loc: str


@dataclass
class Table:
    idx: int
    rows: list            # list[list[str]]
    loc: str


@dataclass
class SectionFmt:
    idx: int
    left: float
    top: float
    right: float
    bottom: float


@dataclass
class DocFacts:
    path: str
    kind: str                         # docx | pdf | xlsx | image (OCR) | scan (scanned PDF, OCR)
    paragraphs: list = field(default_factory=list)
    tables: list = field(default_factory=list)
    sections: list = field(default_factory=list)
    pages: list = field(default_factory=list)          # pdf only: text per page
    pdf_fonts: Counter = field(default_factory=Counter)
    pdf_sizes: Counter = field(default_factory=Counter)
    pdf_margins: list = field(default_factory=list)    # (left, top, right, bottom) inches per page

    @property
    def text(self):
        if self.kind != "docx":
            return "\n".join(self.pages)
        return "\n".join(p.text for p in self.paragraphs if p.where == "body")

    @property
    def all_text(self):
        t = self.text
        tt = "\n".join(" | ".join(r) for tb in self.tables for r in tb.rows)
        ht = "\n".join(p.text for p in self.paragraphs if p.where != "body")
        return "\n".join([t, tt, ht])


def parse_number(s: str):
    """'(1,234)' -> -1234.0 ; '-' / dash -> 0.0 ; non-numeric -> None."""
    if s is None:
        return None
    t = s.strip().replace("AED", "").replace("USD", "").replace("\u00a0", "").replace(" ", "")
    if t in {"-", "\u2013", "\u2014"}:
        return 0.0
    if not t or not NUM_RE.match(t):
        return None
    neg = t.startswith("(") and t.endswith(")")
    try:
        v = float(t.strip("()").replace(",", ""))
    except ValueError:
        return None
    return -abs(v) if neg else v


# ---------------------------------------------------------------- DOCX
def _style_chain(style):
    while style is not None:
        yield style
        style = style.base_style


def _doc_default(doc, tag):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    el = doc.styles.element.find(f".//w:docDefaults/w:rPrDefault/w:rPr/w:{tag}", ns)
    return el


def _eff_font(run, para, doc):
    if run.font.name:
        return run.font.name
    for st in _style_chain(para.style):
        if st.font is not None and st.font.name:
            return st.font.name
    el = _doc_default(doc, "rFonts")
    if el is not None:
        W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        name = el.get(W + "ascii")
        if name:
            return name
    return None


def _eff_size(run, para, doc):
    if run.font.size:
        return run.font.size.pt
    for st in _style_chain(para.style):
        if st.font is not None and st.font.size:
            return st.font.size.pt
    el = _doc_default(doc, "sz")
    if el is not None:
        W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        v = el.get(W + "val")
        if v:
            return int(v) / 2
    return None


def _jc_raw(pf_owner):
    """Raw w:jc value ('both', 'left', 'start', ...) - python-docx's enum rejects 'start'/'end' written by LibreOffice."""
    el = pf_owner._element if hasattr(pf_owner, "_element") else None
    ppr = getattr(el, "pPr", None)
    jc = getattr(ppr, "jc", None) if ppr is not None else None
    if jc is None:
        return None
    from docx.oxml.ns import qn
    return jc.get(qn("w:val"))


def _is_justified(para):
    v = _jc_raw(para)
    if v is not None:
        return v in ("both", "distribute")
    for st in _style_chain(para.style):
        v = _jc_raw(st)
        if v is not None:
            return v in ("both", "distribute")
    return False


def load_docx(path) -> DocFacts:
    import docx
    doc = docx.Document(str(path))
    facts = DocFacts(path=str(path), kind="docx")
    n = 0

    def add_para(p, where, loc_prefix):
        nonlocal n
        runs = [Run(r.text, _eff_font(r, p, doc), _eff_size(r, p, doc), r.bold) for r in p.runs if r.text.strip()]
        facts.paragraphs.append(Para(n, p.text, p.style.name if p.style is not None else "", _is_justified(p), runs, where,
                                     f"{loc_prefix} para {n + 1}: '{p.text.strip()[:50]}'"))
        n += 1

    for p in doc.paragraphs:
        add_para(p, "body", "Body")
    for si, s in enumerate(doc.sections):
        facts.sections.append(SectionFmt(si + 1, s.left_margin.inches, s.top_margin.inches,
                                         s.right_margin.inches, s.bottom_margin.inches))
        for kind, part in (("header", s.header), ("footer", s.footer)):
            for p in part.paragraphs:
                if p.text.strip():
                    add_para(p, kind, f"Section {si + 1} {kind}")
    for ti, t in enumerate(doc.tables):
        rows = []
        for r in t.rows:
            cells = []
            seen = set()
            for c in r.cells:           # merged cells repeat; keep one
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                cells.append(c.text.strip())
                for p in c.paragraphs:
                    runs = [Run(rn.text, _eff_font(rn, p, doc), _eff_size(rn, p, doc), rn.bold) for rn in p.runs if rn.text.strip()]
                    if runs:
                        facts.paragraphs.append(Para(n, p.text, p.style.name, False, runs, "table",
                                                     f"Table {ti + 1}: '{p.text.strip()[:40]}'"))
                        n += 1
            rows.append(cells)
        facts.tables.append(Table(ti + 1, rows, f"Table {ti + 1}"))
    return facts


# ---------------------------------------------------------------- PDF
def load_pdf(path) -> DocFacts:
    import pdfplumber
    facts = DocFacts(path=str(path), kind="pdf")
    with pdfplumber.open(str(path)) as pdf:
        for pi, page in enumerate(pdf.pages):
            facts.pages.append(page.extract_text() or "")
            chars = page.chars
            if chars:
                for c in chars:
                    facts.pdf_fonts[c["fontname"]] += 1
                    facts.pdf_sizes[round(c["size"] * 2) / 2] += 1
                xs0 = min(c["x0"] for c in chars)
                xs1 = max(c["x1"] for c in chars)
                top = min(c["top"] for c in chars)
                bot = max(c["bottom"] for c in chars)
                facts.pdf_margins.append((xs0 / 72, top / 72, (page.width - xs1) / 72, (page.height - bot) / 72))
            for ti, tb in enumerate(page.extract_tables() or []):
                rows = [[(c or "").replace("\n", " ").strip() for c in r] for r in tb]
                facts.tables.append(Table(len(facts.tables) + 1, rows, f"Page {pi + 1} table {ti + 1}"))
    return facts


_TESS_WIN = (r"C:\\Program Files\\Tesseract-OCR\\tesseract.exe", r"C:\\Program Files (x86)\\Tesseract-OCR\\tesseract.exe")
OCR_NOTES: list = []          # low-confidence warnings collected during a run (shown by the app / doctor)


def tesseract_setup():
    """Locate Tesseract even when it is not on PATH (default Windows folders) and pick languages. Returns (exe, langs) or raises."""
    import shutil
    import pytesseract
    exe = shutil.which("tesseract") or next((p for p in _TESS_WIN if Path(p).exists()), None)
    if exe is None:
        raise FileNotFoundError("tesseract not found")
    pytesseract.pytesseract.tesseract_cmd = exe
    try:
        have = set(pytesseract.get_languages(config=""))
    except Exception:
        have = {"eng"}
    langs = "eng+ara" if {"eng", "ara"} <= have else "eng"
    return exe, langs


def _ocr_image(img, source=""):
    where = f" for '{Path(source).name}'" if source else ""
    try:
        import pytesseract
    except ImportError:
        raise RuntimeError(f"This file is a picture or scanned PDF, so it needs OCR{where}, but the Python package 'pytesseract' is not installed. "
                           "Fix: run  python -m pip install pytesseract Pillow  and install the Tesseract program "
                           "(Windows: https://github.com/UB-Mannheim/tesseract/wiki). "
                           "No OCR needed if you upload the Word file or a PDF saved from Word (text PDF).")
    try:
        exe, langs = tesseract_setup()
        w, h = img.size
        if max(w, h) < 1500:                      # small scans misread digits: enlarge
            k = 1500 / max(w, h)
            img = img.resize((int(w * k), int(h * k)))
        txt = pytesseract.image_to_string(img, lang=langs, config="--psm 6")
        try:
            conf = [float(c) for c in pytesseract.image_to_data(img, lang=langs, config="--psm 6", output_type=pytesseract.Output.DICT)["conf"] if str(c) not in ("-1", "-1.0")]
            if conf and sum(conf) / len(conf) < 70:
                OCR_NOTES.append(f"Low OCR confidence ({sum(conf) / len(conf):.0f}%) in '{Path(source).name or 'image'}': verify every figure and date against the original scan.")
        except Exception:
            pass
        OCR_NOTES.append(f"'{Path(source).name or 'image'}' was read by OCR (scan/picture): digits such as 6/8 or 0/9 can be misread - agree every figure, date and registration number to the original before relying on a finding.")
        return txt
    except FileNotFoundError:
        raise RuntimeError(f"The Python package is present but the Tesseract program is not found{where}. Install it "
                           "(Windows: https://github.com/UB-Mannheim/tesseract/wiki; default folder C:\\Program Files\\Tesseract-OCR), then restart the app. "
                           "Or upload the Word / text-PDF version.")
    except Exception as e:
        raise RuntimeError(f"OCR failed{where}: {type(e).__name__}: {str(e)[:200]}")


def _table_from_text(text, start_idx=1, loc="OCR"):
    """Rebuild table rows from OCR lines: split on 2+ spaces so label | numbers separate."""
    rows = [re.split(r"\s{2,}", ln.strip()) for ln in text.splitlines() if ln.strip()]
    rows = [r for r in rows if len(r) > 1]
    return [Table(start_idx, rows, loc)] if rows else []


def load_image(path) -> DocFacts:
    from PIL import Image
    img = Image.open(str(path)).convert("L")
    txt = _ocr_image(img, str(path))
    f = DocFacts(path=str(path), kind="image", pages=[txt])
    f.tables = _table_from_text(txt, 1, "Image OCR")
    return f


def load_scanned_pdf(path) -> DocFacts:
    import pdfplumber
    f = DocFacts(path=str(path), kind="scan")
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages):
            txt = _ocr_image(page.to_image(resolution=250).original.convert("L"), str(path))
            f.pages.append(txt)
            f.tables += _table_from_text(txt, len(f.tables) + 1, f"Page {i + 1} OCR")
    return f


def load_xlsx_doc(path) -> DocFacts:
    """FS pack prepared in Excel: every sheet becomes text lines + a table."""
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    f = DocFacts(path=str(path), kind="xlsx")
    for ws in wb.worksheets:
        rows, lines = [], [ws.title]
        for r in ws.iter_rows(values_only=True):
            cells = ["" if v is None else (f"{v:,.0f}" if isinstance(v, (int, float)) and abs(v) >= 1000 else str(v)) for v in r]
            if any(c.strip() for c in cells):
                while cells and not cells[-1].strip():
                    cells.pop()
                rows.append(cells)
                lines.append("  ".join(c for c in cells if c.strip()))
        f.pages.append("\n".join(lines))
        f.tables.append(Table(len(f.tables) + 1, rows, f"Sheet '{ws.title}'"))
    return f


IMAGE_EXT = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp"}


def load_document(path) -> DocFacts:
    from .legacy import ensure_modern
    path = ensure_modern(path)           # .doc/.xls/.rtf -> .docx/.xlsx (offline)
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        return load_docx(path)
    if ext == ".pdf":
        f = load_pdf(path)
        if sum(len(p.strip()) for p in f.pages) < 50:        # no text layer = scanned PDF
            return load_scanned_pdf(path)
        return f
    if ext in (".xlsx", ".xlsm"):
        return load_xlsx_doc(path)
    if ext in IMAGE_EXT:
        return load_image(path)
    raise ValueError(f"Unsupported format {ext}. Use Word (.docx/.doc), PDF (text or scanned), Excel (.xlsx/.xls) or an image (JPG/PNG/TIF).")


# ---------------------------------------------------------------- TB
def load_tb(path):
    """Trial balance XLSX with columns: Account, Class (Asset/Liability/Equity/Income/Expense), Debit, Credit."""
    import openpyxl
    from .legacy import ensure_modern
    path = ensure_modern(path)
    if Path(path).suffix.lower() == ".csv":
        import csv
        with open(str(path), newline="", encoding="utf-8-sig") as fh:
            rows_iter = list(csv.reader(fh))
    else:
        wb = openpyxl.load_workbook(str(path), data_only=True)
        rows_iter = wb.worksheets[0].iter_rows(values_only=True)
    header = None
    out = []
    for row in rows_iter:
        if header is None:
            low = [str(c).strip().lower() if c is not None else "" for c in row]
            if "debit" in low and "credit" in low:
                header = {name: i for i, name in enumerate(low)}
            continue
        if all(c is None for c in row):
            continue
        g = lambda k: row[header[k]] if k in header and header[k] < len(row) else None
        num = lambda v: float(v) if isinstance(v, (int, float)) else (parse_number(str(v)) or 0.0 if v is not None else 0.0)
        out.append({"account": g("account"), "class": str(g("class") or "").strip().lower(),
                    "debit": num(g("debit")), "credit": num(g("credit"))})
    if header is None:
        raise ValueError("TB header row with 'Debit' and 'Credit' columns not found")
    return out


# ---------------------------------------------------------------- table helpers
def row_label_and_numbers(row):
    label = ""
    nums = []
    for c in row:
        v = parse_number(c)
        if v is None:
            if not label and c.strip():
                label = c.strip()
        else:
            nums.append(v)
    return label, nums


def find_row(tables, label_regex, table_filter=None):
    """Return (label, nums, loc) for first row whose label matches; nums = numeric cells."""
    rx = re.compile(label_regex, re.I)
    for t in tables:
        if table_filter and not table_filter(t):
            continue
        for ri, r in enumerate(t.rows):
            label, nums = row_label_and_numbers(r)
            if label and rx.search(label) and nums:
                return label, nums, f"{t.loc}, row {ri + 1}"
    return None


def table_has(t, regex):
    rx = re.compile(regex, re.I)
    return any(rx.search(c) for r in t.rows for c in r)
