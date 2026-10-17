"""Module C part 2 - the remaining TP checklist items.

Class A  : automated rules (PASS / EXC / NOTRUN).
Class B/C: ASSISTED - the tool extracts the evidence and a prompt; status REVIEW, severity Info. The reviewer ticks or raises an exception.
"""
from __future__ import annotations
import re
from collections import Counter

from .rules_fs import NotRun
from .rules_tp import (TPCtx, _ex, _norm, _comp_rows, _printed_page_map, _num, quartiles, report_stats, tp21_annexures, tp17_index)

CUR = ("AED", "USD", "EUR", "GBP", "INR", "SAR")


def _rv(cid, title, detail, loc="Report", exp="", found=""):
    return _ex(cid, title, "Info", detail, loc, exp, found)


def _pdf(c):
    return c.doc.kind in ("pdf", "scan")


def _fs_text(c):
    if not c.fs:
        raise NotRun("Audited FS not supplied (or could not be read - scanned PDF without OCR)")
    return re.sub(r"[ \t]+", " ", c.fs.all_text)


def _pdf_pages(c):
    import pdfplumber
    return pdfplumber.open(str(c.path))


def _annex_start_page(c):
    for i, p in enumerate(c.doc.pages):
        top = "\n".join([l for l in p.splitlines() if l.strip()][:4])
        if re.search(r"^\s*(appendix|annexure)\b", top, re.I | re.M) and i > 2:
            return i
    return None


# ------------------------------------------------------------------ class A
def tp13_pagenos(c: TPCtx):
    if c.doc.kind == "docx":
        import zipfile
        z = zipfile.ZipFile(str(c.path))
        xml = z.read("word/document.xml").decode("utf8", "ignore")
        restarts = re.findall(r'<w:pgNumType[^>]*w:start="(\d+)"', xml)
        if len(restarts) > 1 or (restarts and restarts[0] != "1"):
            return [_ex("TP-13", "Page numbering restarts or starts at a number other than 1", "Medium", f"pgNumType start values {restarts}", "Section settings", "continuous 1..n", str(restarts))]
        return []
    pm = _printed_page_map(c)
    if not pm:
        raise NotRun("Printed page numbers not readable in the PDF")
    seq = [pm[k] for k in sorted(pm)]
    out, phys = [], sorted(pm)
    for a, b in zip(phys, phys[1:]):
        if pm[b] != pm[a] + (b - a):
            out.append(_ex("TP-13", "Page number sequence broken", "Medium", f"PDF page {a} shows {pm[a]}; page {b} shows {pm[b]}", f"PDF pages {a}-{b}", str(pm[a] + (b - a)), str(pm[b])))
    miss = [p for p in range(1, len(c.doc.pages) + 1) if p > 2 and p not in pm]
    if miss:
        out.append(_ex("TP-13", "Pages without a readable page number", "Low", f"PDF pages {miss[:10]}", "Header/footer", "number on every body page", "missing"))
    return out[:6]


def tp22_annex_sheet(c: TPCtx):
    if not c.wb:
        raise NotRun("Benchmarking workbook not supplied")
    import openpyxl
    from .legacy import ensure_modern
    names = []
    if getattr(c, "wb_path", None):
        _wbk = openpyxl.load_workbook(str(ensure_modern(c.wb_path)), read_only=True)
        names = _wbk.sheetnames
        _wbk.close()
    if not names:
        raise NotRun("Workbook sheet names not available")
    t = c.doc.all_text
    ann = re.findall(r"Annexure\s+(\d{1,2}|[A-Z])\s*[:\-–]\s*([^\n]{3,80})", t)
    if not ann:
        raise NotRun("No annexure list ('Annexure n: title') in the report")
    out = []
    for n, title in ann:
        if re.search(r"screenshot|write ?up", title, re.I):
            continue                                     # database screenshots are pictures, not sheets
        words = [w for w in re.findall(r"[a-z]{4,}", title.lower()) if w not in ("search", "jurisdiction", "summary")]
        hit = [s for s in names if re.search(rf"\b{re.escape(n)}\b", s) or any(w in s.lower() for w in words)]
        if not hit:
            out.append(_ex("TP-22", "Excel annexure has no matching worksheet in the workbook", "Medium", f"Annexure {n}: {title.strip()} - sheets {names[:8]}", "Appendix", "sheet named for the annexure", "none"))
    return out


def tp24_workbook(c: TPCtx):
    if not getattr(c, "wb_path", None):
        raise NotRun("Benchmarking workbook not supplied")
    from pathlib import Path
    nm = Path(c.wb_path).name.lower()
    out = []
    key = re.sub(r"[^a-z]", "", c.name.lower())[:6]
    if key not in re.sub(r"[^a-z]", "", nm) and not str(c.cy) in nm:
        out.append(_ex("TP-24", "Workbook file name does not identify the client or year", "Low", f"'{Path(c.wb_path).name}'", "Engagement file", f"{c.name} / FY{c.cy}", Path(c.wb_path).name))
    if not (c.wb and c.wb.comps):
        out.append(_ex("TP-24", "Workbook supplied but no comparables/PLI table recognised", "Medium", "Reader expects a company column and a PLI/amount column (check layout)", Path(c.wb_path).name, "recognised table", "none"))
    return out


def tp25_prior(c: TPCtx):
    t = _norm(c)
    if not re.search(r"financial update", t[:3000], re.I):
        raise NotRun("Applies to Financial Update reports only")
    n = len(re.findall(r"Transfer Pricing Benchmark(?:ing)? Analysis", t))
    if n < 3:
        return [_ex("TP-25", "Previous benchmarking report is not attached after the update", "High", "No second report (benchmark title/cover) found after the update", "Report", "previous report attached", "absent")]
    return []


def tp27_annex_border(c: TPCtx):
    if c.doc.kind == "docx":
        secs = c.extras()
        bad = [s["i"] for s in secs if not s["border"]]
        return [_ex("TP-27", "Section without page border (annexure pages)", "Low", f"Sections {bad}", "Sections", "border on all pages", "none")] if bad else []
    st = _annex_start_page(c)
    if st is None:
        raise NotRun("Annexure/appendix pages not found")
    out = []
    with _pdf_pages(c) as pdf:
        for i in range(st, len(pdf.pages)):
            pg = pdf.pages[i]
            big = [r for r in pg.rects if r["width"] > pg.width * 0.85 and r["height"] > pg.height * 0.8]
            edges = [l for l in pg.lines if l["width"] > pg.width * 0.85 or l["height"] > pg.height * 0.8]
            if not big and len(edges) < 4 and not _pixel_frame(pg):
                out.append(_ex("TP-27", "Annexure/appendix page has no page border", "Low", f"PDF page {i + 1}", f"Page {i + 1}", "border", "none"))
    return out[:6]


def tp28_annex_figs(c: TPCtx):
    if not (c.wb and getattr(c.wb, "all_numbers", None)):
        raise NotRun("Benchmarking workbook numbers not available")
    wn = c.wb.all_numbers
    rows = _comp_rows(c)
    if not rows:
        raise NotRun("No annexure/comparables table in the report")
    miss = [r for r in rows if not any(abs(r[4] - v) <= max(1, abs(r[4]) * 0.0005) or abs(r[4] / 1000 - v) < 0.6 for v in wn)]
    return [_ex("TP-28", "Table figures not found in the source workbook", "High", f"{len(miss)} of {len(rows)} rows, e.g. {miss[0][0]}: {miss[0][4]:,.0f}", "Comparables table", "agrees to workbook", f"{len(miss)} unmatched")] if miss else []


def tp29_fs_year(c: TPCtx):
    t = _fs_text(c)
    out = []
    if not re.search(r"independent auditor|auditor['’]s report|report on the audit", t, re.I):
        out.append(_ex("TP-29", "FS supplied has no auditor's report - may not be the finalised audited FS", "High", "No 'Independent auditor's report' found", "FS", "signed audited FS", "not found"))
    if re.search(r"\b(draft|unaudited)\b", t[:4000], re.I):
        out.append(_ex("TP-29", "FS marked draft/unaudited", "High", "Word 'draft' or 'unaudited' near the start", "FS", "final audited", "draft"))
    yrs = {int(y) for y in re.findall(r"(?:year|period) ended[^.\n]{0,25}?(20\d{2})", t, re.I)}
    if yrs and c.cy not in yrs:
        out.append(_ex("TP-29", "FS year does not agree to the engagement year", "High", f"FS years {sorted(yrs)}; engagement FY {c.cy}", "FS", str(c.cy), str(sorted(yrs))))
    if not yrs:
        out.append(_ex("TP-29", "Year of the FS could not be identified", "Low", "No 'year ended' wording found", "FS", "identifiable year", "unclear"))
    return out


def _fs_amount(t, rx):
    for m in re.finditer(rx + r"[^\n\d(-]{0,40}\(?(-?\d[\d,]*\.?\d*)\)?", t, re.I):
        v = float(m.group(1).replace(",", ""))
        if v:
            return v
    return None


def tp31_opcost(c: TPCtx):
    t = _fs_text(c)
    L = c.wb.labels if c.wb else {}
    cogs = _fs_amount(t, r"(?:cost of (?:sales|revenue|goods sold)|direct costs?)")
    opex = _fs_amount(t, r"(?:general and administrative(?: expenses)?|administrative expenses|operating expenses|selling and distribution(?: expenses)?)")
    if cogs is None or opex is None:
        raise NotRun("Cost of sales / operating expenses lines not found in the FS text")
    mine = cogs + opex
    out, tested = [], (L.get("oc") or (None,))[0]
    if tested is None:
        m = re.search(r"(?:total )?operating costs?[^.\d]{0,40}?(\d[\d,]{4,})", _norm(c), re.I)
        tested = float(m.group(1).replace(",", "")) if m else None
    if tested is None:
        raise NotRun("Tested-party operating cost not found in workbook or report")
    if abs(mine - tested) > max(1, 0.001 * mine):
        out.append(_ex("TP-31", "Operating costs differ from cost of sales + operating expenses per FS", "High", f"COGS {cogs:,.0f} + opex {opex:,.0f} = {mine:,.0f} vs {tested:,.0f}; difference {tested - mine:,.0f}. Check non-operating items / omissions / double counting.", "FS vs workbook", f"{mine:,.0f}", f"{tested:,.0f}"))
    nonop = re.findall(r"(finance (?:cost|income)|interest (?:expense|income)|exchange (?:gain|loss)|gain on disposal|impairment|income tax|zakat)", t, re.I)
    if nonop and out:
        out[0].detail += f" Non-operating captions in FS: {sorted({n.lower() for n in nonop})[:5]}."
    return out


def tp33_units(c: TPCtx):
    rt = _norm(c)
    out = []
    cur_r = Counter(re.findall(r"\b(AED|USD|EUR|GBP|INR|SAR)\b", rt))
    if c.fs:
        cur_f = Counter(re.findall(r"\b(AED|USD|EUR|GBP|INR|SAR)\b", c.fs.all_text))
        if cur_f and cur_r and cur_f.most_common(1)[0][0] != cur_r.most_common(1)[0][0]:
            out.append(_ex("TP-33", "Currency differs between report and FS", "High", f"Report {dict(cur_r.most_common(3))}; FS {dict(cur_f.most_common(3))}", "Report / FS", "same currency", "different"))
        sf = bool(re.search(r"in thousands|AED ?['’]?000|'000", c.fs.all_text, re.I))
        sr = bool(re.search(r"in thousands|AED ?['’]?000|'000", rt, re.I))
        if sf != sr:
            out.append(_ex("TP-33", "Scale (units vs thousands) differs between FS and report", "High", f"FS in thousands: {sf}; report in thousands: {sr}", "Report / FS", "same scale", "different"))
        yrs_f = {int(y) for y in re.findall(r"(?:year|period) ended[^.\n]{0,25}?(20\d{2})", c.fs.all_text, re.I)}
        if yrs_f and c.cy not in yrs_f:
            out.append(_ex("TP-33", "Period per FS differs from engagement year", "High", f"FS {sorted(yrs_f)} vs FY {c.cy}", "FS", str(c.cy), str(sorted(yrs_f))))
    if not c.fs and not out:
        out.extend(tp42_conventions(c))
    if not c.fs and not out:
        raise NotRun("FS not supplied - only report-internal consistency (TP-42) was tested")
    return out


def tp35_verified(c: TPCtx):
    L = c.wb.labels if c.wb else {}
    if not L:
        raise NotRun("Tested-party figures not found in the workbook (labels: operating profit, operating costs, revenue, PLI)")
    t = c.doc.all_text.replace(" ", " ")
    out = []
    for k, (v, src) in L.items():
        if k == "tested_pli":
            vs = [f"{v * (100 if c.wb.fraction else 1):.2f}", f"{v * (100 if c.wb.fraction else 1):.1f}"]
        else:
            vs = [f"{v:,.0f}", f"{v / 1000:,.0f}", f"{v / 1e6:,.1f}", f"{v:,.2f}"]
        if not any(x in t for x in vs):
            out.append(_ex("TP-35", f"Workbook figure for '{k}' not found in the report", "High", f"{v:,.2f} ({src}) - report does not state it (tried {vs[:3]})", src, f"{v:,.2f}", "absent"))
    return out


def tp40_pdf_annex(c: TPCtx):
    t = c.doc.all_text
    refs = sorted({m for m in re.findall(r"Annexure\s+(\d{1,2}|[A-Z])\b", t, re.I)})
    if not refs:
        raise NotRun("No annexures referred to")
    out = []
    ids = {h[0].upper() for h in re.findall(r"^\s*Annexure\s+(\d{1,2}|[A-Z])\s*(?:[:\-–].*)?$", t, re.M | re.I) for h in [h]} if False else set(re.findall(r"^\s*Annexure\s+(\d{1,2}|[A-Z])\s*(?:[:\-–].*)?$", t, re.M | re.I))
    if _pdf(c):
        st = _annex_start_page(c)
        if st is not None:
            tail = c.doc.pages[st:]
            empty = [st + i + 1 for i, p in enumerate(tail) if len(p.strip()) < 40]
            if empty:
                out.append(_ex("TP-40", "Annexure page has no readable text (picture)", "Low", f"PDF pages {empty[:8]}", "Annexures", "text-searchable annexure", "image only"))
    if c.doc.kind == "docx":
        import zipfile
        z = zipfile.ZipFile(str(c.path))
        emb = [n for n in z.namelist() if n.startswith("word/embeddings/")]
        if emb:
            out.append(_ex("TP-40", "Annexures are embedded objects", "Medium", f"{len(emb)} embedded object(s): {emb[:3]}", "Annexures", "visible annexures", "embedded"))
    return out


def tp42_conventions(c: TPCtx):
    t = _norm(c)
    out = []
    scales = {k for k, rx in {"mm": r"\bAED\s*mm\b|\(AED mm\)", "million": r"\bmillion\b|AED ?m\b", "thousand": r"thousand|'000|000s"}.items() if re.search(rx, t, re.I)}
    if len(scales) > 1:
        out.append(_ex("TP-42", "Mixed units of measurement in the report", "Medium", f"Scales used: {sorted(scales)}", "Report", "one convention", str(sorted(scales))))
    fy = Counter(int(y) for y in re.findall(r"\[?FY\s*(20\d{2})\]?", t))
    yo = Counter(int(y) for y in re.findall(r"financial year(?: ending| ended)?[^.\d]{0,15}(20\d{2})", t, re.I))
    ys = set(fy) | set(yo)
    if len(ys) > 1 and abs(max(ys) - min(ys)) >= 1 and not (c.cy in ys and (c.cy - 1) in ys and len(ys) == 2 and False):
        out.append(_ex("TP-42", "More than one financial year is referred to as the analysis period", "Medium", f"'FY' labels {dict(fy)}; 'financial year' sentences {dict(yo)}", "Report", "one period", str(sorted(ys))))
    curs = {k for k in CUR if re.search(rf"\b{k}\b", t)}
    if len(curs) > 1:
        out.append(_ex("TP-42", "More than one currency appears", "Low", f"{sorted(curs)}", "Report", "AED (state conversion basis if converted)", str(sorted(curs))))
    if re.search(r"historical rate", t, re.I) and not re.search(r"(exchange|conversion|converted)", t, re.I):
        out.append(_ex("TP-42", "Amounts converted at 'historical rate' but the conversion basis/rate is not explained", "Low", "Add the source and rate/date used for currency conversion", "Section 3.2", "stated basis", "not stated"))
    return out


def tp44_conclusion(c: TPCtx):
    t = c.doc.all_text
    if not re.search(r"^\s*(?:\d+\.?\s*)?(results? and conclusions?|conclusions?)\b", t, re.I | re.M):
        return [_ex("TP-44", "No 'Results and Conclusion' section found", "High", "Designated section missing or titled differently", "Report", "Results and Conclusion", "absent")]
    ok = any(re.search(r"lower quartile", " ".join(" ".join(r) for r in tb.rows), re.I) and re.search(r"upper quartile", " ".join(" ".join(r) for r in tb.rows), re.I)
             and re.search(r"actual|tested", " ".join(" ".join(r) for r in tb.rows), re.I) for tb in c.doc.tables)
    if not ok:
        return [_ex("TP-44", "Results table (Q1, median, Q3 and tested-party actual) not found", "High", "A table with lower/upper quartile and the actual figure is required in the conclusion", "Results section", "results table", "absent")]
    return []


# ------------------------------------------------------------------ class B / C : assisted
def _assist(fn):
    fn.assist = True
    return fn


@_assist
def tp04_cover(c: TPCtx):
    pic = False
    if c.doc.kind == "docx":
        import zipfile
        z = zipfile.ZipFile(str(c.path))
        pic = any(n.startswith("word/media/") for n in z.namelist())
    else:
        with _pdf_pages(c) as pdf:
            pic = len(pdf.pages[0].images) > 0
    return [_rv("TP-04", "Cover picture", ("A picture is on the cover/first pages. " if pic else "No picture detected on the cover. ") +
                "Confirm it is relevant to the client's activity (construction materials trading here) and not an AI image with garbled text.", "Cover", "relevant, clean picture", "picture present" if pic else "none")]


@_assist
def tp05_abbrev(c: TPCtx):
    t = _norm(c)
    defined = dict((m.group(2), m.group(1)) for m in re.finditer(r"([A-Z][A-Za-z\- ]{4,60}?)\s*\(\s*[“\"']?([A-Z]{2,6})[”\"']?\s*\)", t))
    used = Counter(re.findall(r"\b([A-Z]{2,6})\b", t))
    undefined = [(a, n) for a, n in used.most_common(40) if a not in defined and a not in ("AED", "UAE", "OECD", "LLC", "FZE", "FTA", "VAT", "PDF", "FY", "DRAFT", "USD", "EUR", "CEO", "CFO", "MD", "COO", "SVP", "KMP")]
    ev = f"Defined: {sorted(defined)}. Used but not defined in the text: {[a for a, _ in undefined[:12]]}."
    return [_rv("TP-05", "Abbreviation list for the reviewer", ev, "Report", "defined at first use", f"{len(undefined)} not defined")]


def _style_summary(c):
    if c.doc.kind != "docx":
        return None
    import docx
    d = docx.Document(str(c.path))
    al, sp = Counter(), Counter()
    for p in d.paragraphs:
        if p.text.strip():
            al[str(p.alignment)] += 1
            sp[str(p.paragraph_format.line_spacing)] += 1
    heads = Counter((p.style.name) for p in d.paragraphs if p.style.name.lower().startswith("heading"))
    return al, sp, heads


@_assist
def tp08_headings(c: TPCtx):
    if _pdf(c):
        sizes = Counter()
        with _pdf_pages(c) as pdf:
            for pg in pdf.pages[2:12]:
                for ch in pg.chars:
                    sizes[round(ch["size"] * 2) / 2] += 1
        big = sorted([s for s in sizes if s >= 12], reverse=True)
        return [_rv("TP-08", "Heading sizes seen (PDF)", f"Font sizes ≥12pt seen: {big}; expected main 19.5 pt, sub 12 pt. Check bold, numbering (1., 1.1) and hierarchy against the template.", "Report", "19.5 / 12 pt", str(big))]
    al, sp, heads = _style_summary(c)
    return [_rv("TP-08", "Heading styles (Word)", f"Heading styles used {dict(heads)}; paragraph alignment mix {dict(al)}", "Report")]


@_assist
def tp09_paragraph(c: TPCtx):
    s = _style_summary(c)
    if not s:
        return [_rv("TP-09", "Paragraph format", "PDF supplied: alignment/spacing cannot be read. Check justified body text, line spacing and spacing before/after visually or supply the Word file.", "Report")]
    al, sp, _ = s
    return [_rv("TP-09", "Paragraph format (Word)", f"Alignment mix {dict(al)}; line spacing values {dict(sp)}. Body is normally justified; one dominant spacing expected.", "Report")]


@_assist
def tp10_tables(c: TPCtx):
    n = len({tb.loc for tb in c.doc.tables if len(tb.rows) > 2 and len(tb.rows[0]) > 2})
    dec = Counter()
    for tb in c.doc.tables:
        for r in tb.rows:
            for x in r:
                if re.fullmatch(r"-?\d[\d,]*\.\d+", x.strip()):
                    dec[len(x.split(".")[1])] += 1
    return [_rv("TP-10", "Tables for formatting review", f"{n} data tables. Decimal places used in numeric cells: {dict(dec)} (one convention expected per column). Check header shading, borders and units (AED / AED mm).", "Tables", "one convention", str(dict(dec)))]


@_assist
def tp18_index_updated(c: TPCtx):
    try:
        f = tp17_index(c)
        ev = f"Index vs headings test (TP-17): {len(f)} exception(s). " + ("Index entries/pages disagree - update the index after final pagination." if f else "Entries and page numbers agree.")
    except NotRun as e:
        ev = f"Index could not be tested: {e}"
    return [_rv("TP-18", "Index currency", ev + " Word: right-click the index > Update field > entire table before PDF conversion.", "Index")]


@_assist
def tp19_continued(c: TPCtx):
    if c.doc.kind == "docx":
        txt = [p.text.strip() for p in c.doc.paragraphs if p.text.strip()]
        n = sum(1 for x in txt if "(continued)" in x.lower())
        return [_rv("TP-19", "Continuation labels", f"{n} headings carry '(continued)'. Confirm each follows a heading of the same name and none appears on a first page.", "Report")]
    pages, out = c.doc.pages, []
    for i, p in enumerate(pages):
        for m in re.finditer(r"^\s*(.{5,70}?)\s*\(continued\)\s*$", p, re.I | re.M):
            base = m.group(1).strip().lower()
            if i == 0 or base not in pages[i - 1].lower():
                out.append(f"p.{i + 1} '{m.group(1).strip()}'")
    n_all = sum(len(re.findall(r"\(continued\)", p, re.I)) for p in pages)
    return [_rv("TP-19", "Continuation labels", f"{n_all} '(continued)' labels. Without a matching start on the previous page: {out[:6] or 'none'}. Check no section that continues lacks the label.", "Report", "label only on continued headings", str(len(out)))]


@_assist
def tp26_objects(c: TPCtx):
    if c.doc.kind == "docx":
        import zipfile
        z = zipfile.ZipFile(str(c.path))
        emb = [n for n in z.namelist() if n.startswith(("word/embeddings/", "word/media/"))]
        return [_rv("TP-26", "Inserted objects", f"{len(emb)} embedded object(s)/image(s) {emb[:5]}. Open each to confirm it is visible, bordered and not clipped.", "Annexures")]
    with _pdf_pages(c) as pdf:
        pg = [i + 1 for i, p in enumerate(pdf.pages) if p.images]
    return [_rv("TP-26", "Inserted objects", f"Pictures/objects on PDF pages {pg}. Confirm readable, bordered and fit the page; embedded files (zip icon) are not part of a PDF.", "Annexures")]


@_assist
def tp32_staff(c: TPCtx):
    t = _fs_text(c)
    hits = []
    for m in re.finditer(r"(staff costs?|employee (?:benefits|costs?)|salaries[^\n]{0,20}|directors?['’]? remuneration|managerial remuneration|key management[^\n]{0,30}|end of service[^\n]{0,15}|allowances)[^\n\d(-]{0,30}\(?(-?\d[\d,]{3,})", t, re.I):
        hits.append(f"{m.group(1).strip()} = {m.group(2)}")
    return [_rv("TP-32", "Employee cost lines found in the FS", "; ".join(hits[:8]) or "no staff-cost line found with an amount - look in the notes", "FS notes", "all staff-related cost in employee cost", f"{len(hits)} lines")]


@_assist
def tp34_reclass(c: TPCtx):
    t = _norm(c)
    sn = re.findall(r"[^.]{0,80}\b(reclassif\w+|adjust(?:ed|ment)s? (?:to|for) (?:the )?(?:financ|operating)\w*|aggregat\w+)[^.]{0,80}\.", t, re.I)
    recon = bool(re.search(r"reconcil", t, re.I))
    return [_rv("TP-34", "Reclassifications / adjustments in the report", f"{len(sn)} statement(s) found; reconciliation mentioned: {recon}. " + ("; ".join(s[:100] for s in sn[:2]) if sn else "None stated - confirm no reclassifications were made."), "Financial analysis")]


@_assist
def tp36_trace(c: TPCtx):
    t = _fs_text(c)
    pages = [re.sub(r"\s+", " ", p) for p in c.fs.pages]
    figs = []
    for k, (v, src) in (c.wb.labels.items() if c.wb else []):
        if k == "tested_pli":
            continue
        s = f"{v:,.0f}"
        pg = [i + 1 for i, p in enumerate(pages) if s in p]
        figs.append(f"{k} {s}: FS page {pg[:3] if pg else 'NOT FOUND'}")
    if not figs:
        for tb in c.doc.tables:
            for r in tb.rows:
                for x in r:
                    if re.fullmatch(r"\d{1,3}(,\d{3}){2,}", x.strip()) and len(figs) < 8:
                        pg = [i + 1 for i, p in enumerate(pages) if x.strip() in p]
                        figs.append(f"{x.strip()}: FS page {pg[:3] if pg else 'not found'}")
    return [_rv("TP-36", "Figure-to-FS trace", "; ".join(figs) or "no tested-party figures found to trace", "Report vs FS", "every figure traced", f"{sum(1 for f in figs if 'NOT' in f.upper())} untraced")]


@_assist
def tp41_titles(c: TPCtx):
    rows = _comp_rows(c)
    if not rows:
        raise NotRun("No comparables table with job titles recognised")
    sample = "; ".join(f"{r[0]} | {r[2]} | {r[4]:,.0f}" for r in rows[:6])
    return [_rv("TP-41", "Job titles and values for tick-back to annexures", f"{len(rows)} rows. First rows: {sample}. Annexures are not readable here (zip/embedded): tick every row to Annexure 7/9.", "Final set", "title and value agree to annexure", f"{len(rows)} rows to tick")]


@_assist
def tp48_narrative(c: TPCtx):
    t = _norm(c)
    sents = re.findall(r"[^.]{0,120}(?:within|outside|above|below|exceeds?|not aligned)[^.]{0,80}(?:range|benchmark)[^.]*\.", t, re.I)
    ind = re.findall(r"(?:industry|sector)[^.]{0,60}", t, re.I)
    return [_rv("TP-48", "Conclusion statements to confirm against the results", "Position statements: " + " | ".join(s.strip()[:160] for s in sents[:3]) + f". Industry descriptions: {sorted({i.strip()[:50] for i in ind})[:4]}. Check wording agrees with the numbers and method.", "Conclusion")]


@_assist
def tp51_charts(c: TPCtx):
    if _pdf(c):
        with _pdf_pages(c) as pdf:
            pg = [i + 1 for i, p in enumerate(pdf.pages) if p.images]
        return [_rv("TP-51", "Charts to check", f"Picture/chart objects on PDF pages {pg}. Agree each bar/line value to the table beside it (e.g. screening funnel 63 / 36 / 27) and check axis, labels and period.", "Charts")]
    import zipfile
    z = zipfile.ZipFile(str(c.path))
    ch = [n for n in z.namelist() if n.startswith("word/charts/")]
    return [_rv("TP-51", "Charts to check", f"{len(ch)} chart part(s). Agree values to the workings; check axis scale, labels, period.", "Charts")]


@_assist
def tp56_corroborate(c: TPCtx):
    rows = _comp_rows(c)
    if not rows:
        raise NotRun("No comparables table recognised")
    co = sorted({r[1] for r in rows})
    return [_rv("TP-56", "Companies to corroborate with the database", f"{len(co)} companies / {len(rows)} entries to tick to S&P Capital IQ: {', '.join(co[:10])}{'…' if len(co) > 10 else ''}. Check descriptions, financials and the executive compensation year.", "Final set", "agree to database", f"{len(co)} companies")]


@_assist
def tp39_corrob(c: TPCtx):
    t = _norm(c)
    m = re.findall(r"[^.]{0,100}corroborat\w+[^.]{0,100}\.", t, re.I)
    return [_rv("TP-39", "Corroborative method", ("Statement: " + m[0].strip()[:220]) if m else "The report does not mention a corroborative (secondary) method. Confirm whether the firm's methodology requires one for this transaction (e.g. a second method or internal comparable).", "Section 3", "stated and applied", "stated" if m else "not mentioned")]


@_assist
def tp43_differences(c: TPCtx):
    return [_rv("TP-43", "Differences for investigation", "See exceptions raised by TP-33, TP-35, TP-42, TP-46, TP-50, TP-64 in this paper: for each, record whether the cause is rounding, units or currency and whether it is resolved or escalated.", "Review paper")]


@_assist
def tp52_final(c: TPCtx):
    return [_rv("TP-52", "Final file formats", f"This review used a {c.doc.kind.upper()} file. " + ("Formatting checks on borders, fonts, headers need the Word file; run again with the Word version and compare with the PDF." if c.doc.kind != "docx" else "Compare the PDF made from this Word file page by page (pagination, borders, objects)."), "Final files")]


@_assist
def tp53_open(c: TPCtx):
    return [_rv("TP-53", "Exceptions status", "Every exception in this paper must be marked Cleared / Waived (Partner) / Disputed in the exception log before sign-off.", "Review paper")]


@_assist
def tp54_files(c: TPCtx):
    from pathlib import Path
    nm = Path(c.path).name
    flags = []
    if re.search(r"draft|v\d|copy|\(\d\)|review|final_final", nm, re.I):
        flags.append("file name contains draft/version/copy markers")
    if re.sub(r"[^a-z]", "", c.name.lower())[:6] not in re.sub(r"[^a-z]", "", nm.lower()):
        flags.append("client name not in the file name")
    return [_rv("TP-54", "Saved file name", f"'{nm}': " + ("; ".join(flags) if flags else "naming looks consistent") + ". Confirm saving per the document management procedure.", "Engagement file")]


@_assist
def tp59_downward(c: TPCtx):
    t = _norm(c)
    kw = {"rationale": r"rational", "benchmarking": r"benchmark", "FS reconciliation": r"reconcil", "corresponding adjustment": r"corresponding adjustment"}
    have = {k: bool(re.search(v, t, re.I)) for k, v in kw.items()}
    dn = bool(re.search(r"downward", t, re.I)) or (c.ct_adj is not None and c.ct_adj < 0)
    return [_rv("TP-59", "Downward adjustment (CTP011)", ("Downward adjustment indicated. " if dn else "No downward adjustment indicated (upward/none) - applies only if the return reduces income. ") + f"Prerequisites mentioned in the report: {have}.", "Return / TP file")]


ASSIST_REGISTRY = {"TP-04": tp04_cover, "TP-05": tp05_abbrev, "TP-08": tp08_headings, "TP-09": tp09_paragraph, "TP-10": tp10_tables,
                   "TP-18": tp18_index_updated, "TP-19": tp19_continued, "TP-26": tp26_objects, "TP-32": tp32_staff, "TP-34": tp34_reclass,
                   "TP-36": tp36_trace, "TP-39": tp39_corrob, "TP-41": tp41_titles, "TP-43": tp43_differences, "TP-48": tp48_narrative,
                   "TP-51": tp51_charts, "TP-52": tp52_final, "TP-53": tp53_open, "TP-54": tp54_files, "TP-56": tp56_corroborate, "TP-59": tp59_downward}
A_REGISTRY = {"TP-13": tp13_pagenos, "TP-22": tp22_annex_sheet, "TP-24": tp24_workbook, "TP-25": tp25_prior, "TP-27": tp27_annex_border,
              "TP-28": tp28_annex_figs, "TP-29": tp29_fs_year, "TP-31": tp31_opcost, "TP-33": tp33_units, "TP-35": tp35_verified,
              "TP-40": tp40_pdf_annex, "TP-42": tp42_conventions, "TP-44": tp44_conclusion}


# ------------------------------------------------------------------ PDF versions of the formatting checks (so a PDF report is checked, not 'Not run')
def _pdf_bands(c):
    """per body page (index >= 2): header words, footer words, page width/height."""
    out = []
    with _pdf_pages(c) as pdf:
        for i, pg in enumerate(pdf.pages):
            if i < 2:
                continue
            ws = pg.extract_words(keep_blank_chars=False, use_text_flow=False)
            hdr = [w for w in ws if w["top"] < pg.height * 0.055]
            ftr = [w for w in ws if w["bottom"] > pg.height * 0.945]
            if ftr:                                  # footer = the last text line only
                lt = max(w["top"] for w in ftr)
                ftr = [w for w in ftr if w["top"] >= lt - 3]
            out.append((i + 1, hdr, ftr, pg.width, pg))
    return out


def _txt(ws):
    return " ".join(w["text"] for w in sorted(ws, key=lambda w: (round(w["top"]), w["x0"])))


def pdf_font(c):
    bad = Counter()
    for k, v in c.doc.pdf_fonts.items():
        name = re.sub(r"^[A-Z]{6}\+", "", k)
        if re.search(r"symbol|wingding|dingbat|bullet|marlett|math", name, re.I):
            continue
        if not re.search(r"times", name, re.I):
            bad[name] += v
    tot = sum(c.doc.pdf_fonts.values()) or 1
    return [_ex("TP-06", f"Font other than Times New Roman in the PDF: {n}", "Low", f"{v} characters ({v / tot:.0%})", "PDF fonts", "Times New Roman", n) for n, v in bad.most_common(4) if v / tot > 0.005]


def pdf_size(c):
    sz = c.doc.pdf_sizes
    if not sz:
        raise NotRun("No text layer to measure")
    body = sz.most_common(1)[0][0]
    exp = c.p["heading_sizes"]["body"]
    out = []
    if abs(body - exp) > 0.6:
        out.append(_ex("TP-07", f"Dominant body size is {body:g} pt (standard {exp:g} pt)", "Low", "Check the PDF was not scaled when printed/converted", "PDF text", f"{exp:g} pt", f"{body:g} pt"))
    big = [s for s, n in sz.items() if s >= 16 and n >= 20]
    if big and all(abs(s - c.p["heading_sizes"]["main"]) > 0.6 for s in big):
        out.append(_ex("TP-07", f"Main heading size {sorted(big)} pt (standard {c.p['heading_sizes']['main']:g} pt)", "Low", "Check scaling", "PDF headings", f"{c.p['heading_sizes']['main']:g} pt", str(sorted(big))))
    return out


def pdf_header(c):
    out = []
    for pno, hdr, ftr, w, pg in _pdf_bands(c):
        t = _txt(hdr).lower()
        if not any(x.lower() in t for x in c.p["titles"]):
            out.append(_ex("TP-11", "Header does not show the prescribed title", "Low", f"'{_txt(hdr)[:60]}'", f"PDF p.{pno}", "; ".join(c.p["titles"][:2]), _txt(hdr)[:40] or "(blank)"))
    return out[:6]


def pdf_pageno(c):
    out = []
    for pno, hdr, ftr, w, pg in _pdf_bands(c):
        nums = [x for x in hdr if re.fullmatch(r"\d{1,3}", x["text"])]
        if not nums:
            out.append(_ex("TP-12", "No page number in the header", "Low", "", f"PDF p.{pno}", "number at right of header", "none"))
        elif nums[-1]["x0"] < w * 0.8:
            out.append(_ex("TP-12", "Page number is not at the right of the header", "Low", f"x={nums[-1]['x0']:.0f} of {w:.0f}", f"PDF p.{pno}", "right", "left/centre"))
    return out[:6]


def pdf_footer(c):
    out = []
    for pno, hdr, ftr, w, pg in _pdf_bands(c):
        t = _txt(ftr)
        if c.name.lower() not in t.lower():
            out.append(_ex("TP-14", "Footer does not show the company name as per the trade licence", "Medium", f"'{t[:60]}'", f"PDF p.{pno}", c.name, t[:40] or "(blank)"))
        elif min(x["x0"] for x in ftr) < w * 0.5:
            out.append(_ex("TP-14", "Company name is not at the right of the footer", "Low", f"'{t[:60]}'", f"PDF p.{pno}", "right", "left/centre"))
    return out[:6]


def pdf_consistency(c):
    h, f = {}, {}
    for pno, hdr, ftr, w, pg in _pdf_bands(c):
        h.setdefault(re.sub(r"\d+", "#", _txt(hdr)), []).append(pno)
        f.setdefault(re.sub(r"\d+", "#", _txt(ftr)), []).append(pno)
    out = []
    for lab, d in (("Header", h), ("Footer", f)):
        if len(d) > 1:
            out.append(_ex("TP-15", f"{lab} text differs between pages", "Low", " | ".join(f"'{k[:40]}' on pages {v[:4]}" for k, v in d.items())[:240], "PDF pages", "identical", f"{len(d)} variants"))
    return out


def _pixel_frame(pg):
    """frame drawn as paths/images: look for a long dark vertical run near the left edge and a horizontal run near the top."""
    try:
        im = pg.to_image(resolution=40).original.convert("L")
    except Exception:
        return False
    w, h = im.size
    px = im.load()
    v = any(sum(1 for y in range(h) if px[x, y] < 110) > h * 0.8 for x in range(int(w * 0.02), int(w * 0.12)))
    hz = any(sum(1 for x in range(w) if px[x, y] < 110) > w * 0.8 for y in range(int(h * 0.02), int(h * 0.12)))
    return v and hz


def pdf_border(c):
    out = []
    for pno, hdr, ftr, w, pg in _pdf_bands(c):
        big = [r for r in pg.rects if r["width"] > pg.width * 0.85 and r["height"] > pg.height * 0.8]
        edges = [l for l in pg.lines if l["width"] > pg.width * 0.85 or l["height"] > pg.height * 0.8]
        if not big and len(edges) < 4 and not _pixel_frame(pg):
            out.append(_ex("TP-16", "No page border on this page", "Low", "", f"PDF p.{pno}", "1.5 pt border", "none"))
        else:
            lw = (big[0].get("linewidth") if big else (edges[0].get("linewidth") if edges else None))
            if lw and abs(lw - 1.5) > 0.4 and pno == 3:
                out.append(_ex("TP-16", f"Border line thickness {lw:.2f} pt (standard 1.5 pt)", "Low", "measured on page 3", f"PDF p.{pno}", "1.5 pt", f"{lw:.2f} pt"))
    return out[:6]


def _dual(docx_fn, pdf_fn):
    def run(c):
        return docx_fn(c) if c.doc.kind == "docx" else pdf_fn(c)
    return run


def install_pdf_variants(reg):
    for cid, pf in (("TP-06", pdf_font), ("TP-07", pdf_size), ("TP-11", pdf_header), ("TP-12", pdf_pageno), ("TP-14", pdf_footer), ("TP-15", pdf_consistency), ("TP-16", pdf_border)):
        reg[cid] = _dual(reg[cid], pf)
    return reg
