"""Engagement fact sheet / liquidation tracker loaders: Excel (primary), YAML (still supported).
Excel layout: column A = field key, column B = value, column C = guidance. Lists are separated by ';'.
Both formats produce the same dict, so the rules engine does not care which was used."""
from __future__ import annotations
from datetime import date, datetime
from pathlib import Path

from .models import load_yaml

ENG_FIELDS = [  # key, example, guidance, kind
    ("fy", 2025, "Current reporting year", "int"),
    ("py", 2024, "Prior year", "int"),
    ("client.name", "GOLDEN CREST TRADING DMCC", "Exactly as on the trade licence, incl. capitalisation", "str"),
    ("client.legal_status", "FZC", "FZC | FZE | LLC | SOLE | LLP | CIVIL | OTHER", "str"),
    ("client.authority", "DMCC", "Licensing authority named in the auditor's report", "str"),
    ("client.emirate", "Dubai", "", "str"),
    ("client.year_end", "31 December", "Day Month", "str"),
    ("client.shareholders", "Full Name 1; Full Name 2", "Per MOA - separate with ;", "list"),
    ("client.signatories_legal", "Full Name 1", "Exactly as passport/MOA - separate with ;", "list"),
    ("fs_signatories", "Full Name 1", "Who actually signed the FS - separate with ;", "list"),
    ("br_on_file", "No", "Yes if a Board Resolution authorises a signatory who is not all shareholders", "bool"),
    ("allowed_years", "", "e.g. incorporation year so it is not flagged as a stray year - separate with ;", "intlist"),
    ("review_log.preparer.done", "Yes", "Workflow status", "bool"),
    ("review_log.preparer.date", "2026-03-01", "dd/mm/yyyy or yyyy-mm-dd", "datestr"),
    ("review_log.supervisor.done", "Yes", "", "bool"),
    ("review_log.supervisor.date", "2026-03-05", "", "datestr"),
    ("review_log.manager.done", "No", "", "bool"),
    ("review_log.manager.date", "", "", "datestr"),
    ("review_log.ready_before_print.done", "No", "", "bool"),
    ("review_log.ready_before_print.date", "", "", "datestr"),
]

LIQ_FIELDS = [
    ("company", "COMPANY NAME", "As on trade licence", "str"),
    ("authority", "dmcc", "mainland | difc | dso | dmcc | rakez | jafza | dda", "str"),
    ("br.filed_with_authority", "No", "Board Resolution filed with the authority?", "bool"),
    ("br.liquidator_named_correctly", "No", "Liquidator named correctly in the BR?", "bool"),
    ("dates.board_resolution", "", "Leave blank until it happens", "date"),
    ("dates.liquidator_appointment", "", "", "date"),
    ("dates.fta_notified", "", "", "date"),
    ("dates.vat_dereg_initiated", "", "", "date"),
    ("dates.newspaper_ads", "", "Every publication date, separate with ;", "datelist"),
    ("dates.report_issue_planned", "", "", "date"),
    ("dates.report_submitted", "", "", "date"),
    ("dates.trade_licence_expiry", "", "", "date"),
    ("dates.bank_closure", "", "", "date"),
    ("partner_approval_short_gap", "No", "DSO: Yes if the Partner approved a report inside the 14-day gap", "bool"),
    # --- answers used by LQ-09, LQ-12, LQ-15 to LQ-18 (leave blank = not yet confirmed -> shown as Not checked)
    ("partners.all_informed", "", "Yes / No: all partners/shareholders aware of the liquidation", "str"),
    ("partners.dispute", "", "Yes / No: any legal dispute between partners", "str"),
    ("dates.ct_dereg_initiated", "", "Date CT deregistration was initiated on EmaraTax", "date"),
    ("vat.returns_filed_up_to_date", "", "Yes / No", "str"),
    ("vat.penalties_outstanding", "", "Yes / No", "str"),
    ("ct.returns_filed_up_to_date", "", "Yes / No", "str"),
    ("ct.penalties_outstanding", "", "Yes / No", "str"),
    ("auth.physical_submission_done", "", "DMCC: Yes / No", "str"),
    ("auth.portal_balance_cleared", "", "DIFC: Yes / No", "str"),
    ("auth.psa_refund_cleared", "", "DIFC: Yes / No", "str"),
    ("auth.advert_one_file", "", "Advertisement uploaded by liquidator as one file: Yes / No", "str"),
    ("report.built_from_tb", "", "Report built from TB/GL/SOFP/income statement: Yes / No", "str"),
    ("report.adjustments_posted", "", "Adjustments posted before the report: Yes / No", "str"),
    ("fta.final_letter_date", "", "Date of the FTA final communication", "date"),
    ("fta.final_letter_status", "", "cleared | pending | penalty", "str"),
    ("fta.status_in_report", "", "Status stated in the liquidation report: cleared | pending | penalty", "str"),
    ("legal.open_claims", "", "Yes / No: legal claims or contingent liabilities open", "str"),
    ("legal.settled_or_provided", "", "Yes / No: settled or provided for", "str"),
    ("liquidator.name", "", "Name of the liquidator (checked against the BR, SOA and report)", "str"),
    ("shareholders", "", "Shareholders / partners, separate with ;  (checked against the BR)", "list"),
]
OPTIONAL_PREFIX = ("liquidator.", "shareholders", "partners.", "vat.", "ct.", "auth.", "report.", "fta.", "legal.")   # added later: older trackers still load
NOCS = ["telecom", "utilities", "chiller", "landlord", "po_box", "vat_deregistration", "ct_deregistration",
        "authority_portal", "labour_visa", "establishment_card", "esr_aml", "customs"]


# ---------------------------------------------------------------- conversions
def _blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _to_date(v):
    if _blank(v):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    raise ValueError(f"Cannot read '{s}' as a date. Use dd/mm/yyyy or yyyy-mm-dd.")


def _convert(v, kind):
    if kind == "str":
        return "" if _blank(v) else str(v).strip()
    if kind == "int":
        return int(float(v))
    if kind == "bool":
        return str(v).strip().lower() in ("yes", "y", "true", "1")
    if kind == "datestr":
        d = _to_date(v)
        return d.isoformat() if d else None
    if kind == "date":
        return _to_date(v)
    parts = [] if _blank(v) else [x.strip() for x in str(v).replace("\n", ";").split(";") if x.strip()]
    if kind == "list":
        return parts
    if kind == "intlist":
        return [int(float(x)) for x in parts] if not isinstance(v, (int, float)) else [int(v)]
    if kind == "datelist":
        if isinstance(v, (date, datetime)):
            return [_to_date(v)]
        return [_to_date(x) for x in parts]
    raise ValueError(kind)


def _set(d, dotted, val):
    keys = dotted.split(".")
    for k in keys[:-1]:
        d = d.setdefault(k, {})
    d[keys[-1]] = val


def _sheet_dict(ws):
    return {str(r[0]).strip(): r[1] for r in ws.iter_rows(min_row=2, values_only=True)
            if r and r[0] is not None and str(r[0]).strip()}


def _load_kv(path, fields, sheet0):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb[sheet0] if sheet0 in wb.sheetnames else wb.worksheets[0]
    raw = _sheet_dict(ws)
    out, missing = {}, []
    for key, _, _, kind in fields:
        if key not in raw and kind not in ("date", "datelist", "datestr", "intlist") and not key.startswith(OPTIONAL_PREFIX):
            missing.append(key)
        _set(out, key, _convert(raw.get(key), kind))
    if missing:
        raise ValueError("Fact sheet is missing field(s): " + ", ".join(missing) + ". Use the latest template.")
    return out, wb


def load_engagement_xlsx(path):
    eng, _ = _load_kv(path, ENG_FIELDS, "Fact Sheet")
    for k, v in eng["review_log"].items():
        if not v.get("date"):
            v.pop("date", None)
    if not eng["client"]["name"] or not eng["fy"]:
        raise ValueError("Client name and FY are mandatory in the fact sheet")
    eng.setdefault("py", eng["fy"] - 1)
    return eng


def load_liquidation_xlsx(path):
    inp, wb = _load_kv(path, LIQ_FIELDS, "Liquidation")
    inp["authority"] = inp["authority"].lower()
    inp["nocs"] = {}
    if "NOCs" in wb.sheetnames:
        raw = _sheet_dict(wb["NOCs"])
        for n in NOCS:
            v = raw.get(n)
            inp["nocs"][n] = "na" if str(v).strip().lower() in ("na", "n/a") else _to_date(v)
    inp["bank_accounts"] = []
    if "Banks" in wb.sheetnames:
        for r in wb["Banks"].iter_rows(min_row=2, values_only=True):
            if r and not _blank(r[0]):
                inp["bank_accounts"].append({"bank": str(r[0]).strip(), "location": str(r[1] or "UAE").strip(),
                                             "no_liability_letter": _to_date(r[2] if len(r) > 2 else None),
                                             "closure_letter": _to_date(r[3] if len(r) > 3 else None)})
    return inp


def load_input(path, kind):
    """kind = 'engagement' | 'liquidation'; accepts .xlsx/.xlsm or .yaml/.yml."""
    ext = Path(path).suffix.lower()
    if ext in (".yaml", ".yml"):
        return load_yaml(path)
    from .legacy import ensure_modern
    path = ensure_modern(path)
    ext = Path(path).suffix.lower()
    if kind == "engagement":
        if ext in (".xlsx", ".xlsm"):
            try:
                return load_engagement_xlsx(path)
            except ValueError:                      # not the template layout: read label/value pairs
                return load_engagement_document(path)
        if ext in (".docx", ".pdf", ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp"):
            return load_engagement_document(path)
    elif ext in (".xlsx", ".xlsm"):
        return load_liquidation_xlsx(path)
    raise ValueError(f"Unsupported file {ext} for the {kind} input")


# ---------------------------------------------------------------- template writers
def _kv_sheet(ws, fields, title):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws.append(["Field", "Value", "Guidance"])
    for k, ex, g, _ in fields:
        ws.append([k, ex, g])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0B1F3A")
    for r in ws.iter_rows(min_row=2):
        r[0].font = Font(bold=True, color="0B1F3A")
        r[1].fill = PatternFill("solid", fgColor="FFF4DB")
        r[2].font = Font(italic=True, color="666666")
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 38; ws.column_dimensions["B"].width = 42; ws.column_dimensions["C"].width = 70
    ws.freeze_panes = "A2"


def make_templates(folder):
    import openpyxl
    folder = Path(folder); folder.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Fact Sheet"
    _kv_sheet(ws, ENG_FIELDS, "Fact Sheet")
    wb.save(folder / "Engagement_Fact_Sheet.xlsx")

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Liquidation"
    _kv_sheet(ws, LIQ_FIELDS, "Liquidation")
    n = wb.create_sheet("NOCs")
    n.append(["NOC", "Date obtained / NA", "Guidance"])
    for k in NOCS:
        n.append([k, "", "Date = obtained; NA = not applicable; blank = pending"])
    b = wb.create_sheet("Banks")
    b.append(["Bank", "Location", "No-liability letter (date)", "Closure letter (date)"])
    b.append(["Bank name", "UAE", "", ""])
    from openpyxl.styles import Font, PatternFill
    for sh in (n, b):
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0B1F3A")
        for col, w in zip("ABCD", (26, 24, 28, 24)):
            sh.column_dimensions[col].width = w
    wb.save(folder / "Liquidation_Tracker.xlsx")


if __name__ == "__main__":
    make_templates(Path(__file__).resolve().parent.parent / "templates")


def _flat(d, prefix=""):
    for k, v in d.items():
        if isinstance(v, dict):
            yield from _flat(v, f"{prefix}{k}.")
        else:
            yield f"{prefix}{k}", v


def _cell(v):
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, list):
        return "; ".join(x.isoformat() if isinstance(x, date) else str(x) for x in v)
    return "" if v is None else v


def dict_to_xlsx(d, path, kind):
    """Write a dict (e.g. loaded from a YAML sample) into the Excel template layout."""
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.title = "Fact Sheet" if kind == "engagement" else "Liquidation"
    ws.append(["Field", "Value", "Guidance"])
    flat = dict(_flat({k: v for k, v in d.items() if k not in ("nocs", "bank_accounts")}))
    fields = ENG_FIELDS if kind == "engagement" else LIQ_FIELDS
    for k, *_ in fields:
        ws.append([k, _cell(flat.get(k))])
    if kind == "liquidation":
        n = wb.create_sheet("NOCs"); n.append(["NOC", "Date obtained / NA"])
        for k in NOCS:
            v = d.get("nocs", {}).get(k)
            n.append([k, "NA" if str(v).lower() == "na" else _cell(v)])
        b = wb.create_sheet("Banks"); b.append(["Bank", "Location", "No-liability letter (date)", "Closure letter (date)"])
        for r in d.get("bank_accounts", []):
            b.append([r["bank"], r.get("location", "UAE"), r.get("no_liability_letter"), r.get("closure_letter")])
    wb.save(str(path))


# ---------------------------------------------------------------- engagement fact sheet in ANY format
import re as _re

ENG_LABELS = {
    "fy": "Financial year", "py": "Prior year", "client.name": "Client name (as on trade licence)",
    "client.legal_status": "Legal status", "client.authority": "Licensing authority", "client.emirate": "Emirate",
    "client.year_end": "Year end", "client.shareholders": "Shareholders", "client.signatories_legal": "Legal signatories",
    "fs_signatories": "FS signatories", "br_on_file": "Board resolution on file", "allowed_years": "Allowed years",
    "review_log.preparer.done": "Preparer done", "review_log.preparer.date": "Preparer date",
    "review_log.supervisor.done": "Supervisor done", "review_log.supervisor.date": "Supervisor date",
    "review_log.manager.done": "Manager done", "review_log.manager.date": "Manager date",
    "review_log.ready_before_print.done": "Ready before print done", "review_log.ready_before_print.date": "Ready before print date",
}


def _norm(s):
    return _re.sub(r"[^a-z0-9]+", " ", str(s).lower().replace("(as on trade licence)", "")).strip()


_ALIAS = {}
for _k, _lab, *_ in [(k, ENG_LABELS[k]) for k in ENG_LABELS]:
    _ALIAS[_norm(_k)] = _k
    _ALIAS[_norm(_lab)] = _k
_ALIAS.update({"company name": "client.name", "company": "client.name", "name": "client.name", "fy": "fy", "year": "fy",
               "reporting year": "fy", "authority": "client.authority", "board resolution": "br_on_file"})
_KIND = {k: kind for k, _, _, kind in ENG_FIELDS}


def _pairs_from_document(doc):
    """(label, value) pairs from table rows (col1|col2) and 'Label: value' / 'Label - value' text lines."""
    pairs = []
    for t in doc.tables:
        for r in t.rows:
            cells = [c.strip() for c in r if c and c.strip()]
            if len(cells) >= 2:
                pairs.append((cells[0], cells[1]))
    for ln in doc.text.splitlines():
        m = _re.match(r"^\s*([A-Za-z][A-Za-z0-9 ./()_-]{1,60}?)\s*[:=]\s*(.+?)\s*$", ln)
        if m:
            pairs.append((m.group(1), m.group(2)))
    return pairs


def load_engagement_document(path):
    """Fact sheet supplied as Word / PDF / image (OCR) / Excel in the standard label layout."""
    from .loaders import load_document
    doc = load_document(path)
    found = {}
    for lab, val in _pairs_from_document(doc):
        k = _ALIAS.get(_norm(lab))
        if k and k not in found:
            found[k] = val
    need = [k for k in ("fy", "client.name", "client.legal_status", "client.authority") if k not in found]
    if need:
        raise ValueError("Could not find these fact-sheet fields in " + Path(path).name + ": "
                         + ", ".join(ENG_LABELS[k] for k in need)
                         + ". Use the supplied fact-sheet layout (Field | Value, or 'Field: value' lines).")
    out = {}
    for key, _, _, kind in ENG_FIELDS:
        if key == "py" and key not in found:
            continue
        _set(out, key, _convert(found.get(key), kind))
    for v in out["review_log"].values():
        if not v.get("date"):
            v.pop("date", None)
    out.setdefault("py", out["fy"] - 1)
    out["_source"] = {"format": doc.kind, "ocr": doc.kind in ("image", "scan")}
    return out


def make_word_template(path):
    from docx import Document
    d = Document()
    d.add_heading("Engagement Fact Sheet", 1)
    d.add_paragraph("Fill the Value column from the trade licence, MOA/AOA, passports/EID and Board Resolution - not from the FS under review. "
                    "Separate lists with ';'. Flags: Yes / No. Dates: dd/mm/yyyy.")
    t = d.add_table(rows=1, cols=2); t.style = "Table Grid"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Field", "Value"
    for k, ex, g, _ in ENG_FIELDS:
        r = t.add_row().cells
        r[0].text = ENG_LABELS[k]; r[1].text = str(ex)
    d.save(str(path))
