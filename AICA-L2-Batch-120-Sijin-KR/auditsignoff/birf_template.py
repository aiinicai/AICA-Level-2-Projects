"""Builds the BIRF (Benchmarking Information Request Form) Excel template that the TP review reads back."""
from __future__ import annotations
from pathlib import Path


def build_birf_template(path):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BIRF"
    navy, grey = "1F3A5F", "EEF2F7"
    thin = Side(style="thin", color="B8C2CC")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "BENCHMARKING INFORMATION REQUEST FORM (BIRF) - Managerial remuneration / payments to connected persons"
    ws["A1"].font = Font(bold=True, size=13, color=navy)
    for i, lab in enumerate(("Company name (as per trade licence)", "Financial year / period", "Report type (Full benchmarking / Financial update)", "Completed by / date"), start=3):
        ws.cell(i, 1, lab).font = Font(bold=True)
        ws.cell(i, 3).fill = PatternFill("solid", fgColor="FFF8DC")
        ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=6)
    hdr = ["Recipient name", "Designation", "Connected person type", "Role / time commitment (%)", "Salary", "Bonus", "Allowances",
           "End-of-service benefits", "Share-based pay", "Benefits in kind", "Total (AED)"]
    h = 8
    ws.cell(h - 1, 1, "Amounts in AED for the financial year. One row per recipient. Leave a cell blank if not applicable.").font = Font(italic=True, color="666666")
    for j, t in enumerate(hdr, 1):
        c = ws.cell(h, j, t)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=navy); c.border = box
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[h].height = 34
    first, last = h + 1, h + 12
    for r in range(first, last + 1):
        for j in range(1, 12):
            c = ws.cell(r, j); c.border = box
            if j >= 5:
                c.number_format = "#,##0"
        ws.cell(r, 11, f'=IF(COUNT(E{r}:J{r})=0,"",SUM(E{r}:J{r}))').number_format = "#,##0"
        ws.cell(r, 11).font = Font(bold=True)
    tr = last + 1
    ws.cell(tr, 1, "Total").font = Font(bold=True)
    for j in range(5, 12):
        col = openpyxl.utils.get_column_letter(j)
        c = ws.cell(tr, j, f"=SUM({col}{first}:{col}{last})"); c.font = Font(bold=True); c.number_format = "#,##0"
    for j in range(1, 12):
        ws.cell(tr, j).fill = PatternFill("solid", fgColor=grey); ws.cell(tr, j).border = box
    dv = DataValidation(type="list", formula1='"Director,Shareholder,Relative of shareholder,Relative of director,Other"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"C{first}:C{last}")
    for col, w in zip("ABCDEFGHIJK", (28, 22, 24, 18, 13, 13, 13, 16, 14, 14, 14)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["A"].width = 34
    ws.freeze_panes = ws.cell(h + 1, 1)
    ins = wb.create_sheet("How to fill")
    for i, t in enumerate([
        "How to fill the BIRF",
        "1. Keep the header row (row 8) exactly as it is - the review tool reads the column names.",
        "2. One row per person who receives managerial remuneration or a benefit from the company.",
        "3. Enter the amount of each component for the financial year being benchmarked, in AED.",
        "4. The Total column and the Total row add up automatically; do not type over them.",
        "5. Total managerial remuneration must agree to the notes to the audit report and to the TP report.",
        "6. Upload this file on the Transfer pricing tab (BIRF box). The tool then checks: remuneration split vs the report, the total, and the persons named.",
        "7. Save in Excel before uploading so that the totals are calculated.",
    ], 1):
        ins.cell(i, 1, t).font = Font(bold=(i == 1), size=12 if i == 1 else 11)
    ins.column_dimensions["A"].width = 120
    wb.save(str(path))
    return path


if __name__ == "__main__":
    import sys
    build_birf_template(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "templates" / "BIRF_Template.xlsx")
