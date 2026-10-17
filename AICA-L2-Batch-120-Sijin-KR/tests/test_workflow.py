import json

import pytest

from auditsignoff.workflow import Workflow, WorkflowError


def _js(p, finds):
    p.write_text(json.dumps({"results": [{"check": "X", "status": "EXC", "findings": finds}]})); return p


def _f(i, sev="High"):
    return {"check_id": f"TP-{i}", "title": f"Issue {i}", "severity": sev, "detail": "d", "location": "p.3", "expected": "", "found": ""}


def test_full_flow_l1_l2_l3(tmp_path):
    w = Workflow(tmp_path / "w.db")
    w.create("E1", "TP", "ABC", "2025", "P")
    assert w.ingest("E1", _js(tmp_path / "a.json", [_f(1), _f(2, "Low")]), "S", "L1")["new"] == 2
    w.decide(1, "Cleared", "ok", "S", "L1")
    assert "Returned" in w.complete_round("E1", "S", "L1")           # low point still open -> back to preparer
    with pytest.raises(WorkflowError):
        w.submit_for_review("E1", "P", "preparer")                   # unanswered
    w.respond(2, "Corrected", "", "P", "preparer"); w.submit_for_review("E1", "P", "preparer")
    r = w.ingest("E1", _js(tmp_path / "b.json", []), "S", "L1")
    assert r["auto_cleared"] == 1
    assert "Manager" in w.complete_round("E1", "S", "L1", signed=True)
    with pytest.raises(WorkflowError):
        w.ingest("E1", _js(tmp_path / "c.json", []), "S", "L1")      # L1 cannot do the manager round
    w.ingest("E1", _js(tmp_path / "c.json", []), "M", "L2"); assert "Partner" in w.complete_round("E1", "M", "L2", signed=True)
    w.ingest("E1", _js(tmp_path / "d.json", []), "Pt", "L3"); assert "sign-off" in w.complete_round("E1", "Pt", "L3", signed=True)
    assert w.state_text("E1").startswith("Signed off")
    with pytest.raises(WorkflowError):
        w.ingest("E1", _js(tmp_path / "e.json", []), "Pt", "L3")


def test_claimed_correction_still_found_is_reopened(tmp_path):
    w = Workflow(tmp_path / "w.db"); w.create("E", "TP", "A", "2025", "P")
    w.ingest("E", _js(tmp_path / "a.json", [_f(1)]), "S", "L1"); w.complete_round("E", "S", "L1")
    w.respond(1, "Corrected", "", "P", "preparer"); w.submit_for_review("E", "P", "preparer")
    assert w.ingest("E", _js(tmp_path / "b.json", [_f(1)]), "S", "L1")["reopened"] == 1
    assert w.points("E")[0]["status"] == "REOPENED"


def test_dispute_needs_comment_and_waive_needs_role(tmp_path):
    w = Workflow(tmp_path / "w.db"); w.create("E", "TP", "A", "2025", "P")
    w.ingest("E", _js(tmp_path / "a.json", [_f(1)]), "S", "L1"); w.complete_round("E", "S", "L1")
    with pytest.raises(WorkflowError):
        w.respond(1, "Disputed", "", "P", "preparer")
    w.respond(1, "Disputed", "immaterial", "P", "preparer"); w.submit_for_review("E", "P", "preparer")
    w.ingest("E", _js(tmp_path / "b.json", [_f(1)]), "S", "L1")
    with pytest.raises(WorkflowError):
        w.decide(1, "Waived", "accept", "S", "L1")                   # L1 may not waive
    with pytest.raises(WorkflowError):
        w.decide(1, "Reopen", "", "S", "L1")                         # reason required


def test_trail_export(tmp_path):
    w = Workflow(tmp_path / "w.db"); w.create("E", "TP", "A", "2025", "P")
    w.ingest("E", _js(tmp_path / "a.json", [_f(1)]), "S", "L1")
    import openpyxl
    wb = openpyxl.load_workbook(w.export_trail("E", tmp_path / "t.xlsx"))
    assert wb.sheetnames == ["Flow", "Review points", "Rounds", "Files", "Audit log"]


def test_assignment_queue_attachment_and_message(tmp_path):
    w = Workflow(tmp_path / "w.db"); w.create("E", "TP", "A", "2025", "P")
    w.assign("E", "L1", "Senior 1", "P", "preparer"); w.assign("E", "preparer", "Preparer 1", "P", "preparer")
    assert [q["ref"] for q in w.my_queue("Senior 1")] == ["E"] and w.my_queue("Preparer 1") == []
    f = tmp_path / "fs.docx"; f.write_bytes(b"x" * 300)
    w.attach("E", f, "P", "preparer", "Updated file", "v2")
    assert w.attachments("E")[0]["filename"] == "fs.docx" and (tmp_path / "files" / "E").exists()
    w.ingest("E", _js(tmp_path / "a.json", [_f(1)]), "S", "L1"); w.complete_round("E", "S", "L1")
    assert [q["ref"] for q in w.my_queue("Preparer 1")] == ["E"]
    assert "Preparer 1" in w.handoff_message("E")


def test_staff_list_typed_and_uploaded(tmp_path):
    import openpyxl
    w = Workflow(tmp_path / "w.db")
    assert w.users()                                  # falls back to the sample names
    w.add_staff("Asha", "Preparer"); w.add_staff("Ravi", "first reviewer")
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Name", "Role", "Email"]); ws.append(["Meera", "Manager", "m@x.com"]); ws.append(["Omar", "Partner", ""]); ws.append(["Bad", "Cook", ""])
    f = tmp_path / "s.xlsx"; wb.save(f)
    r = w.import_staff(f)
    assert r["added"] == 2 and len(r["rejected"]) == 1
    assert w.users() == {"Omar": "L3", "Meera": "L2", "Ravi": "L1", "Asha": "preparer"}
    w.remove_staff("Asha"); assert "Asha" not in w.users()


def test_signoff_required_flow_and_stamp(tmp_path):
    import docx
    w = Workflow(tmp_path / "w.db"); w.create("E", "TP", "ABC", "2025", "Asha")
    w.assign("E", "L1", "Ravi", "Asha", "preparer"); w.assign("E", "L2", "Meera", "Asha", "preparer")
    fl = w.flow("E")
    assert fl["with"] == "Ravi" and fl["next"] == "Manager review" and "PENDING" in fl["headline"]
    assert fl["rows"][0]["status"] == "Signed off"           # preparer signed on creation
    w.ingest("E", _js(tmp_path / "a.json", []), "Ravi", "L1")
    with pytest.raises(WorkflowError):
        w.complete_round("E", "Ravi", "L1")                  # must sign
    w.complete_round("E", "Ravi", "L1", signed=True, comment="no issues")
    fl = w.flow("E")
    assert [r["status"] for r in fl["rows"]] == ["Signed off", "Signed off", "Pending - awaiting review", "Pending"]
    assert fl["with"] == "Meera" and "Manager review" in fl["headline"]
    # stamp onto a review paper that has the report's sign-off table
    d = docx.Document(); d.add_paragraph("Final Review Paper")
    t = d.add_table(rows=4, cols=5)
    for i, h in enumerate(["", "Prepared by", "First review", "Manager review", "Partner review"]):
        t.rows[0].cells[i].text = h
    for i, h in enumerate(["Name", "Date", "Sign"], 1):
        t.rows[i].cells[0].text = h
    src = tmp_path / "p.docx"; d.save(src)
    out = w.stamp_docx("E", src, tmp_path / "stamped.docx")
    txt = " | ".join(c.text for tb in docx.Document(str(out)).tables for r in tb.rows for c in r.cells)
    assert "Ravi" in txt and "Asha" in txt
    assert "REVIEW STATUS" in docx.Document(str(out)).paragraphs[0].text


def test_stamp_has_datetime_trail_and_is_idempotent(tmp_path):
    import docx
    w = Workflow(tmp_path / "w.db"); w.create("E", "Liquidation", "ACMEROCKET FZE", "2026", "Lina")
    for rl, nm in (("preparer", "Lina"), ("L1", "Ravi"), ("L2", "Meera"), ("L3", "Omar")):
        w.assign("E", rl, nm, "Lina", "preparer")
    w.ingest("E", _js(tmp_path / "a.json", []), "Ravi", "L1"); w.complete_round("E", "Ravi", "L1", signed=True)
    d = docx.Document(); d.add_paragraph("Final Review Paper")
    t = d.add_table(rows=4, cols=5)
    for i, h in enumerate(["", "Prepared by", "First review", "Manager review", "Partner review"]):
        t.rows[0].cells[i].text = h
    for i, h in enumerate(["Name", "Date", "Sign"], 1):
        t.rows[i].cells[0].text = h
    src = tmp_path / "p.docx"; d.save(src)
    out = w.stamp_docx("E", src, tmp_path / "s1.docx")
    w.ingest("E", _js(tmp_path / "b.json", []), "Meera", "L2"); w.complete_round("E", "Meera", "L2", signed=True)
    out2 = w.stamp_docx("E", out, tmp_path / "s2.docx")              # re-stamp the stamped file
    dd = docx.Document(str(out2))
    assert sum(1 for p in dd.paragraphs if p.text.startswith("REVIEW STATUS")) == 1
    txt = " | ".join(c.text for tb in dd.tables for r in tb.rows for c in r.cells)
    import re
    assert re.search(r"\d{2}-[A-Z][a-z]{2}-\d{4} \d{2}:\d{2}", txt)      # date AND time
    assert "Completed First review" in txt and "Manager review - Meera" in txt and "Partner review - Omar" in txt
    sign = [tb for tb in dd.tables if tb.rows[1].cells[0].text == "Name"][0]
    assert sign.rows[1].cells[2].text == "Ravi" and sign.rows[1].cells[3].text == "Meera" and sign.rows[2].cells[3].text


def test_auto_engagement_created_and_paper_stamped(samples, tmp_path):
    import docx
    from auditsignoff import cli
    from auditsignoff.workflow import Workflow
    cli.set_team("Noor", "Hana", "Meera", "Omar"); cli.set_engagement(tmp_path / "w.db", None, auto=True)
    try:
        res, ms, paper = cli.review_liq(str(samples / "liq_dmcc_clean.yaml"), tmp_path / "o", __import__("datetime").date(2026, 10, 14))
        w = Workflow(tmp_path / "w.db")
        e = w.engagements()
        assert len(e) == 1 and w.assignees(e[0]["ref"])["L1"] == "Hana"
        txt = docx.Document(str(paper))
        assert txt.paragraphs[0].text.startswith("REVIEW STATUS")
        cells = " | ".join(c.text for t in txt.tables for r in t.rows for c in r.cells)
        assert "Noor" in cells and "Review movement trail" in " ".join(p.text for p in txt.paragraphs)
        cli.review_liq(str(samples / "liq_dmcc_clean.yaml"), tmp_path / "o", __import__("datetime").date(2026, 10, 14))   # run again: no duplicate engagement
        assert len(w.engagements()) == 1
    finally:
        cli.set_engagement(); cli.set_team()
