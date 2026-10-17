import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent))
from fa_fixtures import make_gl, make_masters
from tp_fixtures import make_report, make_workbook, stats

from auditsignoff.inputs import load_input
from auditsignoff.models import load_catalogue, load_params
from auditsignoff.rules_ct import REGISTRY as CT_REG, run_ct
from auditsignoff.rules_fa2 import FA2_REGISTRY
from auditsignoff.rules_fa import REGISTRY as FA_REG, run_fa
from auditsignoff.rules_tp import REGISTRY as TP_REG, run_tp

S = Path(__file__).resolve().parent.parent / "samples"
P = load_params()
ENG = load_input(S / "engagement_clean.xlsx", "engagement")


def _ids(res, status):
    return {r.check_id for r in res if r.status == status}


def test_registry_matches_catalogue_flags():
    cat = load_catalogue()
    from auditsignoff.rules_tp2 import A_REGISTRY, ASSIST_REGISTRY
    from auditsignoff.tp_checklist import CHECKS, LEGACY_USED
    assert LEGACY_USED == set(TP_REG) | set(A_REGISTRY) | set(ASSIST_REGISTRY), "every engine test must sit under a checklist row"
    for mod, reg in (("C", {c["id"] for c in CHECKS if c["impl"]}), ("D", CT_REG), ("E", {**FA_REG, **FA2_REGISTRY, "FA-02": 0, "FA-09": 0, "FA-10": 0})):
        assert {c["id"] for c in cat if c["mod"] == mod and c["impl"]} == set(reg), mod


# ---------------------------------------------------------------- TP
def test_tp_clean_report_has_no_exceptions(tmp_path):
    make_workbook(tmp_path / "GOLDEN_CREST_FY2025_benchmark.xlsx"); make_report(tmp_path / "r.docx")
    res, _ = run_tp(tmp_path / "r.docx", tmp_path / "GOLDEN_CREST_FY2025_benchmark.xlsx", ENG, P, facts={"rp_aggregate": 50_000_000})
    assert not _ids(res, "EXC"), [(r.check_id, [f.title for f in r.findings]) for r in res if r.findings]


def test_tp_errors_detected(tmp_path):
    make_workbook(tmp_path / "w.xlsx")
    st = dict(stats()); st["median"] += 0.5; st["q1"] = stats("EXC")["q1"]
    make_report(tmp_path / "r.docx", name="GOLDEN CREST TRADING DMC", fy=2024, st=st, count=7, border=False, bad_font=True,
                header_title="Report", page_field=False, annex=("A", "C"))
    res, _ = run_tp(tmp_path / "r.docx", tmp_path / "w.xlsx", ENG, P)
    exc = _ids(res, "EXC")
    assert {"TP-01", "TP-03", "TP-17", "TP-21", "TP-22", "TP-23", "TP-25", "TP-45", "TP-46"} <= exc, exc
    titles = " ".join(f.title for r in res for f in r.findings)
    assert "Median in the report does not agree" in titles and "other quartile method" in titles


def test_tp_pli_recompute(tmp_path):
    make_workbook(tmp_path / "w.xlsx", tested_pli=3.2, op=320000, oc=9000000)      # 320k/9m = 3.56% != 3.2%
    make_report(tmp_path / "r.docx")
    res, _ = run_tp(tmp_path / "r.docx", tmp_path / "w.xlsx", ENG, P)
    assert "TP-33" in _ids(res, "EXC")


# ---------------------------------------------------------------- CT (needs LibreOffice to calculate a workbook never opened in Excel)
needs_lo = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")


@needs_lo
def test_ct_detects_workbook_defects(tmp_path):
    wb = openpyxl.load_workbook(S / "ct_working_sheet_blank.xlsx")
    wb["01_CLIENT_INFO"]["C9"] = datetime(2025, 12, 31)
    wb["03_INCOME_SCHEDULE"]["D9"] = 2_000_000; wb["03_INCOME_SCHEDULE"]["D10"] = 500_000
    wb["04_ADJUSTMENTS"]["D9"] = -100_000; wb["04_ADJUSTMENTS"]["D13"] = -50_000
    wb["07_TP_SCHEDULE"]["D8"] = 1_000_000
    wb["09_TAX_LIABILITY"]["D16"] = datetime(2026, 12, 10)
    wb.save(tmp_path / "ct.xlsx")
    res, c = run_ct(tmp_path / "ct.xlsx", P)
    assert c.ti_before_losses == 1_350_000 and round(c.payable) == 87_750
    exc = _ids(res, "EXC")
    assert {"CT-01", "CT-02", "CT-03", "CT-07"} <= exc, exc
    titles = " ".join(f.title for r in res for f in r.findings)
    assert "sign error" in titles and "A5" in titles and "FULL transaction value" in titles


@needs_lo
def test_ct_return_tieout(tmp_path):
    wb = openpyxl.load_workbook(S / "ct_working_sheet_blank.xlsx")
    wb["03_INCOME_SCHEDULE"]["D9"] = 2_000_000; wb["03_INCOME_SCHEDULE"]["D10"] = 500_000
    wb.save(tmp_path / "ct.xlsx")
    res, c = run_ct(tmp_path / "ct.xlsx", P, return_pairs={"Taxable income": "1,500,000", "Corporate tax payable": "100,000"})
    assert c.ti == 1_500_000 and round(c.payable) == 101_250
    assert "CT-05" in _ids(res, "EXC")


# ---------------------------------------------------------------- Forensic
def test_forensic_gl_issues(tmp_path):
    make_gl(tmp_path / "gl.xlsx"); make_masters(tmp_path)
    res, _ = run_fa(P, tmp_path / "gl.xlsx", "2025-12-31", [2025], tmp_path / "vendors.xlsx", tmp_path / "employees.xlsx", ["Mohammed Al Hashmi"])
    assert {"FA-05", "FA-06", "FA-07"} <= _ids(res, "EXC")
    t = " ".join(f.title for r in res for f in r.findings)
    for k in ("weekends", "backdated", "duplicate", "prepared and approved", "just below", "share the same bank"):
        assert k in t, k


def test_forensic_clean_gl_passes(tmp_path):
    make_gl(tmp_path / "gl.xlsx", issues=False)
    res, _ = run_fa(P, tmp_path / "gl.xlsx", "2025-12-31", [2025])
    assert {"FA-01", "FA-05"} <= _ids(res, "PASS")


def test_forensic_evidence_and_authenticity(tmp_path):
    from reportlab.pdfgen import canvas
    import docx
    ev = tmp_path / "ev"; ev.mkdir()
    cv = canvas.Canvas(str(ev / "bank_statement.pdf")); cv.setProducer("iLovePDF"); cv.drawString(72, 700, "Statement"); cv.save()
    d = docx.Document(); d.core_properties.author = "Accounts"; d.core_properties.last_modified_by = "Someone Else"; d.save(ev / "letter.docx")
    reg = openpyxl.Workbook(); ws = reg.active
    ws.append(["Item", "Description", "Source", "Date received", "Received by", "SHA-256"])
    ws.append([1, "bank_statement.pdf", "Bank", datetime(2025, 5, 1), "Amal", "0" * 64])
    ws.append([2, "other.pdf", "", None, "", ""])
    reg.save(tmp_path / "register.xlsx")
    res, _ = run_fa(P, None, None, [2025], register=tmp_path / "register.xlsx", evidence_dir=ev)
    t = " ".join(f.title for r in res for f in r.findings)
    assert "Ilovepdf" in t and "different user than the author" in t
    assert "hash differs" in t and "not in the register" in t and "no source" in t


@needs_lo
def test_ct_demo_has_no_known_defects(tmp_path):
    wb = openpyxl.load_workbook(S / "ct_demo_locked.xlsx")
    wb["01_CLIENT_INFO"]["C9"] = datetime(2025, 12, 31)
    wb["03_INCOME_SCHEDULE"]["D9"] = 2_000_000; wb["03_INCOME_SCHEDULE"]["D10"] = 500_000
    wb["04_ADJUSTMENTS"]["D9"] = -100_000; wb["04_ADJUSTMENTS"]["D13"] = -50_000
    t = wb["07_TP_SCHEDULE"]
    t["D8"] = 1_000_000; t["I8"] = 1_200_000; t["D9"] = 500_000; t["I9"] = 400_000
    wb.save(tmp_path / "ct.xlsx")
    res, c = run_ct(tmp_path / "ct.xlsx", P)
    assert c.ti_before_losses == 1_650_000 and round(c.payable) == 114_750 and c.tp_adj == 300_000
    assert not {"CT-01", "CT-02", "CT-03"} & _ids(res, "EXC")
