"""Synthetic TP report + benchmarking workbook used by tests and the sample set."""
from pathlib import Path
import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

COMPS = [("Alpha Trading Ltd", 4.1), ("Beta Supplies Inc", 5.2), ("Gamma Distribution SA", 6.0), ("Delta Wholesale GmbH", 6.8),
         ("Epsilon Goods BV", 7.5), ("Zeta Commerce Ltd", 8.3), ("Eta Traders SpA", 9.1), ("Theta Imports NV", 10.4)]


def stats(method="INC"):
    from auditsignoff.rules_tp import quartiles
    return quartiles([v for _, v in COMPS], method)


def make_workbook(path, tested_pli=3.2, op=320000, oc=10000000):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Comparables"
    ws.append(["Company", "Operating margin (OP/OC)"])
    for n, v in COMPS:
        ws.append([n, v / 100])
    ws2 = wb.create_sheet("Tested party")
    ws2.append(["Tested party operating profit", op]); ws2.append(["Tested party operating costs", oc]); ws2.append(["Tested party PLI", tested_pli / 100])
    wb.create_sheet("Checks")["A1"] = "=SUM(Comparables!B2:B3)"      # a live formula, as in a real workbook
    wb.save(path)


def _field(par, instr):
    for t, txt in (("begin", None), (None, instr), ("end", None)):
        r = par.add_run()
        if t:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), t); r._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.text = txt; r._r.append(it)


def make_report(path, name="GOLDEN CREST TRADING DMCC", fy=2025, st=None, count=8, border=True, bad_font=False, header_title="Transfer Pricing Benchmark Analysis",
                inside=False, page_field=True, annex=("A", "B")):
    st = st or stats()
    d = Document()
    d.styles["Normal"].font.name = "Arial" if bad_font else "Times New Roman"
    d.styles["Normal"].font.size = Pt(11)
    sec = d.sections[0]
    hp = sec.header.paragraphs[0]; hp.text = header_title + "\t"; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if page_field:
        _field(hp, "PAGE")
    else:
        hp.add_run(" 1")
    fp = sec.footer.paragraphs[0]; fp.text = name; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if border:
        pb = OxmlElement("w:pgBorders")
        for e in ("top", "left", "bottom", "right"):
            x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "12"); pb.append(x)
        sec._sectPr.append(pb)
    d.add_paragraph(name); d.add_paragraph("Transfer Pricing Benchmark Analysis"); d.add_paragraph(f"For the financial year ended 31 December {fy}")
    d.add_page_break()
    d.add_paragraph(f"The analysis uses {count} comparable companies. The tested party PLI is "
                    + ("within the arm's length range." if inside else "outside the arm's length range, so an adjustment to the median is made."))
    d.add_paragraph(f"The median is {st['median']:.2f}%.")
    t = d.add_table(rows=0, cols=2)
    for lab, k in (("Minimum", "min"), ("Lower quartile", "q1"), ("Median", "median"), ("Upper quartile", "q3"), ("Maximum", "max")):
        r = t.add_row().cells; r[0].text = lab; r[1].text = f"{st[k]:.2f}%"
    d.add_paragraph("Tested party operating profit is 320,000 on operating costs of 10,000,000, a PLI of 3.20%.")
    d.add_paragraph("RESULTS AND CONCLUSION")
    t2 = d.add_table(rows=0, cols=4)
    r = t2.add_row().cells
    for i, h in enumerate(("PLI", "Lower quartile", "Upper quartile", "Actual (tested party)")):
        r[i].text = h
    r = t2.add_row().cells
    for i, h in enumerate(("OP/OC", f"{st['q1']:.2f}%", f"{st['q3']:.2f}%", "3.20%")):
        r[i].text = h
    d.add_paragraph("Comparable companies: " + "; ".join(n for n, _ in COMPS))
    d.add_paragraph("The company is licensed by DMCC and owned by Mohammed Al Hashmi, Fatima Khan and Rajan Pillai.")
    d.add_paragraph("Related party transactions exceed the thresholds; the Transfer Pricing Disclosure Form is filed with the return.")
    d.add_paragraph("Industry analysis: the trading sector is stable. Company profile: the entity trades building materials. Controlled transaction: remuneration paid to the connected person.")
    d.add_paragraph("Functional analysis: the functions performed are sales and procurement; the assets used are limited to inventory; the risks borne are market and credit risk. "
                    "The company is the tested party because it is the least complex entity and reliable data is available.")
    d.add_paragraph("Method selection: the TNMM is the most appropriate method. The OP/OC PLI was selected because it is least affected by differences in functions. "
                    "The CUP is not appropriate as no reliable uncontrolled transactions exist. The resale price method is not appropriate for lack of data. "
                    "The cost plus method is not suitable. The profit split method is not applicable. Comparability analysis used a screening of listed companies; single-year data is used.")
    d.add_paragraph("Conclusion: the arm's length range is the interquartile range.")
    d.add_paragraph("Refer Annexure A for the search strategy and Annexure B for the financials.")
    for a in annex:
        d.add_paragraph(f"Annexure {a}")
        d.add_paragraph("Annexure content.")
    d.save(path)
