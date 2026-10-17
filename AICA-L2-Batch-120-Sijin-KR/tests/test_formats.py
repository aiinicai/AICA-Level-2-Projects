import datetime
import shutil
import pytest
from pathlib import Path
from auditsignoff.cli import review_fs, review_liq
from auditsignoff.loaders import load_document, load_tb

S = Path(__file__).resolve().parent.parent / "samples"


def _st(res):
    return sorted((r.check_id, r.status) for r in res)


@pytest.mark.parametrize("case", ["clean", "errors"])
def test_excel_factsheet_same_as_yaml(case, tmp_path):
    a, _ = review_fs(S / f"engagement_{case}.yaml", S / f"fs_{case}.docx", S / f"tb_{case}.xlsx", tmp_path / "a")
    b, _ = review_fs(S / f"engagement_{case}.xlsx", S / f"fs_{case}.docx", S / f"tb_{case}.xlsx", tmp_path / "b")
    assert _st(a) == _st(b)


@pytest.mark.parametrize("case", ["difc_breach", "dmcc_clean", "dso_short_gap"])
def test_excel_tracker_same_as_yaml(case, tmp_path):
    t = datetime.date(2026, 10, 14)
    a = review_liq(S / f"liq_{case}.yaml", tmp_path / "a", t)[0]
    b = review_liq(S / f"liq_{case}.xlsx", tmp_path / "b", t)[0]
    assert _st(a) == _st(b)


def test_csv_tb(tmp_path):
    import openpyxl, csv
    ws = openpyxl.load_workbook(S / "tb_clean.xlsx", data_only=True).worksheets[0]
    p = tmp_path / "tb.csv"
    with open(p, "w", newline="") as fh:
        csv.writer(fh).writerows(ws.iter_rows(values_only=True))
    assert load_tb(p) == load_tb(S / "tb_clean.xlsx")


def test_excel_fs_pack_skips_formatting(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Statement of financial position"]); ws.append(["Total assets", 1000000, 900000])
    p = tmp_path / "fs.xlsx"; wb.save(p)
    d = load_document(p)
    assert d.kind == "xlsx" and "Total assets" in d.text
    res, _ = review_fs(S / "engagement_clean.xlsx", p, None, tmp_path / "o")
    st = {r.check_id: r.status for r in res}
    assert st["SA-13"] == "NOTRUN" and st["SA-15"] == "NOTRUN" and st["SA-16"] == "NOTRUN"


@pytest.mark.skipif(not shutil.which("tesseract"), reason="Tesseract not installed")
def test_image_and_scanned_pdf_ocr(tmp_path):
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1400, 500), "white"); d = ImageDraw.Draw(img)
    f = ImageFont.load_default(size=36)
    d.text((40, 40), "Statement of financial position", fill="black", font=f)
    d.text((40, 120), "Total assets        1,250,000        1,100,000", fill="black", font=f)
    jp, pdf = tmp_path / "fs.jpg", tmp_path / "fs_scan.pdf"
    img.save(jp); img.save(pdf)
    for p, kind in ((jp, "image"), (pdf, "scan")):
        doc = load_document(p)
        assert doc.kind == kind and "financial position" in doc.text.lower()


def _eng_cases(tmp_path):
    import openpyxl
    from docx import Document
    from auditsignoff.inputs import ENG_LABELS, ENG_FIELDS, load_input
    ref = load_input(S / "engagement_errors.xlsx", "engagement")
    return ref, ENG_LABELS


def test_factsheet_word_and_free_xlsx(tmp_path):
    from docx import Document
    from auditsignoff.inputs import load_input, ENG_LABELS
    ref = load_input(S / "engagement_clean.xlsx", "engagement")
    d = Document(); t = d.add_table(rows=0, cols=2)
    vals = {"fy": "2025", "py": "2024", "client.name": ref["client"]["name"], "client.legal_status": "FZC",
            "client.authority": "DMCC", "client.emirate": "Dubai", "client.year_end": "31 December",
            "client.shareholders": "; ".join(ref["client"]["shareholders"]),
            "client.signatories_legal": "; ".join(ref["client"]["signatories_legal"]),
            "fs_signatories": "; ".join(ref["fs_signatories"]), "br_on_file": "No"}
    for k, v in vals.items():
        r = t.add_row().cells; r[0].text = ENG_LABELS[k]; r[1].text = v
    p = tmp_path / "fact.docx"; d.save(p)
    got = load_input(p, "engagement")
    assert got["client"] == ref["client"] and got["fs_signatories"] == ref["fs_signatories"]
    # 'Label: value' lines in a Word file
    d2 = Document()
    for k, v in vals.items():
        d2.add_paragraph(f"{ENG_LABELS[k]}: {v}")
    p2 = tmp_path / "lines.docx"; d2.save(p2)
    assert load_input(p2, "engagement")["client"] == ref["client"]


def test_factsheet_missing_fields_clear_error(tmp_path):
    from docx import Document
    from auditsignoff.inputs import load_input
    d = Document(); d.add_paragraph("Hello"); p = tmp_path / "x.docx"; d.save(p)
    with pytest.raises(ValueError, match="Could not find"):
        load_input(p, "engagement")


@pytest.mark.skipif(not shutil.which("tesseract"), reason="Tesseract not installed")
def test_factsheet_image_ocr(tmp_path):
    from PIL import Image, ImageDraw, ImageFont
    from auditsignoff.inputs import load_input
    img = Image.new("RGB", (1500, 700), "white"); dr = ImageDraw.Draw(img); f = ImageFont.load_default(size=34)
    for i, ln in enumerate(["Financial year: 2025", "Client name: GOLDEN CREST TRADING DMCC", "Legal status: FZC",
                            "Licensing authority: DMCC", "Shareholders: Fatima Khan; Rajan Pillai"]):
        dr.text((40, 40 + i * 70), ln, fill="black", font=f)
    p = tmp_path / "fact.jpg"; img.save(p)
    got = load_input(p, "engagement")
    assert got["fy"] == 2025 and got["client"]["legal_status"] == "FZC" and got["_source"]["ocr"]
