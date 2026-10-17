from datetime import date

import docx
import openpyxl

from auditsignoff.cli import review_liq
from auditsignoff.models import load_yaml

TODAY = date(2026, 10, 14)


def _doc(path, lines):
    d = docx.Document()
    for l in lines:
        d.add_paragraph(l)
    d.save(path)
    return str(path)


def _inp(samples):
    inp = load_yaml(samples / "liq_dmcc_clean.yaml")
    return inp


def _bundle(tmp_path, name, br_date, good=True):
    br = _doc(tmp_path / "br.docx", [f"Board Resolution of {name}", f"Dated {br_date}",
              f"IT IS RESOLVED to voluntarily liquidate {name} and to appoint Mr Ravi Kumar as liquidator.", "Signed by all shareholders: Alpha Holdings."])
    soa = _doc(tmp_path / "soa.docx", [f"Statement of Affairs of {name}", "As at 30 June 2026", "Total assets 500,000", "Total liabilities 100,000" if good else "Total liabilities 900,000",
               "Net surplus 400,000" if good else "Net deficiency 400,000", "VAT payable nil", "Liquidator Ravi Kumar"])
    rep = _doc(tmp_path / "rep.docx", [f"Liquidator's report of {name}", "Assets were realised and sold.", "Liabilities settled and creditors paid.", "Distribution to shareholders made.",
               "Newspaper advertisement published.", "No claims received.", "Bank account closure completed.", "Books and records will be kept.",
               "Corporate tax final return filed; deregistration requested." if good else "Fine.", "Total assets 500,000", "Total liabilities 100,000"])
    ct = _doc(tmp_path / "ct.docx", [f"Corporate Tax registration certificate {name}", "Corporate Tax TRN 100123456789012"])
    return {"br": br, "soa": soa, "report": rep, "ctvat": [ct] if good else []}


def test_docs_missing_gives_notrun(samples, tmp_path):
    p = samples / "liq_dmcc_clean.yaml"
    res, ms, _ = review_liq(str(p), tmp_path, TODAY)
    st = {r.check_id: r.status for r in res}
    assert st["LQ-21"] == st["LQ-22"] == st["LQ-25"] == "NOTRUN"
    assert "Upload" in next(r for r in res if r.check_id == "LQ-22").note


def test_docs_flawed_flag_exceptions(samples, tmp_path):
    p = samples / "liq_dmcc_clean.yaml"
    inp = load_yaml(p)
    docs = _bundle(tmp_path, inp["company"], "1 June 2026", good=False)
    res, ms, _ = review_liq(str(p), tmp_path / "o", TODAY, docs=docs)
    r = {x.check_id: x for x in res}
    titles = " | ".join(f.title for x in res if x.check_id in ("LQ-22", "LQ-23", "LQ-25") for f in x.findings)
    assert r["LQ-22"].status == "EXC" and "deficiency" in titles.lower()
    assert r["LQ-25"].status == "EXC" and "Corporate Tax" in titles
    assert r["LQ-23"].status == "EXC"


def test_docs_good_no_registration_exception(samples, tmp_path):
    p = samples / "liq_dmcc_clean.yaml"
    inp = load_yaml(p)
    docs = _bundle(tmp_path, inp["company"], "1 June 2026", good=True)
    res, ms, _ = review_liq(str(p), tmp_path / "o", TODAY, docs=docs)
    r = {x.check_id: x for x in res}
    assert r["LQ-22"].status == "PASS"
    assert not any("No evidence the company is registered" in f.title for f in r["LQ-25"].findings)


def test_tb_unbalanced(samples, tmp_path):
    p = samples / "liq_dmcc_clean.yaml"
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Account", "Class", "Debit", "Credit"]); ws.append(["Cash", "Asset", 100, 0]); ws.append(["Capital", "Equity", 0, 90])
    tb = tmp_path / "tb.xlsx"; wb.save(tb)
    res, _, _ = review_liq(str(p), tmp_path / "o", TODAY, docs={"tb": str(tb)})
    assert next(x for x in res if x.check_id == "LQ-24").status == "EXC"


def test_combined_soa_and_report_in_one_file(samples, tmp_path):
    p = samples / "liq_dmcc_clean.yaml"
    inp = load_yaml(p)
    one = _doc(tmp_path / "combined.docx", [f"Liquidator's report of {inp['company']}", "Statement of Affairs", "As at 30 June 2026", "Total assets 500,000", "Total liabilities 100,000",
                                            "Net surplus 400,000", "Assets realised and sold. Liabilities settled. Distribution made. Newspaper advertisement. No claims. Bank closure. Books and records kept. Deregistration done."])
    res, _, _ = review_liq(str(p), tmp_path / "o", TODAY, docs={"soa_report": [one]})
    st = {r.check_id: r.status for r in res}
    assert st["LQ-22"] != "NOTRUN" and st["LQ-23"] != "NOTRUN"
