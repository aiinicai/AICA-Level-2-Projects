import docx

from auditsignoff.rules_fs import canon, fuzzy_hits, same_name


def test_dash_case_spacing_are_not_spelling_errors():
    n = "ACMEROCKET - FZE"
    for v in ("ACMEROCKET – FZE", "ACMEROCKET—FZE", "Acmerocket – FZE", "acmerocket-fze", "ACMEROCKET  -  FZE."):
        assert same_name(v, n), v
        hits = fuzzy_hits(f"Report of {v} for the year", n)
        assert hits and hits[0][0] == n
    assert not same_name("ACMEROKET - FZE", n)


def test_real_misspelling_still_found():
    hits = fuzzy_hits("Report of ACMEROKET - FZE for the year", "ACMEROCKET - FZE")
    assert hits and hits[0][0] != "ACMEROCKET - FZE"


def test_liq_consistency_ignores_dash_and_case(samples, tmp_path):
    from auditsignoff.cli import review_liq
    from auditsignoff.models import load_yaml
    inp = load_yaml(samples / "liq_dmcc_clean.yaml")
    n = inp["company"]
    def mk(name, lines):
        d = docx.Document()
        for l in lines: d.add_paragraph(l)
        d.save(tmp_path / name); return str(tmp_path / name)
    var = n.replace(" ", " – ", 1).title()          # different dash + case
    docs = {"br": mk("br.docx", [f"Board Resolution of {n}", "Dated 2 March 2026"]), "report": mk("r.docx", [f"Liquidator's report of {var}", "x"])}
    res, _, _ = review_liq(str(samples / "liq_dmcc_clean.yaml"), tmp_path / "o", __import__("datetime").date(2026, 10, 14), docs=docs)
    assert not [f for r in res if r.check_id == "LQ-26" for f in r.findings if "spelt" in f.title]


def test_team_names_on_signoff(samples, tmp_path):
    from auditsignoff import cli
    cli.set_team("Asha", "Ravi", "Meera", "Omar")
    res, ms, docx_path = cli.review_liq(str(samples / "liq_dmcc_clean.yaml"), tmp_path, __import__("datetime").date(2026, 10, 14))
    cli.set_team()
    d = docx.Document(str(docx_path))
    cells = " | ".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    for n in ("Asha", "Ravi", "Meera", "Omar"):
        assert n in cells
