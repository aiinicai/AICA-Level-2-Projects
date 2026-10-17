import shutil
import subprocess
import pytest
from pathlib import Path
from docx import Document
from auditsignoff.cli import review_fs
from auditsignoff.folder import scan_folder
from auditsignoff.inputs import load_input
from auditsignoff.prefill import prefill, rows_to_engagement, write_fact_sheet_xlsx

S = Path(__file__).resolve().parent.parent / "samples"


def _licence(p):
    d = Document(); d.add_paragraph("Trade Licence"); t = d.add_table(rows=0, cols=2)
    for a, b in [("Company Name", "GOLDEN CREST TRADING DMCC"), ("Legal Type", "Free Zone Company (FZCO)"),
                 ("Licence Number", "DMCC-123456"), ("Issue Date", "14/03/2021"), ("Manager", "Mr. Mohammed Al Hashmi"),
                 ("Issuing Authority", "Dubai Multi Commodities Centre Authority")]:
        r = t.add_row().cells; r[0].text = a; r[1].text = b
    d.save(p)


def _moa(p):
    d = Document(); d.add_paragraph("MEMORANDUM OF ASSOCIATION")
    d.add_paragraph("Shareholders: Mr. Mohammed Al Hashmi (UAE) 50%, Ms. Fatima Khan (India) 30%, Rajan Pillai (India) 20%")
    d.add_paragraph("The financial year of the Company shall end on 31 December each year.")
    d.save(p)


def test_prefill_licence_and_moa(tmp_path):
    _licence(tmp_path / "Trade Licence.docx"); _moa(tmp_path / "MOA.docx")
    rows, conf = prefill([tmp_path / "Trade Licence.docx", tmp_path / "MOA.docx"])
    v = {r["field"]: r["value"] for r in rows}
    assert v["client.name"] == "GOLDEN CREST TRADING DMCC" and v["client.legal_status"] == "FZC"
    assert v["client.authority"] == "DMCC" and v["allowed_years"] == "2021"
    assert v["client.shareholders"] == "Mohammed Al Hashmi; Fatima Khan; Rajan Pillai" and v["client.year_end"] == "31 December"
    eng = rows_to_engagement(rows, 2025)
    x = tmp_path / "fs.xlsx"; write_fact_sheet_xlsx(eng, x)
    back = load_input(x, "engagement")
    assert back["client"]["name"] == "GOLDEN CREST TRADING DMCC" and back["fy"] == 2025 and back["allowed_years"] == [2021]


def test_folder_scan_and_run(tmp_path):
    shutil.copy(S / "fs_clean.docx", tmp_path / "Draft Financial Statements 2025.docx")
    shutil.copy(S / "tb_clean.xlsx", tmp_path / "Trial Balance.xlsx")
    sub = tmp_path / "KYC"; sub.mkdir()
    _licence(sub / "Trade Licence.docx"); _moa(sub / "MOA.docx")
    roles = {r["File"].split("\\")[-1].split("/")[-1]: r["Role"] for r in scan_folder(tmp_path)}
    assert roles["Draft Financial Statements 2025.docx"] == "FS pack" and roles["Trial Balance.xlsx"] == "Trial balance"
    assert roles["Trade Licence.docx"] == "Trade licence" and roles["MOA.docx"] == "MOA / AOA"


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_legacy_doc_and_xls(tmp_path):
    for src, fmt in ((S / "fs_clean.docx", "doc"), (S / "tb_clean.xlsx", "xls")):
        subprocess.run(["soffice", "--headless", "--convert-to", fmt, "--outdir", str(tmp_path), str(src)],
                       check=True, capture_output=True, timeout=180)
    from auditsignoff.loaders import load_document, load_tb
    d = load_document(tmp_path / "fs_clean.doc")
    assert d.kind == "docx" and "financial position" in d.text.lower()
    assert load_tb(tmp_path / "tb_clean.xls") == load_tb(S / "tb_clean.xlsx")
    res, _ = review_fs(S / "engagement_clean.xlsx", tmp_path / "fs_clean.doc", tmp_path / "tb_clean.xls", tmp_path / "o")
    assert res
