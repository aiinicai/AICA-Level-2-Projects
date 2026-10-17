"""The 29 TP checklist rows that were 'rule not built' - each now has a test."""
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import docx
import openpyxl
import pytest
from docx.shared import Pt
from tp_fixtures import make_report, make_workbook
from auditsignoff import tp_checklist as K
from auditsignoff.models import load_params
from auditsignoff.rules_tp import run_tp

P = load_params()
ENG = {"fy": 2025, "client": {"name": "GOLDEN CREST TRADING DMCC", "authority": "DMCC", "shareholders": [], "signatories_legal": []}}


def run(tmp, report, wb=None, **kw):
    res, ctx = run_tp(report, wb, ENG, P, **kw)
    return {r.check_id: r for r in res}, ctx


def titles(r):
    return " | ".join(f.title for f in r.findings)


def _doc(path, paras, tables=(), tnr=True):
    d = docx.Document()
    d.styles["Normal"].font.name = "Times New Roman"; d.styles["Normal"].font.size = Pt(11)
    for kind, text, *rest in paras:
        if kind == "h1":
            h = d.add_heading(text, 1)
            for r in h.runs:
                r.font.name = "Times New Roman"; r.font.size = Pt(19.5)
        elif kind == "h2":
            h = d.add_heading(text, 2)
            for r in h.runs:
                r.font.name = "Times New Roman" if tnr else "Arial"; r.font.size = Pt(12 if tnr else 14)
        else:
            d.add_paragraph(text)
    for rows in tables:
        t = d.add_table(rows=0, cols=len(rows[0]))
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
    d.save(path)


# ------------------------------------------------------------------ 1.2 / 1.4 TOC and headings
def test_toc_rules(tmp_path):
    _doc(tmp_path / "r.docx", [("p", "Table of Contents"), ("p", "1. Introduction\t3"), ("p", "2 Methodology\t4"), ("p", "Annexure A - Search\t9"), ("p", "Annexure B - Missing page\t10"),
                               ("h1", "1. Introduction"), ("h1", "2. Methodology"), ("h1", "3. Conclusion"), ("p", "Annexure A - Search")])
    by, _ = run(tmp_path, tmp_path / "r.docx")
    assert by["TP-05"].status == "EXC" and "Conclusion" in " ".join(f.detail for f in by["TP-05"].findings)            # heading not in TOC
    assert by["TP-06"].status == "EXC" and "numbering" in titles(by["TP-06"])                          # '2 Methodology' vs '2. Methodology'
    assert by["TP-08"].status == "EXC" and "no matching annexure heading" in titles(by["TP-08"])      # Annexure B
    assert by["TP-09"].status in ("PASS", "EXC")


def test_toc_format_flags_font(tmp_path):
    d = docx.Document()
    p = d.add_paragraph(); r = p.add_run("Table of Contents"); r.font.name = "Arial"; r.font.size = Pt(14)
    q = d.add_paragraph(); rr = q.add_run("1. Introduction\t3"); rr.font.name = "Arial"; rr.font.size = Pt(9)
    d.add_heading("1. Introduction", 1)
    d.save(tmp_path / "t.docx")
    by, _ = run(tmp_path, tmp_path / "t.docx")
    assert by["TP-09"].status == "EXC"


def test_subheading_format(tmp_path):
    _doc(tmp_path / "s.docx", [("h2", "1.1 Scope"), ("h2", "Approach")], tnr=False)
    by, _ = run(tmp_path, tmp_path / "s.docx")
    assert by["TP-19"].status == "EXC" and "Arial" in " ".join(f.found for f in by["TP-19"].findings)


# ------------------------------------------------------------------ 1.3 / 1.6 / 1.8 content rules
def test_content_rules_on_a_clean_and_a_bare_report(tmp_path):
    make_workbook(tmp_path / "w.xlsx"); make_report(tmp_path / "r.docx")
    by, _ = run(tmp_path, tmp_path / "r.docx", tmp_path / "w.xlsx")
    for cid in ("TP-15", "TP-28", "TP-39", "TP-40", "TP-68", "TP-56"):
        assert by[cid].status in ("PASS", "REVIEW"), (cid, titles(by[cid]))
    _doc(tmp_path / "bare.docx", [("p", "This is a short note about the company.")])
    by, _ = run(tmp_path, tmp_path / "bare.docx")
    assert by["TP-15"].status == "EXC" and by["TP-28"].status == "EXC" and by["TP-40"].status == "EXC" and by["TP-68"].status == "EXC" and by["TP-39"].status == "EXC"


def test_pli_only_ec_oc_is_flagged(tmp_path):
    _doc(tmp_path / "e.docx", [("p", "The EC/OC indicator is used because it is appropriate.")])
    by, _ = run(tmp_path, tmp_path / "e.docx")
    assert "EC/OC" in titles(by["TP-28"])


def test_profile_vs_fs(tmp_path):
    _doc(tmp_path / "r.docx", [("p", "ABC was incorporated on 5 March 2015 as an FZE. It is wholly owned (100%) by its shareholder. Principal activity: trading in building materials and hardware.")])
    _doc(tmp_path / "fs.docx", [("p", "ABC was incorporated on 9 June 2016 as a limited liability company (L.L.C.). The shareholder owns 51% of the share capital. Principal activity: software consulting services.")])
    by, _ = run(tmp_path, tmp_path / "r.docx", fs_path=tmp_path / "fs.docx")
    t = titles(by["TP-10"])
    assert by["TP-10"].status == "EXC" and "Legal form" in t and "incorporation" in t and "Shareholding" in t


def test_rem_split_vs_birf(tmp_path):
    _doc(tmp_path / "r.docx", [("p", "Total managerial remuneration was AED 1,000,000. Salary AED 600,000.")])
    _doc(tmp_path / "birf.docx", [("p", "Salary 600,000"), ("p", "Bonus 250,000"), ("p", "Allowances 100,000")])
    by, _ = run(tmp_path, tmp_path / "r.docx", birf_path=tmp_path / "birf.docx")
    t = titles(by["TP-16"])
    assert by["TP-16"].status == "EXC" and "bonus" in t.lower() and "add up" in t
    by, _ = run(tmp_path, tmp_path / "r.docx")
    assert by["TP-16"].status == "NOTRUN" and "BIRF" in by["TP-16"].note


def test_rejected_reasons(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Rejected"
    ws.append(["Company", "Rejection reason"]); ws.append(["A Ltd", "Not independent"]); ws.append(["B Ltd", None]); ws.append(["C Ltd", "n/a"])
    wb.save(tmp_path / "w.xlsx"); make_report(tmp_path / "r.docx")
    by, _ = run(tmp_path, tmp_path / "r.docx", tmp_path / "w.xlsx")
    assert by["TP-38"].status == "EXC" and "no recorded reason" in titles(by["TP-38"]) and "not specific" in titles(by["TP-38"])


def test_secondary_conclusion_contradiction(tmp_path):
    _doc(tmp_path / "r.docx", [("p", "The tested party PLI is within the interquartile range under the TNMM. "
                                     "The market salary survey shows the remuneration is above the market range and is not consistent with arm's length.")])
    by, _ = run(tmp_path, tmp_path / "r.docx")
    assert by["TP-43"].status == "EXC" and "contradicts" in titles(by["TP-43"])


# ------------------------------------------------------------------ 1.10 charts
CHART = """<?xml version="1.0"?><c:chartSpace xmlns:c="c" xmlns:a="a"><c:chart><c:plotArea><c:barChart><c:ser><c:cat><c:strRef><c:strCache>%s</c:strCache></c:strRef></c:cat>
<c:val><c:numRef><c:numCache>%s</c:numCache></c:numRef></c:val></c:ser></c:barChart><c:valAx><c:scaling><c:min val="2"/></c:scaling></c:valAx></c:plotArea></c:chart></c:chartSpace>"""


def _with_chart(src, dst, cats, vals):
    cx = CHART % ("".join(f"<c:pt><c:v>{x}</c:v></c:pt>" for x in cats), "".join(f"<c:pt><c:v>{x}</c:v></c:pt>" for x in vals))
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in zi.infolist():
            zo.writestr(i, zi.read(i.filename))
        zo.writestr("word/charts/chart1.xml", cx)


def test_chart_rules(tmp_path):
    make_workbook(tmp_path / "w.xlsx"); make_report(tmp_path / "r0.docx")
    _with_chart(tmp_path / "r0.docx", tmp_path / "r.docx", ["Tested party", "A", "B", "C", "D"], [3.2, 2.0, 4.0, 5.0, 6.0])
    by, _ = run(tmp_path, tmp_path / "r.docx", tmp_path / "w.xlsx")
    t = titles(by["TP-53"])
    assert by["TP-53"].status == "EXC" and "does not start at zero" in t and "axis title missing" in t and "no title" in t
    assert by["TP-52"].status == "EXC" and "different number of points" in titles(by["TP-52"])


# ------------------------------------------------------------------ 1.11 naming
def test_naming_consistency(tmp_path):
    _doc(tmp_path / "r.docx", [("p", "GOLDEN CREST TRADING DMCC is the tested party."), ("p", "GOLDEN CREST TRADNG DMCC earns a margin."),
                               ("p", "The OP/OC is used here and the OP/OR is used there."), ("p", "For the year ended 31 December 2025 and the year ended 31 December 2024.")])
    by, _ = run(tmp_path, tmp_path / "r.docx")
    t = titles(by["TP-56"])
    assert by["TP-56"].status == "EXC" and "named identically" in t and "PLI is named differently" in t and "Period" in t


# ------------------------------------------------------------------ 1.12 financial update
def _update(path, cur_extra=(), prior_extra=(), cur_comps=("Alpha Corp", "Beta Corp", "Gamma Corp"), prior_comps=("Alpha Corp", "Beta Corp", "Delta Corp"),
            cur_q=(2.0, 4.0, 6.0), prior_q=(2.0, 4.0, 6.0), cur_rem="2,000,000", prior_rem="2,000,000"):
    cur = [("p", "Transfer Pricing Financial Update"), ("p", "For the year ended 31 December 2025"), ("p", "Table of Contents"), ("p", "Introduction\t3"), ("h1", "Introduction"),
           ("p", f"Total managerial remuneration is AED {cur_rem}. Database extracted on 15 March 2026. Mr. Ali Hassan is the director."), ("p", "This update must be read with the prior report."),
           ("p", f"Lower quartile {cur_q[0]}%. Median {cur_q[1]}%. Upper quartile {cur_q[2]}%."), *cur_extra]
    prior = [("p", "Transfer Pricing Benchmark Analysis"), ("p", "For the year ended 31 December 2022"), ("p", "Table of Contents"), ("p", "Introduction\t3"), ("h1", "Introduction"),
             ("p", f"Total managerial remuneration is AED {prior_rem}. Mr. Ali Hassan and Mr. Omar Said are directors. Database: S&P Capital IQ. TNMM. OP/OC."),
             ("p", f"Lower quartile {prior_q[0]}%. Median {prior_q[1]}%. Upper quartile {prior_q[2]}%."), *prior_extra]
    d = docx.Document()
    ct = [["Company", "PLI"]] + [[n, f"{3 + i}.0%"] for i, n in enumerate(cur_comps)]
    pt = [["Company", "PLI"]] + [[n, f"{3 + i}.0%"] for i, n in enumerate(prior_comps)]
    for kind, text in cur:
        d.add_paragraph(text) if kind == "p" else d.add_heading(text, 1)
    t = d.add_table(rows=0, cols=2)
    for r in ct:
        c = t.add_row().cells; c[0].text, c[1].text = r
    for kind, text in prior:
        d.add_paragraph(text) if kind == "p" else d.add_heading(text, 1)
    t = d.add_table(rows=0, cols=2)
    for r in pt:
        c = t.add_row().cells; c[0].text, c[1].text = r
    d.save(path)


def test_update_changes_are_flagged(tmp_path):
    _update(tmp_path / "u.docx", cur_q=(3.5, 6.5, 9.0), cur_rem="3,500,000")
    by, ctx = run(tmp_path, tmp_path / "u.docx")
    assert K.report_kind(ctx) == "update"
    assert by["TP-58"].status in ("EXC", "REVIEW") and by["TP-58"].findings
    assert by["TP-63"].status == "REVIEW" and "dropped 1" in by["TP-63"].findings[0].detail
    assert by["TP-64"].status == "EXC" and "Delta" in by["TP-64"].findings[0].detail
    assert by["TP-65"].status == "EXC" and "Gamma" in by["TP-65"].findings[0].detail
    assert by["TP-66"].status == "EXC" and "moved materially" in titles(by["TP-66"])
    assert by["TP-67"].status == "EXC"
    assert by["TP-59"].status == "EXC" and "Omar" in by["TP-59"].findings[0].detail
    assert by["TP-61"].status == "EXC" and "database" in by["TP-61"].findings[0].detail


def test_update_with_reasons_passes_to_review(tmp_path):
    ex = [("p", "Delta Corp was dropped because it was acquired and delisted. Gamma Corp is newly added as it now meets the independence criteria. "
                "The remuneration is unchanged compared with the prior year. The tested party is within the interquartile range. There has been no change in method. The criteria changed: the search was revised.")]
    _update(tmp_path / "u.docx", cur_extra=ex)
    by, _ = run(tmp_path, tmp_path / "u.docx")
    assert by["TP-64"].status == "REVIEW" and by["TP-65"].status == "REVIEW"
    assert by["TP-66"].status in ("REVIEW", "PASS") and by["TP-67"].status in ("REVIEW", "PASS")


def test_validity_cycle(tmp_path):
    _update(tmp_path / "u.docx")
    by, _ = run(tmp_path, tmp_path / "u.docx")
    assert by["TP-58"].status == "EXC" and "refresh cycle" in titles(by["TP-58"])       # prior 2022 vs 2025 = 3 years


def test_rolled_forward_workbook(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(["Company", "FY2023", "FY2024"]); ws.append(["A", 1, 2]); wb.save(tmp_path / "w.xlsx")
    _update(tmp_path / "u.docx", cur_extra=[("p", "Nothing about extraction here.")])
    by, _ = run(tmp_path, tmp_path / "u.docx", tmp_path / "w.xlsx")
    assert by["TP-62"].status in ("EXC", "REVIEW", "PASS")
    wbk = openpyxl.load_workbook(tmp_path / "w.xlsx")
    assert by["TP-62"].status == "EXC" and "rolled forward" in titles(by["TP-62"])


# ------------------------------------------------------------------ 1.13 regulatory
def test_law_references(tmp_path):
    _doc(tmp_path / "r.docx", [("p", "Under Federal Decree-Law No. 47 of 2022 and the OECD Transfer Pricing Guidelines 2017, with Cabinet Decision No. 999 of 2024. "
                                     "The disclosure threshold is AED 10 million for related party transactions.")])
    by, _ = run(tmp_path, tmp_path / "r.docx")
    t = titles(by["TP-70"])
    assert by["TP-70"].status == "EXC" and "Outdated" in t and "not in the firm's reference list" in t and "Threshold quoted" in t


def test_connected_persons_with_birf(tmp_path):
    _doc(tmp_path / "r.docx", [("p", "Remuneration is paid to Mr. Ali Hassan. Related party transactions are disclosed.")])
    _doc(tmp_path / "birf.docx", [("p", "Recipient: Mr. Omar Said, Salary 100,000")])
    by, _ = run(tmp_path, tmp_path / "r.docx", birf_path=tmp_path / "birf.docx")
    assert by["TP-69"].status == "EXC" and "BIRF" in titles(by["TP-69"])


def test_every_row_has_a_test():
    assert all(c["impl"] for c in K.CHECKS)


def test_workbooks_are_closed_after_reading(tmp_path, monkeypatch):
    """Windows keeps a read-only workbook locked, which broke deleting the temp folder (WinError 32)."""
    import openpyxl
    from auditsignoff import rules_tp3
    make_workbook(tmp_path / "w.xlsx")
    opened = []
    real = openpyxl.load_workbook

    def spy(*a, **k):
        wb = real(*a, **k); wb._closed = False; orig = wb.close
        def close(): wb._closed = True; orig()
        wb.close = close; opened.append(wb); return wb
    monkeypatch.setattr(openpyxl, "load_workbook", spy)
    rules_tp3._wb_rows(tmp_path / "w.xlsx")
    assert opened and all(w._closed for w in opened)


def test_birf_template_roundtrip(tmp_path):
    from auditsignoff.birf_template import build_birf_template
    f = tmp_path / "birf.xlsx"
    build_birf_template(f)
    wb = openpyxl.load_workbook(f); ws = wb["BIRF"]
    ws["A9"], ws["B9"], ws["C9"], ws["E9"], ws["F9"], ws["G9"] = "Mr. Ali Hassan", "Managing Director", "Director", 600000, 250000, 100000
    ws["A10"], ws["E10"] = "Ms. Sara Khan", 200000
    wb.save(f)
    _doc(tmp_path / "r.docx", [("p", "Total managerial remuneration was AED 1,150,000. Salary AED 800,000. Bonus AED 250,000. Allowances AED 100,000. Paid to Mr. Ali Hassan and Ms. Sara Khan.")])
    by, _ = run(tmp_path, tmp_path / "r.docx", birf_path=f)
    assert by["TP-16"].status == "PASS", titles(by["TP-16"])
    assert by["TP-69"].status in ("PASS", "REVIEW")
    _doc(tmp_path / "r2.docx", [("p", "Total managerial remuneration was AED 900,000. Salary AED 800,000. Paid to Mr. Ali Hassan.")])
    by, _ = run(tmp_path, tmp_path / "r2.docx", birf_path=f)
    assert by["TP-16"].status == "EXC" and "bonus" in titles(by["TP-16"]).lower() and "add up" in titles(by["TP-16"])
    assert by["TP-69"].status == "EXC" and "not in the report" in titles(by["TP-69"])      # Sara Khan missing
