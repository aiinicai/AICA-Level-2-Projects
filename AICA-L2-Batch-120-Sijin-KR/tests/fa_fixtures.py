import random
from datetime import datetime, timedelta
from pathlib import Path
import openpyxl


def make_gl(path, seed=7, issues=True):
    rnd = random.Random(seed)
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Date", "Created On", "Voucher No", "Account", "Description", "Debit", "Credit", "Posted By", "Approved By", "Vendor"])
    users = ["amal", "bashir", "chitra", "dev"]
    day = datetime(2025, 1, 2)
    n = 0
    for i in range(420):
        d = day + timedelta(days=int(i * 360 / 420))
        while d.weekday() >= 5:
            d += timedelta(days=1)
        amt = round(10 ** rnd.uniform(2.2, 4.5) * rnd.uniform(1, 1.9), 2)
        n += 1
        u = rnd.choice(users)
        v = f"JV{1000 + n}"
        ws.append([d, d, v, "Expenses", "Supplier invoice", amt, 0, u, "approver1", f"Vendor {rnd.randint(1, 30)}"])
        ws.append([d, d, v, "Payables", "Supplier invoice", 0, amt, u, "approver1", f"Vendor {rnd.randint(1, 30)}"])
    if issues:
        sat = datetime(2025, 6, 7)
        for k in range(6):                                                   # round weekend postings
            ws.append([sat, sat, f"JV{5000 + k}", "Consultancy", "Consultancy fee cash", 55000, 0, "dev", "dev", "Mohammed Al Hashmi"])
            ws.append([sat, sat, f"JV{5000 + k}", "Bank", "Consultancy fee cash", 0, 55000, "dev", "dev", "Mohammed Al Hashmi"])
        for k in range(6):                                                   # just below 50,000 limit
            ws.append([datetime(2025, 3, 3 + k, 0, 0) + timedelta(days=0), datetime(2025, 3, 3 + k, 0, 0) + timedelta(days=0), f"JV{5100 + k}", "Expenses", "Services", 48800 + k * 50, 0, "bashir", "approver1", "Vendor 5"])
            ws.append([datetime(2025, 3, 3 + k, 0, 0) + timedelta(days=0), datetime(2025, 3, 3 + k, 0, 0) + timedelta(days=0), f"JV{5100 + k}", "Payables", "Services", 0, 48800 + k * 50, "bashir", "approver1", "Vendor 5"])
        ws.append([datetime(2025, 12, 20), datetime(2026, 2, 14), "JV5200", "Revenue", "Adjustment", 0, 90000, "dev", "approver1", "Vendor 9"])   # backdated
        ws.append([datetime(2025, 12, 20), datetime(2026, 2, 14), "JV5200", "Receivable", "Adjustment", 90000, 0, "dev", "approver1", "Vendor 9"])
        for _ in range(2):                                                    # duplicate
            ws.append([datetime(2025, 4, 8), datetime(2025, 4, 8), "JV5300", "Expenses", "Rent April", 31500, 0, "amal", "approver1", "Vendor 2"])
            ws.append([datetime(2025, 4, 8), datetime(2025, 4, 8), "JV5300", "Payables", "Rent April", 0, 31500, "amal", "approver1", "Vendor 2"])
    wb.save(path)


def make_masters(folder):
    folder = Path(folder)
    v = openpyxl.Workbook(); w = v.active
    w.append(["Vendor Name", "IBAN", "Address", "Phone", "TRN"])
    w.append(["Alpha Supplies LLC", "AE070331234567890123456", "Dubai, Al Quoz", "0501111111", "100111111111113"])
    w.append(["Bright Services FZE", "AE070331234567890123456", "Sharjah, SAIF", "0502222222", "100222222222223"])
    w.append(["Crest Logistics LLC", "AE120260001234567890123", "Dubai, Jebel Ali", "0503333333", "100333333333333"])
    v.save(folder / "vendors.xlsx")
    e = openpyxl.Workbook(); w = e.active
    w.append(["Employee Name", "IBAN", "Address", "Phone"])
    w.append(["Rahul Menon", "AE120260001234567890123", "Dubai, Karama", "0504444444"])
    e.save(folder / "employees.xlsx")
