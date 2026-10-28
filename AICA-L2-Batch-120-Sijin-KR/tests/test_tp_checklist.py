"""Module C against the firm's updated TP checklist (managerial remuneration)."""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import openpyxl
import pytest
from tp_fixtures import make_report, make_workbook
from auditsignoff import tp_checklist as K
from auditsignoff.models import load_catalogue, load_params
from auditsignoff.rules_fs import NotRun
from auditsignoff.rules_tp import run_tp

P = load_params()
ENG = {"fy": 2025, "client": {"name": "GOLDEN CREST TRADING DMCC", "authority": "DMCC", "shareholders": [], "signatories_legal": []}}


def test_catalogue_is_the_new_checklist():
    c = [x for x in load_catalogue() if x["mod"] == "C"]
    assert len(c) == len(K.CHECKS) == 79
    secs = [x["grp"].split()[0] for x in c]
    assert [s for s in dict.fromkeys(secs)] == ["1.1", "1.2", "1.3", "1.4", "1.5", "1.6", "1.7", "1.8", "1.9", "1.10", "1.11", "1.12", "1.13", "1.14"]
    assert sum(1 for x in c if x["grp"].startswith("1.8")) == 9 and sum(1 for x in c if x["grp"].startswith("1.12")) == 10
    assert not any("TPX" in x["src"] for x in c)


def test_full_report_marks_update_rows_na_and_vice_versa(tmp_path):
    make_workbook(tmp_path / "w.xlsx"); make_report(tmp_path / "r.docx")
    res, _ = run_tp(tmp_path / "r.docx", tmp_path / "w.xlsx", ENG, P)
    by = {r.check_id: r for r in res}
    upd = [c["id"] for c in K.CHECKS if c["applies"] == "update" and c["impl"]]
    full = [c["id"] for c in K.CHECKS if c["applies"] == "full" and c["impl"]]
    assert upd and all(by[i].status == "NA" for i in upd)
    assert full and not any(by[i].status == "NA" for i in full)


def test_financial_update_flips_applicability(tmp_path):
    make_workbook(tmp_path / "w.xlsx")
    make_report(tmp_path / "r.docx", header_title="Transfer Pricing Financial Update")
    import docx
    d = docx.Document(tmp_path / "r.docx"); d.paragraphs[0].insert_paragraph_before("Transfer Pricing Financial Update"); d.save(tmp_path / "u.docx")
    res, ctx = run_tp(tmp_path / "u.docx", tmp_path / "w.xlsx", ENG, P)
    assert K.report_kind(ctx) == "update"
    by = {r.check_id: r for r in res}
    assert all(by[c["id"]].status == "NA" for c in K.CHECKS if c["applies"] == "full" and c["impl"])


def test_workbook_integrity(tmp_path):
    make_workbook(tmp_path / "ok.xlsx")
    c = type("C", (), {})(); c.wb_path = tmp_path / "ok.xlsx"
    assert K.workbook_integrity(c) == []
    wb = openpyxl.load_workbook(tmp_path / "ok.xlsx")
    wb["Tested party"].sheet_state = "hidden"; wb["Checks"]["A2"] = "=#REF!+1"
    wb.save(tmp_path / "bad.xlsx"); c.wb_path = tmp_path / "bad.xlsx"
    t = " ".join(f.title for f in K.workbook_integrity(c))
    assert "Hidden" in t and "broken references" in t
    wb = openpyxl.Workbook(); wb.active.append(["x", 1]); wb.save(tmp_path / "hard.xlsx"); c.wb_path = tmp_path / "hard.xlsx"
    assert "no formulas" in K.workbook_integrity(c)[0].title
    c.wb_path = None
    with pytest.raises(NotRun):
        K.workbook_integrity(c)


class _Doc:
    def __init__(self, t): self.all_text = t; self.kind = "docx"; self.pages = []


class _Ctx:
    def __init__(self, rep, fs=None):
        self.doc = _Doc(rep); self.fs = _Doc(fs) if fs is not None else None
    def full_text(self): return self.doc.all_text


def test_remuneration_trace():
    ok = K.rem_trace(_Ctx("Total managerial remuneration was AED 3,000,000.", "Key management compensation 3,000,000"))
    assert ok == []
    bad = K.rem_trace(_Ctx("Director remuneration AED 3,000,000", "Key management compensation 2,700,000"))
    assert bad and bad[0].severity == "High" and "Grey area" in bad[0].detail
    with pytest.raises(NotRun):
        K.rem_trace(_Ctx("no figures here", "remuneration 2,700,000"))
    with pytest.raises(NotRun):
        K.rem_trace(_Ctx("remuneration 3,000,000"))


def test_method_terms_flags_inconsistent_naming():
    f = K.method_terms(_Ctx("The TNMM is applied using OP/OC. Conclusion: Comparable Profits Method gives OP/OR."))[0]
    assert "Comparable Profits Method" in f.detail and "Both OP/OC and OP/OR" in f.detail
