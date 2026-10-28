"""Generate a synthetic FS pack (clean + seeded-error), trial balances, engagement and liquidation inputs.
Everything here is fictitious test data."""
from __future__ import annotations
import sys
from pathlib import Path

import docx
import openpyxl
import yaml
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

NAME = "GOLDEN CREST TRADING DMCC"
SH = ["Mohammed Al Hashmi", "Fatima Khan", "Rajan Pillai"]


def _font(style_or_run, size=None):
    style_or_run.font.name = "Times New Roman"
    if size:
        style_or_run.font.size = Pt(size)


def build_fs(path, errors: bool):
    d = docx.Document()
    st = d.styles["Normal"]
    _font(st, 10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    s = d.sections[0]
    s.left_margin = Inches(1.0 if errors else 1.3)
    s.top_margin, s.right_margin, s.bottom_margin = Inches(1.5), Inches(1.0), Inches(1.0)
    hp = s.header.paragraphs[0]; hp.text = NAME

    def para(text, size=10, bold=False, justify=True, font=None):
        p = d.add_paragraph()
        r = p.add_run(text); r.bold = bold; _font(r, size)
        if font:
            r.font.name = font
        if justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def table(rows, header=None):
        t = d.add_table(rows=0, cols=len(rows[0])); t.style = "Table Grid"
        for r in ([header] if header else []) + rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(v)); _font(run, 10)
                if header and r is header and str(v) == "2025":
                    run.bold = True
        return t

    # cover
    para(NAME, 14, True, False)
    para("Financial statements for the year ended 31 December 2025", 12, True, False)
    # directors report
    para("Directors' report", 12, True, False)
    para(f"The Directors present their report together with the audited financial statements of {NAME} for the year ended 31 December 2025. "
         "The principal activity of the Company is general trading. The Company is owned by three shareholders and is managed by its Directors.",
         justify=not errors)
    para("Signed on behalf of the Directors by Mohamed Al Hashmi." if errors else "Signed on behalf of the Directors by Mohammed Al Hashmi, Fatima Khan and Rajan Pillai.", justify=False)
    # auditor's report
    para("Independent Auditor's Report", 11, True, False)
    para(f"To the Shareholder of {NAME}" if errors else f"To the Shareholders of {NAME}", 11, False, False)
    para("Opinion. We have audited the financial statements of the Company, which comprise the statement of financial position as at "
         "31 December 2025, and the statements of profit or loss, changes in equity and cash flows for the year then ended.", 11)
    para("Report on other legal and regulatory requirements. As required by the applicable laws and regulations of "
         + ("DED" if errors else "DMCC") + ", we report that the Company has maintained proper books of account.", 11)
    para("For the audit firm", 11, True, False)
    if not errors:
        para("Registration No: 1234", 11, False, False)
        para("Ref: RG/AR/2026/045", 11, False, False)
        para("Dubai, 15 March 2026", 11, False, False)
    # SOFP
    para("Statement of financial position", 12, True, False)
    para("as at 31 December 2025" if not errors else "as at 30 December 2025", 11, False, False)
    ta_cy = "1,010,000" if errors else "1,000,000"
    table([
        ["Cash and bank balances", "5", "80,000", "60,000"],
        ["Trade receivables", "6", "420,000", "380,000"],
        ["Inventories", "7", "200,000", "190,000"],
        ["Property, plant and equipment", "8", "300,000", "270,000"],
        ["Total assets", "", ta_cy, "900,000"],
        ["Share capital", "9", "300,000", "300,000"],
        ["Retained earnings", "", "250,000", "190,000"],
        ["Total equity", "", "550,000", "490,000"],
        ["Trade and other payables", "10", "350,000", "330,000"],
        ["Provision for staff terminal benefits", "11", "50,000", "40,000"],
        ["Duties and taxes payable", "12", "50,000", "40,000"],
        ["Total equity and liabilities", "", "1,000,000", "900,000"],
    ], header=["", "Note", "2025", "2024"])
    # P&L
    para("Statement of profit or loss", 12, True, False)
    table([
        ["Revenue", "", "2,000,000", "1,800,000"], ["Cost of revenue", "", "(1,400,000)", "(1,300,000)"],
        ["Gross profit", "", "600,000", "500,000"], ["General and administrative expenses", "", "(500,000)", "(440,000)"],
        ["Finance cost", "", "(40,000)", "(30,000)"], ["Profit for the year", "", "60,000", "30,000"]],
        header=["", "Note", "2025", "2024"])
    # CF
    para("Statement of cash flows", 12, True, False)
    table([
        ["Cash flows from operating activities", "", "", ""], ["Net cash from operating activities", "", "95,000", "70,000"],
        ["Net increase in cash and cash equivalents", "", "20,000", "10,000"],
        ["Cash and cash equivalents at the end of the year", "", "85,000" if errors else "80,000", "60,000"]],
        header=["", "Note", "2024" if errors else "2025", "2023" if errors else "2024"])
    # SOCE
    para("Statement of changes in equity", 12, True, False)
    table([["Balance at 1 January 2025", "300,000", "190,000", "490,000"], ["Profit for the year", "-", "60,000", "60,000"],
           ["Balance at 31 December 2025", "300,000", "250,000", "550,000"]],
          header=["", "Share capital", "Retained earnings", "Total"])
    # notes
    para("Notes to the financial statements", 12, True, False)
    para(f"{'GOLDEN CREST TRADNG DMCC' if errors else NAME} (the 'Company') is a free zone company incorporated in Dubai and licensed by DMCC. "
         "The shareholders are Mohammed Al Hashmi, Fatima Khan and Rajan Pillai.", justify=not errors)
    if errors:
        para("Golden Crest Trading DMCC maintains its records in AED.", justify=False)
        p = para("A paragraph typed in the wrong font and size.", 9, font="Arial")
    for n, t in ((5, "Cash and bank balances"), (7, "Inventories"), (9, "Share capital")):
        para(f"{n}. {t}", 10, True, False)
    para("6. Trade receivables", 10, True, False)
    para("8. Property, plant and equipment", 10, True, False)
    for n in range(10, 16):
        para(f"{n}. Other balance (note {n})", 10, True, False)
    para("Cash and bank balances are held with banks in the UAE.")
    para("Property, plant and equipment are depreciated over their useful lives of 3 to 5 years.")
    if not errors:
        para("Provision for end of service benefits is made in accordance with UAE labour law.")
    para("Trade receivables are stated net of the expected credit loss (IFRS 9 simplified approach). Ageing of trade receivables:")
    table([["Receivables", "0-30 days", "31-60 days", "Over 60 days"], ["Trade receivables", "300,000", "80,000", "40,000"]])
    para("The Company is registered for VAT; duties and taxes payable represent net VAT due to the FTA.")
    para("The Company is subject to UAE corporate tax at the applicable rate; no deferred tax is recognised.")
    std = ("Standards issued but not yet effective: IFRS 19 Subsidiaries without Public Accountability; Amendments to the "
           "Classification and Measurement of Financial Instruments (IFRS 9 and IFRS 7)." if errors else
           "Standards issued but not yet effective: IFRS 18 Presentation and Disclosure in Financial Statements; IFRS 19 Subsidiaries "
           "without Public Accountability; Amendments to the Classification and Measurement of Financial Instruments (IFRS 9 and IFRS 7).")
    para(std)
    d.save(path)


def build_tb(path, errors: bool):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Account", "Class", "Debit", "Credit"])
    rows = [("Cash and bank", "Asset", 80000, 0), ("Trade receivables", "Asset", 420000, 0), ("Inventory", "Asset", 200000, 0),
            ("PPE", "Asset", 300000, 0), ("Trade payables", "Liability", 0, 350000), ("EOSB provision", "Liability", 0, 50000),
            ("VAT payable", "Liability", 0, 50000), ("Share capital", "Equity", 0, 300000), ("Retained earnings b/f", "Equity", 0, 190000),
            ("Revenue", "Income", 0, 2000000), ("Cost of revenue", "Expense", 1400000, 0),
            ("General and administrative", "Expense", 520000 if errors else 500000, 0), ("Finance cost", "Expense", 40000, 0)]
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def engagement(path, errors: bool):
    eng = {"fy": 2025, "py": 2024,
           "client": {"name": NAME, "legal_status": "FZC", "authority": "DMCC", "emirate": "Dubai", "year_end": "31 December",
                      "shareholders": SH, "signatories_legal": SH},
           "fs_signatories": ["Mohamed Al Hashmi"] if errors else SH,
           "br_on_file": False,
           "review_log": {
               "preparer": {"done": True, "date": "2026-03-01"}, "supervisor": {"done": True, "date": "2026-03-05"},
               "manager": {"done": False} if errors else {"done": True, "date": "2026-03-10"},
               "ready_before_print": {"done": False} if errors else {"done": True, "date": "2026-03-12"}}}
    Path(path).write_text(yaml.safe_dump(eng, sort_keys=False), encoding="utf-8")


def liquidation(path, case):
    cases = {
        "difc_breach": {"company": "SAMPLE HOLDINGS LTD", "authority": "difc",
            "br": {"filed_with_authority": True, "liquidator_named_correctly": True},
            "dates": {"board_resolution": "2026-06-01", "liquidator_appointment": "2026-06-05", "fta_notified": None,
                      "vat_dereg_initiated": "2026-08-20", "newspaper_ads": ["2026-09-01", "2026-09-02", "2026-09-03"],
                      "report_issue_planned": "2026-10-05", "bank_closure": "2026-09-20"},
            "nocs": {"telecom": "2026-08-01", "utilities": None, "chiller": "na", "landlord": "2026-08-10", "po_box": "2026-08-12",
                     "vat_deregistration": None, "ct_deregistration": None, "authority_portal": None, "labour_visa": None,
                     "establishment_card": None, "esr_aml": None, "customs": "na"},
            "bank_accounts": [{"bank": "Bank A", "location": "UAE", "no_liability_letter": "2026-09-10", "closure_letter": "2026-09-20"},
                              {"bank": "Bank B", "location": "UK", "no_liability_letter": None, "closure_letter": None}]},
        "dso_short_gap": {"company": "SAMPLE FZCO", "authority": "dso",
            "br": {"filed_with_authority": True, "liquidator_named_correctly": False},
            "dates": {"board_resolution": "2026-10-01", "liquidator_appointment": "2026-10-02", "report_issue_planned": "2026-10-10",
                      "trade_licence_expiry": "2026-08-15"},
            "nocs": {}, "bank_accounts": []},
        "dmcc_clean": {"company": "TIDY TRADING DMCC", "authority": "dmcc",
            "br": {"filed_with_authority": True, "liquidator_named_correctly": True},
            "dates": {"board_resolution": "2026-03-02", "liquidator_appointment": "2026-03-03", "fta_notified": "2026-03-20",
                      "vat_dereg_initiated": "2026-03-15", "report_submitted": "2026-09-15", "bank_closure": "2026-09-10"},
            "nocs": {k: "2026-05-01" for k in ["telecom", "utilities", "landlord", "po_box", "vat_deregistration", "ct_deregistration",
                                                "authority_portal", "labour_visa", "establishment_card", "esr_aml"]} | {"chiller": "na", "customs": "na"},
            "bank_accounts": [{"bank": "Bank A", "location": "UAE", "no_liability_letter": "2026-08-01", "closure_letter": "2026-09-10"}]},
    }
    Path(path).write_text(yaml.safe_dump(cases[case], sort_keys=False), encoding="utf-8")


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "samples")
    out.mkdir(exist_ok=True, parents=True)
    for tag, err in (("clean", False), ("errors", True)):
        build_fs(out / f"fs_{tag}.docx", err)
        build_tb(out / f"tb_{tag}.xlsx", err)
        engagement(out / f"engagement_{tag}.yaml", err)
    for c in ("difc_breach", "dso_short_gap", "dmcc_clean"):
        liquidation(out / f"liq_{c}.yaml", c)
    # forensic (Module E): the ZENITH TRADING FZCO fraud case (bank, signatures, FS, stock) - see make_forensic_case.py
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from make_forensic_case import build
    build(out / "forensic_case", fraud=True)
    print("samples written to", out)
