"""Builds the SYNTHETIC ABC COMPANY transfer-pricing test set (managerial remuneration, external CUP benchmark).

Everything here is invented: company, people, comparables and pay figures. Only the shape of a real study is copied (final set of 27,
interquartile range 1,103,333 / 1,525,809 / 2,775,669). No client data and no database extract is included.

Usage:  python samples/make_abc_tp_sample.py [output_dir]        (default: samples/abc_sample)
Files:  ABC_COMPANY_FY2024_TP_Report.docx                  clean report            ABC_COMPANY_FY2024_TP_Benchmark_Workbook.xlsx          clean workbook (live formulas)
        ABC_FS_Notes.docx, ABC_BIRF.xlsx    audited FS notes and the BIRF        ABC_COMPANY_FY2024_TP_Report_ERRORS.docx / ..._Workbook_ERRORS.xlsx  seeded defects
        engagement_abc.yaml                 fact sheet
Run:    python -m auditsignoff tp --engagement samples/abc_sample/engagement_abc.yaml --report samples/abc_sample/ABC_COMPANY_FY2024_TP_Report.docx \
            --workbook samples/abc_sample/ABC_COMPANY_FY2024_TP_Benchmark_Workbook.xlsx --fs samples/abc_sample/ABC_FS_Notes.docx --birf samples/abc_sample/ABC_BIRF.xlsx
"""
from __future__ import annotations
import random
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl.styles import Font, PatternFill

NAME = "ABC COMPANY"
FY = 2024
KMP = [("Mr. PQR", 1_200_000), ("Mr. XYZ", 1_500_000)]
TOTAL = sum(a for _, a in KMP)

# ---- 27 fictional comparables, sorted values pinned so that QUARTILE.INC gives the report's range ---------------------------------------------
PINNED = {1: 207_793, 7: 1_103_000, 8: 1_103_666, 14: 1_525_809, 20: 2_775_000, 21: 2_776_338, 27: 11_887_150}


def comp_values():
    v = [None] * 28
    for k, x in PINNED.items():
        v[k] = x
    keys = sorted(PINNED)
    for a, b in zip(keys, keys[1:]):
        for i in range(a + 1, b):
            v[i] = round(PINNED[a] + (PINNED[b] - PINNED[a]) * (i - a) / (b - a), -3) + (i % 7) * 113
    return v[1:]


COMPANIES = ["Norvik Building Supplies AB", "Castellan Materials plc", "Harrowgate Timber & Stone Ltd", "Brenner Baustoffe AG", "Valmont Trade SA",
             "Ostrander Construction Goods Oy", "Lindqvist Hardware AB", "Pemberton Merchants plc", "Corvalis Cement Trading SpA", "Hallström Bygg AB",
             "Whitcombe Roofing Supplies plc", "Marnet Materiaux SAS", "Dunmore Builders Merchants Ltd", "Tessaly Aggregates SA", "Rovena Steel Distribution Srl",
             "Kestrel Plumbing & Tile plc", "Almqvist Byggvaror AB", "Sandford Joinery Group plc", "Bellaire Matériaux SA", "Oakridge Building Products Ltd",
             "Verhoeven Bouwgroep NV", "Saltmarsh Timber plc", "Grunwald Handel GmbH", "Ferrant Distribution SA", "Eskilstuna Järn & Bygg AB",
             "Montrose Tiles & Surfaces plc", "Kallio Rakennustarvike Oy"]
ROLES = ["President & CEO", "Managing Director", "MD & Executive Director", "CEO & Director", "Managing Director & CEO", "President & CEO",
         "Group Managing Director", "CEO", "Managing Director", "CEO & President"]


def comparables():
    vals = comp_values()
    rnd = random.Random(2024)
    rows = []
    for i, (c, v) in enumerate(zip(COMPANIES, vals), 1):
        rows.append((f"Executive {i:02d}", c, ROLES[i % len(ROLES)], round(rnd.uniform(40, 2600), 1), int(v)))
    rnd.shuffle(rows)
    return rows


def quartiles(vals):
    import statistics
    from auditsignoff.rules_tp import quartiles as q
    r = q(vals, "INC")
    return {k: round(r[k]) for k in r}


# ------------------------------------------------------------------ workbook
def make_workbook(path, errors=False):
    rows = comparables()
    wb = openpyxl.Workbook()
    hd = Font(bold=True); fill = PatternFill("solid", fgColor="DDEBF7")

    s = wb.active; s.title = "Summary"
    s.append(["ABC COMPANY - Managerial remuneration benchmark (External CUP)  -  SYNTHETIC SAMPLE, FY 2024"]); s["A1"].font = Font(bold=True, size=13)
    s.append([])
    s.append(["Arm's length range (AED)", "Value"])
    n = len(rows) + 1
    rng = f"Final_Set!$E$2:$E${n}"
    for lab, f in (("Lowest value", f"=MIN({rng})"), ("Lower quartile (25%)", f"=_xlfn.QUARTILE.INC({rng},1)"), ("Median", f"=MEDIAN({rng})"),
                   ("Upper quartile (75%)", f"=_xlfn.QUARTILE.INC({rng},3)"), ("Highest value", f"=MAX({rng})"), ("Number of comparables", f"=COUNT({rng})")):
        s.append([lab, 2700000 if (errors and lab.startswith("Upper")) else f])
    s.append([])
    s.append(["Tested party - remuneration paid in FY 2024", "AED"])
    first = s.max_row + 1
    for who, amt in KMP:
        s.append([f"{who} - managerial remuneration", amt])
    s.append(["Total managerial remuneration", f"=SUM(B{first}:B{first + len(KMP) - 1})" if not errors else 2_750_000])
    tot = s.max_row
    s.append([])
    for who, r in zip([w for w, _ in KMP], range(first, first + len(KMP))):
        s.append([f"{who} within interquartile range?", f'=IF(AND(B{r}>=B5,B{r}<=B7),"Within range","Outside range")'])
    for c in ("A3", "B3", f"A{first - 1}", f"B{first - 1}"):
        s[c].font = hd; s[c].fill = fill
    s.column_dimensions["A"].width = 52; s.column_dimensions["B"].width = 18
    for r in s.iter_rows(min_row=4, max_row=s.max_row, min_col=2, max_col=2):
        r[0].number_format = "#,##0"

    f = wb.create_sheet("Final_Set")
    f.append(["Person Name", "Company Name", "Professional Job Functions", "Total Revenue [FY 2024] (AED mm)", "Total Compensation Received/Earned [FY 2024] (AED, Historical rate)"])
    for c in f[1]:
        c.font = hd; c.fill = fill
    rr = list(rows)
    if errors:
        rr[3] = rr[2]                      # duplicate comparable
    for r in rr:
        f.append(list(r))
    for col, w in zip("ABCDE", (18, 36, 30, 22, 30)):
        f.column_dimensions[col].width = w
    for r in f.iter_rows(min_row=2, min_col=5, max_col=5):
        r[0].number_format = "#,##0"

    sd = wb.create_sheet("Annexure 7 Search Results")
    sd.append(["Person Name", "Company Name", "Total Compensation Received/Earned [FY 2024] (AED)"])
    for c in sd[1]:
        c.font = hd; c.fill = fill
    for r in rows:
        sd.append([r[0], r[1], r[4]])
    for j in range(36):
        sd.append([f"Rejected Executive {j + 1:02d}", f"Rejected Candidate {j + 1:02d} Ltd", 90_000 + j * 41_000])
    sd.column_dimensions["B"].width = 36

    a = wb.create_sheet("Annexure 8 Accept-Reject")
    a.append(["#", "Company Name", "Decision", "Reason"])
    for c in a[1]:
        c.font = hd; c.fill = fill
    reasons = ["Role not comparable to a managing director", "Executive is a finance / secretarial officer", "Subsidiary of a listed group - dependent",
               "Insufficient compensation data", "Main activity is manufacturing, not trading", "Not an executive of the company"]
    for i, r in enumerate(rows, 1):
        a.append([i, r[1], "Accepted", "Trading of construction and building materials; MD / CEO level role; independent; data available"])
    for j in range(36):
        a.append([27 + j + 1, f"Rejected Candidate {j + 1:02d} Ltd", "Rejected", reasons[j % len(reasons)]])
    a.column_dimensions["B"].width = 36; a.column_dimensions["D"].width = 80

    q = wb.create_sheet("Search_Strategy")
    for line in ["Database search (synthetic sample - no database extract is included)",
                 "Step 1: SIC Division F Wholesale Trade; Geography UAE; Total compensation FY 2024 > 0; Keyword: building OR construction -> no results",
                 "Step 2: geography extended to the Middle East -> no suitable results",
                 "Step 3: geography extended to Europe OR Middle East",
                 "Search results (search dump): 63 persons. Accepted: 27. Rejected: 36 (reasons in sheet Annexure 8 Accept-Reject)."]:
        q.append([line])
    q.column_dimensions["A"].width = 130
    if errors:
        h = wb.create_sheet("Working_Hidden"); h.append(["scratch", 1]); h.sheet_state = "hidden"
        wb["Summary"]["B5"] = "=#REF!"
    wb.save(path)


def recalc(path):
    """Opens the file in LibreOffice headless and saves it back so the formula results are cached (as they would be after saving in Excel)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("WARNING: LibreOffice not found - formulas have no cached values; open the workbook in Excel and save it once.")
        return
    tmp = Path(tempfile.mkdtemp())
    subprocess.run([soffice, "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(tmp), str(path)], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    shutil.copy(tmp / Path(path).name, path)
    shutil.rmtree(tmp, ignore_errors=True)
    _set_props(path)
    wb = openpyxl.load_workbook(path, read_only=True)
    print("   formulas recalculated; sheets:", [(ws.title, ws.sheet_state) for ws in wb.worksheets])
    wb.close()


def _set_props(path):
    """Sets author / title in docProps/core.xml directly (re-saving with openpyxl would throw away the cached formula results)."""
    import re
    import zipfile
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            b = zi.read(it.filename)
            if it.filename == "docProps/core.xml":
                t = b.decode("utf8")
                t = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator>Capstone Project</dc:creator>", t)
                t = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", "<cp:lastModifiedBy>Capstone Project</cp:lastModifiedBy>", t)
                b = t.encode("utf8")
            zo.writestr(it, b)
    Path(tmp).replace(path)


# ------------------------------------------------------------------ report
def _field(par, instr):
    for t, txt in (("begin", None), (None, instr), ("end", None)):
        r = par.add_run()
        if t:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), t); r._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.text = txt; r._r.append(it)


def make_report(path, errors=False):
    rows = comparables()
    st = quartiles([r[4] for r in rows])
    d = Document()
    d.core_properties.author = "Capstone Project"; d.core_properties.title = "ABC COMPANY synthetic TP report"
    for nm in ("Normal", "Heading 1", "Heading 2"):
        d.styles[nm].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(11)
    d.styles["Heading 1"].font.size = Pt(19.5); d.styles["Heading 2"].font.size = Pt(12)
    for nm in ("Heading 1", "Heading 2"):
        d.styles[nm].font.color.rgb = None

    sec = d.sections[0]
    hp = sec.header.paragraphs[0]; hp.text = "Transfer Pricing Benchmark Analysis\t"; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _field(hp, "PAGE")
    fp = sec.footer.paragraphs[0]; fp.text = NAME; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pb = OxmlElement("w:pgBorders")
    for e in ("top", "left", "bottom", "right"):
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "12"); pb.append(x)
    sec._sectPr.append(pb)

    d.add_paragraph(NAME)
    d.add_paragraph("Transfer Pricing Benchmark Analysis")
    d.add_paragraph(f"For the financial year ended 31 December {FY - 1 if errors else FY}")
    if errors:
        d.add_paragraph("DRAFT")
    d.add_page_break()

    tp_ = d.add_paragraph(style="TOC Heading"); r_ = tp_.add_run("TABLE OF CONTENTS"); r_.font.size = Pt(19.5); r_.font.name = "Times New Roman"
    toc = [("1. INTRODUCTION", 1), ("1.1 ABOUT THE COMPANY", 1), ("1.2 BACKGROUND OF TRANSACTION", 1), ("1.3 AIM OF THE REPORT", 1), ("1.4 EXECUTIVE SUMMARY", 1),
           ("1.5 LIMITATIONS", 1), ("2. ANALYSIS OF THE TRANSACTION", 2), ("2.1 NATURE AND TERMS OF THE IDENTIFIED TRANSACTION", 2), ("2.2 FUNCTIONAL ANALYSIS", 2),
           ("2.3 ASSET ANALYSIS", 2), ("2.4 RISK ANALYSIS", 2), ("2.5 CHARACTERISATION OF THE ENTITIES", 2), ("3. ECONOMIC ANALYSIS", 3),
           ("3.1 SELECTION OF TRANSFER PRICING METHOD", 3), ("3.2 TRANSFER PRICING METHOD APPLIED", 3), ("4. RESULTS AND CONCLUSION", 4), ("5. APPENDIX", 5)]
    toc += [(t, 6 + i) for i, t in enumerate(["Annexure 1: Search strategy screenshot - UAE", "Annexure 2: Search strategy write-up - UAE",
            "Annexure 3: Search strategy screenshot - Middle East", "Annexure 4: Search strategy write-up - Middle East",
            "Annexure 5: Search strategy screenshot - Europe OR Middle East", "Annexure 6: Search strategy write-up - Europe OR Middle East",
            "Annexure 7: Search results", "Annexure 8: Accept-reject matrix"])]
    for t, pg in toc:
        d.add_paragraph(f"{t} ........ {pg}")
    d.add_page_break()

    d.add_heading("1. INTRODUCTION", 1)
    d.add_heading("1.1 ABOUT THE COMPANY", 2)
    d.add_paragraph(f"{NAME}, Dubai Silicon Oasis, Dubai, UAE (\"the company\") is registered as a free zone company with limited liability with the Dubai Silicon Oasis "
                    "Authority (DSOA) and operates under trade licence no. " + ("XXXXXX" if errors else "000000") + ". The principal activity of the company is building and construction material trading. "
                    "The company is managed by Mr. PQR, the manager, and Mr. XYZ, the authorised signatory. TWO LTD is the sole shareholder of the company.")
    d.add_heading("1.2 BACKGROUND OF TRANSACTION", 2)
    kmp_names = " and ".join(("Mr. Mr. " + n[4:]) if errors and i == 1 else n for i, (n, _) in enumerate(KMP))
    d.add_paragraph(f"During the financial year {FY}, {NAME} paid managerial remuneration of AED {(2_750_000 if errors else TOTAL):,} to its key management personnel, "
                    f"{kmp_names} (the \"KMPs\"), who are connected persons under the Corporate Tax Law. The remuneration is the transaction under review.")
    d.add_heading("1.3 AIM OF THE REPORT", 2)
    d.add_paragraph("The aim of this report is to benchmark the managerial remuneration against comparable third-party remuneration and to conclude whether it is consistent with the "
                    "arm's length principle under Federal Decree-Law No. 47 of 2022, Ministerial Decision No. 97 of 2023 and the OECD Transfer Pricing Guidelines.")
    d.add_heading("1.4 EXECUTIVE SUMMARY", 2)
    n_stated = 26 if errors else len(rows)
    d.add_paragraph(f"A benchmarking analysis was performed using {n_stated} comparable third-party executives of independent companies in the construction and building material "
                    "sector. The arm's length range of remuneration is shown below.")
    t = d.add_table(rows=0, cols=2)
    q3 = 2_775_000 if errors else st["q3"]
    for lab, v in (("Lowest value", st["min"]), ("Lower quartile (25%)", st["q1"]), ("Median", st["median"]), ("Upper quartile (75%)", q3), ("Highest value", st["max"])):
        c = t.add_row().cells; c[0].text = lab; c[1].text = f"{v:,}"
    d.add_heading("1.5 LIMITATIONS", 2)
    d.add_paragraph("The analysis relies on information provided by management and on a commercial database; no assurance is given on its accuracy. Conclusions apply to the year reviewed.")

    d.add_heading("2. ANALYSIS OF THE TRANSACTION", 1)
    d.add_heading("2.1 NATURE AND TERMS OF THE IDENTIFIED TRANSACTION", 2)
    d.add_paragraph(f"The KMPs are paid a monthly salary and allowances for managing the company. The managerial remuneration for the year was AED {2_750_000 if errors else TOTAL:,}, "
                    "as disclosed in the related party note to the audited financial statements. The components are:")
    rt = d.add_table(rows=0, cols=3)
    for a_, b_, c_ in (("Component", "Mr. PQR (AED)", "Mr. XYZ (AED)"), ("Salary", "900,000", "1,100,000"), ("Allowances", "300,000", "250,000"),
                       ("End-of-service benefits", "-", "150,000"), ("Total managerial remuneration", "1,200,000", "1,500,000")):
        cc = rt.add_row().cells; cc[0].text, cc[1].text, cc[2].text = a_, b_, c_
    d.add_paragraph(f"Total managerial remuneration: AED {2_750_000 if errors else TOTAL:,} (salary AED 2,000,000, allowances AED 550,000, end-of-service benefits AED 150,000).")
    d.add_heading("2.2 FUNCTIONAL ANALYSIS", 2)
    d.add_paragraph("Functional analysis: the KMPs perform strategic leadership, sales oversight, banking relationships and operational supervision. The assets used are limited to "
                    "office equipment and inventory. The risks borne by the company are market, credit, inventory and compliance risks; the KMPs control these risks.")
    d.add_heading("2.3 ASSET ANALYSIS", 2)
    d.add_paragraph("The company uses working capital and inventory; no intangibles are used in the transaction.")
    d.add_heading("2.4 RISK ANALYSIS", 2)
    d.add_paragraph("Market risk, credit risk and inventory risk are borne by the company. The KMPs do not bear the entrepreneurial risk of the company.")
    d.add_heading("2.5 CHARACTERISATION OF THE ENTITIES", 2)
    d.add_paragraph("Company profile: the company is a trader of building materials. Industry analysis: the trading sector in the UAE is competitive and stable. "
                    "Controlled transaction: payment of managerial remuneration to connected persons. The company is the tested party because it is the least complex party to the transaction and "
                    "reliable data on comparable remuneration is available; the KMPs are the recipients of the payment.")

    d.add_heading("3. ECONOMIC ANALYSIS", 1)
    d.add_heading("3.1 SELECTION OF TRANSFER PRICING METHOD", 2)
    d.add_paragraph("Method selection: the External comparable uncontrolled price (CUP) method is the most appropriate method because reliable third-party remuneration data exist for "
                    "comparable executive roles. The transactional net margin method (TNMM) is not appropriate because the transaction is a payment for services of individuals, not a "
                    "profit margin of the company. The resale price method is not applicable as no goods are resold in the transaction. The cost plus method is not suitable as no "
                    "cost base is marked up. The profit split method is not applicable as the transaction is not highly integrated. "
                    "Profit level indicator: not applicable, as the CUP method compares the remuneration amount directly and no profit-based indicator is used.")
    d.add_heading("3.2 TRANSFER PRICING METHOD APPLIED", 2)
    d.add_paragraph("Data period: single-year data for FY 2024 is used for all comparables. Comparability analysis: the comparables were identified through a database search "
                    "with quantitative screening followed by a qualitative review of every result.")
    d.add_paragraph("Search strategy: first a UAE-only search was run (SIC Division F Wholesale Trade, compensation for FY 2024 greater than zero, keyword building OR construction); it gave "
                    "no results. The search was extended to the Middle East and then to Europe OR Middle East. The searches gave 63 results ('Search dump'). "
                    + ("The search results of 36 comparables are in Annexure 7. " if errors else "")
                    + "After the accept / reject review, the final set has 27 comparable companies and 36 results were rejected; the reasons for each decision are in the accept-reject matrix (Annexure 8).")
    ft = d.add_table(rows=0, cols=2)
    for a_, b_ in (("Total potential comparables", "63"), ("Total manual review rejections", "36"), ("Accepted companies", "27")):
        cc = ft.add_row().cells; cc[0].text, cc[1].text = a_, b_
    d.add_paragraph("Final set of comparables (continued)")
    tb = d.add_table(rows=0, cols=5)
    h = tb.add_row().cells
    for c, x in zip(h, ("Person Name", "Company Name", "Professional Job Functions", "Total Revenue [FY 2024] (AED mm)", "Total Compensation [FY 2024] (AED)")):
        c.text = x
    for r in rows:
        c = tb.add_row().cells
        c[0].text, c[1].text, c[2].text, c[3].text, c[4].text = r[0], r[1], r[2], f"{r[3]:,.1f}", f"{r[4]:,}"
    for lab, v in (("Lowest value", st["min"]), ("Lower quartile (25%)", st["q1"]), ("Median", st["median"]), ("Upper quartile (75%)", q3), ("Highest value", st["max"])):
        c = tb.add_row().cells; c[2].text = lab; c[4].text = f"{v:,}"
    d.add_paragraph(f"The final set of {len(rows)} comparable companies gives an interquartile range of AED {st['q1']:,} (25th percentile) to AED {q3:,} (75th percentile), with a median of "
                    f"AED {st['median']:,}, under the External CUP method using the quartile function QUARTILE.INC.")

    d.add_heading("4. RESULTS AND CONCLUSION", 1)
    d.add_paragraph(f"The managerial remuneration paid to Mr. PQR (AED {KMP[0][1]:,}) and to Mr. XYZ (AED {KMP[1][1]:,}) lies within the interquartile range of AED {st['q1']:,} to AED {q3:,}. "
                    f"Conclusion: the managerial remuneration of AED {(2_750_000 if errors else TOTAL):,} is consistent with the arm's length principle and no transfer pricing adjustment is required.")
    rs = d.add_table(rows=0, cols=5)
    for row_ in (("Recipient", "Lower quartile (AED)", "Median (AED)", "Upper quartile (AED)", "Actual remuneration (AED)"),
                 ("Mr. PQR", f"{st['q1']:,}", f"{st['median']:,}", f"{q3:,}", f"{KMP[0][1]:,}"), ("Mr. XYZ", f"{st['q1']:,}", f"{st['median']:,}", f"{q3:,}", f"{KMP[1][1]:,}")):
        cc = rs.add_row().cells
        for i_, v_ in enumerate(row_):
            cc[i_].text = v_
    d.add_paragraph("The Transfer Pricing Disclosure Form is filed with the Corporate Tax return as related party transactions exceed the AED 40,000,000 threshold.")
    d.add_heading("5. APPENDIX", 1)
    for i, tx in enumerate(["Search strategy screenshot - UAE", "Search strategy write-up - UAE", "Search strategy screenshot - Middle East", "Search strategy write-up - Middle East",
                            "Search strategy screenshot - Europe OR Middle East", "Search strategy write-up - Europe OR Middle East", "Search results", "Accept-reject matrix"], 1):
        d.add_paragraph(f"Annexure {i}: {tx}")
    for i in range(1, 9):
        d.add_page_break()
        d.add_paragraph(f"Annexure {i}")
        d.add_paragraph("Annexure content (synthetic sample).")
    d.save(path)


def make_fs(path, errors=False):
    d = Document()
    d.core_properties.author = "Capstone Project"
    d.add_paragraph(f"{NAME}"); d.add_paragraph(f"Notes to the financial statements for the year ended 31 December {FY}")
    d.add_paragraph("Independent auditor's report to the shareholder of " + NAME)
    d.add_paragraph("Opinion: in our opinion the financial statements present fairly, in all material respects, the financial position of the company as at 31 December "
                    f"{FY} in accordance with IFRS. Basis for opinion: we conducted our audit in accordance with International Standards on Auditing. Signed: Sample Audit Firm, Dubai.")
    d.add_paragraph("Note 20 - Related party transactions")
    d.add_paragraph("Key management personnel remuneration:")
    t = d.add_table(rows=0, cols=3)
    for a, b, c in (("", f"{FY}", f"{FY - 1}"), ("Managerial remuneration - key management personnel", f"{TOTAL:,}", "2,400,000")):
        r = t.add_row().cells; r[0].text, r[1].text, r[2].text = a, b, c
    d.add_paragraph(f"Managerial remuneration paid to key management personnel during the year amounted to AED {TOTAL:,}.")
    d.save(path)


def make_birf(path):
    from auditsignoff.birf_template import build_birf_template
    build_birf_template(path)
    wb = openpyxl.load_workbook(path); ws = wb["BIRF"]
    for r, v in ((3, NAME), (4, f"FY {FY}"), (5, "Full benchmarking"), (6, "Sample / synthetic")):
        ws.cell(r, 3, v)
    for r, row in zip((9, 10), (("Mr. PQR", "Manager", "Director", "100%", 900_000, None, 300_000, None, None, None),
                                ("Mr. XYZ", "Authorised signatory", "Director", "100%", 1_100_000, None, 250_000, 150_000, None, None))):
        for j, v in enumerate(row, 1):
            if v is not None:
                ws.cell(r, j, v)
    wb.properties.creator = "Capstone Project"; wb.save(path)


def engagement(path):
    Path(path).write_text(f"""fy: {FY}
py: {FY - 1}
client:
  name: {NAME}
  legal_status: LLC
  authority: DSOA
  emirate: Dubai
  year_end: 31 December
  shareholders: [TWO LTD]
  signatories_legal: [PQR, XYZ]
""", encoding="utf-8")


def main(out):
    out = Path(out); out.mkdir(parents=True, exist_ok=True)
    engagement(out / "engagement_abc.yaml")
    make_report(out / "ABC_COMPANY_FY2024_TP_Report.docx"); make_report(out / "ABC_COMPANY_FY2024_TP_Report_ERRORS.docx", errors=True)
    make_fs(out / "ABC_FS_Notes.docx"); make_birf(out / "ABC_BIRF.xlsx")
    for nm, e in (("ABC_COMPANY_FY2024_TP_Benchmark_Workbook.xlsx", False), ("ABC_COMPANY_FY2024_TP_Benchmark_Workbook_ERRORS.xlsx", True)):
        make_workbook(out / nm, errors=e); recalc(out / nm)
    print("ABC synthetic TP set written to", out)


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    main(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parent / "abc_sample")
