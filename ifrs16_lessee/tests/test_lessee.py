import datetime as dt
import http.server
import os
import tempfile
import threading
import unittest
from decimal import Decimal as D

from ifrs16lessee.engine import EngineError, build_timeline, generate_payments, present_value, exemption_check
from ifrs16lessee.models import Event, LeaseTerms
from ifrs16lessee.money import RoundingPolicy, add_months, growth_factor
from ifrs16lessee.reports import export_workbook
from ifrs16lessee.sample import load_sample, sample_terms
from ifrs16lessee.store import Store, StoreError
from ifrs16lessee import tally

d = dt.date


def base(**kw):
    t = LeaseTerms(ref='T1', lessor='X', asset='A', commencement=d(2025, 1, 1), end_date=d(2025, 12, 31), payment_amount=D('1000'), discount_rate=D('0.06'))
    for k, v in kw.items():
        setattr(t, k, v)
    return t


class PV(unittest.TestCase):
    def test_single_payment_pv(self):
        t = base(end_date=d(2025, 12, 31), frequency='annual', payment_amount=D('1050'), discount_rate=D('0.05'), timing='arrears')
        tl = build_timeline(t)
        # one payment on 1 Jan 2026, 365 days at 5% annual ACT/365 => 1050/1.05
        self.assertEqual(tl.initial['liability'], D('1000.00'))

    def test_pv_of_annuity_matches_manual(self):
        t = base(frequency='annual', end_date=d(2027, 12, 31), payment_amount=D('1000'), discount_rate=D('0.05'))
        tl = build_timeline(t)
        manual = sum(D(1000) / growth_factor(d(2025, 1, 1), p, D('0.05'), 'ACT/365', 'annual') for p in (d(2026, 1, 1), d(2027, 1, 1), d(2028, 1, 1)))
        self.assertAlmostEqual(float(tl.initial['liability']), float(manual), places=2)

    def test_rou_composition(self):
        t = base(incentives_received=D('500'), initial_direct_costs=D('300'), restoration_obligation=D('200'), prepaid_before_commencement=D('100'))
        tl = build_timeline(t)
        self.assertEqual(tl.initial['rou_cost'], tl.initial['liability'] + D('100') + D('300') + D('200') - D('500'))

    def test_zero_rate_pv_equals_sum(self):
        tl = build_timeline(base(discount_rate=D('0')))
        self.assertEqual(tl.initial['liability'], D('12000.00'))


class AdvanceArrears(unittest.TestCase):
    def test_dates(self):
        a = generate_payments(base(timing='advance'))
        r = generate_payments(base(timing='arrears'))
        self.assertEqual(a[0].date, d(2025, 1, 1))
        self.assertEqual(r[0].date, d(2025, 2, 1))
        self.assertEqual(r[-1].date, d(2026, 1, 1))
        self.assertEqual(len(a), 12)
        self.assertEqual(len(r), 12)

    def test_advance_first_payment_excluded_and_in_rou(self):
        ta, tr = build_timeline(base(timing='advance')), build_timeline(base(timing='arrears'))
        self.assertEqual(ta.initial['payments_at_commencement'], D('1000.00'))
        self.assertGreater(tr.initial['liability'], ta.initial['liability'])
        self.assertEqual(ta.initial['rou_cost'], ta.initial['liability'] + D('1000.00'))

    def test_quarterly_semiannual(self):
        self.assertEqual(len(generate_payments(base(frequency='quarterly'))), 4)
        self.assertEqual(len(generate_payments(base(frequency='semi-annual'))), 2)

    def test_end_of_month_anchor(self):
        t = base(commencement=d(2025, 1, 31), end_date=d(2025, 6, 30), timing='arrears')
        ds = [p.date for p in generate_payments(t)]
        self.assertEqual(ds[0], d(2025, 2, 28))
        self.assertEqual(ds[1], d(2025, 3, 31))


class Schedule(unittest.TestCase):
    def test_runs_to_nil_and_totals(self):
        t = base(end_date=d(2027, 12, 31), useful_life_years=D(0))
        tl = build_timeline(t)
        self.assertEqual(tl.rows[-1].close_liab, D('0.00'))
        self.assertEqual(sum(r.payment for r in tl.rows), D('36000.00'))
        self.assertEqual(sum(r.interest for r in tl.rows), D('36000.00') - tl.initial['liability'])
        self.assertEqual(sum(r.depreciation for r in tl.rows), tl.initial['rou_cost'])
        self.assertEqual(tl.rows[-1].close_rou, D('0.00'))
        for a, b in zip(tl.rows, tl.rows[1:]):
            self.assertEqual(a.close_liab, b.open_liab)
            self.assertEqual(a.close_rou, b.open_rou)
        for r in tl.rows:
            self.assertEqual(r.open_liab + r.interest - r.payment + r.adjustment, r.close_liab)

    def test_rounding_reconciliation_is_visible_and_small(self):
        tl = build_timeline(base(end_date=d(2027, 12, 31), discount_rate=D('0.0731')))
        self.assertLess(abs(tl.rows[-1].rounding_adj), D('0.50'))
        self.assertEqual(tl.rows[-1].close_liab, D('0.00'))

    def test_rounding_policy_zero_decimals(self):
        tl = build_timeline(base(rounding=RoundingPolicy(0, 'ROUND_HALF_UP')))
        self.assertEqual(tl.initial['liability'], tl.initial['liability'].quantize(D('1')))

    def test_irregular_payments_and_non_month_end_dates(self):
        from ifrs16lessee.models import Payment
        t = base(payments=[Payment(d(2025, 3, 17), D('3000')), Payment(d(2025, 8, 9), D('4000')), Payment(d(2026, 1, 1), D('5000'))])
        tl = build_timeline(t)
        self.assertEqual(tl.rows[-1].close_liab, D('0.00'))
        self.assertTrue(any(r.end == d(2025, 3, 17) for r in tl.rows))

    def test_day_count_conventions_differ(self):
        a = build_timeline(base(day_count='ACT/365')).initial['liability']
        b = build_timeline(base(day_count='ACT/360')).initial['liability']
        self.assertNotEqual(a, b)

    def test_declining_balance_depreciates_fully(self):
        tl = build_timeline(base(depreciation_method='declining_balance', end_date=d(2027, 12, 31)))
        self.assertEqual(tl.rows[-1].close_rou, D('0.00'))

    def test_validation_messages(self):
        with self.assertRaises(EngineError):
            build_timeline(base(end_date=d(2024, 1, 1)))
        with self.assertRaises(EngineError):
            build_timeline(base(discount_rate=D('5')))
        with self.assertRaises(EngineError):
            build_timeline(base(purchase_option_certain=True, purchase_option_price=D('100')))

    def test_exemption_path(self):
        t = base(exemption='short_term', exemption_elected=True)
        tl = build_timeline(t)
        self.assertTrue(tl.exempt)
        self.assertEqual(tl.initial['liability'], 0)
        exp = sum((l.debit for j in tl.journals for l in j.lines if l.account == 'LEASE_EXP'), D(0))
        self.assertEqual(exp, D('12000.00'))
        self.assertTrue(exemption_check(base())['short_term_eligible'])
        self.assertFalse(exemption_check(base(end_date=d(2026, 6, 30)))['short_term_eligible'])


class Journals(unittest.TestCase):
    def test_all_balanced_and_refs_unique(self):
        s = Store()
        ref = load_sample(s)
        tl = s.timeline(ref)
        refs = [j.ref for j in tl.journals]
        self.assertEqual(len(refs), len(set(refs)))
        for j in tl.journals:
            self.assertTrue(j.balanced(), j.ref)
            self.assertTrue(j.trace)

    def test_journals_agree_with_schedule(self):
        tl = build_timeline(base(end_date=d(2026, 12, 31)))
        interest = sum((l.debit for j in tl.journals if j.source == 'INTEREST' for l in j.lines if l.account == 'INTEREST_EXP'), D(0))
        dep = sum((l.debit for j in tl.journals if j.source == 'DEPRECIATION' for l in j.lines if l.account == 'DEP_EXP'), D(0))
        self.assertEqual(interest, sum(r.interest for r in tl.rows))
        self.assertEqual(dep, sum(r.depreciation for r in tl.rows))
        bal = {}
        for j in tl.journals:
            for l in j.lines:
                bal[l.account] = bal.get(l.account, D(0)) + l.debit - l.credit
        self.assertEqual(bal['LEASE_LIAB'], D(0))
        self.assertEqual(bal['ROU_COST'] + bal['ROU_ACCDEP'], D(0))


class Events(unittest.TestCase):
    def setUp(self):
        self.t = base(end_date=d(2027, 12, 31))

    def run_ev(self, ev):
        tl = build_timeline(self.t, [ev])
        return tl, tl.events[0]

    def test_index_change_uses_original_rate(self):
        tl, r = self.run_ev(Event('index_change', d(2026, 1, 1), {'new_payment_amount': '1100'}))
        self.assertEqual(r.rate_used, D('0.06'))
        self.assertIn('42', r.rate_reason)
        self.assertGreater(r.after['liability'], r.before['liability'])

    def test_floating_rate_requires_revised_rate(self):
        with self.assertRaises(EngineError):
            self.run_ev(Event('index_change', d(2026, 1, 1), {'new_payment_amount': '1100', 'floating_rate_linked': True}))
        tl, r = self.run_ev(Event('index_change', d(2026, 1, 1), {'new_payment_amount': '1100', 'floating_rate_linked': True, 'revised_rate': '7'}))
        self.assertEqual(r.rate_used, D('0.07'))

    def test_term_reassessment_requires_and_uses_revised_rate(self):
        with self.assertRaises(EngineError):
            self.run_ev(Event('term_reassessment', d(2026, 1, 1), {'new_end_date': '2028-12-31'}))
        tl, r = self.run_ev(Event('term_reassessment', d(2026, 1, 1), {'new_end_date': '2028-12-31', 'revised_rate': '8'}))
        self.assertEqual(r.rate_used, D('0.08'))
        self.assertEqual(tl.rows[-1].end, d(2029, 1, 1))
        self.assertEqual(tl.rows[-1].close_liab, D('0.00'))

    def test_rvg_unchanged_rate(self):
        tl, r = self.run_ev(Event('rvg_change', d(2026, 1, 1), {'new_rvg': '5000'}))
        self.assertEqual(r.rate_used, D('0.06'))

    def test_history_preserved_before_event(self):
        base_tl = build_timeline(self.t)
        tl, r = self.run_ev(Event('index_change', d(2026, 1, 1), {'new_payment_amount': '1100'}))
        old = [(x.start, x.end, x.interest, x.payment, x.depreciation) for x in base_tl.rows if x.end < d(2026, 1, 1)]
        new = [(x.start, x.end, x.interest, x.payment, x.depreciation) for x in tl.rows if x.end < d(2026, 1, 1)]
        self.assertEqual(old, new)
        old_j = {j.ref for j in base_tl.journals if j.date < d(2026, 1, 1)}
        new_j = {j.ref for j in tl.journals if j.date < d(2026, 1, 1)}
        self.assertEqual(old_j, new_j)

    def test_rou_floor_and_profit_or_loss(self):
        # big decrease in payments with a tiny remaining ROU asset: residual goes to P&L
        t = base(end_date=d(2027, 12, 31), useful_life_years=D(0))
        ev = Event('modification', d(2026, 12, 1), {'new_payment_amount': '1', 'revised_rate': '6'})
        t2 = base(end_date=d(2027, 12, 31))
        tl = build_timeline(t2, [Event('modification', d(2027, 11, 1), {'new_payment_amount': '1', 'revised_rate': '6'})])
        r = tl.events[0]
        self.assertGreaterEqual(r.after['rou_carrying'], 0)
        # force a case where reduction exceeds carrying amount: shorten depreciation via short useful life
        t3 = base(end_date=d(2027, 12, 31), useful_life_years=D('1'), ownership_transfers=True)
        tl3 = build_timeline(t3, [Event('modification', d(2026, 6, 1), {'new_payment_amount': '1', 'revised_rate': '6'})])
        r3 = tl3.events[0]
        self.assertEqual(r3.after['rou_carrying'], D('0.00'))
        gl = [l for j in r3.journals for l in j.lines if l.account == 'GAIN_LOSS']
        self.assertTrue(gl and gl[0].credit > 0)
        self.assertTrue(all(j.balanced() for j in r3.journals))

    def test_separate_lease_not_adjusted(self):
        ev = Event('modification', d(2026, 1, 1), {'adds_right_of_use': True, 'added_consideration': '2000', 'standalone_price': '2000'})
        tl, r = self.run_ev(ev)
        self.assertEqual(r.decision, 'separate_lease')
        self.assertEqual(r.journals, [])
        self.assertEqual(r.before['liability'], r.after['liability'])
        self.assertIsNotNone(r.new_lease_template)

    def test_not_separate_if_price_not_commensurate(self):
        ev = Event('modification', d(2026, 1, 1), {'adds_right_of_use': True, 'added_consideration': '500', 'standalone_price': '2000', 'revised_rate': '7', 'new_payment_amount': '1500'})
        tl, r = self.run_ev(ev)
        self.assertEqual(r.decision, 'modification_remeasured')
        self.assertEqual(r.rate_used, D('0.07'))

    def test_modification_requires_revised_rate(self):
        with self.assertRaises(EngineError):
            self.run_ev(Event('modification', d(2026, 1, 1), {'new_payment_amount': '800'}))

    def test_scope_decrease_modification_goes_through_termination_logic(self):
        tl, r = self.run_ev(Event('modification', d(2026, 1, 1), {'scope_decrease_pct': '25'}))
        self.assertEqual(r.decision, 'modification_scope_decrease')

    def test_partial_termination(self):
        tl, r = self.run_ev(Event('termination', d(2026, 1, 1), {'scope_pct': '40', 'termination_payment': '2000'}))
        L = r.before['liability']
        dL = [b for b in r.bridge if b['label'].startswith('Lease liability derecognised')][0]['amount']
        self.assertEqual(dL, (L * D('0.4')).quantize(D('0.01')))
        self.assertAlmostEqual(float(r.after['liability']), float(L * D('0.6')), places=2)
        self.assertAlmostEqual(float(r.after['rou_carrying']), float(r.before['rou_carrying'] * D('0.6')), places=2)
        gain = [b for b in r.bridge if b['label'].startswith('Gain')][0]['amount']
        self.assertEqual(gain, dL - [b for b in r.bridge if b['label'].startswith('ROU asset derecognised')][0]['amount'] - D('2000'))
        self.assertEqual(r.decision, 'partial_termination')
        self.assertEqual(tl.rows[-1].close_liab, D('0.00'))
        self.assertEqual(sum(x.payment for x in tl.rows if x.start >= d(2026, 1, 1) or x.kind == 'day0') + 0, sum(x.payment for x in tl.rows if x.start >= d(2026, 1, 1)))
        self.assertTrue(all(j.balanced() for j in tl.journals))

    def test_partial_termination_with_remeasurement(self):
        with self.assertRaises(EngineError):
            self.run_ev(Event('termination', d(2026, 1, 1), {'scope_pct': '40', 'new_payment_amount': '500'}))
        tl, r = self.run_ev(Event('termination', d(2026, 1, 1), {'scope_pct': '40', 'new_payment_amount': '500', 'revised_rate': '7'}))
        self.assertEqual(r.rate_used, D('0.07'))
        self.assertTrue(any('Remeasurement' in b['label'] for b in r.bridge))

    def test_full_termination_stops_entries(self):
        tl, r = self.run_ev(Event('termination', d(2026, 3, 15), {'scope_pct': '100', 'termination_payment': '500'}))
        self.assertEqual(r.after['liability'], D('0.00'))
        self.assertEqual(r.after['rou_carrying'], D('0.00'))
        self.assertTrue(tl.final['terminated'])
        self.assertTrue(all(j.date <= d(2026, 3, 15) for j in tl.journals))
        self.assertEqual(tl.rows[-1].close_liab, D('0.00'))
        self.assertTrue(all(j.balanced() for j in tl.journals))
        with self.assertRaises(EngineError):
            build_timeline(self.t, [Event('termination', d(2026, 3, 15), {'scope_pct': '100'}), Event('index_change', d(2026, 5, 1), {'new_payment_amount': '9'})])

    def test_invalid_scope(self):
        with self.assertRaises(EngineError):
            self.run_ev(Event('termination', d(2026, 1, 1), {'scope_pct': '0'}))
        with self.assertRaises(EngineError):
            self.run_ev(Event('termination', d(2025, 1, 1), {'scope_pct': '10'}))

    def test_rows_chain_through_events(self):
        s = Store(); load_sample(s)
        rows = s.timeline('SAMPLE-001').rows
        for a, b in zip(rows, rows[1:]):
            self.assertEqual(a.close_liab, b.open_liab)
            self.assertEqual(a.close_rou, b.open_rou)
        for r in rows:
            self.assertEqual(r.open_liab + r.interest - r.payment + r.adjustment, r.close_liab)
            self.assertEqual(r.open_rou - r.depreciation + r.rou_adjustment, r.close_rou)

    def test_mid_month_event_date(self):
        tl, r = self.run_ev(Event('index_change', d(2026, 1, 20), {'new_payment_amount': '1100'}))
        self.assertTrue(any(x.end == d(2026, 1, 20) for x in tl.rows))
        self.assertEqual(tl.rows[-1].close_liab, D('0.00'))


class Persistence(unittest.TestCase):
    def test_survives_restart_and_events_lock(self):
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, 'x.db')
            s = Store(p)
            ref = load_sample(s)
            n_j = len(s.journals(ref))
            fin = s.timeline(ref).final
            s.set_setting('functional_currency', 'USD')
            s.close()
            s2 = Store(p)
            self.assertEqual([l['ref'] for l in s2.list_leases()], [ref])
            self.assertEqual(len(s2.journals(ref)), n_j)
            self.assertEqual(s2.timeline(ref).final, fin)
            self.assertEqual(len(s2.get_events(ref)), 2)
            self.assertEqual(s2.get_setting('functional_currency'), 'USD')
            self.assertEqual(len(s2.versions(ref)), 3)
            self.assertTrue(len(s2.audit_rows(ref)) >= 3)
            with self.assertRaises(StoreError):
                s2.save_lease(sample_terms(ref))

    def test_posted_history_blocks_earlier_events(self):
        s = Store()
        s.save_lease(base(end_date=d(2026, 12, 31)))
        refs = [j['ref'] for j in s.journals('T1') if j['jdate'] <= '2025-06-01']
        s.approve_journals(refs)
        for r in refs:
            s.log_posting(r, 'post', '<x/>', '<RESPONSE><CREATED>1</CREATED></RESPONSE>', 200, 'success', 'v1')
        with self.assertRaises(StoreError):
            s.preview('T1', Event('index_change', d(2025, 3, 1), {'new_payment_amount': '1100'}))
        before = {j['ref']: j['content_hash'] for j in s.journals('T1') if j['posting_state'] == 'posted'}
        s.commit_event('T1', Event('index_change', d(2025, 9, 1), {'new_payment_amount': '1100'}))
        after = {j['ref']: j['content_hash'] for j in s.journals('T1') if j['posting_state'] == 'posted'}
        self.assertEqual(before, after)

    def test_unapproved_future_journals_superseded(self):
        s = Store()
        s.save_lease(base(end_date=d(2026, 12, 31)))
        s.commit_event('T1', Event('termination', d(2025, 7, 1), {'scope_pct': '100'}))
        self.assertTrue(all(j['jdate'] <= '2025-07-01' for j in s.journals('T1')))
        self.assertTrue(any(j['status'] == 'superseded' for j in s.journals('T1', include_superseded=True)))


class MockTally(http.server.BaseHTTPRequestHandler):
    mode = 'ok'
    received = []

    def log_message(self, *a):
        pass

    def do_GET(self):
        self.send_response(200); self.end_headers(); self.wfile.write(b'<RESPONSE>TallyPrime Server is Running</RESPONSE>')

    def do_POST(self):
        body = self.rfile.read(int(self.headers['Content-Length'])).decode()
        MockTally.received.append(body)
        m = MockTally.mode
        if m == 'ok':
            out, code = '<RESPONSE><CREATED>1</CREATED><ALTERED>0</ALTERED><ERRORS>0</ERRORS><LASTVCHID>42</LASTVCHID></RESPONSE>', 200
        elif m == 'lineerror':
            out, code = '<RESPONSE><LINEERROR>Ledger "Bank Account" does not exist!</LINEERROR><CREATED>0</CREATED><ERRORS>1</ERRORS></RESPONSE>', 200
        elif m == 'ambiguous':
            out, code = '<RESPONSE></RESPONSE>', 200
        elif m == 'empty':
            out, code = '', 200
        elif m == 'ledgers':
            out, code = '<ENVELOPE><LEDGER NAME="Right-of-Use Assets"/><LEDGER NAME="Lease Liability"/></ENVELOPE>', 200
        else:
            out, code = 'boom', 500
        self.send_response(code); self.end_headers(); self.wfile.write(out.encode())


class TallyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.srv = http.server.HTTPServer(('127.0.0.1', 0), MockTally)
        cls.port = cls.srv.server_address[1]
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()

    def setUp(self):
        MockTally.received.clear()
        self.s = Store()
        self.s.save_lease(base())
        self.ref = self.s.journals('T1')[0]['ref']
        self.s.approve_journals([self.ref])
        self.cfg = tally.TallyConfig(host='127.0.0.1', port=self.port, company='Demo Co', timeout=3)

    def snapshot(self):
        tl = self.s.timeline('T1')
        return [r.to_dict() for r in tl.rows], [(j['ref'], j['content_hash']) for j in self.s.journals('T1')]

    def test_xml_structure(self):
        x = tally.preview(self.s, self.ref, self.cfg)
        for k in ('<TALLYREQUEST>Import</TALLYREQUEST>', '<ID>Vouchers</ID>', 'SVCURRENTCOMPANY>Demo Co<', 'VCHTYPE="Journal"', 'ACTION="Create"', '<DATE>20250101</DATE>', '<ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>'):
            self.assertIn(k, x)
        self.assertEqual(MockTally.received, [])        # preview sends nothing

    def test_post_requires_confirmation_and_approval(self):
        with self.assertRaises(tally.TallyConfigError):
            tally.post_journal(self.s, self.ref, self.cfg)
        other = [j['ref'] for j in self.s.journals('T1') if j['status'] == 'draft'][0]
        with self.assertRaises(tally.TallyConfigError):
            tally.post_journal(self.s, other, self.cfg, True)
        self.assertEqual(MockTally.received, [])

    def test_success_and_idempotency(self):
        MockTally.mode = 'ok'
        res = tally.post_journal(self.s, self.ref, self.cfg, True)
        self.assertEqual(res['outcome'], 'success')
        j = self.s.journal(self.ref)
        self.assertEqual((j['posting_state'], j['voucher_id']), ('posted', '42'))
        with self.assertRaises(tally.TallyConfigError):
            tally.post_journal(self.s, self.ref, self.cfg, True)
        self.assertEqual(len(MockTally.received), 1)
        row = self.s.posting_rows(self.ref)[0]
        self.assertTrue(row['request_xml'] and row['response_text'] and row['attempted_at'])

    def test_line_error_is_failure_and_calculations_unchanged(self):
        MockTally.mode = 'lineerror'
        a = self.snapshot()
        res = tally.post_journal(self.s, self.ref, self.cfg, True)
        self.assertEqual(res['outcome'], 'failed')
        self.assertIn('does not exist', res['reason'])
        self.assertEqual(self.s.journal(self.ref)['posting_state'], 'failed')
        self.assertEqual(a, self.snapshot())

    def test_http_200_alone_is_not_success(self):
        for mode in ('ambiguous', 'empty'):
            MockTally.mode = mode
            s = Store(); s.save_lease(base()); r = s.journals('T1')[0]['ref']; s.approve_journals([r])
            res = tally.post_journal(s, r, self.cfg, True)
            self.assertEqual(res['outcome'], 'unconfirmed', mode)
            self.assertEqual(s.journal(r)['posting_state'], 'unconfirmed')
            with self.assertRaises(tally.TallyConfigError):      # duplicate guard until resolved
                tally.post_journal(s, r, self.cfg, True)
            s.resolve_unconfirmed(r, False)
            self.assertEqual(s.journal(r)['posting_state'], 'pending')

    def test_http_500(self):
        MockTally.mode = 'err'
        self.assertEqual(tally.post_journal(self.s, self.ref, self.cfg, True)['outcome'], 'failed')

    def test_tally_down_keeps_pending_and_export_fallback(self):
        a = self.snapshot()
        down = tally.TallyConfig(host='127.0.0.1', port=1, company='Demo Co', timeout=1)
        res = tally.post_journal(self.s, self.ref, down, True)
        self.assertEqual(res['outcome'], 'unavailable')
        j = self.s.journal(self.ref)
        self.assertEqual((j['posting_state'], j['status']), ('pending', 'approved'))
        self.assertEqual(a, self.snapshot())
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, 'v.xml')
            n = tally.export_xml_file([j], down, self.s.get_setting('account_map'), p)
            self.assertEqual(n, 1)
            import xml.etree.ElementTree as ET
            ET.parse(p)

    def test_connection_and_ledger_validation(self):
        self.assertTrue(tally.test_connection(self.cfg)['ok'])
        self.assertFalse(tally.test_connection(tally.TallyConfig(port=1, timeout=1))['ok'])
        MockTally.mode = 'ledgers'
        v = tally.validate_ledgers(self.cfg, self.s.get_setting('account_map'))
        self.assertFalse(v['ok'])
        self.assertNotIn('Lease Liability', v['missing'])
        self.assertIn('Bank Account', v['missing'])

    def test_missing_mapping_and_company(self):
        m = self.s.get_setting('account_map'); m['CASH']['ledger'] = ''
        cashj = [j for j in self.s.journals('T1') if any(l['account'] == 'CASH' for l in j['lines'])][0]
        with self.assertRaises(tally.TallyConfigError):
            tally.build_voucher_xml(cashj, self.cfg, m)
        with self.assertRaises(tally.TallyConfigError):
            tally.build_voucher_xml(cashj, tally.TallyConfig(company=''), self.s.get_setting('account_map'))


class Reports(unittest.TestCase):
    def test_workbook(self):
        from openpyxl import load_workbook
        s = Store(); ref = load_sample(s)
        with tempfile.TemporaryDirectory() as td:
            p = export_workbook(s, [ref], os.path.join(td, 'r.xlsx'), 'Demo Co')
            wb = load_workbook(p)
            for n in ('Summary', 'Schedules', 'Liability roll-forward', 'ROU roll-forward', 'Reassessment bridges', 'Termination calcs', 'Journals', 'Posting log', 'Audit history'):
                self.assertIn(n, wb.sheetnames)
            self.assertGreater(wb['Journals'].max_row, 50)

    def test_rollforward_reconciles(self):
        from ifrs16lessee.reports import rollforward
        s = Store(); ref = load_sample(s)
        tl = s.timeline(ref)
        L, R = rollforward(tl)
        self.assertEqual(L[-1][-1], D('0.00'))
        self.assertEqual(R[-1][-1], D('0.00'))
        for a, b in zip(L, L[1:]):
            self.assertEqual(a[-1], b[1])


if __name__ == '__main__':
    unittest.main()
