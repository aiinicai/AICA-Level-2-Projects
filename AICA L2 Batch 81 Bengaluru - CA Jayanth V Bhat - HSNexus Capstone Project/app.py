import json
import os
import re
import threading
import uuid
from datetime import datetime

from flask import Flask, render_template, request, jsonify, send_from_directory, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from extractor import process_pdf, compute_confidence

APP_NAME = "HSNexus"
APP_TAGLINE = "GST Invoice Extractor"
APP_AUTHOR = "Jayanth V Bhat"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)
CACHE_PATH = os.path.join(OUTPUT_DIR, ".process_cache.json")

app = Flask(__name__)

COLUMNS = [
    "Invoice Number",
    "PDF File Name",
    "HSN/SAC Code",
    "Vendor Name",
    "GSTIN of Vendor",
    "Taxable Value",
    "Tax",
    "Confidence",
    "Folder Path",
    "Extraction Notes",
]

# In-memory scan jobs, so the UI can poll for progress on long folder scans
# instead of the browser just hanging on one big request.
SCAN_JOBS = {}
SCAN_JOBS_LOCK = threading.Lock()

# Absolute paths of PDFs from the most recently completed scan - the
# "View PDF" link is only allowed to serve files from this set, so the
# open-pdf endpoint can't be used to read arbitrary files off the disk.
LAST_SCAN_PATHS = set()


# ---------------------------------------------------------------------
# Result cache - skip re-processing (in particular re-OCR'ing) a PDF that
# hasn't changed since the last scan.
# ---------------------------------------------------------------------

def load_cache():
    if os.path.isfile(CACHE_PATH):
        try:
            with open(CACHE_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}


def save_cache(cache):
    try:
        with open(CACHE_PATH, "w", encoding="utf-8") as f:
            json.dump(cache, f)
    except Exception:
        pass


def file_signature(path):
    st = os.stat(path)
    return st.st_mtime, st.st_size


# ---------------------------------------------------------------------
# Folder scan -> flat rows (one per HSN/SAC line item)
# ---------------------------------------------------------------------

def flag_duplicate_invoices(rows):
    """Mark rows whose (Invoice Number, Vendor Name) shows up in more than
    one PDF - usually means the same invoice got saved/scanned twice."""
    files_by_key = {}
    for row in rows:
        inv = (row.get("Invoice Number") or "").strip()
        if not inv:
            continue
        vendor = (row.get("Vendor Name") or "").strip()
        files_by_key.setdefault((inv, vendor), set()).add(row["PDF File Name"])

    for row in rows:
        inv = (row.get("Invoice Number") or "").strip()
        if not inv:
            continue
        vendor = (row.get("Vendor Name") or "").strip()
        other_files = files_by_key.get((inv, vendor), set()) - {row["PDF File Name"]}
        if other_files:
            note = f"Possible duplicate invoice - also found in: {', '.join(sorted(other_files))}"
            row["Extraction Notes"] = f"{row['Extraction Notes']} {note}".strip()
            row["Confidence"] = "Low"


def build_rows(folder_path, progress_cb=None):
    """Walk folder_path (including sub-folders) and extract a row per HSN/SAC."""
    pdf_paths = []
    for root, _dirs, files in os.walk(folder_path):
        for fname in sorted(files):
            if fname.lower().endswith(".pdf"):
                pdf_paths.append((root, fname))

    total = len(pdf_paths)
    cache = load_cache()
    cache_dirty = False

    rows = []
    pdf_count = 0
    error_count = 0
    scanned_paths = []

    for idx, (root, fname) in enumerate(pdf_paths, start=1):
        full_path = os.path.join(root, fname)
        abs_path = os.path.abspath(full_path)
        pdf_count += 1
        scanned_paths.append(abs_path)

        if progress_cb:
            progress_cb(idx, total, fname)

        try:
            mtime, size = file_signature(full_path)
        except OSError:
            mtime, size = None, None

        cached = cache.get(abs_path)
        if cached and cached.get("mtime") == mtime and cached.get("size") == size:
            result = cached["result"]
        else:
            result = process_pdf(full_path)
            cache[abs_path] = {"mtime": mtime, "size": size, "result": result}
            cache_dirty = True

        if result["note"].startswith("Error processing PDF"):
            error_count += 1

        items = result["items"]
        if not items:
            rows.append({
                "Invoice Number": result["invoice_number"] or "",
                "PDF File Name": fname,
                "HSN/SAC Code": "",
                "Vendor Name": result["vendor_name"] or "",
                "GSTIN of Vendor": result["vendor_gstin"] or "",
                "Taxable Value": "",
                "Tax": "",
                "Confidence": "Low",
                "Folder Path": root,
                "Extraction Notes": result["note"],
                "AbsPath": abs_path,
            })
        else:
            for item in items:
                notes = result["note"]
                if item.get("warning"):
                    notes = f"{notes} {item['warning']}".strip() if notes else item["warning"]
                rows.append({
                    "Invoice Number": result["invoice_number"] or "",
                    "PDF File Name": fname,
                    "HSN/SAC Code": item["hsn"],
                    "Vendor Name": result["vendor_name"] or "",
                    "GSTIN of Vendor": result["vendor_gstin"] or "",
                    "Taxable Value": item["taxable"],
                    "Tax": item["tax"],
                    "Confidence": compute_confidence(result, item),
                    "Folder Path": root,
                    "Extraction Notes": notes,
                    "AbsPath": abs_path,
                })

    if cache_dirty:
        save_cache(cache)

    flag_duplicate_invoices(rows)

    return rows, pdf_count, error_count, scanned_paths


# ---------------------------------------------------------------------
# Excel export - branded, colored, one sheet of line items + one summary
# ---------------------------------------------------------------------

BANNER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FILL = PatternFill(start_color="2E5395", end_color="2E5395", fill_type="solid")
BAND_FILL_A = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
BAND_FILL_B = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
FLAG_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
ERROR_FILL = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")

CONFIDENCE_FILLS = {
    "High": PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid"),
    "Medium": PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),
    "Low": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
}

THIN_BORDER = Border(*(Side(style="thin", color="B7C6E0") for _ in range(4)))

NUMERIC_COLUMNS = {"Taxable Value", "Tax"}


def write_banner(ws, n_cols, title, subtitle):
    """No merged cells - "center across selection" alignment makes the
    title visually span the columns while every cell stays independent."""
    ws.append([title] + [""] * (n_cols - 1))
    ws.append([subtitle] + [""] * (n_cols - 1))
    center_continuous = Alignment(horizontal="centerContinuous", vertical="center")
    for i in range(1, n_cols + 1):
        c1 = ws.cell(row=1, column=i)
        c1.fill = BANNER_FILL
        c1.alignment = center_continuous
        c2 = ws.cell(row=2, column=i)
        c2.fill = BANNER_FILL
        c2.alignment = center_continuous
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws.row_dimensions[1].height = 28
    ws["A2"].font = Font(italic=True, size=10, color="D9E2F3")


def autosize_columns(ws, columns, rows):
    for i, col in enumerate(columns, start=1):
        max_len = len(col)
        for row in rows:
            val = row.get(col, "")
            max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 60)


def write_line_items_sheet(wb, rows):
    ws = wb.active
    ws.title = "HSN-SAC Extract"

    n_cols = len(COLUMNS)
    last_col_letter = get_column_letter(n_cols)
    write_banner(ws, n_cols, f"{APP_NAME} – {APP_TAGLINE}", f"Developed by {APP_AUTHOR}")

    header_row_idx = 3
    ws.append([])
    for i, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row_idx, column=i, value=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws.freeze_panes = f"A{header_row_idx + 1}"

    # Data rows, banded by invoice (falls back to PDF file name/folder when
    # the invoice number couldn't be read, so every invoice's HSN/SAC lines
    # still get grouped together visually).
    band_key = None
    band_fill = BAND_FILL_A
    for row in rows:
        key = (row.get("Invoice Number") or row.get("PDF File Name"), row.get("Folder Path"))
        if key != band_key:
            band_key = key
            band_fill = BAND_FILL_B if band_fill is BAND_FILL_A else BAND_FILL_A

        row_idx = ws.max_row + 1
        is_flagged = bool(row.get("Extraction Notes"))
        for i, col in enumerate(COLUMNS, start=1):
            value = row.get(col, "")
            cell = ws.cell(row=row_idx, column=i, value=value)
            cell.border = THIN_BORDER
            if col in NUMERIC_COLUMNS and isinstance(value, (int, float)):
                cell.number_format = "#,##0.00"
            if col == "Confidence" and value in CONFIDENCE_FILLS:
                cell.fill = CONFIDENCE_FILLS[value]
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
            elif col == "Extraction Notes" and is_flagged:
                cell.fill = ERROR_FILL if "Error processing PDF" in value else FLAG_FILL
            else:
                cell.fill = band_fill

    last_row = ws.max_row
    ws.auto_filter.ref = f"A{header_row_idx}:{last_col_letter}{last_row}"
    autosize_columns(ws, COLUMNS, rows)


SUMMARY_COLUMNS = ["Name", "Line Items", "Total Taxable Value", "Total Tax"]


def write_summary_table(ws, start_row, heading, totals):
    n_cols = len(SUMMARY_COLUMNS)
    ws.cell(row=start_row, column=1, value=heading).font = Font(bold=True, size=12, color="1F3864")

    header_row = start_row + 1
    for i, col in enumerate(SUMMARY_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=i, value=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    row_idx = header_row
    band_fill = BAND_FILL_B
    for name, data in sorted(totals.items(), key=lambda kv: -kv[1]["taxable"]):
        row_idx += 1
        band_fill = BAND_FILL_A if band_fill is BAND_FILL_B else BAND_FILL_B
        values = [name, data["rows"], round(data["taxable"], 2), round(data["tax"], 2)]
        for i, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=i, value=val)
            cell.border = THIN_BORDER
            cell.fill = band_fill
            if i >= 3:
                cell.number_format = "#,##0.00"

    for i, col in enumerate(SUMMARY_COLUMNS, start=1):
        width = max(len(col), 20) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 50)

    return row_idx + 2  # next free row, with one blank row of padding


def write_summary_sheet(wb, rows):
    ws = wb.create_sheet("Summary")
    n_cols = len(SUMMARY_COLUMNS)
    write_banner(ws, n_cols, f"{APP_NAME} – Summary", f"Developed by {APP_AUTHOR}")

    vendor_totals = {}
    hsn_totals = {}
    for row in rows:
        taxable = row.get("Taxable Value")
        if not isinstance(taxable, (int, float)):
            continue
        tax = row.get("Tax") if isinstance(row.get("Tax"), (int, float)) else 0.0

        vendor = row.get("Vendor Name") or "(Unknown vendor)"
        vt = vendor_totals.setdefault(vendor, {"taxable": 0.0, "tax": 0.0, "rows": 0})
        vt["taxable"] += taxable
        vt["tax"] += tax
        vt["rows"] += 1

        hsn = row.get("HSN/SAC Code") or "(Unknown)"
        ht = hsn_totals.setdefault(hsn, {"taxable": 0.0, "tax": 0.0, "rows": 0})
        ht["taxable"] += taxable
        ht["tax"] += tax
        ht["rows"] += 1

    next_row = write_summary_table(ws, 4, "Vendor-wise Summary", vendor_totals)
    write_summary_table(ws, next_row, "HSN/SAC-wise Summary", hsn_totals)


def write_excel(rows, filename):
    wb = Workbook()
    write_line_items_sheet(wb, rows)
    write_summary_sheet(wb, rows)

    out_path = os.path.join(OUTPUT_DIR, filename)
    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------------
# Background scan jobs, so the UI can show live progress
# ---------------------------------------------------------------------

def run_scan_job(job_id, folder_path):
    def progress_cb(done, total, current_file):
        with SCAN_JOBS_LOCK:
            job = SCAN_JOBS.get(job_id)
            if job is not None:
                job.update(done=done, total=total, current_file=current_file)

    try:
        rows, pdf_count, error_count, scanned_paths = build_rows(folder_path, progress_cb=progress_cb)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"HSN_SAC_Extract_{timestamp}_{uuid.uuid4().hex[:6]}.xlsx"
        write_excel(rows, filename)

        global LAST_SCAN_PATHS
        LAST_SCAN_PATHS = set(scanned_paths)

        with SCAN_JOBS_LOCK:
            SCAN_JOBS[job_id].update(
                status="done",
                rows=rows,
                columns=COLUMNS,
                excel_file=filename,
                summary={
                    "pdf_count": pdf_count,
                    "row_count": len(rows),
                    "error_count": error_count,
                },
            )
    except Exception as e:
        with SCAN_JOBS_LOCK:
            job = SCAN_JOBS.get(job_id)
            if job is not None:
                job.update(status="error", error=str(e))


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/scan/start", methods=["POST"])
def scan_start():
    data = request.get_json(silent=True) or {}
    folder_path = (data.get("folder_path") or "").strip().strip('"')

    if not folder_path:
        return jsonify({"error": "Please provide a folder path."}), 400
    if not os.path.isdir(folder_path):
        return jsonify({"error": f"Folder not found: {folder_path}"}), 400

    job_id = uuid.uuid4().hex
    with SCAN_JOBS_LOCK:
        SCAN_JOBS[job_id] = {"status": "running", "done": 0, "total": 0, "current_file": ""}

    thread = threading.Thread(target=run_scan_job, args=(job_id, folder_path), daemon=True)
    thread.start()

    return jsonify({"job_id": job_id})


@app.route("/api/scan/status/<job_id>")
def scan_status(job_id):
    with SCAN_JOBS_LOCK:
        job = SCAN_JOBS.get(job_id)
        if not job:
            return jsonify({"error": "Unknown job."}), 404
        return jsonify(dict(job))


@app.route("/api/download/<path:filename>")
def download(filename):
    safe_name = os.path.basename(filename)
    if not re.fullmatch(r"HSN_SAC_Extract_[\d_]+_[0-9a-f]+\.xlsx", safe_name):
        return jsonify({"error": "Invalid file name."}), 400
    if not os.path.isfile(os.path.join(OUTPUT_DIR, safe_name)):
        return jsonify({"error": "File not found."}), 404
    return send_from_directory(OUTPUT_DIR, safe_name, as_attachment=True)


@app.route("/api/open-pdf")
def open_pdf():
    path = request.args.get("path", "")
    abs_path = os.path.abspath(path)
    # Only serve a file that was part of the most recently completed scan -
    # this endpoint must not become a general "read any file" hole.
    if abs_path not in LAST_SCAN_PATHS:
        return jsonify({"error": "File not recognized from the last scan."}), 403
    if not abs_path.lower().endswith(".pdf") or not os.path.isfile(abs_path):
        return jsonify({"error": "File not found."}), 404
    return send_file(abs_path, mimetype="application/pdf")


if __name__ == "__main__":
    app.run(debug=True, port=5000, threaded=True)
