"""Integration tests through the REST API (TestClient): workflow, controls, events, imports, exports, AI reader."""
import io
import os
import shutil
import tempfile
import time
from pathlib import Path

import pytest

TMP = tempfile.mkdtemp(prefix="lease116test_")
os.environ["LEASE116_DATA_DIR"] = TMP
os.environ["LEASE116_DB_URL"] = f"sqlite:///{Path(TMP, 'test.db').as_posix()}"

from fastapi.testclient import TestClient  # noqa: E402

from app.db.base import reset_engine  # noqa: E402

reset_engine(os.environ["LEASE116_DB_URL"])
from app.main import app  # noqa: E402

SAMPLES = Path(__file__).resolve().parents[1] / "samples" / "agreements"


def login(c, user, pwd):
    r = c.post("/api/auth/login", json={"username": user, "password": pwd})
    assert r.status_code == 200, r.text
    return r


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        login(c, "admin", "admin116")
        r = c.post("/api/system/demo-data")
        assert r.status_code == 200, r.text
        yield c
    shutil.rmtree(TMP, ignore_errors=True)


def test_auth_required(client):
    with TestClient(app) as anon:
        assert anon.get("/api/leases").status_code == 401
        assert anon.post("/api/auth/login", json={"username": "admin", "password": "wrong"}).status_code == 401


def test_register_and_dashboard(client):
    login(client, "admin", "admin116")
    r = client.get("/api/leases?page_size=100")
    assert r.status_code == 200 and r.json()["total"] == 12
    d = client.get("/api/dashboard?as_of=2026-03-31").json()
    k = d["kpis"]
    assert float(k["total_liability"]) > 0
    assert abs(float(k["current_liability"]) + float(k["noncurrent_liability"]) - float(k["total_liability"])) < 0.02


def _create_simple_lease(c):
    r = c.post("/api/leases", json={"description": "TEST — simple office", "commencement_date": "2025-04-01",
                                    "contract_end": "2030-03-31", "discount_rate_pct": "9.5", "asset_class_id": 1})
    assert r.status_code == 200, r.text
    lid = r.json()["id"]
    r = c.put(f"/api/leases/{lid}/payment-terms", json={"amount": "100000", "start_date": "2025-04-01", "end_date": "2030-03-31",
                                                         "frequency_months": 1, "timing": "ADVANCE", "alignment": "ANNIVERSARY",
                                                         "escalations": [{"value": "5", "every_months": 12, "compounding": True}]})
    assert r.status_code == 200 and r.json()["generated"] == 60
    return lid


def test_maker_checker_and_calculation(client):
    login(client, "preparer", "Lease@116")
    lid = _create_simple_lease(client)
    r = client.post(f"/api/leases/{lid}/calculate")
    assert r.status_code == 200, r.text
    s = r.json()["summary"]
    assert s["periods"][-1]["liab_close"] == "0.00"
    assert client.post(f"/api/leases/{lid}/workflow/submit").status_code == 200
    # preparer cannot approve
    r = client.post(f"/api/leases/{lid}/workflow/approve")
    assert r.status_code == 403
    login(client, "approver", "Lease@116")
    r = client.post(f"/api/leases/{lid}/workflow/approve", json={"comment": "ok"})
    assert r.status_code == 200 and r.json()["status"] == "Approved"
    # approved lease cannot be edited directly
    r = client.patch(f"/api/leases/{lid}", json={"description": "changed"})
    assert r.status_code in (403, 409)


def test_period_lock_blocks_events(client):
    login(client, "admin", "admin116")
    leases = client.get("/api/leases?search=TEST").json()["rows"]
    lid = leases[0]["id"]
    login(client, "approver", "Lease@116")
    assert client.post("/api/periods/2025-06-30/lock", json={"reason": "Q1 close"}).status_code == 200
    login(client, "preparer", "Lease@116")
    r = client.post(f"/api/leases/{lid}/events", json={"event_type": "IMPAIRMENT", "effective_date": "2025-06-15",
                                                      "details": {"impairment_amount": "100000"}})
    assert r.status_code == 409
    login(client, "approver", "Lease@116")
    assert client.post("/api/periods/2025-06-30/reopen", json={"reason": "x"}).status_code == 403   # only admin
    login(client, "admin", "admin116")
    assert client.post("/api/periods/2025-06-30/reopen", json={"reason": "Correction of IBR"}).status_code == 200


def test_event_preview_and_modification_flow(client):
    login(client, "preparer", "Lease@116")
    lid = client.get("/api/leases?search=TEST").json()["rows"][0]["id"]
    ev = {"event_type": "MODIFICATION", "effective_date": "2026-04-01", "description": "Rent reduced",
          "details": {"revised_rate_pct": "10", "payments_mode": "GENERATE",
                      "new_payment_terms": {"amount": "90000", "end_date": "2030-03-31", "frequency_months": 1, "timing": "ADVANCE"}}}
    r = client.post(f"/api/leases/{lid}/events/preview", json=ev)
    assert r.status_code == 200, r.text
    assert r.json()["event"]["liability_after"] < r.json()["event"]["liability_before"]
    r = client.post(f"/api/leases/{lid}/events", json=ev)
    assert r.status_code == 200, r.text
    assert client.post(f"/api/leases/{lid}/workflow/submit").status_code == 200
    login(client, "reviewer", "Lease@116")
    assert client.post(f"/api/leases/{lid}/workflow/start_review").status_code == 200
    login(client, "approver", "Lease@116")
    r = client.post(f"/api/leases/{lid}/workflow/approve", json={"comment": "Addendum verified"})
    assert r.status_code == 200 and r.json()["status"] == "Modified"
    lease = client.get(f"/api/leases/{lid}").json()
    assert lease["events"][0]["status"] == "Approved"
    assert len([x for x in lease["runs"] if x["status"] in ("Superseded", "Approved")]) >= 2   # history preserved


def test_journals_balance_and_exports(client):
    login(client, "admin", "admin116")
    js = client.get("/api/journals?start=2025-04-01&end=2026-03-31").json()
    assert js and all(j["balanced"] for j in js)
    for fmt in ("csv", "xlsx", "tally", "sap"):
        r = client.get(f"/api/journals/export?format={fmt}&start=2026-03-01&end=2026-03-31")
        assert r.status_code == 200 and len(r.content) > 100
    assert b"<ENVELOPE>" in client.get("/api/journals/export?format=tally&start=2026-03-01&end=2026-03-31").content


def test_reports_and_disclosures(client):
    login(client, "auditor", "Lease@116")
    rep = client.get("/api/reports").json()
    assert len(rep) == 19
    for r in rep:
        x = client.get(f"/api/reports/{r['code']}?as_of=2026-03-31&start=2025-04-01&end=2026-03-31")
        assert x.status_code == 200, (r["code"], x.text)
    d = client.get("/api/disclosures?start=2025-04-01&end=2026-03-31").json()
    assert d["liability_movement"]["reconciliation_difference"] == "0.00"
    assert all(r["reconciliation_difference"] == "0.00" for r in d["rou_movement"])
    for fmt in ("xlsx", "csv", "pdf"):
        assert client.get(f"/api/reports/lease_register/export?format={fmt}&as_of=2026-03-31").status_code == 200


def test_lease_exports(client):
    login(client, "admin", "admin116")
    lid = client.get("/api/leases?search=DEMO-001").json()["rows"][0]["id"]
    for kind in ("workpaper.xlsx", "memo.docx", "schedule.xlsx", "schedule.csv", "schedule.pdf"):
        r = client.get(f"/api/leases/{lid}/export/{kind}")
        assert r.status_code == 200 and len(r.content) > 1000, kind
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get(f"/api/leases/{lid}/export/workpaper.xlsx").content))
    assert {"Index", "PV of payments", "Payment schedule", "Journals"} <= set(wb.sheetnames)
    assert str(wb["PV of payments"]["E6"].value).startswith("=1/(1+$B$2)^D6")


def test_import_workflow(client):
    login(client, "preparer", "Lease@116")
    tpl = client.get("/api/imports/templates/lease_master")
    assert tpl.status_code == 200
    csv_data = ("entity_code,description,asset_class,lessor_name,commencement_date,contract_end,rent_amount,discount_rate_pct\n"
                "DEMO,IMPORTED — Store 1,Buildings,Lessor A,01-04-2025,31-03-2028,50000,9.5\n"
                "NOPE,IMPORTED — Bad row,Buildings,Lessor B,01-04-2025,31-03-2024,,9.5\n")
    r = client.post("/api/imports", data={"kind": "lease_master"}, files={"file": ("leases.csv", csv_data, "text/csv")})
    assert r.status_code == 200
    b = r.json()
    assert b["rows_error"] == 1 and b["status"] == "Errors"
    assert client.get(f"/api/imports/{b['id']}/errors.csv").status_code == 200
    good = csv_data.splitlines()[0] + "\n" + csv_data.splitlines()[1] + "\n"
    r = client.post("/api/imports", data={"kind": "lease_master"}, files={"file": ("leases.csv", good, "text/csv")})
    bid = r.json()["id"]
    assert client.post(f"/api/imports/{bid}/approve").status_code in (403,)     # uploader cannot approve (SoD)
    login(client, "approver", "Lease@116")
    r = client.post(f"/api/imports/{bid}/approve")
    assert r.status_code == 200 and r.json()["imported"] == 1
    assert client.get("/api/leases?search=IMPORTED").json()["total"] == 1


def test_ai_reader_rules_mode_and_confirm(client):
    login(client, "preparer", "Lease@116")
    data = (SAMPLES / "Sample_Leave_and_License_Office.pdf").read_bytes()
    r = client.post("/api/extractions", data={"mode": "rules"}, files={"file": ("office.pdf", data, "application/pdf")})
    assert r.status_code == 200, r.text
    rid = r.json()["extraction_id"]
    for _ in range(60):
        x = client.get(f"/api/extractions/{rid}").json()
        if x["status"] != "Running":
            break
        time.sleep(0.5)
    assert x["status"] == "Completed", x
    f = x["result"]["fields"]
    assert f["rent_amount"]["value"] == "450000"
    assert client.get(f"/api/documents/{x['document_id']}/pages/2.png").status_code == 200
    r = client.post(f"/api/extractions/{rid}/confirm", json={"fields": {}, "answers": {"rc_continue_after_lock_in": True,
                                                                                       "discount_rate": "9.25", "deposit_market_rate": "9"}})
    assert r.status_code == 200, r.text
    lease = r.json()
    assert lease["status"] == "Draft" and len(lease["payments"]) > 60
    calc = client.post(f"/api/leases/{lease['id']}/calculate")
    assert calc.status_code == 200, calc.text
    s = calc.json()["summary"]
    assert s["initial"]["liability"] and s["periods"][-1]["liab_close"] == "0.00"


def test_cloud_ai_requires_permission(client):
    login(client, "preparer", "Lease@116")
    data = (SAMPLES / "Sample_Retail_Lease_Revenue_Share.txt").read_bytes()
    r = client.post("/api/extractions", data={"mode": "claude", "allow_cloud": "true"}, files={"file": ("r.txt", data, "text/plain")})
    assert r.status_code == 400 and "not allowed" in r.text


def test_audit_trail_records_changes(client):
    login(client, "auditor", "Lease@116")
    rows = client.get("/api/audit?limit=2000").json()
    actions = {r["action"] for r in rows}
    assert {"CREATE", "CALCULATE", "APPROVE", "LOCK_PERIOD", "REOPEN_PERIOD", "IMPORT_APPROVE"} <= actions


def test_ledger_integrity_all_demo_leases(client):
    """Posted journals must roll up to the schedule balances for every capitalised lease at every period end
    (liability, ROU cost, accumulated depreciation and impairment) — including FX and sale-and-leaseback leases."""
    from decimal import Decimal
    from sqlalchemy import select
    from app.db.base import session_scope
    from app.db.models import Lease
    from app.services.lease_service import reporting_run
    from app.services.report_service import functional_summary
    checked = 0
    with session_scope() as db:
        for lease in db.scalars(select(Lease)).all():
            run = reporting_run(db, lease)
            if not run or run.summary.get("kind") != "LESSEE":
                continue
            fs = functional_summary(run.summary)
            posts = sorted(run.summary.get("postings", []), key=lambda p: p["date"])
            for row in fs["periods"]:
                pe = row["period_end"]
                bal = {}
                for p in posts:
                    if p["date"] <= pe:
                        for ln in p["lines"]:
                            bal[ln["role"]] = bal.get(ln["role"], Decimal(0)) + Decimal(ln["debit"]) - Decimal(ln["credit"])
                assert -bal.get("LEASE_LIABILITY", Decimal(0)) == Decimal(row["liab_close"]), (lease.lease_code, pe, "liability")
                assert bal.get("ROU_ASSET", Decimal(0)) == Decimal(row["rou_cost_close"]), (lease.lease_code, pe, "rou cost")
                assert -bal.get("ROU_ACC_DEP", Decimal(0)) == Decimal(row["rou_accdep_close"]), (lease.lease_code, pe, "acc dep")
                assert -bal.get("ROU_ACC_IMP", Decimal(0)) == Decimal(row["rou_accimp_close"]), (lease.lease_code, pe, "acc imp")
                checked += 1
    assert checked > 100


def test_portfolio_endpoints(client):
    login(client, "admin", "admin116")
    r = client.get("/api/payments?start=2026-01-01&end=2026-03-31")
    assert r.status_code == 200
    rows = r.json()["rows"]
    assert rows and all("2026-01-01" <= x["date"] <= "2026-03-31" for x in rows)
    assert {"Payment"} <= {x["direction"] for x in rows}
    ev = client.get("/api/events?event_type=MODIFICATION,TERMINATION")
    assert ev.status_code == 200 and {e["event_type"] for e in ev.json()} <= {"MODIFICATION", "TERMINATION"}
    sch = client.get("/api/extraction-schema").json()
    keys = {f["key"] for f in sch}
    assert {"commencement_date", "rent_amount", "rent_frequency"} <= keys
    assert any(f["choices"] for f in sch)


def test_stale_flag_tracks_input_changes_only(client):
    login(client, "preparer", "Lease@116")
    r = client.post("/api/leases", json={"description": "STALE test", "commencement_date": "2025-04-01", "contract_end": "2027-03-31",
                                         "discount_rate_pct": "9", "asset_class_id": 1})
    lid = r.json()["id"]
    client.put(f"/api/leases/{lid}/payment-terms", json={"amount": "100000", "start_date": "2025-04-01", "end_date": "2027-03-31",
                                                         "frequency_months": 1, "timing": "ADVANCE"})
    assert client.post(f"/api/leases/{lid}/calculate").status_code == 200
    assert client.get(f"/api/leases/{lid}").json()["stale"] is False
    time.sleep(0.01)
    client.patch(f"/api/leases/{lid}", json={"discount_rate_pct": "9.5"})
    assert client.get(f"/api/leases/{lid}").json()["stale"] is True
    client.post(f"/api/leases/{lid}/calculate")
    assert client.get(f"/api/leases/{lid}").json()["stale"] is False
    assert client.post(f"/api/leases/{lid}/workflow/submit").status_code == 200
    assert client.get(f"/api/leases/{lid}").json()["stale"] is False


def test_launcher_shutdown_needs_the_launch_token(client, monkeypatch):
    """Only the Lease116.exe launcher that started the server (per-launch secret token) may stop it."""
    monkeypatch.delenv("LEASE116_LAUNCHER_TOKEN", raising=False)
    assert client.post("/api/system/shutdown").status_code == 403           # console start: no token → never
    monkeypatch.setenv("LEASE116_LAUNCHER_TOKEN", "3f9a1c0d5e7b4a6f8c2d1e0f9a8b7c6d")
    assert client.post("/api/system/shutdown").status_code == 403
    assert client.post("/api/system/shutdown", headers={"X-Lease116-Token": "wrong"}).status_code == 403
    r = client.post("/api/system/shutdown", headers={"X-Lease116-Token": "3f9a1c0d5e7b4a6f8c2d1e0f9a8b7c6d"})
    assert r.status_code == 200 and r.json()["status"] == "stopping"
    assert client.get("/api/health").json()["status"] == "ok"               # TestClient has no uvicorn server to stop


def _read(client, data: bytes, name: str) -> dict:
    r = client.post("/api/extractions", data={"mode": "rules"}, files={"file": (name, data, "text/plain")})
    assert r.status_code == 200, r.text
    rid = r.json()["extraction_id"]
    for _ in range(100):
        x = client.get(f"/api/extractions/{rid}").json()
        if x["status"] != "Running":
            return x
        time.sleep(0.2)
    raise AssertionError("read did not finish")


def test_admin_deletes_unused_agreement_reads_but_keeps_lease_evidence(client):
    """Administrator housekeeping: reads not used for a lease can be deleted (read, uploaded copy, previews, audited);
    a read that created a lease is kept as that lease's source evidence."""
    import hashlib

    from app import config
    login(client, "preparer", "Lease@116")
    base = (SAMPLES / "Sample_Retail_Lease_Revenue_Share.txt").read_bytes()
    data = base + f"\n[deletion test {time.time()}]\n".encode()
    stored = config.DOCS_DIR / hashlib.sha256(data).hexdigest()[:2] / f"{hashlib.sha256(data).hexdigest()[:16]}_delete-me.txt"
    a, b = _read(client, data, "delete-me.txt"), _read(client, data, "delete-me.txt")   # same file twice: one stored copy
    kept = _read(client, base + f"\n[kept {time.time()}]\n".encode(), "kept.txt")
    r = client.post(f"/api/extractions/{kept['id']}/confirm", json={"fields": {}, "answers": {
        "rc_continue_after_lock_in": True, "discount_rate": "9.5", "deposit_market_rate": "9"}})
    assert r.status_code == 200, r.text
    assert client.delete(f"/api/extractions/{a['id']}").status_code == 403            # preparer cannot delete
    assert client.post("/api/extractions/delete-unused").status_code == 403

    login(client, "admin", "admin116")
    assert stored.exists()
    r = client.delete(f"/api/extractions/{a['id']}")
    assert r.status_code == 200 and r.json()["deleted"] == 1
    assert client.get(f"/api/extractions/{a['id']}").status_code == 404
    assert client.get(f"/api/documents/{a['document_id']}/file").status_code == 404
    assert stored.exists()                                                               # still used by the second upload
    assert not (config.DOCS_DIR / "previews" / str(a["document_id"])).exists()

    r = client.delete(f"/api/extractions/{kept['id']}")                                  # created a lease → kept
    assert r.status_code == 409 and "source evidence" in r.json()["detail"]["message"]
    assert client.get(f"/api/extractions/{kept['id']}").status_code == 200

    r = client.post("/api/extractions/delete-unused")
    assert r.status_code == 200 and r.json()["deleted"] >= 1
    assert client.get(f"/api/extractions/{b['id']}").status_code == 404
    assert not stored.exists()                                                           # last user of the file gone
    remaining = client.get("/api/extractions").json()
    assert remaining and all(x["lease_id"] for x in remaining)                          # only lease evidence remains
    trail = client.get("/api/audit?limit=5000").json()
    deleted = {(t["object_type"], t["object_id"]) for t in trail if t["action"] == "DELETE"}
    assert ("ExtractionRun", str(a["id"])) in deleted and ("Document", str(a["document_id"])) in deleted
    assert ("ExtractionRun", str(b["id"])) in deleted


def test_refused_upload_leaves_no_copy_on_disk(client):
    import hashlib

    from app import config
    login(client, "preparer", "Lease@116")
    data = f"Refused cloud upload {time.time()}".encode()
    sha = hashlib.sha256(data).hexdigest()
    r = client.post("/api/extractions", data={"mode": "claude", "allow_cloud": "true"}, files={"file": ("refused.txt", data, "text/plain")})
    assert r.status_code == 400
    assert not (config.DOCS_DIR / sha[:2] / f"{sha[:16]}_refused.txt").exists()


# ------------------------------------------------------------------------------------------------ lessor module (Ind AS 116.61–97)
def _lessor_ledger_check(summary: dict) -> int:
    """Posted lessor journals must roll up to the schedule balances at every period end."""
    from decimal import Decimal
    posts = sorted(summary.get("postings", []), key=lambda p: p["date"])
    checked = 0
    for row in summary["rows"]:
        pe = row["period_end"]
        bal = {}
        for p in posts:
            if p["date"] <= pe:
                for ln in p["lines"]:
                    bal[ln["role"]] = bal.get(ln["role"], Decimal(0)) + Decimal(ln["debit"]) - Decimal(ln["credit"])
        g = lambda k: bal.get(k, Decimal(0))  # noqa: E731
        assert g("NET_INVESTMENT_LEASE") + g("NET_INVESTMENT_SUBLEASE") == Decimal(row["ni_close"]), (pe, "net investment")
        assert -g("LOSS_ALLOWANCE_LEASE_RECEIVABLES") == Decimal(row["loss_allowance_close"]), (pe, "loss allowance")
        assert g("ACCRUED_LEASE_INCOME") == Decimal(row["accrued_close"]), (pe, "accrued income")
        assert g("LESSOR_IDC_ASSET") == Decimal(row["idc_close"]), (pe, "IDC")
        assert -g("SECURITY_DEPOSIT_RECEIVED") == Decimal(row["dep_close"]), (pe, "deposit")
        checked += 1
    return checked


def test_lessor_register_filter_and_demo_ledger_integrity(client):
    from decimal import Decimal
    login(client, "admin", "admin116")
    rows = client.get("/api/leases?page_size=100&role=LESSOR").json()["rows"]
    by = {r["lease_code"]: r for r in rows}
    assert {"DEMO-010", "DEMO-012"} <= set(by) and all(r["role"] == "LESSOR" for r in rows)
    assert by["DEMO-012"]["classification"] == "FINANCE" and by["DEMO-010"]["classification"] == "OPERATING"
    checked = 0
    for code in ("DEMO-010", "DEMO-012"):
        run = client.get(f"/api/leases/{by[code]['id']}/result").json()
        s = run["summary"]
        assert s["kind"] == "LESSOR"
        checked += _lessor_ledger_check(s)
    assert checked > 100
    # DEMO-012 at 31-Mar-2026: undiscounted receipts reconciled to the net investment (para 94)
    pos = client.get(f"/api/leases/{by['DEMO-012']['id']}/lessor-position?as_of=2026-03-31").json()
    rc = pos["reconciliation"]
    assert Decimal(rc["difference"]) == 0
    assert (Decimal(rc["undiscounted_lease_payments"]) - Decimal(rc["unearned_finance_income"])
            + Decimal(rc["discounted_unguaranteed_residual"])) == Decimal(rc["net_investment"])


def test_lessor_manual_lease_workpaper_memo_and_position(client):
    from decimal import Decimal

    from docx import Document as Docx
    from openpyxl import load_workbook
    login(client, "preparer", "Lease@116")
    r = client.post("/api/leases", json={"description": "TEST — CNC machine let out on finance lease", "commencement_date": "2025-04-01",
                                        "contract_end": "2030-03-31", "asset_class_id": 1, "role": "LESSOR",
                                        "lessor_details": {"fair_value": "2500000", "carrying_amount": "2100000", "economic_life_months": 72,
                                                           "unguaranteed_residual": "200000", "lessor_idc": "25000"}})
    assert r.status_code == 200, r.text
    lid = r.json()["id"]
    r = client.put(f"/api/leases/{lid}/payment-terms", json={"amount": "150000", "start_date": "2025-04-01", "end_date": "2030-03-31",
                                                           "frequency_months": 3, "timing": "ARREARS", "alignment": "ANNIVERSARY"})
    assert r.status_code == 200 and r.json()["generated"] == 20
    calc = client.post(f"/api/leases/{lid}/calculate")
    assert calc.status_code == 200, calc.text
    s = calc.json()["summary"]
    assert s["kind"] == "LESSOR" and s["classification"] == "FINANCE"
    assert s["net_investment"] == "2525000.00"                                               # FV + IDC
    assert s["derecognition_gain"] == "400000.00" and Decimal(s["selling_profit"]) == 0     # FV - carrying amount (Ind AS 16.68)
    assert abs(float(s["implicit_rate_pct"]) - 9.2717) < 0.0001
    assert s["rows"][-1]["ni_close"] == "0" or Decimal(s["rows"][-1]["ni_close"]) == 0
    assert _lessor_ledger_check(s) == len(s["rows"])
    pos = client.get(f"/api/leases/{lid}/lessor-position?as_of=2026-03-31").json()
    assert pos["available"] and Decimal(pos["reconciliation"]["difference"]) == 0 and Decimal(pos["total_undiscounted"]) == Decimal("2400000")
    # lessor workpaper and memo
    wb = load_workbook(io.BytesIO(client.get(f"/api/leases/{lid}/export/workpaper.xlsx?as_of=2026-03-31").content))
    assert {"Index", "Inputs", "Classification", "PV of receipts", "Schedule", "Maturity & reconciliation", "Journals", "Judgments"} <= set(wb.sheetnames)
    memo = Docx(io.BytesIO(client.get(f"/api/leases/{lid}/export/memo.docx?as_of=2026-03-31").content))
    text = "\n".join(p.text for p in memo.paragraphs)
    assert "Lease accounting memo (lessor)" in text and "Classification (Ind AS 116.61–66)" in text
    for kind in ("schedule.xlsx", "schedule.csv", "schedule.pdf"):
        r = client.get(f"/api/leases/{lid}/export/{kind}")
        assert r.status_code == 200 and len(r.content) > 500, kind


def test_lessor_blocking_messages_never_assume_inputs(client):
    login(client, "preparer", "Lease@116")
    r = client.post("/api/leases", json={"description": "TEST — floor let out, inputs pending", "commencement_date": "2025-04-01",
                                        "contract_end": "2028-03-31", "asset_class_id": 1, "role": "LESSOR"})
    lid = r.json()["id"]
    client.put(f"/api/leases/{lid}/payment-terms", json={"amount": "100000", "start_date": "2025-04-01", "end_date": "2028-03-31",
                                                        "frequency_months": 1, "timing": "ADVANCE", "alignment": "ANNIVERSARY"})
    r = client.post(f"/api/leases/{lid}/calculate")
    assert r.status_code == 400 and "classification cannot be concluded" in r.text.lower()
    client.put(f"/api/leases/{lid}/deposit", json={"amount": "600000", "payment_date": "2025-04-01", "refund_date": "2028-04-01"})
    client.patch(f"/api/leases/{lid}", json={"lessor_details": {"fair_value": "90000000", "carrying_amount": "60000000",
                                                                "economic_life_months": 480}})
    r = client.post(f"/api/leases/{lid}/calculate")
    assert r.status_code == 400 and "market interest rate is required" in r.text
    client.put(f"/api/leases/{lid}/deposit", json={"amount": "600000", "payment_date": "2025-04-01", "refund_date": "2028-04-01",
                                                  "market_rate_pct": "9"})
    r = client.post(f"/api/leases/{lid}/calculate")
    assert r.status_code == 200 and r.json()["summary"]["classification"] == "OPERATING"


def test_lessor_reader_confirm_and_csv_import(client):
    login(client, "preparer", "Lease@116")
    data = (SAMPLES / "Sample_Retail_Lease_Revenue_Share.txt").read_bytes()
    r = client.post("/api/extractions", data={"mode": "rules"}, files={"file": ("retail.txt", data, "text/plain")})
    rid = r.json()["extraction_id"]
    for _ in range(60):
        x = client.get(f"/api/extractions/{rid}").json()
        if x["status"] != "Running":
            break
        time.sleep(0.5)
    assert x["status"] == "Completed", x
    r = client.post(f"/api/extractions/{rid}/confirm", json={"fields": {"payment_timing": "ADVANCE"}, "role": "LESSOR",
                                                           "answers": {"rc_continue_after_lock_in": True, "deposit_market_rate": "9"}})
    assert r.status_code == 200, r.text
    lease = r.json()
    assert lease["role"] == "LESSOR" and lease["lessor"] == "VWX Apparel Private Limited"        # counterparty = the lessee
    assert lease["discount_rate_pct"] in (None, "") and lease["deposit"]["market_rate_pct"] == "9"
    assert client.post(f"/api/leases/{lease['id']}/calculate").status_code == 400                # classification inputs pending
    client.patch(f"/api/leases/{lease['id']}", json={"lessor_details": {"fair_value": "100000000", "carrying_amount": "60000000",
                                                                        "economic_life_months": 600}})
    calc = client.post(f"/api/leases/{lease['id']}/calculate")
    assert calc.status_code == 200, calc.text
    s = calc.json()["summary"]
    assert s["classification"] == "OPERATING" and s["totals"]["total_non_lease_income"] != "0"      # CAM kept out (para 17)
    # CSV import of a lessor lease
    csv_data = ("entity_code,description,asset_class,lessor_name,commencement_date,contract_end,rent_amount,role,fair_value,"
                "carrying_amount,economic_life_months,deposit_amount,deposit_market_rate_pct\n"
                "DEMO,IMPORTED LESSOR — Unit 7,Buildings,Tenant T Private Limited,01-04-2025,31-03-2030,250000,LESSOR,"
                "150000000,90000000,600,1500000,9\n"
                "DEMO,IMPORTED LESSOR — Bad,Buildings,Tenant U,01-04-2025,31-03-2030,250000,LESSOR,,,,,\n")
    r = client.post("/api/imports", data={"kind": "lease_master"}, files={"file": ("lessor.csv", csv_data, "text/csv")})
    assert r.status_code == 200
    bid = r.json()["id"]
    login(client, "approver", "Lease@116")
    r = client.post(f"/api/imports/{bid}/approve")
    assert r.status_code == 200 and r.json()["imported"] == 2
    got = client.get("/api/leases?search=IMPORTED LESSOR&page_size=10").json()["rows"]
    assert len(got) == 2 and all(g["role"] == "LESSOR" for g in got)
    unit7 = next(g for g in got if "Unit 7" in g["description"])
    login(client, "preparer", "Lease@116")
    calc = client.post(f"/api/leases/{unit7['id']}/calculate")
    assert calc.status_code == 200, calc.text
    assert calc.json()["summary"]["classification"] == "OPERATING"
    assert calc.json()["summary"]["deposit"]["market_rate_pct"] == "9"


def test_lessor_event_preview_dashboard_disclosures_and_reports(client):
    from decimal import Decimal
    login(client, "admin", "admin116")
    lid = client.get("/api/leases?search=DEMO-012").json()["rows"][0]["id"]
    r = client.post(f"/api/leases/{lid}/events/preview", json={"event_type": "UGR_REVISION", "effective_date": "2027-04-01",
                                                              "description": "Residual estimate reduced", "details": {"new_unguaranteed_residual": "100000"}})
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["lessor"] and p["event"]["reference"] == "Ind AS 116.77" and Decimal(p["event"]["gain_loss"]) < 0
    assert any(x["event"] == "UGR_REVISION" for x in p["postings"])
    r = client.post(f"/api/leases/{lid}/events/preview", json={"event_type": "REASSESSMENT", "effective_date": "2027-04-01"})
    assert r.status_code == 400 and "does not apply to a lessor lease" in r.text
    d = client.get("/api/dashboard?as_of=2026-03-31").json()["lessor"]
    assert d["count"] >= 2 and d["finance_leases"] >= 1 and d["operating_leases"] >= 1
    assert Decimal(d["kpis"]["net_investment"]) > 0 and Decimal(d["kpis"]["deposits_received"]) > 0
    disc = client.get("/api/disclosures?start=2025-04-01&end=2026-03-31").json()["lessor"]
    assert Decimal(disc["net_investment_movement"]["reconciliation_difference"]) == 0
    assert {x["ref"] for x in disc["income_table"]} == {"90(a)(i)", "90(a)(ii)", "90(a)(iii)", "90(b)"}
    assert disc["maturity_finance"] and disc["maturity_operating"] and disc["reconciliation"]
    pres = disc["presentation"]
    assert Decimal(pres["ni_current"]) + Decimal(pres["ni_noncurrent"]) == Decimal(disc["net_investment_movement"]["closing"])
    for code in ("lessor_schedule", "lessor_maturity"):
        rep = client.get(f"/api/reports/{code}?as_of=2026-03-31&start=2025-04-01&end=2026-03-31").json()
        assert rep["rows"], code
    mat = client.get("/api/reports/lessor_maturity?as_of=2026-03-31").json()
    assert all(Decimal(r.get("check") or 0) == 0 for r in mat["rows"])
    x = client.get("/api/disclosures/export?start=2025-04-01&end=2026-03-31")
    assert x.status_code == 200
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.content))["Lease note"]
    labels = {str(c.value) for row in ws.iter_rows() for c in row if c.value}
    assert any(v.startswith("G. Lessor disclosures") for v in labels)
    assert any(v.startswith("G3. Finance leases — maturity analysis") for v in labels)


def test_runs_from_the_previous_lessor_engine_are_flagged_not_broken(client):
    """A database upgraded from v1.0 may hold lessor runs without the v1.1 structure: every screen must still load, the run is
    kept out of the lessor note / dashboard / reports and listed as an exception until recalculated."""
    from sqlalchemy import select
    from app.db.base import session_scope
    from app.db.models import CalcRun, Lease
    login(client, "admin", "admin116")
    with session_scope() as db:
        lease = db.scalar(select(Lease).where(Lease.lease_code == "DEMO-010"))
        run = db.scalar(select(CalcRun).where(CalcRun.lease_id == lease.id, CalcRun.is_current == True))  # noqa: E712
        s = dict(run.summary)
        run.summary = {"kind": "LESSOR", "classification": "OPERATING", "net_investment": "0", "selling_profit": "0",
                       "rows": [{"period_start": r["period_start"], "period_end": r["period_end"], "lease_income": r["lease_income"],
                                 "receipts": r["receipts"], "accrued_rent_receivable": r["accrued_close"]} for r in s["rows"]],
                       "postings": s["postings"], "flags": [], "explanation": [], "periods": []}
        lid = lease.id
    assert client.get(f"/api/leases/{lid}/lessor-position?as_of=2026-03-31").json()["available"] is False
    assert client.get(f"/api/leases/{lid}/export/workpaper.xlsx").status_code == 409
    for url in ("/api/leases?page_size=100", f"/api/leases/{lid}", "/api/dashboard?as_of=2026-03-31",
                "/api/disclosures?start=2025-04-01&end=2026-03-31", "/api/disclosures/export?start=2025-04-01&end=2026-03-31"):
        assert client.get(url).status_code == 200, url
    for r in client.get("/api/reports").json():
        assert client.get(f"/api/reports/{r['code']}?as_of=2026-03-31&start=2025-04-01&end=2026-03-31").status_code == 200, r["code"]
    note = client.get("/api/disclosures?start=2025-04-01&end=2026-03-31").json()["lessor"]
    assert "DEMO-010" not in [x["lease_code"] for x in note["by_lease"]]
    exc = client.get("/api/reports/exceptions?as_of=2026-03-31").json()["rows"]
    assert any(x["lease_code"] == "DEMO-010" and x["category"] == "Recalculate (lessor engine)" for x in exc)
