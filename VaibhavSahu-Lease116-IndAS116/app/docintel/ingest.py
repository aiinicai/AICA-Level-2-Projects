"""Document ingestion: PDF (digital or scanned), images, Word (.docx) and text.

For each page we keep: text, word/line boxes (for evidence highlighting), whether OCR
was used, a rendered preview image and computer-vision findings (skew, blur, stamps).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from PIL import Image

from . import ocr as ocrmod
from .textfix import normalize_ocr_text
from .vision import VisionReport, image_from_bytes, preprocess, to_png_bytes

try:
    try:
        import pymupdf as fitz  # PyMuPDF >= 1.24.3
    except ImportError:  # pragma: no cover - older PyMuPDF
        import fitz  # type: ignore
    HAS_FITZ = True
except Exception:  # pragma: no cover
    HAS_FITZ = False

SCANNED_TEXT_THRESHOLD = 60   # characters; below this a PDF page is treated as scanned


@dataclass
class Page:
    number: int
    text: str
    source: str                      # pdf-text | ocr | docx | txt
    width: float = 0
    height: float = 0
    lines: list = field(default_factory=list)   # [{"text", "bbox": [x0,y0,x1,y1] normalised 0..1, "conf"}]
    preview_png: Optional[bytes] = None
    vision: Optional[VisionReport] = None
    ocr_engine: str = ""
    page_type: str = "BODY"          # ESTAMP | BODY | SCHEDULE | SIGNATURE | ANNEXURE


@dataclass
class LoadedDocument:
    filename: str
    kind: str                         # pdf | image | docx | txt
    pages: list[Page]
    warnings: list[str] = field(default_factory=list)
    ocr_used: bool = False
    ocr_engine: str = ""

    @property
    def full_text(self) -> str:
        return "\n".join(f"[[PAGE {p.number}]]\n{p.text}" for p in self.pages)


def _classify_page(text: str) -> str:
    t = text.lower()
    if "e-stamp" in t or "estamp" in t or "stamp duty amount" in t or "certificate no" in t and "stamp" in t:
        return "ESTAMP"
    if re.search(r"\bschedule\b", t[:400]) or "description of the premises" in t or "description of the property" in t:
        return "SCHEDULE"
    if re.search(r"\bannexure\b", t[:300]):
        return "ANNEXURE"
    if re.search(r"in witness whereof", t) and len(t) < 2500:
        return "SIGNATURE"
    return "BODY"


def _ocr_page(img: Image.Image, page_no: int, llm=None, ocr_pref: str = "auto") -> Page:
    try:
        proc, rep = preprocess(img)
    except Exception as exc:  # computer-vision failure must not stop OCR — fall back to the original image
        from .vision import VisionReport
        proc, rep = img.convert("RGB"), VisionReport(width=img.width, height=img.height,
                                                     notes=[f"Image clean-up skipped: {type(exc).__name__}"])
    lines, engine = ocrmod.run_ocr(proc, ocr_pref, llm)
    W, H = proc.width, proc.height
    norm_lines = [{"text": normalize_ocr_text(l.text), "conf": round(l.conf, 3),
                   "bbox": [round(l.bbox[0] / W, 4), round(l.bbox[1] / H, 4), round(l.bbox[2] / W, 4), round(l.bbox[3] / H, 4)]}
                  for l in lines]
    text = normalize_ocr_text(ocrmod.lines_to_text(lines))
    return Page(page_no, text, "ocr", W, H, norm_lines, to_png_bytes(img), rep, engine, _classify_page(text))


def load_pdf(data: bytes, filename: str, llm=None, ocr_pref: str = "auto", force_ocr: bool = False,
             render_dpi: int = 110, ocr_dpi: int = 220) -> LoadedDocument:
    pages: list[Page] = []
    warnings: list[str] = []
    ocr_used = False
    engine_used = ""
    if HAS_FITZ:
        doc = fitz.open(stream=data, filetype="pdf")
        if doc.needs_pass:
            raise ValueError("PDF is password-protected — please upload an unlocked copy.")
        for i, pg in enumerate(doc, start=1):
            text = pg.get_text("text") or ""
            if force_ocr or len(text.strip()) < SCANNED_TEXT_THRESHOLD:
                pix = pg.get_pixmap(dpi=ocr_dpi)
                img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
                try:
                    p = _ocr_page(img, i, llm, ocr_pref)
                    ocr_used = True
                    engine_used = p.ocr_engine
                except RuntimeError as exc:
                    warnings.append(f"Page {i}: {exc}")
                    p = Page(i, text, "pdf-text", pg.rect.width, pg.rect.height)
                pages.append(p)
                continue
            W, H = pg.rect.width, pg.rect.height
            lines = []
            for b in pg.get_text("dict")["blocks"]:
                for ln in b.get("lines", []):
                    t = "".join(s["text"] for s in ln.get("spans", [])).strip()
                    if not t:
                        continue
                    x0, y0, x1, y1 = ln["bbox"]
                    lines.append({"text": t, "conf": 1.0, "bbox": [round(x0 / W, 4), round(y0 / H, 4), round(x1 / W, 4), round(y1 / H, 4)]})
            pix = pg.get_pixmap(dpi=render_dpi)
            pages.append(Page(i, text, "pdf-text", W, H, lines, pix.tobytes("png"), None, "", _classify_page(text)))
        return LoadedDocument(filename, "pdf", pages, warnings, ocr_used, engine_used)
    # fallback: pypdf text + pypdfium2 rendering
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    try:
        import pypdfium2 as pdfium  # type: ignore
        pdoc = pdfium.PdfDocument(data)
    except Exception:
        pdoc = None
    for i, pg in enumerate(reader.pages, start=1):
        text = pg.extract_text() or ""
        img = None
        if pdoc is not None:
            img = pdoc[i - 1].render(scale=ocr_dpi / 72).to_pil().convert("RGB")
        if (force_ocr or len(text.strip()) < SCANNED_TEXT_THRESHOLD) and img is not None:
            try:
                p = _ocr_page(img, i, llm, ocr_pref)
                ocr_used = True
                engine_used = p.ocr_engine
                pages.append(p)
                continue
            except RuntimeError as exc:
                warnings.append(f"Page {i}: {exc}")
        pages.append(Page(i, text, "pdf-text", float(pg.mediabox.width), float(pg.mediabox.height), [],
                          to_png_bytes(img) if img else None, None, "", _classify_page(text)))
    return LoadedDocument(filename, "pdf", pages, warnings, ocr_used, engine_used)


def load_image(data: bytes, filename: str, llm=None, ocr_pref: str = "auto") -> LoadedDocument:
    img = Image.open(io.BytesIO(data))
    frames = []
    try:
        i = 0
        while True:
            img.seek(i)
            frames.append(img.copy().convert("RGB"))
            i += 1
    except EOFError:
        pass
    if not frames:
        frames = [image_from_bytes(data)]
    pages = []
    engine = ""
    for n, fr in enumerate(frames, start=1):
        p = _ocr_page(fr, n, llm, ocr_pref)
        engine = p.ocr_engine
        pages.append(p)
    return LoadedDocument(filename, "image", pages, [], True, engine)


def load_docx(data: bytes, filename: str) -> LoadedDocument:
    import docx  # python-docx

    d = docx.Document(io.BytesIO(data))
    chunks: list[list[str]] = [[]]
    from docx.oxml.ns import qn
    for block in d.element.body.iterchildren():
        if block.tag == qn("w:p"):
            para_text = "".join(t.text or "" for t in block.iter(qn("w:t")))
            if any(br.get(qn("w:type")) == "page" for br in block.iter(qn("w:br"))):
                if para_text.strip():
                    chunks[-1].append(para_text)
                chunks.append([])
                continue
            chunks[-1].append(para_text)
        elif block.tag == qn("w:tbl"):
            for row in block.iter(qn("w:tr")):
                cells = []
                for cell in row.iter(qn("w:tc")):
                    cells.append(" ".join((t.text or "") for t in cell.iter(qn("w:t"))).strip())
                chunks[-1].append(" | ".join(cells))
    pages = [Page(i, "\n".join(c), "docx", page_type=_classify_page("\n".join(c))) for i, c in enumerate(chunks, start=1)
             if "\n".join(c).strip()]
    return LoadedDocument(filename, "docx", pages)


def load_text(data: bytes, filename: str) -> LoadedDocument:
    text = data.decode("utf-8", errors="replace")
    parts = re.split(r"\f|\n\s*\[\[PAGE \d+\]\]\s*\n", text)
    pages = [Page(i, p, "txt", page_type=_classify_page(p)) for i, p in enumerate(parts, start=1) if p.strip()]
    return LoadedDocument(filename, "txt", pages)


def load_document(data: bytes, filename: str, llm=None, ocr_pref: str = "auto", force_ocr: bool = False) -> LoadedDocument:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf" or data[:4] == b"%PDF":
        return load_pdf(data, filename, llm, ocr_pref, force_ocr)
    if ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"):
        return load_image(data, filename, llm, ocr_pref)
    if ext == ".docx":
        return load_docx(data, filename)
    if ext in (".txt", ".md"):
        return load_text(data, filename)
    if ext == ".doc":
        raise ValueError("Legacy .doc format is not supported — save the file as .docx or PDF.")
    raise ValueError(f"Unsupported file type: {ext}")
