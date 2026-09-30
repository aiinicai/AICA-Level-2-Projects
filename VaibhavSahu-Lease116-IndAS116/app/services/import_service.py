"""Bulk import: templates, validation, preview, approval and import (spec section 29).

Workflow: Upload -> Validate -> Show errors -> User corrects -> Preview -> Approve -> Import.
Invalid rows are never silently discarded; an error report is available for download.
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..db.models import (AssetClass, Company, Counterparty, DiscountRate, Entity, GLBalance, ImportBatch, Lease,
                         LeaseCost, Deposit, LeaseOptionRow, PaymentScheduleRow)
from ..engine.calendar_utils import add_months, next_day, parse_date, prev_day
from . import audit
from .lease_service import ServiceError, next_lease_code, regenerate_payments

TEMPLATES = {
    "lease_master": [
        ("entity_code", True, "Entity code (as set up in Settings)"), ("description", True, "Lease description"),
        ("asset_class", True, "Asset class name, e.g. Buildings"),
        ("lessor_name", True, "Counterparty: the lessor / licensor — or the lessee (tenant) where role = LESSOR"),
        ("related_party", False, "Y/N"), ("location", False, ""), ("cost_centre", False, ""), ("contract_number", False, ""),
        ("contract_date", False, "DD-MM-YYYY"), ("commencement_date", True, "DD-MM-YYYY"), ("contract_end", True, "DD-MM-YYYY (last day)"),
        ("currency", False, "INR default"), ("rent_amount", True, "Per payment period"), ("frequency", False, "MONTHLY/QUARTERLY/HALF_YEARLY/ANNUAL"),
        ("timing", False, "ADVANCE/ARREARS"), ("due_day", False, "e.g. 7"), ("escalation_pct", False, "e.g. 5"),
        ("escalation_every_months", False, "e.g. 12"), ("rent_free_months", False, ""), ("cam_amount", False, "Non-lease per period"),
        ("discount_rate_pct", False, "IBR % p.a. (required before calculation)"), ("deposit_amount", False, ""),
        ("deposit_market_rate_pct", False, ""), ("lock_in_months", False, "Lessee may terminate after lock-in"),
        ("lock_in_rc_continue", False, "Y/N — reasonably certain to continue after lock-in"), ("renewal_months", False, ""),
        ("renewal_rc", False, "Y/N — reasonably certain to renew (lessee option)"), ("idc_amount", False, "Initial direct costs"),
        ("lease_type", False, "STANDARD/SHORT_TERM/LOW_VALUE (lessee only)"),
        ("role", False, "LESSEE (default) / LESSOR — LESSOR when your entity lets out the asset"),
        ("fair_value", False, "Lessor: fair value of the underlying asset at inception"),
        ("carrying_amount", False, "Lessor: carrying amount of the asset"),
        ("economic_life_months", False, "Lessor: economic life of the asset (months)"),
        ("unguaranteed_residual", False, "Lessor: unguaranteed residual value (finance leases)"),
        ("implicit_rate_pct", False, "Lessor: rate implicit in the lease % (blank = solved)"),
        ("classification", False, "Lessor: FINANCE / OPERATING — only to record a documented judgment"),
        ("classification_rationale", False, "Lessor: rationale (required with classification)"),
    ],
    "payments": [("lease_code", True, "Existing Lease ID"), ("date", True, "DD-MM-YYYY"), ("lease_amount", True, ""),
                 ("non_lease_amount", False, ""), ("category", False, "FIXED/VARIABLE/INDEX_LINKED/RVG/PURCHASE_OPTION/..."),
                 ("description", False, "")],
    "discount_rates": [("currency", True, "INR"), ("tenor_from_months", True, "0"), ("tenor_to_months", True, "60"),
                       ("rate_pct", True, "e.g. 9.25"), ("effective_date", True, "DD-MM-YYYY"), ("entity_code", False, ""),
                       ("security", False, "Secured/Unsecured"), ("source", False, "e.g. bank sanction letter"),
                       ("methodology", False, "Build-up approach")],
    "opening_balances": [("lease_code", True, ""), ("cutover_date", True, "Balances as at start of this date"),
                         ("liability", True, ""), ("rou_cost", True, ""), ("rou_acc_dep", False, ""), ("use_implied_rate", False, "Y/N")],
    "gl_balances": [("account_code", True, ""), ("as_of", True, "DD-MM-YYYY"), ("balance", True, "Debit positive, credit negative"),
                    ("entity_code", False, ""), ("lease_code", False, "")],
}


def template_xlsx(kind: str) -> bytes:
    if kind not in TEMPLATES:
        raise ServiceError("Unknown template")
    wb = Workbook()
    ws = wb.active
    ws.title = kind
    for c, (name, req, hint) in enumerate(TEMPLATES[kind], start=1):
        cell = ws.cell(1, c, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F" if req else "5B7FA6")
        ws.cell(2, c, ("REQUIRED. " if req else "") + hint).font = Font(italic=True, color="666666")
        ws.column_dimensions[cell.column_letter].width = max(16, len(name) + 4)
    ws.cell(3, 1, "← enter data from row 3 (row 2 contains guidance and is ignored)").font = Font(italic=True, color="999999")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _read_rows(data: bytes, filename: str) -> list[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
        out = []
        for i, r in enumerate(rows[1:], start=2):
            if i == 2 and r and isinstance(r[0], str) and (r[0].startswith("REQUIRED") or r[0].startswith("←")):
                continue
            if r is None or all(v in (None, "") for v in r):
                continue
            if isinstance(r[0], str) and r[0].startswith("←"):
                continue
            out.append({"_row": i, **{hdr[j]: r[j] for j in range(min(len(hdr), len(r))) if hdr[j]}})
        return out
    text = data.decode("utf-8-sig", errors="replace")
    rd = csv.DictReader(io.StringIO(text))
    return [{"_row": i, **row} for i, row in enumerate(rd, start=2)
            if any(v not in (None, "") for v in row.values()) and not str(list(row.values())[0]).startswith("REQUIRED")]


def _dec(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return Decimal(repr(v)) if isinstance(v, float) else Decimal(v)
    return Decimal(str(v).replace(",", "").strip())


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return parse_date(v)


def validate(db: Session, kind: str, rows: list[dict]) -> tuple[list[dict], list[dict]]:
    errors, clean = [], []
    spec = TEMPLATES[kind]
    ents = {e.code: e.id for e in db.scalars(select(Entity)).all()}
    leases = {l.lease_code: l.id for l in db.scalars(select(Lease)).all()}
    for r in rows:
        errs = []
        out = {"_row": r["_row"]}
        for name, req, _ in spec:
            v = r.get(name)
            if req and (v is None or str(v).strip() == ""):
                errs.append(f"{name} is required")
                continue
            try:
                if name.endswith("date") or name in ("as_of", "contract_end"):
                    out[name] = _date(v)
                elif name in ("rent_amount", "escalation_pct", "cam_amount", "discount_rate_pct", "deposit_amount",
                              "deposit_market_rate_pct", "idc_amount", "lease_amount", "non_lease_amount", "rate_pct", "liability",
                              "rou_cost", "rou_acc_dep", "balance", "fair_value", "carrying_amount", "unguaranteed_residual",
                              "implicit_rate_pct"):
                    out[name] = _dec(v)
                elif name in ("due_day", "escalation_every_months", "rent_free_months", "lock_in_months", "renewal_months",
                              "tenor_from_months", "tenor_to_months", "economic_life_months"):
                    out[name] = int(Decimal(str(v))) if v not in (None, "") else None
                else:
                    out[name] = str(v).strip() if v not in (None, "") else None
            except (InvalidOperation, ValueError) as exc:
                errs.append(f"{name}: invalid value '{v}' ({exc})")
        if kind == "lease_master":
            if out.get("entity_code") and out["entity_code"] not in ents:
                errs.append(f"entity_code '{out['entity_code']}' not found")
            if out.get("commencement_date") and out.get("contract_end") and out["contract_end"] < out["commencement_date"]:
                errs.append("contract_end precedes commencement_date")
            if out.get("frequency") and out["frequency"] not in ("MONTHLY", "QUARTERLY", "HALF_YEARLY", "ANNUAL"):
                errs.append("frequency must be MONTHLY/QUARTERLY/HALF_YEARLY/ANNUAL")
            if out.get("timing") and out["timing"] not in ("ADVANCE", "ARREARS"):
                errs.append("timing must be ADVANCE/ARREARS")
            if out.get("role"):
                out["role"] = out["role"].upper()
                if out["role"] not in ("LESSEE", "LESSOR"):
                    errs.append("role must be LESSEE or LESSOR")
            if out.get("classification"):
                out["classification"] = out["classification"].upper()
                if out["classification"] not in ("FINANCE", "OPERATING"):
                    errs.append("classification must be FINANCE or OPERATING")
                elif not out.get("classification_rationale"):
                    errs.append("classification_rationale is required when a classification is recorded")
            if out.get("role") == "LESSOR" and out.get("lease_type") in ("SHORT_TERM", "LOW_VALUE"):
                errs.append("recognition exemptions (short-term / low-value) apply to lessees only")
        if kind in ("payments", "opening_balances") and out.get("lease_code") and out["lease_code"] not in leases:
            errs.append(f"lease_code '{out['lease_code']}' not found")
        if errs:
            errors.append({"row": r["_row"], "errors": errs, "data": {k: str(v) for k, v in r.items() if k != "_row"}})
        else:
            clean.append({k: (v.isoformat() if isinstance(v, date) else str(v) if isinstance(v, Decimal) else v) for k, v in out.items()})
    return clean, errors


def upload(db: Session, company: Company, kind: str, data: bytes, filename: str, user) -> ImportBatch:
    if kind not in TEMPLATES:
        raise ServiceError("Unknown import template")
    rows = _read_rows(data, filename)
    clean, errors = validate(db, kind, rows)
    b = ImportBatch(company_id=company.id, template=kind, filename=filename, status="Validated" if not errors else "Errors",
                    uploaded_by=user.id, rows_total=len(rows), rows_valid=len(clean), rows_error=len(errors), errors=errors,
                    preview=clean[:200], payload=clean)
    db.add(b)
    db.flush()
    audit.log(db, user, "IMPORT_UPLOAD", "ImportBatch", b.id, filename, "rows", None, f"{len(rows)} rows; {len(errors)} with errors")
    return b


def error_report_csv(b: ImportBatch) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Row", "Errors", "Data"])
    for e in b.errors:
        w.writerow([e["row"], "; ".join(e["errors"]), str(e["data"])])
    return buf.getvalue().encode("utf-8-sig")


def approve_and_import(db: Session, company: Company, b: ImportBatch, user) -> dict:
    from .settings_service import get_settings
    if b.status not in ("Validated",):
        raise ServiceError("Only fully validated batches can be imported — correct the errors and upload again.")
    if get_settings(company)["controls"].get("segregation_of_duties") and b.uploaded_by == user.id and user.role.code != "ADMIN":
        raise ServiceError("Segregation of duties: the uploader cannot approve the import.", status=403)
    n = 0
    for r in b.payload:
        if b.template == "lease_master":
            _import_lease(db, company, r, user)
        elif b.template == "payments":
            lease = db.scalar(select(Lease).where(Lease.lease_code == r["lease_code"]))
            no = max([p.line_no for p in lease.payments] + [0]) + 1
            lease.payments.append(PaymentScheduleRow(line_no=no, payment_date=parse_date(r["date"]), lease_amount=_dec(r["lease_amount"]),
                                                     non_lease_amount=_dec(r.get("non_lease_amount")) or Decimal(0),
                                                     category=r.get("category") or "FIXED", description=r.get("description"),
                                                     source="MANUAL"))
        elif b.template == "discount_rates":
            ent = db.scalar(select(Entity).where(Entity.code == r.get("entity_code"))) if r.get("entity_code") else None
            db.add(DiscountRate(company_id=company.id, entity_id=ent.id if ent else None, scope="PORTFOLIO", currency=r["currency"],
                                tenor_from_months=int(r["tenor_from_months"]), tenor_to_months=int(r["tenor_to_months"]),
                                rate_pct=_dec(r["rate_pct"]), effective_date=parse_date(r["effective_date"]), security=r.get("security"),
                                source=r.get("source"), methodology=r.get("methodology"), status="Approved", approved_by=user.id,
                                approved_at=datetime.now()))
        elif b.template == "opening_balances":
            lease = db.scalar(select(Lease).where(Lease.lease_code == r["lease_code"]))
            lease.opening_balance = {"cutover_date": r["cutover_date"], "liability": r["liability"], "rou_cost": r["rou_cost"],
                                     "rou_acc_dep": r.get("rou_acc_dep") or "0",
                                     "use_implied_rate": (r.get("use_implied_rate") or "Y").upper().startswith("Y")}
        elif b.template == "gl_balances":
            ent = db.scalar(select(Entity).where(Entity.code == r.get("entity_code"))) if r.get("entity_code") else None
            db.add(GLBalance(company_id=company.id, entity_id=ent.id if ent else None, account_code=str(r["account_code"]),
                             as_of=parse_date(r["as_of"]), balance=_dec(r["balance"]), lease_code=r.get("lease_code"),
                             source=b.filename))
        n += 1
    b.status = "Imported"
    b.approved_by = user.id
    b.approved_at = datetime.now()
    b.imported_at = datetime.now()
    audit.log(db, user, "IMPORT_APPROVE", "ImportBatch", b.id, b.filename, "status", "Validated", "Imported",
              approval_status="Approved")
    return {"imported": n}


def _import_lease(db: Session, company: Company, r: dict, user) -> Lease:
    ent = db.scalar(select(Entity).where(Entity.code == r["entity_code"]))
    ac = db.scalar(select(AssetClass).where(AssetClass.name == r["asset_class"]))
    if ac is None:
        ac = AssetClass(company_id=company.id, code=r["asset_class"][:20].upper().replace(" ", "_"), name=r["asset_class"])
        db.add(ac)
        db.flush()
    cp = db.scalar(select(Counterparty).where(Counterparty.name == r["lessor_name"]))
    if cp is None:
        cp = Counterparty(company_id=company.id, name=r["lessor_name"],
                          related_party=(r.get("related_party") or "N").upper().startswith("Y"))
        db.add(cp)
        db.flush()
    comm, end = parse_date(r["commencement_date"]), parse_date(r["contract_end"])
    freq = {"MONTHLY": 1, "QUARTERLY": 3, "HALF_YEARLY": 6, "ANNUAL": 12}[r.get("frequency") or "MONTHLY"]
    cfg = {"amount": r["rent_amount"], "start_date": comm.isoformat(), "end_date": end.isoformat(), "frequency_months": freq,
           "timing": r.get("timing") or "ADVANCE", "alignment": "CALENDAR" if r.get("due_day") and freq == 1 else "ANNIVERSARY",
           "due_day": r.get("due_day"), "escalations": [], "rent_free": [], "non_lease_amount": r.get("cam_amount") or "0"}
    if r.get("escalation_pct"):
        cfg["escalations"].append({"value": r["escalation_pct"], "every_months": r.get("escalation_every_months") or 12, "compounding": True})
    if r.get("rent_free_months"):
        cfg["rent_free"].append({"start": comm.isoformat(), "end": prev_day(add_months(comm, int(r["rent_free_months"]))).isoformat()})
    role = (r.get("role") or "LESSEE").upper()
    lessor_details = None
    if role == "LESSOR":
        lessor_details = {k: r[k] for k in ("fair_value", "carrying_amount", "economic_life_months", "unguaranteed_residual",
                                            "implicit_rate_pct") if r.get(k) not in (None, "")}
        if r.get("idc_amount"):
            lessor_details["lessor_idc"] = r["idc_amount"]
        if r.get("classification"):
            lessor_details["classification_override"] = r["classification"]
            lessor_details["override_rationale"] = r.get("classification_rationale")
    lease = Lease(lease_code=next_lease_code(db), company_id=company.id, entity_id=ent.id, description=r["description"],
                  asset_class_id=ac.id, counterparty_id=cp.id, location=r.get("location"), cost_centre=r.get("cost_centre"),
                  contract_number=r.get("contract_number"), contract_date=parse_date(r.get("contract_date")), commencement_date=comm,
                  contract_end=end, currency=r.get("currency") or company.functional_currency,
                  payment_frequency=r.get("frequency") or "MONTHLY", payment_timing=r.get("timing") or "ADVANCE",
                  discount_rate_pct=_dec(r.get("discount_rate_pct")) if role == "LESSEE" else None, generator_config=cfg,
                  lease_type=(r.get("lease_type") or "STANDARD") if role == "LESSEE" else "STANDARD", role=role,
                  lessor_details=lessor_details, created_by=user.id, status="Draft")
    db.add(lease)
    db.flush()
    if r.get("lock_in_months") and int(r["lock_in_months"]) > 0:
        rc = r.get("lock_in_rc_continue")
        lease.options.append(LeaseOptionRow(kind="TERMINATION", holder="LESSEE", exercise_date=prev_day(add_months(comm, int(r["lock_in_months"]))),
                                            reasonably_certain=None if rc in (None, "") else str(rc).upper().startswith("Y"),
                                            description=f"Lessee may terminate after {r['lock_in_months']}-month lock-in"))
    if r.get("renewal_months"):
        rc = r.get("renewal_rc")
        lease.options.append(LeaseOptionRow(kind="EXTENSION", holder="LESSEE", exercise_date=next_day(end),
                                            extension_end_date=prev_day(add_months(next_day(end), int(r["renewal_months"]))),
                                            reasonably_certain=None if rc in (None, "") else str(rc).upper().startswith("Y"),
                                            description="Renewal option"))
    if r.get("deposit_amount"):
        lease.deposits.append(Deposit(amount=_dec(r["deposit_amount"]), payment_date=comm, refund_date=next_day(end),
                                      market_rate_pct=_dec(r.get("deposit_market_rate_pct")), interest_bearing=False))
    if r.get("idc_amount") and role == "LESSEE":
        lease.costs.append(LeaseCost(kind="IDC", cost_date=comm, amount=_dec(r["idc_amount"]), description="Initial direct costs (import)"))
    regenerate_payments(db, lease, user)
    audit.log(db, user, "CREATE", "Lease", lease.id, lease.lease_code, None, None, f"Imported from lease master template ({role.title()})")
    return lease
