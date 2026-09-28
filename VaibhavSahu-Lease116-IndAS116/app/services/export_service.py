"""Exports: Excel (incl. ICAI-style audit workpaper with live formulas), CSV, PDF and a Word accounting memo."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..engine.calendar_utils import parse_date
from ..engine.decimal_utils import D, ZERO, fmt_money, q
from ..engine.journals import EVENT_LABELS
from ..engine.references import REFS

INDIAN_FMT = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'
INDIAN_FMT_NEG = '[<=-10000000](##\\,##\\,##\\,##0.00);[<=-100000](##\\,##\\,##0.00);(##,##0.00)'
MONEY_FMT = '#,##0.00;(#,##0.00);"-"'
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(D(v))   # Excel cells are IEEE doubles; exact figures remain in the application and PDF
    except Exception:
        return v


def _cell_value(v, typ):
    if typ in ("money", "num", "pct", "rate"):
        return _num(v)
    if typ == "date" and isinstance(v, str) and len(v) >= 10:
        try:
            return parse_date(v)
        except Exception:
            return v
    return v


def table_xlsx(title: str, columns: list[dict], rows: list[dict], meta: dict | None = None, totals: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31].replace("/", "-")
    r = 1
    ws.cell(r, 1, title).font = Font(bold=True, size=13)
    r += 1
    for k, v in (meta or {}).items():
        ws.cell(r, 1, k).font = BOLD
        ws.cell(r, 2, str(v))
        r += 1
    r += 1
    hdr = r
    for c, col in enumerate(columns, start=1):
        cell = ws.cell(r, c, col["label"])
        cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        r += 1
        for c, col in enumerate(columns, start=1):
            v = _cell_value(row.get(col["key"]), col["type"])
            cell = ws.cell(r, c, v)
            cell.border = BOX
            if col["type"] == "money":
                cell.number_format = MONEY_FMT
            elif col["type"] == "date" and isinstance(v, date):
                cell.number_format = "DD-MMM-YYYY"
            elif col["type"] in ("pct", "rate"):
                cell.number_format = "0.0000"
    if totals:
        r += 1
        ws.cell(r, 1, "Total").font = BOLD
        for c, col in enumerate(columns, start=1):
            if col["key"] in totals and col["type"] == "money":
                first, last = hdr + 1, r - 1
                L = get_column_letter(c)
                cell = ws.cell(r, c, f"=SUBTOTAL(9,{L}{first}:{L}{last})" if last >= first else 0)
                cell.number_format = MONEY_FMT
                cell.font = BOLD
    for c, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(c)].width = 16 if col["type"] in ("money", "date") else min(45, max(10, len(col["label"]) + 4))
    ws.freeze_panes = ws.cell(hdr + 1, 1)
    ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(columns))}{max(hdr, r - (1 if totals else 0))}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table_csv(columns: list[dict], rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c["label"] for c in columns])
    for row in rows:
        w.writerow(["" if row.get(c["key"]) is None else row.get(c["key"]) for c in columns])
    return buf.getvalue().encode("utf-8-sig")


def table_pdf(title: str, columns: list[dict], rows: list[dict], meta: dict | None = None, totals: dict | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=title)
    ss = getSampleStyleSheet()
    small = ss["BodyText"].clone("small")
    small.fontSize = 6.5
    small.leading = 8
    story = [Paragraph(title, ss["Title"])]
    if meta:
        story.append(Paragraph(" &nbsp;|&nbsp; ".join(f"<b>{k}:</b> {v}" for k, v in meta.items()), small))
    story.append(Spacer(1, 4 * mm))

    def fmt(v, typ):
        if v is None or v == "":
            return ""
        if typ == "money":
            return fmt_money(v)
        if typ == "date":
            try:
                return parse_date(v).strftime("%d-%b-%Y")
            except Exception:
                return str(v)
        return str(v)
    data = [[Paragraph(f"<b>{c['label']}</b>", small) for c in columns]]
    for row in rows:
        data.append([Paragraph(fmt(row.get(c["key"]), c["type"]), small) for c in columns])
    if totals:
        data.append([Paragraph("<b>Total</b>", small)] + [Paragraph(f"<b>{fmt(totals.get(c['key']), c['type'])}</b>" if c["key"] in totals else "", small)
                                                         for c in columns[1:]])
    avail = landscape(A4)[0] - 20 * mm
    widths = []
    for c in columns:
        widths.append(2.2 if c["type"] == "money" else 1.6 if c["type"] == "date" else 2.6 if c["type"] == "text" else 1.4)
    tot = sum(widths)
    widths = [w / tot * avail for w in widths]
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F9")])]))
    story.append(t)
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph(f"Generated by Lease116 on {datetime.now():%d-%b-%Y %H:%M}. Amounts in the currency of the report.", small))
    doc.build(story)
    return buf.getvalue()


# --------------------------------------------------------------------------- audit workpaper (ICAI style) with live formulas
def lease_workpaper_xlsx(lease: dict, run: dict, header: dict) -> bytes:
    """Workpaper: Index/header, inputs, PV of payments (live XNPV-style formulas), payment schedule (formulas where the
    lease has no events), period schedules, journals, judgments. Formulas let a reviewer re-perform the calculation."""
    wb = Workbook()
    s = run["summary"]
    init = s.get("initial") or {}
    ws = wb.active
    ws.title = "Index"
    rows = [("Client", header.get("client", "")), ("Period / Financial year", header.get("period", "")),
            ("Workpaper reference", header.get("wp_ref", f"WP-LEASE-{lease['lease_code']}")),
            ("Area", "Leases — Ind AS 116 recomputation"), ("Lease", f"{lease['lease_code']} — {lease['description']}"),
            ("Prepared by / date", f"{header.get('prepared_by', '')} / {header.get('prepared_on', date.today().strftime('%d-%b-%Y'))}"),
            ("Reviewed by / date", f"{header.get('reviewed_by', '')} / {header.get('reviewed_on', '')}"),
            ("Objective", "To recompute the lease liability, right-of-use asset, interest, depreciation and related journals in "
                          "accordance with Ind AS 116 and agree them to the books."),
            ("Source documents", header.get("sources", "Lease agreement; payment schedule; IBR working; client lease register.")),
            ("Procedure", "Inputs agreed to the agreement; lease term judgments reviewed; PV recomputed using live formulas "
                          "(sheet 'PV of payments'); schedules and journals generated by Lease116 (calc run "
                          f"{run.get('run_no')}, status {run.get('status')}, engine {s.get('engine_version', '')})."),
            ("Tick marks", "✓ agreed to source; ∑ cast / cross-cast; Ⓡ recomputed"),
            ("Conclusion", header.get("conclusion", "[To be completed by the preparer after comparison with the books]"))]
    ws.cell(1, 1, "LEASE ACCOUNTING — AUDIT WORKPAPER").font = Font(bold=True, size=13)
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(i, 1, k).font = BOLD
        ws.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    n = len(rows) + 5
    ws.cell(n, 1, "Sheets").font = BOLD
    for j, (sh, desc) in enumerate([("Inputs", "Lease inputs and judgments"), ("PV of payments", "Initial measurement — live formulas"),
                                    ("Payment schedule", "Effective-interest schedule by payment"),
                                    ("Liability schedule", "Monthly roll-forward"), ("ROU schedule", "Monthly roll-forward"),
                                    ("Journals", "Journal entries"), ("Judgments", "Flags requiring reviewer attention")], start=n + 1):
        ws.cell(j, 1, sh)
        ws.cell(j, 2, desc)

    # Inputs
    wi = wb.create_sheet("Inputs")
    items = [("Lease ID", lease["lease_code"]), ("Description", lease["description"]), ("Lessor", lease.get("lessor")),
             ("Asset class", lease.get("asset_class")), ("Commencement date", lease.get("commencement_date")),
             ("Contract end", lease.get("contract_end")), ("Accounting lease term end", (s.get("term") or {}).get("term_end")),
             ("Currency", lease.get("currency")), ("Discount rate % p.a.", init.get("rate_pct")),
             ("Effective annual rate", init.get("effective_annual_rate")), ("Day count", init.get("daycount")),
             ("Rate convention", init.get("convention")), ("Initial lease liability", init.get("liability")),
             ("Initial ROU asset", init.get("rou"))]
    for i, (k, v) in enumerate(items, start=1):
        wi.cell(i, 1, k).font = BOLD
        c = wi.cell(i, 2, _num(v) if k in ("Discount rate % p.a.", "Effective annual rate", "Initial lease liability", "Initial ROU asset") else v)
        if k.startswith("Initial"):
            c.number_format = MONEY_FMT
    r = len(items) + 2
    wi.cell(r, 1, "ROU asset build-up (Ind AS 116.24)").font = BOLD
    for comp in init.get("rou_components", []):
        r += 1
        wi.cell(r, 1, comp[0])
        c = wi.cell(r, 2, _num(comp[1]))
        c.number_format = MONEY_FMT
        wi.cell(r, 3, comp[2])
    r += 2
    wi.cell(r, 1, "Lease term explanation").font = BOLD
    for line in (s.get("term") or {}).get("explanation", []):
        r += 1
        wi.cell(r, 1, line)
    wi.column_dimensions["A"].width = 60
    wi.column_dimensions["B"].width = 22
    wi.column_dimensions["C"].width = 30

    # PV of payments with live formulas
    wp = wb.create_sheet("PV of payments")
    daycount = (init.get("daycount") or "")
    wp.cell(1, 1, "Commencement date").font = BOLD
    wp.cell(1, 2, parse_date(init.get("measurement_date"))).number_format = "DD-MMM-YYYY"
    wp.cell(2, 1, "Effective annual rate (R)").font = BOLD
    wp.cell(2, 2, _num(init.get("effective_annual_rate"))).number_format = "0.0000000000"
    wp.cell(3, 1, "Method").font = BOLD
    wp.cell(3, 2, daycount)
    hdr = ["#", "Payment date", "Amount", "Years (t)", "Discount factor = 1/(1+R)^t", "Present value", "Category", "Engine PV (check)"]
    for c, h in enumerate(hdr, start=1):
        cell = wp.cell(5, c, h)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
    act365 = daycount.startswith("Exact")
    for i, pv in enumerate(init.get("pv_lines", []), start=6):
        wp.cell(i, 1, i - 5)
        wp.cell(i, 2, parse_date(pv["date"])).number_format = "DD-MMM-YYYY"
        wp.cell(i, 3, _num(pv["amount"])).number_format = MONEY_FMT
        if act365:
            wp.cell(i, 4, f"=(B{i}-$B$1)/365").number_format = "0.000000"
        else:
            wp.cell(i, 4, _num(pv["years"])).number_format = "0.000000"
        wp.cell(i, 5, f"=1/(1+$B$2)^D{i}").number_format = "0.0000000000"
        wp.cell(i, 6, f"=C{i}*E{i}").number_format = MONEY_FMT
        wp.cell(i, 7, pv["category"])
        wp.cell(i, 8, _num(pv["present_value"])).number_format = MONEY_FMT
    last = 5 + len(init.get("pv_lines", []))
    tr = last + 1
    wp.cell(tr, 2, "Total").font = BOLD
    wp.cell(tr, 3, f"=SUM(C6:C{last})").number_format = MONEY_FMT
    wp.cell(tr, 6, f"=SUM(F6:F{last})").number_format = MONEY_FMT
    wp.cell(tr, 8, f"=SUM(H6:H{last})").number_format = MONEY_FMT
    wp.cell(tr + 1, 2, "Difference (recomputed − engine)").font = BOLD
    wp.cell(tr + 1, 6, f"=ROUND(F{tr}-H{tr},2)").number_format = MONEY_FMT
    if act365:
        wp.cell(tr + 2, 2, "Cross-check with Excel XNPV")
        wp.cell(tr + 2, 3, "Excel XNPV(R, amounts, dates) discounts to the first date in the range; include the commencement date with a nil amount.")
    for c, w in zip("ABCDEFGH", (5, 14, 16, 12, 22, 18, 20, 18)):
        wp.column_dimensions[c].width = w

    # Payment schedule (formulas if no events)
    wsp = wb.create_sheet("Payment schedule")
    heads = ["#", "Date", "Opening", "Interest", "Payment", "Adjustment", "Principal", "Closing", "Note"]
    for c, h in enumerate(heads, start=1):
        cell = wsp.cell(1, c, h)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
    use_formulas = act365 and not s.get("events")
    wsp.cell(1, 11, "R").font = BOLD
    wsp.cell(1, 12, _num(init.get("effective_annual_rate")))
    wsp.cell(2, 11, "Commencement").font = BOLD
    wsp.cell(2, 12, parse_date(init.get("measurement_date"))).number_format = "DD-MMM-YYYY"
    for i, pr in enumerate(s.get("payment_rows", []), start=2):
        wsp.cell(i, 1, pr["no"])
        wsp.cell(i, 2, parse_date(pr["date"])).number_format = "DD-MMM-YYYY"
        if use_formulas:
            wsp.cell(i, 3, _num(pr["opening"]) if i == 2 else f"=H{i-1}")
            prev_date = "$L$2" if i == 2 else f"B{i-1}"
            wsp.cell(i, 4, f"=ROUND(C{i}*((1+$L$1)^(({'B' + str(i)}-{prev_date})/365)-1),2)")
            wsp.cell(i, 8, f"=C{i}+D{i}-E{i}+F{i}")
        else:
            wsp.cell(i, 3, _num(pr["opening"]))
            wsp.cell(i, 4, _num(pr["interest"]))
            wsp.cell(i, 8, _num(pr["closing"]))
        wsp.cell(i, 5, _num(pr["payment"]))
        wsp.cell(i, 6, _num(pr["adjustment"]))
        wsp.cell(i, 7, f"=E{i}-D{i}")
        wsp.cell(i, 9, pr.get("note"))
        for c in range(3, 9):
            wsp.cell(i, c).number_format = MONEY_FMT
    if use_formulas:
        wsp.cell(3, 11, "Formulas: interest = opening × ((1+R)^(days/365) − 1), rounded to 2 decimals per payment; "
                        "rounding differences vs the monthly schedule are ≤ ₹0.01 per line.")
    for c, w in zip("ABCDEFGHI", (5, 13, 17, 15, 15, 14, 15, 17, 30)):
        wsp.column_dimensions[c].width = w

    # period schedules
    for name, cols in (("Liability schedule", [("period_end", "Period end"), ("liab_open", "Opening"), ("liab_additions", "Additions"),
                                               ("interest", "Interest"), ("payments", "Payments"), ("liab_remeasurement", "Remeasurement"),
                                               ("liab_modification", "Modification"), ("liab_derecognised", "Derecognised"),
                                               ("liab_close", "Closing"), ("liab_current", "Current"), ("liab_noncurrent", "Non-current")]),
                       ("ROU schedule", [("period_end", "Period end"), ("rou_open", "Opening"), ("rou_additions", "Additions"),
                                         ("depreciation", "Depreciation"), ("impairment", "Impairment"),
                                         ("rou_remeasurement", "Remeasurement"), ("rou_modification", "Modification"),
                                         ("rou_derecognised", "Derecognised"), ("rou_close", "Closing"),
                                         ("rou_cost_close", "Gross cost"), ("rou_accdep_close", "Acc. depreciation")])):
        w = wb.create_sheet(name)
        for c, (_, h) in enumerate(cols, start=1):
            cell = w.cell(1, c, h)
            cell.fill, cell.font = HDR_FILL, HDR_FONT
        w.cell(1, len(cols) + 1, "Check (∑)").font = BOLD
        for i, pr in enumerate(s.get("periods", []), start=2):
            for c, (k, _) in enumerate(cols, start=1):
                v = parse_date(pr[k]) if k == "period_end" else _num(pr.get(k))
                cell = w.cell(i, c, v)
                cell.number_format = "DD-MMM-YYYY" if k == "period_end" else MONEY_FMT
            if name == "Liability schedule":
                w.cell(i, len(cols) + 1, f"=ROUND(B{i}+C{i}+D{i}-E{i}+F{i}+G{i}-H{i}-I{i},2)").number_format = MONEY_FMT
            else:
                w.cell(i, len(cols) + 1, f"=ROUND(B{i}+C{i}-D{i}-E{i}+F{i}+G{i}-H{i}-I{i},2)").number_format = MONEY_FMT
        for c in range(1, len(cols) + 2):
            w.column_dimensions[get_column_letter(c)].width = 15
        w.freeze_panes = "B2"

    # journals
    wj = wb.create_sheet("Journals")
    heads = ["Date", "Event", "Narration", "GL role", "Debit", "Credit"]
    for c, h in enumerate(heads, start=1):
        cell = wj.cell(1, c, h)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
    r = 1
    for p in s.get("postings", []):
        for l in p["lines"]:
            r += 1
            wj.cell(r, 1, parse_date(p["date"])).number_format = "DD-MMM-YYYY"
            wj.cell(r, 2, EVENT_LABELS.get(p["event"], p["event"]))
            wj.cell(r, 3, p["narration"])
            wj.cell(r, 4, l["role"])
            wj.cell(r, 5, _num(l["debit"]) or None).number_format = MONEY_FMT
            wj.cell(r, 6, _num(l["credit"]) or None).number_format = MONEY_FMT
    wj.cell(r + 1, 4, "Total").font = BOLD
    wj.cell(r + 1, 5, f"=SUM(E2:E{r})").number_format = MONEY_FMT
    wj.cell(r + 1, 6, f"=SUM(F2:F{r})").number_format = MONEY_FMT
    for c, w in zip("ABCDEF", (13, 22, 70, 28, 16, 16)):
        wj.column_dimensions[c].width = w

    # judgments
    wg = wb.create_sheet("Judgments")
    for c, h in enumerate(["Flag", "Detail", "Reference", "Assumptions"], start=1):
        cell = wg.cell(1, c, h)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
    for i, f in enumerate(s.get("flags", []), start=2):
        wg.cell(i, 1, f.get("title"))
        wg.cell(i, 2, f.get("detail"))
        wg.cell(i, 3, f.get("reference"))
        wg.cell(i, 4, "; ".join(f.get("assumptions") or []))
    for c, w in zip("ABCD", (34, 80, 26, 50)):
        wg.column_dimensions[c].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- accounting memo (Word)
def lease_memo_docx(lease: dict, run: dict, company: str) -> bytes:
    import docx
    from docx.shared import Pt

    s = run["summary"]
    init = s.get("initial") or {}
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    d.add_heading(f"Lease accounting memo — {lease['lease_code']}", 0)
    d.add_paragraph(f"Entity: {lease.get('entity')}    Company: {company}    Framework: Ind AS 116    Prepared: {date.today():%d-%b-%Y}")
    d.add_heading("1. Background and facts", 1)
    d.add_paragraph(f"{lease['description']}. Lessor: {lease.get('lessor') or '—'}. Underlying asset: "
                    f"{lease.get('asset_description') or lease.get('asset_class') or '—'}. Location: {lease.get('location') or '—'}. "
                    f"Contract date {lease.get('contract_date') or '—'}; commencement {lease.get('commencement_date')}; contractual end "
                    f"{lease.get('contract_end')}.")
    d.add_heading("2. Does the contract contain a lease? (Ind AS 116.9–11, B9–B31)", 1)
    a = lease.get("assessment") or {}
    d.add_paragraph(a.get("conclusion") or "The contract conveys the right to control the use of an identified asset for a period of "
                    "time in exchange for consideration. [Preparer to confirm identified asset, substantive substitution rights, "
                    "economic benefits and right to direct use.]")
    d.add_heading("3. Lease term (Ind AS 116.18–21, B34–B41)", 1)
    for line in (s.get("term") or {}).get("explanation", []):
        d.add_paragraph(line, style="List Bullet")
    d.add_heading("4. Lease payments and discount rate (Ind AS 116.26–28)", 1)
    d.add_paragraph(f"Discount rate {init.get('rate_pct')}% p.a. ({lease.get('rate_basis')}); {init.get('convention')}; "
                    f"{init.get('daycount')}. Rate source: {lease.get('rate_source') or '[document IBR methodology and approval]'}.")
    excl = init.get("excluded") or []
    if excl:
        d.add_paragraph(f"{len(excl)} contractual payment line(s) were excluded from the lease liability, e.g.: "
                        + "; ".join(sorted({e['reason'] for e in excl}))[:900])
    d.add_heading("5. Initial measurement (Ind AS 116.23–26)", 1)
    t = d.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Component", "Amount", "Reference"
    for comp in init.get("rou_components", []):
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = comp[0], fmt_money(comp[1]), comp[2]
    row = t.add_row().cells
    row[0].text, row[1].text = "Right-of-use asset at commencement", fmt_money(init.get("rou") or 0)
    d.add_paragraph(f"Lease liability at commencement: {fmt_money(init.get('liability') or 0)} (present value of "
                    f"{len(init.get('pv_lines') or [])} unpaid lease payments).")
    d.add_heading("6. Subsequent measurement and events", 1)
    tot = s.get("totals") or {}
    d.add_paragraph(f"Total interest over the term {fmt_money(tot.get('total_interest') or 0)}; total depreciation "
                    f"{fmt_money(tot.get('total_depreciation') or 0)}; total gains/(losses) on events {fmt_money(tot.get('total_gain_loss') or 0)}.")
    for ev in s.get("events", []):
        d.add_paragraph(f"{ev['type'].title()} effective {ev['effective_date']}: {ev['description']} — liability "
                        f"{fmt_money(ev['liability_before'])} → {fmt_money(ev['liability_after'])}; ROU {fmt_money(ev['rou_before'])} → "
                        f"{fmt_money(ev['rou_after'])}; gain/(loss) {fmt_money(ev['gain_loss'])}. [{ev.get('reference')}]", style="List Bullet")
    if s.get("deposit"):
        dp = s["deposit"]
        d.add_heading("7. Security deposit (Ind AS 109)", 1)
        d.add_paragraph(f"Deposit {fmt_money(dp['amount_paid'])} refundable on {dp['refund_date']}; fair value at market rate "
                        f"{dp['market_rate_pct']}% = {fmt_money(dp['initial_fair_value'])}; difference {fmt_money(dp['difference_prepaid_rent'])} "
                        f"— {dp['treatment']}.")
    d.add_heading("8. Judgments and reviewer attention", 1)
    for f in s.get("flags", []):
        d.add_paragraph(f"{f['title']} — {f['detail']} ({f.get('reference', '')})", style="List Bullet")
    d.add_heading("9. Presentation and disclosure", 1)
    d.add_paragraph("ROU assets presented separately (or disclosed); lease liabilities presented as current/non-current financial "
                    "liabilities (Schedule III, Division II); interest in finance costs; cash flows — principal and interest in "
                    "financing activities (Ind AS 116.50).")
    d.add_heading("10. Approval", 1)
    d.add_paragraph(f"Calculation run {run.get('run_no')} — status {run.get('status')}. Prepared by: ________  Reviewed by: ________  "
                    "Approved by: ________")
    d.add_paragraph("References are to the notified Ind AS 116 text; paragraph descriptions in this memo are summaries.").italic = True
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def disclosure_xlsx(disc: dict, company: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Lease note"
    per = disc["period"]
    ws.cell(1, 1, f"{company} — Notes to financial statements: Leases ({per['framework'].replace('_', ' ')})").font = Font(bold=True, size=12)
    ws.cell(2, 1, f"For the period {per['start']} to {per['end']}")
    r = 4
    ws.cell(r, 1, "A. Amounts recognised (Ind AS 116.53)").font = BOLD
    for x in disc["para53"]:
        r += 1
        ws.cell(r, 1, f"{x['item']} [{x['ref']}]")
        ws.cell(r, 2, _num(x["amount"])).number_format = MONEY_FMT
        for k, v in (x.get("by_class") or {}).items():
            r += 1
            ws.cell(r, 1, f"    – {k}")
            ws.cell(r, 2, _num(v)).number_format = MONEY_FMT
    r += 2
    ws.cell(r, 1, "B. Right-of-use assets — movement by class").font = BOLD
    r += 1
    heads = ["Class", "Opening", "Additions", "Remeasurement / modification", "Derecognition", "Depreciation", "Impairment", "Closing"]
    for c, h in enumerate(heads, start=1):
        cell = ws.cell(r, c, h)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
    for row in disc["rou_movement"]:
        r += 1
        vals = [row["asset_class"], row["opening_nbv"], row["additions"], row["remeasurement_modification"], row["derecognition"],
                row["depreciation"], row["impairment"], row["closing_nbv"]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(r, c, _num(v) if c > 1 else v)
            if c > 1:
                cell.number_format = MONEY_FMT
        if row["asset_class"] == "Total":
            for c in range(1, 9):
                ws.cell(r, c).font = BOLD
    r += 2
    ws.cell(r, 1, "C. Lease liabilities — movement").font = BOLD
    lm = disc["liability_movement"]
    for k, label in (("opening", "Opening balance"), ("additions", "Additions"), ("interest", "Finance cost accrued"),
                     ("payments", "Payments"), ("remeasurement_modification", "Remeasurement / modification"),
                     ("derecognition", "Derecognition on termination"), ("closing", "Closing balance")):
        r += 1
        ws.cell(r, 1, label)
        v = lm.get(k, 0)
        ws.cell(r, 2, -_num(v) if k in ("payments", "derecognition") else _num(v)).number_format = MONEY_FMT
    r += 1
    ws.cell(r, 1, "    Current (Schedule III)")
    ws.cell(r, 2, _num(disc["presentation"]["current"])).number_format = MONEY_FMT
    r += 1
    ws.cell(r, 1, "    Non-current (Schedule III)")
    ws.cell(r, 2, _num(disc["presentation"]["non_current"])).number_format = MONEY_FMT
    r += 2
    ws.cell(r, 1, "D. Maturity analysis — contractual undiscounted cash flows (Ind AS 116.58; Ind AS 107)").font = BOLD
    for m in disc["maturity"]["rows"]:
        r += 1
        ws.cell(r, 1, m["bucket"])
        ws.cell(r, 2, _num(m["undiscounted"])).number_format = MONEY_FMT
    r += 1
    ws.cell(r, 1, "Total undiscounted lease liabilities").font = BOLD
    ws.cell(r, 2, _num(disc["maturity"]["total_undiscounted"])).number_format = MONEY_FMT
    r += 1
    ws.cell(r, 1, "Less: future finance charges")
    ws.cell(r, 2, -_num(disc["maturity"]["future_finance_charges"])).number_format = MONEY_FMT
    r += 1
    ws.cell(r, 1, "Lease liabilities (carrying amount)").font = BOLD
    ws.cell(r, 2, _num(disc["maturity"]["carrying_amount"])).number_format = MONEY_FMT
    r += 2
    ws.cell(r, 1, "E. Cash flow presentation").font = BOLD
    r += 1
    ws.cell(r, 1, disc["cash_flow_classification"])
    r += 2
    ws.cell(r, 1, "F. Qualitative disclosures (Ind AS 116.59–60) — to be completed by the entity").font = BOLD
    for p in disc["qualitative_prompts"]:
        r += 1
        ws.cell(r, 1, f"[{p['ref']}] {p['prompt']}" + (f"  (undiscounted exposure: {p['amount_undiscounted']})" if p.get("amount_undiscounted") else ""))
    ld = disc.get("lessor") or {}
    if ld.get("count"):
        r += 2
        ws.cell(r, 1, "G. Lessor disclosures (Ind AS 116.89–97)" + (" — including subleases (intermediate lessor)" if ld.get("includes_subleases") else "")
                ).font = Font(bold=True, size=12)
        r += 1
        ws.cell(r, 1, "G1. Lease income (para 90, tabular — para 91)").font = BOLD
        for x in ld["income_table"]:
            r += 1
            ws.cell(r, 1, f"{x['item']} [{x['ref']}]")
            ws.cell(r, 2, _num(x["amount"])).number_format = MONEY_FMT
        for x in ld.get("other_items") or []:
            r += 1
            ws.cell(r, 1, f"Not a para 90 item — {x['item']} [{x['ref']}]").font = Font(italic=True)
            ws.cell(r, 2, _num(x["amount"])).number_format = MONEY_FMT
        mv = ld["net_investment_movement"]
        if any(_num(v) for k, v in mv.items() if k != "reconciliation_difference"):
            r += 2
            ws.cell(r, 1, "G2. Net investment in finance leases — movement (para 93)").font = BOLD
            for k, label, sign in (("opening", "Opening balance", 1), ("additions", "Additions — new finance leases", 1),
                                   ("finance_income", "Finance income", 1), ("receipts", "Lease payments received / receivable", -1),
                                   ("residual_returned", "Unguaranteed residual realised — asset returned", -1),
                                   ("remeasurement", "Remeasurements / modifications", 1),
                                   ("derecognised", "Derecognised — terminations / reclassification", -1), ("closing", "Closing balance (gross)", 1)):
                r += 1
                ws.cell(r, 1, label)
                ws.cell(r, 2, sign * (_num(mv[k]) or 0)).number_format = MONEY_FMT
            la = ld["loss_allowance"]
            r += 1
            ws.cell(r, 1, "Loss allowance (Ind AS 109) — closing")
            ws.cell(r, 2, -(_num(la["closing"]) or 0)).number_format = MONEY_FMT
        if ld["maturity_finance"]:
            r += 2
            ws.cell(r, 1, "G3. Finance leases — maturity analysis of lease payments receivable (para 94)").font = BOLD
            for m in ld["maturity_finance"]:
                r += 1
                ws.cell(r, 1, m["bucket"])
                ws.cell(r, 2, _num(m["amount"])).number_format = MONEY_FMT
            rc = ld["reconciliation"]
            for label, key, sign in (("Total undiscounted lease payments receivable", "undiscounted_lease_payments", 1),
                                     ("Less: unearned finance income", "unearned_finance_income", -1),
                                     ("Add: discounted unguaranteed residual value", "discounted_unguaranteed_residual", 1),
                                     ("Net investment in finance leases", "net_investment", 1), ("Less: loss allowance", "loss_allowance", -1),
                                     ("Net investment net of loss allowance", "net_investment_net_of_allowance", 1)):
                r += 1
                ws.cell(r, 1, label).font = BOLD if key in ("net_investment", "net_investment_net_of_allowance") else Font()
                ws.cell(r, 2, sign * (_num(rc[key]) or 0)).number_format = MONEY_FMT
        if ld["maturity_operating"]:
            r += 2
            ws.cell(r, 1, "G4. Operating leases — maturity analysis of lease payments (para 97)").font = BOLD
            for m in ld["maturity_operating"]:
                r += 1
                ws.cell(r, 1, m["bucket"])
                ws.cell(r, 2, _num(m["amount"])).number_format = MONEY_FMT
        pr = ld["presentation"]
        r += 2
        ws.cell(r, 1, "G5. Balance sheet (lessor)").font = BOLD
        for label, key in (("Net investment — current", "ni_current"), ("Net investment — non-current", "ni_noncurrent"),
                           ("Accrued lease income (straight-lining) — asset", "accrued_income_asset"),
                           ("Deferred lease income — liability", "deferred_income_liability"),
                           ("Security deposits received — financial liability (amortised cost)", "deposits_received"),
                           ("Initial direct costs — operating leases (in the carrying amount of the asset)", "idc_carrying")):
            r += 1
            ws.cell(r, 1, label)
            ws.cell(r, 2, _num(pr[key])).number_format = MONEY_FMT
        if ld.get("current_method"):
            r += 1
            ws.cell(r, 1, "Basis: " + ld["current_method"]).font = Font(italic=True)
        r += 2
        ws.cell(r, 1, "G6. Qualitative disclosures (paras 92–96) — to be completed by the entity").font = BOLD
        for p in ld["qualitative_prompts"]:
            r += 1
            ws.cell(r, 1, f"[{p['ref']}] {p['prompt']}")
        if ld.get("operating_lease_assets"):
            r += 1
            ws.cell(r, 1, "Assets subject to operating leases (for the Ind AS 16 / 40 disaggregation): "
                    + "; ".join(f"{a['lease_code']} ({a['asset_class']})" for a in ld["operating_lease_assets"]))
    ws.column_dimensions["A"].width = 80
    for c in "BCDEFGH":
        ws.column_dimensions[c].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- lessor: schedule columns, workpaper, memo
LESSOR_COLS = [("period_end", "Period end", "date"), ("classification", "Classification", "text"),
               ("ni_open", "Opening net investment", "money"), ("ni_additions", "Additions", "money"),
               ("finance_income", "Finance income", "money"), ("ni_receipts", "Receipts applied", "money"),
               ("ni_residual_returned", "Residual asset returned", "money"), ("ni_remeasurement", "Remeasurement", "money"),
               ("ni_derecognised", "Derecognised", "money"), ("ni_close", "Closing net investment", "money"),
               ("ni_current", "Current", "money"), ("ni_noncurrent", "Non-current", "money"),
               ("lease_income", "Operating lease income", "money"), ("lease_payments_due", "Lease payments due", "money"),
               ("accrued_adjustment", "Accrued income adjustment", "money"), ("accrued_close", "Accrued / (deferred) lease income", "money"),
               ("variable_income", "Variable lease income", "money"), ("non_lease_income", "Non-lease revenue", "money"),
               ("receipts", "Lease payments billed", "money"), ("idc_amortisation", "IDC amortisation", "money"),
               ("idc_close", "IDC carrying", "money"), ("dep_unwinding", "Deposit unwinding", "money"),
               ("dep_close", "Deposit received (carrying)", "money"), ("ecl_charge", "ECL charge", "money"),
               ("loss_allowance_close", "Loss allowance", "money"), ("gain_loss", "Gain / (loss) on events", "money")]


def lessor_schedule_columns(rows: list[dict]) -> list[dict]:
    """Only the columns that carry amounts for this lease (a finance lease has no straight-line columns and vice versa)."""
    cols = []
    for k, label, typ in LESSOR_COLS:
        if typ != "money" or any(D(r.get(k) or 0) != 0 for r in rows):
            cols.append({"key": k, "label": label, "type": typ})
    return cols


def _hdr(ws, r, heads, start_col=1):
    for c, h in enumerate(heads, start=start_col):
        cell = ws.cell(r, c, h)
        cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def lessor_workpaper_xlsx(lease: dict, run: dict, header: dict, as_of: date) -> bytes:
    """Lessor workpaper: classification evidence, PV of receipts (live formulas), schedule with arithmetic checks,
    security deposit (Ind AS 109, live formulas), maturity analysis and reconciliation (paras 94 / 97), journals, judgments."""
    from ..engine.lessor import position_at

    s = run["summary"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Index"
    fin = s.get("classification") == "FINANCE"
    rows = [("Client", header.get("client", "")), ("Period / Financial year", header.get("period", "")),
            ("Workpaper reference", header.get("wp_ref", f"WP-LEASE-{lease['lease_code']}")),
            ("Area", "Leases (lessor) — Ind AS 116 recomputation"), ("Lease", f"{lease['lease_code']} — {lease['description']}"),
            ("Prepared by / date", f"{header.get('prepared_by', '')} / {header.get('prepared_on', date.today().strftime('%d-%b-%Y'))}"),
            ("Reviewed by / date", f"{header.get('reviewed_by', '')} / {header.get('reviewed_on', '')}"),
            ("Objective", "To evaluate the classification of the lease (Ind AS 116.61–66) and recompute the "
                          + ("net investment and finance income" if fin else "straight-line lease income")
                          + ", the security deposit received (Ind AS 109) and the related journals, and agree them to the books."),
            ("Source documents", header.get("sources", "Lease agreement; fair value / useful-life evidence; fixed-asset register; "
                                                       "billing records; deposit receipt.")),
            ("Procedure", "Inputs agreed to the agreement and asset records; classification indicators evaluated (sheet 'Classification'); "
                          "PV recomputed with live formulas; schedule arithmetic checked row by row; maturity analysis and reconciliation "
                          f"at {as_of:%d-%b-%Y}; journals generated by Lease116 (calc run {run.get('run_no')}, status {run.get('status')}, "
                          f"lessor engine {s.get('engine_version', '')})."),
            ("Tick marks", "✓ agreed to source; ∑ cast / cross-cast; Ⓡ recomputed"),
            ("Conclusion", header.get("conclusion", "[To be completed by the preparer after comparison with the books]"))]
    ws.cell(1, 1, "LEASE ACCOUNTING (LESSOR) — AUDIT WORKPAPER").font = Font(bold=True, size=13)
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(i, 1, k).font = BOLD
        ws.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    sheets = [("Inputs", "Lease and lessor inputs, lease term"), ("Classification", "Indicators 63(a)–(e), 64(a)–(c); PV and term tests"),
              ("PV of receipts", "Present value of lease payments / net investment — live formulas"),
              ("Schedule", "Monthly schedule with arithmetic check columns")]
    if s.get("deposit"):
        sheets.append(("Deposit", "Security deposit received — fair value and amortised cost (live formulas)"))
    sheets += [("Maturity & reconciliation", f"Paras 94 / 97 at {as_of:%d-%b-%Y}"), ("Journals", "Journal entries"),
               ("Judgments", "Flags and issues requiring reviewer attention")]
    if s.get("events"):
        sheets.append(("Events", "Modifications, terminations, residual revisions, expected credit losses"))
    n = len(rows) + 5
    ws.cell(n, 1, "Sheets").font = BOLD
    for j, (sh, desc) in enumerate(sheets, start=n + 1):
        ws.cell(j, 1, sh)
        ws.cell(j, 2, desc)

    # Inputs
    wi = wb.create_sheet("Inputs")
    det = lease.get("lessor_details") or {}
    term = s.get("term") or {}
    dep = s.get("deposit") or {}
    items = [("Lease ID", lease["lease_code"]), ("Description", lease["description"]), ("Role", "Lessor"),
             ("Lessee (counterparty)", lease.get("lessor")), ("Asset class / asset", f"{lease.get('asset_class') or ''} — {lease.get('asset_description') or ''}"),
             ("Commencement date", parse_date(lease.get("commencement_date"))), ("Contract end", parse_date(lease.get("contract_end"))),
             ("Lease term end", parse_date(term.get("term_end"))), ("Currency", lease.get("currency")),
             ("Fair value of the underlying asset", _num(det.get("fair_value"))), ("Carrying amount of the asset", _num(det.get("carrying_amount"))),
             ("Economic life (months)", _num(det.get("economic_life_months"))), ("Unguaranteed residual value", _num(det.get("unguaranteed_residual"))),
             ("Lessor initial direct costs", _num(det.get("lessor_idc"))),
             ("Rate implicit in the lease % p.a.", _num(s.get("implicit_rate_pct"))), ("Rate source", s.get("rate_source")),
             ("'Substantially all' benchmark % (policy)", _num(det.get("substantially_all_pct") or 90)),
             ("'Major part' benchmark % (policy)", _num(det.get("major_part_pct") or 75)),
             ("Manufacturer / dealer lessor", "Yes" if det.get("manufacturer_dealer") else "No"),
             ("Straight-line basis (operating leases)", "Daily" if s.get("income_method") == "DAILY" else "Equal monthly (part months by days)"),
             ("Classification override", det.get("classification_override") or "None"), ("Override rationale", det.get("override_rationale") or "")]
    if dep:
        items += [("Security deposit received", _num(dep.get("amount_received"))), ("Deposit refund date", parse_date(dep.get("refund_date"))),
                  ("Deposit market rate %", _num(dep.get("market_rate_pct")))]
    for i, (k, v) in enumerate(items, start=1):
        wi.cell(i, 1, k).font = BOLD
        c = wi.cell(i, 2, v)
        if isinstance(v, date):
            c.number_format = "DD-MMM-YYYY"
        elif isinstance(v, float):
            c.number_format = MONEY_FMT
    r = len(items) + 2
    wi.cell(r, 1, "Lease term (Ind AS 116 App. A, 18–21)").font = BOLD
    for line in term.get("explanation", []):
        r += 1
        wi.cell(r, 1, line)
    r += 2
    wi.cell(r, 1, "Lease payments — inclusion (Ind AS 116.70)").font = BOLD
    r += 1
    _hdr(wi, r, ["Date", "Amount", "Category", "Included", "Reason"])
    for p in s.get("payments") or []:
        r += 1
        wi.cell(r, 1, parse_date(p["date"])).number_format = "DD-MMM-YYYY"
        wi.cell(r, 2, _num(p.get("lease_amount"))).number_format = MONEY_FMT
        wi.cell(r, 3, p.get("category"))
        wi.cell(r, 4, "Yes" if p.get("included") else "No")
        wi.cell(r, 5, p.get("inclusion_reason"))
    wi.column_dimensions["A"].width = 44
    wi.column_dimensions["B"].width = 20
    wi.column_dimensions["C"].width = 18
    wi.column_dimensions["D"].width = 10
    wi.column_dimensions["E"].width = 90

    # Classification
    wc = wb.create_sheet("Classification")
    wc.cell(1, 1, "Lease classification — Ind AS 116.61–66").font = Font(bold=True, size=12)
    _hdr(wc, 3, ["Ref", "Indicator", "Met?", "Value %", "Evidence / note"])
    for i, ind in enumerate(s.get("indicators") or [], start=4):
        wc.cell(i, 1, ind["code"])
        wc.cell(i, 2, ind["indicator"])
        wc.cell(i, 3, {True: "Yes", False: "No"}.get(ind.get("met"), "Not determinable"))
        wc.cell(i, 4, _num(ind.get("value")))
        wc.cell(i, 5, ind.get("note") or "")
    r = 4 + len(s.get("indicators") or []) + 1
    wc.cell(r, 1, "Indicator-based suggestion").font = BOLD
    wc.cell(r, 2, (s.get("suggested_classification") or "").title())
    wc.cell(r + 1, 1, "Final classification").font = BOLD
    wc.cell(r + 1, 2, f"{(s.get('classification') or '').title()} lease" + (f" (override: {det.get('override_rationale')})" if det.get("classification_override") else ""))
    r += 3
    wc.cell(r, 1, "Term test (63(c))").font = BOLD
    wc.cell(r + 1, 1, "Lease term (months)")
    wc.cell(r + 1, 2, _num(term.get("term_months")))
    wc.cell(r + 2, 1, "Economic life (months)")
    wc.cell(r + 2, 2, _num(det.get("economic_life_months")))
    wc.cell(r + 3, 1, "Term / economic life")
    wc.cell(r + 3, 2, f"=IFERROR(B{r + 1}/B{r + 2},\"n/a\")").number_format = "0.00%"
    wc.cell(r + 5, 1, "PV test (63(d))").font = BOLD
    wc.cell(r + 6, 1, "Fair value of the underlying asset")
    wc.cell(r + 6, 2, _num(det.get("fair_value"))).number_format = MONEY_FMT
    wc.cell(r + 7, 1, "PV of lease payments (sheet 'PV of receipts')")
    wc.cell(r + 7, 2, "='PV of receipts'!F1").number_format = MONEY_FMT
    wc.cell(r + 8, 1, "PV / fair value")
    wc.cell(r + 8, 2, f"=IFERROR(B{r + 7}/B{r + 6},\"n/a\")").number_format = "0.00%"
    wc.cell(r + 10, 1, "Explanation").font = BOLD
    for j, line in enumerate(s.get("explanation") or [], start=r + 11):
        wc.cell(j, 1, line)
    for col, w in zip("ABCDE", (40, 70, 16, 12, 90)):
        wc.column_dimensions[col].width = w

    # PV of receipts
    wp = wb.create_sheet("PV of receipts")
    rb = s.get("rate_basis") or {}
    act = (rb.get("daycount") or "ACT/365F").startswith("ACT")
    wp.cell(1, 1, "Commencement date").font = BOLD
    wp.cell(1, 2, parse_date(rb.get("origin") or lease.get("commencement_date"))).number_format = "DD-MMM-YYYY"
    wp.cell(1, 5, "PV of lease payments").font = BOLD
    wp.cell(2, 1, "Effective annual rate (R)").font = BOLD
    wp.cell(2, 2, _num(rb.get("effective_annual_rate"))).number_format = "0.0000000000"
    wp.cell(3, 1, "Method").font = BOLD
    wp.cell(3, 2, "Actual/365 (XNPV-consistent)" if act else "Months/12")
    _hdr(wp, 5, ["#", "Date", "Amount", "Years (t)", "Discount factor = 1/(1+R)^t", "Present value", "Kind", "Engine PV (check)"])
    lines = s.get("pv_lines") or []
    if not lines and rb.get("effective_annual_rate"):
        comm = parse_date(rb.get("origin"))
        lines = [{"date": p["date"], "amount": p["lease_amount"], "years": None, "present_value": None, "kind": "Lease payment"}
                 for p in (s.get("payments") or []) if p.get("included") and parse_date(p["date"]) >= comm]
    i = 5
    for i, pv in enumerate(lines, start=6):
        wp.cell(i, 1, i - 5)
        wp.cell(i, 2, parse_date(pv["date"])).number_format = "DD-MMM-YYYY"
        wp.cell(i, 3, _num(pv["amount"])).number_format = MONEY_FMT
        if act or pv.get("years") is None:
            wp.cell(i, 4, f"=MAX(0,(B{i}-$B$1)/365)").number_format = "0.000000"
        else:
            wp.cell(i, 4, _num(pv["years"])).number_format = "0.000000"
        wp.cell(i, 5, f"=1/(1+$B$2)^D{i}").number_format = "0.0000000000"
        wp.cell(i, 6, f"=C{i}*E{i}").number_format = MONEY_FMT
        wp.cell(i, 7, pv.get("kind"))
        if pv.get("present_value") is not None:
            wp.cell(i, 8, _num(pv["present_value"])).number_format = MONEY_FMT
    last = i
    if last >= 6:
        wp.cell(last + 1, 2, "Total").font = BOLD
        wp.cell(last + 1, 3, f"=SUM(C6:C{last})").number_format = MONEY_FMT
        wp.cell(last + 1, 6, f"=SUM(F6:F{last})").number_format = MONEY_FMT
        wp.cell(last + 1, 8, f"=SUM(H6:H{last})").number_format = MONEY_FMT
        day1 = _num(s.get("receivable_at_commencement")) or 0
        pv_lp_formula = f"=SUMIF(G6:G{last},\"Lease payment\",F6:F{last})" + (f"+{day1}" if day1 else "")
        wp.cell(1, 6, pv_lp_formula).number_format = MONEY_FMT
        if day1:
            wp.cell(1, 7, "incl. lease payments received at commencement (not discounted, not part of the net investment)")
        if fin:
            wp.cell(last + 2, 2, "Net investment at commencement (engine)").font = BOLD
            wp.cell(last + 2, 6, _num(s.get("net_investment"))).number_format = MONEY_FMT
            wp.cell(last + 3, 2, "Difference (recomputed − engine)").font = BOLD
            wp.cell(last + 3, 6, f"=ROUND(F{last + 1}-F{last + 2},2)").number_format = MONEY_FMT
    else:
        wp.cell(1, 6, "n/a")
        wp.cell(6, 2, "Rate implicit in the lease not determinable — PV test concluded on undiscounted amounts (see 'Classification').")
    for col, w in zip("ABCDEFGH", (5, 14, 16, 12, 24, 18, 26, 18)):
        wp.column_dimensions[col].width = w

    # Schedule with checks
    wsch = wb.create_sheet("Schedule")
    srows = s.get("rows") or []
    cols = lessor_schedule_columns(srows)
    keys = [c["key"] for c in cols]
    _hdr(wsch, 1, [c["label"] for c in cols] + ["Check — net investment", "Check — accrued income"])
    L = {k: get_column_letter(j) for j, k in enumerate(keys, start=1)}
    for i2, pr in enumerate(srows, start=2):
        for j, cdef in enumerate(cols, start=1):
            k = cdef["key"]
            v = parse_date(pr[k]) if k == "period_end" else (str(pr.get(k) or "").title() if k == "classification" else _num(pr.get(k)))
            cell = wsch.cell(i2, j, v)
            cell.number_format = "DD-MMM-YYYY" if k == "period_end" else MONEY_FMT
        def ref_(k):
            return f"{L[k]}{i2}" if k in L else "0"
        wsch.cell(i2, len(cols) + 1, f"=ROUND({ref_('ni_open')}+{ref_('ni_additions')}+{ref_('finance_income')}-{ref_('ni_receipts')}"
                                     f"-{ref_('ni_residual_returned')}+{ref_('ni_remeasurement')}-{ref_('ni_derecognised')}-{ref_('ni_close')},2)"
                  ).number_format = MONEY_FMT
        prev_acc = f"{L['accrued_close']}{i2 - 1}" if ("accrued_close" in L and i2 > 2) else "0"
        wsch.cell(i2, len(cols) + 2, f"=ROUND({prev_acc}+{ref_('lease_income')}-{ref_('lease_payments_due')}+{ref_('accrued_adjustment')}"
                                     f"-{ref_('accrued_close')},2)").number_format = MONEY_FMT
    tr = len(srows) + 2
    wsch.cell(tr, 1, "Total").font = BOLD
    for j, cdef in enumerate(cols, start=1):
        if cdef["type"] == "money" and cdef["key"] in ("ni_additions", "finance_income", "ni_receipts", "ni_residual_returned",
                                                         "ni_remeasurement", "ni_derecognised", "lease_income", "lease_payments_due",
                                                         "variable_income", "non_lease_income", "receipts", "idc_amortisation",
                                                         "dep_unwinding", "ecl_charge", "gain_loss", "accrued_adjustment"):
            col = get_column_letter(j)
            wsch.cell(tr, j, f"=SUM({col}2:{col}{tr - 1})").number_format = MONEY_FMT
    for j in range(1, len(cols) + 3):
        wsch.column_dimensions[get_column_letter(j)].width = 15
    wsch.freeze_panes = "B2"

    # Deposit
    if dep:
        wd = wb.create_sheet("Deposit")
        wd.cell(1, 1, "Security deposit received — Ind AS 109").font = Font(bold=True, size=12)
        vals = [("Amount received", _num(dep.get("amount_received")), MONEY_FMT),
                ("Recognised on", parse_date(dep.get("receipt_date")), "DD-MMM-YYYY"),
                ("Refund date", parse_date(dep.get("refund_date")), "DD-MMM-YYYY"),
                ("Amount refundable", _num(dep.get("refund_amount")), MONEY_FMT),
                ("Market rate % p.a.", _num(dep.get("market_rate_pct")), "0.0000")]
        for i2, (k, v, fmt) in enumerate(vals, start=3):
            wd.cell(i2, 1, k).font = BOLD
            wd.cell(i2, 2, v).number_format = fmt
        wd.cell(8, 1, "Years to refund (t)").font = BOLD
        if act:
            wd.cell(8, 2, "=(B5-B4)/365").number_format = "0.000000"
        else:
            from ..engine.calendar_utils import months_between_frac
            wd.cell(8, 2, float(months_between_frac(parse_date(dep.get("receipt_date")), parse_date(dep.get("refund_date"))) / 12)
                    ).number_format = "0.000000"
        wd.cell(9, 1, "Fair value = refundable / (1 + r)^t").font = BOLD
        wd.cell(9, 2, "=B6/(1+B7/100)^B8").number_format = MONEY_FMT
        wd.cell(10, 1, "Fair value per engine").font = BOLD
        wd.cell(10, 2, _num(dep.get("initial_fair_value"))).number_format = MONEY_FMT
        wd.cell(11, 1, "Lease payment received in advance = amount − fair value").font = BOLD
        wd.cell(11, 2, "=B3-B10").number_format = MONEY_FMT
        wd.cell(12, 1, "Difference (recomputed − engine fair value)").font = BOLD
        wd.cell(12, 2, "=ROUND(B9-B10,2)").number_format = MONEY_FMT
        _hdr(wd, 14, ["Period end", "Opening", "Recognised", "Unwinding (finance cost)", "Refund", "Closing", "Check"])
        for i2, pr in enumerate([x for x in srows if D(x.get("dep_open") or 0) or D(x.get("dep_close") or 0) or D(x.get("dep_additions") or 0)
                                 or D(x.get("dep_refund") or 0)], start=15):
            wd.cell(i2, 1, parse_date(pr["period_end"])).number_format = "DD-MMM-YYYY"
            for j, k in enumerate(("dep_open", "dep_additions", "dep_unwinding", "dep_refund", "dep_close"), start=2):
                wd.cell(i2, j, _num(pr.get(k))).number_format = MONEY_FMT
            wd.cell(i2, 7, f"=ROUND(B{i2}+C{i2}+D{i2}-E{i2}-F{i2},2)").number_format = MONEY_FMT
        wd.column_dimensions["A"].width = 52
        for col in "BCDEFG":
            wd.column_dimensions[col].width = 18

    # Maturity & reconciliation
    wm = wb.create_sheet("Maturity & reconciliation")
    pos = position_at(s, as_of)
    wm.cell(1, 1, f"Position at {as_of:%d-%b-%Y} — classification: {pos['classification'].title()}").font = Font(bold=True, size=12)
    _hdr(wm, 3, ["Undiscounted lease payments to be received", "Amount"])
    r = 3
    for m in pos["maturity"]:
        r += 1
        wm.cell(r, 1, m["bucket"])
        wm.cell(r, 2, _num(m["amount"])).number_format = MONEY_FMT
    r += 1
    wm.cell(r, 1, "Total undiscounted").font = BOLD
    wm.cell(r, 2, f"=SUM(B4:B{r - 1})" if r > 4 else 0).number_format = MONEY_FMT
    tot_row = r
    rc = pos.get("reconciliation")
    if rc:
        r += 2
        wm.cell(r, 1, "Reconciliation to the net investment (Ind AS 116.94)").font = BOLD
        wm.cell(r + 1, 1, "Undiscounted lease payments receivable")
        wm.cell(r + 1, 2, f"=B{tot_row}").number_format = MONEY_FMT
        wm.cell(r + 2, 1, "Less: unearned finance income")
        wm.cell(r + 2, 2, -_num(rc["unearned_finance_income"])).number_format = MONEY_FMT
        wm.cell(r + 3, 1, "Present value of lease payments receivable")
        wm.cell(r + 3, 2, f"=B{r + 1}+B{r + 2}").number_format = MONEY_FMT
        wm.cell(r + 4, 1, "Add: discounted unguaranteed residual value")
        wm.cell(r + 4, 2, _num(rc["discounted_unguaranteed_residual"])).number_format = MONEY_FMT
        wm.cell(r + 5, 1, "Net investment in the lease").font = BOLD
        wm.cell(r + 5, 2, f"=B{r + 3}+B{r + 4}").number_format = MONEY_FMT
        wm.cell(r + 6, 1, "Net investment per schedule")
        wm.cell(r + 6, 2, _num(rc.get("ni_per_schedule"))).number_format = MONEY_FMT
        wm.cell(r + 7, 1, "Difference").font = BOLD
        wm.cell(r + 7, 2, f"=IF(B{r + 6}=\"\",\"\",ROUND(B{r + 5}-B{r + 6},2))").number_format = MONEY_FMT
        wm.cell(r + 8, 1, "Less: loss allowance (Ind AS 109)")
        wm.cell(r + 8, 2, -_num(pos["loss_allowance"])).number_format = MONEY_FMT
        wm.cell(r + 9, 1, "Net investment net of loss allowance").font = BOLD
        wm.cell(r + 9, 2, f"=B{r + 5}+B{r + 8}").number_format = MONEY_FMT
        wm.cell(r + 10, 1, "Current / non-current (principal recovered within 12 months)")
        wm.cell(r + 10, 2, f"{fmt_money(pos['ni_current'])} / {fmt_money(pos['ni_noncurrent'])}")
    else:
        r += 2
        wm.cell(r, 1, "Accrued / (deferred) lease income (straight-lining)")
        wm.cell(r, 2, _num(pos["accrued_lease_income"])).number_format = MONEY_FMT
    if dep:
        wm.cell(r + 12, 1, "Security deposit received — carrying amount")
        wm.cell(r + 12, 2, _num(pos["deposit_carrying"])).number_format = MONEY_FMT
    wm.column_dimensions["A"].width = 60
    wm.column_dimensions["B"].width = 20

    # Journals
    wj = wb.create_sheet("Journals")
    _hdr(wj, 1, ["Date", "Event", "Narration", "GL role", "Debit", "Credit"])
    r = 1
    for p in s.get("postings", []):
        for l in p["lines"]:
            r += 1
            wj.cell(r, 1, parse_date(p["date"])).number_format = "DD-MMM-YYYY"
            wj.cell(r, 2, EVENT_LABELS.get(p["event"], p["event"]))
            wj.cell(r, 3, p["narration"])
            wj.cell(r, 4, l["role"])
            wj.cell(r, 5, _num(l["debit"]) or None).number_format = MONEY_FMT
            wj.cell(r, 6, _num(l["credit"]) or None).number_format = MONEY_FMT
    wj.cell(r + 1, 4, "Total").font = BOLD
    wj.cell(r + 1, 5, f"=SUM(E2:E{r})").number_format = MONEY_FMT
    wj.cell(r + 1, 6, f"=SUM(F2:F{r})").number_format = MONEY_FMT
    for col, w in zip("ABCDEF", (13, 24, 80, 30, 16, 16)):
        wj.column_dimensions[col].width = w

    # Judgments
    wg = wb.create_sheet("Judgments")
    _hdr(wg, 1, ["Flag / issue", "Detail", "Reference", "Assumptions / severity"])
    r = 1
    for f in s.get("flags", []):
        r += 1
        wg.cell(r, 1, f.get("title"))
        wg.cell(r, 2, f.get("detail"))
        wg.cell(r, 3, f.get("reference"))
        wg.cell(r, 4, "; ".join(f.get("assumptions") or []))
    for iss in s.get("issues", []):
        r += 1
        wg.cell(r, 1, iss.get("code"))
        wg.cell(r, 2, iss.get("message"))
        wg.cell(r, 3, iss.get("reference"))
        wg.cell(r, 4, iss.get("severity"))
    for col, w in zip("ABCD", (34, 90, 26, 50)):
        wg.column_dimensions[col].width = w

    if s.get("events"):
        we = wb.create_sheet("Events")
        _hdr(we, 1, ["Event", "Effective", "Description", "Step", "Amount", "Explanation", "Reference"])
        r = 1
        for ev in s["events"]:
            for st in ev.get("steps") or [["—", None, ""]]:
                r += 1
                we.cell(r, 1, ev["type"].replace("_", " ").title() + (f" — {ev['nature'].replace('_', ' ').lower()}" if ev.get("nature") else ""))
                we.cell(r, 2, parse_date(ev["effective_date"])).number_format = "DD-MMM-YYYY"
                we.cell(r, 3, ev.get("description"))
                we.cell(r, 4, st[0])
                we.cell(r, 5, _num(st[1])).number_format = MONEY_FMT
                we.cell(r, 6, st[2])
                we.cell(r, 7, ev.get("reference"))
        for col, w in zip("ABCDEFG", (34, 13, 50, 50, 16, 70, 28)):
            we.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def lessor_memo_docx(lease: dict, run: dict, company: str, as_of: date) -> bytes:
    """Accounting memo for a lease in which the entity is the LESSOR (Ind AS 116.61–97; Ind AS 109)."""
    import docx
    from docx.shared import Pt

    from ..engine.lessor import position_at

    s = run["summary"]
    det = lease.get("lessor_details") or {}
    fin = s.get("classification") == "FINANCE"
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    d.add_heading(f"Lease accounting memo (lessor) — {lease['lease_code']}", 0)
    d.add_paragraph(f"Entity: {lease.get('entity')}    Company: {company}    Framework: Ind AS 116 (lessor)    Prepared: {date.today():%d-%b-%Y}")
    d.add_heading("1. Background and facts", 1)
    d.add_paragraph(f"{lease['description']}. The entity is the lessor. Lessee: {lease.get('lessor') or '—'}. Underlying asset: "
                    f"{lease.get('asset_description') or lease.get('asset_class') or '—'}. Location: {lease.get('location') or '—'}. "
                    f"Contract date {lease.get('contract_date') or '—'}; commencement {lease.get('commencement_date')}; contractual end "
                    f"{lease.get('contract_end')}.")
    d.add_heading("2. Does the contract contain a lease? (Ind AS 116.9–11, B9–B31)", 1)
    a = lease.get("assessment") or {}
    d.add_paragraph(a.get("conclusion") or "The contract conveys to the lessee the right to control the use of an identified asset for a "
                    "period of time in exchange for consideration. [Preparer to confirm the identified asset, substitution rights and the "
                    "lessee's right to direct the use.]")
    d.add_heading("3. Lease term (Ind AS 116 App. A, 18–21)", 1)
    for line in (s.get("term") or {}).get("explanation", []):
        d.add_paragraph(line, style="List Bullet")
    d.add_heading("4. Classification (Ind AS 116.61–66)", 1)
    t = d.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(("Ref", "Indicator", "Met?", "Evidence")):
        t.rows[0].cells[i].text = h
    for ind in s.get("indicators") or []:
        row = t.add_row().cells
        row[0].text, row[1].text = ind["code"], ind["indicator"]
        row[2].text = {True: "Yes", False: "No"}.get(ind.get("met"), "Not determinable")
        row[3].text = ind.get("note") or ("" if ind.get("value") is None else f"{ind['value']}%")
    d.add_paragraph("")
    for line in s.get("explanation") or []:
        d.add_paragraph(line, style="List Bullet")
    cls_word = (s.get("classification") or "").lower()
    concl = f"Conclusion: the lease is classified as {'an' if cls_word[:1] in 'aeiou' else 'a'} {cls_word} lease"
    if det.get("classification_override"):
        concl += f" (classification recorded by management — {det.get('override_rationale')})"
    d.add_paragraph(concl + ". Classification is made at the inception date and reassessed only on a lease modification (para 66).")
    d.add_heading("5. Lease payments (Ind AS 116.70) and non-lease components (para 17)", 1)
    pays = s.get("payments") or []
    excl = [p for p in pays if not p.get("included")]
    d.add_paragraph(f"{len(pays) - len(excl)} payment line(s) are lease payments. " + (
        f"{len(excl)} line(s) are excluded: " + "; ".join(sorted({p.get('inclusion_reason') or '' for p in excl}))[:900] if excl else
        "No payment lines are excluded."))
    d.add_paragraph("Variable payments not linked to an index or rate are recognised as income when earned and disclosed separately "
                    "(paras 81, 90); non-lease components (services / CAM) are revenue under Ind AS 115 (para 17).")
    d.add_heading("6. Measurement", 1)
    if fin:
        d.add_paragraph(f"Rate implicit in the lease: {q(D(s.get('implicit_rate_pct') or 0), 4)}% p.a. ({s.get('rate_source')}).")
        t = d.add_table(rows=1, cols=2)
        t.style = "Light Grid Accent 1"
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Item", "Amount"
        for label, key in (("Net investment at commencement (para 68)", "net_investment"), ("Gross investment", "gross_investment"),
                           ("Unearned finance income", "unearned_finance_income"), ("PV of unguaranteed residual value", "pv_unguaranteed_residual"),
                           ("Lease payments received at / before commencement", "receivable_at_commencement"),
                           ("Selling profit / (loss) at commencement — manufacturer / dealer (para 71)", "selling_profit"),
                           ("Gain / (loss) on derecognition of the asset — other lessors (Ind AS 16.68, 71)", "derecognition_gain")):
            row = t.add_row().cells
            row[0].text, row[1].text = label, fmt_money(s.get(key) or 0)
        if det.get("manufacturer_dealer"):
            d.add_paragraph(f"Manufacturer / dealer: revenue {fmt_money(s.get('revenue') or 0)}; cost of sale {fmt_money(s.get('cost_of_sale') or 0)} "
                            "(para 71); costs of obtaining the lease expensed at commencement (para 74).")
        d.add_paragraph("At commencement the underlying asset is derecognised and the net investment is presented as a receivable "
                        "(para 67). Finance income is recognised at a constant periodic rate on the net investment (paras 75–76); the "
                        "net investment is subject to the derecognition and impairment requirements of Ind AS 109 (para 77).")
    else:
        tot = s.get("totals") or {}
        d.add_paragraph(f"Lease payments are recognised as income on a straight-line basis over the lease term "
                        f"({'daily' if s.get('income_method') == 'DAILY' else 'equal monthly amounts, part months pro-rata by days'}) — "
                        f"para 81. Total operating lease income over the term: {fmt_money(tot.get('total_lease_income') or 0)}. "
                        "Ind AS 116 does not carry forward the Ind AS 17 exception for escalations in line with expected inflation — "
                        "escalations are straight-lined.")
        d.add_paragraph("The underlying asset remains in property, plant and equipment / investment property of the lessor, is depreciated "
                        "in line with the lessor's normal policy (para 84) and tested for impairment under Ind AS 36 (para 85). Initial direct "
                        "costs are added to the carrying amount of the asset and expensed over the lease term on the same basis as the "
                        "lease income (para 83).")
    if s.get("deposit"):
        dp = s["deposit"]
        d.add_heading("7. Security deposit received (Ind AS 109)", 1)
        d.add_paragraph(f"Deposit {fmt_money(dp['amount_received'])} refundable on {dp['refund_date']}; fair value at the market rate "
                        f"{dp['market_rate_pct']}% = {fmt_money(dp['initial_fair_value'])} (financial liability at amortised cost; unwinding is "
                        f"a finance cost). The excess {fmt_money(dp['lease_payment_element'])} is a lease payment received in advance "
                        f"({'part of the lease payments of the finance lease' if fin else 'recognised as lease income over the lease term'}).")
    d.add_heading("8. Subsequent events", 1)
    evs = s.get("events") or []
    if not evs:
        d.add_paragraph("None recorded.")
    ev_names = {"MODIFICATION": "Modification", "TERMINATION": "Early termination", "UGR_REVISION": "Revision of the unguaranteed residual value",
                "ECL": "Expected credit loss allowance"}
    for ev in evs:
        d.add_paragraph(f"{ev_names.get(ev['type'], ev['type'].title())} effective {ev['effective_date']}: {ev.get('description') or ''} — "
                        f"{ev.get('balance_label') or 'balance'} {fmt_money(ev.get('balance_before') or 0)} → {fmt_money(ev.get('balance_after') or 0)}; "
                        f"gain / (loss) {fmt_money(ev.get('gain_loss') or 0)}. [{ev.get('reference')}]", style="List Bullet")
    d.add_heading("9. Position at the reporting date", 1)
    pos = position_at(s, as_of)
    d.add_paragraph(f"At {as_of:%d-%b-%Y}: classification {pos['classification'].title()}; "
                    + (f"net investment {fmt_money(pos['ni_close'])} (current {fmt_money(pos['ni_current'])}, non-current "
                       f"{fmt_money(pos['ni_noncurrent'])}); loss allowance {fmt_money(pos['loss_allowance'])}." if pos["classification"] == "FINANCE"
                       else f"accrued / (deferred) lease income {fmt_money(pos['accrued_lease_income'])}.")
                    + (f" Security deposit received carried at {fmt_money(pos['deposit_carrying'])}." if s.get("deposit") else ""))
    t = d.add_table(rows=1, cols=2)
    t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Undiscounted lease payments to be received (paras 94 / 97)", "Amount"
    for m in pos["maturity"]:
        row = t.add_row().cells
        row[0].text, row[1].text = m["bucket"], fmt_money(m["amount"])
    rc = pos.get("reconciliation")
    if rc:
        d.add_paragraph(f"Reconciliation (para 94): undiscounted {fmt_money(rc['undiscounted_lease_payments'])} less unearned finance income "
                        f"{fmt_money(rc['unearned_finance_income'])} plus discounted unguaranteed residual "
                        f"{fmt_money(rc['discounted_unguaranteed_residual'])} = net investment {fmt_money(rc['net_investment'])}.")
    d.add_heading("10. Judgments and reviewer attention", 1)
    for f in s.get("flags", []):
        d.add_paragraph(f"{f['title']} — {f['detail']} ({f.get('reference', '')})", style="List Bullet")
    d.add_heading("11. Presentation and disclosure", 1)
    if fin:
        d.add_paragraph("Net investment in finance leases presented as a financial asset (receivable), split into current and non-current "
                        "(Schedule III, Division II); finance income in revenue / other income; loss allowance per Ind AS 109.")
    else:
        d.add_paragraph("Underlying asset presented according to its nature (para 88); accrued lease income (straight-lining) presented as an "
                        "asset (or deferred income as a liability); security deposit received as a financial liability.")
    d.add_paragraph("Disclosures: income by type (para 90, tabular — para 91); qualitative and quantitative information on leasing activities "
                    "and residual-asset risk management (para 92); " + ("significant changes in the net investment (para 93); maturity analysis "
                    "and reconciliation to the net investment (para 94)." if fin else "Ind AS 16 / 40 disclosures for assets subject to "
                    "operating leases (paras 95–96); maturity analysis of lease payments (para 97)."))
    d.add_heading("12. Approval", 1)
    d.add_paragraph(f"Calculation run {run.get('run_no')} — status {run.get('status')}. Prepared by: ________  Reviewed by: ________  "
                    "Approved by: ________")
    d.add_paragraph("References are to the notified Ind AS 116 text; paragraph descriptions in this memo are summaries.").italic = True
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
