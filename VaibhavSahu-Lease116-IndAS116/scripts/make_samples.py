"""Generate fictional sample agreements used for demos and automated tests.

All parties, addresses and figures are fictional placeholders (ABC / XYZ style).
Outputs (samples/agreements/):
  1. Sample_Leave_and_License_Office.pdf        — digital PDF with e-stamp page
  2. Sample_Warehouse_Lease_Deed_SCANNED.pdf     — image-only PDF (tilted, noisy, blue stamp) for OCR/CV
  3. Sample_Equipment_Lease.docx                 — Word equipment lease with purchase option
  4. Sample_Retail_Lease_Revenue_Share.txt       — retail lease with minimum guarantee + revenue share
"""
from __future__ import annotations

import io
import os
import random
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "samples" / "agreements"

OFFICE_PAGES = [
    """e-Stamp
Certificate No. : IN-XX00000000000000X (SAMPLE)
Certificate Issued Date : 15-Mar-2025
Account Reference : SAMPLE ONLINE
Description of Document : Article 36A Leave and Licence
Consideration Price (Rs.) : 2,70,00,000
First Party : ABC REALTY PRIVATE LIMITED
Second Party : XYZ TECHNOLOGIES PRIVATE LIMITED
Stamp Duty Paid By : XYZ TECHNOLOGIES PRIVATE LIMITED
Stamp Duty Amount(Rs.) : 1,18,500
(One Lakh Eighteen Thousand Five Hundred only)
THIS IS A FICTIONAL SAMPLE FOR SOFTWARE TESTING""",
    """LEAVE AND LICENSE AGREEMENT

This Leave and License Agreement is made and entered into at Mumbai on this 15th day of March, 2025.

BETWEEN

ABC Realty Private Limited, a company incorporated under the Companies Act, 2013, having its registered office at 1 Sample Road, Mumbai 400001, hereinafter referred to as the "LICENSOR" (which expression shall unless repugnant to the context include its successors and assigns) of the ONE PART;

AND

XYZ Technologies Private Limited, a company incorporated under the Companies Act, 2013, having its registered office at 2 Demo Street, Pune 411001, hereinafter referred to as the "LICENSEE" (which expression shall unless repugnant to the context include its successors and permitted assigns) of the OTHER PART.

WHEREAS the Licensor is the owner of office premises bearing Unit No. 501, 5th Floor, Sample Business Park, situated at Plot 10, Andheri East, Mumbai 400093 admeasuring 9,000 sq. ft. carpet area (the "Licensed Premises"), more particularly described in the Schedule hereunder written.

NOW THIS AGREEMENT WITNESSETH AS FOLLOWS:

1. LICENSE PERIOD
1.1 The Licensor hereby grants to the Licensee a leave and license to use the Licensed Premises for a period of 60 (sixty) months commencing from 1st April, 2025 (the "Commencement Date") and ending on 31st March, 2030.
1.2 The Licensed Premises shall be handed over to the Licensee on the Commencement Date for carrying out fit-out works.

2. LOCK-IN PERIOD
2.1 Both the parties shall be bound by a lock-in period of 36 (thirty six) months from the Commencement Date, during which neither party shall terminate this Agreement.

3. LICENSE FEE
3.1 The Licensee shall pay to the Licensor a monthly license fee of Rs. 4,50,000/- (Rupees Four Lakh Fifty Thousand Only) per month, being Rs. 50 per sq. ft. of carpet area, exclusive of GST.
3.2 The license fee shall be payable in advance on or before the 7th day of each English calendar month.
3.3 The Licensee shall be entitled to a rent-free fit-out period of 60 (sixty) days from the Commencement Date, during which no license fee shall be payable.
3.4 The license fee shall be escalated by 5% (five percent) every 12 (twelve) months on the last paid license fee.
""",
    """4. SECURITY DEPOSIT
4.1 The Licensee shall pay to the Licensor an interest free refundable security deposit of Rs. 27,00,000/- (Rupees Twenty Seven Lakh Only), equivalent to six months' license fee. The security deposit shall not carry any interest and shall be refunded on expiry or earlier termination against handover of the Licensed Premises.

5. MAINTENANCE CHARGES AND TAXES
5.1 The Licensee shall pay common area maintenance charges of Rs. 55,000/- (Rupees Fifty Five Thousand Only) per month directly to the facility manager.
5.2 GST as applicable shall be payable extra by the Licensee.

6. STAMP DUTY AND REGISTRATION
6.1 The stamp duty and registration charges on this Agreement shall be borne and paid by the Licensee.
6.2 The Licensee has paid brokerage of Rs. 4,50,000/- to the property consultant.

7. RENEWAL
7.1 This Agreement may be renewed for a further period of 60 (sixty) months by mutual consent of the parties on such terms as may be mutually agreed, subject to an increase of 15% on the last paid license fee.

8. TERMINATION
8.1 After the expiry of the lock-in period, the Licensee shall be entitled to terminate this Agreement by giving 3 (three) months' prior written notice to the Licensor.
8.2 The Licensor may terminate this Agreement only in the event of breach by the Licensee of any of its obligations, including non-payment of license fee for two consecutive months, after giving 30 days' notice to remedy the breach.

9. HANDOVER AND REINSTATEMENT
9.1 On expiry or earlier termination the Licensee shall remove all fit-outs and restore the Licensed Premises to its original condition, reasonable wear and tear excepted.

10. SUB-LETTING
10.1 The Licensee shall not sub-let, sub-license or part with possession of the Licensed Premises without the prior written consent of the Licensor.
""",
    """SCHEDULE
(Description of the Licensed Premises)
Office premises bearing Unit No. 501, 5th Floor, Sample Business Park, Plot 10, Andheri East, Mumbai 400093, admeasuring 9,000 sq. ft. carpet area together with 10 car parking spaces.

IN WITNESS WHEREOF the parties have set their hands on the day and year first hereinabove written.

For ABC Realty Private Limited (LICENSOR)            For XYZ Technologies Private Limited (LICENSEE)

Authorised Signatory                                   Authorised Signatory
""",
]

WAREHOUSE_TEXT = """LEASE DEED

This Lease Deed is executed at Bhiwandi on this 2nd day of January, 2025
BETWEEN PQR Logistics Parks LLP, a limited liability partnership having its office at 5 Sample Nagar, Thane,
hereinafter referred to as the "Lessor" AND LMN Retail Private Limited, a company incorporated under the
Companies Act, 2013 having its registered office at 9 Example Lane, Mumbai, hereinafter referred to as the "Lessee".

1. The Lessor hereby leases to the Lessee the warehouse bearing Shed No. B-2 situated at Sample Industrial
Estate, Bhiwandi admeasuring 40,000 sq. ft. for a period of 9 (nine) years commencing from 1st February, 2025.

2. The lease period shall comprise three blocks of three years each. The Lessee shall have the option to renew
the lease for a further period of 5 (five) years at the sole option of the Lessee on the same terms.

3. The Lessee shall pay a monthly rent of Rs. 12,00,000/- (Rupees Twelve Lakh Only) in advance on or before
the 5th day of each month.

4. The rent shall be increased by 15% (fifteen percent) after every 3 (three) years on the last paid rent.

5. The Lessee may terminate this lease after the first three years by giving 6 (six) months' notice.

6. The Lessee shall deposit an interest-free security deposit of Rs. 72,00,000/- (Rupees Seventy Two Lakh Only).

7. On expiry the Lessee shall hand over vacant possession in the same condition as at the commencement.
"""

EQUIPMENT_PARAS = [
    "EQUIPMENT LEASE AGREEMENT",
    "This Equipment Lease Agreement is made on 20th June, 2025 between DEF Leasing Limited, a company incorporated under "
    "the Companies Act, 2013, hereinafter referred to as the \"Lessor\" and GHI Auto Components Private Limited, a company "
    "incorporated under the Companies Act, 2013, hereinafter referred to as the \"Lessee\".",
    "1. EQUIPMENT: Four CNC vertical machining centres (the \"Equipment\") as described in Annexure A.",
    "2. TERM: The lease shall be for a period of 36 (thirty six) months commencing from 1st July, 2025 (the Commencement Date).",
    "3. LEASE RENTAL: The Lessee shall pay a quarterly lease rental of Rs. 6,00,000/- (Rupees Six Lakh Only) per quarter "
    "in arrears at the end of each quarter.",
    "4. PURCHASE OPTION: At the end of the lease term the Lessee shall have the option to purchase the Equipment at a price "
    "of Rs. 5,00,000/- (Rupees Five Lakh Only).",
    "5. The Lessee shall not sub-lease the Equipment.",
    "6. The Lessor may substitute the Equipment with equipment of equivalent specification at any time at its cost.",
]

RETAIL_TEXT = """LEASE AGREEMENT (Retail store — fictional sample)
This Lease Agreement is made on 10th August 2025 between STU Malls Private Limited, hereinafter referred to as the "Lessor"
and VWX Apparel Private Limited, hereinafter referred to as the "Lessee".
1. The Lessor grants the Lessee the retail unit G-14 situated at Sample Mall, Bengaluru admeasuring 2,500 sq. ft. for a
term of 9 years commencing on 1st September, 2025, with a lock-in period of 3 years for the Lessee.
2. Rent: the Lessee shall pay the higher of (a) minimum guaranteed rent of Rs. 3,00,000 per month and (b) revenue share of
12% of the gross sales of the store. The minimum guaranteed rent shall be escalated by 15% every three years.
3. Common area maintenance charges of Rs. 90 per sq. ft. per month shall be payable in addition.
4. Security deposit: Rs. 18,00,000 (Rupees Eighteen Lakh Only), interest free.
5. Either party may terminate this agreement by giving 6 months' notice after the lock-in period.
"""


def make_office_pdf(path: Path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    W, H = A4
    for i, page in enumerate(OFFICE_PAGES):
        y = H - 22 * mm
        if i == 0:
            c.setFillColorRGB(0.95, 0.97, 1)
            c.rect(15 * mm, H - 150 * mm, W - 30 * mm, 135 * mm, fill=1, stroke=0)
            c.setFillColorRGB(0, 0, 0)
        for raw in page.split("\n"):
            lines = _wrap(raw, 95)
            for ln in lines:
                if ln.isupper() and len(ln) > 3:
                    c.setFont("Helvetica-Bold", 11)
                else:
                    c.setFont("Helvetica", 10)
                c.drawString(18 * mm, y, ln)
                y -= 5.2 * mm
                if y < 20 * mm:
                    c.showPage()
                    y = H - 22 * mm
        c.setFont("Helvetica-Oblique", 7)
        c.drawString(18 * mm, 10 * mm, f"Fictional sample agreement for Lease116 testing — page {i + 1}")
        c.showPage()
    c.save()


def _wrap(text: str, width: int) -> list[str]:
    import textwrap
    if not text.strip():
        return [""]
    return textwrap.wrap(text, width=width) or [""]


def make_scanned_pdf(path: Path):
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    import numpy as np

    W, H = 1654, 2339  # A4 @ 200 dpi
    img = Image.new("RGB", (W, H), (250, 249, 244))
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf", 30)
    bold = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf", 40)
    y = 140
    for raw in WAREHOUSE_TEXT.split("\n"):
        if raw.strip() == "LEASE DEED":
            d.text((W // 2 - 140, y), raw, fill=(20, 20, 20), font=bold)
            y += 70
            continue
        for ln in _wrap(raw, 88):
            d.text((120, y), ln, fill=(25, 25, 25), font=font)
            y += 44
    # blue rubber stamp
    cx, cy, r = 1250, 1950, 150
    d.ellipse((cx - r, cy - r, cx + r, cy + r), outline=(40, 60, 170), width=8)
    d.ellipse((cx - r + 25, cy - r + 25, cx + r - 25, cy + r - 25), outline=(40, 60, 170), width=4)
    sfont = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 26)
    d.text((cx - 95, cy - 20), "SAMPLE SEAL", fill=(40, 60, 170), font=sfont)
    d.text((170, 2050), "For PQR Logistics Parks LLP            For LMN Retail Private Limited", fill=(25, 25, 25), font=font)
    img = img.rotate(1.6, resample=Image.BICUBIC, expand=False, fillcolor=(250, 249, 244))
    arr = np.array(img).astype(np.int16)
    rng = np.random.default_rng(7)
    arr += rng.normal(0, 9, arr.shape).astype(np.int16)
    arr = np.clip(arr, 0, 255).astype(np.uint8)
    img = Image.fromarray(arr).filter(ImageFilter.GaussianBlur(0.6))
    img.save(str(path), "PDF", resolution=200.0)


def make_equipment_docx(path: Path):
    import docx
    doc = docx.Document()
    doc.add_heading(EQUIPMENT_PARAS[0], level=1)
    for p in EQUIPMENT_PARAS[1:]:
        doc.add_paragraph(p)
    t = doc.add_table(rows=3, cols=3)
    for r, row in enumerate([("Annexure A", "Equipment", "Qty"), ("1", "CNC vertical machining centre (sample)", "4"),
                             ("", "Value when new (each)", "Rs. 45,00,000")]):
        for cidx, val in enumerate(row):
            t.cell(r, cidx).text = val
    doc.save(str(path))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    make_office_pdf(OUT / "Sample_Leave_and_License_Office.pdf")
    make_scanned_pdf(OUT / "Sample_Warehouse_Lease_Deed_SCANNED.pdf")
    make_equipment_docx(OUT / "Sample_Equipment_Lease.docx")
    (OUT / "Sample_Retail_Lease_Revenue_Share.txt").write_text(RETAIL_TEXT, encoding="utf-8")
    print("Samples written to", OUT)


if __name__ == "__main__":
    main()
