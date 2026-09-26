"""Builds structured Excel templates with dropdowns and input validation."""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.accounting.gst import CGST_RATES, IGST_RATES, STATE_NAMES
from app.excel.template_spec import META_SHEET, TemplateSpec

DATA_ROWS = 1000  # rows that receive dropdowns / validation
LISTS_SHEET = "Lists"

_REQUIRED_FILL = PatternFill("solid", fgColor="1F3864")
_OPTIONAL_FILL = PatternFill("solid", fgColor="4472C4")
_VOUCHER_LEVEL_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="BFBFBF")


@dataclass
class TemplateLists:
    """Values for dropdowns, normally taken from the cached Tally masters."""

    ledgers: list[str] = field(default_factory=list)
    bank_ledgers: list[str] = field(default_factory=list)
    stock_items: list[str] = field(default_factory=list)
    units: list[str] = field(default_factory=list)


def _fmt_rate(r) -> str:
    return format(r.normalize(), "f")


def build_template(spec: TemplateSpec, lists: TemplateLists, *, company_name: str = "",
                   financial_year: str = "", sample_rows: list[dict[str, Any]] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = spec.sheet_name

    # ---- header row -----------------------------------------------------------
    for idx, col in enumerate(spec.columns, start=1):
        cell = ws.cell(row=1, column=idx, value=col.header)
        cell.fill = _REQUIRED_FILL if col.required else _OPTIONAL_FILL
        cell.font = _VOUCHER_LEVEL_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
        tip = ("Required. " if col.required else "Optional. ") + (col.help or "")
        if spec.group_by and col.level == "voucher":
            tip += " (Voucher-level: can be left blank on continuation rows.)"
        cell.comment = Comment(tip.strip(), "BRMCo")
        ws.column_dimensions[get_column_letter(idx)].width = col.width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(spec.columns))}1"

    # ---- lists sheet --------------------------------------------------------------
    lst = wb.create_sheet(LISTS_SHEET)
    list_columns = {
        "ledgers": sorted(set(lists.ledgers), key=str.lower),
        "bank_ledgers": sorted(set(lists.bank_ledgers or lists.ledgers), key=str.lower),
        "stock_items": sorted(set(lists.stock_items), key=str.lower),
        "units": sorted(set(lists.units), key=str.lower),
        "states": STATE_NAMES,
    }
    list_ranges: dict[str, str] = {}
    for c_idx, (name, values) in enumerate(list_columns.items(), start=1):
        letter = get_column_letter(c_idx)
        lst.cell(row=1, column=c_idx, value=name).font = Font(bold=True)
        for r_idx, v in enumerate(values, start=2):
            lst.cell(row=r_idx, column=c_idx, value=v)
        if values:
            list_ranges[name] = f"={LISTS_SHEET}!${letter}$2:${letter}${len(values) + 1}"
    lst.sheet_state = "hidden"

    inline_lists = {
        "yesno": '"Yes,No"',
        "cgst_rates": '"' + ",".join(_fmt_rate(r) for r in sorted(CGST_RATES)) + '"',
        "igst_rates": '"' + ",".join(_fmt_rate(r) for r in sorted(IGST_RATES)) + '"',
    }

    # ---- validations + formats ---------------------------------------------------------
    last = DATA_ROWS + 1
    for idx, col in enumerate(spec.columns, start=1):
        letter = get_column_letter(idx)
        rng = f"{letter}2:{letter}{last}"
        dv: DataValidation | None = None
        if col.choices:
            dv = DataValidation(type="list", formula1='"' + ",".join(col.choices) + '"', allow_blank=True)
        elif col.list_source in inline_lists:
            dv = DataValidation(type="list", formula1=inline_lists[col.list_source], allow_blank=True)
        elif col.list_source in list_ranges:
            dv = DataValidation(type="list", formula1=list_ranges[col.list_source], allow_blank=True)
            dv.errorStyle = "warning"  # masters may be stale; the validator is the final check
        elif col.type == "date":
            dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
            dv.error = "Enter a valid date (DD-MM-YYYY)."
        elif col.type in ("amount", "qty", "rate", "gst_rate"):
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
            dv.error = "Enter a number (0 or more)."
        elif col.type == "signed":
            dv = DataValidation(type="decimal", operator="between", formula1="-100", formula2="100", allow_blank=True)
            dv.error = "Round off should be a small amount between -100 and 100."
        if dv is not None:
            dv.showErrorMessage = True
            dv.errorTitle = col.header
            ws.add_data_validation(dv)
            dv.add(rng)

        number_format = {"date": "DD-MM-YYYY", "amount": "#,##0.00", "signed": "#,##0.00",
                         "qty": "#,##0.###", "text": "@"}.get(col.type)
        if number_format:
            for r in range(2, min(last, 201) + 1):
                ws.cell(row=r, column=idx).number_format = number_format

    # ---- sample data ------------------------------------------------------------------
    for r_idx, row in enumerate(sample_rows or [], start=2):
        for c_idx, col in enumerate(spec.columns, start=1):
            if col.key in row and row[col.key] is not None:
                ws.cell(row=r_idx, column=c_idx, value=row[col.key])

    # ---- instructions ---------------------------------------------------------------------
    ins = wb.create_sheet("Instructions", 1)
    ins.column_dimensions["A"].width = 28
    ins.column_dimensions["B"].width = 100
    ins["A1"] = f"BRMCo Accounting Hub — {spec.title}"
    ins["A1"].font = Font(bold=True, size=14)
    ins["A2"] = f"Template {spec.template_id} v{spec.version}"
    if company_name:
        ins["A3"] = f"Company: {company_name}   Financial year: {financial_year}"
    line = 5
    for note in (
        f"Enter data on the '{spec.sheet_name}' sheet from row 2. Do not rename, reorder or delete the header row.",
        "Dark blue headers are required; lighter blue headers are optional. Hover a header for help.",
        "Dates: DD-MM-YYYY. Amounts: plain numbers without currency symbols.",
        "Ledger and item names must match Tally exactly. Dropdowns come from the last Tally master sync.",
        *spec.notes,
        "Do not delete the hidden sheets — they identify the template version.",
    ):
        ins.cell(row=line, column=1, value="•")
        ins.cell(row=line, column=2, value=note).alignment = Alignment(wrap_text=True)
        line += 1
    line += 1
    ins.cell(row=line, column=1, value="Column").font = Font(bold=True)
    ins.cell(row=line, column=2, value="Description").font = Font(bold=True)
    for col in spec.columns:
        line += 1
        ins.cell(row=line, column=1, value=col.header + (" *" if col.required else ""))
        ins.cell(row=line, column=2, value=col.help or "")

    # ---- metadata ---------------------------------------------------------------------------
    meta = wb.create_sheet(META_SHEET)
    for r, (k, v) in enumerate((
        ("template_id", spec.template_id),
        ("template_version", spec.version),
        ("voucher_kind", spec.kind.value),
        ("generated_at", datetime.now().isoformat(timespec="seconds")),
        ("company", company_name),
        ("financial_year", financial_year),
    ), start=1):
        meta.cell(row=r, column=1, value=k)
        meta.cell(row=r, column=2, value=v)
    meta.sheet_state = "veryHidden"

    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_filename(spec: TemplateSpec, today: date | None = None) -> str:
    today = today or date.today()
    return f"{spec.template_id.lower()}-v{spec.version}-{today:%Y%m%d}.xlsx"
