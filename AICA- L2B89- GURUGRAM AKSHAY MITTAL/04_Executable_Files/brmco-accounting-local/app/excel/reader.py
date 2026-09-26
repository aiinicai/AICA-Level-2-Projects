"""Reads an uploaded workbook into raw rows.

File-level checks only: valid .xlsx, correct template, supported version,
required columns present. Field-level parsing happens in ``excel.validator``.
"""
from __future__ import annotations

import io
import logging
import zipfile
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.accounting.models import ValidationIssue
from app.excel.template_spec import META_SHEET, SUPPORTED_VERSIONS, TemplateSpec

logger = logging.getLogger("brmco.excel")

MAX_ROWS = 10_000


@dataclass
class RawRow:
    row: int
    values: dict[str, Any]   # column key -> raw cell value


@dataclass
class ReadResult:
    rows: list[RawRow] = field(default_factory=list)
    issues: list[ValidationIssue] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not any(i.severity == "error" for i in self.issues)


def _err(message: str, code: str, column: str | None = None) -> ValidationIssue:
    return ValidationIssue(severity="error", message=message, code=code, column=column)


def _norm_header(value: Any) -> str:
    return str(value or "").strip().rstrip("*").strip().lower()


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def read_workbook(content: bytes, spec: TemplateSpec, file_name: str = "") -> ReadResult:
    result = ReadResult()

    if file_name and not file_name.lower().endswith((".xlsx", ".xlsm")):
        result.issues.append(_err(f'"{file_name}" is not an .xlsx file. Please upload the BRMCo Excel template.',
                                  "file.type"))
        return result
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        logger.warning("Unreadable workbook %s: %s", file_name, exc)
        result.issues.append(_err("The file could not be opened as an Excel workbook. "
                                  "Save it as .xlsx in Excel and try again.", "file.invalid"))
        return result

    # ---- template identity ---------------------------------------------------------
    if META_SHEET not in wb.sheetnames:
        result.issues.append(_err(
            "This is not a BRMCo template (identification sheet missing). "
            f"Download a fresh {spec.title} and copy your data into it.", "file.template"))
        return result
    meta = {str(r[0]).strip(): r[1] for r in wb[META_SHEET].iter_rows(values_only=True) if r and r[0]}
    template_id = str(meta.get("template_id") or "")
    version = str(meta.get("template_version") or "")
    if template_id != spec.template_id:
        result.issues.append(_err(
            f'Wrong template: this file is "{template_id or "unknown"}" but {spec.title} '
            f'("{spec.template_id}") is expected here.', "file.template"))
        return result
    if version not in SUPPORTED_VERSIONS:
        result.issues.append(_err(
            f"Template version {version or '?'} is not supported (supported: "
            f"{', '.join(sorted(SUPPORTED_VERSIONS))}). Download the latest template.", "file.version"))
        return result

    if spec.sheet_name not in wb.sheetnames:
        result.issues.append(_err(f'Sheet "{spec.sheet_name}" is missing from the workbook.', "file.sheet"))
        return result
    ws = wb[spec.sheet_name]

    # ---- header mapping ----------------------------------------------------------------
    header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    position: dict[str, int] = {}
    by_header = {_norm_header(c.header): c for c in spec.columns}
    unknown: list[str] = []
    for idx, value in enumerate(header_cells):
        key = _norm_header(value)
        if not key:
            continue
        col = by_header.get(key)
        if col is None:
            unknown.append(str(value))
        elif col.key in position:
            result.issues.append(_err(f'Column "{col.header}" appears more than once.', "file.column", col.header))
        else:
            position[col.key] = idx

    missing = [c.header for c in spec.columns if c.key not in position]
    for header in missing:
        col = next(c for c in spec.columns if c.header == header)
        if col.required:
            result.issues.append(_err(f'Required column "{header}" is missing.', "file.column", header))
        else:
            result.issues.append(ValidationIssue(
                severity="error", code="file.column", column=header,
                message=f'Column "{header}" is missing. Do not delete template columns — leave them blank instead.'))
    for header in unknown:
        result.issues.append(ValidationIssue(severity="warning", code="file.column", column=header,
                                             message=f'Unexpected column "{header}" will be ignored.'))
    if not result.ok:
        return result

    # ---- data rows ------------------------------------------------------------------------
    for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if excel_row - 1 > MAX_ROWS:
            result.issues.append(_err(f"The file has more than {MAX_ROWS} rows. Split it into smaller files.",
                                      "file.size"))
            break
        mapped = {key: (values[pos] if pos < len(values) else None) for key, pos in position.items()}
        if all(_is_blank(v) for v in mapped.values()):
            continue
        result.rows.append(RawRow(excel_row, mapped))

    if not result.rows and result.ok:
        result.issues.append(_err(f'No data found on sheet "{spec.sheet_name}".', "file.empty"))
    logger.info("Read %s: %d data rows", file_name or spec.template_id, len(result.rows))
    return result
