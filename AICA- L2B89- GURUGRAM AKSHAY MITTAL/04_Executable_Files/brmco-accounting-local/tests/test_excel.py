"""Excel template generation, reading and cell validation."""
import io
from datetime import datetime

from openpyxl import load_workbook

from app.accounting.models import VoucherKind
from app.excel.reader import read_workbook
from app.excel.samples import SAMPLE_ROWS
from app.excel.template_spec import SPECS, spec_for
from app.excel.validator import parse_rows
from tests.conftest import workbook


def edit(content: bytes, fn) -> bytes:
    wb = load_workbook(io.BytesIO(content))
    fn(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_every_template_round_trips():
    for kind, spec in SPECS.items():
        content = workbook(kind)
        result = read_workbook(content, spec, "x.xlsx")
        assert result.ok, (kind, result.issues)
        assert len(result.rows) == len(SAMPLE_ROWS[kind])
        parsed, issues = parse_rows(spec, result.rows)
        assert not issues, (kind, issues)


def test_missing_column():
    spec = spec_for(VoucherKind.SALES)
    content = edit(workbook(VoucherKind.SALES), lambda wb: wb[spec.sheet_name].delete_cols(3))
    result = read_workbook(content, spec, "x.xlsx")
    assert not result.ok
    assert any('"Customer Ledger" is missing' in i.message for i in result.issues)


def test_wrong_template_type():
    content = workbook(VoucherKind.JOURNAL)
    result = read_workbook(content, spec_for(VoucherKind.SALES), "x.xlsx")
    assert not result.ok
    assert result.issues[0].code == "file.template"


def test_unsupported_version():
    content = edit(workbook(VoucherKind.SALES), lambda wb: wb["_brmco_meta"].cell(row=2, column=2, value="0.9"))
    result = read_workbook(content, spec_for(VoucherKind.SALES), "x.xlsx")
    assert result.issues[0].code == "file.version"


def test_not_an_excel_file():
    result = read_workbook(b"not a zip", spec_for(VoucherKind.SALES), "x.xlsx")
    assert result.issues[0].code == "file.invalid"
    result = read_workbook(b"a,b", spec_for(VoucherKind.SALES), "x.csv")
    assert result.issues[0].code == "file.type"


def test_invalid_values_and_dates():
    spec = spec_for(VoucherKind.RECEIPT)
    rows = [{**SAMPLE_ROWS[VoucherKind.RECEIPT][0], "amount": "12,34x", "voucher_date": "31-02-2026"}]
    result = read_workbook(workbook(VoucherKind.RECEIPT, rows), spec, "x.xlsx")
    _, issues = parse_rows(spec, result.rows)
    cols = {(i.row, i.column) for i in issues}
    assert (2, "Amount") in cols
    assert (2, "Voucher Date") in cols


def test_invalid_gst_rate_and_gstin():
    spec = spec_for(VoucherKind.SALES)
    rows = [{**SAMPLE_ROWS[VoucherKind.SALES][0], "cgst_rate": 7, "customer_gstin": "27ABCDE1234F1Z9"}]
    result = read_workbook(workbook(VoucherKind.SALES, rows), spec, "x.xlsx")
    _, issues = parse_rows(spec, result.rows)
    assert any(i.column == "CGST Rate" for i in issues)
    assert any(i.column == "Customer GSTIN" for i in issues)


def test_text_dates_are_accepted():
    spec = spec_for(VoucherKind.RECEIPT)
    rows = [{**SAMPLE_ROWS[VoucherKind.RECEIPT][0], "voucher_date": "27/09/2026", "instrument_date": None}]
    result = read_workbook(workbook(VoucherKind.RECEIPT, rows), spec, "x.xlsx")
    parsed, issues = parse_rows(spec, result.rows)
    assert not issues
    assert parsed[0].values["voucher_date"] == datetime(2026, 9, 27).date()


def test_duplicate_entry_in_file(config, masters):
    from app.accounting.services import VoucherAssembler
    from app.accounting.validation import VoucherValidator
    from tests.conftest import TODAY, NoHistory

    spec = spec_for(VoucherKind.RECEIPT)
    first = SAMPLE_ROWS[VoucherKind.RECEIPT][0]
    result = read_workbook(workbook(VoucherKind.RECEIPT, [first, first]), spec, "x.xlsx")
    parsed, _ = parse_rows(spec, result.rows)
    vouchers, _ = VoucherAssembler(spec, config, masters).assemble(parsed)
    issues = VoucherValidator(config, masters, NoHistory(), today=TODAY).validate(vouchers)
    assert any(i.code == "dup.file" and i.row == 3 for i in issues)


def test_conflicting_invoice_rows(config, masters):
    from app.accounting.services import VoucherAssembler

    spec = spec_for(VoucherKind.SALES)
    a, b, c = SAMPLE_ROWS[VoucherKind.SALES]
    c = {**c, "customer_ledger": "ABC Traders"}  # continuation row names a different customer
    result = read_workbook(workbook(VoucherKind.SALES, [a, b, c]), spec, "x.xlsx")
    parsed, _ = parse_rows(spec, result.rows)
    _, issues = VoucherAssembler(spec, config, masters).assemble(parsed)
    assert any(i.code == "row.inconsistent" and i.row == 4 for i in issues)


def test_template_has_dropdowns_and_hidden_meta():
    wb = load_workbook(io.BytesIO(workbook(VoucherKind.SALES, [])))
    ws = wb[spec_for(VoucherKind.SALES).sheet_name]
    assert len(ws.data_validations.dataValidation) > 5
    assert wb["_brmco_meta"].sheet_state == "veryHidden"
