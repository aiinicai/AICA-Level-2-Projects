"""
AICA SOP TOOL - SOP Builder v14

Change in v14
-------------
Generic interface command listings are now removed.

Why they appeared
-----------------
Optical character recognition reads the application ribbon, which is
present on every Excel screen regardless of what the user did. The result
is a sentence that enumerates the ribbon rather than describing an action:

    "Captured Excel commands include Insert, Formulas, Data, Review,
     Home, File, View, Sort, Filter, and Find."

In earlier documents these listings were bundled into the same sentence as
a worksheet tab claim, so removing the tab claim removed the commands as a
side effect. Once the prompt forbade tab claims, the AI began writing the
command list as a separate sentence, which no existing filter matched.

How v14 distinguishes them
--------------------------
A genuine reference names one or two commands tied to the action and
contains no comma separated list:

    "Use the captured Save As and Save commands to save the workbook."   kept

A ribbon listing enumerates many items:

    "The Excel interface includes Insert, Formulas, Data, Review, View."  removed

The test requires a command marker, an enumeration verb, and at least two
commas, so an action reference is never removed.

Retained from v13
-----------------
    - sentences stating only that a detail needs validation are removed
    - claims about visible cell references are removed
    - claims about displayed worksheet tabs are removed
    - no per step folder path box; the narrative states the path inline
    - consolidated Captured Folder Locations table in Prerequisites
    - Markdown horizontal rules discarded
    - readable full page or sliced process flow

Install
-------
    python -m pip install python-docx pillow pywin32

Run
---
    python sop_builder_v14.py

Build EXE
---------
    pyinstaller --onefile --noconsole --clean --name AICA_SOP_Tool_v14 sop_builder_v14.py
"""

import csv
import os
import re
import sys
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

APP_TITLE = "AICA SOP Tool - SOP Builder v14"
TOOL_NAME = "AICA SOP Tool"
BRAND_HEX = "6B001B"
BRAND = RGBColor(107, 0, 27)
GREY = RGBColor(85, 85, 85)
PATH_BLUE = RGBColor(31, 56, 100)
LIGHT_BRAND = "F3E6EA"
SUCCESS = "#1F7A1F"
WARN = "#9A6700"
ERROR = "#C00000"

# ---------------------------------------------------------------------
# Content filters
# ---------------------------------------------------------------------

VALIDATION_PATTERN = re.compile(
    r"(?:needs?|requires?)\s+business\s+validation"
    r"|business\s+validation\s+(?:is\s+)?required",
    re.IGNORECASE,
)

CELL_CLAIM_PATTERN = re.compile(
    r"\bcell\s+reference|\bvisible\s+cell|\breferences?\s+(?:included|include|"
    r"visible)|\bwith\s+cell\s+[A-Z]{1,3}\d",
    re.IGNORECASE,
)

TAB_CLAIM_PATTERN = re.compile(
    r"\b(?:displayed|display|displays|shows?|showed|contains?|contained|"
    r"include[sd]?|evidence[sd]?)\b[^.]{0,60}\btabs?\b"
    r"|\btabs?\b[^.]{0,40}\b(?:include[sd]?|visible|captured|displayed)\b"
    r"|\bcaptured\s+workbook\s+displayed\b"
    r"|\bworkbook\s+context\s+displayed\b",
    re.IGNORECASE,
)

# A sentence that enumerates interface commands rather than instructing.
COMMAND_MARKER = re.compile(r"\b(?:command|interface|ribbon)\w*\b",
                            re.IGNORECASE)
ENUM_VERB = re.compile(
    r"\b(?:include[sd]?|including|captured|displayed|available)\b",
    re.IGNORECASE)

# Minimum comma count for a sentence to count as an enumeration. A genuine
# reference such as "the Save As and Save commands" contains no commas.
COMMAND_MIN_COMMAS = 2

RULE_PATTERN = re.compile(r"^[-*_\s]{3,}$")

MAX_IMAGE_HEIGHT = 8.6
TALL_RATIO = 1.9

SECTION_KEYS = (
    "purpose", "scope", "process_owner", "prerequisites", "inputs",
    "detailed_steps", "outputs", "controls", "risks", "escalation",
)

SECTION_TITLES = {
    "purpose": "Purpose",
    "scope": "Scope",
    "process_owner": "Process Owner and Frequency",
    "prerequisites": "Prerequisites and System Access",
    "inputs": "Inputs",
    "detailed_steps": "Detailed Process Steps",
    "outputs": "Outputs and Deliverables",
    "controls": "Key Controls and Validation Checks",
    "risks": "Risks and Common Errors",
    "escalation": "Escalation Contacts",
}

ALIASES = {
    "purpose": "purpose",
    "scope": "scope",
    "process owner and frequency": "process_owner",
    "process owner": "process_owner",
    "frequency": "process_owner",
    "prerequisites and system access": "prerequisites",
    "prerequisites": "prerequisites",
    "system access": "prerequisites",
    "inputs": "inputs",
    "detailed process steps": "detailed_steps",
    "process steps": "detailed_steps",
    "procedure": "detailed_steps",
    "outputs and deliverables": "outputs",
    "outputs": "outputs",
    "deliverables": "outputs",
    "key controls and validation checks": "controls",
    "key controls": "controls",
    "controls": "controls",
    "validation checks": "controls",
    "risks and common errors": "risks",
    "risks": "risks",
    "common errors": "risks",
    "escalation contacts": "escalation",
    "escalation": "escalation",
}


# ---------------------------------------------------------------------
# Sentence level cleaning
# ---------------------------------------------------------------------

def is_command_listing(sentence):
    """
    True when a sentence merely enumerates interface commands.

    Requires a command marker, an enumeration verb and a comma separated
    list. An action reference naming one or two commands has no commas and
    is therefore kept.
    """
    if not COMMAND_MARKER.search(sentence):
        return False
    if not ENUM_VERB.search(sentence):
        return False
    return sentence.count(",") >= COMMAND_MIN_COMMAS


def split_sentences(text):
    """Split prose into sentences, keeping the terminator attached."""
    if not text:
        return []
    parts = re.split(r"(?<=[.!?])\s+", text.strip())
    return [p for p in parts if p.strip()]


def clean_prose(text, filters):
    """
    Remove sentences that assert content the capture cannot support.

    Returns (cleaned_text, counts).
    """
    counts = {"validation": 0, "cells": 0, "tabs": 0, "commands": 0}
    if not text:
        return "", counts

    kept = []
    for sentence in split_sentences(text):
        if filters.get("validation") and VALIDATION_PATTERN.search(sentence):
            counts["validation"] += 1
            continue
        if filters.get("cells") and CELL_CLAIM_PATTERN.search(sentence):
            counts["cells"] += 1
            continue
        if filters.get("tabs") and TAB_CLAIM_PATTERN.search(sentence):
            counts["tabs"] += 1
            continue
        if filters.get("commands") and is_command_listing(sentence):
            counts["commands"] += 1
            continue
        kept.append(sentence.strip())

    cleaned = " ".join(kept).strip()
    cleaned = re.sub(r"\s*[,;]\s*(?:and|therefore|however)?\s*$", ".", cleaned)
    cleaned = re.sub(r"\s{2,}", " ", cleaned)
    return cleaned.strip(), counts


def cell_is_empty_claim(text, drop_validation=True):
    """True when a table cell states nothing beyond the validation phrase."""
    if not drop_validation or not text:
        return False
    stripped = VALIDATION_PATTERN.sub("", text).strip()
    stripped = stripped.strip(" .,:;-\u2013\u2014")
    return stripped == ""


# ---------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------

def app_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def safe_name(value):
    value = re.sub(r"[^A-Za-z0-9 _-]", "", value or "").strip()
    return re.sub(r"\s+", "_", value) or "Process"


def available_path(directory, stem, ext):
    first = os.path.join(directory, stem + ext)
    if not os.path.exists(first):
        return first
    number = 2
    while True:
        candidate = os.path.join(directory, "{}_v{}{}".format(stem, number, ext))
        if not os.path.exists(candidate):
            return candidate
        number += 1


def is_horizontal_rule(text):
    stripped = (text or "").strip()
    if not stripped:
        return False
    if not RULE_PATTERN.match(stripped):
        return False
    return len(re.sub(r"\s", "", stripped)) >= 3


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def write_text(paragraph, text, size=10.5, bold=False, color=None,
               italic=False, font="Segoe UI"):
    if not text:
        return
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def set_cell(cell, value, bold=False, color=None, size=8.5, mono=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    font = "Consolas" if mono else "Segoe UI"
    write_text(paragraph, value, size=size, bold=bold, color=color, font=font)


# ---------------------------------------------------------------------
# Narrative parsing
# ---------------------------------------------------------------------

def strip_md(text):
    value = (text or "").strip()
    value = value.lstrip("#").strip()
    value = value.replace("**", "").replace("__", "").replace("`", "")
    value = re.sub(r"^[-*\u2022]\s+", "", value)
    return value.strip()


def heading_key(line):
    value = strip_md(line)
    value = re.sub(r"^step\s+\d+\s*[:.)-]?\s*", "", value, flags=re.I)
    value = re.sub(r"^\d+\s*[.)-]\s*", "", value)
    value = value.rstrip(":").strip().lower()
    if value.startswith("process owner"):
        return "process_owner"
    for alias, key in ALIASES.items():
        if value == alias:
            return key
        if value.startswith(alias + " ") and len(value) <= len(alias) + 45:
            return key
    return ""


def parse_sections(narrative):
    sections = {key: [] for key in SECTION_KEYS}
    current = ""
    for raw in narrative.replace("\r\n", "\n").split("\n"):
        if is_horizontal_rule(raw):
            continue
        key = heading_key(raw)
        if key:
            current = key
            continue
        if current:
            sections[current].append(raw.rstrip())
    return sections


def split_row(line):
    sentinel = "\uE000"
    protected = line.replace("\\|", sentinel).strip().strip("|")
    return [strip_md(cell.replace(sentinel, "|")) for cell in protected.split("|")]


def is_separator(cells):
    return bool(cells) and all(
        re.fullmatch(r"\s*:?-{2,}:?\s*", c or "") for c in cells)


def table_rows(lines):
    rows = []
    for line in lines:
        if not line.strip().startswith("|"):
            continue
        cells = split_row(line)
        if not is_separator(cells):
            rows.append(cells)
    return rows


def parse_detail_rows(lines):
    data = {}
    for cells in table_rows(lines):
        try:
            number = int(cells[0])
        except (ValueError, IndexError):
            continue
        data[number] = {
            "activity": cells[1] if len(cells) > 1 else "",
            "instruction": " | ".join(cells[2:]).strip() if len(cells) > 2 else "",
        }
    return data


# ---------------------------------------------------------------------
# Session readers
# ---------------------------------------------------------------------

def read_steps(session):
    path = os.path.join(session, "process_steps.csv")
    if not os.path.exists(path):
        raise FileNotFoundError(
            "process_steps.csv was not found in the selected session.")

    steps = []
    has_path_column = False
    with open(path, "r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        fields = set(reader.fieldnames or [])
        if not {"StepNo", "StepDescription"}.issubset(fields):
            raise ValueError(
                "process_steps.csv requires StepNo and StepDescription columns.")
        has_path_column = "FolderPath" in fields
        for row in reader:
            try:
                number = int(float(row.get("StepNo") or 0))
            except ValueError:
                continue
            if number <= 0:
                continue
            steps.append({
                "StepNo": number,
                "RecordType": (row.get("RecordType") or "Activity").strip(),
                "Application": (row.get("Application") or "").strip(),
                "StepDescription": (row.get("StepDescription") or "").strip(),
                "FolderPath": (row.get("FolderPath") or "").strip(),
                "Screenshot": (row.get("Screenshot") or "").strip(),
            })
    steps.sort(key=lambda item: item["StepNo"])
    if not steps:
        raise ValueError("No usable captured steps were found.")
    return steps, has_path_column


def detect_generation(steps, has_path_column):
    shots = [s["Screenshot"] for s in steps if s["Screenshot"]]
    event_shots = [s for s in shots if s.lower().startswith("evt_")]
    distinct = len(set(shots))
    duplicates = len(shots) - distinct
    paths = len({s["FolderPath"] for s in steps if s["FolderPath"]})
    modern = bool(event_shots) or (has_path_column and paths > 0)
    return {
        "modern": modern, "distinct_shots": distinct,
        "duplicates": duplicates, "paths": paths,
    }


def read_narrative(session):
    for filename in ("sop_text.txt", "sop_text.md", "sop.txt"):
        path = os.path.join(session, filename)
        if os.path.exists(path):
            for encoding in ("utf-8-sig", "utf-8", "cp1252"):
                try:
                    with open(path, "r", encoding=encoding) as handle:
                        return handle.read(), filename
                except UnicodeDecodeError:
                    continue
    return "", ""


def find_flow(session):
    for filename in ("process_flow.png", "process_flow.jpg",
                     "process_flow.jpeg"):
        path = os.path.join(session, filename)
        if os.path.exists(path):
            return path
    return ""


def find_screenshot(session, filename):
    if not filename:
        return ""
    for path in (os.path.join(session, "screenshots", filename),
                 os.path.join(session, filename)):
        if os.path.exists(path):
            return path
    return ""


# ---------------------------------------------------------------------
# Word rendering
# ---------------------------------------------------------------------

def configure_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Segoe UI"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BRAND


def cover(doc, meta, session_name):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("STANDARD OPERATING PROCEDURE")
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = GREY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta["process_name"])
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = BRAND
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TOOL_NAME + " generated process documentation")
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = GREY
    doc.add_paragraph()

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("Process Owner", meta["owner"]),
        ("Department", meta["department"]),
        ("Frequency", meta["frequency"]),
        ("Document Version", meta["version"] or "1.0"),
        ("Effective Date", meta["effective_date"]),
        ("Prepared By", meta["prepared_by"] or TOOL_NAME),
        ("Approved By", meta["approved_by"]),
        ("Capture Session", session_name),
    ]
    for label, value in rows:
        if not value:
            continue
        cells = table.add_row().cells
        set_cell(cells[0], label, True, BRAND, 9.5)
        set_cell(cells[1], value, False, None, 9.5)
        shade(cells[0], LIGHT_BRAND)
    doc.add_page_break()


def header_footer(doc, meta):
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.text = "{}  |  Version {}".format(
        meta["process_name"], meta["version"] or "1.0")
    for run in hp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GREY
    fp = section.footer.paragraphs[0]
    run = fp.add_run(TOOL_NAME + "  |  ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    add_page_number(fp)


def document_control(doc, meta):
    doc.add_heading("Document Control", level=1)
    table = doc.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("Version", "Effective Date", "Prepared By", "Approved By")
    values = (meta["version"] or "1.0", meta["effective_date"],
              meta["prepared_by"], meta["approved_by"])
    for index, value in enumerate(headers):
        set_cell(table.cell(0, index), value, True,
                 RGBColor(255, 255, 255), 8.5)
        shade(table.cell(0, index), BRAND_HEX)
    repeat_header(table.rows[0])
    for index, value in enumerate(values):
        set_cell(table.cell(1, index), value or "-", False, None, 8.5)


def place_image(doc, path, max_width, max_height=5.9):
    if not path:
        return False
    try:
        with Image.open(path) as picture:
            px_w, px_h = picture.size
        if px_w <= 0 or px_h <= 0:
            return False
        ratio = px_h / float(px_w)
        width = max_width
        if width * ratio > max_height:
            width = max_height / ratio
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p.add_run().add_picture(path, width=Inches(width))
        return True
    except Exception:
        return False


def render_flowchart(doc, path, max_width, work_dir):
    if not path:
        return False, 0
    try:
        with Image.open(path) as picture:
            px_w, px_h = picture.size
            ratio = px_h / float(px_w)
            if ratio <= TALL_RATIO:
                doc.add_page_break()
                ok = place_image(doc, path, max_width, MAX_IMAGE_HEIGHT)
                if ok:
                    doc.add_page_break()
                return ok, 1 if ok else 0

            segments = max(2, min(int(ratio / TALL_RATIO) + 1, 6))
            slice_height = px_h // segments
            overlap = int(slice_height * 0.04)
            slice_dir = os.path.join(work_dir, "_flow_slices")
            os.makedirs(slice_dir, exist_ok=True)

            placed = 0
            for index in range(segments):
                top = max(0, index * slice_height - (overlap if index else 0))
                bottom = min(px_h, (index + 1) * slice_height + overlap)
                crop = picture.crop((0, top, px_w, bottom))
                slice_path = os.path.join(
                    slice_dir, "flow_part_{:02d}.png".format(index + 1))
                crop.save(slice_path, "PNG")
                doc.add_page_break()
                heading = doc.add_paragraph()
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = heading.add_run("Process Flow - part {} of {}".format(
                    index + 1, segments))
                run.font.name = "Segoe UI"
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = BRAND
                if place_image(doc, slice_path, max_width, MAX_IMAGE_HEIGHT):
                    placed += 1
            doc.add_page_break()
            return placed > 0, placed
    except Exception:
        return False, 0


def render_table(doc, rows, filters):
    """
    Render a Markdown table.

    A data row is dropped when a cell states nothing beyond the validation
    phrase, because such a row carries no information.
    """
    if not rows:
        return False, 0

    dropped = 0
    kept = [rows[0]]

    for row in rows[1:]:
        cleaned = []
        drop_row = False
        for cell in row:
            if cell_is_empty_claim(cell, filters.get("validation")):
                drop_row = True
                break
            text, _ = clean_prose(cell, filters)
            cleaned.append(text)
        if drop_row:
            dropped += 1
            continue
        if len(cleaned) > 1 and not any(c.strip() for c in cleaned[1:]):
            dropped += 1
            continue
        kept.append(cleaned)

    if len(kept) <= 1:
        return False, dropped

    columns = max(len(row) for row in kept)
    table = doc.add_table(rows=len(kept), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for r_index, row in enumerate(kept):
        for c_index in range(columns):
            value = row[c_index] if c_index < len(row) else ""
            cell = table.cell(r_index, c_index)
            if r_index == 0:
                set_cell(cell, value, True, RGBColor(255, 255, 255), 8.5)
                shade(cell, BRAND_HEX)
            else:
                set_cell(cell, value, False, None, 8.5)
                if r_index % 2 == 0:
                    shade(cell, "F7F3F4")
    repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return True, dropped


def render_section(doc, lines, filters, tally):
    index = 0
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if not stripped or is_horizontal_rule(stripped):
            index += 1
            continue

        if stripped.startswith("|"):
            block = []
            while index < len(lines):
                candidate = lines[index].strip()
                if not candidate:
                    index += 1
                    continue
                if not candidate.startswith("|"):
                    break
                block.append(lines[index])
                index += 1
            _, dropped = render_table(doc, table_rows(block), filters)
            tally["rows"] += dropped
            continue

        text = strip_md(raw)
        if text:
            cleaned, counts = clean_prose(text, filters)
            for key in counts:
                tally[key] += counts[key]
            if cleaned:
                bullet = raw.lstrip().startswith(("-", "*", "\u2022"))
                paragraph = (doc.add_paragraph(style="List Bullet") if bullet
                             else doc.add_paragraph())
                write_text(paragraph, cleaned)
        index += 1


def owner_section(doc, meta):
    doc.add_heading("Process Owner and Frequency", level=1)
    rows = [
        ("Process Owner", meta["owner"]),
        ("Process Performer", meta["performer"] or meta["owner"]),
        ("Frequency", meta["frequency"]),
        ("Reporting Cut-off", meta["cutoff"]),
        ("Review / Approval", meta["review"] or meta["approved_by"]),
    ]
    rows = [(label, value) for label, value in rows if value]
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for index, (label, value) in enumerate(rows):
        set_cell(table.cell(index, 0), label, True, BRAND, 9)
        set_cell(table.cell(index, 1), value, False, None, 9)
        shade(table.cell(index, 0), LIGHT_BRAND)


def captured_paths_block(doc, steps):
    paths = []
    for step in steps:
        value = step["FolderPath"]
        if value and value not in paths:
            paths.append(value)
    if not paths:
        return 0

    doc.add_heading("Captured Folder Locations", level=2)
    doc.add_paragraph(
        "The following folder locations were captured directly from Windows "
        "File Explorer during the recorded session.")

    table = doc.add_table(rows=len(paths) + 1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    set_cell(table.cell(0, 0), "#", True, RGBColor(255, 255, 255), 8.5)
    set_cell(table.cell(0, 1), "Folder path", True,
             RGBColor(255, 255, 255), 8.5)
    shade(table.cell(0, 0), BRAND_HEX)
    shade(table.cell(0, 1), BRAND_HEX)
    repeat_header(table.rows[0])

    for index, value in enumerate(paths, start=1):
        set_cell(table.cell(index, 0), str(index), False, None, 8.5)
        set_cell(table.cell(index, 1), value, False, PATH_BLUE, 8.5, mono=True)
        if index % 2 == 0:
            shade(table.cell(index, 0), "F7F3F4")
            shade(table.cell(index, 1), "F7F3F4")

    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return len(paths)


def details_section(doc, steps, ai_details, session, width, dedupe,
                    filters, tally):
    doc.add_heading("Detailed Process Steps", level=1)
    inserted = 0
    repeated = 0
    absent = 0
    missing_files = []
    first_use = {}

    for step in steps:
        number = step["StepNo"]
        ai = ai_details.get(number, {})
        activity = (ai.get("activity") or step["StepDescription"]
                    or "Captured action")
        instruction = ai.get("instruction") or step["StepDescription"]

        activity, counts = clean_prose(activity, filters)
        for key in counts:
            tally[key] += counts[key]
        if not activity:
            activity = step["StepDescription"] or "Captured action"

        doc.add_heading("Step {}: {}".format(number, activity), level=2)

        cleaned, counts = clean_prose(instruction, filters)
        for key in counts:
            tally[key] += counts[key]

        if cleaned:
            paragraph = doc.add_paragraph()
            write_text(paragraph, cleaned)

        filename = step["Screenshot"]
        shot = find_screenshot(session, filename)

        if not filename:
            absent += 1
            continue
        if not shot:
            missing_files.append(filename)
            continue
        if dedupe and filename in first_use:
            repeated += 1
            continue
        if place_image(doc, shot, width, 5.9):
            first_use[filename] = number
            inserted += 1

    return {
        "inserted": inserted, "repeated": repeated, "absent": absent,
        "missing": sorted(set(missing_files)),
    }


def convert_pdf(docx_path):
    try:
        import pythoncom
        import win32com.client
    except Exception:
        return "", "Install pywin32 for automatic PDF conversion."
    word = None
    pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
    try:
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(os.path.abspath(docx_path))
        document.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
        document.Close(False)
        word.Quit()
        pythoncom.CoUninitialize()
        return pdf_path, ""
    except Exception as error:
        try:
            if word:
                word.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass
        return "", str(error)


# ---------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------

def build(session, meta, options):
    steps, has_path_column = read_steps(session)
    capture = detect_generation(steps, has_path_column)

    narrative, narrative_name = read_narrative(session)
    if not narrative:
        raise ValueError(
            "sop_text.txt was not found. Generate and save the AI narrative "
            "first.")

    sections = parse_sections(narrative)
    ai_details = parse_detail_rows(sections["detailed_steps"])
    flow = find_flow(session)

    filters = {
        "validation": options["drop_validation"],
        "cells": options["drop_cells"],
        "tabs": options["drop_tabs"],
        "commands": options["drop_commands"],
    }
    tally = {"validation": 0, "cells": 0, "tabs": 0,
             "commands": 0, "rows": 0}

    if options["dedupe_mode"] == "auto":
        dedupe = not capture["modern"] and capture["duplicates"] > 0
    elif options["dedupe_mode"] == "always":
        dedupe = True
    else:
        dedupe = False

    doc = Document()
    configure_doc(doc)
    cover(doc, meta, os.path.basename(session.rstrip("\\/")))
    header_footer(doc, meta)
    document_control(doc, meta)

    page = doc.sections[0]
    usable = (page.page_width.inches - page.left_margin.inches
              - page.right_margin.inches)
    step_width = min(options["width"], usable)

    doc.add_heading("Purpose", level=1)
    render_section(doc, sections["purpose"], filters, tally)

    doc.add_heading("Scope", level=1)
    render_section(doc, sections["scope"], filters, tally)

    flow_ok, flow_parts = False, 0
    if options["flow"] and flow:
        doc.add_heading("Process Flow", level=1)
        doc.add_paragraph(
            "The following diagram presents the captured sequence of "
            "application activities and approved user-interaction events.")
        flow_ok, flow_parts = render_flowchart(doc, flow, usable, session)

    owner_section(doc, meta)

    doc.add_heading(SECTION_TITLES["prerequisites"], level=1)
    render_section(doc, sections["prerequisites"], filters, tally)
    listed_paths = 0
    if options["show_paths"]:
        listed_paths = captured_paths_block(doc, steps)

    doc.add_heading(SECTION_TITLES["inputs"], level=1)
    render_section(doc, sections["inputs"], filters, tally)

    shots = details_section(doc, steps, ai_details, session, step_width,
                            dedupe, filters, tally)

    for key in ("outputs", "controls", "risks", "escalation"):
        doc.add_heading(SECTION_TITLES[key], level=1)
        render_section(doc, sections[key], filters, tally)

    doc.add_page_break()
    doc.add_heading("Appendix - Capture Quality Summary", level=1)

    generation = ("Interaction Capture v4 or later (event-level screenshots)"
                  if capture["modern"] else
                  "Interaction Capture v3 (inherited screenshots)")

    removed = (tally["validation"] + tally["cells"]
               + tally["tabs"] + tally["commands"])

    summary_rows = [
        ["Metric", "Result"],
        ["Generated By", TOOL_NAME],
        ["Capture Generation Detected", generation],
        ["Narrative Source", narrative_name],
        ["Captured Steps", str(len(steps))],
        ["AI Detailed Rows Matched", str(len(ai_details))],
        ["Distinct Screens Embedded", str(shots["inserted"])],
        ["Steps With No Captured Screen", str(shots["absent"])],
        ["Folder Paths Captured", str(capture["paths"])],
        ["Interface Listings Removed", str(tally["commands"])],
        ["Other Unverifiable Statements Removed",
         str(tally["validation"] + tally["cells"] + tally["tabs"])],
        ["Flowchart Segments", str(flow_parts) if flow_ok else "0"],
    ]
    render_table(doc, summary_rows,
                 {"validation": False, "cells": False,
                  "tabs": False, "commands": False})

    if shots["missing"]:
        doc.add_paragraph(
            "Missing screenshot files: " + ", ".join(shots["missing"]))

    stem = "SOP_{}".format(safe_name(meta["process_name"]))
    docx_path = available_path(session, stem, ".docx")
    doc.save(docx_path)

    pdf_path, pdf_error = "", ""
    if options["pdf"]:
        pdf_path, pdf_error = convert_pdf(docx_path)

    return {
        "docx": docx_path, "pdf": pdf_path, "pdf_error": pdf_error,
        "flow_parts": flow_parts, "shots": shots, "capture": capture,
        "listed_paths": listed_paths, "tally": tally, "removed": removed,
    }


# ---------------------------------------------------------------------
# GUI
# ---------------------------------------------------------------------

# --- AICA scrollable window helper ---
def attach_scroll(window, bg="white"):
    """
    Return a frame that scrolls vertically inside the given window.

    Every widget placed in the returned frame can be reached by scrolling,
    so content below the fold stays accessible on a smaller display. The
    window itself keeps its own title, geometry and protocol handlers.
    """
    import tkinter as _tk
    from tkinter import ttk as _ttk

    outer = _tk.Frame(window, bg=bg)
    outer.pack(fill="both", expand=True)

    canvas = _tk.Canvas(outer, bg=bg, highlightthickness=0, bd=0)
    bar = _ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
    canvas.configure(yscrollcommand=bar.set)

    canvas.pack(side="left", fill="both", expand=True)
    body = _tk.Frame(canvas, bg=bg)
    holder = canvas.create_window((0, 0), window=body, anchor="nw")

    def _content_changed(_event=None):
        canvas.configure(scrollregion=canvas.bbox("all"))
        needed = body.winfo_reqheight() > canvas.winfo_height()
        if needed and not bar.winfo_ismapped():
            bar.pack(side="right", fill="y")
        elif not needed and bar.winfo_ismapped():
            bar.pack_forget()

    def _width_changed(event):
        canvas.itemconfigure(holder, width=event.width)
        _content_changed()

    body.bind("<Configure>", _content_changed)
    canvas.bind("<Configure>", _width_changed)

    def _wheel(event):
        if body.winfo_reqheight() <= canvas.winfo_height():
            return
        step = -1 if getattr(event, "delta", 0) > 0 else 1
        canvas.yview_scroll(step, "units")

    def _wheel_on(_event=None):
        canvas.bind_all("<MouseWheel>", _wheel)

    def _wheel_off(_event=None):
        canvas.unbind_all("<MouseWheel>")

    canvas.bind("<Enter>", _wheel_on)
    canvas.bind("<Leave>", _wheel_off)

    def _key(event):
        if event.keysym == "Prior":
            canvas.yview_scroll(-1, "pages")
        elif event.keysym == "Next":
            canvas.yview_scroll(1, "pages")
        elif event.keysym == "Home":
            canvas.yview_moveto(0)
        elif event.keysym == "End":
            canvas.yview_moveto(1)

    for key in ("<Prior>", "<Next>", "<Home>", "<End>"):
        window.bind(key, _key)

    return body
# --- end helper ---


class App:

    def __init__(self, root):
        self.root = root
        self.session = ""

        root.title(APP_TITLE)
        root.geometry("950x860")
        root.minsize(900, 400)
        root.configure(bg="white")
        # Scrollable body. Widgets are placed here so content
        # below the fold can be reached by scrolling.
        self._body = attach_scroll(root)


        bottom = tk.Frame(self._body, bg="white")
        bottom.pack(side="bottom", fill="x", padx=22, pady=(6, 12))
        tk.Button(bottom, text="BUILD SOP", command=self.run,
                  bg=SUCCESS, fg="white", relief="flat",
                  font=("Segoe UI", 12, "bold"), width=28,
                  cursor="hand2").pack(side="left")
        self.message = tk.Label(bottom, text="Select a session folder.",
                                bg="white", fg="gray", wraplength=560,
                                justify="left", anchor="w")
        self.message.pack(side="left", padx=14)

        tk.Label(self._body, text=TOOL_NAME.upper(),
                 font=("Segoe UI", 20, "bold"),
                 bg="white", fg="#" + BRAND_HEX).pack(pady=(12, 0))
        tk.Label(self._body, text="SOP Builder | Version 14",
                 font=("Segoe UI", 9), bg="white", fg="gray").pack()

        top = tk.Frame(self._body, bg="white")
        top.pack(fill="x", padx=22, pady=10)
        tk.Button(top, text="SELECT SESSION FOLDER", command=self.select,
                  bg="#" + BRAND_HEX, fg="white", relief="flat", width=24,
                  font=("Segoe UI", 10, "bold"),
                  cursor="hand2").pack(side="left")
        self.folder = tk.Label(top, text="No folder selected",
                               bg="white", fg="gray")
        self.folder.pack(side="left", padx=12)

        validation = tk.LabelFrame(self._body, text=" Session validation ",
                                   bg="white", fg="#" + BRAND_HEX,
                                   font=("Segoe UI", 9, "bold"))
        validation.pack(fill="x", padx=22, pady=4)
        self.checks = {}
        for name in ("process_steps.csv", "sop_text.txt",
                     "process_flow.png", "screenshots"):
            row = tk.Frame(validation, bg="white")
            row.pack(fill="x", padx=12, pady=2)
            tk.Label(row, text=name + ":", width=20, anchor="w",
                     bg="white", font=("Segoe UI", 9)).pack(side="left")
            label = tk.Label(row, text="Not checked", bg="white", fg="gray",
                             font=("Segoe UI", 9))
            label.pack(side="left")
            self.checks[name] = label

        self.capture_label = tk.Label(
            validation, text="Capture generation: not detected",
            bg="white", fg="gray", font=("Segoe UI", 9, "bold"),
            anchor="w", wraplength=880, justify="left")
        self.capture_label.pack(fill="x", padx=12, pady=(4, 8))

        meta_box = tk.LabelFrame(self._body, text=" Document metadata ",
                                 bg="white", fg="#" + BRAND_HEX,
                                 font=("Segoe UI", 9, "bold"))
        meta_box.pack(fill="x", padx=22, pady=6)
        grid = tk.Frame(meta_box, bg="white")
        grid.pack(fill="x", padx=12, pady=8)

        self.entries = {}
        fields = [
            ("process_name", "Process name", "Monthly Revenue Report"),
            ("owner", "Process owner", "Dheeraj Gupta"),
            ("performer", "Process performer", "Dheeraj Gupta"),
            ("department", "Department", "Finance"),
            ("frequency", "Frequency", "Monthly"),
            ("cutoff", "Reporting cut-off", "31-Aug-26"),
            ("review", "Review / approval", "Anand Gupta"),
            ("version", "Version", "1.0"),
            ("effective_date", "Effective date",
             datetime.now().strftime("%d %b %Y")),
            ("prepared_by", "Prepared by", "Dheeraj Gupta"),
            ("approved_by", "Approved by", "Anand Gupta"),
        ]
        for index, (key, label, default) in enumerate(fields):
            r_index, col = index // 2, (index % 2) * 2
            tk.Label(grid, text=label + ":", bg="white").grid(
                row=r_index, column=col, sticky="w", padx=(0, 6), pady=3)
            entry = tk.Entry(grid, width=30)
            entry.grid(row=r_index, column=col + 1, sticky="w",
                       padx=(0, 22), pady=3)
            entry.insert(0, default)
            self.entries[key] = entry

        # ---- content filters ----
        content = tk.LabelFrame(self._body, text=" Remove unverifiable content ",
                                bg="white", fg="#" + BRAND_HEX,
                                font=("Segoe UI", 9, "bold"))
        content.pack(fill="x", padx=22, pady=6)

        self.drop_validation = tk.BooleanVar(value=True)
        self.drop_cells = tk.BooleanVar(value=True)
        self.drop_tabs = tk.BooleanVar(value=True)
        self.drop_commands = tk.BooleanVar(value=True)

        rows = [
            (self.drop_commands,
             "Remove interface command listings",
             "The ribbon is visible on every screen regardless of the "
             "action, so a list such as Insert, Formulas, Data, Review "
             "describes the application rather than the step. A reference "
             "naming one or two commands is kept."),
            (self.drop_validation,
             "Remove sentences stating 'needs Business Validation'",
             "Such a sentence carries no instruction. Table rows whose only "
             "value is the phrase are removed as well."),
            (self.drop_cells,
             "Remove claims about visible cell references",
             "Tokens such as LTO7 and PT053 are data values, not cell "
             "references, so these claims are unreliable."),
            (self.drop_tabs,
             "Remove claims about displayed worksheet tabs",
             "Sheet names are matched anywhere in the screen text, so these "
             "claims often disagree with the screenshot."),
        ]
        for variable, label, detail in rows:
            block = tk.Frame(content, bg="white")
            block.pack(fill="x", padx=12, pady=(6, 0))
            tk.Checkbutton(block, text=label, variable=variable, bg="white",
                           activebackground="white",
                           font=("Segoe UI", 9)).pack(anchor="w")
            tk.Label(block, text=detail, bg="white", fg="gray",
                     font=("Segoe UI", 8), anchor="w", justify="left",
                     wraplength=850).pack(anchor="w", padx=22, pady=(0, 4))

        # ---- build options ----
        opts = tk.LabelFrame(self._body, text=" Build options ", bg="white",
                             fg="#" + BRAND_HEX,
                             font=("Segoe UI", 9, "bold"))
        opts.pack(fill="x", padx=22, pady=6)
        row1 = tk.Frame(opts, bg="white")
        row1.pack(fill="x", padx=12, pady=8)
        self.flow_var = tk.BooleanVar(value=True)
        self.path_var = tk.BooleanVar(value=True)
        self.pdf_var = tk.BooleanVar(value=True)
        tk.Checkbutton(row1, text="Full-page process flow",
                       variable=self.flow_var, bg="white",
                       activebackground="white").pack(side="left")
        tk.Checkbutton(row1, text="List folder paths",
                       variable=self.path_var, bg="white",
                       activebackground="white").pack(side="left", padx=12)
        tk.Checkbutton(row1, text="Create PDF through Word",
                       variable=self.pdf_var, bg="white",
                       activebackground="white").pack(side="left", padx=12)
        tk.Label(row1, text="   Screenshot width:",
                 bg="white").pack(side="left")
        self.width = tk.Spinbox(row1, from_=3.0, to=7.0,
                                increment=0.5, width=5)
        self.width.pack(side="left", padx=(6, 0))
        self.width.delete(0, "end")
        self.width.insert(0, "6.0")

        self.dedupe_var = tk.StringVar(value="auto")

    def initial(self):
        sessions = os.path.join(app_dir(), "sessions")
        return sessions if os.path.isdir(sessions) else app_dir()

    def select(self):
        path = filedialog.askdirectory(title="Select session folder",
                                       initialdir=self.initial())
        if not path:
            return
        self.session = path
        self.folder.config(text=os.path.basename(path), fg="black")

        states = {
            "process_steps.csv": os.path.exists(
                os.path.join(path, "process_steps.csv")),
            "sop_text.txt": bool(read_narrative(path)[0]),
            "process_flow.png": bool(find_flow(path)),
            "screenshots": os.path.isdir(os.path.join(path, "screenshots")),
        }
        for key, found in states.items():
            self.checks[key].config(
                text="Found" if found else "Not found",
                fg=SUCCESS if found else ERROR)

        if states["process_steps.csv"]:
            try:
                steps, has_path = read_steps(path)
                capture = detect_generation(steps, has_path)
                self.capture_label.config(
                    text="Capture: {} step(s), {} distinct screen(s), "
                         "{} folder path(s).".format(
                             len(steps), capture["distinct_shots"],
                             capture["paths"]),
                    fg=SUCCESS)
            except Exception as error:
                self.capture_label.config(
                    text="Capture: could not be read - {}".format(error),
                    fg=ERROR)

        ready = states["process_steps.csv"] and states["sop_text.txt"]
        self.message.config(
            text="Session validated." if ready
            else "process_steps.csv and sop_text.txt are both required.",
            fg=SUCCESS if ready else ERROR)

    def run(self):
        if not self.session:
            messagebox.showwarning(APP_TITLE,
                                   "Select a session folder first.")
            return
        meta = {key: entry.get().strip()
                for key, entry in self.entries.items()}
        try:
            selected = float(self.width.get())
        except ValueError:
            selected = 6.0

        options = {
            "flow": self.flow_var.get(),
            "show_paths": self.path_var.get(),
            "pdf": self.pdf_var.get(),
            "dedupe_mode": self.dedupe_var.get(),
            "width": selected,
            "drop_validation": self.drop_validation.get(),
            "drop_cells": self.drop_cells.get(),
            "drop_tabs": self.drop_tabs.get(),
            "drop_commands": self.drop_commands.get(),
        }

        self.message.config(text="Building SOP...", fg="gray")
        self.root.update_idletasks()
        try:
            result = build(self.session, meta, options)
        except Exception as error:
            self.message.config(text=str(error), fg=ERROR)
            messagebox.showerror(APP_TITLE, str(error))
            return

        created = [os.path.basename(result["docx"])]
        if result["pdf"]:
            created.append(os.path.basename(result["pdf"]))
        tally = result["tally"]
        shots = result["shots"]

        lines = [
            "SOP created.", "",
            "Created:", "- " + "\n- ".join(created), "",
            "Removed from the narrative:",
            "  interface listings     {}".format(tally["commands"]),
            "  validation sentences   {}".format(tally["validation"]),
            "  cell reference claims  {}".format(tally["cells"]),
            "  worksheet tab claims   {}".format(tally["tabs"]),
            "  uninformative rows     {}".format(tally["rows"]),
            "",
            "Screens embedded: {}".format(shots["inserted"]),
            "Folder paths: {}".format(result["capture"]["paths"]),
            "Flow segments: {}".format(result["flow_parts"]),
        ]
        if result["pdf_error"]:
            lines.extend(["", "PDF conversion issue:", result["pdf_error"]])
        lines.extend(["", "Saved in:", self.session])

        self.message.config(text="Created " + ", ".join(created), fg=SUCCESS)
        messagebox.showinfo(APP_TITLE, "\n".join(lines))
        try:
            os.startfile(self.session)
        except Exception:
            pass


if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.mainloop()
