"""AI-Powered Bank Reconciliation & Audit Assistant - deterministic engine.

All numbers are computed here in code. The AI layer (Claude) only explains and recommends,
using summary_for_ai.json produced by this script.

Usage:  python recon_engine.py --data sample_data --out reconciliation.xlsx
"""
import argparse, itertools, json, re, sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# Column mapping - edit these to adapt to another bank/ledger export format
BANK_COLS = dict(date="Date", narration="Narration", ref="Ref_No", debit="Withdrawal", credit="Deposit", balance="Balance")
BOOK_COLS = dict(date="Date", narration="Narration", ref="Voucher_Ref", debit="Payment", credit="Receipt", balance="Balance")

STOP = {"payment", "to", "from", "receipt", "neft", "imps", "upi", "by", "cheque", "chq", "deposit", "towards",
        "ltd", "pvt", "the", "and", "inv", "clearing", "advance", "incl", "gst"}

FLAG_INFO = {  # code: (risk, why it matters, evidence to request)
    "DUPLICATE_PAYMENT": ("High", "Same payment posted twice in books; bank books understated, possible double payment.", "Cheque register, bank clearing proof, voucher copies"),
    "TRANSPOSITION": ("Medium", "Digits transposed between books and bank; books misstated.", "Cheque copy, payment voucher, supplier invoice"),
    "AMOUNT_ERROR": ("Medium", "Books amount differs from bank for the same reference.", "Cheque copy, voucher, supplier confirmation"),
    "UNIDENTIFIED_BANK_ITEM": ("High", "Bank movement with no book entry or explanation.", "Bank advice, authorisation, counterparty details"),
    "STALE_CHEQUE": ("Medium", "Cheque outstanding beyond 90 days; may be stale/void, liability and expense need review.", "Confirmation from payee, cheque register, reissue/cancel decision"),
    "ROUND_SUM": ("Medium", "Large round-sum payment; higher fraud/advance risk.", "Approval, contract/PO, supporting invoice"),
    "NEAR_LIMIT": ("Medium", "Payment just below approval limit; possible limit avoidance.", "Approval matrix, PO, invoice, related payments to same party"),
    "WEEKEND": ("Low", "Entry dated on weekend/holiday; check authority and actual posting date.", "Voucher with approver, system audit log"),
    "SPLIT_PAYMENT": ("High", "Payments to same party split below limit that together exceed it.", "Approvals, invoices, party ledger"),
    "MISSING_NARRATION": ("Low", "No narration; weak audit trail.", "Voucher and supporting document"),
    "MISSING_REF": ("Low", "No cheque/UTR reference in books; matched only on amount, date and narration.", "Bank advice, voucher copy"),
}


def num(x):
    x = str(x).replace(",", "").strip()
    return float(x) if x else 0.0


def similarity(a, b):
    """Token-overlap (Jaccard) on party-name words."""
    t = lambda s: {w for w in re.findall(r"[a-z]+", str(s).lower()) if w not in STOP and len(w) > 2}
    ta, tb = t(a), t(b)
    return len(ta & tb) / len(ta | tb) if ta | tb else 0.0


def standardise(raw, cols, prefix):
    return pd.DataFrame({
        "id": [f"{prefix}{i + 1:03d}" for i in range(len(raw))],
        "date": pd.to_datetime(raw[cols["date"]], format="mixed", dayfirst=True, errors="coerce"),
        "narration": raw[cols["narration"]].str.strip(),
        "ref": raw[cols["ref"]].str.strip(),
        "amount": [num(c) - num(d) for c, d in zip(raw[cols["credit"]], raw[cols["debit"]])],
        "balance": [num(b) if str(b).strip() else float("nan") for b in raw[cols["balance"]]]})


SYNONYMS = {
    "Date": ["date", "txn date", "transaction date", "value date", "tran date", "posting date"],
    "Narration": ["narration", "description", "particulars", "remarks", "transaction remarks", "details", "transaction details"],
    "Ref_No": ["ref_no", "ref no", "ref", "reference", "reference no", "chq/ref no", "chq no", "cheque no", "chq./ref.no.", "utr", "instrument no"],
    "Withdrawal": ["withdrawal", "withdrawals", "withdrawal amt", "withdrawal amount", "debit", "debits", "dr", "debit amount"],
    "Deposit": ["deposit", "deposits", "deposit amt", "deposit amount", "credit", "credits", "cr", "credit amount"],
    "Balance": ["balance", "closing balance", "running balance", "bal"]}


def read_statement(text):
    """Bank statement CSV text -> DataFrame with the canonical columns. Header names are matched loosely."""
    import io
    raw = pd.read_csv(io.StringIO(text.lstrip("\ufeff")), dtype=str).fillna("")
    norm = {re.sub(r"\s+", " ", c.strip().lower()): c for c in raw.columns}
    out, missing = {}, []
    for canon, names in SYNONYMS.items():
        hit = next((norm[n] for n in names if n in norm), None)
        if hit is None and canon not in ("Ref_No", "Balance"):
            missing.append(canon)
        out[canon] = raw[hit].astype(str).str.strip() if hit is not None else pd.Series([""] * len(raw))
    if missing:
        raise ValueError("Statement is missing column(s): " + ", ".join(missing) +
                         ". Expected headers like: Date, Narration, Ref_No, Withdrawal, Deposit, Balance.")
    if not len(raw):
        raise ValueError("The statement has no rows.")
    return pd.DataFrame(out)


def validate(df, name, opening):
    """Data-quality checks: missing fields, duplicates, running balance, opening+net=closing."""
    issues, prev = [], opening
    for r in df.itertuples():
        if pd.isna(r.date):
            issues.append(f"{name} {r.id}: missing/invalid date")
        if not r.narration:
            issues.append(f"{name} {r.id}: blank narration")
        exp = prev + r.amount
        if not pd.isna(r.balance) and abs(exp - r.balance) > 0.01:
            issues.append(f"{name} {r.id}: running balance {r.balance:,.2f} != expected {exp:,.2f}")
        prev = exp if pd.isna(r.balance) else r.balance
    dup = df[df.duplicated(["date", "amount", "ref"], keep=False) & (df.ref != "")]
    for r in dup.itertuples():
        issues.append(f"{name} {r.id}: possible duplicate (same date, amount, ref {r.ref})")
    closing = opening + df.amount.sum()
    return issues, closing


def run_matching(bank, books, tol):
    B, K = bank.set_index("id"), books.set_index("id")
    ub, uk = set(B.index), set(K.index)
    matches, errors = [], []

    def add(kind, conf, bids, kids):
        matches.append(dict(kind=kind, conf=conf, bank_ids=list(bids), book_ids=list(kids)))
        ub.difference_update(bids)
        uk.difference_update(kids)

    def one_to_one(kind, rule, conf_fn):
        for bid in sorted(ub.copy()):
            b = B.loc[bid]
            cands = [k for k in uk if abs(K.loc[k, "amount"] - b["amount"]) < 0.005 and rule(b, K.loc[k])]
            if cands:
                k = min(cands, key=lambda x: (abs((b["date"] - K.loc[x, "date"]).days), x))
                add(kind, conf_fn(b, K.loc[k]), [bid], [k])

    dd = lambda b, k: abs((b["date"] - k["date"]).days)
    one_to_one("Exact (amount + ref + date)", lambda b, k: b["ref"] and b["ref"] == k["ref"] and dd(b, k) == 0, lambda b, k: "High")
    one_to_one("Reference + amount (date lag)", lambda b, k: b["ref"] and b["ref"] == k["ref"] and dd(b, k) <= 7,
               lambda b, k: "High" if dd(b, k) <= 5 else "Medium")
    one_to_one("Near (amount + date + narration)", lambda b, k: dd(b, k) <= tol and similarity(b["narration"], k["narration"]) >= 0.3,
               lambda b, k: "Medium" if similarity(b["narration"], k["narration"]) >= 0.6 else "Low")

    # same reference, same direction, different amount -> amount error
    for bid in sorted(ub.copy()):
        b = B.loc[bid]
        for k in sorted(uk):
            kr = K.loc[k]
            if b["ref"] and b["ref"] == kr["ref"] and b["amount"] * kr["amount"] > 0:
                digits = lambda x: sorted(str(int(round(abs(x)))))
                code = "TRANSPOSITION" if digits(b["amount"]) == digits(kr["amount"]) else "AMOUNT_ERROR"
                errors.append(dict(bank_id=bid, book_id=k, code=code))
                ub.discard(bid)
                uk.discard(k)
                break

    # grouped matches (one bank entry = several book entries, and the reverse)
    def group_pass(single_is_bank):
        S, G = (B, K) if single_is_bank else (K, B)
        sp, gp = (ub, uk) if single_is_bank else (uk, ub)
        for sid in sorted(sp.copy()):
            s = S.loc[sid]
            cands = [g for g in sorted(gp) if G.loc[g, "amount"] * s["amount"] > 0
                     and abs((G.loc[g, "date"] - s["date"]).days) <= 10
                     and similarity(G.loc[g, "narration"], s["narration"]) >= 0.3][:12]
            for n in (2, 3, 4):
                hit = next((c for c in itertools.combinations(cands, n)
                            if abs(sum(G.loc[list(c), "amount"]) - s["amount"]) < 0.005), None)
                if hit:
                    if single_is_bank:
                        add("Many-to-one (1 bank = several books)", "Medium", [sid], hit)
                    else:
                        add("One-to-many (1 books = several bank)", "Medium", hit, [sid])
                    break
    group_pass(True)
    group_pass(False)
    return matches, errors, sorted(ub), sorted(uk)


def classify(bank, books, opening, unb, unk, errors, cfg):
    B, K = bank.set_index("id"), books.set_index("id")
    items = []

    def item(cat, src, r, remark, books_adj=0.0, bank_adj=0.0, code=""):
        items.append(dict(category=cat, source=src, date=r["date"], narration=r["narration"], ref=r["ref"],
                          amount=r["amount"], remark=remark, books_adj=books_adj, bank_adj=bank_adj, code=code))

    unb = list(unb)
    for _, o in opening.iterrows():                      # brought-forward items
        hit = next((b for b in unb if B.loc[b, "ref"] == o["ref"] and abs(B.loc[b, "amount"] - o["amount"]) < 0.005), None)
        if hit:
            unb.remove(hit)                              # cleared in this period
        else:
            item("Outstanding cheque (B/F)", "Books (B/F)", o, "Brought forward; not presented this period",
                 bank_adj=o["amount"], code="OUTSTANDING_CHEQUE_BF")
    for b in unb:
        r = B.loc[b]
        txt = r["narration"].lower()
        if r["amount"] < 0 and re.search(r"charge|chg|fee|gst", txt):
            item("Bank charges not in books", "Bank", r, "Pass journal entry", books_adj=r["amount"], code="BANK_CHARGE")
        elif r["amount"] > 0 and "interest" in txt:
            item("Bank interest not in books", "Bank", r, "Pass journal entry", books_adj=r["amount"], code="BANK_INTEREST")
        else:
            item("Unidentified bank item", "Bank", r, "Investigate", books_adj=r["amount"], code="UNIDENTIFIED_BANK_ITEM")
    for k in unk:
        r = K.loc[k]
        dup = ((books.id != k) & (books.amount == r["amount"]) & (books.ref == r["ref"]) & (books.date == r["date"]) & (r["ref"] != "")).any()
        if dup:
            item("Duplicate posting in books", "Books", r, "Reverse duplicate entry", books_adj=-r["amount"], code="DUPLICATE_PAYMENT")
        elif r["amount"] < 0:
            item("Outstanding cheque", "Books", r, "Issued, not yet presented", bank_adj=r["amount"], code="OUTSTANDING_CHEQUE")
        else:
            item("Deposit in transit", "Books", r, "Recorded, not yet credited", bank_adj=r["amount"], code="DEPOSIT_IN_TRANSIT")
    for e in errors:
        b, k = B.loc[e["bank_id"]], K.loc[e["book_id"]]
        r = dict(b)
        r["narration"] = f"{k['narration']} | books {k['amount']:,.2f} vs bank {b['amount']:,.2f}"
        item("Amount error in books", "Books vs Bank", r, e["code"].title().replace("_", " ") + " - correct books",
             books_adj=b["amount"] - k["amount"], code=e["code"])
    return items


def audit_flags(books, items, cfg, period_end):
    flags, limit = [], cfg["approval_limit"]

    def add(code, r):
        risk, why, ev = FLAG_INFO[code]
        flags.append(dict(code=code, ref=r["ref"], date=r["date"], narration=r["narration"], amount=r["amount"],
                          risk=risk, issue=why, evidence=ev))

    for r in books.to_dict("records"):
        a = abs(r["amount"])
        if a >= 100000 and a % 10000 == 0:
            add("ROUND_SUM", r)
        if r["amount"] < 0 and 0.9 * limit <= a < limit:
            add("NEAR_LIMIT", r)
        if r["date"].weekday() >= 5:
            add("WEEKEND", r)
        if not r["narration"]:
            add("MISSING_NARRATION", r)
        if not r["ref"]:
            add("MISSING_REF", r)
    pays = [r for r in books.to_dict("records") if r["amount"] < 0]
    for x, y in itertools.combinations(pays, 2):
        if (abs((x["date"] - y["date"]).days) <= 3 and x["amount"] != y["amount"] and abs(x["amount"]) < limit
                and abs(y["amount"]) < limit and abs(x["amount"] + y["amount"]) >= limit
                and similarity(x["narration"], y["narration"]) >= 0.5):
            add("SPLIT_PAYMENT", x)
    for it in items:
        if it["category"].startswith("Outstanding cheque") and (period_end - it["date"]).days > cfg["stale_days"]:
            add("STALE_CHEQUE", it)
        if it["code"] in ("DUPLICATE_PAYMENT", "TRANSPOSITION", "AMOUNT_ERROR", "UNIDENTIFIED_BANK_ITEM"):
            add(it["code"], it)
    return flags


def detected_set(matches, items, flags, bank):
    B = bank.set_index("id")
    d = {(f["code"], f["ref"]) for f in flags}
    d |= {(i["code"], i["ref"]) for i in items}
    for m in matches:
        if m["kind"].startswith("Many-to-one"):
            d.add(("MANY_TO_ONE", B.loc[m["bank_ids"][0], "ref"]))
        if m["kind"].startswith("Near"):
            d.add(("NEAR_MATCH", B.loc[m["bank_ids"][0], "ref"]))
    return d


def evaluate(detected, key_path):
    key = pd.read_csv(key_path, dtype=str).fillna("")
    truth = set(zip(key.code, key.ref))
    tp, fn = truth & detected, truth - detected
    scored_codes = {c for c, _ in truth}
    fp = {x for x in detected if x[0] in scored_codes and x not in truth}
    return dict(planted=len(truth), detected=len(tp), missed=sorted(fn), false_positives=sorted(fp),
                recall=len(tp) / len(truth), precision=len(tp) / (len(tp) + len(fp)) if tp or fp else 0)


# ------------------------------------------------------------------ Excel output
HDR = PatternFill("solid", fgColor="1F3864")
F = lambda **k: Font(**{"name": "Arial", "size": 10, **k})
NUM = "#,##0.00;(#,##0.00);-"


def sheet(wb, name, headers, rows, widths, num_cols=(), date_cols=()):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = F(bold=True, color="FFFFFF"), HDR, Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F()
            if c.column in num_cols:
                c.number_format = NUM
            if c.column in date_cols:
                c.number_format = "DD-MM-YYYY"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    return ws


def write_excel(path, cfg, bank, books, matches, items, flags, metrics, bank_close, books_close, entries=None, approval=None):
    B, K = bank.set_index("id"), books.set_index("id")
    wb = Workbook()
    wb.remove(wb.active)
    # Summary
    ws = wb.create_sheet("Summary")
    ws["A1"], ws["A1"].font = f"Bank Reconciliation & Audit Summary - {cfg['entity']} - {cfg['period']}" + (f" - {cfg['bank']}" if cfg.get("bank") else ""), F(bold=True, size=12)
    ws["A2"], ws["A2"].font = "SYNTHETIC DATA ONLY. Figures computed by code; AI commentary is advisory. Blue = input values, black = formulas.", F(italic=True)
    rows = [("Bank transactions", metrics["bank_n"]), ("Books transactions", metrics["books_n"]),
            ("Match groups found", "=COUNTA(Matched!A:A)-1"), ("Bank transactions matched", metrics["bank_matched"]),
            ("Bank match rate", "=B7/B4"), ("Reconciling items (unmatched)", "=COUNTA(Unmatched!A:A)-1"),
            ("Audit exceptions raised", "=COUNTA(Exceptions!A:A)-1"), ("  of which High risk", '=COUNTIF(Exceptions!F:F,"High")'),
            ("Unreconciled difference (Rs)", "=BRS!B20"), ("Status", "=BRS!B21")]
    ws["A3"], ws["B3"] = "Metric", "Value"
    for c in ws[3]:
        c.font, c.fill = F(bold=True, color="FFFFFF"), HDR
    for i, (k, v) in enumerate(rows, 4):
        ws.cell(i, 1, k).font = F()
        c = ws.cell(i, 2, v)
        c.font = F(color="0000FF") if not str(v).startswith("=") else F()
    ws["B8"].number_format, ws["B12"].number_format = "0.0%", NUM
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 38, 28
    # Matched
    rows = []
    for n, m in enumerate(matches, 1):
        bs, ks = [B.loc[i] for i in m["bank_ids"]], [K.loc[i] for i in m["book_ids"]]
        lag = abs((bs[0]["date"] - ks[0]["date"]).days) if len(bs) == len(ks) == 1 else None
        rows.append([f"M{n:03d}", m["kind"], m["conf"], bs[0]["date"], "; ".join(x["ref"] for x in bs),
                     "; ".join(x["narration"] for x in bs), sum(x["amount"] for x in bs),
                     ks[0]["date"], "; ".join(x["ref"] or "(blank)" for x in ks),
                     "; ".join(x["narration"] for x in ks), sum(x["amount"] for x in ks), lag])
    sheet(wb, "Matched", ["Match ID", "Match type", "Confidence", "Bank date", "Bank ref", "Bank narration", "Bank amount",
                          "Books date", "Books ref", "Books narration", "Books amount", "Date gap (days)"],
          rows, [9, 34, 11, 12, 18, 40, 14, 12, 18, 40, 14, 10], num_cols=(7, 11), date_cols=(4, 8))
    # Unmatched
    rows = [[i["source"], i["date"], i["narration"], i["ref"], i["amount"], i["category"], i["remark"], i["books_adj"], i["bank_adj"]] for i in items]
    sheet(wb, "Unmatched", ["Source", "Date", "Narration", "Ref", "Amount (as recorded)", "Category", "Action",
                            "Adjustment to Books", "Adjustment to Bank"], rows, [14, 12, 50, 14, 16, 28, 30, 16, 16],
          num_cols=(5, 8, 9), date_cols=(2,))
    # BRS
    ws = wb.create_sheet("BRS")
    sums = lambda col, cat: f'=SUMIFS(Unmatched!${col}:${col},Unmatched!$F:$F,"{cat}")'
    lines = [(1, "BANK RECONCILIATION STATEMENT", None), (2, f"{cfg['entity']} | {cfg['period']} | " + (f"{cfg['bank']} | " if cfg.get("bank") else "") + "Currency INR", None),
             (4, "A. From bank side", None), (5, "Balance as per Bank Statement", round(bank_close, 2)),
             (6, "Add: Deposits in transit", sums("I", "Deposit in transit")),
             (7, "Less: Outstanding cheques (current period)", sums("I", "Outstanding cheque")),
             (8, "Less: Outstanding cheques (brought forward)", sums("I", "Outstanding cheque (B/F)")),
             (9, "Adjusted Bank Balance", "=SUM(B5:B8)"),
             (11, "B. From books side", None), (12, "Balance as per Bank Books", round(books_close, 2)),
             (13, "Add: Bank interest not in books", sums("H", "Bank interest not in books")),
             (14, "Less: Bank charges not in books", sums("H", "Bank charges not in books")),
             (15, "Add back: Duplicate posting in books", sums("H", "Duplicate posting in books")),
             (16, "Correct: Amount errors in books", sums("H", "Amount error in books")),
             (17, "Unidentified bank items", sums("H", "Unidentified bank item")),
             (18, "Adjusted Books Balance", "=SUM(B12:B17)"),
             (20, "Difference (A - B)", "=B9-B18"),
             (21, "Status", '=IF(ABS(B20)<0.01,"RECONCILED","UNRECONCILED - INVESTIGATE")'),
             (23, "Note: Rows 5 and 12 are closing balances from the input files (blue). All other lines are SUMIFS over the Unmatched sheet.", None)]
    for r, a, b in lines:
        ws.cell(r, 1, a).font = F(bold=r in (1, 4, 9, 11, 18, 20, 21))
        if b is not None:
            c = ws.cell(r, 2, b)
            c.number_format, c.font = NUM, F(color="0000FF") if r in (5, 12) else F(bold=r in (9, 18, 20, 21))
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 52, 26
    # Exceptions
    order = {"High": 0, "Medium": 1, "Low": 2}
    rows = [[f["code"], f["ref"] or "(blank)", f["date"], f["narration"], f["amount"], f["risk"], f["issue"], f["evidence"]]
            for f in sorted(flags, key=lambda x: (order[x["risk"]], x["code"]))]
    sheet(wb, "Exceptions", ["Flag", "Ref", "Date", "Narration", "Amount", "Risk", "Why it matters", "Evidence to request"],
          rows, [22, 16, 12, 50, 14, 9, 60, 50], num_cols=(5,), date_cols=(3,))
    if entries is not None:
        dec = (approval or {}).get("decisions", [])
        sheet(wb, "Journal entries", ["#", "Debit", "Credit", "Amount", "Reason", "Decision"],
              [[n, e["dr"], e["cr"], e["amount"], e["why"], dec[n - 1] if n <= len(dec) else "Pending"] for n, e in enumerate(entries, 1)],
              [5, 34, 34, 16, 50, 12], num_cols=(4,))
        ws = wb["Journal entries"]
        ws.cell(len(entries) + 3, 1, "Proposed entries - subject to CA approval. Only entries marked Approved should be posted.").font = F(italic=True)
    if approval:
        ws = wb.create_sheet("Approval")
        rows = [("Status", "APPROVED"), ("Approved by", approval["reviewer"]), ("Approved on", approval["approved_at"]),
                ("Entity", cfg.get("entity", "")), ("Bank", cfg.get("bank", "")), ("Period", cfg.get("period", "")),
                ("Difference at approval (Rs)", approval["difference"]),
                ("Journal entries approved / rejected", f"{approval['n_approved']} / {approval['n_rejected']}"),
                ("Audit exceptions (High / total)", f"{approval['n_high']} / {approval['n_flags']}"),
                ("Reviewer remarks", approval.get("remarks") or "-")]
        ws.append(["Item", "Value"])
        for c in ws[1]:
            c.font, c.fill = F(bold=True, color="FFFFFF"), HDR
        for r in rows:
            ws.append(list(r))
        for row in ws.iter_rows(min_row=2):
            row[0].font, row[1].font = F(bold=True), F()
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
        ws["B8"].number_format = NUM
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 38, 70
        wb.move_sheet("Approval", offset=-(len(wb.sheetnames) - 1))
    wb.save(path)


def fdate(d):
    return d.strftime("%d-%m-%Y") if not pd.isna(d) else ""


def propose_entries(items, bank_ledger="Bank A/c"):
    """Proposed adjusting journal entries (subject to CA approval) for items that must be booked."""
    out = []
    for i in items:
        a, d = abs(i["books_adj"]), fdate(i["date"])
        c = i["category"]
        if c == "Bank charges not in books":
            out.append(dict(dr="Bank Charges A/c", cr=bank_ledger, amount=a, why=f"Charge {i['ref']} on {d}"))
        elif c == "Bank interest not in books":
            out.append(dict(dr=bank_ledger, cr="Interest Income A/c", amount=a, why=f"Interest {i['ref']} on {d}"))
        elif c == "Duplicate posting in books":
            out.append(dict(dr=bank_ledger, cr="Supplier A/c (reverse duplicate)", amount=a, why=f"{i['ref']}: {i['narration'].split('|')[0].strip()}"))
        elif c == "Amount error in books":
            dr, cr = ("Supplier A/c", bank_ledger) if i["books_adj"] < 0 else (bank_ledger, "Supplier A/c")
            out.append(dict(dr=dr, cr=cr, amount=a, why=f"{i['ref']} {i['code'].title().replace('_', ' ')}"))
        elif c == "Unidentified bank item":
            dr, cr = ("Suspense A/c", bank_ledger) if i["books_adj"] < 0 else (bank_ledger, "Suspense A/c")
            out.append(dict(dr=dr, cr=cr, amount=a, why=f"{i['ref']}: {i['narration']}"))
    return out


def reconcile(bank, books, opening, bank_opening, books_opening, cfg):
    """Core pipeline on standardised frames -> run dict (everything the UI, Excel and AI summary need)."""
    period_end = pd.to_datetime(cfg["period_end"], format="%d-%m-%Y")
    vb, bank_close = validate(bank, "Bank", bank_opening)
    vk, books_close = validate(books, "Books", books_opening)
    matches, errors, unb, unk = run_matching(bank, books, cfg["date_tolerance"])
    items = classify(bank, books, opening, unb, unk, errors, cfg)
    flags = audit_flags(books, items, cfg, period_end)
    bank_adj = bank_close + sum(i["bank_adj"] for i in items)
    books_adj = books_close + sum(i["books_adj"] for i in items)
    sumf = lambda f, cat: sum(i[f] for i in items if i["category"] == cat)
    brs = dict(bank_close=bank_close, books_close=books_close, adj_bank=bank_adj, adj_books=books_adj,
               dit=sumf("bank_adj", "Deposit in transit"), oc=sumf("bank_adj", "Outstanding cheque"),
               ocbf=sumf("bank_adj", "Outstanding cheque (B/F)"), interest=sumf("books_adj", "Bank interest not in books"),
               charges=sumf("books_adj", "Bank charges not in books"), dup=sumf("books_adj", "Duplicate posting in books"),
               err=sumf("books_adj", "Amount error in books"), unid=sumf("books_adj", "Unidentified bank item"))
    matched_bank = sum(len(m["bank_ids"]) for m in matches)
    return dict(cfg=cfg, bank=bank, books=books, opening=opening, matches=matches, items=items, flags=flags,
                issues=vb + vk, brs=brs, diff=round(bank_adj - books_adj, 2),
                metrics=dict(bank_n=len(bank), books_n=len(books), bank_matched=matched_bank),
                entries=propose_entries(items, cfg.get("bank") or "Bank A/c"), detected=detected_set(matches, items, flags, bank))


def ai_summary(run):
    c, b = run["cfg"], run["brs"]
    return dict(entity=c.get("entity", ""), bank=c.get("bank", ""), period=c.get("period", ""),
                note="Figures computed by code. Do not recompute.",
                balances=dict(bank_closing=b["bank_close"], books_closing=b["books_close"], adjusted_bank=b["adj_bank"],
                              adjusted_books=b["adj_books"], difference=run["diff"]),
                metrics=dict(run["metrics"], bank_match_rate=round(run["metrics"]["bank_matched"] / max(run["metrics"]["bank_n"], 1), 3)),
                validation_issues=run["issues"],
                reconciling_items=[dict(category=i["category"], date=fdate(i["date"]), ref=i["ref"], amount=i["amount"],
                                        narration=i["narration"]) for i in run["items"]],
                exceptions=[dict(flag=f["code"], risk=f["risk"], ref=f["ref"], date=fdate(f["date"]), amount=f["amount"],
                                 narration=f["narration"]) for f in run["flags"]])


def to_ui(run):
    """JSON-friendly view of a run for the web app."""
    B, K = run["bank"].set_index("id"), run["books"].set_index("id")
    side = lambda df, ids: [dict(date=fdate(df.loc[i, "date"]), ref=df.loc[i, "ref"], narration=df.loc[i, "narration"],
                                 amount=float(df.loc[i, "amount"])) for i in ids]
    return dict(
        bank=run["cfg"].get("bank", ""), entity=run["cfg"].get("entity", ""), period=run["cfg"].get("period", ""),
        metrics=dict(run["metrics"], items=len(run["items"]), flags=len(run["flags"]),
                     high=sum(f["risk"] == "High" for f in run["flags"])),
        diff=float(run["diff"]), reconciled=bool(abs(run["diff"]) < 0.01), brs={k: float(v) for k, v in run["brs"].items()},
        matches=[dict(id=f"M{n:03d}", kind=m["kind"], conf=m["conf"], bank=side(B, m["bank_ids"]), books=side(K, m["book_ids"]))
                 for n, m in enumerate(run["matches"], 1)],
        items=[dict(category=i["category"], source=i["source"], date=fdate(i["date"]), narration=i["narration"], ref=i["ref"],
                    amount=float(i["amount"]), action=i["remark"], books_adj=float(i["books_adj"]), bank_adj=float(i["bank_adj"]))
               for i in run["items"]],
        flags=[dict(code=f["code"], ref=f["ref"], date=fdate(f["date"]), narration=f["narration"], amount=float(f["amount"]),
                    risk=f["risk"], issue=f["issue"], evidence=f["evidence"]) for f in run["flags"]],
        entries=[dict(e, amount=float(e["amount"])) for e in run["entries"]], issues=run["issues"], ai=ai_summary(run))


def export_run(path, run, approval=None):
    write_excel(path, run["cfg"], run["bank"], run["books"], run["matches"], run["items"], run["flags"], run["metrics"],
                run["brs"]["bank_close"], run["brs"]["books_close"], entries=run["entries"], approval=approval)


def safe_name(n):
    """Folder/file-friendly version of a bank ledger name."""
    return re.sub(r"[^A-Za-z0-9]+", "_", n).strip("_").lower()


def choose_bank(banks, wanted=None):
    """Pick one bank ledger from the Tally export: --bank NAME, auto if only one, else numbered menu."""
    names = [b["name"] for b in banks]
    if wanted:
        hits = [b for b in banks if wanted.lower() == b["name"].lower()] or \
               [b for b in banks if wanted.lower() in b["name"].lower()]
        if len(hits) != 1:
            raise SystemExit(f"--bank '{wanted}' matched {len(hits)} banks. Available: " + "; ".join(names))
        return hits[0]
    if len(banks) == 1:
        return banks[0]
    if not sys.stdin.isatty():
        raise SystemExit("Several banks found; pass --bank NAME. Available: " + "; ".join(names))
    print("\nBank accounts found (Current Assets > Bank Accounts):")
    for i, b in enumerate(banks, 1):
        print(f"  {i}. {b['name']}   (opening {b['opening_balance']:,.2f}, {len(b['rows'])} entries)")
    while True:
        pick = input("Select bank number: ").strip()
        if pick.isdigit() and 1 <= int(pick) <= len(banks):
            return banks[int(pick) - 1]
        print("Invalid choice, try again.")


DEFAULT_CFG = dict(entity="", period="", approval_limit=50000, stale_days=90, date_tolerance=3)


def parse_opening_items(text):
    """B/F outstanding items CSV text (Date, Narration, Ref, Amount) -> DataFrame (may be empty)."""
    import io
    cols = ["date", "narration", "ref", "amount"]
    if not text or not text.strip():
        return pd.DataFrame(columns=cols)
    op = pd.read_csv(io.StringIO(text.lstrip("\ufeff")), dtype=str).fillna("")
    if not {"Date", "Narration", "Ref", "Amount"} <= set(op.columns):
        raise ValueError("B/F items file needs columns: Date, Narration, Ref, Amount.")
    return pd.DataFrame({"date": pd.to_datetime(op.Date, format="mixed", dayfirst=True), "narration": op.Narration,
                         "ref": op.Ref, "amount": [num(x) for x in op.Amount]})


def run_from_inputs(statement_text, books_rows, books_opening, opening_text, cfg, bank_opening=None):
    """Raw inputs -> run dict. books_rows: DataFrame or list of dicts with the bank-books columns."""
    bank = standardise(read_statement(statement_text), BANK_COLS, "B")
    books_df = books_rows if isinstance(books_rows, pd.DataFrame) else pd.DataFrame(books_rows, columns=list(BOOK_COLS.values()))
    books = standardise(books_df.astype(str).fillna(""), BOOK_COLS, "K")
    if bank_opening is None:
        first = bank.balance.iloc[0]
        bank_opening = (first - bank.amount.iloc[0]) if not pd.isna(first) else 0.0
    return reconcile(bank, books, parse_opening_items(opening_text), float(bank_opening), float(books_opening), cfg)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Bank reconciliation & audit engine")
    ap.add_argument("--data", default="sample_data", help="folder with config/statements (and optional tally_bank_books.json)")
    ap.add_argument("--tally", help="Tally export JSON from tally_extract.py (default: <data>/tally_bank_books.json if present)")
    ap.add_argument("--bank", help="Bank ledger name (asked interactively if several and not given)")
    ap.add_argument("--statement", help="Bank statement CSV (default: <data>/banks/<bank>/bank_statement.csv)")
    ap.add_argument("--opening-items", help="B/F outstanding items CSV")
    ap.add_argument("--out", help="Excel output path")
    a = ap.parse_args(argv)
    d = Path(a.data)
    cfg = dict(DEFAULT_CFG)
    if (d / "config.json").exists():
        cfg.update(json.load(open(d / "config.json")))
    tally = Path(a.tally) if a.tally else (d / "tally_bank_books.json" if (d / "tally_bank_books.json").exists() else None)

    if tally:                                            # ---- Bank books auto-extracted from Tally
        tj = json.load(open(tally, encoding="utf-8"))
        chosen = choose_bank(tj["banks"], a.bank)
        sn = safe_name(chosen["name"])
        books_raw = pd.DataFrame(chosen["rows"], columns=list(BOOK_COLS.values())).astype(str).fillna("")
        books_opening = chosen["opening_balance"]
        cfg["bank"] = chosen["name"]
        cfg["entity"] = tj.get("company") or cfg["entity"]
        cfg["period"] = f"{tj['period_from']} to {tj['period_to']}"
        cfg["period_end"] = tj["period_to"]
        stmt = Path(a.statement) if a.statement else next(
            (p for p in (d / "banks" / sn / "bank_statement.csv", d / "bank_statement.csv") if p.exists()), None)
        if stmt is None:
            raise SystemExit(f"Bank statement not found. Put it at {d / 'banks' / sn / 'bank_statement.csv'} or use --statement.")
        oi = Path(a.opening_items) if a.opening_items else d / "banks" / sn / "opening_items.csv"
        out = a.out or f"reconciliation_{sn}.xlsx"
        summary_path = f"summary_for_ai_{sn}.json"
    else:                                                # ---- Legacy: bank_books.csv in data folder
        books_raw = pd.read_csv(d / "bank_books.csv", dtype=str).fillna("")
        books_opening = cfg["books_opening"]
        stmt = Path(a.statement) if a.statement else d / "bank_statement.csv"
        oi = Path(a.opening_items) if a.opening_items else d / "opening_items.csv"
        out = a.out or "reconciliation.xlsx"
        summary_path = "summary_for_ai.json"

    run = run_from_inputs(Path(stmt).read_text(encoding="utf-8-sig"), books_raw, books_opening,
                          oi.read_text(encoding="utf-8-sig") if oi.exists() else "", cfg,
                          bank_opening=None if tally else cfg.get("bank_opening"))
    export_run(out, run)
    Path("last_output.txt").write_text(str(out), encoding="utf-8")
    Path(summary_path).write_text(json.dumps(ai_summary(run), indent=2, default=float), encoding="utf-8")
    key_ok = (d / "answer_key.csv").exists() and cfg.get("answer_key_bank", cfg.get("bank")) == cfg.get("bank")
    ev = evaluate(run["detected"], d / "answer_key.csv") if key_ok else None
    bank_close, books_close, diff = run["brs"]["bank_close"], run["brs"]["books_close"], run["diff"]
    bank_adj, books_adj = run["brs"]["adj_bank"], run["brs"]["adj_books"]
    matches, items, flags = run["matches"], run["items"], run["flags"]
    vb, vk = run["issues"], []

    if cfg.get("bank"):
        print(f"Bank: {cfg['bank']}")
    print(f"Bank closing {bank_close:,.2f} | Books closing {books_close:,.2f}")
    print(f"Adjusted bank {bank_adj:,.2f} | Adjusted books {books_adj:,.2f} | Difference {diff:,.2f}")
    print(f"Matches {len(matches)} | Reconciling items {len(items)} | Exceptions {len(flags)} | Validation issues {len(vb + vk)}")
    if ev:
        print(f"Accuracy vs answer key: detected {ev['detected']}/{ev['planted']} (recall {ev['recall']:.0%}, precision {ev['precision']:.0%})")
        print("Missed:", ev["missed"], "| False positives:", ev["false_positives"])
    print(f"Excel: {out} | AI summary: {summary_path}")
    return dict(diff=diff, evaluation=ev, bank=cfg.get("bank"))


if __name__ == "__main__":
    main()
