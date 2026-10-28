"""
redactor_core.py - format-preserving redaction & restore engine.

redact_file(src, rules_xlsx)  -> writes <name>_REDACTED.<ext> + <name>_MAP.json next to the source
restore_file(src, map_json)   -> writes <name>_RESTORED.<ext>

Formats: .pdf (text PDFs)  .docx  .xlsx  .txt/.csv/.md
"""
import re, json
from pathlib import Path

NBSP = " "
TEXT_EXT = (".txt", ".csv", ".md")
HL_HEX = {"green": "C6EFCE", "blue": "BDD7EE"}
HL_RGB = {"green": (0.78, 0.94, 0.81), "blue": (0.74, 0.84, 0.93)}
GREEN_HEX, YELLOW_HEX = "C6EFCE", "FFF176"      # restored words / filled-in values
COLOUR_SET = {"C6EFCE", "FFF176", "BDD7EE", "FFFF00"}
G_ON, G_OFF, Y_ON, Y_OFF = "\ue000", "\ue001", "\ue002", "\ue003"   # colour marks inside plain text


# ------------------------------------------------------------------ helpers
def _norm(s):
    return re.sub(r"\s+", " ", s).strip().lower()


def _flex(s):
    """regex for s where any whitespace run matches any whitespace run"""
    return r"\s+".join(re.escape(p) for p in s.split())


# ------------------------------------------------------------------ rules
class Rules:
    def __init__(self):
        self.patterns = []   # (label, compiled regex)
        self.manual = []     # (original, replacement, marked)
        self.open = "[["
        self.close = "]]"
        self.pad = False
        self.hl = "green"


def _yes(v, default=True):
    if v is None or str(v).strip() == "":
        return default
    x = str(v).strip().upper()
    if x in ("Y", "YES", "1", "TRUE", "ON"):
        return True
    if x in ("N", "NO", "0", "FALSE", "OFF"):
        return False
    return default


def _cols(ws, wanted):
    """Find columns by header text (so extra columns, e.g. a Category, do no harm). wanted={key:(prefixes,)}"""
    header = [str(c.value or "").strip().lower() for c in ws[1]]
    idx = {}
    for key, prefixes in wanted.items():
        for i, h in enumerate(header):
            if h and any(h.startswith(p) for p in prefixes):
                idx[key] = i
                break
    return idx, len(header)


def _lab(x):
    return re.sub(r"[\s_]+", "", str(x or "")).lower()


DEFAULT_PATTERNS = [
    ("PAN", r"\b[A-Z]{5}[0-9]{4}[A-Z]\b", "Y", "5 letters, 4 digits, 1 letter"),
    ("GSTIN", r"\b[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]\b", "Y", "15 characters"),
    ("TAN", r"\b[A-Z]{4}[0-9]{5}[A-Z]\b", "Y", ""),
    ("AADHAAR", r"(?<!\d)[0-9]{4}\s?[0-9]{4}\s?[0-9]{4}(?!\d)", "Y", "12 digits"),
    ("EMAIL", r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", "Y", ""),
    ("MOBILE", r"(?<!\d)(?:\+91[\s-]?)?[6-9][0-9]{9}(?!\d)", "Y", "10-digit Indian mobile"),
    ("CIN", r"\b[LU][0-9]{5}[A-Z]{2}[0-9]{4}[A-Z]{3}[0-9]{6}\b", "Y", "Company CIN"),
    ("ITBA Ref No", r"\bITBA(?:/[A-Za-z0-9\-]+(?:\([A-Za-z0-9]+\))*){3,6}", "Y",
     "Income-tax order / notice / DIN references like ITBA/AST/S/143(3)/2021-22/1037213701(1)"),
    ("PINCODE", r"(?<!\d)[1-9][0-9]{5}(?!\d)", "N", "Turn on if needed"),
    ("DIN", r"(?<!\d)[0-9]{8}(?!\d)", "N", "Director ID - may clash with other numbers"),
    ("DOCNO", r"\b[A-Z]{2}[0-9]{13,16}\b", "N", "ARN / DIN-type references"),
]

_ACRONYMS = {"PAN", "GSTIN", "TAN", "CIN", "DIN", "ITBA", "ARN", "TDS", "GST", "TRACES", "IFSC"}


def _pretty_label(label):
    l = str(label).replace("_", " ").strip()
    return l.title() if l.isupper() and l not in _ACRONYMS and len(l) > 3 else l


def add_missing_patterns(xlsx):
    """Append standard patterns that are not yet in the user's Patterns sheet (matched by label).
    A backup copy of the rules file is saved first. Returns the labels added."""
    import shutil, time
    from openpyxl import load_workbook
    xlsx = Path(xlsx)
    wb = load_workbook(str(xlsx))
    if "Patterns" not in wb.sheetnames:
        wb.create_sheet("Patterns").append(["Label", "Regex", "Enabled (Y/N)", "Note"])
    ws = wb["Patterns"]
    idx, ncols = _cols(ws, {"label": ("label",), "regex": ("regex",), "enabled": ("enabled",), "note": ("note",)})
    li, ri = idx.get("label", 0), idx.get("regex", 1)
    have = {_lab(ws.cell(row=r, column=li + 1).value) for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=li + 1).value}
    missing = [p for p in DEFAULT_PATTERNS if _lab(p[0]) not in have]
    if not missing:
        return []
    shutil.copy2(str(xlsx), str(xlsx.with_name(f"{xlsx.stem}_backup_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")))
    row = ws.max_row + 1
    for label, regex, enabled, note in missing:
        ws.cell(row=row, column=li + 1, value=label)
        ws.cell(row=row, column=ri + 1, value=regex)
        if "enabled" in idx:
            ws.cell(row=row, column=idx["enabled"] + 1, value=enabled)
        if "note" in idx:
            ws.cell(row=row, column=idx["note"] + 1, value=note)
        row += 1
    wb.save(str(xlsx))
    return [p[0] for p in missing]


def load_rules(xlsx):
    from openpyxl import load_workbook
    wb = load_workbook(str(xlsx), data_only=True)
    r = Rules()

    if "Settings" in wb.sheetnames:
        for row in wb["Settings"].iter_rows(min_row=2, values_only=True):
            k, v = (tuple(row) + (None, None))[:2]
            if not k:
                continue
            k = str(k).strip().lower()
            if k.startswith("marker start") and v is not None:
                r.open = str(v)
            elif k.startswith("marker end") and v is not None:
                r.close = str(v)
            elif k.startswith("keep length"):
                r.pad = _yes(v, False)
            elif k.startswith("highlight") and v is not None:
                x = str(v).strip().lower()
                r.hl = "none" if x in ("", "n", "no", "none", "off", "0", "false") else ("blue" if x.startswith("b") else "green")

    if "Patterns" in wb.sheetnames:
        ws = wb["Patterns"]
        idx, _n = _cols(ws, {"label": ("label",), "regex": ("regex",), "enabled": ("enabled",)})
        li, ri, ei = idx.get("label", 0), idx.get("regex", 1), idx.get("enabled")
        for row in ws.iter_rows(min_row=2, values_only=True):
            row = tuple(row) + (None,) * 8
            label, regex = row[li], row[ri]
            enabled = row[ei] if ei is not None else None
            if not label or not regex or not _yes(enabled):
                continue
            try:
                r.patterns.append((str(label).strip().replace(" ", "_"), re.compile(str(regex).strip())))
            except re.error as e:
                raise ValueError(f"Bad regex for '{label}' in Patterns sheet: {e}")

    if "Replacements" in wb.sheetnames:
        ws = wb["Replacements"]
        idx, _n = _cols(ws, {"orig": ("original",), "repl": ("replace",), "mark": ("mark",)})
        oi, pi, mi = idx.get("orig", 0), idx.get("repl", 1), idx.get("mark")
        for row in ws.iter_rows(min_row=2, values_only=True):
            row = tuple(row) + (None,) * 8
            orig, repl = row[oi], row[pi]
            mark = row[mi] if mi is not None else None
            if orig is None or repl is None or str(orig).strip() == "":
                continue
            r.manual.append((str(orig).strip(), str(repl).strip(), _yes(mark, True)))
    return r


# ------------------------------------------------------------------ items
def build_items(full_text, rules):
    items, seen = [], set()

    def add(original, core, marked, kind):
        k = _norm(original)
        if not k or k in seen:
            return
        seen.add(k)
        items.append({"original": original.strip(), "core": core, "marked": marked, "kind": kind})

    for orig, repl, mark in rules.manual:             # the user's own rules win
        add(orig, repl, mark, "rule")
    groups = {}
    for label, rx in rules.patterns:
        for m in rx.finditer(full_text):
            v = m.group(0).strip()
            k = _norm(v)
            if not k or k in seen:
                continue
            seen.add(k)
            groups.setdefault(label, []).append(v)
    for label, vals in groups.items():                # [[Redacted PAN]]; several different ones -> [[Redacted PAN 1]], [[Redacted PAN 2]]
        nice = _pretty_label(label)
        for n, v in enumerate(vals, 1):
            items.append({"original": v, "core": f"Redacted {nice}" if len(vals) == 1 else f"Redacted {nice} {n}",
                          "marked": True, "kind": "pattern"})
    return items


def display(it, rules):
    return f"{rules.open}{it['core']}{rules.close}" if it["marked"] else it["core"]


def display_like(it, rules, matched):
    """display() but keeps a leading capital if the original had one."""
    d = display(it, rules)
    core = it["core"]
    if matched[:1].isupper() and core[:1].islower():
        i = d.find(core)
        if i >= 0:
            d = d[:i] + core[:1].upper() + core[1:] + d[i + len(core):]
    return d


def compile_items(items):
    if not items:
        return None, {}
    ordered = sorted(items, key=lambda i: -len(i["original"]))
    alts = "|".join(_flex(i["original"]) for i in ordered)
    rx = re.compile(r"(?<!\w)(?:" + alts + r")(?!\w)", re.I)
    return rx, {_norm(i["original"]): i for i in items}


# ------------------------------------------------------------------ text extraction
def _docx_paragraphs(doc):
    from docx.oxml.ns import qn
    roots = [doc.element]
    for s in doc.sections:
        for h in (s.header, s.footer, s.first_page_header, s.first_page_footer,
                  s.even_page_header, s.even_page_footer):
            try:
                if not h.is_linked_to_previous:
                    roots.append(h._element)
            except Exception:
                pass
    for root in roots:
        for p in root.iter(qn("w:p")):
            yield p


def extract_text(path):
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        import fitz
        d = fitz.open(str(path))
        t = "\n".join(pg.get_text() for pg in d)
        d.close()
        return t
    if ext == ".docx":
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(str(path))
        return "\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in _docx_paragraphs(doc))
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(path))
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and not c.value.startswith("="):
                        out.append(c.value)
        return "\n".join(out)
    if ext in TEXT_EXT:
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type '{ext}'. Supported: PDF, DOCX, XLSX, TXT, CSV, MD "
                     f"(for .doc / .xls, save as .docx / .xlsx first).")


# ------------------------------------------------------------------ docx / xlsx / text engines
def _shade(run, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rpr = run.get_or_add_rPr()
    for e in rpr.findall(qn("w:shd")):
        rpr.remove(e)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    rpr.insert_element_before(shd, "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
                              "w:eastAsianLayout", "w:specVanish", "w:oMath")


def _set_t(t, text):
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _unhl(run):
    from docx.oxml.ns import qn
    rpr = run.find(qn("w:rPr"))
    if rpr is not None:
        for h in rpr.findall(qn("w:highlight")):
            rpr.remove(h)


def _apply_segments(t, segs, mark, strip_hl=False):
    """Write text segments [(text, is_replacement)] for w:t, shading only the replaced parts
    (the run is split so that the neighbouring text keeps no colour)."""
    import copy
    from docx.oxml.ns import qn
    segs = [(x, f) for x, f in segs if x != ""]
    run = t.getparent()
    if not segs:
        _set_t(t, "")
        return
    simple = run is not None and len(run.findall(qn("w:t"))) == 1 and \
        all(ch.tag in (qn("w:rPr"), qn("w:t")) for ch in run)
    if not simple or len(segs) == 1:
        _set_t(t, "".join(x for x, _ in segs))
        if run is not None and any(f for _, f in segs) and (not simple or segs[0][1]):
            if strip_hl:
                _unhl(run)
            _shade(run, mark)
        return
    tmpl = copy.deepcopy(run)
    _set_t(t, segs[0][0])
    if segs[0][1]:
        if strip_hl:
            _unhl(run)
        _shade(run, mark)
    prev = run
    for x, f in segs[1:]:
        nr = copy.deepcopy(tmpl)
        _set_t(nr.find(qn("w:t")), x)
        if f:
            if strip_hl:
                _unhl(nr)
            _shade(nr, mark)
        prev.addnext(nr)
        prev = nr


def _replace_in_ts(ts, rx, fn, touched=None, mark=None, strip_hl=False):
    """Replace matches across a paragraph's w:t elements, keeping each run's formatting.
    touched: optional list that receives the w:t elements whose text was really changed.
    mark: optional hex fill - replaced text gets this background colour."""
    texts = [t.text or "" for t in ts]
    full = "".join(texts)
    ms = list(rx.finditer(full))
    if not ms:
        return 0
    reps = [fn(m) for m in ms]
    pos, bounds = 0, []
    for tx in texts:
        bounds.append((pos, pos + len(tx)))
        pos += len(tx)
    for t, (a, b) in zip(ts, bounds):
        segs, cur, hit = [], a, False
        for m, rep in zip(ms, reps):
            s, e = m.start(), m.end()
            if e <= cur or s >= b:
                continue
            segs.append((full[cur:max(s, cur)], False))
            if a <= s < b:
                segs.append((rep, rep != m.group(0)))
            cur = min(e, b)
            if rep != m.group(0):
                hit = True
        segs.append((full[cur:b], False))
        if mark and any(f for _, f in segs):
            _apply_segments(t, segs, mark, strip_hl)
        else:
            _set_t(t, "".join(x for x, _ in segs))
        if hit and touched is not None:
            touched.append(t)
    return len(ms)


def _process_docx(src, out, rx, fn, scrub=False, unhighlight=False, mark=None, strip_hl=False):
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(str(src))
    n = 0
    touched = []
    for p in _docx_paragraphs(doc):
        ts = list(p.iter(qn("w:t")))
        if ts:
            n += _replace_in_ts(ts, rx, fn, touched, mark, strip_hl)
    if unhighlight:                       # filled-in text should not stay yellow
        for t in touched:
            run = t.getparent()
            rpr = run.find(qn("w:rPr")) if run is not None else None
            if rpr is not None:
                for h in rpr.findall(qn("w:highlight")):
                    rpr.remove(h)
    if scrub:
        cp = doc.core_properties
        for a in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
            try:
                setattr(cp, a, "")
            except Exception:
                pass
    doc.save(str(out))
    return n


def _process_xlsx(src, out, rx, fn, scrub=False, mark=None):
    from openpyxl import load_workbook
    wb = load_workbook(str(src))
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and not v.startswith("="):
                    nv, k = rx.subn(fn, v)
                    if k:
                        c.value = nv
                        n += k
                        if mark and nv != v:
                            from openpyxl.styles import PatternFill
                            c.fill = PatternFill("solid", fgColor=mark)
    if scrub:
        for a in ("creator", "lastModifiedBy", "title", "subject", "keywords", "description"):
            try:
                setattr(wb.properties, a, "")
            except Exception:
                pass
    wb.save(str(out))
    return n


def _process_text(src, out, rx, fn, scrub=False):
    t = Path(src).read_text(encoding="utf-8", errors="ignore")
    nt, n = rx.subn(fn, t)
    Path(out).write_text(nt, encoding="utf-8")
    return n


# ------------------------------------------------------------------ PDF engine
def _pdf_font(name, flags):
    n = (name or "").lower()
    bold = bool(flags & 16) or "bold" in n or "black" in n
    ital = bool(flags & 2) or "italic" in n or "oblique" in n
    if "cour" in n or "mono" in n or "consolas" in n:
        fam = ("cour", "cobo", "coit", "cobi")
    elif any(k in n for k in ("times", "serif", "georgia", "garamond", "cambria", "book", "roman", "palatino")) \
            and "sans" not in n:
        fam = ("tiro", "tibo", "tiit", "tibi")
    else:
        fam = ("helv", "hebo", "heit", "hebi")
    return fam[(1 if bold else 0) + (2 if ital else 0)]


def build_page_string(lines):
    """lines: list of lines, each a list of char dicts {c,bbox,size,...}.
    Returns (text, ref) where ref[i] is (line_index, char_dict) or None for newline/synthetic space."""
    s, ref = [], []
    for li, chars in enumerate(lines):
        prev = None
        for ch in chars:
            if prev is not None and prev["c"] != " " and ch["c"] != " " \
                    and ch["bbox"][0] - prev["bbox"][2] > 0.25 * ch["size"]:
                s.append(" "); ref.append(None)
            s.append(ch["c"]); ref.append((li, ch))
            prev = ch
        s.append("\n"); ref.append(None)
    return "".join(s), ref


def match_segments(ref, start, end):
    """One segment per physical line touched by the match: [x0,y0,x1,y1,first_char,last_index]"""
    segs = {}
    for k in range(start, end):
        r = ref[k]
        if r is None or r[1]["c"].isspace():
            continue
        li, ch = r
        x0, y0, x1, y1 = ch["bbox"]
        if li in segs:
            g = segs[li]
            g[0], g[1], g[2], g[3], g[5] = min(g[0], x0), min(g[1], y0), max(g[2], x1), max(g[3], y1), k
        else:
            segs[li] = [x0, y0, x1, y1, ch, k]
    return [segs[k] for k in sorted(segs)]


def distribute(tokens, widths, measure):
    """Spread replacement words over the line-segments of a wrapped match."""
    out, i = [], 0
    for k, w in enumerate(widths):
        if k == len(widths) - 1:
            out.append(" ".join(tokens[i:]))
            break
        cur = []
        while i < len(tokens):
            if measure(" ".join(cur + [tokens[i]])) <= w:
                cur.append(tokens[i]); i += 1
            else:
                break
        out.append(" ".join(cur))
    return out


def _page_lines(page):
    lines = []
    for b in page.get_text("rawdict")["blocks"]:
        if b.get("type") != 0:
            continue
        for l in b["lines"]:
            chars = []
            for sp in l["spans"]:
                for c in sp["chars"]:
                    chars.append({"c": c["c"], "bbox": tuple(c["bbox"]), "size": sp["size"], "color": sp["color"],
                                  "font": sp["font"], "flags": sp["flags"], "oy": c["origin"][1]})
            if chars:
                lines.append(chars)
    return lines


def fit_pdf_text(txt, w, size, measure, tail=""):
    """Make txt fit width w without touching neighbouring text: squeeze sideways, shrink a little,
    and only then shorten. Returns (text, fontsize, x_scale, was_shortened)."""
    fs, sx = size, 1.0
    tw = measure(txt, fs)
    if tw <= w or tw <= 0:
        return txt, fs, sx, False
    sx = max(0.6, w / tw)
    if tw * sx > w:
        fs = max(0.7 * size, size * w / (tw * sx))
    if measure(txt, fs) * sx <= w + 0.01:
        return txt, fs, sx, False
    base = txt[:-len(tail)] if tail and txt.endswith(tail) else txt
    while len(base) > 1 and measure(base + ".." + tail, fs) * sx > w:
        base = base[:-1]
    return base + ".." + tail, fs, sx, True


def _redact_pdf(src, out, rx, lookup, rules, warnings, text_counts):
    import fitz
    doc = fitz.open(str(src))
    hl_rgb = HL_RGB.get(rules.hl)
    hit_counts = {}
    for page in doc:
        lines = _page_lines(page)
        s, ref = build_page_string(lines)
        plans = []
        for m in rx.finditer(s):
            it = lookup.get(_norm(m.group(0)))
            if not it:
                continue
            segs = match_segments(ref, m.start(), m.end())
            if not segs:
                continue
            plans.append((it, segs, m.group(0)))
            k = _norm(m.group(0))
            hit_counts[k] = hit_counts.get(k, 0) + 1
        if not plans:
            continue

        for it, segs, mt in plans:
            for g in segs:
                h = g[3] - g[1]
                page.add_redact_annot(fitz.Rect(g[0] + 0.2, g[1] + h * 0.1, g[2] - 0.2, g[3] - h * 0.1), fill=(1, 1, 1))
        try:
            page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE)
        except Exception:
            page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)

        right_edge = page.rect.x1 - 20
        for it, segs, mt in plans:
            first = segs[0][4]
            fn_ = _pdf_font(first["font"], first["flags"])
            size = first["size"]
            tokens = display_like(it, rules, mt).encode("latin-1", "replace").decode("latin-1").split(" ")
            widths = [g[2] - g[0] for g in segs]
            # last segment may borrow the blank space after it on the same line
            last = segs[-1]
            nxt = None
            li_last = ref[last[5]][0]
            for j in range(last[5] + 1, len(ref)):
                if ref[j] is None:
                    continue
                if ref[j][0] != li_last:
                    break
                if not ref[j][1]["c"].isspace():
                    nxt = ref[j][1]["bbox"][0]
                    break
            room = (nxt - 0.3 * size) if nxt is not None else right_edge
            widths[-1] = max(widths[-1], room - last[0])
            parts = distribute(tokens, widths, lambda t: fitz.get_text_length(t, fontname=fn_, fontsize=size))
            for g, txt, w in zip(segs, parts, widths):
                if not txt:
                    continue
                ch = g[4]
                col = ch["color"]
                color = (((col >> 16) & 255) / 255, ((col >> 8) & 255) / 255, (col & 255) / 255)
                tail = rules.close if it["marked"] else ""
                txt, fs, sx, cut = fit_pdf_text(txt, w, size, lambda t_, f_: fitz.get_text_length(t_, fontname=fn_, fontsize=f_), tail)
                if cut:
                    warnings.append(f"PDF: replacement for '{it['original'][:40]}' is too long for the space in the PDF, so it is "
                                    f"shortened there (the .md copy has the full text). Use a shorter replacement for a cleaner PDF.")
                if hl_rgb:
                    used = fitz.get_text_length(txt, fontname=fn_, fontsize=fs) * sx
                    page.draw_rect(fitz.Rect(g[0], g[1], g[0] + used + 1, g[3]), color=None, fill=hl_rgb, overlay=True)
                pt = fitz.Point(g[0], ch["oy"])
                kw = dict(fontname=fn_, fontsize=fs, color=color)
                if sx < 0.999:
                    kw["morph"] = (pt, fitz.Matrix(sx, 1))
                page.insert_text(pt, txt, **kw)

    for k, c in text_counts.items():
        if hit_counts.get(k, 0) < c:
            warnings.append(f"PDF: '{k}' found {c}x in text but only {hit_counts.get(k, 0)} replaced - please check.")
    doc.set_metadata({})
    try:
        doc.del_xml_metadata()
    except Exception:
        pass
    doc.save(str(out), garbage=4, deflate=True, clean=True)
    doc.close()
    return sum(hit_counts.values())


# ------------------------------------------------------------------ markdown (for the AI)
_LIST_RX = re.compile(r"^(\(?[ivxlcdm]{1,6}\)|\(?[a-zA-Z0-9]{1,2}\)|\d{1,2}[.)]|[•●▪\-–])\s+")


def _md_cell(v):
    return str(v).replace("|", "\\|").replace("\n", " ").strip()


def _rows_to_md(rows):
    rows = [[_md_cell("" if c is None else c) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [r + [""] * (w - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "|".join(["---"] * w) + "|"]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def _join_lines(lines):
    out = []
    for t in lines:
        if out and not _LIST_RX.match(t):
            out[-1] += " " + t
        else:
            out.append(t)
    return "\n".join(out)


def _docx_to_md(path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    doc = Document(str(path))
    out = []
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            txt = p.text.strip()
            if not txt:
                out.append("")
                continue
            st = ((p.style.name if p.style is not None else "") or "").strip()
            if st.lower().startswith("heading"):
                try:
                    lvl = int(st.split()[-1])
                except Exception:
                    lvl = 2
                txt = "#" * min(max(lvl, 1), 6) + " " + txt
            elif st.lower() == "title":
                txt = "# " + txt
            else:
                ppr = el.find(qn("w:pPr"))
                if ppr is not None and ppr.find(qn("w:numPr")) is not None:
                    txt = "- " + txt
            out.append(txt)
        elif el.tag == qn("w:tbl"):
            t = Table(el, doc)
            out.append(_rows_to_md([[c.text for c in r.cells] for r in t.rows]))
    return re.sub(r"\n{3,}", "\n\n", "\n\n".join(out)).strip() + "\n"


def _xlsx_to_md(path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True)
    out = []
    for ws in wb.worksheets:
        md = _rows_to_md([list(r) for r in ws.iter_rows(values_only=True)])
        if md:
            out.append(f"## {ws.title}\n\n{md}")
    return "\n\n".join(out) + "\n"


def _pdf_to_md(path):
    import fitz
    d = fitz.open(str(path))
    sizes = {}
    for pg in d:
        for b in pg.get_text("dict")["blocks"]:
            if b.get("type") != 0:
                continue
            for l in b["lines"]:
                for sp in l["spans"]:
                    k = round(sp["size"])
                    sizes[k] = sizes.get(k, 0) + len(sp["text"])
    body = max(sizes, key=sizes.get) if sizes else 10
    pages = []
    for pg in d:
        tabs = []
        try:
            for t in pg.find_tables().tables:
                rows = t.extract()
                if rows and len(rows) >= 2:
                    tabs.append((fitz.Rect(t.bbox), _rows_to_md(rows)))
        except Exception:
            tabs = []
        parts = []
        pending = sorted(tabs, key=lambda x: x[0].y0)
        for b in pg.get_text("dict")["blocks"]:
            if b.get("type") != 0:
                continue
            r = fitz.Rect(b["bbox"])
            while pending and pending[0][0].y0 <= r.y0:
                parts.append(pending.pop(0)[1])
            if r.get_area() > 0 and any((r & tb[0]).get_area() > 0.5 * r.get_area() for tb in tabs):
                continue
            lines, size = [], 0
            for l in b["lines"]:
                txt = "".join(sp["text"] for sp in l["spans"]).strip()
                if txt:
                    lines.append(txt)
                    size = max(size, max(sp["size"] for sp in l["spans"]))
            if not lines:
                continue
            txt = _join_lines(lines)
            if size >= body * 1.2 and len(lines) <= 3:
                txt = "## " + " ".join(lines)
            parts.append(txt)
        parts.extend(x[1] for x in pending)
        pages.append("\n\n".join(parts))
    d.close()
    return "\n\n---\n\n".join(pages).strip() + "\n"


def to_markdown(path):
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return _pdf_to_md(path)
    if ext == ".docx":
        return _docx_to_md(path)
    if ext == ".xlsx":
        return _xlsx_to_md(path)
    return Path(path).read_text(encoding="utf-8", errors="ignore")


# ------------------------------------------------------------------ fill-in items  {{FILL: ...}}
FILL_RX = re.compile(r"\{\{[^{}]*?\}\}")


def doc_paragraphs(path):
    """Paragraph-level text of a Word / Excel / text file (for showing context)."""
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(str(path))
        return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in _docx_paragraphs(doc)]
    return extract_text(path).split("\n")


def find_fill_items(paragraphs):
    """Unique {{...}} items in order of first appearance: [{key,text,label,hint,count}]
    Item form:  {{FILL: label | HINT: guidance}}   (HINT is optional)"""
    seen, items = {}, []
    for para in paragraphs:
        for m in FILL_RX.finditer(para):
            k = _norm(m.group(0))
            if k in seen:
                seen[k]["count"] += 1
                continue
            inner = m.group(0)[2:-2]
            parts = re.split(r"\|\s*HINT\s*:", inner, maxsplit=1, flags=re.I)
            label = re.sub(r"^\s*fill\s*:?\s*", "", parts[0], flags=re.I).strip() or parts[0].strip()
            hint = parts[1].strip() if len(parts) > 1 else ""
            seen[k] = {"key": k, "text": m.group(0), "label": label, "hint": hint, "count": 1}
            items.append(seen[k])
    return items


def _fill_fn(mapping, counter):
    def fn(m):
        v = mapping.get(_norm(m.group(0)))
        if v is None or str(v).strip() == "":
            return m.group(0)
        counter[0] += 1
        return str(v).strip()
    return fn


def apply_fills(path, mapping, out_path=None):
    """mapping: {normalised item text: value}. Edits the file (or writes out_path). Returns number filled."""
    path = Path(path)
    out = Path(out_path) if out_path else path
    ext = path.suffix.lower()
    cnt = [0]
    fn = _fill_fn(mapping, cnt)
    if ext == ".docx":
        _process_docx(path, out, FILL_RX, fn, mark=YELLOW_HEX, strip_hl=True)
    elif ext == ".xlsx":
        _process_xlsx(path, out, FILL_RX, fn, mark=YELLOW_HEX)
    else:
        _process_text(path, out, FILL_RX, fn)
    return cnt[0]


def fill_text(text, mapping, mark=False):
    cnt = [0]
    f = _fill_fn(mapping, cnt)

    def f2(m):
        r = f(m)
        return Y_ON + r + Y_OFF if (mark and r != m.group(0)) else r
    return FILL_RX.sub(f2, text)


def clear_colours(path, out_path):
    """Remove the green / yellow marking (highlights and our shading) from a Word or Excel file."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".docx":
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(str(path))
        for p in _docx_paragraphs(doc):
            for rpr in p.iter(qn("w:rPr")):
                for h in rpr.findall(qn("w:highlight")):
                    rpr.remove(h)
                for sh in rpr.findall(qn("w:shd")):
                    if (sh.get(qn("w:fill")) or "").upper() in COLOUR_SET:
                        rpr.remove(sh)
        doc.save(str(out_path))
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill
        wb = load_workbook(str(path))
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    rgb = str(c.fill.fgColor.rgb or "")[-6:].upper() if c.fill and c.fill.fill_type else ""
                    if rgb in COLOUR_SET:
                        c.fill = PatternFill(fill_type=None)
        wb.save(str(out_path))
    else:
        raise ValueError("Colours can be cleared in Word (.docx) or Excel (.xlsx) files only.")
    return str(out_path)


# ------------------------------------------------------------------ "Notes for me" + guidance comments
_NOTES_RX = re.compile(r"^\W*notes\s+for\s+me\W*$", re.I)


def split_notes_text(text):
    """-> (main_text, notes_text or '')"""
    m = re.search(r"^\W*notes\s+for\s+me\W*$", text, re.I | re.M)
    if not m:
        return text, ""
    return text[:m.start()].rstrip() + "\n", text[m.start():].strip() + "\n"


def split_notes(path, notes_path):
    """Move the 'Notes for me' section (heading to the end) of a Word file into its own file.
    Returns the notes file path, or None if the reply has no such section."""
    from docx import Document
    from docx.oxml.ns import qn

    def find(doc):
        kids = [k for k in doc.element.body.iterchildren() if k.tag != qn("w:sectPr")]
        for i, k in enumerate(kids):
            if k.tag == qn("w:p") and _NOTES_RX.match("".join(t.text or "" for t in k.iter(qn("w:t"))).strip()):
                return kids, i
        return kids, None

    doc = Document(str(path))
    kids, i = find(doc)
    if i is None:
        return None
    notes = Document(str(path))
    nkids, _ = find(notes)
    for k in nkids[:i]:
        notes.element.body.remove(k)
    for k in kids[i:]:
        doc.element.body.remove(k)
    doc.save(str(path))
    notes.save(str(notes_path))
    return str(notes_path)


def _split_hint(item_text):
    """'{{FILL: a | HINT: b}}' -> ('{{FILL: a}}', 'b'); no hint -> (item_text, '')"""
    m = re.match(r"^\{\{(.*?)\|\s*HINT\s*:(.*)\}\}$", item_text, re.S | re.I)
    if not m:
        return item_text, ""
    return "{{" + m.group(1).rstrip() + "}}", m.group(2).strip()


def hints_to_comments(path, items):
    """For blanks left unfilled: keep just '{{FILL: label}}' in the text and put the AI's HINT in a
    Word / Excel comment next to it. Returns number of comments added."""
    path = Path(path)
    ext = path.suffix.lower()
    todo = {i["key"]: i for i in items if i.get("hint")}
    if not todo:
        return 0
    n = 0
    if ext == ".docx":
        from docx import Document
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph
        doc = Document(str(path))
        if not hasattr(doc, "add_comment"):
            return 0                                      # old python-docx: leave the hint in the text
        stripped = {}

        def fn(m):
            it = todo.get(_norm(m.group(0)))
            if not it:
                return m.group(0)
            st, _ = _split_hint(m.group(0))
            stripped.setdefault(_norm(st), it)
            return st
        body_ps = list(doc.element.body.iter(qn("w:p")))
        for p in body_ps:
            ts = list(p.iter(qn("w:t")))
            if ts:
                _replace_in_ts(ts, FILL_RX, fn, None, YELLOW_HEX, False)
        done = set()
        for p in body_ps:
            para = Paragraph(p, doc._body)
            for run in para.runs:
                m = FILL_RX.search(run.text)
                if not m:
                    continue
                k = _norm(m.group(0))
                if k in stripped and k not in done:
                    done.add(k)
                    doc.add_comment(runs=run, text="HOW TO FILL:  " + stripped[k]["hint"], author="AI guidance", initials="AI")
                    n += 1
        doc.save(str(path))
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.comments import Comment
        wb = load_workbook(str(path))
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and FILL_RX.search(c.value):
                        hints = []

                        def fn(m):
                            it = todo.get(_norm(m.group(0)))
                            if not it:
                                return m.group(0)
                            hints.append(it["hint"])
                            return _split_hint(m.group(0))[0]
                        nv = FILL_RX.sub(fn, c.value)
                        if hints:
                            c.value = nv
                            c.comment = Comment("HOW TO FILL:  " + "\n".join(hints), "AI guidance")
                            c.comment.width, c.comment.height = 340, 140
                            n += 1
        wb.save(str(path))
    return n


# ------------------------------------------------------------------ public API
def _remove_stale(outdir, stem):
    """Delete numbered leftovers of earlier runs (e.g. name_REDACTED (1).pdf, name_REDACTED_2.md)."""
    rx = re.compile(r"^" + re.escape(stem) + r"_(?:REDACTED|MAP)(?:[ _\-]*\(?\d+\)?)\.(?:pdf|docx|xlsx|md|txt|csv|json)$", re.I)
    gone = []
    try:
        for f in Path(outdir).iterdir():
            if f.is_file() and rx.match(f.name):
                try:
                    f.unlink()
                    gone.append(f.name)
                except Exception:
                    pass
    except Exception:
        pass
    return gone


def redact_file(src, rules_xlsx, outdir=None, out_path=None, make_md=True):
    try:
        return _redact_file(src, rules_xlsx, outdir, out_path, make_md)
    except PermissionError as e:
        raise PermissionError(f"Could not save '{Path(getattr(e, 'filename', '') or src).name}' - it is open in another program "
                              f"(Acrobat / Word / Excel / a preview pane). Close it, then click Redact again.") from None


def _redact_file(src, rules_xlsx, outdir=None, out_path=None, make_md=True):
    src = Path(src)
    ext = src.suffix.lower()
    rules = load_rules(rules_xlsx)
    full = extract_text(src)
    if ext == ".pdf" and len(full.strip()) < 30:
        raise ValueError("This PDF has no selectable text (scanned image). Run OCR first, then retry.")
    items = build_items(full, rules)
    rx, lookup = compile_items(items)
    warnings = []

    text_counts = {}
    if rx:
        for m in rx.finditer(full):
            k = _norm(m.group(0))
            text_counts[k] = text_counts.get(k, 0) + 1
    cores = {}
    for it in items:
        if it["core"].lower() in cores and cores[it["core"].lower()] != _norm(it["original"]):
            warnings.append(f"'{it['original']}' and another item both become '{it['core']}' - on restore it will come back as the FIRST one listed. Use different replacements to be safe.")
        cores.setdefault(it["core"].lower(), _norm(it["original"]))

    if out_path:
        out = Path(out_path)
        outdir = out.parent
    else:
        outdir = Path(outdir) if outdir else src.parent
        out = outdir / f"{src.stem}_REDACTED{ext}"
    mapp = outdir / f"{src.stem}_MAP.json"
    replaced = out.exists()
    removed = [] if out_path else _remove_stale(outdir, src.stem)
    pad_ok = ext in (".docx", ".xlsx") + TEXT_EXT

    def fn(m):
        it = lookup[_norm(m.group(0))]
        d = display_like(it, rules, m.group(0))
        if rules.pad and pad_ok and len(d) < len(m.group(0)):
            d += NBSP * (len(m.group(0)) - len(d))
        return d

    n = 0
    if items:
        if ext == ".pdf":
            n = _redact_pdf(src, out, rx, lookup, rules, warnings, text_counts)
        elif ext == ".docx":
            n = _process_docx(src, out, rx, fn, scrub=True, mark=HL_HEX.get(rules.hl))
        elif ext == ".xlsx":
            n = _process_xlsx(src, out, rx, fn, scrub=True, mark=HL_HEX.get(rules.hl))
        else:
            n = _process_text(src, out, rx, fn)
    else:
        warnings.append("No rules matched anything - nothing was redacted.")

    if items and out.exists():
        try:
            chk = [i for i in items if _norm(i["original"]) not in _norm(display(i, rules))]
            rx2, _ = compile_items(chk)
            if rx2:
                for k in sorted({_norm(m.group(0)) for m in rx2.finditer(extract_text(out))}):
                    warnings.append(f"STILL PRESENT in output: '{k}' - not fully replaced, please check.")
        except Exception as e:
            warnings.append(f"Could not run final check: {e}")

    md_path = None
    if items and make_md and ext in (".pdf", ".docx", ".xlsx"):
        try:
            md = to_markdown(src)
            md = rx.sub(lambda m: display_like(lookup[_norm(m.group(0))], rules, m.group(0)), md)
            md_path = out.with_suffix(".md")
            md_path.write_text(md, encoding="utf-8")
            chk = [i for i in items if _norm(i["original"]) not in _norm(display(i, rules))]
            rx3, _ = compile_items(chk)
            if rx3:
                for k in sorted({_norm(m.group(0)) for m in rx3.finditer(md)}):
                    warnings.append(f"STILL PRESENT in Markdown copy: '{k}' - please check.")
        except Exception as e:
            md_path = None
            warnings.append(f"Markdown copy could not be made: {e}")

    mapp.write_text(json.dumps({
        "source": src.name, "open": rules.open, "close": rules.close,
        "items": [{"original": i["original"], "core": i["core"], "marked": i["marked"]} for i in items],
    }, indent=2, ensure_ascii=False), encoding="utf-8")

    return {"out": str(out) if items else None, "md": str(md_path) if md_path else None, "map": str(mapp), "count": n,
            "replaced": replaced, "removed": removed,
            "pairs": [(i["original"], display(i, rules)) for i in items], "warnings": warnings}


def _restore_engine(map_json):
    data = json.loads(Path(map_json).read_text(encoding="utf-8"))
    op, cl, items = data["open"], data["close"], data["items"]
    lookup = {}
    for i in items:
        lookup.setdefault(_norm(i["core"]), re.sub(r"\s+", " ", i["original"]).strip())   # no stray line breaks
    cores = sorted({i["core"] for i in items}, key=len, reverse=True)
    if not cores:
        raise ValueError("The mapping file has no items.")
    rx = re.compile(r"(?<!\w)(?:" + re.escape(op) + r"\s*)?(?:" + "|".join(_flex(c) for c in cores) +
                    r")(?:\s*" + re.escape(cl) + r")?(?!\w)", re.I)

    def fn(m):
        txt = m.group(0)
        if op and txt.startswith(op):
            txt = txt[len(op):]
        if cl and txt.endswith(cl):
            txt = txt[:-len(cl)]
        orig = lookup.get(_norm(txt))
        if orig is None:
            return m.group(0)
        if txt[:1].isupper() and orig[:1].islower():      # LLM capitalised the start of a sentence
            orig = orig[:1].upper() + orig[1:]
        return orig

    return rx, fn, op, cl


def _unknown_markers(text, op, cl):
    if not (op and cl):
        return []
    return [f"Unknown marker left in text (AI may have changed it): {t}"
            for t in sorted(set(re.findall(re.escape(op) + r".{1,80}?" + re.escape(cl), text)))]


def strip_marks(t):
    return re.sub("[\ue000-\ue003]", "", t)


def restore_text(text, map_json, mark=False):
    """Restore plain text (e.g. pasted from the AI). Returns (restored_text, warnings).
    mark=True wraps restored words in colour marks (used when saving as Word)."""
    rx, fn, op, cl = _restore_engine(map_json)

    def f2(m):
        r = fn(m)
        return G_ON + r + G_OFF if (mark and r != m.group(0)) else r
    out = rx.sub(f2, text)
    return out, _unknown_markers(out, op, cl)


def save_text_as_docx(text, path):
    """Plain text -> Word. Colour marks become green / yellow shading; leftover {{ }} blanks are highlighted yellow."""
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    d = Document()
    commented = set()
    for line in text.split("\n"):
        p = d.add_paragraph()
        state = None
        for part in re.split("([\ue000-\ue003])", line):
            if part == G_ON:
                state = GREEN_HEX
            elif part == Y_ON:
                state = YELLOW_HEX
            elif part in (G_OFF, Y_OFF):
                state = None
            elif part:
                if state:
                    _shade(p.add_run(part)._r, state)
                    continue
                pos = 0
                for m in FILL_RX.finditer(part):
                    if m.start() > pos:
                        p.add_run(part[pos:m.start()])
                    shown, hint = _split_hint(m.group(0))
                    if hint and not hasattr(d, "add_comment"):
                        shown, hint = m.group(0), ""
                    run = p.add_run(shown)
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    if hint and _norm(shown) not in commented:
                        commented.add(_norm(shown))
                        d.add_comment(runs=run, text="HOW TO FILL:  " + hint, author="AI guidance", initials="AI")
                    pos = m.end()
                if pos < len(part):
                    p.add_run(part[pos:])
    d.save(str(path))


def restore_file(src, map_json, outdir=None, out_path=None):
    try:
        return _restore_file(src, map_json, outdir, out_path)
    except PermissionError as e:
        raise PermissionError(f"Could not save the restored file - it is open in another program (Word / Excel). "
                              f"Close it, then try again.") from None


def _restore_file(src, map_json, outdir=None, out_path=None):
    src = Path(src)
    ext = src.suffix.lower()
    if ext == ".pdf":
        raise ValueError("Please paste the AI answer into a Word (.docx) or text (.txt) file - PDF answers can't be restored.")
    rx, fn, op, cl = _restore_engine(map_json)
    if out_path:
        out = Path(out_path)
    else:
        outdir = Path(outdir) if outdir else src.parent
        out = outdir / f"{src.stem}_RESTORED{ext}"
    if ext == ".docx":
        n = _process_docx(src, out, rx, fn, mark=GREEN_HEX, strip_hl=True)
    elif ext == ".xlsx":
        n = _process_xlsx(src, out, rx, fn, mark=GREEN_HEX)
    elif ext in TEXT_EXT:
        n = _process_text(src, out, rx, fn)
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Use Word, Excel or text.")
    notes = None
    try:
        if ext == ".docx":
            notes = split_notes(out, out.with_name(out.stem + "_NOTES.docx"))
        elif ext in TEXT_EXT:
            main, nt = split_notes_text(Path(out).read_text(encoding="utf-8"))
            if nt:
                Path(out).write_text(main, encoding="utf-8")
                notes = str(out.with_name(out.stem + "_NOTES" + ext))
                Path(notes).write_text(nt, encoding="utf-8")
    except Exception:
        notes = None
    return {"out": str(out), "notes": notes, "count": n, "warnings": _unknown_markers(extract_text(out), op, cl)}


# ------------------------------------------------------------------ prompt for the AI
PROMPT_VERSION = 5
DEFAULT_PROMPT = """#PROMPT_VERSION 5
I am uploading a REDACTED version of a document connected with Indian tax law. It may be a notice, an order, a show-cause notice, an appeal, an adjudication or litigation paper, or any similar tax document. Please prepare the appropriate response on behalf of the assessee / taxpayer - for example a reply to a notice, a submission, or the statement of facts and grounds for an appeal - as indicated by the document and by my special instructions at the end. If it is not clear what response is expected, assume a reply and say so in "Notes for me". Draft it professionally, as a Chartered Accountant would.

FILES
- The Markdown (.md) file is the document for you to work from. Base your response on it.
- The PDF is the same redacted document, attached only as my own reference copy for later. Please do NOT spend time processing it: look at the PDF only if something in the Markdown is unclear (a table, a layout question). Do not wait on it or comment on it.
- If only one of the two is attached, use whichever is there. Ignore small layout artefacts caused by the Markdown conversion.

HOW THE REDACTION WORKS (important)
- Client details (names, addresses, offices, order and reference numbers, nature of business, etc.) were replaced with substitutes. Every substituted item is enclosed in double square brackets, for example @@OPEN@@Redacted PAN@@CLOSE@@ or @@OPEN@@XYZ Company Ltd@@CLOSE@@.
- Treat each bracketed item as the real name / value. Do not ask me for the real details and do not try to guess or "un-redact" them.

RULES FOR YOUR RESPONSE (so that my software can restore the real details automatically)
1. Wherever your response mentions a bracketed item, copy it EXACTLY as written in the document, including both brackets, with the same spelling, capital / small letters, spacing and punctuation inside the brackets. Do not abbreviate, translate, shorten, re-order or split it across two lines.
2. Keep each bracketed item as one unbroken piece of text: do not apply bold / italic / underline / colour to only part of it, and do not put hyperlinks, line breaks, tabs or special hyphens inside it. (Formatting the whole item is fine.)
3. Do not create new bracketed items of your own, and do not drop, change or add brackets. If the same party is referred to in different ways, use the item exactly as it appears in the document each time.
4. Do not mention that the document was redacted and do not add comments such as "[redacted]".

FILL-INS - ONE FIXED STYLE ONLY
5. Anything I must fill in myself (information not in the document, decisions, amounts, dates, document references, names of signatories, place, enclosures, etc.) must be written in exactly this form:
   {{FILL: what is needed | HINT: how to fill it}}
   i.e. two opening curly braces, the word FILL, a colon, a short precise description of WHAT is needed, a vertical bar, the word HINT, a colon, your guidance to me, two closing curly braces.
6. The description must say exactly WHAT is needed and for WHICH context, because my software will show it to me and ask for the value. The HINT is your instruction to me on HOW to fill it: a suggested value, the options to choose between, or what I should confirm first. My software shows the HINT only in a pop-up (and as a comment if the blank stays unfilled) and removes it from the document text, so keep it practical and to one or two sentences.
7. NATURAL FLOW - very important. The finished text, after I type my answer into each blank, must read as smooth, natural prose, exactly as if you had written it in one go. Therefore:
   - NEVER write a lead-in followed by a blank, such as "Our position is as follows: {{...}}", "Date - {{...}}" or "Status: {{...}}". That reads like a phone-banking voice menu that reads a fixed sentence and then plays a different voice for the value. Write the complete sentence yourself and put the blank INSIDE it, so that the grammar, tense, articles and punctuation work with whatever I type.
   - Keep each blank as short as possible: a date, name, amount, number, or a short phrase.
   - Where I must choose between alternatives, put them in the HINT as complete, ready-to-type wording that fits the blank exactly, and make sure the words around the blank work with EACH alternative. Example: "In response to the notice, we {{FILL: our response to the hearing date | HINT: type one of: 'will attend through our authorised representative on the scheduled date' / 'request an adjournment to <date> on account of <reason>'}}."
   - If a whole paragraph is needed from me (for example, the facts), say so in the HINT and give a suggested draft in the HINT that I can adapt.
   - Put sentence-ending punctuation after the closing braces, outside the blank. Never end a blank with a full stop inside it.
8. Other good examples of the description / HINT: {{FILL: date of this reply | HINT: normally today's date; confirm with me if it is to be backdated}}; {{FILL: date of the notice being replied to | HINT: take it from the notice if shown, otherwise ask me}}. Bad: {{FILL: date}}.
9. Use this style every time and nothing else - no asterisks (** or ***), no underscores, no square brackets, no angle brackets, no "XXXX" and no blank lines like "______". One fill-in per item; do not nest them, do not split them over two lines, do not use curly braces or vertical bars inside the description or the hint. If the same information is needed again, repeat the identical wording (including the hint).
10. Text that is not in brackets or braces is genuine (dates, amounts, section numbers, designations of authorities). Use it as it is.
11. Quote section / rule numbers and case laws only if you are sure; otherwise write {{FILL: verify the reference | HINT: I could not confirm this citation - please check it}}.

OUTPUT FORMAT
@@FORMAT@@
Give only the response document. If you have notes or doubts for me, put them in a short separate section titled "Notes for me" at the very end, after the response (my software saves that section as a separate file, so keep nothing else there).

SPECIAL INSTRUCTIONS FOR THIS DOCUMENT
@@SPECIAL@@
"""

FORMAT_RULES = {
    "word": "Give the reply as a downloadable Word (.docx) file, not a PDF. Use normal paragraph formatting (bold, underline, alignment, numbered points) and keep the layout simple: no text boxes, no images of text, no tracked changes, no tables unless really needed. COLOURS (text highlight colour, applied to the whole item): highlight every double-square-bracket item (including its brackets) in GREEN, and highlight every {{FILL: ...}} item (including its braces) in YELLOW. Use no other highlighting anywhere.",
    "excel": "Give the reply as a downloadable Excel (.xlsx) file. Put text in normal cells (no formulas that build text), one idea per cell, and keep each bracketed item inside a single cell. COLOURS: fill a cell GREEN if it contains a double-square-bracket item and YELLOW if it contains a {{FILL: ...}} item; no other cell colours.",
    "text": "Give the reply as plain text only: one paragraph per block, no markdown symbols (no **, no #, no pipe tables). Use the {{FILL: ... | HINT: ...}} form for every blank.",
}


def ensure_prompt_file(app_dir):
    """Create prompt_fixed.txt, or upgrade an older one (the old file is kept as prompt_fixed_old_<time>.txt)."""
    import time
    p = Path(app_dir) / "prompt_fixed.txt"
    try:
        if p.exists():
            cur = p.read_text(encoding="utf-8")
            m = re.search(r"^#PROMPT_VERSION\s+(\d+)", cur, re.M)
            if m and int(m.group(1)) >= PROMPT_VERSION:
                return p
            (Path(app_dir) / f"prompt_fixed_old_{time.strftime('%Y%m%d_%H%M%S')}.txt").write_text(cur, encoding="utf-8")
        p.write_text(DEFAULT_PROMPT, encoding="utf-8")
    except Exception:
        pass
    return p


def build_prompt(app_dir, fmt, open_marker, close_marker, special):
    p = ensure_prompt_file(app_dir)
    try:
        txt = p.read_text(encoding="utf-8")
    except Exception:
        txt = DEFAULT_PROMPT
    txt = "\n".join(l for l in txt.split("\n") if not l.startswith("#PROMPT_VERSION"))
    return (txt.strip() + "\n").replace("@@OPEN@@", open_marker).replace("@@CLOSE@@", close_marker) \
        .replace("@@FORMAT@@", FORMAT_RULES.get(fmt, FORMAT_RULES["word"])) \
        .replace("@@SPECIAL@@", special.strip() or "(none)")


if __name__ == "__main__":
    import sys
    if len(sys.argv) == 4 and sys.argv[1] == "redact":
        print(redact_file(sys.argv[2], sys.argv[3]))
    elif len(sys.argv) == 4 and sys.argv[1] == "restore":
        print(restore_file(sys.argv[2], sys.argv[3]))
    else:
        print("usage: redactor_core.py redact|restore <file> <rules.xlsx|map.json>")
