import unittest,tempfile,json
from pathlib import Path
from decimal import Decimal as D
from engine import *
from report import export_report
from openpyxl import load_workbook

class TaxTests(unittest.TestCase):
 def profile(self,**kw):
  p=blank();p.update({k:True for k in BLOCKERS});p.update(kw);return p
 def test_new_slabs(self):
  for income,expected in [(0,0),(400000,0),(800000,20000),(1200000,60000),(1600000,120000),(2000000,200000),(2400000,300000),(2500000,330000)]:
   with self.subTest(income=income):self.assertEqual(tax(income,'New',30)['slab_tax'],expected)
 def test_old_age_slabs(self):
  for age,expected in [(35,12500),(60,10000),(79,10000),(80,0)]:
   with self.subTest(age=age):self.assertEqual(tax(500000,'Old',age)['slab_tax'],expected)
 def test_rebate_old(self):self.assertEqual(tax(500000,'Old',35)['tax'],0);self.assertEqual(tax(500010,'Old',35)['tax'],13000)
 def test_new_rebate(self):self.assertEqual(tax(1200000,'New',35)['tax'],0)
 def test_marginal_relief(self):
  r=tax(1210000,'New',35);self.assertEqual(r['slab_tax'],61500);self.assertEqual(r['marginal_relief'],51500);self.assertEqual(r['tax'],10400)
 def test_marginal_ends(self):self.assertEqual(tax(1300000,'New',35)['marginal_relief'],0)
 def test_rounding(self):
  for value,expected in [('14.99',10),('15',20),('24.99',20),('25',30),('-15',-20),('0',0)]:
   with self.subTest(value=value):self.assertEqual(round10(D(value)),expected)
 def test_salary1275(self):self.assertEqual(calculate(self.profile(salary='1275000'))['New']['tax'],0)
 def test_standard_capped(self):
  r=calculate(self.profile(salary='20000'));self.assertEqual(r['Old']['standard_deduction'],20000);self.assertEqual(r['New']['total_income'],0)
 def test_hra_nonmetro(self):self.assertEqual(calculate(self.profile(salary='1800000',hra='500000',hra_salary='900000',rent='450000'))['Old']['hra_exemption'],360000)
 def test_hra_metro(self):self.assertEqual(calculate(self.profile(salary='1800000',hra='500000',hra_salary='900000',rent='600000',metro=True))['Old']['hra_exemption'],450000)
 def test_hra_no_rent(self):self.assertEqual(calculate(self.profile(salary='500000',hra='100000',hra_salary='300000'))['Old']['hra_exemption'],0)
 def test_new_disallows(self):
  r=calculate(self.profile(salary='1500000',c80='150000',nps='50000',home_interest='200000',professional_tax='2500'))['New'];self.assertEqual(r['chapter_via'],0);self.assertEqual(r['house_loss_setoff'],0);self.assertEqual(r['professional_tax'],0)
 def test_claim_caps(self):
  c=calculate(self.profile(salary='2000000',c80='200000',nps='60000',health_self='80000',health_parents='80000'))['Old']['claims']
  self.assertEqual(list(c.values())[:4],[150000,50000,25000,25000])
 def test_senior_claims(self):
  c=calculate(self.profile(age=65,salary='1200000',savings='20000',deposits='90000',health_self='60000',health_parents='60000',parents_senior=True))['Old']['claims'];self.assertEqual(c['80TTB'],50000);self.assertEqual(c['80D self/family'],50000);self.assertEqual(c['80D parents'],50000)
 def test_tta_excludes_fd(self):self.assertEqual(calculate(self.profile(salary='1000000',savings='5000',deposits='90000'))['Old']['claims']['80TTA'],5000)
 def test_gti_cap(self):self.assertEqual(calculate(self.profile(savings='5000',c80='150000'))['Old']['total_income'],0)
 def test_home_cap(self):self.assertEqual(calculate(self.profile(salary='1500000',home_interest='250000',home_cap='30000'))['Old']['house_loss_setoff'],30000)
 def test_loss_blocked(self):self.assertTrue(calculate(self.profile(salary='60000',home_interest='200000'))['blocked'])
 def test_scope_blocks(self):
  for key in BLOCKERS:
   p=self.profile();p[key]=False
   with self.subTest(key=key):self.assertTrue(calculate(p)['blocked'])
 def test_invalid_amounts(self):
  for value in ['NaN','Infinity','-1','abc']:
   with self.subTest(value=value),self.assertRaises(ValueError):calculate(self.profile(salary=value))
 def test_year_blocked(self):
  with self.assertRaises(ValueError):calculate(self.profile(year='2025-26'))
 def test_surcharge_blocked(self):
  with self.assertRaises(ValueError):calculate(self.profile(salary='5000001'))
 def test_hra_consistency(self):
  with self.assertRaises(ValueError):calculate(self.profile(salary='100000',hra='80000',hra_salary='80000'))
 def test_boolean_type(self):
  with self.assertRaises(ValueError):calculate(self.profile(resident='false'))
 def test_age(self):
  with self.assertRaises(ValueError):calculate(self.profile(age='35.5'))
 def test_conditional(self):self.assertIn('Conditional',calculate(self.profile(salary='1500000'))['status'])
 def test_credit_independence(self):
  p=self.profile(salary='1500000');a=calculate(p);p['credits']='200000';b=calculate(p);self.assertEqual(a['winner'],b['winner']);self.assertEqual(b['Old']['balance'],round10(b['Old']['raw_tax']-200000))
 def test_scenarios_no_mutation(self):
  p=self.profile(salary='2000000',c80='100000');q=deepcopy(p);s=scenarios(p);self.assertEqual(p,q);self.assertEqual(s[1]['additional_contribution'],50000);self.assertLessEqual(s[1]['old_tax'],s[0]['old_tax'])
 def test_ai_anonymous(self):
  p=self.profile(case='Private Name',salary='1500000');self.assertNotIn('Private Name',ai_prompt(p,calculate(p)))
 def test_export_and_formula_safety(self):
  p=self.profile(case='=1+1',salary='1500000')
  with tempfile.TemporaryDirectory() as d:
   path=Path(d)/'report.xlsx';export_report(p,path,notes='=HYPERLINK("x")');w=load_workbook(path)
   self.assertEqual(len(w.sheetnames),8);self.assertEqual(w['Comparison']['B2'].data_type,'s');self.assertEqual(w['Comparison']['B2'].value,'=1+1')
 def test_samples(self):
  for path in (Path(__file__).parent/'samples').glob('*.json'):
   with self.subTest(path=path.name):calculate(json.loads(path.read_text()))

if __name__=='__main__':unittest.main()
