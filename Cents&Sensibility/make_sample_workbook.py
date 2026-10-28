"""Creates sample_data/sample_workbook_with_notes.xlsx - a fictional multi-sheet workbook
(statements, notes to accounts, company details) to demo the 'Excel workbook with notes' route."""
from openpyxl import Workbook
from openpyxl.styles import Font

wb = Workbook()
B = Font(bold=True)

def sheet(title, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for r in rows:
        ws.append(r)
    ws["A1"].font = B
    ws.column_dimensions["A"].width = 52
    for c in "BCD":
        ws.column_dimensions[c].width = 16
    return ws

sheet("Company Info", [
    ["Company details"],
    ["Name of the company", "Meridian Alloys Private Limited"],
    ["CIN", "U27100KL2012PTC030000"],
    ["Registered office address", "14/772 Industrial Estate, Kalamassery, Ernakulam - 683 104"],
    ["Directors", "Mr. Rohan Menon (DIN 01234567); Mrs. Latha Pillai (DIN 07654321)"],
    ["Statutory auditors", "Nair & Associates, Chartered Accountants (FRN 004567S)"],
], first=True)

sheet("Balance Sheet", [
    ["Meridian Alloys Private Limited - Balance Sheet (Rs. in Lakhs)", "Note", "31-Mar-2026", "31-Mar-2025"],
    ["EQUITY AND LIABILITIES"],
    ["Share capital", 1, 100, 100],
    ["Reserves and surplus", "", 700, 635],
    ["Long-term borrowings", 2, 300, 315],
    ["Short-term borrowings", 2, 150, 120],
    ["Trade payables", "", 250, 200],
    ["Other current liabilities", "", 100, 80],
    ["TOTAL EQUITY AND LIABILITIES", "", 1600, 1450],
    ["ASSETS"],
    ["Property, plant and equipment", "", 800, 750],
    ["Other non-current assets", "", 100, 90],
    ["Inventories", "", 300, 260],
    ["Trade receivables", "", 250, 200],
    ["Cash and bank balances", "", 100, 90],
    ["Other current assets", "", 50, 60],
    ["TOTAL ASSETS", "", 1600, 1450],
])

sheet("Profit and Loss", [
    ["Meridian Alloys Private Limited - Statement of Profit and Loss (Rs. in Lakhs)", "Note", "FY 2025-26", "FY 2024-25"],
    ["Revenue from operations", "", 3000, 2600],
    ["Other income", "", 40, 30],
    ["Total income", "", 3040, 2630],
    ["Cost of materials consumed", "", 1700, 1500],
    ["Employee benefits expense", "", 450, 400],
    ["Finance costs", "", 70, 65],
    ["Depreciation and amortisation", "", 90, 85],
    ["Other expenses", 3, 640, 520],
    ["Total expenses", "", 2950, 2570],
    ["Profit before tax", "", 90, 60],
    ["Tax expense", "", 25, 20],
    ["Profit for the year", "", 65, 40],
])

sheet("Note 1 Share capital", [
    ["Note 1 - Share capital", "31-Mar-2026", "31-Mar-2025"],
    ["Authorised: 15,00,000 equity shares of Rs.10 each", 150, 150],
    ["Issued, subscribed and paid up: 10,00,000 equity shares of Rs.10 each", 100, 100],
])

sheet("Note 2 Borrowings", [
    ["Note 2 - Borrowings", "31-Mar-2026", "31-Mar-2025"],
    ["Term loans from banks (secured)", 250, 270],
    ["Lease liabilities", 50, 45],
    ["Long-term borrowings", 300, 315],
    ["Working capital loans (secured)", 150, 120],
    ["Short-term borrowings", 150, 120],
])

sheet("Note 3 Other expenses", [
    ["Note 3 - Other expenses", "FY 2025-26", "FY 2024-25"],
    ["Power and fuel", 210, 170],
    ["Repairs and maintenance", 90, 75],
    ["Rent", 60, 55],
    ["Professional fees", 55, 40],
    ["Travelling", 45, 35],
    ["Freight and forwarding", 120, 95],
    ["Miscellaneous expenses", 60, 50],
    ["Total", 640, 520],
])

wb.save("sample_data/sample_workbook_with_notes.xlsx")
print("saved")
