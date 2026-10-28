"""
Privacy layer — local redaction / pseudonymisation before anything reaches an LLM.

What it does
------------
1. Registered names (entity, directors/signatories, organisations, addresses) are
   replaced by stable tokens such as [ENTITY_1] / [PERSON_2] before text leaves the
   machine, and put back locally in whatever comes back.
2. India-specific identifiers are found by pattern and tokenised: PAN, TAN, GSTIN,
   CIN, DIN, ICAI membership no., firm registration no., UDIN, ISIN, BSE scrip
   codes, NSE symbol, e-mail, phone, URL/domain, PIN code.
3. An "egress guard" (assert_clean) re-checks the exact outgoing text and refuses to
   send if anything registered is still in it.
4. PDF text is extracted locally (pypdf) so that the *redacted text*, never the PDF
   file itself, is what a cloud model sees.

What it does NOT do (state this honestly)
-----------------------------------------
Redaction is not anonymisation. Financial figures can re-identify a company to
anyone who knows the sector. Detection is only as good as the list you give it plus
the patterns below, which is why the app shows a preview and asks for approval
before any text is sent.
"""

from __future__ import annotations

import io
import re
from collections import Counter
from dataclasses import dataclass, field

__all__ = [
    "RedactionError", "NoTextLayerError", "Redactor", "RedactionResult",
    "parse_terms", "review_flags", "DEFAULT_LINE_KEYWORDS", "DEFAULT_CUT_MARKERS", "LINE_TOKEN", "describe_patterns",
    "pdf_page_count", "parse_page_range", "extract_pdf_text",
    "excel_sheet_info", "extract_excel_text",
]


class RedactionError(Exception):
    """Raised when text that must be clean still contains masked items."""


class NoTextLayerError(RedactionError):
    """Raised when a PDF has no usable text layer (scanned image)."""


# ---------------------------------------------------------------------------
# Word lists
# ---------------------------------------------------------------------------
HONORIFICS = {
    "mr", "mrs", "ms", "miss", "dr", "shri", "smt", "sri", "prof", "capt",
    "col", "maj", "justice", "hon", "ca", "cs", "cma", "adv",
}
NAME_STOPWORDS = {"and", "of", "the", "for", "co"}
# Words too generic to mask on their own when they are merely a *part* of an
# entity name (the full name is always masked in full regardless).
GENERIC_ENTITY_WORDS = {
    "limited", "ltd", "private", "pvt", "public", "company", "corporation", "corp",
    "india", "indian", "group", "holdings", "holding", "industries", "industry",
    "enterprises", "enterprise", "services", "service", "solutions", "technologies",
    "technology", "systems", "global", "international", "national", "bank",
    "finance", "financial", "capital", "trading", "traders", "infrastructure",
    "infra", "projects", "ventures", "partners", "associates", "llp", "inc",
    "new", "first", "general", "united", "hospital", "hospitals", "health",
    "healthcare", "pharma", "pharmaceuticals", "the", "and", "of", "for",
}
GENERIC_ADDRESS_FRAGMENTS = {"india"}

_LEGAL_SUFFIX = re.compile(
    r"\s+(?:private\s+limited|pvt\.?\s*ltd\.?|limited|ltd\.?|llp|l\.l\.p\.|inc\.?|"
    r"corporation|corp\.?)\s*$",
    re.I,
)


# Lines containing any of these words are removed whole (see Redactor.line_keywords).
DEFAULT_LINE_KEYWORDS = [
    "DIN", "Director Identification Number", "CIN", "Corporate Identity Number",
    "Address", "Registered Office", "Regd. Office", "Corporate Office",
    "PAN", "TAN", "GSTIN", "UDIN", "Firm Registration", "Membership No",
    "Telephone", "Phone", "Mobile", "Fax", "Email", "E-mail", "Website",
]
LINE_TOKEN = "[REDACTED_LINE]"
# From a line containing one of these, the rest of that sheet/page is dropped (the audit-report
# signature block: directors, auditor, membership no., UDIN, place, date).
DEFAULT_CUT_MARKERS = ["As per our report of even date", "For and on behalf of the Board"]
_SECTION_START = re.compile(r"^\[(?:Sheet: |Page )")

# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip())


def parse_terms(text: str) -> list[str]:
    """Turn a textarea ('one per line', ';' also accepted) into a clean list."""
    out, seen = [], set()
    for raw in re.split(r"[\n;]+", text or ""):
        t = _norm(re.sub(r"^[\s\-•*]+", "", raw))
        if len(t) >= 2 and t.lower() not in seen:
            seen.add(t.lower())
            out.append(t)
    return out


def _strip_legal_suffix(name: str) -> str:
    n, prev = _norm(name), None
    while prev != n:
        prev = n
        n = _LEGAL_SUFFIX.sub("", n)
    return n


def _strip_honorifics(name: str) -> str:
    words = _norm(name).split(" ")
    while words and words[0].strip(".").lower() in HONORIFICS:
        words.pop(0)
    return " ".join(words)


def _term_regex(term: str) -> str:
    """Whitespace-flexible, '&'/'and'-tolerant, boundary-guarded regex for a term."""
    pieces = []
    for p in _norm(term).split(" "):
        if p.lower() in ("&", "and"):
            pieces.append(r"(?:&|and)")
        else:
            pieces.append(re.escape(p))
    return r"(?<![A-Za-z0-9])" + r"\s+".join(pieces) + r"(?![A-Za-z0-9])"


def _name_parts(name: str, generic: set[str]) -> list[str]:
    parts = []
    for w in re.split(r"[\s,]+", _norm(name)):
        w = w.strip(".,()&'’-")
        lw = w.lower()
        if (len(w) >= 3 and lw not in generic and lw not in HONORIFICS
                and lw not in NAME_STOPWORDS and not w.isdigit()):
            parts.append(w)
    return parts


# ---------------------------------------------------------------------------
# Identifier patterns (applied after registered names). Order matters: contextual
# and long patterns first, bare patterns later.
# (kind, regex, flags, description)
# Group 1 (if present) is the part that gets replaced; the label stays readable.
# ---------------------------------------------------------------------------
_PATTERN_SPECS = [
    ("UDIN", r"\bUDIN\s*(?:no\.?)?\s*[:\-–]?\s*(\d{8}[A-Z0-9]{10})\b", re.I,
     "UDIN (18 characters)"),
    ("FRN", r"\b(?:firm(?:['’]s)?\s+reg(?:istration|n)?\.?\s*(?:no\.?|number)?|FRN)"
            r"\s*[:\-–]?\s*(\d{4,7}[A-Z](?:\s*/\s*[A-Z]\s*-?\s*\d{5,7})?)", re.I,
     "Firm registration no. / FRN"),
    ("MEMBERSHIP", r"\b(?:membership|member|mem|m)\.?\s*(?:no\.?|number|num\.?)"
                   r"\s*[:\-–.]?\s*(\d{5,7})\b", re.I,
     "ICAI membership no."),
    ("DIN", r"\bDIN\s*(?:no\.?|number)?\s*[:\-–]?\s*(\d{8})\b", re.I,
     "Director Identification Number"),
    ("ISIN", r"\b(IN[EF][A-Z0-9]{8}\d)\b", 0, "ISIN"),
    ("CIN", r"\b([LU]\d{5}[A-Z]{2}\d{4}[A-Z]{3}\d{6})\b", 0, "CIN (21 characters)"),
    ("GSTIN", r"\b(\d{2}[A-Z]{5}\d{4}[A-Z][1-9A-Z]Z[0-9A-Z])\b", 0, "GSTIN"),
    ("PAN", r"\b([A-Z]{5}\d{4}[A-Z])\b", 0, "PAN"),
    ("TAN", r"\b([A-Z]{4}\d{5}[A-Z])\b", 0, "TAN"),
    ("SCRIP", r"\b(?:scrip\s*code|bse\s*code|security\s*code)s?\s*[:\-–]?\s*"
              r"((?:\d{5,6}(?:\s*\([A-Za-z]{1,4}\))?(?:\s*(?:,|&|and)\s*(?=\d))?)+)", re.I,
     "BSE scrip code(s)"),
    ("SYMBOL", r"\b(?i:nse\s*symbol|trading\s*symbol|stock\s*symbol)\s*[:\-–]?\s*"
               r"([A-Z][A-Z0-9&\-]{2,15})\b", 0, "NSE / trading symbol"),
    ("EMAIL", r"\b([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,})\b", 0, "E-mail address"),
    ("URL", r"((?:https?://|www\.)[^\s<>()\[\],;]+)", re.I, "Web address"),
    ("DOMAIN", r"\b([a-z0-9][a-z0-9\-]{1,}(?:\.[a-z0-9\-]{2,})*"
               r"\.(?:com|co\.in|in|org|net|gov\.in|info|biz|edu))\b", re.I, "Domain name"),
    ("PHONE", r"\b(?:tel(?:ephone)?|phone|ph|mobile|mob|fax|contact(?:\s*no\.?)?)"
              r"\s*(?:no\.?|number|nos\.?)?\s*[:.\-–]?\s*(\+?\d[\d\s\-()]{6,}\d)", re.I,
     "Phone / fax (after a label)"),
    ("PHONE", r"(\+91[\s\-]?\(?\d{2,5}\)?[\s\-]?\d{5,8}(?:[\s\-]?\d{2,4})?)", 0,
     "Phone with +91"),
    ("PHONE", r"\b(0\d{2,4}[\s\-]\d{6,8})\b", 0, "Phone with STD code"),
    ("PIN", r"\b(?:pin(?:\s*code)?|pincode|zip)\s*[:\-–]?\s*([1-9]\d{2}\s?\d{3})\b", re.I,
     "PIN code (after a label)"),
    ("PIN", r"(?<=[A-Za-z,\]])\s*[–\-]\s*([1-9]\d{2}\s?\d{3})\b", 0,
     "PIN code (city – 400 001 form)"),
]
_PATTERNS = [(k, re.compile(rx, fl), d) for k, rx, fl, d in _PATTERN_SPECS]


def describe_patterns() -> list[tuple[str, str]]:
    """(kind, description) pairs for the UI, de-duplicated and in order."""
    seen, out = set(), []
    for kind, _, desc in _PATTERNS:
        if (kind, desc) not in seen:
            seen.add((kind, desc))
            out.append((kind, desc))
    return out


# ---------------------------------------------------------------------------
# Result container
# ---------------------------------------------------------------------------
@dataclass
class RedactionResult:
    text: str
    entries: list[dict] = field(default_factory=list)      # hits: kind, token, original, hits
    unmatched: list[dict] = field(default_factory=list)    # registered but never found

    @property
    def total_hits(self) -> int:
        return sum(e["hits"] for e in self.entries)


# ---------------------------------------------------------------------------
# The redactor
# ---------------------------------------------------------------------------
class Redactor:
    """
    Build one per operation from the current privacy settings. Deterministic:
    the same inputs always produce the same tokens.
    """

    def __init__(self, entity: str = "", persons=(), orgs=(), addresses=(),
                 mask_name_parts: bool = True, use_patterns: bool = True,
                 line_keywords=(), cut_markers=()):
        self.cut_markers = [m for m in (_norm(x) for x in cut_markers) if m]
        self._cut_rx = (re.compile("|".join(_term_regex(m)[len(r"(?<![A-Za-z0-9])"):-len(r"(?![A-Za-z0-9])")]
                                    for m in sorted(self.cut_markers, key=len, reverse=True)), re.I)
                        if self.cut_markers else None)
        self.line_keywords = [k for k in (_norm(x) for x in line_keywords) if k]
        self._line_rx = (re.compile(r"(?<![A-Za-z0-9])(?:" + "|".join(
            _term_regex(k)[len(r"(?<![A-Za-z0-9])"):-len(r"(?![A-Za-z0-9])")]
            for k in sorted(self.line_keywords, key=len, reverse=True)) + r")(?![A-Za-z0-9])",
            re.I) if self.line_keywords else None)
        self.mask_name_parts = mask_name_parts
        self.use_patterns = use_patterns
        self._counters: Counter = Counter()
        self._token_by_canon: dict[tuple[str, str], str] = {}
        self._value_by_token: dict[str, str] = {}
        self._kind_by_token: dict[str, str] = {}
        self._variants: list[tuple[str, str, bool]] = []   # (text, token, is_part)
        self._registered: list[tuple[str, str]] = []       # (token, canonical)

        if _norm(entity):
            self._add_entity(entity)
        for p in persons:
            self._add_person(p)
        for o in orgs:
            self._add_org(o)
        for a in addresses:
            self._add_address(a)
        self._build()

    # ---- registration -----------------------------------------------------
    def _new_token(self, kind: str) -> str:
        self._counters[kind] += 1
        return f"[{kind}_{self._counters[kind]}]"

    def _register(self, kind: str, canonical: str) -> str:
        key = (kind, canonical.lower())
        if key in self._token_by_canon:
            return self._token_by_canon[key]
        tok = self._new_token(kind)
        self._token_by_canon[key] = tok
        self._value_by_token[tok] = canonical
        self._kind_by_token[tok] = kind
        self._registered.append((tok, canonical))
        return tok

    def _variant(self, text: str, token: str, is_part: bool):
        text = _norm(text)
        if len(text) >= 2:
            self._variants.append((text, token, is_part))

    def _add_entity(self, name: str):
        name = _norm(name)
        tok = self._register("ENTITY", name)
        base = _strip_legal_suffix(name)
        self._variant(name, tok, False)
        if base and base.lower() != name.lower():
            self._variant(base, tok, False)
        if self.mask_name_parts:
            for w in _name_parts(base, GENERIC_ENTITY_WORDS):
                self._variant(w, tok, True)

    def _add_person(self, name: str):
        name = _norm(name)
        if not name:
            return
        tok = self._register("PERSON", name)
        bare = _strip_honorifics(name)
        self._variant(name, tok, False)
        if bare and bare.lower() != name.lower():
            self._variant(bare, tok, False)
        if self.mask_name_parts:
            for w in _name_parts(bare, set()):
                self._variant(w, tok, True)

    def _add_org(self, name: str):
        name = _norm(name)
        if not name:
            return
        tok = self._register("ORG", name)
        base = _strip_legal_suffix(name)
        self._variant(name, tok, False)
        if base and base.lower() != name.lower():
            self._variant(base, tok, False)
        if self.mask_name_parts:
            for w in _name_parts(base, GENERIC_ENTITY_WORDS):
                self._variant(w, tok, True)

    def _add_address(self, addr: str):
        addr = _norm(addr)
        if not addr:
            return
        tok = self._register("ADDRESS", addr)
        self._variant(addr, tok, False)
        if self.mask_name_parts:
            for frag in addr.split(","):
                frag = _norm(frag)
                if (len(frag) >= 5 and re.search(r"[A-Za-z0-9]", frag)
                        and frag.lower() not in GENERIC_ADDRESS_FRAGMENTS):
                    self._variant(frag, tok, True)

    # ---- compile ----------------------------------------------------------
    def _compile(self, variants):
        if not variants:
            return None, []
        ordered = sorted(variants, key=lambda v: -len(v[0]))       # longest first
        seen, uniq = set(), []
        for text, tok, is_part in ordered:
            k = text.lower()
            if k not in seen:                                       # first registration wins
                seen.add(k)
                uniq.append((text, tok, is_part))
        rx = re.compile("|".join(f"({_term_regex(t)})" for t, _, _ in uniq), re.I)
        return rx, [tok for _, tok, _ in uniq]

    def _build(self):
        self._rx_all, self._tok_all = self._compile(self._variants)
        self._rx_full, self._tok_full = self._compile(
            [v for v in self._variants if not v[2]])

    # ---- pattern tokens ---------------------------------------------------
    def _token_for_value(self, kind: str, value: str) -> str:
        key = (kind, re.sub(r"\s+", "", value).upper())
        if key in self._token_by_canon:
            return self._token_by_canon[key]
        tok = self._new_token(kind)
        self._token_by_canon[key] = tok
        self._value_by_token[tok] = _norm(value)
        self._kind_by_token[tok] = kind
        return tok

    @staticmethod
    def _is_amount_row(line: str) -> bool:
        """True for a statement/note row with at least two amount cells after its label."""
        cells = [c.strip() for c in line.split("|")][1:]
        n = sum(1 for c in cells if re.fullmatch(r"\(?-?[\d,]+(?:\.\d+)?\)?", c or "x"))
        return n >= 2

    # ---- whole-line removal -----------------------------------------------
    def _wipe_lines(self, text: str, line_hits: Counter) -> str:
        """Replace every line containing a trigger word. If such a line ends with
        ':' or '-' (a label whose value wraps onto the next line), the next
        non-empty line is removed too."""
        if self._line_rx is None and self._cut_rx is None:
            return text
        out, carry, cutting = [], False, False
        for line in text.split("\n"):
            if _SECTION_START.match(line.strip()):
                cutting = carry = False
                out.append(line)
                continue
            if cutting:
                continue
            if self._cut_rx is not None:
                cm = self._cut_rx.search(line)
                if cm:
                    line_hits["signature block after \u2018" + cm.group(0).lower() + "\u2019"] += 1
                    out.append(LINE_TOKEN)
                    cutting = True
                    continue
            m = self._line_rx.search(line) if self._line_rx is not None else None
            if m and self._is_amount_row(line):
                m = None                      # e.g. "Telephone expenses | 540 | 540": a financial row
            if m:
                line_hits[m.group(0).strip().upper()] += 1
                out.append(LINE_TOKEN)
                carry = line.rstrip().endswith((":", "-", "\u2013"))
            elif carry and line.strip():
                line_hits["(continuation line)"] += 1
                out.append(LINE_TOKEN)
                carry = False
            else:
                out.append(line)
        return "\n".join(out)

    # ---- public: redact ---------------------------------------------------
    def redact(self, text: str) -> RedactionResult:
        text = text or ""
        hits: Counter = Counter()
        line_hits: Counter = Counter()
        text = self._wipe_lines(text, line_hits)

        if self._rx_all is not None:
            def term_repl(m):
                tok = self._tok_all[m.lastindex - 1]
                hits[tok] += 1
                return tok
            text = self._rx_all.sub(term_repl, text)

        if self.use_patterns:
            for kind, rx, _ in _PATTERNS:
                def pat_repl(m, kind=kind, rx=rx):
                    g = 1 if rx.groups else 0
                    tok = self._token_for_value(kind, m.group(g))
                    hits[tok] += 1
                    if g == 0:
                        return tok
                    s0 = m.start(0)
                    return m.group(0)[: m.start(g) - s0] + tok + m.group(0)[m.end(g) - s0:]
                text = rx.sub(pat_repl, text)

        entries = [
            {"kind": self._kind_by_token[t], "token": t,
             "original": self._value_by_token[t], "hits": n}
            for t, n in hits.items()
        ]
        for kw, n in line_hits.items():
            entries.append({"kind": "LINE", "token": LINE_TOKEN,
                            "original": (kw if kw.startswith("signature block") else f"line containing \u2018{kw}\u2019"),
                            "hits": n})
        entries.sort(key=lambda e: (e["kind"], e["token"]))
        unmatched = [{"kind": self._kind_by_token[t], "token": t, "original": c}
                     for t, c in self._registered if t not in hits]
        return RedactionResult(text=text, entries=entries, unmatched=unmatched)

    # ---- public: restore --------------------------------------------------
    def restore(self, text):
        """Put real values back into text that came back from the model."""
        if not isinstance(text, str) or not self._value_by_token:
            return text
        for tok in sorted(self._value_by_token, key=len, reverse=True):
            name = tok.strip("[]")
            rx = re.compile(r"\[?\b" + re.escape(name) + r"\b\]?", re.I)
            val = self._value_by_token[tok]
            text = rx.sub(lambda m, v=val: v, text)
        return text

    def restore_obj(self, obj):
        if isinstance(obj, str):
            return self.restore(obj)
        if isinstance(obj, dict):
            return {k: self.restore_obj(v) for k, v in obj.items()}
        if isinstance(obj, (list, tuple)):
            return [self.restore_obj(v) for v in obj]
        return obj

    def entity_alias(self) -> str:
        """What the model should call the entity (a token, never the real name)."""
        for tok, kind in self._kind_by_token.items():
            if kind == "ENTITY":
                return tok
        return "the Entity"

    # ---- public: egress guard ---------------------------------------------
    def find_leaks(self, text: str, include_parts: bool = True,
                   include_lines: bool = False) -> list[dict]:
        """Registered terms / identifier patterns still present in `text`."""
        text = text or ""
        found: Counter = Counter()
        rx = self._rx_all if include_parts else self._rx_full
        toks = self._tok_all if include_parts else self._tok_full
        if rx is not None:
            for m in rx.finditer(text):
                found[self._kind_by_token[toks[m.lastindex - 1]]] += 1
        if self.use_patterns:
            for kind, prx, _ in _PATTERNS:
                for _m in prx.finditer(text):
                    found[kind] += 1
        if include_lines:
            for line in text.split("\n"):
                if line.strip() == LINE_TOKEN:
                    continue
                if self._cut_rx is not None and self._cut_rx.search(line):
                    found["LINE"] += 1
                elif (self._line_rx is not None and self._line_rx.search(line)
                      and not self._is_amount_row(line)):
                    found["LINE"] += 1
        return [{"kind": k, "count": n} for k, n in sorted(found.items())]

    def assert_clean(self, text: str, where: str = "outgoing text",
                     include_parts: bool = True, include_lines: bool = False) -> None:
        leaks = self.find_leaks(text, include_parts=include_parts,
                                include_lines=include_lines)
        if leaks:
            summary = ", ".join(f"{l['kind']}×{l['count']}" for l in leaks)
            raise RedactionError(
                f"Blocked: {where} still contains masked items ({summary}). "
                "Nothing was sent. Remove them from the text or add them to the "
                "redaction list and try again."
            )

    # ---- convenience -------------------------------------------------------
    @property
    def n_registered(self) -> int:
        return len(self._registered)


# ---------------------------------------------------------------------------
# Heuristic hints: things that LOOK like names but are not yet masked.
# Deliberately simple regexes — hints for a human, not a detector.
# ---------------------------------------------------------------------------
_FLAG_HONORIFIC = re.compile(
    r"\b(?:Mr|Mrs|Ms|Miss|Dr|Shri|Smt|Sri|Prof|Capt|Justice)\.?\s+"
    r"([A-Z][A-Za-z.'\-]+(?:\s+[A-Z][A-Za-z.'\-]+){0,3})")
_FLAG_ROLE = re.compile(
    r"(?m)^[ \t]*([A-Z][A-Za-z.'\-]+(?:[ \t]+[A-Z][A-Za-z.'\-]+){1,3})[ \t]*\n[ \t]*"
    r"(?:Chairman|Chairperson|Managing Director|Whole[- ]?time Director|Director|"
    r"Chief Financial Officer|Chief Executive Officer|CFO|CEO|Company Secretary|"
    r"Partner|Proprietor)")
_FLAG_FIRM = re.compile(
    r"\b((?:[A-Z][A-Za-z.&']*\s+){1,4}(?:LLP|&\s*Co\.?|and\s+Co\.?|Chartered Accountants))")
_FLAG_PLACE = re.compile(
    r"(?m)^[ \t]*(?:Place|Registered\s+Office|Regd\.?\s+Office)\s*:\s*(.+?)[ \t]*$")


def review_flags(text: str, limit: int = 20) -> list[str]:
    """Possible unmasked names/places in (already redacted) text."""
    flags, seen = [], set()

    def add(label, value):
        value = _norm(value)
        if value and "[" not in value and value.lower() not in seen:
            seen.add(value.lower())
            flags.append(f"{label}: {value}")

    for m in _FLAG_HONORIFIC.finditer(text or ""):
        add("Name after honorific", m.group(0))
    for m in _FLAG_ROLE.finditer(text or ""):
        add("Name above a designation", m.group(1))
    for m in _FLAG_FIRM.finditer(text or ""):
        add("Firm name", m.group(1))
    for m in _FLAG_PLACE.finditer(text or ""):
        add("Place / office line", m.group(1))
    return flags[:limit]


# ---------------------------------------------------------------------------
# Local PDF text extraction
# ---------------------------------------------------------------------------
def _pdf_reader(pdf_bytes: bytes):
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise RedactionError("pypdf is not installed. Run: pip install pypdf") from e
    return PdfReader(io.BytesIO(pdf_bytes))


def pdf_page_count(pdf_bytes: bytes) -> int:
    return len(_pdf_reader(pdf_bytes).pages)


def parse_page_range(spec: str, total: int) -> list[int]:
    """'168-169, 171' -> [168, 169, 171] (1-based). Blank -> all pages."""
    spec = (spec or "").strip()
    if not spec:
        return list(range(1, total + 1))
    spec = re.sub(r"\s*[-–]\s*", "-", spec)
    pages: list[int] = []
    for chunk in re.split(r"[,;\s]+", spec):
        if not chunk:
            continue
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", chunk)
        if not m:
            raise ValueError(f"Cannot read page range '{chunk}'. Use e.g. 168-169, 171.")
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        if a > b:
            a, b = b, a
        if a < 1 or b > total:
            raise ValueError(f"Pages must be between 1 and {total}.")
        pages.extend(range(a, b + 1))
    return sorted(dict.fromkeys(pages))


def extract_pdf_text(pdf_bytes: bytes, pages: list[int] | None = None,
                     min_chars_per_page: int = 120) -> str:
    """
    Local text extraction (nothing leaves the machine). Page markers are added so
    the model can tell pages apart. Raises NoTextLayerError for scanned PDFs.
    """
    reader = _pdf_reader(pdf_bytes)
    total = len(reader.pages)
    pages = pages or list(range(1, total + 1))
    chunks, chars = [], 0
    for p in pages:
        txt = (reader.pages[p - 1].extract_text() or "").strip()
        chars += len(txt)
        chunks.append(f"[Page {p}]\n{txt}")
    if chars / max(len(pages), 1) < min_chars_per_page:
        raise NoTextLayerError(
            "This PDF has no usable text layer (it looks like a scan). It cannot be "
            "redacted locally without OCR. Use the Excel template, or turn Privacy "
            "mode off to send the file as-is."
        )
    return "\n\n".join(chunks)


# ---------------------------------------------------------------------------
# Local Excel workbook reading (multi-sheet, notes to accounts, trial balances)
# ---------------------------------------------------------------------------
def _xl_fmt(v) -> str:
    import datetime as _dt
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v == int(v) and abs(v) < 1e15 else f"{v:.2f}"
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%d-%m-%Y")
    return re.sub(r"\s+", " ", str(v)).strip()


def _xl_open(xlsx_bytes: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise RedactionError("openpyxl is not installed. Run: pip install openpyxl") from e
    try:
        return load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    except Exception as e:
        raise RedactionError(
            "Could not open this file as an .xlsx/.xlsm workbook. Old .xls files must be "
            f"re-saved as .xlsx first. ({e})") from e


def excel_sheet_info(xlsx_bytes: bytes) -> list[dict]:
    """[{name, rows, hidden}] for every sheet, read locally."""
    wb = _xl_open(xlsx_bytes)
    out = []
    for ws in wb.worksheets:
        rows = 0
        for r in ws.iter_rows(values_only=True):
            if any(c not in (None, "") for c in r):
                rows += 1
        out.append({"name": ws.title, "rows": rows,
                    "hidden": getattr(ws, "sheet_state", "visible") != "visible"})
    wb.close()
    return out


def extract_excel_text(xlsx_bytes: bytes, sheets: list[str] | None = None,
                       max_rows_per_sheet: int = 3000) -> str:
    """
    Turn the chosen sheets into plain text, one line per non-empty row, cells joined by
    ' | ', each sheet introduced by '[Sheet: name]'. Nothing leaves the computer here.
    """
    wb = _xl_open(xlsx_bytes)
    chunks, numeric = [], 0
    for ws in wb.worksheets:
        if sheets is not None and ws.title not in sheets:
            continue
        lines = [f"[Sheet: {ws.title}]"]
        n = 0
        for r in ws.iter_rows(values_only=True):
            cells = [_xl_fmt(c) for c in r]
            while cells and not cells[-1]:
                cells.pop()
            if not any(cells):
                continue
            cells = [c for i, c in enumerate(cells)
                     if c or (i > 0 and cells[i - 1] and True)]      # drop runs of empty cells
            while cells and not cells[0]:
                cells.pop(0)
            numeric += sum(1 for c in r if isinstance(c, (int, float)) and not isinstance(c, bool))
            lines.append(" | ".join(cells))
            n += 1
            if n >= max_rows_per_sheet:
                lines.append(f"[... sheet truncated after {max_rows_per_sheet} rows ...]")
                break
        if n:
            chunks.append("\n".join(lines))
    wb.close()
    if not chunks:
        raise RedactionError("The selected sheets contain no readable cells.")
    if numeric == 0:
        raise RedactionError(
            "No numbers were found. If the figures are formulas that were never calculated, "
            "open the file in Excel, save it, and upload it again.")
    return "\n\n".join(chunks)
