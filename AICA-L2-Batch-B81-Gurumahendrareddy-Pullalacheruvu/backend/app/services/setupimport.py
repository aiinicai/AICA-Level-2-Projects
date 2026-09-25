"""The set-up workbook: build it, validate it, commit it.

A CA lives in Excel. An in-app form asking for four hundred figures will not be
finished; a workbook the finance team already knows how to fill will be. So the
same schema drives three things — the blank template, the worked example, and
the validator — which is the only way those three stay in step.

The validation rule is the same one the plan importer follows: **a file that
silently half-imports is worse than one that is rejected.** Every row is either
accepted or rejected with a row number and a sentence saying what is wrong.
Nothing is written until the whole file has been read.
"""
from __future__ import annotations

import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from sqlalchemy.orm import Session

from app.models import (
    BankAccount, Bill, BurnCategory, Commitment, CommitmentType, CostNature,
    Criticality, Customer, Entity, Invoice, LedgerEntry, StatutoryDue,
    StatutoryHead, Vendor,
)
from app.services.common import Ctx, fmt_inr

TRUE_WORDS = {"y", "yes", "true", "1", "t"}
FALSE_WORDS = {"n", "no", "false", "0", "f", ""}
DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d-%b-%y",
                "%d.%m.%Y", "%m/%d/%Y")


# ---------------------------------------------------------------------------
# Schema
# ---------------------------------------------------------------------------
@dataclass
class Col:
    header: str
    key: str
    kind: str = "text"                 # text|number|date|bool|choice
    required: bool = False
    choices: list[str] = field(default_factory=list)
    help: str = ""
    width: int = 18


@dataclass
class Sheet:
    name: str
    stage: int
    title: str
    why: str
    cols: list[Col]


SHEETS: list[Sheet] = [
    Sheet("Bank Accounts", 1,
          "Where the money is",
          "Cash available is the numerator of every runway figure in the tool. "
          "A restricted balance with no reason recorded is treated as available, "
          "so the reason column matters as much as the amount.",
          [Col("Bank / Institution", "institution", "text", True, width=26),
           Col("Account name", "account_name", "text", True, width=26),
           Col("Last 4 digits", "account_masked", "text", help="Last four only — never the full number.", width=13),
           Col("Purpose", "purpose", "choice", False,
               ["Operating", "Payroll", "Deposit", "Collections", "Tax"], width=15),
           Col("Balance", "balance", "number", True, help="In rupees.", width=16),
           Col("As on", "as_on", "date", True, width=13),
           Col("Restricted?", "is_restricted", "bool", help="Y if lien-marked, escrowed or pledged.", width=13),
           Col("Restriction reason", "restriction_reason", "text",
               help="Required if Restricted is Y.", width=34)]),

    Sheet("Receipts & Payments", 1,
          "What has actually moved",
          "Three months of bank movement is what net burn is computed from. "
          "Without it the tool can show a balance but not a runway.",
          [Col("Date", "txn_date", "date", True, width=13),
           Col("Type", "kind", "choice", True, ["Receipt", "Payment"], width=12),
           Col("Party", "party", "text", True, width=28),
           Col("Category", "burn_category", "choice", False, BurnCategory.ALL,
               help="Payments only. Leave blank on receipts.", width=26),
           Col("Amount", "amount", "number", True, help="Positive. The Type column says the direction.", width=16),
           Col("Nature", "cost_nature", "choice", False, CostNature.ALL, width=15),
           Col("One-off?", "is_one_off", "bool",
               help="Y excludes it from normalised burn. State why in the note.", width=12),
           Col("Note", "narration", "text", width=34)]),

    Sheet("Open Invoices", 2,
          "What is owed to you",
          "Drives ageing, DSO, concentration and the weighted collectible figure.",
          [Col("Customer", "customer", "text", True, width=28),
           Col("Invoice no", "invoice_no", "text", True, width=16),
           Col("Invoice date", "invoice_date", "date", True, width=14),
           Col("Due date", "due_date", "date", True, width=14),
           Col("Invoice amount", "amount", "number", True, width=16),
           Col("Still outstanding", "outstanding", "number", True, width=17),
           Col("Agreed terms (days)", "credit_terms_days", "number", width=18),
           Col("Likelihood of collection %", "collection_probability", "number",
               help="0 to 100. Your judgement, not a formula. Blank = 100.", width=22),
           Col("Disputed?", "is_disputed", "bool", width=12),
           Col("Dispute reason", "dispute_reason", "text", width=30)]),

    Sheet("Open Bills", 2,
          "What you owe",
          "Ordered on screen by whether it can wait, not by size.",
          [Col("Vendor", "vendor", "text", True, width=28),
           Col("Bill no", "bill_no", "text", True, width=16),
           Col("Bill date", "bill_date", "date", True, width=14),
           Col("Due date", "due_date", "date", True, width=14),
           Col("Amount", "amount", "number", True, width=16),
           Col("Still outstanding", "outstanding", "number", True, width=17),
           Col("Category", "burn_category", "choice", False, BurnCategory.ALL, width=26),
           Col("Can it wait?", "deferrable", "bool",
               help="N for statutory, payroll and anything with a penalty.", width=14),
           Col("Cost of deferring", "deferral_cost", "number",
               help="Penal interest or late fee, in rupees.", width=17),
           Col("Description", "description", "text", width=30)]),

    Sheet("Statutory Dues", 2,
          "First charge on cash",
          "These carry penal interest and cannot be negotiated, so they get "
          "their own layer rather than sitting among the payables.",
          [Col("Head", "head", "choice", True, StatutoryHead.ALL, width=18),
           Col("Period", "period", "text", True, help='e.g. "Aug-26" or "Q2 FY26-27".', width=16),
           Col("Due date", "due_date", "date", True, width=14),
           Col("Amount", "amount", "number", True, width=16),
           Col("Already earmarked", "earmarked_amount", "number",
               help="Cash set aside for it. Blank = nothing set aside.", width=18),
           Col("Reference", "reference", "text", width=18),
           Col("Notes", "notes", "text", width=30)]),

    Sheet("Committed Not Billed", 3,
          "Money already promised",
          "The number most cash tools miss entirely: spend that is committed but "
          "for which no invoice has arrived, so it is in no ledger anywhere.",
          [Col("Type", "commitment_type", "choice", True, CommitmentType.ALL, width=24),
           Col("Counterparty", "counterparty", "text", True, width=28),
           Col("Description", "description", "text", width=32),
           Col("Total value", "total_value", "number", True, width=16),
           Col("Consumed so far", "consumed_to_date", "number", width=17),
           Col("Cancellable?", "cancellable", "bool", width=14),
           Col("Notice period (days)", "notice_period_days", "number", width=19),
           Col("Cost to exit", "exit_cost", "number", width=15),
           Col("Starts", "starts_on", "date", width=13),
           Col("Ends", "ends_on", "date", width=13),
           Col("Monthly run-rate", "monthly_runrate", "number", width=18),
           Col("Category", "burn_category", "choice", False, BurnCategory.ALL, width=26)]),
]

SHEETS_BY_NAME = {s.name: s for s in SHEETS}
STAGE_LABELS = {
    1: "Stage 1 — enough for a runway number",
    2: "Stage 2 — open items, so the calendar and ageing work",
    3: "Stage 3 — the judgement layer",
}


# ---------------------------------------------------------------------------
# Coercion
# ---------------------------------------------------------------------------
def _as_date(raw) -> date | None:
    if raw in (None, ""):
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _as_number(raw) -> float | None:
    if raw in (None, ""):
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).strip().replace(",", "").replace("₹", "").replace("Rs.", "").replace("Rs", "")
    neg = s.startswith("(") and s.endswith(")")
    if neg:
        s = s[1:-1]
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if neg else v


def _as_bool(raw) -> bool | None:
    if raw is None:
        return False
    if isinstance(raw, bool):
        return raw
    s = str(raw).strip().lower()
    if s in TRUE_WORDS:
        return True
    if s in FALSE_WORDS:
        return False
    return None


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------
def build_template(example_rows: dict[str, list[dict]] | None = None) -> bytes:
    """The workbook. With `example_rows`, the worked example instead of a blank.

    Both come out of this one function on purpose: a sample file that has
    drifted from the template it illustrates is worse than no sample at all.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    _write_readme(wb, bool(example_rows))

    head_fill = PatternFill("solid", fgColor="1E2761")
    req_fill = PatternFill("solid", fgColor="2E3D7A")
    head_font = Font(color="FFFFFF", bold=True, size=10)
    note_font = Font(color="5B6478", italic=True, size=9)

    for sheet in SHEETS:
        ws = wb.create_sheet(sheet.name)

        ws.cell(row=1, column=1, value=sheet.title).font = Font(bold=True, size=13, color="1E2761")
        ws.cell(row=2, column=1, value=sheet.why).font = note_font
        ws.merge_cells(start_row=2, start_column=1, end_row=2,
                       end_column=max(len(sheet.cols), 4))
        ws.cell(row=3, column=1, value=STAGE_LABELS[sheet.stage]).font = note_font

        hdr = 5
        for i, col in enumerate(sheet.cols, start=1):
            c = ws.cell(row=hdr, column=i, value=col.header + (" *" if col.required else ""))
            c.fill = req_fill if col.required else head_fill
            c.font = head_font
            c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = col.width
            if col.help:
                c.comment = None
                ws.cell(row=hdr + 1, column=i, value=col.help).font = note_font
        ws.row_dimensions[hdr].height = 32
        ws.row_dimensions[hdr + 1].height = 26
        ws.freeze_panes = ws.cell(row=hdr + 2, column=1)

        # Dropdowns, so a choice column cannot be filled in with free text.
        for i, col in enumerate(sheet.cols, start=1):
            letter = get_column_letter(i)
            if col.kind == "choice" and col.choices:
                joined = ",".join(col.choices)
                if len(joined) < 250:
                    dv = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=True)
                    ws.add_data_validation(dv)
                    dv.add(f"{letter}{hdr + 2}:{letter}500")
            elif col.kind == "bool":
                dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{letter}{hdr + 2}:{letter}500")

        rows = (example_rows or {}).get(sheet.name, [])
        for r, row in enumerate(rows, start=hdr + 2):
            for i, col in enumerate(sheet.cols, start=1):
                v = row.get(col.key)
                if isinstance(v, bool):
                    v = "Y" if v else "N"
                cell = ws.cell(row=r, column=i, value=v)
                if col.kind == "number":
                    cell.number_format = "#,##0"
                elif col.kind == "date":
                    cell.number_format = "DD-MM-YYYY"

        # Number/date formats on the empty rows too, so typing behaves.
        for r in range(hdr + 2 + len(rows), hdr + 2 + len(rows) + 80):
            for i, col in enumerate(sheet.cols, start=1):
                if col.kind == "number":
                    ws.cell(row=r, column=i).number_format = "#,##0"
                elif col.kind == "date":
                    ws.cell(row=r, column=i).number_format = "DD-MM-YYYY"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_readme(wb, is_example: bool) -> None:
    from openpyxl.styles import Font

    ws = wb.active
    ws.title = "Read me first"
    ws.column_dimensions["A"].width = 108

    lines: list[tuple[str, str]] = [
        ("Cash Runway — set-up workbook" + ("  (worked example)" if is_example else ""), "h1"),
        ("", ""),
        ("Fill in the sheets in order. You do not have to do all of them at once.", ""),
        ("", ""),
        (STAGE_LABELS[1], "h2"),
        ("Bank Accounts and three months of Receipts & Payments. That alone gives you", ""),
        ("cash available, net burn, runway and the cash-out date. Ten minutes of typing", ""),
        ("gets you a number you can defend.", ""),
        ("", ""),
        (STAGE_LABELS[2], "h2"),
        ("Open Invoices, Open Bills and Statutory Dues. These turn on ageing, DSO,", ""),
        ("concentration, the 13-week calendar and the statutory funding gap.", ""),
        ("", ""),
        (STAGE_LABELS[3], "h2"),
        ("Committed Not Billed — spend already promised that no invoice has arrived for.", ""),
        ("Most cash tools miss it entirely, which is how a company with a comfortable", ""),
        ("balance runs out anyway.", ""),
        ("", ""),
        ("Rules", "h2"),
        ("•  Amounts in rupees. Not lakhs, not crores, no formatting — 4520000, not 45.2 L.", ""),
        ("•  Dates as DD-MM-YYYY. Excel date cells are fine too.", ""),
        ("•  Y / N in the yes-no columns.", ""),
        ("•  A column marked * must be filled in. A row missing one is rejected, with the", ""),
        ("   row number and the reason — nothing is imported half-way.", ""),
        ("•  Enter every amount as a positive number. The Type column says the direction.", ""),
        ("•  Do not rename the sheets or move the header row.", ""),
        ("", ""),
        ("What happens when you upload it", "h2"),
        ("The file is read and checked before anything is saved. You get a report listing", ""),
        ("every accepted and every rejected row. Only when you accept that report is", ""),
        ("anything written, and every figure carries your name and the date as its source.", ""),
    ]
    if is_example:
        lines += [
            ("", ""),
            ("About this example", "h2"),
            ("These are the figures for Northwind Robotics Pvt Ltd, the demonstration", ""),
            ("company. Use it to see the shape a filled-in sheet takes, then clear the rows", ""),
            ("and enter your own — or upload it as-is against a test entity.", ""),
        ]

    for i, (text, style) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if style == "h1":
            c.font = Font(bold=True, size=15, color="1E2761")
        elif style == "h2":
            c.font = Font(bold=True, size=11, color="1E2761")


# ---------------------------------------------------------------------------
# Read and validate
# ---------------------------------------------------------------------------
def validate(content: bytes) -> dict:
    """Read every sheet and report row by row. Writes nothing."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        raise ValueError(f"That file could not be opened as a workbook ({e}). "
                         f"Save it as .xlsx and try again.")

    parsed: dict[str, list[dict]] = {}
    errors: list[dict] = []
    summary: list[dict] = []

    for sheet in SHEETS:
        if sheet.name not in wb.sheetnames:
            summary.append({"sheet": sheet.name, "stage": sheet.stage,
                            "accepted": 0, "rejected": 0, "present": False})
            continue

        ws = wb[sheet.name]
        header_row = _find_header(ws, sheet)
        if header_row is None:
            errors.append({"sheet": sheet.name, "row": None,
                           "message": "The header row is missing or was renamed. "
                                      "Download a fresh template and copy your rows in."})
            summary.append({"sheet": sheet.name, "stage": sheet.stage,
                            "accepted": 0, "rejected": 0, "present": True})
            continue

        index = _column_index(ws, header_row, sheet)
        rows, errs = _read_rows(ws, header_row, index, sheet)
        parsed[sheet.name] = rows
        errors.extend(errs)
        summary.append({"sheet": sheet.name, "stage": sheet.stage,
                        "accepted": len(rows), "rejected": len(errs), "present": True})

    cross = _cross_checks(parsed)
    accepted = sum(s["accepted"] for s in summary)
    rejected = sum(s["rejected"] for s in summary)

    return {
        "summary": summary,
        "accepted": accepted,
        "rejected": rejected,
        "errors": errors[:400],
        "errors_truncated": len(errors) > 400,
        "warnings": cross,
        "parsed": parsed,
        "verdict": _verdict(accepted, rejected, summary),
        "stages_covered": sorted({s["stage"] for s in summary if s["accepted"]}),
    }


def _find_header(ws, sheet: Sheet) -> int | None:
    """Locate the header by its first column, rather than trusting a row number."""
    want = sheet.cols[0].header.lower()
    for r in range(1, 12):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().lower().rstrip("*").strip() == want:
            return r
    return None


def _column_index(ws, header_row: int, sheet: Sheet) -> dict[str, int]:
    """Map key → column number by header text, so column order can move."""
    found: dict[str, int] = {}
    headers = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v:
            headers[str(v).strip().lower().rstrip("*").strip()] = c
    for col in sheet.cols:
        n = headers.get(col.header.lower())
        if n:
            found[col.key] = n
    return found


def _read_rows(ws, header_row: int, index: dict[str, int],
               sheet: Sheet) -> tuple[list[dict], list[dict]]:
    rows: list[dict] = []
    errors: list[dict] = []
    seen_keys: set[str] = set()

    for r in range(header_row + 1, ws.max_row + 1):
        raw = {col.key: (ws.cell(row=r, column=index[col.key]).value
                         if col.key in index else None)
               for col in sheet.cols}

        # Skip the help row and any blank row.
        if all(v in (None, "") for v in raw.values()):
            continue
        helps = {c.help for c in sheet.cols if c.help}
        if any(isinstance(v, str) and v in helps for v in raw.values()):
            continue

        rec: dict = {}
        problems: list[str] = []
        for col in sheet.cols:
            v = raw.get(col.key)
            if col.key not in index:
                if col.required:
                    problems.append(f"the column '{col.header}' is missing from the sheet")
                continue

            if col.kind == "date":
                d = _as_date(v)
                if v not in (None, "") and d is None:
                    problems.append(f"'{col.header}' is not a date I can read ({v!r}) — use DD-MM-YYYY")
                rec[col.key] = d
            elif col.kind == "number":
                n = _as_number(v)
                if v not in (None, "") and n is None:
                    problems.append(f"'{col.header}' is not a number ({v!r})")
                rec[col.key] = n
            elif col.kind == "bool":
                b = _as_bool(v)
                if b is None:
                    problems.append(f"'{col.header}' should be Y or N, not {v!r}")
                rec[col.key] = bool(b)
            elif col.kind == "choice":
                s = (str(v).strip() if v not in (None, "") else None)
                if s and col.choices and s not in col.choices:
                    problems.append(f"'{col.header}' must be one of: {', '.join(col.choices)} "
                                    f"— found {s!r}")
                rec[col.key] = s
            else:
                rec[col.key] = (str(v).strip() if v not in (None, "") else None)

            if col.required and rec.get(col.key) in (None, "", 0.0) and col.kind != "bool":
                if not (col.kind == "number" and rec.get(col.key) == 0.0):
                    problems.append(f"'{col.header}' is required and is empty")

        problems.extend(_row_rules(sheet, rec))

        key = _dedupe_key(sheet, rec)
        if key:
            if key in seen_keys:
                problems.append(f"this is a duplicate of an earlier row ({key})")
            seen_keys.add(key)

        if problems:
            errors.append({"sheet": sheet.name, "row": r,
                           "message": "; ".join(problems),
                           "values": {k: str(v) for k, v in raw.items() if v not in (None, "")}})
        else:
            rec["_row"] = r
            rows.append(rec)

    return rows, errors


def _dedupe_key(sheet: Sheet, rec: dict) -> str | None:
    if sheet.name == "Open Invoices":
        return f"invoice {rec.get('invoice_no')}"
    if sheet.name == "Open Bills":
        return f"bill {rec.get('bill_no')}"
    if sheet.name == "Statutory Dues":
        return f"{rec.get('head')} {rec.get('period')}"
    if sheet.name == "Bank Accounts":
        return f"{rec.get('institution')} {rec.get('account_name')}"
    return None


def _row_rules(sheet: Sheet, rec: dict) -> list[str]:
    """Checks that need more than one cell — where the real errors live."""
    p: list[str] = []

    if sheet.name == "Bank Accounts":
        if rec.get("is_restricted") and not rec.get("restriction_reason"):
            p.append("marked restricted but no reason given — a restricted balance "
                     "with no reason is treated as available")
        if (rec.get("balance") or 0) < 0:
            p.append("balance is negative — record an overdraft as a facility, not a bank account")

    elif sheet.name == "Receipts & Payments":
        if (rec.get("amount") or 0) <= 0:
            p.append("amount must be positive — the Type column says the direction")
        if rec.get("kind") == "Payment" and not rec.get("burn_category"):
            p.append("a payment needs a category, or it cannot appear in burn by category")
        if rec.get("is_one_off") and not rec.get("narration"):
            p.append("marked one-off but no note — normalised burn is only trustworthy "
                     "if the exclusions are explained")
        d = rec.get("txn_date")
        if d and d > date.today():
            p.append(f"dated {d:%d-%m-%Y}, which is in the future")

    elif sheet.name == "Open Invoices":
        a, o = rec.get("amount"), rec.get("outstanding")
        if a is not None and o is not None and o > a + 0.5:
            p.append(f"outstanding ({fmt_inr(o)}) is more than the invoice ({fmt_inr(a)})")
        if rec.get("invoice_date") and rec.get("due_date") and rec["due_date"] < rec["invoice_date"]:
            p.append("due date is before the invoice date")
        prob = rec.get("collection_probability")
        if prob is not None and not (0 <= prob <= 100):
            p.append("likelihood of collection must be between 0 and 100")
        if rec.get("is_disputed") and not rec.get("dispute_reason"):
            p.append("marked disputed but no reason given")

    elif sheet.name == "Open Bills":
        a, o = rec.get("amount"), rec.get("outstanding")
        if a is not None and o is not None and o > a + 0.5:
            p.append(f"outstanding ({fmt_inr(o)}) is more than the bill ({fmt_inr(a)})")
        if rec.get("bill_date") and rec.get("due_date") and rec["due_date"] < rec["bill_date"]:
            p.append("due date is before the bill date")

    elif sheet.name == "Statutory Dues":
        if (rec.get("earmarked_amount") or 0) > (rec.get("amount") or 0) + 0.5:
            p.append("more is earmarked than is due")

    elif sheet.name == "Committed Not Billed":
        if (rec.get("consumed_to_date") or 0) > (rec.get("total_value") or 0) + 0.5:
            p.append("consumed so far is more than the total value")
        if rec.get("starts_on") and rec.get("ends_on") and rec["ends_on"] < rec["starts_on"]:
            p.append("ends before it starts")

    return p


def _cross_checks(parsed: dict[str, list[dict]]) -> list[str]:
    """Whole-file observations. These warn; they do not reject."""
    w: list[str] = []

    banks = parsed.get("Bank Accounts", [])
    if banks:
        dates = {b["as_on"] for b in banks if b.get("as_on")}
        if len(dates) > 1:
            w.append(f"The bank balances are as at {len(dates)} different dates "
                     f"({', '.join(d.strftime('%d-%m-%Y') for d in sorted(dates))}). "
                     f"Cash available will be stated as at the earliest of them.")
        if not any(b.get("is_restricted") for b in banks):
            w.append("No balance is marked restricted. If anything is lien-marked, "
                     "escrowed or pledged, say so — otherwise the tool will treat "
                     "the whole balance as spendable.")

    rp = parsed.get("Receipts & Payments", [])
    if rp:
        ds = [r["txn_date"] for r in rp if r.get("txn_date")]
        if ds:
            span = (max(ds) - min(ds)).days
            if span < 60:
                w.append(f"The receipts and payments cover {span} days. Net burn is a "
                         f"three-month average, so with less than that it will be "
                         f"computed on what is here and labelled accordingly.")
        if not any(r.get("kind") == "Receipt" for r in rp):
            w.append("There are no receipts, only payments. Net burn will equal gross burn.")

    inv = parsed.get("Open Invoices", [])
    if inv:
        total = sum(i.get("outstanding") or 0 for i in inv)
        by_customer: dict[str, float] = {}
        for i in inv:
            by_customer[i["customer"]] = by_customer.get(i["customer"], 0) + (i.get("outstanding") or 0)
        if total and by_customer:
            top, amt = max(by_customer.items(), key=lambda kv: kv[1])
            if amt / total > 0.4:
                w.append(f"{top} is {amt / total:.0%} of receivables. That concentration "
                         f"will drive the exposure line on Money Coming In.")

    if not parsed.get("Statutory Dues"):
        w.append("No statutory dues entered. These are first charge on cash and carry "
                 "penal interest — an empty sheet here understates what must be paid.")

    return w


def _verdict(accepted: int, rejected: int, summary: list[dict]) -> str:
    stage1 = [s for s in summary if s["stage"] == 1 and s["accepted"]]
    if not accepted:
        return ("Nothing could be read from that file. Check that the sheets are named "
                "as the template names them.")
    if len(stage1) < 2:
        return (f"{accepted} row(s) can be imported, but Stage 1 is incomplete. Bank "
                f"Accounts and Receipts & Payments are what produce a runway number.")
    if rejected:
        return (f"{accepted} row(s) ready to import, {rejected} rejected. Fix the rejected "
                f"rows and upload again, or import the good rows now and add the rest later.")
    return f"All {accepted} row(s) read cleanly. Nothing has been saved yet."


# ---------------------------------------------------------------------------
# Commit
# ---------------------------------------------------------------------------
def commit(db: Session, entity: Entity, parsed: dict[str, list[dict]],
           user_name: str) -> dict:
    """Write the accepted rows. Called only after a person has seen the report."""
    from app.models.base import DataSource

    src = DataSource.UPLOAD.value
    counts: dict[str, int] = {}

    def prov(obj):
        obj.source = src
        obj.created_by = user_name
        return obj

    # -- banks ---------------------------------------------------------------
    for r in parsed.get("Bank Accounts", []):
        db.add(prov(BankAccount(
            entity_id=entity.id, institution=r["institution"],
            account_name=r["account_name"], account_masked=r.get("account_masked"),
            purpose=r.get("purpose") or "Operating",
            balance=r["balance"], books_balance=r["balance"],
            is_restricted=bool(r.get("is_restricted")),
            restriction_reason=r.get("restriction_reason"),
            as_on=r["as_on"])))
    counts["bank_accounts"] = len(parsed.get("Bank Accounts", []))

    # -- movement ------------------------------------------------------------
    for r in parsed.get("Receipts & Payments", []):
        receipt = r["kind"] == "Receipt"
        amt = abs(r["amount"])
        db.add(LedgerEntry(
            entity_id=entity.id, txn_date=r["txn_date"],
            voucher_type="Receipt" if receipt else "Payment",
            party=r.get("party"), narration=r.get("narration"),
            debit=amt if receipt else 0.0, credit=0.0 if receipt else amt,
            cash_amount=amt if receipt else -amt,
            burn_category=None if receipt else r.get("burn_category"),
            cost_nature=None if receipt else r.get("cost_nature"),
            is_one_off=bool(r.get("is_one_off")),
            one_off_note=r.get("narration") if r.get("is_one_off") else None,
            classified_by=user_name, source="upload"))
    counts["ledger_entries"] = len(parsed.get("Receipts & Payments", []))

    # -- receivables ---------------------------------------------------------
    customers: dict[str, Customer] = {c.name: c for c in db.query(Customer)
                                      .filter(Customer.entity_id == entity.id).all()}
    for r in parsed.get("Open Invoices", []):
        c = customers.get(r["customer"])
        if c is None:
            c = prov(Customer(entity_id=entity.id, name=r["customer"],
                              credit_terms_days=int(r.get("credit_terms_days") or 30)))
            db.add(c)
            db.flush()
            customers[r["customer"]] = c
        prob = r.get("collection_probability")
        db.add(prov(Invoice(
            entity_id=entity.id, customer_id=c.id, invoice_no=r["invoice_no"],
            invoice_date=r["invoice_date"], due_date=r["due_date"],
            amount=r["amount"], outstanding=r["outstanding"],
            status="open" if r["outstanding"] > 0 else "paid",
            collection_probability=(prob / 100.0) if prob is not None else None,
            is_disputed=bool(r.get("is_disputed")),
            dispute_reason=r.get("dispute_reason"))))
    counts["invoices"] = len(parsed.get("Open Invoices", []))
    counts["customers"] = len(customers)

    # -- payables ------------------------------------------------------------
    vendors: dict[str, Vendor] = {v.name: v for v in db.query(Vendor)
                                  .filter(Vendor.entity_id == entity.id).all()}
    for r in parsed.get("Open Bills", []):
        v = vendors.get(r["vendor"])
        if v is None:
            v = prov(Vendor(entity_id=entity.id, name=r["vendor"],
                            criticality=Criticality.MEDIUM,
                            burn_category=r.get("burn_category")))
            db.add(v)
            db.flush()
            vendors[r["vendor"]] = v
        db.add(prov(Bill(
            entity_id=entity.id, vendor_id=v.id, bill_no=r["bill_no"],
            bill_date=r["bill_date"], due_date=r["due_date"],
            amount=r["amount"], outstanding=r["outstanding"],
            description=r.get("description"), burn_category=r.get("burn_category"),
            deferrable=bool(r.get("deferrable")),
            deferral_cost=r.get("deferral_cost") or 0.0,
            status="unpaid" if r["outstanding"] > 0 else "paid")))
    counts["bills"] = len(parsed.get("Open Bills", []))
    counts["vendors"] = len(vendors)

    # -- statutory -----------------------------------------------------------
    for r in parsed.get("Statutory Dues", []):
        ear = r.get("earmarked_amount") or 0.0
        db.add(prov(StatutoryDue(
            entity_id=entity.id, head=r["head"], period=r["period"],
            due_date=r["due_date"], amount=r["amount"],
            earmarked_amount=ear, funded=ear >= r["amount"] - 0.5,
            reference=r.get("reference"), notes=r.get("notes"))))
    counts["statutory_dues"] = len(parsed.get("Statutory Dues", []))

    # -- commitments ---------------------------------------------------------
    for r in parsed.get("Committed Not Billed", []):
        db.add(prov(Commitment(
            entity_id=entity.id, commitment_type=r["commitment_type"],
            counterparty=r["counterparty"], description=r.get("description"),
            total_value=r["total_value"],
            consumed_to_date=r.get("consumed_to_date") or 0.0,
            cancellable=bool(r.get("cancellable")),
            notice_period_days=int(r.get("notice_period_days") or 0),
            exit_cost=r.get("exit_cost") or 0.0,
            starts_on=r.get("starts_on"), ends_on=r.get("ends_on"),
            monthly_runrate=r.get("monthly_runrate") or 0.0,
            burn_category=r.get("burn_category"))))
    counts["commitments"] = len(parsed.get("Committed Not Billed", []))

    db.flush()
    return counts
