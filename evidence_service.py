"""Read supporting client documents in memory with bounded resource use."""
import csv
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from docx import Document
from services.pdf_service import extract_text_from_pdf
from utils.helpers import normalize_text
from utils.validators import AppError

MAX_FILES = 10
MAX_BATCH_BYTES = 50 * 1024 * 1024
MAX_DOCUMENT_CHARS = 150_000


def bounded(text: str) -> str:
    if len(text) > MAX_DOCUMENT_CHARS:
        raise AppError("A supporting document exceeds the text limit. Prepare a smaller relevant document; no text was silently omitted.")
    if not text.strip():
        raise AppError("A supporting document contains no readable text. Please upload a searchable or scanned PDF instead.")
    return normalize_text(text)


def check_office_archive(data: bytes) -> None:
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 60 * 1024 * 1024:
                raise AppError("This Office document is too large when unpacked. Export a smaller PDF.")
    except BadZipFile:
        raise AppError("This Office document is damaged or encrypted. Export an unlocked PDF.") from None


def extract_supporting_document(data: bytes, name: str) -> tuple[str, list[str]]:
    if not data or len(data) > 20 * 1024 * 1024:
        raise AppError("Each supporting document must be non-empty and no larger than 20 MB.")
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".pdf":
            result = extract_text_from_pdf(data, name)
            return bounded(result.text), list(result.warnings)
        if suffix in {".docx", ".xlsx"}:
            check_office_archive(data)
        if suffix == ".docx":
            document = Document(BytesIO(data))
            lines = [p.text for p in document.paragraphs]
            for index, table in enumerate(document.tables, 1):
                lines.append(f"Table {index}")
                lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            return bounded("\n".join(lines)), ["Word body paragraphs and tables were read. Images, text boxes, headers, comments and tracked changes need manual review; export to PDF if material."]
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
            lines, length, cells = [], 0, 0
            try:
                for sheet in workbook.worksheets:
                    lines.append(f"Sheet: {sheet.title}")
                    for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        cells += len(row)
                        if cells > 100_000:
                            raise AppError("This workbook is too large to review safely. Export the relevant sheets to PDF.")
                        if any(value is not None for value in row):
                            line = f"Row {row_index}: " + " | ".join("" if v is None else str(v) for v in row)
                            length += len(line)
                            if length > MAX_DOCUMENT_CHARS:
                                raise AppError("This workbook exceeds the text limit. Export the relevant sheets to PDF.")
                            lines.append(line)
            finally:
                workbook.close()
            return bounded("\n".join(lines)), ["Formula expressions are preserved, not calculated. Charts, drawings and embedded objects are not read. Verify displayed totals and supporting sheets."]
        if suffix in {".txt", ".csv"}:
            text = data.decode("utf-8-sig")
            if suffix == ".csv":
                text = "\n".join(f"Row {i}: " + " | ".join(row) for i, row in enumerate(csv.reader(StringIO(text)), 1))
            return bounded(text), []
        raise AppError("Supporting documents must be PDF, DOCX, XLSX, CSV or UTF-8 TXT. Export images as PDF.")
    except AppError:
        raise
    except Exception:
        raise AppError("A supporting document could not be read. Check its format, remove password protection, or export it to PDF.") from None
