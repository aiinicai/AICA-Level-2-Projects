"""
AI Auditor V8 - Sample Financial Statement Generator
Creates realistic sample Excel and PDF financial statements for manufacturing and trading companies.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def generate_manufacturing_excel(filepath: str):
    wb = openpyxl.Workbook()
    
    # Styles
    navy = "1A365D"
    steel = "2B6CB0"
    font_head = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_norm = Font(name="Segoe UI", size=10)
    fill_head = PatternFill(start_color=navy, end_color=navy, fill_type="solid")
    fill_sub = PatternFill(start_color=steel, end_color=steel, fill_type="solid")
    fill_alt = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
    
    # 1. Balance Sheet
    ws_bs = wb.active
    ws_bs.title = "Balance Sheet"
    ws_bs["A1"] = "BHARAT DYNAMICS & HEAVY ENGINEERING LTD"
    ws_bs["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="1A365D")
    ws_bs["A2"] = "Balance Sheet as at 31st March 2024 (Amount in ₹ Lakhs)"
    ws_bs["A2"].font = Font(name="Segoe UI", size=10, italic=True)

    headers = ["Particulars", "Note No.", "31st March 2024 (FY 2023-24)", "31st March 2023 (FY 2022-23)"]
    for c, h in enumerate(headers, 1):
        cell = ws_bs.cell(row=4, column=c, value=h)
        cell.font = font_head
        cell.fill = fill_head
        cell.alignment = Alignment(horizontal="center", vertical="center")

    bs_data = [
        ("I. EQUITY AND LIABILITIES", "", "", ""),
        ("1. Shareholders' Funds", "", "", ""),
        ("Share Capital", "1", 1250.00, 1250.00),
        ("Reserves and Surplus", "2", 4850.50, 4210.00),
        ("2. Non-Current Liabilities", "", "", ""),
        ("Long-Term Borrowings (Term Loans)", "3", 2100.00, 1650.00),
        ("Deferred Tax Liabilities (Net)", "4", 320.00, 290.00),
        ("3. Current Liabilities", "", "", ""),
        ("Short-Term Borrowings (Cash Credit Limits)", "5", 1850.00, 1200.00),
        ("Trade Payables", "6", 1420.00, 1150.00),
        ("Other Current Liabilities & Provisions", "7", 460.00, 380.00),
        ("TOTAL EQUITY AND LIABILITIES", "", 12250.50, 10130.00),
        ("", "", "", ""),
        ("II. ASSETS", "", "", ""),
        ("1. Non-Current Assets", "", "", ""),
        ("Property, Plant and Equipment (Gross Block)", "8", 5400.00, 4600.00),
        ("Capital Work-in-Progress (CWIP)", "9", 450.00, 210.00),
        ("Non-Current Investments", "10", 820.00, 800.00),
        ("Other Non-Current Assets", "11", 280.50, 220.00),
        ("2. Current Assets", "", "", ""),
        ("Inventories (Raw Material, WIP, FG)", "12", 2350.00, 1850.00),
        ("Trade Receivables (Debtors)", "13", 2210.00, 1720.00),
        ("Cash and Cash Equivalents", "14", 390.00, 410.00),
        ("Other Current Assets & Advances", "15", 350.00, 320.00),
        ("TOTAL ASSETS", "", 12250.50, 10130.00)
    ]

    for r_idx, row in enumerate(bs_data, 5):
        for c_idx, val in enumerate(row, 1):
            cell = ws_bs.cell(row=r_idx, column=c_idx, value=val)
            if isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")
                cell.font = font_bold if "TOTAL" in str(row[0]) else font_norm
            else:
                cell.font = font_bold if (str(val).startswith("I.") or str(val).startswith("1.") or str(val).startswith("2.") or str(val).startswith("3.") or "TOTAL" in str(val)) else font_norm
                
    ws_bs.column_dimensions["A"].width = 45
    ws_bs.column_dimensions["B"].width = 12
    ws_bs.column_dimensions["C"].width = 25
    ws_bs.column_dimensions["D"].width = 25

    # 2. Profit and Loss
    ws_pl = wb.create_sheet(title="Profit and Loss")
    ws_pl["A1"] = "BHARAT DYNAMICS & HEAVY ENGINEERING LTD"
    ws_pl["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="1A365D")
    ws_pl["A2"] = "Statement of Profit and Loss for the Year Ended 31st March 2024 (₹ in Lakhs)"
    ws_pl["A2"].font = Font(name="Segoe UI", size=10, italic=True)

    headers_pl = ["Particulars", "Note No.", "FY 2023-24", "FY 2022-23"]
    for c, h in enumerate(headers_pl, 1):
        cell = ws_pl.cell(row=4, column=c, value=h)
        cell.font = font_head
        cell.fill = fill_head
        cell.alignment = Alignment(horizontal="center", vertical="center")

    pl_data = [
        ("Revenue from Operations (Gross Sales)", "16", 14500.00, 12200.00),
        ("Other Income", "17", 180.00, 140.00),
        ("TOTAL INCOME", "", 14680.00, 12340.00),
        ("EXPENSES:", "", "", ""),
        ("Cost of Materials Consumed", "18", 8200.00, 6900.00),
        ("Changes in Inventories of FG & WIP", "19", -250.00, -180.00),
        ("Employee Benefits Expense", "20", 1850.00, 1550.00),
        ("Finance Costs (Interest Expense)", "21", 420.00, 280.00),
        ("Depreciation and Amortisation Expense", "22", 480.00, 420.00),
        ("Power and Fuel Expenses", "23", 640.00, 520.00),
        ("Freight & Transportation Outward", "24", 390.00, 310.00),
        ("Other Operating Expenses", "25", 1520.00, 1280.00),
        ("TOTAL EXPENSES", "", 13250.00, 11080.00),
        ("PROFIT BEFORE EXCEPTIONAL ITEMS & TAX", "", 1430.00, 1260.00),
        ("Tax Expense (Current Tax + Deferred Tax)", "26", 360.00, 315.00),
        ("PROFIT FOR THE PERIOD (PAT)", "", 1070.00, 945.00)
    ]

    for r_idx, row in enumerate(pl_data, 5):
        for c_idx, val in enumerate(row, 1):
            cell = ws_pl.cell(row=r_idx, column=c_idx, value=val)
            if isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")
                cell.font = font_bold if "TOTAL" in str(row[0]) or "PROFIT" in str(row[0]) else font_norm
            else:
                cell.font = font_bold if ("TOTAL" in str(val) or "PROFIT" in str(val) or "EXPENSES" in str(val)) else font_norm

    ws_pl.column_dimensions["A"].width = 45
    ws_pl.column_dimensions["B"].width = 12
    ws_pl.column_dimensions["C"].width = 25
    ws_pl.column_dimensions["D"].width = 25

    # 3. Cash Flow Statement
    ws_cf = wb.create_sheet(title="Cash Flow Statement")
    ws_cf["A1"] = "BHARAT DYNAMICS & HEAVY ENGINEERING LTD"
    ws_cf["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="1A365D")
    ws_cf["A2"] = "Cash Flow Statement for the Year Ended 31st March 2024 (₹ in Lakhs)"
    ws_cf["A2"].font = Font(name="Segoe UI", size=10, italic=True)

    for c, h in enumerate(headers_pl, 1):
        cell = ws_cf.cell(row=4, column=c, value=h)
        cell.font = font_head
        cell.fill = fill_head
        cell.alignment = Alignment(horizontal="center", vertical="center")

    cf_data = [
        ("A. Cash Flow from Operating Activities", "", 1120.00, 980.00),
        ("B. Cash Flow from / (used in) Investing Activities", "", -1040.00, -780.00),
        ("C. Cash Flow from / (used in) Financing Activities", "", -100.00, -150.00),
        ("Net Increase / (Decrease) in Cash and Cash Equivalents", "", -20.00, 50.00),
        ("Cash and Cash Equivalents at the Beginning of the Year", "", 410.00, 360.00),
        ("Cash and Cash Equivalents at the End of the Year", "", 390.00, 410.00)
    ]

    for r_idx, row in enumerate(cf_data, 5):
        for c_idx, val in enumerate(row, 1):
            cell = ws_cf.cell(row=r_idx, column=c_idx, value=val)
            if isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")
                cell.font = font_bold
            else:
                cell.font = font_bold

    ws_cf.column_dimensions["A"].width = 50
    ws_cf.column_dimensions["B"].width = 12
    ws_cf.column_dimensions["C"].width = 25
    ws_cf.column_dimensions["D"].width = 25

    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb.save(filepath)

def generate_sample_pdf(filepath: str):
    """Generates a professional multi-page sample financial statement PDF."""
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    doc = SimpleDocTemplate(filepath, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#1A365D'),
        alignment=1
    )
    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#4A5568'),
        alignment=1
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#2B6CB0')
    )

    elements.append(Paragraph("APEX GLOBAL TRADING CORPORATION LTD", title_style))
    elements.append(Paragraph("Audited Financial Statements for FY 2023-24 (Amount in ₹ Lakhs)", sub_style))
    elements.append(Spacer(1, 15))

    # Balance Sheet Table
    elements.append(Paragraph("Balance Sheet as at 31st March 2024", h2_style))
    elements.append(Spacer(1, 6))

    bs_table_data = [
        ["Particulars", "FY 2023-24", "FY 2022-23"],
        ["Share Capital", "800.00", "800.00"],
        ["Reserves and Surplus", "2,450.00", "2,100.00"],
        ["Long-Term Borrowings", "1,200.00", "950.00"],
        ["Short-Term Borrowings", "1,650.00", "1,100.00"],
        ["Trade Payables", "1,850.00", "1,320.00"],
        ["Other Current Liabilities", "310.00", "260.00"],
        ["Total Equity & Liabilities", "8,260.00", "6,530.00"],
        ["Property, Plant and Equipment", "2,800.00", "2,400.00"],
        ["Non-Current Investments", "450.00", "400.00"],
        ["Inventories", "2,150.00", "1,550.00"],
        ["Trade Receivables", "2,350.00", "1,680.00"],
        ["Cash and Bank Balances", "290.00", "320.00"],
        ["Other Current Assets", "220.00", "180.00"],
        ["Total Assets", "8,260.00", "6,530.00"]
    ]

    t_bs = Table(bs_table_data, colWidths=[260, 130, 130])
    t_bs.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1A365D')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 9),
        ('ALIGN', (1,0), (-1,-1), 'RIGHT'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E0')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F7FAFC')]),
        ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
        ('FONTNAME', (0,7), (-1,7), 'Helvetica-Bold'),
    ]))
    elements.append(t_bs)
    elements.append(Spacer(1, 15))

    # P&L Table
    elements.append(Paragraph("Statement of Profit and Loss for FY 2023-24", h2_style))
    elements.append(Spacer(1, 6))

    pl_table_data = [
        ["Particulars", "FY 2023-24", "FY 2022-23"],
        ["Revenue from Operations", "18,400.00", "15,200.00"],
        ["Other Income", "120.00", "95.00"],
        ["Cost of Materials / Purchases", "13,200.00", "10,800.00"],
        ["Employee Benefits Expense", "1,950.00", "1,650.00"],
        ["Finance Costs", "480.00", "320.00"],
        ["Depreciation Expense", "340.00", "290.00"],
        ["Other Operating Expenses", "1,650.00", "1,380.00"],
        ["Profit Before Tax (PBT)", "900.00", "855.00"],
        ["Tax Expense", "225.00", "215.00"],
        ["Profit After Tax (PAT)", "675.00", "640.00"]
    ]

    t_pl = Table(pl_table_data, colWidths=[260, 130, 130])
    t_pl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#2B6CB0')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 9),
        ('ALIGN', (1,0), (-1,-1), 'RIGHT'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E0')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F7FAFC')]),
        ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
    ]))
    elements.append(t_pl)

    doc.build(elements)

if __name__ == "__main__":
    out_dir = os.path.join(os.path.dirname(__file__))
    xlsx_path = os.path.join(out_dir, "sample_manufacturing_financial_statements.xlsx")
    pdf_path = os.path.join(out_dir, "sample_trading_financial_statements.pdf")
    generate_manufacturing_excel(xlsx_path)
    generate_sample_pdf(pdf_path)
    print("Sample test files generated successfully!")
