"""Tests for core.word_converter's non-COM-dependent pieces.

Actual Word/LibreOffice conversion isn't exercised here (that needs a real
installed Word or LibreOffice binary and is exercised manually / in CI on a
Windows runner with Office installed) -- these tests cover the parts that
run anywhere: backend discovery and .docx metadata inspection.
"""
from __future__ import annotations

import pytest

pytest.importorskip("docx")

from core.word_converter import WordDocumentInfo, find_libreoffice, inspect_word_document


def _make_docx(tmp_path):
    import docx

    doc = docx.Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph with a few more words in it.")
    table = doc.add_table(rows=2, cols=2)
    path = tmp_path / "sample.docx"
    doc.save(str(path))
    return path


def test_inspect_docx_reports_structure(tmp_path):
    path = _make_docx(tmp_path)
    info = inspect_word_document(path)
    assert info.supported
    assert info.paragraph_count >= 2
    assert info.table_count == 1
    assert info.approx_word_count > 0


def test_inspect_legacy_doc_is_unsupported(tmp_path):
    # python-docx cannot read the legacy binary .doc format.
    fake_doc = tmp_path / "legacy.doc"
    fake_doc.write_bytes(b"not a real doc file")
    info = inspect_word_document(fake_doc)
    assert info == WordDocumentInfo(supported=False)


def test_inspect_corrupt_docx_returns_unsupported(tmp_path):
    fake = tmp_path / "broken.docx"
    fake.write_bytes(b"not actually a zip/docx")
    info = inspect_word_document(fake)
    assert not info.supported


def test_find_libreoffice_returns_none_when_absent():
    # On a machine without LibreOffice installed at all, this should not
    # raise -- it should simply report "not found".
    result = find_libreoffice("Z:/definitely/not/a/real/path/soffice.exe")
    # Result depends on whether LibreOffice happens to be installed on the
    # test machine via PATH/well-known locations; we only assert the call
    # is safe and returns either None or an existing path.
    assert result is None or __import__("pathlib").Path(result).exists()
