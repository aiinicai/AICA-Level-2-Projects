"""
==============================================================================
 FAMILY INVESTMENT MATRIX
 A standalone, password-protected family investment dashboard covering
 Equity (HDFC Securities, ICICI Direct, SBI Securities, Zerodha, etc.)
 and Mutual Funds (CAMS, KFintech, etc.)

 HOW TO RUN:
     Recommended: python family_investment_matrix_rectified_v4.py
     Alternative: streamlit run family_investment_matrix_rectified_v4.py
     Missing packages install automatically on first run.

 This single file contains everything: database schema, authentication,
 data/holdings engine, Excel template generator, PDF extractor (beta),
 news fetcher, and the full Streamlit dashboard UI.
==============================================================================
"""

import binascii
import hashlib
import importlib.util
import math
import os
import re
import sqlite3
import subprocess
import sys
from datetime import date, datetime
from io import BytesIO
from urllib.parse import quote


def _prepare_streamlit_first_run():
    """Prevent Streamlit's first-run email prompt on Windows/macOS/Linux."""
    try:
        config_dir = os.path.join(os.path.expanduser("~"), ".streamlit")
        os.makedirs(config_dir, exist_ok=True)
        credentials = os.path.join(config_dir, "credentials.toml")
        if not os.path.exists(credentials):
            with open(credentials, "w", encoding="utf-8") as fh:
                fh.write('[general]\nemail = ""\n')
    except Exception:
        # Failure to create this optional file must never stop the app.
        pass


def ensure_dependencies():
    requirements = {
        "streamlit": "streamlit>=1.38", "pandas": "pandas>=2.0",
        "openpyxl": "openpyxl>=3.1", "pdfplumber": "pdfplumber>=0.11",
        "plotly": "plotly>=5.20", "requests": "requests>=2.31",
        "yfinance": "yfinance>=0.2.40", "feedparser": "feedparser>=6.0",
    }
    missing = [pkg for module, pkg in requirements.items() if importlib.util.find_spec(module) is None]
    if not missing:
        return
    print("Installing missing packages:", ", ".join(missing))
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--upgrade", *missing])
    except subprocess.CalledProcessError as exc:
        raise SystemExit(
            "\nCould not install required packages. Connect to the internet and run:\n"
            f"{sys.executable} -m pip install streamlit pandas openpyxl pdfplumber plotly requests yfinance feedparser\n"
        ) from exc


ensure_dependencies()

def _inside_streamlit_runtime():
    """True only when Streamlit itself is executing this script."""
    try:
        from streamlit.runtime.scriptrunner import get_script_run_ctx
        return get_script_run_ctx(suppress_warning=True) is not None
    except Exception:
        return False


# When started with ordinary `python file.py`, relaunch through Streamlit BEFORE
# any @st.cache_data decorators are created.  This removes the harmless
# "No runtime found" warning seen in the previous version.  If the user instead
# runs `streamlit run file.py`, the runtime check is already True and no relaunch
# occurs.
if __name__ == "__main__" and not _inside_streamlit_runtime():
    _prepare_streamlit_first_run()
    child_env = os.environ.copy()
    child_env["STREAMLIT_BROWSER_GATHER_USAGE_STATS"] = "false"
    cmd = [
        sys.executable, "-m", "streamlit", "run", os.path.abspath(__file__),
        "--server.headless=false",
        "--browser.gatherUsageStats=false",
    ]
    try:
        raise SystemExit(subprocess.call(cmd, env=child_env))
    except KeyboardInterrupt:
        raise SystemExit(0)

import openpyxl
import pandas as pd
import pdfplumber
import plotly.express as px
import requests
import streamlit as st
import yfinance as yf
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

try:
    import feedparser
except ImportError:
    feedparser = None


# ==============================================================================
# 1. CONFIG / CONSTANTS
# ==============================================================================

APP_NAME = "Family Investment Matrix"
ASSET_TYPES = ["Equity", "Mutual Fund"]
TXN_TYPES = ["BUY", "SELL", "SIP", "DIVIDEND", "BONUS", "SPLIT"]
RELATIONS = ["Self", "Spouse", "Son", "Daughter", "Father", "Mother",
             "Brother", "Sister", "Other"]
DEFAULT_PASSWORD = "123456"

# Only these roles may add new family members today.
# To open this up later (per your future access-control plan), just add
# 'member' to this list — nothing else needs to change.
CAN_ADD_MEMBER_ROLES = ["admin"]


def can_add_member(role: str) -> bool:
    return role in CAN_ADD_MEMBER_ROLES


DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "family_investments.db")


# ==============================================================================
# 2. DATABASE LAYER
# ==============================================================================

def get_conn():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False, timeout=30)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_conn()
    cur = conn.cursor()
    cur.executescript("""
    CREATE TABLE IF NOT EXISTS members (
        member_id   INTEGER PRIMARY KEY AUTOINCREMENT,
        name        TEXT NOT NULL,
        relation    TEXT,
        mobile_no   TEXT UNIQUE NOT NULL,
        dob         TEXT,
        pan         TEXT,
        email       TEXT,
        created_at  TEXT DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS users (
        user_id         INTEGER PRIMARY KEY AUTOINCREMENT,
        mobile_no       TEXT UNIQUE NOT NULL,
        password_hash   TEXT NOT NULL,
        salt            TEXT NOT NULL,
        role            TEXT NOT NULL DEFAULT 'member',
        member_id       INTEGER,
        is_first_login  INTEGER DEFAULT 1,
        created_at      TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (member_id) REFERENCES members(member_id)
    );

    CREATE TABLE IF NOT EXISTS brokers (
        broker_id    INTEGER PRIMARY KEY AUTOINCREMENT,
        broker_name  TEXT UNIQUE NOT NULL,
        broker_type  TEXT NOT NULL
    );

    CREATE TABLE IF NOT EXISTS transactions (
        txn_id           INTEGER PRIMARY KEY AUTOINCREMENT,
        member_id        INTEGER NOT NULL,
        broker_id        INTEGER NOT NULL,
        asset_type       TEXT NOT NULL,
        instrument_name  TEXT NOT NULL,
        identifier       TEXT,
        txn_type         TEXT NOT NULL,
        txn_date         TEXT NOT NULL,
        quantity         REAL NOT NULL DEFAULT 0,
        price            REAL NOT NULL DEFAULT 0,
        amount           REAL NOT NULL DEFAULT 0,
        charges          REAL DEFAULT 0,
        folio_no         TEXT,
        remarks          TEXT,
        upload_source    TEXT DEFAULT 'Manual',
        created_by       INTEGER,
        created_at       TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (member_id) REFERENCES members(member_id),
        FOREIGN KEY (broker_id) REFERENCES brokers(broker_id)
    );

    CREATE TABLE IF NOT EXISTS price_cache (
        identifier   TEXT PRIMARY KEY,
        asset_type   TEXT,
        name         TEXT,
        last_price   REAL,
        price_date   TEXT,
        updated_at   TEXT
    );

    CREATE TABLE IF NOT EXISTS instrument_master (
        identifier       TEXT PRIMARY KEY,
        instrument_name  TEXT,
        asset_type       TEXT DEFAULT 'Equity',
        exchange         TEXT DEFAULT 'NSE',
        trading_symbol   TEXT,
        isin             TEXT,
        hdfc_segment     TEXT,
        hdfc_token       TEXT,
        icici_stock_code TEXT,
        upstox_key       TEXT,
        yahoo_symbol     TEXT,
        updated_at       TEXT
    );

    CREATE TABLE IF NOT EXISTS app_settings (
        setting_key   TEXT PRIMARY KEY,
        setting_value TEXT,
        updated_at    TEXT
    );

    CREATE TABLE IF NOT EXISTS news_cache (
        news_id          INTEGER PRIMARY KEY AUTOINCREMENT,
        instrument_name  TEXT,
        title            TEXT,
        link             TEXT UNIQUE,
        source           TEXT,
        published        TEXT,
        summary          TEXT,
        fetched_at       TEXT
    );
    """)
    conn.commit()

    default_brokers = [
        ("HDFC Securities", "Equity Broker"),
        ("ICICI Direct", "Equity Broker"),
        ("SBI Securities", "Equity Broker"),
        ("Zerodha", "Equity Broker"),
        ("Kotak Securities", "Equity Broker"),
        ("Angel One", "Equity Broker"),
        ("Groww", "Equity Broker"),
        ("Upstox", "Equity Broker"),
        ("CAMS", "MF RTA"),
        ("KFintech", "MF RTA"),
        ("Direct - AMC", "MF RTA"),
        ("Other", "Other"),
    ]
    for name, btype in default_brokers:
        cur.execute("INSERT OR IGNORE INTO brokers (broker_name, broker_type) VALUES (?, ?)", (name, btype))
    conn.commit()
    conn.close()


# ==============================================================================
# 3. AUTHENTICATION
# ==============================================================================

def hash_password(password, salt=None):
    if salt is None:
        salt = binascii.hexlify(os.urandom(16)).decode()
    pwd_hash = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 100_000)
    return binascii.hexlify(pwd_hash).decode(), salt


def verify_password(password, salt, stored_hash):
    pwd_hash, _ = hash_password(password, salt)
    return pwd_hash == stored_hash


def admin_exists():
    conn = get_conn()
    row = conn.execute("SELECT 1 FROM users WHERE role='admin' LIMIT 1").fetchone()
    conn.close()
    return row is not None


def create_admin(name, mobile_no, relation="Self"):
    conn = get_conn()
    cur = conn.cursor()
    cur.execute("INSERT INTO members (name, relation, mobile_no) VALUES (?, ?, ?)", (name, relation, mobile_no))
    member_id = cur.lastrowid
    pwd_hash, salt = hash_password(DEFAULT_PASSWORD)
    cur.execute("""INSERT INTO users (mobile_no, password_hash, salt, role, member_id, is_first_login)
                   VALUES (?, ?, ?, 'admin', ?, 1)""", (mobile_no, pwd_hash, salt, member_id))
    conn.commit()
    conn.close()


def add_member(name, mobile_no, relation, dob=None, pan=None, email=None):
    conn = get_conn()
    cur = conn.cursor()
    existing = cur.execute("SELECT 1 FROM members WHERE mobile_no=?", (mobile_no,)).fetchone()
    if existing:
        conn.close()
        return False, "A member with this mobile number already exists."
    cur.execute("""INSERT INTO members (name, relation, mobile_no, dob, pan, email)
                   VALUES (?, ?, ?, ?, ?, ?)""", (name, relation, mobile_no, dob, pan, email))
    member_id = cur.lastrowid
    pwd_hash, salt = hash_password(DEFAULT_PASSWORD)
    cur.execute("""INSERT INTO users (mobile_no, password_hash, salt, role, member_id, is_first_login)
                   VALUES (?, ?, ?, 'member', ?, 1)""", (mobile_no, pwd_hash, salt, member_id))
    conn.commit()
    conn.close()
    return True, "Member added successfully."


def authenticate(mobile_no, password):
    conn = get_conn()
    row = conn.execute("SELECT * FROM users WHERE mobile_no=?", (mobile_no,)).fetchone()
    conn.close()
    if row is None:
        return None
    if verify_password(password, row["salt"], row["password_hash"]):
        return dict(row)
    return None


def change_password(user_id, new_password):
    pwd_hash, salt = hash_password(new_password)
    conn = get_conn()
    conn.execute("UPDATE users SET password_hash=?, salt=?, is_first_login=0 WHERE user_id=?",
                 (pwd_hash, salt, user_id))
    conn.commit()
    conn.close()


# ==============================================================================
# 4. DATA / HOLDINGS / PRICES / NEWS / VALIDATION
# ==============================================================================

def get_members():
    conn = get_conn()
    df = pd.read_sql_query("SELECT * FROM members ORDER BY name", conn)
    conn.close()
    return df


def get_member_name(member_id):
    conn = get_conn()
    row = conn.execute("SELECT name FROM members WHERE member_id=?", (member_id,)).fetchone()
    conn.close()
    return row["name"] if row else "Unknown"


def get_brokers():
    conn = get_conn()
    df = pd.read_sql_query("SELECT * FROM brokers ORDER BY broker_name", conn)
    conn.close()
    return df


def add_broker(name, btype="Other"):
    conn = get_conn()
    conn.execute("INSERT OR IGNORE INTO brokers (broker_name, broker_type) VALUES (?, ?)", (name, btype))
    conn.commit()
    conn.close()


def load_all_transactions():
    conn = get_conn()
    query = """
        SELECT t.*, m.name AS member_name, b.broker_name, b.broker_type
        FROM transactions t
        JOIN members m ON t.member_id = m.member_id
        JOIN brokers b ON t.broker_id = b.broker_id
    """
    df = pd.read_sql_query(query, conn)
    conn.close()
    return df


def parse_date_safe(val):
    try:
        return pd.to_datetime(val, dayfirst=True).strftime("%Y-%m-%d")
    except Exception:
        return None


def safe_float(val):
    try:
        if val is None or pd.isna(val):
            return 0.0
        result = float(str(val).replace(",", "").strip())
        return result if math.isfinite(result) else 0.0
    except Exception:
        return 0.0


def clean_optional_text(val):
    if val is None:
        return None
    try:
        if pd.isna(val):
            return None
    except Exception:
        pass
    text = str(val).strip()
    return None if (not text or text.lower() in {"nan", "none", "nat"}) else text


def validate_transaction_fields(data):
    errors = []
    if not data.get("member_id"):
        errors.append("Member not recognized / not selected.")
    if not data.get("broker_id"):
        errors.append("Broker/Platform not recognized / not selected.")
    if not data.get("instrument_name"):
        errors.append("Instrument name is required.")
    if data.get("asset_type") not in ("Equity", "Mutual Fund"):
        errors.append("Asset type must be 'Equity' or 'Mutual Fund'.")
    if data.get("txn_type") not in TXN_TYPES:
        errors.append(f"Transaction type must be one of {TXN_TYPES}.")

    txn_date = data.get("txn_date")
    if not txn_date:
        errors.append("Transaction date is missing or invalid.")
    else:
        try:
            d = datetime.strptime(txn_date, "%Y-%m-%d").date()
            if d > date.today():
                errors.append("Transaction date cannot be in the future.")
            if d.year < 1990:
                errors.append("Transaction date looks incorrect (before 1990).")
        except Exception:
            errors.append("Transaction date format is invalid.")

    qty = data.get("quantity") or 0
    amt = data.get("amount") or 0
    price = data.get("price") or 0

    if data.get("txn_type") in ("BUY", "SELL", "SIP", "BONUS") and qty <= 0:
        errors.append("Quantity must be greater than 0 for this transaction type.")
    if amt < 0 or price < 0 or qty < 0:
        errors.append("Quantity, Price and Amount cannot be negative.")
    if data.get("txn_type") in ("BUY", "SELL", "SIP") and qty and price:
        expected = qty * price
        if amt and abs(expected - amt) > max(50, 0.05 * expected):
            errors.append(f"Amount (Rs.{amt:,.2f}) does not match Qty x Price (Rs.{expected:,.2f}). Please verify.")
    return errors


def check_duplicate(member_id, broker_id, instrument_name, txn_date, txn_type, quantity, amount):
    """Flags a possible duplicate if same member+broker+instrument+date+type
    also matches closely on quantity OR amount."""
    conn = get_conn()
    rows = conn.execute(
        """SELECT * FROM transactions
           WHERE member_id=? AND broker_id=? AND instrument_name=? AND txn_date=? AND txn_type=?""",
        (member_id, broker_id, instrument_name, txn_date, txn_type)).fetchall()
    conn.close()
    matches = []
    for r in rows:
        r = dict(r)
        qty_close = abs((r["quantity"] or 0) - (quantity or 0)) < 0.001
        amt_close = abs((r["amount"] or 0) - (amount or 0)) <= max(1.0, 0.01 * abs(amount or 0))
        if qty_close or amt_close:
            matches.append(r)
    return matches


def insert_transaction(data, force=False):
    if not force:
        dup = check_duplicate(data["member_id"], data["broker_id"], data["instrument_name"],
                               data["txn_date"], data["txn_type"], data["quantity"], data["amount"])
        if dup:
            return "DUPLICATE", "Possible duplicate transaction found.", dup

    conn = get_conn()
    conn.execute("""INSERT INTO transactions
        (member_id, broker_id, asset_type, instrument_name, identifier, txn_type, txn_date,
         quantity, price, amount, charges, folio_no, remarks, upload_source, created_by)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (data["member_id"], data["broker_id"], data["asset_type"], data["instrument_name"],
         data.get("identifier"), data["txn_type"], data["txn_date"], data["quantity"],
         data["price"], data["amount"], data.get("charges", 0), data.get("folio_no"),
         data.get("remarks"), data.get("upload_source", "Manual"), data.get("created_by")))
    conn.commit()
    conn.close()
    return "OK", "Transaction saved successfully.", []


def compute_holdings(transactions_df):
    if transactions_df.empty:
        return pd.DataFrame()

    df = transactions_df.copy()

    # Keep identifiers consistent irrespective of whether they came from
    # SQLite, Excel, manual entry, or pandas' NaN/float inference.
    # Examples: 125497.0 -> "125497", RELIANCE -> "RELIANCE", NaN -> "".
    if "identifier" in df.columns:
        df["identifier"] = df["identifier"].apply(_clean_identifier)

    df["txn_date"] = pd.to_datetime(df["txn_date"])
    df = df.sort_values("txn_date")

    group_cols = ["member_id", "member_name", "broker_id", "broker_name",
                  "asset_type", "instrument_name", "identifier"]
    results = []
    for keys, grp in df.groupby(group_cols, dropna=False):
        net_qty = 0.0
        net_invested = 0.0
        realized_gain = 0.0
        dividend_income = 0.0
        for _, row in grp.iterrows():
            ttype = str(row["txn_type"]).upper()
            qty = row["quantity"] or 0
            amt = row["amount"] or 0
            chg = row["charges"] or 0
            if ttype in ("BUY", "SIP"):
                net_qty += qty
                net_invested += amt + chg
            elif ttype in ("BONUS", "SPLIT"):
                net_qty += qty
            elif ttype == "SELL":
                avg_cost = (net_invested / net_qty) if net_qty > 0 else 0
                cost_removed = avg_cost * qty
                realized_gain += (amt - chg) - cost_removed
                net_qty -= qty
                net_invested -= cost_removed
            elif ttype == "DIVIDEND":
                dividend_income += amt

        avg_cost_per_unit = (net_invested / net_qty) if net_qty > 0 else 0
        result = dict(zip(group_cols, keys))
        result.update({
            "net_quantity": round(net_qty, 4),
            "invested_amount": round(max(net_invested, 0), 2),
            "avg_cost_per_unit": round(avg_cost_per_unit, 4),
            "realized_gain": round(realized_gain, 2),
            "dividend_income": round(dividend_income, 2),
        })
        results.append(result)
    return pd.DataFrame(results)


def attach_current_value(holdings_df):
    """Attach cached prices and enforce numeric dtypes used by charts/calculations.

    Recent Plotly/Pandas versions are stricter about mixed dtypes.  SQLite NULLs or
    missing live prices can otherwise turn current_value into an object column and
    make wide-form Plotly charts fail.
    """
    if holdings_df.empty:
        return holdings_df

    merged = holdings_df.copy()

    # Ensure portfolio calculation columns are always numeric.
    for col in ["net_quantity", "invested_amount", "avg_cost_per_unit",
                "realized_gain", "dividend_income"]:
        if col in merged.columns:
            merged[col] = pd.to_numeric(merged[col], errors="coerce").fillna(0.0).astype(float)

    # Pandas merge keys must have the same dtype.  Identifiers can arrive as
    # float64 from Excel/NaN inference (e.g. AMFI 125497.0) while SQLite returns
    # TEXT. Normalize both sides to the same canonical string representation.
    if "identifier" not in merged.columns:
        merged["identifier"] = ""
    merged["identifier"] = merged["identifier"].apply(_clean_identifier).astype("string").fillna("")

    conn = get_conn()
    prices = pd.read_sql_query("SELECT identifier, last_price FROM price_cache", conn)
    conn.close()

    if prices.empty:
        prices = pd.DataFrame({
            "identifier": pd.Series(dtype="string"),
            "last_price": pd.Series(dtype="float64"),
        })
    else:
        prices["identifier"] = prices["identifier"].apply(_clean_identifier).astype("string").fillna("")
        prices["last_price"] = pd.to_numeric(prices["last_price"], errors="coerce")
        # A corrupted/legacy cache should not create duplicate holdings rows.
        prices = prices.drop_duplicates(subset=["identifier"], keep="last")

    merged = merged.merge(prices, on="identifier", how="left", validate="many_to_one")
    merged["current_price"] = pd.to_numeric(merged.get("last_price"), errors="coerce")
    merged["current_value"] = (
        merged["current_price"] * merged["net_quantity"]
    ).astype(float)
    merged["unrealized_gain"] = (
        merged["current_value"] - merged["invested_amount"]
    ).astype(float)

    invested = pd.to_numeric(merged["invested_amount"], errors="coerce").fillna(0.0)
    gain = pd.to_numeric(merged["unrealized_gain"], errors="coerce")
    merged["unrealized_gain_pct"] = 0.0
    mask = invested.ne(0) & gain.notna()
    merged.loc[mask, "unrealized_gain_pct"] = (gain[mask] / invested[mask]) * 100.0

    return merged


def make_invested_current_bar(df, x_col, title):
    """Create a robust grouped bar chart using Plotly long-form data.

    Long-form avoids Plotly Express' wide-form mixed-dtype ValueError and works
    consistently across recent Plotly/Pandas releases.
    """
    if df is None or df.empty:
        return None

    chart_df = df[[x_col, "invested_amount", "current_value"]].copy()
    chart_df[x_col] = chart_df[x_col].fillna("Unknown").astype(str)
    for col in ["invested_amount", "current_value"]:
        chart_df[col] = pd.to_numeric(chart_df[col], errors="coerce").fillna(0.0).astype(float)

    chart_df = chart_df.melt(
        id_vars=[x_col],
        value_vars=["invested_amount", "current_value"],
        var_name="Value Type",
        value_name="Amount",
    )
    chart_df["Value Type"] = chart_df["Value Type"].map({
        "invested_amount": "Invested Amount",
        "current_value": "Current Value",
    })
    chart_df["Amount"] = pd.to_numeric(chart_df["Amount"], errors="coerce").fillna(0.0).astype(float)

    return px.bar(
        chart_df,
        x=x_col,
        y="Amount",
        color="Value Type",
        barmode="group",
        title=title,
        labels={x_col: x_col.replace("_", " ").title(), "Amount": "Amount (₹)"},
    )


def _clean_identifier(identifier):
    """Return a stable text key for equity symbols and AMFI scheme codes."""
    if identifier is None:
        return ""
    try:
        if pd.isna(identifier):
            return ""
    except Exception:
        pass

    value = str(identifier).strip().upper()
    if value in {"", "NAN", "NONE", "NULL", "<NA>"}:
        return ""

    # Excel commonly converts numeric AMFI codes to floats such as 125497.0.
    if value.endswith(".0") and value[:-2].isdigit():
        value = value[:-2]
    return value


def _equity_symbol_candidates(identifier):
    """Build Yahoo Finance candidates for an Indian equity identifier.

    Examples:
      RELIANCE      -> RELIANCE.NS, RELIANCE.BO, RELIANCE
      RELIANCE.NS   -> RELIANCE.NS only first, then sensible fallbacks
      500325.BO     -> 500325.BO first
    """
    symbol = _clean_identifier(identifier).replace(" ", "")
    if not symbol:
        return []

    candidates = []
    def add(x):
        if x and x not in candidates:
            candidates.append(x)

    if symbol.endswith((".NS", ".BO")):
        add(symbol)
        base = symbol.rsplit(".", 1)[0]
        add(base + (".BO" if symbol.endswith(".NS") else ".NS"))
        add(base)
    else:
        # NSE is the normal default for Indian listed shares.
        add(symbol + ".NS")
        add(symbol + ".BO")
        add(symbol)
    return candidates


def _valid_market_price(value):
    try:
        value = float(value)
        return value if math.isfinite(value) and value > 0 else None
    except (TypeError, ValueError):
        return None


def _fetch_yahoo_chart_price(ticker_symbol):
    """Fallback to Yahoo's chart endpoint if yfinance is rate-limited/fails."""
    try:
        url = f"https://query1.finance.yahoo.com/v8/finance/chart/{ticker_symbol}"
        response = requests.get(
            url,
            params={"range": "5d", "interval": "1d"},
            headers={"User-Agent": "Mozilla/5.0"},
            timeout=12,
        )
        response.raise_for_status()
        result = response.json().get("chart", {}).get("result") or []
        if not result:
            return None
        meta = result[0].get("meta", {})
        for key in ("regularMarketPrice", "previousClose", "chartPreviousClose"):
            price = _valid_market_price(meta.get(key))
            if price is not None:
                return price
        closes = (((result[0].get("indicators") or {}).get("quote") or [{}])[0].get("close") or [])
        for value in reversed(closes):
            price = _valid_market_price(value)
            if price is not None:
                return price
    except Exception:
        return None
    return None




def get_app_setting(key, default=""):
    try:
        conn = get_conn()
        row = conn.execute("SELECT setting_value FROM app_settings WHERE setting_key=?", (key,)).fetchone()
        conn.close()
        return (row[0] if row else default) or default
    except Exception:
        return default


def set_app_setting(key, value):
    conn = get_conn()
    conn.execute("""INSERT INTO app_settings(setting_key, setting_value, updated_at)
                    VALUES (?, ?, ?)
                    ON CONFLICT(setting_key) DO UPDATE SET setting_value=excluded.setting_value,
                    updated_at=excluded.updated_at""", (key, str(value), str(datetime.now())))
    conn.commit(); conn.close()


def get_instrument_mapping(identifier):
    ident = _clean_identifier(identifier)
    if not ident:
        return {}
    conn = get_conn()
    row = conn.execute("SELECT * FROM instrument_master WHERE identifier=?", (ident,)).fetchone()
    conn.close()
    return dict(row) if row else {}


def upsert_instrument_mapping(identifier, instrument_name="", asset_type="Equity", exchange="NSE",
                              trading_symbol="", isin="", hdfc_segment="", hdfc_token="",
                              icici_stock_code="", upstox_key="", yahoo_symbol=""):
    ident = _clean_identifier(identifier)
    if not ident:
        return False
    conn = get_conn()
    conn.execute("""INSERT INTO instrument_master
        (identifier,instrument_name,asset_type,exchange,trading_symbol,isin,hdfc_segment,hdfc_token,
         icici_stock_code,upstox_key,yahoo_symbol,updated_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
        ON CONFLICT(identifier) DO UPDATE SET
        instrument_name=excluded.instrument_name, asset_type=excluded.asset_type, exchange=excluded.exchange,
        trading_symbol=excluded.trading_symbol, isin=excluded.isin, hdfc_segment=excluded.hdfc_segment,
        hdfc_token=excluded.hdfc_token, icici_stock_code=excluded.icici_stock_code,
        upstox_key=excluded.upstox_key, yahoo_symbol=excluded.yahoo_symbol, updated_at=excluded.updated_at""",
        (ident,instrument_name,asset_type,exchange,trading_symbol,isin,hdfc_segment,hdfc_token,
         icici_stock_code,upstox_key,yahoo_symbol,str(datetime.now())))
    conn.commit(); conn.close(); return True


def _provider_order():
    preferred = (get_app_setting("preferred_price_provider", "AUTO") or "AUTO").upper()
    standard = ["HDFC", "ICICI", "ZERODHA", "UPSTOX", "YAHOO"]
    if preferred == "AUTO" or preferred not in standard:
        return standard
    return [preferred] + [p for p in standard if p != preferred]


def fetch_hdfc_ltp(identifier):
    """HDFC InvestRight LTP adapter.

    HDFC's official developer portal confirms real-time market insights and LTP market-data support.
    Because the exact authenticated endpoint/payload can evolve, V7 keeps endpoint and header names configurable.
    Set HDFC_LTP_ENDPOINT plus credentials in Settings/environment, and map segment/token in Instrument Master.
    """
    api_key = _runtime_secret("HDFC_API_KEY")
    access_token = _runtime_secret("HDFC_ACCESS_TOKEN")
    endpoint = _runtime_secret("HDFC_LTP_ENDPOINT")
    m = get_instrument_mapping(identifier)
    segment = str(m.get("hdfc_segment") or m.get("exchange") or "NSE").strip()
    token = str(m.get("hdfc_token") or "").strip()
    if not api_key or not access_token:
        return None, None, "HDFC API key/access token not configured."
    if not endpoint:
        return None, None, "HDFC LTP endpoint not configured; copy the current Fetch LTP endpoint from your InvestRight API documentation."
    if not token:
        return None, None, "HDFC instrument token missing in Instrument Master."
    try:
        headers={"Authorization": f"Bearer {access_token}", "X-API-Key": api_key,
                 "Content-Type":"application/json", "Accept":"application/json",
                 "User-Agent":"Family-Fortune-Tracker/7.0"}
        # Generic documented-style adapter; Settings lets the endpoint remain current without source edits.
        r=requests.post(endpoint, json={"segment":segment,"token":token}, headers=headers, timeout=12)
        if r.status_code in (401,403): return None, f"{segment}:{token}", "HDFC authentication failed/expired."
        r.raise_for_status(); payload=r.json()
        def walk(obj):
            if isinstance(obj, dict):
                for k,v in obj.items():
                    if str(k).lower() in {"ltp","last_price","lastprice","lasttradedprice"}:
                        p=_valid_market_price(v)
                        if p is not None: return p
                for v in obj.values():
                    p=walk(v)
                    if p is not None: return p
            elif isinstance(obj,list):
                for v in obj:
                    p=walk(v)
                    if p is not None: return p
            return None
        price=walk(payload)
        return (price, f"{segment}:{token}", None) if price is not None else (None,f"{segment}:{token}","HDFC returned no recognizable LTP field.")
    except Exception as exc:
        return None, f"{segment}:{token}", str(exc)


def fetch_icici_breeze_ltp(identifier):
    """ICICI Direct Breeze official quote adapter using breeze-connect SDK when installed."""
    api_key=_runtime_secret("ICICI_BREEZE_API_KEY")
    api_secret=_runtime_secret("ICICI_BREEZE_API_SECRET")
    session_token=_runtime_secret("ICICI_BREEZE_SESSION_TOKEN")
    m=get_instrument_mapping(identifier)
    stock_code=str(m.get("icici_stock_code") or m.get("trading_symbol") or _clean_identifier(identifier)).strip().upper()
    exchange=str(m.get("exchange") or "NSE").strip().upper()
    if not api_key or not api_secret or not session_token:
        return None, stock_code, "ICICI Breeze API key/secret/session token not configured."
    try:
        from breeze_connect import BreezeConnect
    except Exception:
        return None, stock_code, "breeze-connect package is not installed. Run: pip install breeze-connect"
    try:
        breeze=BreezeConnect(api_key=api_key)
        breeze.generate_session(api_secret=api_secret, session_token=session_token)
        resp=breeze.get_quotes(stock_code=stock_code, exchange_code=exchange, product_type="cash",
                               expiry_date="", right="others", strike_price="0")
        rows=(resp or {}).get("Success") or []
        if isinstance(rows,dict): rows=[rows]
        for row in rows:
            for key in ("ltp","last","last_price","close"):
                p=_valid_market_price((row or {}).get(key))
                if p is not None: return p, f"{exchange}:{stock_code}", None
        return None, f"{exchange}:{stock_code}", str((resp or {}).get("Error") or "ICICI Breeze returned no LTP.")
    except Exception as exc:
        return None, f"{exchange}:{stock_code}", str(exc)


def _runtime_secret(name):
    """Read a market-data credential from Streamlit session state or environment.

    Credentials entered in Settings are session-only and are not written to the
    local database. Environment variables can be used for persistent local setup.
    """
    try:
        value = st.session_state.get(name)
        if value:
            return str(value).strip()
    except Exception:
        pass
    return str(os.environ.get(name, '') or '').strip()


def fetch_zerodha_ltp(identifier):
    """Fetch LTP from Zerodha Kite Connect's official quote API.

    Requires ZERODHA_API_KEY and ZERODHA_ACCESS_TOKEN. The identifier may be
    RELIANCE, RELIANCE.NS, INFY, NSE:INFY, or BSE:500325.
    Returns (price, resolved_symbol, error).
    """
    api_key = _runtime_secret('ZERODHA_API_KEY')
    access_token = _runtime_secret('ZERODHA_ACCESS_TOKEN')
    if not api_key or not access_token:
        return None, None, 'Zerodha API credentials are not configured.'

    ident = _clean_identifier(identifier)
    if not ident:
        return None, None, 'Identifier is blank.'

    if ':' in ident:
        exchange, symbol = ident.split(':', 1)
        exchange = exchange.upper().strip()
        symbol = symbol.strip().upper()
    elif ident.upper().endswith('.BO'):
        exchange, symbol = 'BSE', ident[:-3].upper()
    elif ident.upper().endswith('.NS'):
        exchange, symbol = 'NSE', ident[:-3].upper()
    else:
        exchange, symbol = 'NSE', ident.upper()

    instrument = f'{exchange}:{symbol}'
    try:
        r = requests.get(
            'https://api.kite.trade/quote/ltp',
            params=[('i', instrument)],
            headers={
                'X-Kite-Version': '3',
                'Authorization': f'token {api_key}:{access_token}',
                'User-Agent': 'Family-Fortune-Tracker/1.0',
            },
            timeout=12,
        )
        if r.status_code == 403:
            return None, instrument, 'Zerodha authentication failed or access token expired.'
        r.raise_for_status()
        payload = r.json()
        row = (payload.get('data') or {}).get(instrument)
        if not row:
            return None, instrument, f'Zerodha returned no quote for {instrument}.'
        price = _valid_market_price(row.get('last_price'))
        if price is None:
            return None, instrument, f'Zerodha returned an invalid LTP for {instrument}.'
        return price, instrument, None
    except Exception as exc:
        return None, instrument, str(exc) or 'Unable to contact Zerodha Kite Connect.'


def fetch_upstox_ltp(identifier):
    """Fetch LTP from Upstox's official Market Quote API.

    Requires UPSTOX_ACCESS_TOKEN. Upstox's quote endpoint needs an instrument
    key such as NSE_EQ|INE848E01016. Therefore this provider is used when the
    saved identifier is already an Upstox instrument key or an NSE/BSE ISIN.
    Returns (price, resolved_key, error).
    """
    token = _runtime_secret('UPSTOX_ACCESS_TOKEN')
    if not token:
        return None, None, 'Upstox access token is not configured.'

    ident = _clean_identifier(identifier)
    if not ident:
        return None, None, 'Identifier is blank.'

    upper = ident.upper()
    if upper.startswith(('NSE_EQ|', 'BSE_EQ|')):
        key = ident
    elif upper.startswith('INE') and len(upper) == 12:
        key = f'NSE_EQ|{upper}'
    else:
        return None, None, 'Upstox requires an instrument key/ISIN; a trading symbol alone is not enough.'

    try:
        r = requests.get(
            'https://api.upstox.com/v3/market-quote/ltp',
            params={'instrument_key': key},
            headers={
                'Accept': 'application/json',
                'Authorization': f'Bearer {token}',
                'User-Agent': 'Family-Fortune-Tracker/1.0',
            },
            timeout=12,
        )
        if r.status_code in (401, 403):
            return None, key, 'Upstox authentication failed or access token expired.'
        r.raise_for_status()
        data = r.json().get('data') or {}
        # Response map keys can be display keys rather than the exact instrument key.
        for row in data.values():
            price = _valid_market_price((row or {}).get('last_price'))
            if price is not None:
                return price, key, None
        return None, key, f'Upstox returned no LTP for {key}.'
    except Exception as exc:
        return None, key, str(exc) or 'Unable to contact Upstox.'

def fetch_equity_price(identifier):
    """V7 multi-provider equity price engine with user-selectable priority."""
    ident=_clean_identifier(identifier)
    if not ident:
        return None,None,None,"Identifier is blank."
    m=get_instrument_mapping(ident)
    provider_errors=[]
    for provider in _provider_order():
        if provider == "HDFC" and (_runtime_secret("HDFC_API_KEY") or _runtime_secret("HDFC_ACCESS_TOKEN")):
            p,sym,err=fetch_hdfc_ltp(ident)
            if p is not None: return p,sym,"HDFC Securities InvestRight (official API)",None
            if err: provider_errors.append("HDFC: "+err)
        elif provider == "ICICI" and _runtime_secret("ICICI_BREEZE_API_KEY"):
            p,sym,err=fetch_icici_breeze_ltp(ident)
            if p is not None: return p,sym,"ICICI Direct Breeze (official API)",None
            if err: provider_errors.append("ICICI: "+err)
        elif provider == "ZERODHA" and _runtime_secret('ZERODHA_API_KEY') and _runtime_secret('ZERODHA_ACCESS_TOKEN'):
            p,sym,err=fetch_zerodha_ltp(m.get("trading_symbol") or ident)
            if p is not None: return p,sym,"Zerodha Kite Connect (official API)",None
            if err: provider_errors.append("Zerodha: "+err)
        elif provider == "UPSTOX" and _runtime_secret('UPSTOX_ACCESS_TOKEN'):
            p,sym,err=fetch_upstox_ltp(m.get("upstox_key") or m.get("isin") or ident)
            if p is not None: return p,sym,"Upstox Market Quote (official API)",None
            if err: provider_errors.append("Upstox: "+err)
        elif provider == "YAHOO":
            candidates=[]
            if m.get("yahoo_symbol"): candidates=[m["yahoo_symbol"]]
            candidates += [x for x in _equity_symbol_candidates(m.get("trading_symbol") or ident) if x not in candidates]
            last_error="Price not available from Yahoo Finance."
            for ticker_symbol in candidates:
                try:
                    ticker=yf.Ticker(ticker_symbol)
                    try:
                        fi=ticker.fast_info
                        value=fi.get("last_price") if hasattr(fi,"get") else getattr(fi,"last_price",None)
                        p=_valid_market_price(value)
                        if p is not None: return p,ticker_symbol,"Yahoo Finance (yfinance)",None
                    except Exception as exc: last_error=str(exc) or last_error
                    try:
                        hist=ticker.history(period="5d",interval="1d",auto_adjust=False)
                        if hist is not None and not hist.empty and "Close" in hist.columns:
                            vals=pd.to_numeric(hist["Close"],errors="coerce").dropna()
                            if not vals.empty:
                                p=_valid_market_price(vals.iloc[-1])
                                if p is not None: return p,ticker_symbol,"Yahoo Finance (history)",None
                    except Exception as exc: last_error=str(exc) or last_error
                except Exception as exc: last_error=str(exc) or last_error
                p=_fetch_yahoo_chart_price(ticker_symbol)
                if p is not None: return p,ticker_symbol,"Yahoo Finance (HTTP fallback)",None
            provider_errors.append("Yahoo: "+last_error)
    return None, ident, None, " | ".join(provider_errors) or "No configured price provider returned a quote."

def fetch_mf_nav(scheme_code):
    code = _clean_identifier(scheme_code)
    if not code:
        return None, "AMFI scheme code is blank."
    try:
        r = requests.get(
            f"https://api.mfapi.in/mf/{code}/latest",
            headers={"User-Agent": "Mozilla/5.0"},
            timeout=12,
        )
        r.raise_for_status()
        data = r.json()
        rows = data.get("data") or []
        if rows:
            price = _valid_market_price(rows[0].get("nav"))
            if price is not None:
                return price, None
        return None, "NAV was not returned for this AMFI scheme code."
    except Exception as exc:
        return None, str(exc) or "Unable to contact MFAPI."


def update_price(identifier, asset_type, name):
    identifier = _clean_identifier(identifier)
    result = {
        "identifier": identifier,
        "name": name,
        "asset_type": asset_type,
        "price": None,
        "resolved_symbol": identifier,
        "source": None,
        "error": None,
    }

    if not identifier:
        result["error"] = "Missing NSE/BSE symbol or AMFI scheme code."
        return result

    if asset_type == "Equity":
        price, resolved_symbol, source, error = fetch_equity_price(identifier)
        result.update(price=price, resolved_symbol=resolved_symbol or identifier, source=source, error=error)
    elif asset_type == "Mutual Fund":
        price, error = fetch_mf_nav(identifier)
        result.update(price=price, source="MFAPI / AMFI data", error=error)
    else:
        result["error"] = f"Unsupported asset type: {asset_type}"
        return result

    price = result["price"]
    if price is not None:
        conn = get_conn()
        conn.execute("""INSERT INTO price_cache (identifier, asset_type, name, last_price, price_date, updated_at)
                        VALUES (?, ?, ?, ?, ?, ?)
                        ON CONFLICT(identifier) DO UPDATE SET last_price=excluded.last_price,
                        asset_type=excluded.asset_type, name=excluded.name,
                        price_date=excluded.price_date, updated_at=excluded.updated_at""",
                     (identifier, asset_type, name, float(price), str(date.today()), str(datetime.now())))
        conn.commit()
        conn.close()
    return result


def refresh_all_prices(holdings_df):
    """Refresh unique portfolio identifiers and return success/failure details."""
    results = []
    seen = set()
    if holdings_df is None or holdings_df.empty:
        return results

    for _, row in holdings_df.iterrows():
        identifier = _clean_identifier(row.get("identifier"))
        asset_type = str(row.get("asset_type") or "").strip()
        key = (identifier, asset_type)
        if not identifier or key in seen:
            continue
        seen.add(key)
        results.append(update_price(identifier, asset_type, row.get("instrument_name") or identifier))
    return results


def save_manual_current_price(identifier, asset_type, name, price):
    """Manual fallback for a provider outage or an unrecognized ticker."""
    identifier = _clean_identifier(identifier)
    price = _valid_market_price(price)
    if not identifier or price is None:
        return False
    conn = get_conn()
    conn.execute("""INSERT INTO price_cache (identifier, asset_type, name, last_price, price_date, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?)
                    ON CONFLICT(identifier) DO UPDATE SET last_price=excluded.last_price,
                    asset_type=excluded.asset_type, name=excluded.name,
                    price_date=excluded.price_date, updated_at=excluded.updated_at""",
                 (identifier, asset_type, name, float(price), str(date.today()), str(datetime.now())))
    conn.commit()
    conn.close()
    return True


def fetch_news_for_instrument(instrument_name, max_items=5):
    if feedparser is None:
        return []
    query = quote(f"{instrument_name} share stock")
    url = f"https://news.google.com/rss/search?q={query}&hl=en-IN&gl=IN&ceid=IN:en"
    items = []
    try:
        feed = feedparser.parse(url)
        for entry in feed.entries[:max_items]:
            src = entry.get("source", {})
            src_name = src.get("title", "News") if isinstance(src, dict) else "News"
            items.append({
                "instrument_name": instrument_name,
                "title": entry.get("title"),
                "link": entry.get("link"),
                "source": src_name,
                "published": entry.get("published", ""),
                "summary": entry.get("summary", ""),
            })
    except Exception:
        pass
    return items


def refresh_news_for_portfolio(instrument_names):
    conn = get_conn()
    cur = conn.cursor()
    count = 0
    for name in instrument_names:
        for it in fetch_news_for_instrument(name):
            cur.execute("""INSERT OR IGNORE INTO news_cache
                (instrument_name, title, link, source, published, summary, fetched_at)
                VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (it["instrument_name"], it["title"], it["link"], it["source"],
                 it["published"], it["summary"], str(datetime.now())))
            count += cur.rowcount
    conn.commit()
    conn.close()
    return count


def get_cached_news(limit=200):
    conn = get_conn()
    rows = conn.execute("SELECT * FROM news_cache ORDER BY fetched_at DESC LIMIT ?", (limit,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def news_fetched_today():
    conn = get_conn()
    row = conn.execute("SELECT MAX(fetched_at) as m FROM news_cache").fetchone()
    conn.close()
    if row and row["m"]:
        return row["m"][:10] == str(date.today())
    return False


# ==============================================================================
# 5. EXCEL TEMPLATE GENERATOR / READER
# ==============================================================================

TEMPLATE_COLUMNS = [
    "Member Name*", "Broker/Platform*", "Asset Type*", "Instrument Name*",
    "Identifier (NSE Symbol / AMFI Scheme Code)", "Transaction Type*",
    "Transaction Date* (DD-MM-YYYY)", "Quantity/Units*", "Price per Unit*",
    "Total Amount*", "Charges/Brokerage", "Folio Number (MF only)", "Remarks",
]

RENAME_MAP = {
    "Member Name*": "member_name",
    "Broker/Platform*": "broker_name",
    "Asset Type*": "asset_type",
    "Instrument Name*": "instrument_name",
    "Identifier (NSE Symbol / AMFI Scheme Code)": "identifier",
    "Transaction Type*": "txn_type",
    "Transaction Date* (DD-MM-YYYY)": "txn_date",
    "Quantity/Units*": "quantity",
    "Price per Unit*": "price",
    "Total Amount*": "amount",
    "Charges/Brokerage": "charges",
    "Folio Number (MF only)": "folio_no",
    "Remarks": "remarks",
}


def generate_template(member_names, broker_names):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(TEMPLATE_COLUMNS)
    for i in range(1, len(TEMPLATE_COLUMNS) + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        ws.column_dimensions[cell.column_letter].width = 24

    lst = wb.create_sheet("Lists")
    lst.sheet_state = "hidden"
    member_names = member_names or ["(add a member first)"]
    broker_names = broker_names or ["Other"]
    for i, m in enumerate(member_names, start=1):
        lst.cell(row=i, column=1, value=m)
    for i, b in enumerate(broker_names, start=1):
        lst.cell(row=i, column=2, value=b)
    for i, a in enumerate(["Equity", "Mutual Fund"], start=1):
        lst.cell(row=i, column=3, value=a)
    for i, t in enumerate(["BUY", "SELL", "SIP", "DIVIDEND", "BONUS", "SPLIT"], start=1):
        lst.cell(row=i, column=4, value=t)

    max_rows = 500
    dv_member = DataValidation(type="list", formula1=f"=Lists!$A$1:$A${len(member_names)}", allow_blank=True)
    dv_broker = DataValidation(type="list", formula1=f"=Lists!$B$1:$B${len(broker_names)}", allow_blank=True)
    dv_asset = DataValidation(type="list", formula1="=Lists!$C$1:$C$2", allow_blank=True)
    dv_txn = DataValidation(type="list", formula1="=Lists!$D$1:$D$6", allow_blank=True)
    for dv in (dv_member, dv_broker, dv_asset, dv_txn):
        ws.add_data_validation(dv)
    dv_member.add(f"A2:A{max_rows}")
    dv_broker.add(f"B2:B{max_rows}")
    dv_asset.add(f"C2:C{max_rows}")
    dv_txn.add(f"F2:F{max_rows}")

    ins = wb.create_sheet("Instructions")
    instructions = [
        "FAMILY INVESTMENT MATRIX - UPLOAD TEMPLATE INSTRUCTIONS",
        "",
        "1. Fields marked with * are mandatory.",
        "2. Member Name must exactly match an existing family member already added in the app.",
        "3. Broker/Platform must match one of the listed platforms (or 'Other').",
        "4. Asset Type: 'Equity' for shares, 'Mutual Fund' for MF folios.",
        "5. Identifier: For Equity use the NSE trading symbol (e.g. RELIANCE, TCS, INFY).",
        "   For Mutual Fund use the AMFI Scheme Code (e.g. 125497). Required for live pricing.",
        "6. Transaction Type: BUY, SELL, SIP, DIVIDEND, BONUS or SPLIT.",
        "7. Transaction Date must be in DD-MM-YYYY format.",
        "8. Total Amount should equal Quantity x Price (+/- charges) - the app checks this.",
        "9. The app automatically screens for duplicate rows (same member+broker+share+date+amount)",
        "   and will ask you to confirm before importing any that look suspicious.",
        "10. Save the file and upload it from the 'Upload Investments' page, Excel Upload tab.",
    ]
    for i, line in enumerate(instructions, start=1):
        ins.cell(row=i, column=1, value=line)
    ins.column_dimensions["A"].width = 100

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def read_uploaded_excel(file):
    df = pd.read_excel(file, sheet_name="Transactions")
    df.columns = [str(c).strip() for c in df.columns]
    df = df.rename(columns=RENAME_MAP)
    df = df.dropna(how="all")
    return df.reset_index(drop=True)


# ==============================================================================
# 6. PDF EXTRACTOR (BETA)
# ==============================================================================

DATE_PATTERN = re.compile(r"(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})")
AMOUNT_PATTERN = re.compile(r"([\d,]+\.\d{2})")


def extract_transactions_from_pdf(file):
    rows = []
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    if row and any(row):
                        rows.append(row)
            if not tables:
                text = page.extract_text() or ""
                for line in text.split("\n"):
                    dates = DATE_PATTERN.findall(line)
                    amounts = AMOUNT_PATTERN.findall(line)
                    if dates and amounts:
                        rows.append([dates[0], line.strip(), None, amounts[-1]])

    if not rows:
        return pd.DataFrame()

    max_len = max(len(r) for r in rows)
    norm_rows = [list(r) + [None] * (max_len - len(r)) for r in rows]
    df = pd.DataFrame(norm_rows)
    df.columns = [f"col_{i}" for i in range(df.shape[1])]
    return df


# ==============================================================================
# 7. STREAMLIT UI
# ==============================================================================

def inject_custom_css():
    """Force a clear professional light UI regardless of Streamlit/system theme."""
    st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Poppins:wght@400;500;600;700&display=swap');

    :root {
        --fft-primary: #0B4F8A;
        --fft-primary-dark: #083B68;
        --fft-accent: #D98200;
        --fft-bg: #F4F8FC;
        --fft-card: #FFFFFF;
        --fft-text: #17212B;
        --fft-muted: #4B5D6B;
        --fft-border: #CFDCE8;
    }

    html, body, [class*="css"], .stApp {
        font-family: 'Poppins', sans-serif !important;
    }

    .stApp,
    [data-testid="stAppViewContainer"],
    [data-testid="stMain"] {
        background: var(--fft-bg) !important;
        color: var(--fft-text) !important;
    }

    [data-testid="stHeader"] {
        background: rgba(244,248,252,0.96) !important;
    }

    [data-testid="stSidebar"] {
        background: #EAF2F9 !important;
        border-right: 1px solid var(--fft-border) !important;
    }

    [data-testid="stSidebar"] * {
        color: var(--fft-text) !important;
    }

    h1, h2, h3, h4, h5, h6 {
        color: var(--fft-primary-dark) !important;
        font-weight: 700 !important;
    }

    p, label, span, div, li {
        text-shadow: none !important;
    }

    [data-testid="stMarkdownContainer"] p,
    [data-testid="stCaptionContainer"],
    .stCaption,
    label {
        color: var(--fft-text) !important;
    }

    div[data-testid="stMetric"] {
        background: var(--fft-card) !important;
        border-radius: 14px !important;
        padding: 16px !important;
        border: 1px solid var(--fft-border) !important;
        box-shadow: 0 2px 8px rgba(19,59,92,0.08) !important;
    }

    div[data-testid="stMetric"] * {
        color: var(--fft-text) !important;
    }

    div[data-testid="stMetricValue"] {
        color: var(--fft-primary-dark) !important;
        font-weight: 700 !important;
    }

    .stButton > button,
    .stDownloadButton > button,
    button[kind="primary"] {
        background: var(--fft-primary) !important;
        color: #FFFFFF !important;
        border: 1px solid var(--fft-primary-dark) !important;
        border-radius: 8px !important;
        font-weight: 600 !important;
    }

    .stButton > button:hover,
    .stDownloadButton > button:hover {
        background: var(--fft-primary-dark) !important;
        color: #FFFFFF !important;
    }

    input, textarea,
    [data-baseweb="select"] > div,
    [data-testid="stNumberInput"] input,
    [data-testid="stTextInput"] input,
    [data-testid="stDateInput"] input {
        background: #FFFFFF !important;
        color: #111827 !important;
        border-color: #AFC2D4 !important;
    }

    [data-baseweb="popover"], [role="listbox"], [role="option"] {
        background: #FFFFFF !important;
        color: #111827 !important;
    }

    [data-testid="stDataFrame"],
    [data-testid="stTable"] {
        background: #FFFFFF !important;
        border: 1px solid var(--fft-border) !important;
        border-radius: 8px !important;
    }

    [data-testid="stAlert"] {
        color: #17212B !important;
    }

    hr {
        border-color: #CBD8E5 !important;
    }

    /* Keep Plotly containers visually separate from the page. */
    [data-testid="stPlotlyChart"] {
        background: #FFFFFF !important;
        border: 1px solid var(--fft-border) !important;
        border-radius: 12px !important;
        padding: 6px !important;
        box-shadow: 0 2px 8px rgba(19,59,92,0.06) !important;
    }
    </style>
    """, unsafe_allow_html=True)


def style_plotly_figure(fig):
    """Apply a light, high-contrast chart style independent of Streamlit theme."""
    if fig is None:
        return None
    try:
        fig.update_layout(
            template="plotly_white",
            paper_bgcolor="#FFFFFF",
            plot_bgcolor="#FFFFFF",
            font=dict(color="#17212B", size=13, family="Arial, sans-serif"),
            title_font=dict(color="#083B68", size=18),
            legend=dict(font=dict(color="#17212B"), bgcolor="rgba(255,255,255,0.85)"),
            margin=dict(l=28, r=28, t=60, b=40),
            colorway=["#0B4F8A", "#E07A00", "#2A8F5B", "#7A4FA3", "#C23B3B", "#1F7A8C"],
        )
        fig.update_xaxes(
            showgrid=True, gridcolor="#E5EDF5", zerolinecolor="#C9D6E2",
            tickfont=dict(color="#17212B"), title_font=dict(color="#17212B")
        )
        fig.update_yaxes(
            showgrid=True, gridcolor="#E5EDF5", zerolinecolor="#C9D6E2",
            tickfont=dict(color="#17212B"), title_font=dict(color="#17212B")
        )
    except Exception:
        pass
    return fig


def show_plotly(fig, use_container_width=True):
    if fig is not None:
        st.plotly_chart(style_plotly_figure(fig), use_container_width=use_container_width, theme=None)


@st.cache_data(ttl=300)
def _load_txn_cached():
    return load_all_transactions()


# ---------------------------- LOGIN / BOOTSTRAP -----------------------------

def login_page():
    st.markdown(f"<h1 style='text-align:center'>💰 {APP_NAME}</h1>", unsafe_allow_html=True)
    st.markdown("<p style='text-align:center'>One family. One fortune. Full transparency.</p>",
                unsafe_allow_html=True)

    if not admin_exists():
        st.subheader("🔐 First-Time Setup — Create Admin Account")
        st.caption("The admin logs in with a 10-digit mobile number. Default password is 123456 "
                   "— you will be asked to change it immediately after your first login.")
        with st.form("bootstrap"):
            name = st.text_input("Your Name")
            mobile = st.text_input("Your 10-digit Mobile Number (this becomes your login ID)")
            submitted = st.form_submit_button("Create Admin Account")
            if submitted:
                if not name.strip():
                    st.error("Name is required.")
                elif not (mobile.isdigit() and len(mobile) == 10):
                    st.error("Enter a valid 10-digit mobile number.")
                else:
                    try:
                        create_admin(name.strip(), mobile.strip())
                    except sqlite3.IntegrityError:
                        st.error("This mobile number is already present in the database.")
                    else:
                        st.success("Admin account created. Login with password 123456.")
                        st.rerun()
        return

    st.subheader("🔐 Family Login")
    with st.form("login"):
        mobile = st.text_input("Mobile Number (User ID)")
        password = st.text_input("Password", type="password")
        submitted = st.form_submit_button("Login")
        if submitted:
            user = authenticate(mobile.strip(), password)
            if user:
                st.session_state.logged_in = True
                st.session_state.user = user
                st.session_state.must_change_password = bool(user["is_first_login"])
                st.rerun()
            else:
                st.error("Invalid mobile number or password.")


def force_change_password_page():
    st.title("🔑 Set a New Password")
    st.info("This is your first login. Please set a new password to continue.")
    with st.form("force_change"):
        p1 = st.text_input("New Password", type="password")
        p2 = st.text_input("Confirm New Password", type="password")
        sub = st.form_submit_button("Set Password")
        if sub:
            if len(p1) < 4:
                st.error("Password should be at least 4 characters.")
            elif p1 != p2:
                st.error("Passwords do not match.")
            else:
                change_password(st.session_state.user["user_id"], p1)
                st.session_state.must_change_password = False
                st.session_state.user["is_first_login"] = 0
                st.success("Password updated!")
                st.rerun()


# ------------------------------- DASHBOARD -----------------------------------

def dashboard_page():
    st.title("🏠 Family Fortune — Dashboard")
    txn_df = _load_txn_cached()
    if txn_df.empty:
        st.info("No transactions uploaded yet. Go to '⬆️ Upload Investments' to add your first entry.")
        return

    holdings = attach_current_value(compute_holdings(txn_df))

    col_a, col_b = st.columns([3, 1])
    with col_b:
        if st.button("🔄 Refresh Live Prices"):
            with st.spinner("Fetching latest prices from Yahoo Finance / MFAPI..."):
                results = refresh_all_prices(holdings)
            success = [r for r in results if r.get("price") is not None]
            failed = [r for r in results if r.get("price") is None]
            st.session_state["last_price_refresh"] = {"success": success, "failed": failed}
            st.cache_data.clear()
            if success:
                st.success(f"Updated {len(success)} instrument(s).")
            if failed:
                st.warning(f"Could not update {len(failed)} instrument(s). See details below.")

    refresh_info = st.session_state.get("last_price_refresh")
    if refresh_info:
        failed = refresh_info.get("failed", [])
        success = refresh_info.get("success", [])
        if success:
            with st.expander("✅ Last live-price refresh — successful instruments"):
                ok_df = pd.DataFrame(success)
                show_cols = [c for c in ["name", "identifier", "resolved_symbol", "price", "source"] if c in ok_df.columns]
                st.dataframe(ok_df[show_cols], use_container_width=True, hide_index=True)
        if failed:
            with st.expander("⚠️ Live price could not be fetched — check identifiers"):
                bad_df = pd.DataFrame(failed)
                show_cols = [c for c in ["name", "identifier", "resolved_symbol", "error"] if c in bad_df.columns]
                st.dataframe(bad_df[show_cols], use_container_width=True, hide_index=True)
                st.caption("For NSE shares enter the Yahoo/NSE symbol such as RELIANCE, TCS, INFY. "
                           "For BSE-only shares you may enter the Yahoo symbol such as 500325.BO. "
                           "For mutual funds enter the AMFI scheme code.")

    total_invested = holdings["invested_amount"].sum()
    total_current = holdings["current_value"].fillna(0).sum()
    total_gain = total_current - total_invested
    gain_pct = (total_gain / total_invested * 100) if total_invested else 0

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("💰 Total Invested", f"₹{total_invested:,.0f}")
    k2.metric("📈 Current Value", f"₹{total_current:,.0f}")
    k3.metric("💹 Gain / Loss", f"₹{total_gain:,.0f}", f"{gain_pct:.2f}%")
    k4.metric("👪 Family Members", f"{holdings['member_name'].nunique()}")

    st.divider()
    c1, c2 = st.columns(2)
    with c1:
        fig = px.pie(holdings, names="member_name", values="current_value", hole=0.45,
                     title="Current Value Share — By Member",
                     color_discrete_sequence=px.colors.sequential.Sunset)
        show_plotly(fig)
    with c2:
        fig2 = px.pie(holdings, names="broker_name", values="current_value", hole=0.45,
                      title="Current Value Share — By Platform",
                      color_discrete_sequence=px.colors.sequential.Tealgrn)
        show_plotly(fig2)

    c3, c4 = st.columns(2)
    with c3:
        by_member = holdings.groupby("member_name")[["invested_amount", "current_value"]].sum().reset_index()
        fig3 = make_invested_current_bar(
            by_member, "member_name", "Invested vs Current Value — By Member"
        )
        if fig3 is not None:
            show_plotly(fig3)
    with c4:
        fig4 = px.pie(holdings, names="asset_type", values="current_value", hole=0.45,
                      title="Asset Allocation", color_discrete_sequence=["#D4AF37", "#1F77B4"])
        show_plotly(fig4)

    st.subheader("🌳 Portfolio Map — Member → Platform → Share")
    tree_df = holdings.copy()
    tree_df["current_value"] = tree_df["current_value"].fillna(0)
    fig5 = px.treemap(tree_df, path=["member_name", "broker_name", "instrument_name"],
                       values="current_value", color="unrealized_gain_pct",
                       color_continuous_scale="RdYlGn", title="Click any block to drill down")
    show_plotly(fig5)

    st.subheader("🏆 Top Gainers & Losers (Family-wide)")
    g1, g2 = st.columns(2)
    cols_show = ["member_name", "instrument_name", "invested_amount", "current_value", "unrealized_gain"]
    g1.dataframe(holdings.sort_values("unrealized_gain", ascending=False).head(5)[cols_show])
    g2.dataframe(holdings.sort_values("unrealized_gain").head(5)[cols_show])


# --------------------------- MEMBER VIEW / DRILL-DOWN -------------------------

def member_view_page():
    st.title("👪 Member-wise View")
    txn_df = _load_txn_cached()
    if txn_df.empty:
        st.info("No data yet.")
        return
    holdings = attach_current_value(compute_holdings(txn_df))

    members = ["🌐 Whole Family"] + sorted(holdings["member_name"].unique().tolist())
    choice = st.selectbox("Select View", members)

    if choice == "🌐 Whole Family":
        sub, sub_txn = holdings, txn_df
    else:
        sub = holdings[holdings.member_name == choice]
        sub_txn = txn_df[txn_df.member_name == choice]

    invested = sub["invested_amount"].sum()
    current = sub["current_value"].fillna(0).sum()
    gain = current - invested
    k1, k2, k3 = st.columns(3)
    k1.metric("Invested", f"₹{invested:,.0f}")
    k2.metric("Current", f"₹{current:,.0f}")
    k3.metric("Gain/Loss", f"₹{gain:,.0f}", f"{(gain / invested * 100 if invested else 0):.2f}%")

    if not sub.empty:
        fig = make_invested_current_bar(
            sub, "instrument_name", f"{choice} — Holdings"
        )
        if fig is not None:
            show_plotly(fig)

    st.subheader("📂 Drill Down: Platform → Share → Transactions")
    for broker in sorted(sub["broker_name"].dropna().unique()):
        with st.expander(f"📁 {broker}"):
            bsub = sub[sub.broker_name == broker]
            st.dataframe(bsub[["instrument_name", "net_quantity", "invested_amount",
                                "current_value", "unrealized_gain", "unrealized_gain_pct"]],
                         use_container_width=True)
            instruments = bsub["instrument_name"].unique().tolist()
            if instruments:
                inst_choice = st.selectbox("View transactions for:", instruments,
                                            key=f"inst_{broker}_{choice}")
                t = sub_txn[(sub_txn.broker_name == broker) & (sub_txn.instrument_name == inst_choice)]
                st.dataframe(t[["txn_date", "txn_type", "quantity", "price", "amount",
                                "charges", "remarks"]].sort_values("txn_date"),
                             use_container_width=True)


# ------------------------------ ALL TRANSACTIONS -------------------------------

def all_transactions_page():
    st.title("🧾 All Transactions")
    df = _load_txn_cached()
    if df.empty:
        st.info("No transactions yet.")
        return
    c1, c2, c3 = st.columns(3)
    mem_filter = c1.multiselect("Member", sorted(df["member_name"].unique()))
    brok_filter = c2.multiselect("Platform", sorted(df["broker_name"].unique()))
    type_filter = c3.multiselect("Transaction Type", sorted(df["txn_type"].unique()))

    filtered = df.copy()
    if mem_filter:
        filtered = filtered[filtered.member_name.isin(mem_filter)]
    if brok_filter:
        filtered = filtered[filtered.broker_name.isin(brok_filter)]
    if type_filter:
        filtered = filtered[filtered.txn_type.isin(type_filter)]

    st.dataframe(filtered.sort_values("txn_date", ascending=False), use_container_width=True)
    st.download_button("⬇️ Export as CSV", filtered.to_csv(index=False),
                        "transactions_export.csv", "text/csv")


# ------------------------------- QUERY EXPLORER --------------------------------

STANDARD_QUERIES = [
    "Total portfolio value (Family)",
    "Invested vs Current Value — by Member",
    "Invested vs Current Value — by Platform",
    "Top 5 Gainers (Family)",
    "Top 5 Losers (Family)",
    "Asset Allocation — Equity vs Mutual Fund (by Member)",
    "Dividend Income Received — by Member",
    "Realized Gains from Sales — by Member",
    "All Holdings of a Selected Share (across family)",
    "All Holdings of a Selected Member",
    "All Holdings on a Selected Platform",
    "Transaction History — filter by Date Range",
]


def query_explorer_page():
    st.title("🔍 Query Explorer")
    txn_df = _load_txn_cached()
    if txn_df.empty:
        st.info("No data yet.")
        return
    holdings = attach_current_value(compute_holdings(txn_df))
    q = st.selectbox("Choose a standard query", STANDARD_QUERIES)

    if q == "Total portfolio value (Family)":
        st.metric("Total Current Value", f"₹{holdings['current_value'].fillna(0).sum():,.0f}")
    elif q == "Invested vs Current Value — by Member":
        res = holdings.groupby("member_name")[["invested_amount", "current_value"]].sum().reset_index()
        st.dataframe(res)
        fig = make_invested_current_bar(res, "member_name", "Invested vs Current Value — by Member")
        if fig is not None:
            show_plotly(fig)
    elif q == "Invested vs Current Value — by Platform":
        res = holdings.groupby("broker_name")[["invested_amount", "current_value"]].sum().reset_index()
        st.dataframe(res)
        fig = make_invested_current_bar(res, "broker_name", "Invested vs Current Value — by Platform")
        if fig is not None:
            show_plotly(fig)
    elif q == "Top 5 Gainers (Family)":
        st.dataframe(holdings.sort_values("unrealized_gain", ascending=False).head(5))
    elif q == "Top 5 Losers (Family)":
        st.dataframe(holdings.sort_values("unrealized_gain").head(5))
    elif q == "Asset Allocation — Equity vs Mutual Fund (by Member)":
        res = holdings.groupby(["member_name", "asset_type"])["current_value"].sum().reset_index()
        show_plotly(px.bar(res, x="member_name", y="current_value", color="asset_type", barmode="stack"))
    elif q == "Dividend Income Received — by Member":
        st.dataframe(holdings.groupby("member_name")["dividend_income"].sum().reset_index())
    elif q == "Realized Gains from Sales — by Member":
        st.dataframe(holdings.groupby("member_name")["realized_gain"].sum().reset_index())
    elif q == "All Holdings of a Selected Share (across family)":
        share = st.selectbox("Select Share/Fund", sorted(holdings["instrument_name"].unique()))
        st.dataframe(holdings[holdings.instrument_name == share])
    elif q == "All Holdings of a Selected Member":
        mem = st.selectbox("Select Member", sorted(holdings["member_name"].unique()))
        st.dataframe(holdings[holdings.member_name == mem])
    elif q == "All Holdings on a Selected Platform":
        plat = st.selectbox("Select Platform", sorted(holdings["broker_name"].unique()))
        st.dataframe(holdings[holdings.broker_name == plat])
    elif q == "Transaction History — filter by Date Range":
        d1 = st.date_input("From", value=date(2020, 1, 1))
        d2 = st.date_input("To", value=date.today())
        mask = (pd.to_datetime(txn_df["txn_date"]) >= pd.to_datetime(d1)) & \
               (pd.to_datetime(txn_df["txn_date"]) <= pd.to_datetime(d2))
        st.dataframe(txn_df[mask])


# --------------------------------- NEWS BOARD -----------------------------------

def news_page():
    st.title("📰 Family Portfolio — News Board")
    txn_df = _load_txn_cached()
    if txn_df.empty:
        st.info("Add investments first to see relevant news.")
        return
    instruments = sorted(txn_df["instrument_name"].unique().tolist())

    colA, colB = st.columns([3, 1])
    with colB:
        if st.button("🔄 Fetch Latest News Now"):
            with st.spinner("Searching financial news sites for your portfolio shares..."):
                n = refresh_news_for_portfolio(instruments)
            st.success(f"Fetched {n} new articles.")
    with colA:
        if not news_fetched_today():
            st.info("News hasn't been refreshed today. Click 'Fetch Latest News Now' →")

    filter_share = st.selectbox("Filter by Share/Fund", ["All"] + instruments)
    news_items = get_cached_news(200)
    if filter_share != "All":
        news_items = [n for n in news_items if n["instrument_name"] == filter_share]

    if not news_items:
        st.info("No news cached yet. Click the refresh button above.")
        return

    for item in news_items:
        with st.container(border=True):
            st.markdown(f"**[{item['title']}]({item['link']})**")
            st.caption(f"🏷️ {item['instrument_name']} | 🗞️ {item['source']} | 🕒 {item['published']}")
            with st.expander("Show snippet"):
                st.write(item["summary"])


# ------------------------------- ADD MEMBER --------------------------------------

def add_member_page(user):
    st.title("➕ Add Family Member")
    if not can_add_member(user["role"]):
        st.warning("Only the Admin can add new family members currently "
                   "(this restriction can be changed in a future version).")
        return
    with st.form("add_member_form"):
        name = st.text_input("Full Name*")
        relation = st.selectbox("Relation*", RELATIONS)
        mobile = st.text_input("10-digit Mobile Number* (used as login ID)")
        dob = st.date_input("Date of Birth", value=None)
        pan = st.text_input("PAN (optional)")
        email = st.text_input("Email (optional)")
        sub = st.form_submit_button("Add Member")
        if sub:
            if not name.strip():
                st.error("Name is required.")
            elif not (mobile.isdigit() and len(mobile) == 10):
                st.error("Enter a valid 10-digit mobile number.")
            else:
                ok, msg = add_member(name.strip(), mobile.strip(), relation,
                                      str(dob) if dob else None, pan.strip() or None, email.strip() or None)
                if ok:
                    st.success(f"{msg} Default password is '123456' — they must change it on first login.")
                else:
                    st.error(msg)


# -------------------------------- UPLOAD PAGE -------------------------------------

def manual_entry_tab(user, members_df, brokers_df):
    # Clear widget state only BEFORE the widgets are created. Streamlit does not
    # permit changing a widget's session-state value after that widget has been
    # instantiated in the same run.
    if st.session_state.pop("manual_reset_requested", False):
        for key in [
            "manual_instrument", "manual_identifier", "manual_quantity", "manual_price",
            "manual_charges", "manual_folio", "manual_remarks",
        ]:
            st.session_state.pop(key, None)

    st.subheader("Add a Single Transaction")
    st.caption("Total Amount is calculated automatically as Quantity × Price per Unit.")

    # These widgets intentionally sit OUTSIDE st.form so Streamlit reruns as
    # Quantity/Price changes and the transaction total updates immediately.
    c1, c2, c3 = st.columns(3)
    member_name = c1.selectbox("Family Member*", members_df["name"].tolist(), key="manual_member")
    broker_name = c2.selectbox("Broker/Platform*", brokers_df["broker_name"].tolist(), key="manual_broker")
    asset_type = c3.selectbox("Asset Type*", ASSET_TYPES, key="manual_asset")

    c4, c5 = st.columns(2)
    instrument_name = c4.text_input("Instrument Name* (Company / Fund Name)", key="manual_instrument")
    identifier = c5.text_input(
        "Identifier (NSE Symbol / Yahoo Symbol / AMFI Scheme Code)",
        key="manual_identifier",
        help="Equity examples: RELIANCE, TCS, INFY or 500325.BO. Mutual Fund: AMFI scheme code.",
    )

    c6, c7, c8 = st.columns(3)
    txn_type = c6.selectbox("Transaction Type*", TXN_TYPES, key="manual_txn_type")
    txn_date = c7.date_input("Transaction Date*", value=date.today(), max_value=date.today(), key="manual_date")
    quantity = c8.number_input(
        "Quantity / Units*", min_value=0.0, step=1.0, format="%.4f", key="manual_quantity"
    )

    c9, c10, c11 = st.columns(3)
    price = c9.number_input(
        "Price per Unit*", min_value=0.0, step=0.01, format="%.4f", key="manual_price"
    )
    calculated_amount = round(float(quantity or 0) * float(price or 0), 2)
    with c10:
        st.metric("Total Amount (Auto = Qty × Price)", f"₹{calculated_amount:,.2f}")
    charges = c11.number_input(
        "Charges / Brokerage", min_value=0.0, step=0.01, format="%.2f", key="manual_charges"
    )

    if txn_type in ("BUY", "SIP"):
        net_cash = calculated_amount + float(charges or 0)
        st.info(f"Purchase value: ₹{calculated_amount:,.2f}  |  Including charges: ₹{net_cash:,.2f}")
    elif txn_type == "SELL":
        net_cash = max(calculated_amount - float(charges or 0), 0.0)
        st.info(f"Sale value: ₹{calculated_amount:,.2f}  |  Net after charges: ₹{net_cash:,.2f}")

    c12, c13 = st.columns(2)
    folio_no = c12.text_input("Folio Number (MF only)", key="manual_folio")
    remarks = c13.text_input("Remarks", key="manual_remarks")

    submitted = st.button("💾 Save Transaction", type="primary", key="manual_save")

    if submitted:
        member_id = int(members_df[members_df.name == member_name]["member_id"].iloc[0])
        broker_id = int(brokers_df[brokers_df.broker_name == broker_name]["broker_id"].iloc[0])
        data = dict(
            member_id=member_id,
            broker_id=broker_id,
            asset_type=asset_type,
            instrument_name=instrument_name.strip(),
            identifier=(_clean_identifier(identifier) or None),
            txn_type=txn_type,
            txn_date=str(txn_date),
            quantity=float(quantity or 0),
            price=float(price or 0),
            amount=float(calculated_amount),
            charges=float(charges or 0),
            folio_no=(folio_no.strip() or None),
            remarks=(remarks.strip() or None),
            upload_source="Manual",
            created_by=user["user_id"],
        )
        errors = validate_transaction_fields(data)
        if errors:
            for e in errors:
                st.error(e)
        else:
            status, msg, dupes = insert_transaction(data, force=False)
            if status == "DUPLICATE":
                st.session_state.pending_txn = data
                st.session_state.pending_dupes = dupes
            else:
                st.success(msg)
                st.cache_data.clear()
                st.session_state.manual_reset_requested = True
                st.rerun()

    if st.session_state.get("pending_txn"):
        st.warning("⚠️ A similar transaction already exists. Please review before confirming:")
        st.dataframe(pd.DataFrame(st.session_state.pending_dupes)
                     [["txn_date", "instrument_name", "txn_type", "quantity", "price", "amount"]])
        colx, coly = st.columns(2)
        if colx.button("✅ Yes, Save Anyway (Not a duplicate)", key="manual_save_duplicate"):
            insert_transaction(st.session_state.pending_txn, force=True)
            st.success("Transaction saved.")
            st.session_state.pending_txn = None
            st.session_state.pending_dupes = None
            st.cache_data.clear()
            st.session_state.manual_reset_requested = True
            st.rerun()
        if coly.button("❌ No, Discard this Entry", key="manual_discard_duplicate"):
            st.info("Discarded.")
            st.session_state.pending_txn = None
            st.session_state.pending_dupes = None
            st.rerun()


def excel_upload_tab(user, members_df, brokers_df):
    st.subheader("Bulk Upload via Excel Template")
    template_bytes = generate_template(members_df["name"].tolist(), brokers_df["broker_name"].tolist())
    st.download_button("⬇️ Download Excel Template", data=template_bytes,
                        file_name="Family_Investment_Upload_Template.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    uploaded = st.file_uploader("Upload filled template", type=["xlsx"])
    if uploaded:
        df = read_uploaded_excel(uploaded)
        st.write(f"Found {len(df)} rows.")
        member_map = dict(zip(members_df["name"], members_df["member_id"]))
        broker_map = dict(zip(brokers_df["broker_name"], brokers_df["broker_id"]))

        review_rows = []
        for idx, row in df.iterrows():
            data = {
                "member_id": member_map.get(str(row.get("member_name", "")).strip()),
                "broker_id": broker_map.get(str(row.get("broker_name", "")).strip()),
                "asset_type": str(row.get("asset_type", "")).strip(),
                "instrument_name": str(row.get("instrument_name", "")).strip(),
                "identifier": clean_optional_text(row.get("identifier")),
                "txn_type": str(row.get("txn_type", "")).strip().upper(),
                "txn_date": parse_date_safe(row.get("txn_date")),
                "quantity": safe_float(row.get("quantity")),
                "price": safe_float(row.get("price")),
                "amount": safe_float(row.get("amount")),
                "charges": safe_float(row.get("charges")) or 0,
                "folio_no": clean_optional_text(row.get("folio_no")),
                "remarks": clean_optional_text(row.get("remarks")),
                "upload_source": "Excel", "created_by": user["user_id"],
            }
            errors = validate_transaction_fields(data)
            status = "ERROR" if errors else "OK"
            if not errors:
                dupes = check_duplicate(data["member_id"], data["broker_id"], data["instrument_name"],
                                         data["txn_date"], data["txn_type"], data["quantity"], data["amount"])
                if dupes:
                    status = "DUPLICATE"

            review_rows.append({
                "Row": idx + 2, "Member": row.get("member_name"), "Instrument": row.get("instrument_name"),
                "Date": data["txn_date"], "Type": data["txn_type"], "Amount": data["amount"],
                "Status": status, "Issues": "; ".join(errors), "_data": data,
            })

        review_df = pd.DataFrame(review_rows)
        st.dataframe(review_df.drop(columns=["_data"]), use_container_width=True)

        ok_count = (review_df["Status"] == "OK").sum()
        dup_count = (review_df["Status"] == "DUPLICATE").sum()
        err_count = (review_df["Status"] == "ERROR").sum()
        st.write(f"✅ {ok_count} Ready | ⚠️ {dup_count} Possible Duplicates | ❌ {err_count} Errors")

        force_dupes = st.checkbox("Import possible-duplicate rows anyway (please verify manually first)")
        if st.button("📥 Import Valid Rows"):
            saved, skipped = 0, 0
            for _, r in review_df.iterrows():
                if r["Status"] == "OK" or (r["Status"] == "DUPLICATE" and force_dupes):
                    insert_transaction(r["_data"], force=True)
                    saved += 1
                else:
                    skipped += 1
            st.success(f"Imported {saved} transactions. Skipped {skipped}.")
            st.cache_data.clear()


def pdf_upload_tab(user, members_df, brokers_df):
    st.subheader("Extract Transactions from PDF (Beta)")
    st.caption("Works best with tabular contract notes/statements. Broker PDF formats vary widely, "
               "so please review and correct the extracted data below before saving.")
    member_for_pdf = st.selectbox("This PDF belongs to Member*", members_df["name"].tolist(), key="pdf_member")
    broker_for_pdf = st.selectbox("Broker/Platform*", brokers_df["broker_name"].tolist(), key="pdf_broker")
    asset_type_pdf = st.selectbox("Asset Type*", ASSET_TYPES, key="pdf_asset")
    pdf_file = st.file_uploader("Upload PDF Statement", type=["pdf"])

    if pdf_file:
        raw_df = extract_transactions_from_pdf(pdf_file)
        if raw_df.empty:
            st.warning("Could not detect tabular data in this PDF. Please use Manual Entry or Excel Upload instead.")
            return

        st.write("Extracted raw data — please verify / edit:")
        edited = st.data_editor(raw_df, num_rows="dynamic", use_container_width=True)

        st.info("Map the columns below to the required fields:")
        cols = ["--"] + list(edited.columns)
        c1, c2, c3, c4 = st.columns(4)
        col_date = c1.selectbox("Date column", cols)
        col_name = c2.selectbox("Instrument Name column", cols)
        col_qty = c3.selectbox("Quantity column", cols)
        col_amt = c4.selectbox("Amount column", cols)
        col_type = st.selectbox("Default Transaction Type for all rows", TXN_TYPES)

        if st.button("Preview Mapped Transactions"):
            member_id = int(members_df[members_df.name == member_for_pdf]["member_id"].iloc[0])
            broker_id = int(brokers_df[brokers_df.broker_name == broker_for_pdf]["broker_id"].iloc[0])
            mapped = []
            for _, r in edited.iterrows():
                qty = safe_float(r.get(col_qty)) if col_qty != "--" else 0
                amt = safe_float(r.get(col_amt)) if col_amt != "--" else 0
                data = {
                    "member_id": member_id, "broker_id": broker_id, "asset_type": asset_type_pdf,
                    "instrument_name": str(r.get(col_name, "")).strip() if col_name != "--" else "",
                    "identifier": None, "txn_type": col_type,
                    "txn_date": parse_date_safe(r.get(col_date)) if col_date != "--" else None,
                    "quantity": qty, "price": round(amt / qty, 4) if qty else 0, "amount": amt,
                    "charges": 0, "folio_no": None, "remarks": "Imported from PDF",
                    "upload_source": "PDF", "created_by": user["user_id"],
                }
                mapped.append(data)
            st.session_state.pdf_mapped_rows = mapped
            st.dataframe(pd.DataFrame(mapped))

        if st.session_state.get("pdf_mapped_rows") and st.button("💾 Save All Mapped Transactions"):
            saved = 0
            for data in st.session_state.pdf_mapped_rows:
                if not validate_transaction_fields(data):
                    insert_transaction(data, force=True)
                    saved += 1
            st.success(f"Saved {saved} transactions from PDF.")
            st.session_state.pdf_mapped_rows = None
            st.cache_data.clear()


def upload_page(user):
    st.title("⬆️ Upload Investments")
    members_df = get_members()
    brokers_df = get_brokers()
    if members_df.empty:
        st.warning("Please add at least one family member first (➕ Add Family Member).")
        return
    tab1, tab2, tab3 = st.tabs(["📝 Manual Entry", "📊 Excel Upload", "📄 PDF Extract (Beta)"])
    with tab1:
        manual_entry_tab(user, members_df, brokers_df)
    with tab2:
        excel_upload_tab(user, members_df, brokers_df)
    with tab3:
        pdf_upload_tab(user, members_df, brokers_df)


# ---------------------------------- SETTINGS --------------------------------------

def settings_page(user):
    st.title("⚙️ Settings")
    st.subheader("🔑 Change Password")
    with st.form("change_pw"):
        old = st.text_input("Current Password", type="password")
        new1 = st.text_input("New Password", type="password")
        new2 = st.text_input("Confirm New Password", type="password")
        sub = st.form_submit_button("Update Password")
        if sub:
            check = authenticate(user["mobile_no"], old)
            if not check:
                st.error("Current password is incorrect.")
            elif len(new1) < 4:
                st.error("New password too short.")
            elif new1 != new2:
                st.error("Passwords do not match.")
            else:
                change_password(user["user_id"], new1)
                st.success("Password updated successfully.")

    st.divider()
    st.subheader("🔌 Official Market Data Connection")
    st.caption(
        "Recommended: use an official broker API for current share prices instead of scraping NSE/BSE webpages. "
        "Credentials entered here remain only in this running Streamlit session and are not saved in the database."
    )
    provider_labels = ["AUTO - Best available", "HDFC Securities", "ICICI Direct Breeze", "Zerodha", "Upstox", "Yahoo Finance"]
    stored_pref = get_app_setting("preferred_price_provider", "AUTO")
    pref_map = {"AUTO - Best available":"AUTO","HDFC Securities":"HDFC","ICICI Direct Breeze":"ICICI","Zerodha":"ZERODHA","Upstox":"UPSTOX","Yahoo Finance":"YAHOO"}
    reverse_pref = {v:k for k,v in pref_map.items()}
    selected_pref = st.selectbox("Preferred share-price provider", provider_labels,
                                 index=provider_labels.index(reverse_pref.get(stored_pref,"AUTO - Best available")),
                                 key="preferred_price_provider_select")
    if st.button("Save Provider Preference", key="save_provider_preference"):
        set_app_setting("preferred_price_provider", pref_map[selected_pref])
        st.success("Price-provider priority saved locally.")

    with st.expander("Configure HDFC Securities InvestRight", expanded=False):
        st.write("Map each security's HDFC segment/token in Instrument Master. The endpoint is configurable so the app can follow the current InvestRight documentation without hard-coding a potentially changing URL.")
        hk=st.text_input("HDFC API Key", value=_runtime_secret("HDFC_API_KEY"), key="hdfc_api_key_input")
        ht=st.text_input("HDFC Access Token", value=_runtime_secret("HDFC_ACCESS_TOKEN"), type="password", key="hdfc_access_token_input")
        he=st.text_input("HDFC Fetch LTP Endpoint", value=_runtime_secret("HDFC_LTP_ENDPOINT"), key="hdfc_ltp_endpoint_input")
        if st.button("Use HDFC for this session", key="save_hdfc_session_credentials"):
            st.session_state["HDFC_API_KEY"]=hk.strip(); st.session_state["HDFC_ACCESS_TOKEN"]=ht.strip(); st.session_state["HDFC_LTP_ENDPOINT"]=he.strip()
            st.success("HDFC credentials loaded for this session.")

    with st.expander("Configure ICICI Direct Breeze", expanded=False):
        st.write("Breeze requires API Key, Secret Key and a session token. ICICI states the session token is generated daily/valid until midnight.")
        ik=st.text_input("ICICI Breeze API Key", value=_runtime_secret("ICICI_BREEZE_API_KEY"), key="icici_api_key_input")
        ise=st.text_input("ICICI Breeze Secret Key", value=_runtime_secret("ICICI_BREEZE_API_SECRET"), type="password", key="icici_secret_input")
        ist=st.text_input("ICICI Breeze Session Token", value=_runtime_secret("ICICI_BREEZE_SESSION_TOKEN"), type="password", key="icici_session_input")
        if st.button("Use ICICI Breeze for this session", key="save_icici_session_credentials"):
            st.session_state["ICICI_BREEZE_API_KEY"]=ik.strip(); st.session_state["ICICI_BREEZE_API_SECRET"]=ise.strip(); st.session_state["ICICI_BREEZE_SESSION_TOKEN"]=ist.strip()
            st.success("ICICI Breeze credentials loaded for this session.")

    with st.expander("Configure Zerodha Kite Connect (recommended if you use Zerodha)", expanded=False):
        st.write("Enter your Kite Connect API key and the current access token. The app uses only the LTP quote endpoint; it does not place orders.")
        zk = st.text_input("Zerodha API Key", value=_runtime_secret('ZERODHA_API_KEY'), key="zerodha_api_key_input")
        zt = st.text_input("Zerodha Access Token", value=_runtime_secret('ZERODHA_ACCESS_TOKEN'), type="password", key="zerodha_access_token_input")
        cza, czb = st.columns(2)
        if cza.button("Use Zerodha for this session", key="save_zerodha_session_credentials"):
            st.session_state['ZERODHA_API_KEY'] = zk.strip()
            st.session_state['ZERODHA_ACCESS_TOKEN'] = zt.strip()
            st.success("Zerodha credentials loaded for this session. Use Dashboard → Refresh Live Prices.")
        if czb.button("Test Zerodha Connection", key="test_zerodha_connection"):
            st.session_state['ZERODHA_API_KEY'] = zk.strip()
            st.session_state['ZERODHA_ACCESS_TOKEN'] = zt.strip()
            sample = _load_txn_cached()
            equities = sample[(sample['asset_type'] == 'Equity') & sample['identifier'].notna()] if not sample.empty else pd.DataFrame()
            if equities.empty:
                st.warning("Add at least one equity with a trading symbol before testing.")
            else:
                test_id = _clean_identifier(equities.iloc[0]['identifier'])
                pr, sym, err = fetch_zerodha_ltp(test_id)
                if pr is not None:
                    st.success(f"Connected successfully: {sym} LTP ₹{pr:,.2f}")
                else:
                    st.error(err or "Zerodha test failed.")

    with st.expander("Configure Upstox Market Quote API", expanded=False):
        st.write("Upstox uses instrument keys such as NSE_EQ|INE848E01016. If your portfolio stores only NSE trading symbols, Zerodha is easier to integrate.")
        ut = st.text_input("Upstox Access Token", value=_runtime_secret('UPSTOX_ACCESS_TOKEN'), type="password", key="upstox_access_token_input")
        if st.button("Use Upstox for this session", key="save_upstox_session_credentials"):
            st.session_state['UPSTOX_ACCESS_TOKEN'] = ut.strip()
            st.success("Upstox token loaded for this session. Use Dashboard → Refresh Live Prices.")

    st.info(
        "For persistent local configuration, you may set environment variables ZERODHA_API_KEY, "
        "ZERODHA_ACCESS_TOKEN and/or UPSTOX_ACCESS_TOKEN before starting the app. Avoid hard-coding secrets into the .py file."
    )

    st.divider()
    st.subheader("🧭 Instrument Master")
    st.caption("Map your normal portfolio identifier once to provider-specific codes. This makes provider switching seamless.")
    txm=_load_txn_cached()
    if not txm.empty:
        choices=sorted({_clean_identifier(x) for x in txm["identifier"].dropna().tolist() if _clean_identifier(x)})
        if choices:
            im_id=st.selectbox("Portfolio Identifier", choices, key="im_identifier")
            existing=get_instrument_mapping(im_id)
            cols=st.columns(3)
            im_name=cols[0].text_input("Instrument Name", value=str(existing.get("instrument_name") or ""), key="im_name")
            im_exchange=cols[1].selectbox("Exchange", ["NSE","BSE"], index=0 if str(existing.get("exchange") or "NSE").upper()!="BSE" else 1, key="im_exchange")
            im_symbol=cols[2].text_input("Trading Symbol", value=str(existing.get("trading_symbol") or im_id), key="im_symbol")
            cols2=st.columns(3)
            im_isin=cols2[0].text_input("ISIN", value=str(existing.get("isin") or ""), key="im_isin")
            im_hseg=cols2[1].text_input("HDFC Segment", value=str(existing.get("hdfc_segment") or im_exchange), key="im_hseg")
            im_htoken=cols2[2].text_input("HDFC Token", value=str(existing.get("hdfc_token") or ""), key="im_htoken")
            cols3=st.columns(3)
            im_icici=cols3[0].text_input("ICICI Breeze Stock Code", value=str(existing.get("icici_stock_code") or im_symbol), key="im_icici")
            im_up=cols3[1].text_input("Upstox Instrument Key", value=str(existing.get("upstox_key") or ""), key="im_up")
            im_y=cols3[2].text_input("Yahoo Symbol", value=str(existing.get("yahoo_symbol") or ""), key="im_y")
            if st.button("Save Instrument Mapping", key="save_instrument_mapping"):
                upsert_instrument_mapping(im_id,im_name,"Equity",im_exchange,im_symbol,im_isin,im_hseg,im_htoken,im_icici,im_up,im_y)
                st.success("Instrument mapping saved.")

    st.divider()
    st.subheader("💹 Manual Current Price / NAV Override")
    st.caption("Use this only if the live provider is temporarily unavailable or does not recognize an instrument. "
               "The value is saved in the local price cache and is used by the dashboard until the next successful live refresh.")
    txn_df_for_prices = _load_txn_cached()
    if not txn_df_for_prices.empty:
        price_options_df = txn_df_for_prices[["instrument_name", "identifier", "asset_type"]].copy()
        price_options_df["identifier"] = price_options_df["identifier"].apply(_clean_identifier)
        price_options_df = price_options_df[price_options_df["identifier"] != ""].drop_duplicates()
        if not price_options_df.empty:
            option_labels = {
                f"{r.instrument_name} | {r.identifier} | {r.asset_type}": (r.instrument_name, r.identifier, r.asset_type)
                for r in price_options_df.itertuples(index=False)
            }
            selected_label = st.selectbox("Instrument", list(option_labels.keys()), key="manual_price_override_instrument")
            selected_name, selected_identifier, selected_asset_type = option_labels[selected_label]
            manual_current = st.number_input("Current Price / NAV", min_value=0.0, step=0.01, format="%.4f",
                                             key="manual_current_price_override")
            if st.button("Save Manual Current Price", key="save_manual_current_price"):
                if save_manual_current_price(selected_identifier, selected_asset_type, selected_name, manual_current):
                    st.success("Manual current price saved successfully.")
                    st.cache_data.clear()
                    st.rerun()
                else:
                    st.error("Enter a current price greater than zero.")
        else:
            st.info("No investment with a usable identifier is available yet.")
    else:
        st.info("Add investments before using the manual price override.")

    st.divider()
    st.subheader("🏦 Manage Broker / Platform List")
    st.dataframe(get_brokers())
    with st.form("add_broker"):
        name = st.text_input("New Broker/Platform Name")
        btype = st.selectbox("Type", ["Equity Broker", "MF RTA", "Other"])
        sub2 = st.form_submit_button("Add Platform")
        if sub2 and name.strip():
            add_broker(name.strip(), btype)
            st.success("Platform added.")
            st.rerun()


# ------------------------------- MAIN APP SHELL ------------------------------------

def main_app():
    user = st.session_state.user
    with st.sidebar:
        st.markdown(f"### 👋 Welcome, {get_member_name(user['member_id'])}")
        st.caption(f"Role: {user['role'].title()}")
        page = st.radio("Navigate", [
            "🏠 Dashboard", "👪 Member View", "⬆️ Upload Investments",
            "📰 News Board", "🧾 All Transactions", "🔍 Query Explorer",
            "➕ Add Family Member", "⚙️ Settings",
        ])
        st.divider()
        if st.button("🚪 Logout"):
            st.session_state.logged_in = False
            st.session_state.user = None
            st.rerun()

    pages = {
        "🏠 Dashboard": dashboard_page,
        "👪 Member View": member_view_page,
        "⬆️ Upload Investments": lambda: upload_page(user),
        "📰 News Board": news_page,
        "🧾 All Transactions": all_transactions_page,
        "🔍 Query Explorer": query_explorer_page,
        "➕ Add Family Member": lambda: add_member_page(user),
        "⚙️ Settings": lambda: settings_page(user),
    }
    pages[page]()


def main():
    st.set_page_config(page_title=APP_NAME, page_icon="💰", layout="wide")
    inject_custom_css()
    init_db()

    for key, default in [("logged_in", False), ("must_change_password", False),
                          ("pending_txn", None), ("pending_dupes", None),
                          ("pdf_mapped_rows", None)]:
        if key not in st.session_state:
            st.session_state[key] = default

    if not st.session_state.logged_in:
        login_page()
        return
    if st.session_state.must_change_password:
        force_change_password_page()
        return
    main_app()



if __name__ == "__main__":
    # At this point we are already inside the Streamlit child process.
    main()
