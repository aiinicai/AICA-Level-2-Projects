"""Core logic for the MSME Payables Compliance Tracker.

Covers:
  * reading/validating uploaded Vendor Master, Bills and Payments files
  * statutory due-date computation (MSMED Act s.15: agreed days capped at 45, else 15)
  * bill-level ledger as of any date (paid on time / late / outstanding / overdue)
  * Section 43B(h) disallowance for a financial year
  * MSME Form 1 (half-yearly return) summary
  * pay-by list and month-end trend
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

import numpy as np
import pandas as pd

STATUTORY_MAX_DAYS = 45
NO_AGREEMENT_DAYS = 15
COVERED_CATEGORIES = ("Micro", "Small")
CATEGORIES = ["Micro", "Small", "Medium", "Not MSME"]
AGE_BUCKETS = ["Not yet due", "1-15 days", "16-30 days", "31-45 days", "46-90 days", "90+ days"]

# --------------------------------------------------------------------------- #
# Schemas (also drive the downloadable templates)
# --------------------------------------------------------------------------- #

@dataclass(frozen=True)
class Col:
    name: str
    required: bool
    kind: str  # text | date | number | category
    description: str
    example: str


SCHEMAS: dict[str, list[Col]] = {
    "vendor_master": [
        Col("Vendor Name", True, "text", "Supplier name exactly as in your books. Used to link Bills and Payments.", "Sri Lakshmi Plastics"),
        Col("MSME Category", True, "category", "Micro, Small, Medium or Not MSME (as per Udyam certificate).", "Micro"),
        Col("Vendor ID", False, "text", "Your internal vendor code (informational).", "V0001"),
        Col("PAN", False, "text", "Supplier PAN (printed on Form 1).", "ABCDE1234F"),
        Col("GSTIN", False, "text", "Supplier GSTIN (informational).", "33ABCDE1234F1Z5"),
        Col("Udyam Number", False, "text", "Udyam registration number.", "UDYAM-TN-02-0012345"),
        Col("Udyam Registration Date", False, "date", "Bills accepted before this date are not treated as covered. Leave blank if unknown.", "01-04-2021"),
        Col("Agreed Credit Days", False, "number", "Credit days in a WRITTEN agreement. Blank = no agreement (15 days applies). Values above 45 are capped to 45.", "30"),
    ],
    "bills": [
        Col("Bill No", True, "text", "Supplier invoice / bill number.", "INV-2026-0101"),
        Col("Vendor Name", True, "text", "Must match a Vendor Name in the Vendor Master.", "Sri Lakshmi Plastics"),
        Col("Bill Date", True, "date", "Date of the bill.", "12-07-2026"),
        Col("Bill Amount", True, "number", "Amount payable to the vendor (as you want it tracked, e.g. incl. GST). Must be > 0.", "118000"),
        Col("Acceptance Date", False, "date", "Date of acceptance/deemed acceptance of goods or services. Blank = Bill Date is used.", "14-07-2026"),
        Col("Description", False, "text", "Free text (informational).", "Moulded housings"),
        Col("Branch", False, "text", "Branch / cost centre for slicing the dashboard (informational).", "Chennai"),
    ],
    "payments": [
        Col("Vendor Name", True, "text", "Must match a Vendor Name in the Vendor Master.", "Sri Lakshmi Plastics"),
        Col("Payment Date", True, "date", "Date of payment / settlement.", "20-08-2026"),
        Col("Amount Paid", True, "number", "Amount paid. Record credit notes / adjustments as payments too.", "50000"),
        Col("Bill No", False, "text", "Bill this payment settles. Blank = on-account; applied to the vendor's oldest open bills (FIFO by due date).", "INV-2026-0101"),
        Col("Reference", False, "text", "Cheque / UTR / voucher no (informational).", "UTR123456"),
        Col("Payment Mode", False, "text", "TReDS or Other (blank = Other). Drives Form 1 column (d)(i) vs (d)(ii).", "TReDS"),
    ],
    "opening": [
        Col("Vendor Name", True, "text", "Must match a Vendor Name in the Vendor Master.", "Sri Lakshmi Plastics"),
        Col("Outstanding Amount", True, "number", "Amount unpaid to this vendor at the opening date. Party-wise: one total per vendor. Bill-wise: the unpaid part of each bill.", "250000"),
        Col("Bill No", False, "text", "Bill-wise mode: required. Party-wise mode: leave blank.", "INV-2026-0044"),
        Col("Bill Date", False, "date", "Bill-wise mode: required. Party-wise mode: optional - the date the dues are assumed to date from (blank = opening date).", "15-03-2026"),
        Col("Acceptance Date", False, "date", "Bill-wise only. Blank = Bill Date is used.", "17-03-2026"),
        Col("Description", False, "text", "Free text (informational).", "Moulded housings"),
        Col("Branch", False, "text", "Branch / cost centre (informational).", "Chennai"),
    ],
}

SHEET_NAMES = {"vendor_master": "Vendor Master", "bills": "Bills", "payments": "Payments", "opening": "Opening Balances"}

_ALIASES = {
    "vendorname": "Vendor Name", "vendor": "Vendor Name", "supplier": "Vendor Name",
    "suppliername": "Vendor Name", "partyname": "Vendor Name", "party": "Vendor Name",
    "msmecategory": "MSME Category", "category": "MSME Category", "enterprisetype": "MSME Category",
    "msmetype": "MSME Category", "udyamcategory": "MSME Category",
    "vendorid": "Vendor ID", "vendorcode": "Vendor ID",
    "pan": "PAN", "panno": "PAN", "pannumber": "PAN",
    "gstin": "GSTIN", "gstno": "GSTIN",
    "udyamnumber": "Udyam Number", "udyamno": "Udyam Number", "udyam": "Udyam Number",
    "udyamregistrationnumber": "Udyam Number",
    "udyamregistrationdate": "Udyam Registration Date", "udyamdate": "Udyam Registration Date",
    "agreedcreditdays": "Agreed Credit Days", "creditdays": "Agreed Credit Days",
    "creditperiod": "Agreed Credit Days", "agreeddays": "Agreed Credit Days",
    "billno": "Bill No", "billnumber": "Bill No", "invoiceno": "Bill No",
    "invoicenumber": "Bill No", "invno": "Bill No", "billref": "Bill No",
    "billdate": "Bill Date", "invoicedate": "Bill Date",
    "billamount": "Bill Amount", "invoiceamount": "Bill Amount", "amount": "Bill Amount",
    "acceptancedate": "Acceptance Date", "dateofacceptance": "Acceptance Date",
    "description": "Description", "narration": "Description",
    "branch": "Branch", "location": "Branch", "costcentre": "Branch",
    "paymentdate": "Payment Date", "paiddate": "Payment Date", "datepaid": "Payment Date",
    "amountpaid": "Amount Paid", "paidamount": "Amount Paid", "paymentamount": "Amount Paid",
    "reference": "Reference", "utr": "Reference", "chequeno": "Reference",
    "openingbalance": "Outstanding Amount", "outstandingamount": "Outstanding Amount", "outstanding": "Outstanding Amount",
    "balance": "Outstanding Amount", "balanceoutstanding": "Outstanding Amount", "closingbalance": "Outstanding Amount",
    "paymentmode": "Payment Mode", "modeofpayment": "Payment Mode", "mode": "Payment Mode",
}


def _hkey(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def _txt(v):
    """Cell -> stripped str, or None when blank/NaN (pandas may hold missing text as NaN)."""
    if v is None or v is pd.NA or (isinstance(v, float) and np.isnan(v)):
        return None
    t = str(v).strip()
    return None if t in ("", "nan", "None", "<NA>") else t


def vendor_key(name) -> str:
    """Normalised vendor key for matching across files."""
    s = re.sub(r"\s+", " ", str(name)).strip().casefold()
    s = re.sub(r"[.,]+", "", s)
    return s


def bill_key(no) -> str:
    return re.sub(r"\s+", "", str(no)).casefold()


# --------------------------------------------------------------------------- #
# Reading + standardising files
# --------------------------------------------------------------------------- #

def read_table(data: bytes, filename: str, kind: str) -> pd.DataFrame:
    """Read csv / xlsx / xls bytes. For workbooks, prefers the sheet named for `kind`."""
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "csv":
        for enc in ("utf-8-sig", "latin-1"):
            try:
                return pd.read_csv(io.BytesIO(data), dtype=object, encoding=enc)
            except UnicodeDecodeError:
                continue
        raise ValueError("Could not decode the CSV file.")
    if ext in ("xlsx", "xlsm", "xls"):
        engine = "xlrd" if ext == "xls" else "openpyxl"
        xf = pd.ExcelFile(io.BytesIO(data), engine=engine)
        want = SHEET_NAMES[kind].casefold()
        sheet = next((s for s in xf.sheet_names if s.casefold() == want), xf.sheet_names[0])
        return xf.parse(sheet, dtype=object)
    raise ValueError(f"Unsupported file type '.{ext}'. Use .xlsx, .xls or .csv.")


def _to_num(s: pd.Series) -> pd.Series:
    cleaned = s.astype(str).str.replace(",", "", regex=False).str.strip()
    cleaned = cleaned.replace({"nan": np.nan, "None": np.nan, "<NA>": np.nan, "": np.nan})
    return pd.to_numeric(cleaned, errors="coerce")


def _to_date(s: pd.Series) -> pd.Series:
    return pd.to_datetime(s, errors="coerce", dayfirst=True, format="mixed").dt.normalize()


def _norm_category(v) -> str | None:
    if v is None or (isinstance(v, float) and np.isnan(v)) or str(v).strip() == "":
        return None
    t = str(v).strip().casefold()
    if t.startswith("micro"):
        return "Micro"
    if t.startswith("small"):
        return "Small"
    if t.startswith("medium"):
        return "Medium"
    if t in ("not msme", "non msme", "non-msme", "other", "others", "large", "na", "n/a", "no", "not applicable"):
        return "Not MSME"
    return "?"


def standardise(df: pd.DataFrame, kind: str) -> tuple[pd.DataFrame, list[str]]:
    """Validate and clean an uploaded table. Returns (clean_df, warnings).

    Raises ValueError if required columns are missing.
    """
    spec = SCHEMAS[kind]
    issues: list[str] = []
    ren = {}
    for c in df.columns:
        canon = _ALIASES.get(_hkey(c))
        if kind == "opening" and _hkey(c) in ("amount", "billamount"):
            canon = "Outstanding Amount"
        if canon and canon not in ren.values():
            ren[c] = canon
    d = df.rename(columns=ren).copy()
    missing = [c.name for c in spec if c.required and c.name not in d.columns]
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(missing)}. Download the template to see the expected columns.")
    for c in spec:
        if c.name not in d.columns:
            d[c.name] = np.nan
    d = d[[c.name for c in spec]]
    d = d.dropna(how="all")
    d["_row"] = d.index + 2  # spreadsheet row number (header is row 1)

    for c in spec:
        if c.kind in ("text", "category"):
            d[c.name] = pd.Series([_txt(v) for v in d[c.name]], index=d.index, dtype=object)
        elif c.kind == "date":
            d[c.name] = _to_date(d[c.name])
        elif c.kind == "number":
            d[c.name] = _to_num(d[c.name])

    def drop(mask: pd.Series, why: str):
        nonlocal d
        if mask.any():
            rows = ", ".join(map(str, d.loc[mask, "_row"].tolist()[:15]))
            more = "…" if mask.sum() > 15 else ""
            issues.append(f"{int(mask.sum())} row(s) excluded – {why} (file rows: {rows}{more})")
            d = d[~mask]

    d["vkey"] = d["Vendor Name"].map(lambda v: vendor_key(v) if isinstance(v, str) and v else None)
    drop(d["vkey"].isna(), "Vendor Name is blank")

    if kind == "vendor_master":
        d["MSME Category"] = d["MSME Category"].map(_norm_category)
        bad = d["MSME Category"] == "?"
        if bad.any():
            issues.append(f"{int(bad.sum())} vendor(s) have an unrecognised MSME Category and were set to 'Not MSME': "
                          + ", ".join(d.loc[bad, "Vendor Name"].head(10)))
            d.loc[bad, "MSME Category"] = "Not MSME"
        drop(d["MSME Category"].isna(), "MSME Category is blank")
        dup = d.duplicated("vkey", keep="last")
        if dup.any():
            issues.append(f"{int(dup.sum())} duplicate vendor name(s) – last occurrence kept: "
                          + ", ".join(d.loc[dup, "Vendor Name"].head(10)))
            d = d[~dup]
        over = d["Agreed Credit Days"] > STATUTORY_MAX_DAYS
        if over.any():
            issues.append(f"{int(over.sum())} vendor(s) have Agreed Credit Days above {STATUTORY_MAX_DAYS}; they will be capped at {STATUTORY_MAX_DAYS}.")
        neg = d["Agreed Credit Days"] < 0
        d.loc[neg, "Agreed Credit Days"] = np.nan

    elif kind == "bills":
        d["bkey"] = d["Bill No"].map(lambda v: bill_key(v) if isinstance(v, str) and v else None)
        drop(d["bkey"].isna(), "Bill No is blank")
        drop(d["Bill Date"].isna(), "Bill Date is missing or not a valid date")
        drop(d["Bill Amount"].isna(), "Bill Amount is missing or not a number")
        drop(d["Bill Amount"] <= 0, "Bill Amount is zero/negative (record credit notes in the Payments file)")
        early = d["Acceptance Date"].notna() & (d["Acceptance Date"] < d["Bill Date"])
        if early.any():
            issues.append(f"{int(early.sum())} bill(s) have Acceptance Date before Bill Date – Bill Date used instead.")
            d.loc[early, "Acceptance Date"] = pd.NaT
        dup = d.duplicated(["vkey", "bkey"], keep=False)
        if dup.any():
            issues.append(f"{int(dup.sum())} row(s) share the same Vendor + Bill No. They are tracked separately but payments referring to that Bill No "
                          "will be applied to them oldest-first.")

    elif kind == "opening":
        d["bkey"] = d["Bill No"].map(lambda v: bill_key(v) if isinstance(v, str) and v else None)
        drop(d["Outstanding Amount"].isna(), "Outstanding Amount is missing or not a number")
        drop(d["Outstanding Amount"] <= 0, "Outstanding Amount is zero/negative")

    elif kind == "payments":
        d["bkey"] = d["Bill No"].map(lambda v: bill_key(v) if isinstance(v, str) and v else None)
        drop(d["Payment Date"].isna(), "Payment Date is missing or not a valid date")
        drop(d["Amount Paid"].isna(), "Amount Paid is missing or not a number")
        drop(d["Amount Paid"] == 0, "Amount Paid is zero")
        d["Payment Mode"] = d["Payment Mode"].map(lambda v: "TReDS" if v and "tred" in str(v).casefold() else "Other")

    d = d.drop(columns=["_row"]).reset_index(drop=True)
    return d, issues


def opening_to_bills(op: pd.DataFrame, mode: str, opening_date) -> tuple[pd.DataFrame, list[str]]:
    """Turn standardised opening balances into bill rows (raw bills layout, ready for standardise(..., 'bills')).

    mode 'bill'  : one row per unpaid bill (Bill No and Bill Date required).
    mode 'party' : one 'Opening balance' bill per vendor (amounts summed); Bill Date = given date or the opening date.
    The unpaid amount becomes the Bill Amount; later payments then reduce it like any other bill.
    """
    od = pd.Timestamp(opening_date).normalize()
    issues: list[str] = []
    cols = ["Bill No", "Vendor Name", "Bill Date", "Bill Amount", "Acceptance Date", "Description", "Branch"]
    if op is None or not len(op):
        return pd.DataFrame(columns=cols), issues
    d = op.copy()
    if mode == "bill":
        bad = d["Bill No"].isna() | d["Bill Date"].isna()
        if bad.any():
            issues.append(f"{int(bad.sum())} opening row(s) excluded – bill-wise mode needs both Bill No and Bill Date: "
                          + ", ".join(d.loc[bad, "Vendor Name"].astype(str).head(8)))
            d = d[~bad]
        late = d["Bill Date"] > od
        if late.any():
            issues.append(f"{int(late.sum())} opening bill(s) are dated after the opening date ({od:%d-%m-%Y}) – check they belong in the opening balance.")
        out = pd.DataFrame({"Bill No": d["Bill No"], "Vendor Name": d["Vendor Name"], "Bill Date": d["Bill Date"],
                            "Bill Amount": d["Outstanding Amount"], "Acceptance Date": d["Acceptance Date"],
                            "Description": d["Description"].where(d["Description"].notna(), "Opening balance b/f"),
                            "Branch": d["Branch"]})
    else:
        g = d.groupby("vkey", sort=False)
        out = pd.DataFrame({"Vendor Name": g["Vendor Name"].first(), "Bill Amount": g["Outstanding Amount"].sum(),
                            "Bill Date": g["Bill Date"].min(), "Branch": g["Branch"].first()}).reset_index(drop=True)
        out["Bill Date"] = out["Bill Date"].fillna(od)
        out["Bill No"] = "Opening balance"
        out["Acceptance Date"] = pd.NaT
        out["Description"] = "Opening balance b/f (party-wise)"
        n_dated = int(d.groupby("vkey")["Bill Date"].min().notna().sum())
        issues.append("Party-wise opening balances have no bill dates, so ageing, due dates and interest run from "
                      f"{'the date you entered' if n_dated else 'the opening date ' + format(od, '%d-%m-%Y')}. "
                      "Dues that are really older will be shown as younger than they are – use bill-wise opening balances for accuracy.")
        out = out[cols]
    return out[cols].reset_index(drop=True), issues


# --------------------------------------------------------------------------- #
# Ledger
# --------------------------------------------------------------------------- #

def _allocate(bills: pd.DataFrame, payments: pd.DataFrame):
    """Apply payments to bills. Returns (allocations df, unmatched df, unapplied df).

    allocations: one row per (bill index, payment date, amount)
    """
    alloc: list[tuple[int, pd.Timestamp, float, str]] = []
    remaining = bills["Bill Amount"].astype(float).copy()
    lookup: dict[tuple[str, str], list[int]] = {}
    for i, r in bills.sort_values(["due_date", "accept_date"]).iterrows():
        lookup.setdefault((r["vkey"], r["bkey"]), []).append(i)

    unmatched, unapplied = [], []

    def apply(idx_list, pdate, amt, mode):
        left = amt
        for i in idx_list:
            if left <= 0.005:
                break
            take = min(left, remaining[i])
            if take > 0.005:
                alloc.append((i, pdate, take, mode if mode == "TReDS" else "Other"))
                remaining[i] -= take
                left -= take
        return left

    pay = payments.sort_values("Payment Date")
    specific = pay[pay["bkey"].notna()]
    onacct = pay[pay["bkey"].isna()]

    for _, p in specific.iterrows():
        idxs = lookup.get((p["vkey"], p["bkey"]))
        if not idxs:
            unmatched.append(p)
            continue
        left = apply(idxs, p["Payment Date"], float(p["Amount Paid"]), p.get("Payment Mode"))
        if left > 0.005:
            q = p.copy()
            q["Amount Paid"] = left
            unapplied.append(q)

    by_vendor = {k: g.sort_values(["due_date", "accept_date"]).index.tolist() for k, g in bills.groupby("vkey")}
    for _, p in onacct.iterrows():
        idxs = [i for i in by_vendor.get(p["vkey"], []) if bills.at[i, "Bill Date"] <= p["Payment Date"]]
        if not idxs:
            unapplied.append(p)
            continue
        left = apply(idxs, p["Payment Date"], float(p["Amount Paid"]), p.get("Payment Mode"))
        if left > 0.005:
            q = p.copy()
            q["Amount Paid"] = left
            unapplied.append(q)

    a = pd.DataFrame(alloc, columns=["idx", "Payment Date", "Amount", "Mode"])
    return a, pd.DataFrame(unmatched), pd.DataFrame(unapplied)


@dataclass
class LedgerResult:
    ledger: pd.DataFrame
    unmatched_payments: pd.DataFrame   # payments naming a bill that does not exist
    unapplied_payments: pd.DataFrame   # advances / excess payments
    unknown_vendors: list[str]         # in bills but not in the vendor master
    as_of: pd.Timestamp
    allocations: pd.DataFrame | None = None   # idx (ledger row), Payment Date, Amount, Mode


def build_ledger(bills: pd.DataFrame, payments: pd.DataFrame | None, master: pd.DataFrame, as_of) -> LedgerResult:
    as_of = pd.Timestamp(as_of).normalize()
    b = bills[bills["Bill Date"] <= as_of].copy().reset_index(drop=True)

    m = master.drop_duplicates("vkey", keep="last").set_index("vkey")
    b["Category"] = b["vkey"].map(m["MSME Category"])
    unknown = sorted(b.loc[b["Category"].isna(), "Vendor Name"].unique().tolist())
    b["Category"] = b["Category"].fillna("Unclassified")
    b["PAN"] = b["vkey"].map(m["PAN"])
    b["Udyam Number"] = b["vkey"].map(m["Udyam Number"])
    reg = b["vkey"].map(m["Udyam Registration Date"])
    agreed = b["vkey"].map(m["Agreed Credit Days"])

    b["accept_date"] = b["Acceptance Date"].fillna(b["Bill Date"])
    days = np.where(agreed.notna() & (agreed > 0), np.minimum(agreed, STATUTORY_MAX_DAYS), NO_AGREEMENT_DAYS)
    b["Days Allowed"] = days.astype(int)
    b["due_date"] = b["accept_date"] + pd.to_timedelta(b["Days Allowed"], unit="D")
    b["Covered"] = b["Category"].isin(COVERED_CATEGORIES) & (reg.isna() | (b["accept_date"] >= reg))

    if payments is not None and len(payments):
        p = payments[payments["Payment Date"] <= as_of].copy()
    else:
        p = pd.DataFrame(columns=["vkey", "bkey", "Payment Date", "Amount Paid", "Reference", "Payment Mode"])
    if len(p):
        alloc, unmatched, unapplied = _allocate(b, p)
    else:
        alloc = pd.DataFrame(columns=["idx", "Payment Date", "Amount", "Mode"])
        unmatched = unapplied = pd.DataFrame()

    b["Paid"] = 0.0
    b["Paid On Time"] = 0.0
    b["Paid Late"] = 0.0
    b["Last Payment Date"] = pd.NaT
    if len(alloc):
        alloc["due"] = alloc["idx"].map(b["due_date"])
        alloc["ontime"] = alloc["Payment Date"] <= alloc["due"]
        tot = alloc.groupby("idx")["Amount"].sum()
        on = alloc[alloc["ontime"]].groupby("idx")["Amount"].sum()
        late = alloc[~alloc["ontime"]].groupby("idx")["Amount"].sum()
        last = alloc.groupby("idx")["Payment Date"].max()
        b["Paid"] = b.index.map(tot).fillna(0.0)
        b["Paid On Time"] = b.index.map(on).fillna(0.0)
        b["Paid Late"] = b.index.map(late).fillna(0.0)
        b["Last Payment Date"] = b.index.map(last)

    b["Outstanding"] = (b["Bill Amount"] - b["Paid"]).clip(lower=0).round(2)
    settled = b["Outstanding"] <= 0.005
    b["Days Overdue"] = np.where(~settled, (as_of - b["due_date"]).dt.days.clip(lower=0), 0)
    b["Status"] = np.select(
        [settled & (b["Paid Late"] <= 0.005), settled, b["Days Overdue"] > 0],
        ["Paid on time", "Paid late", "Overdue"],
        default="Within limit",
    )
    od = b["Days Overdue"]
    b["Age Bucket"] = np.select(
        [settled | (od <= 0), od <= 15, od <= 30, od <= 45, od <= 90],
        [AGE_BUCKETS[0], AGE_BUCKETS[1], AGE_BUCKETS[2], AGE_BUCKETS[3], AGE_BUCKETS[4]],
        default=AGE_BUCKETS[5],
    )
    b.loc[settled, "Age Bucket"] = "Settled"
    b = b.rename(columns={"accept_date": "Accepted On", "due_date": "Due Date"})
    cols = ["Bill No", "Vendor Name", "Category", "Covered", "Branch", "Bill Date", "Accepted On", "Days Allowed",
            "Due Date", "Bill Amount", "Paid", "Paid On Time", "Paid Late", "Outstanding", "Last Payment Date",
            "Days Overdue", "Status", "Age Bucket", "PAN", "Udyam Number", "vkey", "bkey"]
    return LedgerResult(b[cols], unmatched, unapplied, unknown, as_of, alloc)


# --------------------------------------------------------------------------- #
# Reports
# --------------------------------------------------------------------------- #

def fy_bounds(fy_start_year: int) -> tuple[pd.Timestamp, pd.Timestamp]:
    return pd.Timestamp(fy_start_year, 4, 1), pd.Timestamp(fy_start_year + 1, 3, 31)


def fy_start_year_of(d) -> int:
    d = pd.Timestamp(d)
    return d.year if d.month >= 4 else d.year - 1


def half_year_end(d) -> pd.Timestamp:
    """Most recent half-year end (30 Sep / 31 Mar) on or before d."""
    d = pd.Timestamp(d)
    cands = [pd.Timestamp(y, m, dd) for y in (d.year - 1, d.year) for m, dd in ((3, 31), (9, 30))]
    return max(c for c in cands if c <= d)


def form1(ledger: pd.DataFrame, as_of) -> tuple[pd.DataFrame, pd.DataFrame]:
    """MSME Form 1: dues to Micro/Small suppliers outstanding beyond 45 days from acceptance.

    Returns (vendor_summary, bill_annexure).
    """
    as_of = pd.Timestamp(as_of).normalize()
    L = ledger
    mask = L["Covered"] & (L["Outstanding"] > 0.005) & ((as_of - L["Accepted On"]).dt.days > STATUTORY_MAX_DAYS)
    ann = L[mask].copy()
    ann["Days Since Acceptance"] = (as_of - ann["Accepted On"]).dt.days
    if ann.empty:
        empty = pd.DataFrame(columns=["Sl No", "Supplier Name", "PAN", "Udyam Number", "Category", "Amount Due",
                                      "Date From Which Due", "Reason For Delay"])
        return empty, ann
    g = (ann.groupby(["vkey"], as_index=False)
         .agg(**{"Supplier Name": ("Vendor Name", "first"), "PAN": ("PAN", "first"),
                 "Udyam Number": ("Udyam Number", "first"), "Category": ("Category", "first"),
                 "Amount Due": ("Outstanding", "sum"), "Date From Which Due": ("Due Date", "min")})
         .sort_values("Amount Due", ascending=False).drop(columns="vkey"))
    g.insert(0, "Sl No", range(1, len(g) + 1))
    g["Reason For Delay"] = ""
    ann = ann.sort_values(["Vendor Name", "Due Date"])
    return g.reset_index(drop=True), ann


def disallowance_43bh(bills, payments, master, fy_start_year: int) -> dict:
    """Section 43B(h) view for a financial year.

    * disallowed: Micro/Small bills accepted in the FY, unpaid at 31-Mar and already past their due date.
    * reversal: amounts disallowed in the prior FY that were paid during this FY (deductible now).
    * not_yet_due: unpaid at 31-Mar but due date falls after year end (not disallowed; pay in time).
    """
    start, end = fy_bounds(fy_start_year)
    cur = build_ledger(bills, payments, master, end).ledger
    in_fy = cur[(cur["Accepted On"] >= start) & (cur["Accepted On"] <= end) & cur["Covered"]]
    out = in_fy[in_fy["Outstanding"] > 0.005]
    disallowed = out[out["Due Date"] <= end].copy()
    not_yet_due = out[out["Due Date"] > end].copy()

    pstart, pend = fy_bounds(fy_start_year - 1)
    prev = build_ledger(bills, payments, master, pend).ledger
    prev_dis = prev[(prev["Accepted On"] >= pstart) & prev["Covered"] & (prev["Outstanding"] > 0.005) & (prev["Due Date"] <= pend)]
    if len(prev_dis):
        cm = cur.set_index(["vkey", "bkey", "Bill Date"])["Outstanding"]
        rev = prev_dis.copy()
        keys = list(zip(rev["vkey"], rev["bkey"], rev["Bill Date"]))
        now_out = pd.Series([cm.get(k, np.nan) for k in keys], index=rev.index)
        rev["Outstanding Now"] = now_out
        rev["Paid In FY"] = (rev["Outstanding"] - rev["Outstanding Now"]).clip(lower=0)
        reversal = rev[rev["Paid In FY"] > 0.005].copy()
    else:
        reversal = prev_dis.copy()
        reversal["Paid In FY"] = []
    # late payments within the FY itself (allowed in year of payment; shown for interest awareness)
    paid_late = in_fy[in_fy["Paid Late"] > 0.005].copy()
    return {"fy_start": start, "fy_end": end, "disallowed": disallowed, "not_yet_due": not_yet_due,
            "reversal": reversal, "paid_late": paid_late}


def pay_by_list(ledger: pd.DataFrame, as_of, horizon_days: int = 15) -> pd.DataFrame:
    as_of = pd.Timestamp(as_of).normalize()
    L = ledger
    m = L["Covered"] & (L["Outstanding"] > 0.005) & (L["Due Date"] <= as_of + pd.Timedelta(days=horizon_days))
    out = L[m].copy()
    out["Days To Due"] = (out["Due Date"] - as_of).dt.days
    return out.sort_values("Due Date")[["Vendor Name", "Bill No", "Due Date", "Days To Due", "Outstanding", "Status", "Branch"]]


def month_end_trend(bills, payments, master, as_of, months: int = 12) -> pd.DataFrame:
    as_of = pd.Timestamp(as_of).normalize()
    ends = [e for e in pd.date_range(end=as_of, periods=months + 1, freq="ME")]
    if not ends or ends[-1] < as_of:
        ends.append(as_of)
    rows = []
    for e in ends[-months:]:
        L = build_ledger(bills, payments, master, e).ledger
        c = L[L["Covered"] & (L["Outstanding"] > 0.005)]
        rows.append({"Month End": e, "Outstanding (Micro+Small)": c["Outstanding"].sum(),
                     "Overdue (Micro+Small)": c.loc[c["Days Overdue"] > 0, "Outstanding"].sum()})
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------- #
# Indian number formatting + styled Excel output
# --------------------------------------------------------------------------- #

# Excel number format with Indian digit grouping (12,34,56,789.00). Conditional sections cover up to crores;
# negatives and values below one lakh use the standard 3-digit grouping, which is identical there.
INR_XL = r"[>=10000000]##\,##\,##\,##0.00;[>=100000]##\,##\,##0.00;##,##0.00"


def fmt_inr(x, decimals: int = 2) -> str:
    """12345678.5 -> '1,23,45,678.50' (Indian grouping). Blank for missing values."""
    if x is None or (isinstance(x, float) and np.isnan(x)) or x is pd.NA or x is pd.NaT:
        return ""
    x = float(x)
    neg = x < 0
    ip, _, fp = f"{abs(x):.{decimals}f}".partition(".")
    if len(ip) > 3:
        head, tail = ip[:-3], ip[-3:]
        ip = re.sub(r"(\d)(?=(\d\d)+$)", r"\1,", head) + "," + tail
    return ("-" if neg else "") + ip + (f".{fp}" if decimals else "")


_AMOUNT_RE = re.compile(r"(amount|outstanding|paid|principal|interest payable)", re.I)
_SUM_EXTRA = {"Bills"}
_NO_TOTAL_SHEETS = {"Vendor Master", "Bank rates"}
_STATUS_COLOURS = {"Overdue": ("FFC7CE", "9C0006"), "Open": ("FFC7CE", "9C0006"), "Paid late": ("FFEB9C", "9C5700"),
                   "Paid on time": ("C6EFCE", "006100"), "Within limit": ("DDEBF7", "1F4E78")}


def is_amount_col(name, series: pd.Series) -> bool:
    """Rupee-amount column: numeric and named like an amount (not days / counts / rates)."""
    nm = str(name).strip().lower()
    counted = nm.endswith("no.") or nm.endswith(" no") or "days" in nm or "%" in nm      # counts, days, rates are not rupee amounts
    return (pd.api.types.is_numeric_dtype(series) and not pd.api.types.is_bool_dtype(series)
            and bool(_AMOUNT_RE.search(nm)) and not counted)


def _style_sheet(ws, df: pd.DataFrame, add_totals: bool) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    n, ncol = len(df), len(df.columns)
    if ncol == 0:
        return
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    band = PatternFill("solid", fgColor="F2F7FC")

    kinds = []
    for c in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[c]):
            kinds.append("date")
        elif is_amount_col(c, df[c]):
            kinds.append("amount")
        elif pd.api.types.is_numeric_dtype(df[c]) and not pd.api.types.is_bool_dtype(df[c]):
            kinds.append("number")
        else:
            kinds.append("text")

    for j, c in enumerate(df.columns, start=1):
        h = ws.cell(row=1, column=j)
        h.font, h.fill, h.border = Font(bold=True, color="FFFFFF"), head_fill, box
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32

    for i in range(n):
        r = i + 2
        for j, kind in enumerate(kinds, start=1):
            cell = ws.cell(row=r, column=j)
            cell.border = box
            if i % 2 == 1:
                cell.fill = band
            if kind == "amount":
                cell.number_format = INR_XL
            elif kind == "date":
                cell.number_format = "DD-MM-YYYY"
                cell.alignment = Alignment(horizontal="center")
            elif kind == "number":
                cell.number_format = "General"
            if df.columns[j - 1] == "Status" and cell.value in _STATUS_COLOURS:
                bg, fg = _STATUS_COLOURS[cell.value]
                cell.fill, cell.font = PatternFill("solid", fgColor=bg), Font(color=fg, bold=True)
                cell.alignment = Alignment(horizontal="center")

    sum_cols = [j for j, (c, k) in enumerate(zip(df.columns, kinds), start=1) if k == "amount" or c in _SUM_EXTRA]
    if add_totals and n and sum_cols:
        tr = n + 2
        label_col = next((j for j in range(1, ncol + 1) if j not in sum_cols), 1)
        top = Border(left=thin, right=thin, top=Side(style="medium", color="1F4E78"), bottom=Side(style="double", color="1F4E78"))
        tfill = PatternFill("solid", fgColor="DDEBF7")
        for j in range(1, ncol + 1):
            cell = ws.cell(row=tr, column=j)
            cell.border, cell.fill, cell.font = top, tfill, Font(bold=True)
        ws.cell(row=tr, column=label_col, value="Total")
        for j in sum_cols:
            L = get_column_letter(j)
            cell = ws.cell(row=tr, column=j, value=f"=SUBTOTAL(9,{L}2:{L}{n + 1})")   # respects filters
            cell.number_format = INR_XL if kinds[j - 1] == "amount" else "#,##0"

    for j, (c, kind) in enumerate(zip(df.columns, kinds), start=1):
        if kind == "amount":
            body = [len(fmt_inr(v)) for v in df[c].head(500)] + ([len(fmt_inr(df[c].sum()))] if n else [])
        elif kind == "date":
            body = [10]
        else:
            body = [len(str(v)) for v in df[c].head(500) if v is not None]
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(max([len(str(c)) * 0.95 + 2] + [b + 2 for b in body]), 45))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{n + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def export_workbook(sheets: dict[str, pd.DataFrame], totals: bool = True) -> bytes:
    """Write DataFrames to a formatted .xlsx: bold header, banded rows, Indian-grouped amounts, a Total row
    (live SUBTOTAL formulas) under amount columns, filters and frozen header."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="DD-MM-YYYY", date_format="DD-MM-YYYY") as xw:
        for name, df in sheets.items():
            df = df.drop(columns=[c for c in ("vkey", "bkey") if c in df.columns]).reset_index(drop=True)
            sn = name[:31]
            df.to_excel(xw, sheet_name=sn, index=False)
            _style_sheet(xw.sheets[sn], df, totals and name not in _NO_TOTAL_SHEETS)
        xw.book.calculation.fullCalcOnLoad = True
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# MSME Form 1 – supplier-wise table in the prescribed column layout
# --------------------------------------------------------------------------- #

F1_COLS = [
    "S.No.", "Name of MSE Supplier", "PAN of the Supplier",
    "Paid ≤45d TReDS - No.", "Paid ≤45d TReDS - Amount (Rs.)",
    "Paid ≤45d Other mode - No.", "Paid ≤45d Other mode - Amount (Rs.)",
    "Paid after 45d - No.", "Paid after 45d - Amount (Rs.)",
    "Outstanding ≤45d - No.", "Outstanding ≤45d - Amount (Rs.)",
    "Outstanding >45d - No.", "Outstanding >45d - Amount (Rs.)",
    "Reason for delay in payment / amount outstanding",
]
_F1_BUCKETS = ["d_treds", "d_other", "e", "f", "g"]


def half_year_start(as_of) -> pd.Timestamp:
    """First day of the half-year (1 Apr / 1 Oct) containing as_of."""
    d = pd.Timestamp(as_of)
    return pd.Timestamp(d.year, 4, 1) if 4 <= d.month <= 9 else (
        pd.Timestamp(d.year, 10, 1) if d.month >= 10 else pd.Timestamp(d.year - 1, 10, 1))


def form1_table(res: "LedgerResult", as_of, period_start=None) -> pd.DataFrame:
    """Supplier-wise Form 1 (columns a–h).

    * (d)/(e) Paid: payments made in [period_start, as_of] against covered (Micro/Small) bills,
      classed by days from acceptance to payment (<=45 within, >45 after); (d) split by TReDS / other mode.
    * (f)/(g) Outstanding at as_of, classed by days since acceptance (<=45 / >45).
    * "No." = number of distinct bills; Amount = rupees. A bill settled in parts can appear in more than one column.
    """
    as_of = pd.Timestamp(as_of).normalize()
    start = pd.Timestamp(period_start) if period_start is not None else half_year_start(as_of)
    L = res.ledger
    cov = L[L["Covered"]]
    rows: list[pd.DataFrame] = []

    A = res.allocations
    if A is not None and len(A):
        A = A[A["idx"].isin(cov.index) & (A["Payment Date"] >= start) & (A["Payment Date"] <= as_of)].copy()
        if len(A):
            A["days"] = (A["Payment Date"] - A["idx"].map(L["Accepted On"])).dt.days
            A["bucket"] = np.where(A["days"] > STATUTORY_MAX_DAYS, "e", np.where(A["Mode"] == "TReDS", "d_treds", "d_other"))
            A["vkey"] = A["idx"].map(L["vkey"])
            rows.append(A[["vkey", "idx", "bucket", "Amount"]])
    O = cov[cov["Outstanding"] > 0.005].copy()
    if len(O):
        O["bucket"] = np.where((as_of - O["Accepted On"]).dt.days > STATUTORY_MAX_DAYS, "g", "f")
        O["idx"] = O.index
        rows.append(O.rename(columns={"Outstanding": "Amount"})[["vkey", "idx", "bucket", "Amount"]])
    if not rows:
        return pd.DataFrame(columns=F1_COLS)
    T = pd.concat(rows, ignore_index=True)
    g = T.groupby(["vkey", "bucket"]).agg(n=("idx", "nunique"), amt=("Amount", "sum")).reset_index()
    names = cov.drop_duplicates("vkey").set_index("vkey")[["Vendor Name", "PAN"]]
    out = []
    for vkey, sub in g.groupby("vkey"):
        rec = {"Name of MSE Supplier": names.at[vkey, "Vendor Name"], "PAN of the Supplier": names.at[vkey, "PAN"]}
        vals = {r.bucket: (int(r.n), round(float(r.amt), 2)) for r in sub.itertuples()}
        for col_n, col_a, b in zip(F1_COLS[3:13:2], F1_COLS[4:13:2], _F1_BUCKETS):
            rec[col_n], rec[col_a] = vals.get(b, (0, 0.0))
        out.append(rec)
    df = pd.DataFrame(out).sort_values("Name of MSE Supplier", key=lambda s: s.str.casefold()).reset_index(drop=True)
    df.insert(0, "S.No.", range(1, len(df) + 1))
    df["Reason for delay in payment / amount outstanding"] = ""
    return df[F1_COLS]


def form1_workbook(table: pd.DataFrame, annexure: pd.DataFrame, company: str, as_of, period_start) -> bytes:
    """Excel with the Form 1 two-level header (a)–(h), totals row, and a bill annexure sheet."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    base = export_workbook({"Bill annexure": annexure})
    wb = load_workbook(io.BytesIO(base))
    ws = wb.create_sheet("Form 1", 0)

    ws["A1"] = f"MSME Form 1 – {company}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Period: {pd.Timestamp(period_start):%d-%m-%Y} to {pd.Timestamp(as_of):%d-%m-%Y}"

    fill = PatternFill("solid", fgColor="DDEBF7")
    thin = Side(style="thin", color="7F7F7F")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def head(rng: str, text: str):
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value, c.alignment, c.font = text, center, Font(bold=True)
        for row in ws[rng]:
            for cell in row:
                cell.fill, cell.border = fill, border

    for col, text in (("A", "(a)\nS.No."), ("B", "(b)\nName of MSE Supplier"), ("C", "(c)\nPAN of the Supplier"),
                      ("N", "(h)\nReason for delay in payment / amount outstanding")):
        head(f"{col}4:{col}6", text)
    head("D4:G4", "(d)\nPaid within 45 days")
    head("D5:E5", "Through TReDS (i)")
    head("F5:G5", "Other mode of payments (ii)")
    head("H4:I5", "(e)\nPaid after 45 days")
    head("J4:K5", "(f)\nOutstanding for 45 days or less")
    head("L4:M5", "(g)\nOutstanding for more than 45 days")
    for i, col in enumerate("DEFGHIJKLM"):
        head(f"{col}6:{col}6", "No." if i % 2 == 0 else "Amount (Rs.)")

    r0 = 7
    for i, row in enumerate(table.itertuples(index=False), start=r0):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=None if (isinstance(v, float) and pd.isna(v)) else v)
            c.border = border
            if j in (5, 7, 9, 11, 13):
                c.number_format = INR_XL
    tr = r0 + len(table)
    ws.cell(row=tr, column=2, value="Total").font = Font(bold=True)
    tot_fill = PatternFill("solid", fgColor="DDEBF7")
    for j in range(1, 15):
        ws.cell(row=tr, column=j).fill = tot_fill
        ws.cell(row=tr, column=j).border = border
    ws.cell(row=tr, column=2).font = Font(bold=True)
    for j in range(4, 14):
        L = chr(64 + j)
        c = ws.cell(row=tr, column=j, value=f"=SUM({L}{r0}:{L}{tr - 1})" if len(table) else 0)
        c.font, c.border = Font(bold=True), border
        c.number_format = INR_XL if j % 2 == 1 else "0"
    for col, w in zip("ABCDEFGHIJKLMN", (6, 34, 14, 7, 16, 7, 16, 7, 16, 7, 16, 7, 16, 40)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "D7"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
