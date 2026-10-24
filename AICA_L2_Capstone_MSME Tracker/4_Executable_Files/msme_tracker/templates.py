"""Downloadable input templates (xlsx workbook with instructions, or per-file csv/xlsx)."""
from __future__ import annotations

import io

import pandas as pd

from engine import SCHEMAS, SHEET_NAMES


def _header(kind: str) -> list[str]:
    return [c.name for c in SCHEMAS[kind]]


def instructions_df(kind: str) -> pd.DataFrame:
    return pd.DataFrame([{"Column": c.name, "Required?": "Yes" if c.required else "Optional", "Type": c.kind,
                          "What to enter": c.description, "Example": c.example} for c in SCHEMAS[kind]])


def single_template(kind: str, fmt: str) -> tuple[bytes, str, str]:
    """Returns (bytes, filename, mime) for one input file's template. fmt: 'xlsx' or 'csv'."""
    base = f"MSME_{SHEET_NAMES[kind].replace(' ', '_')}_template"
    empty = pd.DataFrame(columns=_header(kind))
    if fmt == "csv":
        return empty.to_csv(index=False).encode("utf-8-sig"), f"{base}.csv", "text/csv"
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        empty.to_excel(xw, sheet_name=SHEET_NAMES[kind], index=False)
        instructions_df(kind).to_excel(xw, sheet_name="Instructions", index=False)
        _style(xw, SHEET_NAMES[kind], _header(kind), kind)
        _width(xw.sheets["Instructions"], instructions_df(kind))
    return buf.getvalue(), f"{base}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def combined_template() -> tuple[bytes, str, str]:
    """One workbook with all three input sheets + an Instructions sheet."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for kind in SCHEMAS:
            pd.DataFrame(columns=_header(kind)).to_excel(xw, sheet_name=SHEET_NAMES[kind], index=False)
            _style(xw, SHEET_NAMES[kind], _header(kind), kind)
        rows = []
        for kind in SCHEMAS:
            d = instructions_df(kind)
            d.insert(0, "Sheet", SHEET_NAMES[kind])
            rows.append(d)
        ins = pd.concat(rows, ignore_index=True)
        ins.to_excel(xw, sheet_name="Instructions", index=False)
        _width(xw.sheets["Instructions"], ins)
    return buf.getvalue(), "MSME_Tracker_Input_Template.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _style(xw, sheet: str, header: list[str], kind: str) -> None:
    from openpyxl.styles import Font, PatternFill
    ws = xw.sheets[sheet]
    req = {c.name for c in SCHEMAS[kind] if c.required}
    for i, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78" if name in req else "7F7F7F")
        ws.column_dimensions[cell.column_letter].width = max(16, len(name) + 4)
    ws.freeze_panes = "A2"


def _width(ws, df: pd.DataFrame) -> None:
    for i, col in enumerate(df.columns, start=1):
        w = max([len(str(col))] + [len(str(v)) for v in df[col]]) + 2
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(w, 70)
