import io
import sys
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import engine as E  # noqa: E402


def master():
    raw = pd.DataFrame({
        "Vendor Name": ["Micro Co", "Small Co", "Medium Co", "Late Reg Co"],
        "MSME Category": ["micro", "Small", "Medium", "Micro"],
        "Agreed Credit Days": [None, 30, None, None],
        "Udyam Registration Date": [None, None, None, "01-08-2026"],
    })
    m, issues = E.standardise(raw, "vendor_master")
    assert not issues
    return m


def bills(rows):
    raw = pd.DataFrame(rows, columns=["Bill No", "Vendor Name", "Bill Date", "Bill Amount"])
    b, issues = E.standardise(raw, "bills")
    assert not issues, issues
    return b


def pays(rows):
    raw = pd.DataFrame(rows, columns=["Vendor Name", "Payment Date", "Amount Paid", "Bill No"])
    p, issues = E.standardise(raw, "payments")
    assert not issues, issues
    return p


def test_due_date_rules():
    b = bills([["1", "Micro Co", "01-07-2026", 100], ["2", "Small Co", "01-07-2026", 100], ["3", "Medium Co", "01-07-2026", 100]])
    L = E.build_ledger(b, None, master(), "2026-09-30").ledger.set_index("Bill No")
    assert L.loc["1", "Days Allowed"] == 15 and L.loc["1", "Due Date"] == pd.Timestamp("2026-07-16")   # no agreement
    assert L.loc["2", "Days Allowed"] == 30 and L.loc["2", "Due Date"] == pd.Timestamp("2026-07-31")   # agreement
    assert bool(L.loc["1", "Covered"]) and bool(L.loc["2", "Covered"]) and not bool(L.loc["3", "Covered"])


def test_agreed_days_capped_at_45():
    m = master()
    m.loc[m["Vendor Name"] == "Small Co", "Agreed Credit Days"] = 90
    b = bills([["1", "Small Co", "01-07-2026", 100]])
    L = E.build_ledger(b, None, m, "2026-09-30").ledger
    assert L.loc[0, "Days Allowed"] == 45


def test_acceptance_date_overrides_bill_date():
    raw = pd.DataFrame([["1", "Micro Co", "01-07-2026", 100, "05-07-2026"]],
                       columns=["Bill No", "Vendor Name", "Bill Date", "Bill Amount", "Acceptance Date"])
    b, _ = E.standardise(raw, "bills")
    L = E.build_ledger(b, None, master(), "2026-09-30").ledger
    assert L.loc[0, "Due Date"] == pd.Timestamp("2026-07-20")


def test_status_paid_on_time_late_and_overdue():
    b = bills([["1", "Micro Co", "01-07-2026", 100], ["2", "Micro Co", "01-07-2026", 100], ["3", "Micro Co", "01-07-2026", 100]])
    p = pays([["Micro Co", "10-07-2026", 100, "1"], ["Micro Co", "01-08-2026", 100, "2"]])
    L = E.build_ledger(b, p, master(), "2026-09-30").ledger.set_index("Bill No")
    assert L.loc["1", "Status"] == "Paid on time"
    assert L.loc["2", "Status"] == "Paid late" and L.loc["2", "Paid Late"] == 100
    assert L.loc["3", "Status"] == "Overdue" and L.loc["3", "Days Overdue"] == (pd.Timestamp("2026-09-30") - pd.Timestamp("2026-07-16")).days


def test_partial_payment_split_on_time_and_late():
    b = bills([["1", "Micro Co", "01-07-2026", 100]])
    p = pays([["Micro Co", "10-07-2026", 40, "1"], ["Micro Co", "01-09-2026", 30, "1"]])
    r = E.build_ledger(b, p, master(), "2026-09-30").ledger.iloc[0]
    assert r["Paid On Time"] == 40 and r["Paid Late"] == 30 and r["Outstanding"] == 30 and r["Status"] == "Overdue"


def test_payments_after_as_of_ignored():
    b = bills([["1", "Micro Co", "01-07-2026", 100]])
    p = pays([["Micro Co", "10-10-2026", 100, "1"]])
    r = E.build_ledger(b, p, master(), "2026-09-30").ledger.iloc[0]
    assert r["Outstanding"] == 100


def test_onaccount_payment_applied_fifo_by_due_date():
    b = bills([["A", "Micro Co", "01-07-2026", 100], ["B", "Micro Co", "10-07-2026", 100]])
    p = pays([["Micro Co", "20-07-2026", 130, None]])
    L = E.build_ledger(b, p, master(), "2026-09-30").ledger.set_index("Bill No")
    assert L.loc["A", "Outstanding"] == 0 and L.loc["B", "Outstanding"] == 70


def test_unmatched_and_unapplied_payments_reported():
    b = bills([["A", "Micro Co", "01-07-2026", 100]])
    p = pays([["Micro Co", "20-07-2026", 50, "ZZZ"], ["Micro Co", "20-07-2026", 150, "A"]])
    res = E.build_ledger(b, p, master(), "2026-09-30")
    assert len(res.unmatched_payments) == 1
    assert res.unapplied_payments["Amount Paid"].sum() == 50   # excess over bill A


def test_unknown_vendor_unclassified():
    b = bills([["A", "Ghost Ltd", "01-07-2026", 100]])
    res = E.build_ledger(b, None, master(), "2026-09-30")
    assert res.unknown_vendors == ["Ghost Ltd"]
    assert res.ledger.loc[0, "Category"] == "Unclassified" and not res.ledger.loc[0, "Covered"]


def test_udyam_registration_date_excludes_earlier_bills():
    b = bills([["A", "Late Reg Co", "01-07-2026", 100], ["B", "Late Reg Co", "01-09-2026", 100]])
    L = E.build_ledger(b, None, master(), "2026-09-30").ledger.set_index("Bill No")
    assert not bool(L.loc["A", "Covered"]) and bool(L.loc["B", "Covered"])


def test_form1_only_covered_outstanding_beyond_45_days():
    b = bills([["A", "Micro Co", "01-07-2026", 100],     # accepted 1 Jul -> >45 days at 30 Sep, outstanding
               ["B", "Micro Co", "20-08-2026", 100],     # 41 days -> excluded
               ["C", "Medium Co", "01-07-2026", 100],    # not covered
               ["D", "Micro Co", "01-06-2026", 100]])    # paid
    p = pays([["Micro Co", "10-07-2026", 100, "D"]])
    f1, ann = E.form1(E.build_ledger(b, p, master(), "2026-09-30").ledger, "2026-09-30")
    assert list(ann["Bill No"]) == ["A"]
    assert f1.loc[0, "Supplier Name"] == "Micro Co" and f1.loc[0, "Amount Due"] == 100
    assert f1.loc[0, "Date From Which Due"] == pd.Timestamp("2026-07-16")


def test_43bh_disallowance_not_yet_due_and_reversal():
    # FY 2026-27: 1 Apr 2026 - 31 Mar 2027
    b = bills([["P", "Micro Co", "01-12-2025", 500],      # prior FY, unpaid at 31-Mar-26, due 16-Dec-25 -> disallowed in FY25
               ["X", "Micro Co", "01-01-2027", 100],      # FY26, unpaid at 31-Mar-27, due 16-Jan-27 -> disallowed
               ["Y", "Micro Co", "25-03-2027", 100],      # due 9-Apr-27 -> not yet due
               ["Z", "Micro Co", "01-01-2027", 100],      # paid on time
               ["M", "Medium Co", "01-01-2027", 999]])    # not covered
    p = pays([["Micro Co", "10-06-2026", 300, "P"], ["Micro Co", "05-01-2027", 100, "Z"]])
    d = E.disallowance_43bh(b, p, master(), 2026)
    assert list(d["disallowed"]["Bill No"]) == ["X"] and d["disallowed"]["Outstanding"].sum() == 100
    assert list(d["not_yet_due"]["Bill No"]) == ["Y"]
    # P is also disallowed in FY25-26 but 300 of it was paid in FY26-27 -> deductible now. 200 remains carried.
    assert d["reversal"]["Paid In FY"].sum() == 300
    # FY25-26 view: P disallowed 500
    d0 = E.disallowance_43bh(b, p, master(), 2025)
    assert d0["disallowed"]["Outstanding"].sum() == 500


def test_standardise_flexible_headers_and_errors():
    raw = pd.DataFrame({"Invoice No": ["1", None], "Supplier": ["A", "B"], "Invoice Date": ["01/07/2026", "x"], "Amount": ["1,000.50", "5"]})
    b, issues = E.standardise(raw, "bills")
    assert len(b) == 1 and b.loc[0, "Bill Amount"] == 1000.5 and issues
    with pytest.raises(ValueError, match="Missing required"):
        E.standardise(pd.DataFrame({"Foo": [1]}), "bills")


def test_read_table_csv_xlsx_and_multisheet():
    df = pd.DataFrame({"Bill No": ["1"], "Vendor Name": ["A"], "Bill Date": ["01-07-2026"], "Bill Amount": [10]})
    assert len(E.read_table(df.to_csv(index=False).encode(), "b.csv", "bills")) == 1
    buf = io.BytesIO()
    with pd.ExcelWriter(buf) as xw:
        pd.DataFrame({"x": [1]}).to_excel(xw, sheet_name="Other", index=False)
        df.to_excel(xw, sheet_name="Bills", index=False)
    got = E.read_table(buf.getvalue(), "all.xlsx", "bills")
    assert list(got.columns)[0] == "Bill No"


def test_half_year_end():
    assert E.half_year_end("2026-10-20") == pd.Timestamp("2026-09-30")
    assert E.half_year_end("2026-03-31") == pd.Timestamp("2026-03-31")
    assert E.half_year_end("2026-02-01") == pd.Timestamp("2025-09-30")


def test_form1_table_layout_and_buckets():
    b = bills([["A", "Micro Co", "01-07-2026", 100],    # paid in 20 days via TReDS      -> (d)(i)
               ["B", "Micro Co", "01-07-2026", 200],    # paid in 20 days other mode     -> (d)(ii)
               ["C", "Micro Co", "01-07-2026", 300],    # paid on day 60                 -> (e)
               ["D", "Micro Co", "01-09-2026", 400],    # outstanding, 29 days           -> (f)
               ["F", "Micro Co", "01-07-2026", 500],    # outstanding, 91 days           -> (g)
               ["M", "Medium Co", "01-07-2026", 999]])  # not an MSE supplier -> excluded
    raw = pd.DataFrame([["Micro Co", "21-07-2026", 100, "A", "TReDS"], ["Micro Co", "21-07-2026", 200, "B", None],
                        ["Micro Co", "30-08-2026", 300, "C", "Other"]],
                       columns=["Vendor Name", "Payment Date", "Amount Paid", "Bill No", "Payment Mode"])
    p, _ = E.standardise(raw, "payments")
    res = E.build_ledger(b, p, master(), "2026-09-30")
    t = E.form1_table(res, "2026-09-30")
    assert list(t.columns) == E.F1_COLS and len(t) == 1
    r = t.iloc[0]
    assert (r["Paid ≤45d TReDS - No."], r["Paid ≤45d TReDS - Amount (Rs.)"]) == (1, 100)
    assert (r["Paid ≤45d Other mode - No."], r["Paid ≤45d Other mode - Amount (Rs.)"]) == (1, 200)
    assert (r["Paid after 45d - No."], r["Paid after 45d - Amount (Rs.)"]) == (1, 300)
    assert (r["Outstanding ≤45d - No."], r["Outstanding ≤45d - Amount (Rs.)"]) == (1, 400)
    assert (r["Outstanding >45d - No."], r["Outstanding >45d - Amount (Rs.)"]) == (1, 500)


def test_form1_period_filter_and_workbook():
    from openpyxl import load_workbook
    b = bills([["A", "Micro Co", "01-03-2026", 100]])
    raw = pd.DataFrame([["Micro Co", "10-03-2026", 100, "A"]], columns=["Vendor Name", "Payment Date", "Amount Paid", "Bill No"])
    p, _ = E.standardise(raw, "payments")
    res = E.build_ledger(b, p, master(), "2026-09-30")
    assert len(E.form1_table(res, "2026-09-30")) == 0          # paid before period start (1 Apr) -> not in this return
    t = E.form1_table(res, "2026-09-30", "2026-03-01")
    assert t.iloc[0]["Paid ≤45d Other mode - Amount (Rs.)"] == 100
    wb = load_workbook(io.BytesIO(E.form1_workbook(t, pd.DataFrame(), "Acme", "2026-09-30", "2026-03-01")))
    ws = wb["Form 1"]
    assert wb.sheetnames[0] == "Form 1" and "(d)" in ws["D4"].value and ws["B7"].value == "Micro Co" and ws["B8"].value == "Total"


def test_payments_mixing_linked_and_onaccount_rows():
    """Regression: blank Bill No in a column that also has values must mean 'on account', not a bill called 'nan'."""
    b = bills([["A", "Micro Co", "01-07-2026", 100], ["B", "Micro Co", "10-07-2026", 100]])
    raw = pd.DataFrame([["Micro Co", "20-07-2026", 100, "A"], ["Micro Co", "25-07-2026", 30, None]],
                       columns=["Vendor Name", "Payment Date", "Amount Paid", "Bill No"])
    p, _ = E.standardise(raw, "payments")
    assert p["bkey"].isna().tolist() == [False, True]
    res = E.build_ledger(b, p, master(), "2026-09-30")
    assert len(res.unmatched_payments) == 0
    assert res.ledger.set_index("Bill No").loc["B", "Outstanding"] == 70


# ---- opening balances ----
def _op(rows):
    import pandas as pd
    raw = pd.DataFrame(rows, columns=["Vendor Name", "Outstanding Amount", "Bill No", "Bill Date", "Acceptance Date", "Description", "Branch"])
    return E.standardise(raw, "opening")[0]


def _master():
    import pandas as pd
    m = pd.DataFrame({"Vendor Name": ["Alpha Ltd"], "MSME Category": ["Micro"], "Agreed Credit Days": [30]})
    return E.standardise(m, "vendor_master")[0]


def test_opening_billwise_flows_through_ledger_and_payments():
    import pandas as pd
    op = _op([["Alpha Ltd", 100000, "OB-1", "15-02-2026", None, None, None], ["Alpha Ltd", 50000, "OB-2", "20-03-2026", None, None, None]])
    ob, _ = E.opening_to_bills(op, "bill", "2026-03-31")
    bills = E.standardise(ob, "bills")[0]
    assert list(bills["Bill Amount"]) == [100000, 50000] and (bills["Description"] == "Opening balance b/f").all()
    pays = E.standardise(pd.DataFrame({"Vendor Name": ["Alpha Ltd"], "Payment Date": ["10-04-2026"], "Amount Paid": [40000], "Bill No": ["OB-1"]}), "payments")[0]
    res = E.build_ledger(bills, pays, _master(), pd.Timestamp("2026-06-30"))
    L = res.ledger.set_index("Bill No")
    assert L.loc["OB-1", "Outstanding"] == 60000 and L.loc["OB-2", "Outstanding"] == 50000
    # OB-1: dated 15-02 + 30 days = 17-03 -> overdue at 30-06
    assert L.loc["OB-1", "Due Date"] == pd.Timestamp("2026-03-17")


def test_opening_partywise_sums_per_vendor_and_dates_from_opening_date():
    import pandas as pd
    op = _op([["Alpha Ltd", 70000, None, None, None, None, None], ["alpha ltd", 30000, None, None, None, None, "Pune"]])
    ob, issues = E.opening_to_bills(op, "party", "2026-03-31")
    assert len(ob) == 1 and ob.loc[0, "Bill Amount"] == 100000 and ob.loc[0, "Bill No"] == "Opening balance"
    assert ob.loc[0, "Bill Date"] == pd.Timestamp("2026-03-31") and any("Party-wise" in i for i in issues)


def test_opening_billwise_needs_bill_no_and_date_and_alias_amount():
    import pandas as pd
    op = E.standardise(pd.DataFrame({"Party": ["Alpha Ltd", "Alpha Ltd"], "Amount": ["1,00,000", "5000"], "Bill No": ["B1", None],
                                     "Bill Date": ["01-03-2026", "02-03-2026"]}), "opening")[0]
    ob, issues = E.opening_to_bills(op, "bill", "2026-03-31")
    assert len(ob) == 1 and ob.loc[0, "Bill Amount"] == 100000 and any("excluded" in i for i in issues)
