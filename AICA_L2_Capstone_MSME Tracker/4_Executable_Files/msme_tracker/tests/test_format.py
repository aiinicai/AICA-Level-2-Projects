import io
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import engine as E  # noqa: E402


@pytest.mark.parametrize("val,exp", [(0, "0.00"), (999.5, "999.50"), (1000, "1,000.00"), (99999.99, "99,999.99"),
                                     (100000, "1,00,000.00"), (1234567.5, "12,34,567.50"), (12345678.9, "1,23,45,678.90"),
                                     (123456789, "12,34,56,789.00"), (-1234567.891, "-12,34,567.89"), (None, ""), (float("nan"), "")])
def test_fmt_inr(val, exp):
    assert E.fmt_inr(val) == exp


def test_fmt_inr_no_decimals():
    assert E.fmt_inr(3055000, 0) == "30,55,000" and E.fmt_inr(1e7, 0) == "1,00,00,000"


def test_is_amount_col():
    df = pd.DataFrame({"Bill Amount": [1.0], "Outstanding": [1.0], "Paid Late": [1.0], "Interest Payable": [1.0], "Days Overdue": [3],
                       "Days Allowed": [15], "Bank Rate %": [5.5], "Bills": [2], "Vendor Name": ["x"], "Covered": [True],
                       "Paid ≤45d TReDS - Amount (Rs.)": [1.0], "Paid ≤45d TReDS - No.": [1]})
    got = {c for c in df.columns if E.is_amount_col(c, df[c])}
    assert got == {"Bill Amount", "Outstanding", "Paid Late", "Interest Payable", "Paid ≤45d TReDS - Amount (Rs.)"}


def _sample():
    return pd.DataFrame({"Vendor Name": ["A", "B", "C"], "Bill Date": pd.to_datetime(["2026-07-01", "2026-07-02", "2026-07-03"]),
                         "Bill Amount": [1234567.5, 100000.0, 250.0], "Outstanding": [1000000.0, 0.0, 250.0],
                         "Days Overdue": [10, 0, 5], "Status": ["Overdue", "Paid on time", "Paid late"]})


def test_export_styling_formats_and_totals():
    wb = load_workbook(io.BytesIO(E.export_workbook({"Ledger": _sample()})))
    ws = wb["Ledger"]
    assert all(c.font.bold for c in ws[1]) and ws["A1"].fill.fgColor.rgb.endswith("1F4E78")     # bold header on dark band
    assert ws["C2"].number_format == E.INR_XL and ws["D2"].number_format == E.INR_XL            # Indian amounts
    assert ws["B2"].number_format == "DD-MM-YYYY"                                               # dates
    assert ws["E2"].number_format == "General"                                                  # days stay plain
    assert ws["A5"].value == "Total" and ws["A5"].font.bold
    assert ws["C5"].value == "=SUBTOTAL(9,C2:C4)" and ws["D5"].value == "=SUBTOTAL(9,D2:D4)"
    assert ws["E5"].value is None                                                               # no total for days
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:F4"                            # filter excludes the total row
    assert ws["F2"].fill.fgColor.rgb.endswith("FFC7CE")                                         # status colour


def test_no_totals_for_master_empty_and_text_only_sheets():
    wb = load_workbook(io.BytesIO(E.export_workbook({"Vendor Master": pd.DataFrame({"Vendor Name": ["A"], "Agreed Credit Days": [30]}),
                                                     "Empty": pd.DataFrame(), "Names": pd.DataFrame({"x": ["a", "b"]}),
                                                     "NoRows": pd.DataFrame({"Outstanding": pd.Series(dtype=float)})})))
    assert wb["Vendor Master"].max_row == 2 and wb["Names"].max_row == 3 and wb["NoRows"].max_row == 1
    assert all(c.font.bold for c in wb["Vendor Master"][1])


def test_form1_workbook_indian_format_and_totals():
    t = pd.DataFrame([["1", "A", "P1"] + [1, 123456.0] * 5 + [""]], columns=E.F1_COLS)
    t["S.No."] = 1
    wb = load_workbook(io.BytesIO(E.form1_workbook(t, pd.DataFrame(), "Co", "2026-09-30", "2026-04-01")))
    ws = wb["Form 1"]
    assert ws["E7"].number_format == E.INR_XL and ws["D7"].number_format == "General"
    assert ws["E8"].value == "=SUM(E7:E7)" and ws["E8"].number_format == E.INR_XL and ws["E8"].font.bold
    assert ws["D4"].font.bold


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_totals_recalculate_to_the_right_numbers():
    d = Path(tempfile.mkdtemp())
    (d / "in.xlsx").write_bytes(E.export_workbook({"Ledger": _sample()}))
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(d / "out"),
                    str(d / "in.xlsx")], capture_output=True, timeout=120)
    ws = load_workbook(d / "out" / "in.xlsx", data_only=True)["Ledger"]
    assert ws["C5"].value == pytest.approx(1334817.5) and ws["D5"].value == pytest.approx(1000250.0)
