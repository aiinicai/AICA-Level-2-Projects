# -*- coding: utf-8 -*-
"""
Ind AS (IGAAP) vs US GAAP - Multi-Standard Comparison
------------------------------------------------------
Purpose:
    Compare an Indian company's Ind AS accounting against US GAAP across
    several standards where the two frameworks genuinely differ, driven off
    a "journal extract" style Excel upload (one sheet per standard).

    Pick one or more applicable standards on Frame 1. Each selected
    standard's sheet is read, both frameworks' treatment is computed, and
    Frame 3 shows transaction detail, side-by-side journal entries, and a
    reconciliation that starts with the profit (P&L) impact.

SCOPE AND AN HONEST LIMITATION
-------------------------------
"All standards where there is a difference" is not something a generic
journal extract can drive — differences in leases, financial instruments,
income taxes, and business combinations depend on judgement-heavy inputs
(discount rates, classification tests, fair-value hierarchies) that a flat
transaction sheet cannot supply without fabricating numbers. This tool
therefore covers five standards where the Ind AS vs US GAAP difference is
mechanical and can be computed directly from transaction-level data:

    1. Inventory                          Ind AS 2   vs  ASC 330
    2. Property, Plant & Equipment        Ind AS 16  vs  ASC 360-10   (revaluation model)
    3. Impairment of Assets               Ind AS 36  vs  ASC 360-10 / ASC 350  (reversal)
    4. Intangible Assets                  Ind AS 38  vs  ASC 730 / ASC 985-20 (development costs)
    5. Employee Benefits                  Ind AS 19  vs  ASC 715      (actuarial gains/losses)

Each standard's calculation involves simplifications, disclosed in that
standard's "Note" column and in Frame 2 — e.g. the employee benefits
amortisation is a first-year straight-line approximation, not a multi-year
corridor/runoff model. Treat these as illustrative, not audit-ready.

Excel input (a "Journal Extract" workbook, one sheet per standard):
    Use "Create Excel Template" on Frame 1 for the exact column layout for
    all five standards - it is the same layout this tool reads back in.
"""

# -------------------- AUTO INSTALL REQUIRED LIBRARIES --------------------
import sys
import subprocess
import importlib.util

REQUIRED = {
    "PyQt6": "PyQt6",
    "pandas": "pandas",
    "openpyxl": "openpyxl",
}

def ensure_packages():
    missing = []
    for module, package in REQUIRED.items():
        if importlib.util.find_spec(module) is None:
            missing.append(package)
    if missing:
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", "--disable-pip-version-check", *missing]
        )
    # Manual one-liner if auto-install is ever blocked in your environment:
    #   pip install PyQt6 pandas openpyxl

ensure_packages()

# -------------------- IMPORTS --------------------
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from PyQt6.QtCore import Qt, pyqtSignal
from PyQt6.QtGui import QFont, QCursor, QColor
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QGridLayout, QLabel, QLineEdit, QPushButton, QFileDialog,
    QMessageBox, QTableWidget, QTableWidgetItem, QHeaderView,
    QSplitter, QTabWidget, QGroupBox, QComboBox, QScrollArea, QFrame,
    QListWidget, QListWidgetItem, QStackedWidget, QStatusBar, QCheckBox
)

APP_TITLE = "Ind AS (IGAAP) vs US GAAP - Multi-Standard Comparison"

IND_AS_INVENTORY_METHODS = ["FIFO", "Weighted Average"]
US_GAAP_INVENTORY_METHODS = ["FIFO", "LIFO", "Weighted Average"]


# ==================== SHARED HELPERS ====================
def clean_number(value, default=0.0):
    if pd.isna(value) or value is None or str(value).strip() == "":
        return default
    return float(value)

def clean_date(value):
    if pd.isna(value):
        return ""
    try:
        return pd.to_datetime(value).strftime("%d-%b-%Y")
    except Exception:
        return str(value)

def to_yes_no(value):
    return str(value).strip().lower() in ("y", "yes", "true", "1")

def require_columns(df, cols, sheet_name):
    missing = [c for c in cols if c not in df.columns]
    if missing:
        raise ValueError(
            f"Sheet '{sheet_name}' is missing column(s): {', '.join(missing)}. "
            "Use 'Create Excel Template' for the exact layout."
        )

def mk_card(category, title, ind_label, ind_value, us_label, us_value,
            adj_label="Add / (Less): US GAAP vs Ind AS adjustment"):
    """category is one of 'PL' (profit), 'OCI', 'BS' (balance sheet). The
    adjustment is always computed as us_value - ind_value so every card is
    internally consistent."""
    return {
        "category": category, "title": title,
        "ind_label": ind_label, "ind_value": ind_value,
        "adj_label": adj_label, "adj_value": us_value - ind_value,
        "us_label": us_label, "us_value": us_value,
    }

def tb_adj(section, line_item, adjustment):
    """One line of a standard's contribution to the full Trial Balance /
    Financial Statement (US GAAP amount minus Ind AS amount for that TB line).
    section is one of PL_INCOME, PL_EXPENSE, BS_ASSET, BS_LIABILITY, BS_EQUITY."""
    return {"section": section, "line_item": line_item, "adjustment": adjustment}


# ==================== STANDARD 1: INVENTORY (Ind AS 2 vs ASC 330) ====================
def prepare_inventory(df):
    required = ["Date", "Transaction", "Quantity", "Unit Cost", "Unit Selling Price"]
    require_columns(df, required, "Inventory")
    work = df.copy()
    if "Item" not in work.columns:
        work["Item"] = "General Inventory"
    work["Transaction"] = work["Transaction"].astype(str).str.strip().str.lower()
    work["Quantity"] = work["Quantity"].apply(clean_number)
    work["Unit Cost"] = work["Unit Cost"].apply(clean_number)
    work["Unit Selling Price"] = work["Unit Selling Price"].apply(clean_number)
    work["Item"] = work["Item"].astype(str)
    if not work["Transaction"].isin(["purchase", "sale"]).all():
        raise ValueError("Inventory sheet: Transaction must be either 'Purchase' or 'Sale'.")
    if (work["Quantity"] <= 0).any():
        raise ValueError("Inventory sheet: Quantity must be greater than zero for every transaction.")
    return work

def run_inventory_costing(work, method):
    key = method.strip().lower().replace(" ", "_")
    fifo_lifo_layers, wavg_state = {}, {}
    rows, cogs_total, revenue_total = [], 0.0, 0.0

    def inventory_value():
        if key in ("fifo", "lifo"):
            return sum(q * c for layers in fifo_lifo_layers.values() for q, c in layers)
        return sum(c for _, c in wavg_state.values())

    for excel_row, r in work.iterrows():
        txn, item = r["Transaction"], r["Item"]
        qty, cost, sell = float(r["Quantity"]), float(r["Unit Cost"]), float(r["Unit Selling Price"])
        date = clean_date(r["Date"])
        cogs, revenue = 0.0, 0.0

        if txn == "purchase":
            if cost <= 0:
                raise ValueError(f"Inventory sheet, row {excel_row + 2}: Unit Cost is required for a Purchase.")
            if key in ("fifo", "lifo"):
                fifo_lifo_layers.setdefault(item, []).append([qty, cost])
            else:
                q0, c0 = wavg_state.get(item, (0.0, 0.0))
                wavg_state[item] = (q0 + qty, c0 + qty * cost)
        else:
            if sell <= 0:
                raise ValueError(f"Inventory sheet, row {excel_row + 2}: Unit Selling Price is required for a Sale.")
            if key == "fifo":
                layers = fifo_lifo_layers.get(item, [])
                if qty > sum(q for q, _ in layers) + 1e-9:
                    raise ValueError(f"Inventory sheet, row {excel_row + 2}: Sale exceeds available inventory for '{item}' (FIFO).")
                remaining, new_layers = qty, []
                for q, c in layers:
                    if remaining <= 1e-12:
                        new_layers.append([q, c]); continue
                    take = min(q, remaining)
                    cogs += take * c; q -= take; remaining -= take
                    if q > 1e-12:
                        new_layers.append([q, c])
                fifo_lifo_layers[item] = new_layers
            elif key == "lifo":
                layers = fifo_lifo_layers.get(item, [])
                if qty > sum(q for q, _ in layers) + 1e-9:
                    raise ValueError(f"Inventory sheet, row {excel_row + 2}: Sale exceeds available inventory for '{item}' (LIFO).")
                remaining, stack = qty, [x[:] for x in layers]
                while remaining > 1e-12:
                    q, c = stack[-1]
                    take = min(q, remaining)
                    cogs += take * c; q -= take; remaining -= take
                    if q <= 1e-12:
                        stack.pop()
                    else:
                        stack[-1][0] = q
                fifo_lifo_layers[item] = stack
            else:
                q0, c0 = wavg_state.get(item, (0.0, 0.0))
                if qty > q0 + 1e-9:
                    raise ValueError(f"Inventory sheet, row {excel_row + 2}: Sale exceeds available inventory for '{item}' (Weighted Average).")
                avg_cost = (c0 / q0) if q0 > 1e-12 else 0.0
                cogs = qty * avg_cost
                q0 -= qty; c0 -= cogs
                if q0 < 1e-9:
                    q0, c0 = 0.0, 0.0
                wavg_state[item] = (q0, c0)
            revenue = qty * sell
            revenue_total += revenue
            cogs_total += cogs

        rows.append({
            "Date": date, "Transaction": "Purchase" if txn == "purchase" else "Sale", "Item": item,
            "Quantity": qty, "Unit Cost": cost, "Unit Selling Price": sell, "Revenue": revenue,
            "COGS": cogs, "Inventory": inventory_value(),
            "Entry": (f"Dr Cost of Goods Sold {cogs:,.2f} | Cr Inventory {cogs:,.2f}" if txn == "sale" else ""),
        })
    return {"rows": rows, "cogs_total": cogs_total, "revenue_total": revenue_total, "inventory_final": inventory_value()}

def calc_inventory_standard(df, ind_method, us_method):
    work = prepare_inventory(df)
    ind, us = run_inventory_costing(work, ind_method), run_inventory_costing(work, us_method)
    ind_col, us_col = f"Ind AS ({ind_method})", f"US GAAP ({us_method})"

    trans_rows = []
    for i in range(len(ind["rows"])):
        ir, ur = ind["rows"][i], us["rows"][i]
        trans_rows.append({
            "Date": ir["Date"], "Transaction": ir["Transaction"], "Item": ir["Item"],
            "Quantity": ir["Quantity"], "Unit Cost": ir["Unit Cost"], "Unit Selling Price": ir["Unit Selling Price"],
            "Revenue": ir["Revenue"], f"{ind_col} COGS": ir["COGS"], f"{us_col} COGS": ur["COGS"],
            f"{ind_col} Inventory": ir["Inventory"], f"{us_col} Inventory": ur["Inventory"],
        })
    transactions = pd.DataFrame(trans_rows)

    def build_journal(result):
        journal_rows = []
        for r in result["rows"]:
            if r["Transaction"] == "Purchase":
                journal_rows.append([r["Date"], "Purchase", "Dr Inventory", "Cr Trade Payables / Cash", r["Quantity"] * r["Unit Cost"]])
            else:
                journal_rows.append([r["Date"], "Sale", "Dr Trade Receivable / Cash", "Cr Revenue", r["Revenue"]])
                journal_rows.append([r["Date"], "Cost of Goods Sold", "Dr Cost of Goods Sold", "Cr Inventory", r["COGS"]])
        return pd.DataFrame(journal_rows, columns=["Date", "Transaction", "Debit", "Credit", "Amount"])

    ind_journal, us_journal = build_journal(ind), build_journal(us)
    ind_gp, us_gp = ind["revenue_total"] - ind["cogs_total"], us["revenue_total"] - us["cogs_total"]

    financial_lines = pd.DataFrame([
        ["Revenue", ind["revenue_total"], us["revenue_total"], us["revenue_total"] - ind["revenue_total"]],
        ["Cost of Goods Sold", ind["cogs_total"], us["cogs_total"], us["cogs_total"] - ind["cogs_total"]],
        ["Gross Profit", ind_gp, us_gp, us_gp - ind_gp],
        ["Closing Inventory", ind["inventory_final"], us["inventory_final"], us["inventory_final"] - ind["inventory_final"]],
    ], columns=["Line Item", "Ind AS", "US GAAP", "US GAAP less Ind AS"])

    recon_cards = [
        mk_card("PL", "Profit Reconciliation — Inventory (Ind AS 2 vs ASC 330)",
                f"Gross Profit — Ind AS ({ind_method})", ind_gp, f"Gross Profit — US GAAP ({us_method})", us_gp),
        mk_card("BS", "Closing Inventory Reconciliation",
                f"Closing Inventory — Ind AS ({ind_method})", ind["inventory_final"],
                f"Closing Inventory — US GAAP ({us_method})", us["inventory_final"]),
    ]

    return {
        "id": "inventory", "label": f"Inventory ({ind_method} vs {us_method})",
        "method_caption": f"Method selected — Ind AS: {ind_method}  |  US GAAP: {us_method}",
        "transactions": transactions,
        "number_cols_transactions": ["Unit Cost", "Unit Selling Price", "Revenue",
                                      f"{ind_col} COGS", f"{us_col} COGS", f"{ind_col} Inventory", f"{us_col} Inventory"],
        "ind_journal": ind_journal, "us_journal": us_journal,
        "financial_lines": financial_lines, "recon_cards": recon_cards,
        "tb_adjustments": [
            tb_adj("PL_EXPENSE", "Cost of Goods Sold", us["cogs_total"] - ind["cogs_total"]),
            tb_adj("BS_ASSET", "Inventories", us["inventory_final"] - ind["inventory_final"]),
        ],
    }


# ==================== STANDARD 2: PPE REVALUATION (Ind AS 16 vs ASC 360-10) ====================
def calc_ppe_revaluation(df):
    cols = ["Date", "Asset", "Carrying Amount (Cost Model)", "Fair Value on Revaluation Date", "Remaining Useful Life (Years)"]
    require_columns(df, cols, "PPE Revaluation")

    rows, ind_journal, us_journal = [], [], []
    total_surplus = total_extra_dep = total_ind_ppe = total_us_ppe = 0.0

    for _, r in df.iterrows():
        date, asset = clean_date(r["Date"]), str(r["Asset"])
        carrying = clean_number(r["Carrying Amount (Cost Model)"])
        fair_value = clean_number(r["Fair Value on Revaluation Date"])
        life = clean_number(r["Remaining Useful Life (Years)"])
        surplus = fair_value - carrying
        extra_dep, note = 0.0, ""

        if surplus > 0:
            if life <= 0:
                note = "Useful life not provided - incremental depreciation not computed."
            else:
                extra_dep = surplus / life
            ind_journal.append([date, "Revaluation", "Dr Property, Plant & Equipment", "Cr Revaluation Surplus (OCI)", surplus])
            if extra_dep > 0:
                ind_journal.append([date, "Incremental Depreciation", "Dr Depreciation Expense (P&L)", "Cr Accumulated Depreciation", extra_dep])
            us_journal.append([date, "Revaluation - not recognised", "No entry (revaluation model not permitted under US GAAP)", "-", 0.0])
        else:
            note = "Downward or nil revaluation - requires a separate impairment assessment (not computed here)."

        ind_ppe_after = carrying + max(surplus, 0.0) - extra_dep
        us_ppe_after = carrying
        total_surplus += max(surplus, 0.0)
        total_extra_dep += extra_dep
        total_ind_ppe += ind_ppe_after
        total_us_ppe += us_ppe_after

        rows.append({
            "Date": date, "Asset": asset, "Carrying Amount": carrying, "Fair Value": fair_value,
            "Revaluation Surplus (Ind AS only)": max(surplus, 0.0), "Remaining Useful Life (Years)": life,
            "Incremental Depreciation p.a. (Ind AS only)": extra_dep, "Note": note,
        })

    financial_lines = pd.DataFrame([
        ["PP&E Carrying Amount (post-revaluation)", total_ind_ppe, total_us_ppe, total_us_ppe - total_ind_ppe],
        ["Revaluation Surplus (OCI / Equity)", total_surplus, 0.0, -total_surplus],
        ["Incremental Depreciation (P&L expense)", total_extra_dep, 0.0, -total_extra_dep],
    ], columns=["Line Item", "Ind AS", "US GAAP", "US GAAP less Ind AS"])

    recon_cards = [
        mk_card("PL", "Profit Reconciliation — PP&E Revaluation (Ind AS 16 vs ASC 360-10)",
                "P&L effect under Ind AS (incremental depreciation on revaluation surplus)", -total_extra_dep,
                "P&L effect under US GAAP (no revaluation, no incremental depreciation)", 0.0),
        mk_card("OCI", "OCI / Equity Reconciliation — PP&E Revaluation",
                "Revaluation Surplus recognised in OCI under Ind AS", total_surplus,
                "Revaluation Surplus under US GAAP (revaluation not permitted)", 0.0),
        mk_card("BS", "PP&E Carrying Value Reconciliation",
                "PP&E carrying amount under Ind AS", total_ind_ppe,
                "PP&E carrying amount under US GAAP", total_us_ppe),
    ]

    return {
        "id": "ppe_revaluation", "label": "Property, Plant & Equipment (Revaluation)", "method_caption": None,
        "transactions": pd.DataFrame(rows),
        "number_cols_transactions": ["Carrying Amount", "Fair Value", "Revaluation Surplus (Ind AS only)",
                                      "Incremental Depreciation p.a. (Ind AS only)"],
        "ind_journal": pd.DataFrame(ind_journal, columns=["Date", "Transaction", "Debit", "Credit", "Amount"]),
        "us_journal": pd.DataFrame(us_journal, columns=["Date", "Transaction", "Debit", "Credit", "Amount"]),
        "financial_lines": financial_lines, "recon_cards": recon_cards,
        "tb_adjustments": [
            tb_adj("PL_EXPENSE", "Depreciation and Amortisation Expense", -total_extra_dep),
            tb_adj("BS_ASSET", "Property, Plant & Equipment", total_us_ppe - total_ind_ppe),
            tb_adj("BS_EQUITY", "Other Equity (Reserves & OCI)", -total_surplus),
        ],
        "oci_line": {
            "label": "Revaluation Surplus on Property, Plant & Equipment (Ind AS 16)",
            "reclassified": False, "ind": total_surplus, "us": 0.0,
        },
    }


# ==================== STANDARD 3: IMPAIRMENT REVERSAL (Ind AS 36 vs ASC 360-10/350) ====================
def calc_impairment_reversal(df):
    cols = ["Date", "Asset", "Impairment Loss Recognised in Prior Period", "Carrying Amount Before Reversal", "Recoverable Amount Now"]
    require_columns(df, cols, "Impairment Reversal")

    rows, ind_journal, us_journal = [], [], []
    total_reversal = ind_bs = us_bs = 0.0

    for _, r in df.iterrows():
        date, asset = clean_date(r["Date"]), str(r["Asset"])
        prior_loss = clean_number(r["Impairment Loss Recognised in Prior Period"])
        carrying = clean_number(r["Carrying Amount Before Reversal"])
        recoverable = clean_number(r["Recoverable Amount Now"])
        indicated = recoverable - carrying
        reversal = max(min(indicated, prior_loss), 0.0)
        note = "" if reversal > 0 else "No reversal indicated (recoverable amount not above carrying amount)."

        if reversal > 0:
            ind_journal.append([date, "Impairment Reversal", "Dr Asset", "Cr Impairment Reversal Gain (P&L)", reversal])
            us_journal.append([date, "Impairment Reversal - not recognised", "No entry (reversal prohibited under US GAAP)", "-", 0.0])

        total_reversal += reversal
        ind_bs += carrying + reversal
        us_bs += carrying

        rows.append({
            "Date": date, "Asset": asset, "Carrying Amount Before Reversal": carrying,
            "Recoverable Amount Now": recoverable, "Impairment Reversal (Ind AS only)": reversal, "Note": note,
        })

    financial_lines = pd.DataFrame([
        ["Asset Carrying Amount (post-reversal)", ind_bs, us_bs, us_bs - ind_bs],
        ["Impairment Reversal Gain (P&L)", total_reversal, 0.0, -total_reversal],
    ], columns=["Line Item", "Ind AS", "US GAAP", "US GAAP less Ind AS"])

    recon_cards = [
        mk_card("PL", "Profit Reconciliation — Impairment Reversal (Ind AS 36 vs ASC 360-10/350)",
                "Impairment reversal gain recognised under Ind AS", total_reversal,
                "Impairment reversal under US GAAP (prohibited)", 0.0),
        mk_card("BS", "Asset Carrying Value Reconciliation — Impairment",
                "Carrying amount under Ind AS", ind_bs, "Carrying amount under US GAAP", us_bs),
    ]

    return {
        "id": "impairment_reversal", "label": "Impairment of Assets (Reversal)", "method_caption": None,
        "transactions": pd.DataFrame(rows),
        "number_cols_transactions": ["Carrying Amount Before Reversal", "Recoverable Amount Now", "Impairment Reversal (Ind AS only)"],
        "ind_journal": pd.DataFrame(ind_journal, columns=["Date", "Transaction", "Debit", "Credit", "Amount"]),
        "us_journal": pd.DataFrame(us_journal, columns=["Date", "Transaction", "Debit", "Credit", "Amount"]),
        "financial_lines": financial_lines, "recon_cards": recon_cards,
        "tb_adjustments": [
            tb_adj("PL_INCOME", "Other Income", -total_reversal),
            tb_adj("BS_ASSET", "Property, Plant & Equipment", us_bs - ind_bs),
        ],
    }


# ==================== STANDARD 4: DEVELOPMENT COSTS (Ind AS 38 vs ASC 730/985-20) ====================
def calc_development_costs(df):
    cols = ["Date", "Project", "Development Cost Incurred", "Meets Ind AS Capitalisation Criteria (Yes/No)", "Useful Life (Years)"]
    require_columns(df, cols, "Development Costs")

    rows, ind_journal, us_journal = [], [], []
    total_cost_all = total_cost_cap = total_amort = 0.0

    for _, r in df.iterrows():
        date, project = clean_date(r["Date"]), str(r["Project"])
        cost = clean_number(r["Development Cost Incurred"])
        meets = to_yes_no(r["Meets Ind AS Capitalisation Criteria (Yes/No)"])
        life = clean_number(r["Useful Life (Years)"])
        total_cost_all += cost
        us_journal.append([date, "Development Cost", "Dr Research & Development Expense (P&L)", "Cr Cash / Payables", cost])

        if meets:
            amort = cost / life if life > 0 else cost
            total_cost_cap += cost
            total_amort += amort
            ind_journal.append([date, "Development Cost - Capitalised", "Dr Intangible Asset", "Cr Cash / Payables", cost])
            ind_journal.append([date, "Amortisation", "Dr Amortisation Expense (P&L)", "Cr Accumulated Amortisation", amort])
            note = ""
            ind_expense_row = amort
        else:
            ind_journal.append([date, "Development Cost - Expensed", "Dr Research & Development Expense (P&L)", "Cr Cash / Payables", cost])
            note = "Capitalisation criteria not met - expensed under both frameworks (no GAAP difference)."
            ind_expense_row = cost

        rows.append({
            "Date": date, "Project": project, "Development Cost Incurred": cost,
            "Capitalised under Ind AS": "Yes" if meets else "No", "Useful Life (Years)": life if meets else "",
            "Ind AS P&L Expense": ind_expense_row, "US GAAP P&L Expense (Full Cost)": cost, "Note": note,
        })

    ind_expense = total_amort + (total_cost_all - total_cost_cap)
    us_expense = total_cost_all
    ind_intangible_bs = total_cost_cap - total_amort
    us_intangible_bs = 0.0

    financial_lines = pd.DataFrame([
        ["Development Cost — P&L Expense", ind_expense, us_expense, us_expense - ind_expense],
        ["Intangible Asset — Carrying Amount", ind_intangible_bs, us_intangible_bs, us_intangible_bs - ind_intangible_bs],
    ], columns=["Line Item", "Ind AS", "US GAAP", "US GAAP less Ind AS"])

    recon_cards = [
        mk_card("PL", "Profit Reconciliation — Development Costs (Ind AS 38 vs ASC 730/985-20)",
                "P&L expense under Ind AS (amortisation of capitalised cost, plus any non-qualifying cost expensed)", -ind_expense,
                "P&L expense under US GAAP (cost expensed as incurred)", -us_expense),
        mk_card("BS", "Intangible Asset Carrying Value Reconciliation",
                "Intangible asset under Ind AS", ind_intangible_bs, "Intangible asset under US GAAP", us_intangible_bs),
    ]

    return {
        "id": "development_costs", "label": "Intangible Assets (Development Costs)", "method_caption": None,
        "transactions": pd.DataFrame(rows),
        "number_cols_transactions": ["Development Cost Incurred", "Ind AS P&L Expense", "US GAAP P&L Expense (Full Cost)"],
        "ind_journal": pd.DataFrame(ind_journal, columns=["Date", "Transaction", "Debit", "Credit", "Amount"]),
        "us_journal": pd.DataFrame(us_journal, columns=["Date", "Transaction", "Debit", "Credit", "Amount"]),
        "financial_lines": financial_lines, "recon_cards": recon_cards,
        "tb_adjustments": [
            tb_adj("PL_EXPENSE", "Other Expenses", us_expense - ind_expense),
            tb_adj("BS_ASSET", "Intangible Assets", us_intangible_bs - ind_intangible_bs),
        ],
    }


# ==================== STANDARD 5: EMPLOYEE BENEFITS (Ind AS 19 vs ASC 715) ====================
def calc_employee_benefits(df):
    cols = ["Date", "Plan", "Net Actuarial Gain/(Loss) for the Period", "Average Remaining Service Period (Years)"]
    require_columns(df, cols, "Employee Benefits")

    rows, ind_journal, us_journal = [], [], []
    total_actuarial = total_amort = 0.0

    for _, r in df.iterrows():
        date, plan = clean_date(r["Date"]), str(r["Plan"])
        amt = clean_number(r["Net Actuarial Gain/(Loss) for the Period"])
        life = clean_number(r["Average Remaining Service Period (Years)"])
        amort = amt / life if life > 0 else 0.0
        total_actuarial += amt
        total_amort += amort

        ind_journal.append([date, "Remeasurement of Defined Benefit Plan", "Dr/Cr Plan Asset or Obligation", "Cr/Dr Other Comprehensive Income", amt])
        us_journal.append([date, "Remeasurement - initial", "Dr/Cr Plan Asset or Obligation", "Cr/Dr Accumulated OCI", amt])
        if abs(amort) > 1e-9:
            us_journal.append([date, "Amortisation of Actuarial Loss/(Gain)", "Dr/Cr Net Periodic Benefit Cost (P&L)", "Cr/Dr Accumulated OCI", amort])

        note = ("US GAAP amortisation shown is a simplified first-year straight-line approximation; "
                "a full corridor/runoff model needs the opening unamortised balance, which this sheet does not carry."
                if life > 0 else "Average remaining service period not provided - amortisation not computed.")
        rows.append({
            "Date": date, "Plan": plan, "Net Actuarial Gain/(Loss)": amt,
            "Average Remaining Service Period (Years)": life,
            "Ind AS - Recognised in P&L": 0.0, "US GAAP - Amortised to P&L this period": amort, "Note": note,
        })

    ind_pl, us_pl = 0.0, total_amort
    ind_oci, us_oci = total_actuarial, total_actuarial - total_amort

    financial_lines = pd.DataFrame([
        ["Actuarial Gain/(Loss) — Recognised in P&L", ind_pl, us_pl, us_pl - ind_pl],
        ["Actuarial Gain/(Loss) — Balance in OCI", ind_oci, us_oci, us_oci - ind_oci],
    ], columns=["Line Item", "Ind AS", "US GAAP", "US GAAP less Ind AS"])

    recon_cards = [
        mk_card("PL", "Profit Reconciliation — Employee Benefits (Ind AS 19 vs ASC 715)",
                "P&L effect under Ind AS (actuarial gains/losses never recycled to P&L)", ind_pl,
                "P&L effect under US GAAP (amortisation of actuarial loss/gain)", us_pl),
        mk_card("OCI", "OCI Reconciliation — Employee Benefits",
                "OCI balance under Ind AS (permanent)", ind_oci,
                "OCI balance under US GAAP (net of amount reclassified to P&L)", us_oci),
    ]

    return {
        "id": "employee_benefits", "label": "Employee Benefits (Actuarial Gains/Losses)", "method_caption": None,
        "transactions": pd.DataFrame(rows),
        "number_cols_transactions": ["Net Actuarial Gain/(Loss)", "Ind AS - Recognised in P&L", "US GAAP - Amortised to P&L this period"],
        "ind_journal": pd.DataFrame(ind_journal, columns=["Date", "Transaction", "Debit", "Credit", "Amount"]),
        "us_journal": pd.DataFrame(us_journal, columns=["Date", "Transaction", "Debit", "Credit", "Amount"]),
        "financial_lines": financial_lines, "recon_cards": recon_cards,
        "tb_adjustments": [
            tb_adj("PL_EXPENSE", "Employee Benefit Expense", -(us_pl - ind_pl)),
            tb_adj("BS_EQUITY", "Other Equity (Reserves & OCI)", us_oci - ind_oci),
        ],
        "oci_line": {
            "label": "Remeasurement of Defined Benefit Plans (Ind AS 19)",
            "reclassified": False, "ind": ind_oci, "us": us_oci,
        },
    }


# ==================== STANDARD REGISTRY ====================
STANDARDS = [
    {
        "id": "inventory", "label": "Inventory", "ind_as": "Ind AS 2", "us_gaap": "ASC 330",
        "sheet": "Inventory",
        "checklist_label": "Inventory — Ind AS 2 vs US GAAP ASC 330",
        "has_method_choice": True,
        "template_columns": ["Date", "Transaction", "Item", "Quantity", "Unit Cost", "Unit Selling Price"],
        "template_rows": [
            ["01-04-2026", "Purchase", "A100", 100, 100, 0],
            ["05-04-2026", "Purchase", "A100", 100, 120, 0],
            ["10-04-2026", "Sale", "A100", 120, 0, 200],
            ["15-04-2026", "Purchase", "A100", 50, 140, 0],
            ["20-04-2026", "Sale", "A100", 25, 0, 205],
        ],
        "ind_desc": ("Ind AS 2 permits FIFO or weighted-average; LIFO is not permitted. The cost formula "
                     "chosen determines both cost of goods sold and closing inventory value."),
        "us_desc": ("ASC 330 permits FIFO, LIFO, weighted-average or specific identification. LIFO shifts "
                     "the most recent (often higher) costs into COGS, understating inventory relative to "
                     "current cost during periods of rising prices."),
        "ind_basis": "{ind_method} closing stock", "us_basis": "{us_method} closing stock",
    },
    {
        "id": "ppe_revaluation", "label": "Property, Plant & Equipment (Revaluation)",
        "ind_as": "Ind AS 16", "us_gaap": "ASC 360-10", "sheet": "PPE Revaluation",
        "checklist_label": "Property, Plant & Equipment (Revaluation) — Ind AS 16 vs US GAAP ASC 360-10",
        "has_method_choice": False,
        "template_columns": ["Date", "Asset", "Carrying Amount (Cost Model)", "Fair Value on Revaluation Date", "Remaining Useful Life (Years)"],
        "template_rows": [["01-04-2026", "Plant A", 500000, 650000, 10]],
        "ind_desc": ("Ind AS 16 permits the revaluation model as an accounting policy choice: PP&E can be "
                     "carried at fair value, with the upward surplus taken to OCI/Revaluation Reserve. "
                     "The revalued carrying amount is then depreciated over remaining useful life, so "
                     "future depreciation is higher than under the cost model."),
        "us_desc": ("US GAAP does not permit upward revaluation of PP&E; assets are carried at "
                     "depreciated historical cost only. There is therefore no revaluation surplus and no "
                     "related incremental depreciation."),
        "ind_basis": "Revaluation model", "us_basis": "Cost model (no revaluation)",
    },
    {
        "id": "impairment_reversal", "label": "Impairment of Assets (Reversal)",
        "ind_as": "Ind AS 36", "us_gaap": "ASC 360-10 / ASC 350", "sheet": "Impairment Reversal",
        "checklist_label": "Impairment of Assets (Reversal) — Ind AS 36 vs US GAAP ASC 360-10/350",
        "has_method_choice": False,
        "template_columns": ["Date", "Asset", "Impairment Loss Recognised in Prior Period", "Carrying Amount Before Reversal", "Recoverable Amount Now"],
        "template_rows": [["01-04-2026", "Machine B", 200000, 800000, 950000]],
        "ind_desc": ("Ind AS 36 requires a previously recognised impairment loss (other than on goodwill) "
                     "to be reversed through profit or loss if the recoverable amount subsequently "
                     "increases, capped at the carrying amount that would have applied had no impairment "
                     "been recognised."),
        "us_desc": ("US GAAP prohibits reversal of an impairment loss once recognised, for both "
                     "long-lived assets held for use and goodwill. The asset stays at its written-down "
                     "carrying amount even if its recoverable amount later recovers."),
        "ind_basis": "Reversal recognised", "us_basis": "Reversal prohibited",
    },
    {
        "id": "development_costs", "label": "Intangible Assets (Development Costs)",
        "ind_as": "Ind AS 38", "us_gaap": "ASC 730 / ASC 985-20", "sheet": "Development Costs",
        "checklist_label": "Intangible Assets (Development Costs) — Ind AS 38 vs US GAAP ASC 730/985-20",
        "has_method_choice": False,
        "template_columns": ["Date", "Project", "Development Cost Incurred", "Meets Ind AS Capitalisation Criteria (Yes/No)", "Useful Life (Years)"],
        "template_rows": [["01-04-2026", "Project Zeta", 1000000, "Yes", 5]],
        "ind_desc": ("Ind AS 38 requires development expenditure to be capitalised as an intangible asset "
                     "once technical feasibility and the other specified recognition criteria are met, "
                     "then amortised over its useful life."),
        "us_desc": ("ASC 730 generally requires research and development costs to be expensed as "
                     "incurred (software development costs after technological feasibility under "
                     "ASC 985-20 being a narrow exception not modelled here), so the full cost hits "
                     "profit immediately rather than being amortised."),
        "ind_basis": "Costs capitalised", "us_basis": "Costs expensed as incurred",
    },
    {
        "id": "employee_benefits", "label": "Employee Benefits (Actuarial Gains/Losses)",
        "ind_as": "Ind AS 19", "us_gaap": "ASC 715", "sheet": "Employee Benefits",
        "checklist_label": "Employee Benefits (Actuarial Gains/Losses) — Ind AS 19 vs US GAAP ASC 715",
        "has_method_choice": False,
        "template_columns": ["Date", "Plan", "Net Actuarial Gain/(Loss) for the Period", "Average Remaining Service Period (Years)"],
        "template_rows": [["31-03-2027", "Gratuity Plan", -150000, 8]],
        "ind_desc": ("Ind AS 19 (revised) requires actuarial gains and losses (remeasurements) on defined "
                     "benefit plans to be recognised in OCI and never reclassified to profit or loss in "
                     "later periods."),
        "us_desc": ("ASC 715 also recognises actuarial gains/losses in OCI initially, but then amortises "
                     "them into net periodic benefit cost (P&L) over the average remaining service period "
                     "of active employees, effectively recycling part of the OCI balance into profit over time."),
        "ind_basis": "Remeasurement in OCI (no recycling)", "us_basis": "Remeasurement amortised to P&L",
    },
]
STANDARDS_BY_ID = {s["id"]: s for s in STANDARDS}

CALC_FUNCS = {
    "ppe_revaluation": calc_ppe_revaluation,
    "impairment_reversal": calc_impairment_reversal,
    "development_costs": calc_development_costs,
    "employee_benefits": calc_employee_benefits,
}

# ==================== TRIAL BALANCE -> FULL FINANCIAL STATEMENTS ====================
# The Trial Balance is the company's Ind AS books (P&L + Balance Sheet line
# items, as recorded). US GAAP financials are derived by applying each
# selected standard's tb_adjustments on top of it. Every P&L adjustment also
# feeds Retained Earnings on the Balance Sheet (ordinary double-entry), so if
# the Trial Balance itself balances under Ind AS, the derived US GAAP
# statement balances too — this is checked and shown, not assumed.
TB_REQUIRED_COLUMNS = ["Line Item", "Statement", "Type", "Amount"]

TB_TEMPLATE_ROWS = [
    ["Revenue from Operations", "P&L", "Income", 5000000],
    ["Other Income", "P&L", "Income", 50000],
    ["Cost of Goods Sold", "P&L", "Expense", 2500000],
    ["Employee Benefit Expense", "P&L", "Expense", 800000],
    ["Depreciation and Amortisation Expense", "P&L", "Expense", 300000],
    ["Finance Costs", "P&L", "Expense", 150000],
    ["Other Expenses", "P&L", "Expense", 400000],
    ["Tax Expense", "P&L", "Expense", 225000],
    ["Property, Plant & Equipment", "Balance Sheet", "Asset", 3000000],
    ["Intangible Assets", "Balance Sheet", "Asset", 200000],
    ["Inventories", "Balance Sheet", "Asset", 900000],
    ["Trade Receivables", "Balance Sheet", "Asset", 700000],
    ["Cash & Cash Equivalents", "Balance Sheet", "Asset", 75000],
    ["Other Assets", "Balance Sheet", "Asset", 100000],
    ["Equity Share Capital", "Balance Sheet", "Equity", 1000000],
    ["Retained Earnings", "Balance Sheet", "Equity", 1475000],
    ["Other Equity (Reserves & OCI)", "Balance Sheet", "Equity", 500000],
    ["Borrowings", "Balance Sheet", "Liability", 1000000],
    ["Trade Payables", "Balance Sheet", "Liability", 400000],
    ["Other Liabilities", "Balance Sheet", "Liability", 600000],
]

# Fixed classification for every line item this tool ever produces (the TB
# template's own lines, plus every line item the five standards' adjustments
# can add) so both frameworks can be laid out in their real, native format
# rather than one generic shared table. Tuple = (pl_bucket or bs_bucket, current/non-current).
LINE_ITEM_CLASSIFICATION = {
    "revenue from operations": ("income", "operating"),
    "other income": ("income", "other"),
    "cost of goods sold": ("expense", "cogs"),
    "employee benefit expense": ("expense", "operating"),
    "depreciation and amortisation expense": ("expense", "operating"),
    "finance costs": ("expense", "finance"),
    "other expenses": ("expense", "operating"),
    "tax expense": ("expense", "tax"),
    "property, plant & equipment": ("asset", "non_current"),
    "intangible assets": ("asset", "non_current"),
    "inventories": ("asset", "current"),
    "trade receivables": ("asset", "current"),
    "cash & cash equivalents": ("asset", "current"),
    "other assets": ("asset", "current"),
    "equity share capital": ("equity", "share_capital"),
    "retained earnings": ("equity", "retained_earnings"),
    "other equity (reserves & oci)": ("equity", "reserves_oci"),
    "borrowings": ("liability", "non_current"),
    "trade payables": ("liability", "current"),
    "other liabilities": ("liability", "current"),
}

def classify_line(line_item, fallback_section):
    """fallback_section is one of PL_INCOME/PL_EXPENSE/BS_ASSET/BS_LIABILITY/BS_EQUITY,
    used only for a line item this tool doesn't already know by name (e.g. a
    custom line added straight into the Trial Balance)."""
    bucket = LINE_ITEM_CLASSIFICATION.get(line_item.strip().lower())
    if bucket:
        return bucket
    defaults = {
        "PL_INCOME": ("income", "other"), "PL_EXPENSE": ("expense", "operating"),
        "BS_ASSET": ("asset", "current"), "BS_LIABILITY": ("liability", "current"),
        "BS_EQUITY": ("equity", "reserves_oci"),
    }
    return defaults[fallback_section]

def tb_section_code(statement, type_):
    statement, type_ = str(statement).strip().lower(), str(type_).strip().lower()
    if statement.startswith("p"):
        return "PL_INCOME" if type_.startswith("inc") else "PL_EXPENSE"
    if type_.startswith("asset"):
        return "BS_ASSET"
    if type_.startswith("liab"):
        return "BS_LIABILITY"
    return "BS_EQUITY"

SECTION_LABELS = {
    "PL_INCOME": "Income", "PL_EXPENSE": "Expense",
    "BS_ASSET": "Assets", "BS_LIABILITY": "Liabilities", "BS_EQUITY": "Equity",
}

# Cosmetic label swap for the US GAAP statements only — same underlying figure,
# the terminology real US GAAP filings use.
US_GAAP_LABEL_MAP = {
    "revenue from operations": "Net Revenue",
    "equity share capital": "Common Stock",
    "trade receivables": "Accounts Receivable",
    "trade payables": "Accounts Payable",
    "borrowings": "Long-Term Debt",
    "other equity (reserves & oci)": "Accumulated Other Comprehensive Income",
    "finance costs": "Interest Expense",
}

def us_label(name):
    return US_GAAP_LABEL_MAP.get(name.strip().lower(), name)

def gather(lines_by_key, section_code, subclass, side):
    """Every line item (Trial Balance or standard-added) in this section and
    sub-classification bucket, as (label, amount), in first-seen order."""
    out = []
    for (sec, _), v in lines_by_key.items():
        if sec != section_code:
            continue
        _, cls_sub = classify_line(v["line_item"], sec)
        if cls_sub != subclass:
            continue
        amt = v["ind"] if side == "ind" else v["ind"] + v["us_extra"]
        out.append((v["line_item"], amt))
    return out

def _row(label, amount=None, indent=0, bold=False):
    return {"Line Item": ("    " * indent) + label, "Amount": amount, "_bold": bold}

def build_ind_as_pl(lines_by_key, side):
    """Statement of Profit and Loss, Ind AS / Schedule III style."""
    revenue = gather(lines_by_key, "PL_INCOME", "operating", side)
    other_income = gather(lines_by_key, "PL_INCOME", "other", side)
    total_income = sum(a for _, a in revenue + other_income)

    cogs = gather(lines_by_key, "PL_EXPENSE", "cogs", side)
    operating_exp = gather(lines_by_key, "PL_EXPENSE", "operating", side)
    finance_exp = gather(lines_by_key, "PL_EXPENSE", "finance", side)
    tax_exp = gather(lines_by_key, "PL_EXPENSE", "tax", side)
    total_expenses = sum(a for _, a in cogs + operating_exp + finance_exp)
    total_tax = sum(a for _, a in tax_exp)

    pbt = total_income - total_expenses
    pat = pbt - total_tax

    rows = [_row("I. Revenue from Operations", indent=0, bold=True)]
    for n, a in revenue:
        rows.append(_row(n, a, 1))
    rows.append(_row("II. Other Income", bold=True))
    for n, a in other_income:
        rows.append(_row(n, a, 1))
    rows.append(_row("III. Total Income (I + II)", total_income, bold=True))
    rows.append(_row("IV. Expenses", bold=True))
    for n, a in cogs + operating_exp + finance_exp:
        rows.append(_row(n, a, 1))
    rows.append(_row("Total Expenses", total_expenses, bold=True))
    rows.append(_row("V. Profit Before Tax (III − IV)", pbt, bold=True))
    rows.append(_row("VI. Tax Expense", bold=True))
    for n, a in tax_exp:
        rows.append(_row(n, a, 1))
    rows.append(_row("VII. Profit for the Period", pat, bold=True))
    return pd.DataFrame(rows)

def build_ind_as_bs(lines_by_key, side):
    """Balance Sheet, Ind AS / Schedule III style (Equity & Liabilities, then Assets)."""
    eq_share = gather(lines_by_key, "BS_EQUITY", "share_capital", side)
    retained = gather(lines_by_key, "BS_EQUITY", "retained_earnings", side)
    reserves = gather(lines_by_key, "BS_EQUITY", "reserves_oci", side)
    total_equity = sum(a for _, a in eq_share + retained + reserves)

    liab_nc = gather(lines_by_key, "BS_LIABILITY", "non_current", side)
    liab_c = gather(lines_by_key, "BS_LIABILITY", "current", side)
    total_liab = sum(a for _, a in liab_nc + liab_c)

    asset_nc = gather(lines_by_key, "BS_ASSET", "non_current", side)
    asset_c = gather(lines_by_key, "BS_ASSET", "current", side)
    total_assets = sum(a for _, a in asset_nc + asset_c)

    rows = [_row("I. EQUITY AND LIABILITIES", bold=True)]
    rows.append(_row("Equity", indent=1, bold=True))
    for n, a in eq_share + retained + reserves:
        rows.append(_row(n, a, 2))
    rows.append(_row("Total Equity", total_equity, 1, bold=True))
    rows.append(_row("Liabilities", indent=1, bold=True))
    rows.append(_row("Non-current liabilities", indent=2, bold=True))
    for n, a in liab_nc:
        rows.append(_row(n, a, 3))
    rows.append(_row("Current liabilities", indent=2, bold=True))
    for n, a in liab_c:
        rows.append(_row(n, a, 3))
    rows.append(_row("Total Liabilities", total_liab, 1, bold=True))
    rows.append(_row("TOTAL EQUITY AND LIABILITIES", total_equity + total_liab, bold=True))
    rows.append(_row("II. ASSETS", bold=True))
    rows.append(_row("Non-current assets", indent=1, bold=True))
    for n, a in asset_nc:
        rows.append(_row(n, a, 2))
    rows.append(_row("Current assets", indent=1, bold=True))
    for n, a in asset_c:
        rows.append(_row(n, a, 2))
    rows.append(_row("TOTAL ASSETS", total_assets, bold=True))
    return pd.DataFrame(rows)

def build_us_gaap_pl(lines_by_key, side):
    """Multi-step Income Statement, US GAAP style."""
    revenue = gather(lines_by_key, "PL_INCOME", "operating", side)
    cogs = gather(lines_by_key, "PL_EXPENSE", "cogs", side)
    total_revenue = sum(a for _, a in revenue)
    total_cogs = sum(a for _, a in cogs)
    gross_profit = total_revenue - total_cogs

    operating_exp = gather(lines_by_key, "PL_EXPENSE", "operating", side)
    total_opex = sum(a for _, a in operating_exp)
    operating_income = gross_profit - total_opex

    other_income = gather(lines_by_key, "PL_INCOME", "other", side)
    finance_exp = gather(lines_by_key, "PL_EXPENSE", "finance", side)
    total_other = sum(a for _, a in other_income) - sum(a for _, a in finance_exp)
    income_before_tax = operating_income + total_other

    tax_exp = gather(lines_by_key, "PL_EXPENSE", "tax", side)
    total_tax = sum(a for _, a in tax_exp)
    net_income = income_before_tax - total_tax

    rows = []
    for n, a in revenue:
        rows.append(_row(us_label(n), a, 0, bold=True))
    for n, a in cogs:
        rows.append(_row(us_label(n), a, 1))
    rows.append(_row("Gross Profit", gross_profit, bold=True))
    rows.append(_row("Operating Expenses", bold=True))
    for n, a in operating_exp:
        rows.append(_row(us_label(n), a, 1))
    rows.append(_row("Total Operating Expenses", total_opex, 1, bold=True))
    rows.append(_row("Operating Income", operating_income, bold=True))
    rows.append(_row("Other Income (Expense)", bold=True))
    for n, a in other_income:
        rows.append(_row(us_label(n), a, 1))
    for n, a in finance_exp:
        rows.append(_row(us_label(n), -a, 1))
    rows.append(_row("Income Before Income Taxes", income_before_tax, bold=True))
    rows.append(_row("Income Tax Expense", bold=True))
    for n, a in tax_exp:
        rows.append(_row(us_label(n), a, 1))
    rows.append(_row("Net Income", net_income, bold=True))
    return pd.DataFrame(rows)

def append_oci_section(pl_df, oci_lines, side, net_profit, bottom_line_label="Net Income"):
    """Ind AS 1 / ASC 220 both require the P&L (or an immediately adjoining
    statement) to continue past the profit line into Other Comprehensive
    Income and a Total Comprehensive Income figure. oci_lines is the list of
    {"label","reclassified","ind","us"} dicts collected from whichever
    standards actually touch OCI (currently PPE Revaluation and Employee
    Benefits) — if none are selected, no OCI section is added."""
    if not oci_lines:
        return pl_df
    not_reclassified = [o for o in oci_lines if not o.get("reclassified")]
    reclassified = [o for o in oci_lines if o.get("reclassified")]
    rows = [_row("Other Comprehensive Income", bold=True)]
    total_oci = 0.0
    if not_reclassified:
        rows.append(_row("Items that will not be reclassified to profit or loss", indent=1, bold=True))
        for o in not_reclassified:
            v = o["ind"] if side == "ind" else o["us"]
            total_oci += v
            rows.append(_row(o["label"], v, 2))
    if reclassified:
        rows.append(_row("Items that will be reclassified to profit or loss", indent=1, bold=True))
        for o in reclassified:
            v = o["ind"] if side == "ind" else o["us"]
            total_oci += v
            rows.append(_row(o["label"], v, 2))
    rows.append(_row("Total Other Comprehensive Income", total_oci, bold=True))
    rows.append(_row(f"Total Comprehensive Income for the Period ({bottom_line_label} + OCI)", net_profit + total_oci, bold=True))
    return pd.concat([pl_df, pd.DataFrame(rows)], ignore_index=True)

def build_us_gaap_bs(lines_by_key, side):
    """Classified Balance Sheet, US GAAP style (Assets, then Liabilities & Stockholders' Equity)."""
    asset_c = gather(lines_by_key, "BS_ASSET", "current", side)
    asset_nc = gather(lines_by_key, "BS_ASSET", "non_current", side)
    total_ca = sum(a for _, a in asset_c)
    total_nca = sum(a for _, a in asset_nc)

    liab_c = gather(lines_by_key, "BS_LIABILITY", "current", side)
    liab_nc = gather(lines_by_key, "BS_LIABILITY", "non_current", side)
    total_cl = sum(a for _, a in liab_c)
    total_ncl = sum(a for _, a in liab_nc)

    eq_share = gather(lines_by_key, "BS_EQUITY", "share_capital", side)
    retained = gather(lines_by_key, "BS_EQUITY", "retained_earnings", side)
    reserves = gather(lines_by_key, "BS_EQUITY", "reserves_oci", side)
    total_equity = sum(a for _, a in eq_share + retained + reserves)

    rows = [_row("ASSETS", bold=True)]
    rows.append(_row("Current Assets", indent=1, bold=True))
    for n, a in asset_c:
        rows.append(_row(us_label(n), a, 2))
    rows.append(_row("Total Current Assets", total_ca, 1, bold=True))
    rows.append(_row("Non-Current Assets", indent=1, bold=True))
    for n, a in asset_nc:
        rows.append(_row(us_label(n), a, 2))
    rows.append(_row("Total Non-Current Assets", total_nca, 1, bold=True))
    rows.append(_row("TOTAL ASSETS", total_ca + total_nca, bold=True))

    rows.append(_row("LIABILITIES", bold=True))
    rows.append(_row("Current Liabilities", indent=1, bold=True))
    for n, a in liab_c:
        rows.append(_row(us_label(n), a, 2))
    rows.append(_row("Total Current Liabilities", total_cl, 1, bold=True))
    rows.append(_row("Non-Current Liabilities", indent=1, bold=True))
    for n, a in liab_nc:
        rows.append(_row(us_label(n), a, 2))
    rows.append(_row("Total Non-Current Liabilities", total_ncl, 1, bold=True))
    rows.append(_row("TOTAL LIABILITIES", total_cl + total_ncl, bold=True))

    rows.append(_row("STOCKHOLDERS' EQUITY", bold=True))
    for n, a in eq_share:
        rows.append(_row(us_label(n), a, 1))
    for n, a in retained:
        rows.append(_row(us_label(n), a, 1))
    for n, a in reserves:
        rows.append(_row(us_label(n), a, 1))
    rows.append(_row("TOTAL STOCKHOLDERS' EQUITY", total_equity, bold=True))
    rows.append(_row("TOTAL LIABILITIES AND STOCKHOLDERS' EQUITY", total_cl + total_ncl + total_equity, bold=True))
    return pd.DataFrame(rows)

def build_financial_statements(tb_df, results):
    """tb_df: the Trial Balance sheet (Ind AS books). results: dict of
    standard_id -> calc result (only standards actually computed). Returns
    a dict of P&L and Balance Sheet lines, Ind AS and US GAAP side by side,
    plus a balance check for each framework."""
    require_columns(tb_df, TB_REQUIRED_COLUMNS, "Trial Balance")

    lines = {}  # (section, line_item_lower) -> {"line_item", "section", "ind": float, "us_extra": float, "from_tb": bool}
    for _, r in tb_df.iterrows():
        line_item = str(r["Line Item"]).strip()
        if not line_item or line_item.lower() == "nan":
            continue
        section = tb_section_code(r["Statement"], r["Type"])
        amount = clean_number(r["Amount"])
        key = (section, line_item.lower())
        lines[key] = {"line_item": line_item, "section": section, "ind": amount, "us_extra": 0.0, "from_tb": True}

    pl_adjustment_total = 0.0
    for sid, res in results.items():
        for adj in res.get("tb_adjustments", []):
            key = (adj["section"], adj["line_item"].lower())
            if key not in lines:
                lines[key] = {"line_item": adj["line_item"], "section": adj["section"], "ind": 0.0, "us_extra": 0.0, "from_tb": False}
            lines[key]["us_extra"] += adj["adjustment"]
        # Every standard's own "PL" reconciliation card is its net profit impact —
        # reuse it directly so the Retained Earnings plug always matches the P&L.
        for c in res.get("recon_cards", []):
            if c["category"] == "PL":
                pl_adjustment_total += c["adj_value"]

    # Retained Earnings picks up the cumulative profit adjustment (ordinary
    # double-entry: a P&L difference flows to equity via Retained Earnings).
    re_key = None
    for key, v in lines.items():
        if v["section"] == "BS_EQUITY" and v["line_item"].strip().lower() == "retained earnings":
            re_key = key
            break
    if re_key is None:
        re_key = ("BS_EQUITY", "retained earnings")
        lines[re_key] = {"line_item": "Retained Earnings", "section": "BS_EQUITY", "ind": 0.0, "us_extra": 0.0, "from_tb": False}
    lines[re_key]["us_extra"] += pl_adjustment_total

    def section_total(code, side):
        return sum((v["ind"] if side == "ind" else v["ind"] + v["us_extra"])
                   for v in lines.values() if v["section"] == code)

    net_profit_ind = section_total("PL_INCOME", "ind") - section_total("PL_EXPENSE", "ind")
    net_profit_us = section_total("PL_INCOME", "us") - section_total("PL_EXPENSE", "us")
    asset_ind, asset_us = section_total("BS_ASSET", "ind"), section_total("BS_ASSET", "us")
    liab_ind, liab_us = section_total("BS_LIABILITY", "ind"), section_total("BS_LIABILITY", "us")
    equity_ind, equity_us = section_total("BS_EQUITY", "ind"), section_total("BS_EQUITY", "us")

    oci_lines = [res["oci_line"] for res in results.values() if "oci_line" in res]
    ind_pl = append_oci_section(build_ind_as_pl(lines, "ind"), oci_lines, "ind", net_profit_ind, "Profit for the Period")
    us_pl = append_oci_section(build_us_gaap_pl(lines, "us"), oci_lines, "us", net_profit_us, "Net Income")

    return {
        "ind_as_pl": ind_pl, "ind_as_bs": build_ind_as_bs(lines, "ind"),
        "us_gaap_pl": us_pl, "us_gaap_bs": build_us_gaap_bs(lines, "us"),
        "balance_check_ind": asset_ind - (liab_ind + equity_ind),
        "balance_check_us": asset_us - (liab_us + equity_us),
        "net_profit_ind": net_profit_ind, "net_profit_us": net_profit_us,
    }


# -------------------- EXCEL EXPORT --------------------
def style_workbook(path):
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in ws.columns:
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 12), 45)
    wb.save(path)

def sheet_name(text, suffix=""):
    name = f"{text}{suffix}"
    return name[:31] if len(name) > 31 else name

def export_single_sheet(df, title, default_name):
    start_path = str(Path.home() / default_name)
    path, _ = QFileDialog.getSaveFileName(
        None, f"Export {title}", start_path, "Excel Workbook (*.xlsx)",
        options=QFileDialog.Option.DontUseNativeDialog
    )
    if not path:
        return
    try:
        df.to_excel(path, index=False)
        style_workbook(path)
        QMessageBox.information(None, "Export Complete", f"{title} exported successfully.")
    except Exception as e:
        QMessageBox.critical(None, "Export Error", str(e))


# -------------------- DESIGN TOKENS (single source of truth for styling) --------------------
class Tokens:
    BG = "#0b0f15"
    SURFACE = "#0f141b"
    SURFACE_RAISED = "#131a23"
    BORDER = "#263140"
    DIVIDER = "#1c2530"
    TEXT_PRIMARY = "#e6ecf3"
    TEXT_SECONDARY = "#b6c3d2"
    TEXT_MUTED = "#97a6b8"
    TEXT_DISABLED = "#5f6e80"
    IND_ACCENT = "#38976f"
    IND_ACCENT_LIGHT = "#7ecaa5"
    US_ACCENT = "#5b84c4"
    US_ACCENT_LIGHT = "#9bb8e0"
    WARNING = "#d9a441"
    FONT_SANS = "'Segoe UI', 'SF Pro Display', Arial, sans-serif"
    FONT_MONO = "Consolas, 'Courier New', monospace"
    SIZE_TITLE = 25
    SIZE_SECTION = 11
    SIZE_BODY = 14
    SIZE_CAPTION = 12
    SP = {n: n for n in (4, 8, 12, 16, 20, 24)}  # spacing multiples of 4 — use Tokens.SP[n]
    RADIUS_CONTROL = 6
    RADIUS_CONTAINER = 8
    RADIUS_PILL = 999

def amp(text):
    """PyQt/PySide treats a bare '&' as a mnemonic accelerator and swallows it.
    Double it so it renders literally, in any widget text (buttons, labels,
    group titles, tab labels)."""
    return text.replace("&", "&&") if isinstance(text, str) else text

def money_inr(v):
    """Indian digit grouping, ₹ prefix, whole rupees (no decimals), negatives
    in parentheses — used only in the Frame 2 comparison rows and the Frame 3
    reconciliation table. Data tables elsewhere keep the existing money()
    formatting with decimals."""
    neg = v < 0
    integer_part = f"{round(abs(v)):,d}".replace(",", "")
    if len(integer_part) <= 3:
        grouped = integer_part
    else:
        last3, rest, parts = integer_part[-3:], integer_part[:-3], []
        while len(rest) > 2:
            parts.insert(0, rest[-2:])
            rest = rest[:-2]
        if rest:
            parts.insert(0, rest)
        grouped = ",".join(parts) + "," + last3
    text = f"₹{grouped}"
    return f"({text})" if neg else text

def clamped_text(text, limit=150):
    """~2 lines of prose with a trailing ellipsis when collapsed."""
    if len(text) <= limit:
        return text
    return text[:limit].rsplit(" ", 1)[0] + "…"

def profit_impact_phrase(adj_value):
    """adj_value = US GAAP amount − Ind AS amount for a standard's profit-impact
    card. Returns (phrase, side) where side is 'ind', 'us', or 'none'."""
    if abs(adj_value) < 0.005:
        return "No difference", "none"
    if adj_value > 0:
        return f"{money_inr(adj_value)} higher under US GAAP", "us"
    return f"{money_inr(abs(adj_value))} higher under Ind AS", "ind"

# -------------------- UI HELPERS --------------------
def money(v):
    return f"{round(v):,d}"

def set_table(table, df, number_cols=None):
    bold_col = "_bold"
    has_bold = bold_col in df.columns
    display_cols = [c for c in df.columns if c != bold_col]
    table.clear()
    table.setRowCount(len(df))
    table.setColumnCount(len(display_cols))
    table.setHorizontalHeaderLabels([str(c) for c in display_cols])
    number_cols = set(number_cols or [])
    bold_font = QFont()
    bold_font.setBold(True)
    for i, (_, row) in enumerate(df.iterrows()):
        is_bold = has_bold and bool(row[bold_col])
        for j, col in enumerate(display_cols):
            value = row[col]
            is_number = False
            if value is None or (isinstance(value, float) and pd.isna(value)):
                text = ""
            elif col in number_cols and isinstance(value, (int, float)):
                text = money(float(value))
                is_number = True
            elif isinstance(value, float):
                text = f"{value:,.4f}".rstrip("0").rstrip(".")
                is_number = True
            else:
                text = str(value)
            item = QTableWidgetItem(text)
            if is_number:
                item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            if is_bold:
                item.setFont(bold_font)
            table.setItem(i, j, item)
    table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
    table.horizontalHeader().setStretchLastSection(True)
    table.verticalHeader().setVisible(False)

def clear_layout(layout):
    while layout.count():
        item = layout.takeAt(0)
        w = item.widget()
        if w is not None:
            w.deleteLater()
        elif item.layout() is not None:
            clear_layout(item.layout())

def build_recon_card(card, emphasise_title=False):
    box = QGroupBox(card["title"])
    v = QVBoxLayout(box)
    v.setSpacing(10)

    def row(label_text, value_text, emphasise=False):
        lab = QLabel(label_text)
        lab.setWordWrap(True)
        val = QLabel(value_text)
        val.setAlignment(Qt.AlignmentFlag.AlignRight)
        val.setStyleSheet(f"font-weight:{'800' if emphasise else '600'}; color:#F8FAFC; font-size:14px;")
        lab.setStyleSheet("color:#94A3B8; font-size:12px;")
        cell = QVBoxLayout()
        cell.addWidget(lab)
        cell.addWidget(val)
        return cell

    v.addLayout(row(card["ind_label"], money(card["ind_value"])))
    for _ in range(1):
        sep = QFrame(); sep.setFrameShape(QFrame.Shape.HLine); sep.setStyleSheet("color:#263244;")
        v.addWidget(sep)
    v.addLayout(row(card["adj_label"], money(card["adj_value"])))
    sep2 = QFrame(); sep2.setFrameShape(QFrame.Shape.HLine); sep2.setStyleSheet("color:#263244;")
    v.addWidget(sep2)
    v.addLayout(row(card["us_label"], money(card["us_value"]), emphasise=True))
    box.setMinimumWidth(280)
    if emphasise_title:
        box.setStyleSheet("QGroupBox { border: 1px solid #2563EB; } QGroupBox::title { color: #93C5FD; font-weight: 800; }")
    return box

def build_entries_split(ind_df, us_df):
    split = QSplitter(Qt.Orientation.Horizontal)
    ind_box = QGroupBox("Ind AS (IGAAP) — Journal Entries")
    iv = QVBoxLayout(ind_box)
    ind_table = QTableWidget()
    set_table(ind_table, ind_df, ["Amount"])
    iv.addWidget(ind_table)

    us_box = QGroupBox("US GAAP — Journal Entries")
    uv = QVBoxLayout(us_box)
    us_table = QTableWidget()
    set_table(us_table, us_df, ["Amount"])
    uv.addWidget(us_table)

    split.addWidget(ind_box)
    split.addWidget(us_box)
    return split


# -------------------- STEPPER (replaces the plain QTabWidget as page navigation) --------------------
class StepItem(QWidget):
    clicked = pyqtSignal()

    def __init__(self, number, label):
        super().__init__()
        self.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.setSpacing(8)
        row = QHBoxLayout()
        row.setSpacing(8)
        self.circle = QLabel(str(number))
        self.circle.setFixedSize(24, 24)
        self.circle.setAlignment(Qt.AlignmentFlag.AlignCenter)
        row.addWidget(self.circle)
        self.text_label = QLabel(label)
        row.addWidget(self.text_label)
        self.stale_badge = QLabel("STALE")
        self.stale_badge.setVisible(False)
        row.addWidget(self.stale_badge)
        row.addStretch()
        v.addLayout(row)
        self.underline = QFrame()
        self.underline.setFixedHeight(2)
        v.addWidget(self.underline)
        self.set_active(False)

    def mousePressEvent(self, event):
        self.clicked.emit()
        super().mousePressEvent(event)

    def set_active(self, active):
        t = Tokens
        if active:
            self.circle.setStyleSheet(f"background:{t.US_ACCENT}; color:white; border-radius:12px; font-weight:700; font-size:{t.SIZE_CAPTION}px;")
            self.text_label.setStyleSheet(f"color:{t.TEXT_PRIMARY}; font-weight:700; font-size:{t.SIZE_BODY}px;")
            self.underline.setStyleSheet(f"background:{t.US_ACCENT}; border:none;")
        else:
            self.circle.setStyleSheet(f"background:{t.DIVIDER}; color:{t.TEXT_MUTED}; border-radius:12px; font-weight:700; font-size:{t.SIZE_CAPTION}px;")
            self.text_label.setStyleSheet(f"color:{t.TEXT_MUTED}; font-weight:600; font-size:{t.SIZE_BODY}px;")
            self.underline.setStyleSheet("background:transparent; border:none;")

    def set_stale(self, stale):
        self.stale_badge.setVisible(stale)
        if stale:
            self.stale_badge.setStyleSheet(
                f"background:{Tokens.WARNING}; color:#241c02; border-radius:{Tokens.RADIUS_PILL}px; "
                f"padding:1px 8px; font-size:10px; font-weight:700;"
            )

class StepperBar(QWidget):
    stepClicked = pyqtSignal(int)

    def __init__(self, labels):
        super().__init__()
        layout = QHBoxLayout(self)
        layout.setContentsMargins(0, 4, 0, 4)
        layout.setSpacing(32)
        self.items = []
        for i, lab in enumerate(labels):
            item = StepItem(i + 1, lab)
            item.clicked.connect(lambda idx=i: self.stepClicked.emit(idx))
            layout.addWidget(item)
            self.items.append(item)
        layout.addStretch()

    def set_active_index(self, idx):
        for i, item in enumerate(self.items):
            item.set_active(i == idx)

    def set_stale(self, idx, stale):
        self.items[idx].set_stale(stale)


class StandardRow(QWidget):
    """One selectable standard row: checkbox + name + code pair. The whole
    row toggles, not just the checkbox (checkbox is display-only)."""
    changed = pyqtSignal()

    def __init__(self, standard, checked=False):
        super().__init__()
        self.standard = standard
        self.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        row = QHBoxLayout(self)
        row.setContentsMargins(12, 8, 12, 8)
        row.setSpacing(12)
        self.checkbox = QCheckBox()
        self.checkbox.setChecked(checked)
        self.checkbox.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents, True)
        self.checkbox.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        row.addWidget(self.checkbox)
        name = QLabel(standard["label"])
        name.setStyleSheet(f"font-size:{Tokens.SIZE_BODY}px; color:{Tokens.TEXT_PRIMARY}; font-weight:600;")
        row.addWidget(name)
        row.addStretch()
        codes = QLabel(f"{standard['ind_as']} vs {standard['us_gaap']}")
        codes.setStyleSheet(f"font-family:{Tokens.FONT_MONO}; color:{Tokens.TEXT_MUTED}; font-size:{Tokens.SIZE_CAPTION}px;")
        codes.setMinimumWidth(240)
        codes.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        row.addWidget(codes)
        self.setStyleSheet(f"StandardRow {{ background: transparent; border-radius: {Tokens.RADIUS_CONTROL}px; }}")

    def mousePressEvent(self, event):
        self.checkbox.setChecked(not self.checkbox.isChecked())
        self.changed.emit()
        super().mousePressEvent(event)

    def enterEvent(self, event):
        self.setStyleSheet(f"StandardRow {{ background: {Tokens.SURFACE_RAISED}; border-radius: {Tokens.RADIUS_CONTROL}px; }}")
        super().enterEvent(event)

    def leaveEvent(self, event):
        self.setStyleSheet(f"StandardRow {{ background: transparent; border-radius: {Tokens.RADIUS_CONTROL}px; }}")
        super().leaveEvent(event)

    def is_checked(self):
        return self.checkbox.isChecked()

    def set_checked(self, value):
        self.checkbox.setChecked(value)


# -------------------- MAIN WINDOW --------------------
class ComparisonApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.company = ""
        self.file_path = ""
        self.sheets = None          # dict: sheet_name -> DataFrame, read once on Load
        self.results = {}           # dict: standard_id -> result dict
        self.skipped = []           # standards selected but sheet missing/errored
        self.financials_fs = None   # last built full Financial Statements (P&L + BS), or None
        self.is_stale = False       # True once selection/method changes after a completed calculation
        self.is_calculating = False

        self.setWindowTitle(APP_TITLE)
        self.resize(1600, 1000)
        self.setMinimumSize(1280, 840)
        self.apply_theme()
        self.build_ui()

    def apply_theme(self):
        t = Tokens
        self.setStyleSheet(f"""
            QMainWindow, QWidget {{ background: {t.BG}; color: {t.TEXT_PRIMARY}; font-size: {t.SIZE_BODY}px; }}
            QFrame#card {{
                background: {t.SURFACE}; border: 1px solid {t.BORDER}; border-radius: {t.RADIUS_CONTAINER}px;
            }}
            QLabel#title {{ color: {t.TEXT_PRIMARY}; font-size: {t.SIZE_TITLE}px; font-weight: 800; }}
            QLabel#subtitle {{ color: {t.TEXT_MUTED}; font-size: {t.SIZE_CAPTION}px; }}
            QLabel#sectionHeading {{
                color: {t.TEXT_MUTED}; font-size: {t.SIZE_SECTION}px; font-weight: 700; letter-spacing: 1.5px;
            }}
            QLineEdit, QComboBox {{
                background: {t.SURFACE_RAISED}; border: 1px solid {t.BORDER}; border-radius: {t.RADIUS_CONTROL}px;
                padding: 10px; color: {t.TEXT_PRIMARY}; font-size: {t.SIZE_BODY}px;
            }}
            QLineEdit:focus, QComboBox:focus {{ border: 1px solid {t.US_ACCENT}; }}
            QLineEdit:disabled, QComboBox:disabled {{ color: {t.TEXT_DISABLED}; border-color: {t.DIVIDER}; }}
            QPushButton {{
                background: transparent; border: none; border-radius: {t.RADIUS_CONTROL}px;
                padding: 10px 18px; color: {t.TEXT_PRIMARY}; font-weight: 700; font-size: {t.SIZE_BODY}px;
            }}
            QPushButton#primary {{ background: {t.US_ACCENT}; color: white; }}
            QPushButton#primary:hover {{ background: {t.US_ACCENT_LIGHT}; }}
            QPushButton#primary:disabled {{ background: {t.DIVIDER}; color: {t.TEXT_DISABLED}; }}
            QPushButton#ghost {{ background: transparent; border: 1px solid {t.BORDER}; color: {t.TEXT_SECONDARY}; }}
            QPushButton#ghost:hover {{ border-color: {t.US_ACCENT}; color: {t.TEXT_PRIMARY}; }}
            QPushButton#secondary {{ background: {t.SURFACE_RAISED}; border: 1px solid {t.BORDER}; color: {t.TEXT_SECONDARY}; }}
            QPushButton#secondary:hover {{ border-color: {t.US_ACCENT}; color: {t.TEXT_PRIMARY}; }}
            QPushButton#linkButton {{ background: transparent; color: {t.US_ACCENT_LIGHT}; padding: 2px 6px; font-weight: 700; }}
            QPushButton#linkButton:hover {{ color: {t.TEXT_PRIMARY}; text-decoration: underline; }}
            QTableWidget {{
                background: {t.SURFACE}; alternate-background-color: {t.SURFACE_RAISED}; gridline-color: {t.DIVIDER};
                border: 1px solid {t.BORDER}; border-radius: {t.RADIUS_CONTAINER}px; color: {t.TEXT_PRIMARY};
                font-size: {t.SIZE_BODY}px;
            }}
            QHeaderView::section {{
                background: {t.SURFACE_RAISED}; color: {t.TEXT_SECONDARY}; padding: 8px; border: 0px;
                font-weight: 700;
            }}
            QTabWidget::pane {{ border: 1px solid {t.BORDER}; border-radius: {t.RADIUS_CONTAINER}px; }}
            QTabBar::tab {{
                background: {t.SURFACE}; color: {t.TEXT_MUTED}; padding: 10px 20px;
                border: 1px solid {t.BORDER}; border-bottom: none; font-size: {t.SIZE_BODY}px;
            }}
            QTabBar::tab:selected {{ color: {t.TEXT_PRIMARY}; background: {t.SURFACE_RAISED}; }}
            QScrollArea {{ border: none; background: transparent; }}
            QGroupBox {{
                background: {t.SURFACE}; border: 1px solid {t.BORDER}; border-radius: {t.RADIUS_CONTAINER}px;
                margin-top: 10px; font-size: {t.SIZE_BODY}px;
            }}
            QGroupBox::title {{
                subcontrol-origin: margin; left: 12px; padding: 0 6px;
                color: {t.TEXT_SECONDARY}; font-weight: 700;
            }}
            QStatusBar {{ background: {t.SURFACE}; border-top: 1px solid {t.BORDER}; color: {t.TEXT_SECONDARY}; }}
            QStatusBar QLabel {{ color: {t.TEXT_SECONDARY}; }}
            QCheckBox {{ spacing: 8px; font-size: {t.SIZE_BODY}px; }}
            QCheckBox::indicator {{
                width: 18px; height: 18px; border-radius: 4px;
                border: 2px solid {t.BORDER}; background: {t.SURFACE};
            }}
            QCheckBox::indicator:checked {{
                background: {t.US_ACCENT}; border-color: {t.US_ACCENT};
            }}
        """)

    def build_ui(self):
        t = Tokens
        root = QWidget()
        main = QVBoxLayout(root)
        main.setContentsMargins(20, 20, 20, 0)
        main.setSpacing(16)

        header = QHBoxLayout()
        title_box = QVBoxLayout()
        title_box.setSpacing(4)
        title = QLabel("Ind AS (IGAAP) vs US GAAP")
        title.setObjectName("title")
        subtitle = QLabel("Multi-standard comparison · profit-first, side-by-side reconciliation")
        subtitle.setObjectName("subtitle")
        title_box.addWidget(title)
        title_box.addWidget(subtitle)
        header.addLayout(title_box)
        header.addStretch()
        main.addLayout(header)

        self.stepper = StepperBar(["Company & standards", "Accounting treatment", "Reconciliation"])
        self.stepper.stepClicked.connect(self.go_to_step)
        main.addWidget(self.stepper)
        divider = QFrame()
        divider.setFixedHeight(1)
        divider.setStyleSheet(f"background:{t.DIVIDER}; border:none;")
        main.addWidget(divider)

        self.stack = QStackedWidget()
        main.addWidget(self.stack, 1)

        self.frame1 = self.build_frame1()
        self.frame2 = self.build_frame2()
        self.frame3 = self.build_frame3()
        self.stack.addWidget(self.frame1)
        self.stack.addWidget(self.frame2)
        self.stack.addWidget(self.frame3)

        self.setCentralWidget(root)
        self.rebuild_frame2()

        # ---- status bar: coloured dot + mono status line (left), Back/Next (right) ----
        bar = QStatusBar()
        self.setStatusBar(bar)
        self.status_dot = QLabel("●")
        self.status_dot.setStyleSheet(f"color:{t.TEXT_DISABLED}; font-size:14px;")
        self.status_text = QLabel("No file loaded")
        self.status_text.setStyleSheet(f"font-family:{t.FONT_MONO}; color:{t.TEXT_SECONDARY}; font-size:{t.SIZE_CAPTION}px;")
        self.status_text.setMaximumWidth(700)
        self.status_text.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        bar.addWidget(self.status_dot)
        bar.addWidget(self.status_text)

        self.btn_back = QPushButton("Back")
        self.btn_back.setObjectName("secondary")
        self.btn_back.clicked.connect(self.go_back)
        self.btn_next = QPushButton("Next")
        self.btn_next.setObjectName("secondary")
        self.btn_next.clicked.connect(self.go_next)
        bar.addPermanentWidget(self.btn_back)
        bar.addPermanentWidget(self.btn_next)

        self.go_to_step(0)

    def go_to_step(self, idx):
        self.stack.setCurrentIndex(idx)
        self.stepper.set_active_index(idx)
        self.btn_back.setEnabled(idx > 0)
        self.btn_next.setEnabled(idx < 2)

    def go_back(self):
        self.go_to_step(max(0, self.stack.currentIndex() - 1))

    def go_next(self):
        self.go_to_step(min(2, self.stack.currentIndex() + 1))

    def set_status(self, text, dot_color=None):
        t = Tokens
        self.status_text.setText(text)
        self.status_text.setToolTip(text)
        self.status_dot.setStyleSheet(f"color:{dot_color or t.TEXT_DISABLED}; font-size:14px;")

    def mark_stale(self, stale):
        self.is_stale = stale
        self.stepper.set_stale(2, stale)
        if stale:
            self.set_status("Selection changed — reconciliation is out of date", Tokens.WARNING)

    # ---------------- FRAME 1 — Company & standards ----------------
    def build_frame1(self):
        t = Tokens
        outer = QWidget()
        outer_layout = QVBoxLayout(outer)
        outer_layout.setContentsMargins(0, 0, 0, 0)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        w = QWidget()
        layout = QVBoxLayout(w)
        layout.setContentsMargins(4, 20, 4, 20)
        layout.setSpacing(20)

        # ---- Company & source data card ----
        card = QFrame()
        card.setObjectName("card")
        card_v = QVBoxLayout(card)
        card_v.setContentsMargins(20, 20, 20, 20)
        card_v.setSpacing(16)

        grid = QGridLayout()
        grid.setHorizontalSpacing(12)
        grid.setVerticalSpacing(12)
        grid.setColumnMinimumWidth(0, 150)

        name_label = QLabel("Company name")
        name_label.setStyleSheet(f"color:{t.TEXT_SECONDARY};")
        grid.addWidget(name_label, 0, 0)
        self.company_edit = QLineEdit()
        self.company_edit.setPlaceholderText("Enter company name")
        self.company_edit.textChanged.connect(self.update_load_button_state)
        grid.addWidget(self.company_edit, 0, 1, 1, 2)

        file_label = QLabel("Journal extract")
        file_label.setStyleSheet(f"color:{t.TEXT_SECONDARY};")
        grid.addWidget(file_label, 1, 0)
        file_box = QFrame()
        file_box.setStyleSheet(
            f"background:{t.SURFACE_RAISED}; border:1px solid {t.BORDER}; border-radius:{t.RADIUS_CONTROL}px;"
        )
        file_box_h = QHBoxLayout(file_box)
        file_box_h.setContentsMargins(10, 6, 10, 6)
        file_box_h.setSpacing(8)
        self.file_name_label = QLabel("No file selected")
        self.file_name_label.setStyleSheet(f"font-family:{t.FONT_MONO}; color:{t.TEXT_PRIMARY};")
        self.file_dir_label = QLabel("")
        self.file_dir_label.setStyleSheet(f"color:{t.TEXT_MUTED}; font-size:{t.SIZE_CAPTION}px;")
        file_box_h.addWidget(self.file_name_label)
        file_box_h.addWidget(self.file_dir_label)
        file_box_h.addStretch()
        grid.addWidget(file_box, 1, 1)
        browse = QPushButton("Browse…")
        browse.setObjectName("secondary")
        browse.clicked.connect(self.browse_excel)
        grid.addWidget(browse, 1, 2)

        card_v.addLayout(grid)

        btn_divider = QFrame()
        btn_divider.setFixedHeight(1)
        btn_divider.setStyleSheet(f"background:{t.DIVIDER}; border:none;")
        card_v.addWidget(btn_divider)

        button_row = QHBoxLayout()
        template_btn = QPushButton("Create Excel template")
        template_btn.setObjectName("ghost")
        template_btn.clicked.connect(self.create_template)
        button_row.addWidget(template_btn)
        button_row.addStretch()
        self.load_btn = QPushButton(amp("Load & calculate"))
        self.load_btn.setObjectName("primary")
        self.load_btn.setEnabled(False)
        self.load_btn.setCursor(QCursor(Qt.CursorShape.ArrowCursor))
        self.load_btn.clicked.connect(self.load_excel)
        button_row.addWidget(self.load_btn)
        self.loading_caption = QLabel("")
        self.loading_caption.setStyleSheet(f"color:{t.TEXT_MUTED}; font-size:{t.SIZE_CAPTION}px;")
        button_row.addWidget(self.loading_caption)
        card_v.addLayout(button_row)

        layout.addWidget(card)

        # ---- Standards selector card ----
        std_card = QFrame()
        std_card.setObjectName("card")
        std_v = QVBoxLayout(std_card)
        std_v.setContentsMargins(20, 20, 20, 20)
        std_v.setSpacing(12)

        std_header = QHBoxLayout()
        std_title = QLabel("Applicable accounting standards")
        std_title.setStyleSheet(f"color:{t.TEXT_PRIMARY}; font-weight:700;")
        std_header.addWidget(std_title)
        std_header.addStretch()
        self.selected_count_label = QLabel("")
        self.selected_count_label.setStyleSheet(f"color:{t.TEXT_MUTED}; font-size:{t.SIZE_CAPTION}px;")
        std_header.addWidget(self.selected_count_label)
        select_all_btn = QPushButton("Select all")
        select_all_btn.setObjectName("linkButton")
        select_all_btn.clicked.connect(self.select_all_standards)
        std_header.addWidget(select_all_btn)
        clear_btn = QPushButton("Clear")
        clear_btn.setObjectName("linkButton")
        clear_btn.clicked.connect(self.clear_all_standards)
        std_header.addWidget(clear_btn)
        std_v.addLayout(std_header)

        self.standard_rows = {}
        for s in STANDARDS:
            row = StandardRow(s, checked=(s["id"] == "inventory"))
            row.changed.connect(self.on_standard_row_changed)
            self.standard_rows[s["id"]] = row
            std_v.addWidget(row)
            if s is not STANDARDS[-1]:
                d = QFrame()
                d.setFixedHeight(1)
                d.setStyleSheet(f"background:{t.DIVIDER}; border:none;")
                std_v.addWidget(d)

        method_grid = QGridLayout()
        method_grid.setColumnStretch(0, 1)
        method_grid.setColumnStretch(1, 1)
        method_grid.setHorizontalSpacing(20)
        self.ind_method_label = QLabel("Inventory — Ind AS method")
        self.ind_method_label.setStyleSheet(f"color:{t.TEXT_SECONDARY};")
        method_grid.addWidget(self.ind_method_label, 0, 0)
        self.us_method_label = QLabel("Inventory — US GAAP method")
        self.us_method_label.setStyleSheet(f"color:{t.TEXT_SECONDARY};")
        method_grid.addWidget(self.us_method_label, 0, 1)
        self.ind_method_combo = QComboBox()
        self.ind_method_combo.addItems(IND_AS_INVENTORY_METHODS)
        self.ind_method_combo.currentTextChanged.connect(self.on_method_changed)
        method_grid.addWidget(self.ind_method_combo, 1, 0)
        self.us_method_combo = QComboBox()
        self.us_method_combo.addItems(US_GAAP_INVENTORY_METHODS)
        self.us_method_combo.setCurrentText("LIFO")
        self.us_method_combo.currentTextChanged.connect(self.on_method_changed)
        method_grid.addWidget(self.us_method_combo, 1, 1)
        std_v.addLayout(method_grid)

        self.inventory_disabled_hint = QLabel("Enable the Inventory standard above to set cost formulas.")
        self.inventory_disabled_hint.setStyleSheet(f"color:{t.TEXT_MUTED}; font-size:{t.SIZE_CAPTION}px;")
        std_v.addWidget(self.inventory_disabled_hint)

        layout.addWidget(std_card)

        # ---- inline warnings (missing sheets), populated after Load & calculate ----
        self.warnings_layout = QVBoxLayout()
        self.warnings_layout.setSpacing(8)
        layout.addLayout(self.warnings_layout)

        layout.addStretch()
        self.update_selected_count()
        self.update_inventory_controls_enabled()
        scroll.setWidget(w)
        outer_layout.addWidget(scroll)
        return outer

    def selected_standard_ids(self):
        return [sid for sid, row in self.standard_rows.items() if row.is_checked()]

    def update_selected_count(self):
        n, m = len(self.selected_standard_ids()), len(STANDARDS)
        self.selected_count_label.setText(f"{n} of {m} selected")

    def select_all_standards(self):
        for row in self.standard_rows.values():
            row.set_checked(True)
        self.on_standard_row_changed()

    def clear_all_standards(self):
        for row in self.standard_rows.values():
            row.set_checked(False)
        self.on_standard_row_changed()

    def update_inventory_controls_enabled(self):
        enabled = "inventory" in self.selected_standard_ids()
        for widget in (self.ind_method_combo, self.us_method_combo, self.ind_method_label, self.us_method_label):
            widget.setEnabled(enabled)
        self.inventory_disabled_hint.setVisible(not enabled)

    def update_load_button_state(self):
        ready = bool(self.company_edit.text().strip()) and bool(self.file_path)
        self.load_btn.setEnabled(ready)
        self.load_btn.setCursor(QCursor(Qt.CursorShape.PointingHandCursor if ready else Qt.CursorShape.ForbiddenCursor))

    def rebuild_warnings(self):
        clear_layout(self.warnings_layout)
        for item in self.skipped:
            t = Tokens
            box = QFrame()
            box.setStyleSheet(f"background:{t.SURFACE}; border:1px solid {t.WARNING}; border-radius:{t.RADIUS_CONTROL}px;")
            bl = QHBoxLayout(box)
            bl.setContentsMargins(12, 8, 12, 8)
            text = QLabel(item["message"])
            text.setWordWrap(True)
            text.setStyleSheet(f"color:{t.TEXT_PRIMARY};")
            bl.addWidget(text)
            self.warnings_layout.addWidget(box)

    # ---------------- FRAME 2 (reacts to Frame 1) ----------------
    def build_frame2(self):
        w = QWidget()
        outer = QVBoxLayout(w)
        outer.setContentsMargins(4, 20, 4, 20)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        inner = QWidget()
        self.frame2_layout = QVBoxLayout(inner)
        self.frame2_layout.setSpacing(0)
        scroll.setWidget(inner)
        outer.addWidget(scroll)
        return w

    def rebuild_frame2(self):
        clear_layout(self.frame2_layout)
        selected = self.selected_standard_ids()
        if not selected:
            empty = QLabel("No standards selected yet.")
            empty.setStyleSheet(f"color:{Tokens.TEXT_MUTED}; padding:12px;")
            self.frame2_layout.addWidget(empty)
            self.frame2_layout.addStretch()
            return

        header_grid = QGridLayout()
        header_grid.setColumnStretch(0, 1)
        header_grid.setColumnStretch(1, 1)
        header_grid.setHorizontalSpacing(20)
        ind_h = QLabel("IND AS (IGAAP)")
        ind_h.setObjectName("sectionHeading")
        us_h = QLabel("US GAAP")
        us_h.setObjectName("sectionHeading")
        header_grid.addWidget(ind_h, 0, 0)
        header_grid.addWidget(us_h, 0, 1)
        self.frame2_layout.addLayout(header_grid)

        for i, sid in enumerate(selected):
            divider = QFrame()
            divider.setFixedHeight(1)
            divider.setStyleSheet(f"background:{Tokens.DIVIDER}; border:none;")
            self.frame2_layout.addWidget(divider)
            self.frame2_layout.addWidget(self.build_frame2_row(sid))
        self.frame2_layout.addStretch()

    def build_frame2_row(self, sid):
        t = Tokens
        s = STANDARDS_BY_ID[sid]
        result = self.results.get(sid)
        pl_card = next((c for c in result["recon_cards"] if c["category"] == "PL"), None) if result else None

        row_widget = QWidget()
        v = QVBoxLayout(row_widget)
        v.setContentsMargins(0, 16, 0, 16)
        v.setSpacing(10)

        header = QHBoxLayout()
        name_lbl = QLabel(s["label"])
        name_lbl.setStyleSheet(f"font-size:15px; font-weight:700; color:{t.TEXT_PRIMARY};")
        header.addWidget(name_lbl)
        header.addStretch()
        impact_caption = QLabel("Profit impact")
        impact_caption.setStyleSheet(f"color:{t.TEXT_SECONDARY}; font-size:{t.SIZE_CAPTION}px;")
        header.addWidget(impact_caption)
        if pl_card:
            phrase, side = profit_impact_phrase(pl_card["adj_value"])
            color = {"ind": t.IND_ACCENT_LIGHT, "us": t.US_ACCENT_LIGHT, "none": t.TEXT_MUTED}[side]
        else:
            phrase, color = "Load & calculate to see", t.TEXT_MUTED
        delta_lbl = QLabel(phrase)
        delta_lbl.setStyleSheet(f"font-family:{t.FONT_MONO}; font-weight:700; color:{color};")
        header.addWidget(delta_lbl)
        v.addLayout(header)

        ind_method, us_method = self.ind_method_combo.currentText(), self.us_method_combo.currentText()
        ind_basis = s["ind_basis"].format(ind_method=ind_method, us_method=us_method) if sid == "inventory" else s["ind_basis"]
        us_basis = s["us_basis"].format(ind_method=ind_method, us_method=us_method) if sid == "inventory" else s["us_basis"]
        ind_amount_txt = money_inr(pl_card["ind_value"]) if pl_card else "—"
        us_amount_txt = money_inr(pl_card["us_value"]) if pl_card else "—"
        ind_is_winner = bool(pl_card and pl_card["ind_value"] > pl_card["us_value"])
        us_is_winner = bool(pl_card and pl_card["us_value"] > pl_card["ind_value"])

        cells = QGridLayout()
        cells.setColumnStretch(0, 1)
        cells.setColumnStretch(1, 1)
        cells.setHorizontalSpacing(20)

        def make_cell(code, accent, accent_light, amount_text, basis, prose, is_winner):
            cell = QWidget()
            h = QHBoxLayout(cell)
            h.setContentsMargins(0, 0, 0, 0)
            h.setSpacing(10)
            bar = QFrame()
            bar.setFixedWidth(3)
            bar.setStyleSheet(f"background:{accent if is_winner else t.BORDER}; border:none;")
            h.addWidget(bar)
            content = QVBoxLayout()
            content.setSpacing(4)
            code_row = QHBoxLayout()
            code_lbl = QLabel(code)
            code_lbl.setStyleSheet(f"font-family:{t.FONT_MONO}; font-weight:700; color:{accent_light}; font-size:{t.SIZE_CAPTION}px; letter-spacing:0.5px;")
            code_row.addWidget(code_lbl)
            code_row.addStretch()
            if is_winner:
                pill = QLabel("HIGHER PROFIT")
                pill.setStyleSheet(f"background:{accent}; color:white; border-radius:{t.RADIUS_PILL}px; padding:1px 8px; font-size:10px; font-weight:700;")
                code_row.addWidget(pill)
            content.addLayout(code_row)
            amt_lbl = QLabel(amount_text)
            amt_lbl.setStyleSheet(f"font-family:{t.FONT_MONO}; font-weight:700; font-size:20px; color:{t.TEXT_PRIMARY};")
            content.addWidget(amt_lbl)
            basis_lbl = QLabel(basis)
            basis_lbl.setWordWrap(True)
            basis_lbl.setStyleSheet(f"color:{t.TEXT_MUTED}; font-size:{t.SIZE_CAPTION}px;")
            content.addWidget(basis_lbl)
            prose_lbl = QLabel(clamped_text(prose))
            prose_lbl.setWordWrap(True)
            prose_lbl.setStyleSheet(f"color:{t.TEXT_SECONDARY}; font-size:{t.SIZE_BODY}px; line-height:1.55;")
            content.addWidget(prose_lbl)
            h.addLayout(content)
            return cell, prose_lbl

        ind_cell, ind_prose_lbl = make_cell(s["ind_as"], t.IND_ACCENT, t.IND_ACCENT_LIGHT, ind_amount_txt, ind_basis, s["ind_desc"], ind_is_winner)
        us_cell, us_prose_lbl = make_cell(s["us_gaap"], t.US_ACCENT, t.US_ACCENT_LIGHT, us_amount_txt, us_basis, s["us_desc"], us_is_winner)
        cells.addWidget(ind_cell, 0, 0)
        cells.addWidget(us_cell, 0, 1)
        v.addLayout(cells)

        toggle_btn = QPushButton("Why this differs")
        toggle_btn.setObjectName("linkButton")
        state = {"expanded": False}

        def toggle():
            state["expanded"] = not state["expanded"]
            if state["expanded"]:
                ind_prose_lbl.setText(s["ind_desc"])
                us_prose_lbl.setText(s["us_desc"])
                toggle_btn.setText("Hide detail")
            else:
                ind_prose_lbl.setText(clamped_text(s["ind_desc"]))
                us_prose_lbl.setText(clamped_text(s["us_desc"]))
                toggle_btn.setText("Why this differs")

        toggle_btn.clicked.connect(toggle)
        v.addWidget(toggle_btn)
        return row_widget

    def on_standard_row_changed(self):
        self.update_selected_count()
        self.update_inventory_controls_enabled()
        self.update_load_button_state()
        self.rebuild_frame2()
        if self.results:
            self.mark_stale(True)

    def on_method_changed(self):
        self.rebuild_frame2()
        if self.sheets is not None and "inventory" in self.selected_standard_ids():
            self.recalculate()
        elif self.results:
            self.mark_stale(True)

    # ---------------- FRAME 3 ----------------
    def build_frame3(self):
        t = Tokens
        w = QWidget()
        layout = QVBoxLayout(w)
        layout.setContentsMargins(4, 20, 4, 20)
        layout.setSpacing(16)

        self.outer_tabs = QTabWidget()
        layout.addWidget(self.outer_tabs)

        # Tab: Reconciliation — one bordered table, profit-after-tax bridge
        self.recon_tab = QWidget()
        recon_layout = QVBoxLayout(self.recon_tab)
        recon_container = QFrame()
        recon_container.setObjectName("card")
        recon_container_v = QVBoxLayout(recon_container)
        recon_container_v.setContentsMargins(0, 0, 0, 0)
        self.recon_table = QTableWidget()
        self.recon_table.setColumnCount(2)
        self.recon_table.setHorizontalHeaderLabels(["LINE", "AMOUNT (₹)"])
        self.recon_table.verticalHeader().setVisible(False)
        self.recon_table.setShowGrid(False)
        self.recon_table.setStyleSheet(
            f"QTableWidget {{ border: none; }} "
            f"QHeaderView::section {{ font-size:{t.SIZE_SECTION}px; letter-spacing:1px; }} "
            f"QTableWidget::item {{ padding-right: 16px; }}"
        )
        recon_container_v.addWidget(self.recon_table)
        recon_layout.addWidget(recon_container)
        self.outer_tabs.addTab(self.recon_tab, amp("Reconciliation"))

        # Tab: Financial Statements (Trial Balance + standard adjustments -> full Ind AS / US GAAP financials,
        # each laid out in its own framework's real presentation format, side by side)
        self.financials_tab = QWidget()
        fin_layout = QVBoxLayout(self.financials_tab)
        self.financials_status = QLabel(
            "Add a 'Trial Balance' sheet to the workbook (see 'Create Excel Template') to generate full "
            "financial statements here. Without it, the standard-wise tabs still work independently."
        )
        self.financials_status.setWordWrap(True)
        self.financials_status.setStyleSheet("color:#94A3B8; padding:6px;")
        fin_layout.addWidget(self.financials_status)

        fw_split = QSplitter(Qt.Orientation.Horizontal)

        ind_panel = QGroupBox("Ind AS — Schedule III Format")
        ind_panel_v = QVBoxLayout(ind_panel)
        ind_split = QSplitter(Qt.Orientation.Vertical)
        ind_pl_box = QGroupBox("Statement of Profit and Loss")
        ind_pl_v = QVBoxLayout(ind_pl_box)
        self.ind_pl_table = QTableWidget()
        ind_pl_v.addWidget(self.ind_pl_table)
        ind_bs_box = QGroupBox("Balance Sheet")
        ind_bs_v = QVBoxLayout(ind_bs_box)
        self.ind_bs_table = QTableWidget()
        ind_bs_v.addWidget(self.ind_bs_table)
        ind_split.addWidget(ind_pl_box)
        ind_split.addWidget(ind_bs_box)
        ind_panel_v.addWidget(ind_split)

        us_panel = QGroupBox("US GAAP — Classified Format")
        us_panel_v = QVBoxLayout(us_panel)
        us_split = QSplitter(Qt.Orientation.Vertical)
        us_pl_box = QGroupBox("Income Statement")
        us_pl_v = QVBoxLayout(us_pl_box)
        self.us_pl_table = QTableWidget()
        us_pl_v.addWidget(self.us_pl_table)
        us_bs_box = QGroupBox("Balance Sheet")
        us_bs_v = QVBoxLayout(us_bs_box)
        self.us_bs_table = QTableWidget()
        us_bs_v.addWidget(self.us_bs_table)
        us_split.addWidget(us_pl_box)
        us_split.addWidget(us_bs_box)
        us_panel_v.addWidget(us_split)

        fw_split.addWidget(ind_panel)
        fw_split.addWidget(us_panel)
        fin_layout.addWidget(fw_split)
        self.outer_tabs.addTab(self.financials_tab, "Financial Statements")

        # dynamic per-standard tabs go after these two fixed tabs
        self.dynamic_tab_count = 0

        exports = QHBoxLayout()
        b1 = QPushButton("Extract Consolidated Reconciliation")
        b1.clicked.connect(self.export_reconciliation)
        b2 = QPushButton("Extract Ind AS Financials")
        b2.clicked.connect(self.export_indas)
        b3 = QPushButton("Extract US GAAP Financials")
        b3.clicked.connect(self.export_usgaap)
        b4 = QPushButton("Extract Full Workbook")
        b4.setObjectName("success")
        b4.clicked.connect(self.export_full)
        for b in [b1, b2, b3, b4]:
            exports.addWidget(b)
        layout.addLayout(exports)

        return w

    # ---------------- ACTIONS ----------------
    def browse_excel(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select Journal Extract Excel", str(Path.home()), "Excel Workbook (*.xlsx *.xls)",
            options=QFileDialog.Option.DontUseNativeDialog
        )
        if path:
            self.file_path = path
            p = Path(path)
            self.file_name_label.setText(amp(p.name))
            self.file_dir_label.setText(amp(str(p.parent)))
            self.file_name_label.setToolTip(path)
            self.file_dir_label.setToolTip(path)
            self.update_load_button_state()

    def create_template(self):
        start_path = str(Path.home() / "Journal_Extract_Template.xlsx")
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Excel Template", start_path, "Excel Workbook (*.xlsx)",
            options=QFileDialog.Option.DontUseNativeDialog
        )
        if not path:
            return
        try:
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                readme_rows = [["Sheet", "Ind AS", "US GAAP", "What differs"]]
                for s in STANDARDS:
                    readme_rows.append([s["sheet"], s["ind_as"], s["us_gaap"], s["label"]])
                readme_rows.append(["Trial Balance", "-", "-", "Base financial statement figures (Ind AS books); full Ind AS/US GAAP financials are generated from this plus the standard sheets"])
                pd.DataFrame(readme_rows[1:], columns=readme_rows[0]).to_excel(writer, sheet_name="Read Me", index=False)
                pd.DataFrame(TB_TEMPLATE_ROWS, columns=TB_REQUIRED_COLUMNS).to_excel(writer, sheet_name="Trial Balance", index=False)
                for s in STANDARDS:
                    pd.DataFrame(s["template_rows"], columns=s["template_columns"]).to_excel(
                        writer, sheet_name=sheet_name(s["sheet"]), index=False
                    )
            style_workbook(path)
            QMessageBox.information(
                self, "Template Created",
                "Template created with a 'Read Me' sheet, a 'Trial Balance' sheet (the company's Ind AS "
                "books - the base for the full financial statements), and one sheet per standard. Fill in "
                "the Trial Balance plus only the standard sheets applicable to your company, then upload "
                "the same file here. Sheets you leave blank are simply skipped."
            )
        except Exception as e:
            QMessageBox.critical(self, "Template Error", str(e))

    def load_excel(self):
        company = self.company_edit.text().strip()
        if not company or not self.file_path:
            return
        selected = self.selected_standard_ids()
        if not selected:
            QMessageBox.warning(self, "Standard Required", "Please select at least one applicable standard.")
            return
        filename = Path(self.file_path).name
        self.is_calculating = True
        self.load_btn.setEnabled(False)
        self.loading_caption.setText("Reading workbook…")
        self.set_status(f"Reading {filename}…", Tokens.US_ACCENT)
        QApplication.processEvents()
        try:
            self.sheets = pd.read_excel(self.file_path, sheet_name=None)
        except Exception as e:
            QMessageBox.critical(self, "File Error", str(e))
            self.is_calculating = False
            self.loading_caption.setText("")
            self.update_load_button_state()
            self.set_status(f"{filename} · failed to read", Tokens.WARNING)
            return
        self.company = company
        self.recalculate()
        self.is_calculating = False
        self.loading_caption.setText("")
        self.update_load_button_state()
        if self.results:
            self.mark_stale(False)
            n_selected, n_loaded = len(selected), len(self.results)
            self.set_status(f"{filename} · loaded {n_loaded} of {n_selected} standards · {n_selected} selected", Tokens.IND_ACCENT)
            self.go_to_step(1)

    def recalculate(self):
        if self.sheets is None:
            return
        selected = self.selected_standard_ids()
        self.results, self.skipped = {}, []

        for sid in selected:
            s = STANDARDS_BY_ID[sid]
            sheet_df = self.sheets.get(s["sheet"])
            if sheet_df is None:
                self.skipped.append({
                    "label": s["label"],
                    "message": f"{s['label']}: expected a sheet named \"{s['sheet']}\" — not found in the workbook.",
                })
                continue
            try:
                if sid == "inventory":
                    result = calc_inventory_standard(sheet_df, self.ind_method_combo.currentText(), self.us_method_combo.currentText())
                else:
                    result = CALC_FUNCS[sid](sheet_df)
                self.results[sid] = result
            except Exception as e:
                self.skipped.append({"label": s["label"], "message": f"{s['label']}: {e}"})

        self.rebuild_warnings()

        if self.results:
            self.populate_results()
            self.rebuild_frame2()

    def populate_results(self):
        # ---- reconciliation: one table, profit after tax bridge (item 17) ----
        self.populate_recon_table()

        # ---- full financial statements, from Trial Balance + standards' adjustments ----
        tb_df = None
        if self.sheets is not None:
            for name, df in self.sheets.items():
                if name.strip().lower() == "trial balance":
                    tb_df = df
                    break
        if tb_df is None:
            self.financials_status.setText(
                "No 'Trial Balance' sheet found in the uploaded workbook. Add one (see 'Create Excel "
                "Template') to generate full financial statements here — the standard-wise tabs above "
                "still reflect every applied standard regardless."
            )
            self.financials_status.setStyleSheet("color:#FBBF24; padding:6px;")
            empty = pd.DataFrame(columns=["Line Item", "Amount"])
            set_table(self.ind_pl_table, empty)
            set_table(self.ind_bs_table, empty)
            set_table(self.us_pl_table, empty)
            set_table(self.us_bs_table, empty)
            self.financials_fs = None
        else:
            try:
                fs = build_financial_statements(tb_df, self.results)
                set_table(self.ind_pl_table, fs["ind_as_pl"], ["Amount"])
                set_table(self.ind_bs_table, fs["ind_as_bs"], ["Amount"])
                set_table(self.us_pl_table, fs["us_gaap_pl"], ["Amount"])
                set_table(self.us_bs_table, fs["us_gaap_bs"], ["Amount"])
                self.financials_fs = fs
                check_msg = ""
                if abs(fs["balance_check_ind"]) > 1.0:
                    check_msg += f" Ind AS Trial Balance does not balance (Assets − Liabilities − Equity = {money(fs['balance_check_ind'])})."
                if abs(fs["balance_check_us"]) > 1.0:
                    check_msg += f" US GAAP statement does not balance ({money(fs['balance_check_us'])})."
                if check_msg:
                    self.financials_status.setText("Balance check failed:" + check_msg + " Review the Trial Balance figures.")
                    self.financials_status.setStyleSheet("color:#F87171; padding:6px;")
                else:
                    self.financials_status.setText(
                        f"Generated from the Trial Balance plus {len(self.results)} applied standard(s). "
                        "Balance sheet balances under both frameworks (Assets = Liabilities + Equity, checked)."
                    )
                    self.financials_status.setStyleSheet("color:#34D399; padding:6px;")
                self.populate_recon_table()  # net_profit_ind/us now available — refresh the bridge rows
            except Exception as e:
                self.financials_status.setText(f"Could not build financial statements: {e}")
                self.financials_status.setStyleSheet("color:#F87171; padding:6px;")

        # ---- dynamic per-standard tabs ----
        while self.outer_tabs.count() > 2:
            w = self.outer_tabs.widget(2)
            self.outer_tabs.removeTab(2)
            if w is not None:
                w.deleteLater()

        for sid, r in self.results.items():
            std_widget = QWidget()
            std_layout = QVBoxLayout(std_widget)
            if r.get("method_caption"):
                cap = QLabel(r["method_caption"])
                cap.setStyleSheet("color:#93C5FD; padding:4px 0;")
                std_layout.addWidget(cap)
            inner_tabs = QTabWidget()

            trans_table = QTableWidget()
            set_table(trans_table, r["transactions"], r["number_cols_transactions"])
            inner_tabs.addTab(trans_table, amp("Transactions"))

            inner_tabs.addTab(build_entries_split(r["ind_journal"], r["us_journal"]), amp("Accounting Entries"))

            fin_table = QTableWidget()
            set_table(fin_table, r["financial_lines"], ["Ind AS", "US GAAP", "US GAAP less Ind AS"])
            inner_tabs.addTab(fin_table, amp("Financial Extract"))

            std_layout.addWidget(inner_tabs)
            self.outer_tabs.addTab(std_widget, amp(r["label"][:28]))

    def populate_recon_table(self):
        """Item 17: 'Profit after tax — Ind AS' (needs a Trial Balance; flagged
        rather than invented when absent), one signed 'Adjustment — <standard>'
        row per selected standard, then 'Profit after tax — US GAAP'."""
        t = Tokens
        rows = []  # (label, value_or_None, style)
        ind_profit = self.financials_fs["net_profit_ind"] if self.financials_fs else None
        us_profit = self.financials_fs["net_profit_us"] if self.financials_fs else None

        rows.append(("Profit after tax — Ind AS", ind_profit, "semibold"))
        for sid, r in self.results.items():
            pl_card = next((c for c in r["recon_cards"] if c["category"] == "PL"), None)
            if pl_card is not None:
                rows.append((f"Adjustment — {r['label']}", pl_card["adj_value"], "normal"))
        rows.append(("Profit after tax — US GAAP", us_profit, "raised"))

        self.recon_table.setRowCount(len(rows))
        mono_bold = QFont(t.FONT_MONO.split(",")[0].strip())
        mono_bold.setBold(True)
        mono_normal = QFont(t.FONT_MONO.split(",")[0].strip())
        sans_bold = QFont()
        sans_bold.setBold(True)

        for i, (label, value, style) in enumerate(rows):
            label_item = QTableWidgetItem(label)  # QTableWidgetItem doesn't parse mnemonics — no amp() needed
            label_item.setFont(sans_bold if style in ("semibold", "raised") else QFont())
            value_text = money_inr(value) if value is not None else "Upload a Trial Balance to compute"
            value_item = QTableWidgetItem(value_text)
            value_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            value_item.setFont(mono_bold if style in ("semibold", "raised") else mono_normal)
            if value is None:
                value_item.setForeground(QColor(t.TEXT_MUTED))
            self.recon_table.setItem(i, 0, label_item)
            self.recon_table.setItem(i, 1, value_item)
            if style == "raised":
                self.recon_table.item(i, 0).setBackground(QColor(t.SURFACE_RAISED))
                self.recon_table.item(i, 1).setBackground(QColor(t.SURFACE_RAISED))

        self.recon_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        self.recon_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Fixed)
        self.recon_table.setColumnWidth(1, 220)

    def export_reconciliation(self):
        if not self.results:
            QMessageBox.warning(self, "No Data", "Load and calculate a journal extract first.")
            return
        rows = []
        for sid, r in self.results.items():
            for c in r["recon_cards"]:
                rows.append([r["label"], c["category"], c["ind_label"], c["ind_value"], c["adj_label"], c["adj_value"], c["us_label"], c["us_value"]])
        df = pd.DataFrame(rows, columns=["Standard", "Category", "Ind AS Item", "Ind AS Amount", "Adjustment Label", "Adjustment Amount", "US GAAP Item", "US GAAP Amount"])
        export_single_sheet(df, "Consolidated reconciliation", f"{self.company}_IndAS_to_USGAAP_Reconciliation.xlsx")

    def export_indas(self):
        if not self.results:
            QMessageBox.warning(self, "No Data", "Load and calculate a journal extract first.")
            return
        if self.financials_fs:
            df = pd.concat([self.financials_fs["ind_as_pl"], self.financials_fs["ind_as_bs"]], ignore_index=True).drop(columns=["_bold"])
        else:
            rows = []
            for sid, r in self.results.items():
                fl = r["financial_lines"][["Line Item", "Ind AS"]].copy()
                fl.insert(0, "Standard", r["label"])
                rows.append(fl)
            df = pd.concat(rows, ignore_index=True)
        export_single_sheet(df, "Ind AS financial statement", f"{self.company}_IndAS_Financials.xlsx")

    def export_usgaap(self):
        if not self.results:
            QMessageBox.warning(self, "No Data", "Load and calculate a journal extract first.")
            return
        if self.financials_fs:
            df = pd.concat([self.financials_fs["us_gaap_pl"], self.financials_fs["us_gaap_bs"]], ignore_index=True).drop(columns=["_bold"])
        else:
            rows = []
            for sid, r in self.results.items():
                fl = r["financial_lines"][["Line Item", "US GAAP"]].copy()
                fl.insert(0, "Standard", r["label"])
                rows.append(fl)
            df = pd.concat(rows, ignore_index=True)
        export_single_sheet(df, "US GAAP financial statement", f"{self.company}_USGAAP_Financials.xlsx")

    def export_full(self):
        if not self.results:
            QMessageBox.warning(self, "No Data", "Load and calculate a journal extract first.")
            return
        start_path = str(Path.home() / f"{self.company}_IndAS_vs_USGAAP.xlsx")
        path, _ = QFileDialog.getSaveFileName(
            self, "Export Full Workbook", start_path, "Excel Workbook (*.xlsx)",
            options=QFileDialog.Option.DontUseNativeDialog
        )
        if not path:
            return
        try:
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                if self.financials_fs:
                    self.financials_fs["ind_as_pl"].drop(columns=["_bold"]).to_excel(writer, sheet_name="Ind AS - P&L", index=False)
                    self.financials_fs["ind_as_bs"].drop(columns=["_bold"]).to_excel(writer, sheet_name="Ind AS - Balance Sheet", index=False)
                    self.financials_fs["us_gaap_pl"].drop(columns=["_bold"]).to_excel(writer, sheet_name="US GAAP - Income Stmt", index=False)
                    self.financials_fs["us_gaap_bs"].drop(columns=["_bold"]).to_excel(writer, sheet_name="US GAAP - Balance Sheet", index=False)
                recon_rows = []
                for sid, r in self.results.items():
                    for c in r["recon_cards"]:
                        recon_rows.append([r["label"], c["category"], c["ind_label"], c["ind_value"], c["adj_label"], c["adj_value"], c["us_label"], c["us_value"]])
                pd.DataFrame(recon_rows, columns=["Standard", "Category", "Ind AS Item", "Ind AS Amount", "Adjustment Label", "Adjustment Amount", "US GAAP Item", "US GAAP Amount"]).to_excel(
                    writer, sheet_name="Consolidated Reconciliation", index=False)
                for sid, r in self.results.items():
                    short = STANDARDS_BY_ID[sid]["label"]
                    r["transactions"].to_excel(writer, sheet_name=sheet_name(short, " Txn"), index=False)
                    r["ind_journal"].to_excel(writer, sheet_name=sheet_name(short, " Ind Jnl"), index=False)
                    r["us_journal"].to_excel(writer, sheet_name=sheet_name(short, " US Jnl"), index=False)
                    r["financial_lines"].to_excel(writer, sheet_name=sheet_name(short, " Extract"), index=False)
            style_workbook(path)
            QMessageBox.information(self, "Export Complete", "Full workbook exported with the financial statements (if a Trial Balance was supplied), the consolidated reconciliation, and every applied standard's detail.")
        except Exception as e:
            QMessageBox.critical(self, "Export Error", str(e))


# -------------------- ENTRY POINT --------------------
if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setApplicationName(APP_TITLE)
    window = ComparisonApp()
    window.show()
    sys.exit(app.exec())
