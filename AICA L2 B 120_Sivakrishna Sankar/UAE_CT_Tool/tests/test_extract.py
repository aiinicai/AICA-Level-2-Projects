import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook

from helpers import GEN, ensure_sample_data
from ct_tool.extract import parse_gl, parse_tb

ensure_sample_data()


def test_tb_xlsx_extracts_all_accounts_and_balances():
    tb = parse_tb(str(GEN / "Sample_TB_USD.xlsx"), currency="USD")
    assert len(tb.accounts) == 43
    assert abs(sum(a.net for a in tb.accounts)) < 0.01
    assert not tb.warnings


def test_tb_csv_matches_xlsx():
    a = parse_tb(str(GEN / "Sample_TB_USD.xlsx"), currency="USD")
    b = parse_tb(str(GEN / "Sample_TB_USD.csv"), currency="USD")
    assert [(x.code, round(x.net, 2)) for x in a.accounts] == [(x.code, round(x.net, 2)) for x in b.accounts]


def _same(a, b):
    assert len(b.accounts) == len(a.accounts)
    am = {x.code: x.net for x in a.accounts}
    for x in b.accounts:
        assert abs(am[x.code] - x.net) < 0.01, (x.code, x.net, am[x.code])


def test_tb_pdf_with_table_lines_matches_xlsx():
    a = parse_tb(str(GEN / "Sample_TB_USD.xlsx"), currency="USD")
    _same(a, parse_tb(str(GEN / "Sample_TB_USD.pdf"), currency="USD", ocr="off"))


def test_tb_plain_text_pdf_gets_debit_and_credit_right_by_column_position():
    a = parse_tb(str(GEN / "Sample_TB_USD.xlsx"), currency="USD")
    b = parse_tb(str(GEN / "Sample_TB_USD_plain.pdf"), currency="USD", ocr="off")
    assert b.extraction_method == "pdf-text-positional"
    _same(a, b)


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="Tesseract not installed")
def test_tb_scanned_pdf_via_ocr_reads_signs_correctly_and_warns():
    a = parse_tb(str(GEN / "Sample_TB_USD.xlsx"), currency="USD")
    b = parse_tb(str(GEN / "Sample_TB_USD_scanned.pdf"), currency="USD", ocr="auto")
    assert b.extraction_method == "pdf-ocr"
    assert any("OCR" in w for w in b.warnings), "OCR output must carry a review warning"
    _same(a, b)          # holds for this clean 300 dpi synthetic scan; real scans can differ - hence the warning


def test_gl_extract_reconciles_to_tb_net_per_account():
    tb = parse_tb(str(GEN / "Sample_TB_USD.xlsx"), currency="USD")
    gl = parse_gl(str(GEN / "Sample_GL_USD.xlsx"))
    assert len(gl.lines) == 220
    net = {}
    for l in gl.lines:
        net[l.account_code] = net.get(l.account_code, 0) + l.debit - l.credit
    for a in tb.accounts:
        assert abs(net.get(a.code, 0) - a.net) < 0.02, a.code


def test_single_signed_column_and_title_rows(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ABC Trading LLC"
    ws["A2"] = "Trial balance as at 31 Dec 2025"
    ws.append([])
    ws.append(["Code", "Account", "Balance"])
    for r in [("1000", "Cash at bank", 500), ("2000", "Trade payables", -200), ("3000", "Share capital", -300)]:
        ws.append(list(r))
    f = tmp_path / "tb.xlsx"
    wb.save(f)
    tb = parse_tb(str(f), currency="AED")
    assert [round(a.net) for a in tb.accounts] == [500, -200, -300]
    assert "ABC Trading" in tb.title


def test_unbalanced_tb_is_reported_not_hidden(tmp_path):
    from ct_tool import validate
    wb = Workbook()
    ws = wb.active
    ws.append(["Code", "Account", "Debit", "Credit"])
    ws.append(["1000", "Cash at bank", 1000, None])
    ws.append(["3000", "Share capital", None, 900])
    f = tmp_path / "tb.xlsx"
    wb.save(f)
    tb = parse_tb(str(f), currency="AED")
    res = validate.check_tb(tb)
    assert any(c.status == "FAIL" for c in res)


def test_source_files_are_never_modified():
    import hashlib
    def h(p):
        return hashlib.sha256(Path(p).read_bytes()).hexdigest()
    files = [GEN / "Sample_TB_USD.xlsx", GEN / "Sample_GL_USD.xlsx", GEN / "Sample_TB_USD.pdf"]
    before = [h(p) for p in files]
    from helpers import make_project, scenario_answers
    from ct_tool.pipeline import run_analysis
    from ct_tool.service import export_all
    import tempfile
    p = make_project("USD")
    p.answers = scenario_answers()
    export_all(run_analysis(p), tempfile.mkdtemp())
    assert before == [h(p) for p in files]
