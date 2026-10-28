import shutil
import tempfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

import xl_verify as V
from helpers import TEMPLATE, confirm_all, make_project, make_scaled_project, scenario_answers
from ct_tool.excel_out import ClientWorkbook
from ct_tool.pipeline import run_analysis
from ct_tool.review_xlsx import ReviewWorkbook
from ct_tool.word_out import build_report

LO = pytest.mark.skipif(V.soffice_path() is None, reason="LibreOffice not available for formula recalculation")


def build_client(an):
    out = Path(tempfile.mkdtemp()) / "client.xlsx"
    ClientWorkbook(an, str(TEMPLATE)).build(str(out))
    return out


def check_against_python(an, tolerance=1.0):
    """Recalculate the workbook with LibreOffice and compare the key cells with the Python computation."""
    path = build_client(an)
    calc = V.recalc(path)
    assert V.errors(calc) == []
    wb = V.values(calc)
    ws = wb["Calculation of CT"]
    c = an.comp
    assert ws["E11"].value == pytest.approx(c.A, abs=tolerance)
    assert ws["E72"].value == pytest.approx(c.payable, abs=tolerance)
    if c.basis == "ordinary":
        assert ws["E25"].value == pytest.approx(c.B, abs=tolerance)
        assert ws["E50"].value == pytest.approx(c.C, abs=tolerance)
        assert ws["E57"].value == pytest.approx(c.D, abs=tolerance)
        assert ws["E63"].value == pytest.approx(c.E, abs=tolerance)
    # statements agree with the Python statements and the TB
    assert wb["Profit and Loss"]["D30"].value == pytest.approx(an.st.profit_for_year, abs=tolerance)
    bs = wb["Balance Sheet"]
    checks = [bs.cell(r, 3).value for r in range(1, bs.max_row + 1) if bs.cell(r, 3).value in ("Balances",) or str(bs.cell(r, 3).value).startswith("DOES NOT")]
    return wb, checks, path


@LO
def test_usd_scope_a_all_confirmed_workbook_matches_python():
    p = make_project("USD")
    p.answers = scenario_answers("A")
    an = confirm_all(p)
    wb, checks, _ = check_against_python(an)
    assert checks == [] or checks == ["Balances"]
    assert "FX Rates" in wb.sheetnames and "Adjustments" in wb.sheetnames


@LO
def test_nothing_confirmed_workbook_omits_inapplicable_sheets():
    p = make_project("USD")
    p.answers = scenario_answers("none")
    p.answers["foreign_taxes"] = []          # no foreign tax reported -> no Taxes Paid Abroad sheet
    an = run_analysis(p)
    wb, _, path = check_against_python(an)
    names = load_workbook(path).sheetnames
    for gone in ("Adjustments", "Disallowed Expenses", "UGUL", "Freezone", "Tax Loss Carried Forward", "Taxes Paid Abroad", "DTA & DTL Working", "Notes"):
        assert gone not in names, gone
    assert {"Assumptions", "Calculation of CT", "Profit and Loss", "Balance Sheet", "TB (AED)"} <= set(names)


@LO
def test_aed_tb_has_no_fx_sheet():
    p = make_project("AED")
    p.answers = scenario_answers("none")
    an = run_analysis(p)
    wb, _, path = check_against_python(an)
    assert "FX Rates" not in load_workbook(path).sheetnames


@LO
def test_eur_translation_workbook_balances_and_shows_translation_difference():
    p = make_project("EUR", rates={"pl": 4.0, "closing": 4.2, "equity": 3.9})
    p.answers = scenario_answers("none")
    an = confirm_all(p)
    wb, _, path = check_against_python(an)
    assert "FX Rates" in wb.sheetnames
    tb = wb["TB (AED)"]
    texts = " ".join(str(c.value) for row in tb.iter_rows() for c in row if c.value is not None)
    assert "translation" in texts.lower()
    bs = wb["Balance Sheet"]
    diff_cells = [bs.cell(r, 4).value for r in range(1, bs.max_row + 1) if str(bs.cell(r, 2).value or "").startswith("Check")]
    assert diff_cells and abs(diff_cells[0]) < 0.5


@LO
def test_loss_relief_workbook_matches_python():
    p = make_project("USD")
    p.answers = scenario_answers("A")
    p.answers["opening_losses"] = 6_000_000.0
    an = confirm_all(p)
    wb, _, path = check_against_python(an)
    assert "Tax Loss Carried Forward" in load_workbook(path).sheetnames
    assert wb["Calculation of CT"]["E60"].value == pytest.approx(-an.comp.relief_cf, abs=1)


@LO
def test_small_business_relief_workbook_matches_python():
    p = make_scaled_project(tempfile.mkdtemp(), "AED", 0.03)
    p.answers = {**scenario_answers("none"), "sbr_elect": "Yes", "mne_group": "No", "revenue_previous": 0.0}
    an = confirm_all(p)
    assert an.comp.basis == "SBR"
    check_against_python(an)


@LO
def test_qfzp_workbook_matches_python():
    p = make_project("AED")
    p.answers = scenario_answers("none")
    p.answers.update(entity_status="Free Zone", free_zone_name="Test FZ", qfzp_claimed="Yes", qfzp_substance="Yes", qfzp_audited="Yes", qfzp_domestic_pe="No",
                     qfzp_activities=["Trading of qualifying commodities"],
                     income_classification=[{"account": a, "class": "Qualifying income"} for a in ("4000", "4100", "4500")] +
                                           [{"account": "4200", "class": "Exempt income"}, {"account": "4300", "class": "Non-qualifying income"},
                                            {"account": "4600", "class": "Non-qualifying income"}])
    an = confirm_all(p)
    assert an.comp.basis in ("QFZP", "ordinary")      # outcome depends on the de-minimis test; the workbook must agree either way
    wb, _, path = check_against_python(an)
    if an.comp.basis == "QFZP":
        assert "Freezone" in load_workbook(path).sheetnames


@LO
def test_dta_sheet_only_when_requested():
    p = make_project("USD")
    p.answers = scenario_answers("none")
    p.settings.include_dta = True
    an = run_analysis(p)
    _, _, path = check_against_python(an)
    assert "DTA & DTL Working" in load_workbook(path).sheetnames


def test_client_workbook_has_no_stale_defined_names_or_external_links():
    p = make_project("USD")
    p.answers = scenario_answers("A")
    an = confirm_all(p)
    path = build_client(an)
    wb = load_workbook(path)
    assert not list(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else True
    import zipfile
    names = zipfile.ZipFile(path).namelist()
    assert not [n for n in names if "externalLink" in n]


def test_review_workbook_has_no_empty_sheets_and_keeps_proposed_and_confirmed_separate():
    p = make_project("USD")
    p.answers = scenario_answers("A")
    an = run_analysis(p)
    out = Path(tempfile.mkdtemp()) / "review.xlsx"
    ReviewWorkbook(an).build(str(out))
    wb = load_workbook(out)
    for ws in wb:
        assert ws.max_row > 5, f"{ws.title} is empty"
    reg = wb["Adjustments Register"]
    statuses = {reg.cell(r, 7).value for r in range(6, reg.max_row + 1)}
    assert statuses == {"Proposed"}
    assert "Rules & Sources" in wb.sheetnames and "Reconciliation Checks" in wb.sheetnames


def test_word_report_structure_status_vocabulary_and_disclaimer():
    p = make_project("USD")
    p.answers = scenario_answers("A")
    an = run_analysis(p)
    out = Path(tempfile.mkdtemp()) / "r.docx"
    build_report(an, str(out))
    d = Document(out)
    heads = [x.text for x in d.paragraphs if x.style.name.startswith("Heading 1")]
    for n in range(1, 16):
        assert any(h.startswith(f"{n}. ") for h in heads), f"section {n} missing"
    text = "\n".join(x.text for x in d.paragraphs) + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    for w in ("Confirmed", "Potential adjustment", "Information required", "Technical review required"):
        assert w in text
    assert "does not replace a complete professional review" in text
    assert "DRAFT" in text


def test_word_report_never_calls_a_proposed_item_disallowed():
    p = make_project("USD")
    p.answers = scenario_answers("A")
    an = run_analysis(p)          # nothing confirmed
    out = Path(tempfile.mkdtemp()) / "r.docx"
    build_report(an, str(out))
    d = Document(out)
    import re
    pieces = [x.text for x in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
    # the only sentences allowed to contain "is disallowed" are the explicit negations in the disclaimer / wording note
    allowed = ("NOT conclusions that an expense is disallowed", "described as 'disallowed' only where its status is Confirmed")
    for t in pieces:
        if any(a in t for a in allowed):
            continue
        assert not re.search(r"\b(is|are|was|were|be)\s+(hereby\s+)?(wholly\s+)?disallowed", t, re.I), t[:200]


def test_word_report_uses_supplied_letterhead():
    lh = Document()
    lh.sections[0].header.paragraphs[0].text = "ACME ADVISORY - LETTERHEAD TEST"
    lh.add_paragraph("body text that must be removed")
    path = Path(tempfile.mkdtemp()) / "lh.docx"
    lh.save(path)
    p = make_project("USD")
    p.answers = scenario_answers("A")
    an = run_analysis(p)
    out = Path(tempfile.mkdtemp()) / "r.docx"
    build_report(an, str(out), str(path))
    d = Document(out)
    assert "LETTERHEAD TEST" in d.sections[0].header.paragraphs[0].text
    assert "must be removed" not in " ".join(x.text for x in d.paragraphs)


def test_project_roundtrip_gives_identical_results():
    from ct_tool.project import Project
    p = make_project("USD")
    p.answers = scenario_answers("A")
    an = confirm_all(p)
    f = Path(tempfile.mkdtemp()) / "p.ctproj"
    p.save(str(f))
    an2 = run_analysis(Project.load(str(f)))
    assert an2.comp.payable == an.comp.payable and an2.comp.E == an.comp.E
    assert [a.status for a in an2.adjs] == [a.status for a in an.adjs]


def test_cli_end_to_end(tmp_path):
    from ct_tool.cli import main
    from helpers import GEN
    rc = main(["analyse", "--tb", str(GEN / "Sample_TB_USD.xlsx"), "--gl", str(GEN / "Sample_GL_USD.xlsx"), "--currency", "USD",
               "--period-start", "2025-01-01", "--period-end", "2025-12-31", "--out", str(tmp_path)])
    assert rc == 0
    got = sorted(x.suffix for x in tmp_path.iterdir())
    assert got == [".ctproj", ".docx", ".xlsx", ".xlsx"]


def test_cli_refuses_non_peg_currency_without_rates(tmp_path, capsys):
    from ct_tool.cli import main
    from helpers import GEN
    rc = main(["analyse", "--tb", str(GEN / "Sample_TB_EUR.xlsx"), "--currency", "EUR", "--out", str(tmp_path)])
    assert rc == 2 and not list(tmp_path.iterdir())
