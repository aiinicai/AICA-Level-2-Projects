"""Marjan Gulf Trading LLC (sample): AED trial balance, year ended 31 Dec 2025, with disallowed items.
Expected figures below were worked out BY HAND, independently of the tool (see docs/MARJAN_WORKED_EXAMPLE.md)."""
import csv
import json
import os
import sys

import pytest

from helpers import ROOT, TEMPLATE, confirm_all  # noqa: F401
sys.path.insert(0, str(ROOT / "sample_data"))
import make_marjan_tb as M  # noqa: E402
import xl_verify as V  # noqa: E402

from ct_tool.extract import parse_tb, to_num  # noqa: E402
from ct_tool.fx import default_settings  # noqa: E402
from ct_tool.pipeline import run_analysis  # noqa: E402
from ct_tool.project import Project  # noqa: E402
from ct_tool.questionnaire import RB_OPTIONS  # noqa: E402
from ct_tool.service import safe_name  # noqa: E402

ANSWERS = json.loads((ROOT / "sample_data" / "marjan" / "Marjan_answers.json").read_text(encoding="utf-8"))
LO = pytest.mark.skipif(V.soffice_path() is None, reason="LibreOffice not available")

# hand calculation (AED): profit 2,990,000
#   less exempt dividend                      (400,000)   participation exemption, conditions answered as met
#   add back 50% of client entertainment        70,000    50% x 140,000
#   add back donations (not to a QPBE)          75,000
#   add back fines and penalties                28,000
#   add back director personal expenses         42,000
#   taxable income                          2,805,000     tax = 9% x (2,805,000 - 375,000) = 218,700
PROFIT, TAXABLE, TAX = 2_990_000.00, 2_805_000.00, 218_700.00


@pytest.fixture(scope="module")
def tb_path(tmp_path_factory):
    d = tmp_path_factory.mktemp("marjan")
    M.write(d / "tb.xlsx")
    return str(d / "tb.xlsx")


def project(tb_path, **over):
    p = Project(name="Marjan")
    p.tb = parse_tb(tb_path, currency="AED")
    p.fx = default_settings("AED", "2025-01-01", "2025-12-31")
    p.fx_confirmed = True
    p.answers = {**ANSWERS, **over}
    return p


def test_tb_reads_cleanly_balances_and_maps_every_account(tb_path):
    p = project(tb_path)
    assert len(p.tb.accounts) == 46 and p.tb.warnings == []
    assert abs(sum(a.net for a in p.tb.accounts)) < 0.005
    an = run_analysis(p)
    assert [m.account.name for m in an.mapped if m.method != "rule"] == []          # nothing unmapped, nothing guessed from a default
    assert next(m for m in an.mapped if m.account.name.startswith("Emirates NBD")).line == "A_CASH"
    assert [c.id for c in an.checks if c.status == "FAIL"] == []


def test_nothing_is_counted_until_confirmed_and_the_four_expected_items_are_proposed(tb_path):
    an = run_analysis(project(tb_path))
    assert an.comp.A == PROFIT and an.comp.E == PROFIT                                  # proposed items are NOT in the computation
    assert {a.adj_id.split(":")[0] for a in an.adjs} == {"d_participating", "a_entertainment", "a_donations", "a_non_business", "a_fines"}
    assert all(a.status == "Proposed" for a in an.adjs)


def test_all_confirmed_matches_the_hand_calculation(tb_path):
    an = confirm_all(project(tb_path))
    assert an.comp.A == PROFIT and an.comp.E == TAXABLE and an.comp.payable == TAX
    amounts = {a.adj_id.split(":")[0]: a.amount for a in an.adjs}
    assert amounts == {"d_participating": 400_000, "a_entertainment": 70_000, "a_donations": 75_000, "a_non_business": 42_000, "a_fines": 28_000}


@pytest.mark.parametrize("over, taxable", [
    ({"entertainment_staff_pct": 50.0}, 2_770_000.0),                     # only the client half is restricted: 35,000
    ({"fines_compensation": 8_000.0}, 2_797_000.0),                       # compensation for damages is not disallowed
    ({"donations_qpbe": "All to qualifying entities"}, 2_730_000.0),
    ({"dividends": [{"account": "4200", "payer": "X", "jurisdiction": "Free Zone person", "ownership": 100, "cost": 1_500_000, "months": 6,
                     "subject_tax": "Yes"}]}, 3_205_000.0),               # holding period not met: dividend stays taxable
    ({"opening_losses": 1_000_000.0}, 1_805_000.0),
])
def test_answers_change_the_computation_as_calculated_by_hand(tb_path, over, taxable):
    an = confirm_all(project(tb_path, **over))
    assert an.comp.E == pytest.approx(taxable, abs=0.01)
    assert an.comp.payable == pytest.approx(round(0.09 * max(taxable - 375_000, 0), 2), abs=0.01)


def test_realisation_basis_scope_a_deducts_the_fair_value_gain_and_adds_back_the_impairment(tb_path):
    an = confirm_all(project(tb_path, realisation_basis=RB_OPTIONS[1], first_period="First tax period", first_period_start="2025-01-01"))
    assert an.comp.E == pytest.approx(TAXABLE - 120_000 + 95_000, abs=0.01)


@LO
def test_workbook_formulas_recalculate_to_the_same_figures(tb_path):
    import tempfile
    from pathlib import Path
    from ct_tool.excel_out import ClientWorkbook
    an = confirm_all(project(tb_path))
    out = Path(tempfile.mkdtemp()) / "c.xlsx"
    ClientWorkbook(an, str(TEMPLATE)).build(str(out))
    calc = V.recalc(out)
    assert V.errors(calc) == []
    ws = V.values(calc)["Calculation of CT"]
    assert ws["E11"].value == pytest.approx(PROFIT) and ws["E63"].value == pytest.approx(TAXABLE) and ws["E72"].value == pytest.approx(TAX)


def test_post_closing_tb_is_refused_loudly_not_computed_as_nil(tb_path):
    p = project(tb_path)
    profit = sum(a.net for a in p.tb.accounts if a.code[:1] in "456")
    p.tb.accounts = [a for a in p.tb.accounts if a.code[:1] in "123"]
    re_acc = next(a for a in p.tb.accounts if a.code == "3100")
    re_acc.net, re_acc.credit = re_acc.net + profit, -(re_acc.net + profit)
    an = run_analysis(p)
    assert any(c.id == "FS-04" and c.status == "FAIL" for c in an.checks)


# ------------------------------------------------------------------ file-reading bugs found in the audit
def _write_csv(path, style, delim, title=True):
    def fmt(v):
        a = f"{abs(v):,.2f}"
        return {"us": a, "eu": a.replace(",", "X").replace(".", ",").replace("X", "."), "space": a.replace(",", " ").replace(".", ","),
                "swiss": a.replace(",", "'")}[style]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=delim)
        if title:
            w.writerow(["Marjan Gulf Trading LLC (sample)"])
            w.writerow(["Trial balance 31 December 2025"])
            w.writerow([])
        w.writerow(["Account Code", "Account Name", "Debit", "Credit"])
        for c, n, v in M._ordered(round(-sum(x for _, _, x in M.LINES), 2)):
            w.writerow([c, n, fmt(v) if v >= 0 else "", fmt(v) if v < 0 else ""])


@pytest.mark.parametrize("style, delim", [("us", ","), ("us", ";"), ("us", "\t"), ("eu", ";"), ("space", ";"), ("swiss", ",")])
def test_csv_with_title_rows_and_any_number_style_reads_exactly(tmp_path, style, delim):
    p = tmp_path / "tb.csv"
    _write_csv(p, style, delim)
    tb = parse_tb(str(p), currency="AED")
    assert len(tb.accounts) == 46 and abs(sum(a.net for a in tb.accounts)) < 0.005
    assert -sum(a.net for a in tb.accounts if a.code[:1] in "456") == PROFIT


def test_number_parser_handles_regional_formats():
    for text, val in [("1.234,56", 1234.56), ("1 234,56", 1234.56), ("1'234.56", 1234.56), ("1,234,567", 1234567.0), ("1234,5", 1234.5), ("1.234.567", 1234567.0),
                      ("(1.234,56)", -1234.56), ("−1,234.50", -1234.5), ("EUR 1.234,56", 1234.56), ("1,234.50 Cr", -1234.5)]:
        assert to_num(text) == pytest.approx(val), text


def test_manual_header_row_override_works(tb_path):
    tb = parse_tb(tb_path, "Trial Balance", "AED", header_row=4)             # row 5 in Excel numbering
    assert len(tb.accounts) == 46
    with pytest.raises(ValueError, match="does not look like a heading row"):
        parse_tb(tb_path, "Trial Balance", "AED", header_row=1)


def test_damaged_or_empty_files_give_a_plain_message(tmp_path):
    (tmp_path / "bad.xlsx").write_bytes(b"not a workbook")
    (tmp_path / "empty.csv").write_text("")
    for name in ("bad.xlsx", "empty.csv"):
        with pytest.raises(ValueError):
            parse_tb(str(tmp_path / name), currency="AED")


def test_file_names_are_valid_on_windows_and_keep_arabic_names():
    assert safe_name("Marjan Gulf Trading LLC (sample)") == "Marjan_Gulf_Trading_LLC_sample"
    assert safe_name("شركة مرجان") == "شركة_مرجان"
    assert safe_name("CON") == "_CON" and safe_name("a/b:c*?") == "a_b_c" and safe_name("") == "Entity"


def test_review_workbook_shows_accepted_answers_as_confirmed(tb_path, tmp_path):
    from openpyxl import load_workbook
    from ct_tool.review_xlsx import ReviewWorkbook
    p = project(tb_path)
    p.answers.pop("mne_group")
    p.answers.setdefault("_meta", {})["mne_group"] = {"source": "default", "confirmed": True, "note": "accepted by user"}
    an = run_analysis(p)
    ReviewWorkbook(an).build(str(tmp_path / "r.xlsx"))
    ws = load_workbook(tmp_path / "r.xlsx")["Questionnaire"]
    rows = {str(r[1].value): str(r[3].value) for r in ws.iter_rows() if r[1].value}
    q = next(v for k, v in rows.items() if "Multinational" in k)
    assert q.startswith("Answered")
