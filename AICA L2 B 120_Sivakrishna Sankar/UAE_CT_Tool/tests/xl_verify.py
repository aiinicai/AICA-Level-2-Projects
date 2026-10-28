"""Recalculate a workbook with LibreOffice (headless) and read the calculated values back.
openpyxl writes formulas but no cached values, so this is how formula results are verified."""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl


def soffice_path():
    return shutil.which("soffice") or shutil.which("libreoffice")


def recalc(path) -> Path:
    exe = soffice_path()
    if not exe:
        raise RuntimeError("LibreOffice not available")
    out = Path(tempfile.mkdtemp(prefix="recalc_"))
    subprocess.run([exe, "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(path)], check=True, capture_output=True, timeout=180)
    return out / Path(path).name


def errors(path_calculated) -> list:
    wb = openpyxl.load_workbook(path_calculated, data_only=True)
    bad = []
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:")):
                    bad.append((ws.title, c.coordinate, c.value))
    return bad


def values(path_calculated):
    return openpyxl.load_workbook(path_calculated, data_only=True)
