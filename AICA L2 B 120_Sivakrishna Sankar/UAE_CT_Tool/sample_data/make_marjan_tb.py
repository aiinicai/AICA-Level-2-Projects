"""Builds ONE synthetic trial balance for a fictional company: 'Marjan Gulf Trading LLC (sample)', AED, year ended 31 December 2025.
Fictional - not a real entity. Designed to be read by the CT tool without manual fixes and to contain typical disallowed / adjustable items.

Run:  python sample_data/make_marjan_tb.py [output_dir]
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ENTITY = "Marjan Gulf Trading LLC (sample)"

# (code, name, balance: debit +, credit -)   Retained earnings (opening) is the balancing figure and is added below.
LINES = [
    # ---- assets
    ("1010", "Cash on hand", 15_000.00),
    ("1020", "Emirates NBD - current account", 2_850_000.00),
    ("1100", "Trade receivables", 3_200_000.00),
    ("1110", "Allowance for doubtful debts", -180_000.00),
    ("1200", "Due from related party - Marjan Holdings LLC", 750_000.00),
    ("1300", "Prepayments and deposits", 240_000.00),
    ("1400", "Investment in subsidiary - Marjan Logistics FZE", 1_500_000.00),
    ("1410", "Investments at fair value through profit or loss - quoted equities", 900_000.00),
    ("1500", "Furniture and equipment - cost", 1_400_000.00),
    ("1510", "Accumulated depreciation", -560_000.00),
    ("1600", "Inventory", 2_100_000.00),
    # ---- liabilities
    ("2000", "Trade payables", -1_900_000.00),
    ("2100", "Accrued expenses", -320_000.00),
    ("2200", "Loan from shareholder - Marjan Holdings LLC", -1_800_000.00),
    ("2300", "Bank term loan", -1_000_000.00),
    ("2400", "Provision for employees' end of service benefits", -310_000.00),
    ("2500", "VAT payable", -120_000.00),
    # ---- equity
    ("3000", "Share capital", -1_000_000.00),
    # ---- income
    ("4000", "Sales - goods", -16_500_000.00),
    ("4100", "Service revenue", -3_000_000.00),
    ("4200", "Dividend income - Marjan Logistics FZE", -400_000.00),
    ("4300", "Unrealised gain on fair value of investments", -120_000.00),
    ("4400", "Interest income - bank deposits", -35_000.00),
    ("4500", "Foreign exchange gain", -22_000.00),
    # ---- expenses
    ("5000", "Cost of goods sold", 11_800_000.00),
    ("6000", "Salaries and wages", 2_900_000.00),
    ("6010", "Office rent", 480_000.00),
    ("6020", "Client entertainment and hospitality", 140_000.00),
    ("6030", "Donations", 75_000.00),
    ("6040", "Fines and penalties", 28_000.00),
    ("6050", "Professional and legal fees", 160_000.00),
    ("6060", "Depreciation expense", 210_000.00),
    ("6070", "Interest expense - shareholder loan", 90_000.00),
    ("6080", "Interest expense - bank term loan", 60_000.00),
    ("6085", "Bank charges", 24_000.00),
    ("6090", "Impairment of trade receivables - provision for doubtful debts", 95_000.00),
    ("6100", "Repairs and maintenance", 70_000.00),
    ("6110", "Staff end of service benefit expense", 85_000.00),
    ("6120", "Management fee - Marjan Holdings LLC", 300_000.00),
    ("6130", "Foreign exchange loss", 18_000.00),
    ("6140", "Director personal expenses", 42_000.00),
    ("6150", "Marketing and advertising", 220_000.00),
    ("6160", "Telephone and utilities", 96_000.00),
    ("6170", "Insurance", 64_000.00),
    ("6180", "Travel and accommodation", 130_000.00),
]
RE = ("3100", "Retained earnings - opening")


def lines():
    total = round(sum(v for _, _, v in LINES), 2)
    re_bal = round(-total, 2)                       # credit (negative) balancing figure
    assert re_bal < 0, "retained earnings should be a credit balance"
    return _ordered(re_bal)


def _ordered(re_bal):
    out = []
    for c, n, v in LINES:
        out.append((c, n, v))
        if c == "3000":
            out.append((RE[0], RE[1], re_bal))
    return out


def write(path: Path):
    tb = lines()
    assert abs(sum(v for _, _, v in tb)) < 0.005
    wb = Workbook()
    ws = wb.active
    ws.title = "Trial Balance"
    ws["A1"] = ENTITY
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Trial Balance as at 31 December 2025"
    ws["A3"] = "Year ended 31 December 2025   |   Currency: UAE Dirham (AED)   |   Status: draft, unaudited"
    hdr = ["Account Code", "Account Name", "Debit", "Credit"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(5, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="283A4A")
        c.alignment = Alignment(horizontal="center" if i > 2 else "left")
    r, td, tc = 6, 0.0, 0.0
    for code, name, v in tb:
        ws.cell(r, 1, code)
        ws.cell(r, 2, name)
        if v >= 0:
            ws.cell(r, 3, v)
            td += v
        else:
            ws.cell(r, 4, -v)
            tc += -v
        for col in (3, 4):
            ws.cell(r, col).number_format = "#,##0.00"
        r += 1
    ws.cell(r, 2, "Total").font = Font(bold=True)
    for col, val in ((3, td), (4, tc)):
        c = ws.cell(r, col, round(val, 2))
        c.number_format = "#,##0.00"
        c.font = Font(bold=True)
    for k, w in zip("ABCD", (14, 66, 18, 18)):
        ws.column_dimensions[k].width = w
    wb.save(path)
    return tb, td, tc


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "marjan"
    out.mkdir(parents=True, exist_ok=True)
    tb, td, tc = write(out / "Marjan_Gulf_Trading_TB_31Dec2025.xlsx")
    pl = sum(v for c, _, v in tb if c[0] in "456")
    print(f"written; debits {td:,.2f} credits {tc:,.2f}; profit before tax {-pl:,.2f}; opening retained earnings {[v for c,_,v in tb if c=='3100'][0]:,.2f}")
