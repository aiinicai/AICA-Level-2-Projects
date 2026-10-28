"""Shared Excel styling that matches the supplied computation template (Arial 10, navy headers,
amber banners, accounting number formats) plus colour conventions for review status."""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY, AMBER, GREEN, ORANGE = "283A4A", "FDC50A", "92D050", "FF3300"
PROPOSED_FILL, CONFIRMED_FILL, INFO_FILL, REJECT_FILL = "FFF2CC", "E2EFDA", "FCE4D6", "EDEDED"
INPUT_BLUE = "0000FF"

NUM0 = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
NUM2 = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'
RATE = "0.000000"
PCT = "0.0%"

thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TOPBOT = Border(top=thin, bottom=thin)


def font(bold=False, color=None, size=10, italic=False):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def put(ws, ref, value=None, bold=False, color=None, fmt=None, fill_=None, align=None, wrap=False, border=None, size=10,
        italic=False, valign="center"):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font(bold, color, size, italic)
    if fmt:
        c.number_format = fmt
    if fill_:
        c.fill = fill(fill_)
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    if border:
        c.border = border
    return c


def header_row(ws, row, labels, start_col=1, height=None):
    for i, lab in enumerate(labels):
        ref = f"{get_column_letter(start_col + i)}{row}"
        put(ws, ref, lab, bold=True, color="FFFFFF", fill_=NAVY, align="center", wrap=True, border=BOX)
    if height:
        ws.row_dimensions[row].height = height


def banner(ws, row, text, c1=1, c2=6):
    for c in range(c1, c2 + 1):
        ref = f"{get_column_letter(c)}{row}"
        put(ws, ref, None, bold=True, fill_=AMBER, border=BOX)
    ws[f"{get_column_letter(c1)}{row}"].value = text


def widths(ws, mapping: dict):
    for k, v in mapping.items():
        ws.column_dimensions[k].width = v


def status_fill(status: str):
    s = (status or "").lower()
    if s.startswith("confirmed"):
        return CONFIRMED_FILL
    if s.startswith("rejected"):
        return REJECT_FILL
    if "information" in s or "technical" in s:
        return INFO_FILL
    return PROPOSED_FILL


def page(ws, landscape=True, area=None, rows_repeat=None):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if area:
        ws.print_area = area
    if rows_repeat:
        ws.print_title_rows = rows_repeat
    ws.sheet_view.showGridLines = False
