"""Client workbook: the supplied computation template, populated in AED.

 * The template's worksheets, headings, fonts, colours and the layout of the 'Calculation of CT' sheet are preserved.
 * Inapplicable worksheets are removed (see engine.applicability) and listed in an applicability log on the Assumptions sheet.
 * Only CONFIRMED adjustments enter the computation; the workbook is stamped DRAFT while items are pending.
 * Calculations are live Excel formulas wherever practical (blue font = typed input, black = formula).
 * Stale content carried inside the template (about 18,800 defined names left over from unrelated files) is removed and logged.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from . import engine
from .engine import ROWS, DISALLOWED_SHEET_ROWS, ans, params, _num, _table
from .mapping import fs_lines
from .questionnaire import RB_OPTIONS, RB_CODE, YES, NO, UNK
from .xl_style import *   # noqa

CALC = "Calculation of CT"
TBS = "TB (AED)"


def fmt_date(iso: str) -> str:
    try:
        return datetime.strptime(iso[:10], "%Y-%m-%d").strftime("%d %B %Y").lstrip("0")
    except Exception:
        return iso or ""


def _rsum(col: str, first: int, r: int):
    """Total formula for rows first..r-1; a plain 0 when there are no data rows (avoids a self-referencing SUM)."""
    return f"=SUM({col}{first}:{col}{r - 1})" if r > first else 0


def _reset(ws, from_row: int = 1):
    """Clear cells, merges and validations at/after from_row (keeps column widths and sheet setup)."""
    for mr in list(ws.merged_cells.ranges):
        if mr.max_row >= from_row:
            ws.unmerge_cells(str(mr))
    for row in ws.iter_rows(min_row=from_row, max_row=max(ws.max_row, from_row)):
        for c in row:
            c.value = None
            c.style = "Normal"
    if ws.data_validations:
        ws.data_validations.dataValidation = []
    for r in range(from_row, ws.max_row + 1):
        ws.row_dimensions[r].height = None


def _titles(ws, col="A", title_cell_calc=True):
    ws[f"{col}1"].value = f"='{CALC}'!$C$2"
    ws[f"{col}2"].value = f"='{CALC}'!$C$3"
    put(ws, f"{col}1", bold=True, size=11)
    put(ws, f"{col}2", bold=True, size=11)


def _acct_names(accts, limit=2) -> str:
    names = [a[1] or a[0] for a in accts if (a[1] or a[0])]
    s = "; ".join(names[:limit])
    return s + (" ..." if len(names) > limit else "")


class ClientWorkbook:
    def __init__(self, an, template_path: str):
        self.an = an
        self.wb = openpyxl.load_workbook(template_path)
        self.log = []
        self.P = params()
        self.conf = [a for a in an.adjs if a.status == "Confirmed"]
        self.applic = an.applic
        self.cell = {}                       # logical name -> "'Sheet'!$X$n"
        self.adj_last = 0

    # ----------------------------------------------------------------- helpers
    def sheet(self, name):
        return self.wb[name]

    def _hygiene(self):
        n = len(self.wb.defined_names)
        for ws in self.wb.worksheets:
            n += len(ws.defined_names)
            for k in list(ws.defined_names.keys()):
                del ws.defined_names[k]
        for k in list(self.wb.defined_names.keys()):
            del self.wb.defined_names[k]
        self.log.append(f"Removed {n:,} stale defined names carried inside the template (leftovers from unrelated workbooks, mostly #REF!).")

    # ----------------------------------------------------------------- main
    def build(self, out_path: str) -> list:
        an = self.an
        self._hygiene()
        keep = {n for n, (inc, _) in self.applic.items() if inc}
        # template sheets that are never kept as-is
        for nm in ("Notes",):
            if nm in self.wb.sheetnames:
                del self.wb[nm]
        for nm in ("UGUL", "Provision", "DTA & DTL Working", "Tax Loss Carried Forward", "Disallowed Expenses", "Freezone", "Taxes Paid Abroad"):
            if nm not in keep and nm in self.wb.sheetnames:
                del self.wb[nm]
                self.log.append(f"Sheet '{nm}' omitted: {self.applic[nm][1]}.")
        # TB sheet is rebuilt in place
        tbws = self.wb["TB 2025"]
        tbws.title = TBS
        self._write_tb(tbws)
        if "Adjustments" in keep:
            self._write_adjustments()
        self._write_pl_bs()
        if "FX Rates" in keep:
            self._write_fx()
        if "Disallowed Expenses" in keep:
            self._write_disallowed()
        if "UGUL" in keep:
            self._write_ugul()
        if "Provision" in keep:
            self._write_provision()
        if "Freezone" in keep:
            self._write_freezone()
        if "Tax Loss Carried Forward" in keep:
            self._write_losses()
        if "Taxes Paid Abroad" in keep:
            self._write_foreign_tax()
        if "DTA & DTL Working" in keep:
            self._write_dta()
        self._write_assumptions()
        self._write_calc()
        self._order_sheets()
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.properties.title = f"{an.entity} - Corporate Tax computation (AED)"
        self.wb.properties.creator = an.project.settings.preparer or "UAE CT Computation Tool"
        Path(out_path).parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(out_path)
        return self.log

    def _order_sheets(self):
        order = ["Assumptions", CALC, "Profit and Loss", "Balance Sheet", "Adjustments", "Disallowed Expenses", "Freezone", "UGUL", "Provision",
                 "Tax Loss Carried Forward", "Taxes Paid Abroad", "DTA & DTL Working", "FX Rates", TBS]
        self.wb._sheets = [self.wb[n] for n in order if n in self.wb.sheetnames]
        self.wb.active = 0

    # ----------------------------------------------------------------- TB (AED)
    def _write_tb(self, ws):
        an = self.an
        _reset(ws)
        for col in "ABCDEFGHIJKLMNO":
            ws.column_dimensions[col].hidden = False
        widths(ws, {"A": 12, "B": 46, "C": 14, "D": 13, "E": 36, "F": 12, "G": 16, "H": 16, "I": 16, "J": 12, "K": 24, "L": 18, "M": 16, "N": 16, "O": 46, "P": 4})
        ccy = an.project.fx.tb_currency
        put(ws, "A1", f"='{CALC}'!$C$2", bold=True, size=11)
        put(ws, "A2", f"Trial balance - source currency ({ccy}) and AED translation", bold=True, size=11)
        put(ws, "A3", "Blue = typed from the source TB; black = formula. Mapped line keys drive the P&L and balance sheet (SUMIFS on column D).", italic=True, size=9)
        hdr = ["Code", "Account", "Source class.", "Line key", "Mapped to", "Method", f"Debit ({ccy})", f"Credit ({ccy})", f"Net ({ccy})", "Rate to AED",
               "Rate basis", "Net (AED)", "Debit (AED)", "Credit (AED)", "Review flags"]
        header_row(ws, 5, hdr, height=30)
        r = 6
        lines = fs_lines()
        gl_based = False
        for a in an.tr.accounts:
            acc = a.account
            put(ws, f"A{r}", acc.code, color=INPUT_BLUE)
            put(ws, f"B{r}", acc.name, color=INPUT_BLUE)
            put(ws, f"C{r}", acc.hint or "", color=INPUT_BLUE)
            put(ws, f"D{r}", a.line)
            put(ws, f"E{r}", lines[a.line]["label"])
            put(ws, f"F{r}", a.mapped.method)
            put(ws, f"G{r}", acc.debit, color=INPUT_BLUE, fmt=NUM2)
            put(ws, f"H{r}", acc.credit, color=INPUT_BLUE, fmt=NUM2)
            put(ws, f"I{r}", f"=G{r}-H{r}" if abs(acc.debit - acc.credit - acc.net) < 0.005 else acc.net, fmt=NUM2,
                color=None if abs(acc.debit - acc.credit - acc.net) < 0.005 else INPUT_BLUE)
            put(ws, f"K{r}", a.rate_basis)
            if "GL" in a.rate_basis:
                gl_based = True
                put(ws, f"J{r}", a.rate, fmt=RATE, color=INPUT_BLUE)
                put(ws, f"L{r}", a.net_aed, fmt=NUM2, color=INPUT_BLUE)
            else:
                put(ws, f"J{r}", a.rate if an.project.fx.is_aed else None, fmt=RATE, color=INPUT_BLUE) if an.project.fx.is_aed else None
                put(ws, f"L{r}", f"=I{r}*J{r}", fmt=NUM2)
            put(ws, f"M{r}", f"=MAX(L{r},0)", fmt=NUM2)
            put(ws, f"N{r}", f"=MAX(-L{r},0)", fmt=NUM2)
            flags = list(a.mapped.flags) + list(a.flags) + list(acc.notes)
            put(ws, f"O{r}", " | ".join(flags), size=9, wrap=False)
            for col in "ABCDEFGHIJKLMNO":
                ws[f"{col}{r}"].border = BOX
            r += 1
        last = r - 1
        # rate cells: link to the FX sheet where applicable (set after FX sheet is built) - store rows to patch
        self._tb_first, self._tb_last = 6, last
        derived_row = None
        if an.tr.diff_applicable:
            derived_row = r
            put(ws, f"B{r}", "Currency translation difference (derived - not in the source TB; requires approval)", bold=True, fill_=INFO_FILL)
            put(ws, f"D{r}", "DERIVED_CTR")
            put(ws, f"E{r}", "Currency translation difference (derived)")
            put(ws, f"L{r}", f"=-SUM(L6:L{last})", fmt=NUM2, bold=True)
            put(ws, f"M{r}", f"=MAX(L{r},0)", fmt=NUM2)
            put(ws, f"N{r}", f"=MAX(-L{r},0)", fmt=NUM2)
            put(ws, f"O{r}", "Arises because P&L, balance sheet and equity are translated at different rates (see FX Rates).", size=9)
            for col in "ABCDEFGHIJKLMNO":
                ws[f"{col}{r}"].border = BOX
            r += 1
        put(ws, f"B{r}", "Total", bold=True)
        for col in "GHILMN":
            put(ws, f"{col}{r}", f"=SUM({col}6:{col}{r - 1})", bold=True, fmt=NUM2, border=TOPBOT)
        self._tb_total_row = r
        put(ws, f"B{r + 1}", "Check: AED net must be nil", italic=True)
        put(ws, f"L{r + 1}", f'=IF(ABS(L{r})<0.5,"OK","CHECK")', bold=True)
        ws.freeze_panes = "C6"
        page(ws, True, f"A1:O{r + 1}", "5:5")
        self.tb_range = (6, r - 1 if derived_row else last)
        self.tb_last_incl_derived = r - 1
        self.cell["tb_total"] = f"'{TBS}'!$L${r}"

    def _tb_sumifs(self, key: str, sign: int = 1) -> str:
        lo, hi = 6, self.tb_last_incl_derived
        f = f"SUMIFS('{TBS}'!$L${lo}:$L${hi},'{TBS}'!$D${lo}:$D${hi},\"{key}\")"
        return f"=-{f}" if sign < 0 else f"={f}"

    # ----------------------------------------------------------------- P&L and balance sheet
    def _stmt_sheet(self, name):
        if name in self.wb.sheetnames:
            del self.wb[name]
        return self.wb.create_sheet(name)

    def _write_pl_bs(self):
        an, st = self.an, self.an.st
        lines = fs_lines()
        # ---------- P&L
        ws = self._stmt_sheet("Profit and Loss")
        widths(ws, {"A": 3, "B": 58, "C": 60, "D": 20})
        put(ws, "B1", f"='{CALC}'!$C$2", bold=True, size=11)
        put(ws, "B2", f"='{CALC}'!$C$3", bold=True, size=11)
        put(ws, "B3", "Statement of profit or loss (AED) - prepared from the trial balance; income positive, expenses negative", italic=True, size=9)
        header_row(ws, 5, ["Line", "Source accounts (see TB (AED) sheet)", "AED"], start_col=2, height=22)
        r = 6
        refs = {}
        sect_cells = {"OTHER_INCOME": [], "OPEX": [], "DEPRECIATION": [], "FINANCE": [], "FX": [], "TAX": []}
        for row in st.pl_rows:
            if row.kind == "header":
                put(ws, f"B{r}", row.label, bold=True)
                r += 1
                continue
            if row.kind == "line":
                put(ws, f"B{r}", row.label, border=BOX)
                names = "; ".join((x[1] or x[0]) for x in row.accounts[:3]) + (" ..." if len(row.accounts) > 3 else "")
                put(ws, f"C{r}", names, size=9, border=BOX)
                put(ws, f"D{r}", self._tb_sumifs(row.line_key, -1), fmt=NUM0, border=BOX)
                sec = lines[row.line_key]["section"]
                if row.line_key == "P_REV":
                    refs["rev"] = f"D{r}"
                elif row.line_key == "P_COS":
                    refs["cos"] = f"D{r}"
                if sec in sect_cells:
                    sect_cells[sec].append(f"D{r}")
                r += 1
                continue
            lab = row.label
            if lab == "Gross profit":
                f = f"={refs['rev']}+{refs['cos']}"
                refs["gross"] = f"D{r}"
            elif lab == "Total other income":
                f = "=" + ("+".join(sect_cells["OTHER_INCOME"]) or "0")
                refs["oi"] = f"D{r}"
            elif lab == "Total operating expenses":
                f = "=" + ("+".join(sect_cells["OPEX"]) or "0")
                refs["opex"] = f"D{r}"
            elif lab.startswith("Profit/(loss) before tax"):
                parts = [refs["gross"]] + ([refs["oi"]] if "oi" in refs else []) + ([refs["opex"]] if "opex" in refs else []) + \
                        sect_cells["DEPRECIATION"] + sect_cells["FINANCE"] + sect_cells["FX"]
                f = "=" + "+".join(parts)
                refs["pbt"] = f"D{r}"
            else:   # profit for the year
                f = "=" + "+".join([refs["pbt"]] + sect_cells["TAX"])
                refs["pfy"] = f"D{r}"
            put(ws, f"B{r}", lab, bold=True, fill_="D9E1F2" if row.kind == "total" else None)
            put(ws, f"D{r}", f, bold=True, fmt=NUM0, border=TOPBOT, fill_="D9E1F2" if row.kind == "total" else None)
            r += 1
        if st.unmapped_accounts:
            put(ws, f"B{r + 1}", "Unmapped accounts are excluded from this statement - see TB (AED).", color="C00000", bold=True)
        self.cell["profit"] = f"'Profit and Loss'!$D${refs['pfy'].lstrip('D')}"
        self.cell["pbt"] = f"'Profit and Loss'!{refs['pbt']}"
        self.pl_refs = refs
        page(ws, False, f"B1:D{r + 2}")
        # ---------- Balance sheet
        ws = self._stmt_sheet("Balance Sheet")
        widths(ws, {"A": 3, "B": 66, "C": 60, "D": 20})
        put(ws, "B1", f"='{CALC}'!$C$2", bold=True, size=11)
        put(ws, "B2", "Statement of financial position (AED) - at " + fmt_date(an.period_end), bold=True, size=11)
        put(ws, "B3", "Assets are not split into current / non-current: the TB does not support that classification.", italic=True, size=9)
        header_row(ws, 5, ["Line", "Source accounts (see TB (AED) sheet)", "AED"], start_col=2, height=22)
        r = 6
        asset_cells, liab_cells, eq_cells = [], [], []
        cur = None
        tot = {}
        for row in st.bs_rows:
            if row.kind == "header":
                put(ws, f"B{r}", row.label, bold=True, fill_="D9E1F2")
                put(ws, f"C{r}", None, fill_="D9E1F2")
                put(ws, f"D{r}", None, fill_="D9E1F2")
                cur = row.label
                r += 1
                continue
            if row.kind == "note":
                put(ws, f"B{r}", row.label, color="C00000", bold=True)
                put(ws, f"D{r}", self._tb_sumifs("UNMAPPED", 1), fmt=NUM0)
                r += 1
                continue
            if row.kind == "line":
                put(ws, f"B{r}", row.label, border=BOX)
                names = "; ".join((x[1] or x[0]) for x in row.accounts[:3]) + (" ..." if len(row.accounts) > 3 else "")
                put(ws, f"C{r}", names, size=9, border=BOX)
                if row.line_key == "PROFIT_YEAR":
                    f = f"={self.cell['profit']}"
                elif row.line_key == "DERIVED_CTR":
                    f = self._tb_sumifs("DERIVED_CTR", -1)
                else:
                    sec = lines[row.line_key]["section"]
                    f = self._tb_sumifs(row.line_key, 1 if sec == "ASSETS" else -1)
                put(ws, f"D{r}", f, fmt=NUM0, border=BOX, fill_=INFO_FILL if row.line_key == "DERIVED_CTR" else None)
                (asset_cells if cur == "ASSETS" else liab_cells if cur == "LIABILITIES" else eq_cells).append(f"D{r}")
                r += 1
                continue
            lab = row.label
            if lab == "Total assets":
                f = "=" + "+".join(asset_cells); tot["ta"] = f"D{r}"
            elif lab == "Total liabilities":
                f = "=" + ("+".join(liab_cells) or "0"); tot["tl"] = f"D{r}"
            elif lab == "Total equity":
                f = "=" + "+".join(eq_cells); tot["te"] = f"D{r}"
            else:
                f = f"={tot['tl']}+{tot['te']}"; tot["tle"] = f"D{r}"
            put(ws, f"B{r}", lab, bold=True)
            put(ws, f"D{r}", f, bold=True, fmt=NUM0, border=TOPBOT)
            r += 1
        r += 1
        put(ws, f"B{r}", "Check: total assets less total liabilities and equity", italic=True)
        put(ws, f"D{r}", f"={tot['ta']}-{tot['tle']}", fmt=NUM2)
        unm = round(an.st.unmapped_net, 2) if an.st.unmapped_accounts else 0
        extra = f' - difference equals the unmapped accounts ({unm:,.2f}); map them' if unm and abs(abs(an.st.bs_difference) - abs(unm)) < 1 else ""
        put(ws, f"C{r}", f'=IF(ABS(D{r})<0.5,"Balances","DOES NOT BALANCE - review{extra}")', bold=True)
        if an.tr.diff_applicable:
            put(ws, f"B{r + 1}", "The 'Currency translation difference' line is derived (not a TB figure) and requires approval - see FX Rates.", color="C00000", size=9)
        self.cell["bs_check"] = f"'Balance Sheet'!$D${r}"
        page(ws, False, f"B1:D{r + 2}")

    # ----------------------------------------------------------------- FX
    def _write_fx(self):
        an = self.an
        fx, rec = an.project.fx, an.tr.method_record
        if "FX Rates" in self.wb.sheetnames:
            del self.wb["FX Rates"]
        ws = self.wb.create_sheet("FX Rates")
        widths(ws, {"A": 3, "B": 44, "C": 24, "D": 70})
        put(ws, "B1", f"='{CALC}'!$C$2", bold=True, size=11)
        put(ws, "B2", "Conversion of the trial balance to UAE dirhams (AED)", bold=True, size=11)
        banner(ws, 4, "Conversion record (FTA Decision No. 13 of 2023, Art. 2)", 2, 4)
        rows = [("TB / source currency", fx.tb_currency), ("Functional currency (as stated)", fx.functional_currency or fx.tb_currency),
                ("Method", rec.get("method_description", "")), ("Reason for the method", fx.method_reason or "NOT RECORDED - complete before filing"),
                ("Rate source", fx.rate_source or "NOT RECORDED - complete before filing"),
                ("Used GL transactions for P&L conversion", "Yes" if rec.get("gl_used") else "No - single P&L rate applied")]
        r = 5
        for k, v in rows:
            put(ws, f"B{r}", k, bold=True, border=BOX)
            put(ws, f"C{r}", v, border=BOX, wrap=True, color=INPUT_BLUE)
            ws.merge_cells(f"C{r}:D{r}")
            r += 1
        r += 1
        banner(ws, r, "Rates applied", 2, 4)
        r += 1
        header_row(ws, r, ["Purpose", "Rate (AED per 1 " + fx.tb_currency + ")", "Basis"], start_col=2)
        r += 1
        first_rate_row = r
        for row in an.tr.rate_table:
            put(ws, f"B{r}", row["purpose"], border=BOX)
            put(ws, f"C{r}", row["rate"], fmt=RATE, color=INPUT_BLUE, border=BOX)
            put(ws, f"D{r}", row["basis"], border=BOX, size=9)
            r += 1
        self.cell["fx_pl_rate"] = f"'FX Rates'!$C${first_rate_row}"
        self.cell["fx_close_rate"] = f"'FX Rates'!$C${first_rate_row + 1}"
        self.cell["fx_eq_rate"] = f"'FX Rates'!$C${first_rate_row + 2}"
        r += 1
        put(ws, f"B{r}", "Currency translation difference (derived)", bold=True)
        put(ws, f"C{r}", self._tb_sumifs("DERIVED_CTR", 1) if an.tr.diff_applicable else 0, fmt=NUM2)
        put(ws, f"D{r}", ("Derived by formula as the balancing amount in equity; it arises from using different rates for P&L, balance sheet and equity. "
                          "It is NOT a TB figure and needs approval." if an.tr.diff_applicable else "Nil - a single rate (or the peg) applies to all lines."), size=9, wrap=True)
        ws.row_dimensions[r].height = 40
        # patch TB rate cells so that the rates are linked, not typed twice
        tb = self.wb[TBS]
        if not fx.is_aed:
            for rr in range(6, self._tb_last + 1):
                basis = tb[f"K{rr}"].value or ""
                if "GL" in basis:
                    continue
                key = {"P&L rate": "fx_pl_rate", "Closing rate": "fx_close_rate", "Historical equity rate": "fx_eq_rate"}.get(basis)
                if key:
                    put(tb, f"J{rr}", "=" + self.cell[key], fmt=RATE)
        else:
            pass
        page(ws, True)

    # ----------------------------------------------------------------- Adjustments (confirmed)
    def _write_adjustments(self):
        if "Adjustments" in self.wb.sheetnames:
            del self.wb["Adjustments"]
        ws = self.wb.create_sheet("Adjustments")
        widths(ws, {"A": 30, "B": 30, "C": 8, "D": 54, "E": 18, "F": 36, "G": 60, "H": 10, "I": 28, "J": 28, "K": 18})
        put(ws, "A1", f"='{CALC}'!$C$2", bold=True, size=11)
        put(ws, "A2", f"='{CALC}'!$C$3", bold=True, size=11)
        banner(ws, 4, "Schedule of confirmed tax adjustments (AED) - feeds the Calculation of CT sheet", 1, 11)
        header_row(ws, 5, ["Ref", "Row key", "Calc. row", "Description", "Amount (AED)", "Source account(s)", "Basis / reason", "Rule ref", "Supporting reference",
                           "Management response", "Reviewer"], height=30)
        r = 6
        for a in sorted(self.conf, key=lambda x: ROWS[x.row][0]):
            vals = [a.adj_id, a.row, ROWS[a.row][0], a.description, a.amount, _acct_names(a.accounts, 3), a.reason, a.rule_ref,
                    a.support_ref or "; ".join(a.docs_needed[:2]), a.mgmt_response, a.reviewer + (f" [override: {a.override_note}]" if a.override_note else "")]
            for i, v in enumerate(vals):
                put(ws, f"{get_column_letter(i + 1)}{r}", v, border=BOX, wrap=(i in (3, 5, 6)), fmt=NUM0 if i == 4 else None, size=9 if i in (0, 1, 6, 7) else 10,
                    color=INPUT_BLUE if i == 4 else None, fill_=CONFIRMED_FILL if i == 0 else None)
            r += 1
        self.adj_last = max(r - 1, 6)
        put(ws, f"D{r}", "Total (Σ with sign by row type - see Calculation of CT)", bold=True)
        ws.freeze_panes = "A6"
        page(ws, True, f"A1:K{r}", "5:5")

    def _adj_sumif(self, key: str):
        if not any(a.row == key for a in self.conf):
            return 0
        return f"=SUMIF(Adjustments!$B$6:$B${self.adj_last},\"{key}\",Adjustments!$E$6:$E${self.adj_last})"

    # ----------------------------------------------------------------- Disallowed expenses
    def _write_disallowed(self):
        ws = self.wb["Disallowed Expenses"]
        _reset(ws, 3)
        ws.column_dimensions["B"].hidden = False
        widths(ws, {"A": 12, "B": 34, "C": 40, "D": 18, "E": 18, "F": 34})
        _titles(ws)
        put(ws, "A4", "Details of Non-Deductible Expenses (confirmed adjustments)", bold=True, fill_=AMBER)
        for c in "BCDEF":
            put(ws, f"{c}4", None, fill_=AMBER)
        r = 6
        totals = []
        for key in DISALLOWED_SHEET_ROWS:
            items = [a for a in self.conf if a.row == key]
            if not items:
                continue
            put(ws, f"A{r}", ROWS[key][2], bold=True, wrap=True, fill_=AMBER)
            for c in "BCDEF":
                put(ws, f"{c}{r}", None, fill_=AMBER)
            ws.merge_cells(f"A{r}:F{r}")
            r += 1
            header_row(ws, r, ["Account ID", "G/L account", "Description", "Disallowed (AED)", "Treatment", "Reference"])
            r += 1
            first = r
            for a in items:
                code = a.accounts[0][0] if a.accounts else ""
                put(ws, f"A{r}", code, border=BOX)
                put(ws, f"B{r}", _acct_names(a.accounts, 2), border=BOX)
                put(ws, f"C{r}", a.description, border=BOX, wrap=True)
                put(ws, f"D{r}", a.amount, fmt=NUM0, border=BOX, color=INPUT_BLUE)
                put(ws, f"E{r}", {"a_entertainment": "50% disallowed", "a_donations": "100% disallowed", "a_fines": "100% disallowed"}.get(key, "Disallowed"), border=BOX)
                put(ws, f"F{r}", a.rule_ref + (" / " + a.support_ref if a.support_ref else ""), border=BOX, size=9)
                r += 1
            put(ws, f"A{r}", "Sub total", bold=True)
            put(ws, f"D{r}", _rsum("D", first, r), bold=True, fmt=NUM0, border=TOPBOT)
            totals.append(f"D{r}")
            r += 2
        put(ws, f"A{r}", "Total disallowed expenses on this sheet", bold=True)
        put(ws, f"D{r}", "=" + "+".join(totals), bold=True, fmt=NUM0, border=TOPBOT)
        keys = {a.row for a in self.conf if a.row in DISALLOWED_SHEET_ROWS}
        calc_refs = "+".join([f"'{CALC}'!E{ROWS[k][0]}" for k in sorted(keys, key=lambda k: ROWS[k][0])])
        put(ws, f"A{r + 1}", "Check: agrees to Calculation of CT", italic=True)
        put(ws, f"D{r + 1}", f'=IF(ABS(D{r}-({calc_refs}))<0.5,"OK","CHECK")', bold=True)
        page(ws, False, f"A1:F{r + 1}")

    # ----------------------------------------------------------------- UGUL
    def _write_ugul(self):
        ws = self.wb["UGUL"]
        _reset(ws, 3)
        widths(ws, {"A": 47, "B": 18, "C": 24, "D": 22, "E": 24, "F": 24, "G": 4, "H": 4})
        put(ws, "A1", f"='{CALC}'!$C$2", bold=True, size=11)
        put(ws, "A2", f"='{CALC}'!$C$3", bold=True, size=11)
        banner(ws, 4, "Unrealised gains / (losses) of previous tax periods realised in the current tax period (confirmed)", 1, 6)
        header_row(ws, 5, ["Description of asset/liability", "Date of realisation event", "Tax period(s) in which the gain/(loss) was deferred",
                           "Previously recognised through the income statement?", "Amount of previously unrealised gains realised (AED)",
                           "Amount of previously unrealised losses realised (AED)"], height=52)
        r = 6
        rows = _table(self.an.answers, "prev_unrealised")
        conf_ids = {a.adj_id for a in self.conf}
        first = r
        for i, row in enumerate(rows):
            g, l = _num(row.get("gain")), _num(row.get("loss"))
            gconf = f"o_prev_gain:prev{i}" in conf_ids
            lconf = f"o_prev_loss:prevl{i}" in conf_ids
            put(ws, f"A{r}", row.get("asset", ""), border=BOX, color=INPUT_BLUE)
            put(ws, f"B{r}", row.get("date", ""), border=BOX, color=INPUT_BLUE)
            put(ws, f"C{r}", row.get("period", ""), border=BOX, color=INPUT_BLUE)
            put(ws, f"D{r}", row.get("through_pl", ""), border=BOX, color=INPUT_BLUE)
            put(ws, f"E{r}", g if gconf else 0, border=BOX, fmt=NUM0, color=INPUT_BLUE)
            put(ws, f"F{r}", l if lconf else 0, border=BOX, fmt=NUM0, color=INPUT_BLUE)
            r += 1
        put(ws, f"A{r}", "Total", bold=True)
        put(ws, f"E{r}", _rsum("E", first, r), bold=True, fmt=NUM0, border=TOPBOT)
        put(ws, f"F{r}", _rsum("F", first, r), bold=True, fmt=NUM0, border=TOPBOT)
        put(ws, f"A{r + 1}", "Check: agrees to Calculation of CT rows 53-54", italic=True)
        put(ws, f"E{r + 1}", f'=IF(ABS(E{r}-\'{CALC}\'!E53)+ABS(F{r}-\'{CALC}\'!E54)<0.5,"OK","CHECK")', bold=True)
        r += 3
        banner(ws, r, "Unrealised gains / (losses) of the current period deferred under the realisation basis (to be tracked until realised)", 1, 6)
        r += 1
        header_row(ws, r, ["Description of asset/liability", "Row in computation", "Treatment", "Source account(s)", "Amount deferred (AED)", "Gain / (loss)"], height=30)
        r += 1
        for a in self.conf:
            if a.category == "Realisation basis" and a.row in ("d_unrealised_capital", "d_unrealised_revenue", "a_unreal_loss_capital", "a_unreal_loss_revenue"):
                put(ws, f"A{r}", a.description, border=BOX, wrap=True)
                put(ws, f"B{r}", ROWS[a.row][0], border=BOX)
                put(ws, f"C{r}", "Deducted" if a.row.startswith("d_") else "Added back", border=BOX)
                put(ws, f"D{r}", _acct_names(a.accounts, 2), border=BOX, wrap=True)
                put(ws, f"E{r}", a.amount, border=BOX, fmt=NUM0, color=INPUT_BLUE)
                put(ws, f"F{r}", "Gain" if a.row.startswith("d_") else "Loss", border=BOX)
                r += 1
        page(ws, True, f"A1:F{r}")

    # ----------------------------------------------------------------- Provision
    def _write_provision(self):
        ws = self.wb["Provision"]
        _reset(ws, 3)
        ws.column_dimensions["C"].hidden = True
        widths(ws, {"A": 5.2, "B": 43.8, "D": 16, "E": 16.2, "F": 14.3, "G": 16, "H": 16.8})
        put(ws, "B1", f"='{CALC}'!$C$2", bold=True, size=11)
        put(ws, "B2", f"='{CALC}'!$C$3", bold=True, size=11)
        put(ws, "B3", "Movement in provisions", bold=True)
        put(ws, "H4", "(Amounts in AED)", italic=True, size=9)
        header_row(ws, 5, ["S/N", "Provision", "Ref.", "Balance at the beginning", "Provisions made", "Provisions utilised", "Adjustments / reversals",
                           "Balance at the end"], height=40)
        rows = _table(self.an.answers, "provisions")
        r = 6
        first = r
        for i, row in enumerate(rows):
            put(ws, f"A{r}", i + 1, border=BOX)
            put(ws, f"B{r}", row.get("name", ""), border=BOX, color=INPUT_BLUE)
            for col, key in (("D", "opening"), ("E", "made"), ("F", "utilised"), ("G", "adjustments")):
                put(ws, f"{col}{r}", _num(row.get(key)), border=BOX, fmt=NUM0, color=INPUT_BLUE)
            put(ws, f"H{r}", f"=D{r}+E{r}-F{r}-G{r}", border=BOX, fmt=NUM0)
            r += 1
        last = max(r - 1, first)
        put(ws, f"B{r}", "Total", bold=True)
        for col in "DEFGH":
            put(ws, f"{col}{r}", _rsum(col, first, r), bold=True, fmt=NUM0, border=TOPBOT)
        if not rows:
            put(ws, f"B{r + 1}", "No provision schedule was provided; provision accounts exist in the TB - complete the schedule.", color="C00000")
        r += 3
        put(ws, f"B{r}", "Calculation of disallowed provisions (provisions other than end-of-service / indemnity)", bold=True)
        r += 2
        header_row(ws, r, ["S/N", "Provision", "Ref.", "Provisions made", "Adjustments", "Disallowed"], height=30)
        r += 1
        f0 = r
        for i, row in enumerate(rows):
            nm = str(row.get("name", "")).lower()
            excl = any(w in nm for w in ("indemn", "end of service", "gratuity", "eosb"))
            put(ws, f"A{r}", i + 1, border=BOX)
            put(ws, f"B{r}", row.get("name", "") + (" (excluded)" if excl else ""), border=BOX)
            put(ws, f"D{r}", 0 if excl else f"=E{first + i}", border=BOX, fmt=NUM0)
            put(ws, f"E{r}", 0 if excl else f"=G{first + i}", border=BOX, fmt=NUM0)
            put(ws, f"F{r}", f"=D{r}-E{r}", border=BOX, fmt=NUM0)
            r += 1
        put(ws, f"B{r}", "Total disallowed (Option 1 / scope A)", bold=True)
        put(ws, f"F{r}", _rsum("F", f0, r), bold=True, fmt=NUM0, border=TOPBOT)
        put(ws, f"B{r + 1}", "Applied in the computation only where the realisation basis (scope A) is elected and the adjustment is confirmed.", italic=True, size=9)
        page(ws, False, f"A1:H{r + 1}")

    # ----------------------------------------------------------------- Freezone (QFZP)
    def _write_freezone(self):
        an = self.an
        ws = self.wb["Freezone"]
        _reset(ws, 3)
        ws.column_dimensions["E"].hidden = False
        widths(ws, {"A": 15.7, "B": 56, "C": 26, "D": 20, "E": 28})
        _titles(ws)
        put(ws, "A4", "Qualifying Free Zone Person - analysis of qualifying and non-qualifying income", bold=True, fill_=AMBER)
        for c in "BCDE":
            put(ws, f"{c}4", None, fill_=AMBER)
        q = an.comp.qfzp
        status = engine.qfzp_status(an.answers)
        r = 6
        put(ws, f"A{r}", "Status", bold=True)
        if an.comp.basis == "QFZP":
            put(ws, f"B{r}", "QFZP basis applied: all stated conditions are confirmed (subject to review).", color="006100", bold=True)
        else:
            put(ws, f"B{r}", "QFZP claimed but NOT applied in the computation - conditions not evidenced: " + "; ".join(q.get("missing_conditions", status["missing"])) +
                ". The ordinary basis is used; the QFZP result below is an alternative only.", color="C00000", bold=True, wrap=True)
            ws.row_dimensions[r].height = 42
        r += 2
        cls = {x.get("account"): x.get("class", UNK) for x in _table(an.answers, "income_classification")}
        groups = {"Qualifying income": [], "Non-qualifying income": [], "Excluded activity": [], "Exempt income": [], UNK: []}
        ded = {}
        for a in self.conf:
            if ROWS[a.row][1] == "B":
                for code, name in a.accounts:
                    ded[code or name] = ded.get(code or name, 0.0) + a.amount / max(len(a.accounts), 1)
        for code, name, amt in an.ctx.get("income_accounts", []):
            c = cls.get(code) or cls.get(name) or UNK
            groups[c].append((code, name, amt, ded.get(code, 0.0) + (ded.get(name, 0.0) if name != code else 0.0)))
        sums = {}
        for title, key in (("Qualifying income", "Qualifying income"), ("Non-qualifying income", "Non-qualifying income"),
                           ("Excluded activities (non-qualifying)", "Excluded activity"), ("Exempt income", "Exempt income"),
                           ("Unclassified income (treated as non-qualifying until classified)", UNK)):
            banner(ws, r, title, 1, 5)
            r += 1
            header_row(ws, r, ["Account ID", "Description", "Reference", "Income (AED)", "Less: removed by confirmed adjustments"])
            r += 1
            f = r
            for code, name, amt, d in groups[key]:
                put(ws, f"A{r}", code, border=BOX)
                put(ws, f"B{r}", name, border=BOX)
                put(ws, f"C{r}", "P&L / TB (AED)", border=BOX, size=9)
                put(ws, f"D{r}", amt, border=BOX, fmt=NUM0, color=INPUT_BLUE)
                put(ws, f"E{r}", d, border=BOX, fmt=NUM0, color=INPUT_BLUE)
                r += 1
            put(ws, f"A{r}", "Total", bold=True)
            put(ws, f"D{r}", _rsum("D", f, r), bold=True, fmt=NUM0, border=TOPBOT)
            put(ws, f"E{r}", _rsum("E", f, r), bold=True, fmt=NUM0, border=TOPBOT)
            sums[key] = r
            r += 2
        S = lambda k: f"(D{sums[k]}-E{sums[k]})"
        banner(ws, r, "De minimis test (parameters from the rule register - verify)", 1, 5)
        r += 1
        put(ws, f"B{r}", "Non-qualifying revenue (incl. excluded activities and unclassified)", border=BOX)
        put(ws, f"D{r}", f"=D{sums['Non-qualifying income']}+D{sums['Excluded activity']}+D{sums[UNK]}", fmt=NUM0, border=BOX)
        nqrev = r
        r += 1
        put(ws, f"B{r}", "Total revenue / income for the test", border=BOX)
        put(ws, f"D{r}", f"=D{sums['Qualifying income']}+D{nqrev}+D{sums['Exempt income']}", fmt=NUM0, border=BOX)
        totrev = r
        r += 1
        put(ws, f"B{r}", f"Limit: lower of AED {self.P['qfzp_de_minimis_amount']:,.0f} and {self.P['qfzp_de_minimis_pct']:.0%} of total revenue", border=BOX)
        put(ws, f"D{r}", f"=MIN({self.P['qfzp_de_minimis_amount']},{self.P['qfzp_de_minimis_pct']}*D{totrev})", fmt=NUM0, border=BOX)
        lim = r
        r += 1
        put(ws, f"B{r}", "Result", bold=True, border=BOX)
        put(ws, f"D{r}", f'=IF(D{nqrev}<=D{lim},"Within limit","EXCEEDS LIMIT")', bold=True, border=BOX)
        r += 2
        banner(ws, r, "Allocation of expenses and split of taxable income", 1, 5)
        r += 1
        put(ws, f"B{r}", "Taxable income before loss relief (Calculation of CT, line D)", border=BOX)
        put(ws, f"D{r}", f"='{CALC}'!E57", fmt=NUM0, border=BOX)
        ti = r
        r += 1
        qi_net = f"({S('Qualifying income')})"
        nq_net = f"({S('Non-qualifying income')}+{S('Excluded activity')}+{S(UNK)}+{S('Exempt income')})"
        put(ws, f"B{r}", "Qualifying income after adjustments", border=BOX); put(ws, f"D{r}", f"={qi_net}", fmt=NUM0, border=BOX); qi_r = r; r += 1
        put(ws, f"B{r}", "Non-qualifying income after adjustments (exempt income not removed stays taxable)", border=BOX); put(ws, f"D{r}", f"={nq_net}", fmt=NUM0, border=BOX); nq_r = r; r += 1
        put(ws, f"B{r}", "Total expenses and other adjustments (income less taxable income)", border=BOX); put(ws, f"D{r}", f"=D{qi_r}+D{nq_r}-D{ti}", fmt=NUM0, border=BOX); ex_r = r; r += 1
        put(ws, f"B{r}", "Directly attributable to qualifying income", border=BOX); put(ws, f"D{r}", _num(ans(an.answers, "qfzp_direct_q", 0.0)), fmt=NUM0, border=BOX, color=INPUT_BLUE); dq_r = r; r += 1
        put(ws, f"B{r}", "Directly attributable to non-qualifying income", border=BOX); put(ws, f"D{r}", _num(ans(an.answers, "qfzp_direct_nq", 0.0)), fmt=NUM0, border=BOX, color=INPUT_BLUE); dn_r = r; r += 1
        put(ws, f"B{r}", "Indirect expenses (allocated by income ratio - assumption, technical review)", border=BOX); put(ws, f"D{r}", f"=D{ex_r}-D{dq_r}-D{dn_r}", fmt=NUM0, border=BOX); ind_r = r; r += 1
        put(ws, f"B{r}", "Qualifying share of indirect expenses", border=BOX); put(ws, f"D{r}", f"=IF(D{qi_r}+D{nq_r}>0,D{qi_r}/(D{qi_r}+D{nq_r}),0)", fmt="0.00%", border=BOX); rat_r = r; r += 1
        put(ws, f"B{r}", "Taxable income - qualifying (0%)", bold=True, border=BOX); put(ws, f"D{r}", f"=D{qi_r}-D{dq_r}-D{ind_r}*D{rat_r}", fmt=NUM0, bold=True, border=BOX); qt_r = r; r += 1
        put(ws, f"B{r}", "Taxable income - non-qualifying, before loss relief", bold=True, border=BOX); put(ws, f"D{r}", f"=D{nq_r}-D{dn_r}-D{ind_r}*(1-D{rat_r})", fmt=NUM0, bold=True, border=BOX); nqt_r = r; r += 1
        self.cell["fz_nq_pre"] = f"Freezone!$D${nqt_r}"
        self.cell["fz_q"] = f"Freezone!$D${qt_r}"
        put(ws, f"B{r}", "Check: split equals taxable income", italic=True); put(ws, f"D{r}", f'=IF(ABS(D{qt_r}+D{nqt_r}-D{ti})<0.5,"OK","CHECK")', bold=True); r += 1
        # loss relief and tax
        self.cell["fz_nq_after_row"] = r
        put(ws, f"B{r}", "Less: tax loss relief applied to non-qualifying income", border=BOX)
        put(ws, f"D{r}", "=0", fmt=NUM0, border=BOX)       # patched in _write_losses when the loss sheet exists
        self._fz_loss_cell = f"D{r}"
        r += 1
        put(ws, f"B{r}", "Taxable income - non-qualifying, after loss relief", bold=True, border=BOX); put(ws, f"D{r}", f"=D{nqt_r}-D{r - 1}", fmt=NUM0, bold=True, border=BOX); nqa_r = r; r += 1
        put(ws, f"B{r}", "Corporate Tax at 9% on non-qualifying income (no AED 375,000 band for a QFZP)", bold=True, border=BOX)
        put(ws, f"D{r}", f"=ROUND(MAX(D{nqa_r},0)*{self.P['standard_rate']},0)", fmt=NUM0, bold=True, border=BOX)
        self.cell["fz_tax"] = f"Freezone!$D${r}"
        self.cell["fz_q_income"] = f"Freezone!$D${qt_r}"
        r += 2
        banner(ws, r, "Conditions (management responses)", 1, 5)
        r += 1
        for lab, qid in (("Adequate substance evidenced", "qfzp_substance"), ("Audited financial statements available", "qfzp_audited"),
                         ("Domestic permanent establishment / natural-person income", "qfzp_domestic_pe")):
            put(ws, f"B{r}", lab, border=BOX)
            put(ws, f"D{r}", ans(an.answers, qid, UNK), border=BOX, color=INPUT_BLUE)
            r += 1
        acts = ans(an.answers, "qfzp_activities", []) or []
        put(ws, f"B{r}", "Qualifying activities undertaken: " + ("; ".join(acts) if acts else "none stated"), wrap=True)
        ws.row_dimensions[r].height = 28
        r += 2
        banner(ws, r, "EBITDA (accounting) - supports the interest-limitation test", 1, 5)
        r += 1
        it = an.comp.interest_test
        for lab, f in (("Profit for the year before tax", f"={self.cell['pbt']}"),
                       ("Add: finance costs", it.get("interest_expense", 0.0)),
                       ("Less: interest income", -it.get("interest_income", 0.0)),
                       ("Add: depreciation and amortisation", sum(a.net_aed for a in an.tr.accounts if a.line == "P_DEP"))):
            put(ws, f"B{r}", lab, border=BOX)
            put(ws, f"D{r}", f, fmt=NUM0, border=BOX, color=None if isinstance(f, str) else INPUT_BLUE)
            r += 1
        put(ws, f"B{r}", "EBITDA (approximation - the legislated 'tax EBITDA' differs)", bold=True)
        put(ws, f"D{r}", f"=SUM(D{r - 4}:D{r - 1})", bold=True, fmt=NUM0, border=TOPBOT)
        page(ws, False, f"A1:E{r}")

    # ----------------------------------------------------------------- Tax losses
    def _write_losses(self):
        an, comp = self.an, self.an.comp
        ws = self.wb["Tax Loss Carried Forward"]
        _reset(ws, 3)
        widths(ws, {"A": 14, "B": 22, "C": 24, "D": 26, "E": 24, "F": 4})
        put(ws, "A1", f"='{CALC}'!$C$2", bold=True, size=11)
        put(ws, "A2", f"='{CALC}'!$C$3", bold=True, size=11)
        put(ws, "A4", "Tax losses carried forward (AED)", bold=True, fill_=AMBER)
        for c in "BCDE":
            put(ws, f"{c}4", None, fill_=AMBER)
        header_row(ws, 6, ["Tax period", "Tax loss brought forward", "Taxable income / (loss) of the period (before relief)",
                           f"Tax loss relief (up to {self.P['loss_utilisation_cap']:.0%} of taxable income)", "Tax loss carried forward"], height=55)
        yr = (an.period_end or "")[:4]
        put(ws, "A7", yr, border=BOX)
        put(ws, "B7", comp.loss_opening, fmt=NUM0, border=BOX, color=INPUT_BLUE)
        base = self.cell.get("fz_nq_pre") if comp.basis == "QFZP" else f"'{CALC}'!E57"
        put(ws, "C7", f"={base}", fmt=NUM0, border=BOX)
        sbr = comp.basis == "SBR"
        put(ws, "D7", 0 if sbr else f"=MIN(B7,{self.P['loss_utilisation_cap']}*MAX(C7,0))", fmt=NUM0, border=BOX)
        put(ws, "E7", f"=B7-D7+MAX(-C7,0)" if not sbr else "=B7", fmt=NUM0, border=BOX)
        put(ws, "A9", "Tax losses received from a group company", border=BOX)
        put(ws, "B9", _num(ans(an.answers, "loss_transfer_in", 0.0)), fmt=NUM0, border=BOX, color=INPUT_BLUE)
        put(ws, "D9", 0 if sbr else f"=MIN(B9,MAX({self.P['loss_utilisation_cap']}*MAX(C7,0)-D7,0))", fmt=NUM0, border=BOX)
        put(ws, "E9", "=B9-D9" if not sbr else "=B9", fmt=NUM0, border=BOX)
        put(ws, "A11", "Relief claimed against taxable income (feeds Calculation of CT rows 60-61)", bold=True)
        put(ws, "D11", "=D7+D9", fmt=NUM0, bold=True, border=TOPBOT)
        if sbr:
            put(ws, "A12", "Small Business Relief is applied: no loss relief is used and losses of this period are not carried forward.", color="C00000", size=9)
        if comp.basis == "QFZP":
            put(ws, "A13", "QFZP basis: relief is applied against non-qualifying income - technical review of loss treatment required.", color="C00000", size=9)
        self.cell["loss_cf"] = "'Tax Loss Carried Forward'!$D$7"
        self.cell["loss_group"] = "'Tax Loss Carried Forward'!$D$9"
        if "fz_nq_pre" in self.cell:
            fz = self.wb["Freezone"]
            fz[self._fz_loss_cell].value = "='Tax Loss Carried Forward'!D11"
        page(ws, False, "A1:E13")

    # ----------------------------------------------------------------- Foreign tax
    def _write_foreign_tax(self):
        ws = self.wb["Taxes Paid Abroad"]
        _reset(ws, 3)
        widths(ws, {"A": 8, "B": 20, "C": 26, "D": 18, "E": 10, "F": 18, "G": 16, "H": 18, "I": 20})
        put(ws, "A1", f"='{CALC}'!$C$2", bold=True, size=11)
        put(ws, "A2", f"='{CALC}'!$C$3", bold=True, size=11)
        put(ws, "A4", "Details of taxes paid abroad", bold=True, fill_=AMBER)
        for c in "BCDEFGHI":
            put(ws, f"{c}4", None, fill_=AMBER)
        header_row(ws, 6, ["S/N", "Country", "Nature of income", "Income (AED)", "Tax rate", "Tax paid (AED)", "Proof of payment available?",
                           "Income exempt in the UAE?", "Creditable (AED)"], height=40)
        rows = _table(self.an.answers, "foreign_taxes")
        r = 7
        f = r
        for i, row in enumerate(rows):
            put(ws, f"A{r}", i + 1, border=BOX)
            put(ws, f"B{r}", row.get("country", ""), border=BOX, color=INPUT_BLUE)
            put(ws, f"C{r}", row.get("nature", ""), border=BOX, color=INPUT_BLUE)
            put(ws, f"D{r}", _num(row.get("income")), border=BOX, fmt=NUM0, color=INPUT_BLUE)
            put(ws, f"E{r}", _num(row.get("rate")) / 100.0, border=BOX, fmt=PCT, color=INPUT_BLUE)
            put(ws, f"F{r}", _num(row.get("paid")), border=BOX, fmt=NUM0, color=INPUT_BLUE)
            put(ws, f"G{r}", row.get("proof", UNK), border=BOX, color=INPUT_BLUE)
            put(ws, f"H{r}", row.get("exempt_income", UNK), border=BOX, color=INPUT_BLUE)
            put(ws, f"I{r}", f'=IF(AND(G{r}="Yes",H{r}="No"),IF(D{r}>0,MIN(F{r},{self.P['standard_rate']}*D{r}),F{r}),0)', border=BOX, fmt=NUM0)
            r += 1
        l = max(r - 1, f)
        put(ws, f"B{r}", "Total creditable (proof of payment AND income not exempt)", bold=True)
        put(ws, f"F{r}", f"=SUM(I{f}:I{l})", bold=True, fmt=NUM0, border=TOPBOT)
        put(ws, f"B{r + 1}", "Credit claimed (limited to Corporate Tax payable)", bold=True)
        self.cell["ftc_total_row"] = r
        self.cell["ftc"] = f"'Taxes Paid Abroad'!$F${r + 1}"
        self._ftc_cell_row = r + 1
        put(ws, f"F{r + 1}", f"=MIN(F{r},'{CALC}'!E67+'{CALC}'!E68)", bold=True, fmt=NUM0, border=TOPBOT)
        put(ws, f"B{r + 3}", "Each row is capped at the standard rate x the income shown and the total at Corporate Tax payable. No credit is given without proof of payment and confirmation that the income is not exempt. Technical review still required.", italic=True, size=9)
        page(ws, False, f"A1:I{r + 3}")

    # ----------------------------------------------------------------- DTA/DTL (optional)
    def _write_dta(self):
        ws = self.wb["DTA & DTL Working"]
        an = self.an
        put(ws, "D4", f"='{CALC}'!E9")
        mapping = {7: "A_PPE", 8: "A_INV", 9: "A_RP", 10: "A_TRADE_REC", 11: "A_CASH", 12: "E_CAPITAL", 13: "E_RESERVE", 14: "E_RE", 15: "E_SHCA",
                   16: "L_EOSB", 17: "L_BORR", 18: "L_RP", 19: "L_TRADE_PAY"}
        for row, key in mapping.items():
            sign = 1 if key.startswith("A_") else -1
            ws[f"C{row}"].value = self._tb_sumifs(key, sign)
        ws["G23"].value = "=SUM(G7:G22)"
        loss_cell = f"-'{CALC}'!E63" if an.comp.E < 0 else "0"
        ws["D21"].value = f"=MAX({loss_cell},0)"
        put(ws, "B25", "Tax bases equal carrying values until temporary differences are entered in column D - complete before relying on this working. "
                       "(Sum range corrected to include rows 7-22.)", italic=True, size=9, color="C00000")
        self.log.append("DTA & DTL Working: carrying values linked to the TB; total formula corrected to include rows 7-22 (template summed rows 11-22 only).")

    # ----------------------------------------------------------------- Assumptions
    def _write_assumptions(self):
        an = self.an
        ws = self.wb["Assumptions"]
        A = an.answers
        ws["B2"].value = an.entity
        ws["B3"].value = f"Year ended: {fmt_date(an.period_end)}"
        pend = sum(1 for a in an.adjs if a.status == "Proposed")
        if pend or an.open_items:
            put(ws, "B4", f"DRAFT - {pend} proposed adjustment(s) and {len(an.open_items)} open item(s) awaiting confirmation. Figures are preliminary.",
                bold=True, color="C00000")
        rb = RB_CODE.get(ans(A, "realisation_basis", RB_OPTIONS[3]), "unknown")
        q = engine.qfzp_status(A)
        txt = {}
        if q["claimed"] and q["ok"]:
            txt[8] = "The management has elected to avail the Qualifying Free Zone Person (QFZP) regime under the UAE Corporate Tax Law and has confirmed that substance, audited financial statements and the other stated conditions are met."
        elif q["claimed"]:
            txt[8] = ("The management has indicated that it intends to avail the QFZP regime. The following conditions are not yet evidenced: " + "; ".join(q["missing"]) +
                      ". Accordingly the computation applies the ordinary Corporate Tax rates.")
        elif ans(A, "entity_status") == "Free Zone":
            txt[8] = "The company is registered in a Free Zone but has not confirmed QFZP status; the computation applies the ordinary Corporate Tax rates."
        else:
            txt[8] = "The company is not claiming Qualifying Free Zone Person status."
        txt[9] = {"A": "The management has opted for the realisation basis - Option 1 (all assets and liabilities subject to fair value or impairment accounting).",
                  "B": "The management has opted for the realisation basis - Option 2 (assets and liabilities held on capital account only).",
                  "none": "The management has not elected the realisation basis; unrealised gains and losses are taken into account as recorded (current basis, Option 3).",
                  "unknown": "The realisation-basis election has not been confirmed by management; the computation is presented on the current basis (Option 3) pending confirmation."}[rb]
        if rb in ("A", "B") and ans(A, "realisation_evidence") != YES:
            txt[9] += " Evidence that the election was made in the first tax period is outstanding."
        tr_ = ans(A, "transitional", UNK)
        txt[10] = {YES: "The management has opted for Transitional Rules under the Corporate Tax Law.",
                   NO: "The management has not claimed Transitional Rules.", UNK: "Whether Transitional Rules are claimed has not been confirmed; none has been applied."}.get(tr_)
        yn = lambda qid: ans(A, qid, UNK)
        txt[11] = " The management has been informed of the requirement to maintain proper supporting documents for all expenses incurred by the company and recorded in the statement of profit and loss." if yn("conf_documentation") == YES else \
            " The management has not yet confirmed that it has been informed of the requirement to maintain supporting documents for all expenses."
        txt[12] = " The management has confirmed that all expenses debited to the profit and loss account during the relevant tax period have been incurred wholly and exclusively for business purposes." if yn("conf_business_purpose") == YES else \
            " Management confirmation that all expenses were incurred wholly and exclusively for business purposes is outstanding."
        txt[13] = " The management has confirmed that no expenses of a disallowed nature have been recorded in the books of accounts, other than those already disallowed in the tax computation." if yn("conf_disallowed") == YES else \
            " Management confirmation that no other expenses of a disallowable nature are recorded is outstanding."
        txt[14] = " The management has confirmed the exempt-income position reflected in this computation." if yn("conf_exempt_income") == YES else \
            " Management confirmation of the exempt-income position is outstanding."
        txt[15] = ("The management has confirmed that all related party transactions disclosed in the financial statements have been conducted at arm's length in accordance with applicable transfer pricing regulations. "
                   "We have not reviewed transfer pricing documentation. Management has been made aware of the consequences of non-compliance with transfer pricing regulations and has agreed to maintain the necessary supporting documentation.") if yn("conf_arms_length") == YES else \
            "Management confirmation that related party transactions are at arm's length is outstanding; no transfer pricing adjustment has been made."
        txt[16] = "The management has confirmed the position on investments in subsidiaries; the computation has been prepared accordingly." if yn("conf_subsidiaries") == YES else \
            "The position on investments in subsidiaries (including the basis of preparation of their accounts under Ministerial Decision No. 134 of 2023 and IFRS) has not been confirmed by management."
        for r, t in txt.items():
            ws[f"C{r}"].value = t
            ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(25.5, 13 * (len(t) // 95 + 1))
        # currency note (row 17 is free in the template)
        fx = an.project.fx
        if not fx.is_aed:
            put(ws, "B17", "Currency", bold=True, border=BOX)
            ws["B17"].alignment = Alignment(wrap_text=True, vertical="top")
            put(ws, "C17", f"The trial balance is maintained in {fx.tb_currency} and has been converted to UAE dirhams (AED) using the method recorded on the 'FX Rates' sheet "
                           f"({an.tr.method_record.get('method_description', '')}). All amounts in the statements and computation are in AED.", wrap=True, border=BOX)
            ws.row_dimensions[17].height = 45
        # applicability log below the disclaimer
        r = 21
        put(ws, f"B{r}", "Sheet applicability log", bold=True, fill_=AMBER)
        put(ws, f"C{r}", None, fill_=AMBER)
        r += 1
        for name, (inc, reason) in self.applic.items():
            put(ws, f"B{r}", name, border=BOX, bold=inc)
            put(ws, f"C{r}", ("Included - " if inc else "Not applicable, not produced - ") + reason, border=BOX, size=9, color=None if inc else "7F7F7F")
            r += 1
        r += 1
        put(ws, f"B{r}", "Template changes", bold=True, fill_=AMBER)
        put(ws, f"C{r}", None, fill_=AMBER)
        r += 1
        for line in self.log:
            put(ws, f"B{r}", "-", border=BOX)
            put(ws, f"C{r}", line, border=BOX, size=9, wrap=True)
            ws.row_dimensions[r].height = 26
            r += 1

    # ----------------------------------------------------------------- Calculation of CT
    def _write_calc(self):
        an, comp = self.an, self.an.comp
        A = an.answers
        ws = self.wb[CALC]
        # clear stale values of the Option 2/3 columns (hidden in the template) and other-client comments
        for r in range(9, 75):
            for col in "FG":
                ws[f"{col}{r}"].value = None
        for r in range(11, 24):
            ws[f"I{r}"].value = None
        if ws.data_validations:
            ws.data_validations.dataValidation = []
        rb = RB_CODE.get(ans(A, "realisation_basis", RB_OPTIONS[3]), "unknown")
        ws["D5"].value = "Accrual Basis" if ans(A, "accounting_basis", "Accrual") == "Accrual" else "Cash Basis"
        ws["D6"].value = ans(A, "accounting_framework", "IFRS")
        ws["D7"].value = "At Arm's Length" if ans(A, "conf_arms_length") == YES else "Not confirmed"
        ws["E9"].value = {"A": "Realisation Basis\nOption 1\n(AED)", "B": "Realisation Basis\nOption 2\n(AED)", "none": "Current Basis\nOption 3\n(AED)",
                          "unknown": "Current Basis\n(election not confirmed)\n(AED)"}[rb]
        pend = sum(1 for a in an.adjs if a.status == "Proposed")
        put(ws, "C8", f"DRAFT - {pend} proposed adjustment(s) await approval; {len(an.open_items)} open item(s). Only confirmed adjustments are included below." if (pend or an.open_items) else "Final - all adjustments confirmed.",
            bold=True, color="C00000" if (pend or an.open_items) else "006100")
        ws["E11"].value = f"={self.cell['profit']}"
        ws["I11"].value = {"A": "Realisation Basis - Option 1", "B": "Realisation Basis - Option 2", "none": "Current basis - no election", "unknown": "Election not confirmed - current basis shown"}[rb]
        ws["I12"].value = {YES: "Transitional Rule - Yes", NO: "Transitional Rule - No", UNK: "Transitional Rule - not confirmed"}.get(ans(A, "transitional", UNK))
        ws["I13"].value = f"Source: {an.project.tb.source_file} ({an.project.fx.tb_currency}" + ("" if an.project.fx.is_aed else " converted to AED") + ")"
        for key, (row, sec, label) in ROWS.items():
            ws[f"E{row}"].value = self._adj_sumif(key)
            ws[f"E{row}"].number_format = NUM0
            ws[f"E{row}"].font = font(color=None if str(ws[f"E{row}"].value).startswith("=") else None)
            items = [a for a in self.conf if a.row == key]
            if items:
                ws[f"I{row}"].value = "; ".join(a.description[:55] for a in items[:2]) + (" ..." if len(items) > 2 else "")
                ws[f"I{row}"].font = font(size=9)
        ws["E24"].value = 0
        # losses
        if "loss_cf" in self.cell:
            ws["E60"].value = f"=-{self.cell['loss_cf']}"
            ws["E61"].value = f"=-{self.cell['loss_group']}"
        else:
            ws["E60"].value = 0
            ws["E61"].value = 0
        # tax
        if comp.basis == "QFZP":
            ws["C67"].value = "0% on Qualifying Income"
            ws["E67"].value = 0
            ws["I67"].value = f"Qualifying Income taxed at 0% (see Freezone sheet): AED {comp.q_taxable or 0:,.0f}"
            ws["C68"].value = "9% on Non-Qualifying Income"
            ws["E68"].value = f"=ROUND(MAX({self.cell['fz_nq_pre']}-" + (f"{self.cell['loss_cf']}-{self.cell['loss_group']}" if 'loss_cf' in self.cell else "0") + ",0)*" + f"{self.P['standard_rate']},0)"
        elif comp.basis == "SBR":
            ws["C67"].value = "Small Business Relief - taxable income treated as nil"
            ws["E67"].value = 0
            ws["C68"].value = "9% on taxable income"
            ws["E68"].value = 0
            ws["I67"].value = "Small Business Relief elected; statutory conditions confirmed by management."
        else:
            ws["C67"].value = f"0% on taxable income up to AED {self.P['nil_rate_band']:,.0f}"
            ws["E67"].value = 0
            ws["C68"].value = f"9% on taxable income above AED {self.P['nil_rate_band']:,.0f}"
            ws["E68"].value = f"=IF(E63<={self.P['nil_rate_band']},0,ROUND((E63-{self.P['nil_rate_band']})*{self.P['standard_rate']},0))"
            if ans(A, "qfzp_claimed") == YES:
                ws["I67"].value = "QFZP claimed but not supported - see Freezone sheet and Assumptions."
        if "ftc" in self.cell:
            ws["E70"].value = f"=-{self.cell['ftc']}"
        else:
            ws["E70"].value = 0
        ws["E72"].value = "=E67+E68+E70"
        ws["E73"].value = "=IFERROR(E72/E11,0)"
        ws["E73"].number_format = "0.0%"
        ws["E57"].value = "=((E11-E25)+E50)+E53-E54"
        ws["E25"].value = "=SUM(E14:E24)"
        ws["E50"].value = "=SUM(E28:E48)"
        ws["E63"].value = "=E57+E60+E61"
        ws.print_area = "A1:I74"
        ws.column_dimensions["I"].width = 60
        ws.column_dimensions["I"].hidden = False
