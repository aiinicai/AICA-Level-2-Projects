"""Internal REVIEW workbook (separate from the client computation workbook).

Holds everything a reviewer needs: summary, adjustments register (proposed AND confirmed), high-risk accounts, GL tests,
exempt income/reliefs, losses, open items and questionnaire, reconciliation checks, account mapping, scenarios, rule
register and audit trail. A sheet that would be empty is not produced (user requirement).
Values are written as numbers (this is a review record, not a live model); the live model is the client workbook.
"""
from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .engine import ROWS
from .mapping import fs_lines
from .models import load_json
from .questionnaire import SECTIONS, build_questions, ans
from .xl_style import (AMBER, BOX, NAVY, NUM0, NUM2, PCT, banner, font, header_row, page, put, status_fill, widths)


def _fmt_date(s):
    try:
        return datetime.fromisoformat(s).strftime("%d %B %Y")
    except Exception:
        return s or ""


class ReviewWorkbook:
    def __init__(self, an):
        self.an = an
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.omitted = []

    # ------------------------------------------------------------------ helpers
    def _sheet(self, name, title, subtitle=""):
        ws = self.wb.create_sheet(name)
        put(ws, "A1", self.an.entity, bold=True, size=12)
        put(ws, "A2", f"{title}" + (f" - {subtitle}" if subtitle else ""), bold=True)
        put(ws, "A3", f"Period: {_fmt_date(self.an.period_start)} to {_fmt_date(self.an.period_end)}   |   Status: DRAFT / PRELIMINARY - for review, not a filed position",
            italic=True, size=9, color="C00000")
        return ws

    def _table(self, name, title, headers, rows, col_widths, num_cols=(), pct_cols=(), wrap_cols=(), status_col=None, subtitle="", start_row=5):
        if not rows:
            self.omitted.append((name, "nothing to report"))
            return None
        ws = self._sheet(name, title, subtitle)
        header_row(ws, start_row, headers, height=32)
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for r, row in enumerate(rows, start_row + 1):
            for c, v in enumerate(row, 1):
                ref = f"{get_column_letter(c)}{r}"
                put(ws, ref, v, border=BOX, fmt=NUM0 if c in num_cols else (PCT if c in pct_cols else None), wrap=c in wrap_cols, valign="top", size=9)
            if status_col:
                ws[f"{get_column_letter(status_col)}{r}"].fill = _fill(status_fill(str(row[status_col - 1])))
        ws.freeze_panes = ws[f"A{start_row + 1}"]
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
        page(ws, True, rows_repeat=f"{start_row}:{start_row}")
        return ws

    # ------------------------------------------------------------------ sheets
    def summary(self):
        an, c = self.an, self.an.comp
        ws = self._sheet("Summary", "Review summary")
        widths(ws, {"A": 46, "B": 22, "C": 70})
        r = 5
        banner(ws, r, "Basis of preparation", 1, 3)
        r += 1
        mr = an.tr.method_record
        rows = [
            ("Entity", an.entity, ""),
            ("Tax period", f"{_fmt_date(an.period_start)} to {_fmt_date(an.period_end)}", ""),
            ("TB currency", an.project.fx.tb_currency, "All statements and the computation are in AED."),
            ("AED conversion method", mr.get("method_description", ""), mr.get("reason", "")),
            ("Rate source", mr.get("rate_source", ""), mr.get("decision", "")),
            ("Source files", ", ".join(x for x in [an.project.tb.source_file if an.project.tb else "", an.project.gl.source_file if an.project.gl else ""] if x), ""),
            ("Financial statements status", ans(an.answers, "tb_status", "Not stated"), ""),
        ]
        for a, b, d in rows:
            put(ws, f"A{r}", a, bold=True, border=BOX)
            put(ws, f"B{r}", b, border=BOX, wrap=True)
            put(ws, f"C{r}", d, border=BOX, wrap=True, size=9)
            r += 1
        r += 1
        banner(ws, r, "Computation (CONFIRMED adjustments only)", 1, 3)
        r += 1
        for a, v, d in [
            ("Accounting profit/(loss) - AED", c.A, "From the profit and loss statement built from the TB."),
            ("Total deductions - AED", c.B, ""), ("Total additions (non-deductible) - AED", c.C, ""),
            ("Taxable income/(loss) before losses - AED", c.D, ""), ("Net taxable income/(loss) - AED", c.E, ""),
            ("Corporate tax payable - AED", c.payable, f"Basis: {c.basis}. Preliminary - see open items."),
        ]:
            put(ws, f"A{r}", a, bold=True, border=BOX)
            put(ws, f"B{r}", v, border=BOX, fmt=NUM0)
            put(ws, f"C{r}", d, border=BOX, size=9, wrap=True)
            r += 1
        pi = an.pending_impact()
        r += 1
        banner(ws, r, "Not yet in the computation", 1, 3)
        r += 1
        put(ws, f"A{r}", "Proposed adjustments awaiting a decision", bold=True, border=BOX)
        put(ws, f"B{r}", pi["count"], border=BOX)
        put(ws, f"C{r}", "Shown separately; they do not change the figures above until confirmed.", border=BOX, size=9)
        r += 1
        if pi["count"]:
            put(ws, f"A{r}", "Effect on taxable income if ALL were confirmed - AED", bold=True, border=BOX)
            put(ws, f"B{r}", pi["delta_taxable"], border=BOX, fmt=NUM0)
            put(ws, f"C{r}", f"Tax would change by AED {pi['delta_tax']:,.0f} (illustrative only - no decision has been taken).", border=BOX, size=9)
            r += 1
        put(ws, f"A{r}", "Open items (information / technical review)", bold=True, border=BOX)
        put(ws, f"B{r}", len(an.open_items), border=BOX)
        r += 2
        banner(ws, r, "Risk and check summary", 1, 3)
        r += 1
        from collections import Counter
        rc = Counter(f.risk for f in an.review.flags)
        for lab, v in [("High-risk indicators", rc.get("High", 0)), ("Medium-risk indicators", rc.get("Medium", 0)), ("Low-risk indicators", rc.get("Low", 0)),
                       ("GL transactions flagged", len(an.review.txn_flags)),
                       ("Reconciliation checks failed", sum(1 for x in an.checks if x.status == "FAIL")),
                       ("Reconciliation checks with warnings", sum(1 for x in an.checks if x.status == "WARN"))]:
            put(ws, f"A{r}", lab, bold=True, border=BOX)
            put(ws, f"B{r}", v, border=BOX)
            r += 1
        r += 1
        banner(ws, r, "Sheets omitted from the client workbook as not applicable", 1, 3)
        r += 1
        for sh, (inc, why) in an.applic.items():
            if not inc:
                put(ws, f"A{r}", sh, border=BOX)
                put(ws, f"B{r}", "Omitted", border=BOX)
                put(ws, f"C{r}", why, border=BOX, size=9)
                r += 1
        r += 1
        put(ws, f"A{r}", "Preliminary review only. No tax position has been taken; items marked 'Information required' or 'Technical review required' "
                         "need a qualified adviser's decision before filing.", italic=True, size=9, color="C00000")
        page(ws, False)

    def adjustments(self):
        rows = []
        for a in self.an.adjs:
            spec = ROWS.get(a.row, (0, "", a.row))
            rows.append([a.adj_id, spec[2], a.category, a.description, "; ".join(f"{x[0]} {x[1]}".strip() for x in a.accounts), a.amount, a.status,
                         a.report_status, a.blocking, a.reason, a.rule_ref, "; ".join(a.docs_needed), a.mgmt_response, a.reviewer, a.support_ref,
                         a.override_note, ", ".join(a.gl_refs[:8])])
        self._table("Adjustments Register", "Adjustments register (proposed, confirmed and rejected kept separate)",
                    ["ID", "Computation row", "Category", "Description", "Accounts", "Amount (AED)", "Decision", "Report status", "Blocking issue", "Reason",
                     "Rule ref", "Documents needed", "Management response", "Reviewer", "Support ref", "Override note", "GL references"], rows,
                    [22, 34, 16, 40, 30, 15, 12, 18, 30, 44, 8, 30, 26, 14, 14, 26, 24], num_cols=(6,), wrap_cols=(2, 4, 5, 9, 10, 12, 13, 16), status_col=7)

    def risk(self):
        rows = [[f.flag_id, f.risk, f.category_label, f.code, f.name, f.amount_aed, f.reason, f.tax_implication, f.status, "; ".join(f.documents),
                 f.proposed_adjustment if f.proposed_adjustment is not None else "", f.mgmt_response, f.reviewer_conclusion, ", ".join(f.gl_refs[:8])]
                for f in self.an.review.flags]
        self._table("High-Risk Accounts", "Account-level risk indicators (indicators, not conclusions)",
                    ["Flag ID", "Risk", "Category", "Code", "Account", "Balance (AED)", "Why flagged", "Possible tax implication", "Status", "Documents to request",
                     "Proposed adj. (AED)", "Management response", "Reviewer conclusion", "GL references"], rows,
                    [20, 8, 28, 9, 38, 15, 44, 50, 20, 34, 15, 26, 26, 24], num_cols=(6, 11), wrap_cols=(3, 5, 7, 8, 10, 12, 13))

    def gl_tests(self):
        rows = [[t.line_id, t.date, t.code, t.name, t.amount_aed, t.description, t.counterparty, t.reference, "; ".join(t.tests), t.risk]
                for t in self.an.review.txn_flags]
        ws = self._table("GL Review", "General-ledger transaction tests", ["GL line", "Date", "Code", "Account", "Amount (AED)", "Narration", "Counterparty",
                         "Reference", "Tests triggered", "Risk"], rows, [11, 11, 9, 34, 15, 36, 26, 11, 54, 8], num_cols=(5,), wrap_cols=(4, 6, 9))
        if ws is not None and self.an.review.notes:
            put(ws, "A4", " | ".join(self.an.review.notes), italic=True, size=9)
        elif not self.an.project.gl:
            self.omitted.append(("GL Review", "no GL provided - transaction tests not run"))

    def reliefs(self):
        an, c = self.an, self.an.comp
        rows = []
        for a in an.adjs:
            if a.category in ("Exempt income", "Realisation basis", "Foreign tax", "Interest limitation"):
                rows.append([a.category, a.description, a.amount, a.status, a.rule_ref, a.reason, a.blocking])
        extra = []
        if c.interest_test:
            it = c.interest_test
            extra.append(["Interest limitation test", f"Net interest AED {it.get('net_interest', 0):,.0f}; approx. tax EBITDA AED {it.get('approx_tax_ebitda', 0):,.0f}; "
                          f"threshold AED {it.get('threshold', 0):,.0f}", it.get("excess", 0), "Test", "INT-01",
                          "Within the limit" if it.get("within") else "Excess over the limit", "Tax EBITDA is approximated - confirm."])
        if c.sbr:
            extra.append(["Small Business Relief test", " ".join(c.sbr.get("failures", []) + c.sbr.get("open", [])) or "All tested conditions met",
                          "", "Test", "SBR-01", "Potentially eligible" if c.sbr.get("potential") else "Not eligible on the data provided", "Election is not applied automatically."])
        if c.qfzp:
            extra.append(["QFZP assessment", "; ".join(f"{k}: {v}" for k, v in c.qfzp.items() if isinstance(v, (str, int, float, bool))), "", "Test", "QFZP-01", "", ""])
        rows += extra
        self._table("Exempt Income & Reliefs", "Exempt income, elections and reliefs", ["Area", "Item", "Amount (AED)", "Decision", "Rule ref", "Basis / result", "Open point"],
                    rows, [22, 62, 15, 12, 9, 50, 40], num_cols=(3,), wrap_cols=(2, 6, 7), status_col=4)

    def losses(self):
        c = self.an.comp
        if not (c.loss_opening or c.relief_cf or c.relief_group or c.loss_closing):
            self.omitted.append(("Tax Losses", "no tax losses involved"))
            return
        rows = [["Tax losses brought forward", c.loss_opening], ["Losses used in this period (capped at 75% of taxable income)", c.relief_cf],
                ["Losses received from group (capped)", c.relief_group], ["Tax losses carried forward (computed)", c.loss_closing]]
        self._table("Tax Losses", "Tax losses and attributes", ["Item", "AED"], rows, [64, 18], num_cols=(2,),
                    subtitle="Opening losses are as stated by management and are not verified against prior returns.")

    def open_items(self):
        rows = [[i + 1, o.get("topic", ""), o.get("item", ""), o.get("status", ""), "; ".join(o.get("documents", []))] for i, o in enumerate(self.an.open_items)]
        self._table("Open Items", "Open items - information required / technical review", ["#", "Topic", "Item", "Status", "Documents / evidence"], rows,
                    [5, 22, 90, 22, 44], wrap_cols=(3, 5), status_col=4)

    def questionnaire(self):
        an = self.an
        ctx = an.ctx
        rows = []
        for q in build_questions():
            if q.relevant and not q.relevant(an.answers, ctx):
                continue
            v = an.answers.get(q.id)
            if isinstance(v, dict) and "value" in v:
                v = v["value"]
            if isinstance(v, list):
                v = "; ".join(str(x) for x in v) if q.kind == "multi" else f"{len(v)} row(s)"
            user_meta = an.project.answers.get("_meta", {}).get(q.id, {})
            confirmed = (q.id in an.project.answers) or bool(user_meta.get("confirmed"))
            src = an.answers.get("_meta", {}).get(q.id, {}).get("source")
            status = ("Answered / confirmed by the user" if confirmed else
                      "Prefilled from the files - not yet confirmed" if src == "prefilled" else "Default value - not yet confirmed")
            rows.append([SECTIONS.get(q.section, q.section), q.text, "" if v is None else v, status, q.evidence, q.why])
        self._table("Questionnaire", "Questionnaire answers and assumptions", ["Section", "Question", "Answer", "Status", "Evidence normally expected", "Why it matters"], rows,
                    [30, 62, 34, 28, 34, 54], wrap_cols=(1, 2, 3, 5, 6))

    def checks(self):
        rows = [[x.id, x.description, x.status, x.detail] for x in self.an.checks]
        ws = self._table("Reconciliation Checks", "Reconciliation and data-quality checks", ["ID", "Check", "Result", "Detail"], rows, [10, 46, 10, 110], wrap_cols=(2, 4))
        if ws:
            for r in range(6, 6 + len(rows)):
                s = ws[f"C{r}"].value
                ws[f"C{r}"].fill = _fill({"PASS": "E2EFDA", "FAIL": "F8CBAD", "WARN": "FFF2CC", "INFO": "DDEBF7"}.get(s, "FFFFFF"))
                ws[f"C{r}"].font = font(bold=True, size=9)

    def mapping(self):
        lines = fs_lines()
        rows = []
        for m, a in zip(self.an.mapped, self.an.tr.accounts):
            rows.append([m.account.code, m.account.name, m.account.net, a.rate, a.net_aed, lines[m.line]["label"], m.method, m.rule_id, "; ".join(m.flags)])
        ws = self._table("Account Mapping", "Trial balance mapping to financial statement lines (override in the application)",
                         ["Code", "Account", f"Net ({self.an.project.fx.tb_currency}; Dr +)", "Rate", "Net (AED; Dr +)", "Mapped to", "Method", "Rule", "Review flags"], rows,
                         [9, 46, 17, 10, 17, 44, 12, 8, 60], num_cols=(3, 5), wrap_cols=(2, 6, 9))
        if ws:
            for r in range(6, 6 + len(rows)):
                ws[f"D{r}"].number_format = "0.000000"
                ws[f"C{r}"].number_format = NUM2
                ws[f"E{r}"].number_format = NUM2

    def scenarios(self):
        rows = [[s["option"], s["deductions"], s["additions"], s["taxable_income"], s["tax_payable"]] for s in self.an.scenarios]
        self._table("Scenarios", "Realisation-basis scenarios (what-if only; the election is made once, in the first tax period)",
                    ["Option", "Deductions (AED)", "Additions (AED)", "Taxable income (AED)", "Tax payable (AED)"], rows, [34, 18, 18, 20, 18], num_cols=(2, 3, 4, 5),
                    subtitle="Includes only realisation-basis items; other proposed adjustments are excluded. Not advice on which option to elect.")

    def rules(self):
        reg = load_json("rules.json")
        rows = [[r["id"], r["topic"], r["summary"], r["source"], r.get("reference", ""), r.get("applicable_from") or "", r.get("applicable_to") or "",
                 r.get("date_checked", ""), r["verification"], r.get("note", "")] for r in reg["rules"]]
        ws = self._table("Rules & Sources", "Tax rules relied on, with verification status (rules file version " + reg.get("version", "") + ")",
                         ["ID", "Topic", "Rule as implemented", "Source", "Reference", "From", "To", "Date checked", "Verification", "Note"], rows,
                         [9, 32, 70, 40, 16, 11, 11, 12, 14, 56], wrap_cols=(2, 3, 4, 10))
        if ws:
            for r in range(6, 6 + len(rows)):
                v = ws[f"I{r}"].value
                ws[f"I{r}"].fill = _fill("E2EFDA" if v == "PRIMARY_READ" else ("FFF2CC" if v == "SECONDARY" else "F8CBAD"))

    def audit(self):
        p = self.an.project
        rows = [[x["time"], x["action"], x["detail"]] for x in p.audit_log]
        for k, v in p.decisions.items():
            rows.append(["", "Decision stored", f"{k}: {v.get('status')} by {v.get('reviewer', '')} ({v.get('support_ref', '')})"])
        self._table("Audit Trail", "Audit trail (actions recorded in the project file)", ["Time", "Action", "Detail"], rows, [20, 28, 110], wrap_cols=(3,))

    # ------------------------------------------------------------------ build
    def build(self, path: str):
        self.summary()
        self.adjustments()
        self.risk()
        self.gl_tests()
        self.reliefs()
        self.losses()
        self.open_items()
        self.questionnaire()
        self.checks()
        self.mapping()
        self.scenarios()
        self.rules()
        self.audit()
        if self.omitted:
            ws = self.wb["Summary"]
            r = ws.max_row + 2
            put(ws, f"A{r}", "Review sheets not produced", bold=True)
            for n, why in self.omitted:
                r += 1
                put(ws, f"A{r}", n, border=BOX)
                put(ws, f"B{r}", why, border=BOX, size=9)
        self.wb.save(path)
        return path


def _fill(h):
    from .xl_style import fill
    return fill(h)
