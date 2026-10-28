"""Extraction of Trial Balance and General Ledger data from Excel / CSV / PDF (text or scanned).

Design rules:
 * Source files are only ever read, never modified.
 * Nothing is guessed silently: every heuristic that could be wrong adds a warning that is shown
   in the extraction preview and must be confirmed by the user.
"""
from __future__ import annotations

import io
import logging
import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import pandas as pd

from .models import Account, TrialBalance, GLLine, GeneralLedger

log = logging.getLogger("ct_tool.extract")

# --------------------------------------------------------------------------------------
# Generic helpers
# --------------------------------------------------------------------------------------

_NUM_RE = re.compile(r"^\(?-?[\d,]*\.?\d+\)?$")


def clean_text(v) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    s = str(v).replace("\n", " ").replace("\r", " ").replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip()


def _normalise_separators(s: str) -> str:
    """Make '1,234.50', '1.234,50', '1 234,50', "1'234.50" and '1234,5' all readable as plain floats."""
    s = s.replace("\u2212", "-").replace("\u00a0", " ").replace("\u202f", " ")
    s = re.sub(r"(?<=\d)[ '\u2019](?=\d{3}(?!\d))", "", s)            # space / apostrophe used as thousands separator
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        if re.fullmatch(r"-?\(?\d{1,3}(,\d{3})+\)?", s):
            s = s.replace(",", "")                                        # 1,234 or 1,234,567
        elif re.fullmatch(r"-?\(?\d+,\d{1,2}\)?", s):
            s = s.replace(",", ".")                                       # 1234,50 (decimal comma)
        else:
            s = s.replace(",", "")
    elif s.count(".") > 1:
        s = s.replace(".", "")                                            # 1.234.567
    return s


def to_num(v) -> Optional[float]:
    """Parse a number from a cell. Handles 1,234.50  1.234,50  1 234,50  (1,234.50)  -1234  'EUR 1,234.50'  '1,234.50 Cr' and blanks/dashes.
    A trailing Cr/CR returns a negative number, Dr a positive one. Returns None if not numeric."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isnan(v):
            return None
        return float(v)
    s = clean_text(v)
    if s in ("", "-", "--", "—", "–", "nil", "Nil", "NIL"):
        return None
    sign = 1.0
    low = s.lower()
    if low.endswith(" cr") or low.endswith("cr") and re.search(r"\d\s*cr$", low):
        sign = -1.0
        s = re.sub(r"\s*cr$", "", s, flags=re.I)
    elif low.endswith(" dr") or re.search(r"\d\s*dr$", low):
        s = re.sub(r"\s*dr$", "", s, flags=re.I)
    s = re.sub(r"^[A-Za-z]{3}\s*(?=[\d(\-.\u2212])", "", s.strip())      # leading currency code, e.g. 'EUR 1,200.00'
    s = re.sub(r"(?<=\d)\s*[A-Za-z]{3}$", "", s)                          # trailing currency code
    s = _normalise_separators(s.strip())
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg = True
        s = s[1:-1]
    if s.endswith("-") and s[:-1].replace(".", "").isdigit():   # trailing minus
        neg = True
        s = s[:-1]
    try:
        x = float(s)
    except ValueError:
        return None
    if neg:
        x = -x
    return sign * x


def code_to_str(v) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return clean_text(v)


def to_iso_date(v) -> Optional[str]:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = clean_text(v)
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y", "%d-%b-%Y", "%d-%b-%y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    try:
        return pd.to_datetime(s, dayfirst=True).date().isoformat()
    except Exception:
        return None


# --------------------------------------------------------------------------------------
# Reading raw tables
# --------------------------------------------------------------------------------------

def list_sheets(path: str) -> list:
    p = Path(path)
    ext = p.suffix.lower()
    try:
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            names = list(wb.sheetnames)
            wb.close()
            return names
        if ext == ".xls":
            return list(pd.ExcelFile(path).sheet_names)
    except Exception as e:      # corrupt / password-protected / not really an Excel file
        raise ValueError(f"Could not open {p.name} as an Excel workbook ({type(e).__name__}). It may be damaged, password-protected, or not really an Excel file - "
                         "open it in Excel and save a copy, or export the trial balance to CSV.") from e
    return [p.stem]


def _formula_cells_without_cache(path: str, sheet: str) -> int:
    """Count formula cells whose cached value is missing (file saved by a library, never opened in Excel)."""
    try:
        import openpyxl
        wf = openpyxl.load_workbook(path, data_only=False)
        wv = openpyxl.load_workbook(path, data_only=True)
        n = 0
        for row in wf[sheet].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and wv[sheet][c.coordinate].value is None:
                    n += 1
        return n
    except Exception:
        return 0


def _read_csv_rows(path: str) -> list:
    """Read a delimited text file whose rows may have different numbers of fields (title lines, totals...). Delimiter is detected, not assumed."""
    import csv
    raw = Path(path).read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin1"):
        try:
            text = raw.decode(enc)
            if enc == "utf-16" and "\x00" in text[:200]:
                continue
            break
        except UnicodeError:
            continue
    if text is None or not text.strip():
        raise ValueError("The file is empty or could not be decoded.")
    lines = text.splitlines()
    sample = [ln for ln in lines[:200] if ln.strip()]
    best, best_score = ",", -1
    for d in (",", ";", "\t", "|"):
        counts = [len(next(csv.reader([ln], delimiter=d))) for ln in sample]
        multi = [c for c in counts if c > 1]
        if not multi:
            continue
        mode = max(set(multi), key=multi.count)            # the usual width of a data row
        score = multi.count(mode) * 1000 + mode
        if score > best_score:
            best, best_score = d, score
    rows = [r for r in csv.reader(lines, delimiter=best)]
    return rows


def read_raw(path: str, sheet: Optional[str] = None) -> pd.DataFrame:
    """Return the sheet as a DataFrame with no header interpretation (all cells as objects)."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm", ".xls"):
        try:
            return pd.read_excel(path, sheet_name=sheet if sheet is not None else 0, header=None, dtype=object)
        except ValueError:
            raise
        except Exception as e:     # corrupt / password-protected / wrong extension
            raise ValueError(f"Could not open {p.name} as an Excel workbook ({type(e).__name__}). It may be damaged, password-protected, or not really an Excel file - "
                             "open it in Excel and save a copy, or export the trial balance to CSV.") from e
    if ext in (".csv", ".txt", ".tsv"):
        rows = _read_csv_rows(path)
        width = max((len(r) for r in rows), default=0)
        if not width:
            raise ValueError("The file contains no data.")
        return pd.DataFrame([r + [None] * (width - len(r)) for r in rows], dtype=object).replace({"": None})
    raise ValueError(f"Unsupported file type: {ext}")


# --------------------------------------------------------------------------------------
# Header detection
# --------------------------------------------------------------------------------------

def _role_of(header: str, kind: str) -> Optional[str]:
    h = clean_text(header).lower()
    if not h:
        return None
    h_ns = re.sub(r"[^a-z/ ]", "", h)
    has = lambda *w: any(x in h for x in w)
    if re.fullmatch(r"(dr\s*/\s*cr|cr\s*/\s*dr|d\s*/\s*c|dr\s*or\s*cr|side|balance type|dr/cr indicator|debit\s*/\s*credit)", h_ns.strip()):
        return "side"
    if kind == "tb":
        if has("opening", "brought forward", "b/f", "beginning"):
            if has("debit") or re.search(r"\bdr\b", h):
                return "opening_dr"
            if has("credit") or re.search(r"\bcr\b", h):
                return "opening_cr"
            return "opening"
        if has("closing", "ending", "cumulative", "c/f"):
            if has("debit") or re.search(r"\bdr\b", h):
                return "closing_dr"
            if has("credit") or re.search(r"\bcr\b", h):
                return "closing_cr"
            return "closing"
        if has("debit") or re.search(r"\bdr\b", h):
            return "debit"
        if has("credit") or re.search(r"\bcr\b", h):
            return "credit"
        if re.fullmatch(r"(net|balance|net balance|amount|total)", h_ns.strip()):
            return "net"
        if has("classif", "category", "nature", "group", "account type", "section") or h_ns.strip() in ("type",):
            return "hint"
        if re.fullmatch(r"(account\s*)?(code|no|number|id|#|ref)\.?|a/c\s*(code|no)\.?|gl code", h_ns.strip()) or h.strip() in ("#",):
            return "code"
        if re.fullmatch(r"(account(\s*(name|title|description))?|description|particulars|ledger(\s*name)?|name|gl account|account/description)", h_ns.strip()):
            return "name"
        if re.search(r"\b(account|ledger|particulars?|head|description|name)\b|a/c", h_ns) and not has("type", "group", "class", "no.", "number", "code"):
            return "name"
        return None
    # GL
    if (has("date") and not has("due")) or h_ns.strip() in ("dt", "txn dt", "trans dt", "posting dt", "vch dt"):
        return "date"
    if re.fullmatch(r"(account\s*)?(code|no|number|id)\.?|a/c\s*(code|no)\.?|gl code|acct\.?\s*(no|code)", h_ns.strip()):
        return "code"
    if re.fullmatch(r"(account(\s*name|\s*title)?|ledger(\s*name)?|gl account|account/description|a/c(\s*name)?|account description)", h_ns.strip()):
        return "name"
    if has("debit") or re.search(r"\bdr\b", h):
        return "debit"
    if has("credit") or re.search(r"\bcr\b", h):
        return "credit"
    if re.fullmatch(r"(amount|net|net amount|value)", h_ns.strip()):
        return "amount"
    if has("narration", "description", "memo", "particulars", "details", "remarks", "text") or h_ns.strip() in ("narr", "narr.", "desc", "desc."):
        return "description"
    if has("vendor", "supplier", "customer", "payee", "counterpart", "party", "name of", "contact"):
        return "counterparty"
    if has("reference", "ref", "voucher", "document", "journal", "entry no", "invoice", "txn", "transaction"):
        return "reference"
    return None


_GENERIC_SUB = {"debit", "credit", "dr", "cr", "amount", "balance", "net", "total", "debit amount", "credit amount"}


def _combined_labels(df: pd.DataFrame, i: int):
    """Header made of two stacked rows (e.g. 'Closing Balance' over 'Debit' / 'Credit'). Returns a list of labels or None."""
    if i < 1:
        return None
    top, cur = df.iloc[i - 1].tolist(), df.iloc[i].tolist()
    if not any(clean_text(c).lower() in _GENERIC_SUB for c in cur):
        return None
    if not any(clean_text(t) and to_num(t) is None for t in top):
        return None
    ff, last = [], ""
    for t in top:                      # a merged heading covers the blank cells to its right
        txt = clean_text(t)
        last = txt if txt else last
        ff.append(last)
    labels = []
    for j, c in enumerate(cur):
        lo = clean_text(c)
        if lo.lower() in _GENERIC_SUB:
            labels.append(f"{ff[j]} {lo}".strip())
        elif lo:
            labels.append(lo)
        else:
            labels.append(clean_text(top[j]))
    return labels


def find_header_row(df: pd.DataFrame, kind: str = "tb", max_scan: int = 40):
    """Return (row_index, colmap, duplicate_cols). colmap: role -> first column index with that role.
    A heading that spans two rows is combined ('Closing Balance' + 'Debit' -> closing debit); the lower row is returned."""
    best = (-1, 0, {}, [])
    for i in range(min(max_scan, len(df))):
        cands = [df.iloc[i].tolist()]
        comb = _combined_labels(df, i)
        if comb:
            cands.append(comb)
        for labels in cands:
            roles = {}
            dups = []
            for j, v in enumerate(labels):
                r = _role_of(v, kind)
                if r:
                    if r in roles:
                        dups.append((j, r))
                    else:
                        roles[r] = j
            if kind == "tb":
                has_amt = any(k in roles for k in ("debit", "credit", "net", "closing", "closing_dr", "opening", "opening_dr"))
                ok = has_amt and ("name" in roles or "code" in roles)
            else:
                ok = ("debit" in roles or "amount" in roles) and ("date" in roles or "description" in roles or "name" in roles)
            score = len(roles) + (3 if ok else 0)
            if ok and score > best[1]:
                best = (i, score, roles, dups)
    if best[0] < 0:
        return None, {}, []
    return best[0], best[2], best[3]


def colmap_for_row(df: pd.DataFrame, row: int, kind: str = "tb"):
    """Column roles for a header row chosen by the user (also tries the two-row form). Returns (colmap, duplicate_cols)."""
    best = ({}, [])
    cands = [df.iloc[row].tolist()] if 0 <= row < len(df) else []
    if 0 <= row < len(df):
        comb = _combined_labels(df, row)
        if comb:
            cands.append(comb)
    for labels in cands:
        roles, dups = {}, []
        for j, v in enumerate(labels):
            r = _role_of(v, kind)
            if r:
                if r in roles:
                    dups.append((j, r))
                else:
                    roles[r] = j
        if len(roles) > len(best[0]):
            best = (roles, dups)
    return best


def _looks_like_code(v) -> bool:
    s = clean_text(v)
    return bool(s) and bool(re.fullmatch(r"[\d][\d.\-/ ]*", s))


def _separate_code_and_name(df: pd.DataFrame, header_row: int, colmap: dict, dups: list):
    """If the first 'account' column holds codes/segments (101.000.1000.0000) and another column holds the description, use them as code and name."""
    if "name" not in colmap:
        return colmap, dups
    vals = [df.iat[i, colmap["name"]] for i in range(header_row + 1, min(len(df), header_row + 60)) if colmap["name"] < df.shape[1]]
    vals = [v for v in vals if clean_text(v)]
    if len(vals) >= 3 and sum(_looks_like_code(v) for v in vals) >= 0.7 * len(vals):
        alt = [j for j, r in dups if r == "name"]
        if alt and "code" not in colmap:
            colmap = dict(colmap)
            colmap["code"], colmap["name"] = colmap["name"], alt[0]
            dups = [(j, r) for j, r in dups if not (r == "name" and j == alt[0])]
        elif "code" not in colmap:
            colmap = dict(colmap)
            colmap["code"] = colmap.pop("name")
    return colmap, dups


def auto_pick_sheet(path: str, kind: str = "tb"):
    """Try every sheet and return (sheet_name, score) of the best candidate for a TB or GL."""
    scores = []
    for s in list_sheets(path):
        try:
            df = read_raw(path, s)
            i, roles, _ = find_header_row(df, kind)
            if i is not None:
                scores.append((s, len(roles) + (len(df) / 1000.0)))
        except Exception:
            continue
    if not scores:
        return None
    return sorted(scores, key=lambda x: -x[1])[0][0]


_LEGAL = re.compile(r"\b(llc|l\.l\.c|fze|fzco|fz-llc|ltd|limited|inc|pjsc|psc|co\.|company|group|holdings?|trading|establishment|est\.|corporation|corp|plc|gmbh|pvt|private)\b", re.I)
_NOT_ENTITY = re.compile(r"trial balance|general ledger|oracle|ledger vouchers|tally|^ledger:|period:|currency|run date|all amounts|amounts in|^\(?amounts|^page|closing pack|prepared by|^sheets:", re.I)


_HARD_NOT = re.compile(r"period:|currency|run date|prepared by|all amounts|amounts in|sheets:|closing pack|^ledger:", re.I)


def _guess_entity(title: str) -> str:
    parts = [clean_text(x) for x in re.split(r" \| ", title or "") if clean_text(x)]
    if len(parts) > 1:
        legal = [x for x in parts if _LEGAL.search(x) and not _HARD_NOT.search(x)]
        plain = [x for x in parts if not _NOT_ENTITY.search(x)]
        title = (legal or plain or parts)[0]
    t = clean_text(title)
    t = re.split(r"\b(revised|draft|final|trial balance|tb\b|as (at|on)|for the year|shared)", t, flags=re.I)[0]
    return t.strip(" -:,")


# --------------------------------------------------------------------------------------
# Trial balance
# --------------------------------------------------------------------------------------

_SKIP_NAME = re.compile(r"^\s*(grand\s+)?(sub)?totals?\b|^\s*report total|^\s*\*{3}|^\s*page \d+|^\s*difference\b|^\s*check\b|^\s*net (profit|loss)\b|^\s*balance\b", re.I)


def _rows_to_accounts(rows: list, colmap: dict, source: str, tb: TrialBalance, row_offset: int = 0, totals: Optional[list] = None) -> None:
    """Convert raw rows (list of lists) into Account objects using colmap. 'totals' collects (debit, credit) printed on total rows."""
    g = lambda r, role: (r[colmap[role]] if role in colmap and colmap[role] < len(r) else None)

    # Decide how the credit column is signed by looking at the data
    cred_vals = [to_num(g(r, "credit")) for r in rows]
    cred_vals = [v for v in cred_vals if v is not None and abs(v) > 0]
    deb_vals = [to_num(g(r, "debit")) for r in rows]
    deb_vals = [v for v in deb_vals if v is not None and abs(v) > 0]
    credit_signed = bool(cred_vals) and all(v < 0 for v in cred_vals)
    debit_has_neg = any(v < 0 for v in deb_vals)
    if "credit" in colmap:
        if credit_signed:
            tb.sign_convention = "Credit column holds negative numbers (Dr positive / Cr negative); net = debit + credit."
        else:
            tb.sign_convention = "Debit and credit columns hold positive numbers; net = debit - credit."
        if debit_has_neg:
            tb.warnings.append("Some values in the Debit column are negative - treated as credit balances. Please confirm in the preview.")
    skipped_groups = 0
    for k, r in enumerate(rows):
        name = clean_text(g(r, "name"))
        code = code_to_str(g(r, "code"))
        if not name and code and not isinstance(g(r, "code"), (int, float)):
            name, code = code, ""
        if not name and not code:
            continue
        if _SKIP_NAME.match(name):
            if totals is not None and re.match(r"^\s*(grand\s+)?totals?\b|^\s*report total", name, re.I):
                tv = {role: to_num(g(r, role)) for role in ("debit", "credit", "closing_dr", "closing_cr")}
                dr = tv["closing_dr"] if tv["closing_dr"] is not None else tv["debit"]
                cr = tv["closing_cr"] if tv["closing_cr"] is not None else tv["credit"]
                if dr is not None or cr is not None:
                    totals.append((abs(dr or 0.0), abs(cr or 0.0)))
            continue
        vals = {role: to_num(g(r, role)) for role in
                ("debit", "credit", "net", "closing", "closing_dr", "closing_cr", "opening", "opening_dr", "opening_cr")}
        if all(v is None for v in vals.values()):
            skipped_groups += 1      # heading row without numbers
            continue
        d = vals["debit"] or 0.0
        c = vals["credit"] or 0.0
        notes = []
        side = clean_text(g(r, "side")).lower() if "side" in colmap else ""
        sgn = -1.0 if side.startswith("c") else (1.0 if side.startswith("d") else None)
        # net balance
        if vals["closing_dr"] is not None or vals["closing_cr"] is not None:
            net = (vals["closing_dr"] or 0.0) - abs(vals["closing_cr"] or 0.0)
            debit_amt, credit_amt = max(net, 0.0), max(-net, 0.0)
        elif vals["closing"] is not None:
            net = vals["closing"]
            if sgn is not None:
                net = abs(net) * sgn
            debit_amt, credit_amt = max(net, 0.0), max(-net, 0.0)
        elif "debit" in colmap or "credit" in colmap:
            net = (d + c) if (credit_signed or debit_has_neg) else (d - abs(c))
            debit_amt, credit_amt = abs(d), abs(c)
            if vals["net"] is not None and abs(vals["net"] - net) > 0.01:
                notes.append(f"Net column ({vals['net']:,.2f}) differs from debit/credit ({net:,.2f}); net column used.")
                net = vals["net"]
        elif vals["net"] is not None:
            net = vals["net"]
            if sgn is not None:
                net = abs(net) * sgn
            debit_amt, credit_amt = max(net, 0.0), max(-net, 0.0)
        else:
            continue
        movement = None
        if ("debit" in colmap or "credit" in colmap) and (vals["closing_dr"] is not None or vals["closing_cr"] is not None or vals["closing"] is not None) \
                and (vals["opening"] is not None or vals["opening_dr"] is not None or vals["opening_cr"] is not None):
            movement = d - abs(c)
        opening = None
        if vals["opening"] is not None:
            opening = vals["opening"]
        elif vals["opening_dr"] is not None or vals["opening_cr"] is not None:
            opening = (vals["opening_dr"] or 0.0) - abs(vals["opening_cr"] or 0.0)
        acc = Account(code=code, name=name, debit=round(debit_amt, 6), credit=round(credit_amt, 6),
                      net=round(net, 6), opening=opening, movement=movement,
                      hint=clean_text(g(r, "hint")) or None,
                      source_ref=f"{source} row {k + 1 + row_offset}", notes=notes)
        tb.accounts.append(acc)
    if skipped_groups:
        tb.warnings.append(f"{skipped_groups} heading row(s) without amounts were skipped.")


def _strip_group_subtotals(tb: TrialBalance, totals: list) -> None:
    """Some systems (Tally 'group-wise' exports) print each group's total above its own ledgers. If the printed Grand Total (or the
    debit = credit test) shows the lines are double counted, find the rows that equal the sum of the rows beneath them, exclude them and
    use the top-level group as the classification of the ledgers. Accepted only when the remaining ledgers agree to the printed total."""
    accs = tb.accounts
    n = len(accs)
    if n < 4:
        return
    dr_sum = sum(max(a.net, 0.0) for a in accs)
    net_sum = sum(a.net for a in accs)
    if totals:
        if any(abs(dr_sum - t[0]) <= 0.5 for t in totals):
            return
    elif abs(net_sum) <= 0.5:
        return
    nets = [a.net for a in accs]
    memo: dict = {}

    def span(i):
        if i in memo:
            return memo[i]
        memo[i] = None
        if abs(nets[i]) < 0.005:
            return None
        j, acc = i + 1, 0.0
        while j < n:
            e = span(j)
            acc += nets[j]
            j = e if e else j + 1
            if abs(acc - nets[i]) <= 0.01:
                memo[i] = j
                return j
        return None

    is_group = [span(i) is not None for i in range(n)]
    leaves = [a for a, g in zip(accs, is_group) if not g]
    if len(leaves) == n or not leaves:
        return
    l_dr = sum(max(a.net, 0.0) for a in leaves)
    l_net = sum(a.net for a in leaves)
    # the printed Grand Total adds up the top-level groups (each already net of its contents), so compare with that as well
    top_dr, depth_end = 0.0, 0
    for i, a in enumerate(accs):
        if is_group[i] and i >= depth_end:
            top_dr += max(a.net, 0.0)
            depth_end = memo[i]
    ok = abs(l_net) <= 0.5 and (not totals or any(abs(l_dr - t[0]) <= 0.5 or abs(top_dr - t[0]) <= 0.5 for t in totals))
    if not ok:
        tb.warnings.append("The debit/credit totals do not agree to the lines and group sub-totals may be included, but they could not be isolated "
                           "reliably - check the preview against the source file.")
        return
    # classification of each ledger from its top-level group
    stack = []          # (end, name)
    for i, a in enumerate(accs):
        while stack and i >= stack[-1][0]:
            stack.pop()
        if is_group[i]:
            stack.append((memo[i], a.name))
            continue
        if stack:
            a.group = " > ".join(s_[1] for s_ in stack)
            if not a.hint:
                a.hint = stack[0][1]
    names = [a.name for a, g in zip(accs, is_group) if g]
    tb.accounts = leaves
    tb.warnings.append(f"{len(names)} group sub-total row(s) in the source (e.g. {', '.join(names[:4])}) were excluded to avoid double counting; "
                       "the ledgers beneath them agree to the printed total. The top-level group is used as each ledger's classification. Please confirm in the preview.")


def parse_tb(path: str, sheet: Optional[str] = None, currency: str = "", header_row: Optional[int] = None,
             colmap: Optional[dict] = None, ocr: str = "auto") -> TrialBalance:
    """Read a trial balance from Excel/CSV. Pass header_row/colmap to override auto-detection."""
    p = Path(path)
    if p.suffix.lower() == ".pdf":
        return parse_tb_pdf(path, currency=currency, ocr=ocr)
    if sheet is None and p.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        sheet = auto_pick_sheet(path, "tb") or list_sheets(path)[0]
    df = read_raw(path, sheet)
    tb = TrialBalance(source_file=p.name, sheet=sheet or "", currency=currency)
    if header_row is not None and colmap is None:
        colmap, _d = colmap_for_row(df, header_row, "tb")
        if not colmap:
            raise ValueError(f"Row {header_row + 1} does not look like a heading row (no account / debit / credit columns found in it).")
    if header_row is None or colmap is None:
        hr, cm, dups = find_header_row(df, "tb")
        if hr is None:
            raise ValueError("Could not find a header row with account and debit/credit/balance columns. "
                             "Pick the header row and columns manually in the extraction preview.")
        cm, dups = _separate_code_and_name(df, hr, cm, dups)
        header_row, colmap = hr, cm
        if dups:
            hdr = df.iloc[hr].tolist()
            tb.ignored_columns = [f"column {j + 1} '{clean_text(hdr[j])}' (duplicate {r})" for j, r in dups]
            tb.warnings.append("Repeated column headings found - only the first occurrence of each was used and the rest were ignored "
                               "(they usually hold translated amounts): " + "; ".join(tb.ignored_columns))
    # title: any text above the header row
    above = [clean_text(v) for i in range(header_row) for v in df.iloc[i].tolist() if clean_text(v)]
    tb.title = " | ".join(above)
    rows = df.iloc[header_row + 1:].values.tolist()
    totals: list = []
    _rows_to_accounts(rows, colmap, f"{p.name}/{sheet}" if sheet else p.name, tb, row_offset=header_row + 1, totals=totals)
    _strip_group_subtotals(tb, totals)
    if p.suffix.lower() in (".xlsx", ".xlsm") and sheet:
        n = _formula_cells_without_cache(path, sheet)
        if n:
            tb.warnings.append(f"{n} formula cell(s) have no stored value (the file was saved by a program that did not calculate). "
                               "Open the file in Excel, save it, and re-upload for reliable numbers.")
    if not tb.accounts:
        tb.warnings.append("No account rows were extracted.")
    return tb


def entity_guess(tb: TrialBalance) -> str:
    return _guess_entity(tb.title)


# --------------------------------------------------------------------------------------
# PDF (text or scanned)
# --------------------------------------------------------------------------------------

_LINE_NUM = r"\(?-?[\d,]+(?:\.\d+)?\)?(?:\s*(?:Dr|Cr|DR|CR))?"
_LINE_RE = re.compile(rf"^(?P<head>.*?[A-Za-z].*?)\s+(?P<nums>(?:{_LINE_NUM}\s*){{1,6}})$")


_AMT_TOKEN = re.compile(r"^\(?-?\d[\d,]*(?:\.\d+)?\)?-?$")
_ROLE_WORDS = {"debit": "debit", "debits": "debit", "dr": "debit", "credit": "credit", "credits": "credit", "cr": "credit",
               "balance": "net", "net": "net", "closing": "closing", "opening": "opening"}


def _pdf_words(path: str):
    """Words with positions from the text layer: list per page of dicts(text,x0,x1,top)."""
    import pdfplumber
    pages = []
    with pdfplumber.open(path) as pdf:
        for pg in pdf.pages:
            ws = pg.extract_words(keep_blank_chars=False, use_text_flow=False) or []
            pages.append([{"text": w["text"], "x0": w["x0"], "x1": w["x1"], "top": w["top"], "bottom": w["bottom"], "conf": 100.0} for w in ws])
    return pages


def _ocr_words(path: str, dpi: int = 300):
    import pypdfium2 as pdfium
    import pytesseract
    pdf = pdfium.PdfDocument(path)
    pages, confs = [], []
    for i in range(len(pdf)):
        img = pdf[i].render(scale=dpi / 72).to_pil().convert("L")
        d = pytesseract.image_to_data(img, config="--psm 6", output_type=pytesseract.Output.DICT)
        ws = []
        for t, cf, l, tp, w, h in zip(d["text"], d["conf"], d["left"], d["top"], d["width"], d["height"]):
            t = clean_text(t).strip("|[]{}")
            if not t:
                continue
            try:
                cf = float(cf)
            except ValueError:
                cf = -1.0
            if cf >= 0:
                confs.append(cf)
            ws.append({"text": t, "x0": l, "x1": l + w, "top": tp, "bottom": tp + h, "conf": cf})
        pages.append(ws)
    return pages, (sum(confs) / len(confs) if confs else 0.0)


def _group_lines(words: list):
    words = sorted(words, key=lambda w: (w["top"], w["x0"]))
    lines, cur, cur_top = [], [], None
    for w in words:
        h = max(w["bottom"] - w["top"], 1)
        if cur and abs(w["top"] - cur_top) <= 0.6 * h:
            cur.append(w)
        else:
            if cur:
                lines.append(sorted(cur, key=lambda x: x["x0"]))
            cur, cur_top = [w], w["top"]
    if cur:
        lines.append(sorted(cur, key=lambda x: x["x0"]))
    return lines


def _positional_rows(pages_words: list):
    """Column-aware parsing: the header words 'Debit' / 'Credit' / 'Balance' fix the x position of each amount column, and every
    amount is assigned to the nearest column. Returns (title_text, rows, unparsed_count, found_header)."""
    title, rows, unparsed = [], [], 0
    centers = None
    for words in pages_words:
        for ln in _group_lines(words):
            texts = [w["text"] for w in ln]
            roles = [(_ROLE_WORDS[t.lower().strip(".:")], (w["x0"] + w["x1"]) / 2) for t, w in zip(texts, ln) if t.lower().strip(".:") in _ROLE_WORDS]
            if roles and not any(_AMT_TOKEN.match(t) for t in texts) and len(texts) <= 10:
                # a header line (also repeated on later pages). 'Closing Balance' -> one column
                cols = {}
                for r, x in roles:
                    key = "closing" if r in ("closing",) else r
                    if r == "net" and "closing" in cols:
                        continue
                    cols[key] = x
                if centers is None or len(cols) >= len(centers):
                    centers = cols
                continue
            if centers is None:
                title.append(" ".join(texts))
                continue
            xs = sorted(centers.values())
            spacing = min([b - a for a, b in zip(xs, xs[1:])] or [250.0])
            tol = max(spacing * 0.6, 40.0)
            amounts, left = [], []
            for w in ln:
                cx = (w["x0"] + w["x1"]) / 2
                near = min(centers.items(), key=lambda kv: abs(kv[1] - cx))
                if _AMT_TOKEN.match(w["text"]) and abs(near[1] - cx) <= tol and (w["x0"] > min(xs) - tol * 2):
                    amounts.append((near[0], to_num(w["text"]), w))
                else:
                    left.append(w)
            if not amounts:
                txt = " ".join(x["text"] for x in left)
                if re.search(r"[A-Za-z]{3,}", txt) and not _SKIP_NAME.match(txt):
                    unparsed += 1
                continue
            name_words = [w["text"] for w in left]
            code = ""
            if name_words and re.match(r"^\d[\d.\-]*$", name_words[0]) and len(name_words) > 1:
                code, name_words = name_words[0], name_words[1:]
            name = clean_text(" ".join(name_words))
            if not name or _SKIP_NAME.match(name):
                continue
            row = {"code": code, "name": name, "conf": min(a[2]["conf"] for a in amounts)}
            for role, val, _ in amounts:
                row[role] = (row.get(role) or 0.0) + (val or 0.0)
            rows.append(row)
    return " ".join(title), rows, unparsed, centers is not None


def _pdf_page_text_and_tables(path: str):
    import pdfplumber
    pages = []
    with pdfplumber.open(path) as pdf:
        for pg in pdf.pages:
            try:
                tables = pg.extract_tables()
            except Exception:
                tables = []
            pages.append({"text": pg.extract_text() or "", "tables": tables})
    return pages


def _ocr_pdf(path: str, dpi: int = 300):
    """Render pages with pypdfium2 (no poppler needed) and OCR with Tesseract. Returns (list_of_text, mean_conf)."""
    import pypdfium2 as pdfium
    import pytesseract
    pdf = pdfium.PdfDocument(path)
    texts, confs = [], []
    for i in range(len(pdf)):
        img = pdf[i].render(scale=dpi / 72).to_pil()
        data = pytesseract.image_to_data(img, config="--psm 6", output_type=pytesseract.Output.DICT)
        lines = {}
        for t, conf, b, p_, l in zip(data["text"], data["conf"], data["block_num"], data["par_num"], data["line_num"]):
            if clean_text(t):
                lines.setdefault((b, p_, l), []).append(clean_text(t))
                try:
                    cf = float(conf)
                    if cf >= 0:
                        confs.append(cf)
                except ValueError:
                    pass
        texts.append("\n".join(" ".join(v) for v in lines.values()))
    return texts, (sum(confs) / len(confs) if confs else 0.0)


def parse_tb_pdf(path: str, currency: str = "", ocr: str = "auto") -> TrialBalance:
    p = Path(path)
    tb = TrialBalance(source_file=p.name, currency=currency, extraction_method="pdf-text")
    pages = _pdf_page_text_and_tables(path)
    all_text = "\n".join(pg["text"] for pg in pages)
    use_ocr = ocr == "force" or (ocr == "auto" and len(re.sub(r"\s", "", all_text)) < 30)
    table_rows = []
    if not use_ocr:
        for pg in pages:
            for t in pg["tables"]:
                if t and max(len(r) for r in t) >= 3:
                    table_rows.extend(t)
    if table_rows and not use_ocr:
        df = pd.DataFrame(table_rows, dtype=object)
        hr, cm, dups = find_header_row(df, "tb")
        if hr is not None:
            tb.extraction_method = "pdf-table"
            tb.title = " ".join(clean_text(v) for i in range(hr) for v in df.iloc[i].tolist() if clean_text(v))
            _rows_to_accounts(df.iloc[hr + 1:].values.tolist(), cm, p.name, tb)
            tb.warnings.append("Extracted from PDF tables - please check the preview carefully against the PDF.")
            return tb
    # column-aware parsing (text layer or OCR) - decides debit/credit by the x position of each amount
    try:
        if use_ocr:
            pw, conf = _ocr_words(path)
            tb.extraction_method = "pdf-ocr"
        else:
            pw, conf = _pdf_words(path), 100.0
    except Exception as e:                       # noqa
        pw, conf = None, 0.0
        if use_ocr:
            tb.warnings.append(f"The PDF has no text layer and OCR failed ({e}). Install Tesseract OCR or upload Excel/CSV.")
            return tb
    if pw is not None:
        title, prow, unp, found = _positional_rows(pw)
        if found and prow:
            tb.title = title
            if not use_ocr:
                tb.extraction_method = "pdf-text-positional"
            for r in prow:
                net = r.get("closing", r.get("net"))
                if net is None:
                    net = (r.get("debit") or 0.0) - abs(r.get("credit") or 0.0)
                tb.accounts.append(Account(code=r["code"], name=r["name"], debit=max(net, 0), credit=max(-net, 0), net=net,
                                           source_ref=f"{p.name} (PDF{', OCR' if use_ocr else ''})",
                                           notes=([f"OCR confidence {r['conf']:.0f}%"] if use_ocr else [])))
            if use_ocr:
                tb.warnings.append(f"Scanned PDF read by OCR (mean confidence {conf:.0f}%). OCR can misread digits - verify every amount against the PDF; "
                                   "the debit/credit totals below must agree.")
                if conf < 75:
                    tb.warnings.append("Low OCR confidence - results are unreliable. Prefer an Excel/CSV export.")
                low = [a for a, r in zip(tb.accounts, prow) if r["conf"] < 80]
                if low:
                    tb.warnings.append(f"{len(low)} line(s) were read with low confidence: " + ", ".join(a.name[:25] for a in low[:6]) + ("..." if len(low) > 6 else ""))
            else:
                tb.warnings.append("Read from the PDF text layer by column position - please check the preview against the PDF.")
            if unp:
                tb.warnings.append(f"{unp} line(s) of the PDF could not be read as accounts and were skipped.")
            return tb
    # line parsing (text layer or OCR) - fallback when no column headings were found
    if use_ocr:
        try:
            texts, conf = _ocr_pdf(path)
        except Exception as e:           # noqa
            tb.warnings.append(f"The PDF has no text layer and OCR failed ({e}). Install Tesseract OCR or upload Excel/CSV.")
            return tb
        text = "\n".join(texts)
        tb.extraction_method = "pdf-ocr"
        tb.warnings.append(f"Scanned PDF read by OCR (mean confidence {conf:.0f}%). OCR errors are possible - verify every amount.")
        if conf < 75:
            tb.warnings.append("Low OCR confidence - results are unreliable. Prefer an Excel/CSV export.")
    else:
        text = all_text
    lines = [clean_text(x) for x in text.splitlines() if clean_text(x)]
    header_roles = None
    unparsed = 0
    for ln in lines:
        low = ln.lower()
        if header_roles is None and ("debit" in low or "credit" in low or "balance" in low):
            order = []
            for w in re.findall(r"opening|debit|credit|closing|balance|net|dr\b|cr\b", low):
                order.append({"dr": "debit", "cr": "credit", "balance": "net"}.get(w, w))
            if order:
                header_roles = order
                tb.title = tb.title or ""
                continue
        if header_roles is None:
            tb.title = (tb.title + " " + ln).strip() if len(tb.title) < 200 else tb.title
            continue
        m = _LINE_RE.match(ln)
        if not m or _SKIP_NAME.match(m.group("head")):
            if not _SKIP_NAME.match(ln):
                unparsed += 1
            continue
        head = m.group("head").strip()
        nums = re.findall(_LINE_NUM, m.group("nums"))
        nums = [to_num(n) for n in nums]
        code = ""
        mm = re.match(r"^(\d[\d.\-]*)\s+(.*)$", head)
        if mm:
            code, head = mm.group(1), mm.group(2)
        row = {"code": code, "name": head}
        if len(nums) == len(header_roles):
            for role, v in zip(header_roles, nums):
                row[role] = v
        else:
            # ambiguous: a blank cell makes the column of each number uncertain
            row_notes = f"PDF line has {len(nums)} amount(s) but {len(header_roles)} column(s): column assignment is a guess - verify."
            if len(nums) == 1:
                row["net"] = nums[0]
            elif len(nums) == 2 and "debit" in header_roles and "credit" in header_roles:
                row["debit"], row["credit"] = nums[0], nums[1]
            else:
                tb.warnings.append(f"Line '{ln[:60]}' could not be assigned to columns and was skipped.")
                continue
            row["_note"] = row_notes
        d, c, net = row.get("debit"), row.get("credit"), row.get("closing", row.get("net"))
        if net is None:
            net = (d or 0.0) - abs(c or 0.0)
        acc = Account(code=code, name=head, debit=max(net, 0), credit=max(-net, 0), net=net,
                      source_ref=f"{p.name} (PDF)", notes=[row["_note"]] if "_note" in row else [])
        tb.accounts.append(acc)
    if unparsed:
        tb.warnings.append(f"{unparsed} line(s) of the PDF could not be read as accounts and were skipped.")
    if not tb.accounts:
        tb.warnings.append("No accounts could be read from this PDF.")
    return tb


# --------------------------------------------------------------------------------------
# General ledger
# --------------------------------------------------------------------------------------

def parse_gl(path: str, sheet: Optional[str] = None, header_row: Optional[int] = None,
             colmap: Optional[dict] = None) -> GeneralLedger:
    p = Path(path)
    if p.suffix.lower() == ".pdf":
        raise ValueError("GL PDFs are not supported in this version: please export the GL to Excel or CSV. "
                         "(TB PDFs, including scanned ones, are supported.)")
    if sheet is None and p.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        sheet = auto_pick_sheet(path, "gl") or list_sheets(path)[0]
    df = read_raw(path, sheet)
    gl = GeneralLedger(source_file=p.name, sheet=sheet or "")
    if header_row is not None and colmap is None:
        colmap, _d = colmap_for_row(df, header_row, "gl")
        if not colmap:
            raise ValueError(f"Row {header_row + 1} does not look like a heading row (no date / account / debit / credit columns found in it).")
    if header_row is None or colmap is None:
        hr, cm, dups = find_header_row(df, "gl")
        if hr is None:
            raise ValueError("Could not find the GL header row (need at least Date, Debit/Amount and Description/Account columns).")
        cm, dups = _separate_code_and_name(df, hr, cm, dups)
        header_row, colmap = hr, cm
    g = lambda r, role: (r[colmap[role]] if role in colmap and colmap[role] < len(r) else None)
    grouped = "name" not in colmap and "code" not in colmap
    cur_code, cur_name = "", ""
    if grouped:       # the first ledger's heading sits directly above the first column-heading row
        for i in range(header_row - 1, max(-1, header_row - 4), -1):
            cells = [clean_text(v) for v in df.iloc[i].tolist() if clean_text(v)]
            if len(cells) == 1 and re.match(r"^(account|ledger|a/c)\b", cells[0], re.I):
                m = re.match(r"^(?:(?:account|ledger|a/c)\s*[:\-]?\s*)?(\d[\w.\-]*)\s*[-–:]?\s+(.+)$", cells[0], re.I)
                cur_code, cur_name = (m.group(1), m.group(2)) if m else ("", re.sub(r"^(account|ledger|a/c)(\s*name)?\s*[:\-]?\s*", "", cells[0], flags=re.I))
                break
    n_nodate = n_open = 0
    rows = df.iloc[header_row + 1:].values.tolist()
    for k, r in enumerate(rows):
        first = clean_text(r[0]) if len(r) else ""
        dt = to_iso_date(g(r, "date")) if "date" in colmap else None
        deb = to_num(g(r, "debit")) or 0.0
        cre = to_num(g(r, "credit")) or 0.0
        if "amount" in colmap and "debit" not in colmap:
            amt = to_num(g(r, "amount")) or 0.0
            sd = clean_text(g(r, "side")).lower() if "side" in colmap else ""
            if sd.startswith("c"):
                deb, cre = 0.0, abs(amt)
            elif sd.startswith("d"):
                deb, cre = abs(amt), 0.0
            else:
                deb, cre = (amt, 0.0) if amt >= 0 else (0.0, -amt)
        has_amt = abs(deb) > 0 or abs(cre) > 0
        if grouped:
            # account heading rows: text in one cell and no amount/date
            if first and not has_amt and dt is None and sum(1 for v in r if clean_text(v)) <= 2 and not _SKIP_NAME.match(first):
                m = re.match(r"^(?:(?:account|ledger|a/c)\s*[:\-]?\s*)?(\d[\w.\-]*)\s*[-–:]?\s+(.+)$", first, re.I)
                cur_code, cur_name = (m.group(1), m.group(2)) if m else ("", re.sub(r"^(account|ledger|a/c)(\s*name)?\s*[:\-]?\s*", "", first, flags=re.I))
                continue
            code, name = cur_code, cur_name
        else:
            code = code_to_str(g(r, "code"))
            name = clean_text(g(r, "name"))
            if not code and name:
                m = re.match(r"^(\d[\w.\-]*)\s*[-–:]\s*(.+)$", name) or re.match(r"^(\d{3,}[\d.\-]*)\s+(\D.*)$", name)
                if m:
                    code, name = m.group(1), m.group(2)
        if not has_amt:
            continue
        desc = clean_text(g(r, "description"))
        if _SKIP_NAME.match(desc or name or first):
            continue
        if re.match(r"^\s*(closing balance|closing|current total|total|balance c/?[fd])\b", desc or "", re.I):
            continue      # running / closing balance printed by the system - not a transaction
        is_open = bool(re.search(r"opening balance|balance (b/?f|brought)|b/f\b", desc or "", re.I) and dt is None)
        if is_open:
            n_open += 1
        elif dt is None:
            n_nodate += 1
        gl.lines.append(GLLine(line_id=f"GL{len(gl.lines) + 1:06d}", date=dt, account_code=code, account_name=name,
                               description=desc, debit=abs(deb), credit=abs(cre),
                               counterparty=clean_text(g(r, "counterparty")), reference="OPENING" if is_open else clean_text(g(r, "reference")),
                               source_ref=f"{p.name}/{sheet} row {header_row + k + 2}"))
    if n_nodate:
        gl.warnings.append(f"{n_nodate} GL line(s) have no readable date.")
    if n_open:
        gl.warnings.append(f"{n_open} opening-balance line(s) without a date were kept (reference OPENING) so that the GL agrees to the closing TB; "
                           "they are not treated as transactions in the risk tests.")
    if not gl.lines:
        gl.warnings.append("No GL transactions were extracted.")
    return gl


def _norm_label(x: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (x or "").lower()).strip()


def align_gl_to_tb(gl, tb) -> int:
    """Give every GL line the same account code / name as its trial-balance account (match by code, else by name), so that reconciliation and
    the transaction tests look the same account up the same way. Returns the number of distinct GL accounts that could not be matched."""
    if not gl or not tb or not tb.accounts:
        return 0
    by_code = {a.code.strip(): a for a in tb.accounts if a.code and a.code.strip()}
    by_name = {}
    for a in tb.accounts:
        by_name.setdefault(_norm_label(a.name), a)
        if a.code:
            by_name.setdefault(_norm_label(f"{a.code} {a.name}"), a)
    cache, unmatched, changed = {}, [], 0
    for l in gl.lines:
        k = (l.account_code, l.account_name)
        if k not in cache:
            acc = by_code.get((l.account_code or "").strip()) or by_name.get(_norm_label(l.account_name)) or by_name.get(_norm_label(l.account_code))
            cache[k] = acc
            if acc is None:
                unmatched.append(l.account_name or l.account_code)
        acc = cache[k]
        if acc is not None and (l.account_code != acc.code or l.account_name != acc.name):
            l.account_code, l.account_name = acc.code, acc.name
            changed += 1
    msg = f"GL account labels were aligned to the trial balance for {changed} line(s) (matched by account code, otherwise by name)."
    if changed and msg not in gl.warnings:
        gl.warnings.append(msg)
    return len(set(unmatched))
