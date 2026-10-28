"""Word tax-risk and review report (brief section 9: fifteen sections, four status words, preliminary-review statement).

Language rule: a flagged item is never described as 'disallowed' unless it is a CONFIRMED adjustment. Everything else is
'potential', 'information required' or 'technical review required'.
If the user supplies a letterhead .docx, it is used as the base document (its header/footer/styles are kept; any body text in it is removed).
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .engine import ROWS
from .mapping import fs_lines
from .models import load_json
from .questionnaire import ans, unanswered

NAVY = RGBColor(0x28, 0x3A, 0x4A)
STATUS_COL = {"Confirmed": "E2EFDA", "Potential adjustment": "FFF2CC", "Information required": "FCE4D6", "Technical review required": "DDEBF7",
              "Rejected": "EDEDED", "High": "F8CBAD", "Medium": "FFF2CC", "Low": "E2EFDA"}
PRELIM = ("This report is the output of an automated, preliminary review of the trial balance and general ledger supplied. It does not replace a complete "
          "professional review of the company's records, contracts, elections and supporting documentation. Figures are draft, in AED, and no tax position "
          "has been taken. Items shown as potential adjustments are NOT conclusions that an expense is disallowed or that income is exempt.")


def _d(s):
    try:
        return datetime.fromisoformat(s).strftime("%d %B %Y")
    except Exception:
        return s or "not stated"


def _n(x, dp=0):
    if x is None or x == "":
        return ""
    if abs(x) < 0.5 * 10 ** -dp:
        x = 0.0                      # avoid '-0' and '(0)'
    return f"{x:,.{dp}f}" if x >= 0 else f"({-x:,.{dp}f})"


def _shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


class ReportBuilder:
    def __init__(self, an, letterhead: str = ""):
        self.an = an
        lh = letterhead or an.project.settings.letterhead_path
        if lh and Path(lh).exists() and Path(lh).suffix.lower() == ".docx":
            self.doc = Document(lh)
            body = self.doc.element.body
            for el in list(body):
                if el.tag != qn("w:sectPr"):
                    body.remove(el)
            self.letterhead_used = True
        else:
            self.doc = Document()
            self.letterhead_used = False
        self._setup()

    # ------------------------------------------------------------------ basics
    def _setup(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
        st = d.styles["Normal"]
        st.font.name = "Arial"
        st.font.size = Pt(9.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        for name, size in (("Heading 1", 14), ("Heading 2", 11.5), ("Heading 3", 10)):
            if name in [s.name for s in d.styles]:
                h = d.styles[name]
                h.font.name = "Arial"
                h.font.size = Pt(size)
                h.font.bold = True
                h.font.color.rgb = NAVY
                rf = h.element.rPr.rFonts
                for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                    if rf.get(qn(att)) is not None:
                        del rf.attrib[qn(att)]
                rf.set(qn("w:ascii"), "Arial")
                rf.set(qn("w:hAnsi"), "Arial")
                rf.set(qn("w:eastAsia"), "Arial")
        if not self.letterhead_used:
            fp = sec.footer.paragraphs[0]
            fp.text = f"{self.an.entity} - Corporate Tax review report - DRAFT / PRELIMINARY"
            fp.runs[0].font.size = Pt(7.5)
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._page_number(sec)

    def _page_number(self, sec):
        p = sec.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Page ")
        run.font.size = Pt(7.5)
        for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            r = p.add_run()
            r.font.size = Pt(7.5)
            if t:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), t)
                r._r.append(fc)
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = txt
                r._r.append(it)

    def h(self, text, level=1):
        return self.doc.add_heading(text, level)

    def p(self, text="", bold=False, italic=False, size=None, color=None, after=4, align=None):
        para = self.doc.add_paragraph()
        r = para.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        para.paragraph_format.space_after = Pt(after)
        if align:
            para.alignment = align
        return para

    def bullets(self, items):
        for t in items:
            para = self.doc.add_paragraph(style="List Bullet")
            para.add_run(t)
            para.paragraph_format.space_after = Pt(2)

    def note(self, text):
        self.p(text, italic=True, size=8.5, color="595959")

    def table(self, headers, rows, widths_cm, status_col=None, num_cols=(), font=8):
        if not rows:
            return None
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        hdr = t.rows[0]
        _repeat_header(hdr)
        for i, hd in enumerate(headers):
            c = hdr.cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(hd)
            r.bold = True
            r.font.size = Pt(font)
            r.font.color.rgb = RGBColor(255, 255, 255)
            _shade(c, "283A4A")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                run = para.add_run("" if v is None else str(v))
                run.font.size = Pt(font)
                if i in num_cols:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if status_col is not None and i == status_col and str(v) in STATUS_COL:
                    _shade(cells[i], STATUS_COL[str(v)])
                    run.bold = True
        for i, w in enumerate(widths_cm):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def kv(self, rows, w=(5.2, 12.2)):
        t = self.doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        t.autofit = False
        for k, v in rows:
            cells = t.add_row().cells
            cells[0].text, cells[1].text = "", ""
            r = cells[0].paragraphs[0].add_run(k)
            r.bold = True
            r.font.size = Pt(8.5)
            _shade(cells[0], "EDEDED")
            cells[1].paragraphs[0].add_run("" if v is None else str(v)).font.size = Pt(8.5)
            cells[0].width, cells[1].width = Cm(w[0]), Cm(w[1])
        t.columns[0].width, t.columns[1].width = Cm(w[0]), Cm(w[1])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ------------------------------------------------------------------ content
    def build(self, path: str):
        an = self.an
        d = self.doc
        title = d.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = title.add_run("UAE Corporate Tax\nTax Risk and Review Report")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(20), NAVY
        self.p(f"{an.entity}", bold=True, size=13)
        self.p(f"Tax period: {_d(an.period_start)} to {_d(an.period_end)}   |   All amounts in AED   |   Prepared {datetime.now().strftime('%d %B %Y')}", size=9)
        self.p("DRAFT - PRELIMINARY REVIEW", bold=True, color="C00000", size=10)
        self.p(PRELIM, italic=True, size=8.5, color="C00000")
        self.h("Status key", 2)
        self.table(["Status", "Meaning"], [
            ["Confirmed", "Supported by sufficient evidence and an applicable rule; included in the computation."],
            ["Potential adjustment", "Requires further evidence or confirmation; NOT included in the computation."],
            ["Information required", "Cannot be concluded from the supplied records."],
            ["Technical review required", "Requires professional judgement or additional research."]], [4.5, 12.9], status_col=0)
        self.s1_company()
        self.s2_scope()
        self.s3_files()
        self.s4_profit()
        self.s5_computation()
        self.s6_disallowed()
        self.s7_exempt()
        self.s8_reliefs()
        self.s9_risk_accounts()
        self.s10_gl()
        self.s11_missing()
        self.s12_questions()
        self.s13_impact()
        self.s14_technical()
        self.s15_next()
        self.appendices()
        self.h("Important notice", 2)
        self.p(PRELIM, italic=True, size=8.5)
        d.save(path)
        return path

    def s1_company(self):
        an = self.an
        self.h("1. Company and tax period", 1)
        A = an.answers
        self.kv([("Taxable person", an.entity), ("Registration status", ans(A, "entity_status", "Not stated")),
                 ("Tax period", f"{_d(an.period_start)} to {_d(an.period_end)}"),
                 ("First Corporate Tax period?", ans(A, "first_period", "Not stated")),
                 ("Accounting framework / basis", f"{ans(A, 'accounting_framework', 'Not stated')} / {ans(A, 'accounting_basis', 'Not stated')}"),
                 ("Status of the financial statements", ans(A, "tb_status", "Not stated")),
                 ("Reporting currency of the books", an.project.fx.tb_currency), ("Currency of this report and computation", "UAE dirham (AED)")])

    def s2_scope(self):
        an = self.an
        p = an.project
        self.h("2. Scope of data reviewed", 1)
        bullets = [f"Trial balance: {len(p.tb.accounts)} account(s)." if p.tb else "No trial balance."]
        bullets.append(f"General ledger: {len(p.gl.lines):,} line(s)." if p.gl else "No general ledger was supplied - transaction-level tests were NOT performed.")
        if p.prior_tb:
            bullets.append(f"Prior-period trial balance supplied ({len(p.prior_tb.accounts)} accounts) - used for movement analysis only.")
        bullets.append("Supporting documents (invoices, contracts, agreements, share registers, tax certificates) were not reviewed by the tool: "
                       + (f"{len(p.supporting_files)} file(s) are listed in the project." if p.supporting_files else "none were listed in the project."))
        self.bullets(bullets)
        mr = an.tr.method_record
        self.h("Conversion to AED", 3)
        if p.fx.is_aed:
            self.p("The trial balance is in AED; no conversion was required.")
        else:
            self.kv([("Books currency", mr.get("currency", "")), ("Method", mr.get("method_description", "")), ("Reason recorded", mr.get("reason", "")),
                     ("Rate source", mr.get("rate_source", "")), ("Basis", mr.get("decision", ""))])
            self.table(["Purpose", "Rate (AED per unit)", "Basis"], [[x["purpose"], f"{x['rate']:.6f}", x["basis"]] for x in an.tr.rate_table], [7, 3.4, 7])
            if an.tr.diff_applicable:
                self.p(f"A derived currency translation difference of AED {_n(an.tr.translation_diff)} arises in equity. It is a computed balancing figure, not a "
                       "trial-balance amount, and needs approval.", color="C00000", size=9)

    def s3_files(self):
        an = self.an
        p = an.project
        self.h("3. Files and worksheets analysed", 1)
        rows = []
        if p.tb:
            rows.append(["Trial balance", p.tb.source_file, p.tb.sheet or "-", p.tb.extraction_method, f"{len(p.tb.accounts)} accounts", "; ".join(p.tb.warnings) or "-"])
        if p.gl:
            rows.append(["General ledger", p.gl.source_file, p.gl.sheet or "-", p.gl.extraction_method, f"{len(p.gl.lines):,} lines", "; ".join(p.gl.warnings) or "-"])
        if p.prior_tb:
            rows.append(["Prior-period TB", p.prior_tb.source_file, p.prior_tb.sheet or "-", p.prior_tb.extraction_method, f"{len(p.prior_tb.accounts)} accounts", "-"])
        self.table(["Item", "File", "Worksheet", "Extraction", "Records", "Extraction warnings"], rows, [2.6, 4.2, 2.4, 2.0, 2.2, 4.0])
        self.note("Source files are read only and are never modified by the tool.")
        fails = [c for c in an.checks if c.status in ("FAIL", "WARN")]
        if fails:
            self.h("Reconciliation exceptions", 3)
            self.table(["ID", "Check", "Result", "Detail"], [[c.id, c.description, c.status, c.detail] for c in fails], [1.8, 4.6, 1.5, 9.5])

    def s4_profit(self):
        an = self.an
        st = an.st
        self.h("4. Accounting profit or loss", 1)
        rows = []
        for r in st.pl_rows:
            if r.amount is None:
                continue
            tag = "" if r.kind == "line" else "  "
            rows.append([f"{tag}{r.label}" + (" (subtotal)" if r.kind == "subtotal" else ""), _n(r.amount)])
        self.table(["Line", "AED"], rows, [12.5, 4.9], num_cols=(1,))
        self.p(f"Profit/(loss) for the year per the AED statement of profit or loss prepared from the trial balance: AED {_n(st.profit_for_year)}.")
        if st.unmapped_accounts:
            self.p(f"{len(st.unmapped_accounts)} unmapped account(s) (net AED {_n(st.unmapped_net)}) are excluded until mapped - the statements are incomplete.", color="C00000")

    def s5_computation(self):
        an, c = self.an, self.an.comp
        self.h("5. Preliminary Corporate Tax computation", 1)
        self.p("Only CONFIRMED adjustments are included. Proposed adjustments are listed in sections 6, 7 and 13 and are not in these figures.", italic=True, size=9)
        rows = [["Accounting profit/(loss)", _n(c.A)], ["Less: deductions / exempt income (confirmed)", _n(-c.B)], ["Add: non-deductible expenses (confirmed)", _n(c.C)]]
        if c.o_gain or c.o_loss:
            rows.append(["Add/(deduct): previously unrealised amounts now realised", _n(c.o_gain - c.o_loss)])
        rows += [["Taxable income/(loss)", _n(c.D)], ["Less: tax losses utilised", _n(-(c.relief_cf + c.relief_group))], ["Net taxable income/(loss)", _n(c.E)],
                 ["Corporate Tax at 0% / 9% (basis: " + c.basis + ")", _n(c.tax_9 + c.tax_0)], ["Less: foreign tax credit", _n(-c.credit)], ["Corporate Tax payable (preliminary)", _n(c.payable)]]
        self.table(["Item", "AED"], rows, [12.5, 4.9], num_cols=(1,))
        for n in c.notes:
            self.note(n)
        pi = an.pending_impact()
        if pi["count"]:
            self.p(f"If all {pi['count']} proposed adjustment(s) were confirmed, taxable income would change by AED {_n(pi['delta_taxable'])} and tax by about AED {_n(pi['delta_tax'])}. "
                   "This is illustrative; no decision has been taken.", size=9)

    def _adj_rows(self, adjs):
        rows = []
        for a in adjs:
            acc = "; ".join(f"{x[0]} {x[1]}".strip() for x in a.accounts)
            rows.append([acc, a.description, _n(a.amount), a.report_status, (a.blocking or a.reason), "; ".join(a.docs_needed), a.mgmt_response or "-"])
        return rows

    def s6_disallowed(self):
        an = self.an
        self.h("6. Potentially disallowed expenses", 1)
        adjs = [a for a in an.adjs if ROWS[a.row][1] == "C" and a.status != "Rejected"]
        if adjs:
            self.table(["Account", "Item", "AED", "Status", "Observation", "Documents required", "Mgmt response"], self._adj_rows(adjs),
                       [3.1, 3.4, 1.9, 2.1, 3.3, 2.4, 1.6], status_col=3, num_cols=(2,), font=7)
        else:
            self.p("No potentially disallowed expenses were identified from the account names and ledger.")
        covered = {x[0] for a in adjs for x in a.accounts}
        other = [f for f in an.review.flags if f.category in ("MISC", "PROF", "REPAIR", "COMM", "GIFT") and f.code not in covered]
        if other:
            self.h("Accounts requiring review (no amount proposed)", 3)
            self.table(["Account", "Balance (AED)", "Observation", "Status"], [[f"{f.code} {f.name}", _n(f.amount_aed), f.tax_implication, f.status] for f in other],
                       [4.5, 2.4, 8.0, 2.5], status_col=3, num_cols=(1,), font=7.5)
        self.note("Wording: an item is described as 'disallowed' only where its status is Confirmed. All other items are potential adjustments pending evidence.")

    def s7_exempt(self):
        an = self.an
        self.h("7. Potentially exempt income", 1)
        adjs = [a for a in an.adjs if a.category == "Exempt income" and a.status != "Rejected"]
        if adjs:
            self.table(["Account", "Item", "AED", "Status", "Observation", "Documents required", "Mgmt response"], self._adj_rows(adjs),
                       [3.1, 3.4, 1.9, 2.1, 3.3, 2.4, 1.6], status_col=3, num_cols=(2,), font=7)
        oi = [o for o in an.open_items if o["topic"] == "Exempt income"]
        if oi:
            self.table(["Open point", "Status"], [[o["item"], o["status"]] for o in oi], [13.5, 3.9], status_col=1, font=8)
        if not adjs and not oi:
            self.p("No dividend, share-disposal or other potentially exempt income accounts were identified.")
        self.note("Exemption conditions (ownership/cost, holding period, subject-to-tax) are applied only to the facts entered in the questionnaire; the tool does not verify them.")

    def s8_reliefs(self):
        an, c = self.an, self.an.comp
        self.h("8. QFZP and other relief considerations", 1)
        A = an.answers
        rows = [["Realisation basis election", ans(A, "realisation_basis", "Not confirmed"), "Applied only as stated; never elected automatically."],
                ["Small Business Relief", "Potentially available" if c.sbr.get("potential") else "Not indicated on the data provided",
                 " ".join(c.sbr.get("failures", []) + c.sbr.get("open", [])) or "Election not made."],
                ["QFZP", ans(A, "qfzp_claimed", "Unknown"), "; ".join(f"{k}: {v}" for k, v in (c.qfzp or {}).items() if isinstance(v, (str, int, float, bool)))[:400] or "Not claimed."],
                ["Interest limitation", "Within limit" if c.interest_test.get("within", True) else "Excess identified",
                 f"Net interest AED {_n(c.interest_test.get('net_interest', 0))} vs. threshold AED {_n(c.interest_test.get('threshold', 0))} (tax EBITDA is approximated)."],
                ["Tax losses", f"Brought forward AED {_n(c.loss_opening)}", "Opening losses are per management and not verified against prior returns."],
                ["Foreign tax credit", f"AED {_n(c.credit)} claimed", "Requires proof of payment and confirmation that the income is not exempt in the UAE."]]
        self.table(["Area", "Position", "Comment"], rows, [4.0, 4.4, 9.0], font=8)
        if an.scenarios:
            self.h("Realisation-basis scenarios (what-if only)", 3)
            self.table(["Option", "Taxable income (AED)", "Tax payable (AED)"], [[s["option"], _n(s["taxable_income"]), _n(s["tax_payable"])] for s in an.scenarios],
                       [7.5, 5, 4.9], num_cols=(1, 2))
            self.note("The election is made once, in the first tax period, and is irrevocable without FTA approval. This comparison is not a recommendation.")

    def s9_risk_accounts(self):
        an = self.an
        self.h("9. High-risk accounts and transactions", 1)
        flags = [f for f in an.review.flags if f.risk in ("High", "Medium")]
        if flags:
            self.table(["Code", "Account", "AED", "Risk", "Observation / possible implication", "Documents required", "Mgmt response"],
                       [[f.code, f.name, _n(f.amount_aed), f.risk, f"{f.category_label}: {f.tax_implication}", "; ".join(f.documents), f.mgmt_response or "-"] for f in flags],
                       [1.2, 3.2, 1.9, 1.4, 5.0, 3.0, 1.7], status_col=3, num_cols=(2,), font=7)
        low = [f for f in an.review.flags if f.risk == "Low"]
        if low:
            self.p(f"{len(low)} low-risk indicator(s) are listed in the review workbook.", size=9)
        if not an.review.flags:
            self.p("No account-level risk indicators were triggered.")

    def s10_gl(self):
        an = self.an
        self.h("10. Unusual or material GL entries", 1)
        if not an.project.gl:
            self.p("No general ledger was supplied; transaction-level tests (round sums, vague narrations, year-end entries, related-party counterparties) were not performed.",
                   color="C00000")
            return
        tf = an.review.txn_flags
        top = [t for t in tf if t.risk != "Low"][:25]
        self.p(f"{len(tf)} ledger line(s) triggered at least one test; the highest-priority {len(top)} are shown. Tests indicate where to ask questions - they are not findings.", size=9)
        if top:
            self.table(["Date", "Ref", "Account", "AED", "Narration / counterparty", "Tests triggered", "Risk"],
                       [[t.date or "", t.reference or t.line_id, f"{t.code} {t.name}", _n(t.amount_aed), f"{t.description} / {t.counterparty}".strip(" /"), "; ".join(t.tests), t.risk] for t in top],
                       [1.8, 1.7, 3.2, 2.0, 3.6, 3.7, 1.4], status_col=6, num_cols=(3,), font=7)
        for n in an.review.notes:
            self.note(n)

    def s11_missing(self):
        an = self.an
        self.h("11. Missing documents and information", 1)
        by = {}
        for f in an.review.flags:
            for dname in f.documents:
                by.setdefault(dname, set()).add(f"{f.code} {f.name}".strip())
        for a in an.adjs:
            for dname in a.docs_needed:
                for x in a.accounts:
                    by.setdefault(dname, set()).add(f"{x[0]} {x[1]}".strip())
        for o in an.open_items:
            for dname in o.get("documents", []):
                by.setdefault(dname, set()).add(o["topic"])
        rows = [[dname, "; ".join(sorted(v))[:300]] for dname, v in sorted(by.items())]
        if rows:
            self.table(["Document / information", "Needed for"], rows, [6.0, 11.4], font=7.5)
        else:
            self.p("No outstanding document requests were generated.")

    def s12_questions(self):
        an = self.an
        self.h("12. Questions requiring management confirmation", 1)
        qs = unanswered(an.answers, an.ctx)
        rows = [[i + 1, q.text, q.why or "-", q.evidence or "-"] for i, q in enumerate(qs)]
        if rows:
            self.table(["#", "Question", "Why it matters", "Evidence expected"], rows, [0.9, 7.3, 5.2, 4.0], font=7.5)
        else:
            self.p("All required questions have an answer; answers marked as defaults or prefilled still need management's confirmation.")
        confs = [(k, ans(an.answers, k)) for k in ("conf_documentation", "conf_business_purpose", "conf_disallowed", "conf_exempt_income", "conf_arms_length", "conf_subsidiaries")]
        self.note("Management confirmations recorded: " + "; ".join(f"{k.replace('conf_', '')}={v}" for k, v in confs))

    def s13_impact(self):
        an = self.an
        self.h("13. Potential tax adjustments and estimated impact", 1)
        rate = load_json("rules.json")["params"]["standard_rate"]
        prop = [a for a in an.adjs if a.status == "Proposed"]
        rows = []
        for a in prop:
            sign = 1 if ROWS[a.row][1] == "C" else -1
            eff = sign * a.amount * rate
            rows.append([a.description, _n(a.amount), "Increase" if sign > 0 else "Decrease", _n(eff), a.report_status])
        if rows:
            self.table(["Proposed adjustment", "Amount (AED)", "Effect on taxable income", f"Indicative tax at {rate:.0%} (AED)", "Status"], rows,
                       [6.6, 2.4, 2.5, 3.0, 2.9], status_col=4, num_cols=(1, 3), font=7.5)
            self.note("Indicative tax = adjustment x standard rate. It ignores the nil-rate band, losses, credits and interactions between items, and is not a computed liability.")
            pi = an.pending_impact()
            self.p(f"Combined effect if all were confirmed: taxable income AED {_n(pi['delta_taxable'])}; tax AED {_n(pi['delta_tax'])} (full recomputation).", size=9)
        else:
            self.p("There are no proposed adjustments awaiting a decision.")
        conf = [a for a in an.adjs if a.status == "Confirmed"]
        if conf:
            self.h("Confirmed adjustments included in the computation", 3)
            self.table(["Adjustment", "Amount (AED)", "Evidence / support", "Reviewer"], [[a.description, _n(a.amount), a.support_ref or "-", a.reviewer or "-"] for a in conf],
                       [8.0, 2.8, 4.2, 2.4], num_cols=(1,), font=7.5)

    def s14_technical(self):
        an = self.an
        self.h("14. Unresolved technical issues and assumptions", 1)
        tech = [o for o in an.open_items if o["status"] in ("Technical review required", "Information required")]
        if tech:
            self.table(["Topic", "Issue", "Status"], [[o["topic"], o["item"], o["status"]] for o in tech], [3.0, 11.0, 3.4], status_col=2, font=7.5)
        self.h("Rules relied on - verification status", 3)
        reg = load_json("rules.json")
        rows = [[r["id"], r["topic"], r["source"], r["verification"], r.get("date_checked", "")] for r in reg["rules"]]
        self.table(["ID", "Rule", "Source", "Verification", "Checked"], rows, [1.6, 5.4, 6.0, 2.5, 1.9], font=7)
        self.note("PRIMARY_READ = official text read; SECONDARY = reported by a secondary source only; UNVERIFIED = recalled, not checked. "
                  "Anything not PRIMARY_READ must be verified against the official source before the figures are relied on.")
        self.h("Assumptions and limitations", 3)
        self.bullets([
            "Mapping of accounts to statement lines is rule-based; accounts marked for review may be misclassified.",
            "Opening tax losses, prior-period unrealised amounts and elections are as stated by management and were not verified.",
            "Entertainment, donation, fine and similar adjustments are derived from account names and narrations; actual composition may differ.",
            "Tax EBITDA for the interest limitation is approximated from the accounts.",
            "The tool does not file returns, apply elections or reliefs automatically, or contact the FTA.",
        ])

    def s15_next(self):
        an = self.an
        self.h("15. Recommended next steps", 1)
        steps = []
        if an.st.unmapped_accounts:
            steps.append("Map the unmapped trial balance accounts and re-run the analysis.")
        if any(c.status == "FAIL" for c in an.checks):
            steps.append("Resolve the failed reconciliation checks listed in section 3.")
        if not an.project.gl:
            steps.append("Provide the general ledger so that transaction-level tests can be performed.")
        steps.append("Obtain the documents listed in section 11, starting with the high-risk accounts in section 9.")
        steps.append("Obtain management's answers to the questions in section 12 and record their responses against each flagged item.")
        steps.append("Decide each proposed adjustment (confirm / reject) with the supporting reference; items awaiting evidence stay outside the computation.")
        steps.append("Have a qualified adviser review the technical issues in section 14 and verify the rules marked as secondary or unverified.")
        steps.append("Update the computation, re-issue the workbook, and complete the final professional review before any filing.")
        for i, s in enumerate(steps, 1):
            para = self.doc.add_paragraph(style="List Number")
            para.add_run(s)
            para.paragraph_format.space_after = Pt(2)

    def appendices(self):
        an = self.an
        self.h("Appendix A - Reconciliation checks", 1)
        self.table(["ID", "Check", "Result", "Detail"], [[c.id, c.description, c.status, c.detail] for c in an.checks], [1.8, 4.6, 1.4, 9.6], font=7)
        rev = [(m, a) for m, a in zip(an.mapped, an.tr.accounts) if m.needs_review]
        if rev:
            lines = fs_lines()
            self.h("Appendix B - Account mappings needing review", 1)
            self.table(["Code", "Account", "Mapped to", "Why review"], [[m.account.code, m.account.name, lines[m.line]["label"], "; ".join(m.flags) or m.method] for m, _ in rev],
                       [1.5, 5.0, 4.9, 6.0], font=7)


def build_report(an, path: str, letterhead: str = "") -> str:
    return ReportBuilder(an, letterhead).build(path)
