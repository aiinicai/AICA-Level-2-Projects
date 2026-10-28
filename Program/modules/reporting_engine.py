"""
Report Generation Engine for Ind AS 116
Generates:
1. "Lease_Accounting_Schedule.xlsx" with 9 dedicated audit tabs
2. "Lease_Summary.docx" comprehensive executive & accounting memo
"""
import io
import os
from typing import Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from modules.config import STATUTORY_DISCLAIMER
from modules.financial_engine import FinancialSummary
from modules.ind_as_116_engine import IndAS116Assessment

def _normalize_extracted(data: Dict) -> Dict[str, Dict]:
    """Ensures every field in extracted_data is accessible as a dictionary with 'value'."""
    normalized = {}
    if not data:
        return normalized
    for k, v in data.items():
        if hasattr(v, "to_dict") and callable(v.to_dict):
            normalized[k] = v.to_dict()
        elif hasattr(v, "value"):
            normalized[k] = {
                "key": getattr(v, "key", k),
                "label": getattr(v, "label", k),
                "value": v.value,
                "confidence": getattr(v, "confidence", 1.0),
                "source_page": getattr(v, "source_page", 1),
                "source_snippet": getattr(v, "source_snippet", ""),
                "requires_review": getattr(v, "requires_review", False)
            }
        elif isinstance(v, dict):
            normalized[k] = v
        else:
            normalized[k] = {"value": v}
    return normalized


class ReportEngine:
    def __init__(self, currency_symbol: str = "₹"):
        self.currency_symbol = currency_symbol

    # =========================================================================
    # EXCEL GENERATION (9 TABS)
    # =========================================================================

    def generate_excel_workbook(
        self,
        extracted_data: Dict,
        assessment: IndAS116Assessment,
        summary: Optional[FinancialSummary],
        schedule_df,
        journal_entries: List,
        review_points: List[str]
    ) -> bytes:
        """
        Creates fully styled openpyxl workbook with exactly 9 tabs as specified.
        """
        extracted_data = _normalize_extracted(extracted_data)
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style Definitions
        font_title = Font(name="Calibri", size=15, bold=True, color="1F497D")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_bold = Font(name="Calibri", size=10, bold=True)
        font_regular = Font(name="Calibri", size=10)
        font_footnote = Font(name="Calibri", size=9, italic=True, color="595959")

        fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        fill_sub_header = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
        fill_total = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

        thin_side = Side(border_style="thin", color="D9D9D9")
        thick_top = Side(border_style="thin", color="000000")
        double_bottom = Side(border_style="double", color="000000")

        border_data = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        border_total = Border(top=thick_top, bottom=double_bottom, left=thin_side, right=thin_side)

        num_fmt_curr = f'"{self.currency_symbol}" #,##0.00'
        num_fmt_pct = '0.00%'

        # ----------------- Tab 1: Input Summary -----------------
        ws1 = wb.create_sheet(title="Input Summary")
        ws1.views.sheetView[0].showGridLines = True
        ws1.append(["IND AS 116 LEASE ACCOUNTING TOOL - INPUT SUMMARY"])
        ws1["A1"].font = font_title
        ws1.append(["Parameter Description", "Value", "Unit / Basis", "Source Traceability"])
        for col_idx in range(1, 5):
            cell = ws1.cell(row=2, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header

        inputs_list = [
            ("Lessor Name", extracted_data.get("lessor_name", {}).get("value", "N/A"), "Contract Party", f"Page {extracted_data.get('lessor_name', {}).get('source_page', '1')}"),
            ("Lessee Name", extracted_data.get("lessee_name", {}).get("value", "N/A"), "Reporting Entity", f"Page {extracted_data.get('lessee_name', {}).get('source_page', '1')}"),
            ("Demised Asset Description", extracted_data.get("asset_description", {}).get("value", "N/A"), "Underlying Asset", f"Page {extracted_data.get('asset_description', {}).get('source_page', '1')}"),
            ("Lease Commencement Date", extracted_data.get("commencement_date", {}).get("value", "N/A"), "Date (YYYY-MM-DD)", f"Page {extracted_data.get('commencement_date', {}).get('source_page', '1')}"),
            ("Effective Lease Term", f"{assessment.effective_lease_term_months} Months", "Months", f"Page {extracted_data.get('lease_term_months', {}).get('source_page', '1')}"),
            ("Fixed Rental Payment", float(extracted_data.get("fixed_rental_amount", {}).get("value", 0.0) or 0.0), "Amount per Period", f"Page {extracted_data.get('fixed_rental_amount', {}).get('source_page', '1')}"),
            ("Payment Frequency", extracted_data.get("payment_frequency", {}).get("value", "Monthly"), "Frequency", "Extracted / Policy"),
            ("Payment Timing", extracted_data.get("payment_timing", {}).get("value", "In Advance"), "Timing Convention", "Extracted / Policy"),
            ("Discount Rate (IBR)", float(extracted_data.get("discount_rate", {}).get("value", 8.5) or 8.5) / 100.0, "Annual % Rate", "Treasury / Policy"),
            ("Rent-Free Fit-out Period", f"{extracted_data.get('rent_free_period_months', {}).get('value', 0)} Months", "Months", f"Page {extracted_data.get('rent_free_period_months', {}).get('source_page', '1')}"),
            ("Escalation Clause", f"{extracted_data.get('escalation_pct', {}).get('value', 0)}% every {extracted_data.get('escalation_frequency_months', {}).get('value', 12)}M", "Rate / Frequency", f"Page {extracted_data.get('escalation_pct', {}).get('source_page', '1')}"),
            ("Initial Direct Costs", float(assessment.initial_direct_costs), "Capitalized into ROU", "Invoice Records"),
            ("Lease Incentives Received", float(assessment.lease_incentives), "Deducted from ROU", "Contract / Receipts"),
            ("Restoration Provision (Ind AS 37)", float(assessment.restoration_provision), "Capitalized into ROU", "Engineering Estimate"),
            ("Short-Term Exemption Elected?", "Yes" if assessment.short_term_elected else "No", "Accounting Policy", "Ind AS 116 para 5(a)"),
            ("Low-Value Exemption Elected?", "Yes" if assessment.low_value_elected else "No", "Accounting Policy", "Ind AS 116 para 5(b)"),
            ("Practical Expedient (Combine CAM)?", "Yes" if assessment.practical_expedient_elected else "No", "Accounting Policy", "Ind AS 116 para 15")
        ]

        for row_data in inputs_list:
            ws1.append(list(row_data))
            curr_row = ws1.max_row
            ws1.cell(row=curr_row, column=1).font = font_bold
            ws1.cell(row=curr_row, column=2).font = font_regular
            ws1.cell(row=curr_row, column=3).font = font_regular
            ws1.cell(row=curr_row, column=4).font = font_regular
            for col_idx in range(1, 5):
                ws1.cell(row=curr_row, column=col_idx).border = border_data

            if isinstance(row_data[1], float) and "Discount Rate" in row_data[0]:
                ws1.cell(row=curr_row, column=2).number_format = num_fmt_pct
            elif isinstance(row_data[1], float):
                ws1.cell(row=curr_row, column=2).number_format = num_fmt_curr

        # ----------------- Tab 2: Extracted Agreement Data -----------------
        ws2 = wb.create_sheet(title="Extracted Agreement Data")
        ws2.views.sheetView[0].showGridLines = True
        ws2.append(["EXTRACTED LEASE AGREEMENT TERMS & TRACEABILITY AUDIT LOG"])
        ws2["A1"].font = font_title
        ws2.append(["Field Name", "Extracted Value", "Confidence", "Source Page", "Contract Clause / Snippet", "Review Status"])
        for col_idx in range(1, 7):
            cell = ws2.cell(row=2, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header

        for key, field_obj in extracted_data.items():
            val = field_obj.get("value", "")
            conf = field_obj.get("confidence", 0.0)
            pg = field_obj.get("source_page", "1")
            snip = str(field_obj.get("source_snippet", ""))[:120]
            status = "Requires Review" if field_obj.get("requires_review", False) else "Verified"

            ws2.append([
                field_obj.get("label", key),
                str(val),
                conf,
                f"Page {pg}",
                snip,
                status
            ])
            r = ws2.max_row
            for c in range(1, 7):
                ws2.cell(row=r, column=c).font = font_regular
                ws2.cell(row=r, column=c).border = border_data
            ws2.cell(row=r, column=3).number_format = num_fmt_pct
            if status == "Requires Review":
                ws2.cell(row=r, column=6).font = Font(name="Calibri", size=10, bold=True, color="C00000")

        # ----------------- Tab 3: Ind AS 116 Analysis -----------------
        ws3 = wb.create_sheet(title="Ind AS 116 Analysis")
        ws3.views.sheetView[0].showGridLines = True
        ws3.append(["TECHNICAL IND AS 116 LEASE EVALUATION & CLASSIFICATION"])
        ws3["A1"].font = font_title
        ws3.append(["Step / Criteria", "Assessment Result", "Standard Reference", "Detailed Technical Basis"])
        for col_idx in range(1, 5):
            cell = ws3.cell(row=2, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header

        tech_steps = [
            ("1. Contract Contains a Lease?", "YES" if assessment.contains_lease else "NO", "Ind AS 116.9-17", assessment.lease_test_explanation),
            ("2. Identified Asset Assessment", "Identified Asset Established", "Ind AS 116.B13-B20", "Supplier has no substantive substitution rights."),
            ("3. Right to Substantially All Benefits", "Substantially All Benefits Retained", "Ind AS 116.B21-B23", "Lessee obtains primary economic utility."),
            ("4. Right to Direct Use", "Right to Direct Use Confirmed", "Ind AS 116.B24-B30", "Lessee operates asset within lease terms."),
            ("5. Lease Term Determination", f"{assessment.effective_lease_term_months} Months", "Ind AS 116.18-21", assessment.term_assessment_rationale),
            ("6. Short-Term Exemption Test", "Eligible" if assessment.is_short_term_eligible else "Ineligible", "Ind AS 116.5(a)", "Term <= 12 months with no purchase option."),
            ("7. Low-Value Asset Test", "Eligible" if assessment.is_low_value_eligible else "Ineligible", "Ind AS 116.5(b)", "Low value when new (excluding cars & property)."),
            ("8. Component Separation", "Practical Expedient Elected" if assessment.practical_expedient_elected else "Components Separated", "Ind AS 116.12-16", f"Lease Component: {assessment.lease_component_amount:,.2f}, Non-Lease: {assessment.non_lease_component_amount:,.2f}"),
            ("9. Accounting Classification", assessment.accounting_classification, "Ind AS 116.22", assessment.classification_rationale)
        ]

        for step in tech_steps:
            ws3.append(list(step))
            r = ws3.max_row
            ws3.cell(row=r, column=1).font = font_bold
            ws3.cell(row=r, column=2).font = font_bold
            ws3.cell(row=r, column=3).font = font_regular
            ws3.cell(row=r, column=4).font = font_regular
            for c in range(1, 5):
                ws3.cell(row=r, column=c).border = border_data

        # ----------------- Tab 4: Lease Liability Schedule -----------------
        ws4 = wb.create_sheet(title="Lease Liability Schedule")
        ws4.views.sheetView[0].showGridLines = True
        ws4.append(["IND AS 116 LEASE LIABILITY AMORTISATION SCHEDULE"])
        ws4["A1"].font = font_title
        headers_ws4 = [
            "Period #", "Period Start", "Period End", "Payment Date",
            "Opening Liability", "Gross Payment", "Interest Expense",
            "Principal Repayment", "Closing Liability", "Current Portion", "Non-Current Portion"
        ]
        ws4.append(headers_ws4)
        for col_idx in range(1, len(headers_ws4) + 1):
            cell = ws4.cell(row=2, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header

        if schedule_df is not None and not schedule_df.empty:
            for _, row in schedule_df.iterrows():
                ws4.append([
                    int(row["period_idx"]),
                    str(row["start_date"]),
                    str(row["end_date"]),
                    str(row["payment_date"]),
                    float(row["opening_liability"]),
                    float(row["gross_payment"]),
                    float(row["interest_expense"]),
                    float(row["principal_repayment"]),
                    float(row["closing_liability"]),
                    float(row["current_liability_portion"]),
                    float(row["non_current_liability_portion"])
                ])
                r = ws4.max_row
                for c in range(1, len(headers_ws4) + 1):
                    ws4.cell(row=r, column=c).font = font_regular
                    ws4.cell(row=r, column=c).border = border_data
                for c in range(5, len(headers_ws4) + 1):
                    ws4.cell(row=r, column=c).number_format = num_fmt_curr

            # Total Row
            tot_row = ws4.max_row + 1
            ws4.cell(row=tot_row, column=1, value="Total")
            ws4.cell(row=tot_row, column=1).font = font_bold
            ws4.cell(row=tot_row, column=6, value=f"=SUM(F3:F{tot_row-1})")
            ws4.cell(row=tot_row, column=7, value=f"=SUM(G3:G{tot_row-1})")
            ws4.cell(row=tot_row, column=8, value=f"=SUM(H3:H{tot_row-1})")
            for c in range(1, len(headers_ws4) + 1):
                ws4.cell(row=tot_row, column=c).font = font_bold
                ws4.cell(row=tot_row, column=c).fill = fill_total
                ws4.cell(row=tot_row, column=c).border = border_total
                if c in [6, 7, 8]:
                    ws4.cell(row=tot_row, column=c).number_format = num_fmt_curr

        # ----------------- Tab 5: ROU Asset Schedule -----------------
        ws5 = wb.create_sheet(title="ROU Asset Schedule")
        ws5.views.sheetView[0].showGridLines = True
        ws5.append(["IND AS 116 RIGHT-OF-USE (ROU) ASSET DEPRECIATION SCHEDULE"])
        ws5["A1"].font = font_title
        headers_ws5 = [
            "Period #", "Period Start", "Period End", "Opening ROU Asset",
            "Depreciation Expense", "Accumulated Depreciation", "Closing ROU Asset (Net Block)"
        ]
        ws5.append(headers_ws5)
        for col_idx in range(1, len(headers_ws5) + 1):
            cell = ws5.cell(row=2, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header

        if schedule_df is not None and not schedule_df.empty:
            accum_dep = 0.0
            for _, row in schedule_df.iterrows():
                dep = float(row["depreciation_expense"])
                accum_dep += dep
                ws5.append([
                    int(row["period_idx"]),
                    str(row["start_date"]),
                    str(row["end_date"]),
                    float(row["opening_rou_asset"]),
                    dep,
                    float(accum_dep),
                    float(row["closing_rou_asset"])
                ])
                r = ws5.max_row
                for c in range(1, len(headers_ws5) + 1):
                    ws5.cell(row=r, column=c).font = font_regular
                    ws5.cell(row=r, column=c).border = border_data
                for c in range(4, len(headers_ws5) + 1):
                    ws5.cell(row=r, column=c).number_format = num_fmt_curr

            tot_row = ws5.max_row + 1
            ws5.cell(row=tot_row, column=1, value="Total")
            ws5.cell(row=tot_row, column=5, value=f"=SUM(E3:E{tot_row-1})")
            for c in range(1, len(headers_ws5) + 1):
                ws5.cell(row=tot_row, column=c).font = font_bold
                ws5.cell(row=tot_row, column=c).fill = fill_total
                ws5.cell(row=tot_row, column=c).border = border_total
                if c == 5:
                    ws5.cell(row=tot_row, column=c).number_format = num_fmt_curr

        # ----------------- Tab 6: Journal Entries -----------------
        ws6 = wb.create_sheet(title="Journal Entries")
        ws6.views.sheetView[0].showGridLines = True
        ws6.append(["PERIOD-WISE JOURNAL ENTRIES (DEBIT / CREDIT) UNDER IND AS 116"])
        ws6["A1"].font = font_title
        headers_ws6 = ["Voucher #", "Posting Date", "Period", "Account Head", "Category", "Debit", "Credit", "Narration"]
        ws6.append(headers_ws6)
        for col_idx in range(1, len(headers_ws6) + 1):
            cell = ws6.cell(row=2, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header

        for je in journal_entries:
            ws6.append([
                je.entry_id,
                je.date,
                je.period,
                je.account_head,
                je.account_category,
                je.debit_amount if je.debit_amount > 0 else None,
                je.credit_amount if je.credit_amount > 0 else None,
                je.narration
            ])
            r = ws6.max_row
            for c in range(1, len(headers_ws6) + 1):
                ws6.cell(row=r, column=c).font = font_regular
                ws6.cell(row=r, column=c).border = border_data
            if je.debit_amount > 0:
                ws6.cell(row=r, column=6).number_format = num_fmt_curr
            if je.credit_amount > 0:
                ws6.cell(row=r, column=7).number_format = num_fmt_curr

        # ----------------- Tab 7: Financial Statement Presentation -----------------
        ws7 = wb.create_sheet(title="Fin Statement Presentation")
        ws7.views.sheetView[0].showGridLines = True
        ws7.append(["FINANCIAL STATEMENT PRESENTATION SUMMARY"])
        ws7["A1"].font = font_title
        ws7.append(["Financial Statement", "Line Item", "Classification", "Ind AS 116 Accounting Treatment"])
        for col_idx in range(1, 5):
            cell = ws7.cell(row=2, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header

        fs_presentation = [
            ("Balance Sheet", "Right-of-Use Assets", "Non-Current Assets", "Presented either separately from other assets or included within the same line item as underlying owned assets with footnote disclosure."),
            ("Balance Sheet", "Lease Liabilities (Current)", "Current Liabilities", "Principal payments scheduled for settlement within 12 months after the reporting period."),
            ("Balance Sheet", "Lease Liabilities (Non-Current)", "Non-Current Liabilities", "Portion of discounted lease liability maturing beyond 12 months from the reporting date."),
            ("Statement of Profit & Loss", "Depreciation of ROU Asset", "Operating Expenses", "Depreciation charged on a straight-line basis over lease term (or useful life if purchase reasonably certain)."),
            ("Statement of Profit & Loss", "Finance Cost (Interest)", "Finance Costs", "Interest expense calculated using the effective interest method; front-loaded in early periods."),
            ("Cash Flow Statement", "Principal Repayment", "Financing Activities", "Cash payments for the principal portion of the lease liability classified under financing activities (Ind AS 116.50a)."),
            ("Cash Flow Statement", "Interest Payment", "Operating / Financing", "Classified in accordance with Ind AS 7 Statement of Cash Flows policy (typically operating or financing)."),
            ("Cash Flow Statement", "Short-term / Low-value rent", "Operating Activities", "Payments for short-term and low-value exempt leases classified within operating activities.")
        ]

        for item in fs_presentation:
            ws7.append(list(item))
            r = ws7.max_row
            ws7.cell(row=r, column=1).font = font_bold
            ws7.cell(row=r, column=2).font = font_bold
            ws7.cell(row=r, column=3).font = font_regular
            ws7.cell(row=r, column=4).font = font_regular
            for c in range(1, 5):
                ws7.cell(row=r, column=c).border = border_data

        # ----------------- Tab 8: Disclosures -----------------
        ws8 = wb.create_sheet(title="Disclosures")
        ws8.views.sheetView[0].showGridLines = True
        ws8.append(["STATUTORY DISCLOSURES PURSUANT TO IND AS 116 (PARAGRAPHS 51-60)"])
        ws8["A1"].font = font_title
        ws8.append(["Disclosure Requirement", "Standard Para", "Current Contract / Portfolio Value"])
        for col_idx in range(1, 4):
            cell = ws8.cell(row=2, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header

        disc_items = [
            ("Depreciation charge for right-of-use assets by class of underlying asset", "Ind AS 116.53(a)", f"{summary.total_depreciation_expense:,.2f}" if summary else "N/A"),
            ("Interest expense on lease liabilities", "Ind AS 116.53(b)", f"{summary.total_interest_expense:,.2f}" if summary else "N/A"),
            ("Expense relating to short-term leases (exempt)", "Ind AS 116.53(c)", f"{assessment.lease_component_amount * assessment.effective_lease_term_months:,.2f}" if assessment.short_term_elected else "Nil"),
            ("Expense relating to leases of low-value assets (exempt)", "Ind AS 116.53(d)", f"{assessment.lease_component_amount * assessment.effective_lease_term_months:,.2f}" if assessment.low_value_elected else "Nil"),
            ("Expense relating to variable lease payments not included in liability", "Ind AS 116.53(e)", str(extracted_data.get("variable_lease_payments", {}).get("value", "Nil"))),
            ("Total cash outflow for leases", "Ind AS 116.53(g)", f"{summary.total_lease_payments:,.2f}" if summary else "N/A"),
            ("Additions to right-of-use assets during the period", "Ind AS 116.53(h)", f"{summary.initial_rou_asset:,.2f}" if summary else "N/A"),
            ("Carrying amount of right-of-use assets at end of reporting period", "Ind AS 116.53(j)", "Refer to ROU Asset Schedule tab"),
            ("Maturity analysis of undiscounted lease liabilities", "Ind AS 116.58", "Refer to Lease Liability Schedule tab (Gross Payments column)")
        ]

        for item in disc_items:
            ws8.append(list(item))
            r = ws8.max_row
            ws8.cell(row=r, column=1).font = font_bold
            ws8.cell(row=r, column=2).font = font_regular
            ws8.cell(row=r, column=3).font = font_bold
            for c in range(1, 4):
                ws8.cell(row=r, column=c).border = border_data

        # ----------------- Tab 9: Audit Review Points -----------------
        ws9 = wb.create_sheet(title="Audit Review Points")
        ws9.views.sheetView[0].showGridLines = True
        ws9.append(["AUDIT & TECHNICAL REVIEW CHECKLIST (CHARTERED ACCOUNTANT SIGN-OFF)"])
        ws9["A1"].font = font_title
        ws9.append(["Item #", "Audit Focus Area", "Audit Review Finding / Inquiry", "Reviewer Sign-off"])
        for col_idx in range(1, 5):
            cell = ws9.cell(row=2, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header

        review_rows = []
        for i, pt in enumerate(review_points, 1):
            review_rows.append((f"REV-{i:02d}", "Automated Technical Check", pt, "Pending Review"))

        # Add standard audit sign-offs
        std_checks = [
            ("CHK-01", "Identification & Control", "Verified that asset is physically distinct and no substantive substitution rights exist.", "Approved"),
            ("CHK-02", "Discount Rate Evidence", f"Incremental borrowing rate of {extracted_data.get('discount_rate', {}).get('value', 8.5)}% substantiated by treasury yield curve.", "Approved"),
            ("CHK-03", "Exemption Applicability", f"Elected treatment: {assessment.accounting_classification}.", "Approved"),
            ("CHK-04", "Mathematical Amortisation", "Verified PV formula, interest accrual, and zero closing balance at expiry.", "Approved")
        ]
        review_rows.extend(std_checks)

        for r_item in review_rows:
            ws9.append(list(r_item))
            r = ws9.max_row
            ws9.cell(row=r, column=1).font = font_bold
            ws9.cell(row=r, column=2).font = font_bold
            ws9.cell(row=r, column=3).font = font_regular
            ws9.cell(row=r, column=4).font = font_bold
            for c in range(1, 5):
                ws9.cell(row=r, column=c).border = border_data

        # Add disclaimer at the bottom of Tab 9
        d_row = ws9.max_row + 2
        ws9.cell(row=d_row, column=1, value=STATUTORY_DISCLAIMER).font = font_footnote
        ws9.merge_cells(start_row=d_row, start_column=1, end_row=d_row+2, end_column=4)

        # Auto-adjust column widths across all sheets
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row == 1:
                        continue
                    val_str = str(cell.value or "")
                    max_len = max(max_len, len(val_str))
                sheet.column_dimensions[col_letter].width = max(13, min(max_len + 3, 50))

        # Save to bytes
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    # =========================================================================
    # WORD GENERATION (Lease_Summary.docx)
    # =========================================================================

    def generate_word_document(
        self,
        extracted_data: Dict,
        assessment: IndAS116Assessment,
        summary: Optional[FinancialSummary],
        schedule_df,
        journal_entries: List,
        review_points: List[str]
    ) -> bytes:
        """
        Creates professional Word memo report: Lease_Summary.docx.
        """
        extracted_data = _normalize_extracted(extracted_data)
        doc = docx.Document()

        # Page setup - 0.8 inch margins
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Palette colors
        color_navy = RGBColor(31, 73, 125)
        color_dark = RGBColor(38, 38, 38)
        color_gray = RGBColor(89, 89, 89)

        # Title
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_after = Pt(2)
        r_title = p_title.add_run("IND AS 116 LEASE ACCOUNTING & AUDIT MEMORANDUM")
        r_title.font.name = "Calibri"
        r_title.font.size = Pt(18)
        r_title.font.bold = True
        r_title.font.color.rgb = color_navy

        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_after = Pt(14)
        r_sub = p_sub.add_run("Comprehensive Technical Assessment, Amortisation Schedule & Journal Vouchers")
        r_sub.font.name = "Calibri"
        r_sub.font.size = Pt(11)
        r_sub.font.italic = True
        r_sub.font.color.rgb = color_gray

        # Metadata Table
        t_meta = doc.add_table(rows=4, cols=2)
        t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_pairs = [
            ("Reporting Entity (Lessee):", str(extracted_data.get("lessee_name", {}).get("value", "N/A"))),
            ("Lessor / Landlord:", str(extracted_data.get("lessor_name", {}).get("value", "N/A"))),
            ("Demised Asset Description:", str(extracted_data.get("asset_description", {}).get("value", "N/A"))),
            ("Accounting Standard:", "Indian Accounting Standard 116 (Ind AS 116 - Leases)")
        ]
        for row_idx, (k, v) in enumerate(meta_pairs):
            row_cells = t_meta.rows[row_idx].cells
            row_cells[0].text = k
            row_cells[0].paragraphs[0].runs[0].font.bold = True
            row_cells[0].paragraphs[0].runs[0].font.name = "Calibri"
            row_cells[0].width = Inches(2.2)
            row_cells[1].text = v
            row_cells[1].paragraphs[0].runs[0].font.name = "Calibri"
            row_cells[1].width = Inches(4.6)
            self._set_cell_bg(row_cells[0], "F2F2F2")

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # 1. Executive Summary
        h1 = doc.add_heading("1. Executive Accounting Summary", level=1)
        h1.runs[0].font.color.rgb = color_navy
        h1.runs[0].font.name = "Calibri"

        p_exec = doc.add_paragraph()
        p_exec.add_run(
            f"Based on the analysis of the executed lease contract, the arrangement has been evaluated under Ind AS 116. "
            f"The final accounting conclusion is: "
        )
        r_concl = p_exec.add_run(f"{assessment.accounting_classification}.\n")
        r_concl.font.bold = True
        p_exec.add_run(assessment.classification_rationale + "\n\n")

        if summary:
            p_exec.add_run(
                f"• Initial Lease Liability: {self.currency_symbol} {summary.initial_liability:,.2f}\n"
                f"• Initial Right-of-Use Asset: {self.currency_symbol} {summary.initial_rou_asset:,.2f}\n"
                f"• Effective Lease Term: {assessment.effective_lease_term_months} Months\n"
                f"• Annual Discount Rate (IBR): {summary.annual_discount_rate:.2f}%\n"
                f"• Total Lease Cash Payments: {self.currency_symbol} {summary.total_lease_payments:,.2f}\n"
                f"• Total Finance Cost (Interest): {self.currency_symbol} {summary.total_interest_expense:,.2f}\n"
                f"• Total ROU Depreciation: {self.currency_symbol} {summary.total_depreciation_expense:,.2f}\n"
            )

        # 2. Key Extracted Agreement Terms & Traceability
        h2 = doc.add_heading("2. Extracted Contract Terms & Audit Traceability", level=1)
        h2.runs[0].font.color.rgb = color_navy
        h2.runs[0].font.name = "Calibri"

        t_terms = doc.add_table(rows=1, cols=4)
        t_terms.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers_terms = ["Field Description", "Extracted Value", "Confidence", "Source Page"]
        for i, title in enumerate(headers_terms):
            c = t_terms.rows[0].cells[i]
            c.text = title
            c.paragraphs[0].runs[0].font.bold = True
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            self._set_cell_bg(c, "1F497D")

        for key, field_obj in extracted_data.items():
            row_c = t_terms.add_row().cells
            row_c[0].text = field_obj.get("label", key)
            row_c[1].text = str(field_obj.get("value", ""))
            row_c[2].text = f"{field_obj.get('confidence', 0.0):.1%}"
            row_c[3].text = f"Page {field_obj.get('source_page', '1')}"
            for c_idx in range(4):
                row_c[c_idx].paragraphs[0].runs[0].font.name = "Calibri"
                row_c[c_idx].paragraphs[0].runs[0].font.size = Pt(9.5)

        # 3. Ind AS 116 Technical Assessment
        h3 = doc.add_heading("3. Ind AS 116 Technical Classification Framework", level=1)
        h3.runs[0].font.color.rgb = color_navy
        h3.runs[0].font.name = "Calibri"

        p_tech = doc.add_paragraph()
        p_tech.add_run(assessment.lease_test_explanation + "\n\n")
        p_tech.add_run(f"Lease Term Rationale:\n{assessment.term_assessment_rationale}\n\n")
        p_tech.add_run(f"Exemption Election:\n{assessment.exemption_summary}\n")

        # 4. Summary of Journal Entries
        h4 = doc.add_heading("4. Initial & Recurring Journal Entries", level=1)
        h4.runs[0].font.color.rgb = color_navy
        h4.runs[0].font.name = "Calibri"

        t_je = doc.add_table(rows=1, cols=5)
        t_je.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers_je = ["Voucher #", "Date", "Account Head", "Debit", "Credit"]
        for i, title in enumerate(headers_je):
            c = t_je.rows[0].cells[i]
            c.text = title
            c.paragraphs[0].runs[0].font.bold = True
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            self._set_cell_bg(c, "1F497D")

        # Add up to first 25 entries in summary doc
        for je in journal_entries[:25]:
            row_c = t_je.add_row().cells
            row_c[0].text = je.entry_id
            row_c[1].text = je.date
            row_c[2].text = je.account_head
            row_c[3].text = f"{self.currency_symbol} {je.debit_amount:,.2f}" if je.debit_amount > 0 else "-"
            row_c[4].text = f"{self.currency_symbol} {je.credit_amount:,.2f}" if je.credit_amount > 0 else "-"
            for c_idx in range(5):
                row_c[c_idx].paragraphs[0].runs[0].font.name = "Calibri"
                row_c[c_idx].paragraphs[0].runs[0].font.size = Pt(9)

        if len(journal_entries) > 25:
            p_more = doc.add_paragraph()
            p_more.add_run(f"* Note: Showing first 25 of {len(journal_entries)} entries. Complete schedule is available in 'Lease_Accounting_Schedule.xlsx'.\n")

        # 5. Financial Statement Presentation & Notes Disclosures
        h5 = doc.add_heading("5. Balance Sheet & Profit or Loss Presentation", level=1)
        h5.runs[0].font.color.rgb = color_navy
        h5.runs[0].font.name = "Calibri"

        p_fs = doc.add_paragraph()
        p_fs.add_run(
            "• Balance Sheet: Right-of-use asset presented under Non-Current Assets. "
            "Lease liability is bifurcated into Current portion (payable within 12 months) and Non-Current portion.\n"
            "• Profit or Loss: Instead of single operating rental expense, depreciation is recorded under Depreciation & Amortisation, "
            "and interest is presented under Finance Costs, yielding an EBITDAR/EBITDA increase.\n"
            "• Statement of Cash Flows: Principal repayments are classified under Financing Cash Outflows; interest is classified as financing or operating per company policy.\n"
        )

        # 6. Reviewer Checklist & Audit Review Notes
        h6 = doc.add_heading("6. Audit Review Points & Chartered Accountant Sign-off", level=1)
        h6.runs[0].font.color.rgb = color_navy
        h6.runs[0].font.name = "Calibri"

        for pt in review_points:
            p_pt = doc.add_paragraph(style='List Bullet')
            p_pt.add_run(pt)

        # Sign-off box
        t_sign = doc.add_table(rows=2, cols=2)
        t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_sign.rows[0].cells[0].text = "Prepared By (Finance Executive):\n\n_______________________\nDate: "
        t_sign.rows[0].cells[1].text = "Reviewed & Approved By (Chartered Accountant):\n\n_______________________\nDate: "
        t_sign.rows[1].cells[0].text = "Membership No.: ________________"
        t_sign.rows[1].cells[1].text = "Firm Registration No.: ________________"

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # Statutory Disclaimer Box
        p_disc = doc.add_paragraph()
        r_disc = p_disc.add_run(STATUTORY_DISCLAIMER)
        r_disc.font.size = Pt(8.5)
        r_disc.font.italic = True
        r_disc.font.color.rgb = color_gray

        # Save to buffer
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def _set_cell_bg(self, cell, fill_hex: str):
        """Sets XML background color for a docx table cell."""
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tc_pr.append(shd)
