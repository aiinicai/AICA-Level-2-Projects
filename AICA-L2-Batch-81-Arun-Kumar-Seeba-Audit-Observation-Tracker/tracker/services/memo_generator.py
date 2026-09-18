import os
from pathlib import Path
from django.conf import settings
from django.utils import timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from tracker.models import Engagement, Query, ExceptionApproval, Memo, Staff
from tracker.services.audit_service import log_audit


def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def check_memo_gate(engagement):
    """
    FR-11 Gate: Check if all queries are in terminal state.
    Returns (is_allowed, open_queries, has_ex03_override).
    """
    open_queries = engagement.queries.exclude(status__in=['Closed', 'Closed - Exception Approved', 'Withdrawn'])
    
    if not open_queries.exists():
        return True, [], False

    # Check if there is an approved EX-03 exception
    ex03 = engagement.exceptions.filter(exception_type='EX-03', decision='Approved').first()
    if ex03:
        return True, list(open_queries), True

    return False, list(open_queries), False


def generate_docx_memo(engagement, generated_by_staff, is_draft=True, request=None):
    """
    Generates a high-quality, structured Audit Summary Memo in .docx format.
    """
    is_allowed, open_queries, has_ex03 = check_memo_gate(engagement)
    if not is_allowed:
        raise ValueError(f"Cannot generate memo: {len(open_queries)} queries remain open without an approved EX-03 exception.")

    # Determine version
    existing_memos_count = engagement.memos.count()
    version_no = f"v1.{existing_memos_count}" if is_draft else f"v1.{existing_memos_count} Final"

    doc = Document()

    # Set margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Header / Footer
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hdr_run = p_hdr.add_run(f"{engagement.engagement_code} | {engagement.financial_year} | {version_no}")
        hdr_run.font.size = Pt(8.5)
        hdr_run.font.color.rgb = RGBColor(128, 128, 128)

        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if is_draft:
            ftr_run = p_ftr.add_run("DRAFT — SUBJECT TO REVIEW AND SIGN-OFF. NOT FOR EXTERNAL ISSUE.")
            ftr_run.font.bold = True
            ftr_run.font.color.rgb = RGBColor(180, 40, 40)
            ftr_run.font.size = Pt(9)
        else:
            ftr_run = p_ftr.add_run("CONFIDENTIAL — AUDIT MEMORANDUM (FINAL SIGNED COPY)")
            ftr_run.font.bold = True
            ftr_run.font.color.rgb = RGBColor(40, 80, 140)
            ftr_run.font.size = Pt(9)

    # 1. FIRM LETTERHEAD
    p_firm = doc.add_paragraph()
    p_firm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_firm = p_firm.add_run("AICA & ASSOCIATES\n")
    run_firm.font.size = Pt(18)
    run_firm.font.bold = True
    run_firm.font.color.rgb = RGBColor(26, 54, 93)

    run_firm_sub = p_firm.add_run("CHARTERED ACCOUNTANTS\nICAI Firm Registration No: 104522W | Peer Review Certified Firm\n")
    run_firm_sub.font.size = Pt(9.5)
    run_firm_sub.font.color.rgb = RGBColor(74, 85, 104)

    # Divider Line
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div_run = p_div.add_run("―" * 45)
    p_div_run.font.color.rgb = RGBColor(200, 200, 200)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("AUDIT OBSERVATION & QUERY SUMMARY MEMO\n")
    run_title.font.size = Pt(15)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    if is_draft:
        p_watermark = doc.add_paragraph()
        p_watermark.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_wm = p_watermark.add_run("[ DRAFT MEMORANDUM — PENDING FORMAL PARTNER SIGN-OFF ]")
        run_wm.font.size = Pt(10)
        run_wm.font.bold = True
        run_wm.font.color.rgb = RGBColor(197, 48, 48)

    doc.add_paragraph() # Spacing

    # ENGAGEMENT METADATA TABLE
    table_meta = doc.add_table(rows=6, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.autofit = False

    lead_partner = engagement.get_lead_partner()
    reviewer_partner = engagement.get_reviewer_partner()
    manager_assignments = engagement.team_members.filter(role_on_engagement='Manager', is_active=True).first()
    manager_name = manager_assignments.staff.full_name if manager_assignments else "N/A"

    meta_data = [
        ("Client Name & Code:", f"{engagement.client.client_name} ({engagement.client.client_code})"),
        ("Financial Year & Audit Type:", f"{engagement.financial_year} | {engagement.audit_type}"),
        ("Audit Period Under Review:", f"{engagement.period_from.strftime('%d-%b-%Y')} to {engagement.period_to.strftime('%d-%b-%Y')}"),
        ("Engagement Code & Status:", f"{engagement.engagement_code} ({engagement.get_status_display()})"),
        ("Engagement Team Leadership:", f"Lead Partner: {lead_partner.full_name if lead_partner else 'Not Assigned'} | Manager: {manager_name}"),
        ("Memo Date & Version:", f"{timezone.now().strftime('%d-%b-%Y')} | Version {version_no}")
    ]

    for i, (label, val) in enumerate(meta_data):
        row = table_meta.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.3)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F1F5F9")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)

    doc.add_paragraph()

    # SECTION 1: BASIS OF PREPARATION AND SCOPE
    h1 = doc.add_heading("1. Scope & Basis of Preparation", level=1)
    h1.runs[0].font.color.rgb = RGBColor(26, 54, 93)
    
    p_scope = doc.add_paragraph(
        f"This Audit Observation & Query Summary Memorandum is prepared by AICA & Associates in connection with our "
        f"{engagement.audit_type} of {engagement.client.client_name} for the financial year {engagement.financial_year}. "
        f"The primary objective of this tracker and memo is to systematically document all queries, information requests, "
        f"and audit observations identified during the audit engagement, record management explanations and remediation evidence, "
        f"and summarize the final disposition of each matter in accordance with Standards on Auditing (SA 230: Audit Documentation "
        f"and SA 260: Communication with Those Charged with Governance) issued by the Institute of Chartered Accountants of India (ICAI)."
    )
    p_scope.paragraph_format.line_spacing = 1.15

    # SECTION 2: EXECUTIVE SUMMARY STATISTICS
    h2 = doc.add_heading("2. Query Lifecycle & Disposition Summary", level=1)
    h2.runs[0].font.color.rgb = RGBColor(26, 54, 93)

    queries = engagement.queries.all()
    total_queries = queries.count()
    closed_normal = queries.filter(status='Closed').count()
    closed_exception = queries.filter(status='Closed - Exception Approved').count()
    withdrawn = queries.filter(status='Withdrawn').count()
    open_count = len(open_queries)

    table_stats = doc.add_table(rows=6, cols=3)
    table_stats.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Disposition Category", "Count of Queries", "% of Total Portfolio"]
    for j, text in enumerate(headers):
        cell = table_stats.rows[0].cells[j]
        set_cell_background(cell, "1E293B")
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    stats_rows = [
        ("Closed Normally with Management Evidence", str(closed_normal), f"{(closed_normal/total_queries*100):.1f}%" if total_queries else "0%"),
        ("Closed via Audit Partner Exception Approval", str(closed_exception), f"{(closed_exception/total_queries*100):.1f}%" if total_queries else "0%"),
        ("Withdrawn / Superseded by Audit Team", str(withdrawn), f"{(withdrawn/total_queries*100):.1f}%" if total_queries else "0%"),
        ("Open Items (Accepted via EX-03 Exception)", str(open_count), f"{(open_count/total_queries*100):.1f}%" if total_queries else "0%"),
        ("Total Audit Queries & Observations Raised", str(total_queries), "100.0%")
    ]

    for i, (cat, count, pct) in enumerate(stats_rows, start=1):
        row = table_stats.rows[i]
        for j, val in enumerate([cat, count, pct]):
            cell = row.cells[j]
            if i == 5:
                set_cell_background(cell, "F1F5F9")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if i == 5:
                r.font.bold = True

    doc.add_paragraph()

    # SECTION 3: ANNEXURE A - FULL QUERY REGISTER
    h_ann_a = doc.add_heading("Annexure A — Master Audit Query Register", level=1)
    h_ann_a.runs[0].font.color.rgb = RGBColor(26, 54, 93)

    table_reg = doc.add_table(rows=1, cols=6)
    table_reg.alignment = WD_TABLE_ALIGNMENT.CENTER
    reg_headers = ["Query No", "Area", "Type", "Priority", "Status", "Closure Remarks"]
    for j, th in enumerate(reg_headers):
        cell = table_reg.rows[0].cells[j]
        set_cell_background(cell, "334155")
        p = cell.paragraphs[0]
        r = p.add_run(th)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)

    for q in queries.order_by('serial_no', 'id'):
        row = table_reg.add_row()
        q_cols = [
            q.query_no or "Draft",
            q.get_area_display().split(' (')[0],
            q.query_type,
            q.priority,
            q.get_status_display(),
            (q.closure_remarks[:120] + "...") if len(q.closure_remarks) > 120 else (q.closure_remarks or "―")
        ]
        for j, val in enumerate(q_cols):
            cell = row.cells[j]
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)

    doc.add_paragraph()

    # SECTION 4: ANNEXURE B - STRUCTURED OBSERVATIONS (5-ELEMENT FINDINGS)
    h_ann_b = doc.add_heading("Annexure B — Key Observations & Control Deficiencies (5-Element Structure)", level=1)
    h_ann_b.runs[0].font.color.rgb = RGBColor(26, 54, 93)

    obs_queries = queries.filter(query_type__in=['Observation', 'Control Deficiency'], carry_to_memo=True)
    if not obs_queries.exists():
        doc.add_paragraph("No major internal control deficiencies or formal observations recorded for this engagement.")
    else:
        for idx, obs in enumerate(obs_queries, start=1):
            h_obs = doc.add_heading(f"Observation {idx}: {obs.query_no or 'Draft'} — {obs.title}", level=2)
            h_obs.runs[0].font.color.rgb = RGBColor(30, 41, 59)
            h_obs.runs[0].font.size = Pt(11)

            t_obs = doc.add_table(rows=6, cols=2)
            t_obs.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            # Fetch latest client response
            latest_client_resp = obs.responses.filter(is_client_submission=True).order_by('-responded_on').first()
            resp_text = latest_client_resp.response_text if latest_client_resp else "No formal client response recorded."

            elements = [
                ("1. Condition (What was found):", obs.condition or obs.description),
                ("2. Criteria (What should be / Standard):", obs.criteria or "Applicable Accounting Standard / Internal Policy"),
                ("3. Cause (Root Cause Analysis):", obs.cause or "Operational process variance"),
                ("4. Effect / Risk (Financial / Compliance Impact):", obs.effect_risk or f"Risk Rating: {obs.risk_rating}"),
                ("5. Audit Recommendation:", obs.recommendation or "Strengthen internal review and reconciliation controls."),
                ("6. Client Management Response:", resp_text)
            ]

            for row_idx, (el_title, el_body) in enumerate(elements):
                row = t_obs.rows[row_idx]
                c0, c1 = row.cells[0], row.cells[1]
                c0.width = Inches(2.2)
                c1.width = Inches(4.6)
                set_cell_background(c0, "F8FAFC")
                set_cell_margins(c0, 40, 40, 80, 80)
                set_cell_margins(c1, 40, 40, 80, 80)

                p0 = c0.paragraphs[0]
                r0 = p0.add_run(el_title)
                r0.font.bold = True
                r0.font.size = Pt(8.5)

                p1 = c1.paragraphs[0]
                r1 = p1.add_run(el_body)
                r1.font.size = Pt(8.5)

            doc.add_paragraph()

    # SECTION 5: ANNEXURE C - PARTNER EXCEPTION APPROVALS
    h_ann_c = doc.add_heading("Annexure C — Audit Partner Exception Register", level=1)
    h_ann_c.runs[0].font.color.rgb = RGBColor(26, 54, 93)

    exceptions = engagement.exceptions.filter(decision='Approved')
    if not exceptions.exists():
        doc.add_paragraph("No partner exceptions were requested or approved during this engagement.")
    else:
        table_exc = doc.add_table(rows=1, cols=6)
        table_exc.alignment = WD_TABLE_ALIGNMENT.CENTER
        exc_headers = ["Ref No", "Exception Type", "Query / Scope", "Justification", "Approved By", "Date"]
        for j, th in enumerate(exc_headers):
            cell = table_exc.rows[0].cells[j]
            set_cell_background(cell, "475569")
            p = cell.paragraphs[0]
            r = p.add_run(th)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8.5)

        for exc in exceptions:
            row = table_exc.add_row()
            exc_cols = [
                exc.reference_no,
                exc.exception_type,
                exc.query.query_no if exc.query else "Engagement Scope",
                (exc.justification[:100] + "...") if len(exc.justification) > 100 else exc.justification,
                exc.approved_by.full_name if exc.approved_by else "Partner",
                exc.approved_on.strftime('%d-%b-%Y') if exc.approved_on else "―"
            ]
            for j, val in enumerate(exc_cols):
                cell = row.cells[j]
                p = cell.paragraphs[0]
                r = p.add_run(val)
                r.font.size = Pt(8)

    doc.add_paragraph()

    # SECTION 6: ANNEXURE D - REPEAT OBSERVATIONS
    h_ann_d = doc.add_heading("Annexure D — Repeat Observations from Prior Year", level=1)
    h_ann_d.runs[0].font.color.rgb = RGBColor(26, 54, 93)

    repeat_queries = queries.filter(is_repeat_observation=True)
    if not repeat_queries.exists():
        doc.add_paragraph("No repeat observations from prior audit cycles identified on this engagement.")
    else:
        for r_q in repeat_queries:
            p_rep = doc.add_paragraph()
            r_rep = p_rep.add_run(f"• {r_q.query_no or 'Draft'}: {r_q.title} (Prior Year Ref: {r_q.prior_year_query_ref.query_no if r_q.prior_year_query_ref else 'Prior FY Audit'}) — Current Status: {r_q.get_status_display()}")
            r_rep.font.size = Pt(9)

    doc.add_paragraph()

    # SECTION 7: SIGN-OFF AND QUALITY REVIEW BLOCK
    h_sign = doc.add_heading("3. Engagement Sign-Off & Quality Review", level=1)
    h_sign.runs[0].font.color.rgb = RGBColor(26, 54, 93)

    table_sign = doc.add_table(rows=4, cols=3)
    table_sign.alignment = WD_TABLE_ALIGNMENT.CENTER

    sign_headers = ["Prepared By (Senior/Manager)", "Reviewed By (Manager/EQCR)", "Approved By (Audit Partner)"]
    for j, th in enumerate(sign_headers):
        cell = table_sign.rows[0].cells[j]
        set_cell_background(cell, "F1F5F9")
        p = cell.paragraphs[0]
        r = p.add_run(th)
        r.font.bold = True
        r.font.size = Pt(9)

    # Names
    table_sign.rows[1].cells[0].paragraphs[0].add_run(f"Name: {generated_by_staff.full_name}")
    table_sign.rows[1].cells[1].paragraphs[0].add_run(f"Name: {reviewer_partner.full_name if reviewer_partner else manager_name}")
    table_sign.rows[1].cells[2].paragraphs[0].add_run(f"Name: {lead_partner.full_name if lead_partner else 'Engagement Partner'}")

    # Designations & Membership
    table_sign.rows[2].cells[0].paragraphs[0].add_run(f"Designation: {generated_by_staff.designation}")
    table_sign.rows[2].cells[1].paragraphs[0].add_run(f"Designation: {'EQCR Reviewer' if reviewer_partner else 'Audit Manager'}")
    table_sign.rows[2].cells[2].paragraphs[0].add_run(f"Membership No: {lead_partner.membership_no if (lead_partner and lead_partner.membership_no) else 'ICAI Qualified'}")

    # Date & Signature Status
    table_sign.rows[3].cells[0].paragraphs[0].add_run(f"Date: {timezone.now().strftime('%d-%b-%Y')}")
    table_sign.rows[3].cells[1].paragraphs[0].add_run(f"Date: {timezone.now().strftime('%d-%b-%Y')}")
    if is_draft:
        table_sign.rows[3].cells[2].paragraphs[0].add_run("Status: [Draft - Pending Sign-Off]")
    else:
        table_sign.rows[3].cells[2].paragraphs[0].add_run(f"Signed Off on: {timezone.now().strftime('%d-%b-%Y %H:%M')}")

    for row in table_sign.rows:
        for cell in row.cells:
            set_cell_margins(cell, 40, 40, 60, 60)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)

    # Ensure output directory exists
    memos_dir = Path(settings.MEDIA_ROOT) / 'memos' / f"eng_{engagement.id}"
    memos_dir.mkdir(parents=True, exist_ok=True)
    
    file_name = f"Audit_Memo_{engagement.engagement_code.replace('/', '_')}_{version_no.replace(' ', '_')}.docx"
    file_path = memos_dir / file_name
    doc.save(str(file_path))

    relative_path = f"memos/eng_{engagement.id}/{file_name}"

    memo_record = Memo.objects.create(
        engagement=engagement,
        version_no=version_no,
        generated_by=generated_by_staff,
        generated_on=timezone.now(),
        file_path=relative_path,
        status='Draft' if is_draft else 'Signed',
        signed_by=lead_partner if not is_draft else None,
        signed_on=timezone.now() if not is_draft else None,
        open_query_count_at_generation=open_count,
        exception_count=exceptions.count()
    )

    engagement.memo_generated_on = timezone.now()
    engagement.memo_version = version_no
    if not is_draft:
        engagement.status = 'Memo Issued'
        engagement.memo_signed_by = lead_partner
    engagement.save()

    log_audit(
        request=request,
        entity_type='Memo',
        entity_id=f"{engagement.engagement_code}-{version_no}",
        action='Memo Generate' if is_draft else 'Memo Sign-off',
        field_changed='status',
        old_value='None' if is_draft else 'Draft',
        new_value='Draft' if is_draft else 'Signed',
        user=request.user if request else None
    )

    return memo_record, str(file_path)
