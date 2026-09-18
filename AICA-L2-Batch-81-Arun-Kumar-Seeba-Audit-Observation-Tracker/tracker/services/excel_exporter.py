import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_queries_to_excel(queries_queryset):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Audit Queries"

    # Header styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "Query No", "Engagement Code", "Client Name", "Audit Type", "FY",
        "Audit Area", "Query Type", "Title", "Priority", "Risk Rating",
        "Raised By", "Raised On", "Due Date", "Status", "Days Open",
        "Ageing Bucket", "Amount Involved (₹)", "Closed By", "Closed On", "Closure Remarks"
    ]

    ws.append(headers)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_font = Font(name="Calibri", size=10)

    for q in queries_queryset:
        row_data = [
            q.query_no or "Draft",
            q.engagement.engagement_code,
            q.engagement.client.client_name,
            q.engagement.audit_type,
            q.engagement.financial_year,
            q.get_area_display(),
            q.query_type,
            q.title,
            q.priority,
            q.risk_rating,
            q.raised_by.full_name if q.raised_by else "―",
            q.raised_on.strftime('%d-%b-%Y') if q.raised_on else "―",
            q.due_date.strftime('%d-%b-%Y') if q.due_date else "―",
            q.get_status_display(),
            q.days_open,
            q.ageing_bucket,
            float(q.amount_involved) if q.amount_involved else None,
            q.closed_by.full_name if q.closed_by else "―",
            q.closed_on.strftime('%d-%b-%Y') if q.closed_on else "―",
            q.closure_remarks or "―"
        ]
        ws.append(row_data)

    # Styling data cells and auto-adjust widths
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = row_font
            cell.border = thin_border
            if cell.column == 17 and cell.value: # Amount
                cell.number_format = '#,##,##0.00'

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
