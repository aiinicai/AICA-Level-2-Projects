"""Run: python test_cases.py. Core tests need only Pandas and OpenPyXL.
This does NOT open a live Tkinter window or verify a Windows EXE.
"""
from io import BytesIO
from pathlib import Path
import ast
import platform
import sys

import pandas as pd
import openpyxl
from openpyxl import Workbook, load_workbook
from app import (REQUIRED, REPORT_COLUMNS, analyze, excel_bytes, filter_report,
                 paise, read_trial_balance, sample_data)

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "testing_data"
DATA.mkdir(exist_ok=True)
results = []


def workbook(frame):
    return excel_bytes({"Trial Balance": frame})


def assert_true(condition, message="Unexpected result"):
    if not condition:
        raise AssertionError(message)


def rejected(raw, fragment):
    try:
        read_trial_balance(raw)
    except ValueError as exc:
        assert_true(fragment.casefold() in str(exc).casefold(), str(exc))
        return str(exc)
    raise AssertionError("Invalid input was accepted")


def test(number, scenario, expected, action):
    try:
        actual = action()
        status = "PASS"
    except Exception as exc:
        actual, status = f"{type(exc).__name__}: {exc}", "FAIL"
    results.append({"Test No.": number, "Test Scenario": scenario,
                    "Expected Result": expected, "Actual Result": str(actual),
                    "Status – PASS/FAIL": status})
    print(f"Test {number}: {status} — {scenario}")


sample = sample_data()
valid = workbook(sample)
(ROOT / "sample_trial_balance.xlsx").write_bytes(valid)
(DATA / "valid_24_ledgers.xlsx").write_bytes(valid)
missing = workbook(sample.drop(columns=["Credit"]))
(DATA / "missing_credit_column.xlsx").write_bytes(missing)
empty_book = BytesIO()
Workbook().save(empty_book)
empty = empty_book.getvalue()
(DATA / "empty_workbook.xlsx").write_bytes(empty)
(DATA / "headers_only.xlsx").write_bytes(workbook(sample.iloc[:0]))
(DATA / "zero_byte.xlsx").write_bytes(b"")
bad_number = sample.copy().astype(object)
bad_number.loc[0, "Debit"] = "not a number"
(DATA / "invalid_numeric.xlsx").write_bytes(workbook(bad_number))
blank = sample.copy().astype(object)
blank.loc[0, "Closing Balance"] = None
(DATA / "blank_closing_balance.xlsx").write_bytes(workbook(blank))
parsed = read_trial_balance(valid)
report, summary = analyze(parsed)


def contains(kind, ledger):
    selected = report[(report["Exception Type"] == kind) & (report["Ledger Name"] == ledger)]
    assert_true(not selected.empty, f"{ledger}: {kind} not found")
    return f"Identified {ledger}: {kind}; {len(selected)} matching report row(s)."


def valid_test():
    assert_true(len(parsed) == 24 and list(parsed.columns) == REQUIRED)
    return "24 ledgers parsed successfully by the upload handler; native Open dialog not exercised."


def mismatch_test():
    assert_true(summary["Total Debit"] == 2449084.0, "Debit total incorrect")
    assert_true(summary["Total Credit"] == 1553074.0, "Credit total incorrect")
    assert_true(summary["Difference"] == 896010.0, "Difference incorrect")
    contains("Trial Balance Difference Identified", "[Overall Trial Balance]")
    return "Debit 2,449,084.00; Credit 1,553,074.00; Difference +896,010.00; TB-level flag generated."


def duplicate_test():
    subset = report[report["Exception Type"] == "Duplicate Ledger Name"]
    assert_true(len(subset) == 2 and set(subset["Ledger Name"]) == {"Office Supplies"})
    return "Both Office Supplies source rows identified (2 flags)."


def export_test():
    content = excel_bytes({"Exceptions": report})
    (ROOT / "sample_audit_exception_report.xlsx").write_bytes(content)
    restored = pd.read_excel(BytesIO(content), engine="openpyxl")
    pd.testing.assert_frame_equal(restored, report, check_dtype=False)
    assert_true(list(restored.columns) == REPORT_COLUMNS)
    return f"Valid XLSX bytes produced and all {len(restored)} report rows round-tripped exactly. Native Save dialog not exercised."


test(1, "Upload valid Excel file (handler)", "Trial Balance loads successfully with 24 ledgers.", valid_test)
test(2, "Required Credit column missing", "Clear validation error; handler rejects safely.", lambda: rejected(missing, "Missing required columns: Credit"))
test(3, "Mismatched debit and credit totals", "TB difference of +896,010.00 identified.", mismatch_test)
test(4, "Negative closing balance", "Cash in Hand flagged under Negative Balance.", lambda: contains("Negative Balance", "Cash in Hand"))
test(5, "Duplicate ledger names", "Both Office Supplies rows identified.", duplicate_test)
test(6, "Materiality 500,000; closing 750,000", "Plant and Machinery flagged High Value.", lambda: contains("High Value", "Plant and Machinery"))
test(7, "Round-figure debit or credit total", "Rent Expense flagged for aggregate review.", lambda: contains("Round Figure Transaction – Review Required", "Rent Expense"))
test(8, "Expense with credit closing balance", "Travel Expense flagged as Unusual Balance.", lambda: contains("Unusual Balance – Review Required", "Travel Expense"))
test(9, "Export exception report (generation/round trip)", "Excel bytes contain the exact report data; native Save dialog requires local verification.", export_test)
test(10, "Empty Excel workbook", "Understandable empty worksheet error.", lambda: rejected(empty, "first worksheet is empty"))
test(11, "Headers with no data rows", "No ledger rows error.", lambda: rejected(workbook(sample.iloc[:0]), "No ledger rows"))
test(12, "Zero-byte file", "Empty file error.", lambda: rejected(b"", "uploaded file is empty"))
test(13, "Blank / NaN amount", "Row and column-specific error; no silent zero fill.", lambda: rejected(workbook(blank), "Excel row 2, Closing Balance: blank"))
test(14, "Non-numeric amount", "Numeric conversion error with row and column.", lambda: rejected(workbook(bad_number), "Excel row 2, Debit"))
test(15, "Transactions with zero closing", "Clearing Account flagged; dormant account excluded.", lambda: contains("Zero Closing Balance", "Clearing Account"))


def round_zero_test():
    subset = report[report["Ledger Name"] == "Dormant Account"]
    assert_true(subset.empty)
    return "Dormant Account has no exceptions; zeros are excluded from multiples and movement checks."


def boundary_test():
    frame = pd.DataFrame([
        ["At boundary", "Assets", 500000, 0, 0, 500000],
        ["Above", "Assets", 500000.01, 0, 0, 500000.01],
        ["Credit above", "Liabilities", -500000.01, 0, 0, -500000.01],
    ], columns=REQUIRED)
    rows, _ = analyze(read_trial_balance(workbook(frame)))
    names = set(rows.loc[rows["Exception Type"] == "High Value", "Ledger Name"])
    assert_true(names == {"Above", "Credit above"})
    return "Exactly 500,000 excluded; +/-500,000.01 included."


def precise_test():
    frame = pd.DataFrame([["A", "Assets", 0, 0.10, 0.30, -0.20],
                          ["B", "Assets", 0, 0.20, 0, 0.20]], columns=REQUIRED)
    rows, sums = analyze(read_trial_balance(workbook(frame)), round_multiple=0.10)
    assert_true(sums["Difference"] == 0)
    assert_true("Trial Balance Difference Identified" not in set(rows["Exception Type"]))
    assert_true(len(rows[rows["Exception Type"].str.startswith("Round Figure")]) == 2)
    return "0.10 + 0.20 exactly equals 0.30 using paise; 0.10 multiples detected."


def custom_rules_test():
    rows, _ = analyze(parsed, rules={"Expenses": "Ignore", "Liabilities": "Debit"})
    unusual = set(rows.loc[rows["Exception Type"] == "Unusual Balance – Review Required", "Ledger Name"])
    assert_true(unusual == {"Bank Loan"})
    return "Ignore suppresses expense checks; Debit for liabilities flags Bank Loan; unmapped groups ignored."


def filters_test():
    selected = filter_report(report, ["High Value"], ["Assets"])
    assert_true(set(selected["Ledger Name"]) == {"Plant and Machinery"})
    assert_true(filter_report(report, [], []).equals(report))
    restored = pd.read_excel(BytesIO(excel_bytes({"Exceptions": selected, "All Exceptions": report})))
    pd.testing.assert_frame_equal(restored, selected, check_dtype=False)
    return "AND filters produce Plant and Machinery; blank filters retain all; filtered export verified."


def mutate_error(column, value, message):
    frame = sample.copy().astype(object)
    frame.loc[0, column] = value
    return rejected(workbook(frame), message)


def duplicate_headers_test():
    book = load_workbook(BytesIO(valid))
    book.active.cell(1, 7, "Debit")
    buffer = BytesIO()
    book.save(buffer)
    return rejected(buffer.getvalue(), "headers must not be duplicated")


def formula_test():
    book = load_workbook(BytesIO(valid))
    book.active["F2"] = "=C2+D2-E2"
    buffer = BytesIO()
    book.save(buffer)
    return rejected(buffer.getvalue(), "Formula cells")


def literal_export_test():
    name = '=HYPERLINK("https://example.invalid","text")'
    content = excel_bytes({"Exceptions": pd.DataFrame([[name]], columns=["Ledger Name"])})
    book = load_workbook(BytesIO(content), data_only=False)
    assert_true(book.active["A2"].data_type == "s" and book.active["A2"].value == name)
    return "Formula-like ledger name exported as literal string, not executable Excel formula."


def no_exceptions_test():
    frame = pd.DataFrame([["Normal Asset", "Assets", 1, 0, 0, 1]], columns=REQUIRED)
    rows, sums = analyze(read_trial_balance(workbook(frame)))
    assert_true(rows.empty and list(rows.columns) == REPORT_COLUMNS and sums["Total Exceptions"] == 0)
    restored = pd.read_excel(BytesIO(excel_bytes({"Exceptions": rows})))
    assert_true(restored.empty and list(restored.columns) == REPORT_COLUMNS)
    return "Empty exception result retains all five headers and exports successfully."


def normalized_duplicate_test():
    frame = sample.copy()
    frame.loc[21, "Ledger Name"] = "  office   SUPPLIES "
    rows, _ = analyze(read_trial_balance(workbook(frame)))
    assert_true(len(rows[rows["Exception Type"] == "Duplicate Ledger Name"]) == 2)
    return "Case, outer spaces and repeated internal spaces ignored for duplicate matching."


def syntax_test():
    for filename in ("app.py", "test_cases.py"):
        text = (ROOT / filename).read_text(encoding="utf-8")
        ast.parse(text, filename=filename)
        compile(text, filename, "exec")
    return "Both Python sources compile. Pandas/OpenPyXL imports verified; live Tkinter window not tested here."


def settings_test():
    for kwargs in ({"materiality": -1}, {"round_multiple": 0}, {"round_multiple": -1}, {"rules": {"Assets": "Invalid"}}):
        try:
            analyze(parsed, **kwargs)
        except ValueError:
            continue
        raise AssertionError("Invalid setting accepted")
    return "Negative materiality, non-positive multiple and invalid expected nature rejected."


def blank_text_test():
    return mutate_error("Ledger Name", "   ", "Ledger Name: blank")


test(16, "Zero movement and zero closing", "Dormant account not flagged.", round_zero_test)
test(17, "Strict absolute threshold boundary", "Only absolute amounts greater than threshold flagged.", boundary_test)
test(18, "Decimal precision and configurable multiple", "No float mismatch; decimal multiples work.", precise_test)
test(19, "Configurable accounting rules", "User-selected nature and Ignore respected.", custom_rules_test)
test(20, "Both filters and filtered export", "Type/group filters combine; exported rows match.", filters_test)
test(21, "Negative period debit", "Debit must be non-negative.", lambda: mutate_error("Debit", -1, "must be non-negative"))
test(22, "Too many decimal places", "Reject rather than silently round.", lambda: mutate_error("Debit", 1.001, "two decimal places"))
test(23, "Invalid/corrupted workbook", "Understandable invalid Excel error.", lambda: rejected(b"not an xlsx", "Cannot read"))
test(24, "Duplicate required header", "Reject ambiguous columns.", duplicate_headers_test)
test(25, "Formula input", "Require pasted values instead of cached formulas.", formula_test)
test(26, "Formula-like text export", "Text remains text in generated Excel.", literal_export_test)
test(27, "No exceptions", "Empty report exports with headers; zero count.", no_exceptions_test)
test(28, "Normalized duplicate names", "Both case/whitespace variants flagged.", normalized_duplicate_test)
test(29, "Syntax and available imports", "Python sources compile without syntax errors.", syntax_test)
test(30, "Invalid review settings", "Clear settings errors.", settings_test)
test(31, "Blank ledger name", "Reject whitespace-only name.", blank_text_test)
test(32, "Infinite amount", "Reject non-finite numeric input.", lambda: mutate_error("Debit", float("inf"), "numeric amount"))

# Desktop-specific tests use mocks for dialogs/widgets, not a live Tk window.
from unittest.mock import Mock
from app import AnalyzerWindow, DEFAULT_RULES, report_bytes


def full_export_test():
    content = report_bytes(report, summary, DEFAULT_RULES, 500000, 10000,
                           ["High Value"], ["Assets"], "valid_24_ledgers.xlsx")
    sheets = pd.read_excel(BytesIO(content), sheet_name=None)
    assert_true(set(sheets) == {"Exceptions", "All Exceptions", "Summary", "Settings", "Expected Nature"})
    assert_true(len(sheets["Exceptions"]) == 1 and len(sheets["All Exceptions"]) == 27)
    values = dict(sheets["Settings"].itertuples(index=False, name=None))
    assert_true(values["Materiality (INR)"] == 500000 and values["Ledger Group filter"] == "Assets")
    (ROOT / "sample_audit_exception_report.xlsx").write_bytes(
        report_bytes(report, summary, DEFAULT_RULES, 500000, 10000, source="sample_trial_balance.xlsx"))
    return "Five workbook sheets verified, filtered/full rows and settings agree; full sample export saved."


def mock_window():
    window = AnalyzerWindow.__new__(AnalyzerWindow)
    window.root = Mock()
    window.dialog = Mock()
    window.message = Mock()
    window.data = parsed
    window.dirty = False
    window.status = Mock()
    window.export_button = Mock()
    return window


def dirty_state_test():
    window = mock_window()
    window.mark_dirty()
    assert_true(window.dirty)
    window.export_button.configure.assert_called_with(state="disabled")
    window.export()
    window.message.showinfo.assert_called_once()
    window.dialog.asksaveasfilename.assert_not_called()
    return "Mocked controller disables export when settings change; stale export blocked."


def invalid_open_test():
    window = mock_window()
    original = window.data
    window.dialog.askopenfilename.return_value = str(DATA / "missing_credit_column.xlsx")
    window.open_file()
    assert_true(window.data is original)
    window.message.showerror.assert_called_once()
    assert_true("Missing required columns" in window.message.showerror.call_args.args[1])
    return "Actual invalid workbook rejected; mocked error dialog called and prior data preserved."


def save_cancel_test():
    window = mock_window()
    window.dialog.asksaveasfilename.return_value = ""
    producer = Mock()
    window.save_workbook("unused.xlsx", producer)
    producer.assert_not_called()
    window.message.showerror.assert_not_called()
    return "Cancelled mocked save dialog does not generate or write a file."


def save_failure_test():
    window = mock_window()
    # Simulate a locked workbook without changing the fixture.
    window.dialog.asksaveasfilename.return_value = str(DATA / "headers_only.xlsx")
    producer = Mock(side_effect=PermissionError("Simulated file lock"))
    window.save_workbook("report.xlsx", producer)
    window.message.showerror.assert_called_once()
    assert_true("Simulated file lock" in window.message.showerror.call_args.args[1])
    return "Simulated save permission error caught and understandable error dialog requested."


test(33, "Desktop report with all export sheets", "Filtered/full rows, summary, rules and settings agree.", full_export_test)
test(34, "Changed settings / stale export guard (mock)", "Export disabled until recalculated.", dirty_state_test)
test(35, "Invalid replacement workbook (mock dialog)", "Clear error; previous valid data preserved.", invalid_open_test)
test(36, "Cancel Save dialog (mock)", "No file generation and no error.", save_cancel_test)
test(37, "Save failure (mock)", "File write failure reported without crashing callback.", save_failure_test)

try:
    import tkinter
    tk_status = "Tkinter import succeeded; a live window was not exercised by this test runner."
except ImportError as exc:
    tk_status = f"Tkinter runtime unavailable: {exc}"

frame = pd.DataFrame(results)
frame.to_csv(ROOT / "test_results.csv", index=False, encoding="utf-8-sig")
frame[["Test No.", "Test Scenario", "Expected Result"]].to_csv(ROOT / "test_cases.csv", index=False, encoding="utf-8-sig")
context = pd.DataFrame([
    ("Scope", "Excel handling, calculations, exceptions, export, syntax and mocked controller checks; no live Tkinter window."),
    ("Python", platform.python_version()), ("Platform", platform.system()),
    ("Pandas", pd.__version__), ("OpenPyXL", openpyxl.__version__),
    ("Tkinter runtime", tk_status),
    ("UI status", "NOT RUN — rendered widgets, real native file dialogs and clicks require local verification."),
    ("EXE status", "NOT BUILT / NOT RUN — Windows build environment required."),
], columns=["Item", "Detail"])
(ROOT / "test_results.xlsx").write_bytes(excel_bytes({"Test Results": frame, "Test Scope": context}))
passed = int((frame["Status – PASS/FAIL"] == "PASS").sum())
lines = ["# Test result report — Tkinter desktop edition", "", f"**Core/controller results: {passed}/{len(frame)} PASS.**", "",
         "Scope: actual Excel/core tests and mocked controller checks. Live Tkinter widgets, native file dialogs and Windows EXE execution were NOT RUN.",
         "Tests 1 and 9 exercise the same open/export functions used by the UI, not GUI clicks. Tests 34–37 use mocks, not actual widgets.",
         tk_status, "",
         f"Environment: Python {platform.python_version()}, Pandas {pd.__version__}, OpenPyXL {openpyxl.__version__}.", "",
         "| " + " | ".join(frame.columns) + " |", "| " + " | ".join(["---"] * len(frame.columns)) + " |"]
for row in frame.itertuples(index=False, name=None):
    lines.append("| " + " | ".join(str(cell).replace("|", "/").replace("\n", " ") for cell in row) + " |")
(ROOT / "TEST_RESULTS.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
print(f"Completed: {passed}/{len(frame)} core/controller tests passed; reports and fixtures generated.")
sys.exit(0 if passed == len(frame) else 1)
