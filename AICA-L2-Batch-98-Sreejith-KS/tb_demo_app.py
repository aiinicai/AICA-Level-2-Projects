"""
Financial Statements from Tally - capstone application (Streamlit)

Flow:  Tally trial balance  ->  previous year  ->  classification  ->  depreciation  ->  statements  ->  checks  ->  Excel

Frameworks:  AS (Schedule III, Division I)  |  Ind AS (Schedule III, Division II)  |  Non-corporate entity (ICAI framework)
Previous year:  Tally (previous period)  |  uploaded trial balance  |  last year's workbook produced by this app
Depreciation: Schedule II (Companies Act, SLM or WDV)  |  Income Tax Act (WDV, block of assets)  |  deferred tax (information)

Run it with:  python run_demo.py      (or:  python -m streamlit run tb_demo_app.py)
All data stays on this computer. The app only talks to Tally on localhost.
"""
import io
import re
from datetime import date, timedelta
from xml.sax.saxutils import escape
import xml.etree.ElementTree as ET

import pandas as pd
import requests
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

DEFAULT_URL = "http://localhost:9000"

# ====================================================================== heads ==
CAP = "Capital"
HEAD_LIST = [
    "Capital", "Reserves and Surplus", "Long-term Borrowings", "Other Non-current Liabilities", "Short-term Borrowings", "Trade Payables",
    "Other Current Liabilities (Statutory Dues)", "Other Current Liabilities", "Short-term Provisions",
    "Property, Plant & Equipment", "Investments", "Other Non-Current Assets", "Inventories", "Trade Receivables",
    "Cash & Cash Equivalents", "Short-term Loans and Advances", "Other Current Assets",
    "Revenue from Operations", "Other Income", "Purchases of Stock-in-Trade", "Employee Benefits Expense",
    "Finance Costs", "Depreciation & Amortisation", "Changes in Inventories", "Tax Expense", "Other Expenses",
]
CAP_NAME = {"Company": "Share Capital", "Firm (partnership or LLP)": "Partners' Capital Accounts",
            "Proprietorship": "Proprietor's Capital Account"}
CR_HEADS = {"Reserves and Surplus", "Long-term Borrowings", "Other Non-current Liabilities", "Short-term Borrowings", "Trade Payables",
            "Other Current Liabilities (Statutory Dues)", "Other Current Liabilities", "Short-term Provisions",
            "Revenue from Operations", "Other Income"} | set(CAP_NAME.values())
PL_HEADS = {"Revenue from Operations", "Other Income", "Purchases of Stock-in-Trade", "Employee Benefits Expense",
            "Finance Costs", "Depreciation & Amortisation", "Changes in Inventories", "Tax Expense", "Other Expenses"}
EXPORT_ORDER = (list(CAP_NAME.values()) + HEAD_LIST[1:10] + ["Opening Stock"] + HEAD_LIST[10:] + ["Unclassified"])
DEP_GROUP = "Depreciation calculator"

GROUP_HEAD = {
    "Sundry Debtors": "Trade Receivables", "Sundry Creditors": "Trade Payables",
    "Bank Accounts": "Cash & Cash Equivalents", "Cash-in-Hand": "Cash & Cash Equivalents",
    "Stock-in-Hand": "Inventories", "Fixed Assets": "Property, Plant & Equipment", "Investments": "Investments",
    "Deposits (Asset)": "Short-term Loans and Advances", "Loans & Advances (Asset)": "Short-term Loans and Advances",
    "Current Assets": "Other Current Assets", "Misc. Expenses (ASSET)": "Other Non-Current Assets",
    "Duties & Taxes": "Other Current Liabilities (Statutory Dues)", "Provisions": "Short-term Provisions",
    "Current Liabilities": "Other Current Liabilities", "Bank OD A/c": "Short-term Borrowings",
    "Secured Loans": "Long-term Borrowings", "Unsecured Loans": "Long-term Borrowings",
    "Loans (Liability)": "Long-term Borrowings", "Reserves & Surplus": "Reserves and Surplus",
    "Capital Account": "Capital", "Sales Accounts": "Revenue from Operations", "Direct Incomes": "Other Income",
    "Indirect Incomes": "Other Income", "Purchase Accounts": "Purchases of Stock-in-Trade",
    "Direct Expenses": "Other Expenses", "Indirect Expenses": "Other Expenses",
}

FRAMEWORKS = ["AS (Schedule III, Division I)", "Ind AS (Schedule III, Division II)", "Non-corporate entity (ICAI framework)"]
FRAME_TEXT = {
    FRAMEWORKS[0]: "(Schedule III, Division I - Accounting Standards, Companies Act 2013) | Amounts in Rupees",
    FRAMEWORKS[1]: "(Schedule III, Division II - Ind AS, Companies Act 2013) | Amounts in Rupees",
    FRAMEWORKS[2]: "(Non-corporate entity, ICAI framework) | Amounts in Rupees",
}


def head_sign(head):
    return 1 if head in CR_HEADS else -1


def keyword_head(ledger, head):
    """Finer split of expense ledgers by name."""
    n = ledger.lower()
    if head == "Other Current Assets" and re.search(r"closing stock|stock-in-trade|inventor", n):
        return "Inventories"          # Tally does not allow a plain ledger under Stock-in-Hand in journals
    if head != "Other Expenses":
        return head
    if "depreciation" in n:
        return "Depreciation & Amortisation"
    if "changes in inventor" in n or "increase/decrease in stock" in n or "(increase) / decrease in stock" in n:
        return "Changes in Inventories"
    if re.search(r"salar|wage|bonus|staff|gratuity|\bpf\b|provident|\besi\b|labour", n):
        return "Employee Benefits Expense"
    if "interest" in n:
        return "Finance Costs"
    return head


def is_plac(series):
    return series.str.strip().str.lower().isin(["profit & loss a/c", "profit and loss a/c"])


def inr(x):
    x = int(round(float(x)))
    s = str(abs(x))
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        s = re.sub(r"(\d)(?=(\d\d)+$)", r"\1,", head) + "," + tail
    return ("-" if x < 0 else "") + s


def ordinal(n):
    return "%d%s" % (n, "th" if 10 <= n % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th"))


def long_date(d):
    return "%s %s %d" % (ordinal(d.day), d.strftime("%B"), d.year)


def period_dates(year_end):
    try:
        prev_end = date(year_end.year - 1, year_end.month, year_end.day)
    except ValueError:
        prev_end = date(year_end.year - 1, 2, 28)
    return prev_end + timedelta(days=1), prev_end


# ===================================================================== tally ==
def sanitize(text):
    """Remove characters Tally sometimes emits that are illegal in XML."""
    text = re.sub(r"&#(?:[0-8]|11|12|1[4-9]|2[0-9]|3[01]);", "", text)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    return re.sub(r"^<\?xml[^>]*\?>", "", text.strip())


def tally_collection(url, coll_type, methods, company="", period=None, timeout=240):
    """Ask Tally for a collection. One NATIVEMETHOD tag per field (reliable)."""
    svc = "<SVCURRENTCOMPANY>%s</SVCURRENTCOMPANY>" % escape(company) if company else ""
    if period:
        svc += "<SVFROMDATE>%s</SVFROMDATE><SVTODATE>%s</SVTODATE>" % (period[0].strftime("%Y%m%d"), period[1].strftime("%Y%m%d"))
    native = "".join("<NATIVEMETHOD>%s</NATIVEMETHOD>" % m for m in methods)
    body = (
        "<ENVELOPE><HEADER><VERSION>1</VERSION><TALLYREQUEST>Export</TALLYREQUEST>"
        "<TYPE>Collection</TYPE><ID>DemoColl</ID></HEADER><BODY><DESC><STATICVARIABLES>"
        "<SVEXPORTFORMAT>$$SysName:XML</SVEXPORTFORMAT>%s</STATICVARIABLES><TDL><TDLMESSAGE>"
        '<COLLECTION NAME="DemoColl" ISMODIFY="No"><TYPE>%s</TYPE>%s</COLLECTION>'
        "</TDLMESSAGE></TDL></DESC></BODY></ENVELOPE>" % (svc, coll_type, native)
    )
    try:
        r = requests.post(url, data=body.encode("utf-8"), headers={"Content-Type": "text/xml"}, timeout=timeout)
    except requests.exceptions.ReadTimeout:
        raise RuntimeError("Tally did not answer the %s request within %d seconds%s. Go to Tally, press Esc until you are at the Gateway of Tally "
                           "(no pop-up or report waiting for a key), and try again." % (coll_type, timeout, " for the selected period" if period else ""))
    raw = r.content
    text = raw.decode("utf-16", errors="ignore") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else raw.decode("utf-8", errors="ignore")
    root = ET.fromstring(sanitize(text))
    rows = []
    for el in root.iter(coll_type.upper()):
        d = {c.tag.upper(): (c.text or "").strip() for c in el}
        if not d.get("NAME"):
            d["NAME"] = el.get("NAME", "")
        rows.append(d)
    return rows


def parse_amount(s):
    """Return (number, 'Dr'/'Cr'/None)."""
    if not s:
        return 0.0, None
    t = s.replace(",", "").strip()
    m = re.search(r"-?\d+(?:\.\d+)?", t)
    if not m:
        return 0.0, None
    tag = "Dr" if re.search(r"\bDr\b", t, re.I) else ("Cr" if re.search(r"\bCr\b", t, re.I) else None)
    return float(m.group()), tag


def classify(ledger, parent, groups):
    if ledger.strip().lower() in ("profit & loss a/c", "profit and loss a/c"):
        return "Reserves and Surplus", "Profit & Loss A/c"
    chain, node, seen = [], parent, set()
    while node and node != "Primary" and node not in seen:
        seen.add(node)
        chain.append(node)
        node = groups.get(node, "")
    head = ""
    for g in chain:
        if g in GROUP_HEAD:
            head = GROUP_HEAD[g]
            break
    return keyword_head(ledger, head), (chain[-1] if chain else parent)


def load_trial_balance(url, company, period=None):
    groups_rows = tally_collection(url, "Group", ["Name", "Parent"], company, period)
    groups = {g["NAME"]: g.get("PARENT", "") for g in groups_rows}
    led = tally_collection(url, "Ledger", ["Name", "Parent", "ClosingBalance"], company, period)
    if not led:
        raise ValueError("Tally returned no ledgers. Check that a company is open in Tally.")
    recs = []
    for r in led:
        val, tag = parse_amount(r.get("CLOSINGBALANCE", ""))
        if abs(val) < 0.005:
            continue
        head, top = classify(r["NAME"], r.get("PARENT", ""), groups)
        recs.append({"Ledger": r["NAME"], "Group": r.get("PARENT", ""), "Top group": top,
                     "Head": head, "_val": val, "_tag": tag})
    if not recs:
        raise ValueError("Tally returned ledgers but none has a balance for this period. If this is the previous year, the data may be in a different company; use the upload option instead.")
    df = pd.DataFrame(recs)
    cs_items = None
    try:
        stock = tally_collection(url, "StockItem", ["Name", "ClosingValue"], company, period, timeout=25)
        if stock:
            cs_items = abs(sum(parse_amount(r.get("CLOSINGVALUE", ""))[0] for r in stock))
    except Exception:
        cs_items = None

    def dr_pos(row, flip):  # debit-positive amount; Tally normally shows debits as negative numbers
        if row["_tag"] == "Dr":
            return abs(row["_val"])
        if row["_tag"] == "Cr":
            return -abs(row["_val"])
        return row["_val"] if flip else -row["_val"]

    df["Dr_pos"] = df.apply(lambda r: dr_pos(r, False), axis=1)
    asset_heads = {"Trade Receivables", "Cash & Cash Equivalents", "Property, Plant & Equipment", "Inventories"}
    note = ""
    if df.loc[df["Head"].isin(asset_heads), "Dr_pos"].sum() < 0:
        df["Dr_pos"] = df.apply(lambda r: dr_pos(r, True), axis=1)
        note = "Sign convention was reversed automatically (assets came out as credits)."
    df = df.drop(columns=["_val", "_tag"])
    df.attrs["groups"] = groups
    return df, note, cs_items




# ============================================================ previous year ==
def norm_name(s):
    s = str(s).lower().replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", "", s)


EXPORT_TO_INTERNAL = {v: "Capital" for v in CAP_NAME.values()}


def py_classify(py, cy, authoritative=False, groups=None):
    """Give every previous-year ledger a head and say where it came from.
    Order: head given by last year's workbook; same ledger name in the current year (after the user's fixes);
    group rule (if the file has a group column); similar name in the current year (flagged); otherwise blank."""
    py = py.copy()
    cy_map = {}
    for l, h in zip(cy["Ledger"], cy["Head"]):
        if h:
            cy_map[norm_name(l)] = (l, h)
    names = list(cy_map)
    import difflib
    heads, srcs = [], []
    for _, r in py.iterrows():
        given = r.get("Head", "") if "Head" in py.columns else ""
        key = norm_name(r["Ledger"])
        if authoritative and given:
            heads.append(given); srcs.append("Last year's workbook"); continue
        if key in cy_map:
            heads.append(cy_map[key][1]); srcs.append("Same ledger in current year"); continue
        grp = r.get("Group", "") if "Group" in py.columns else ""
        if grp and groups is not None:
            h, _t = classify(r["Ledger"], grp, groups)
            if h:
                heads.append(h); srcs.append("Group rule"); continue
        m = difflib.get_close_matches(key, names, n=1, cutoff=0.88)
        if m:
            heads.append(cy_map[m[0]][1]); srcs.append("Similar to: " + cy_map[m[0]][0]); continue
        heads.append(given or ""); srcs.append("Group rule" if given else "Not classified")
    py["Head"], py["Source"] = heads, srcs
    return py


HEAD_COLS = ("head", "schedule iii head", "statement head", "classification", "mapped head", "mapping", "schedule iii classification",
             "schedule iii mapping", "bs / pl head", "grouping")
LED_COLS = ("ledger", "particulars", "ledger name", "account", "account name", "name of ledger")


PRIMARY_MAP = {
    "share capital": "Capital", "equity share capital": "Capital", "reserves and surplus": "Reserves and Surplus",
    "share application money pending allotment": "Other Current Liabilities",
    "long-term borrowings": "Long-term Borrowings", "long term borrowings": "Long-term Borrowings",
    "other long-term liabilities": "Other Non-current Liabilities", "long-term provisions": "Other Non-current Liabilities",
    "deferred tax liabilities (net)": "Other Non-current Liabilities",
    "short-term borrowings": "Short-term Borrowings", "trade payables": "Trade Payables",
    "other current liabilities": "Other Current Liabilities", "short-term provisions": "Short-term Provisions",
    "property, plant and equipment": "Property, Plant & Equipment", "intangible assets": "Property, Plant & Equipment",
    "capital work-in-progress": "Property, Plant & Equipment", "intangible assets under development": "Property, Plant & Equipment",
    "non-current investments": "Investments", "current investments": "Investments", "investments": "Investments",
    "deferred tax assets (net)": "Other Non-Current Assets", "long-term loans and advances": "Other Non-Current Assets",
    "other non-current assets": "Other Non-Current Assets", "inventories": "Inventories", "trade receivables": "Trade Receivables",
    "cash and bank balances": "Cash & Cash Equivalents", "cash and cash equivalents": "Cash & Cash Equivalents",
    "short-term loans and advances": "Short-term Loans and Advances", "other current assets": "Other Current Assets",
    "revenue from operations": "Revenue from Operations", "other income": "Other Income",
    "cost of materials consumed": "Purchases of Stock-in-Trade", "purchases of stock-in-trade": "Purchases of Stock-in-Trade",
    "purchase of stock-in-trade": "Purchases of Stock-in-Trade",
    "change in inventories": "Changes in Inventories", "changes in inventories": "Changes in Inventories",
    "changes in inventories of finished goods, work-in-progress and stock-in-trade": "Changes in Inventories",
    "employee benefits expense": "Employee Benefits Expense", "employee benefit expense": "Employee Benefits Expense",
    "finance costs": "Finance Costs", "depreciation and amortisation expense": "Depreciation & Amortisation",
    "depreciation and amortization expense": "Depreciation & Amortisation", "other expenses": "Other Expenses",
    "tax expenses": "Tax Expense", "tax expense": "Tax Expense", "exceptional items": "Other Expenses",
}


def map_head(text):
    """Map a head written in someone's workbook to this app's vocabulary (blank if not recognised)."""
    t = str(text).strip()
    if not t or t.lower() == "none":
        return ""
    low = t.lower()
    if low in PRIMARY_MAP:
        return PRIMARY_MAP[low]
    if low.startswith("user_defined_expense"):
        return "Other Expenses"
    k = norm_name(t)
    if k in {norm_name(v) for v in CAP_NAME.values()} or k in ("capital", "sharecapital", "equitysharecapital"):
        return "Capital"
    if k in ("openingstock", "stockintrade", "closingstock"):
        return "Inventories"
    for h in HEAD_LIST:
        if norm_name(h) == k:
            return h
    for key, v in PRIMARY_MAP.items():
        if norm_name(key) == k:
            return v
    return ""


def _grid_to_tb(grid):
    hdr = None
    for i in range(min(25, len(grid))):
        cells = [str(c).strip().lower() for c in grid.iloc[i]]
        if any(c in LED_COLS for c in cells) and any(("debit" in c or c in ("dr", "closing balance", "balance", "amount")) for c in cells):
            hdr = i
            break
    if hdr is None:
        raise ValueError("Could not find a header row with a ledger column and Debit/Credit (or Closing Balance) columns. "
                         "First row seen: %s" % [str(c) for c in grid.iloc[0]][:8])
    cols = [str(c).strip().lower() for c in grid.iloc[hdr]]
    led_i = next(i for i, c in enumerate(cols) if c in LED_COLS)
    grp_i = next((i for i, c in enumerate(cols) if c in ("group", "parent", "under", "tally group", "primary group")), None)
    head_i = next((i for i, c in enumerate(cols) if c in HEAD_COLS), None)
    deb = [i for i, c in enumerate(cols) if "debit" in c]
    cre = [i for i, c in enumerate(cols) if "credit" in c]
    clo_d = [i for i in deb if "clos" in cols[i]]
    clo_c = [i for i in cre if "clos" in cols[i]]
    di = (clo_d or deb[-1:] or [None])[0]
    ci = (clo_c or cre[-1:] or [None])[0]
    bal_i = next((i for i, c in enumerate(cols) if c in ("closing balance", "balance", "amount")), None)
    if di is None and ci is None and bal_i is None:
        raise ValueError("No Debit/Credit or Closing Balance column. Columns seen: %s" % cols[:10])
    recs, signed = [], False
    for i in range(hdr + 1, len(grid)):
        row = grid.iloc[i]
        led = str(row.iloc[led_i]).strip()
        if not led or led.lower() in ("grand total", "total", "difference in opening balances", "diff in opening balances"):
            continue
        if di is not None or ci is not None:
            d = parse_amount(str(row.iloc[di]))[0] if di is not None else 0.0
            c = parse_amount(str(row.iloc[ci]))[0] if ci is not None else 0.0
            val = d - c
        else:
            amt, tag = parse_amount(str(row.iloc[bal_i]))
            val = amt if tag == "Dr" else -amt if tag == "Cr" else amt
            signed = signed or tag is None
        if abs(val) < 0.005:
            continue
        recs.append({"Ledger": led, "Group": str(row.iloc[grp_i]).strip() if grp_i is not None else "",
                     "Head": map_head(row.iloc[head_i]) if head_i is not None else "", "Dr_pos": val, "_raw": str(row.iloc[head_i]) if head_i is not None else ""})
    if not recs:
        raise ValueError("No ledger balances found under the header row. Columns seen: %s" % cols[:10])
    out = pd.DataFrame(recs)
    note = "Read %d ledgers (ledger column '%s'%s%s)." % (len(out), cols[led_i], ", group column found" if grp_i is not None else ", no group column",
                                                         ", head column '%s'" % cols[head_i] if head_i is not None else "")
    return out, note, signed, head_i is not None


def parse_tb_upload(raw, filename):
    """Read a trial balance exported from Tally (Excel or CSV). Returns (df[Ledger, Group, Head, Dr_pos], note)."""
    import io as _io
    try:
        if filename.lower().endswith(".csv"):
            grid = pd.read_csv(_io.BytesIO(raw), header=None, dtype=str, keep_default_na=False)
        else:
            grid = pd.read_excel(_io.BytesIO(raw), header=None, dtype=str).fillna("")
    except Exception as e:
        raise ValueError("Could not open the file (%s). Use .xlsx or .csv." % e)
    out, note, _s, _h = _grid_to_tb(grid)
    return out.drop(columns="_raw"), note


def _parse_vcfo_tb(wb, period_end):
    """The firm's VCFO workbook: a TB sheet with GL code / GL Name, three years of Books-Adjustments-Final columns
    under a row of year-end dates, and Type / Primary Group / Secondary Group columns. Debits are positive."""
    ws = wb["TB"]
    hdr = None
    for r in range(1, min(15, ws.max_row) + 1):
        vals = [str(ws.cell(row=r, column=c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
        if "gl name" in vals and "primary group" in vals:
            hdr = r
            break
    if hdr is None:
        return None
    heads = [str(ws.cell(row=hdr, column=c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
    name_c, prim_c, sec_c = heads.index("gl name") + 1, heads.index("primary group") + 1, (heads.index("secondary group") + 1 if "secondary group" in heads else None)
    dates = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=hdr - 1, column=c).value
        if hasattr(v, "date"):
            dates[c] = v.date()
    if not dates:
        raise ValueError("The TB sheet has no year-end dates above the columns, so the previous year cannot be located.")
    want = period_end
    col = None
    if want is not None:
        for c, d in dates.items():
            if d == want:
                col = c
                break
    if col is None:
        raise ValueError("The TB sheet has year-ends %s but the previous year-end needed is %s. Check the 'Financial year ends on' date."
                         % (", ".join(sorted(d.strftime("%d-%b-%Y") for d in dates.values())), want.strftime("%d-%b-%Y") if want else "n/a"))
    fin = next((c for c in range(col, col + 3) if heads[c - 1] == "final"), col + 2)
    recs = []
    for r in range(hdr + 1, ws.max_row + 1):
        led = ws.cell(row=r, column=name_c).value
        v = ws.cell(row=r, column=fin).value
        if not led or not isinstance(v, (int, float)) or abs(v) < 0.005:
            continue
        raw_head = str(ws.cell(row=r, column=prim_c).value or "")
        recs.append({"Ledger": str(led).strip(), "Group": str(ws.cell(row=r, column=sec_c).value or "") if sec_c else "",
                     "Head": map_head(raw_head), "Dr_pos": float(v), "_raw": raw_head})
    if not recs:
        raise ValueError("No balances found in the TB sheet for %s." % want)
    out = pd.DataFrame(recs)
    tot = out["Dr_pos"].sum()
    unk = sorted(set(out.loc[(out["Head"] == "") & (out["_raw"].str.strip().str.lower() != "none"), "_raw"]))
    msg = ("TB sheet, year ended %s, 'Final' column: %d ledgers, heads from 'Primary Group'. Trial balance difference %s." %
           (want.strftime("%d-%b-%Y"), len(out), inr(tot)))
    if unk:
        msg += " Not recognised (pick them in the Previous year tab): %s." % ", ".join(unk[:6])
    return out.drop(columns="_raw"), None, 0.0, msg


def parse_workbook(raw, period_end=None):
    """Read the Trial Balance Mapping sheet of a workbook produced by this app (last year's). Returns
    (df[Ledger, Group, Head, Dr_pos], closing_stock_or_None, tax, note)."""
    import io as _io
    from openpyxl import load_workbook
    wb = load_workbook(_io.BytesIO(raw), data_only=True)
    ws = wb["Trial Balance Mapping"] if "Trial Balance Mapping" in wb.sheetnames else None
    if ws is None:
        if "TB" in wb.sheetnames:
            got = _parse_vcfo_tb(wb, period_end)
            if got is not None:
                return got
        return _parse_other_workbook(raw, wb.sheetnames)
    recs, closing, tax = [], None, 0.0
    for r in range(6, ws.max_row + 1):
        a, b, c, e, f = (ws.cell(row=r, column=k).value for k in (1, 2, 3, 5, 6))
        if not a:
            continue
        t = str(a)
        if t.startswith("Opening reserves"):
            continue
        if t.startswith("Closing stock"):
            closing = -float(e or 0)
            continue
        if t.startswith("Income-tax provision"):
            tax = float(f or 0)
            continue
        if t.startswith(("Check", "Memo")):
            continue
        if not isinstance(e, (int, float)) or not c:
            continue
        head = "Capital" if c in CAP_NAME.values() else "Inventories" if c == "Opening Stock" else ("" if c == "Unclassified" else c)
        recs.append({"Ledger": t, "Group": b or "", "Head": head, "Dr_pos": -float(e)})
    if not recs:
        raise ValueError("No ledger rows found in the Trial Balance Mapping sheet.")
    return pd.DataFrame(recs), closing, tax, "Read %d ledgers with their heads from last year's workbook." % len(recs)


# ============================================================= depreciation ==
SCH2_DEFAULTS = [  # Schedule II, Part C - useful lives (years). Confirm each asset before relying on them.
    ("Buildings - RCC structure", 60), ("Buildings - other than RCC (incl. factory)", 30),
    ("Plant and machinery - general", 15), ("Furniture and fittings", 10), ("Office equipment", 5),
    ("Electrical installations and equipment", 10), ("Motor vehicles - cars", 8),
    ("Computers - end-user devices (desktops, laptops)", 3), ("Computers - servers and networks", 6),
    ("Land (not depreciated)", 0),
]
IT_DEFAULTS = [  # Income-tax depreciation rates, block of assets (WDV). Confirm against the current rate table.
    ("Buildings - residential", 5), ("Buildings - other (office, factory)", 10), ("Furniture and fittings", 10),
    ("Plant and machinery - general", 15), ("Motor cars (not used on hire)", 15), ("Computers and software", 40),
    ("Intangible assets (know-how, patents, licences)", 25), ("Land (not depreciated)", 0),
]
S2_COLS = ["Tally ledger", "Asset class", "Useful life (years)", "Remaining life override (years)",
           "Opening gross block", "Opening accumulated depreciation", "Additions (cost)", "Date put to use (additions)",
           "Disposals (original cost)", "Accumulated depreciation on disposals", "Sale proceeds", "Date of sale"]
S2_NUM = ["Useful life (years)", "Remaining life override (years)", "Opening gross block", "Opening accumulated depreciation",
          "Additions (cost)", "Disposals (original cost)", "Accumulated depreciation on disposals", "Sale proceeds"]
S2_DATES = ["Date put to use (additions)", "Date of sale"]
S2_OUT = ["Tally ledger", "Asset class", "Depreciation for the year", "Closing gross block",
          "Closing accumulated depreciation", "Closing net block", "Profit / (loss) on disposal"]
IT_COLS = ["Tally ledger", "Block of assets", "Rate (%)", "Opening WDV", "Additions (used 180 days or more)",
           "Additions (used under 180 days)", "Sale proceeds / deletions"]
IT_OUT = ["Tally ledger", "Block of assets", "Depreciation for the year", "Closing WDV", "Short-term capital gain (block ceases)"]


def guess_class(name, table):
    n = name.lower()
    pairs = [("land", "Land"), ("building", "Buildings"), ("factory", "Buildings"), ("shed", "Buildings"),
             ("furniture", "Furniture"), ("computer", "Computers"), ("laptop", "Computers"), ("server", "Computers - servers"),
             ("vehicle", "Motor"), ("car", "Motor"), ("truck", "Motor"), ("electric", "Electrical"),
             ("office", "Office"), ("machine", "Plant"), ("plant", "Plant")]
    for key, label in pairs:
        if key in n:
            for cls, val in table:
                if cls.startswith(label) or label.lower() in cls.lower():
                    return cls, val
    for cls, val in table:
        if cls.startswith("Plant"):
            return cls, val
    return table[0]


def prefill_dep_tables(df, py=None):
    """Prefill from the fixed-asset ledgers. With a previous year, opening values come from last year's closing balance
    of the same ledger; otherwise they are the current closing balances (which the user must edit)."""
    fa = df[(df["Head"] == "Property, Plant & Equipment") & (df["Dr_pos"].abs() > 0.005)]
    pymap = {}
    if py is not None:
        pymap = {norm_name(l): v for l, v, h in zip(py["Ledger"], py["Dr_pos"], py["Head"]) if h == "Property, Plant & Equipment"}
    s2, it = [], []
    for _, r in fa.iterrows():
        c2, life = guess_class(r["Ledger"], SCH2_DEFAULTS)
        ci, rate = guess_class(r["Ledger"], IT_DEFAULTS)
        v = float(pymap.get(norm_name(r["Ledger"]), 0.0)) if pymap else float(r["Dr_pos"])
        gb, ad = (v, 0.0) if v >= 0 else (0.0, -v)
        s2.append([r["Ledger"], c2, life, 0.0, round(gb, 2), round(ad, 2), 0.0, None, 0.0, 0.0, 0.0, None])
        it.append([r["Ledger"], ci, rate, round(v, 2), 0.0, 0.0, 0.0])
    if not s2:
        s2 = [["", SCH2_DEFAULTS[2][0], SCH2_DEFAULTS[2][1], 0.0, 0.0, 0.0, 0.0, None, 0.0, 0.0, 0.0, None]]
        it = [["", IT_DEFAULTS[3][0], IT_DEFAULTS[3][1], 0.0, 0.0, 0.0, 0.0]]
    return pd.DataFrame(s2, columns=S2_COLS), pd.DataFrame(it, columns=IT_COLS)


def _num(df, cols):
    out = df.copy()
    for c in cols:
        out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0.0)
    return out


def _date(v):
    try:
        if v is None or (isinstance(v, float) and v != v):
            return None
        t = pd.to_datetime(v)
        return None if pd.isna(t) else t.date()
    except Exception:
        return None


def days_in_year(start, end):
    return (end - start).days + 1


def days_added(v, start, end):
    d = _date(v)
    return 0 if d is None else max(0, min(days_in_year(start, end), (end - d).days + 1))


def days_sold(v, start, end):
    d = _date(v)
    return 0 if d is None else max(0, min(days_in_year(start, end), (d - start).days + 1))


def dep_sch2(df, method, residual_pct, start, end):
    """Schedule II depreciation for the year.
    SLM: (cost - residual) / life on the cost still held, never below residual; additions and disposals pro rata by days.
    WDV: rate = 1 - residual^(1/life) on the written-down value.
    Remaining-life override (SLM only): (written-down value - residual) / remaining life.
    Depreciation on assets sold is charged up to the date of sale and the asset's accumulated depreciation leaves the block."""
    d = _num(df, S2_NUM)
    dy = days_in_year(start, end)
    res = max(0.0, residual_pct) / 100.0
    out = []
    for i in range(len(d)):
        r = d.iloc[i]
        life, rem = r["Useful life (years)"], r["Remaining life override (years)"]
        gb, ad, add = r["Opening gross block"], r["Opening accumulated depreciation"], r["Additions (cost)"]
        dis, adis, proceeds = r["Disposals (original cost)"], r["Accumulated depreciation on disposals"], r["Sale proceeds"]
        da = days_added(df.iloc[i]["Date put to use (additions)"], start, end)
        dd = days_sold(df.iloc[i]["Date of sale"], start, end)
        g, a = gb - dis, ad - adis
        if life <= 0:
            d_rem = d_dis = d_add = 0.0
        elif method == "SLM":
            f = (1 - res) / life
            if rem > 0:
                amt = max(0.0, g - a - g * res)
                d_rem = min(amt / rem, amt)
            else:
                d_rem = min(g * f, max(0.0, g * (1 - res) - a))
            d_dis = min(dis * f * dd / dy, max(0.0, dis * (1 - res) - adis))
            d_add = add * f * da / dy
        else:
            rate = 1 - res ** (1.0 / life)
            d_rem, d_dis, d_add = rate * (g - a), rate * (dis - adis) * dd / dy, rate * add * da / dy
        dep = max(0.0, d_rem + d_dis + d_add)
        cg = gb + add - dis
        ca = ad - adis + dep - d_dis
        pl = (proceeds - ((dis - adis) - d_dis)) if dis > 0 else 0.0
        out.append([df.iloc[i]["Tally ledger"], df.iloc[i]["Asset class"], round(dep, 2), round(cg, 2), round(ca, 2),
                    round(cg - ca, 2), round(pl, 2)])
    return pd.DataFrame(out, columns=S2_OUT)


def dep_it(df):
    """Income-tax depreciation on blocks of assets (WDV). Half rate on additions used under 180 days.
    If sale proceeds exceed the block, the block ceases and the excess is a short-term capital gain (s.50)."""
    d = _num(df, IT_COLS[2:])
    out = []
    for i in range(len(d)):
        r = d.iloc[i]
        rate, op, a1, a2, sale = r["Rate (%)"], r["Opening WDV"], r["Additions (used 180 days or more)"], \
            r["Additions (used under 180 days)"], r["Sale proceeds / deletions"]
        total = op + a1 + a2 - sale
        if total <= 0:
            dep, closing, stcg = 0.0, 0.0, max(0.0, -total)
        else:
            dep = min(max(0.0, op + a1 - sale) * rate / 100.0 + a2 * rate / 200.0, total)
            closing, stcg = total - dep, 0.0
        out.append([df.iloc[i]["Tally ledger"], df.iloc[i]["Block of assets"], round(dep, 2), round(closing, 2), round(stcg, 2)])
    return pd.DataFrame(out, columns=IT_OUT)


def _parse_other_workbook(raw, sheetnames):
    """A workbook not made by this app (for example last year's financials prepared elsewhere): look through every sheet for a
    ledger list that has a head / classification column, and use the sheet whose heads we recognise best."""
    import io as _io
    best, tried = None, []
    for name in sheetnames:
        try:
            grid = pd.read_excel(_io.BytesIO(raw), sheet_name=name, header=None, dtype=str).fillna("")
            out, note, signed, has_head = _grid_to_tb(grid)
        except Exception as e:
            tried.append("%s: %s" % (name, str(e)[:60]))
            continue
        score = (out["Head"] != "").sum() if has_head else 0
        if best is None or score > best[0]:
            best = (score, name, out, note, signed, has_head)
    if best is None:
        raise ValueError("No ledger list found. Sheets looked at: %s. This option needs either a workbook produced by this application "
                         "or a sheet with a Ledger column, a balance (Debit/Credit or Closing Balance) and a head column. "
                         "Otherwise use 'Upload trial balance'." % "; ".join(tried)[:400])
    score, name, out, note, signed, has_head = best
    if signed:   # one signed column with no Dr/Cr marker: Tally shows debits as negative; flip if assets come out as credits
        out["Dr_pos"] = -out["Dr_pos"]
        a = out.loc[out["Head"].isin({"Trade Receivables", "Cash & Cash Equivalents", "Property, Plant & Equipment", "Inventories"}), "Dr_pos"].sum()
        if a < 0:
            out["Dr_pos"] = -out["Dr_pos"]
    if not has_head:
        return out.drop(columns="_raw"), None, 0.0, "Sheet '%s': %s No head column was found, so heads are taken from the current year and groups." % (name, note)
    unk = out.loc[(out["Head"] == "") & (out["_raw"].str.strip() != ""), "_raw"].unique()
    msg = "Sheet '%s': %s %d of %d heads recognised." % (name, note, score, len(out))
    if len(unk):
        msg += " Not recognised (pick them in the Previous year tab): %s." % ", ".join(list(unk)[:6])
    return out.drop(columns="_raw"), None, 0.0, msg


# ============================================================ mapping & rules ==
ADJ_EXP = "Depreciation per calculator - difference to books (expense)"
ADJ_FA = "Depreciation per calculator - difference to books (fixed assets)"
SGN = lambda head: 1 if head in CR_HEADS else -1


def export_head(head, entity, has_cs):
    if head == "":
        return "Unclassified"
    if head == "Capital":
        return CAP_NAME[entity]
    if head == "Inventories" and has_cs:
        return "Opening Stock"
    return head


RECLASS = {("Trade Receivables", -1): "Other Current Liabilities", ("Trade Payables", 1): "Short-term Loans and Advances",
           ("Other Current Liabilities (Statutory Dues)", 1): "Other Current Assets", ("Other Current Liabilities", 1): "Other Current Assets",
           ("Short-term Loans and Advances", -1): "Other Current Liabilities", ("Other Current Assets", -1): "Other Current Liabilities",
           ("Cash & Cash Equivalents", -1): "Short-term Borrowings"}


def build_mapping(df, entity, cs_items=None, accept=False, adj=0.0, force_closing=None, reclass=False, force_note=None):
    """Turn the classified trial balance into mapping rows plus the derived rows
    (opening reserves, closing stock). Returns a dict used by both the screen and the Excel file."""
    plac = is_plac(df["Ledger"])
    plac_cr = float(-df.loc[plac, "Dr_pos"].sum()) if plac.any() else None
    led = df[~plac].copy()
    unclassified = int((led["Head"] == "").sum())
    closing, status, msg = None, "none", ""
    or0 = float(led["Dr_pos"].sum())
    p_l = -float(led.loc[led["Head"].isin(PL_HEADS), "Dr_pos"].sum())
    os_ = float(led.loc[led["Head"] == "Inventories", "Dr_pos"].sum())
    delta = None
    if force_closing is not None:
        status, closing = "given", round(float(force_closing), 2)
        msg = force_note or ("Closing stock %s taken from last year's workbook." % inr(closing))
    elif plac_cr is None:
        msg = ("No Profit & Loss A/c ledger was found. Statements are built from the ledger balances; "
               "reserves are the balancing figure.")
    elif unclassified:
        status = "unclassified"
        msg = "Classify every ledger first. Reserves and closing stock cannot be derived until then."
    elif abs(plac_cr - (or0 + p_l)) < 1:
        status = "ok"
        msg = ("Tally's Profit & Loss A/c ledger already includes this year's profit. Opening reserves are derived "
               "by excluding it, so profit is counted once.")
    elif abs(plac_cr - or0) < 1:
        status = "ok_open"
        msg = "Tally's Profit & Loss A/c ledger holds opening reserves only. This year's profit comes from the P&L ledgers."
    else:
        delta = plac_cr - or0 - p_l           # = closing stock less stock in ledgers
        verified = cs_items is not None and abs((cs_items - os_) - delta) <= max(5.0, 0.0005 * abs(cs_items))
        if verified:
            status, closing = "verified", round(float(cs_items), 2)
            msg = ("Closing stock %s taken from Tally's stock items; it agrees to the %s implied by the books."
                   % (inr(closing), inr(os_ + delta)))
        elif accept:
            status, closing = "accepted", round(os_ + delta, 2)
            msg = ("Closing stock %s is derived from the books and accepted by you. Compare it with the Stock Summary."
                   % inr(closing))
        else:
            status = "pending"
            msg = ("The books imply closing stock of %s against %s in the stock ledgers (difference %s). "
                   "Compare with Tally's Stock Summary, then tick 'Accept derived closing stock' in the sidebar."
                   % (inr(os_ + delta), inr(os_), inr(delta)))
    has_cs = closing is not None
    rows, rec_amt = [], 0.0
    for _, r in led.iterrows():
        eh, grp = export_head(r["Head"], entity, has_cs), r["Group"]
        if reclass and (eh, 1 if r["Dr_pos"] > 0 else -1) in RECLASS:
            eh, grp = RECLASS[(eh, 1 if r["Dr_pos"] > 0 else -1)], "%s [reclassified]" % grp
            rec_amt += abs(float(r["Dr_pos"]))
        rows.append((r["Ledger"], grp, eh, SGN(eh), round(-float(r["Dr_pos"]), 2)))
    if adj and abs(adj) >= 0.005:
        rows.append((ADJ_EXP, DEP_GROUP, "Depreciation & Amortisation", -1, round(-adj, 2)))
        rows.append((ADJ_FA, DEP_GROUP, "Property, Plant & Equipment", -1, round(adj, 2)))
    rows.sort(key=lambda x: (EXPORT_ORDER.index(x[2]) if x[2] in EXPORT_ORDER else 99, x[0].lower()))
    or_cr = round(-sum(x[4] for x in rows), 2)
    ht = {}
    for _, _, h, s, e in rows:
        ht[h] = ht.get(h, 0.0) + s * e
    ht["Reserves and Surplus"] = ht.get("Reserves and Surplus", 0.0) + or_cr
    if has_cs:
        ht["Inventories"] = ht.get("Inventories", 0.0) + closing
    return {"rows": rows, "or_cr": or_cr, "closing": closing, "status": status, "msg": msg, "plac_cr": plac_cr,
            "os": os_, "delta": delta, "adj": float(adj or 0.0), "ht": ht, "entity": entity,
            "unclassified": unclassified, "p_l": p_l, "reclass_amt": rec_amt}


# ================================================================ statements ==
def H(h, c=1): return (c, "h", h)
def R(k, c=1): return (c, "r", k)
def Pp(k, c=1): return (c, "p", k)
def X(k, c=1): return (c, "x", k)


def row(kind, key=None, label="", note=False, terms=None, indent=0):
    return {"kind": kind, "key": key, "label": label, "note": note, "terms": list(terms or []),
            "indent": indent, "no": None}


def _pl(framework, entity):
    ind = framework == FRAMEWORKS[1]
    nc = framework == FRAMEWORKS[2]
    if nc:
        pat_label = {"Firm (partnership or LLP)": "Net profit / (loss) transferred to partners' accounts (before apportionment)",
                     "Proprietorship": "Net profit / (loss) transferred to proprietor's capital account"}[entity]
        return [
            row("line", "rev", "Sales / revenue from operations", True, [H("Revenue from Operations")], 1),
            row("line", "oth", "Other income", True, [H("Other Income")], 1),
            row("total", "tinc", "Total Income", False, [R("rev"), R("oth")]),
            row("blank"),
            row("head", None, "Expenses"),
            row("line", "pur", "Purchases of stock-in-trade", True, [H("Purchases of Stock-in-Trade")], 1),
            row("line", "chg", "Changes in stock-in-trade (opening less closing)", True,
                [H("Opening Stock"), X("closing", -1), H("Changes in Inventories")], 1),
            row("line", "emp", "Salaries, wages and staff costs", True, [H("Employee Benefits Expense")], 1),
            row("line", "fin", "Interest and finance charges", True, [H("Finance Costs")], 1),
            row("line", "dep", "Depreciation", True, [H("Depreciation & Amortisation")], 1),
            row("line", "oex", "Other expenses", True, [H("Other Expenses")], 1),
            row("total", "texp", "Total Expenses", False, [R(k) for k in ("pur", "chg", "emp", "fin", "dep", "oex")]),
            row("blank"),
            row("profit", "pbt", "Net profit / (loss) before tax", False, [R("tinc"), R("texp", -1)]),
            row("line", "tax", "Provision for income tax", True, [H("Tax Expense"), X("tax")], 1),
            row("grand", "pat", pat_label, False, [R("pbt"), R("tax", -1)]),
        ]
    rows = [
        row("line", "rev", "I. Revenue from operations", True, [H("Revenue from Operations")], 1),
        row("line", "oth", "II. Other income", True, [H("Other Income")], 1),
        row("total", "tinc", "III. Total Income (I + II)", False, [R("rev"), R("oth")]),
        row("blank"),
        row("head", None, "IV. Expenses"),
        row("line", "pur", "Purchases of stock-in-trade", True, [H("Purchases of Stock-in-Trade")], 1),
        row("line", "chg", "Changes in inventories of stock-in-trade", True, [H("Opening Stock"), X("closing", -1), H("Changes in Inventories")], 1),
        row("line", "emp", "Employee benefits expense", True, [H("Employee Benefits Expense")], 1),
        row("line", "fin", "Finance costs", True, [H("Finance Costs")], 1),
        row("line", "dep", "Depreciation and amortisation expense", True, [H("Depreciation & Amortisation")], 1),
        row("line", "oex", "Other expenses", True, [H("Other Expenses")], 1),
        row("total", "texp", "Total Expenses (IV)", False, [R(k) for k in ("pur", "chg", "emp", "fin", "dep", "oex")]),
        row("blank"),
        row("profit", "pbe", "V. Profit / (Loss) before exceptional items and tax (III - IV)", False, [R("tinc"), R("texp", -1)]),
        row("line", "exc", "VI. Exceptional items", False, [], 1),
        row("profit", "pbt", "VII. Profit / (Loss) before tax (V - VI)", False, [R("pbe"), R("exc", -1)]),
        row("line", "tax", "VIII. Tax expense", True, [H("Tax Expense"), X("tax")], 1),
        row("grand", "pat", "IX. Profit / (Loss) for the year (VII - VIII)", False, [R("pbt"), R("tax", -1)]),
    ]
    if ind:
        rows += [row("line", "oci", "X. Other comprehensive income (not available from Tally; nil)", False, [], 1),
                 row("grand", "tci", "XI. Total comprehensive income for the year (IX + X)", False, [R("pat"), R("oci")])]
    rows += [row("blank"), row("text", "eps", "Earnings per share (Rs) - Basic / Diluted", True)]
    return rows


def _bs(framework, entity):
    cap = CAP_NAME[entity]
    ind = framework == FRAMEWORKS[1]
    nc = framework == FRAMEWORKS[2]
    if nc:
        funds = "Partners' funds" if entity.startswith("Firm") else "Proprietor's funds"
        capl = "Partners' capital accounts" if entity.startswith("Firm") else "Proprietor's capital account"
        return [
            row("section", None, "I. CAPITAL AND LIABILITIES"),
            row("sub", None, "(1) " + funds),
            row("line", "cap", "(a) " + capl, True, [H(cap)], 2),
            row("line", "res", "(b) Reserves and surplus (incl. current-year profit)", True,
                [H("Reserves and Surplus"), Pp("pat")], 2),
            row("total", "fund", "Total " + funds, False, [R("cap"), R("res")]),
            row("sub", None, "(2) Loan funds"),
            row("line", "ltb", "(a) Term loans and other long-term borrowings", True, [H("Long-term Borrowings")], 2),
            row("line", "stb", "(b) Bank overdraft, cash credit and short-term loans", True, [H("Short-term Borrowings")], 2),
            row("total", "loans", "Total Loan funds", False, [R("ltb"), R("stb")]),
            row("sub", None, "(3) Current liabilities and provisions"),
            row("line", "tp", "(a) Sundry creditors", True, [H("Trade Payables")], 2),
            row("line", "ocl", "(b) Statutory dues and other current liabilities", True,
                [H("Other Current Liabilities (Statutory Dues)"), H("Other Current Liabilities")], 2),
            row("line", "prov", "(c) Provisions", True, [H("Short-term Provisions"), X("tax")], 2),
            row("line", "oncl", "(d) Other long-term liabilities and provisions", True, [H("Other Non-current Liabilities")], 2),
            row("total", "cl", "Total Current liabilities and provisions", False, [R("tp"), R("ocl"), R("prov"), R("oncl")]),
            row("grand", "tel", "TOTAL CAPITAL AND LIABILITIES", False, [R("fund"), R("loans"), R("cl")]),
            row("blank"),
            row("section", None, "II. PROPERTY AND ASSETS"),
            row("line", "ppe", "(1) Fixed assets", True, [H("Property, Plant & Equipment")], 1),
            row("line", "inv", "(2) Investments", True, [H("Investments")], 1),
            row("line", "onca", "(3) Other non-current assets", True, [H("Other Non-Current Assets")], 1),
            row("sub", None, "(4) Current assets, loans and advances"),
            row("line", "stock", "(a) Stock-in-trade", True, [H("Inventories")], 2),
            row("line", "tr", "(b) Sundry debtors", True, [H("Trade Receivables")], 2),
            row("line", "cash", "(c) Cash and bank balances", True, [H("Cash & Cash Equivalents")], 2),
            row("line", "sla", "(d) Loans and advances", True, [H("Short-term Loans and Advances")], 2),
            row("line", "oca", "(e) Other current assets", True, [H("Other Current Assets")], 2),
            row("total", "ca", "Total Current assets, loans and advances", False,
                [R(k) for k in ("stock", "tr", "cash", "sla", "oca")]),
            row("grand", "tas", "TOTAL PROPERTY AND ASSETS", False, [R("ppe"), R("inv"), R("onca"), R("ca")]),
        ]
    if ind:
        return [
            row("section", None, "I. ASSETS"),
            row("sub", None, "(1) Non-current assets"),
            row("line", "ppe", "(a) Property, plant and equipment", True, [H("Property, Plant & Equipment")], 2),
            row("line", "inv", "(b) Financial assets - Investments", True, [H("Investments")], 2),
            row("line", "onca", "(c) Other non-current assets", True, [H("Other Non-Current Assets")], 2),
            row("total", "nca", "Total Non-current assets", False, [R("ppe"), R("inv"), R("onca")]),
            row("sub", None, "(2) Current assets"),
            row("line", "stock", "(a) Inventories", True, [H("Inventories")], 2),
            row("line", "tr", "(b) Financial assets - Trade receivables", True, [H("Trade Receivables")], 2),
            row("line", "cash", "(c) Financial assets - Cash and cash equivalents", True, [H("Cash & Cash Equivalents")], 2),
            row("line", "sla", "(d) Financial assets - Loans", True, [H("Short-term Loans and Advances")], 2),
            row("line", "oca", "(e) Other current assets", True, [H("Other Current Assets")], 2),
            row("total", "ca", "Total Current assets", False, [R(k) for k in ("stock", "tr", "cash", "sla", "oca")]),
            row("grand", "tas", "TOTAL ASSETS", False, [R("nca"), R("ca")]),
            row("blank"),
            row("section", None, "II. EQUITY AND LIABILITIES"),
            row("sub", None, "(1) Equity"),
            row("line", "cap", "(a) Equity share capital", True, [H(cap)], 2),
            row("line", "res", "(b) Other equity (incl. total comprehensive income)", True,
                [H("Reserves and Surplus"), Pp("tci")], 2),
            row("total", "fund", "Total Equity", False, [R("cap"), R("res")]),
            row("sub", None, "(2) Non-current liabilities"),
            row("line", "ltb", "(a) Financial liabilities - Borrowings", True, [H("Long-term Borrowings")], 2),
            row("line", "oncl", "(b) Other non-current liabilities and provisions", True, [H("Other Non-current Liabilities")], 2),
            row("total", "ncl", "Total Non-current liabilities", False, [R("ltb"), R("oncl")]),
            row("sub", None, "(3) Current liabilities"),
            row("line", "stb", "(a) Financial liabilities - Borrowings", True, [H("Short-term Borrowings")], 2),
            row("line", "tp", "(b) Financial liabilities - Trade payables", True, [H("Trade Payables")], 2),
            row("line", "ocl", "(c) Other current liabilities", True,
                [H("Other Current Liabilities (Statutory Dues)"), H("Other Current Liabilities")], 2),
            row("line", "prov", "(d) Provisions", True, [H("Short-term Provisions"), X("tax")], 2),
            row("total", "cl", "Total Current liabilities", False, [R(k) for k in ("stb", "tp", "ocl", "prov")]),
            row("grand", "tel", "TOTAL EQUITY AND LIABILITIES", False, [R("fund"), R("ncl"), R("cl")]),
        ]
    return [
        row("section", None, "I. EQUITY AND LIABILITIES"),
        row("sub", None, "(1) Shareholders' Funds"),
        row("line", "cap", "(a) Share capital", True, [H(cap)], 2),
        row("line", "res", "(b) Reserves and surplus", True, [H("Reserves and Surplus"), Pp("pat")], 2),
        row("total", "fund", "Total Shareholders' Funds", False, [R("cap"), R("res")]),
        row("sub", None, "(2) Non-current liabilities"),
        row("line", "ltb", "(a) Long-term borrowings", True, [H("Long-term Borrowings")], 2),
        row("line", "oncl", "(b) Other long-term liabilities and provisions", True, [H("Other Non-current Liabilities")], 2),
        row("total", "ncl", "Total Non-current Liabilities", False, [R("ltb"), R("oncl")]),
        row("sub", None, "(3) Current Liabilities"),
        row("line", "stb", "(a) Short-term borrowings", True, [H("Short-term Borrowings")], 2),
        row("line", "tp", "(b) Trade payables", True, [H("Trade Payables")], 2),
        row("line", "ocl", "(c) Other current liabilities", True,
            [H("Other Current Liabilities (Statutory Dues)"), H("Other Current Liabilities")], 2),
        row("line", "prov", "(d) Short-term provisions", True, [H("Short-term Provisions"), X("tax")], 2),
        row("total", "cl", "Total Current Liabilities", False, [R(k) for k in ("stb", "tp", "ocl", "prov")]),
        row("grand", "tel", "TOTAL EQUITY AND LIABILITIES", False, [R("fund"), R("ncl"), R("cl")]),
        row("blank"),
        row("section", None, "II. ASSETS"),
        row("sub", None, "(1) Non-current assets"),
        row("line", "ppe", "(a) Property, plant and equipment", True, [H("Property, Plant & Equipment")], 2),
        row("line", "inv", "(b) Investments", True, [H("Investments")], 2),
        row("line", "onca", "(c) Other non-current assets", True, [H("Other Non-Current Assets")], 2),
        row("total", "nca", "Total Non-current Assets", False, [R("ppe"), R("inv"), R("onca")]),
        row("sub", None, "(2) Current assets"),
        row("line", "stock", "(a) Inventories", True, [H("Inventories")], 2),
        row("line", "tr", "(b) Trade receivables", True, [H("Trade Receivables")], 2),
        row("line", "cash", "(c) Cash and cash equivalents", True, [H("Cash & Cash Equivalents")], 2),
        row("line", "sla", "(d) Short-term loans and advances", True, [H("Short-term Loans and Advances")], 2),
        row("line", "oca", "(e) Other current assets", True, [H("Other Current Assets")], 2),
        row("total", "ca", "Total Current Assets", False, [R(k) for k in ("stock", "tr", "cash", "sla", "oca")]),
        row("grand", "tas", "TOTAL ASSETS", False, [R("nca"), R("ca")]),
    ]


def make_specs(framework, entity):
    pl, bs = _pl(framework, entity), _bs(framework, entity)
    n = 0
    for r in bs + pl:                      # notes are numbered in order: Balance Sheet first, then Profit and Loss
        if r["note"]:
            n += 1
            r["no"] = n
    return pl, bs


def evaluate(rows, ht, ext, other=None):
    vals = {}
    for r in rows:
        if r["key"] is None or r["kind"] in ("section", "sub", "blank", "head", "text"):
            continue
        t = 0.0
        for c, k, ref in r["terms"]:
            v = ht.get(ref, 0.0) if k == "h" else vals[ref] if k == "r" else other[ref] if k == "p" else ext[ref]
            t += c * v
        vals[r["key"]] = t
    return vals


def build_statements(M, framework, tax=0.0):
    pl, bs = make_specs(framework, M["entity"])
    ext = {"closing": M["closing"] or 0.0, "tax": float(tax or 0.0)}
    plv = evaluate(pl, M["ht"], ext)
    bsv = evaluate(bs, M["ht"], ext, plv)
    return {"pl": pl, "bs": bs, "plv": plv, "bsv": bsv, "ext": ext, "framework": framework,
            "diff": bsv["tas"] - bsv["tel"]}


def fx(terms, sif, rc, pc, ex):
    parts = []
    for c, k, ref in terms:
        s = sif(ref) if k == "h" else rc(ref) if k == "r" else pc(ref) if k == "p" else ex[ref]
        parts.append(("+" if c > 0 else "-") + s)
    if not parts:
        return "=0"
    f = "".join(parts)
    return "=" + (f[1:] if f[0] == "+" else f)


# ============================================================ HTML rendering ==
CSS = """
<style>
.block-container{padding-top:1.4rem;max-width:1250px}
.app-title{font-size:1.65rem;font-weight:700;color:#1F3864;margin:0}
@media (prefers-color-scheme: dark){.app-title{color:#9DB7E8}}
.app-sub{opacity:.75;margin:0 0 .6rem 0;font-size:.92rem}
.fs{width:100%;border-collapse:collapse;font-size:.9rem}
.fs th{background:#1F3864;color:#fff;padding:7px 10px;text-align:right;font-weight:600}
.fs th:first-child{text-align:left}
.fs td{padding:4px 10px;border-bottom:1px solid rgba(128,128,128,.18)}
.fs td.n{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
.fs td.k{text-align:center;opacity:.7;width:48px}
.fs tr.section td{font-weight:700;color:#1F3864;padding-top:10px}
@media (prefers-color-scheme: dark){.fs tr.section td{color:#9DB7E8}}
.fs tr.sub td{font-weight:600;font-style:italic}
.fs tr.head td{font-weight:700;padding-top:8px}
.fs tr.total td{font-weight:700;border-top:1px solid rgba(128,128,128,.6)}
.fs tr.profit td{font-weight:700}
.fs tr.grand td{font-weight:700;background:rgba(128,128,128,.14);border-top:1px solid rgba(128,128,128,.7);border-bottom:3px double rgba(128,128,128,.7)}
.pill{display:inline-block;padding:2px 10px;border-radius:12px;font-size:.78rem;font-weight:700;color:#fff;margin-right:8px;min-width:56px;text-align:center}
.pill.pass{background:#1E8449}.pill.fail{background:#C0392B}.pill.warn{background:#B9770E}.pill.info{background:#2E6DA4}
.chk{padding:7px 0;border-bottom:1px solid rgba(128,128,128,.18)}
.chk b{margin-right:6px}.chk span.d{opacity:.8}
.stmt-title{font-weight:700;font-size:1.05rem;margin:.2rem 0 .1rem 0}
.stmt-sub{opacity:.7;font-size:.8rem;margin-bottom:.4rem}
</style>
"""


def html_statement(rows, vals, heading, pvals=None):
    from html import escape as esc
    py = pvals is not None
    out = ['<div class="stmt-title">%s</div>' % esc(heading[0]), '<div class="stmt-sub">%s</div>' % esc(heading[1]),
           '<table class="fs"><tr><th>Particulars</th><th>Note</th><th>%s</th>%s</tr>' % (
               esc(heading[2]) if len(heading) > 2 else "Amount (Rs)", ("<th>%s</th>" % esc(heading[3])) if py else "")]
    span = 4 if py else 3
    for r in rows:
        k = r["kind"]
        if k == "blank":
            out.append('<tr><td colspan="%d" style="border:none">&nbsp;</td></tr>' % span)
            continue
        pad = "padding-left:%dpx" % (10 + 18 * r["indent"])
        note = "" if not r["no"] else str(r["no"])
        if k in ("section", "sub", "head"):
            out.append('<tr class="%s"><td colspan="%d" style="%s">%s</td></tr>' % (k, span, pad, esc(r["label"])))
            continue
        if k == "text":
            cells = '<td class="n">n.a.</td>' + ('<td class="n">n.a.</td>' if py else "")
        else:
            cells = '<td class="n">%s</td>' % inr(vals[r["key"]])
            if py:
                cells += '<td class="n">%s</td>' % inr(pvals[r["key"]])
        out.append('<tr class="%s"><td style="%s">%s</td><td class="k">%s</td>%s</tr>' % (k, pad, esc(r["label"]), note, cells))
    out.append("</table>")
    return "".join(out)


# =============================================================== excel export ==
NAVY = "1F3864"
NF0 = '#,##0;(#,##0);"-"'
NF2 = '#,##0.00;(#,##0.00);"-"'
PLS = "'Profit & Loss'"
DEPS = "'Depreciation Schedule'"

REMARK = {
    "pur": "", "emp": "", "fin": "", "oex": "", "rev": "", "oth": "",
    "dep": " Method and rates to be confirmed (see Depreciation Schedule sheet when the calculator was applied).",
    "ltb": " Repayment terms and security to be confirmed.",
    "stb": " Security, interest rate and sanctioned limit to be confirmed with the bank.",
    "tp": " Ageing not analysed (MSME disclosure pending).",
    "ocl": " Confirm against GST, TDS and other statutory returns.",
    "ppe": " Gross block and accumulated depreciation to be confirmed from the fixed-asset register.",
    "tr": " Ageing and provision for doubtful debts not analysed.",
    "cash": " Bank reconciliation and physical cash count to be confirmed.",
    "stock": " Valuation basis (cost / NRV) to be confirmed.",
}


def _clean(label):
    t = re.sub(r"^(\([a-z0-9ivx]+\)|[IVX]+\.)\s*", "", label).strip()
    t = re.sub(r"\s*\(.*?\)\s*$", "", t) if t.lower().startswith("profit") else t
    return t[:1].upper() + t[1:]


def _py(v):
    """Plain Python value for a cell (numpy numbers and empty dates cleaned)."""
    if v is None:
        return None
    if hasattr(v, "item"):
        try:
            v = v.item()
        except Exception:
            pass
    if isinstance(v, float) and v != v:
        return None
    if hasattr(v, "date") and not isinstance(v, date):
        return v.date()
    return v


def default_profile():
    return {"firm": "", "frn": "", "partner": "", "mno": "", "udin": "", "place": "", "rdate": "",
            "cin": "", "pan": "", "address": "", "bdate": "", "eq_shares": 0, "face_value": 10, "pref_div": 0,
            "signs": [{"name": "", "id": "", "desig": ""}, {"name": "", "id": "", "desig": ""}]}


def sig_lines(profile, entity, company):
    if not profile:
        return [], []
    ca = ["As per our report of even date attached"]
    if profile.get("firm"):
        ca += ["For " + profile["firm"], "Chartered Accountants"]
    if profile.get("frn"):
        ca.append("Firm Registration No.: " + profile["frn"])
    ca += ["", ""]
    if profile.get("partner"):
        ca += [profile["partner"], "Partner"]
    if profile.get("mno"):
        ca.append("Membership No.: " + profile["mno"])
    if profile.get("udin"):
        ca.append("UDIN: " + profile["udin"])
    if profile.get("place"):
        ca.append("Place: " + profile["place"])
    if profile.get("rdate"):
        ca.append("Date: " + profile["rdate"])
    who = {"Company": "the Board of Directors", "Firm (partnership or LLP)": "the Partners", "Proprietorship": "the Proprietor"}[entity]
    cl = ["For and on behalf of %s of %s" % (who, company)]
    for s in profile.get("signs", []):
        if s.get("name"):
            cl += ["", "", s["name"], (s.get("desig") or "") + (("  |  " + s["id"]) if s.get("id") else "")]
    if profile.get("bdate"):
        cl += ["", "Date of approval: " + profile["bdate"]]
    return ca, cl


def make_excel(M, framework, company, year_end, tax=0.0, dep=None, policy_dep="", M_py=None, tax_py=0.0, profile=None):
    entity = M["entity"]
    S = build_statements(M, framework, tax)
    pl, bs = S["pl"], S["bs"]
    has_py = M_py is not None
    start, prev_end = period_dates(year_end)
    d_fmt = lambda d: d.strftime("%d-%b-%Y")
    frame = FRAME_TEXT[framework]
    prow = {r["key"]: 6 + i for i, r in enumerate(pl) if r["key"]}
    brow = {r["key"]: 6 + i for i, r in enumerate(bs) if r["key"]}

    f_bold, f_head = Font(bold=True), Font(bold=True, color="FFFFFF")
    fill_head = PatternFill("solid", fgColor=NAVY)
    fill_tot = PatternFill("solid", fgColor="F2F2F2")
    thin, dbl = Side(style="thin"), Side(style="double")
    wb = Workbook()
    ws_bs = wb.active
    ws_bs.title = "Balance Sheet"
    ws_pl = wb.create_sheet("Profit & Loss")
    ws_pol = wb.create_sheet("Accounting Policies")
    ws_nb = wb.create_sheet("Notes - Balance Sheet")
    ws_fa = wb.create_sheet("Fixed Assets Schedule")
    ws_np = wb.create_sheet("Notes - Profit and Loss")
    ws_rt = wb.create_sheet("Notes - Ratios") if entity == "Company" else None
    MAPN, MAPPN = "Trial Balance Mapping", "Trial Balance Mapping PY"
    ws_mp = wb.create_sheet(MAPN)
    ws_mpp = wb.create_sheet(MAPPN) if has_py else None
    ws_dp = wb.create_sheet("Depreciation Schedule") if dep else None

    def setup(ws, widths, h2, h3, merge_to):
        ws.sheet_view.showGridLines = False
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        ws["A1"] = company
        ws["A1"].font = Font(size=14, bold=True)
        ws.merge_cells("A1:%s1" % merge_to)
        ws["A2"] = h2
        ws["A2"].font = Font(bold=True)
        if h3:
            ws["A3"] = h3
        ws.freeze_panes = "A6"

    def header(ws, labels):
        for i, t in enumerate(labels):
            c = ws.cell(row=5, column=i + 1, value=t)
            c.font, c.fill = f_head, fill_head
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[5].height = 32

    # ---------------------------------------------------------- mapping sheets
    def write_map(ws, Mx, taxx, sheet, pl_col, label_period, adj_sumifs=True):
        rows = Mx["rows"]
        has_cs = Mx["closing"] is not None
        first, last_led = 6, 5 + len(rows)
        r_or = last_led + 1
        r_cs = r_or + 1 if has_cs else None
        last = r_cs if has_cs else r_or
        r_tax = last + 2
        q = "'%s'" % sheet
        RF, RC, RB = ("%s!$%s$%d:$%s$%d" % (q, c, first, c, last) for c in "FCB")
        info = {"q": q, "RF": RF, "RC": RC, "RB": RB, "first": first, "last": last, "r_or": r_or, "r_cs": r_cs, "r_tax": r_tax,
                "rows": rows, "M": Mx, "has_cs": has_cs}
        info["sif"] = lambda h: 'SUMIFS(%s,%s,"%s")' % (RF, RC, h)
        info["adj"] = 'SUMIFS(%s,%s,"Depreciation & Amortisation",%s,"%s")' % (RF, RC, RB, DEP_GROUP)
        info["ex"] = {"closing": ("%s!$F$%d" % (q, r_cs)) if has_cs else "0", "tax": "%s!$F$%d" % (q, r_tax)}
        setup(ws, {"A": 52, "B": 22, "C": 42, "D": 7, "E": 22, "F": 22, "G": 9},
              "Ledger-to-statement mapping (%s) | Source: %s" % (label_period, "Tally Prime / uploaded trial balance"),
              "Sign: -1 for Dr-natured heads (assets, expenses), +1 for Cr-natured heads (capital, liabilities, income).", "D")
        ws["A3"].font = Font(size=9, italic=True)
        header(ws, ["Ledger", "Tally Group", "Statement Head", "Sign", "Closing Balance", "Presented Amount", "Dr / Cr"])

        def mrow(r, ledger, group, head, sign, bal):
            ws.cell(row=r, column=1, value=ledger)
            ws.cell(row=r, column=2, value=group)
            ws.cell(row=r, column=3, value=head)
            ws.cell(row=r, column=4, value=sign)
            e = ws.cell(row=r, column=5, value=bal)
            f = ws.cell(row=r, column=6, value="=E%d*D%d" % (r, r))
            ws.cell(row=r, column=7, value='=IF(E%d<0,"Dr","Cr")' % r)
            e.number_format = f.number_format = NF2

        for i, (led, grp, head, sign, bal) in enumerate(rows):
            mrow(first + i, led, grp, head, sign, bal)
        mrow(r_or, "Opening reserves (derived: balancing figure of the ledgers above)", "Profit & Loss A/c",
             "Reserves and Surplus", 1, "=-SUM(E%d:E%d)" % (first, last_led))
        if has_cs:
            mrow(r_cs, "Closing stock (from Stock Summary / derived; not a Tally ledger)", "Stock-in-trade", "Inventories", -1,
                 round(-Mx["closing"], 2))
        ws.cell(row=r_tax, column=1, value="Income-tax provision for the year (entered in the application; not computed)")
        t = ws.cell(row=r_tax, column=6, value=round(float(taxx or 0.0), 2))
        t.number_format = NF2
        r = r_tax + 2
        ws.cell(row=r, column=1, value="Check: sum of ledger closing balances above (excl. P&L A/c). Its negative is the opening reserves.")
        e = ws.cell(row=r, column=5, value="=SUM(E%d:E%d)" % (first, last_led))
        e.number_format, e.font = NF2, f_bold
        if Mx["plac_cr"] is not None:
            ws.cell(row=r + 1, column=1, value="Memo: Profit & Loss A/c ledger closing balance (credit; not in totals)")
            e2 = ws.cell(row=r + 1, column=5, value=round(Mx["plac_cr"], 2))
            e2.number_format = NF2
            ws.cell(row=r + 2, column=1, value="Check: P&L A/c ledger = opening reserves + profit before tax (+ calculator adjustment)")
            ws.cell(row=r + 2, column=6, value=(
                '=IF(ABS(E%d-(F%d+%s!%s%d-%s+%s))<1,"OK",IF(ABS(E%d-F%d)<1,"OK (opening only)","CHECK"))'
                % (r + 1, r_or, PLS, pl_col, prow["pbt"], info["sif"]("Tax Expense"), info["adj"], r + 1, r_or)))
        return info

    cy = write_map(ws_mp, M, tax, MAPN, "C", "%s to %s" % (d_fmt(start), d_fmt(year_end)))
    py = write_map(ws_mpp, M_py, tax_py, MAPPN, "D", "%s to %s" % (d_fmt(period_dates(prev_end)[0]), d_fmt(prev_end))) if has_py else None

    # ---------------------------------------------------------- statements
    cols4 = {"A": 66, "B": 8, "C": 22, "D": 22}

    def write_statement(ws, spec, own):
        for i, r in enumerate(spec):
            rr, k = 6 + i, r["kind"]
            if k == "blank":
                continue
            a = ws.cell(row=rr, column=1, value=r["label"])
            b = ws.cell(row=rr, column=2, value=r["no"])
            c = ws.cell(row=rr, column=3)
            d = ws.cell(row=rr, column=4)
            b.alignment = Alignment(horizontal="center")
            a.alignment = Alignment(indent=2 * r["indent"])
            c.number_format = d.number_format = NF0
            if k in ("section", "sub", "head"):
                a.font = Font(bold=True, color=NAVY) if k == "section" else Font(bold=True, italic=(k == "sub"))
                continue
            for cell, col, ctx in ((c, "C", cy), (d, "D", py)):
                if k == "text":
                    cell.value, cell.alignment = "n.a.", Alignment(horizontal="right")
                elif ctx is None:
                    cell.value, cell.alignment = "-", Alignment(horizontal="right")
                else:
                    cell.value = fx(r["terms"], ctx["sif"], lambda ref, col=col: "%s%d" % (col, own[ref]),
                                    lambda ref, col=col: "%s!%s%d" % (PLS, col, prow[ref]), ctx["ex"])
            if k in ("total", "grand", "profit"):
                a.font = c.font = d.font = f_bold
            if k in ("total", "grand"):
                c.border = Border(top=thin, bottom=dbl if k == "grand" else None)
                d.border = Border(top=thin, bottom=dbl if k == "grand" else None)
            if k == "grand":
                for x in (a, b, c, d):
                    x.fill = fill_tot

    ca_lines, cl_lines = sig_lines(profile, entity, company)

    def ident(ws):
        if profile and (profile.get("cin") or profile.get("pan") or profile.get("address")):
            bits = []
            if profile.get("cin"):
                bits.append("CIN / LLPIN: " + profile["cin"])
            if profile.get("pan"):
                bits.append("PAN: " + profile["pan"])
            if profile.get("address"):
                bits.append(profile["address"])
            ws["A4"] = " | ".join(bits)
            ws["A4"].font = Font(size=9, italic=True)

    def signature(ws, r0):
        for k in range(max(len(ca_lines), len(cl_lines))):
            if k < len(ca_lines):
                c = ws.cell(row=r0 + k, column=1, value=ca_lines[k] or None)
                c.font = Font(bold=(k in (0, 1)))
            if k < len(cl_lines):
                c = ws.cell(row=r0 + k, column=3, value=cl_lines[k] or None)
                c.font = Font(bold=(k == 0))

    py_hdr = lambda lead, d: "%s\n%s (Rs)" % (lead, d_fmt(d))
    setup(ws_pl, cols4, "Statement of Profit and Loss for the year ended %s" % long_date(year_end), frame, "D")
    ident(ws_pl)
    header(ws_pl, ["Particulars", "Note", py_hdr("Year ended", year_end), py_hdr("Year ended", prev_end)])
    write_statement(ws_pl, pl, prow)
    nr = 6 + len(pl) + 1
    signature(ws_pl, nr + 3)

    setup(ws_bs, cols4, "Balance Sheet as at %s" % long_date(year_end), frame, "D")
    ident(ws_bs)
    header(ws_bs, ["Particulars", "Note", py_hdr("As at", year_end), py_hdr("As at", prev_end)])
    write_statement(ws_bs, bs, brow)
    rc = 6 + len(bs) + 1
    ws_bs["F%d" % brow["tas"]] = "Balance check (not printed)"
    ws_bs["F%d" % brow["tas"]].font = Font(size=8, italic=True, color="888888")
    for col in ("C", "D") if has_py else ("C",):
        cc = ws_bs["%s%d" % ("G" if col == "C" else "H", brow["tas"])]
        cc.value = "=ROUND(%s%d-%s%d,2)" % (col, brow["tas"], col, brow["tel"])
        cc.number_format = NF0
        cc.font = Font(size=8, bold=True, color="007A33")
    signature(ws_bs, rc + 5)

    # ---------------------------------------------------------- notes
    ctx = dict(ws_pol=ws_pol, ws_nb=ws_nb, ws_fa=ws_fa, ws_np=ws_np, ws_rt=ws_rt, M=M, M_py=M_py, cy=cy, py=py, pl=pl, bs=bs, prow=prow, brow=brow, framework=framework,
                      company=company, year_end=year_end, prev_end=prev_end, tax=tax, tax_py=tax_py, dep=dep, policy_dep=policy_dep,
                      profile=profile, setup=setup, f_bold=f_bold, f_head=f_head, fill_head=fill_head, thin=thin, dbl=dbl,
                      has_py=has_py, entity=entity, frame=frame, S=S)
    _write_notes(ctx)

    if dep:
        _dep_sheet(ws_dp, dep, setup, f_bold, cy, start, year_end)
    areas = dict(ctx.get("print_areas", {}))
    areas["Balance Sheet"] = "A1:D%d" % ws_bs.max_row
    areas["Profit & Loss"] = "A1:D%d" % ws_pl.max_row
    areas[ws_fa.title] = ctx.get("fa_area")
    land = {ws_fa.title, "Depreciation Schedule", MAPN, MAPPN} | ({ws_rt.title} if ws_rt is not None else set())
    work = {MAPN, MAPPN, "Depreciation Schedule"}
    for ws in wb.worksheets:
        ps = ws.page_setup
        ps.paperSize = 9                                  # A4
        ps.orientation = "landscape" if (ws.title in land and not (ws.title == ws_fa.title and not (dep and dep["basis"] == "sch2" and len(dep["table"])))) else "portrait"
        ps.fitToWidth, ps.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_margins = PageMargins(left=0.6, right=0.6, top=0.7, bottom=0.8)
        ws.print_options.horizontalCentered = True
        ws.oddFooter.left.text = "%s | %s" % (company, ws.title)
        ws.oddFooter.left.size = 8
        ws.oddFooter.right.text = "Page &P of &N"
        ws.oddFooter.right.size = 8
        if areas.get(ws.title):
            ws.print_area = areas[ws.title]
        if ws.title in ("Balance Sheet", "Profit & Loss"):
            ws.print_title_rows = "5:5"
        if ws.title in work:
            ws.sheet_properties.tabColor = "999999"       # working papers, not part of the signed set
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _dep_sheet(ws, dep, setup, f_bold, cy, start, end):
    t = dep["table"]
    s2 = dep["basis"] == "sch2"
    RF, RC, RB, adj_x = cy["RF"], cy["RC"], cy["RB"], cy["adj"]
    setup(ws, {"A": 34, "B": 38, **{chr(67 + i): 16 for i in range(0, 18)}},
          "Depreciation schedule - " + ("Schedule II, Companies Act 2013" if s2 else "Income Tax Act, 1961 (block of assets, WDV)"),
          "Inputs (blue) are editable; the depreciation columns are formulas.", "D")
    ws["A3"].font = Font(size=9, italic=True)
    labels = ["Method", "Residual value (% of cost)", "Days in the year", "Period start", "Period end"]
    for i, lb in enumerate(labels):
        ws.cell(row=4 + i, column=1, value=lb)
    if s2:
        ws["C4"], ws["C5"] = dep["method"], dep["res"]
    else:
        ws["C4"], ws["C5"] = "WDV (block of assets)", "n.a."
    ws["C6"] = "=C8-C7+1"
    ws["C7"], ws["C8"] = start, end
    ws["C7"].number_format = ws["C8"].number_format = "dd-mmm-yyyy"
    ws.freeze_panes = "A11"
    blue = Font(color="0000FF")
    hdr_row = 10
    cols = (S2_COLS + ["Days in use (additions)", "Days in use (disposals)", "Depreciation on assets sold", "Depreciation for the year",
                       "Closing gross block", "Closing accumulated depreciation", "Closing net block", "Profit / (loss) on disposal"]) if s2 \
        else (IT_COLS + ["Total (opening + additions - sale)", "Depreciation for the year", "Closing WDV",
                         "Short-term capital gain (block ceases)"])
    for i, tx in enumerate(cols):
        c = ws.cell(row=hdr_row, column=i + 1, value=tx)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[hdr_row].height = 48
    r0 = hdr_row + 1
    n = len(t)
    for i in range(n):
        r = r0 + i
        for j, val in enumerate(list(t.iloc[i])):
            v = _py(val)
            c = ws.cell(row=r, column=j + 1, value=v)
            if j >= 2:
                c.font, c.number_format = blue, NF2
        if s2:
            res, dy = "($C$5/100)", "$C$6"
            f = "((1-%s)/C%d)" % (res, r)
            rate = "(1-%s^(1/C%d))" % (res, r)
            for col in (3, 4, 8, 12):
                ws.cell(row=r, column=col).number_format = "0" if col in (3, 4) else "dd-mmm-yyyy"
            ws.cell(row=r, column=13, value="=IF(OR(G{r}=0,H{r}=\"\"),0,MAX(0,MIN({dy},$C$8-H{r}+1)))".format(r=r, dy=dy))
            ws.cell(row=r, column=14, value="=IF(OR(I{r}=0,L{r}=\"\"),0,MAX(0,MIN({dy},L{r}-$C$7+1)))".format(r=r, dy=dy))
            if dep["method"] == "SLM":
                o = "=IF(C{r}<=0,0,MIN(I{r}*{f}*N{r}/{dy},MAX(0,I{r}*(1-{res})-J{r})))".format(r=r, f=f, dy=dy, res=res)
                amt = "MAX(0,(E{r}-I{r})-(F{r}-J{r})-(E{r}-I{r})*{res})".format(r=r, res=res)
                p = ("=IF(C{r}<=0,0,ROUND(MAX(0,IF(D{r}>0,MIN({amt}/D{r},{amt}),MIN((E{r}-I{r})*{f},MAX(0,(E{r}-I{r})*(1-{res})-(F{r}-J{r}))))"
                     "+O{r}+G{r}*{f}*M{r}/{dy}),2))").format(r=r, amt=amt, f=f, res=res, dy=dy)
            else:
                o = "=IF(C{r}<=0,0,{rate}*(I{r}-J{r})*N{r}/{dy})".format(r=r, rate=rate, dy=dy)
                p = ("=IF(C{r}<=0,0,ROUND(MAX(0,{rate}*((E{r}-I{r})-(F{r}-J{r}))+O{r}+{rate}*G{r}*M{r}/{dy}),2))"
                     ).format(r=r, rate=rate, dy=dy)
            ws.cell(row=r, column=15, value=o)
            ws.cell(row=r, column=16, value=p)
            ws.cell(row=r, column=17, value="=E%d+G%d-I%d" % (r, r, r))
            ws.cell(row=r, column=18, value="=F%d-J%d+P%d-O%d" % (r, r, r, r))
            ws.cell(row=r, column=19, value="=Q%d-R%d" % (r, r))
            ws.cell(row=r, column=20, value="=IF(I%d>0,K%d-((I%d-J%d)-O%d),0)" % (r, r, r, r, r))
            for cc in range(15, 21):
                ws.cell(row=r, column=cc).number_format = NF2
        else:
            ws.cell(row=r, column=8, value="=D%d+E%d+F%d-G%d" % (r, r, r, r))
            ws.cell(row=r, column=9, value="=IF(H{r}<=0,0,ROUND(MIN(MAX(0,D{r}+E{r}-G{r})*C{r}/100+F{r}*C{r}/200,H{r}),2))".format(r=r))
            ws.cell(row=r, column=10, value="=IF(H{r}<=0,0,ROUND(H{r}-I{r},2))".format(r=r))
            ws.cell(row=r, column=11, value="=MAX(0,-H%d)" % r)
            for cc in range(8, 12):
                ws.cell(row=r, column=cc).number_format = NF2
    rt = r0 + n
    ws.cell(row=rt, column=1, value="Total").font = f_bold
    sumcols = [5, 6, 7, 9, 10, 11, 15, 16, 17, 18, 19, 20] if s2 else [4, 5, 6, 7, 8, 9, 10, 11]
    for cc in sumcols:
        L = chr(64 + cc)
        c = ws.cell(row=rt, column=cc, value="=SUM(%s%d:%s%d)" % (L, r0, L, rt - 1))
        c.number_format, c.font = NF2, f_bold
    dcol = "P" if s2 else "I"
    r = rt + 2
    ws.cell(row=r, column=1, value="Depreciation per calculator").font = f_bold
    ws.cell(row=r, column=3, value="=%s%d" % (dcol, rt))
    ws.cell(row=r + 1, column=1, value="Booked in Tally (ledgers under Depreciation & Amortisation)")
    ws.cell(row=r + 1, column=3, value='=SUMIFS(%s,%s,"Depreciation & Amortisation",%s,"<>%s")' % (RF, RC, RB, DEP_GROUP))
    ws.cell(row=r + 2, column=1, value="Difference (calculator less books)").font = f_bold
    ws.cell(row=r + 2, column=3, value="=C%d-C%d" % (r, r + 1))
    ws.cell(row=r + 3, column=1, value="Adjustment carried in the statements (mapping sheet)")
    ws.cell(row=r + 3, column=3, value="=" + adj_x)
    ws.cell(row=r + 4, column=1, value="Status")
    ws.cell(row=r + 4, column=3, value='=IF(ABS(C%d)<0.5,"Books agree to calculator",IF(ABS(C%d-C%d)<0.5,"Applied to statements","NOT applied to statements"))'
            % (r + 2, r + 2, r + 3))
    for k in range(4):
        ws.cell(row=r + k, column=3).number_format = NF2
    ws.cell(row=r + 4, column=3).font = f_bold
    nxt = r + 6
    if s2:
        ws.cell(row=nxt, column=1, value="Fixed assets per inputs, before this year's depreciation (opening + additions - disposals, net)")
        ws.cell(row=nxt, column=3, value="=E%d-F%d+G%d-(I%d-J%d)" % (rt, rt, rt, rt, rt))
        ws.cell(row=nxt + 1, column=1, value="Fixed assets per Tally before this calculator + depreciation booked in Tally")
        ws.cell(row=nxt + 1, column=3, value='=SUMIFS(%s,%s,"Property, Plant & Equipment",%s,"<>%s")+C%d' % (RF, RC, RB, DEP_GROUP, r + 1))
        ws.cell(row=nxt + 2, column=1, value="Unexplained difference (should be nil when the inputs are complete)").font = f_bold
        ws.cell(row=nxt + 2, column=3, value="=C%d-C%d" % (nxt, nxt + 1))
        for k in range(3):
            ws.cell(row=nxt + k, column=3).number_format = NF2
        nxt += 4
    ws.cell(row=nxt, column=1, value=("Basis: " + dep["note"])).font = Font(size=9, italic=True)

# ================================================================== notes ==
def _policies(framework, entity, policy_dep):
    nc = framework == FRAMEWORKS[2]
    ind = framework == FRAMEWORKS[1]
    std = {FRAMEWORKS[0]: "the Accounting Standards notified under section 133 of the Companies Act, 2013 read with the Companies (Accounting Standards) Rules, and Schedule III (Division I)",
           FRAMEWORKS[1]: "the Indian Accounting Standards (Ind AS) notified under section 133 of the Companies Act, 2013 read with the Companies (Indian Accounting Standards) Rules, and Schedule III (Division II)",
           FRAMEWORKS[2]: "the Accounting Standards issued by the Institute of Chartered Accountants of India as applicable to non-corporate entities"}[framework]
    ent = {"Company": "the Company", "Firm (partnership or LLP)": "the Firm", "Proprietorship": "the Proprietorship"}[entity]
    pol = [
        ("Basis of preparation",
         "The financial statements have been prepared on the historical cost convention, on the accrual basis of accounting and on a going-concern basis, "
         "in accordance with %s. The figures are taken from the books of account maintained in Tally Prime." % std),
        ("Use of estimates",
         "The preparation of financial statements requires management to make estimates and assumptions that affect the reported amounts of assets, "
         "liabilities, income and expenses. Differences between actual results and estimates are recognised in the period in which the results are known."),
        ("Property, plant and equipment and depreciation",
         "Property, plant and equipment are stated at cost less accumulated depreciation and impairment, if any. "
         + (policy_dep or "Depreciation is provided as per Schedule II of the Companies Act, 2013 on the basis of the useful lives prescribed therein."
                          if entity == "Company" else
                          (policy_dep or "Depreciation is provided on the written down value method at the rates adopted by the %s." % ent.replace("the ", "")))),
        ("Inventories",
         "Inventories are valued at the lower of cost and net realisable value. Cost is determined on the basis followed consistently by %s "
         "(first-in-first-out or weighted average)." % ent),
        ("Revenue recognition",
         ("Revenue from contracts with customers is recognised when control of the goods or services is transferred, at the amount of the transaction price, net of GST and returns."
          if ind else
          "Revenue from sale of goods is recognised when the significant risks and rewards of ownership pass to the buyer; revenue from services is recognised as the services are rendered. "
          "Revenue is stated net of GST and returns.")),
        ("Employee benefits",
         "Short-term employee benefits are recognised as an expense at the undiscounted amount in the period in which the related service is rendered. "
         "Post-employment benefits, where applicable, are provided as per the applicable standard."),
        ("Taxes on income",
         "Current tax is the amount of income tax payable determined in accordance with the Income-tax Act, 1961. Deferred tax is recognised on timing differences "
         "between book and taxable income, subject to the prudence considerations of the applicable standard."),
        ("Provisions and contingent liabilities",
         "A provision is recognised when there is a present obligation as a result of a past event, an outflow of resources is probable and a reliable estimate can be made. "
         "Contingent liabilities are disclosed and not provided for."),
        ("Cash and cash equivalents",
         "Cash and cash equivalents comprise cash on hand and balances with banks in current and deposit accounts with an original maturity of three months or less."),
    ]
    if entity == "Company":
        pol.append(("Earnings per share",
                    "Basic earnings per share is computed by dividing the profit for the year attributable to equity shareholders by the number of equity shares outstanding."))
    return pol



def _grp_label(g):
    return ("Reclassified from " + g[:-len(" [reclassified]")]) if str(g).endswith(" [reclassified]") else g


def _no_freeze(ws):
    from openpyxl.worksheet.views import Selection
    ws.freeze_panes = None
    ws.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]


class _W:
    """Small cursor-based writer for the note sheets."""
    def __init__(self, ws, ncols):
        self.ws, self.r, self.ncols = ws, 6, ncols

    def put(self, col, val, fmt=None, font=None, fill=None, align=None, border=None):
        cell = self.ws.cell(row=self.r, column=col, value=val)
        if fmt:
            cell.number_format = fmt
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if align:
            cell.alignment = align
        if border:
            cell.border = border
        return cell

    def bar(self, no, title):
        nav = PatternFill("solid", fgColor=NAVY)
        for k in range(1, self.ncols + 1):
            self.put(k, None, fill=nav)
        self.put(1, no, font=Font(bold=True, color="FFFFFF"), fill=nav, align=Alignment(horizontal="center"))
        self.put(2, title, font=Font(bold=True, color="FFFFFF"), fill=nav)
        self.r += 1

    def para(self, title, text, last_col=7):
        if title:
            self.put(2, title, font=Font(bold=True))
            self.r += 1
        self.put(2, text, align=Alignment(wrap_text=True, vertical="top"))
        self.ws.merge_cells(start_row=self.r, start_column=2, end_row=self.r, end_column=last_col)
        self.ws.row_dimensions[self.r].height = max(15, 15 * (1 + len(text) // 105))
        self.r += 1


def _write_notes(c):
    M, Mp, cy, py = c["M"], c["M_py"], c["cy"], c["py"]
    pl, bs, prow, brow = c["pl"], c["bs"], c["prow"], c["brow"]
    framework, entity, has_py = c["framework"], c["entity"], c["has_py"]
    year_end, prev_end, dep, profile = c["year_end"], c["prev_end"], c["dep"], c["profile"] or {}
    f_bold, thin, dbl, setup, frame, nm = c["f_bold"], c["thin"], c["dbl"], c["setup"], c["frame"], c["company"]
    d_fmt = lambda d: d.strftime("%d-%b-%Y")
    GREY = PatternFill("solid", fgColor="F2F2F2")
    BS_, PL_ = "'Balance Sheet'", PLS
    title_notes = "Notes forming part of the Financial Statements for the year ended %s" % long_date(year_end)
    wide = {"A": 8, "B": 58, **{chr(67 + i): 18 for i in range(0, 10)}}

    # ------------------------------------------------ accounting policies sheet
    ws = c["ws_pol"]
    setup(ws, {"A": 8, "B": 110}, title_notes, frame, "B")
    _no_freeze(ws)
    w = _W(ws, 2)
    w.bar("A.", "General information")
    kind = {"Company": "a company incorporated under the Companies Act", "Firm (partnership or LLP)": "a partnership firm / limited liability partnership",
            "Proprietorship": "a proprietorship concern"}[entity]
    gi = "%s is %s." % (nm, kind)
    if profile.get("cin"):
        gi += " CIN / LLPIN: %s." % profile["cin"]
    if profile.get("address"):
        gi += " Registered office / place of business: %s." % profile["address"]
    gi += " These financial statements are for the year ended %s." % long_date(year_end)
    w.para(None, gi, 2)
    w.r += 1
    w.bar("B.", "Significant accounting policies")
    for t, tx in _policies(framework, entity, c["policy_dep"]):
        w.para(t, tx, 2)

    # ------------------------------------------------ numbered notes
    def mrow_cells(info, Mx):
        out = {}
        for i, (led, grp, head, sg, bal) in enumerate(Mx["rows"]):
            out.setdefault((norm_name(led), head), (led, _grp_label(grp), head, info["first"] + i, sg * bal))
        return out

    cyr = mrow_cells(cy, M)
    pyr = mrow_cells(py, Mp) if has_py else {}
    cell_ref = lambda info, row: "%s!F%d" % (info["q"], row)
    sgn = lambda cf: "-" if cf < 0 else ""

    def components(r, own, sheet):
        out = []
        multi = len(r["terms"]) > 1
        for cf, k, ref in r["terms"]:
            pre = "Less: " if (cf < 0 and multi) else ""
            if k == "h":
                keys = sorted({kk for kk in list(cyr) + list(pyr) if kk[1] == ref})
                items = []
                for kk in keys:
                    a, b = cyr.get(kk), pyr.get(kk)
                    ref_ = a or b
                    items.append({"label": ref_[0], "group": ref_[1], "amt": abs(a[4]) if a else 0,
                                  "cyf": "=%s%s" % (sgn(cf), cell_ref(cy, a[3])) if a else None,
                                  "pyf": "=%s%s" % (sgn(cf), cell_ref(py, b[3])) if b else None})
                if ref == "Reserves and Surplus":
                    items.append({"label": "Surplus in Statement of Profit and Loss: opening balance", "group": "~", "amt": 0,
                                  "cyf": "=%s%s" % (sgn(cf), cell_ref(cy, cy["r_or"])),
                                  "pyf": ("=%s%s" % (sgn(cf), cell_ref(py, py["r_or"]))) if has_py else None})
                if ref == "Inventories" and cy["has_cs"]:
                    items.append({"label": "Closing stock", "group": "~", "amt": 0,
                                  "cyf": "=%s%s" % (sgn(cf), cell_ref(cy, cy["r_cs"])),
                                  "pyf": ("=%s%s" % (sgn(cf), cell_ref(py, py["r_cs"]))) if has_py and py["has_cs"] else None})
                items.sort(key=lambda x: (x["group"] == "~", x["group"], -x["amt"]))
                if pre and len(items) == 1:
                    items[0]["label"] = pre + items[0]["label"]
                groups = {it["group"] for it in items if it["group"] != "~"}
                last_g = None
                for it in items:
                    if len(groups) > 1 and it["group"] not in ("~", last_g):
                        out.append({"sub": it["group"]})
                        last_g = it["group"]
                    out.append(it)
            elif k == "p":
                pc = lambda col, ref=ref, cf=cf: "=%s%s!%s%d" % (sgn(cf), PL_, col, prow[ref])
                out.append({"label": pre + "Profit for the year", "cyf": pc("C"), "pyf": pc("D") if has_py else None})
            elif k == "x":
                if ref == "closing" and not cy["has_cs"] and not (has_py and py["has_cs"]):
                    continue                  # closing stock is held in a ledger, so there is no separate row to show
                lab = {"tax": "Provision for income tax", "closing": "Closing stock"}.get(ref, ref)
                out.append({"label": pre + lab, "cyf": "=%s%s" % (sgn(cf), cy["ex"][ref]),
                            "pyf": ("=%s%s" % (sgn(cf), py["ex"][ref])) if has_py else None})
            else:
                out.append({"label": pre + ref, "cyf": "=%s%s!C%d" % (sgn(cf), sheet, own[ref]),
                            "pyf": ("=%s%s!D%d" % (sgn(cf), sheet, own[ref])) if has_py else None})
        return out

    def head_row(w, is_bs):
        lead = "As at" if is_bs else "Year ended"
        for k, t in ((2, "Particulars"), (3, "%s\n%s" % (lead, d_fmt(year_end))), (4, "%s\n%s" % (lead, d_fmt(prev_end)))):
            w.put(k, t, font=f_bold, fill=GREY, align=Alignment(horizontal="center" if k > 2 else "left", wrap_text=True))
        w.put(8, "Check (not printed)", font=Font(size=8, italic=True, color="888888"), align=Alignment(horizontal="center"))
        w.ws.row_dimensions[w.r].height = 30
        w.r += 1

    notes = sorted([r for r in pl + bs if r["no"]], key=lambda r: r["no"])
    ws_b, ws_p = c["ws_nb"], c["ws_np"]
    for wsx in (ws_b, ws_p):
        setup(wsx, wide, title_notes, frame, "G")
        _no_freeze(wsx)
        wsx.column_dimensions["H"].width = 20
    wb_, wp_ = _W(ws_b, 4), _W(ws_p, 4)
    for L_ in "EFG":
        for w_ in (ws_b, ws_p):
            w_.column_dimensions[L_].width = 2.5        # unused on these sheets; keeps the reviewer check column close
    ppe_no = None
    blocks = []
    for r in notes:
        is_bs = r in bs
        w = wb_ if is_bs else wp_
        own, sheet = (brow, BS_) if is_bs else (prow, PL_)
        key = r["key"]
        start_r = w.r
        w.bar("%d." % r["no"], _clean(r["label"]))
        if key == "eps":
            _eps_block(c, w, f_bold, thin, dbl, lambda: head_row(w, False))
            w.r += 2
            blocks.append((w.ws, start_r, w.r - 1))
            continue
        head_row(w, is_bs)
        comps = components(r, own, sheet)
        first = w.r
        if not any("label" in x for x in comps):
            w.put(2, "No balances")
            w.put(3, 0, NF0)
            w.put(4, 0 if has_py else "-", NF0, align=Alignment(horizontal="right"))
            w.r += 1
        for x in comps:
            if "sub" in x:
                w.put(2, x["sub"], font=Font(bold=True, italic=True))
                w.r += 1
                continue
            w.put(2, x["label"], align=Alignment(indent=1))
            w.put(3, x["cyf"] or 0, NF0)
            w.put(4, (x["pyf"] if has_py and x["pyf"] else (0 if has_py else "-")), NF0,
                  align=Alignment(horizontal="right") if not has_py else None)
            w.r += 1
        last, tot = w.r - 1, w.r
        w.put(2, "Total", font=f_bold, align=Alignment(horizontal="right"))
        for col, L in ((3, "C"), (4, "D")):
            if col == 4 and not has_py:
                w.put(4, "-", align=Alignment(horizontal="right"), font=f_bold, border=Border(top=thin, bottom=dbl))
                continue
            w.put(col, "=SUM(%s%d:%s%d)" % (L, first, L, last), NF0, f_bold, border=Border(top=thin, bottom=dbl))
        chk = ('=IF(AND(ABS(C%d-%s!C%d)<1,ABS(D%d-%s!D%d)<1),"Agrees","CHECK")' % (tot, sheet, own[key], tot, sheet, own[key])) if has_py else \
              ('=IF(ABS(C%d-%s!C%d)<1,"Agrees","CHECK")' % (tot, sheet, own[key]))
        w.put(8, chk, font=Font(size=8, color="007A33"), align=Alignment(horizontal="center"))
        w.r += 1
        if key == "ppe":
            ppe_no = r["no"]
            ppe_tot = tot
            if dep and dep["basis"] == "sch2" and len(dep["table"]):
                w.put(2, "Refer the Fixed Assets Schedule for gross block, depreciation and net block.", font=Font(size=9, italic=True))
                w.r += 1
        w.r += 2
        blocks.append((w.ws, start_r, w.r - 1))
    ratio_area = None
    if entity == "Company":
        ws_r = c["ws_rt"]
        setup(ws_r, {"A": 8, "B": 36, "C": 28, "D": 28, "E": 14, "F": 14, "G": 12}, title_notes, frame, "G")
        _no_freeze(ws_r)
        wr_ = _W(ws_r, 7)
        _ratio_block(c, wr_, f_bold, len(notes) + 1)
        ratio_area = "A1:G%d" % (wr_.r + 1)
    _paginate(blocks)

    # ------------------------------------------------ fixed assets schedule sheet
    _fa_sheet(c, c["ws_fa"], ppe_no, ppe_tot, cyr, pyr, f_bold, thin, dbl)
    c["print_areas"] = {c["ws_pol"].title: "A1:B%d" % c["ws_pol"].max_row, ws_b.title: "A1:D%d" % ws_b.max_row,
                        ws_p.title: "A1:D%d" % ws_p.max_row}
    if ratio_area:
        c["print_areas"]["Notes - Ratios"] = ratio_area


def _paginate(blocks, first_page=680.0, page=760.0):
    """Insert page breaks so that no note is split across two printed pages."""
    from openpyxl.worksheet.pagebreak import Break
    for ws in {b[0] for b in blocks}:
        used, cap = 0.0, first_page
        for w_, s, e in blocks:
            if w_ is not ws:
                continue
            h = sum((ws.row_dimensions[r].height or 15.0) for r in range(s, e + 1))
            if used > 0 and used + h > cap:
                ws.row_breaks.append(Break(id=s - 1))
                used, cap = 0.0, page
            used += h


def _eps_block(c, w, f_bold, thin, dbl, head_row):
    prow, has_py = c["prow"], c["has_py"]
    prof = c["profile"] or {}
    head_row()
    shares = float(prof.get("eq_shares") or 0)
    fv = float(prof.get("face_value") or 10)
    pref = float(prof.get("pref_div") or 0)
    r0 = w.r
    w.put(2, "Profit for the year")
    w.put(3, "=%s!C%d" % (PLS, prow["pat"]), NF0)
    w.put(4, ("=%s!D%d" % (PLS, prow["pat"])) if has_py else "-", NF0, align=Alignment(horizontal="right"))
    w.r += 1
    w.put(2, "Less: Preference dividend")
    w.put(3, pref, NF0)
    w.put(4, pref if has_py else "-", NF0, align=Alignment(horizontal="right"))
    w.r += 1
    w.put(2, "Earnings attributable to equity shareholders", font=f_bold)
    w.put(3, "=C%d-C%d" % (r0, r0 + 1), NF0, f_bold, border=Border(top=thin))
    w.put(4, ("=D%d-D%d" % (r0, r0 + 1)) if has_py else "-", NF0, f_bold, border=Border(top=thin), align=Alignment(horizontal="right"))
    w.r += 1
    w.put(2, "Number of equity shares outstanding")
    w.put(3, shares, '#,##0')
    w.put(4, shares if has_py else "-", '#,##0', align=Alignment(horizontal="right"))
    w.r += 1
    w.put(2, "Face value per equity share (Rs)")
    w.put(3, fv, '#,##0.00')
    w.put(4, fv if has_py else "-", '#,##0.00', align=Alignment(horizontal="right"))
    w.r += 1
    w.put(2, "Basic and diluted earnings per share (Rs)", font=f_bold)
    w.put(3, '=IF(C%d>0,C%d/C%d,"n.a.")' % (r0 + 3, r0 + 2, r0 + 3), '#,##0.00', f_bold, border=Border(top=thin, bottom=dbl),
          align=Alignment(horizontal="right"))
    w.put(4, ('=IF(D%d>0,D%d/D%d,"n.a.")' % (r0 + 3, r0 + 2, r0 + 3)) if has_py else "-", '#,##0.00', f_bold,
          border=Border(top=thin, bottom=dbl), align=Alignment(horizontal="right"))
    w.r += 1


def _fa_sheet(c, ws, ppe_no, ppe_tot, cyr, pyr, f_bold, thin, dbl):
    dep, has_py = c["dep"], c["has_py"]
    year_end, prev_end = c["year_end"], c["prev_end"]
    d_fmt = lambda d: d.strftime("%d-%b-%Y")
    nav = PatternFill("solid", fgColor=NAVY)
    sch2 = bool(dep and dep["basis"] == "sch2" and len(dep["table"]))
    c["setup"](ws, {"A": 8, "B": 44, **{chr(67 + i): 16 for i in range(0, 10)}},
               "Property, Plant and Equipment: schedule as at %s" % long_date(year_end), c["frame"], "L")
    _no_freeze(ws)
    w = _W(ws, 12 if sch2 else 4)
    w.bar("%d." % ppe_no if ppe_no else "", "Property, Plant and Equipment")
    if not sch2:
        for k, t in ((2, "Particulars"), (3, "Net block\nas at %s" % d_fmt(year_end)), (4, "Net block\nas at %s" % d_fmt(prev_end))):
            w.put(k, t, font=f_bold, fill=PatternFill("solid", fgColor="F2F2F2"), align=Alignment(horizontal="center" if k > 2 else "left", wrap_text=True))
        ws.row_dimensions[w.r].height = 30
        w.r += 1
        first = w.r
        keys = sorted({kk for kk in list(cyr) + list(pyr) if kk[1] == "Property, Plant & Equipment"})
        for kk in keys:
            a, b = cyr.get(kk), pyr.get(kk)
            ref_ = a or b
            w.put(2, ref_[0])
            w.put(3, ("=%s!F%d" % (c["cy"]["q"], a[3])) if a else 0, NF0)
            w.put(4, ("=%s!F%d" % (c["py"]["q"], b[3])) if b else (0 if has_py else "-"), NF0, align=Alignment(horizontal="right") if not has_py else None)
            w.r += 1
        w.put(2, "Total", font=f_bold, align=Alignment(horizontal="right"))
        w.put(3, "=SUM(C%d:C%d)" % (first, w.r - 1), NF0, f_bold, border=Border(top=thin, bottom=dbl))
        w.put(4, ("=SUM(D%d:D%d)" % (first, w.r - 1)) if has_py else "-", NF0, f_bold, border=Border(top=thin, bottom=dbl),
              align=Alignment(horizontal="right"))
        c["print_areas"] = c.get("print_areas", {})
        c["fa_area"] = "A1:D%d" % w.r
        return
    t = dep["table"]
    n = len(t)
    r0, rl = 11, 11 + n - 1
    classes = []
    for v in t["Asset class"].tolist():
        v = str(v or "Other")
        if v not in classes:
            classes.append(v)
    hdr = ["Asset class", "Gross block\nas at %s" % d_fmt(prev_end), "Additions", "Deductions", "Gross block\nas at %s" % d_fmt(year_end),
           "Depreciation\nas at %s" % d_fmt(prev_end), "For the year", "On deductions", "Depreciation\nas at %s" % d_fmt(year_end),
           "Net block\nas at %s" % d_fmt(year_end), "Net block\nas at %s" % d_fmt(prev_end)]
    for i, tx in enumerate(hdr):
        w.put(2 + i, tx, font=Font(bold=True, color="FFFFFF"), fill=nav,
              align=Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True))
    ws.row_dimensions[w.r].height = 46
    w.r += 1
    first = w.r

    def sif(col, r):
        return "SUMIFS(%s!$%s$%d:$%s$%d,%s!$B$%d:$B$%d,$B%d)" % (DEPS, col, r0, col, rl, DEPS, r0, rl, r)

    for cl in classes:
        r = w.r
        w.put(2, cl)
        w.put(3, "=" + sif("E", r), NF0)
        w.put(4, "=" + sif("G", r), NF0)
        w.put(5, "=" + sif("I", r), NF0)
        w.put(6, "=C%d+D%d-E%d" % (r, r, r), NF0)
        w.put(7, "=" + sif("F", r), NF0)
        w.put(8, "=" + sif("P", r), NF0)
        w.put(9, "=%s+%s" % (sif("J", r), sif("O", r)), NF0)
        w.put(10, "=G%d+H%d-I%d" % (r, r, r), NF0)
        w.put(11, "=F%d-J%d" % (r, r), NF0)
        w.put(12, "=C%d-G%d" % (r, r), NF0)
        w.r += 1
    last, r = w.r - 1, w.r
    w.put(2, "Total", font=f_bold)
    for col in range(3, 13):
        L = chr(64 + col)
        w.put(col, "=SUM(%s%d:%s%d)" % (L, first, L, last), NF0, f_bold, border=Border(top=thin, bottom=dbl))
    w.r += 2
    w.put(2, "Depreciation for the year includes depreciation on assets sold up to the date of sale; it is removed with the accumulated depreciation of those assets under 'On deductions'.",
          font=Font(size=9, italic=True))
    w.r += 1
    w.put(2, dep["note"], font=Font(size=9, italic=True))
    # reviewer check, outside the print area
    ws.cell(row=r, column=14, value="Net block less Balance Sheet note").font = Font(size=8, italic=True, color="888888")
    ck = ws.cell(row=r, column=15, value="=K%d-%s!C%d" % (r, "'Notes - Balance Sheet'", ppe_tot))
    ck.number_format, ck.font = NF0, Font(size=8, color="888888")
    c["fa_area"] = "A1:L%d" % w.r


def _ratio_block(c, w, f_bold, no):
    brow, prow, has_py = c["brow"], c["prow"], c["has_py"]
    w.ncols = 7
    w.bar("%d." % no, "Analytical ratios")
    w.ncols = 4
    for i, tx in enumerate(["Ratio", "Numerator", "Denominator", "Current year", "Previous year", "% variance"]):
        w.put(2 + i, tx, font=f_bold, fill=PatternFill("solid", fgColor="F2F2F2"), align=Alignment(horizontal="center", wrap_text=True))
    w.r += 1
    B = lambda k, col: "%s!%s%d" % ("'Balance Sheet'", col, brow[k])
    P = lambda k, col: "%s!%s%d" % (PLS, col, prow[k])
    avg = lambda k, col: ("AVERAGE(%s,%s)" % (B(k, "C"), B(k, "D"))) if (col == "C" and has_py) else B(k, col)
    avgnc = lambda col: ("AVERAGE(%s-%s,%s-%s)" % (B("ca", "C"), B("cl", "C"), B("ca", "D"), B("cl", "D"))) if (col == "C" and has_py) \
        else "(%s-%s)" % (B("ca", col), B("cl", col))
    R = [
        ("Current ratio (times)", "Current assets", "Current liabilities", lambda s: (B("ca", s), B("cl", s)), 1, '0.00'),
        ("Debt-equity ratio (times)", "Total borrowings", "Shareholders' funds", lambda s: ("(%s+%s)" % (B("ltb", s), B("stb", s)), B("fund", s)), 1, '0.00'),
        ("Return on equity (%)", "Profit for the year", "Average shareholders' funds", lambda s: (P("pat", s), avg("fund", s)), 100, '0.00'),
        ("Inventory turnover (times)", "Cost of goods sold", "Average inventories", lambda s: ("(%s+%s)" % (P("pur", s), P("chg", s)), avg("stock", s)), 1, '0.00'),
        ("Trade receivables turnover (times)", "Revenue from operations", "Average trade receivables", lambda s: (P("rev", s), avg("tr", s)), 1, '0.00'),
        ("Trade payables turnover (times)", "Purchases", "Average trade payables", lambda s: (P("pur", s), avg("tp", s)), 1, '0.00'),
        ("Net capital turnover (times)", "Revenue from operations", "Average working capital", lambda s: (P("rev", s), avgnc(s)), 1, '0.00'),
        ("Net profit ratio (%)", "Profit for the year", "Revenue from operations", lambda s: (P("pat", s), P("rev", s)), 100, '0.00'),
        ("Return on capital employed (%)", "Profit before tax + finance costs", "Capital employed",
         lambda s: ("(%s+%s)" % (P("pbt", s), P("fin", s)), "(%s+%s+%s)" % (B("fund", s), B("ltb", s), B("stb", s))), 100, '0.00'),
    ]
    for name, nt, dt, fn, mult, fmt in R:
        r = w.r
        w.put(2, name)
        w.put(3, nt, font=Font(size=9))
        w.put(4, dt, font=Font(size=9))
        for col, s in ((5, "C"), (6, "D")):
            if s == "D" and not has_py:
                w.put(col, "-", align=Alignment(horizontal="right"))
                continue
            n_, d_ = fn(s)
            w.put(col, '=IFERROR(%s/%s%s,"n.a.")' % (n_, d_, "*100" if mult == 100 else ""), fmt, align=Alignment(horizontal="right"))
        w.put(7, ('=IFERROR((E%d-F%d)/ABS(F%d)*100,"n.a.")' % (r, r, r)) if has_py else "-", '0.0', align=Alignment(horizontal="right"))
        w.r += 1


# ========================================================== sample (synthetic) data ==
# Every name and number below is invented for demonstration. No real client data is used.
SAMPLE_NAME = "Sample Engineering Private Limited"
SAMPLE_GROUPS = {
    "Capital Account": "Primary", "Reserves & Surplus": "Primary", "Secured Loans": "Primary", "Unsecured Loans": "Primary",
    "Bank OD A/c": "Primary", "Current Liabilities": "Primary", "Duties & Taxes": "Current Liabilities", "Provisions": "Current Liabilities",
    "Sundry Creditors": "Current Liabilities", "Fixed Assets": "Primary", "Investments": "Primary", "Current Assets": "Primary",
    "Stock-in-Hand": "Current Assets", "Sundry Debtors": "Current Assets", "Cash-in-Hand": "Current Assets", "Bank Accounts": "Current Assets",
    "Deposits (Asset)": "Current Assets", "Loans & Advances (Asset)": "Current Assets", "Sales Accounts": "Primary",
    "Purchase Accounts": "Primary", "Direct Expenses": "Primary", "Indirect Expenses": "Primary", "Indirect Incomes": "Primary",
}
# ledger, Tally group, current-year closing (debit positive), previous-year closing (debit positive)
_SAMPLE = [
    ("Equity Share Capital", "Capital Account", -2500000, -2500000),
    ("Term Loan - Sample Bank", "Secured Loans", -3200000, -3900000),
    ("Loan from Director", "Unsecured Loans", -600000, -450000),
    ("Cash Credit - Sample Bank", "Bank OD A/c", -1450000, -1100000),
    ("Plant and Machinery", "Fixed Assets", 5625000, 5300000),
    ("Furniture and Fixtures", "Fixed Assets", 338000, 365000),
    ("Computers and Printers", "Fixed Assets", 215000, 148000),
    ("Vehicles", "Fixed Assets", 790000, 900000),
    ("Alpha Components Pvt Ltd", "Sundry Creditors", -1380000, -1190000),
    ("Beta Metals Traders", "Sundry Creditors", -920000, -760000),
    ("Gamma Logistics", "Sundry Creditors", -340000, -280000),
    ("GST Output Payable", "Duties & Taxes", -410000, -350000),
    ("TDS Payable", "Duties & Taxes", -86000, -71000),
    ("PF Payable", "Duties & Taxes", -54000, -47000),
    ("Provision for Gratuity", "Provisions", -180000, -140000),
    ("Audit Fee Payable", "Current Liabilities", -60000, -50000),
    ("Salary Payable", "Current Liabilities", -230000, -205000),
    ("Advance from Customers", "Current Liabilities", -275000, -120000),
    ("Delta Industries Ltd", "Sundry Debtors", 2150000, 1820000),
    ("Epsilon Engineering", "Sundry Debtors", 1640000, 1330000),
    ("Zeta Retail Co", "Sundry Debtors", 880000, 740000),
    ("Eta Customer (advance received)", "Sundry Debtors", -90000, 0),
    ("Cash in Hand", "Cash-in-Hand", 45000, 38000),
    ("Rent Deposit", "Deposits (Asset)", 300000, 300000),
    ("Advance to Suppliers", "Loans & Advances (Asset)", 260000, 190000),
    ("Staff Advances", "Loans & Advances (Asset)", 75000, 60000),
    ("GST Input Credit", "Current Assets", 390000, 310000),
    ("Prepaid Insurance", "Current Assets", 48000, 41000),
    ("Opening Stock", "Stock-in-Hand", 1900000, 1500000),
    ("Sales - Products", "Sales Accounts", -17800000, -15200000),
    ("Sales - Services", "Sales Accounts", -2600000, -2100000),
    ("Interest Received", "Indirect Incomes", -64000, -52000),
    ("Discount Received", "Indirect Incomes", -22000, -18000),
    ("Purchase of Raw Materials", "Purchase Accounts", 9600000, 8300000),
    ("Freight Inward", "Direct Expenses", 420000, 370000),
    ("Salaries and Wages", "Indirect Expenses", 3250000, 2800000),
    ("PF Employer Contribution", "Indirect Expenses", 310000, 270000),
    ("Staff Welfare", "Indirect Expenses", 120000, 98000),
    ("Interest on Term Loan", "Indirect Expenses", 380000, 450000),
    ("Interest on Cash Credit", "Indirect Expenses", 165000, 130000),
    ("Bank Charges", "Indirect Expenses", 38000, 31000),
    ("Rent", "Indirect Expenses", 720000, 660000),
    ("Electricity", "Indirect Expenses", 410000, 370000),
    ("Repairs and Maintenance", "Indirect Expenses", 265000, 215000),
    ("Insurance", "Indirect Expenses", 96000, 88000),
    ("Professional Fees", "Indirect Expenses", 185000, 160000),
    ("Audit Fees", "Indirect Expenses", 60000, 50000),
    ("Telephone and Internet", "Indirect Expenses", 72000, 66000),
    ("Travelling Expenses", "Indirect Expenses", 238000, 196000),
    ("Office Expenses", "Indirect Expenses", 134000, 119000),
    ("Advertisement", "Indirect Expenses", 150000, 110000),
    ("Depreciation", "Indirect Expenses", 700000, 820000),
]
_SAMPLE_CS = {"cy": 2250000.0, "py": 1900000.0}          # closing stock (valued), not a ledger
_SAMPLE_OPEN_RES_PY = 1850000.0                           # reserves at the start of the previous year
_NOMINAL = {"Sales Accounts", "Indirect Incomes", "Purchase Accounts", "Direct Expenses", "Indirect Expenses"}


def _sample_profit(idx):
    nominal = sum(r[idx] for r in _SAMPLE if r[1] in _NOMINAL)
    open_stock = [r[idx] for r in _SAMPLE if r[0] == "Opening Stock"][0]
    cs = _SAMPLE_CS["cy" if idx == 2 else "py"]
    return -nominal + cs - open_stock


def sample_trial_balance(prior=False):
    """Synthetic trial balance in the layout load_trial_balance() returns: (df, note, closing_stock_value)."""
    idx = 3 if prior else 2
    open_res = _SAMPLE_OPEN_RES_PY if prior else _SAMPLE_OPEN_RES_PY + _sample_profit(3)      # reserves at the start of the year
    plac_cr = open_res + _sample_profit(idx)           # Tally's P&L A/c ledger carries the year's profit as well
    groups = dict(SAMPLE_GROUPS)
    recs = []
    for r in _SAMPLE:
        if abs(r[idx]) < 0.005:
            continue
        head, top = classify(r[0], r[1], groups)
        recs.append({"Ledger": r[0], "Group": r[1], "Top group": top, "Head": head, "Dr_pos": float(r[idx])})
    bank = open_res - sum(x["Dr_pos"] for x in recs)    # ledgers other than the P&L A/c sum to the opening reserves
    head, top = classify("Sample Bank Current Account", "Bank Accounts", groups)
    recs.append({"Ledger": "Sample Bank Current Account", "Group": "Bank Accounts", "Top group": top, "Head": head, "Dr_pos": float(bank)})
    head, top = classify("Profit & Loss A/c", "Primary", groups)
    recs.append({"Ledger": "Profit & Loss A/c", "Group": "Primary", "Top group": top, "Head": head, "Dr_pos": float(-plac_cr)})
    df = pd.DataFrame(recs)
    df.attrs["groups"] = groups
    return df, "Synthetic sample data (no real client information).", _SAMPLE_CS["py" if prior else "cy"]


def sample_profile():
    p = default_profile()
    p.update({"firm": "Sample and Associates", "frn": "000000S", "partner": "A. Sample", "mno": "000000", "udin": "(to be generated)",
              "place": "Kochi", "rdate": "28-Oct-2026", "cin": "U00000KL2015PTC000000", "pan": "AAAAA0000A",
              "address": "Sample Industrial Estate, Kochi, Kerala", "bdate": "28-Oct-2026",
              "eq_shares": 250000, "face_value": 10, "pref_div": 0,
              "signs": [{"name": "Director One", "id": "DIN 00000001", "desig": "Director"},
                        {"name": "Director Two", "id": "DIN 00000002", "desig": "Director"}]})
    return p


# Additions during the sample year (synthetic): they make the fixed-asset inputs agree to the sample ledgers.
_SAMPLE_ADDITIONS = {"Plant and Machinery": (700000.0, date(2025, 6, 15)), "Computers and Printers": (120000.0, date(2025, 10, 1)),
                     "Vehicles": (135000.0, date(2026, 1, 20))}


def sample_dep_inputs(s2, it):
    s2, it = s2.copy(), it.copy()
    for i, r in s2.iterrows():
        net = float(r["Opening gross block"])                      # previous year-end ledger balance (written-down value)
        if net > 0:
            s2.at[i, "Opening gross block"] = round(net * 1.30, 2)         # invented split of cost and accumulated depreciation
            s2.at[i, "Opening accumulated depreciation"] = round(net * 0.30, 2)
        if r["Tally ledger"] in _SAMPLE_ADDITIONS:
            amt, d = _SAMPLE_ADDITIONS[r["Tally ledger"]]
            s2.at[i, "Additions (cost)"] = amt
            s2.at[i, "Date put to use (additions)"] = d
    for i, r in it.iterrows():
        if r["Tally ledger"] in _SAMPLE_ADDITIONS:
            amt, d = _SAMPLE_ADDITIONS[r["Tally ledger"]]
            it.at[i, "Additions (used 180 days or more)" if (date(2026, 3, 31) - d).days >= 180 else "Additions (used under 180 days)"] = amt
    return s2, it


# ========================================================= auditor action list ==
AUDIT_CHECK = "Auditor action list cleared"
_CHECK_ACTION = {
    "Trial balance as pulled": "Agree the trial balance total to Tally's Trial Balance report for the period; investigate any difference.",
    "All ledgers classified": "Review every ledger with no head; decide the Schedule III head and re-run the application.",
    "Balance Sheet balances": "Do not sign until the Balance Sheet balances. Trace the difference to unclassified or reclassified ledgers.",
    "Tally Profit & Loss A/c tie-out": "Agree Tally's Profit & Loss A/c ledger to opening reserves plus the year's profit; explain any difference.",
    "Closing stock": "Agree closing stock to the stock valuation, physical verification report and Tally's Stock Summary.",
    "Fixed assets agree to depreciation calculator": "Agree gross block, additions, disposals and depreciation to the fixed-asset register.",
    "Depreciation calculator": "Decide whether the calculator result should be adopted; if so apply it and have the entry booked in Tally.",
    "Depreciation inputs agree to Tally fixed assets": "Complete the opening values, additions and disposals so the calculator explains the Tally balance.",
    "Opposite balances shown gross": "Confirm that the reclassified balances (advances from customers, debit balances in creditors, overdrafts) are genuine and correctly grouped.",
    "Income-tax provision": "Obtain the income-tax computation; book the current tax and deferred tax and re-run the application.",
    "Previous year: all ledgers classified": "Classify the remaining previous-year ledgers.",
    "Previous year: classification source": "Confirm ledgers classified by similar name against the signed previous-year statements.",
    "Previous year: Balance Sheet balances": "Resolve the previous-year difference before relying on the comparatives.",
    "Opening equity agrees to previous year closing": "Agree opening reserves to last year's signed Balance Sheet; explain capital introduced or withdrawn, dividends and prior-period items.",
    "Same ledgers classified the same way in both years": "Review ledgers whose head changed; disclose any reclassification of comparatives.",
}


def _standard_items(framework, entity):
    co = entity == "Company"
    ind = framework == FRAMEWORKS[1]
    std = "Ind AS" if ind else "AS"
    I = [
        ("Engagement", "Engagement terms and independence", "Confirm the engagement letter, independence, and the auditor's appointment and eligibility.",
         "SA 210, SA 220, Companies Act s.141" if co else "SA 210, SA 220"),
        ("Engagement", "Planning, materiality and risk assessment", "Document materiality, significant risks and the audit approach before fieldwork.", "SA 300, SA 315, SA 320"),
        ("Revenue", "Revenue recognition and cut-off", "Test sales near the year-end for cut-off; agree sales to invoices, dispatches and GST returns (GSTR-1 / 3B); reconcile turnover to the GST and tax returns.",
         "SA 240, SA 330, " + ("Ind AS 115" if ind else "AS 9")),
        ("Receivables", "Trade receivables", "Send balance confirmations; test ageing, subsequent receipts, disputed and doubtful balances; review the provision.",
         "SA 505, SA 500" + (", Ind AS 109 (expected credit loss)" if ind else "")),
        ("Payables", "Trade payables and MSME", "Confirm major creditors; search for unrecorded liabilities; identify MSME suppliers and check interest and disclosure.",
         "SA 505, MSMED Act s.15 and s.22" if co else "SA 505, MSMED Act"),
        ("Inventory", "Inventory existence and valuation", "Attend or review the physical count; test valuation at lower of cost and net realisable value; check slow-moving items.",
         "SA 501, " + ("Ind AS 2" if ind else "AS 2")),
        ("Fixed assets", "Fixed assets and depreciation", "Vouch additions and disposals; inspect assets; check capitalisation, useful lives and the depreciation charge to the Fixed Assets Schedule.",
         ("Schedule II, " if co else "") + ("Ind AS 16" if ind else "AS 10 and AS 6")),
        ("Borrowings", "Borrowings and finance cost", "Obtain bank and lender confirmations and sanction letters; check security, covenants, classification of current maturities and interest accrued.",
         "SA 505" + (", Schedule III disclosures" if co else "")),
        ("Cash and bank", "Cash and bank balances", "Agree bank balances to confirmations and reconciliations; count cash; review unusual or old reconciling items.", "SA 505, SA 500"),
        ("Statutory dues", "GST, TDS, PF, ESI and other dues", "Reconcile the ledgers to returns and challans; check payment dates, interest and any dues outstanding over six months.",
         ("Schedule III, " if co else "") + "SA 250"),
        ("Employees", "Employee costs and benefits", "Agree payroll to registers and statutory returns; check provision for gratuity and leave encashment.", "SA 500, " + ("Ind AS 19" if ind else "AS 15")),
        ("Taxation", "Income tax and deferred tax", "Review the tax computation, advance tax and TDS credits, and deferred tax on depreciation and disallowances.",
         "Income-tax Act 1961, " + ("Ind AS 12" if ind else "AS 22")),
        ("Related parties", "Related-party transactions", "Obtain the list of related parties; test transactions and balances; check disclosure and arm's-length support.",
         "SA 550, " + ("Ind AS 24" if ind else "AS 18") + (", Companies Act s.188" if co else "")),
        ("Contingencies", "Contingent liabilities and commitments", "Obtain legal and tax litigation details; check provisions and disclosure.", "SA 501, SA 580, " + ("Ind AS 37" if ind else "AS 29")),
        ("Subsequent events", "Events after the reporting date", "Review events up to the date of the report for adjustment or disclosure.", "SA 560"),
        ("Going concern", "Going concern", "Assess management's going-concern assessment, cash flows and borrowing headroom.", "SA 570"),
        ("Representations", "Management representation letter", "Obtain a signed representation letter covering completeness, related parties, litigation and going concern.", "SA 580"),
    ]
    if co:
        I += [
            ("Share capital", "Share capital and shareholding", "Agree authorised and paid-up capital, shares issued and shareholding pattern to the statutory registers and ROC filings; add the share capital details to the notes.",
             "Schedule III, Companies Act s.43 and s.62"),
            ("Reporting", "Companies (Auditor's Report) Order 2020", "Determine applicability of CARO 2020 and prepare the clause-by-clause report.", "CARO 2020"),
            ("Reporting", "Financial ratios and other Schedule III disclosures", "Review the analytical ratios, explain variations above 25 per cent, and complete the other Schedule III disclosures (ageing schedules, benami property, wilful defaulter, CSR).",
             "Schedule III"),
            ("Reporting", "Internal financial controls", "Assess whether reporting on internal financial controls over financial reporting applies.", "Companies Act s.143(3)(i), SA 265"),
            ("Reporting", "Audit report, UDIN and signing", "Finalise the report; obtain Board approval of the statements; generate the UDIN; sign in the order: Board, then auditor.",
             "SA 700, Companies Act s.134 and s.143"),
        ]
    else:
        I += [
            ("Capital", "Partners' / proprietor's capital and drawings", "Agree capital, drawings and profit sharing to the partnership deed and partners' confirmations.", "SA 500, SA 505"),
            ("Reporting", "Tax audit and reporting", "Check applicability of tax audit under section 44AB and Form 3CD; prepare the report; generate the UDIN.", "Income-tax Act s.44AB, SA 700"),
        ]
    return I


def make_audit_excel(checks, M, S, dep, tax, py, framework, entity, company, year_end, profile, df):
    from datetime import datetime
    from openpyxl.worksheet.datavalidation import DataValidation
    f_bold, f_head = Font(bold=True), Font(bold=True, color="FFFFFF")
    nav, grey = PatternFill("solid", fgColor=NAVY), PatternFill("solid", fgColor="F2F2F2")
    pri_fill = {"High": PatternFill("solid", fgColor="F8D7DA"), "Medium": PatternFill("solid", fgColor="FFF3CD"), "Low": PatternFill("solid", fgColor="E8F0FA")}
    wb = Workbook()
    ws = wb.active
    ws.title = "Auditor Action List"
    ws.sheet_view.showGridLines = False
    widths = {"A": 7, "B": 17, "C": 34, "D": 62, "E": 28, "F": 10, "G": 12, "H": 18, "I": 30}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws["A1"] = company
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = "Auditor action list for the year ended %s" % long_date(year_end)
    ws["A2"].font = Font(bold=True)
    ws["A3"] = ("Framework: %s | Prepared on %s from the application's checks and standard audit procedures. A working aid for the auditor; "
                "it does not form part of the financial statements. The auditor completes Status, Done by / date and Working paper / remarks for every line, "
                "signs off at the bottom, and uploads the file in the application before filing."
                % (framework, datetime.now().strftime("%d-%b-%Y")))
    ws["A3"].font = Font(size=9, italic=True)
    hdr = ["Ref", "Area", "Item", "Action for the auditor", "Reference", "Priority", "Status", "Done by / date", "Working paper / remarks"]
    r = 5

    def header(r):
        for i, t in enumerate(hdr):
            c = ws.cell(row=r, column=i + 1, value=t)
            c.font, c.fill = f_head, nav
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 30

    def section(r, title):
        for i in range(1, len(hdr) + 1):
            ws.cell(row=r, column=i).fill = grey
        ws.cell(row=r, column=1, value=title).font = f_bold
        return r + 1

    def item(r, ref, area, what, action, refer, pri):
        vals = [ref, area, what, action, refer, pri, "Open", "", ""]
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=i + 1, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=Side(style="hair"))
        ws.cell(row=r, column=6).fill = pri_fill[pri]
        ws.row_dimensions[r].height = max(30, 15 * (1 + max(len(action) // 58, len(what) // 32)))

    header(r)
    r += 1
    first = r
    n = 0
    # ---- A. from this run
    r = section(r, "A. Findings from this run of the application (specific to this client and year)")
    for name, status, detail in checks:
        if status == "pass" or name == AUDIT_CHECK:
            continue
        n += 1
        item(r, "A%d" % n, "Application check", "%s: %s" % (name, detail), _CHECK_ACTION.get(name, "Review and resolve."), "Checks and export tab",
             {"fail": "High", "warn": "Medium"}.get(status, "Low"))
        r += 1
    extra = []
    if M["status"] == "accepted":
        extra.append(("Closing stock", "Closing stock was derived from the books and accepted by the user, not taken from a stock valuation.",
                      "Obtain the stock valuation report and agree quantities and rates.", "High"))
    if entity == "Company" and not float((profile or {}).get("eq_shares") or 0):
        extra.append(("Share capital / EPS", "The number of equity shares was not entered, so earnings per share shows 'n.a.'.",
                      "Obtain the number of shares and shareholding pattern from the statutory registers and update the CA and signatories tab.", "Medium"))
    if py is None:
        extra.append(("Comparatives", "No previous year was supplied; comparatives are blank.", "Obtain the previous-year trial balance or signed statements and re-run.", "High"))
    has_dep_check = any(c_[0].startswith("Depreciation") for c_ in checks)
    if (not has_dep_check) and dep and dep.get("applied") is False and len(dep["table"]) and abs(dep["calc"] - dep["booked"]) >= 1:
        extra.append(("Depreciation", "Calculator depreciation Rs %s differs from Rs %s booked and has not been applied." % (inr(dep["calc"]), inr(dep["booked"])),
                      "Decide which figure is correct; if the calculator is right, post the difference in Tally.", "Medium"))
    if (not has_dep_check) and not (dep and len(dep["table"])):
        extra.append(("Depreciation", "The depreciation calculator was not used, so the Fixed Assets Schedule shows net block only.",
                      "Obtain the fixed-asset register and run the calculator to produce the gross block and depreciation schedule.", "Medium"))
    extra.append(("Ageing and MSME", "Ageing of receivables and payables and the MSME disclosure need bill-wise data, which the trial balance does not contain.",
                  "Obtain bill-wise outstanding from Tally, prepare the ageing schedules and MSME disclosure, and add them to the notes.", "Medium"))
    extra.append(("Loan terms", "Repayment terms, interest rates and security for borrowings are not available from Tally.",
                  "Obtain sanction letters and prepare the disclosure of terms and security.", "Medium"))
    for area, what, action, pri in extra:
        n += 1
        item(r, "A%d" % n, area, what, action, "Notes sheets", pri)
        r += 1
    # ---- B. standard
    r = section(r, "B. Standard audit procedures for this type of entity")
    for i, (area, what, action, refer) in enumerate(_standard_items(framework, entity), 1):
        item(r, "B%d" % i, area, what, action, refer, "Medium" if area in ("Revenue", "Receivables", "Inventory", "Fixed assets", "Reporting") else "Low")
        r += 1
    last = r - 1
    dv1 = DataValidation(type="list", formula1='"Open,In progress,Done,Not applicable"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add("G%d:G%d" % (first, last))
    dv2.add("F%d:F%d" % (first, last))
    ws.freeze_panes = "A6"
    r += 1
    ws.cell(row=r, column=3, value="Open items remaining").font = f_bold
    ws.cell(row=r, column=4, value='=COUNTIF(G%d:G%d,"Open")+COUNTIF(G%d:G%d,"In progress")' % (first, last, first, last))
    ws.cell(row=r + 2, column=3, value="Reviewed by (name, membership no.)").font = f_bold
    ws.cell(row=r + 3, column=3, value="Date of review").font = f_bold
    for k_ in (2, 3):
        c_ = ws.cell(row=r + k_, column=4)
        c_.fill = PatternFill("solid", fgColor="FFF9E5")
        c_.border = Border(bottom=Side(style="thin"), top=Side(style="thin"), left=Side(style="thin"), right=Side(style="thin"))

    # ---- Tool checks sheet
    w2 = wb.create_sheet("Application Checks")
    w2.sheet_view.showGridLines = False
    for k, v in {"A": 46, "B": 11, "C": 100}.items():
        w2.column_dimensions[k].width = v
    w2["A1"] = "Checks run by the application"
    w2["A1"].font = Font(size=13, bold=True)
    for i, t in enumerate(["Check", "Result", "Detail"]):
        c = w2.cell(row=3, column=i + 1, value=t)
        c.font, c.fill = f_head, nav
    colour = {"pass": "E3F4E8", "warn": "FFF3CD", "fail": "F8D7DA", "info": "E8F0FA"}
    for i, (name, status, detail) in enumerate(checks):
        for j, v in enumerate([name, {"pass": "Pass", "warn": "Review", "fail": "Fail", "info": "Info"}[status], detail]):
            c = w2.cell(row=4 + i, column=j + 1, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        w2.cell(row=4 + i, column=2).fill = PatternFill("solid", fgColor=colour[status])

    # ---- confirmations list
    w3 = wb.create_sheet("Confirmations List")
    w3.sheet_view.showGridLines = False
    for k, v in {"A": 40, "B": 24, "C": 26, "D": 18, "E": 8, "F": 16, "G": 16, "H": 16, "I": 30}.items():
        w3.column_dimensions[k].width = v
    w3["A1"] = "Balances for confirmation and vouching (largest first)"
    w3["A1"].font = Font(size=13, bold=True)
    h3 = ["Ledger", "Tally group", "Statement head", "Balance (Rs)", "Dr/Cr", "Request sent on", "Reply received on", "Difference (Rs)", "Remarks"]
    for i, t in enumerate(h3):
        c = w3.cell(row=3, column=i + 1, value=t)
        c.font, c.fill = f_head, nav
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    wanted = ["Trade Receivables", "Trade Payables", "Cash & Cash Equivalents", "Long-term Borrowings", "Short-term Borrowings",
              "Short-term Loans and Advances", "Other Non-Current Assets"]
    rr = 4
    for hd in wanted:
        sub = [(l, g, h, sg * b) for l, g, h, sg, b in M["rows"] if h == hd and abs(b) >= 0.5]
        sub.sort(key=lambda x: -abs(x[3]))
        for l, g, h, v in sub:
            nature = "Cr" if hd in CR_HEADS else "Dr"
            side = nature if v >= 0 else ("Dr" if nature == "Cr" else "Cr")
            for j, val in enumerate([l, g, h, abs(v), side]):
                w3.cell(row=rr, column=j + 1, value=val)
            w3.cell(row=rr, column=4).number_format = NF0
            w3.cell(row=rr, column=8, value="=IF(F%d=\"\",\"\",D%d-H%d)" % (rr, rr, rr)).number_format = NF0
            rr += 1
    w3.freeze_panes = "A4"
    for wsx in wb.worksheets:
        wsx.page_setup.orientation = "landscape"
        wsx.page_setup.paperSize = 9
        wsx.page_setup.fitToWidth, wsx.page_setup.fitToHeight = 1, 0
        wsx.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ================================================================ PDF for signing ==
def _fmt(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if abs(v) < 0.5:
        return "-"
    t = inr(abs(v))
    return "(%s)" % t if v < 0 else t


def make_pdf(M, S, py, framework, entity, company, year_end, dep, profile, tax):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (BaseDocTemplate, Frame, NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table,
                                    TableStyle, KeepTogether)
    has_py = py is not None
    start, prev_end = period_dates(year_end)
    d_fmt = lambda d: d.strftime("%d-%b-%Y")
    frame_txt = FRAME_TEXT[framework].split("|")[0].strip()
    navy = colors.HexColor("#" + NAVY)
    ps = lambda **k: ParagraphStyle("x", fontName=k.pop("fn", "Helvetica"), fontSize=k.pop("fs", 8.5), leading=k.pop("ld", 11), **k)
    S_T, S_H, S_C, S_N, S_NB, S_SM = ps(fn="Helvetica-Bold", fs=14, ld=17), ps(fn="Helvetica-Bold", fs=10.5, ld=13), ps(fs=8, ld=10, textColor=colors.grey), \
        ps(), ps(fn="Helvetica-Bold"), ps(fs=7.5, ld=9.5, textColor=colors.grey)
    pol_style = ps(fs=8.5, ld=11.5, alignment=4)
    pw, ph = A4
    lw, lh = landscape(A4)
    story = []

    def footer(canv, doc):
        canv.saveState()
        canv.setFont("Helvetica", 7.5)
        canv.setFillColor(colors.grey)
        w_, h_ = doc.pagesize
        canv.drawString(15 * mm, 9 * mm, company)
        canv.drawRightString(w_ - 15 * mm, 9 * mm, "Page %d" % doc.page)
        canv.restoreState()

    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=A4, title="%s - Financial Statements %s" % (company, long_date(year_end)), author=(profile or {}).get("firm", "") or "")
    fp = Frame(15 * mm, 16 * mm, pw - 30 * mm, ph - 30 * mm, id="p")
    fl = Frame(15 * mm, 16 * mm, lw - 30 * mm, lh - 30 * mm, id="l")
    doc.addPageTemplates([PageTemplate("portrait", [fp], onPage=footer, pagesize=A4),
                          PageTemplate("landscape", [fl], onPage=footer, pagesize=landscape(A4))])
    W = pw - 30 * mm

    def head(title, sub):
        story.append(Paragraph(company, S_T))
        story.append(Paragraph(title, S_H))
        story.append(Paragraph(sub, S_C))
        story.append(Spacer(1, 5))

    def numtable(rows, colw, bold_rows=(), grand_rows=(), sub_rows=(), hdr_rows=(0,), right_from=1):
        t = Table(rows, colWidths=colw, repeatRows=1)
        st = [("FONT", (0, 0), (-1, -1), "Helvetica", 8.5), ("VALIGN", (0, 0), (-1, -1), "TOP"),
              ("ALIGN", (right_from, 0), (-1, -1), "RIGHT"), ("TOPPADDING", (0, 0), (-1, -1), 2.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
              ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
              ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#D5D9E0"))]
        for r_ in bold_rows:
            st += [("FONT", (0, r_), (-1, r_), "Helvetica-Bold", 8.5), ("LINEABOVE", (right_from, r_), (-1, r_), 0.6, colors.black)]
        for r_ in grand_rows:
            st += [("FONT", (0, r_), (-1, r_), "Helvetica-Bold", 8.5), ("BACKGROUND", (0, r_), (-1, r_), colors.HexColor("#F2F2F2")),
                   ("LINEABOVE", (0, r_), (-1, r_), 0.8, colors.black), ("LINEBELOW", (0, r_), (-1, r_), 1.4, colors.black)]
        for r_ in sub_rows:
            st += [("FONT", (0, r_), (-1, r_), "Helvetica-BoldOblique", 8.5)]
        t.setStyle(TableStyle(st))
        return t

    def sig_block():
        ca, cl = sig_lines(profile, entity, company)
        n_ = max(len(ca), len(cl))
        if not n_:
            return Spacer(1, 1)
        rows = [[Paragraph(ca[i] if i < len(ca) else "", S_NB if i in (0, 1) else S_N), Paragraph(cl[i] if i < len(cl) else "", S_NB if i == 0 else S_N)]
                for i in range(n_)]
        t = Table(rows, colWidths=[W * 0.5, W * 0.5])
        t.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
        return KeepTogether([Spacer(1, 14), t])

    def statement(spec, vals, vals_py, title, sub_hdr, lead):
        head(title, frame_txt)
        if profile and (profile.get("cin") or profile.get("pan") or profile.get("address")):
            story.append(Paragraph(" | ".join(x for x in [("CIN / LLPIN: " + profile["cin"]) if profile.get("cin") else "",
                                                          ("PAN: " + profile["pan"]) if profile.get("pan") else "", profile.get("address", "")] if x), S_SM))
            story.append(Spacer(1, 3))
        rows = [["Particulars", "Note", "%s\n%s" % (lead, d_fmt(year_end)), "%s\n%s" % (lead, d_fmt(prev_end))]]
        bold, grand, sub = [], [], []
        for r in spec:
            k = r["kind"]
            if k == "blank":
                rows.append(["", "", "", ""])
                continue
            ind = "    " * r["indent"]
            lab = ind + r["label"]
            idx = len(rows)
            if k in ("section", "sub", "head"):
                rows.append([lab, "", "", ""])
                sub.append(idx)
                continue
            if k == "text":
                rows.append([lab, "", "n.a.", "n.a."])
                continue
            rows.append([lab, str(r["no"] or ""), _fmt(vals.get(r["key"], 0.0)), _fmt(vals_py.get(r["key"], 0.0)) if has_py else "-"])
            if k in ("total", "profit"):
                bold.append(idx)
            if k == "grand":
                grand.append(idx)
        t = numtable(rows, [W * 0.52, W * 0.08, W * 0.20, W * 0.20], bold, grand, sub)
        t.setStyle(TableStyle([("ALIGN", (1, 0), (1, -1), "CENTER"), ("ALIGN", (2, 0), (-1, -1), "RIGHT")]))
        story.append(t)
        story.append(sig_block())

    pl, bs = S["pl"], S["bs"]
    pyS = py["S"] if has_py else None
    statement(bs, S["bsv"], pyS["bsv"] if has_py else {}, "Balance Sheet as at %s" % long_date(year_end), "", "As at")
    story.append(PageBreak())
    statement(pl, S["plv"], pyS["plv"] if has_py else {}, "Statement of Profit and Loss for the year ended %s" % long_date(year_end), "", "Year ended")

    # ---------------- notes values
    def cells_of(Mx):
        out = {}
        for led, grp, hd, sg, bal in Mx["rows"]:
            out.setdefault((norm_name(led), hd), (led, _grp_label(grp), sg * bal))
        return out

    cyv = cells_of(M)
    pyv = cells_of(py["M"]) if has_py else {}
    ext_cy = {"closing": M["closing"] or 0.0, "tax": float(tax or 0.0)}
    ext_py = {"closing": (py["M"]["closing"] or 0.0) if has_py else 0.0, "tax": float(py["tax"]) if has_py else 0.0}

    def comp_vals(r, own_cy, own_py, ocv, ocp):
        out = []
        multi = len(r["terms"]) > 1
        for cf, k, ref in r["terms"]:
            pre = "Less: " if (cf < 0 and multi) else ""
            if k == "h":
                items = []
                for kk in sorted({x for x in list(cyv) + list(pyv) if x[1] == ref}):
                    a, b = cyv.get(kk), pyv.get(kk)
                    z = a or b
                    items.append((z[0], z[1], cf * a[2] if a else 0.0, cf * b[2] if b else 0.0, abs(a[2]) if a else 0))
                if ref == "Reserves and Surplus":
                    items.append(("Surplus in Statement of Profit and Loss: opening balance", "~", cf * M["or_cr"], cf * py["M"]["or_cr"] if has_py else 0.0, 0))
                if ref == "Inventories" and M["closing"] is not None:
                    items.append(("Closing stock", "~", cf * M["closing"], cf * (py["M"]["closing"] or 0.0) if has_py else 0.0, 0))
                items.sort(key=lambda x: (x[1] == "~", x[1], -x[4]))
                if pre and len(items) == 1:
                    items[0] = (pre + items[0][0],) + items[0][1:]
                groups = {x[1] for x in items if x[1] != "~"}
                lg = None
                for it in items:
                    if len(groups) > 1 and it[1] not in ("~", lg):
                        out.append(("sub", it[1]))
                        lg = it[1]
                    out.append(("row", it[0], it[2], it[3]))
            elif k == "p":
                out.append(("row", pre + "Profit for the year", cf * S["plv"][ref], cf * pyS["plv"][ref] if has_py else 0.0))
            elif k == "x":
                lab = {"tax": "Provision for income tax", "closing": "Closing stock"}.get(ref, ref)
                out.append(("row", pre + lab, cf * ext_cy[ref], cf * ext_py[ref]))
            else:
                out.append(("row", pre + ref, cf * ocv[ref], cf * ocp.get(ref, 0.0)))
        return out

    def note_table(r, comps, lead):
        rows = [["Particulars", "%s\n%s" % (lead, d_fmt(year_end)), "%s\n%s" % (lead, d_fmt(prev_end))]]
        sub, tc, tp = [], 0.0, 0.0
        any_row = False
        for c_ in comps:
            if c_[0] == "sub":
                sub.append(len(rows))
                rows.append([c_[1], "", ""])
            else:
                any_row = True
                rows.append(["    " + c_[1], _fmt(c_[2]), _fmt(c_[3]) if has_py else "-"])
                tc += c_[2]
                tp += c_[3]
        if not any_row:
            rows.append(["    No balances", "-", "-"])
        rows.append(["Total", _fmt(tc), _fmt(tp) if has_py else "-"])
        return numtable(rows, [W * 0.60, W * 0.20, W * 0.20], bold_rows=[len(rows) - 1], sub_rows=sub)

    # ---------------- accounting policies
    story.append(NextPageTemplate("portrait"))
    story.append(PageBreak())
    head("Notes forming part of the Financial Statements", "for the year ended %s" % long_date(year_end))
    story.append(Paragraph("A. General information", S_H))
    kind = {"Company": "a company incorporated under the Companies Act", "Firm (partnership or LLP)": "a partnership firm / limited liability partnership",
            "Proprietorship": "a proprietorship concern"}[entity]
    gi = "%s is %s." % (company, kind)
    if (profile or {}).get("cin"):
        gi += " CIN / LLPIN: %s." % profile["cin"]
    if (profile or {}).get("address"):
        gi += " Registered office / place of business: %s." % profile["address"]
    gi += " These financial statements are for the year ended %s." % long_date(year_end)
    story.append(Paragraph(gi, pol_style))
    story.append(Spacer(1, 6))
    story.append(Paragraph("B. Significant accounting policies", S_H))
    for t_, tx in _policies(framework, entity, ""):
        story.append(Spacer(1, 3))
        story.append(Paragraph("<b>%s</b>" % t_, S_N))
        story.append(Paragraph(tx, pol_style))
    notes = sorted([r for r in pl + bs if r["no"]], key=lambda r: r["no"])
    ppe_done = False

    def fixed_asset_pages():
        story.append(NextPageTemplate("landscape"))
        story.append(PageBreak())
        head("Property, Plant and Equipment", "Schedule as at %s" % long_date(year_end))
        t = dep["table"]
        res = dep["result"]
        classes = []
        for v in t["Asset class"].tolist():
            v = str(v or "Other")
            if v not in classes:
                classes.append(v)
        rows = [["Asset class", "Gross block\nopening", "Additions", "Deductions", "Gross block\nclosing", "Depreciation\nopening", "For the year",
                 "On deductions", "Depreciation\nclosing", "Net block\nclosing", "Net block\nopening"]]
        tot = [0.0] * 10
        for cl in classes:
            m_ = (t["Asset class"].fillna("Other").astype(str) == cl).values
            og, ad, de = float(t.loc[m_, "Opening gross block"].sum()), float(t.loc[m_, "Additions (cost)"].sum()), float(t.loc[m_, "Disposals (original cost)"].sum())
            oa = float(t.loc[m_, "Opening accumulated depreciation"].sum())
            rs = res[m_] if len(res) == len(t) else res[res["Asset class"] == cl]
            cg, ca_, yr = float(rs["Closing gross block"].sum()), float(rs["Closing accumulated depreciation"].sum()), float(rs["Depreciation for the year"].sum())
            nb = float(rs["Closing net block"].sum())
            vals = [og, ad, de, cg, oa, yr, oa + yr - ca_, ca_, nb, og - oa]
            tot = [a + b for a, b in zip(tot, vals)]
            rows.append([cl] + [_fmt(v) for v in vals])
        rows.append(["Total"] + [_fmt(v) for v in tot])
        cw = [lw * 0 + (lw - 30 * mm) * 0.19] + [(lw - 30 * mm) * 0.081] * 10
        story.append(numtable(rows, cw, bold_rows=[len(rows) - 1]))
        story.append(Spacer(1, 5))
        story.append(Paragraph("Depreciation for the year includes depreciation on assets sold up to the date of sale; it is removed with the accumulated "
                               "depreciation of those assets under 'On deductions'. " + dep["note"], S_SM))
        story.append(NextPageTemplate("portrait"))

    def emit(group, title_, lead_is_bs):
        story.append(PageBreak())
        head(title_, "for the year ended %s" % long_date(year_end))
        for r in group:
            key = r["key"]
            own_cy = S["bsv"] if r in bs else S["plv"]
            own_py = (pyS["bsv"] if r in bs else pyS["plv"]) if has_py else {}
            blk = [Spacer(1, 8), Paragraph("%d.  %s" % (r["no"], _clean(r["label"])), S_H), Spacer(1, 2)]
            if key == "eps":
                prof = profile or {}
                sh, fv, pf = float(prof.get("eq_shares") or 0), float(prof.get("face_value") or 10), float(prof.get("pref_div") or 0)
                a, b = S["plv"]["pat"], (pyS["plv"]["pat"] if has_py else 0.0)
                eps = lambda x: ("%.2f" % ((x - pf) / sh)) if sh > 0 else "n.a."
                rows = [["Particulars", "Year ended\n%s" % d_fmt(year_end), "Year ended\n%s" % d_fmt(prev_end)],
                        ["Profit for the year", _fmt(a), _fmt(b) if has_py else "-"], ["Less: Preference dividend", _fmt(pf), _fmt(pf) if has_py else "-"],
                        ["Earnings attributable to equity shareholders", _fmt(a - pf), _fmt(b - pf) if has_py else "-"],
                        ["Number of equity shares outstanding", "{:,.0f}".format(sh), "{:,.0f}".format(sh) if has_py else "-"],
                        ["Face value per equity share (Rs)", "%.2f" % fv, "%.2f" % fv if has_py else "-"],
                        ["Basic and diluted earnings per share (Rs)", eps(a), eps(b) if has_py else "-"]]
                blk.append(numtable(rows, [W * 0.60, W * 0.20, W * 0.20], bold_rows=[3, 6]))
            else:
                comps = comp_vals(r, own_cy, own_py, own_cy, own_py)
                blk.append(note_table(r, comps, "As at" if r in bs else "Year ended"))
                if key == "ppe" and dep and dep["basis"] == "sch2" and len(dep["table"]):
                    blk.append(Paragraph("Refer the Fixed Assets Schedule for gross block, depreciation and net block.", S_SM))
            story.append(KeepTogether(blk))
            if key == "ppe" and dep and dep["basis"] == "sch2" and len(dep["table"]):
                nonlocal_flag.append(1)

    nonlocal_flag = []
    emit([r for r in notes if r in bs], "Notes to the Balance Sheet", True)
    if dep and dep["basis"] == "sch2" and len(dep["table"]):
        fixed_asset_pages()
    emit([r for r in notes if r in pl], "Notes to the Statement of Profit and Loss", False)

    # ---------------- ratios (companies)
    if entity == "Company":
        B = lambda k, p=False: (pyS["bsv"] if p else S["bsv"])[k]
        P = lambda k, p=False: (pyS["plv"] if p else S["plv"])[k]

        def rat(p):
            av = (lambda k: (B(k) + B(k, True)) / 2) if (has_py and not p) else (lambda k, p=p: B(k, p))
            wc = ((B("ca") - B("cl")) + (B("ca", True) - B("cl", True))) / 2 if (has_py and not p) else (B("ca", p) - B("cl", p))
            dv = lambda a, b, m=1: (a / b * m) if b else None
            return [dv(B("ca", p), B("cl", p)), dv(B("ltb", p) + B("stb", p), B("fund", p)), dv(P("pat", p), av("fund"), 100),
                    dv(P("pur", p) + P("chg", p), av("stock")), dv(P("rev", p), av("tr")), dv(P("pur", p), av("tp")), dv(P("rev", p), wc),
                    dv(P("pat", p), P("rev", p), 100), dv(P("pbt", p) + P("fin", p), B("fund", p) + B("ltb", p) + B("stb", p), 100)]
        names = [("Current ratio (times)", "Current assets", "Current liabilities"), ("Debt-equity ratio (times)", "Total borrowings", "Shareholders' funds"),
                 ("Return on equity (%)", "Profit for the year", "Average shareholders' funds"), ("Inventory turnover (times)", "Cost of goods sold", "Average inventories"),
                 ("Trade receivables turnover (times)", "Revenue from operations", "Average trade receivables"),
                 ("Trade payables turnover (times)", "Purchases", "Average trade payables"), ("Net capital turnover (times)", "Revenue from operations", "Average working capital"),
                 ("Net profit ratio (%)", "Profit for the year", "Revenue from operations"), ("Return on capital employed (%)", "Profit before tax + finance costs", "Capital employed")]
        c_, p_ = rat(False), (rat(True) if has_py else [None] * 9)
        rows = [["Ratio", "Numerator", "Denominator", "Current\nyear", "Previous\nyear", "% variance"]]
        f2 = lambda x: "n.a." if x is None else "%.2f" % x
        for (n_, nu, de), a, b in zip(names, c_, p_):
            var = "-" if (not has_py or a is None or not b) else "%.1f" % ((a - b) / abs(b) * 100)
            rows.append([n_, nu, de, f2(a), f2(b) if has_py else "-", var])
        story.append(Spacer(1, 8))
        story.append(KeepTogether([Paragraph("%d.  Analytical ratios" % (len(notes) + 1), S_H), Spacer(1, 2),
                                   numtable(rows, [W * 0.28, W * 0.22, W * 0.22, W * 0.09, W * 0.09, W * 0.10], right_from=3)]))
    doc.build(story)
    return buf.getvalue()


# ======================================================== auditor responses / closure ==
import re as _re


def read_audit_responses(raw):
    """Read a completed Auditor Action List. Returns (rows DataFrame, reviewer, review_date, header text)."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    if "Auditor Action List" not in wb.sheetnames:
        raise ValueError("This is not an Auditor Action List produced by the application (sheet 'Auditor Action List' not found).")
    ws = wb["Auditor Action List"]
    rows, reviewer, rdate = [], "", ""
    for r in ws.iter_rows(min_row=1, values_only=True):
        v = list(r) + [None] * 9
        ref = str(v[0] or "").strip()
        if _re.match(r"^[AB]\d+$", ref):
            rows.append({"Ref": ref, "Area": v[1] or "", "Item": v[2] or "", "Priority": str(v[5] or "").strip(),
                         "Status": str(v[6] or "").strip(), "Done by / date": str(v[7] or "").strip(), "Remarks": str(v[8] or "").strip()})
        if str(v[2] or "").startswith("Reviewed by"):
            reviewer = str(v[3] or "").strip()
        if str(v[2] or "").startswith("Date of review"):
            rdate = str(v[3] or "").strip()
    head = "%s | %s" % (ws["A1"].value or "", ws["A2"].value or "")
    return pd.DataFrame(rows, columns=["Ref", "Area", "Item", "Priority", "Status", "Done by / date", "Remarks"]), reviewer, rdate, head


def assess_responses(resp, reviewer, rdate, current_refs, current_head, file_head):
    """Decide whether the auditor's responses are complete and clean."""
    ex = []
    for _, r in resp.iterrows():
        st_ = r["Status"]
        if st_ in ("", "Open"):
            ex.append((r["Ref"], r["Item"], "Not answered (status is %s)" % (st_ or "blank"), r["Priority"]))
        elif st_ == "In progress":
            ex.append((r["Ref"], r["Item"], "Still in progress", r["Priority"]))
        elif st_ == "Done":
            miss = [k for k in ("Done by / date", "Remarks") if not r[k]]
            if miss:
                ex.append((r["Ref"], r["Item"], "Marked Done but %s not filled" % " and ".join(miss), r["Priority"]))
        elif st_ == "Not applicable":
            if not r["Remarks"]:
                ex.append((r["Ref"], r["Item"], "Marked Not applicable without a reason in remarks", r["Priority"]))
        else:
            ex.append((r["Ref"], r["Item"], "Unrecognised status '%s'" % st_, r["Priority"]))
    got = set(resp["Ref"])
    missing = sorted(current_refs - got, key=lambda x: (x[0], int(x[1:])))
    for m in missing:
        ex.append((m, "(new item since the list was issued)", "Not in the uploaded file; the data has changed. Download the list again", "High"))
    problems = []
    if not reviewer:
        problems.append("Reviewer name is not filled at the bottom of the list")
    if not rdate:
        problems.append("Date of review is not filled")
    if current_head and file_head and current_head.split("|")[0].strip() != file_head.split("|")[0].strip():
        problems.append("The file is for a different entity ('%s'), not '%s'" % (file_head.split("|")[0].strip(), current_head.split("|")[0].strip()))
    n = len(resp)
    closed = int(((resp["Status"] == "Done") | (resp["Status"] == "Not applicable")).sum()) - len([e for e in ex if "Marked" in e[2]])
    return {"total": n, "closed": max(closed, 0), "exceptions": ex, "problems": problems,
            "clean": (not ex) and (not problems) and n > 0, "reviewer": reviewer, "date": rdate}


_SAMPLE_REMARKS = {
    "Application check": "Reviewed the check and the supporting ledgers; difference explained and agreed to books. WP-A",
    "Depreciation": "Agreed additions to purchase invoices and the fixed-asset register; depreciation recomputed under Schedule II. WP-FA1",
    "Ageing and MSME": "Bill-wise outstanding obtained from Tally; ageing and MSME status prepared; no MSME interest payable. WP-AG1",
    "Loan terms": "Sanction letter obtained; repayment terms and security disclosed in the notes. WP-BR1",
    "Share capital / EPS": "Share register agreed; number of shares entered in the application. WP-SC1",
    "Comparatives": "Previous-year closing agreed to signed statements. WP-PY1",
    "Closing stock": "Stock valuation and count report agreed to closing stock. WP-INV1",
}


def sample_audit_answers(raw, with_gaps=False):
    """Fill an Auditor Action List with sample answers for the demonstration."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Auditor Action List"]
    k = 0
    last_row = 0
    for r in range(1, ws.max_row + 1):
        ref = str(ws.cell(row=r, column=1).value or "")
        if _re.match(r"^[AB]\d+$", ref):
            k += 1
            area = str(ws.cell(row=r, column=2).value or "")
            ws.cell(row=r, column=7, value="Done")
            ws.cell(row=r, column=8, value="A. Sample / %s" % ("27-Oct-2026" if k % 3 else "28-Oct-2026"))
            ws.cell(row=r, column=9, value=_SAMPLE_REMARKS.get(area, "Procedure performed as per %s; no exceptions noted. WP-%s%d" % (area.lower(), ref[0], k)))
            last_row = r
        v = str(ws.cell(row=r, column=3).value or "")
        if v.startswith("Reviewed by"):
            ws.cell(row=r, column=4, value="A. Sample, M.No. 000000")
        if v.startswith("Date of review"):
            ws.cell(row=r, column=4, value="28-Oct-2026")
    if with_gaps:
        n = 0
        for r in range(1, last_row + 1):
            ref = str(ws.cell(row=r, column=1).value or "")
            if _re.match(r"^[AB]\d+$", ref):
                n += 1
                if n == 3:
                    ws.cell(row=r, column=7, value="Open")
                    ws.cell(row=r, column=8, value="")
                    ws.cell(row=r, column=9, value="")
                if n == 5:
                    ws.cell(row=r, column=9, value="")
                if n == 8:
                    ws.cell(row=r, column=7, value="In progress")
                if n == 11:
                    ws.cell(row=r, column=7, value="Not applicable")
                    ws.cell(row=r, column=9, value="")
        for r in range(1, ws.max_row + 1):
            if str(ws.cell(row=r, column=3).value or "").startswith("Date of review"):
                ws.cell(row=r, column=4).value = None
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ===================================================================== checks ==
def clean_dep(tbl, kind):
    t = tbl.copy()
    cols = S2_COLS if kind == "s2" else IT_COLS
    t[cols[0]] = t[cols[0]].fillna("").astype(str)
    t[cols[1]] = t[cols[1]].fillna("").astype(str)
    nums = S2_NUM if kind == "s2" else IT_COLS[2:]
    for c in nums:
        t[c] = pd.to_numeric(t[c], errors="coerce").fillna(0.0)
    if kind == "s2":
        for c in S2_DATES:
            t[c] = [_date(v) for v in t[c]]
    keep_cols = S2_NUM[2:] if kind == "s2" else IT_COLS[3:]
    keep = (t[cols[0]].str.strip() != "") | (t[keep_cols].abs().sum(axis=1) > 0)
    return t[keep].reset_index(drop=True)


def run_checks(df, M, S, dep, tax, py=None):
    out = []
    total = float(df["Dr_pos"].sum())
    out.append(("Trial balance as pulled", "pass" if abs(total) < 1 else "info",
                "Debits less credits = %s%s" % (inr(total), "" if abs(total) < 1 else
                 ". Expected when Tally's Profit & Loss A/c already holds this year's profit; see the P&L A/c check.")))
    out.append(("All ledgers classified", "pass" if M["unclassified"] == 0 else "fail",
                "%d ledger(s) without a head" % M["unclassified"]))
    d = S["diff"]
    out.append(("Balance Sheet balances", "pass" if abs(d) < 1 else "fail",
                "Assets %s | Equity and liabilities %s | Difference %s" % (inr(S["bsv"]["tas"]), inr(S["bsv"]["tel"]), inr(d))))
    st_ = M["status"]
    if M["plac_cr"] is None:
        out.append(("Tally Profit & Loss A/c tie-out", "info", "No Profit & Loss A/c ledger found in the pull."))
    elif st_ == "unclassified":
        out.append(("Tally Profit & Loss A/c tie-out", "warn", "Classify all ledgers first."))
    elif st_ == "pending":
        out.append(("Tally Profit & Loss A/c tie-out", "fail", M["msg"]))
    else:
        exp = M["or_cr"] + S["plv"]["pbt"] - M["ht"].get("Tax Expense", 0.0) + M["adj"]
        ok = abs(M["plac_cr"] - exp) < 1
        ok_open = abs(M["plac_cr"] - M["or_cr"]) < 1
        out.append(("Tally Profit & Loss A/c tie-out", "pass" if (ok or ok_open) else "fail",
                    ("Agrees to opening reserves + profit before tax" if ok else
                     "Agrees to opening reserves only" if ok_open else
                     "Does not agree: Tally %s vs %s" % (inr(M["plac_cr"]), inr(exp)))))
    cs = {"verified": ("pass", "Closing stock %s agrees to Tally's stock items." % inr(M["closing"] or 0)),
          "given": ("pass", M["msg"]),
          "accepted": ("warn", "Closing stock %s derived from the books and accepted by you. Confirm with the Stock Summary." % inr(M["closing"] or 0)),
          "pending": ("fail", "Closing stock not resolved. Tick 'Accept derived closing stock' after comparing with the Stock Summary."),
          }.get(st_, ("info", "Stock taken as per ledger balances; no separate closing stock value needed."))
    out.append(("Closing stock", cs[0], cs[1]))
    if dep and dep["applied"]:
        ppe = S["bsv"].get("ppe", 0.0)
        diff = dep["net"] - ppe
        out.append(("Fixed assets agree to depreciation calculator", "pass" if abs(diff) < 1 else "warn",
                    "Calculator net block %s | Statements %s | Difference %s%s" % (
                        inr(dep["net"]), inr(ppe), inr(diff),
                        "" if abs(diff) < 1 else ". Check opening values, additions and ledgers not in the table.")))
    elif dep and len(dep["table"]):
        out.append(("Depreciation calculator", "info", "Calculated %s against %s booked. Not applied to the statements." % (
            inr(dep["calc"]), inr(dep["booked"]))))
    else:
        out.append(("Depreciation calculator", "info", "Not used. Depreciation is as booked in Tally."))
    if dep and dep.get("unexplained") is not None and len(dep["table"]):
        u = dep["unexplained"]
        out.append(("Depreciation inputs agree to Tally fixed assets", "pass" if abs(u) < 1 else "warn",
                    "Unexplained difference %s%s" % (inr(u), "" if abs(u) < 1 else
                     ". Opening values, additions or disposals entered do not explain the Tally balance.")))
    if M.get("reclass_amt"):
        out.append(("Opposite balances shown gross", "info",
                    "Rs %s of balances in the opposite direction to their head were reclassified (see '[reclassified]' in the mapping sheet)." % inr(M["reclass_amt"])))
    out.append(("Income-tax provision", "pass" if tax else "warn",
                ("Rs %s entered" % inr(tax)) if tax else "Nil. Profit is before tax; enter the provision in the sidebar once computed."))
    if py:
        Mp, Sp = py["M"], py["S"]
        out.append(("Previous year: all ledgers classified", "pass" if Mp["unclassified"] == 0 else "fail",
                    "%d ledger(s) without a head" % Mp["unclassified"]))
        n_sim = int(py["cls"]["Source"].str.startswith("Similar").sum())
        out.append(("Previous year: classification source", "pass" if n_sim == 0 else "warn",
                    "%d ledger(s) classified by similar name; confirm them in the Previous year tab." % n_sim if n_sim else
                    "Every ledger matched by name, group rule or last year's workbook."))
        out.append(("Previous year: Balance Sheet balances", "pass" if abs(Sp["diff"]) < 1 else "fail",
                    "Difference %s%s" % (inr(Sp["diff"]), "" if Mp["status"] != "pending" else ". Closing stock for the previous year is not resolved.")))
        cur_open = S["bsv"]["fund"] - S["plv"].get("tci", S["plv"]["pat"])
        rf = cur_open - Sp["bsv"]["fund"]
        out.append(("Opening equity agrees to previous year closing", "pass" if abs(rf) < 1 else "warn",
                    "Current-year opening %s | Previous-year closing %s | Difference %s%s" % (
                        inr(cur_open), inr(Sp["bsv"]["fund"]), inr(rf), "" if abs(rf) < 1 else
                        ". Explain: capital introduced or withdrawn, dividends, prior-period items or reclassification.")))
        nd = len(py["diff"])
        out.append(("Same ledgers classified the same way in both years", "pass" if nd == 0 else "warn",
                    "%d ledger(s) carry a different head in the two years." % nd))
    return out


def pills(checks):
    from html import escape as esc
    lab = {"pass": "PASS", "fail": "FAIL", "warn": "REVIEW", "info": "INFO"}
    return "".join('<div class="chk"><span class="pill %s">%s</span><b>%s</b><span class="d">%s</span></div>'
                   % (s, lab[s], esc(n), esc(d)) for n, s, d in checks)


PROFILE_FILE = "client_profile.json"


def _profile_path():
    import os
    try:
        return os.path.join(os.path.dirname(os.path.abspath(__file__)), PROFILE_FILE)
    except Exception:
        return PROFILE_FILE


def load_profile():
    import json
    p = default_profile()
    try:
        with open(_profile_path(), encoding="utf-8") as fh:
            p.update(json.load(fh))
    except Exception:
        pass
    return p


def save_profile(p):
    import json
    with open(_profile_path(), "w", encoding="utf-8") as fh:
        json.dump(p, fh, indent=2)


def load_sample():
    """Fill the application with synthetic data so it can be demonstrated without Tally."""
    ss = st.session_state
    tb, note, cs = sample_trial_balance(False)
    ss["tb"], ss["note"], ss["cs_items"] = tb, note, cs
    ss["groups"] = tb.attrs.get("groups", {})
    ss["ver"] = ss.get("ver", 0) + 1
    d0, n0, cs0 = sample_trial_balance(True)
    ss["py_tally"] = {"df": d0, "closing": None, "cs_items": cs0, "tax": 0.0, "mode": "tally", "note": "Synthetic previous-year data."}
    ss["pyver"] = ss.get("pyver", 0) + 1
    ss["py_src"] = "Tally (previous period)"
    ss["company_name"] = SAMPLE_NAME
    ss["profile"] = sample_profile()
    ss["sample_mode"] = True


# ======================================================================== app ==
def main():
    st.set_page_config(page_title="Financial Statements from Tally", layout="wide")
    st.markdown(CSS, unsafe_allow_html=True)
    st.markdown('<p class="app-title">Financial Statements from Tally</p>'
                '<p class="app-sub">Trial balance, classification, previous year, depreciation, Balance Sheet, Profit and Loss, notes and Excel. '
                'Data stays on this computer; the app talks only to Tally on localhost.</p>', unsafe_allow_html=True)
    ss = st.session_state
    PYSRC = ["None", "Tally (previous period)", "Upload trial balance", "Upload last year's workbook"]

    # ------------------------------------------------------------------ sidebar
    with st.sidebar:
        st.markdown("### 1. Connect to Tally")
        url = st.text_input("Tally address", DEFAULT_URL)
        if st.button("Find companies in Tally"):
            try:
                ss["companies"] = [c["NAME"] for c in tally_collection(url, "Company", ["Name"]) if c.get("NAME")]
            except Exception as e:
                st.warning("Could not list companies (%s). Leave the company as '(company open in Tally)'." % e)
        choice = st.selectbox("Company", ["(company open in Tally)"] + ss.get("companies", []))
        st.button("Load sample company (synthetic data, no Tally needed)", on_click=load_sample,
                  help="Loads an invented company with two years of data so the whole application can be demonstrated.")
        st.markdown("### 2. Reporting basis")
        framework = st.selectbox("Reporting framework", FRAMEWORKS,
                                 help="AS and Ind AS are for companies (Schedule III). Non-corporate is for firms and proprietorships.")
        if framework == FRAMEWORKS[2]:
            entity = st.radio("Entity type", ["Firm (partnership or LLP)", "Proprietorship"])
        else:
            entity = "Company"
            st.caption("Entity type: Company")
        if not ss.get("company_name") and not choice.startswith("("):
            ss["company_name"] = choice
        company_name = st.text_input("Name for the statements", key="company_name")
        year_end = st.date_input("Financial year ends on", value=date(2026, 3, 31))
        st.markdown("### 3. Previous year")
        py_src = st.radio("Previous-year source", PYSRC, key="py_src",
                          help="Use the upload options when last year's data is in a different Tally company or file.")
        py_up, py_cs_in, py_accept, py_tax_in, prefer = None, 0.0, False, 0.0, True
        if py_src == PYSRC[1]:
            if st.button("Pull previous year"):
                try:
                    with st.spinner("Reading the previous period from Tally..."):
                        d0, n0, cs0 = load_trial_balance(url, "" if choice.startswith("(") else choice, (period_dates(period_dates(year_end)[1])[0], period_dates(year_end)[1]))   # previous year: 1-Apr-2024 to 31-Mar-2025
                    ss["py_tally"] = {"df": d0, "closing": None, "cs_items": cs0, "tax": 0.0, "mode": "tally", "note": n0}
                    ss["pyver"] = ss.get("pyver", 0) + 1
                except requests.exceptions.ConnectionError:
                    st.error("Could not reach Tally. Open Tally and try again.")
                except Exception as e:
                    st.error("Previous period could not be read: %s" % e)
        elif py_src == PYSRC[2]:
            f = st.file_uploader("Previous-year trial balance (.xlsx or .csv)", type=["xlsx", "csv"], key="py_file_tb")
            if f is not None:
                key = ("tb", f.name, f.size)
                if ss.get("py_key") != key:
                    try:
                        d0, n0 = parse_tb_upload(f.getvalue(), f.name)
                        ss["py_up"] = {"df": d0, "closing": None, "cs_items": None, "tax": 0.0, "mode": "upload", "note": n0}
                        ss["py_key"] = key
                        ss["pyver"] = ss.get("pyver", 0) + 1
                    except Exception as e:
                        ss.pop("py_up", None)
                        st.error(str(e))
        elif py_src == PYSRC[3]:
            f = st.file_uploader("Last year's Excel produced by this application", type=["xlsx"], key="py_file_wb")
            if f is not None:
                key = ("wb", f.name, f.size, str(year_end))
                if ss.get("py_key") != key:
                    try:
                        d0, c0, t0, n0 = parse_workbook(f.getvalue(), period_dates(year_end)[1])
                        ss["py_wb"] = {"df": d0, "closing": c0, "cs_items": None, "tax": t0, "mode": "workbook", "note": n0}
                        ss["py_key"] = key
                        ss["pyver"] = ss.get("pyver", 0) + 1
                    except Exception as e:
                        ss.pop("py_wb", None)
                        st.error(str(e))
            prefer = st.checkbox("Use last year's classification for matching current-year ledgers", value=True)
        if py_src != PYSRC[0]:
            py_cs_in = st.number_input("Closing stock at previous year-end (Rs), if known (not needed for workbooks made by this app)", min_value=0.0, value=0.0, step=1000.0,
                                       help="From last year's signed accounts. Leave 0 if stock is inside the ledgers.")
            py_accept = st.checkbox("Accept derived previous-year closing stock", value=False)
            py_tax_in = st.number_input("Previous-year income-tax provision not in ledgers (Rs)", min_value=0.0, value=0.0, step=1000.0)
        st.markdown("### 4. Adjustments")
        accept = st.checkbox("Accept derived closing stock", value=False, key="accept_cs",
                             help="Use only after comparing with the closing stock in Tally's Stock Summary.")
        gross = st.checkbox("Show opposite balances gross", value=True,
                            help="Schedule III: debtors with credit balances go to liabilities, creditors with debit balances to advances, bank overdrafts to borrowings.")
        tax = st.number_input("Income-tax provision for the year (Rs)", min_value=0.0, value=0.0, step=1000.0,
                              help="Enter the figure from your tax computation. The app does not compute tax.")
        c1, c2 = st.columns(2)
        run = c1.button("Pull from Tally", type="primary")
        reset = c2.button("Reset")

    if reset:
        keep = {k: ss[k] for k in ("profile",) if k in ss}
        ss.clear()
        ss.update(keep)
        st.rerun()
    if run:
        try:
            with st.spinner("Reading ledgers from Tally..."):
                tb, note, cs_items = load_trial_balance(url, "" if choice.startswith("(") else choice)
            ss["tb"], ss["note"], ss["cs_items"] = tb, note, cs_items
            ss["groups"] = tb.attrs.get("groups", {})
            ss["ver"] = ss.get("ver", 0) + 1
        except requests.exceptions.ConnectionError:
            st.error("Could not reach Tally at %s. Open Tally Prime, load the company, then press F1 (Help), Settings, "
                     "Connectivity, Client/Server configuration, and set the port to 9000." % url)
            return
        except Exception as e:
            st.error("Something went wrong while reading Tally: %s" % e)
            return
    if "profile" not in ss:
        ss["profile"] = load_profile()
    prof = ss["profile"]
    if ss.get("sample_mode"):
        st.info("Sample mode: the company, ledgers, previous year and signatories are invented for demonstration. "
                "Nothing is read from Tally and nothing real is stored.")

    tabs = st.tabs(["Overview", "Trial balance", "Classification", "Previous year", "Depreciation", "Statements",
                    "CA and signatories", "Checks and export"])
    t_ov, t_tb, t_cl, t_py, t_dp, t_st, t_ca, t_ck = tabs

    # ------------------------------------------- CA and signatories (needs no data)
    with t_ca:
        st.subheader("CA firm and client signatories")
        st.caption("These details appear on the Balance Sheet and Profit and Loss sheets of the Excel file. They are saved only in a small "
                   "file on this computer (client_profile.json beside the app). The UDIN is generated on the UDIN portal; type it here.")
        a, b = st.columns(2)
        with a:
            st.markdown("**Chartered Accountant (auditor / reporting)**")
            prof["firm"] = st.text_input("Firm name", prof["firm"])
            prof["frn"] = st.text_input("Firm Registration Number (FRN)", prof["frn"])
            prof["partner"] = st.text_input("Partner / proprietor name", prof["partner"])
            prof["mno"] = st.text_input("Membership number", prof["mno"])
            prof["udin"] = st.text_input("UDIN", prof["udin"])
            prof["place"] = st.text_input("Place", prof["place"])
            prof["rdate"] = st.text_input("Date of report", prof["rdate"], placeholder="e.g. 28-Oct-2026")
        with b:
            st.markdown("**Client**")
            prof["cin"] = st.text_input("CIN / LLPIN (if any)", prof["cin"])
            prof["pan"] = st.text_input("PAN", prof["pan"])
            prof["address"] = st.text_area("Registered office / address", prof["address"], height=80)
            prof["bdate"] = st.text_input("Date of approval by Board / Partners", prof["bdate"], placeholder="e.g. 28-Oct-2026")
            st.markdown("**For the notes (companies)**")
            prof["eq_shares"] = st.number_input("Number of equity shares outstanding", min_value=0, value=int(prof.get("eq_shares") or 0), step=1000,
                                                help="Used for earnings per share in the notes.")
            prof["face_value"] = st.number_input("Face value per equity share (Rs)", min_value=0.0, value=float(prof.get("face_value") or 10), step=1.0)
            prof["pref_div"] = st.number_input("Preference dividend for the year (Rs)", min_value=0.0, value=float(prof.get("pref_div") or 0), step=1000.0)
        st.markdown("**Client signatories** (directors, partners or proprietor)")
        sg = pd.DataFrame(prof["signs"] if prof["signs"] else default_profile()["signs"]).reindex(columns=["name", "id", "desig"]).fillna("")
        sg = sg.rename(columns={"name": "Name", "id": "DIN / PAN", "desig": "Designation"})
        sg_ed = st.data_editor(sg, num_rows="dynamic", hide_index=True, key="signs_ed")
        prof["signs"] = [{"name": str(r["Name"] or ""), "id": str(r["DIN / PAN"] or ""), "desig": str(r["Designation"] or "")}
                         for _, r in sg_ed.iterrows()]
        if ss.get("sample_mode"):
            st.warning("These are sample details for the demonstration. Saving is switched off so they cannot replace your real details. "
                       "Press Reset in the sidebar to go back to your own details.")
        if st.button("Save these details on this computer", disabled=bool(ss.get("sample_mode"))):
            try:
                save_profile(prof)
                st.success("Saved.")
            except Exception as e:
                st.error("Could not save: %s" % e)

    if "tb" not in ss:
        with t_ov:
            st.info("Open the company in Tally, choose the reporting framework in the sidebar, then click **Pull from Tally**. "
                    "To see the whole application without Tally, click **Load sample company** in the sidebar.")
            st.markdown("**How it works**\n\n1. Pull the trial balance from Tally.\n2. Add the previous year (Tally, an uploaded trial balance, or last year's workbook).\n"
                        "3. Check the classification and fix any ledger.\n4. Run the depreciation calculator (Schedule II or Income Tax Act).\n"
                        "5. Fill in the CA and signatory details.\n6. Review the statements and checks, then download the Excel file.")
        for t in (t_tb, t_cl, t_py, t_dp, t_st, t_ck):
            with t:
                st.caption("Pull the trial balance from Tally first.")
        return

    base = ss["tb"].copy()
    ver = ss.get("ver", 0)
    if ss.get("note"):
        st.caption(ss["note"])

    # ----------------------------------------------------- previous-year raw data
    raw = {PYSRC[1]: ss.get("py_tally"), PYSRC[2]: ss.get("py_up"), PYSRC[3]: ss.get("py_wb")}.get(py_src)
    pyver = ss.get("pyver", 0)
    base["Source"] = "Group rule"
    if raw is not None and raw["mode"] == "workbook" and prefer:
        memo = {norm_name(l): h for l, h in zip(raw["df"]["Ledger"], raw["df"]["Head"]) if h}
        hit = base["Ledger"].map(lambda l: memo.get(norm_name(l)))
        use = hit.notna() & ~is_plac(base["Ledger"])
        base.loc[use, "Head"] = hit[use]
        base.loc[use, "Source"] = "Last year's workbook"

    # ------------------------------------------------------- classification tab
    plac_mask = is_plac(base["Ledger"])
    cls_src = base[~plac_mask].copy()
    cls_src["_u"] = (cls_src["Head"] != "")
    cls_src = cls_src.sort_values(["_u", "Ledger"]).drop(columns="_u").reset_index(drop=True)
    with t_cl:
        st.subheader("Classification")
        n_un = int((cls_src["Head"] == "").sum())
        if n_un:
            st.warning("%d ledger(s) have no head (listed first). Pick a head for each." % n_un)
        else:
            st.success("Every ledger is classified. You can still change any head below.")
        show = cls_src[["Ledger", "Group", "Dr_pos", "Head", "Source"]].rename(columns={"Dr_pos": "Dr (+) / Cr (-)"})
        edited = st.data_editor(
            show, hide_index=True, key="cls_%d_%d" % (ver, pyver), disabled=["Ledger", "Group", "Dr (+) / Cr (-)", "Source"], height=420,
            column_config={"Head": st.column_config.SelectboxColumn("Head", options=[""] + HEAD_LIST, width="large"),
                           "Dr (+) / Cr (-)": st.column_config.NumberColumn(format="%.2f")})
        if plac_mask.any():
            st.caption("The Profit & Loss A/c ledger is handled automatically (opening reserves are derived from the other ledgers).")
    df = base.copy()
    new_head = dict(zip(edited["Ledger"], edited["Head"].fillna("")))
    df["Head"] = [new_head.get(l, h) for l, h in zip(df["Ledger"], df["Head"])]

    # ------------------------------------------------------------ previous year
    py = None
    with t_py:
        st.subheader("Previous year")
        if raw is None:
            st.info("No previous year loaded. Choose a source in the sidebar (section 3).")
            st.markdown(
                "**How the previous year is classified**\n\n"
                "- *Last year's workbook* (from this app): its heads are authoritative, and matching current-year ledgers inherit them.\n"
                "- *Uploaded trial balance or Tally pull*: each ledger takes the head of the same ledger in the current year; otherwise the "
                "group rule; otherwise a similar name (flagged for your review); otherwise you pick it here.\n"
                "- Names are matched ignoring case, spaces, '&' and punctuation.\n"
                "- Whatever the source, the checks compare the two years: balance, opening equity against last year's closing equity, and "
                "ledgers classified differently.")
        else:
            st.caption(raw["note"] or "")
            pdf = raw["df"].copy()
            if "Head" not in pdf.columns:
                pdf["Head"] = ""
            cls = py_classify(pdf, df, authoritative=(raw["mode"] == "workbook"), groups=ss.get("groups"))
            cls_ui = cls[~is_plac(cls["Ledger"])].copy()
            cls_ui["_u"] = cls_ui["Head"] != ""
            cls_ui = cls_ui.sort_values(["_u", "Ledger"]).drop(columns="_u").reset_index(drop=True)
            n0 = int((cls_ui["Head"] == "").sum())
            if n0:
                st.warning("%d previous-year ledger(s) have no head (listed first). Pick a head for each." % n0)
            else:
                st.success("Every previous-year ledger is classified. Review the Source column.")
            shw = cls_ui[["Ledger", "Group", "Dr_pos", "Head", "Source"]].rename(columns={"Dr_pos": "Dr (+) / Cr (-)"})
            ed = st.data_editor(shw, hide_index=True, key="pycls_%d" % pyver, height=380,
                                disabled=["Ledger", "Group", "Dr (+) / Cr (-)", "Source"],
                                column_config={"Head": st.column_config.SelectboxColumn("Head", options=[""] + HEAD_LIST, width="large"),
                                               "Dr (+) / Cr (-)": st.column_config.NumberColumn(format="%.2f")})
            nh = dict(zip(ed["Ledger"], ed["Head"].fillna("")))
            cls["Head"] = [nh.get(l, h) for l, h in zip(cls["Ledger"], cls["Head"])]
            cmap = {norm_name(l): h for l, h in zip(df["Ledger"], df["Head"]) if h and not is_plac(pd.Series([l])).iloc[0]}
            diff = [(l, cmap[norm_name(l)], h) for l, h in zip(cls["Ledger"], cls["Head"])
                    if h and norm_name(l) in cmap and cmap[norm_name(l)] != h and not is_plac(pd.Series([l])).iloc[0]]
            cs_py = py_cs_in or raw.get("cs_items") or None
            M_py = build_mapping(cls, entity, cs_py, py_accept, 0.0, raw.get("closing"), gross)
            tax_py = raw["tax"] or py_tax_in
            S_py = build_statements(M_py, framework, tax_py)
            py = {"cls": cls, "M": M_py, "S": S_py, "diff": diff, "tax": tax_py, "raw": raw}
            st.markdown("**Source of classification**")
            st.dataframe(cls.groupby("Source").size().reset_index(name="Ledgers"), hide_index=True)
            if diff:
                st.warning("These ledgers carry a different head in the two years. Decide which is right.")
                st.dataframe(pd.DataFrame(diff, columns=["Ledger", "Current year head", "Previous year head"]), hide_index=True)
            if M_py["msg"] and M_py["status"] in ("pending", "unclassified"):
                st.warning("Previous year: " + M_py["msg"])
            k1, k2, k3 = st.columns(3)
            k1.metric("Previous-year total assets", inr(S_py["bsv"]["tas"]))
            k2.metric("Previous-year profit for the year", inr(S_py["plv"]["pat"]))
            k3.metric("Previous-year Balance Sheet difference", inr(S_py["diff"]))

    # ------------------------------------------------------ depreciation tab
    booked = float(df.loc[df["Head"] == "Depreciation & Amortisation", "Dr_pos"].sum())
    ppe_tb = float(df.loc[df["Head"] == "Property, Plant & Equipment", "Dr_pos"].sum())
    start, pye = period_dates(year_end)
    end = year_end
    is_sample_co = (company_name or "").strip() == SAMPLE_NAME and py is not None
    sig = (ver, pyver if py is not None else 0, bool(ss.get("sample_mode")), is_sample_co)
    if ss.get("dep_sig") != sig:
        s2b, itb = prefill_dep_tables(base, py["cls"] if py is not None else None)
        if ss.get("sample_mode") or is_sample_co:
            s2b, itb = sample_dep_inputs(s2b, itb)
        for c in S2_DATES:
            s2b[c] = pd.to_datetime(s2b[c])
        ss["s2_base"], ss["it_base"], ss["dep_sig"] = s2b, itb, sig
        ss["dep_ver"] = ss.get("dep_ver", 0) + 1
    dv = ss["dep_ver"]
    dep, adj = None, 0.0
    with t_dp:
        st.subheader("Depreciation calculator")
        st.caption("Period %s to %s (%d days). Opening values %s. Edit the tables before relying on the result." % (
            long_date(start), long_date(end), days_in_year(start, end),
            "come from the previous year's closing balances" if py is not None else
            "are prefilled with the current Tally closing balances, so you must enter the true opening values"))
        if booked:
            st.warning("Tally already has depreciation of Rs %s booked, so fixed-asset balances in Tally are after that charge." % inr(booked))
        sub = st.tabs(["Companies Act - Schedule II", "Income Tax Act - block of assets", "Deferred tax (information)"])
        with sub[0]:
            a, b = st.columns(2)
            method = a.radio("Method", ["SLM", "WDV"], horizontal=True, key="m_%d" % dv)
            res_in = b.number_input("Residual value (% of cost)", min_value=0.0, max_value=50.0, value=5.0, step=1.0, key="r_%d" % dv)
            res_eff = res_in
            if method == "WDV" and res_in <= 0:
                st.error("WDV needs a residual value above zero. Using 5% for the calculation.")
                res_eff = 5.0
            tbl = st.data_editor(
                ss["s2_base"], key="s2_%d" % dv, num_rows="dynamic", hide_index=True, height=280,
                column_config={
                    "Asset class": st.column_config.SelectboxColumn("Asset class", options=[x for x, _ in SCH2_DEFAULTS], width="large"),
                    "Useful life (years)": st.column_config.NumberColumn(min_value=0, step=1, format="%d"),
                    "Remaining life override (years)": st.column_config.NumberColumn(
                        min_value=0.0, format="%.1f", help="SLM only. For an old asset: (written-down value less residual) / remaining life."),
                    "Date put to use (additions)": st.column_config.DateColumn(format="DD-MMM-YYYY"),
                    "Date of sale": st.column_config.DateColumn(format="DD-MMM-YYYY"),
                    **{k: st.column_config.NumberColumn(min_value=0.0, format="%.2f") for k in S2_NUM[2:]}})
            t2 = clean_dep(tbl, "s2")
            res2 = dep_sch2(t2, method, res_eff, start, end)
            w = []
            if ((t2["Additions (cost)"] > 0) & t2["Date put to use (additions)"].isna()).any():
                w.append("Some additions have no date put to use, so no depreciation is charged on them.")
            if ((t2["Disposals (original cost)"] > 0) & t2["Date of sale"].isna()).any():
                w.append("Some disposals have no date of sale, so no depreciation is charged for the year of disposal.")
            if ((t2["Useful life (years)"] <= 0) & ~t2["Asset class"].str.startswith("Land")).any():
                w.append("A non-land asset has a useful life of zero, so it is not depreciated.")
            if method == "WDV" and (t2["Remaining life override (years)"] > 0).any():
                w.append("The remaining-life override applies to SLM only and is ignored under WDV.")
            for x in w:
                st.warning(x)
            calc2, net2 = float(res2["Depreciation for the year"].sum()), float(res2["Closing net block"].sum())
            unexpl = float((t2["Opening gross block"] - t2["Opening accumulated depreciation"] + t2["Additions (cost)"]
                            - (t2["Disposals (original cost)"] - t2["Accumulated depreciation on disposals"])).sum()) - (ppe_tb + booked)
            st.markdown("**Result**")
            if len(res2):
                st.dataframe(res2, hide_index=True)
            else:
                st.caption("No fixed-asset rows. Add rows to the table above.")
            m = st.columns(4)
            m[0].metric("Depreciation per calculator", inr(calc2))
            m[1].metric("Booked in Tally", inr(booked))
            m[2].metric("Difference (calculator less books)", inr(calc2 - booked))
            m[3].metric("Inputs vs Tally fixed assets", inr(unexpl), help="Opening + additions - disposals (net) less Tally fixed assets and depreciation booked. Should be nil.")
            if abs(unexpl) >= 1 and len(res2):
                st.info("The values entered do not reconcile to the fixed assets in Tally (difference %s). Check opening balances, additions and disposals." % inr(unexpl))
            if float(res2["Profit / (loss) on disposal"].abs().sum()) > 0:
                st.caption("Profit / (loss) on disposal is shown for information. It is assumed to be booked in Tally already; if not, book it.")
            with st.expander("Schedule II useful lives used as defaults"):
                st.dataframe(pd.DataFrame(SCH2_DEFAULTS, columns=["Asset class", "Useful life (years)"]), hide_index=True)
        with sub[1]:
            tbl = st.data_editor(
                ss["it_base"], key="it_%d" % dv, num_rows="dynamic", hide_index=True, height=240,
                column_config={
                    "Block of assets": st.column_config.SelectboxColumn("Block of assets", options=[x for x, _ in IT_DEFAULTS], width="large"),
                    "Rate (%)": st.column_config.NumberColumn(min_value=0.0, max_value=100.0, format="%.1f"),
                    **{k: st.column_config.NumberColumn(min_value=0.0, format="%.2f") for k in IT_COLS[3:]}})
            tI = clean_dep(tbl, "it")
            resI = dep_it(tI)
            calcI, netI = float(resI["Depreciation for the year"].sum()), float(resI["Closing WDV"].sum())
            st.markdown("**Result**")
            if len(resI):
                st.dataframe(resI, hide_index=True)
            else:
                st.caption("No blocks. Add rows to the table above.")
            m = st.columns(2)
            m[0].metric("Income-tax depreciation", inr(calcI))
            m[1].metric("Closing tax WDV", inr(netI))
            st.caption("Full rate on opening WDV plus additions used 180 days or more less sale proceeds; half rate on additions used under 180 days. "
                       "Additional depreciation is not included. Opening WDV is the tax WDV, which usually differs from the book value.")
            with st.expander("Income-tax rates used as defaults"):
                st.dataframe(pd.DataFrame(IT_DEFAULTS, columns=["Block of assets", "Rate (%)"]), hide_index=True)
        with sub[2]:
            st.caption("Timing difference between book (Schedule II) and tax depreciation. For information; it is not posted to the statements.")
            rate_dt = st.number_input("Tax rate (%)", min_value=0.0, max_value=60.0, value=25.168, step=0.5, key="dt_%d" % dv)
            opening_dt = st.number_input("Opening deferred tax liability / (asset) per books (Rs)", value=0.0, step=1000.0, key="odt_%d" % dv)
            tdiff = net2 - netI
            dtl = tdiff * rate_dt / 100.0
            k = st.columns(4)
            k[0].metric("Book net block (Schedule II)", inr(net2))
            k[1].metric("Tax WDV", inr(netI))
            k[2].metric("Closing deferred tax liability / (asset)", inr(dtl))
            k[3].metric("Charge / (credit) for the year", inr(dtl - opening_dt))
            st.caption("A positive figure is a liability. Recognise a deferred tax asset only where there is reasonable (AS 22) or virtual (Ind AS 12) certainty "
                       "of future taxable profit. Both tables above must be complete for this to be meaningful.")
        st.markdown("---")
        basis_txt = st.radio("Basis to apply to the statements", ["Companies Act 2013 - Schedule II", "Income Tax Act 1961 - block of assets"],
                             horizontal=True, key="basis_%d" % dv)
        sch2 = basis_txt.startswith("Companies")
        calc, net, t, res_t = (calc2, net2, t2, res2) if sch2 else (calcI, netI, tI, resI)
        allow = sch2 or entity != "Company"
        key_on = "dep_on_%d_%s" % (dv, "s2" if sch2 else "it")
        if not allow:
            st.caption("Income-tax depreciation is used to compute taxable income. A company must charge Schedule II depreciation "
                       "in its books, so applying it to company statements is switched off.")
        elif not len(res_t):
            st.info("Enter the fixed-asset rows above, then apply the depreciation.")
        else:
            diff_dep = calc - booked
            c1, c2, c3 = st.columns(3)
            c1.metric("Depreciation to be charged", inr(calc))
            c2.metric("Already booked in Tally", inr(booked))
            c3.metric("Change to profit before tax", inr(-diff_dep))
            if not ss.get(key_on):
                st.button("Apply depreciation to Balance Sheet and Profit and Loss", type="primary", key="dep_btn_" + key_on,
                          on_click=lambda: ss.__setitem__(key_on, True),
                          help="Replaces the depreciation booked in Tally with the calculated figure. Fixed assets, profit and reserves in the statements, notes and fixed-asset schedule all follow.")
                st.caption("Check the opening values, additions, disposals and rates in the tables above before applying. "
                           "Until you apply, the statements show depreciation as booked in Tally.")
            else:
                st.success("Applied. The statements now show depreciation of Rs %s (Tally had Rs %s): profit before tax %s by Rs %s and the net block of fixed assets by the same amount. "
                           "Open the Statements tab to see the effect." % (inr(calc), inr(booked), "falls" if diff_dep >= 0 else "rises", inr(abs(diff_dep))))
                st.button("Remove (go back to depreciation as booked)", key="dep_rm_" + key_on, on_click=lambda: ss.__setitem__(key_on, False))
        applied = bool(allow and ss.get(key_on) and len(res_t))
        adj = (calc - booked) if applied else 0.0
        note_txt = (("Schedule II, %s method, residual value %.1f%% of cost, %d days. Additions and disposals are depreciated pro rata by days; "
                     "depreciation on assets sold runs to the date of sale." % (method, res_eff, days_in_year(start, end))) if sch2 else
                    "Income Tax Act block of assets, WDV. Half rate on additions used under 180 days. Additional depreciation not included.")
        dep = {"basis": "sch2" if sch2 else "it", "method": method, "res": res_eff, "table": t, "result": res_t, "applied": applied,
               "calc": calc, "booked": booked, "net": net, "note": note_txt, "unexplained": unexpl if sch2 else None}

    # ------------------------------------------------------ build the statements
    man = ss.get("cs_manual")
    M = build_mapping(df, entity, ss.get("cs_items"), accept, adj, man, gross,
                      ("Closing stock Rs %s entered by the user (from the stock summary / valuation report)." % inr(man)) if man is not None else None)
    S = build_statements(M, framework, tax)
    checks = run_checks(df, M, S, dep, tax, py)
    ye = year_end
    nm = company_name.strip() or "Company Name"

    # ---- auditor action list and the response file (clearance before filing)
    prof_a = ss.get("profile") or {}
    has_prof_a = any(prof_a.get(k) for k in ("firm", "frn", "partner", "cin", "pan", "address")) or any(s.get("name") for s in prof_a.get("signs", []))
    dep_a = dep if dep and len(dep["table"]) else None
    ax, ax_err = None, None
    try:
        ax = make_audit_excel(checks, M, S, dep_a, tax, py, framework, entity, nm, ye, prof_a if has_prof_a else None, df)
    except Exception as e:
        ax_err = str(e)
    aud = None
    aud_src = None
    if ax is not None:
        try:
            cur_rows, _, _, cur_head = read_audit_responses(ax)
            cur_refs = set(cur_rows["Ref"])
            raw = None
            if ss.get("aud_mode") in ("clean", "gaps"):
                raw = sample_audit_answers(ax, with_gaps=(ss["aud_mode"] == "gaps"))
                aud_src = "Sample answers (%s)" % ("complete" if ss["aud_mode"] == "clean" else "with deliberate gaps")
            elif ss.get("aud_file"):
                raw = ss["aud_file"]["bytes"]
                aud_src = ss["aud_file"]["name"]
            if raw is not None:
                resp, rv, rd, fh = read_audit_responses(raw)
                aud = assess_responses(resp, rv, rd, cur_refs, cur_head, fh)
        except Exception as e:
            aud = {"error": str(e)}
    if aud is None:
        a_stat, a_msg = "warn", "No completed auditor action list uploaded yet. Download the list, have it completed, and upload it in 'Checks and export'."
    elif "error" in aud:
        a_stat, a_msg = "fail", "The uploaded file could not be read: %s" % aud["error"]
    elif aud["clean"]:
        a_stat, a_msg = "pass", "All %d items closed. Reviewed by %s on %s." % (aud["total"], aud["reviewer"], aud["date"])
    else:
        a_stat, a_msg = "fail", "%d of %d items closed; %d exception(s)%s." % (
            aud["closed"], aud["total"], len(aud["exceptions"]), (" and %d sign-off issue(s)" % len(aud["problems"])) if aud["problems"] else "")
    checks_all = checks + [(AUDIT_CHECK, a_stat, a_msg)]

    def stock_fix_panel(tag):
        """Guided fix for the closing-stock check, shown where the failure is reported."""
        if ss.get("cs_manual") is not None:
            st.info("Closing stock of Rs %s entered by you is in use." % inr(ss["cs_manual"]))
            st.button("Remove the entered value", key="clr_" + tag, on_click=lambda: ss.pop("cs_manual", None))
        if M["status"] != "pending":
            return
        derived = M["os"] + M["delta"]
        st.markdown("##### Fix: closing stock")
        st.write("Tally's Profit & Loss A/c ledger already includes a closing-stock figure, but the stock ledgers do not hold it. "
                 "Choose how to bring closing stock into the statements:")
        a, b, c_ = st.columns(3)
        a.metric("Stock held in ledgers", inr(M["os"]))
        b.metric("Closing stock implied by the books", inr(derived))
        c_.metric("Stock items value in Tally", inr(ss["cs_items"]) if ss.get("cs_items") else "not available")
        if derived < 0:
            st.warning("The implied closing stock is negative, so the books need review before it can be accepted. Enter the value from the stock summary instead.")
        x, y = st.columns(2)
        with x:
            st.markdown("**Option 1.** Accept the value implied by the books")
            st.button("Accept %s" % inr(derived), key="acc_" + tag, type="primary", disabled=derived < 0,
                      on_click=lambda: ss.__setitem__("accept_cs", True),
                      help="Use only after comparing with Tally's Stock Summary. The Profit & Loss tie-out then agrees by construction.")
        with y:
            st.markdown("**Option 2.** Enter closing stock from the stock summary or valuation report")
            v = st.number_input("Closing stock (Rs)", min_value=0.0, value=float(max(derived, 0.0)), step=1000.0, key="man_" + tag)
            st.button("Use this value", key="use_" + tag, on_click=lambda v=v: ss.__setitem__("cs_manual", float(v)))
            st.caption("If your value differs from the books, the Profit & Loss tie-out will show the difference so you can investigate it.")

    with t_ov:
        pv = lambda key, src: (inr(S[src][key] - py["S"][src][key]) if py else None)
        k = st.columns(5)
        k[0].metric("Revenue from operations", inr(S["plv"]["rev"]), pv("rev", "plv"))
        k[1].metric("Profit before tax", inr(S["plv"]["pbt"]), pv("pbt", "plv"))
        k[2].metric("Profit for the year", inr(S["plv"]["pat"]), pv("pat", "plv"))
        k[3].metric("Total assets", inr(S["bsv"]["tas"]), pv("tas", "bsv"))
        k[4].metric("Balance Sheet difference", inr(S["diff"]))
        if py:
            st.caption("Change shown against the previous year.")
        fails = [c for c in checks if c[1] == "fail"]
        if fails:
            st.error("%d check(s) failed. The fix for each is shown below." % len(fails))
            stock_fix_panel("ov")
        else:
            st.success("All checks passed or need only your review. Open 'Checks and export' to download the Excel file.")
        st.markdown(pills(checks_all), unsafe_allow_html=True)
        st.caption("Basis: %s | Entity: %s | Year ended %s" % (framework, entity, long_date(ye)))

    with t_tb:
        st.subheader("Trial balance from Tally")
        view = df[["Ledger", "Group", "Head"]].copy()
        view["Debit"] = df["Dr_pos"].apply(lambda x: inr(x) if x > 0 else "")
        view["Credit"] = df["Dr_pos"].apply(lambda x: inr(-x) if x < 0 else "")
        st.dataframe(view, hide_index=True, height=480)
        st.write("Total debit **%s**   |   Total credit **%s**" % (
            inr(df.loc[df["Dr_pos"] > 0, "Dr_pos"].sum()), inr(-df.loc[df["Dr_pos"] < 0, "Dr_pos"].sum())))

    with t_st:
        left, right = st.columns(2)
        cyl, pyl = ye.strftime("%d-%b-%Y"), pye.strftime("%d-%b-%Y")
        bs_head = ("Balance Sheet as at %s" % long_date(ye), FRAME_TEXT[framework], "As at " + cyl, "As at " + pyl)
        pl_head = ("Statement of Profit and Loss for the year ended %s" % long_date(ye), FRAME_TEXT[framework],
                   "Year ended " + cyl, "Year ended " + pyl)
        with left:
            st.markdown(html_statement(S["bs"], S["bsv"], bs_head, py["S"]["bsv"] if py else None), unsafe_allow_html=True)
        with right:
            st.markdown(html_statement(S["pl"], S["plv"], pl_head, py["S"]["plv"] if py else None), unsafe_allow_html=True)
        st.caption("Amounts are rounded to the nearest rupee for display; totals are computed on unrounded amounts. Notes, the fixed-asset schedule and the print-ready PDF are in 'Checks and export'.%s"
                   % ("" if py else " No previous year is loaded."))
        if framework == FRAMEWORKS[1]:
            st.caption("Ind AS presentation here is simplified: other comprehensive income is nil and the statement of changes in equity is not prepared.")

    with t_ck:
        st.subheader("Checks")
        stock_fix_panel("ck")
        st.markdown(pills(checks_all), unsafe_allow_html=True)
        if M["msg"]:
            (st.warning if M["status"] in ("pending", "unclassified") else st.info)(M["msg"])
        with st.expander("Group-wise totals (to investigate a difference)"):
            g = df.groupby("Top group")["Dr_pos"].sum().reset_index()
            g["Net (Dr +, Cr -)"] = g["Dr_pos"].apply(inr)
            st.dataframe(g[["Top group", "Net (Dr +, Cr -)"]], hide_index=True)
        st.subheader("Auditor review before filing")
        st.write("1. Download the auditor action list. 2. The auditor completes Status, Done by / date and Remarks for every line and signs off at the bottom. "
                 "3. Upload the completed file here. The application checks that nothing is open before the statements are cleared for filing.")
        u1, u2 = st.columns([1, 1])
        with u1:
            if ax is not None:
                st.download_button("1. Download auditor action list (Excel)", ax, file_name="Auditor_Action_List.xlsx", key="dl_ax_top",
                                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            else:
                st.error("Could not build the auditor list: %s" % ax_err)
        with u2:
            def _got_file():
                f_ = ss.get("aud_up")
                ss["aud_mode"] = None
                ss["aud_file"] = {"name": f_.name, "bytes": f_.getvalue()} if f_ is not None else None
            st.file_uploader("2. Upload the completed list (Excel)", type=["xlsx"], key="aud_up", on_change=_got_file)
        if ss.get("sample_mode") or (nm == SAMPLE_NAME):
            st.caption("Demonstration: use sample answers instead of uploading a file, or download a list already filled with sample answers and upload it above.")
            if ax is not None:
                g1, g2 = st.columns(2)
                g1.download_button("Download list with sample answers (complete)", sample_audit_answers(ax), key="dl_sa_ok",
                                   file_name="Auditor_Action_List_SAMPLE_COMPLETED.xlsx",
                                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                g2.download_button("Download list with sample answers (with gaps)", sample_audit_answers(ax, with_gaps=True), key="dl_sa_gap",
                                   file_name="Auditor_Action_List_SAMPLE_WITH_GAPS.xlsx",
                                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            s1, s2, s3 = st.columns(3)
            s1.button("Use sample answers (complete)", key="aud_s1", on_click=lambda: ss.__setitem__("aud_mode", "clean"))
            s2.button("Use sample answers (with gaps)", key="aud_s2", on_click=lambda: ss.__setitem__("aud_mode", "gaps"))
            s3.button("Clear responses", key="aud_s3", on_click=lambda: (ss.__setitem__("aud_mode", None), ss.__setitem__("aud_file", None)))
        elif ss.get("aud_file"):
            st.button("Clear responses", key="aud_s3b", on_click=lambda: ss.__setitem__("aud_file", None))
        if aud is None:
            st.info("No response file loaded yet.")
        elif "error" in aud:
            st.error(a_msg)
        else:
            st.caption("Source: %s" % aud_src)
            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Items on the list", aud["total"])
            m2.metric("Closed", aud["closed"])
            m3.metric("Exceptions", len(aud["exceptions"]))
            m4.metric("Reviewer sign-off", "Complete" if not aud["problems"] else "Incomplete")
            for pr in aud["problems"]:
                st.warning(pr)
            if aud["clean"]:
                st.success("Cleared for filing: every item is closed with a recorded reviewer, date and remarks. Reviewed by %s on %s." % (aud["reviewer"], aud["date"]))
            else:
                st.error("Not cleared for filing. Resolve the exceptions below, update the Excel list and upload it again.")
                ex_df = pd.DataFrame(aud["exceptions"], columns=["Ref", "Item", "Problem", "Priority"])
                st.dataframe(ex_df, hide_index=True)
                st.download_button("Exceptions (CSV)", ex_df.to_csv(index=False), file_name="Auditor_Exceptions.csv", mime="text/csv", key="dl_ex")
        st.subheader("Download")
        policy_dep = ""
        if dep and dep["applied"]:
            policy_dep = ("Depreciation: calculated under %s; see the Depreciation Schedule sheet." % (
                "Schedule II of the Companies Act 2013 (%s, residual %.1f%%)" % (dep["method"], dep["res"]) if dep["basis"] == "sch2"
                else "the Income Tax Act 1961 block of assets (WDV)"))
        has_prof = any(prof.get(k) for k in ("firm", "frn", "partner", "cin", "pan", "address")) or any(s.get("name") for s in prof["signs"])
        n_fail = len([c_ for c_ in checks if c_[1] == "fail"])
        if n_fail:
            st.error("%d check(s) failed. Resolve them before the statements are signed." % n_fail)
        if a_stat != "pass":
            st.warning("Auditor action list is not cleared yet (see 'Auditor review before filing' above). The files below are drafts until it is.")
        else:
            st.success("Auditor action list cleared. The statements below are ready for signing.")
        st.markdown("**For signing**")
        d1, d2, d5 = st.columns(3)
        dep_in = dep if dep and len(dep["table"]) else None
        try:
            xl = make_excel(M, framework, nm, ye, tax, dep_in, policy_dep, py["M"] if py else None, py["tax"] if py else 0.0, prof if has_prof else None)
            d1.download_button("Financial statements (Excel)", xl, file_name="Financial_Statements.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary",
                               help="Balance Sheet, Profit and Loss, policies, notes and fixed-asset schedule, set up to print on A4. The grey tabs are working papers.")
        except Exception as e:
            d1.error("Could not build the Excel file: %s" % e)
        try:
            pdf = make_pdf(M, S, py, framework, entity, nm, ye, dep_in, prof if has_prof else None, tax)
            d2.download_button("Financial statements (PDF for signing)", pdf, file_name="Financial_Statements.pdf", mime="application/pdf", type="primary",
                               help="A4 PDF of the statements, notes and fixed-asset schedule with the signature blocks.")
        except ImportError:
            d2.warning("PDF needs the 'reportlab' library. Close the launcher, run:  python -m pip install reportlab  and start again.")
        except Exception as e:
            d2.error("Could not build the PDF: %s" % e)
        if ax is not None:
            d5.download_button("Auditor action list (Excel)", ax, file_name="Auditor_Action_List.xlsx", key="dl_ax_bottom",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               help="What the auditor must do before signing. Not part of the statements.")
        st.markdown("**Working data**")
        d2, d3, d4 = st.columns(3)
        tbcsv = df[["Ledger", "Group", "Head", "Dr_pos"]].rename(columns={"Dr_pos": "Dr (+) / Cr (-)"}).to_csv(index=False)
        d2.download_button("Trial balance (CSV)", tbcsv, file_name="Trial_Balance_Classified.csv", mime="text/csv")
        if dep and len(dep["result"]):
            d3.download_button("Depreciation (CSV)", dep["result"].to_csv(index=False),
                               file_name="Depreciation_Calculation.csv", mime="text/csv")
        stmt = []
        for nm_, spec, vals, pvals in (("Balance Sheet", S["bs"], S["bsv"], py["S"]["bsv"] if py else {}),
                                       ("Profit and Loss", S["pl"], S["plv"], py["S"]["plv"] if py else {})):
            for r in spec:
                if r["kind"] != "blank":
                    stmt.append({"Statement": nm_, "Particulars": r["label"], "Note": r["no"] or "",
                                 "Current year": round(vals[r["key"]], 2) if r["key"] in vals else "",
                                 "Previous year": round(pvals[r["key"]], 2) if r["key"] in pvals else ""})
        d4.download_button("Statements (CSV)", pd.DataFrame(stmt).to_csv(index=False),
                           file_name="Statements.csv", mime="text/csv")


if __name__ == "__main__":
    main()
