"""
Turns the JSON returned by the GST portal into clean pandas tables and
formatted Excel workbooks. No network code here - easy to test offline.
"""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

import pandas as pd

from portal_config import MONTHS, prd_label

TAX_COLS = ["Taxable value", "IGST", "CGST", "SGST", "Cess"]
_KEYMAP = {"txval": "Taxable value", "iamt": "IGST", "camt": "CGST", "samt": "SGST", "csamt": "Cess",
           "igst": "IGST", "cgst": "CGST", "sgst": "SGST", "cess": "Cess"}


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _taxrow(d: dict | None) -> dict:
    d = d or {}
    return {c: _num(d.get(k)) for k, c in (("txval", "Taxable value"), ("iamt", "IGST"),
                                          ("camt", "CGST"), ("samt", "SGST"), ("csamt", "Cess"))}


def _data(js):
    """Portal responses are {"status":1,"data":{...}} or {"data":{...},"chksum":..}."""
    if isinstance(js, dict) and "data" in js and isinstance(js["data"], (dict, list)):
        return js["data"]
    return js


def parse_dmy(s):
    if not s:
        return None
    s = str(s).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


# =============================================================== GSTR-3B
OSUP = [("osup_det", "3.1(a) Outward taxable supplies"), ("osup_zero", "3.1(b) Zero rated supplies"),
        ("osup_nil_exmp", "3.1(c) Nil rated / exempt"), ("isup_rev", "3.1(d) Inward supplies (reverse charge)"),
        ("osup_nongst", "3.1(e) Non-GST outward supplies")]
ITC_AVL = {"IMPG": "4A(1) Import of goods", "IMPS": "4A(2) Import of services",
           "ISRC": "4A(3) Inward supplies liable to RCM", "ISD": "4A(4) Inward supplies from ISD",
           "OTH": "4A(5) All other ITC"}
ITC_REV = {"RUL": "4B(1) Reversed - Rules 38, 42, 43 & sec 17(5)", "OTH": "4B(2) Reversed - Others"}
ITC_INELG = {"RUL": "4D(1) ITC reclaimed / ineligible (RUL)", "OTH": "4D(2) Ineligible ITC (OTH)"}


def gstr3b_frames(data_by_prd: dict[str, dict]):
    """Returns (detail_long_df, monthly_summary_df)."""
    rows, summ = [], []
    for prd, js in sorted(data_by_prd.items(), key=lambda kv: prd_label(kv[0])):
        d = _data(js) or {}
        r = d.get("r3b", d)
        per = prd_label(prd)
        arn, arndt = d.get("arn", ""), d.get("arnDt", "")
        sup = r.get("sup_details", {}) or {}
        for key, desc in OSUP:
            rows.append({"Period": per, "Table": desc, **_taxrow(sup.get(key))})
        itc = r.get("itc_elg", {}) or {}
        for it in itc.get("itc_avl", []) or []:
            rows.append({"Period": per, "Table": ITC_AVL.get(it.get("ty"), "4A " + str(it.get("ty"))), **_taxrow(it)})
        for it in itc.get("itc_rev", []) or []:
            rows.append({"Period": per, "Table": ITC_REV.get(it.get("ty"), "4B " + str(it.get("ty"))), **_taxrow(it)})
        rows.append({"Period": per, "Table": "4C Net ITC available", **_taxrow(itc.get("itc_net"))})
        for it in itc.get("itc_inelg", []) or []:
            rows.append({"Period": per, "Table": ITC_INELG.get(it.get("ty"), "4D " + str(it.get("ty"))), **_taxrow(it)})
        il = r.get("intr_ltfee", {}) or {}
        rows.append({"Period": per, "Table": "5.1 Interest", **_taxrow(il.get("intr_details"))})
        rows.append({"Period": per, "Table": "5.1 Late fee", **_taxrow(il.get("ltfee_details"))})
        for k, lbl in (("unreg_details", "3.2 Inter-state to unregistered"), ("comp_details", "3.2 Inter-state to composition"),
                       ("uin_details", "3.2 Inter-state to UIN holders")):
            for it in (r.get("inter_sup", {}) or {}).get(k, []) or []:
                rows.append({"Period": per, "Table": f"{lbl} (POS {it.get('pos', '')})", **_taxrow(it)})

        a = _taxrow(sup.get("osup_det"))
        z = _taxrow(sup.get("osup_zero"))
        rc = _taxrow(sup.get("isup_rev"))
        avl = [_taxrow(x) for x in itc.get("itc_avl", []) or []]
        rev = [_taxrow(x) for x in itc.get("itc_rev", []) or []]
        net = _taxrow(itc.get("itc_net"))
        tax = lambda t: t["IGST"] + t["CGST"] + t["SGST"] + t["Cess"]
        tt = r.get("tt_val", {}) or {}
        nil_return = all(tax(_taxrow(sup.get(k))) == 0 and _taxrow(sup.get(k))["Taxable value"] == 0 for k, _ in OSUP) \
            and sum(tax(x) for x in avl) == 0
        summ.append({
            "Period": per,
            "Outward taxable value 3.1(a)": a["Taxable value"],
            "Output tax 3.1(a)": tax(a),
            "Zero rated value 3.1(b)": z["Taxable value"],
            "RCM tax 3.1(d)": tax(rc),
            "ITC available 4A": sum(tax(x) for x in avl),
            "ITC reversed 4B": sum(tax(x) for x in rev),
            "Net ITC 4C": tax(net),
            "Interest paid": tax(_taxrow(il.get("intr_details"))),
            "Late fee paid": tax(_taxrow(il.get("ltfee_details"))),
            "Tax paid in cash": _num(tt.get("tt_csh_pd")),
            "Tax paid through ITC": _num(tt.get("tt_itc_pd")),
            "Nil return": "Yes" if nil_return else "No",
            "ARN": arn, "ARN date": arndt,
        })
    return pd.DataFrame(rows), pd.DataFrame(summ)


# =============================================================== GSTR-1
def gstr1_frames(data_by_prd: dict[str, dict]):
    """Returns (section_long_df, pivot_taxable_df, totals_df)."""
    rows = []
    for prd, js in sorted(data_by_prd.items(), key=lambda kv: prd_label(kv[0])):
        d = _data(js) or {}
        per = prd_label(prd)
        if str(d.get("isnil", "")).upper() == "Y":
            rows.append({"Period": per, "Section": "NIL RETURN", "Level": "Main", "Records": 0,
                         "Taxable value": 0, "IGST": 0, "CGST": 0, "SGST": 0, "Cess": 0, "Invoice value": 0})
            continue
        for s in d.get("sec_sum", []) or []:
            def row(x, name, level):
                out = {"Period": per, "Section": name, "Level": level, "Records": _num(x.get("ttl_rec")),
                       "Taxable value": _num(x.get("ttl_tax")), "IGST": _num(x.get("ttl_igst")),
                       "CGST": _num(x.get("ttl_cgst")), "SGST": _num(x.get("ttl_sgst")),
                       "Cess": _num(x.get("ttl_cess")), "Invoice value": _num(x.get("ttl_val"))}
                for k in ("ttl_expt_amt", "ttl_ngsup_amt", "ttl_nilsup_amt", "ttl_doc_issued",
                          "ttl_doc_cancelled", "net_doc_issued"):
                    if k in x:
                        out[k.replace("ttl_", "").replace("_", " ").title()] = _num(x.get(k))
                return out
            name = s.get("sec_nm", "")
            rows.append(row(s, name, "Main"))
            for sub in s.get("sub_sections", []) or []:
                rows.append(row(sub, f"{name} > {sub.get('sec_nm') or sub.get('typ', '')}", "Sub"))
    df = pd.DataFrame(rows)
    if df.empty:
        return df, df, df
    main = df[df["Level"] == "Main"]
    pivot = main.pivot_table(index="Section", columns="Period", values="Taxable value", aggfunc="sum", fill_value=0)
    pivot = pivot.loc[(pivot != 0).any(axis=1)].reset_index()
    tot = main[main["Section"] == "TTL_LIAB"][["Period"] + TAX_COLS + ["Invoice value"]].copy()
    nil = main[main["Section"] == "NIL RETURN"][["Period"]]
    if not nil.empty:
        z = nil.assign(**{c: 0.0 for c in TAX_COLS + ["Invoice value"]})
        tot = pd.concat([tot, z], ignore_index=True)
    return df, pivot, tot.sort_values("Period").reset_index(drop=True)


# =============================================================== GSTR-1 B2B (counterparty-wise)
def _find_list(obj, key):
    """Return the first list stored under `key` anywhere in a nested dict/list."""
    if isinstance(obj, dict):
        if isinstance(obj.get(key), list):
            return obj[key]
        for v in obj.values():
            r = _find_list(v, key)
            if r is not None:
                return r
    elif isinstance(obj, list):
        for v in obj:
            r = _find_list(v, key)
            if r is not None:
                return r
    return None


def _inv_tax(inv):
    """Sum item-level tax for one invoice/note across its itms list."""
    t = {c: 0.0 for c in TAX_COLS}
    for it in inv.get("itms", []) or []:
        det = it.get("itm_det", it) or {}
        for k, c in (("txval", "Taxable value"), ("iamt", "IGST"), ("camt", "CGST"),
                     ("samt", "SGST"), ("csamt", "Cess")):
            t[c] += _num(det.get(k))
    return t


def gstr1_b2b_partywise(json_objs: list) -> pd.DataFrame:
    """Recipient-GSTIN-wise taxable value + tax reported in GSTR-1, from one or more
    GSTR-1 JSON objects (live or offline). Nets B2B invoices, B2B amendments (b2ba) and
    credit/debit notes (cdnr/cdnra): credit notes reduce, debit notes add.
    Returns columns: Party GSTIN, GSTR-1 invoices, Taxable value, IGST, CGST, SGST, Cess."""
    agg: dict[str, dict] = {}

    def add(ctin, taxd, n_inv=1, sign=1):
        ctin = str(ctin or "").strip().upper()
        if not ctin:
            return
        row = agg.setdefault(ctin, {"Party GSTIN": ctin, "GSTR-1 invoices": 0,
                                    **{c: 0.0 for c in TAX_COLS}})
        row["GSTR-1 invoices"] += n_inv
        for c in TAX_COLS:
            row[c] += sign * taxd[c]

    for js in json_objs or []:
        d = _data(js) or {}
        for sec in ("b2b", "b2ba"):                       # invoices + amendments (add)
            for party in _find_list(d, sec) or []:
                invs = party.get("inv", []) or []
                for inv in invs:
                    add(party.get("ctin"), _inv_tax(inv), 1, +1)
        for sec in ("cdnr", "cdnra"):                     # credit/debit notes
            for party in _find_list(d, sec) or []:
                for nt in party.get("nt", []) or []:
                    # ntty: 'C' credit note (reduces), 'D' debit note (adds)
                    sign = -1 if str(nt.get("ntty", "")).upper().startswith("C") else +1
                    add(party.get("ctin"), _inv_tax(nt), 0, sign)

    if not agg:
        return pd.DataFrame(columns=["Party GSTIN", "GSTR-1 invoices"] + TAX_COLS)
    out = pd.DataFrame(list(agg.values()))
    return out.sort_values("Taxable value", ascending=False).reset_index(drop=True)


def load_gstr1_b2b(folder) -> pd.DataFrame:
    """Scan a year folder for GSTR-1 JSON files (the downloaded full-JSON offline files and
    anything extracted from their zips) and build the counterparty-wise GSTR-1 table."""
    folder = Path(folder)
    objs, seen = [], set()
    patterns = ["*GSTR1*json*.json", "*GSTR1*.json", "**/*GSTR1*.json",
                "GSTR1_*extracted/**/*.json", "**/returns*.json"]
    files = []
    for pat in patterns:
        files.extend(folder.glob(pat))
    for f in files:
        rp = str(f.resolve())
        if rp in seen or not f.is_file():
            continue
        seen.add(rp)
        try:
            objs.append(json.loads(f.read_text(encoding="utf-8", errors="ignore")))
        except Exception:
            continue
    return gstr1_b2b_partywise(objs)


# =============================================================== GSTR-2B
_2B_COLS = {"ctin": "Supplier GSTIN", "trdnm": "Supplier name", "supprd": "Supplier period",
            "supfildt": "Supplier filing date", "inum": "Invoice no", "ntnum": "Note no", "docnum": "Document no",
            "dt": "Date", "docdt": "Date", "typ": "Type", "suptyp": "Supply type", "val": "Invoice value",
            "txval": "Taxable value", "igst": "IGST", "cgst": "CGST", "sgst": "SGST", "cess": "Cess",
            "pos": "Place of supply", "rev": "Reverse charge", "itcavl": "ITC available", "itcelg": "ITC eligible",
            "rsn": "Reason", "diffprcnt": "Diff %", "imsStatus": "IMS status", "boenum": "BoE no",
            "boedt": "BoE date", "portcode": "Port code", "refdt": "Ref date", "isamd": "Amended",
            "oinum": "Orig. invoice no", "oidt": "Orig. date", "ontnum": "Orig. note no", "ontdt": "Orig. note date"}


def _flatten_section(records, period):
    out = []
    for rec in records or []:
        if not isinstance(rec, dict):
            continue
        parent = {k: v for k, v in rec.items() if not isinstance(v, (list, dict))}
        children = [v for v in rec.values() if isinstance(v, list) and v and isinstance(v[0], dict)]
        if children:
            for lst in children:
                for ch in lst:
                    flat = {k: v for k, v in ch.items() if not isinstance(v, (list, dict))}
                    # nested item lists (rate-wise) -> sum numeric tax fields
                    for v in ch.values():
                        if isinstance(v, list) and v and isinstance(v[0], dict):
                            for it in v:
                                det = it.get("itm_det", it)
                                for k in ("txval", "igst", "cgst", "sgst", "cess"):
                                    if k in det and k not in flat:
                                        flat[k] = 0
                                    if k in det:
                                        flat[k] = _num(flat.get(k)) + _num(det.get(k))
                    out.append({"2B period": period, **parent, **flat})
        else:
            out.append({"2B period": period, **parent})
    return out


def _walk_tax(node, path, rows, period):
    if not isinstance(node, dict):
        return
    if any(k in node for k in ("igst", "cgst", "sgst", "cess")):
        rows.append({"2B period": period, "Head": path,
                     "Taxable value": _num(node.get("txval")), "IGST": _num(node.get("igst")),
                     "CGST": _num(node.get("cgst")), "SGST": _num(node.get("sgst")), "Cess": _num(node.get("cess"))})
    for k, v in node.items():
        if isinstance(v, dict):
            _walk_tax(v, f"{path}.{k}" if path else k, rows, period)


def gstr2b_frames(data_by_prd: dict[str, dict]):
    """Returns dict of sheet name -> DataFrame, plus monthly ITC totals df."""
    sections: dict[str, list] = {}
    itc_rows = []
    for prd, js in sorted(data_by_prd.items(), key=lambda kv: prd_label(kv[0])):
        d = _data(js) or {}
        per = prd_label(prd)
        for sec, recs in (d.get("docdata") or {}).items():
            sections.setdefault(sec.upper(), []).extend(_flatten_section(recs, per))
        _walk_tax(d.get("itcsumm") or {}, "", itc_rows, per)

    sheets = {}
    for sec, rows in sections.items():
        df = pd.DataFrame(rows).rename(columns=_2B_COLS)
        df = df.loc[:, ~df.columns.duplicated()]
        for c in TAX_COLS + ["Invoice value"]:
            if c in df:
                df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)
        sheets[f"2B_{sec}"] = df

    itc = pd.DataFrame(itc_rows)
    totals = pd.DataFrame()
    if not itc.empty:
        itc["Head"] = itc["Head"].str.replace("itcavl", "ITC available", regex=False) \
                                 .str.replace("itcunavl", "ITC not available", regex=False) \
                                 .str.replace("itcrev", "ITC reversal", regex=False)
        sheets["2B_ITC_Summary"] = itc
        lvl1 = itc[itc["Head"].str.count(r"\.") == 1]
        avl = lvl1[lvl1["Head"].str.startswith("ITC available")].copy()
        avl["Component"] = avl["Head"].str.split(".").str[1]
        if not avl.empty:
            avl["Tax"] = avl[["IGST", "CGST", "SGST", "Cess"]].sum(axis=1)
            totals = avl.pivot_table(index="2B period", columns="Component", values="Tax", aggfunc="sum", fill_value=0)
            add = [c for c in totals.columns if c != "othersup"]
            totals["2B ITC (excl. othersup)"] = totals[add].sum(axis=1)
            totals = totals.reset_index().rename(columns={"2B period": "Period"})

    b2b = sheets.get("2B_B2B")
    if b2b is not None and not b2b.empty and "Supplier GSTIN" in b2b:
        grp = b2b.groupby(["Supplier GSTIN", "Supplier name"], dropna=False)
        sup = grp.agg(Documents=("Taxable value", "size"),
                      **{c: (c, "sum") for c in TAX_COLS if c in b2b}).reset_index()
        sup["Total tax"] = sup[[c for c in ("IGST", "CGST", "SGST", "Cess") if c in sup]].sum(axis=1)
        sheets["2B_Supplierwise"] = sup.sort_values("Total tax", ascending=False)
    return sheets, totals


# =============================================================== GSTR-9
def gstr9_frame(js):
    d = _data(js) or {}
    rows = []
    for tbl, items in d.items():
        if not isinstance(items, dict):
            continue
        for item, vals in items.items():
            if not isinstance(vals, dict):
                continue
            row = {"Table": tbl.replace("table", "Table "), "Item": item}
            for k, v in vals.items():
                row[_KEYMAP.get(k, k)] = _num(v)
            rows.append(row)
    return pd.DataFrame(rows)


# =============================================================== RECONCILIATION
def reconciliation(g1_tot: pd.DataFrame, g3b_sum: pd.DataFrame, g2b_tot: pd.DataFrame):
    out = pd.DataFrame()
    if g3b_sum is not None and not g3b_sum.empty:
        out = g3b_sum[["Period", "Outward taxable value 3.1(a)", "Zero rated value 3.1(b)", "Output tax 3.1(a)",
                       "ITC available 4A"]].copy()
    if g1_tot is not None and not g1_tot.empty:
        g1 = g1_tot.copy()
        g1["GSTR-1 tax"] = g1[["IGST", "CGST", "SGST", "Cess"]].sum(axis=1)
        g1 = g1.rename(columns={"Taxable value": "GSTR-1 taxable value"})[["Period", "GSTR-1 taxable value", "GSTR-1 tax"]]
        out = g1 if out.empty else out.merge(g1, on="Period", how="outer")
    if g2b_tot is not None and not g2b_tot.empty:
        g2 = g2b_tot[["Period", "2B ITC (excl. othersup)"]]
        out = g2 if out.empty else out.merge(g2, on="Period", how="outer")
    if out.empty:
        return out
    out = out.fillna(0).sort_values("Period").reset_index(drop=True)
    if {"GSTR-1 taxable value", "Outward taxable value 3.1(a)"} <= set(out):
        out["Diff taxable (GSTR-1 - 3B)"] = out["GSTR-1 taxable value"] - out["Outward taxable value 3.1(a)"] \
            - out.get("Zero rated value 3.1(b)", 0)
        out["Diff tax (GSTR-1 - 3B)"] = out["GSTR-1 tax"] - out["Output tax 3.1(a)"]
    if {"2B ITC (excl. othersup)", "ITC available 4A"} <= set(out):
        out["Excess ITC in 3B over 2B"] = (out["ITC available 4A"] - out["2B ITC (excl. othersup)"]).round(2)

    def flag(r):
        f = []
        if abs(r.get("Diff tax (GSTR-1 - 3B)", 0)) > 1:
            f.append("Output tax mismatch")
        if r.get("Excess ITC in 3B over 2B", 0) > 1:
            f.append("3B ITC exceeds 2B")
        return "; ".join(f) or "OK"
    out["Flag"] = out.apply(flag, axis=1)
    return out


# =============================================================== NOTICES & ORDERS
def _fy_of(d):
    """date -> '2025-26'."""
    if not d:
        return ""
    y = d.year if d.month >= 4 else d.year - 1
    return f"{y}-{str(y + 1)[2:]}"


def notices_frame(notices: list, cases: list, today: date | None = None):
    """
    notices: /services/auth/api/get/notices response (list)
    cases:   /litserv/auth/api/case/task/get response (list of case tasks)
    Returns (dataframe, summary dict). 'Pending' = unread, OPEN case, or a future
    action due date.
    """
    today = today or date.today()
    rows = []
    for n in notices or []:
        if not isinstance(n, dict):
            continue
        di = parse_dmy(n.get("dtOfIssue"))
        due = parse_dmy(n.get("dueDate")) if str(n.get("dueDate", "")).upper() != "NA" else None
        unread = str(n.get("isRead", "")).upper() == "N"
        pending = unread or (due is not None and due >= today)
        rows.append({
            "FY": _fy_of(di), "Type": n.get("type", ""), "Reference": n.get("noticeOrderId", ""),
            "Description": n.get("descr", ""), "Issued by": n.get("issuedBy", ""),
            "Date of issue": di, "Due date": due,
            "Status": "Pending" if pending else "Seen",
            "Read": "No" if unread else "Yes",
            "Amount": n.get("amount", ""), "Form": n.get("applnCd", ""),
        })
    for c in cases or []:
        if not isinstance(c, dict):
            continue
        asg = c.get("assignmentDt")
        di = date.fromtimestamp(asg / 1000) if isinstance(asg, (int, float)) and asg else None
        st = str(c.get("status", "")).upper()
        rows.append({
            "FY": _fy_of(di), "Type": "Case: " + (c.get("caseTypeName") or c.get("caseTpeCd") or ""),
            "Reference": c.get("refId") or c.get("arn", ""), "Description": c.get("taskDesc", ""),
            "Issued by": "Tax authority", "Date of issue": di, "Due date": None,
            "Status": "Pending" if st in ("OPEN", "PENDING", "") else "Closed",
            "Read": "No" if str(c.get("isRead", "")).upper() == "N" else "Yes",
            "Amount": "", "Form": c.get("caseTpeCd", ""),
        })
    df = pd.DataFrame(rows, columns=["FY", "Type", "Reference", "Description", "Issued by",
                                     "Date of issue", "Due date", "Status", "Read", "Amount", "Form"])
    if not df.empty:
        df = df.sort_values("Date of issue", ascending=False, na_position="last").reset_index(drop=True)
    summ = {"total": int(len(df)), "pending": int((df["Status"] == "Pending").sum()) if not df.empty else 0,
            "unread": int((df["Read"] == "No").sum()) if not df.empty else 0,
            "orders": int((df["Type"] == "Order").sum()) if not df.empty else 0}
    return df, summ


# =============================================================== EXCEL WRITER
def write_excel(path: Path, sheets: dict[str, pd.DataFrame], title: str = ""):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    sheets = {k: v for k, v in sheets.items() if v is not None and not v.empty}
    if not sheets:
        return None
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            sname = name[:31]
            df.to_excel(xw, sheet_name=sname, index=False, startrow=2)
            ws = xw.sheets[sname]
            ws["A1"] = f"{title} - {name}" if title else name
            ws["A1"].font = Font(bold=True, size=12, color="1F4E79")
            for c in ws[3]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = hdr_fill
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, col in enumerate(df.columns, 1):
                width = max([len(str(col))] + [len(f"{v:,.2f}") if isinstance(v, float) else len(str(v))
                                               for v in df[col].head(300)]) + 2
                ws.column_dimensions[get_column_letter(i)].width = min(max(width, 9), 48)
                if pd.api.types.is_numeric_dtype(df[col]):
                    for row in ws.iter_rows(min_row=4, min_col=i, max_col=i):
                        row[0].number_format = "#,##0.00"
            ws.freeze_panes = "B4"
            ws.auto_filter.ref = f"A3:{get_column_letter(len(df.columns))}{len(df) + 3}"
    return path


def save_json(path: Path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=1, ensure_ascii=False), encoding="utf-8")
