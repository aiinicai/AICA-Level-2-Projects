"""Run with:  python test_core.py   (or: python -m pytest)
Developed by CA Sapan"""
import os
import tempfile
from datetime import date, timedelta

import core

TODAY = date(2026, 10, 17)


def fresh():
    return core.connect(":memory:")


def expect_error(fn, *a, **k):
    try:
        fn(*a, **k)
    except ValueError:
        return
    raise AssertionError(f"{fn.__name__} should have raised ValueError")


def test_task_validation():
    c = fresh()
    expect_error(core.add_task, c, "  ", "Internal", TODAY)
    expect_error(core.add_task, c, "x", "Nonsense", TODAY)
    expect_error(core.add_task, c, "x", "Internal", TODAY, planned_min=-5)


def test_additional_flag_auto_set():
    c = fresh()
    tid = core.add_task(c, "Ad-hoc", "Additional / ad-hoc", TODAY)
    assert core.tasks_df(c).iloc[0]["is_additional"] == 1
    tid2 = core.add_task(c, "Client job flagged extra", "Client work", TODAY, additional=True)
    assert core.tasks_df(c).set_index("id").loc[tid2, "is_additional"] == 1


def test_carry_forward_keeps_original_date():
    c = fresh()
    tid = core.add_task(c, "T", "Internal", TODAY)
    core.carry_forward(c, tid, TODAY + timedelta(days=3))
    r = core.tasks_df(c).iloc[0]
    assert r["original_date"] == "2026-10-17"
    assert r["scheduled_date"] == "2026-10-20"
    assert r["carried_count"] == 1
    assert r["flag"] == "Carried forward"
    expect_error(core.carry_forward, c, tid, TODAY)          # not later
    core.set_status(c, tid, "Done", TODAY + timedelta(days=4))
    expect_error(core.carry_forward, c, tid, TODAY + timedelta(days=9))  # done


def test_age_days_uses_completion_date():
    c = fresh()
    tid = core.add_task(c, "T", "Internal", TODAY - timedelta(days=10))
    core.set_status(c, tid, "Done", TODAY - timedelta(days=4))
    assert core.tasks_df(c).iloc[0]["age_days"] == 6


def test_rollover_only_open_and_past():
    c = fresh()
    a = core.add_task(c, "old open", "Internal", TODAY - timedelta(days=2))
    b = core.add_task(c, "old done", "Internal", TODAY - timedelta(days=2))
    d = core.add_task(c, "today", "Internal", TODAY)
    core.set_status(c, b, "Done", TODAY - timedelta(days=1))
    assert core.rollover_overdue(c, TODAY) == 1
    df = core.tasks_df(c).set_index("id")
    assert df.loc[a, "scheduled_date"] == "2026-10-17" and df.loc[a, "carried_count"] == 1
    assert df.loc[b, "carried_count"] == 0 and df.loc[d, "carried_count"] == 0
    assert core.rollover_overdue(c, TODAY) == 0               # idempotent


def test_time_logging_and_actuals():
    c = fresh()
    tid = core.add_task(c, "T", "Client work", TODAY, planned_min=60)
    core.log_time(c, tid, TODAY, 30)
    core.log_time(c, tid, TODAY, 45)
    assert core.tasks_df(c).iloc[0]["actual_min"] == 75
    expect_error(core.log_time, c, tid, TODAY, 0)
    expect_error(core.log_time, c, 999, TODAY, 10)


def test_delete_task_cascades():
    c = fresh()
    tid = core.add_task(c, "T", "Internal", TODAY)
    core.log_time(c, tid, TODAY, 10)
    core.delete_task(c, tid)
    assert core.time_df(c).empty


def test_kpi_validation():
    c = fresh()
    expect_error(core.add_kpi, c, "", "Monthly", 2026, 10, 5, "Tasks completed")
    expect_error(core.add_kpi, c, "k", "Weekly", 2026, 10, 5, "Tasks completed")
    expect_error(core.add_kpi, c, "k", "Monthly", 2026, None, 5, "Tasks completed")
    expect_error(core.add_kpi, c, "k", "Monthly", 2026, 13, 5, "Tasks completed")
    expect_error(core.add_kpi, c, "k", "Monthly", 2026, 10, 0, "Tasks completed")
    expect_error(core.add_kpi, c, "k", "Monthly", 2026, 10, 5, "Bogus")


def test_kpi_tasks_completed_and_period_boundaries():
    c = fresh()
    core.add_kpi(c, "Done in Oct", "Monthly", 2026, 10, 4, "Tasks completed")
    for day, in_oct in [(date(2026, 9, 30), False), (date(2026, 10, 1), True),
                        (date(2026, 10, 31), True), (date(2026, 11, 1), False)]:
        tid = core.add_task(c, f"t{day}", "Internal", day)
        core.set_status(c, tid, "Done", day)
    p = core.kpi_progress(c, 2026, 10)
    assert p.iloc[0]["Actual"] == 2 and p.iloc[0]["Achievement %"] == 50.0


def test_kpi_hours_by_category_and_december():
    c = fresh()
    core.add_kpi(c, "Client hrs Dec", "Monthly", 2026, 12, 10, "Hours logged", "Client work")
    t1 = core.add_task(c, "a", "Client work", date(2026, 12, 31))
    t2 = core.add_task(c, "b", "Internal", date(2026, 12, 31))
    core.log_time(c, t1, date(2026, 12, 31), 120)   # counted, last day of year
    core.log_time(c, t2, date(2026, 12, 31), 600)   # wrong category
    core.log_time(c, t1, date(2027, 1, 1), 600)     # wrong period
    p = core.kpi_progress(c, 2026, 12)
    assert p.iloc[0]["Actual"] == 2.0 and p.iloc[0]["Achievement %"] == 20.0


def test_kpi_manual_and_yearly_visibility():
    c = fresh()
    m = core.add_kpi(c, "Calls", "Monthly", 2026, 10, 8, "Manual count")
    core.add_kpi(c, "Yearly tasks", "Yearly", 2026, None, 100, "Tasks completed")
    core.add_kpi(c, "Other month", "Monthly", 2026, 9, 3, "Tasks completed")
    core.add_kpi_manual(c, m, date(2026, 10, 5), 3)
    core.add_kpi_manual(c, m, date(2026, 10, 9), 2)
    core.add_kpi_manual(c, m, date(2026, 9, 30), 99)   # outside month
    p = core.kpi_progress(c, 2026, 10)
    assert set(p["KPI"]) == {"Calls", "Yearly tasks"}
    assert p.set_index("KPI").loc["Calls", "Actual"] == 5.0
    other = core.add_kpi(c, "Count", "Monthly", 2026, 10, 1, "Tasks completed")
    expect_error(core.add_kpi_manual, c, other, TODAY, 1)


def test_summary_additional_percentage():
    c = fresh()
    a = core.add_task(c, "regular", "Client work", TODAY)
    b = core.add_task(c, "extra", "Additional / ad-hoc", TODAY)
    core.log_time(c, a, TODAY, 180)
    core.log_time(c, b, TODAY, 60)
    s = core.summary(c, TODAY, TODAY)
    assert s["total_hours"] == 4.0 and s["additional_hours"] == 1.0 and s["additional_pct"] == 25.0


def test_summary_empty_database_does_not_crash():
    c = fresh()
    s = core.summary(c, TODAY, TODAY)
    assert s["total_hours"] == 0 and s["additional_pct"] == 0.0 and s["tasks_done"] == 0
    assert core.kpi_progress(c, 2026, 10).empty


def test_exports_with_data_and_empty():
    for seeded in (True, False):
        c = fresh()
        if seeded:
            core.seed_sample(c, TODAY)
        s, e = date(2026, 10, 1), date(2026, 10, 31)
        with tempfile.TemporaryDirectory() as d:
            x = core.export_excel(c, os.path.join(d, "a.xlsx"), s, e)
            w = core.export_word(c, os.path.join(d, "a.docx"), s, e)
            assert os.path.getsize(x) > 3000 and os.path.getsize(w) > 3000
            from openpyxl import load_workbook
            wb = load_workbook(x)
            assert {"Summary", "KPIs", "Tasks", "Time entries", "About"} <= set(wb.sheetnames)
            from docx import Document
            text = "\n".join(p.text for p in Document(w).paragraphs)
            assert core.CREDIT in text


def test_seed_sample_is_idempotent_and_sensible():
    c = fresh()
    assert core.seed_sample(c, TODAY) is True
    assert core.seed_sample(c, TODAY) is False
    td = core.tasks_df(c)
    assert len(td) == 10 and (td["status"] == "Done").sum() == 7
    open_old = td[(td["status"] != "Done")]
    assert (open_old["scheduled_date"] >= "2026-10-17").all()
    assert open_old["carried_count"].max() >= 1
    assert core.summary(c, date(2026, 10, 1), date(2026, 10, 31))["additional_pct"] > 0


if __name__ == "__main__":
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for t in tests:
        try:
            t(); print("PASS", t.__name__)
        except Exception as ex:
            failed += 1; print("FAIL", t.__name__, "->", repr(ex))
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    raise SystemExit(1 if failed else 0)
