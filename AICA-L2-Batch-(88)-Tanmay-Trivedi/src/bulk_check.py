"""
bulk_check.py
---------------
Runs the same Section 17(5) rules engine (src/rules.py) over a whole Excel
sheet of expense line items in one pass - for when you're reviewing a full
purchase/expense ledger rather than checking one item at a time.

Since a bulk pass can't ask follow-up questions interactively, the Excel
template carries the follow-up answers as extra Y/N columns. Any condition
left blank simply can't be resolved, so that row is flagged "NEEDS REVIEW"
rather than guessed at.

Each row also gets an independent "RCM Alert" column (src/rcm.py) flagging
possible reverse-charge liability under Section 9(3)/9(4) - a separate
question from ITC eligibility, so it's shown alongside the Section 17(5)
verdict rather than folded into it.
"""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .rcm import match_rcm
from .rules import BLOCKED, ELIGIBLE, NEEDS_INPUT, evaluate, match_categories

# Excel column header -> internal condition key
CONDITION_COLUMNS = {
    "Seating capacity over 13? (Y/N)": "seating_over_13",
    "Used for further supply / passenger transport / driver training? (Y/N)": "used_for_exception",
    "Same-category outward supply or composite/mixed supply? (Y/N)": "same_category_outward_supply",
    "Obligatory for employer under law? (Y/N)": "obligatory_under_law",
    "For plant and machinery, not a building/civil structure? (Y/N)": "for_plant_and_machinery",
    "Input service for further supply of works contract? (Y/N)": "further_supply_of_works_contract",
    "Goods imported by the non-resident taxable person? (Y/N)": "imported_goods",
}

TEMPLATE_COLUMNS = ["Description"] + list(CONDITION_COLUMNS.keys())

NAVY = "0A2540"
GREEN = "C6EFCE"
RED = "FFC7CE"
AMBER = "FFF2CC"
WHITE = "FFFFFF"


def _normalise_yn(val) -> str | None:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    s = str(val).strip().upper()
    if s in ("Y", "YES", "TRUE", "1"):
        return "Y"
    if s in ("N", "NO", "FALSE", "0"):
        return "N"
    return None


def make_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bulk Check Template"
    for i, col in enumerate(TEMPLATE_COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=col)
        c.font = Font(bold=True, color=WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    example_rows = [
        ["Car purchased for use by the managing director", "N", "N", "", "", "", "", ""],
        ["Cab purchased by a taxi rental business", "N", "Y", "", "", "", "", ""],
        ["Canteen food for staff, mandated under Factories Act", "", "", "N", "Y", "", "", ""],
        ["Building construction for own office", "", "", "", "", "N", "", ""],
        ["Office stationery purchase", "", "", "", "", "", "", ""],
    ]
    for r, row in enumerate(example_rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r, column=c_idx, value=val)
    ws.column_dimensions["A"].width = 45
    for i in range(2, len(TEMPLATE_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_bulk_check(df: pd.DataFrame) -> pd.DataFrame:
    if "Description" not in df.columns:
        raise ValueError("The uploaded sheet needs a 'Description' column.")

    results = []
    for _, row in df.iterrows():
        desc = str(row.get("Description") or "").strip()
        if not desc:
            continue

        answers = {}
        for col_header, key in CONDITION_COLUMNS.items():
            if col_header in df.columns:
                answers[key] = _normalise_yn(row.get(col_header))

        matches = match_categories(desc)
        rcm_matches = match_rcm(desc)
        rcm_display = (
            "; ".join(f"{c.section} - {c.title}" for c in rcm_matches) if rcm_matches else "-"
        )

        if not matches:
            results.append({
                "Description": desc,
                "Matched Clause": "-",
                "Matched Category": "No Section 17(5) category matched",
                "Verdict": ELIGIBLE,
                "Reasoning": "No blocked-credit category matched this description - treated as eligible "
                             "subject to the general Section 16 conditions. Verify manually, especially if "
                             "the description was brief.",
                "RCM Alert": rcm_display,
            })
            continue

        note_prefix = ""
        if len(matches) > 1:
            titles = "; ".join(f"{c.clause} {c.title}" for c in matches)
            note_prefix = f"Multiple categories matched ({titles}) - showing the first; please verify manually. "

        category = matches[0]
        result = evaluate(category, answers)
        verdict = result["verdict"]
        if verdict == NEEDS_INPUT:
            verdict_display = "NEEDS REVIEW"
            reasoning = note_prefix + result["reasoning"] + f" Unanswered: \"{result['next_question']}\""
        else:
            verdict_display = verdict
            reasoning = note_prefix + result["reasoning"]

        results.append({
            "Description": desc,
            "Matched Clause": category.clause,
            "Matched Category": category.title,
            "Verdict": verdict_display,
            "Reasoning": reasoning,
            "RCM Alert": rcm_display,
        })

    return pd.DataFrame(results)


def to_excel_bytes(result_df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Blocked Credit Report"
    headers = list(result_df.columns)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color=WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"

    for r_idx, record in enumerate(result_df.itertuples(index=False), start=2):
        row_dict = dict(zip(headers, record))
        verdict = str(row_dict.get("Verdict", ""))
        if "BLOCKED" in verdict:
            fill = RED
        elif "NEEDS REVIEW" in verdict:
            fill = AMBER
        elif "ELIGIBLE" in verdict:
            fill = GREEN
        else:
            fill = WHITE
        for c_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=row_dict[h])
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = {"A": 45, "B": 14, "C": 40, "D": 16, "E": 60, "F": 40}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
