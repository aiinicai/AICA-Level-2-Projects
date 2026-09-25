"""
generate_sample_ledger.py
----------------------------
Creates a small, entirely fictional sample expense ledger exercising most
Section 17(5) categories - including a couple of rows with the optional
condition columns deliberately left blank, to demonstrate the "NEEDS
REVIEW" flag in the bulk-check report. Two rows (legal fees, GTA freight)
are also worded to trip the independent RCM Alert (src/rcm.py) while
remaining ITC ELIGIBLE under Section 17(5) - showing the two checks are
independent of each other.
"""

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .bulk_check import TEMPLATE_COLUMNS

OUT_PATH = Path(__file__).resolve().parent.parent / "sample_data" / "Sample_Expense_Ledger.xlsx"

# Description, seating>13, used_for_exception, same_cat_supply, obligatory_law, plant&machinery, further_wc_supply, imported_goods
ROWS = [
    ["Toyota Innova car purchased for the CFO's personal use", "N", "N", "", "", "", "", ""],
    ["Tempo Traveller (17-seater) purchased for staff transport", "Y", "", "", "", "", "", ""],
    ["Motor insurance premium for the above CFO's car", "", "N", "", "", "", "", ""],
    ["Outdoor catering for the annual client dinner", "", "", "N", "N", "", "", ""],
    ["Canteen food supply, mandatory under the Factories Act for this factory", "", "", "N", "Y", "", "", ""],
    ["Group health insurance premium for all employees", "", "", "N", "N", "", "", ""],
    ["Corporate membership of a golf club for the MD", "", "", "", "", "", "", ""],
    ["Leave travel concession paid to employees for their vacation", "", "", "", "", "", "", ""],
    ["Works contract service for installing new factory plant and machinery", "", "", "", "", "Y", "", ""],
    ["Works contract for constructing the new administrative office building", "", "", "", "", "N", "N", ""],
    ["Goods purchased while registered as a non-resident taxable person, imported from abroad", "", "", "", "", "", "", "Y"],
    ["Inventory written off after being damaged in a warehouse leak", "", "", "", "", "", "", ""],
    ["Diwali gift hampers distributed to employees", "", "", "", "", "", "", ""],
    ["Office stationery and printing purchase", "", "", "", "", "", "", ""],
    ["Legal fees paid to an advocate for a GST assessment matter", "", "", "", "", "", "", ""],
    ["Business-class air tickets for a director's client visit", "", "", "", "", "", "", ""],
    ["Freight paid to a goods transport agency for outward dispatch of finished goods", "", "", "", "", "", "", ""],
]

NAVY = "0A2540"
WHITE = "FFFFFF"


def build():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expense Ledger"
    for i, col in enumerate(TEMPLATE_COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=col)
        c.font = Font(bold=True, color=WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(ROWS, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r, column=c_idx, value=val)
    ws.column_dimensions["A"].width = 55
    for i in range(2, len(TEMPLATE_COLUMNS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 22
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    p = build()
    print(f"Wrote: {p}")
