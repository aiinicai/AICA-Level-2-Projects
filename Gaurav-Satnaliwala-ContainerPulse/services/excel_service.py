from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

REQUIRED_COLUMNS = ["Container Number", "Shipping Line"]


def read_container_file(source) -> pd.DataFrame:
    frame = pd.read_excel(source)
    normalized = {str(column).strip().lower().replace("_", " "): column for column in frame.columns}
    missing = [column for column in REQUIRED_COLUMNS if column.lower() not in normalized]
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(missing)}")
    frame = frame.rename(columns={normalized[column.lower()]: column for column in REQUIRED_COLUMNS})[REQUIRED_COLUMNS].copy()
    frame["Container Number"] = frame["Container Number"].astype(str).str.strip().str.upper()
    frame["Shipping Line"] = frame["Shipping Line"].astype(str).str.strip().str.upper().replace({
        "MAERSK LINE": "MAERSK",
        "MEDITERRANEAN SHIPPING COMPANY": "MSC",
    })
    frame = frame[frame["Container Number"].ne("") & frame["Container Number"].ne("NAN")].reset_index(drop=True)
    if frame.empty:
        raise ValueError("The uploaded file does not contain any valid container rows.")
    return frame


def results_to_excel(results: pd.DataFrame) -> bytes:
    output = BytesIO()
    export = results.copy()
    with pd.ExcelWriter(output, engine="openpyxl", datetime_format="DD-MMM-YYYY HH:MM") as writer:
        export.to_excel(writer, index=False, sheet_name="Tracking Results")
        sheet = writer.book["Tracking Results"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="16324F")
            cell.font = Font(color="FFFFFF", bold=True)
        for index, column in enumerate(export.columns, start=1):
            values = [len(str(value)) + 2 for value in export[column].fillna("").head(100)]
            sheet.column_dimensions[get_column_letter(index)].width = min(max([len(str(column)) + 3, *values]), 52)
        result_letter = get_column_letter(export.columns.get_loc("Tracking Result") + 1)
        whole_range = f"A2:{get_column_letter(len(export.columns))}{len(export) + 1}"
        sheet.conditional_formatting.add(
            whole_range,
            FormulaRule(formula=[f'${result_letter}2<>"Success"'], fill=PatternFill("solid", fgColor="FFF3CD")),
        )
    return output.getvalue()
