from __future__ import annotations

import hashlib
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from services.connector_setup import choose_button, choose_field_selectors, choose_input
from services.excel_service import read_container_file, results_to_excel
from services.registry_service import ShippingLineRegistry
from services.status_service import classify_container_state
from services.tracking_service import run_tracking_batch
from tracker.common import TrackingResult
from tracker.generic import security_restriction
from tracker.maersk import _parse_maersk_visible_text
from tracker.msc import _parse_msc_rendered_text
from tracker.runner import CONNECTOR_REGISTRY, track_container


ROOT = Path(__file__).resolve().parents[1]
DUMMY_CONTAINERS = ["DEMO0000001", "DEMO0000002"]
EXPECTED_CONNECTOR_HASHES = {
    "msc.py": "9567603d6ac55e18c25f68370b77da2c67f929cd8128819c0ba8564de9d0f730",
    "maersk.py": "fc5da36e9a0a0129cb519e96d1525ee864303c352231505a8b894ae4083dfe7e",
}


class ContainerPulseCapstoneTests(unittest.TestCase):
    def test_dummy_sample_and_master_workbooks(self):
        for workbook in [
            ROOT / "sample_data" / "sample_containers.xlsx",
            ROOT / "sample_containerpulse_template.xlsx",
            ROOT / "data" / "master_containers.xlsx",
        ]:
            frame = read_container_file(workbook)
            self.assertEqual(frame.columns.tolist(), ["Container Number", "Shipping Line"])
            self.assertEqual(frame["Container Number"].tolist(), DUMMY_CONTAINERS)
            self.assertEqual(frame["Shipping Line"].tolist(), ["MSC", "MAERSK"])

    def test_current_only_tracking_batch(self):
        frame = read_container_file(ROOT / "sample_data" / "sample_containers.xlsx")

        def fake_tracker(container, carrier, headless=False):
            return TrackingResult(
                container,
                carrier,
                True,
                True,
                True,
                None,
                "30 Oct 2099 06:00",
                "In Transit",
                "Demo Port",
                "DEMO VESSEL",
                "Demo tracking event",
            )

        results = run_tracking_batch(frame, tracker=fake_tracker)
        expected = [
            "Container Number", "Shipping Line", "Current Status", "Current Location",
            "Vessel", "Latest ETA", "Last Tracking Event", "Checked At",
            "Tracking Result", "Error",
        ]
        self.assertEqual(results.columns.tolist(), expected)
        self.assertTrue(results["Tracking Result"].eq("Success").all())
        self.assertFalse(any(column.startswith("Previous") for column in results.columns))

    def test_dashboard_classification(self):
        cases = [
            ({"Tracking Result": "Success", "Current Status": "Full Transshipment Discharged"}, "In Transit"),
            ({"Tracking Result": "Success", "Current Status": "Estimated Time of Arrival"}, "In Transit"),
            ({"Tracking Result": "Success", "Current Status": "Shipment delivered"}, "Arrived / Completed"),
            ({"Tracking Result": "Failed", "Current Status": "Shipment delivered"}, "Tracking Error"),
        ]
        for row, expected in cases:
            with self.subTest(row=row):
                self.assertEqual(classify_container_state(row), expected)

    def test_unsupported_carrier_clean_error(self):
        result = track_container("DEMO0000003", "UNSUPPORTED LINE", headless=False)
        self.assertFalse(result.success)
        self.assertEqual(result.error, "Shipping line not currently supported")

    def test_registry_crud_and_statuses(self):
        with tempfile.TemporaryDirectory() as temp:
            registry = ShippingLineRegistry(Path(temp) / "registry.json")
            self.assertEqual(registry.ready_names(), ["MAERSK", "MSC"])
            registry.add("DEMO LINE", "https://example.com/tracking")
            self.assertEqual(registry.get("DEMO LINE")["connector_status"], "Not Configured")
            registry.save_setup_result("DEMO LINE", False, None, "Demo setup failure")
            self.assertEqual(registry.get("DEMO LINE")["connector_status"], "Setup Failed")
            config = {"input_selector": "#container", "button_selector": "#track", "field_selectors": {}}
            registry.save_setup_result("DEMO LINE", True, config)
            self.assertEqual(registry.get("DEMO LINE")["connector_status"], "Ready")
            registry.delete("DEMO LINE")
            self.assertIsNone(registry.get("DEMO LINE"))

    def test_setup_discovery_helpers_and_security_detection(self):
        records = [
            {"tag": "input", "type": "text", "id": "container-number", "name": "", "placeholder": "Enter container number", "aria": "", "dataTest": "", "text": "", "selector": "#container-number", "nextSelector": "", "nextText": ""},
            {"tag": "button", "type": "submit", "id": "track", "name": "", "placeholder": "", "aria": "", "dataTest": "", "text": "Track", "selector": "#track", "nextSelector": "", "nextText": ""},
            {"tag": "dt", "type": "", "id": "", "name": "", "placeholder": "", "aria": "", "dataTest": "", "text": "Current Status", "selector": "dl > dt", "nextSelector": "dl > dd", "nextText": "In Transit"},
        ]
        self.assertEqual(choose_input(records)["selector"], "#container-number")
        self.assertEqual(choose_button(records)["selector"], "#track")
        self.assertEqual(choose_field_selectors(records)["status"], "dl > dd")
        self.assertEqual(security_restriction("Please verify you are human"), "human-verification challenge displayed")
        self.assertEqual(security_restriction("Login required to view tracking"), "login is required by the carrier website")

    def test_synthetic_carrier_page_parsers(self):
        msc_text = """Latest move
Demo Port
POD ETA
30/10/2099
Date
20/09/2099
Demo Port
Gate Out
DEMO VESSEL DM123A
"""
        msc = _parse_msc_rendered_text(msc_text)
        self.assertEqual(msc["status"], "Gate Out")
        self.assertEqual(msc["location"], "Demo Port")
        self.assertEqual(msc["vessel"], "DEMO VESSEL")
        self.assertEqual(msc["eta"], "30/10/2099")

        maersk_text = """DEMO0000002
ORIGINPORT
Demo Terminal
Vessel departure (DEMO VESSEL / 001A)
30 Oct 2099 06:00
DESTPORT
Demo Terminal
Vessel arrival (DEMO VESSEL / 001A)
07 Nov 2099 06:00
"""
        maersk = _parse_maersk_visible_text(maersk_text, "DEMO0000002")
        self.assertTrue(maersk["accepted"])
        self.assertEqual(maersk["vessel"], "DEMO VESSEL")
        self.assertEqual(maersk["location"], "ORIGINPORT (planned origin)")
        self.assertEqual(maersk["eta"], "07 Nov 2099 06:00")

    def test_excel_export(self):
        result = pd.DataFrame([{
            "Container Number": "DEMO0000001", "Shipping Line": "MSC",
            "Current Status": "In Transit", "Current Location": "Demo Port",
            "Vessel": "DEMO VESSEL", "Latest ETA": pd.Timestamp("2099-10-30 06:00"),
            "Last Tracking Event": "Demo event", "Checked At": pd.Timestamp("2099-09-20 12:00"),
            "Tracking Result": "Success", "Error": "",
        }])
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "results.xlsx"
            output.write_bytes(results_to_excel(result))
            sheet = load_workbook(output)["Tracking Results"]
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(sheet["A2"].value, "DEMO0000001")

    def test_msc_and_maersk_connector_sources_are_unchanged(self):
        self.assertEqual(set(CONNECTOR_REGISTRY), {"MSC", "MAERSK"})
        for filename, expected_hash in EXPECTED_CONNECTOR_HASHES.items():
            digest = hashlib.sha256((ROOT / "tracker" / filename).read_bytes()).hexdigest()
            self.assertEqual(digest, expected_hash)

    def test_streamlit_application_starts(self):
        from streamlit.testing.v1 import AppTest

        app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=20).run()
        self.assertFalse(app.exception)
        self.assertIn("Container source", [item.value for item in app.subheader])


if __name__ == "__main__":
    unittest.main()
