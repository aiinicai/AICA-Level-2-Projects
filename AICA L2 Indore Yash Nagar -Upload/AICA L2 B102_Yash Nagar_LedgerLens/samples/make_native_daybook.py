"""Write the sample company's Day Book in TallyPrime's own Excel export layout (detailed view, FY 2025-26).

It copies the traits of a real export:
  - company, address, 'Day Book' and period rows, then a header split over two rows (Debit over Amount)
  - every voucher starts on its first ledger line, which carries the date, voucher type and number
  - the remaining ledger lines follow, then the narration on a row of its own
  - TallyPrime Educational Mode writes the word 'Educational' in place of every date except the 1st, 2nd and 31st

Usage: python samples/make_native_daybook.py
"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

import make_sample

HERE = Path(__file__).parent
SHOWN_DAYS = (1, 2, 31)      # the only days Educational Mode will show


def build(educational: bool = True) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Day Book"
    company = "Rajasthan Precision Industries Private Limited"
    ws["A1"], ws["A2"], ws["A3"], ws["A4"] = company, "Plot No. 45, Sitapura Industrial Area, Jaipur", "Day Book", "1-Apr-25 to 31-Mar-26"
    for r in range(1, 5):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    for c, h in zip("ABEFGH", ["Date", "Particulars", "Vch Type", "Vch No.", "Debit", "Credit"]):
        ws[f"{c}5"] = h
    ws["G6"], ws["H6"] = "Amount", "Amount"
    for row in ws["A5:H6"]:
        for c in row:
            c.font = Font(bold=True)

    r = 7
    for d, vtype, no, narr, lines in make_sample.numbered():
        for i, (ledger, amt, side) in enumerate(lines):
            if i == 0:
                shown = d.day in SHOWN_DAYS or not educational
                ws.cell(r, 1, datetime(d.year, d.month, d.day) if shown else "Educational")
                ws.cell(r, 5, vtype)
                ws.cell(r, 6, str(no))
                ws.cell(r, 1).number_format = "dd-mmm-yy"
            else:
                for col in (1, 3, 4, 5, 6):
                    ws.cell(r, col, "")
            ws.cell(r, 2, ledger)
            ws.cell(r, 7 if side == "Dr" else 8, float(amt))
            r += 1
        if narr:
            ws.cell(r, 2, narr)
            r += 1
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 14, 46
    return wb


if __name__ == "__main__":
    out = HERE / "tally_daybook_export.xlsx"
    build().save(out)
    print(f"Wrote {out}")
