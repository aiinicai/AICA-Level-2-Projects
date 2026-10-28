import io
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "samples"))
sys.path.insert(0, str(Path(__file__).parent))

import make_native_tb  # noqa: E402
from fsbuilder import pipeline as P  # noqa: E402
from fsbuilder.mapping import map_ledgers  # noqa: E402
from fsbuilder.model import build_model  # noqa: E402
from fsbuilder.tally_export import infer_group, parse_tally_trial_balance  # noqa: E402
from fsbuilder.tb import normalise  # noqa: E402
from test_builder_page import _run_upload  # noqa: E402


def native_bytes(**kw) -> bytes:
    buf = io.BytesIO()
    make_native_tb.build(**kw).save(buf)
    return buf.getvalue()


def read(data: bytes, **kw):
    return P.read_trial_balance(io.BytesIO(data), "export.xlsx", **kw)


def pipeline(tb):
    tb = normalise(tb)
    mp = map_ledgers(tb)
    return tb, mp, build_model(tb, mp)


# -------------------------------------------------------------- reading Tally's own layout
def test_native_export_is_read_with_company_period_and_sides_from_number_formats():
    n = parse_tally_trial_balance(native_bytes())
    assert n.company == "Rajasthan Precision Industries Private Limited"
    assert (n.start, n.end) == (date(2024, 4, 1), date(2025, 3, 31))
    df = n.df.set_index("name")
    assert df.loc["Share Capital", "closing_dr_cr"] == "Cr" and df.loc["Cash", "closing_dr_cr"] == "Dr"
    assert df.loc["Plant & Machinery", ["period_credit", "closing_amount"]].tolist() == [123258.96, 1726741.04]
    assert df.loc["Opening Stock", "closing_amount"] == 1_075_000


def test_native_export_balances_so_the_sides_were_read_correctly():
    tb = read(native_bytes())
    assert abs(sum(tb.ledgers["opening"])) < 0.01 and abs(sum(normalise(tb).ledgers["closing"])) < 0.01


def test_statements_from_a_native_export_match_the_known_books():
    tb, mp, m = pipeline(read(native_bytes()))
    tot = {r.label: r for r in m.bs if r.kind == "total"}
    assert tot["TOTAL ASSETS"].cur == tot["TOTAL EQUITY AND LIABILITIES"].cur == 15_086_416
    assert tot["TOTAL ASSETS"].prev == tot["TOTAL EQUITY AND LIABILITIES"].prev == 6_080_000
    assert m.profit == 2_489_416 and not tb.warnings and not any("does not" in b for b in m.basis)
    assert (mp["source"] == "unmapped").sum() == 0


def test_inference_agrees_with_tallys_real_groups_for_every_ledger():
    inferred, real = pipeline(read(native_bytes()))[1].set_index("ledger"), map_ledgers(normalise(P.load_sample())).set_index("ledger")
    for ledger in inferred.index.intersection(real.index):
        assert (inferred.loc[ledger, "line"], inferred.loc[ledger, "sub"]) == (real.loc[ledger, "line"], real.loc[ledger, "sub"]), ledger


def test_period_and_company_come_from_the_file_and_differences_are_reported():
    tb = read(native_bytes(), company="ABC Private Limited", start=date(2023, 4, 1), end=date(2024, 3, 31))
    assert (tb.from_date, tb.to_date) == (date(2024, 4, 1), date(2025, 3, 31)) and tb.company == "ABC Private Limited"
    text = " ".join(tb.adjustments)
    assert "taken from the file (01/04/2024 to 31/03/2025)" in text and "'Rajasthan Precision Industries Private Limited'" in text


def test_blank_company_uses_the_name_in_the_file():
    assert read(native_bytes()).company == "Rajasthan Precision Industries Private Limited"


def test_a_single_day_trial_balance_is_refused_with_instructions():
    data = native_bytes(period_text="For 1-Apr-26")
    with pytest.raises(ValueError) as e:
        read(data, start=date(2025, 4, 1), end=date(2026, 3, 31))
    msg = str(e.value)
    assert "single day (01/04/2026)" in msg and "F2" in msg and "01/04/2025 to 31/03/2026" in msg


# -------------------------------------------------------------- ledger groups
def test_guesses_are_flagged_only_where_they_are_really_guesses():
    tb, mp, m = pipeline(read(native_bytes()))
    flagged = {f.ledger for f in m.flags if f.kind == "Classification" and "guessed" in f.text}
    assert flagged == {"Apex Motors Pvt Ltd", "Jaipur Machine Tools", "Metro Auto Components Ltd", "National Fabricators",
                       "Sunrise Engineering Works", "Bright Bearings Co.", "Jaipur Packaging Works",
                       "Precision Tools Suppliers", "Rajasthan Alloys Pvt Ltd", "Sharma Steel Traders"}
    assert any("worked out from the ledger names" in a and "10 of them are guesses" in a for a in tb.adjustments)
    assert "Cash" not in flagged and "Share Capital" not in flagged


def test_a_ledger_list_replaces_guessing_with_the_real_groups():
    rows = [["Name", "Under"], ["Apex Motors Pvt Ltd", "Sundry Debtors"], ["Bright Bearings Co.", "Sundry Creditors"],
            ["Jaipur Machine Tools", "Sundry Debtors"]]
    wb = Workbook()
    wb.active.title = "Ledgers"
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    tb = read(native_bytes(), ledger_list_file=buf, ledger_list_filename="ledgers.xlsx")
    assert "Apex Motors Pvt Ltd" not in tb.inferred and "Cash" in tb.inferred
    assert tb.ledgers.set_index("name").loc["Apex Motors Pvt Ltd", "group"] == "Sundry Debtors"
    m = pipeline(tb)[2]
    assert not any(f.ledger == "Apex Motors Pvt Ltd" for f in m.flags)


def test_opening_stock_line_is_handled_and_flagged():
    m = pipeline(read(native_bytes()))[2]
    note = next(n for n in m.notes if n.title == "Inventories")
    assert note.rows[0][0] == "Stock in hand (as shown in the trial balance)" and note.total[1] == 1_075_000
    assert any("does not include the closing stock" in f.text for f in m.flags)


@pytest.mark.parametrize("name,side,group", [
    ("Jaipur Machine Tools", "Dr", "Sundry Debtors"), ("Salary & Wages - Staff", "Dr", "Indirect Expenses"),
    ("Term Loan - HDFC Bank", "Cr", "Secured Loans"), ("Machinery Repairs", "Dr", "Indirect Expenses"),
    ("Rent - Factory Premises", "Dr", "Indirect Expenses"), ("Factory Wages", "Dr", "Direct Expenses"),
    ("Vehicle Insurance", "Dr", "Indirect Expenses"), ("HDFC Bank OD A/c", "Cr", "Bank OD A/c"),
    ("Plant & Machinery", "Dr", "Fixed Assets"), ("Interest Received - FD", "Cr", "Indirect Incomes"),
    ("Interest on Term Loan", "Dr", "Indirect Expenses"), ("Share Capital", "Cr", "Capital Account")])
def test_group_inference_cases(name, side, group):
    assert infer_group(name, side)[0] == group


def test_unknown_names_fall_back_to_parties_by_side_and_say_so():
    assert infer_group("Zenith Corp", "Dr") == ("Sundry Debtors", "low")
    assert infer_group("Zenith Corp", "Cr") == ("Sundry Creditors", "low")


# -------------------------------------------------------------- other shapes and failures
def test_separate_debit_and_credit_closing_columns_are_understood():
    wb = Workbook()
    ws = wb.active
    for r in [["Acme Ltd"], ["Trial Balance"], ["For 1-Apr-24 to 31-Mar-25"], ["Particulars", "Debit", "Credit"],
              ["Cash", 5000, None], ["Share Capital", None, 5000], ["Grand Total", 5000, 5000]]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    n = parse_tally_trial_balance(buf.getvalue())
    assert n.df.set_index("name").loc["Cash", ["closing_amount", "closing_dr_cr"]].tolist() == [5000, "Dr"]
    assert n.df.set_index("name").loc["Share Capital", "closing_dr_cr"] == "Cr" and len(n.df) == 2


def test_cells_without_a_dr_cr_marker_give_a_clear_error():
    wb = make_native_tb.build()
    wb.active["B9"].number_format = "0.00"           # remove the marker from the Opening Stock opening balance
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError) as e:
        read(buf.getvalue())
    assert "B9" in str(e.value) and "Dr/Cr" in str(e.value) and "Opening Stock" in str(e.value)


def test_reversed_markers_are_caught_by_the_balance_check():
    wb = make_native_tb.build()
    ws = wb.active
    for r in range(9, ws.max_row + 1):                # turn every Dr into Cr: still arithmetic, but now wrong
        c = ws.cell(r, 5)
        if c.value is not None and c.number_format.endswith('Dr"'):
            c.number_format = '""0.00" Cr"'
    buf = io.BytesIO()
    wb.save(buf)
    assert normalise(read(buf.getvalue())).warnings           # trial balance no longer agrees: reported, not hidden


def test_template_layout_files_still_load_after_this_change():
    tb = read(P.sample_xlsx_bytes())
    assert len(tb.ledgers) == 68 and not tb.inferred


def test_a_workbook_that_is_neither_layout_explains_both_options():
    wb = Workbook()
    wb.active.append(["random", "data"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError) as e:
        read(buf.getvalue())
    assert "header row" in str(e.value) and "template" in str(e.value)


# -------------------------------------------------------------- through the page
def test_native_export_uploads_end_to_end_through_the_page(tmp_path, monkeypatch):
    f = tmp_path / "TrialBal.xlsx"
    f.write_bytes(native_bytes())
    at = _run_upload(monkeypatch, f, "TrialBal.xlsx", company="")
    assert not at.error, [e.value for e in at.error]
    m = {x.label: x.value for x in at.metric}
    assert m["Ledgers"] == "67" and m["Tie-out checks passed"].endswith(" of " + m["Tie-out checks passed"].split(" of ")[1])
    assert "failed" not in " ".join(e.value for e in at.error)
    assert any("Rajasthan Precision" in s.value and "01/04/2024 to 31/03/2025" in s.value for s in at.success)
    assert any("worked out from the ledger names" in i.value for i in at.info)


def test_single_day_export_shows_the_instructions_on_the_page(tmp_path, monkeypatch):
    f = tmp_path / "TrialBal.xlsx"
    f.write_bytes(native_bytes(period_text="For 1-Apr-26"))
    at = _run_upload(monkeypatch, f, "TrialBal.xlsx", company="")
    assert any("single day" in e.value and "F2" in e.value for e in at.error)
    assert not at.metric
