import sys
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from daybook_auditor.fs_format import check_file  # noqa: E402
from fsbuilder.format import inr, rnd, rs  # noqa: E402
from fsbuilder.mapping import map_ledgers  # noqa: E402
from fsbuilder.model import build_model  # noqa: E402
from fsbuilder.render import write_all  # noqa: E402
from fsbuilder.tb import load_snapshot, normalise  # noqa: E402

CO = "Rajasthan Precision Industries Private Limited"


def _tb():
    return normalise(load_snapshot(ROOT / "samples" / "rpi_trial_balance.csv", ROOT / "samples" / "rpi_groups.csv",
                                   CO, date(2024, 4, 1), date(2025, 3, 31)))


@pytest.fixture(scope="module")
def model():
    tb = _tb()
    return build_model(tb, map_ledgers(tb))


def test_number_formatting():
    assert inr(1234567) == "12,34,567" and inr(-1234567) == "(12,34,567)" and inr(0) == "-" and inr(None) == "-"
    assert inr(999) == "999" and inr(100000) == "1,00,000"
    assert rnd(2.5) == 3 and rnd(-2.5) == -3 and rnd(0.4999) == 0
    assert rs(2489415.93, 2) == "Rs 24,89,415.93" and rs(-1) == "Rs (1)"


def test_statements_foot_and_agree(model):
    tot = {r.label: r for r in model.bs if r.kind == "total"}
    a, e = tot["TOTAL ASSETS"], tot["TOTAL EQUITY AND LIABILITIES"]
    assert a.cur == e.cur == 15_086_416 and a.prev == e.prev == 6_080_000
    # every figure on the face of the balance sheet foots: items sum to the totals
    items = [r for r in model.bs if r.kind in ("item", "subitem") and r.cur is not None]
    assert sum(r.cur for r in items if r.label.split(") ")[-1] in
               {"Trade receivables", "Cash and bank balances", "Inventories", "Property, plant and equipment"}) == a.cur
    assert model.profit == 2_489_416
    pl = {r.label: r for r in model.pl}
    assert pl["VII. Profit / (loss) for the year (V - VI)"].cur == model.profit == pl["V. Profit / (loss) before tax (III - IV)"].cur


def test_every_noted_line_ties_to_its_note(model):
    notes = {n.no: n for n in model.notes}
    for r in model.bs + model.pl:
        if r.note:
            n = notes[r.note]
            assert n.total[1] == r.cur, r.label
            if r.prev is not None and len(n.total) > 2:
                assert n.total[2] == r.prev, r.label
            assert sum(row[1] for row in n.rows if len(row) > 1 and row[1] is not None and not row[0].startswith(("Opening", "Add"))) == n.total[1], r.label


def test_ppe_movement_rolls_forward(model):
    ppe = next(n for n in model.notes if n.title == "Property, plant and equipment")
    for cls, opening, add, ded, closing in ppe.movement["rows"]:
        assert opening + add + ded == closing, cls


def test_flags_are_referenced_and_zero_balance_ledgers_not_flagged(model):
    refs = [f.ref for f in model.flags]
    assert len(refs) == len(set(refs)) and {"R1", "O1", "O2", "O3"} <= set(refs)
    assert "Audit Fees" not in {f.ledger for f in model.flags}
    on_face = {x for r in model.bs + model.pl for x in r.flags}
    assert {f.ref for f in model.flags if f.kind == "Classification"} <= on_face


def test_written_files_carry_flags_and_pass_our_own_format_checker(model, tmp_path):
    docx, xlsx = write_all(model, tmp_path)
    for f in (docx, xlsx):
        findings, _ = check_file(f.name, f.read_bytes())
        assert findings.empty, findings.to_string()
    text = "\n".join(p.text for p in Document(docx).paragraphs) + "\n".join(
        c.text for t in Document(docx).tables for row in t.rows for c in row.cells)
    for f in model.flags:
        assert f.ref in text
    assert "31/03/2025" in text and "Balance Sheet as at" in text
    wb = load_workbook(xlsx, data_only=True)
    assert wb.sheetnames == ["Balance Sheet", "P&L", "Cash Flow", "Notes", "Disclosures", "Open items", "Mapping"]
    bs = wb["Balance Sheet"]
    assert any(c.value == 0 for row in bs.iter_rows() for c in row if c.number_format.startswith("#,##0")), "check row"
    assert "O2" in {c.value for row in wb["Open items"].iter_rows() for c in row}
    assert wb["Mapping"].max_row == 3 + 68


def test_rounding_difference_is_absorbed_and_disclosed():
    tb = _tb()
    L = tb.ledgers.copy()
    L.loc[L["name"] == "Cash", "closing"] += 0.31       # moves rounded assets by 1 ...
    L.loc[L["name"] == "Apex Motors Pvt Ltd", "closing"] -= 0.31   # ... while the books still balance exactly
    m = build_model(replace(tb, ledgers=L), map_ledgers(replace(tb, ledgers=L)))
    tot = {r.label: r for r in m.bs if r.kind == "total"}
    assert tot["TOTAL ASSETS"].cur == tot["TOTAL EQUITY AND LIABILITIES"].cur
    assert m.rounding and "absorbed" in m.rounding[0] and not any("does not agree" in b for b in m.basis)
