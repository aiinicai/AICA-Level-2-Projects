import re
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from fsbuilder import pipeline as P  # noqa: E402
from fsbuilder.mapping import map_ledgers  # noqa: E402
from fsbuilder.model import build_model  # noqa: E402
from fsbuilder.render import write_all  # noqa: E402
from fsbuilder.tb import normalise  # noqa: E402


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    tb = normalise(P.load_sample())
    m = build_model(tb, map_ledgers(tb))
    docx, xlsx = write_all(m, tmp_path_factory.mktemp("out"))
    return m, docx, xlsx


class Calc:
    """A small independent spreadsheet engine for the formula shapes we write: refs, SUM, + and -."""

    def __init__(self, path, overrides=None):
        self.wb = load_workbook(path)                    # formulas as written
        self.overrides, self.memo = overrides or {}, {}

    def value(self, sheet, ref):
        key = (sheet, ref)
        if key in self.overrides:
            return self.overrides[key]
        if key not in self.memo:
            v = self.wb[sheet][ref].value
            if isinstance(v, str) and v.startswith("="):
                v = self.evaluate(sheet, v[1:])
            elif not isinstance(v, (int, float)) or isinstance(v, bool):
                v = 0.0                                     # text such as '-'
            self.memo[key] = v
        return self.memo[key]

    def _split(self, token, sheet):
        if "!" in token:
            sh, rest = token.rsplit("!", 1)
            return sh.strip("'"), rest
        return sheet, token

    def _cells(self, token, sheet):
        sh, rest = self._split(token, sheet)
        if ":" in rest:
            c1, r1, c2, r2 = range_boundaries(rest)
            return [self.value(sh, f"{self.wb[sh].cell(r, c).coordinate}") for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        return [self.value(sh, rest)]

    def evaluate(self, sheet, expr):
        expr = re.sub(r"SUM\(([^)]*)\)", lambda m: repr(sum(v for a in m.group(1).split(",") for v in self._cells(a.strip(), sheet))), expr)
        expr = re.sub(r"(?:'[^']+'|[A-Za-z_]+)?!?\$?[A-Z]{1,2}\$?\d+",
                      lambda m: repr(self.value(*self._split(m.group(0), sheet))), expr)
        assert re.fullmatch(r"[\d\.\+\-\(\) eE]+", expr), f"unsupported formula: {expr}"
        return eval(expr)                                   # noqa: S307  (digits and + - only, checked above)


def find_row(ws, text, col="A", startswith=False):
    for row in ws.iter_rows(min_col=1, max_col=1):
        v = row[0].value
        if isinstance(v, str) and (v.startswith(text) if startswith else text in v):
            return row[0].row
    raise AssertionError(f"'{text}' not found in {ws.title}")


# ---------------------------------------------------------------- Excel
def test_every_formula_agrees_with_an_independent_recalculation_and_the_stored_value(built):
    _, _, xlsx = built
    calc, cached = Calc(xlsx), load_workbook(xlsx, data_only=True)
    n = 0
    for ws in calc.wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n += 1
                    stored = cached[ws.title][c.coordinate].value
                    assert stored is not None, f"{ws.title}!{c.coordinate} has no stored value"
                    assert calc.value(ws.title, c.coordinate) == pytest.approx(stored, abs=0.001), f"{ws.title}!{c.coordinate} {c.value}"
    assert n >= 70         # the statements are formulas, not typed numbers


def test_every_figure_with_a_note_reads_the_total_of_that_note(built):
    m, _, xlsx = built
    wb = load_workbook(xlsx)
    notes = wb["Notes"]
    for sheet, rows in (("Balance Sheet", m.bs), ("P&L", m.pl)):
        ws = wb[sheet]
        for row in rows:
            if not row.note:
                continue
            r = find_row(ws, row.label, startswith=True)
            formula = ws[f"C{r}"].value
            assert re.fullmatch(r"=Notes!B\d+", formula), (sheet, row.label, formula)
            target = int(formula.split("B")[-1])
            title_row = find_row(notes, f"Note {row.note}:", startswith=True)
            next_title = min([x.row for x in notes["A"] if isinstance(x.value, str) and x.value.startswith("Note ") and x.row > title_row] or [10 ** 6])
            assert title_row < target < next_title, f"{row.label} points outside Note {row.note}"


def test_a_change_in_a_note_flows_into_the_statements(built):
    m, _, xlsx = built
    base = Calc(xlsx)
    wb = base.wb
    bs, notes = wb["Balance Sheet"], wb["Notes"]
    ta, tel = find_row(bs, "TOTAL ASSETS"), find_row(bs, "TOTAL EQUITY AND LIABILITIES")
    assert base.value("Balance Sheet", f"C{ta}") == base.value("Balance Sheet", f"C{tel}")

    # raise one trade-payable balance by 1,000 in the note
    t = find_row(notes, "Note 4: Trade payables", startswith=True) + 2
    changed = Calc(xlsx, overrides={("Notes", f"B{t}"): base.value("Notes", f"B{t}") + 1000})
    line = find_row(bs, "Trade payables")
    assert changed.value("Balance Sheet", f"C{line}") == base.value("Balance Sheet", f"C{line}") + 1000
    assert changed.value("Balance Sheet", f"C{tel}") == base.value("Balance Sheet", f"C{tel}") + 1000
    check = find_row(bs, "Check: total assets less", startswith=True)
    assert changed.value("Balance Sheet", f"C{check}") == -1000            # the check row now shouts

    # raise an expense by 500 in its note: profit, the cash flow's starting profit and the reserves all follow
    pl, cf = wb["P&L"], wb["Cash Flow"]
    e = find_row(notes, "Note 16: Other expenses", startswith=True) + 2
    changed = Calc(xlsx, overrides={("Notes", f"B{e}"): base.value("Notes", f"B{e}") + 500})
    profit = find_row(pl, "VII. Profit", startswith=True)
    assert changed.value("P&L", f"C{profit}") == base.value("P&L", f"C{profit}") - 500
    pbt = find_row(cf, "Net profit / (loss) before tax", startswith=True)
    assert changed.value("Cash Flow", f"C{pbt}") == base.value("Cash Flow", f"C{pbt}") - 500
    close = find_row(cf, "Cash and cash equivalents at the end", startswith=True)
    assert changed.value("Cash Flow", f"C{close}") == base.value("Cash Flow", f"C{close}") - 500


def test_note_numbers_are_clickable_and_lead_to_the_right_note_and_back(built):
    m, _, xlsx = built
    wb = load_workbook(xlsx)
    seen = 0
    for sheet in ("Balance Sheet", "P&L"):
        ws = wb[sheet]
        for r in range(6, ws.max_row + 1):
            c = ws[f"B{r}"]
            if c.hyperlink is not None and c.value:
                seen += 1
                sh, ref = c.hyperlink.location.split("!")
                assert sh == "'Notes'" and wb["Notes"][ref].value.startswith(f"Note {c.value}:")
    linkable = [n for n in m.notes if n.kind != "text"]
    assert seen == len(linkable)
    back = [c for row in wb["Notes"].iter_rows(min_col=6, max_col=6) for c in row if c.hyperlink is not None]
    assert len(back) == len(linkable)
    for c in back:
        sh, ref = c.hyperlink.location.split("!")
        label = wb[sh.strip("'")][ref].value
        title = wb["Notes"][f"A{c.row}"].value.split(": ", 1)[1]
        assert title.lower() in label.lower(), (title, label)


def test_numbers_show_even_in_a_viewer_that_does_not_calculate(built):
    m, _, xlsx = built
    cached = load_workbook(xlsx, data_only=True)["Balance Sheet"]
    ta = find_row(cached, "TOTAL ASSETS")
    assert cached[f"C{ta}"].value == 15_086_416 and cached[f"D{ta}"].value == 6_080_000


def test_the_workbook_recalculates_when_opened(built):
    assert load_workbook(built[2]).calculation.fullCalcOnLoad is True


def test_cash_flow_closing_cash_is_checked_against_the_components(built):
    _, _, xlsx = built
    calc = Calc(xlsx)
    cf = calc.wb["Cash Flow"]
    assert calc.value("Cash Flow", f"C{find_row(cf, 'Check: closing cash', startswith=True)}") == 0


# ---------------------------------------------------------------- Word
def test_word_note_references_and_notes_link_to_each_other(built):
    m, docx, _ = built
    xml = zipfile.ZipFile(docx).read("word/document.xml").decode()
    anchors = re.findall(r'<w:hyperlink [^>]*w:anchor="([^"]+)"', xml)
    marks = re.findall(r'<w:bookmarkStart [^>]*w:name="([^"]+)"', xml)
    assert len(marks) == len(set(marks)) and set(anchors) <= set(marks)
    for n in (n for n in m.notes if n.kind != "text"):
        assert f"note_{n.no}" in anchors and f"note_{n.no}" in marks      # statement -> note
        assert f"face_{n.no}" in anchors and f"face_{n.no}" in marks      # note -> statement


def test_word_note_numbers_are_still_visible_text(built):
    from docx import Document
    m, docx, _ = built
    cells = [c.text for t in Document(docx).tables for row in t.rows for c in row.cells]
    assert all(str(n.no) in cells for n in m.notes if n.kind != "text")
