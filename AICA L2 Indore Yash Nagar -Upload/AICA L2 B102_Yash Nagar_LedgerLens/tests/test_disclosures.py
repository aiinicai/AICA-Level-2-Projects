import sys
from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from daybook_auditor import narration as N  # noqa: E402
from daybook_auditor.catalog import BASE, catalog_frame  # noqa: E402
from daybook_auditor.rules import FINDING_COLS  # noqa: E402
from fsbuilder.mapping import map_ledgers  # noqa: E402
from fsbuilder.model import build_model  # noqa: E402
from fsbuilder.render import write_all  # noqa: E402
from fsbuilder.tb import load_snapshot, normalise  # noqa: E402


@pytest.fixture(scope="module")
def parts():
    tb = normalise(load_snapshot(ROOT / "samples" / "rpi_trial_balance.csv", ROOT / "samples" / "rpi_groups.csv",
                                 "RPI", date(2024, 4, 1), date(2025, 3, 31)))
    return tb, map_ledgers(tb)


@pytest.fixture(scope="module")
def model(parts):
    return build_model(*parts)


def section(m, no):
    return next(d for d in m.disclosures if d.no == no)


def test_every_statutory_section_is_present(model):
    assert [d.no for d in model.disclosures] == [f"D{i}" for i in range(1, 12)]
    titles = " ".join(d.title for d in model.disclosures)
    for word in ("auditors", "Managerial remuneration", "Related-party", "Earnings per share", "social responsibility", "Contingent",
                 "micro and small", "receivables ageing", "payables ageing", "Additional regulatory information"):
        assert word in titles


def test_rows_fit_their_headers_and_widths(model):
    for d in model.disclosures:
        assert len(d.widths) == len(d.header), d.no
        assert all(len(r) == len(d.header) for r in d.rows), d.no


def test_facts_come_from_the_books(model, parts):
    tb, mapped = parts
    d1 = section(model, "D1")
    assert "40,000" in d1.comment and "no audit fee is charged" in d1.comment        # fee payable, none charged
    d2 = {r[0]: r[1] for r in section(model, "D2").rows}
    assert d2["Remuneration to directors booked in the year"] == "90,190"
    assert d2["Booked remuneration within the ceiling?"] == "Yes"
    names = [r[0] for r in section(model, "D3").rows]
    assert "Unsecured Loan - Director" in names and "Directors Remuneration" in names
    d11 = section(model, "D11").rows[0]
    assert "Unsecured Loan - Director" in d11[2] and "declaration" in d11[3]
    d8 = section(model, "D8").rows
    assert d8[5][-1] == "86,12,376" and d8[-1][-1] == "6,95,000"                      # both years agree to trade receivables
    d9 = section(model, "D9").rows
    assert d9[5][-1] == "65,01,803" and d9[-1][-1] == "5,95,000"


def test_csr_thresholds_are_computed_not_assumed(model):
    d5 = section(model, "D5")
    assert [r[2] for r in d5.rows[:3]] == ["No", "No", "No"] and "does not appear to apply" in d5.comment


def test_csr_applies_when_profit_is_large(parts):
    tb, mapped = parts
    m = build_model(tb, mapped)
    m.figures["pl"]["pbt"] = 6e7                    # Rs 6 crore
    from fsbuilder.disclosures import disclosure_sections
    d5 = next(d for d in disclosure_sections(m) if d.no == "D5")
    assert d5.rows[2][2] == "Yes" and "constitute the CSR committee" in d5.comment


def test_nothing_is_declared_nil_on_management_s_behalf(model):
    d10 = section(model, "D10")
    assert len(d10.rows) >= 15
    assert all(r[3].startswith(("[Management", "Explain")) for r in d10.rows)
    text = " ".join(" ".join(r) for d in model.disclosures for r in d.rows).lower()
    assert "nil" not in text.replace("[management", "")


def test_a_ledger_that_matters_changes_the_disclosure(parts):
    tb, mapped = parts
    m = build_model(tb, mapped)
    m.mapping = pd.concat([m.mapping, m.mapping.iloc[[0]].assign(ledger="Bitcoin wallet", line="CA.OCA", closing=5000.0)], ignore_index=True)
    from fsbuilder.disclosures import disclosure_sections
    d10 = next(d for d in disclosure_sections(m) if d.no == "D10")
    crypto = next(r for r in d10.rows if r[0].startswith("Trade or investment in crypto"))
    assert "Bitcoin wallet" in crypto[2]


def test_disclosures_reach_word_and_excel(model, tmp_path):
    docx, xlsx = write_all(model, tmp_path)
    doc = Document(docx)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Additional regulatory information and other statutory disclosures" in text and "D10. Additional regulatory information" in text
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "Unsecured Loan - Director" in cells and "[Management to confirm]" in cells
    wb = load_workbook(xlsx)
    ws = wb["Disclosures"]
    body = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "D3. Related-party disclosures" in body and "86,12,376" in body


def test_catalog_lists_every_test_even_when_nothing_is_found():
    empty = pd.DataFrame(columns=FINDING_COLS)
    cat = catalog_frame(empty)
    assert len(cat) == len(BASE) + len(N.THEMES)
    assert (cat["Findings"] == 0).all() and (cat["Result"] == "Ran: nothing found").all()
    assert cat.iloc[0]["Test"] == "999 verification"


def test_catalog_counts_the_999_test():
    rec = {"voucher_id": "V1", "vch_type": "Payment", "total": 29_999.0, "narration": "", "vch_no": "1", "date": pd.Timestamp(2025, 5, 1)}
    f = pd.DataFrame(N.threshold_findings([rec], type("C", (), {"near_threshold_pct": 2.0})(),
                                          lambda rule, sev, r, detail, ledger, amount, ref: dict(
                                              zip(FINDING_COLS, [sev, rule, r["date"], r["vch_type"], r["vch_no"], r["voucher_id"], ledger,
                                                                 r["total"], detail, ref, "Rule"])))[:1], columns=FINDING_COLS)
    row = catalog_frame(f).iloc[0]
    assert row["Test"] == "999 verification" and row["Findings"] == 1 and row["Vouchers"] == 1


def test_every_amount_section_has_a_previous_year(model):
    for no in ("D1", "D2", "D3", "D4", "D5", "D6", "D7"):
        header = " ".join(section(model, no).header)
        assert "Previous year" in header, no
    assert [r[-1] for r in section(model, "D7").rows[:1]] == ["5,95,000"]               # opening payables
    d3 = {r[0]: r for r in section(model, "D3").rows}
    assert d3["Unsecured Loan - Director"][-2:] == ["3,00,000", "3,00,000"]
    assert d3["Directors Remuneration"][-1] == "Not loaded"                              # a year's expense has no opening figure
    d5 = section(model, "D5").rows[0]
    assert d5[3] == "37,75,000" and d5[4] == "No"                                       # net worth at the start of the year


def test_signature_blocks_follow_each_statement_and_the_disclosures(model, tmp_path):
    import dataclasses
    from datetime import date as D
    from fsbuilder.model import Signatories
    m = dataclasses.replace(model, sign=Signatories(firm="Mehta and Co Chartered Accountants", frn="012345C", partner="R. Mehta", membership="401234",
                                                    place="Jaipur", sign_date=D(2025, 5, 20), directors=[("A. Sharma", "01234567"), ("B. Verma", "07654321")],
                                                    board_date=D(2025, 5, 20)))
    docx, xlsx = write_all(m, tmp_path)
    doc = Document(docx)
    sig = [t for t in doc.tables if "Firm Registration No." in t.rows[0].cells[0].text]
    assert len(sig) == 4                                                                 # balance sheet, P&L, cash flow, disclosures
    left, right = sig[0].rows[0].cells[0].text, sig[0].rows[0].cells[1].text
    for want in ("For Mehta and Co Chartered Accountants", "012345C", "R. Mehta", "401234", "Date of signing: 20/05/2025"):
        assert want in left
    for want in ("RPI", "A. Sharma", "01234567", "B. Verma", "07654321", "Date of signing: 20/05/2025"):
        assert want in right
    wb = load_workbook(xlsx)
    for sheet in ("Balance Sheet", "P&L", "Cash Flow", "Disclosures"):
        cells = [str(c.value) for row in wb[sheet].iter_rows() for c in row if c.value]
        assert any("Firm Registration No.: 012345C" == c for c in cells), sheet
        assert any(c == "DIN: 07654321" for c in cells), sheet


def test_unfilled_signatories_print_blanks_not_invented_names(model, tmp_path):
    docx, _ = write_all(model, tmp_path)
    text = " ".join(c.text for t in Document(docx).tables for row in t.rows for c in row.cells)
    assert "Firm Registration No.: ______________" in text and "DIN: ______________" in text and "For XYZ Chartered Accountants" in text


# ------------------------------------------------------------ AI remarks travel in the Excel files
def _reviewed(parts, llm):
    import dataclasses
    from fsbuilder.reviewer import run_review
    m = build_model(*parts)
    return dataclasses.replace(m, review=run_review(m, parts[1], llm))


def test_statements_excel_carries_the_ai_analysis(parts, tmp_path):
    from test_reviewer import FakeLLM
    m = _reviewed(parts, FakeLLM())
    _, xlsx = write_all(m, tmp_path)
    ws = load_workbook(xlsx)["AI analysis"]
    body = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "AI review run" in body and "Current ratio" in body
    assert "AI draft" in body                                     # the source of each drafted reason is stated
    assert any(c.value == "Yes" for row in ws.iter_rows() for c in row)
    assert "Analytical observations" in body and len(m.review.ai.observations) >= 1
    for o in m.review.ai.observations:
        assert o.title in body


def test_statements_excel_says_so_when_the_ai_was_not_run(parts, tmp_path):
    m = _reviewed(parts, None)
    _, xlsx = write_all(m, tmp_path)
    ws = load_workbook(xlsx)["AI analysis"]
    body = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "AI review not run" in body and "Not run." in body


def test_files_without_a_review_have_no_ai_sheet(model, tmp_path):
    _, xlsx = write_all(model, tmp_path)
    assert "AI analysis" not in load_workbook(xlsx).sheetnames


def test_build_files_passes_review_and_signatures_through(parts):
    import io
    from fsbuilder import pipeline as P
    from fsbuilder.model import Signatories
    from fsbuilder.reviewer import run_review
    m = build_model(*parts)
    files = P.build_files(m, run_review(m, parts[1], None), Signatories(frn="FRN-77", directors=[("A", "111"), ("B", "222")]))
    wb = load_workbook(io.BytesIO(files["Statements (Excel)"][1]))
    assert "AI analysis" in wb.sheetnames
    cells = [str(c.value) for row in wb["Balance Sheet"].iter_rows() for c in row if c.value]
    assert "Firm Registration No.: FRN-77" in cells and "DIN: 222" in cells


def test_day_book_excel_has_an_ai_sheet_with_the_rule_context():
    import io
    from daybook_auditor.report import ai_sheet, build_excel
    cols = FINDING_COLS
    mk = lambda sev, rule, vid, src, amt=1000.0: dict(zip(cols, [sev, rule, pd.Timestamp(2025, 5, 1), "Payment", "7", vid, "", amt, "detail " + vid, "ref", src]))  # noqa: E731
    f = pd.DataFrame([mk("High", "Cash payment above limit", "V1", "Rule"), mk("High", "AI: Personal expense", "V1", "AI"),
                      mk("Medium", "AI: Wrong ledger", "V2", "AI")], columns=cols)
    a = ai_sheet(f, 40)
    assert list(a["Also flagged by rules"]) == ["Yes", "No: raised only by the AI"]
    assert a.iloc[0]["Rule findings on the same voucher"] == "Cash payment above limit"
    assert "not run" in ai_sheet(f.iloc[:1], 0).iloc[0, 0]
    assert "found nothing" in ai_sheet(f.iloc[:1], 25).iloc[0, 0]
    wb = load_workbook(io.BytesIO(build_excel(f, pd.DataFrame({"a": [1]}), {}, 40)))
    assert "AI analysis" in wb.sheetnames
