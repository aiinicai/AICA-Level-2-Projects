import io
import sys
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from daybook_auditor.parser import load_daybook  # noqa: E402
from daybook_auditor.rules import Config, run_rules  # noqa: E402

TITLE = [["Rajasthan Precision Industries Private Limited"], ["Plot No. 45, Sitapura Industrial Area, Jaipur"], ["Day Book"],
         ["1-Apr-24 to 31-Mar-25"], []]
HDR = ["Date", "Particulars", "Vch Type", "Vch No.", "Debit Amount", "Credit Amount"]
AMT_HDR = ["Date", "Particulars", "Vch Type", "Vch No.", "Amount"]
CR, DR = '""0.00" Cr"', '""0.00" Dr"'


def book(rows, fmts=None, merges=()):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    for rng in merges:
        ws.merge_cells(rng)
    for ref, f in (fmts or {}).items():
        ws[ref].number_format = f
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def load(data, name="export.xlsx"):
    return load_daybook(name, data)


def summary(df):
    return [(r.vch_type, r.vch_no, r.date.date(), r.ledger, r.side, round(r.amount, 2)) for r in df.itertuples()]


# ---------------------------------------------------------------- layouts Tally produces
def test_title_block_merged_cells_and_grand_total_row_are_ignored():
    data = book(TITLE + [HDR,
                         ["1-Apr-24", "Sharma Steel Traders", "Purchase", "17", None, 134816.16],
                         ["2-Apr-24", "Metro Auto Components Ltd", "Sales", "5", 357623.57, None],
                         ["", "Grand Total", "", "", 357623.57, 134816.16]], merges=["A1:C1", "A3:C3"])
    df, meta = load(data)
    assert meta["vouchers"] == 2 and not meta["ledger_level"]
    assert summary(df) == [("Purchase", "17", date(2024, 4, 1), "Sharma Steel Traders", "Cr", 134816.16),
                           ("Sales", "5", date(2024, 4, 2), "Metro Auto Components Ltd", "Dr", 357623.57)]


def test_dr_cr_carried_in_the_number_format_of_a_single_amount_column():
    data = book(TITLE + [AMT_HDR,
                         ["1-Apr-24", "Sharma Steel Traders", "Purchase", "17", 134816.16],
                         ["2-Apr-24", "HDFC Bank - Current A/c", "Receipt", "9", 138198.54],
                         ["3-Apr-24", "Cash", "Contra", "5", 14213.93]],
                fmts={"E7": CR, "E8": CR, "E9": DR})
    df, _ = load(data)
    assert [(r.ledger, r.side) for r in df.itertuples()] == [("Sharma Steel Traders", "Cr"), ("HDFC Bank - Current A/c", "Cr"), ("Cash", "Dr")]


def test_dr_cr_as_a_text_suffix_is_understood_too():
    data = book(TITLE + [AMT_HDR, ["1-Apr-24", "Sharma Steel Traders", "Purchase", "17", "1,34,816.16 Cr"],
                         ["2-Apr-24", "Cash", "Contra", "5", "14213.93 Dr"]])
    df, _ = load(data)
    assert summary(df)[0][-2:] == ("Cr", 134816.16) and summary(df)[1][-2:] == ("Dr", 14213.93)


def test_an_amount_without_any_dr_cr_marker_is_refused_not_guessed():
    with pytest.raises(ValueError) as e:
        load(book(TITLE + [AMT_HDR, ["1-Apr-24", "X", "Purchase", "1", 100.0]]))
    assert "Row 7" in str(e.value) and "Dr/Cr marker" in str(e.value)


def test_date_shown_only_once_per_day_does_not_merge_vouchers():
    data = book(TITLE + [HDR,
                         ["1-Apr-24", "Sharma Steel Traders", "Purchase", 17.0, None, 134816.16],
                         [None, "Metro Auto Components Ltd", "Sales", 5.0, 357623.57, None],
                         [None, "HDFC Bank - Current A/c", "Receipt", 9.0, None, 138198.54],
                         ["2-Apr-24", "Cash", "Contra", 5.0, 14213.93, None]])
    df, meta = load(data)
    assert meta["vouchers"] == 4
    assert [(s[0], s[1], s[2]) for s in summary(df)] == [("Purchase", "17", date(2024, 4, 1)), ("Sales", "5", date(2024, 4, 1)),
                                                         ("Receipt", "9", date(2024, 4, 1)), ("Contra", "5", date(2024, 4, 2))]


def test_real_date_cells_and_numeric_voucher_numbers():
    data = book(TITLE + [HDR, [datetime(2024, 4, 1), "Sharma Steel Traders", "Purchase", 17, None, 134816.16]])
    df, _ = load(data)
    assert summary(df)[0][:3] == ("Purchase", "17", date(2024, 4, 1))


def test_detailed_voucher_lines_with_narration_balance_and_keep_the_narration():
    data = book(TITLE + [HDR,
                         ["1-Apr-24", "Sharma Steel Traders", "Purchase", "17", None, 150000],
                         [None, "Purchase - Raw Material", None, None, 127000, None],
                         [None, "CGST Input", None, None, 11500, None],
                         [None, "SGST Input", None, None, 11500, None],
                         [None, "Sharma Steel Traders", None, None, None, 150000],
                         [None, "Being purchase of raw material", None, None, None, None],
                         ["2-Apr-24", "Cash", "Contra", "5", 5000, None],
                         [None, "Cash", None, None, 5000, None], [None, "HDFC Bank - Current A/c", None, None, None, 5000]])
    df, meta = load(data)
    assert meta["ledger_level"] and meta["vouchers"] == 2
    v1 = df[df["voucher_no"] == "17"] if "voucher_no" in df else df[df["vch_no"] == "17"]
    assert v1["amount"].where(v1["side"] == "Dr").sum() == v1["amount"].where(v1["side"] == "Cr").sum() == 150000
    assert v1["narration"].iloc[0] == "Being purchase of raw material"


def test_the_workbook_may_have_several_sheets_and_the_day_book_is_found():
    wb = Workbook()
    wb.active.title = "Cover"
    wb.active.append(["Day Book for review"])
    ws = wb.create_sheet("Day Book")
    for r in TITLE + [HDR, ["1-Apr-24", "Cash", "Contra", "5", 5000, None]]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    assert load(buf.getvalue())[1]["vouchers"] == 1


def test_csv_export_still_works():
    csv = "Date,Particulars,Vch Type,Vch No.,Debit Amount,Credit Amount\n1-Apr-24,Cash,Contra,5,5000,\n"
    df, meta = load_daybook("daybook.csv", csv.encode())
    assert meta["vouchers"] == 1 and summary(df)[0][-2:] == ("Dr", 5000.0)


def test_a_file_that_is_not_a_day_book_says_what_is_expected():
    wb = Workbook()
    wb.active.append(["Particulars", "Opening Balance", "Closing Balance"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError) as e:
        load(buf.getvalue())
    assert "Date" in str(e.value) and "Particulars" in str(e.value)


def test_a_corrupt_workbook_fails_politely():
    with pytest.raises(ValueError, match="Could not open"):
        load(b"not a workbook")


def test_voucher_sequence_check_works_on_condensed_numbers_that_arrived_as_floats():
    rows = [["1-Apr-24", f"P{n}", "Payment", float(n), 6000.0, None] for n in (1, 2, 3, 5, 6)]    # 4 is missing
    df, meta = load(book(TITLE + [HDR] + rows))
    f = run_rules(df, Config(date(2024, 4, 1), date(2025, 3, 31)), ledger_level=meta["ledger_level"])
    gap = f[f["rule"] == "Voucher number gap"]
    assert len(gap) == 1 and "4" in gap.iloc[0]["detail"]
