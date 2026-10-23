"""Presentation checks on draft financial statements in Word (.docx):
  1. font family and font size consistency
  2. date format consistency (dd/mm/yyyy)
  3. negative numbers shown in brackets, not with a minus sign

Deterministic: no AI. Every finding names where it is and what was expected.
"""
from __future__ import annotations

import calendar
import io
import re
from collections import Counter, defaultdict
from datetime import date, datetime

import pandas as pd
from docx import Document
from docx.oxml.ns import qn

COLUMNS = ["severity", "check", "location", "found", "expected", "context"]

# ---- reading the document ------------------------------------------------------------


def _default_font(doc):
    rpr = doc.styles.element.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
    name, size = None, None
    if rpr is not None:
        f = rpr.find(qn("w:rFonts"))
        if f is not None:
            name = f.get(qn("w:ascii")) or (f"theme:{f.get(qn('w:asciiTheme'))}" if f.get(qn("w:asciiTheme")) else None)
        sz = rpr.find(qn("w:sz"))
        if sz is not None:
            size = int(sz.get(qn("w:val"))) / 2
    return name or "Default", size or 10.0


def _effective(run, para, defaults):
    name, size = run.font.name, run.font.size.pt if run.font.size else None
    style = para.style
    while style is not None and (name is None or size is None):
        name = name or style.font.name
        size = size or (style.font.size.pt if style.font.size else None)
        style = style.base_style
    return name or defaults[0], size or defaults[1]


def _is_heading(para) -> bool:
    """Heading style, or a paragraph whose text is entirely bold (titles are often formatted by hand)."""
    n = para.style.name if para.style is not None else ""
    runs = [r for r in para.runs if r.text.strip()]
    return n.startswith(("Heading", "Title")) or (bool(runs) and all(r.bold for r in runs))


def _units(doc):
    """Yield (location, kind, paragraph) for every body paragraph and table cell paragraph."""
    for i, p in enumerate(doc.paragraphs, start=1):
        if p.text.strip():
            yield f"Paragraph {i}", "heading" if _is_heading(p) else "body", p
    for t, table in enumerate(doc.tables, start=1):
        for r, row in enumerate(table.rows, start=1):
            seen = set()
            for c, cell in enumerate(row.cells, start=1):
                if id(cell._tc) in seen:  # merged cells repeat
                    continue
                seen.add(id(cell._tc))
                for p in cell.paragraphs:
                    if p.text.strip():
                        yield f"Table {t}, row {r}, col {c}", "table", p


def _snip(text: str, n: int = 70) -> str:
    text = " ".join(text.split())
    return text if len(text) <= n else text[: n - 1] + "…"


# ---- 1. fonts ------------------------------------------------------------------------


def check_fonts(doc) -> tuple[list[dict], dict]:
    defaults = _default_font(doc)
    runs = []  # (location, kind, font, size, chars, text)
    for loc, kind, p in _units(doc):
        for r in p.runs:
            if r.text.strip():
                f, s = _effective(r, p, defaults)
                runs.append((loc, kind, f, s, len(r.text.strip()), p.text))

    fam = Counter()
    for _, _, f, _, n, _ in runs:
        fam[f] += n
    size_by_kind: dict[str, Counter] = defaultdict(Counter)
    for _, kind, _, s, n, _ in runs:
        size_by_kind[kind][s] += n

    out = []
    if fam:
        main = fam.most_common(1)[0][0]
        seen = set()
        for loc, _, f, _, _, text in runs:
            if f != main and (loc, f) not in seen:
                seen.add((loc, f))
                out.append(dict(severity="Medium", check="Font type", location=loc, found=f,
                                expected=main, context=_snip(text)))
    for kind, counter in size_by_kind.items():
        if kind == "heading":  # heading levels legitimately differ in size
            continue
        main = counter.most_common(1)[0][0]
        seen = set()
        for loc, k, _, s, _, text in runs:
            if k == kind and s != main and (loc, s) not in seen:
                seen.add((loc, s))
                out.append(dict(severity="Medium", check="Font size", location=loc, found=f"{s:g} pt",
                                expected=f"{main:g} pt ({kind} text)", context=_snip(text)))
    summary = {"Dominant font": fam.most_common(1)[0][0] if fam else "-",
               "Fonts in use": ", ".join(f"{f} ({n})" for f, n in fam.most_common()),
               **{f"Size, {k}": ", ".join(f"{s:g} pt ({n})" for s, n in c.most_common()) for k, c in size_by_kind.items()}}
    return out, summary


# ---- 2. dates ------------------------------------------------------------------------

_MON = "Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|June?|July?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?"
_DATE_RE = re.compile(
    rf"""(?<![\d/.\-])(?:
        (?P<iso>\d{{4}}-\d{{2}}-\d{{2}})
      | (?P<num>\d{{1,2}}[/\-.]\d{{1,2}}[/\-.]\d{{2,4}})
      | (?P<dmy>\d{{1,2}}(?:st|nd|rd|th)?[\s\-]+(?:{_MON})\.?,?[\s\-]+\d{{2,4}})
      | (?P<mdy>(?:{_MON})\.?\s+\d{{1,2}}(?:st|nd|rd|th)?,?\s+\d{{4}})
    )(?![\d/\-])""", re.X | re.I)
_MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_abbr) if m}
# "GN 8.9.13", "para 6.15", "Note 6.2.1": a clause number, not a date, even though it looks like d.m.yy.
_REFERENCE_BEFORE = re.compile(r"(?:\bGN|\bpara(?:s|graphs?)?|\bclauses?|\bNote|\bAnnexure|\bSec(?:tion)?|\bAS|\bRule)\s*[A-Za-z]?$", re.I)


def _is_reference(text: str, m: re.Match) -> bool:
    return bool(_REFERENCE_BEFORE.search(text[max(0, m.start() - 14):m.start()]))


def _classify(m: re.Match) -> tuple[str, str] | None:
    """Return (severity, description of what is wrong) or None if the date is dd/mm/yyyy."""
    if m.group("iso"):
        return "Medium", "yyyy-mm-dd"
    if m.group("dmy"):
        return "Medium", "day-month-name style (e.g. 31 March 2026 / 31-Mar-26)"
    if m.group("mdy"):
        return "Medium", "month-first style (e.g. March 31, 2026)"
    s = m.group("num")
    sep = re.search(r"[/\-.]", s).group()
    a, b, y = re.split(r"[/\-.]", s)
    if sep == "/" and len(a) == 2 and len(b) == 2 and len(y) == 4:
        if int(b) > 12 and int(a) <= 12:
            return "High", "looks like mm/dd/yyyy (month and day swapped)"
        if int(a) > 31 or int(b) > 12 or int(a) < 1 or int(b) < 1:
            return "High", "not a valid calendar date"
        try:
            date(int(y), int(b), int(a))
        except ValueError:
            return "High", "not a valid calendar date"
        return None
    why = []
    if sep != "/":
        why.append(f"'{sep}' separator")
    if len(a) < 2 or len(b) < 2:
        why.append("day/month not 2 digits")
    if len(y) != 4:
        why.append("2-digit year")
    return "Medium", ", ".join(why)


def check_dates(doc) -> tuple[list[dict], dict]:
    out, ok = [], 0
    for loc, _, p in _units(doc):
        for m in _DATE_RE.finditer(p.text):
            if _is_reference(p.text, m):
                continue
            res = _classify(m)
            if res is None:
                ok += 1
            else:
                out.append(dict(severity=res[0], check="Date format", location=loc, found=f"{m.group(0)}  ({res[1]})",
                                expected="dd/mm/yyyy", context=_snip(p.text)))
    return out, {"Dates in dd/mm/yyyy": ok, "Dates in another format": len(out)}


# ---- 3. negative numbers -------------------------------------------------------------

_MINUS_NUM = re.compile(r"(?<![\w/.\-–−])[-−–]\d[\d,]*(?:\.\d+)?(?![\w/%]|\s?%)")
_BRACKET_NUM = re.compile(r"\(\s?(?:\d{1,3}(?:,\d{2,3})+|\d{4,})(?:\.\d+)?\s?\)")


def check_negatives(doc) -> tuple[list[dict], dict]:
    out, brackets = [], 0
    for loc, _, p in _units(doc):
        brackets += len(_BRACKET_NUM.findall(p.text))
        for m in _MINUS_NUM.finditer(p.text):
            out.append(dict(severity="Medium", check="Negative number", location=loc, found=m.group(0).strip(),
                            expected="(" + re.sub(r"^[-−–]", "", m.group(0)).strip() + ")", context=_snip(p.text)))
    return out, {"Negatives in brackets": brackets, "Negatives with a minus sign": len(out)}


# ---- Excel (.xlsx) -------------------------------------------------------------------

_DATE_FMT_OK = re.compile(r"dd/mm/yyyy", re.I)
_DATE_FMT_SWAPPED = re.compile(r"^(mm|m)/(dd|d)/(yyyy|yy)$", re.I)
_MAX_PER_CHECK = 200


def _clean_fmt(fmt: str) -> str:
    fmt = re.sub(r"\[\$-[^\]]*\]|\[[^\]]*\]|\"[^\"]*\"|\.", "", fmt or "")  # locale tags, colours, literals
    return fmt.strip().lower()


def _negatives_bracketed(fmt: str) -> bool:
    """True if the number format shows negatives in brackets (second section contains '(')."""
    if "%" in fmt:
        return True  # percentages are out of scope
    sections = (fmt or "").split(";")
    return len(sections) >= 2 and "(" in sections[1]


def _xlsx_cells(wb_vals):
    for ws in wb_vals.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and type(c).__name__ != "MergedCell":
                    yield ws, c


def check_xlsx(data: bytes) -> tuple[pd.DataFrame, dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)  # cached results of formulas
    cells = list(_xlsx_cells(wb))
    loc = lambda ws, c: f"{ws.title}!{c.coordinate}"  # noqa: E731
    out: list[dict] = []

    # 1. fonts: dominant family overall, dominant size among non-bold cells
    fam, size = Counter(), Counter()
    for _, c in cells:
        fam[c.font.name or "Default"] += 1
        if not c.font.b:
            size[c.font.sz or 11.0] += 1
    main_f = fam.most_common(1)[0][0] if fam else None
    main_s = size.most_common(1)[0][0] if size else None
    for ws, c in cells:
        name = c.font.name or "Default"
        if name != main_f:
            out.append(dict(severity="Medium", check="Font type", location=loc(ws, c), found=name,
                            expected=main_f, context=_snip(str(c.value))))
        if not c.font.b and (c.font.sz or 11.0) != main_s:
            out.append(dict(severity="Medium", check="Font size", location=loc(ws, c), found=f"{c.font.sz:g} pt",
                            expected=f"{main_s:g} pt", context=_snip(str(c.value))))

    # 2 & 3. dates and negatives, by cell type
    ok_dates = brackets = minus = 0
    loose_formats: dict[tuple, list] = defaultdict(list)
    for ws, c in cells:
        v, fmt = c.value, c.number_format
        if isinstance(v, (datetime, date)):
            f = _clean_fmt(fmt)
            shown = v.strftime("%d/%m/%Y")
            if _DATE_FMT_OK.fullmatch(f):
                ok_dates += 1
            elif _DATE_FMT_SWAPPED.match(f):
                out.append(dict(severity="High", check="Date format", location=loc(ws, c),
                                found=f"{shown} displayed as '{fmt}' (month-first)", expected="dd/mm/yyyy",
                                context=_snip(str(c.value))))
            else:
                out.append(dict(severity="Medium", check="Date format", location=loc(ws, c),
                                found=f"{shown} displayed as '{fmt}'", expected="dd/mm/yyyy",
                                context=_snip(str(c.value))))
        elif isinstance(v, (int, float)) and not isinstance(v, bool):
            if v < 0:
                if _negatives_bracketed(fmt):
                    brackets += 1
                else:
                    minus += 1
                    out.append(dict(severity="Medium", check="Negative number", location=loc(ws, c),
                                    found=f"{v:,.0f} displayed with a minus (format '{fmt}')",
                                    expected="brackets, e.g. #,##0_);(#,##0)", context=""))
            elif fmt != "General" and not _negatives_bracketed(fmt):
                loose_formats[(ws.title, fmt)].append(c.coordinate)
        elif isinstance(v, str):
            for m in _DATE_RE.finditer(v):
                if _is_reference(v, m):
                    continue
                res = _classify(m)
                if res is None:
                    ok_dates += 1
                else:
                    out.append(dict(severity=res[0], check="Date format", location=loc(ws, c),
                                    found=f"{m.group(0)}  ({res[1]})", expected="dd/mm/yyyy", context=_snip(v)))
            brackets += len(_BRACKET_NUM.findall(v))
            for m in _MINUS_NUM.finditer(v):
                minus += 1
                out.append(dict(severity="Medium", check="Negative number", location=loc(ws, c), found=m.group(0).strip(),
                                expected="(" + re.sub(r"^[-−–]", "", m.group(0)).strip() + ")", context=_snip(v)))
    for (sheet, fmt), coords in loose_formats.items():
        out.append(dict(severity="Low", check="Negative number format", location=f"{sheet}!{coords[0]}" + (f" (+{len(coords) - 1} more)" if len(coords) > 1 else ""),
                        found=f"format '{fmt}' has no bracket style for negatives", expected="#,##0_);(#,##0)",
                        context="Positive today, but a negative here would show a minus sign"))

    dropped = 0
    capped = []
    per = Counter()
    for f in out:
        per[f["check"]] += 1
        if per[f["check"]] <= _MAX_PER_CHECK:
            capped.append(f)
        else:
            dropped += 1
    summary = {"Dominant font": main_f or "-", "Fonts in use": ", ".join(f"{k} ({n})" for k, n in fam.most_common()),
               "Sizes, non-bold cells": ", ".join(f"{k:g} pt ({n})" for k, n in size.most_common()),
               "Dates in dd/mm/yyyy": ok_dates, "Negatives in brackets": brackets,
               "Negatives with a minus sign": minus}
    if dropped:
        summary["Findings not shown (cap per check)"] = dropped
    return _frame(capped), summary


def _frame(findings: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(findings, columns=COLUMNS)
    order = {"High": 0, "Medium": 1, "Low": 2}
    if len(df):
        df = df.sort_values("severity", key=lambda s: s.map(order), kind="stable").reset_index(drop=True)
    return df


# ---- entry points --------------------------------------------------------------------


def check_docx(data: bytes) -> tuple[pd.DataFrame, dict]:
    doc = Document(io.BytesIO(data))
    findings, summary = [], {}
    for fn in (check_fonts, check_dates, check_negatives):
        f, s = fn(doc)
        findings += f
        summary.update(s)
    return _frame(findings), summary


def check_file(name: str, data: bytes) -> tuple[pd.DataFrame, dict]:
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "docx":
        return check_docx(data)
    if ext in ("xlsx", "xlsm"):
        return check_xlsx(data)
    raise ValueError(f"Unsupported file type '.{ext}'. Upload a .docx or .xlsx file.")
