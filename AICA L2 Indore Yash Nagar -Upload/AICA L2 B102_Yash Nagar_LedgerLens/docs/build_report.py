"""Builds docs/LedgerLens_Project_Report.docx with python-docx.
The report text lives in docs/report_content.js (plain call syntax) and is evaluated here.
Run from the project root: python docs/build_report.py"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK, TEAL, LINE = RGBColor(0x1F, 0x2A, 0x44), RGBColor(0x0F, 0x76, 0x6E), "CBD5E1"
GREYTXT = RGBColor(0x64, 0x74, 0x8B)
ROOT = Path(__file__).resolve().parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)

st = doc.styles
st["Normal"].font.name, st["Normal"].font.size = "Arial", Pt(10.5)
st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st["Normal"].paragraph_format.space_after = Pt(5)
st["Normal"].paragraph_format.line_spacing = 1.18
for name, size, colour, before, after in (("Heading 1", 15, INK, 14, 6), ("Heading 2", 12, TEAL, 10, 4)):
    s = st[name]
    s.font.name, s.font.size, s.font.bold, s.font.color.rgb = "Arial", Pt(size), True, colour
    s.element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    s.element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
    s.paragraph_format.keep_with_next = True
st["List Bullet"].font.name = "Arial"

body = []


def _runs(par, parts):
    for text, is_bold in parts:
        r = par.add_run(text)
        r.bold = bool(is_bold)


def _norm(x):
    if isinstance(x, str):
        return [(x, False)]
    return [x] if isinstance(x, tuple) else list(x)


def run(t): return (t, False)
def b(t): return (t, True)
def h1(t): return ("h1", t)
def h2(t): return ("h2", t)
def p(parts, opts=None, **o): return ("p", _norm(parts), {**(opts or {}), **o})
def bullet(parts): return ("bullet", _norm(parts))
def gap(): return ("gap",)
def table(widths, header, rows): return ("table", widths, header, rows)


def shade(cell, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(sh)


def borders(t):
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), LINE)
        el.append(e)
    t._tbl.tblPr.append(el)


def emit(item):
    kind = item[0]
    if kind == "h1":
        if item[1][:2] in ("3.", "4.", "5."):       # each of the later sections starts on its own page, so no table header is stranded
            doc.add_page_break()
        doc.add_heading(item[1], 1)
    elif kind == "h2":
        doc.add_heading(item[1], 2)
    elif kind == "p":
        par = doc.add_paragraph()
        _runs(par, item[1])
        o = item[2]
        if o.get("center"):
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if o.get("before"):
            par.paragraph_format.space_before = Pt(o["before"])
        if o.get("after") is not None:
            par.paragraph_format.space_after = Pt(o["after"])
    elif kind == "bullet":
        par = doc.add_paragraph(style="List Bullet")
        _runs(par, item[1])
        par.paragraph_format.space_after = Pt(2)
    elif kind == "gap":
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
    elif kind == "table":
        _, widths, header, rows = item
        t = doc.add_table(rows=1, cols=len(widths))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        borders(t)
        total = sum(widths)

        def fill(row, vals, head):
            for i, v in enumerate(vals):
                c = row.cells[i]
                c.width = Cm(17 * widths[i] / total)
                for j, line in enumerate(str(v).split("\n")):
                    par = c.paragraphs[0] if j == 0 else c.add_paragraph()
                    par.paragraph_format.space_after = Pt(2)
                    par.paragraph_format.line_spacing = 1.08
                    r = par.add_run(line)
                    r.font.size = Pt(9.5)
                    r.bold = head
                    if head:
                        r.font.color.rgb = RGBColor(255, 255, 255)
                if head:
                    shade(c, "1F2A44")
        fill(t.rows[0], header, True)
        h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true")
        t.rows[0]._tr.get_or_add_trPr().append(h)
        for r in rows:
            row = t.add_row()
            fill(row, r, False)
            cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(cs)
        sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.line_spacing = 0.5
    elif kind == "picture":
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(str(ROOT / "architecture.png"), width=Cm(15.5))


# ---- title block (no cover page, no contents list: the report is five sections on five pages)
par = doc.add_paragraph()
r = par.add_run("LedgerLens"); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = INK
r = par.add_run("   Tally to reviewed Schedule III statements"); r.font.size = Pt(11.5); r.font.color.rgb = TEAL
bd = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
for k, v in (("val", "single"), ("sz", "10"), ("space", "4"), ("color", "0F766E")):
    e.set(qn(f"w:{k}"), v)
bd.append(e); par._p.get_or_add_pPr().append(bd)
par.paragraph_format.space_after = Pt(2)
emit(p([b("Project report"), run("  |  AICA Level 2 batch submission")], after=4))

# ---- content: JS call syntax is valid Python once object literals are quoted (strings are left alone)
src = (ROOT / "report_content.js").read_text(encoding="utf-8")
src = re.sub(r"^\s*//.*$", "", src, flags=re.M)
parts = re.split(r'("(?:[^"\\]|\\.)*")', src)
for i in range(0, len(parts), 2):
    s = parts[i].replace("true", "True").replace("false", "False")
    s = re.sub(r"\{\s*(\w+):", r'{"\1":', s)
    s = re.sub(r",\s*(\w+):\s", r', "\1": ', s)
    parts[i] = s
ns = dict(run=run, b=b, h1=h1, h2=h2, p=p, bullet=bullet, gap=gap, table=table, body=body, picture=lambda: ("picture",))
exec("".join(parts), ns)
for item in body:
    emit(item)

# ---- footer with page number
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("LedgerLens project report  |  page "); r.font.size = Pt(8.5); r.font.color.rgb = GREYTXT
r2 = fp.add_run(); r2.font.size = Pt(8.5); r2.font.color.rgb = GREYTXT
for typ, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if typ:
        f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), typ); r2._r.append(f)
    else:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r2._r.append(it)

doc.core_properties.title = "LedgerLens project report"
doc.core_properties.author = "LedgerLens"
out = ROOT / "LedgerLens_Project_Report.docx"
doc.save(out)
print("written", out)
