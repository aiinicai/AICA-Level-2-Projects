"""Review report as Word (memo) and Excel (full detail)."""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from openpyxl import Workbook

from ..format import dmy, inr
from ..render import FONT, SIZE, _cell, _para, _run, _shade, _sheet_setup, _table, _x, safe_name
from . import Review
from .checklist import STATUSES
from .ratios import fmt_value, fmt_variance

FILL = {"Complete": "E2EFDA", "Partly provided": "FFF2CC", "Information required": "FFE699", "Observation": "F8CBAD",
        "Applies (likely)": "DDEBF7", "Exempt (likely)": "DDEBF7", "Not applicable": "EDEDED",
        "Pass": "E2EFDA", "Fail": "F8CBAD", "High": "F8CBAD", "Medium": "FFE699", "Low": "DDEBF7"}
AI_NOT_RUN = ("The AI review was not run (no Gemini API key). Ratio explanations are left for management and no "
              "analytical observations were drafted.")


def _num(v) -> str:
    return inr(round(v)) if v is not None else "-"


def build_review_docx(r: Review, path: Path) -> Path:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.8)
        s.top_margin = s.bottom_margin = Inches(0.8)
        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(hp, "Reviewer's report: internal working paper, not part of the financial statements.", italic=True, size=8)
    n = doc.styles["Normal"]
    n.font.name, n.font.size = FONT, Pt(SIZE)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    _para(doc, f"Reviewer's report: {r.company}", True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, f"Draft financial statements as at {dmy(r.as_at)}", True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, "Casting and tie-outs, ratios and the disclosure checklist are computed from the books and the draft. "
               "Items labelled 'AI draft' were written by a language model from a table of facts and must be verified "
               "by a person before use.", italic=True)

    sm = r.summary
    _para(doc, "1. Summary", True, space_after=2)
    w = [3.4, 3.0]
    t = _table(doc, w)
    rows = [("Casting and tie-out checks", f"{sm['checks']} run: {sm['passed']} passed, {sm['failed']} failed"),
            ("Schedule III ratios", f"{len(r.ratios)} computed; {sm['ratios_flagged']} moved more than 25%; "
                                    f"{sm['ratios_unmeasured']} cannot be tested without the previous year's profit and loss"),
            ("Disclosure checklist", ", ".join(f"{k}: {v}" for k, v in sorted(sm['checklist'].items(), key=lambda kv: STATUSES.index(kv[0])))),
            ("AI analytical review", f"{sm['ai_observations']} observation(s), {sm['ai_ungrounded']} with figures not from the books"
             if r.ai_requested else "Not run")]
    for a, b in rows:
        c = t.add_row().cells
        _cell(c[0], a, w[0], bold=True)
        _cell(c[1], b, w[1])
    _para(doc, "")

    _para(doc, "2. Casting and tie-outs", True, space_after=2)
    by_area: dict = {}
    for c in r.tieouts:
        a = by_area.setdefault(c.area, [0, 0])
        a[0] += 1
        a[1] += c.status == "Fail"
    w = [2.6, 1.2, 1.2, 1.4]
    t = _table(doc, w)
    for cell, h, ww in zip(t.add_row().cells, ["Area", "Checks", "Passed", "Failed"], w):
        _cell(cell, h, ww, bold=True, right=h != "Area")
        _shade(cell, "D9E2F3")
    for area, (tot, bad) in by_area.items():
        c = t.add_row().cells
        _cell(c[0], area, w[0])
        _cell(c[1], str(tot), w[1], right=True)
        _cell(c[2], str(tot - bad), w[2], right=True)
        _cell(c[3], str(bad), w[3], right=True)
        if bad:
            _shade(c[3], FILL["Fail"])
    fails = [c for c in r.tieouts if c.status == "Fail"]
    if fails:
        _para(doc, "")
        _para(doc, "Differences found", True, space_after=2)
        w = [0.6, 3.3, 0.9, 0.9, 0.7]
        t = _table(doc, w)
        for cell, h, ww in zip(t.add_row().cells, ["Ref", "Check", "Expected", "Actual", "Difference"], w):
            _cell(cell, h, ww, bold=True)
            _shade(cell, "D9E2F3")
        for c in fails:
            cells = t.add_row().cells
            for cell, v, ww, right in zip(cells, [c.id, c.description, _num(c.expected), _num(c.actual), _num(c.difference)], w,
                                          (False, False, True, True, True)):
                _cell(cell, v, ww, right=right)
    else:
        _para(doc, "Every statement foots, every note agrees to its face, and the statements agree to the Tally books "
                   "(within rounding). The full list of checks is in the Excel workbook.", italic=True)
    _para(doc, "")

    doc.add_page_break()
    _para(doc, "3. Schedule III ratios", True, space_after=2)
    w = [1.9, 0.8, 0.8, 0.8, 2.1]
    t = _table(doc, w)
    for cell, h, ww in zip(t.add_row().cells, ["Ratio", "Current", "Previous", "Change", "Explanation / note"], w):
        _cell(cell, h, ww, bold=True, right=h in ("Current", "Previous", "Change"))
        _shade(cell, "D9E2F3")
    for x in r.ratios:
        c = t.add_row().cells
        _cell(c[0], x.name, w[0])
        _cell(c[1], fmt_value(x, x.cur), w[1], right=True)
        _cell(c[2], fmt_value(x, x.prev), w[2], right=True)
        _cell(c[3], fmt_variance(x.variance), w[3], right=True, bold=x.flagged)
        text = (f"{x.explanation} ({x.explanation_source})" if x.flagged else (x.note or ""))
        if not x.flagged and x.prev is None and x.cur is not None:
            text = (text + " " if text else "") + "Previous-year value needs the previous year's profit and loss."
        _cell(c[4], text, w[4], italic=x.flagged and x.explanation_source != "AI draft")
        if x.flagged:
            _shade(c[3], FILL["Observation"])
    _para(doc, "")
    _para(doc, "Definitions used", True, space_after=1)
    for x in r.ratios:
        _para(doc, f"• {x.name} = {x.formula}", space_after=1)
    _para(doc, "")

    doc.add_page_break()
    _para(doc, "4. Disclosure checklist and applicability screens", True, space_after=2)
    w = [0.5, 2.3, 1.1, 2.5]
    t = _table(doc, w)
    for cell, h, ww in zip(t.add_row().cells, ["Ref", "Requirement", "Status", "Evidence / what to obtain"], w):
        _cell(cell, h, ww, bold=True)
        _shade(cell, "D9E2F3")
    area = None
    for i in r.checklist:
        if i.area != area:
            area = i.area
            c = t.add_row().cells
            _cell(c[0], area, w[0], bold=True)
            m = c[0].merge(c[3])
            m.width = Inches(sum(w))
            _shade(m, "F2F2F2")
        c = t.add_row().cells
        _cell(c[0], i.id, w[0])
        _cell(c[1], i.requirement, w[1])
        _cell(c[2], i.status, w[2])
        _shade(c[2], FILL.get(i.status, "FFFFFF"))
        _cell(c[3], i.evidence + (f" Obtain: {i.ask}." if i.ask else ""), w[3])
    _para(doc, "")

    doc.add_page_break()
    _para(doc, "5. Analytical observations (AI draft)", True, space_after=2)
    if not r.ai_requested:
        _para(doc, AI_NOT_RUN, italic=True)
    else:
        for e in r.ai.errors:
            _para(doc, f"Problem: {e}", italic=True)
        if not r.ai.observations and not r.ai.errors:
            _para(doc, "The model raised no further observations beyond those already listed.", italic=True)
        for o in r.ai.observations:
            _para(doc, f"{o.title} [{o.severity}]" + ("" if o.grounded else "  (contains figures not taken from the books: verify)"),
                  True, space_after=1)
            _para(doc, f"Why it matters: {o.why}", space_after=1)
            _para(doc, f"Procedure: {o.procedure}", space_after=1)
            cited = "; ".join(f"{r.facts.get(i).label}: {r.facts.get(i).text}" for i in o.fact_ids if r.facts.get(i))
            _para(doc, f"Based on: {cited or 'no cited facts'}", italic=True, space_after=6)
    doc.save(path)
    return path


def build_review_xlsx(r: Review, path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Summary")
    _sheet_setup(ws, [40, 70])
    _x(ws, "A1", f"Reviewer's report: {r.company}", True, size=14)
    _x(ws, "A2", f"Draft financial statements as at {dmy(r.as_at)}", True, size=12)
    sm = r.summary
    rows = [("Tie-out checks run", sm["checks"]), ("Passed", sm["passed"]), ("Failed", sm["failed"]),
            ("Ratios moved more than 25%", sm["ratios_flagged"]), ("AI observations", sm["ai_observations"] if r.ai_requested else "not run"),
            *[(f"Checklist: {k}", v) for k, v in sorted(sm["checklist"].items(), key=lambda kv: STATUSES.index(kv[0]))]]
    for i, (a, b) in enumerate(rows, start=4):
        _x(ws, f"A{i}", a)
        _x(ws, f"B{i}", b, right=False)

    ws = wb.create_sheet("Tie-outs")
    _sheet_setup(ws, [10, 18, 80, 16, 16, 14, 10])
    for col, h in zip("ABCDEFG", ["Ref", "Area", "Check", "Expected", "Actual", "Difference", "Result"]):
        _x(ws, f"{col}1", h, True, fill="D9E2F3", right=col in "DEF")
    for i, c in enumerate(r.tieouts, start=2):
        for col, v in zip("ABCDEFG", [c.id, c.area, c.description, c.expected, c.actual, c.difference, c.status]):
            _x(ws, f"{col}{i}", v, num=col in "DEF" and v is not None)
        if c.status == "Fail":
            _x(ws, f"G{i}", c.status, True, fill=FILL["Fail"])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Ratios")
    _sheet_setup(ws, [34, 12, 12, 12, 10, 60, 70, 40])
    for col, h in zip("ABCDEFGH", ["Ratio", "Current", "Previous", "Change", "Flag", "Formula", "Explanation", "Source / note"]):
        _x(ws, f"{col}1", h, True, fill="D9E2F3", right=col in "BCD")
    for i, x in enumerate(r.ratios, start=2):
        pct = x.unit == "%"
        _x(ws, f"A{i}", x.name)
        for col, v in (("B", x.cur), ("C", x.prev)):
            c = _x(ws, f"{col}{i}", v if v is not None else "n/a", right=True)
            if v is not None:
                c.number_format = '0.00%_);(0.00%)' if pct else '0.00"x"_);(0.00"x")'
        c = _x(ws, f"D{i}", x.variance if x.variance is not None else "n/a", right=True)
        if x.variance is not None:
            c.number_format = '0.0%_);(0.0%)'
        _x(ws, f"E{i}", "More than 25%" if x.flagged else "")
        _x(ws, f"F{i}", x.formula, wrap=True)
        _x(ws, f"G{i}", x.explanation, wrap=True)
        _x(ws, f"H{i}", x.explanation_source or x.note, wrap=True)
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Checklist")
    _sheet_setup(ws, [8, 28, 70, 20, 80, 50, 40])
    for col, h in zip("ABCDEFG", ["Ref", "Area", "Requirement", "Status", "Evidence", "What to obtain", "Reference"]):
        _x(ws, f"{col}1", h, True, fill="D9E2F3")
    for i, it in enumerate(r.checklist, start=2):
        for col, v in zip("ABCDEFG", [it.id, it.area, it.requirement, it.status, it.evidence, it.ask, it.ref]):
            _x(ws, f"{col}{i}", v, wrap=col in "CEFG")
        _x(ws, f"D{i}", it.status, fill=FILL.get(it.status))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(r.checklist) + 1}"

    ws = wb.create_sheet("AI observations")
    _sheet_setup(ws, [40, 10, 60, 50, 60, 14])
    for col, h in zip("ABCDEF", ["Title", "Severity", "Why it matters", "Procedure", "Based on", "Grounded"]):
        _x(ws, f"{col}1", h, True, fill="D9E2F3")
    if not r.ai_requested:
        _x(ws, "A2", AI_NOT_RUN, italic=True)
    for i, o in enumerate(r.ai.observations, start=2):
        cited = "; ".join(f"{r.facts.get(f).label}: {r.facts.get(f).text}" for f in o.fact_ids if r.facts.get(f))
        for col, v in zip("ABCDEF", [o.title, o.severity, o.why, o.procedure, cited, "Yes" if o.grounded else "No: verify"]):
            _x(ws, f"{col}{i}", v, wrap=col in "CDE")
    ws = wb.create_sheet("Facts")
    _sheet_setup(ws, [8, 70, 24])
    for col, h in zip("ABC", ["ID", "Fact", "Value"]):
        _x(ws, f"{col}1", h, True, fill="D9E2F3")
    for i, f in enumerate(r.facts.items, start=2):
        _x(ws, f"A{i}", f.id)
        _x(ws, f"B{i}", f.label)
        _x(ws, f"C{i}", f.text)
    wb.save(path)
    return path


def write_review(r: Review, out_dir) -> tuple[Path, Path]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    stem = f"{safe_name(r.company)}_Review_{r.as_at:%Y%m%d}"
    return build_review_docx(r, out / f"{stem}.docx"), build_review_xlsx(r, out / f"{stem}.xlsx")
