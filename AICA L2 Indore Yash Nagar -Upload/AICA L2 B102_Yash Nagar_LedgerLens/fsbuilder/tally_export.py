"""Read TallyPrime's own Trial Balance export (Excel), exactly as Tally writes it.

Tally's layout: title rows (company, address, "Trial Balance", period), then a header block
    Particulars | Opening Balance | Transactions (Debit, Credit) | Closing Balance
then one row per ledger. Whether a balance is Dr or Cr is not in the value: it is in the cell's number
format (`0.00" Dr"` / `0.00" Cr"`). The export has no ledger group, so groups are inferred from the ledger
name and balance side unless the user also supplies a ledger list.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime

import pandas as pd

_MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
_DATE = re.compile(r"(\d{1,2})[-\s]([A-Za-z]{3})[A-Za-z]*[-\s,]*(\d{2,4})")
_DATE_ONLY = re.compile(r"\s*\d{1,2}[-\s][A-Za-z]{3,9}[-\s,]*\d{2,4}\s*")     # a header cell that is just a date
_SIDE_FMT = re.compile(r"\b(Dr|Cr)\b", re.I)
_SIDE_TEXT = re.compile(r"^\s*([\d,]*\.?\d+)\s*(Dr|Cr)\s*$", re.I)


@dataclass
class NativeTB:
    df: pd.DataFrame          # name, opening_amount, opening_dr_cr, period_debit, period_credit, closing_amount, closing_dr_cr
    company: str
    start: date | None
    end: date | None
    sheet: str


def _parse_dates(text: str) -> list[date]:
    out = []
    for d, mon, y in _DATE.findall(text or ""):
        m = _MONTHS.get(mon.lower())
        if not m:
            continue
        year = int(y) + (2000 if len(y) == 2 else 0)
        try:
            out.append(date(year, m, int(d)))
        except ValueError:
            pass
    return out


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _side_of(cell) -> tuple[float, str] | None:
    """(amount, 'Dr'/'Cr') for a balance cell, or None if empty. Raises ValueError if the side cannot be told."""
    v = cell.value
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, str):
        m = _SIDE_TEXT.match(v)
        if not m:
            return None
        return float(m.group(1).replace(",", "")), m.group(2).capitalize()
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        return None
    if v == 0:
        return 0.0, ""
    m = _SIDE_FMT.search((cell.number_format or "").replace('"', " "))
    if not m:
        raise ValueError(f"cannot tell whether {cell.coordinate} is a debit or a credit: the cell has no Dr/Cr marker")
    return abs(float(v)), m.group(1).capitalize()


def _column_roles(ws, first_header_row: int, last_header_row: int, particulars_col: int):
    """{column index: (group, kind)} with group in opening/transactions/closing and kind in balance/debit/credit."""
    merged = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = ws.cell(rng.min_row, rng.min_col).value
    roles = {}
    for c in range(particulars_col + 1, ws.max_column + 1):
        words = []
        for r in range(first_header_row, last_header_row + 1):
            v = merged.get((r, c), ws.cell(r, c).value)
            words.append(_text(v).lower())
        label = " ".join(words)
        group = ("opening" if "opening" in label else "closing" if "closing" in label
                 else "transactions" if "transaction" in label else None)
        kind = "debit" if "debit" in label else "credit" if "credit" in label else "balance" if "balance" in label else None
        if kind and not group:
            group = "closing"      # a bare Debit / Credit pair is the closing balance
        if group and kind:
            roles[c] = (group, kind)
    return roles


def parse_tally_trial_balance(data: bytes) -> NativeTB | None:
    """Return the parsed trial balance, or None if the workbook is not in Tally's layout."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        return None
    for ws in wb.worksheets:
        pr = pc = None
        for r in range(1, min(ws.max_row, 30) + 1):
            for c in range(1, min(ws.max_column, 6) + 1):
                if _text(ws.cell(r, c).value).lower() == "particulars":
                    pr, pc = r, c
                    break
            if pr:
                break
        if not pr:
            continue
        # header block: from the Particulars row until the first row that has a ledger name and a number
        data_row = None
        for r in range(pr + 1, min(ws.max_row, pr + 8) + 1):
            name = _text(ws.cell(r, pc).value)
            has_num = any(isinstance(ws.cell(r, c).value, (int, float)) for c in range(pc + 1, ws.max_column + 1))
            if name and has_num:
                data_row = r
                break
        if not data_row:
            continue
        roles = _column_roles(ws, pr, data_row - 1, pc)
        if not any(g == "closing" for g, _ in roles.values()):
            continue

        rows = []
        for r in range(data_row, ws.max_row + 1):
            name = _text(ws.cell(r, pc).value)
            low = name.lower()
            if not name or low.startswith(("grand total", "total")) or low == "particulars":
                continue
            rec = {"name": name, "opening_amount": 0.0, "opening_dr_cr": "", "closing_amount": 0.0, "closing_dr_cr": "",
                   "period_debit": 0.0, "period_credit": 0.0}
            net = {"opening": 0.0, "closing": 0.0}
            for c, (grp, kind) in roles.items():
                cell = ws.cell(r, c)
                if grp == "transactions":
                    v = cell.value
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        rec["period_debit" if kind == "debit" else "period_credit"] += abs(float(v))
                    continue
                try:
                    got = _side_of(cell) if kind == "balance" else (
                        (abs(float(cell.value)), "Dr" if kind == "debit" else "Cr")
                        if isinstance(cell.value, (int, float)) and cell.value else None)
                except ValueError as e:
                    raise ValueError(f"Row '{name}': {e}. Export the trial balance from Tally again, keeping the "
                                     "default Excel format (Alt+E, Excel (Spreadsheet)).") from None
                if got and got[0]:
                    net[grp] += got[0] if got[1] == "Dr" else -got[0]
            for grp in ("opening", "closing"):
                rec[f"{grp}_amount"] = round(abs(net[grp]), 2)
                rec[f"{grp}_dr_cr"] = "" if abs(net[grp]) < 0.005 else ("Dr" if net[grp] > 0 else "Cr")
            rows.append(rec)
        if not rows:
            continue

        header_text = [_text(ws.cell(r, c).value) for r in range(1, pr) for c in range(1, ws.max_column + 1)]
        company = next((t for t in header_text if t), "")
        dates = []
        for t in header_text + [_text(ws.cell(pr, c).value) for c in range(1, ws.max_column + 1)]:
            if t and (t.lower().startswith(("for ", "as on", "as at")) or " to " in t.lower() or _DATE_ONLY.fullmatch(t)):
                dates = _parse_dates(t)
                if dates:
                    break
        as_on = any(t.lower().startswith(("as on", "as at")) for t in header_text if t)
        start = dates[0] if dates and not as_on else None      # "As on 31-Mar-26" gives a closing date, not a period
        end = dates[-1] if dates else None
        return NativeTB(pd.DataFrame(rows), company, start, end, ws.title)
    return None


# ---- ledger group inference (used only when the file has no group column) --------------------------------

def _c(p):
    return re.compile(p, re.I)


# (pattern, group, side the balance should normally have or None). First match wins; all are high confidence.
_RULES = [
    (_c(r"^profit\s*(&|and)\s*loss"), "Primary", None),
    (_c(r"^opening stock$|stock.in.hand|closing stock|finished goods stock|raw material stock|^stock\b"), "Stock-in-Hand", None),
    (_c(r"^cash\b|cash.in.hand|petty cash"), "Cash-in-Hand", None),
    (_c(r"\b(od|overdraft|cash credit)\b.*(a/c|account)?|bank od"), "Bank OD A/c", None),
    (_c(r"fixed deposit|\bfd\b|term deposit"), "Investments", "Dr"),
    (_c(r"\b(cgst|sgst|igst|utgst|gst)\b|\btds\b|\btcs\b|professional tax|\bpf payable\b|\besi payable\b"), "Duties & Taxes", None),
    (_c(r"share capital|equity capital|capital account|^capital\b|proprietor|partners? capital"), "Capital Account", "Cr"),
    (_c(r"reserve|surplus|securities premium"), "Reserves & Surplus", "Cr"),
    (_c(r"director|promoter|shareholder|relative|unsecured loan|from (a )?friend|inter.?corporate"), "Unsecured Loans", "Cr"),
    (_c(r"term loan|\bloan\b.*\b(bank|hdfc|icici|sbi|axis|kotak)\b|\b(bank|hdfc|icici|sbi|axis|kotak)\b.*\bloan\b|vehicle loan|car loan|secured loan"), "Secured Loans", "Cr"),
    # Banks come after the loan rules so that "Term Loan - HDFC Bank" is a loan, not a bank account.
    (_c(r"\bbank\b.*(current|savings|a/c|account)|(current|savings).*\bbank\b|\b(hdfc|icici|sbi|axis|kotak|pnb|yes|idfc|canara|bob)\b.*(a/c|account|bank)"), "Bank Accounts", None),
    (_c(r"^outstanding\b|^accrued\b|provision for|^provisions?\b"), "Provisions", "Cr"),
    # Words like repairs or insurance make it an expense even if it names an asset ("Machinery Repairs").
    (_c(r"^(?!.*(repair|maintenance|insurance|depreciation|expense|rent\b|hire|charges|running|written off)).*"
        r"(\bplant\b|machinery|furniture|fixture|computer|office equipment|vehicle|\bcar\b|building|premises|\bland\b|equipment|"
        r"software|goodwill|patent|trademark|capital work)"), "Fixed Assets", "Dr"),
    (_c(r"^sales?\b|sales -|sale of|^revenue|turnover"), "Sales Accounts", "Cr"),
    (_c(r"^purchases?\b|purchase -|purchase of"), "Purchase Accounts", "Dr"),
    (_c(r"interest (received|income)|discount received|commission received|dividend|rent received|misc(ellaneous)? income|"
        r"other income|profit on sale|exchange gain|scrap sales?"), "Indirect Incomes", "Cr"),
    (_c(r"factory\s*(wages|labou?r|expenses|salar|overhead)|job work|power\b|fuel\b|freight inward|carriage inward|^wages\b|manufactur|processing charges|import duty"), "Direct Expenses", "Dr"),
    (_c(r"deposit"), "Deposits (Asset)", "Dr"),
    (_c(r"advance|loans? (to|&)|prepaid|staff loan|employee loan"), "Loans & Advances (Asset)", "Dr"),
    (_c(r"expenses?\b|charges\b|\bfees?\b|rent\b|salar|insurance|telephone|travel|conveyance|printing|stationery|repairs?|"
        r"advert|commission|discount allowed|interest (paid|on)|depreciation|bank charges|electricity|legal|professional|"
        r"consult|audit|postage|courier|subscription|donation|bad debts?|office|general|staff welfare|remuneration|"
        r"freight outward|carriage outward|miscellaneous"), "Indirect Expenses", "Dr"),
]


def infer_group(name: str, closing_side: str, opening_side: str = "") -> tuple[str, str]:
    """(Tally group, confidence). 'high' when the name says what it is; 'low' when it is a guess from the balance side."""
    side = closing_side or opening_side
    for pattern, group, expected in _RULES:
        if pattern.search(name) and (expected is None or not side or side == expected):
            return group, "high"
    # Anything else with a balance: parties. Debit balances are debtors, credit balances are creditors.
    if side == "Dr":
        return "Sundry Debtors", "low"
    if side == "Cr":
        return "Sundry Creditors", "low"
    return "Indirect Expenses", "low"
