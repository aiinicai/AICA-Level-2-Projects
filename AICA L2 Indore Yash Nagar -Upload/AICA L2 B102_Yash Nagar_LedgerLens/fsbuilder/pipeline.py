"""Non-UI logic behind the Streamlit page, kept here so it can be tested without a browser."""
from __future__ import annotations

import copy
import hashlib
import io
import json
import tempfile
from dataclasses import replace
from datetime import date
from pathlib import Path

import pandas as pd

from .format import allowed_units, as_unit, choose_unit, inr, rnd
from .tally_export import infer_group, parse_tally_trial_balance
from .mapping import Rulebook
from .model import Model, ratio_note
from .render import write_all
from .reviewer import Review
from .reviewer.report import write_review
from .tb import DEFAULT_GROUPS, TrialBalance, _build, _signed

ROOT = Path(__file__).resolve().parents[1]
SAMPLE_TB, SAMPLE_GROUPS = ROOT / "samples" / "rpi_trial_balance.csv", ROOT / "samples" / "rpi_groups.csv"
SAMPLE_COMPANY, SAMPLE_FROM, SAMPLE_TO = "Rajasthan Precision Industries Private Limited", date(2024, 4, 1), date(2025, 3, 31)
REQUIRED = ["name", "group", "opening_amount", "opening_dr_cr", "closing_amount", "closing_dr_cr"]
MIME = {"docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
UNMAPPED_OPTION = "UNMAPPED · (not mapped)"


def last_financial_year(today: date | None = None) -> tuple[date, date]:
    today = today or date.today()
    end_year = today.year if today >= date(today.year, 4, 1) else today.year - 1
    return date(end_year - 1, 4, 1), date(end_year, 3, 31)


def _drcr(v) -> str:
    s = "" if pd.isna(v) else str(v).strip().lower()
    return {"dr": "Dr", "debit": "Dr", "d": "Dr", "cr": "Cr", "credit": "Cr", "c": "Cr", "": ""}.get(s, s)


# Header spellings people actually use, mapped to our column names.
_HEADER_ALIASES = {
    "name": {"name", "ledger", "ledger_name", "ledger_account", "account", "account_name", "particulars"},
    "group": {"group", "under", "parent", "group_name", "tally_group", "account_group"},
    "opening_amount": {"opening_amount", "opening_balance", "opening_bal", "opening", "opening_amt"},
    "opening_dr_cr": {"opening_dr_cr", "opening_dr/cr", "opening_drcr", "opening_dr_or_cr", "opening_side", "opening_type"},
    "closing_amount": {"closing_amount", "closing_balance", "closing_bal", "closing", "closing_amt"},
    "closing_dr_cr": {"closing_dr_cr", "closing_dr/cr", "closing_drcr", "closing_dr_or_cr", "closing_side", "closing_type"},
    # Deliberately NOT plain "debit"/"credit": in an Excel trial balance those usually mean the closing balance.
    "period_debit": {"period_debit", "period_dr", "debit_during_period", "debits_during_period"},
    "period_credit": {"period_credit", "period_cr", "credit_during_period", "credits_during_period"},
}
_ALIAS_LOOKUP = {a: k for k, names in _HEADER_ALIASES.items() for a in names}
_GROUP_ALIASES = {"name": {"name", "group", "group_name"}, "parent": {"parent", "under", "parent_group"},
                  "nature": {"nature", "type"}}
_GROUP_LOOKUP = {a: k for k, names in _GROUP_ALIASES.items() for a in names}


def _norm_header(c, lookup=None) -> str:
    s = "" if pd.isna(c) else str(c).strip().lower().replace("‑", "-")
    s = "_".join(s.replace("-", " ").replace(".", " ").split())
    s = s.replace("(", "").replace(")", "")
    return (lookup or _ALIAS_LOOKUP).get(s, s)


def _load_table(file, filename: str, what: str, needed: set, lookup=None) -> pd.DataFrame:
    """Read a CSV or an Excel workbook into a frame whose columns are our canonical names.

    For Excel the first sheet that has a header row with every column in `needed` is used, wherever
    that row sits (title rows above it are fine)."""
    ext = (filename or "").lower().rsplit(".", 1)[-1]
    if ext in ("xlsx", "xlsm"):
        try:
            xls = pd.ExcelFile(file)
        except Exception as e:
            raise ValueError(f"Could not open {what} as an Excel workbook: {e}") from e
        for sheet in xls.sheet_names:
            raw = pd.read_excel(xls, sheet_name=sheet, header=None, dtype=object)
            for i in range(min(len(raw), 40)):
                header = [_norm_header(c, lookup) for c in raw.iloc[i].tolist()]
                if needed <= set(header):
                    df = raw.iloc[i + 1:].copy()
                    df.columns = header
                    df = df.loc[:, [c for c in df.columns if c]]
                    df = df.dropna(how="all").reset_index(drop=True)
                    df.attrs["sheet"] = sheet
                    return df
        raise ValueError(f"No sheet in {what} has a header row with the columns: {', '.join(sorted(needed))}. "
                         f"Sheets looked at: {', '.join(xls.sheet_names)}. Download the template from the page to see the layout.")
    if ext not in ("csv", "txt", ""):
        raise ValueError(f"{what} must be a .csv or .xlsx file, not .{ext}.")
    try:
        df = pd.read_csv(file)
    except Exception as e:
        raise ValueError(f"Could not read {what} as CSV: {e}") from e
    df.columns = [_norm_header(c, lookup) for c in df.columns]
    return df


def _to_number(col: pd.Series) -> pd.Series:
    """Numbers, or numbers typed as text such as '1,23,456.00' or 'Rs 500'."""
    def one(v):
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return float("nan")
        if isinstance(v, (int, float)):
            return float(v)
        s = str(v).replace(",", "").replace("₹", "").replace("Rs.", "").replace("Rs", "").strip()
        neg = s.startswith("(") and s.endswith(")")
        s = s.strip("()")
        try:
            return -float(s) if neg else float(s)
        except ValueError:
            return float("nan")
    return col.map(one)


def _native_export(file, filename):
    """Tally's own Excel trial balance export, or None if the file is in some other layout."""
    if (filename or "").lower().rsplit(".", 1)[-1] not in ("xlsx", "xlsm"):
        return None
    data = file.getvalue() if hasattr(file, "getvalue") else file.read()
    if hasattr(file, "seek"):
        file.seek(0)
    return parse_tally_trial_balance(data)


def _prepare_native(native, company, start, end, ledger_list_file, ledger_list_filename):
    """Turn a parsed Tally export into the frame the common reader expects. Returns (raw, company, start, end, inferred, notes)."""
    notes = []
    if native.end and not native.start:       # "As on <date>": the closing date is the file's, the start is as entered
        end = native.end
    if native.start and native.end:
        if native.start == native.end:
            fy = (start, end) if (start and end and (end - start).days >= 300) else last_financial_year()
            raise ValueError(
                f"This trial balance covers a single day ({native.start:%d/%m/%Y}), so it holds no year of transactions. "
                f"In TallyPrime open the Trial Balance, press F2 (Period), set it to the whole financial year "
                f"({fy[0]:%d/%m/%Y} to {fy[1]:%d/%m/%Y} for example) and export to Excel again.")
        if (start, end) != (native.start, native.end) and start and end:
            notes.append(f"The period was taken from the file ({native.start:%d/%m/%Y} to {native.end:%d/%m/%Y}), "
                         f"not from the dates entered ({start:%d/%m/%Y} to {end:%d/%m/%Y}).")
        start, end = native.start, native.end
    entered = (company or "").strip()
    if native.company and entered and " ".join(entered.lower().split()) != " ".join(native.company.lower().split()):
        notes.append(f"The file is for '{native.company}'; the statements carry the name entered, '{entered}'.")
    company = entered or native.company

    raw = native.df.copy()
    given = {}
    if ledger_list_file is not None:
        lst = _load_table(ledger_list_file, ledger_list_filename, "the ledger list", {"name", "group"})
        given = {str(n).strip(): str(g).strip() for n, g in zip(lst["name"], lst["group"]) if pd.notna(n) and pd.notna(g)}
    groups, inferred = [], {}
    for r in raw.itertuples(index=False):
        if r.name in given:
            groups.append(given[r.name])
        else:
            g, conf = infer_group(r.name, r.closing_dr_cr, r.opening_dr_cr)
            groups.append(g)
            inferred[r.name] = (g, conf)
    raw["group"] = groups
    if inferred:
        live = [k for k, v in inferred.items() if v[1] == "low" and
                abs(float(raw.loc[raw["name"] == k, ["opening_amount", "closing_amount"]].abs().sum().sum())) >= 0.005]
        notes.append(f"The trial balance has no ledger groups, so groups were worked out from the ledger names and balance "
                     f"sides for {len(inferred)} ledger(s); {len(live)} of them are guesses and are flagged for confirmation. "
                     "Upload the ledger list with groups from Tally to remove the guesswork.")
    return raw, company, start, end, inferred, notes


def read_trial_balance(file, filename="trial_balance.csv", groups_file=None, groups_filename="groups.csv",
                       company="", start=None, end=None, ledger_list_file=None, ledger_list_filename="") -> TrialBalance:
    """Read an uploaded trial balance: Tally's own Excel export, or the template layout (CSV or Excel).
    Raises ValueError with an actionable message."""
    native = _native_export(file, filename)
    inferred, notes = {}, []
    if native is not None:
        raw, company, start, end, inferred, notes = _prepare_native(native, company, start, end, ledger_list_file,
                                                                    ledger_list_filename)
    else:
        raw = _load_table(file, filename, "the trial balance", set(REQUIRED[:2]))
    missing = [c for c in REQUIRED if c not in raw.columns]
    if missing:
        sheet = f" (sheet '{raw.attrs['sheet']}')" if raw.attrs.get("sheet") else ""
        raise ValueError(f"The trial balance{sheet} is missing column(s): " + ", ".join(missing) +
                         ". Expected: " + ", ".join(REQUIRED) + " (period_debit and period_credit are optional).")
    names = raw["name"].astype(str).str.strip()
    is_total = names.str.lower().str.match(r"^(grand\s+)?totals?\b")      # a totals row is not a ledger
    raw = raw[raw["name"].notna() & (names != "") & ~is_total].reset_index(drop=True)
    if raw.empty:
        raise ValueError("The trial balance has no rows.")
    for c in ("opening_dr_cr", "closing_dr_cr"):
        raw[c] = raw[c].map(_drcr)
        bad = sorted(set(raw[c]) - {"Dr", "Cr", ""})
        if bad:
            raise ValueError(f"Column {c} must contain Dr, Cr or be blank; found: {', '.join(map(str, bad))}.")
    for c in ("opening_amount", "closing_amount"):
        raw[c] = _to_number(raw[c])
        side = raw[c.replace("amount", "dr_cr")]
        if raw[c].isna().any() and (side != "").any():
            rows = raw.index[raw[c].isna() & (side != "")].tolist()
            raise ValueError(f"Column {c} has non-numeric values in row(s) {[r + 2 for r in rows[:5]]}.")
    groups = dict(DEFAULT_GROUPS)
    if groups_file is not None:
        try:
            g = _load_table(groups_file, groups_filename, "the groups file", {"name", "parent"}, _GROUP_LOOKUP).fillna("")
            if "nature" not in g.columns:
                g["nature"] = ""
            groups.update({r["name"]: (r["parent"], r["nature"]) for _, r in g.iterrows()})
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Could not read the groups file (needs name, parent, nature): {e}") from e
    raw = raw.rename(columns={"period_debit": "period_dr", "period_credit": "period_cr"})
    for c in ("period_dr", "period_cr"):
        if c in raw:
            raw[c] = _to_number(raw[c])
    raw = raw.assign(
        opening=[_signed(a, d) for a, d in zip(raw["opening_amount"], raw["opening_dr_cr"])],
        closing=[_signed(a, d) for a, d in zip(raw["closing_amount"], raw["closing_dr_cr"])])
    known = set(groups) | {"Primary"}
    unknown = sorted(set(raw["group"].fillna("").astype(str)) - known)
    if unknown:
        raise ValueError("These ledger groups are not Tally's standard groups and no groups file was given: " +
                         ", ".join(unknown) + ". Upload a groups file (name, parent, nature) as well.")
    tb = _build(raw, groups, company, start, end)
    return replace(tb, adjustments=list(tb.adjustments) + notes, inferred=inferred)


def read_trial_balance_csv(tb_file, groups_file=None, company="", start=None, end=None) -> TrialBalance:
    """Kept for existing callers: the CSV-only entry point."""
    return read_trial_balance(tb_file, "trial_balance.csv", groups_file, "groups.csv", company, start, end)


def sample_xlsx_bytes() -> bytes:
    """An Excel template with the sample company's trial balance, plus a sheet explaining each column."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    df = pd.read_csv(SAMPLE_TB)
    wb = Workbook()
    ws = wb.active
    ws.title = "Trial balance"
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in row])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJ", (36, 22, 22, 15, 15, 15, 15, 15, 15, 15)):
        ws.column_dimensions[col].width = w
    notes = wb.create_sheet("Read me")
    for r in [("Column", "Meaning"),
              ("name", "Ledger name exactly as in Tally"),
              ("group", "The ledger's immediate group in Tally (for example Sundry Debtors)"),
              ("opening_amount", "Balance at the start of the period, as a positive number"),
              ("opening_dr_cr", "Dr or Cr for the opening balance (blank if nil)"),
              ("closing_amount", "Balance at the end of the period, as a positive number"),
              ("closing_dr_cr", "Dr or Cr for the closing balance (blank if nil)"),
              ("period_debit, period_credit", "Optional: transactions in the period; enables the PPE movement table"),
              ("Tip", "Keep the header row; title rows above it are fine. Use the 'Trial balance' sheet or any sheet with these headers.")]:
        notes.append(list(r))
    for c in notes[1]:
        c.font = Font(bold=True)
    notes.column_dimensions["A"].width, notes.column_dimensions["B"].width = 30, 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def load_sample() -> TrialBalance:
    from .tb import load_snapshot
    return load_snapshot(SAMPLE_TB, SAMPLE_GROUPS, SAMPLE_COMPANY, SAMPLE_FROM, SAMPLE_TO)


def line_options(rb: Rulebook) -> list[str]:
    ordered = sorted(rb.lines, key=lambda k: rb.lines[k]["order"])
    return [UNMAPPED_OPTION] + [f"{k} · {rb.lines[k]['label']}" for k in ordered]


def option_for(rb: Rulebook, line: str) -> str:
    return UNMAPPED_OPTION if line not in rb.lines else f"{line} · {rb.lines[line]['label']}"


def editor_frame(mapped: pd.DataFrame, rb: Rulebook) -> pd.DataFrame:
    """The table the user edits: only line, sub-head and reason are editable."""
    return pd.DataFrame({
        "Ledger": mapped["ledger"], "Tally group": mapped["group"],
        "Closing Dr / (Cr)": [inr(rnd(x)) for x in mapped["closing"]],
        "Schedule III line": [option_for(rb, l) for l in mapped["line"]],
        "Sub-head": mapped["sub"], "Source": mapped["source"],
        "Review note": mapped["review"], "Override reason": ""})


def overrides_from_edits(rule_mapped: pd.DataFrame, edited: pd.DataFrame, previous: dict | None = None) -> dict:
    """Any ledger whose line or sub-head differs from what the rules gave becomes an override."""
    base = rule_mapped.set_index("ledger")[["line", "sub"]]
    prev = previous or {}
    out = {}
    for row in edited.itertuples(index=False):
        ledger = row[0]
        line = str(row[3]).split(" · ")[0]
        sub = "" if pd.isna(row[4]) else str(row[4])
        if line == "UNMAPPED":
            continue
        b = base.loc[ledger]
        if line != b["line"] or sub != b["sub"]:
            reason = str(row[7]).strip() if not pd.isna(row[7]) else ""
            out[ledger] = {"line": line, "sub": sub, "reason": reason or prev.get(ledger, {}).get("reason", "") or "set in the app"}
    return out


def parse_overrides_json(data: bytes, rb: Rulebook) -> dict:
    try:
        d = json.loads(data.decode("utf-8"))
    except Exception as e:
        raise ValueError(f"Overrides file is not valid JSON: {e}") from e
    if not isinstance(d, dict):
        raise ValueError("Overrides file must be a JSON object: {ledger: {line, sub, reason}}.")
    for k, v in d.items():
        if not isinstance(v, dict) or v.get("line") not in rb.lines:
            raise ValueError(f"Override for '{k}' needs a valid 'line' (one of: {', '.join(rb.lines)}).")
    return d


def statement_frame(rows, prev_col=True, note_col=True, unit=None) -> pd.DataFrame:
    u = as_unit(unit)
    data = []
    for r in rows:
        if r.kind == "blank":
            continue
        label = r.label + (f"  [{', '.join(r.flags)}]" if r.flags else "")
        no_amount = r.kind in ("head", "sub") or (r.cur is None and r.prev is None and not r.note)
        rec = {"Particulars": label}
        if note_col:
            rec["Note"] = str(r.note) if r.note else ""
        rec["Current year"] = "" if no_amount else u.fmt(r.cur)
        rec["Previous year"] = "" if no_amount else (u.fmt(r.prev) if prev_col else "-")
        data.append(rec)
    return pd.DataFrame(data)


def total_income(mapped: pd.DataFrame) -> float:
    """Revenue from operations plus other income, in rupees: what decides the rounding units Schedule III allows."""
    from .statements import line_totals
    t = line_totals(mapped)
    return float(t.loc[t["line"].isin(["REV.OPS", "REV.OI"]), "closing"].sum())


def unit_choices(mapped: pd.DataFrame):
    """(permitted units, default) for this company."""
    income = total_income(mapped)
    return allowed_units(income), choose_unit(income)


def signature(model: Model) -> str:
    """Changes whenever the numbers or flags change, so stale AI text or files can be detected."""
    payload = json.dumps([[(r.label, r.cur, r.prev) for r in model.bs + model.pl + model.cashflow],
                          [f.text for f in model.flags], model.unit.key], default=str)
    return hashlib.sha1(payload.encode()).hexdigest()[:12]


def with_ratio_reasons(model: Model, review: Review) -> Model:
    """The model with any AI-drafted reasons for ratio changes over 25% written into the ratios note, marked as drafts.
    Reasons that are still blank stay as 'management to explain'."""
    reasons = {r.id: f"{r.explanation} ({r.explanation_source})" for r in review.ratios
               if r.flagged and r.explanation and r.explanation_source.startswith("AI draft")}
    if not reasons:
        return model
    out = copy.copy(model)
    out.notes = list(model.notes)
    i = next(i for i, n in enumerate(out.notes) if n.title == "Analytical ratios")
    out.notes[i] = ratio_note(model, out.notes[i].no, reasons)
    return out


def build_files(model: Model, review: Review, sign=None) -> dict:
    """{label: (file name, bytes, mime)} for the four deliverables. `sign` holds the signatories for the signature blocks."""
    model = copy.copy(with_ratio_reasons(model, review))
    model.review, model.sign = review, sign        # the statements Excel carries the AI analysis; every statement the signatures
    with tempfile.TemporaryDirectory() as d:
        fs_docx, fs_xlsx = write_all(model, d)
        rv_docx, rv_xlsx = write_review(review, d)
        return {
            "Statements (Word)": (fs_docx.name, fs_docx.read_bytes(), MIME["docx"]),
            "Statements (Excel)": (fs_xlsx.name, fs_xlsx.read_bytes(), MIME["xlsx"]),
            "Review (Word)": (rv_docx.name, rv_docx.read_bytes(), MIME["docx"]),
            "Review (Excel)": (rv_xlsx.name, rv_xlsx.read_bytes(), MIME["xlsx"]),
        }


def sample_csv_bytes() -> bytes:
    return SAMPLE_TB.read_bytes()


def overrides_bytes(overrides: dict) -> bytes:
    return json.dumps(overrides, indent=2, ensure_ascii=False).encode("utf-8")
