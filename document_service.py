"""Generate a professional Word reply from exactly the current editor text."""
from io import BytesIO
import re
from docx import Document
from docx.shared import Inches, Pt
from utils.validators import AppError


def generate_docx(reply_text: str) -> bytes:
    if not reply_text.strip():
        raise AppError("Enter a reply before downloading a Word document.")
    try:
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(0.85)
        section.left_margin = section.right_margin = Inches(0.9)
        normal = document.styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.1
        document.core_properties.author = ""
        document.core_properties.title = "Tax Notice Reply"
        for line in reply_text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            if line.lstrip().lower().startswith("subject:") or re.match(r"^\d+[.)]\s", line):
                run.bold = True
                paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.widow_control = True
            if not line:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 0.5
        output = BytesIO()
        document.save(output)
        return output.getvalue()
    except AppError:
        raise
    except Exception:
        raise AppError("The Word reply could not be created. Remove unusual control characters and try again.") from None
