"""
word_export.py
----------------
A Word (.docx) version of the flux report — useful when the commentary
needs to be pasted into a broader close memo or tracked-changes reviewed
by a controller in Word directly.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..legal import PROJECT_NOTICE, DRAFT_WATERMARK, FINAL_WATERMARK


def build_docx(run: dict) -> bytes:
    meta = run["meta"]
    rows = list(run["rows"].values())
    ready, n_material, n_approved = _readiness(run)
    watermark = FINAL_WATERMARK if ready else DRAFT_WATERMARK

    doc = Document()

    title = doc.add_heading(f"{meta.get('statement_label', '')} Flux Analysis", level=1)

    wm = doc.add_paragraph()
    wm_run = wm.add_run(watermark)
    wm_run.bold = True
    wm_run.font.color.rgb = RGBColor(0x1E, 0x7A, 0x46) if ready else RGBColor(0xB0, 0x00, 0x00)

    sub = doc.add_paragraph(f"{meta.get('comparison_label', '')} | Materiality threshold: {meta.get('threshold_pct', '')}% (min. $1,000)")
    sub.runs[0].font.size = Pt(9)
    sub.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    legal = doc.add_paragraph(PROJECT_NOTICE)
    legal.runs[0].italic = True
    legal.runs[0].font.size = Pt(7)
    legal.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_heading("Material Flux Lines", level=2)

    material_rows = [r for r in rows if r["data"]["is_material"]]
    if not material_rows:
        doc.add_paragraph("No lines breached the materiality threshold for this comparison.")
    else:
        table = doc.add_table(rows=1, cols=9)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        headers = ["Entity", "Line Item", "Prior ($)", "Current ($)", "Var ($)", "Var (%)", "Status",
                   "AI-Generated Commentary (reference only)", "Controller Final Commentary (exported)"]
        for i, h in enumerate(headers):
            hdr[i].text = h
            for p in hdr[i].paragraphs:
                for r in p.runs:
                    r.bold = True

        for r in material_rows:
            d = r["data"]
            cells = table.add_row().cells
            cells[0].text = d["entity"]
            cells[1].text = d["line_item"]
            cells[2].text = f"{d['prior_amount']:,.0f}"
            cells[3].text = f"{d['current_amount']:,.0f}"
            cells[4].text = f"{d['variance_usd']:,.0f}"
            cells[5].text = f"{d['variance_pct']:,.1f}%"
            cells[6].text = r["status"].upper()
            cells[7].text = r["ai_commentary"]
            for p in cells[7].paragraphs:
                for docx_run in p.runs:
                    docx_run.italic = True
                    docx_run.font.color.rgb = RGBColor(0x5A, 0x64, 0x72)
            cells[8].text = r["controller_commentary"]

    doc.add_heading("Controller Sign-Off", level=2)
    signoff = run["signoff"]
    ts = signoff.get("signed_off_at")
    ts_str = datetime.fromtimestamp(ts, tz=timezone.utc).strftime("%Y-%m-%d %H:%M UTC") if ts else "(pending)"
    status_line = "FINAL — all material items controller-approved" if ready else f"DRAFT — {n_approved}/{n_material} material items approved"
    doc.add_paragraph(f"Status: {status_line}")
    doc.add_paragraph(f"Controller: {signoff.get('controller_name') or '(pending)'} — {signoff.get('controller_title') or '(pending)'}")
    doc.add_paragraph(f"Signed off: {ts_str}")
    note = doc.add_paragraph(
        "The AI-Generated Commentary column is a system-drafted reference note and is never treated as "
        "final. The Controller Final Commentary column is the one authoritative, editable field, and this "
        "document only carries the FINAL watermark once every material line's status reads APPROVED and "
        "the controller has signed off. Final judgement on each flagged item rests with the controller, "
        "not with the rule engine or any AI-assisted wording."
    )
    note.runs[0].italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _readiness(run):
    rows = list(run["rows"].values())
    material_rows = [r for r in rows if r["data"]["is_material"]]
    approved = [r for r in material_rows if r["status"] == "approved"]
    signed = run["signoff"]["controller_name"] is not None
    ready = signed and len(material_rows) > 0 and len(approved) == len(material_rows)
    return ready, len(material_rows), len(approved)
