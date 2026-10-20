"""Recipient list (clients and employees): Excel/CSV import with validation, and a template.

Required columns: Name (or Client Name), Email ID. Optional: Group (Client / Employee),
Salutation, Company, Consent. Header names are matched flexibly (e.g. "E-mail", "Email Address").
Rows with Consent = No / Opted out are kept out of the send list. A blank Group means Client.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

_EMAIL_RE = re.compile(r"^[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?"
                       r"(?:\.[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?)+$")

HEADER_ALIASES = {
    "name": ("client name", "name", "client", "contact name", "full name", "recipient name", "employee name",
             "staff name"),
    "email": ("email id", "email", "e-mail", "email address", "e-mail id", "mail id", "email-id"),
    "salutation": ("salutation", "title", "prefix"),
    "company": ("company", "organisation", "organization", "firm", "company name", "department"),
    "consent": ("consent", "opt in", "opt-in", "subscribed", "send updates"),
    "group": ("group", "category", "recipient type", "type", "audience", "client / employee"),
}
GROUP_CLIENT, GROUP_EMPLOYEE = "Client", "Employee"
GROUPS = (GROUP_CLIENT, GROUP_EMPLOYEE)
_EMPLOYEE_WORDS = {"employee", "employees", "staff", "internal", "team", "colleague", "emp", "e", "litigation team"}
_CLIENT_WORDS = {"", "client", "clients", "customer", "external", "c"}


def normalise_group(value: str) -> str | None:
    """'Staff' -> 'Employee', '' -> 'Client'; None if not recognised."""
    v = _norm(value)
    if v in _EMPLOYEE_WORDS:
        return GROUP_EMPLOYEE
    if v in _CLIENT_WORDS:
        return GROUP_CLIENT
    return None
_NO_CONSENT = {"no", "n", "false", "0", "opted out", "opt out", "opt-out", "unsubscribed"}
MAX_RECIPIENTS = 1000


def is_valid_email(value: str) -> bool:
    value = (value or "").strip()
    return len(value) <= 254 and bool(_EMAIL_RE.match(value))


@dataclass
class Recipient:
    name: str
    email: str
    salutation: str = ""
    company: str = ""
    row: int = 0
    selected: bool = True
    group: str = GROUP_CLIENT


@dataclass
class RecipientImport:
    recipients: list[Recipient] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)      # rows rejected
    skipped: list[str] = field(default_factory=list)     # opted out / duplicates


class RecipientFileError(Exception):
    pass


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower().replace("_", " "))


def _map_headers(headers: list) -> dict[str, int]:
    mapping: dict[str, int] = {}
    normalised = [_norm(h) for h in headers]
    for key, aliases in HEADER_ALIASES.items():
        for idx, h in enumerate(normalised):
            if h in aliases and key not in mapping:
                mapping[key] = idx
    return mapping


def _read_rows(path: Path) -> list[list]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise RecipientFileError(f"Could not open the Excel file. Make sure it is a valid .xlsx file "
                                     f"and is not password-protected. ({exc.__class__.__name__})") from exc
        try:
            ws = wb.worksheets[0]
            return [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    if suffix == ".csv":
        for enc in ("utf-8-sig", "cp1252"):
            try:
                with path.open(newline="", encoding=enc) as fh:
                    return [row for row in csv.reader(fh)]
            except UnicodeDecodeError:
                continue
        raise RecipientFileError("Could not read the CSV file's text encoding. Save it as 'CSV UTF-8'.")
    if suffix == ".xls":
        raise RecipientFileError("Old .xls files are not supported. In Excel, use File > Save As > "
                                 "'Excel Workbook (.xlsx)' and load that file.")
    raise RecipientFileError("Please choose an Excel (.xlsx) or CSV file.")


def load_recipients(path: str | Path, default_group: str = "Client") -> RecipientImport:
    """Load a recipient list. ``default_group`` applies to rows without a Group column / value
    (a client list loads as Client, an employee list as Employee)."""
    path = Path(path)
    if not path.is_file():
        raise RecipientFileError("The selected file could not be found.")
    rows = _read_rows(path)
    # The header is the first row that contains a recognisable name + email column.
    header_idx, mapping = None, {}
    for i, row in enumerate(rows[:10]):
        mapping = _map_headers(row)
        if "name" in mapping and "email" in mapping:
            header_idx = i
            break
    if header_idx is None:
        raise RecipientFileError("Could not find the 'Name' (or 'Client Name') and 'Email ID' columns. "
                                 "Use the template (Download template) and keep its header row.")

    result = RecipientImport()
    seen: set[str] = set()

    def cell(row, key):
        idx = mapping.get(key)
        if idx is None or idx >= len(row) or row[idx] is None:
            return ""
        return str(row[idx]).strip()

    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(str(c).strip() for c in row if c is not None):
            continue  # blank row
        name, email = cell(row, "name"), cell(row, "email").replace("mailto:", "")
        if not name and not email:
            continue
        if not email:
            result.errors.append(f"Row {offset}: email is missing ({name or 'no name'}).")
            continue
        if not is_valid_email(email):
            result.errors.append(f"Row {offset}: '{email}' is not a valid email address.")
            continue
        if not name:
            result.errors.append(f"Row {offset}: name is missing ({email}).")
            continue
        group = normalise_group(cell(row, "group")) if cell(row, "group") else default_group
        if group is None:
            result.errors.append(f"Row {offset}: group '{cell(row, 'group')}' not recognised - use Client or "
                                 f"Employee ({email}).")
            continue
        if _norm(cell(row, "consent")) in _NO_CONSENT:
            result.skipped.append(f"Row {offset}: {name} has opted out (Consent = No).")
            continue
        key = email.lower()
        if key in seen:
            result.skipped.append(f"Row {offset}: duplicate email {email} ignored.")
            continue
        seen.add(key)
        result.recipients.append(Recipient(name=name, email=email, salutation=cell(row, "salutation"),
                                           company=cell(row, "company"), row=offset, group=group))
        if len(result.recipients) >= MAX_RECIPIENTS:
            result.skipped.append(f"Only the first {MAX_RECIPIENTS} recipients were loaded.")
            break
    if not result.recipients and not result.errors:
        raise RecipientFileError("The file has no recipient rows under the header.")
    return result


HEADER_FILL, HEADER_FONT_COLOR, BORDER_COLOR, BAND_FILL = "0B1F3A", "FFFFFF", "9AA3B5", "F3F5FA"


def style_table(ws, n_cols: int, data_rows: int, widths=(), example_rows: int = 0, wrap_cols=()) -> None:
    """Table look for an Excel sheet: dark bold header, thin borders on every cell, light banding,
    column widths, frozen header and filter buttons. Used for the recipient templates and the audit log."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor=HEADER_FILL)
    band = PatternFill("solid", fgColor=BAND_FILL)
    for c in ws[1][:n_cols]:
        c.font = Font(bold=True, color=HEADER_FONT_COLOR, size=11)
        c.fill = head_fill
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 24
    for r in range(2, data_rows + 2):
        for col in range(1, n_cols + 1):
            cell = ws.cell(r, col)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=col in wrap_cols)
            if r % 2 == 1:
                cell.fill = band
            if r <= example_rows + 1:
                cell.font = Font(italic=True, color="6B7285")
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, n_cols).column_letter}{data_rows + 1}"


def create_template(path: str | Path, group: str | None = None) -> Path:
    """Write a formatted Excel template with example (fictitious) rows.

    ``group`` = "Client" or "Employee" gives a list for that group only (no Group column);
    None gives a combined list with a Group column."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)
    wb = Workbook()
    ws = wb.active
    if group == GROUP_EMPLOYEE:
        ws.title = "Employees"
        headers = ["Employee Name", "Email ID", "Salutation", "Department", "Consent"]
        rows = [["Example Colleague", "colleague@example.com", "Mr", "Litigation", "Yes"],
                ["Second Colleague", "colleague2@example.com", "Ms", "Direct Tax", "Yes"]]
    elif group == GROUP_CLIENT:
        ws.title = "Clients"
        headers = ["Client Name", "Email ID", "Salutation", "Company", "Consent"]
        rows = [["Example Client", "client@example.com", "Mr", "Example Pvt Ltd", "Yes"],
                ["Second Example", "second@example.com", "Ms", "Sample Co", "Yes"]]
    else:
        ws.title = "Recipients"
        headers = ["Name", "Email ID", "Group", "Salutation", "Company", "Consent"]
        rows = [["Example Client", "client@example.com", "Client", "Mr", "Example Pvt Ltd", "Yes"],
                ["Second Example", "second@example.com", "Client", "Ms", "Sample Co", "Yes"],
                ["Example Colleague", "colleague@example.com", "Employee", "", "Our firm", "Yes"]]
    ws.append(headers)
    for r in rows:
        ws.append(r)
    style_table(ws, len(headers), data_rows=max(50, len(rows)),
                widths=(28, 34, 12, 30, 12, 10)[:len(headers)], example_rows=len(rows))
    consent_col = ws.cell(1, len(headers)).column_letter
    consent_dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(consent_dv)
    consent_dv.add(f"{consent_col}2:{consent_col}1000")
    if group is None:
        group_dv = DataValidation(type="list", formula1='"Client,Employee"', allow_blank=True)
        ws.add_data_validation(group_dv)
        group_dv.add("C2:C1000")
    ws.freeze_panes = "A2"
    notes = wb.create_sheet("Instructions")
    who = {GROUP_CLIENT: "client", GROUP_EMPLOYEE: "employee"}.get(group, "recipient")
    for line in (
        f"Fill in one {who} per row on the first sheet. Replace the example rows.",
        "Required: name and Email ID. Optional: Salutation (Mr/Ms/Dr) and the other columns.",
        "Consent: 'No' means the person will NOT receive emails. Blank or 'Yes' = may receive.",
        "Keep the header row unchanged. Save as .xlsx. Use only contact details you are permitted to use.",
    ):
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    notes.insert_rows(1)
    notes["A1"] = "How to fill this list"
    notes["A1"].font = Font(bold=True, size=13, color="0B1F3A")
    wb.save(path)
    return path


# ---------------------------------------------------------- saved lists ---
class RecipientStore:
    """Saved client and employee lists, so they need not be loaded from Excel every time.

    Kept in a small SQLite file in this user's app data folder (recipient_lists.db) - deliberately
    separate from the precedent register, which may sit in a shared folder: contact details stay on the
    computer of the person who sends the emails. Loading a new Excel list replaces the saved list for
    that group; 'Forget' deletes it. Only people who passed the import checks (valid address, consent not
    'No', no duplicates) are saved."""

    def __init__(self, path: str | Path | None = None):
        import sqlite3

        if path is None:
            from .config import user_config_dir
            path = user_config_dir() / "recipient_lists.db"
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._conn = sqlite3.connect(str(self.path), timeout=10)
        self._conn.executescript(
            "CREATE TABLE IF NOT EXISTS lists (grp TEXT PRIMARY KEY, file_name TEXT, saved_at TEXT, cnt INTEGER);"
            "CREATE TABLE IF NOT EXISTS recipients (grp TEXT, pos INTEGER, name TEXT, email TEXT, salutation TEXT,"
            " company TEXT);")
        self._conn.commit()

    def close(self) -> None:
        self._conn.close()

    def save(self, group: str, recipients: list[Recipient], file_name: str = "") -> None:
        from datetime import datetime

        with self._conn:
            self._conn.execute("DELETE FROM recipients WHERE grp = ?", (group,))
            self._conn.executemany(
                "INSERT INTO recipients (grp, pos, name, email, salutation, company) VALUES (?, ?, ?, ?, ?, ?)",
                [(group, i, r.name, r.email, r.salutation, r.company) for i, r in enumerate(recipients)])
            self._conn.execute("INSERT OR REPLACE INTO lists (grp, file_name, saved_at, cnt) VALUES (?, ?, ?, ?)",
                               (group, file_name, datetime.now().strftime("%d-%m-%Y %H:%M"), len(recipients)))

    def load(self, group: str) -> tuple[list[Recipient], dict]:
        info = self._conn.execute("SELECT file_name, saved_at, cnt FROM lists WHERE grp = ?", (group,)).fetchone()
        if not info:
            return [], {}
        rows = self._conn.execute("SELECT name, email, salutation, company FROM recipients WHERE grp = ? "
                                  "ORDER BY pos", (group,)).fetchall()
        people = [Recipient(name=n, email=e, salutation=s or "", company=c or "", row=i + 1, group=group)
                  for i, (n, e, s, c) in enumerate(rows)]
        return people, {"file_name": info[0] or "", "saved_at": info[1] or "", "count": info[2] or 0}

    def forget(self, group: str) -> None:
        with self._conn:
            self._conn.execute("DELETE FROM recipients WHERE grp = ?", (group,))
            self._conn.execute("DELETE FROM lists WHERE grp = ?", (group,))
