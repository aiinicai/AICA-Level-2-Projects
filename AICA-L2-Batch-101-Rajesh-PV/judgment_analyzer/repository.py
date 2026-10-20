"""Precedent register: a searchable SQLite database of analysed judgments.

Each approved analysis is stored with its key facts (court, date, assessment
year, sections, issues, outcome, in whose favour), the full report, the Source
Check and Consistency Check results, and a reviewer sign-off. Only entries
marked "Verified" should be relied on by the team.

Design notes
* SQLite needs no server: the register is a single file. By default it lives in
  the app data folder. A team can point every PC to the same file in a shared
  (OneDrive / SharePoint-synced / network) folder; writes are short transactions
  with a lock timeout, but heavy simultaneous editing from many PCs is not what
  SQLite on a file share is designed for - one person (the reviewer) should do
  most of the writing.
* The same PDF is recognised by its SHA-256 fingerprint, so a judgment that has
  already been analysed and verified is not silently added twice.
* A copy of the PDF can be stored beside the database, so a team member who
  reads a summary email can open the original judgment from the register.
"""

from __future__ import annotations

import hashlib
import json
import logging
import shutil
import sqlite3
from dataclasses import dataclass, field, fields
from datetime import datetime
from pathlib import Path

log = logging.getLogger(__name__)

# Two-level sign-off (maker / checker):
#  Draft     - approved by the preparer in "Review & Approve" and saved; not yet checked by a reviewer
#  Returned  - the reviewer sent it back for correction (reason recorded); the preparer corrects it with
#              "Update details…", which puts it back to Draft for the reviewer
#  Verified  - checked by a reviewer (ideally a different, senior person) - the team may rely on it
STATUS_DRAFT, STATUS_RETURNED, STATUS_VERIFIED = "Draft", "Returned", "Verified"
STATUSES = (STATUS_DRAFT, STATUS_RETURNED, STATUS_VERIFIED)
# Fields that can be corrected later without a new review (a citation found after the order is reported, tags)
MINOR_FIELDS = ("citation", "issue_tags")
SCHEMA_VERSION = 2
ACT_EITHER, ACT_1961, ACT_2025 = "Either Act", "1961 Act", "2025 Act"

_SCHEMA = """
CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS judgments (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ref TEXT UNIQUE NOT NULL,
    pdf_sha256 TEXT,
    file_name TEXT, pdf_copy TEXT,
    court TEXT, bench TEXT, case_title TEXT, citation TEXT, case_number TEXT,
    judgment_date TEXT, assessment_year TEXT, appellant TEXT, respondent TEXT,
    sections TEXT, issues TEXT, issue_tags TEXT,
    outcome TEXT, in_favour_of TEXT, ratio TEXT,
    executive_summary TEXT, report_md TEXT,
    source_check_score REAL, source_check_summary TEXT, consistency_summary TEXT,
    model TEXT, analysed_at TEXT,
    status TEXT NOT NULL DEFAULT 'Draft', reviewed_by TEXT, reviewed_at TEXT, review_notes TEXT,
    created_by TEXT, created_at TEXT, updated_at TEXT
);
CREATE INDEX IF NOT EXISTS ix_judgments_sha ON judgments(pdf_sha256);
CREATE INDEX IF NOT EXISTS ix_judgments_status ON judgments(status);
"""


class RegisterError(Exception):
    pass


@dataclass
class JudgmentRecord:
    id: int | None = None
    ref: str = ""
    pdf_sha256: str = ""
    file_name: str = ""
    pdf_copy: str = ""
    court: str = ""
    bench: str = ""
    case_title: str = ""
    citation: str = ""
    case_number: str = ""
    judgment_date: str = ""
    assessment_year: str = ""
    appellant: str = ""
    respondent: str = ""
    sections: list[str] = field(default_factory=list)        # Income-tax Act, 1961
    sections_2025: list[str] = field(default_factory=list)   # Income-tax Act, 2025
    issues: list[str] = field(default_factory=list)
    issue_tags: list[str] = field(default_factory=list)
    outcome: str = ""
    in_favour_of: str = ""
    ratio: str = ""
    executive_summary: str = ""
    report_md: str = ""
    source_check_score: float | None = None
    source_check_summary: str = ""
    consistency_summary: str = ""
    model: str = ""
    analysed_at: str = ""
    status: str = STATUS_DRAFT
    reviewed_by: str = ""
    reviewed_at: str = ""
    review_notes: str = ""
    created_by: str = ""
    created_at: str = ""
    updated_at: str = ""
    citation_updated_by: str = ""
    citation_updated_at: str = ""
    updated_by: str = ""
    returned_by: str = ""
    returned_at: str = ""
    return_reason: str = ""
    return_fields: list[str] = field(default_factory=list)   # what the reviewer asked to correct (labels)

    @property
    def sections_display(self) -> str:
        parts = []
        if self.sections:
            parts.append(", ".join(self.sections) + (" (1961)" if self.sections_2025 else ""))
        if self.sections_2025:
            parts.append(", ".join(self.sections_2025) + " (2025)")
        return "; ".join(parts)

    @property
    def label(self) -> str:
        bits = [self.case_title or self.file_name, self.citation, self.judgment_date]
        return " - ".join(b for b in bits if b and b.strip().lower() != "not stated")


_LIST_FIELDS = ("sections", "sections_2025", "issues", "issue_tags", "return_fields")
_ADDED_COLUMNS = {"sections_2025": "TEXT", "citation_updated_by": "TEXT", "citation_updated_at": "TEXT",
                  "updated_by": "TEXT", "returned_by": "TEXT", "returned_at": "TEXT", "return_reason": "TEXT",
                  "return_fields": "TEXT"}
_COLUMNS = [f.name for f in fields(JudgmentRecord) if f.name != "id"]


def file_sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


class Repository:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            self._conn = sqlite3.connect(str(self.path), timeout=15, check_same_thread=False)
            self._conn.row_factory = sqlite3.Row
            self._conn.execute("PRAGMA busy_timeout = 15000")
            self._conn.executescript(_SCHEMA)
            have = {r[1] for r in self._conn.execute("PRAGMA table_info(judgments)")}
            for col, typ in _ADDED_COLUMNS.items():          # upgrade a register created by an older version
                if col not in have:
                    self._conn.execute(f"ALTER TABLE judgments ADD COLUMN {col} {typ}")
            self._conn.execute("UPDATE meta SET value = ? WHERE key = 'schema_version'", (str(SCHEMA_VERSION),))
            self._conn.execute("INSERT OR IGNORE INTO meta(key, value) VALUES ('schema_version', ?)",
                               (str(SCHEMA_VERSION),))
            self._conn.commit()
        except sqlite3.DatabaseError as exc:
            raise RegisterError(f"The register file could not be opened ({exc}). If it is in a shared folder, "
                                "check that the folder is available and the file is not damaged.") from exc

    # ------------------------------------------------------------ housekeeping
    def close(self) -> None:
        self._conn.close()

    @property
    def pdf_folder(self) -> Path:
        return self.path.parent / "register_judgments"

    def _next_ref(self) -> str:
        year = datetime.now().year
        prefix = f"JIA-{year}-"
        row = self._conn.execute("SELECT ref FROM judgments WHERE ref LIKE ? ORDER BY ref DESC LIMIT 1",
                                 (prefix + "%",)).fetchone()
        n = int(row["ref"].rsplit("-", 1)[1]) + 1 if row else 1
        return f"{prefix}{n:04d}"

    @staticmethod
    def _to_row(rec: JudgmentRecord) -> dict:
        data = {c: getattr(rec, c) for c in _COLUMNS}
        for k in _LIST_FIELDS:
            data[k] = json.dumps(list(data[k] or []), ensure_ascii=False)
        return data

    @staticmethod
    def _from_row(row: sqlite3.Row) -> JudgmentRecord:
        data = dict(row)
        for k in _LIST_FIELDS:
            try:
                data[k] = json.loads(data.get(k) or "[]")
            except ValueError:
                data[k] = [x.strip() for x in str(data.get(k) or "").split(",") if x.strip()]
        known = {f.name for f in fields(JudgmentRecord)}
        return JudgmentRecord(**{k: ("" if v is None and k != "source_check_score" else v)
                                 for k, v in data.items() if k in known})

    # ------------------------------------------------------------------ write
    def add(self, rec: JudgmentRecord, pdf_path: str | Path | None = None, copy_pdf: bool = True) -> JudgmentRecord:
        rec.ref = rec.ref or self._next_ref()
        rec.created_at = rec.created_at or _now()
        rec.updated_at = _now()
        if pdf_path and copy_pdf:
            try:
                self.pdf_folder.mkdir(parents=True, exist_ok=True)
                dest = self.pdf_folder / f"{rec.ref}.pdf"
                shutil.copy2(pdf_path, dest)
                rec.pdf_copy = dest.name
            except OSError as exc:
                log.warning("Could not copy the PDF into the register folder: %s", exc)
        data = self._to_row(rec)
        cols = ", ".join(data)
        marks = ", ".join("?" for _ in data)
        with self._conn:
            cur = self._conn.execute(f"INSERT INTO judgments ({cols}) VALUES ({marks})", list(data.values()))
        rec.id = cur.lastrowid
        log.info("Register: added %s", rec.ref)
        return rec

    def update(self, rec: JudgmentRecord) -> None:
        if rec.id is None:
            raise RegisterError("Record has no id.")
        rec.updated_at = _now()
        data = self._to_row(rec)
        sets = ", ".join(f"{k} = ?" for k in data)
        with self._conn:
            self._conn.execute(f"UPDATE judgments SET {sets} WHERE id = ?", [*data.values(), rec.id])

    def update_details(self, rec: JudgmentRecord, updated_by: str) -> str:
        """Save corrections made later (key facts, summary text, citation, tags). Returns the new status.

        * Returned entry: saving the correction resubmits it - it goes back to Draft for the reviewer.
        * Verified entry: correcting anything other than the citation or tags changes checked content,
          so it goes back to Draft; citation / tags only keep it Verified."""
        if rec.id is None:
            raise RegisterError("Record has no id.")
        if not updated_by.strip():
            raise RegisterError("Enter your name or initials.")
        old = self.get(rec.id)
        if old is None:
            raise RegisterError("This entry no longer exists.")
        changed = [c for c in _COLUMNS if getattr(old, c) != getattr(rec, c)]
        if "citation" in changed:
            rec.citation_updated_by, rec.citation_updated_at = updated_by.strip(), _now()
        major = [c for c in changed if c not in MINOR_FIELDS and not c.startswith(("citation_updated", "updated"))]
        if old.status == STATUS_RETURNED:
            # The preparer's correction of a returned entry is a resubmission, whatever was corrected
            # (even only the citation or tags): it goes back to the reviewer as Draft.
            rec.status = STATUS_DRAFT
        elif major and old.status == STATUS_VERIFIED:
            rec.status = STATUS_DRAFT
        else:
            rec.status = old.status
        rec.updated_by = updated_by.strip()
        self.update(rec)
        return rec.status

    def set_returned(self, record_id: int, reviewer: str, reason: str, fields_to_fix: list[str] | None = None) -> None:
        """Reviewer does not accept the entry: send it back for correction, saying what to correct."""
        if not reviewer.strip():
            raise RegisterError("Enter the reviewer's name or initials.")
        if not reason.strip() and not fields_to_fix:
            raise RegisterError("Say what needs to be corrected.")
        with self._conn:
            self._conn.execute("UPDATE judgments SET status = ?, returned_by = ?, returned_at = ?, return_reason = ?, "
                               "return_fields = ?, updated_at = ? WHERE id = ?",
                               (STATUS_RETURNED, reviewer.strip(), _now(), reason.strip(),
                                json.dumps(list(fields_to_fix or []), ensure_ascii=False), _now(), record_id))

    def set_verified(self, record_id: int, reviewer: str, notes: str = "") -> None:
        if not reviewer.strip():
            raise RegisterError("Enter the reviewer's name or initials.")
        with self._conn:
            self._conn.execute("UPDATE judgments SET status = ?, reviewed_by = ?, reviewed_at = ?, review_notes = ?, "
                               "updated_at = ? WHERE id = ?",
                               (STATUS_VERIFIED, reviewer.strip(), _now(), notes.strip(), _now(), record_id))

    def update_citation(self, record_id: int, citation: str, updated_by: str) -> None:
        """Record a reported citation found later (judgments are often unreported when first summarised)."""
        if not citation.strip():
            raise RegisterError("Enter the citation.")
        if not updated_by.strip():
            raise RegisterError("Enter your name or initials.")
        with self._conn:
            self._conn.execute("UPDATE judgments SET citation = ?, citation_updated_by = ?, citation_updated_at = ?, "
                               "updated_at = ? WHERE id = ?",
                               (citation.strip(), updated_by.strip(), _now(), _now(), record_id))

    def set_draft(self, record_id: int) -> None:
        with self._conn:
            self._conn.execute("UPDATE judgments SET status = ?, updated_at = ? WHERE id = ?",
                               (STATUS_DRAFT, _now(), record_id))

    def delete(self, record_id: int) -> None:
        rec = self.get(record_id)
        with self._conn:
            self._conn.execute("DELETE FROM judgments WHERE id = ?", (record_id,))
        if rec and rec.pdf_copy:
            try:
                (self.pdf_folder / rec.pdf_copy).unlink(missing_ok=True)
            except OSError:
                pass

    # ------------------------------------------------------------------- read
    def get(self, record_id: int) -> JudgmentRecord | None:
        row = self._conn.execute("SELECT * FROM judgments WHERE id = ?", (record_id,)).fetchone()
        return self._from_row(row) if row else None

    def get_by_ref(self, ref: str) -> JudgmentRecord | None:
        row = self._conn.execute("SELECT * FROM judgments WHERE ref = ?", (ref,)).fetchone()
        return self._from_row(row) if row else None

    def find_by_sha(self, sha: str) -> list[JudgmentRecord]:
        rows = self._conn.execute("SELECT * FROM judgments WHERE pdf_sha256 = ? ORDER BY id", (sha,)).fetchall()
        return [self._from_row(r) for r in rows]

    def search(self, text: str = "", section: str = "", status: str = "", favour: str = "",
               limit: int = 1000, act: str = ACT_EITHER) -> list[JudgmentRecord]:
        where, args = [], []
        for word in (text or "").split():
            like = f"%{word}%"
            where.append("(case_title LIKE ? OR citation LIKE ? OR court LIKE ? OR issues LIKE ? OR issue_tags LIKE ? "
                         "OR ratio LIKE ? OR appellant LIKE ? OR respondent LIKE ? OR case_number LIKE ? "
                         "OR assessment_year LIKE ? OR ref LIKE ?)")
            args += [like] * 11
        if section.strip():
            # Sections are stored as a JSON list: match an exact element, e.g. "14A" but not "114A".
            cols = {ACT_1961: ["sections"], ACT_2025: ["sections_2025"]}.get(act, ["sections", "sections_2025"])
            s = section.strip().replace('"', "")
            where.append("(" + " OR ".join(f"{c} LIKE ? OR {c} LIKE ?" for c in cols) + ")")
            for _c in cols:
                args += [f'%"{s}"%', f'%"{s}(%']
        if status:
            where.append("status = ?")
            args.append(status)
        if favour:
            where.append("in_favour_of = ?")
            args.append(favour)
        sql = "SELECT * FROM judgments" + (" WHERE " + " AND ".join(where) if where else "") + \
              " ORDER BY id DESC LIMIT ?"
        rows = self._conn.execute(sql, [*args, limit]).fetchall()
        return [self._from_row(r) for r in rows]

    def stats(self) -> dict:
        c = self._conn
        total = c.execute("SELECT COUNT(*) FROM judgments").fetchone()[0]
        verified = c.execute("SELECT COUNT(*) FROM judgments WHERE status = ?", (STATUS_VERIFIED,)).fetchone()[0]
        returned = c.execute("SELECT COUNT(*) FROM judgments WHERE status = ?", (STATUS_RETURNED,)).fetchone()[0]
        favour = {r[0] or "Not stated": r[1] for r in
                  c.execute("SELECT in_favour_of, COUNT(*) FROM judgments GROUP BY in_favour_of")}
        counts: dict[str, int] = {}
        for s61, s25 in c.execute("SELECT sections, sections_2025 FROM judgments"):
            for raw, suffix in ((s61, ""), (s25, " (2025)")):
                try:
                    for s in json.loads(raw or "[]"):
                        counts[s + suffix] = counts.get(s + suffix, 0) + 1
                except ValueError:
                    pass
        top = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))[:8]
        return {"total": total, "verified": verified, "returned": returned, "draft": total - verified - returned,
                "favour": favour,
                "top_sections": top}


# ----------------------------------------------------------------- export ---
EXPORT_COLUMNS = [
    ("Ref", "ref"), ("Status", "status"), ("Approved by (preparer)", "created_by"), ("Reviewed by", "reviewed_by"),
    ("Reviewed at", "reviewed_at"), ("Returned - reason", "return_reason"),
    ("Case title", "case_title"), ("Court / tribunal", "court"), ("Bench", "bench"), ("Citation", "citation"),
    ("Case number", "case_number"), ("Date of judgment", "judgment_date"), ("Assessment year", "assessment_year"),
    ("Appellant", "appellant"), ("Respondent", "respondent"), ("Sections - 1961 Act", "sections"),
    ("Sections - 2025 Act", "sections_2025"), ("Issues", "issues"),
    ("Tags", "issue_tags"), ("Outcome", "outcome"), ("In favour of", "in_favour_of"), ("Ratio", "ratio"),
    ("Source check score", "source_check_score"), ("Consistency", "consistency_summary"),
    ("PDF copy", "pdf_copy"), ("Model", "model"), ("Analysed at", "analysed_at"),
]


def export_xlsx(records: list[JudgmentRecord], path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Precedent register"
    ws.append([c for c, _ in EXPORT_COLUMNS])
    for rec in records:
        row = []
        for _, attr in EXPORT_COLUMNS:
            v = getattr(rec, attr)
            if isinstance(v, list):
                v = "; ".join(v)
            if attr == "source_check_score" and v is not None and v != "":
                v = round(float(v), 2)
            row.append(v)
        ws.append(row)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B1F3A")
    widths = {"Case title": 40, "Issues": 60, "Ratio": 60, "Court / tribunal": 30, "Citation": 34}
    for i, (name, _) in enumerate(EXPORT_COLUMNS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(name, 16)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    score_col = [a for _, a in EXPORT_COLUMNS].index("source_check_score") + 1
    for row in ws.iter_rows(min_row=2, min_col=score_col, max_col=score_col):
        for c in row:
            c.number_format = "0%"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return path


def record_from_result(result, facts: dict, issue_tags: list[str], created_by: str = "") -> JudgmentRecord:
    """Build a register record from an AnalysisResult and the reviewer-confirmed key facts."""
    report = result.report
    v = getattr(result, "verification", None)
    c = getattr(result, "consistency", None)
    info = result.document.info
    try:
        sha = file_sha256(info.path)
    except OSError:
        sha = ""
    return JudgmentRecord(
        pdf_sha256=sha, file_name=info.file_name,
        court=facts.get("court", ""), bench=facts.get("bench", ""), case_title=facts.get("case_title", ""),
        citation=facts.get("citation", ""), case_number=facts.get("case_number", ""),
        judgment_date=facts.get("judgment_date", ""), assessment_year=facts.get("assessment_year", ""),
        appellant=facts.get("appellant", ""), respondent=facts.get("respondent", ""),
        sections=list(facts.get("sections") or []), sections_2025=list(facts.get("sections_2025") or []),
        issues=list(facts.get("issues") or []),
        issue_tags=[t.strip() for t in issue_tags if t.strip()],
        outcome=facts.get("outcome", ""), in_favour_of=facts.get("in_favour_of", ""), ratio=facts.get("ratio", ""),
        executive_summary=report.section_body(2), report_md=report.to_markdown(),
        source_check_score=v.score if v else None, source_check_summary=v.summary_line() if v else "",
        consistency_summary=c.summary_line() if c else "Consistency check not run.",
        model=result.model, analysed_at=_now(), created_by=created_by,
    )
