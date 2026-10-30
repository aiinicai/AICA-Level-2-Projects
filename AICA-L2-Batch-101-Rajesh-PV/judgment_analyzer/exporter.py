"""Prepare the analysis for export: plain text, Word (.docx) and PDF.

A small Markdown subset (headings, bullets, numbered lists, quotes, bold and
italic) is converted into a neutral block model that every exporter - and the
GUI renderer - shares.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from .prompts import DISCLAIMER_TITLE

log = logging.getLogger(__name__)

# ----------------------------------------------------------- block model ---
_BULLET_RE = re.compile(r"^(\s*)[-*+•]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^(\s*)(\d{1,3})[.)]\s+(.*)$")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*$")
_INLINE_RE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|__.+?__|(?<![\w*])\*(?!\s).+?(?<!\s)\*(?![\w*])|`.+?`)")


@dataclass
class Block:
    kind: str  # heading | bullet | numbered | quote | paragraph | blank
    text: str = ""
    level: int = 0  # heading level, or list indent level
    number: str = ""  # for numbered items


def parse_blocks(markdown: str) -> list[Block]:
    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush():
        if paragraph:
            blocks.append(Block("paragraph", " ".join(s.strip() for s in paragraph)))
            paragraph.clear()

    for raw in (markdown or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush()
            if blocks and blocks[-1].kind != "blank":
                blocks.append(Block("blank"))
            continue
        if line.strip().startswith("```") or re.fullmatch(r"\s*([-*_])\1{2,}\s*", line):
            flush()  # drop code fences / horizontal rules
            continue
        if m := _HEADING_RE.match(line.strip()):
            flush()
            blocks.append(Block("heading", m.group(2).strip(), len(m.group(1))))
        elif m := _BULLET_RE.match(line):
            flush()
            blocks.append(Block("bullet", m.group(2).strip(), len(m.group(1).expandtabs(4)) // 2))
        elif m := _NUMBERED_RE.match(line):
            flush()
            blocks.append(Block("numbered", m.group(3).strip(), len(m.group(1).expandtabs(4)) // 2, m.group(2)))
        elif line.lstrip().startswith(">"):
            flush()
            blocks.append(Block("quote", line.lstrip()[1:].strip()))
        else:
            paragraph.append(line)
    flush()
    while blocks and blocks[-1].kind == "blank":
        blocks.pop()
    return blocks


def split_inline(text: str) -> list[tuple[str, bool, bool]]:
    """Split text into runs of (text, bold, italic)."""
    runs: list[tuple[str, bool, bool]] = []
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], False, False))
        token = m.group(0)
        if token.startswith("***"):
            runs.append((token[3:-3], True, True))
        elif token.startswith(("**", "__")):
            runs.append((token[2:-2], True, False))
        elif token.startswith("`"):
            runs.append((token[1:-1], False, False))
        else:
            runs.append((token[1:-1], False, True))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], False, False))
    return [r for r in runs if r[0]]


def strip_inline(text: str) -> str:
    return "".join(run[0] for run in split_inline(text))


def markdown_to_plain(markdown: str) -> str:
    """Readable plain text (used for copy-to-clipboard and .txt export)."""
    out: list[str] = []
    for b in parse_blocks(markdown):
        text = strip_inline(b.text)
        if b.kind == "heading":
            if out and out[-1] != "":
                out.append("")
            if b.level <= 2:
                out += [text.upper(), "=" * min(len(text), 80)]
            else:
                out += [text, "-" * min(len(text), 80)]
        elif b.kind == "bullet":
            out.append(f"{'  ' * (b.level + 1)}• {text}")
        elif b.kind == "numbered":
            out.append(f"{'  ' * (b.level + 1)}{b.number}. {text}")
        elif b.kind == "quote":
            out.append(f"    \"{text}\"")
        elif b.kind == "paragraph":
            out.append(text)
        elif b.kind == "blank":
            out.append("")
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n"


# --------------------------------------------------------- export model ---
@dataclass
class ExportDocument:
    title: str
    metadata: list[tuple[str, str]] = field(default_factory=list)
    sections: list[tuple[str, str]] = field(default_factory=list)  # (heading, markdown body)
    notices: list[str] = field(default_factory=list)
    author: str = ""

    def to_markdown(self) -> str:
        lines = [f"# {self.title}", ""]
        lines += [f"- **{k}:** {v}" for k, v in self.metadata]
        if self.notices:
            lines += ["", "**Processing notes:**"] + [f"- {n}" for n in self.notices]
        for heading, body in self.sections:
            lines += ["", f"## {heading}", "", body.strip()]
        return "\n".join(lines).strip() + "\n"

    def to_plain_text(self) -> str:
        return markdown_to_plain(self.to_markdown())


def build_export_document(
    sections: list[tuple[str, str]],
    file_name: str,
    page_count: int | None = None,
    model: str = "",
    mode: str = "",
    notices: list[str] | None = None,
    generated_at: datetime | None = None,
) -> ExportDocument:
    """Assemble an export from ordered sections plus processing metadata."""
    if not sections:
        raise ValueError("There is no analysis to export.")
    generated_at = generated_at or datetime.now()
    meta = [("Source document", file_name or "Unknown")]
    if page_count:
        meta.append(("Pages", str(page_count)))
    if model:
        meta.append(("AI model", model))
    if mode:
        meta.append(("Processing mode", mode))
    meta.append(("Generated", generated_at.strftime("%d %b %Y %H:%M")))
    if not any(h.strip().lower() == DISCLAIMER_TITLE.lower() for h, _ in sections):
        raise ValueError("Export must include the disclaimer section.")
    return ExportDocument(
        title="Judgement Insights Analysis", metadata=meta, sections=list(sections), notices=list(notices or [])
    )


def safe_filename(stem: str, suffix: str) -> str:
    stem = re.sub(r"[^\w\- ]+", "_", stem or "analysis").strip(" _") or "analysis"
    return f"{stem[:80]}_analysis{suffix}"


# ------------------------------------------------------------- writers ---
def save_text(doc: ExportDocument, path: str | Path) -> Path:
    path = Path(path)
    path.write_text(doc.to_plain_text(), encoding="utf-8")
    log.info("Saved text export: %s", path.name)
    return path


def save_docx(doc: ExportDocument, path: str | Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    navy, purple, gold = RGBColor(0x0B, 0x1F, 0x3A), RGBColor(0x4B, 0x2E, 0x83), RGBColor(0x9C, 0x7A, 0x12)
    path = Path(path)
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = d.add_heading(doc.title, level=0)
    for run in title.runs:
        run.font.color.rgb = navy
    for key, value in doc.metadata:
        p = d.add_paragraph()
        r = p.add_run(f"{key}: ")
        r.bold = True
        p.add_run(value)
        p.paragraph_format.space_after = Pt(0)
    if doc.notices:
        p = d.add_paragraph()
        p.add_run("Processing notes").bold = True
        for n in doc.notices:
            d.add_paragraph(n, style="List Bullet")

    def add_runs(paragraph, text):
        for content, bold, italic in split_inline(text):
            run = paragraph.add_run(content)
            run.bold, run.italic = bold, italic
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for heading, body in doc.sections:
        h = d.add_heading(heading, level=1)
        for run in h.runs:
            run.font.color.rgb = gold if heading == DISCLAIMER_TITLE else navy
        for b in parse_blocks(body):
            if b.kind == "heading":
                sub = d.add_heading(strip_inline(b.text), level=min(b.level, 4) if b.level > 2 else 2)
                for run in sub.runs:
                    run.font.color.rgb = purple
            elif b.kind in ("bullet", "numbered"):
                style_name = "List Bullet" if b.kind == "bullet" else "List Number"
                if b.level >= 1:
                    style_name += " 2"
                try:
                    p = d.add_paragraph(style=style_name)
                except KeyError:
                    p = d.add_paragraph(style="List Bullet")
                add_runs(p, b.text)
            elif b.kind == "quote":
                p = d.add_paragraph(style="Intense Quote" if "Intense Quote" in d.styles else None)
                add_runs(p, b.text)
            elif b.kind == "paragraph":
                add_runs(d.add_paragraph(), b.text)

    d.save(path)
    log.info("Saved Word export: %s", path.name)
    return path


_FONT_CANDIDATES = [
    # (regular, bold) - first pair found is used so non-ASCII text renders.
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/Library/Fonts/Arial Unicode.ttf", "/Library/Fonts/Arial Unicode.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
]


def _register_pdf_fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for regular, bold in _FONT_CANDIDATES:
        if Path(regular).is_file() and Path(bold).is_file():
            try:
                pdfmetrics.registerFont(TTFont("JIA-Regular", regular))
                pdfmetrics.registerFont(TTFont("JIA-Bold", bold))
                from reportlab.lib.fonts import addMapping

                addMapping("JIA-Regular", 0, 0, "JIA-Regular")
                addMapping("JIA-Regular", 1, 0, "JIA-Bold")
                addMapping("JIA-Regular", 0, 1, "JIA-Regular")
                addMapping("JIA-Regular", 1, 1, "JIA-Bold")
                return "JIA-Regular", "JIA-Bold"
            except Exception as exc:  # pragma: no cover - font file issues
                log.warning("Could not register font %s: %s", regular, exc)
    return "Helvetica", "Helvetica-Bold"


def _rl_markup(text: str) -> str:
    """Convert inline Markdown to ReportLab paragraph markup (escaped)."""
    from xml.sax.saxutils import escape

    out = []
    for content, bold, italic in split_inline(text):
        piece = escape(content)
        if italic:
            piece = f"<i>{piece}</i>"
        if bold:
            piece = f"<b>{piece}</b>"
        out.append(piece)
    return "".join(out)


def save_pdf(doc: ExportDocument, path: str | Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    path = Path(path)
    regular, bold = _register_pdf_fonts()
    navy, purple, gold = colors.HexColor("#0B1F3A"), colors.HexColor("#4B2E83"), colors.HexColor("#9C7A12")
    base = getSampleStyleSheet()
    from reportlab.lib.enums import TA_JUSTIFY

    body = ParagraphStyle("body", parent=base["BodyText"], fontName=regular, fontSize=10, leading=14,
                          alignment=TA_JUSTIFY)
    title_s = ParagraphStyle("title", parent=body, fontName=bold, fontSize=20, leading=24, textColor=navy,
                             spaceAfter=8)
    h1 = ParagraphStyle("h1", parent=body, fontName=bold, fontSize=14, leading=18, textColor=navy,
                        spaceBefore=12, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=body, fontName=bold, fontSize=11.5, leading=15, textColor=purple,
                        spaceBefore=6, spaceAfter=2)
    meta_s = ParagraphStyle("meta", parent=body, fontSize=9, leading=12, textColor=colors.HexColor("#333333"))
    quote_s = ParagraphStyle("quote", parent=body, leftIndent=12, textColor=colors.HexColor("#444444"))

    story = [Paragraph(_rl_markup(doc.title), title_s)]
    for k, v in doc.metadata:
        story.append(Paragraph(f"<b>{_rl_markup(k)}:</b> {_rl_markup(v)}", meta_s))
    for n in doc.notices:
        story.append(Paragraph(f"<i>Note:</i> {_rl_markup(n)}", meta_s))

    def flush_list(items, kind):
        if items:
            story.append(ListFlowable(
                [ListItem(Paragraph(_rl_markup(t), body), leftIndent=12 + 12 * lvl) for t, lvl in items],
                bulletType="1" if kind == "numbered" else "bullet", start="1" if kind == "numbered" else "•",
                leftIndent=14,
                bulletFontName=regular, bulletFontSize=9,
            ))

    for heading, md in doc.sections:
        hstyle = ParagraphStyle("hd", parent=h1, textColor=gold) if heading == DISCLAIMER_TITLE else h1
        story.append(Paragraph(_rl_markup(heading), hstyle))
        pending: list[tuple[str, int]] = []
        pending_kind = ""
        for b in parse_blocks(md):
            if b.kind in ("bullet", "numbered"):
                if pending and b.kind != pending_kind:
                    flush_list(pending, pending_kind)
                    pending = []
                pending.append((b.text, b.level))
                pending_kind = b.kind
                continue
            flush_list(pending, pending_kind)
            pending = []
            if b.kind == "heading":
                story.append(Paragraph(_rl_markup(strip_inline(b.text)), h2))
            elif b.kind == "quote":
                story.append(Paragraph(_rl_markup(b.text), quote_s))
            elif b.kind == "paragraph":
                story.append(Paragraph(_rl_markup(b.text), body))
            elif b.kind == "blank":
                story.append(Spacer(1, 3))
        flush_list(pending, pending_kind)

    def on_page(canvas, document):
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.setFillColor(colors.grey)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {document.page}")
        canvas.restoreState()

    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm,
                      bottomMargin=18 * mm, title=doc.title, author=doc.author).build(
        story, onFirstPage=on_page, onLaterPages=on_page)
    log.info("Saved PDF export: %s", path.name)
    return path


# ------------------------------------------------------- approved output ---
def citation_line(facts: dict) -> str:
    """Reported citation if known; otherwise identify the order by appeal number, forum and date."""
    citation = (facts.get("citation") or "").strip()
    if citation and citation.lower() not in ("not stated", "not reported", "n/a", "-"):
        return citation
    bits = [facts.get("case_number", ""), facts.get("court", ""),
            f"order dated {facts['judgment_date']}" if facts.get("judgment_date") else ""]
    ident = ", ".join(b.strip() for b in bits if b and b.strip() and b.strip().lower() != "not stated")
    return f"Not reported in the source copy{' - ' + ident if ident else ''}"


def build_final_document(report, facts: dict, generated_at: datetime | None = None,
                         approved_by: str = "", status_note: str = "") -> ExportDocument:
    """The approved analysis as it goes to Word / PDF / text: case name and citation first, then the
    report sections with internal labels and page references removed. No AI tags, no processing notes,
    no disclaimer section (the content has been reviewed and approved by the user)."""
    from .email_composer import clean_for_client

    generated_at = generated_at or datetime.now()
    title = (facts.get("case_title") or "").strip() or "Judgment Analysis"
    meta = [("Generated on", generated_at.strftime("%d-%m-%Y")), ("Case", title), ("Citation", citation_line(facts))]
    if status_note:   # e.g. not yet verified by a reviewer: the document says so
        meta.insert(0, ("Status", status_note))
    sections = []
    for heading, body in report.ordered_sections():
        if heading.strip().lower() == DISCLAIMER_TITLE.lower():
            continue
        cleaned = clean_for_client(body)
        if cleaned.strip():
            sections.append((heading, cleaned))
    return ExportDocument(title=title, metadata=meta, sections=sections, author=approved_by)


# --------------------------------------------------------------------- JSON ---
JSON_FACT_KEYS = ("court", "bench", "case_title", "citation", "case_number", "judgment_date", "assessment_year",
                  "appellant", "respondent", "sections", "sections_2025", "issues", "outcome", "in_favour_of",
                  "ratio")


def build_json(doc: ExportDocument, facts: dict, approved_by: str = "", approved_at: str = "",
               register_ref: str = "") -> dict:
    """Structured copy of the approved analysis (for other tools, a knowledge base or later import).
    Same cleaned content as Word / PDF, plus the confirmed key facts."""
    from . import APP_NAME

    meta = dict(doc.metadata)
    key_facts = {k: facts.get(k, [] if k in ("sections", "sections_2025", "issues") else "") for k in JSON_FACT_KEYS}
    key_facts["sections_income_tax_act_1961"] = key_facts.pop("sections")
    key_facts["sections_income_tax_act_2025"] = key_facts.pop("sections_2025")
    return {
        "generated_on": meta.get("Generated on", ""),
        "case": meta.get("Case", doc.title),
        "citation": meta.get("Citation", ""),
        "key_facts": key_facts,
        "sections": [{"heading": h, "text": body} for h, body in doc.sections],
        "approved_by": approved_by,
        "approved_at": approved_at,
        "register_ref": register_ref,
        "app": APP_NAME,
    }


def save_json(data: dict, path: str | Path) -> Path:
    import json

    path = Path(path)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    log.info("Saved JSON export: %s", path.name)
    return path
