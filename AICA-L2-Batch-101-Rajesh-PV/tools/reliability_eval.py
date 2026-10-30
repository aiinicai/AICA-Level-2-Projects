"""Reliability test kit: measure how consistent and how accurate the AI key-fact extraction is.

Answers the question "if the same judgment is analysed again, do we get the same answer, and is it right?"
with numbers from YOUR test set instead of a generic claim.

Usage (from the app folder, with GEMINI_API_KEY set or the key saved in the app):

    # 1. create an answer-key template listing the PDFs in a folder
    python tools/reliability_eval.py --pdf-dir test_judgments --make-gold gold.xlsx

    # 2. fill in gold.xlsx by hand (the correct facts, from reading each judgment)

    # 3. run every judgment N times and score agreement and accuracy
    python tools/reliability_eval.py --pdf-dir test_judgments --runs 3 --gold gold.xlsx --out results.xlsx

    # optional: also run the full analysis once per judgment and record the Source Check score
    python tools/reliability_eval.py --pdf-dir test_judgments --runs 3 --gold gold.xlsx --full-report

Metrics
* Run-to-run agreement: share of key facts on which all N runs agree (by the app's comparison rules).
* Accuracy: share of key facts that match the human answer key (exact match after normalising spacing
  and punctuation; sections compared as sets). Only fields filled in the answer key are scored.
* Source Check score (with --full-report): share of quotations, citations, sections, dates, amounts and
  paragraph/page references in the report that were found in the judgment.

Use public judgments only. Free-tier requests may be used by Google to improve its products.
"""

from __future__ import annotations

import argparse
import os
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from judgment_analyzer.config import AppSettings, effective_temperature  # noqa: E402
from judgment_analyzer.consistency import (AGREE, DIFFERS, FIELDS, compare_runs, extract_key_facts,  # noqa: E402
                                           scope_text, _norm)
from judgment_analyzer.gemini_client import GeminiClient  # noqa: E402
from judgment_analyzer.key_store import load_api_key  # noqa: E402
from judgment_analyzer.pdf_processor import extract_text, validate_pdf  # noqa: E402

SCORED = [k for k, _l, kind in FIELDS if kind not in ("issues", "prose")]  # free text is not auto-scored
LIST_KEYS = [k for k, _l, kind in FIELDS if kind == "list"]


def make_gold(pdf_dir: Path, out: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Answer key"
    ws.append(["File", *SCORED])
    for pdf in sorted(pdf_dir.glob("*.pdf")):
        ws.append([pdf.name] + [""] * len(SCORED))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B1F3A")
    notes = wb.create_sheet("How to fill")
    for line in ("One row per judgment. Fill in the correct value of each fact from reading the judgment.",
                 "sections (1961 Act) and sections_2025 (2025 Act): numbers separated by ';' e.g. 68; 133(6)",
                 "judgment_date: DD-MM-YYYY.  outcome: Allowed / Partly allowed / Dismissed / Remanded / Withdrawn",
                 "in_favour_of: Assessee / Revenue / Partly assessee / Not applicable",
                 "Leave a cell blank to exclude that fact from the accuracy score."):
        notes.append([line])
    wb.save(out)
    print(f"Answer-key template written to {out}")


def load_gold(path: Path) -> dict[str, dict]:
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "").strip() for h in rows[0]]
    out = {}
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        rec = {head[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(row) if i < len(head)}
        for k in LIST_KEYS:
            rec[k] = [s.strip() for s in rec.get(k, "").split(";") if s.strip()]
        out[str(row[0]).strip()] = rec
    return out


def accuracy(run: dict, gold: dict) -> tuple[int, int]:
    ok = total = 0
    for key in SCORED:
        expected = gold.get(key)
        if not expected:
            continue
        total += 1
        got = run.get(key)
        if key in LIST_KEYS:
            ok += {_norm(x) for x in got or []} == {_norm(x) for x in expected}
        else:
            ok += _norm(got) == _norm(expected)
    return ok, total


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--pdf-dir", required=True, type=Path)
    ap.add_argument("--runs", type=int, default=3)
    ap.add_argument("--model", default=os.environ.get("GEMINI_MODEL", "gemini-3.8-flash"))
    ap.add_argument("--gold", type=Path)
    ap.add_argument("--make-gold", type=Path)
    ap.add_argument("--out", type=Path, default=Path("reliability_results.xlsx"))
    ap.add_argument("--full-report", action="store_true", help="also run the full analysis once (Source Check)")
    ap.add_argument("--pause", type=float, default=4.0, help="seconds between requests (free-tier rate limits)")
    args = ap.parse_args()

    if args.make_gold:
        make_gold(args.pdf_dir, args.make_gold)
        return 0
    key, _src = load_api_key()
    if not key:
        print("No Gemini API key: set GEMINI_API_KEY or save the key in the app.")
        return 2
    settings = AppSettings(model=args.model, ocr_mode="Off", consistency_check=False)
    client = GeminiClient(key, settings.model, effective_temperature(settings), 4096, 300, 4)
    gold = load_gold(args.gold) if args.gold else {}

    from openpyxl import Workbook

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    detail = wb.create_sheet("Per field")
    detail.append(["File", "Field", *[f"Run {i + 1}" for i in range(args.runs)], "Runs agree?", "Answer key",
                   "Run 1 correct?"])
    per_file = [["File", "Runs", "Agreement %", "Accuracy % (run 1)", "Accuracy % (all runs)", "Source Check %",
                 "Errors"]]
    tot_agree = tot_fields = tot_ok = tot_scored = 0
    for pdf in sorted(args.pdf_dir.glob("*.pdf")):
        print(f"• {pdf.name}")
        errors = []
        try:
            doc = extract_text(validate_pdf(pdf), ocr_mode="Off")
        except Exception as exc:  # noqa: BLE001
            per_file.append([pdf.name, 0, "", "", "", "", f"PDF: {exc}"])
            continue
        source = doc.full_text(with_page_markers=False)
        text, _scope = scope_text(doc.full_text(with_page_markers=True))
        runs = []
        for i in range(args.runs):
            try:
                runs.append(extract_key_facts(client, text))
            except Exception as exc:  # noqa: BLE001
                errors.append(f"run {i + 1}: {exc}")
            time.sleep(args.pause)
        if len(runs) < 2:
            per_file.append([pdf.name, len(runs), "", "", "", "", "; ".join(errors)])
            continue
        comps = compare_runs(runs, source)
        agree = sum(1 for c in comps if c.status != DIFFERS)
        tot_agree += agree
        tot_fields += len(comps)
        g = gold.get(pdf.name, {})
        ok1, n1 = accuracy(runs[0], g)
        all_ok = sum(accuracy(r, g)[0] for r in runs)
        all_n = sum(accuracy(r, g)[1] for r in runs)
        tot_ok += all_ok
        tot_scored += all_n
        for c in comps:
            gv = g.get(c.key, "")
            correct = ""
            if gv and c.key in SCORED:
                correct = "Yes" if accuracy({c.key: runs[0].get(c.key)}, {c.key: gv})[0] else "No"
            detail.append([pdf.name, c.label, *c.display_values(), "Yes" if c.status == AGREE else c.status,
                           "; ".join(gv) if isinstance(gv, list) else gv, correct])
        sc = ""
        if args.full_report:
            from judgment_analyzer.analyzer import Analyzer

            try:
                res = Analyzer(settings, client).run(validate_pdf(pdf))
                sc = round((res.verification.score or 0) * 100, 1)
            except Exception as exc:  # noqa: BLE001
                errors.append(f"full report: {exc}")
        per_file.append([pdf.name, len(runs), round(agree / len(comps) * 100, 1),
                         round(ok1 / n1 * 100, 1) if n1 else "", round(all_ok / all_n * 100, 1) if all_n else "",
                         sc, "; ".join(errors)])

    summary.append(["Judgement Insights Analyser - reliability test", time.strftime("%d %b %Y %H:%M")])
    summary.append(["Model", args.model])
    summary.append(["Runs per judgment", args.runs])
    summary.append(["Overall run-to-run agreement %",
                    round(tot_agree / tot_fields * 100, 1) if tot_fields else ""])
    summary.append(["Overall accuracy vs answer key %", round(tot_ok / tot_scored * 100, 1) if tot_scored else
                    "(no answer key)"])
    summary.append([])
    for row in per_file:
        summary.append(row)
    wb.save(args.out)
    print(f"\nResults written to {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
