"""Tests for the formatted template and audit log, SMTP sign-in diagnosis."""

import csv

from openpyxl import load_workbook

from judgment_analyzer.mailer import AUDIT_FIELDS, SmtpMailer, append_audit, audit_xlsx_path, write_audit_xlsx
from judgment_analyzer.recipients import create_template, load_recipients


def test_template_is_a_bordered_table_and_still_loads(tmp_path):
    for group in ("Client", "Employee"):
        path = create_template(tmp_path / f"{group}.xlsx", group)
        ws = load_workbook(path).worksheets[0]
        head = ws["A1"]
        assert head.font.bold and head.fill.fgColor.rgb.endswith("0B1F3A") and head.border.left.style == "thin"
        assert ws["B30"].border.bottom.style == "thin"          # empty rows ready to fill, with borders
        assert ws.auto_filter.ref and ws.freeze_panes == "A2"
        res = load_recipients(path, group)
        assert len(res.recipients) == 2 and not res.errors      # bordered blank rows are ignored


def test_audit_log_gets_formatted_excel_copy(tmp_path):
    row = {k: "" for k in AUDIT_FIELDS}
    row.update({"Timestamp": "2026-10-18 10:00", "Recipient email": "a@example.com", "Status": "OK"})
    bad = dict(row, Status="FAILED", Error="sign-in failed")
    csv_path = append_audit([row, bad], tmp_path / "email_audit_log.csv")
    out = write_audit_xlsx(csv_path)
    assert out == audit_xlsx_path(csv_path) and out.exists()
    ws = load_workbook(out).active
    assert [c.value for c in ws[1]] == AUDIT_FIELDS and ws["A1"].font.bold
    status_col = AUDIT_FIELDS.index("Status") + 1
    assert ws.cell(3, status_col).value == "FAILED" and ws.cell(3, status_col).font.color.rgb.endswith("C62828")
    assert ws.cell(2, 1).border.top.style == "thin"
    with csv_path.open(encoding="utf-8-sig") as fh:
        assert len(list(csv.reader(fh))) == 3                    # CSV stays the master record


def test_diagnosis_text_format():
    text = SmtpMailer.diagnosis_text([("Find the server", True, "ok"), ("Port 465 (SSL): sign in", False, "cut")])
    assert text.splitlines()[1].startswith("FAIL")


def test_gmail_app_password_normalising():
    from judgment_analyzer.mailer import looks_like_gmail_app_password, normalise_password

    assert normalise_password("smtp.gmail.com", "abcd efgh ijkl​mnop ") == "abcdefghijklmnop"
    assert looks_like_gmail_app_password("abcd efgh ijkl mnop")
    assert not looks_like_gmail_app_password("MyGmail@2024")
    assert normalise_password("smtp.office365.com", "pa ss") == "pa ss"


def test_recipient_lists_are_saved_and_reloaded(tmp_path):
    from judgment_analyzer.recipients import Recipient, RecipientStore

    store = RecipientStore(tmp_path / "lists.db")
    store.save("Client", [Recipient("A Client", "a@example.com", "Mr", "A Ltd"),
                          Recipient("B Client", "b@example.com")], "Clients.xlsx")
    store.save("Employee", [Recipient("Colleague", "c@example.com")], "Team.xlsx")
    store.close()
    store = RecipientStore(tmp_path / "lists.db")                 # next session
    people, info = store.load("Client")
    assert [p.email for p in people] == ["a@example.com", "b@example.com"]
    assert people[0].salutation == "Mr" and people[0].group == "Client" and info["file_name"] == "Clients.xlsx"
    store.save("Client", [Recipient("New", "n@example.com")], "Clients2.xlsx")   # replace, not append
    assert [p.email for p in store.load("Client")[0]] == ["n@example.com"]
    store.forget("Employee")
    assert store.load("Employee") == ([], {})


V1_HEADER = ["Timestamp", "Approved by", "Sender email", "Judgment file", "Case title", "Subject", "Recipient name",
             "Recipient email", "Delivery", "Status", "Error", "Content fingerprint"]


def test_mixed_version_audit_log_is_realigned(tmp_path):
    from judgment_analyzer.mailer import read_audit_rows, repair_audit_log

    path = tmp_path / "email_audit_log.csv"
    v1_row = ["16-10-2026 10:00", "R", "r@x.com", "a.pdf", "A v. B", "Subj 1", "Old Person", "o@x.com",
              "Outlook draft", "OK", "", "f1"]
    v2_row = ["17-10-2026 15:56", "R", "r@x.com", "b.pdf", "C v. D", "JIA-2026-0006", "Subj 2", "Example Client",
              "c@x.com", "Client", "Outlook draft", "OK", "", "f2"]
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerows([V1_HEADER, v1_row, v2_row])
    rows = read_audit_rows(path)
    assert rows[0]["Subject"] == "Subj 1" and rows[0].get("Register ref", "") == ""
    assert rows[1]["Register ref"] == "JIA-2026-0006" and rows[1]["Subject"] == "Subj 2"
    assert rows[1]["Recipient group"] == "Client" and rows[1]["Status"] == "OK"
    assert repair_audit_log(path) is True
    with path.open(newline="", encoding="utf-8-sig") as fh:
        fixed = list(csv.reader(fh))
    assert fixed[0] == AUDIT_FIELDS and all(len(r) == len(AUDIT_FIELDS) for r in fixed)
    assert list(tmp_path.glob("email_audit_log_before_repair_*.csv"))          # original kept
    assert repair_audit_log(path) is False                                      # already current
    append_audit([dict.fromkeys(AUDIT_FIELDS, "x")], path)
    ws = load_workbook(write_audit_xlsx(path)).active
    assert [c.value for c in ws[1]] == AUDIT_FIELDS and ws.cell(3, 6).value == "JIA-2026-0006"
