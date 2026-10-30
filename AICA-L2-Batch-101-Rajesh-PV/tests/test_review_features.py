"""Tests for the core review features: paragraph markers, Source Check, Consistency Check, precedent register,
comparison prompt, recipient groups, team email version, SMTP delivery and temperature handling.
No network calls; Gemini and the mail server are faked."""

from __future__ import annotations

import csv
import json
import smtplib
from types import SimpleNamespace

import pytest
from conftest import SAMPLE_REPORT, _make_text_pdf

from judgment_analyzer import gemini_client as gc
from judgment_analyzer.analyzer import Analyzer
from judgment_analyzer.compare import compare_prompt, compare_records
from judgment_analyzer.config import AppSettings, effective_temperature
from judgment_analyzer.consistency import (AGREE, DIFFERS, WORDING, compare_runs, parse_json_object,
                                           run_consistency)
from judgment_analyzer.email_composer import (EmailContent, SenderProfile, build_internal_body, clean_for_client,
                                              render_email)
from judgment_analyzer.mailer import MODE_SMTP, MODE_TEST, SmtpMailer, SmtpSettings, deliver_batch
from judgment_analyzer.paragraphs import add_para_markers, detect_paragraphs
from judgment_analyzer.pdf_processor import PageText, extract_text, validate_pdf
from judgment_analyzer.recipients import load_recipients, normalise_group
from judgment_analyzer.repository import (STATUS_VERIFIED, JudgmentRecord, Repository, export_xlsx,
                                          file_sha256, record_from_result)
from judgment_analyzer.response_parser import parse_report
from judgment_analyzer.verifier import CLOSE, FOUND, NOT_FOUND, verify_report

ORDER_PAGES = [
    "IN THE INCOME TAX APPELLATE TRIBUNAL\nITA No. 999/EXP/2025\nAssessment Year: 2018-19\n"
    "Demo Traders Pvt Ltd ... Appellant\nVs.\nDCIT ... Respondent\nDate of pronouncement: 15.09.2026\nORDER\n"
    "1. This appeal by the assessee is against the order dated 10.03.2025.\n"
    "2. The only ground is the addition of Rs. 1,25,00,000 made under section 68 of the Act.",
    "3. Notices under section 133(6) were not answered by the lenders.\n"
    "4. The assessee filed confirmations and bank statements of both lenders.\n"
    "5. The mere non-response to notices under section 133(6) cannot by itself be a ground for addition when "
    "the documents on record are not shown to be false.\n"
    "6. In the result, the appeal of the assessee is allowed.",
]


@pytest.fixture
def order_pdf(tmp_path):
    return _make_text_pdf(tmp_path / "order.pdf", ORDER_PAGES)


def _report(body_7: str = "", body_2: str = "Addition deleted (para 6).") -> str:
    return SAMPLE_REPORT.replace("The appeal was allowed (p. 3).", body_2).replace(
        "[Court Finding] Section 9 does not apply (p. 3).", body_7 or "[Court Finding] Held for the assessee.")


# ------------------------------------------------------------ paragraphs ---
def test_detect_paragraphs_follows_running_sequence():
    pages = [PageText(1, "Heading\n1. First para\n2. Second\n1. a sub-list item\n3. Third"),
             PageText(2, "4. Fourth\n2023. not a para\n6. Sixth (5 missed by extraction)")]
    marks = detect_paragraphs(pages)
    assert [m.number for m in marks] == [1, 2, 3, 4, 6]
    marked = add_para_markers(pages, marks)
    assert marked[0].text.splitlines()[1].startswith("[Para 1] 1. First")
    assert "[Para 6]" in marked[1].text and pages[1].text.count("[Para") == 0  # originals untouched


def test_detect_paragraphs_needs_three():
    assert detect_paragraphs([PageText(1, "1. Only\n2. Two numbered lines")]) == []


# ------------------------------------------------------------ verifier ---
def test_source_check_finds_and_flags(order_pdf):
    doc = extract_text(validate_pdf(order_pdf), ocr_mode="Off")
    source = doc.full_text(with_page_markers=False)
    body7 = ('[Court Finding] "The mere non-response to notices under section 133(6) cannot by itself be a ground '
             'for addition when the documents on record are not shown to be false." (para 5) '
             '"The assessee had no business at all during the relevant year." (para 14) '
             'The Tribunal referred to ITA No. 999/EXP/2025 and to [2024] 160 taxmann.com 1 dated 01.01.2020; '
             'amount Rs. 1,25,00,000 and Rs. 9,99,999 under section 68 and section 271(1)(c).')
    report = parse_report(_report(body7))
    result = verify_report(report, source, para_numbers={1, 2, 3, 4, 5, 6}, page_count=2)
    by_text = {i.text: i for i in result.items}
    status = lambda prefix: next(i.status for t, i in by_text.items() if t.startswith(prefix))  # noqa: E731
    assert status("“The mere non-response") == FOUND
    assert status("“The assessee had no business") == NOT_FOUND
    assert status("ITA No. 999/EXP/2025") == FOUND
    assert status("[2024] 160 taxmann.com 1") == NOT_FOUND
    assert status("Rs. 1,25,00,000") == FOUND and status("Rs. 9,99,999") == NOT_FOUND
    assert status("Section 68") == FOUND and status("Section 271(1)(c)") == NOT_FOUND
    assert status("para 5") == FOUND and status("para 14") == NOT_FOUND
    assert status("01-01-2020") == NOT_FOUND
    assert 0 < result.score < 1
    assert result.items[0].status == NOT_FOUND  # problems listed first
    md = result.to_markdown()
    assert "Not found in the judgment" in md and "[Uncertain]" in md


def test_source_check_close_match_and_no_items():
    report = parse_report(_report('"the documents on record were not shown to be false by the revenue at all '
                                  'in this case before us"'))
    source = "the documents on record are not shown to be false by the revenue at all in this case before us"
    items = verify_report(report, source).items
    assert any(i.kind == "Quotation" and i.status == CLOSE for i in items)
    empty = verify_report(parse_report("## 2. Executive Summary\nNothing checkable here."), "text")
    assert empty.score is None and "no quotations" in empty.summary_line()


# --------------------------------------------------------- consistency ---
FACTS = {"court": "ITAT", "bench": "A (JM)", "case_title": "Demo Traders v. DCIT", "citation": "Not stated",
         "case_number": "ITA No. 999/EXP/2025", "judgment_date": "15.09.2026", "assessment_year": "2018-19",
         "appellant": "Demo Traders", "respondent": "DCIT", "sections": ["Section 68", "133(6)"],
         "issues": ["Whether addition under section 68 is sustainable"], "outcome": "Allowed",
         "in_favour_of": "Assessee", "ratio": "Non-response to notices alone cannot justify an addition."}


def test_parse_json_object_handles_fences_and_normalises():
    data = parse_json_object("```json\n" + json.dumps(FACTS) + "\n```")
    assert data["sections"] == ["68", "133(6)"]
    assert data["judgment_date"] == "15-09-2026"
    with pytest.raises(ValueError):
        parse_json_object("no json here")


def test_compare_runs_statuses():
    a = parse_json_object(json.dumps(FACTS))
    b = dict(a, assessment_year="2019-20", case_title="Demo Traders v DCIT",
             ratio="Non-response to notices alone does not justify an addition.")
    fields = {f.key: f for f in compare_runs([a, b], source_text="ITA No. 999/EXP/2025 AY 2018-19 15.09.2026")}
    assert fields["assessment_year"].status == DIFFERS
    assert fields["case_title"].status == AGREE  # punctuation-only difference
    assert fields["ratio"].status == WORDING
    assert fields["sections"].status == AGREE
    assert fields["case_number"].in_source is True and fields["judgment_date"].in_source is True


def _fake_client(responses):
    models = SimpleNamespace(calls=[])

    def generate_content(model, contents, config=None):
        models.calls.append({"contents": contents, "config": config})
        text = responses.pop(0) if responses else SAMPLE_REPORT
        return SimpleNamespace(text=text, prompt_feedback=None, candidates=[SimpleNamespace(finish_reason="STOP")],
                               usage_metadata=None)

    models.generate_content = generate_content
    models.get = lambda model: SimpleNamespace(input_token_limit=1_000_000)
    return gc.GeminiClient("A" * 39, "gemini-3.8-flash", client=SimpleNamespace(models=models)), models


def test_run_consistency_two_runs_and_json_mode():
    client, models = _fake_client([json.dumps(FACTS), json.dumps(dict(FACTS, outcome="Dismissed"))])
    result = run_consistency(client, "judgment text", "judgment text")
    assert result.ok and [f.key for f in result.disagreements] == ["outcome"]
    assert models.calls[0]["config"].response_mime_type == "application/json"
    assert "disagree on Outcome" in result.summary_line()


def test_run_consistency_survives_bad_output():
    client, _ = _fake_client(["not json", json.dumps(FACTS)])
    result = run_consistency(client, "t", "t")
    assert not result.ok and result.errors


def test_analyzer_runs_quality_checks(order_pdf):
    client, models = _fake_client([_report(), json.dumps(FACTS), json.dumps(FACTS)])
    result = Analyzer(AppSettings(ocr_mode="Off"), client).run(validate_pdf(order_pdf))
    assert result.para_numbers == {1, 2, 3, 4, 5, 6}
    assert "[Para 1]" in models.calls[0]["contents"]
    assert result.verification is not None and result.consistency.ok
    assert not result.needs_key_fact_review
    assert result.key_facts["case_number"] == "ITA No. 999/EXP/2025"


def test_analyzer_flags_disagreement(order_pdf):
    client, _ = _fake_client([_report(), json.dumps(FACTS), json.dumps(dict(FACTS, in_favour_of="Revenue"))])
    result = Analyzer(AppSettings(ocr_mode="Off"), client).run(validate_pdf(order_pdf))
    assert result.needs_key_fact_review
    assert any("differ" in w for w in result.warnings)


# ----------------------------------------------------------- temperature ---
def test_effective_temperature_follows_model_family():
    assert effective_temperature(AppSettings(model="gemini-3.8-flash")) is None
    assert effective_temperature(AppSettings(model="gemini-2.5-flash")) == 0.2
    assert effective_temperature(AppSettings(model="gemini-3.8-flash", temperature_mode="fixed")) == 0.2
    client, models = _fake_client(["ok"])
    client.generate("p", "s")
    assert models.calls[0]["config"].temperature is None


# -------------------------------------------------------------- register ---
def test_register_add_search_verify_export(tmp_path, order_pdf):
    repo = Repository(tmp_path / "reg" / "register.db")
    a = repo.add(JudgmentRecord(case_title="Demo Traders v. DCIT", sections=["68", "133(6)"], in_favour_of="Assessee",
                                issue_tags=["cash credit"], pdf_sha256=file_sha256(order_pdf)),
                 pdf_path=order_pdf)
    repo.add(JudgmentRecord(case_title="Other Ltd v. ITO", sections=["168"], in_favour_of="Revenue"))
    assert a.ref.startswith("JIA-") and (repo.pdf_folder / a.pdf_copy).exists()
    assert [r.case_title for r in repo.search(section="68")] == ["Demo Traders v. DCIT"]  # not "168"
    assert [r.case_title for r in repo.search("cash")] == ["Demo Traders v. DCIT"]
    assert len(repo.search(favour="Revenue")) == 1
    assert repo.find_by_sha(file_sha256(order_pdf))[0].ref == a.ref
    repo.set_verified(a.id, "RR")
    assert repo.get(a.id).status == STATUS_VERIFIED and repo.stats()["verified"] == 1
    out = export_xlsx(repo.search(), tmp_path / "reg.xlsx")
    from openpyxl import load_workbook
    ws = load_workbook(out).active
    assert ws["A1"].value == "Ref" and ws.max_row == 3
    repo.delete(a.id)
    assert repo.get(a.id) is None and not (repo.pdf_folder / a.pdf_copy).exists()


def test_record_from_result_and_compare(order_pdf):
    client, models = _fake_client([_report(), json.dumps(FACTS), json.dumps(FACTS)])
    result = Analyzer(AppSettings(ocr_mode="Off"), client).run(validate_pdf(order_pdf))
    rec = record_from_result(result, result.key_facts, ["s68"], "RR")
    assert rec.sections == ["68", "133(6)"] and rec.source_check_summary.startswith("Source check")
    rec.ref, other = "JIA-2026-0001", JudgmentRecord(ref="JIA-2026-0002", case_title="B", report_md=rec.report_md)
    prompt = compare_prompt(rec, other)
    assert "JIA-2026-0001" in prompt and "Where the rulings differ or conflict" in prompt
    client2, _ = _fake_client(["## 1. Snapshot\n- x"])
    assert compare_records(client2, rec, other).startswith("## 1. Snapshot")


# ------------------------------------------------------------------ email ---
def test_recipient_groups(tmp_path):
    p = tmp_path / "r.csv"
    with p.open("w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows([["Name", "Email ID", "Group"], ["A", "a@example.com", "Client"],
                                  ["B", "b@example.com", "Staff"], ["C", "c@example.com", ""],
                                  ["D", "d@example.com", "Vendor"]])
    res = load_recipients(p)
    assert [(r.name, r.group) for r in res.recipients] == [("A", "Client"), ("B", "Employee"), ("C", "Client")]
    assert any("Vendor" in e for e in res.errors)
    assert normalise_group("internal") == "Employee" and normalise_group("xyz") is None


def test_team_version_and_clean_output():
    report = parse_report(SAMPLE_REPORT)
    team = build_internal_body(report, [7, 2], "JIA-2026-0007")
    assert "[Court Finding]" in team and "JIA-2026-0007" in team and "(p." not in team
    assert clean_for_client("A point (para 4; p. 2) [AI Observation] here.") == "A point (para 4) here."
    assert clean_for_client("[AI Generated] x [Para 3] y") == "x y"
    content = EmailContent("S", "I", team, SenderProfile("R", email="r@example.com"), disclaimer="Firm note")
    msg = render_email(content, "A", "a@example.com")
    for leftover in ("[Court Finding]", "[AI", "##", "(p. "):
        assert leftover not in msg.text and leftover not in msg.html
    assert "Firm note" in msg.html


class FakeSMTP:
    instances: list = []

    def __init__(self, fail_login=False):
        self.sent, self.fail_login, self.closed = [], fail_login, False
        FakeSMTP.instances.append(self)

    def login(self, user, password):
        if self.fail_login:
            raise smtplib.SMTPAuthenticationError(535, b"bad")

    def send_message(self, mime):
        self.sent.append(mime)

    def quit(self):
        self.closed = True


def _msg(n=1):
    content = EmailContent("Subject", "Intro", "## 1. Case Details\nBody", SenderProfile("R", email="r@example.com"))
    return [render_email(content, f"P{i}", f"p{i}@example.com") for i in range(n)]


def test_smtp_delivery_and_audit(tmp_path):
    settings = SmtpSettings(username="r@example.com")
    mailer = SmtpMailer(settings, "app-password", "R", "r@example.com", factory=lambda s: FakeSMTP())
    audit = tmp_path / "audit.csv"
    results = deliver_batch(_msg(2), MODE_SMTP, {"Approved by": "R"}, smtp=mailer, audit_path=audit,
                            groups=["Client", "Employee"], pause_s=0)
    assert all(r.ok for r in results)
    fake = FakeSMTP.instances[-1]
    assert len(fake.sent) == 2 and fake.sent[0]["To"].endswith("<p0@example.com>") and fake.closed
    rows = list(csv.DictReader(audit.open(encoding="utf-8-sig")))
    assert [r["Recipient group"] for r in rows] == ["Client", "Employee"]
    assert rows[0]["Delivery"] == MODE_SMTP


def test_smtp_login_failure_stops_batch(tmp_path):
    settings = SmtpSettings(host="smtp.gmail.com", username="r@gmail.com")
    mailer = SmtpMailer(settings, "wrong", "R", "r@gmail.com", factory=lambda s: FakeSMTP(fail_login=True))
    results = deliver_batch(_msg(3), MODE_TEST, {}, smtp=mailer, audit_path=tmp_path / "a.csv", pause_s=0)
    assert len(results) == 1 and not results[0].ok and "App Password" in results[0].error


def test_smtp_settings_validation():
    assert SmtpSettings(host="", username="").validate()
    assert SmtpSettings(username="me@example.com").validate() == []


# -------------------------------------------------------- clean outputs ---
def test_final_document_is_clean_and_justified(tmp_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    from judgment_analyzer.exporter import build_final_document, citation_line, save_docx, save_pdf

    report = parse_report(SAMPLE_REPORT)
    facts = {"case_title": "Alpha Traders v. Beta", "citation": "Not stated", "case_number": "Civil Appeal No. 123",
             "court": "High Court of Exampleland", "judgment_date": "01-03-2026"}
    doc = build_final_document(report, facts)
    meta = dict(doc.metadata)
    assert set(meta) == {"Generated on", "Case", "Citation"}
    assert meta["Citation"].startswith("Not reported in the source copy - Civil Appeal No. 123")
    assert all(h != "Disclaimer" for h, _ in doc.sections)
    md = doc.to_markdown()
    assert "[Court Finding]" not in md and "(p. " not in md and "AI Observation" not in md
    path = save_docx(doc, tmp_path / "x.docx")
    d = Document(path)
    text = "\n".join(p.text for p in d.paragraphs)
    assert "AI-generated" not in text and "not legal advice" not in d.sections[0].footer.paragraphs[0].text
    assert any(p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY for p in d.paragraphs)
    save_pdf(doc, tmp_path / "x.pdf")
    assert citation_line({"citation": "[2026] 1 taxmann.com 2"}) == "[2026] 1 taxmann.com 2"


def test_approval_tied_to_content(order_pdf):
    client, _ = _fake_client([_report(), json.dumps(FACTS), json.dumps(FACTS)])
    result = Analyzer(AppSettings(ocr_mode="Off"), client).run(validate_pdf(order_pdf))
    assert result.uncertain_lines()  # SAMPLE_REPORT has an [Uncertain] point
    assert not result.is_approved
    result.approve("RR")
    assert result.is_approved and result.approved_by == "RR"
    result.report.sections[2] += " Edited."
    assert not result.is_approved  # any change after approval needs re-approval


def test_smtp_falls_back_to_other_port(monkeypatch):
    calls = []

    def fake_open(self, host, port, security, helo):
        calls.append((port, security, helo))
        if port == 465:
            raise smtplib.SMTPServerDisconnected("Connection unexpectedly closed")
        return FakeSMTP()

    monkeypatch.setattr(SmtpMailer, "_open", fake_open)
    mailer = SmtpMailer(SmtpSettings(username="me@gmail.com"), "abcd efgh ijkl mnop", "Me", "me@gmail.com")
    assert "587" in mailer.test()
    assert calls[0][:2] == (465, "SSL") and (587, "STARTTLS", None) in calls


def test_smtp_reports_blocked_network(monkeypatch):
    def fake_open(self, host, port, security, helo):
        raise smtplib.SMTPServerDisconnected("Connection unexpectedly closed")

    monkeypatch.setattr(SmtpMailer, "_open", fake_open)
    mailer = SmtpMailer(SmtpSettings(username="me@gmail.com"), "pw", "Me", "me@gmail.com")
    with pytest.raises(Exception) as exc:
        mailer.test()
    assert "465" in str(exc.value) and "587" in str(exc.value) and "firewall" in str(exc.value)


def test_citation_lookup_parsing_and_guard():
    from judgment_analyzer.citation_lookup import lookup_citation, parse_lookup

    assert parse_lookup("CITATION: [2026] 1 taxmann.com 9 (Chennai - Trib.)\nCONFIDENCE: High\nNOTE: x") == (
        "[2026] 1 taxmann.com 9 (Chennai - Trib.)", "high", "x")
    assert parse_lookup("CITATION: NONE\nCONFIDENCE: low")[0] == ""

    class C:
        def __init__(self, sources):
            self.sources = sources

        def generate(self, *a, **k):
            assert k.get("web_search") is True
            return gc.GenerationResult(text="CITATION: [2026] 1 taxmann.com 9\nCONFIDENCE: high\nNOTE: n",
                                       sources=self.sources)

    assert lookup_citation(C([("Taxmann", "https://example.com/x")]), {}).citation == "[2026] 1 taxmann.com 9"
    assert lookup_citation(C([]), {}).citation == ""  # no web source -> discarded


def test_separate_lists_get_group(tmp_path):
    from judgment_analyzer.recipients import create_template

    emp = create_template(tmp_path / "emp.xlsx", "Employee")
    res = load_recipients(emp, default_group="Employee")
    assert {r.group for r in res.recipients} == {"Employee"}
