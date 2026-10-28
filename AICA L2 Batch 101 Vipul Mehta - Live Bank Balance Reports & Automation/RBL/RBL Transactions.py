"""
RBL Corporate NetBanking — Transactions Checker
=================================================
Logs into corporate account, fetches recent transactions for all accounts,
saves everything to an Excel file with an account name column.

Requirements:
    pip install playwright openpyxl
    python -m playwright install chromium

Usage:
    python rbl_transactions.py
"""

import os
import sys
import time
from datetime import datetime
import openpyxl
from pathlib import Path
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

# Support running both from base folder and from RBL subfolder
for _p in [Path(__file__).parent.resolve(), Path(__file__).parent.parent.resolve()]:
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

# ─── Configuration ────────────────────────────────────────────────────────────
from config import RBL_TRANSACTIONS_FILE as OUTPUT_FILE, RBL_BALANCES_FILE as BALANCE_FILE, CHROME_PATH, load_bank_accounts
HEADLESS    = False

# Accounts to process — loaded dynamically from bank_accounts.json
def get_accounts():
    try:
        data = load_bank_accounts()
        accs = [
            {"name": a.get("code") or a.get("name"), "user_id": a.get("id") or a.get("user_id")}
            for a in data.get("RBL", []) if a.get("active", True)
        ]
        if accs:
            return accs
    except Exception:
        pass
    return [

    ]

ACCOUNTS = get_accounts()
# ──────────────────────────────────────────────────────────────────────────────

def get_password():
    from config import get_password as _get_password
    print("\n" + "="*55)
    print("  RBL Corporate NetBanking Transactions Checker")
    print("="*55)
    print("Accounts to fetch: " + ", ".join(a["name"] for a in ACCOUNTS))
    return _get_password("Enter your RBL NetBanking Password (hidden): ")

# ─── Excel helpers ────────────────────────────────────────────────────────────

def to_float(val):
    """Safely converts string currencies to float for Excel."""
    if not val:
        return ""
    try:
        return float(str(val).replace(",", "").replace("₹", "").strip())
    except ValueError:
        return val

def parse_date_value(date_str):
    """Parses various date strings into a datetime.date object for Excel custom formatting."""
    if not date_str:
        return ""
    if isinstance(date_str, datetime):
        return date_str.date()
    
    date_str = str(date_str).strip()
    
    for fmt in ("%d/%m/%y", "%d/%m/%Y", "%d %b %Y", "%d-%b-%Y", "%d-%b-%y"):
        try:
            return datetime.strptime(date_str, fmt).date()
        except ValueError:
            continue
            
    return date_str

def save_balances(rows: list):
    if not rows:
        return
    
    file_exists = os.path.isfile(BALANCE_FILE)
    
    if file_exists:
        wb = openpyxl.load_workbook(BALANCE_FILE)
        ws = wb.active
        ws.append([]) # Add blank row separator if file already has data
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Balances"
        # Write headers
        ws.append(["date", "time", "account_name", "account_type", "account_number", "balance"])

    for row in rows:
        ws.append([
            row["date"], 
            row["time"], 
            row["account_name"], 
            row["account_type"], 
            row["account_number"], 
            to_float(row["balance"]) # Save balance as a number in Excel
        ])

    ws.append([]) # Add blank line after each account's rows
    wb.save(BALANCE_FILE)
    print(f"   ✅ Balances saved → {BALANCE_FILE}")

def save_transactions(rows: list):
    if not rows:
        return
    
    file_exists = os.path.isfile(OUTPUT_FILE)
    
    if file_exists:
        wb = openpyxl.load_workbook(OUTPUT_FILE)
        ws = wb.active
        ws.append([]) # blank separator between account groups
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Transactions"
        ws.append(["fetched_date", "fetched_time", "account_name",
                   "account_number_masked", "txn_date", "description",
                   "debit", "credit", "balance"])
                   
    for row in rows:
        parsed_fetched_date = parse_date_value(row.get("fetched_date", ""))
        parsed_txn_date = parse_date_value(row.get("txn_date", ""))
        
        ws.append([
            parsed_fetched_date,
            row.get("fetched_time", ""),
            row.get("account_name", ""),
            row.get("account_number_masked", ""),
            parsed_txn_date,
            row.get("description", ""),
            to_float(row.get("debit")),
            to_float(row.get("credit")),
            to_float(row.get("balance"))
        ])
        
        row_idx = ws.max_row
        
        c1 = ws.cell(row=row_idx, column=1)
        if hasattr(parsed_fetched_date, 'year'):
            c1.number_format = 'dd-mmm-yy'
            
        c5 = ws.cell(row=row_idx, column=5)
        if hasattr(parsed_txn_date, 'year'):
            c5.number_format = 'dd-mmm-yy'
        
    wb.save(OUTPUT_FILE)
    print(f"   ✅ Transactions saved → {OUTPUT_FILE}")

# ─── Playwright helpers ───────────────────────────────────────────────────────

def quick_fill(page, selectors, value, label, timeout=1000):
    for sel in selectors:
        try:
            el = page.wait_for_selector(sel, timeout=timeout)
            el.fill(value)
            print(f"   [OK] Filled '{label}' -> {sel}")
            return True
        except Exception:
            continue
    print(f"   [FAIL] Could not fill '{label}'")
    return False

def quick_click(page, selectors, label, timeout=1000, use_js=False):
    combined_sel = ", ".join(selectors)
    try:
        el = page.wait_for_selector(combined_sel, timeout=timeout, state="attached")
        if use_js:
            el.evaluate("node => node.click()")
        else:
            el.click()
        print(f"   ✅ Clicked '{label}'")
        return True
    except Exception:
        print(f"   ⚠️  Could not click '{label}'")
        return False

def handle_interstitial_popups(page, max_wait_sec=15):
    ad_close_selectors = [
        "input[id='CLOSE']", "input[name='Action.CLOSE']", "button[aria-label*='close' i]", "a[aria-label*='close' i]",
        "[id*='close' i]", "[class*='close' i]", "button:has-text('✕')",
        "button:has-text('×')", "span:has-text('✕')", "span:has-text('×')",
        "img[src*='close' i]", "a:has-text('Skip')", "button:has-text('Skip')",
        "a:has-text('Close')", "button:has-text('Close')",
        "a:has-text('Remind me later')", "button:has-text('Later')",
        "button:has-text('Proceed to Dashboard')", "a:has-text('Continue to Dashboard')",
        "a:has-text('Proceed to Accounts')", "button:has-text('Dismiss')",
    ]

    target_selectors = ["a#ID_CACTS", "a[name='HREF_Accounts']", "a#menu-button.menu-box"]

    start_time = time.time()
    while time.time() - start_time < max_wait_sec:
        for target_sel in target_selectors:
            try:
                if page.locator(target_sel).first.is_visible():
                    return True
            except Exception:
                pass

        clicked = False
        for sel in ad_close_selectors:
            try:
                loc = page.locator(sel).first
                if loc.is_visible(timeout=500):
                    loc.click(timeout=1000)
                    print(f"   ✅ Dismissed advert / popup via selector: {sel}")
                    page.wait_for_timeout(1000)
                    clicked = True
                    break
            except Exception:
                continue

        if not clicked:
            page.wait_for_timeout(500)

    return False

# ─── Transaction Scraping ─────────────────────────────────────────────────────

def parse_accounts(page) -> list:
    print("      ⏳ Waiting for account table to load...")
    
    try:
        page.wait_for_selector("a.searchsimpletext_wrapnew.bluelink, a[id*='OPR_ACCOUNT_NUMBER_ARRAY']", timeout=15000)
    except Exception:
        pass 
    
    page.wait_for_timeout(2000)
    
    # 1. Grab Account Numbers
    acc_elements = page.locator("a.searchsimpletext_wrapnew.bluelink").all()
    if not acc_elements:
        acc_elements = page.locator("a[id*='OPR_ACCOUNT_NUMBER_ARRAY']").all()
        
    # 2. Grab Account Types ("Current", "Term Deposits", etc.)
    type_elements = page.locator("span[id*='PRODUCT_CATEGORY_ARRAY']").all()
        
    # 3. Grab Balances
    bal_elements = page.locator("span.customNewUIAmount").all()

    numeric_balances = []
    for b in bal_elements:
        text = b.inner_text().strip().replace('INR', '').replace('₹', '').replace(',', '').strip()
        if text and any(char.isdigit() for char in text):
            numeric_balances.append(text)

    print(f"      🔍 Found {len(acc_elements)} account numbers, {len(type_elements)} account types, and {len(numeric_balances)} numeric balances.")

    raw_accounts = []
    total_fd_amount = 0.0

    for i in range(len(acc_elements)):
        try:
            acct_number = acc_elements[i].inner_text().strip()
            balance_str = numeric_balances[i] if i < len(numeric_balances) else ""
            acct_type = type_elements[i].inner_text().strip() if i < len(type_elements) else "Corporate"

            if balance_str:
                bal_float = float(balance_str)
                # Check if this row is a Fixed Deposit / Term Deposit
                is_fd = any(keyword in acct_type.lower() for keyword in ["term deposit", "fixed deposit", "fd"])
                if is_fd:
                    total_fd_amount += bal_float

                raw_accounts.append({
                    "account_type":   acct_type,
                    "account_number": acct_number,
                    "raw_balance":    bal_float,
                    "is_fd":          is_fd
                })
            else:
                print(f"      ⚠️  Skipped account {acct_number} because no valid balance was found.")
        except Exception as e:
            print(f"      ⚠️  Error parsing account {i}: {e}")

    final_accounts = []
    for acc in raw_accounts:
        if not acc["is_fd"]:
            # Subtract FD amount from the main/current account balance
            net_balance = round(acc["raw_balance"] - total_fd_amount, 2)
            print(f"      💵  {acc['account_type']:<14} {acc['account_number']}   Net: {net_balance:.2f} (Gross: {acc['raw_balance']:.2f} - FD: {total_fd_amount:.2f})")
            final_accounts.append({
                "account_type":   acc["account_type"],
                "account_number": acc["account_number"],
                "balance":        str(net_balance),
            })
        else:
            # Keep FD row recorded as-is for reference
            print(f"      💵  {acc['account_type']:<14} {acc['account_number']}   FD Amount: {acc['raw_balance']:.2f}")
            final_accounts.append({
                "account_type":   acc["account_type"],
                "account_number": acc["account_number"],
                "balance":        str(acc["raw_balance"]),
            })

    return final_accounts

def navigate_to_account_statement(page):
    print("   📂 Navigating to Accounts -> Account Statement...")
    
    # Check if menu toggle is visible (for small screens)
    try:
        menu_toggle = page.locator("a#menu-button.menu-box").first
        if menu_toggle.is_visible(timeout=2000):
            print("   📱 Small screen detected. Clicking hamburger menu...")
            menu_toggle.click(force=True)
            page.wait_for_timeout(1000)
    except Exception:
        pass
        
    print("   🔍 Looking for 'Accounts' link...")
    acc_clicked = quick_click(page, ["a#ID_CACTS", "a[name='HREF_Accounts']", "a:has-text('Accounts')"], "Accounts", timeout=8000, use_js=True)
    if not acc_clicked:
        print("   ⚠️  Could not click 'Accounts'. Proceeding anyway to see if 'Account Statement' is visible.")
        
    page.wait_for_load_state("domcontentloaded", timeout=5000)
    page.wait_for_timeout(1500)
    
    print("   🔍 Looking for 'Account Statement' link...")
    stmt_clicked = quick_click(page, ["a#Account-Statement_Account-Statement", "a:has-text('Account Statement')"], "Account Statement", timeout=8000, use_js=True)
    if not stmt_clicked:
        print("   ⚠️  Could not click 'Account Statement'. It might not be visible.")
        
    page.wait_for_load_state("domcontentloaded", timeout=5000)
    page.wait_for_timeout(2000)

def extract_transactions(page, account_name, account_number_masked, fetched_date, fetched_time):
    all_transactions = []
    
    try:
        page.wait_for_selector("#HWListTable5134192", timeout=15000)
    except PlaywrightTimeoutError:
        print("   ⚠️  Transaction table #HWListTable5134192 not found.")
        return all_transactions

    page_num = 1
    while True:
        print(f"      📄 Scraping page {page_num}...")
        page.wait_for_timeout(1500)
        
        rows = page.locator("#HWListTable5134192 tr.listwhiterow, #HWListTable5134192 tr.listgreyrow").all()
        for row in rows:
            try:
                txn_date_el = row.locator("span[id*='HREF_TransactionHistoryFG.TXN_DATE_ARRAY']").first
                desc_el = row.locator("span[id*='HREF_txnRemarksArray']").first
                amount_el = row.locator("span[id*='HREF_TransactionHistoryFG.AMOUNT_ARRAY']").first
                amount_type_el = row.locator("span[id*='HREF_amountType']").first
                balance_el = row.locator("span[id*='HREF_TransactionHistoryFG.ACC_BALANCE_ARRAY']").first
                
                txn_date = txn_date_el.inner_text().strip() if txn_date_el.count() > 0 else ""
                
                desc = ""
                if desc_el.count() > 0:
                    desc = desc_el.get_attribute("title") or desc_el.inner_text()
                    desc = desc.strip()
                    
                amount = amount_el.inner_text().strip() if amount_el.count() > 0 else ""
                amount_type = amount_type_el.inner_text().strip() if amount_type_el.count() > 0 else ""
                balance = balance_el.inner_text().strip() if balance_el.count() > 0 else ""
                
                debit = amount if amount_type == "Dr" else ""
                credit = amount if amount_type == "Cr" else ""
                
                if txn_date or desc or amount:
                    all_transactions.append({
                        "fetched_date": fetched_date,
                        "fetched_time": fetched_time,
                        "account_name": account_name,
                        "account_number_masked": account_number_masked,
                        "txn_date": txn_date,
                        "description": desc,
                        "debit": debit,
                        "credit": credit,
                        "balance": balance
                    })
            except Exception as e:
                pass
                
        next_selectors = ["input.widgetpaginationnext_new[title='Next']", "input[name='Action.OpTransactionListing.GOTO_NEXT__']"]
        next_clicked = False
        for nsel in next_selectors:
            try:
                btn = page.locator(nsel).first
                if btn.is_visible(timeout=1000) and not btn.is_disabled():
                    btn.click()
                    page.wait_for_load_state("domcontentloaded", timeout=10000)
                    page_num += 1
                    next_clicked = True
                    break
            except Exception:
                continue
                
        if not next_clicked:
            break
            
    print(f"      ✅ Total transactions extracted for account: {len(all_transactions)}")
    return all_transactions

def process_all_accounts(page, user_id, account_name):
    now = datetime.now()
    date_str = now.strftime("%d/%m/%y")
    time_str = now.strftime("%H:%M")
    
    all_results = []
    
    navigate_to_account_statement(page)
    
    try:
        page.wait_for_selector("a[name^='HREF_ViewAccountBalancesSummaryFG.ACCOUNT_NUMBER_ARRAY'], a[title='Account Number']", timeout=15000)
        account_links = page.locator("a[name^='HREF_ViewAccountBalancesSummaryFG.ACCOUNT_NUMBER_ARRAY'], a[title='Account Number']").all()
        num_accounts = len(account_links)
        print(f"   🔍 Found {num_accounts} accounts to process.")
    except PlaywrightTimeoutError:
        print("   ⚠️  No account links found on Account Statement page.")
        return all_results
        
    for i in range(num_accounts):
        if i > 0:
            navigate_to_account_statement(page)
            
        try:
            account_links = page.locator("a[name^='HREF_ViewAccountBalancesSummaryFG.ACCOUNT_NUMBER_ARRAY'], a[title='Account Number']").all()
            if i >= len(account_links):
                break
                
            link = account_links[i]
            acc_number = link.inner_text().strip()
            masked_number = "..." + acc_number[-4:] if len(acc_number) >= 4 else acc_number
            
            print(f"   🏦 Processing Account {i+1}/{num_accounts}: {masked_number}")
            link.click()
            page.wait_for_load_state("domcontentloaded", timeout=15000)
            
            txns = extract_transactions(page, account_name, masked_number, date_str, time_str)
            all_results.extend(txns)
            
            # Small delay before returning to list
            page.wait_for_timeout(1500)
            
        except Exception as e:
            print(f"   ❌ Error processing account {i+1}: {e}")
            
    return all_results

def login_and_fetch(page, user_id: str, account_name: str, password: str, is_last: bool = False):
    bal_results = []
    txn_results = []
    
    print(f"\n{'='*55}")
    print(f"  Account: {account_name}  (User ID: {user_id})")
    print(f"{'='*55}")

    try:
        print("🔗 Loading login page...")
        page.goto("https://online.rbl.bank.in/CorpBank/", timeout=30000, wait_until="domcontentloaded")

        print("🔑 Entering Corporate User ID...")
        quick_fill(page, ["input[type='text']"], user_id, "User ID", timeout=10000)

        quick_click(page, ["input[value='Next']"], "Next/Continue", timeout=3000)
        page.wait_for_timeout(2000)

        print("🖼️  Checking 'I confirm the image' checkbox...")
        try:
            page.locator("input[id='AuthenticationFG.TARGET_CHECKBOX']").check(force=True, timeout=3000)
            print("   ✅ Checked using Playwright check(force=True)")
        except Exception:
            try:
                page.evaluate('''() => { 
                    const cb = document.getElementById('AuthenticationFG.TARGET_CHECKBOX'); 
                    if(cb) { cb.click(); cb.checked = true; cb.dispatchEvent(new Event('change', {bubbles: true})); } 
                }''')
                print("   ✅ Checked using JavaScript event dispatch")
            except Exception:
                print("   ⚠️ Could not check the box.")

        print("🔒 Entering password...")
        quick_fill(page, [
            "input[id='AuthenticationFG.ACCESS_CODE']",
            "input[type='password']",
            "input[id*='password']", "input[id*='Password']",
            "input[placeholder*='Password']",
        ], password, "Password")

        quick_click(page, [
            "input[id='VALIDATE_STU_CREDENTIALS_UX']",
            "input[type='submit']",
            "button:has-text('Login')", "button:has-text('LOGIN')",
            "button:has-text('Sign In')",
        ], "Login", use_js=True)

        print("⏳ Waiting for post-login home page (handling potential popups)...")
        page.wait_for_load_state("domcontentloaded", timeout=10000)
        handle_interstitial_popups(page, max_wait_sec=12)

        # Scrape balances
        now = datetime.now()
        date_str = now.strftime("%d/%m/%y")
        time_str = now.strftime("%H:%M")

        print("💰 Reading balances from Home Page...")
        accounts = parse_accounts(page)
        for acc in accounts:
            bal_results.append({
                "date":           date_str,
                "time":           time_str,
                "account_name":   account_name,
                "account_type":   acc["account_type"],
                "account_number": acc["account_number"],
                "balance":        acc["balance"],
            })

        if not accounts:
            print("   ⚠️  No balances found on the home page.")

        # Scrape transactions
        txn_results = process_all_accounts(page, user_id, account_name)

    except PlaywrightTimeoutError as e:
        print(f"   ❌ Timeout: {e}")
    except Exception as e:
        print(f"   ❌ Error: {e}")

    print("🔒 Logging out...")
    for sel in [
        "a:has-text('Logout')",   "a:has-text('logout')",
        "a:has-text('Log Out')",  "a:has-text('Log out')",
        "button:has-text('Logout')", "button:has-text('Log Out')",
        "[aria-label*='logout']", "[aria-label*='Logout']",
        "[class*='logout']", "a[href*='logout']"
    ]:
        try:
            el = page.locator(sel).first
            if el.is_visible(timeout=1500):
                el.click()
                page.wait_for_load_state("domcontentloaded", timeout=8000)
                print("   ✅ Logged out.")
                break
        except Exception:
            continue

    if not is_last:
        try:
            page.goto("https://online.rbl.bank.in/CorpBank/", timeout=10000, wait_until="domcontentloaded")
        except:
            pass

    return bal_results, txn_results

def main():
    password = get_password()
    all_balances = []
    all_transactions = []

    with sync_playwright() as p:
        print("\n🌐 Opening Google Chrome...")
        try:
            browser = p.chromium.launch(executable_path=CHROME_PATH, headless=HEADLESS)
        except Exception:
            print("   ⚠️  Chrome not found, falling back to Chromium.")
            browser = p.chromium.launch(headless=HEADLESS)

        page = browser.new_page()

        for i, account in enumerate(ACCOUNTS):
            is_last = (i == len(ACCOUNTS) - 1)
            bals, txns = login_and_fetch(page, account["user_id"], account["name"], password, is_last=is_last)
            if bals:
                save_balances(bals)
                all_balances.extend(bals)
            if txns:
                save_transactions(txns)
                all_transactions.extend(txns)

        browser.close()

    print(f"\n{'='*55}")
    print("  Final Summary")
    print(f"{'='*55}")
    
    if all_balances:
        print(f"\n  Balances ({len(all_balances)} accounts):")
        for r in all_balances:
            print(f"    {r['account_name']:<6} {r['account_type']:<14} {r['account_number']:<15} {r['balance']:>15}")
        print(f"\n    📂 {BALANCE_FILE}")
    else:
        print("\n  ⚠️  No balance data was captured.")

    if all_transactions:
        print(f"\n  Total transactions extracted: {len(all_transactions)}")
        print(f"    📂 {OUTPUT_FILE}")
    else:
        print("\n  ⚠️  No transaction data was captured.")
    
    print("\n✅ Done!")


if __name__ == "__main__":
    main()
