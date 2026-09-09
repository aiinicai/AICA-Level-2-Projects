"""
CAPITAL GAINS TAX COMPARISON CALCULATOR (12.5% vs 20%)
Professional Report Generator (PDF, Excel, Word)
Author: Senior Python Developer & Tax-Audit Software Architect
"""

import os
from datetime import datetime
from typing import Optional, List, Dict, Any

# PDF ReportLab Imports
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
    HRFlowable,
)
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Excel openpyxl Imports
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Word python-docx Imports
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from tax_calculator import TaxComparisonResult
from utils import format_inr

# Register Arial font if available for proper Unicode (₹) support in ReportLab
PDF_FONT_NAME = "Helvetica"
PDF_FONT_BOLD = "Helvetica-Bold"

try:
    if os.path.exists(r"C:\Windows\Fonts\arial.ttf"):
        pdfmetrics.registerFont(TTFont("Arial", r"C:\Windows\Fonts\arial.ttf"))
        if os.path.exists(r"C:\Windows\Fonts\arialbd.ttf"):
            pdfmetrics.registerFont(TTFont("Arial-Bold", r"C:\Windows\Fonts\arialbd.ttf"))
            PDF_FONT_NAME = "Arial"
            PDF_FONT_BOLD = "Arial-Bold"
        else:
            PDF_FONT_NAME = "Arial"
            PDF_FONT_BOLD = "Arial"
except Exception:
    PDF_FONT_NAME = "Helvetica"
    PDF_FONT_BOLD = "Helvetica-Bold"


class NumberedCanvas(canvas.Canvas):
    """Adds Page X of Y and CA office confidentiality header/footer to each page."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_header_footer(num_pages)
            super().showPage()
        super().save()

    def draw_header_footer(self, page_count: int):
        self.saveState()
        self.setFont(PDF_FONT_NAME, 8)
        self.setFillColor(colors.HexColor("#64748B"))

        # Top Header
        self.drawString(40, 815, "TAX AUDIT ADVISORY MEMORANDUM | CAPITAL GAINS TAX (SEC. 112)")
        self.drawRightString(555, 815, "STRICTLY PRIVATE & CONFIDENTIAL")
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(40, 810, 555, 810)

        # Bottom Footer
        self.line(40, 45, 555, 45)
        self.drawString(40, 32, "Prepared via Capital Gains Tax Comparison Calculator (Finance Act (No. 2) 2024 Engine)")
        self.drawRightString(555, 32, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()


# ==============================================================================
# 1. PDF REPORT GENERATOR
# ==============================================================================
def generate_pdf_report(comparison: TaxComparisonResult, output_path: str) -> str:
    """Generates an audit-grade CA report in PDF format using ReportLab."""
    cg = comparison.cg_result
    m12 = comparison.method_12_5
    m20 = comparison.method_20

    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=50,
        bottomMargin=55,
    )

    styles = getSampleStyleSheet()

    # Custom typography styles
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName=PDF_FONT_BOLD,
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1E3A8A"),
        alignment=0,
    )
    subtitle_style = ParagraphStyle(
        "DocSubTitle",
        parent=styles["Normal"],
        fontName=PDF_FONT_NAME,
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#475569"),
        alignment=0,
    )
    section_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Normal"],
        fontName=PDF_FONT_BOLD,
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#1E3A8A"),
        spaceBefore=10,
        spaceAfter=4,
    )
    cell_text = ParagraphStyle(
        "CellText",
        parent=styles["Normal"],
        fontName=PDF_FONT_NAME,
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1E293B"),
    )
    cell_bold = ParagraphStyle(
        "CellBold",
        parent=styles["Normal"],
        fontName=PDF_FONT_BOLD,
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#0F172A"),
    )
    cell_right = ParagraphStyle(
        "CellRight",
        parent=styles["Normal"],
        fontName=PDF_FONT_NAME,
        fontSize=8.5,
        leading=11,
        alignment=2,
        textColor=colors.HexColor("#1E293B"),
    )
    cell_right_bold = ParagraphStyle(
        "CellRightBold",
        parent=styles["Normal"],
        fontName=PDF_FONT_BOLD,
        fontSize=8.5,
        leading=11,
        alignment=2,
        textColor=colors.HexColor("#0F172A"),
    )
    rec_box_text = ParagraphStyle(
        "RecBoxText",
        parent=styles["Normal"],
        fontName=PDF_FONT_BOLD,
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#065F46"),
        alignment=1,
    )
    rec_sub_text = ParagraphStyle(
        "RecSubText",
        parent=styles["Normal"],
        fontName=PDF_FONT_NAME,
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor("#047857"),
        alignment=1,
    )
    disclaimer_style = ParagraphStyle(
        "Disclaimer",
        parent=styles["Normal"],
        fontName=PDF_FONT_NAME,
        fontSize=7.5,
        leading=10.5,
        textColor=colors.HexColor("#64748B"),
    )

    story = []

    # Title & Metadata
    story.append(Paragraph("CAPITAL GAINS TAX COMPARISON REPORT", title_style))
    story.append(Paragraph(
        f"Comparative Analysis: 12.5% Without Indexation vs. 20% With Indexation | Finance (No. 2) Act, 2024 Regime<br/>"
        f"Generated on: {datetime.now().strftime('%d-%b-%Y at %I:%M %p')} | Assessment Year: {cg.assessment_year}",
        subtitle_style,
    ))
    story.append(Spacer(1, 10))

    # Highlight Card: Recommended Option & Tax Saving
    if comparison.is_20_applicable and comparison.tax_saving > 0:
        rec_data = [
            [Paragraph(f"⭐ RECOMMENDED METHOD: {comparison.recommended_method}", rec_box_text)],
            [Paragraph(
                f"Estimated Tax Saving: <b>{format_inr(comparison.tax_saving, show_paise=True)}</b><br/>"
                f"Tax under 12.5% Method: {format_inr(m12.total_tax_liability)} | Tax under 20% Indexed Method: {format_inr(m20.total_tax_liability if m20 else 0)}",
                rec_sub_text,
            )],
        ]
        rec_table = Table(rec_data, colWidths=[515])
        rec_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#ECFDF5")),
            ("BOX", (0, 0), (-1, -1), 1.5, colors.HexColor("#059669")),
            ("PADDING", (0, 0), (-1, -1), 8),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ]))
        story.append(rec_table)
    elif not comparison.is_20_applicable:
        rec_data = [
            [Paragraph(f"RECOMMENDED METHOD: {comparison.recommended_method}", ParagraphStyle("WarnHead", parent=rec_box_text, textColor=colors.HexColor("#B45309")))],
            [Paragraph(f"Note: 20% Indexed Method Not Applicable. {comparison.ineligibility_reason}", ParagraphStyle("WarnSub", parent=rec_sub_text, textColor=colors.HexColor("#92400E")))],
        ]
        rec_table = Table(rec_data, colWidths=[515])
        rec_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFFBEB")),
            ("BOX", (0, 0), (-1, -1), 1.5, colors.HexColor("#D97706")),
            ("PADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(rec_table)
    else:
        rec_data = [
            [Paragraph(f"RESULT: {comparison.recommended_method}", rec_box_text)],
            [Paragraph(f"Both methods result in the exact same tax liability of {format_inr(comparison.lower_tax_liability)}.", rec_sub_text)],
        ]
        rec_table = Table(rec_data, colWidths=[515])
        rec_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EFF6FF")),
            ("BOX", (0, 0), (-1, -1), 1.5, colors.HexColor("#2563EB")),
            ("PADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(rec_table)

    story.append(Spacer(1, 10))

    # SECTION 1: ASSESSEE & TRANSACTION PROFILE
    story.append(Paragraph("1. Assessee & Transaction Particulars", section_style))
    p_data = [
        [
            Paragraph("Name of Assessee:", cell_bold), Paragraph(cg.assessee_name, cell_text),
            Paragraph("Permanent Account No. (PAN):", cell_bold), Paragraph(cg.pan if cg.pan else "Not Disclosed", cell_text),
        ],
        [
            Paragraph("Status & Residential:", cell_bold), Paragraph(f"{cg.assessee_type} | {cg.residential_status}", cell_text),
            Paragraph("Assessment Year / FY:", cell_bold), Paragraph(f"AY {cg.assessment_year} (FY {cg.financial_year_transfer})", cell_text),
        ],
        [
            Paragraph("Asset Category:", cell_bold), Paragraph(cg.asset_type, cell_text),
            Paragraph("Holding Period:", cell_bold), Paragraph(f"{cg.holding_period_months} months ({cg.holding_period_days} days) - {cg.asset_classification}", cell_text),
        ],
        [
            Paragraph("Date of Acquisition:", cell_bold), Paragraph(f"{cg.date_of_acquisition} (FY {cg.acq_fy}, CII: {cg.acq_cii})", cell_text),
            Paragraph("Date of Sale / Transfer:", cell_bold), Paragraph(f"{cg.date_of_sale} (FY {cg.financial_year_transfer}, CII: {cg.sale_cii})", cell_text),
        ],
    ]
    p_table = Table(p_data, colWidths=[125, 132, 138, 120])
    p_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(p_table)

    story.append(Spacer(1, 10))

    # SECTION 2: SIDE-BY-SIDE STATUTORY COMPARISON
    story.append(Paragraph("2. Side-by-Side Capital Gains & Tax Computation", section_style))

    comp_rows = [
        [
            Paragraph("Particulars", cell_bold),
            Paragraph("12.5% Without Indexation", cell_right_bold),
            Paragraph("20% With Indexation", cell_right_bold),
        ],
        [
            Paragraph("A. Gross Sale Consideration", cell_text),
            Paragraph(format_inr(cg.gross_sale_price), cell_right),
            Paragraph(format_inr(cg.gross_sale_price), cell_right),
        ],
        [
            Paragraph("   Less: Transfer / Selling Expenses", cell_text),
            Paragraph(format_inr(cg.transfer_expenses), cell_right),
            Paragraph(format_inr(cg.transfer_expenses), cell_right),
        ],
        [
            Paragraph("B. Net Sale Consideration", cell_bold),
            Paragraph(format_inr(cg.net_sale_price), cell_right_bold),
            Paragraph(format_inr(cg.net_sale_price), cell_right_bold),
        ],
        [
            Paragraph(f"C. Cost of Acquisition ({'FMV Adopted' if cg.adopted_cost_acq > cg.actual_cost_acq else 'Actual'})", cell_text),
            Paragraph(format_inr(cg.adopted_cost_acq), cell_right),
            Paragraph(format_inr(cg.indexed_cost_acq) if comparison.is_20_applicable else "—", cell_right),
        ],
        [
            Paragraph(f"D. Cost of Improvement ({len(cg.improvements)} entry/entries)", cell_text),
            Paragraph(format_inr(cg.total_actual_improvement), cell_right),
            Paragraph(format_inr(cg.total_indexed_improvement) if comparison.is_20_applicable else "—", cell_right),
        ],
        [
            Paragraph("E. Taxable Long-Term Capital Gain", cell_bold),
            Paragraph(format_inr(max(0.0, cg.ltcg_12_5)), cell_right_bold),
            Paragraph(format_inr(max(0.0, cg.ltcg_20 if cg.ltcg_20 is not None else 0)) if comparison.is_20_applicable else "N/A", cell_right_bold),
        ],
        [
            Paragraph("F. Statutory Tax Rate", cell_text),
            Paragraph(f"{m12.tax_rate_percent}%", cell_right),
            Paragraph(f"{m20.tax_rate_percent}%" if m20 else "N/A", cell_right),
        ],
        [
            Paragraph("G. Basic Capital Gains Tax", cell_text),
            Paragraph(format_inr(m12.basic_tax), cell_right),
            Paragraph(format_inr(m20.basic_tax) if m20 else "N/A", cell_right),
        ],
        [
            Paragraph(f"H. Surcharge (Rate: {m12.surcharge_rate_percent}%)", cell_text),
            Paragraph(format_inr(m12.surcharge_amount), cell_right),
            Paragraph(format_inr(m20.surcharge_amount) if m20 else "N/A", cell_right),
        ],
        [
            Paragraph(f"I. Health & Education Cess ({m12.cess_rate_percent}%)", cell_text),
            Paragraph(format_inr(m12.cess_amount), cell_right),
            Paragraph(format_inr(m20.cess_amount) if m20 else "N/A", cell_right),
        ],
        [
            Paragraph("TOTAL TAX LIABILITY (Rounded u/s 288B)", cell_bold),
            Paragraph(format_inr(m12.total_tax_liability), cell_right_bold),
            Paragraph(format_inr(m20.total_tax_liability) if m20 else "N/A", cell_right_bold),
        ],
    ]

    comp_table = Table(comp_rows, colWidths=[245, 135, 135])
    comp_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
        ("TOPPADDING", (0, 0), (-1, 0), 5),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("BACKGROUND", (0, 3), (-1, 3), colors.HexColor("#F1F5F9")),
        ("BACKGROUND", (0, 6), (-1, 6), colors.HexColor("#F8FAFC")),
        ("BACKGROUND", (0, 11), (-1, 11), colors.HexColor("#FEF3C7")),  # Highlight total row
        ("TOPPADDING", (0, 1), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 3.5),
    ]))
    story.append(comp_table)

    story.append(Spacer(1, 10))

    # SECTION 3: COST OF IMPROVEMENT BREAKDOWN (IF ANY)
    if cg.improvements:
        story.append(Paragraph("3. Cost of Improvement Detailed Indexation Schedule", section_style))
        imp_rows = [
            [
                Paragraph("Particulars", cell_bold),
                Paragraph("Date", cell_bold),
                Paragraph("FY", cell_bold),
                Paragraph("Actual Cost", cell_right_bold),
                Paragraph("CII", cell_bold),
                Paragraph("Indexed Cost", cell_right_bold),
                Paragraph("Status / Remarks", cell_bold),
            ]
        ]
        for imp in cg.improvements:
            imp_rows.append([
                Paragraph(imp.particulars, cell_text),
                Paragraph(imp.date_of_improvement, cell_text),
                Paragraph(imp.fy_of_improvement, cell_text),
                Paragraph(format_inr(imp.amount), cell_right),
                Paragraph(str(imp.cii_improvement) if imp.cii_improvement > 0 else "—", cell_text),
                Paragraph(format_inr(imp.indexed_amount) if imp.is_eligible else "—", cell_right),
                Paragraph("Eligible" if imp.is_eligible else imp.note, cell_text),
            ])

        imp_table = Table(imp_rows, colWidths=[120, 65, 55, 75, 40, 80, 80])
        imp_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(imp_table)
        story.append(Spacer(1, 8))

    # SECTION 4: STATUTORY OPINION & AUDIT NOTES
    story.append(Paragraph("4. Statutory Grounds & Recommendation Rationale", section_style))
    story.append(Paragraph(comparison.recommendation_summary, cell_text))
    story.append(Spacer(1, 4))
    story.append(Paragraph(comparison.detailed_statutory_note, cell_text))

    if cg.pre_2001_notes:
        story.append(Spacer(1, 4))
        story.append(Paragraph(f"<b>Legacy Asset Note (Section 55(2)(b)):</b> {cg.pre_2001_notes}", cell_text))

    story.append(Spacer(1, 15))

    # SECTION 5: DISCLAIMER & AUDIT SIGN-OFF
    disclaimer_text = (
        "<b>LEGAL DISCLAIMER:</b> This report is a calculation aid prepared based on the inputs furnished and "
        "the provisions of Section 112, Section 48, and Section 55 of the Income-tax Act, 1961 as amended by "
        "the Finance (No. 2) Act, 2024. This document should be reviewed with reference to the applicable statutory provisions, "
        "CBDT notifications, circulars, and judicial precedents for the relevant Assessment Year before filing income tax returns."
    )
    story.append(Paragraph(disclaimer_text, disclaimer_style))

    # Build document
    doc.build(story, canvasmaker=NumberedCanvas)
    return output_path


# ==============================================================================
# 2. EXCEL REPORT GENERATOR
# ==============================================================================
def generate_excel_report(comparison: TaxComparisonResult, output_path: str) -> str:
    """Generates an audit-ready multi-tab formatted Excel workbook using openpyxl."""
    cg = comparison.cg_result
    m12 = comparison.method_12_5
    m20 = comparison.method_20

    wb = openpyxl.Workbook()

    # Style Definitions
    font_title = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    font_sub = Font(name="Calibri", size=10, italic=True, color="64748B")
    font_sec = Font(name="Calibri", size=12, bold=True, color="1E3A8A")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="000000")
    font_regular = Font(name="Calibri", size=11, color="1E293B")
    font_rec = Font(name="Calibri", size=12, bold=True, color="065F46")
    font_saving = Font(name="Calibri", size=14, bold=True, color="059669")

    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_sub_header = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    fill_highlight = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_rec = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    thick_bottom = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="double", color="1E3A8A"),
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # SHEET 1: TAX COMPARISON
    ws1 = wb.active
    ws1.title = "Tax Comparison"
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:D1")
    ws1["A1"] = "CAPITAL GAINS TAX COMPARISON REPORT"
    ws1["A1"].font = font_title
    ws1["A1"].alignment = align_left

    ws1.merge_cells("A2:D2")
    ws1["A2"] = f"Comparative Analysis under Finance (No. 2) Act, 2024 | AY {cg.assessment_year} | Generated on {datetime.now().strftime('%d-%m-%Y')}"
    ws1["A2"].font = font_sub

    # Recommended Option Card
    ws1.merge_cells("A4:D4")
    ws1["A4"] = f"RECOMMENDED OPTION: {comparison.recommended_method}"
    ws1["A4"].font = font_rec
    ws1["A4"].fill = fill_rec
    ws1["A4"].alignment = align_center

    ws1.merge_cells("A5:D5")
    ws1["A5"] = f"ESTIMATED TAX SAVING: {format_inr(comparison.tax_saving, show_paise=True)}"
    ws1["A5"].font = font_saving
    ws1["A5"].fill = fill_rec
    ws1["A5"].alignment = align_center

    # Assessee Info Block
    ws1["A7"] = "ASSESSEE & ASSET PARTICULARS"
    ws1["A7"].font = font_sec

    info_data = [
        ("Name of Assessee:", cg.assessee_name, "Permanent Account Number:", cg.pan or "N/A"),
        ("Assessee Category:", cg.assessee_type, "Residential Status:", cg.residential_status),
        ("Capital Asset Type:", cg.asset_type, "Classification / Period:", f"{cg.asset_classification} ({cg.holding_period_months}M)"),
        ("Date of Acquisition:", cg.date_of_acquisition, "Acquisition FY / CII:", f"FY {cg.acq_fy} (CII: {cg.acq_cii})"),
        ("Date of Transfer / Sale:", cg.date_of_sale, "Transfer FY / CII:", f"FY {cg.financial_year_transfer} (CII: {cg.sale_cii})"),
    ]

    r = 8
    for row in info_data:
        ws1[f"A{r}"] = row[0]
        ws1[f"A{r}"].font = font_bold
        ws1[f"B{r}"] = row[1]
        ws1[f"B{r}"].font = font_regular
        ws1[f"C{r}"] = row[2]
        ws1[f"C{r}"].font = font_bold
        ws1[f"D{r}"] = row[3]
        ws1[f"D{r}"].font = font_regular
        for col_l in ["A", "B", "C", "D"]:
            ws1[f"{col_l}{r}"].border = thin_border
        r += 1

    r += 1
    # Side-by-Side Table
    ws1[f"A{r}"] = "SIDE-BY-SIDE CAPITAL GAINS TAX COMPUTATION"
    ws1[f"A{r}"].font = font_sec
    r += 1

    headers = ["Particulars", "12.5% Without Indexation", "20% With Indexation", "Variance / Notes"]
    for idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=r, column=idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_right if idx in (2, 3) else align_left
        cell.border = thin_border

    r += 1
    calc_rows = [
        ("Gross Sale Consideration", cg.gross_sale_price, cg.gross_sale_price, ""),
        ("Less: Transfer Expenses", cg.transfer_expenses, cg.transfer_expenses, ""),
        ("Net Sale Consideration", cg.net_sale_price, cg.net_sale_price, ""),
        ("Cost of Acquisition (Actual / Adopted)", cg.adopted_cost_acq, "—", cg.pre_2001_notes or ""),
        ("Indexed Cost of Acquisition", "—", cg.indexed_cost_acq if comparison.is_20_applicable else "N/A", f"Indexed @ {cg.sale_cii}/{cg.acq_cii}"),
        ("Actual Cost of Improvement", cg.total_actual_improvement, "—", f"{len(cg.improvements)} improvements"),
        ("Indexed Cost of Improvement", "—", cg.total_indexed_improvement if comparison.is_20_applicable else "N/A", ""),
        ("Taxable Long-Term Capital Gain", max(0.0, cg.ltcg_12_5), max(0.0, cg.ltcg_20 if cg.ltcg_20 is not None else 0) if comparison.is_20_applicable else "N/A", ""),
        ("Statutory Tax Rate", f"{m12.tax_rate_percent}%", f"{m20.tax_rate_percent}%" if m20 else "N/A", ""),
        ("Basic Capital Gains Tax", m12.basic_tax, m20.basic_tax if m20 else "N/A", ""),
        (f"Surcharge ({m12.surcharge_rate_percent}%)", m12.surcharge_amount, m20.surcharge_amount if m20 else "N/A", "Capped at 15%"),
        (f"Health & Education Cess ({m12.cess_rate_percent}%)", m12.cess_amount, m20.cess_amount if m20 else "N/A", ""),
        ("TOTAL TAX LIABILITY (SEC. 288B)", m12.total_tax_liability, m20.total_tax_liability if m20 else "N/A", "Lower tax selected"),
    ]

    for item in calc_rows:
        ws1.cell(row=r, column=1, value=item[0]).font = font_bold if "TOTAL" in item[0] or "Net" in item[0] or "Taxable" in item[0] else font_regular
        c2 = ws1.cell(row=r, column=2, value=item[1])
        c3 = ws1.cell(row=r, column=3, value=item[2])
        c4 = ws1.cell(row=r, column=4, value=item[3])

        for col_idx, cell in enumerate([ws1.cell(row=r, column=1), c2, c3, c4], 1):
            cell.border = thick_bottom if "TOTAL" in item[0] else thin_border
            if "TOTAL" in item[0]:
                cell.fill = fill_highlight
                cell.font = font_bold

            if col_idx in (2, 3) and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##,##0.00"
                cell.alignment = align_right
            elif col_idx in (2, 3):
                cell.alignment = align_right

        r += 1

    # Add Rationale & Disclaimer
    r += 1
    ws1[f"A{r}"] = "RATIONALE & STATUTORY NOTES:"
    ws1[f"A{r}"].font = font_sec
    r += 1
    ws1.merge_cells(f"A{r}:D{r+1}")
    ws1[f"A{r}"] = f"{comparison.recommendation_summary}\n{comparison.detailed_statutory_note}"
    ws1[f"A{r}"].font = font_regular
    ws1[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="top")

    r += 3
    ws1.merge_cells(f"A{r}:D{r}")
    ws1[f"A{r}"] = "Note: This report is a calculation aid and should be reviewed with reference to the applicable provisions of the Income-tax Act, Rules, notifications and circulars for the relevant Assessment Year."
    ws1[f"A{r}"].font = font_sub

    # SHEET 2: IMPROVEMENTS & EXPENSES
    ws2 = wb.create_sheet(title="Improvements & Expenses")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = "COST OF IMPROVEMENT SCHEDULE"
    ws2["A1"].font = font_sec
    imp_headers = ["Particulars", "Date of Improvement", "FY", "Actual Amount (INR)", "CII of Year", "Indexed Amount (INR)", "Status / Notes"]
    for idx, h in enumerate(imp_headers, 1):
        cell = ws2.cell(row=2, column=idx, value=h)
        cell.font = font_header
        cell.fill = fill_sub_header
        cell.alignment = align_center
        cell.border = thin_border

    r_imp = 3
    if cg.improvements:
        for imp in cg.improvements:
            ws2.cell(row=r_imp, column=1, value=imp.particulars).font = font_regular
            ws2.cell(row=r_imp, column=2, value=imp.date_of_improvement).font = font_regular
            ws2.cell(row=r_imp, column=3, value=imp.fy_of_improvement).font = font_regular

            c_amt = ws2.cell(row=r_imp, column=4, value=imp.amount)
            c_amt.font = font_regular
            c_amt.number_format = "#,##,##0.00"

            ws2.cell(row=r_imp, column=5, value=imp.cii_improvement).font = font_regular

            c_idx = ws2.cell(row=r_imp, column=6, value=imp.indexed_amount if imp.is_eligible else 0.0)
            c_idx.font = font_regular
            c_idx.number_format = "#,##,##0.00"

            ws2.cell(row=r_imp, column=7, value="Eligible" if imp.is_eligible else imp.note).font = font_regular

            for c in range(1, 8):
                ws2.cell(row=r_imp, column=c).border = thin_border
            r_imp += 1

        # Totals
        ws2.cell(row=r_imp, column=1, value="TOTAL IMPROVEMENTS").font = font_bold
        c_tot1 = ws2.cell(row=r_imp, column=4, value=cg.total_actual_improvement)
        c_tot1.font = font_bold
        c_tot1.number_format = "#,##,##0.00"
        c_tot2 = ws2.cell(row=r_imp, column=6, value=cg.total_indexed_improvement)
        c_tot2.font = font_bold
        c_tot2.number_format = "#,##,##0.00"
        for c in range(1, 8):
            ws2.cell(row=r_imp, column=c).border = thick_bottom
    else:
        ws2.cell(row=r_imp, column=1, value="No Cost of Improvement incurred.").font = font_regular

    # Auto-fit columns
    for sheet in [ws1, ws2]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(output_path)
    return output_path


# ==============================================================================
# 3. WORD (DOCX) REPORT GENERATOR
# ==============================================================================
def generate_word_report(comparison: TaxComparisonResult, output_path: str) -> str:
    """Generates a formal Tax Advisory Memorandum in Microsoft Word (.docx) format."""
    cg = comparison.cg_result
    m12 = comparison.method_12_5
    m20 = comparison.method_20

    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Title
    p_title = doc.add_paragraph()
    run_title = p_title.add_run("CAPITAL GAINS TAX COMPARISON REPORT")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138)

    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run(
        f"Statutory Comparative Analysis: 12.5% Without Indexation vs. 20% With Indexation\n"
        f"Assessment Year: {cg.assessment_year} | Financial Year: {cg.financial_year_transfer} | Date: {datetime.now().strftime('%d-%b-%Y')}"
    )
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Recommendation Callout Box
    rec_p = doc.add_paragraph()
    rec_run = rec_p.add_run(f"⭐ RECOMMENDED METHOD: {comparison.recommended_method}\n")
    rec_run.bold = True
    rec_run.font.size = Pt(12)
    rec_run.font.color.rgb = RGBColor(6, 95, 70)

    if comparison.tax_saving > 0:
        sav_run = rec_p.add_run(f"Estimated Tax Saving: {format_inr(comparison.tax_saving, show_paise=True)}\n")
        sav_run.bold = True
        sav_run.font.size = Pt(13)
        sav_run.font.color.rgb = RGBColor(5, 150, 105)

    note_run = rec_p.add_run(comparison.recommendation_summary)
    note_run.font.size = Pt(10)
    note_run.font.color.rgb = RGBColor(4, 120, 87)

    # Section 1: Assessee & Asset Details
    doc.add_heading("1. Assessee & Transaction Profile", level=2)
    table_p = doc.add_table(rows=4, cols=4)
    table_p.alignment = WD_TABLE_ALIGNMENT.CENTER

    info = [
        ("Name of Assessee", cg.assessee_name, "PAN", cg.pan or "Not Disclosed"),
        ("Category & Status", f"{cg.assessee_type} | {cg.residential_status}", "AY / FY", f"AY {cg.assessment_year} (FY {cg.financial_year_transfer})"),
        ("Capital Asset", cg.asset_type, "Holding Period", f"{cg.holding_period_months} Months ({cg.asset_classification})"),
        ("Acquisition Date & CII", f"{cg.date_of_acquisition} (CII: {cg.acq_cii})", "Sale Date & CII", f"{cg.date_of_sale} (CII: {cg.sale_cii})"),
    ]

    for r_idx, row_vals in enumerate(info):
        for c_idx, val in enumerate(row_vals):
            cell = table_p.cell(r_idx, c_idx)
            cell.text = str(val)
            p = cell.paragraphs[0]
            p.runs[0].font.name = "Calibri"
            p.runs[0].font.size = Pt(9.5)
            if c_idx in (0, 2):
                p.runs[0].font.bold = True

    doc.add_paragraph()

    # Section 2: Side-by-Side Calculation Table
    doc.add_heading("2. Side-by-Side Capital Gains Tax Computation", level=2)
    table_comp = doc.add_table(rows=1, cols=3)
    table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table_comp.rows[0].cells
    hdr_cells[0].text = "Particulars"
    hdr_cells[1].text = "12.5% Without Indexation"
    hdr_cells[2].text = "20% With Indexation"
    for cell in hdr_cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(30, 58, 138)

    rows_data = [
        ("Gross Sale Consideration", format_inr(cg.gross_sale_price), format_inr(cg.gross_sale_price)),
        ("Less: Transfer Expenses", format_inr(cg.transfer_expenses), format_inr(cg.transfer_expenses)),
        ("Net Sale Consideration", format_inr(cg.net_sale_price), format_inr(cg.net_sale_price)),
        ("Cost of Acquisition (Adopted)", format_inr(cg.adopted_cost_acq), format_inr(cg.indexed_cost_acq) if comparison.is_20_applicable else "—"),
        ("Cost of Improvement", format_inr(cg.total_actual_improvement), format_inr(cg.total_indexed_improvement) if comparison.is_20_applicable else "—"),
        ("Taxable Long-Term Capital Gain", format_inr(max(0.0, cg.ltcg_12_5)), format_inr(max(0.0, cg.ltcg_20 if cg.ltcg_20 is not None else 0)) if comparison.is_20_applicable else "N/A"),
        ("Statutory Tax Rate", f"{m12.tax_rate_percent}%", f"{m20.tax_rate_percent}%" if m20 else "N/A"),
        ("Basic Capital Gains Tax", format_inr(m12.basic_tax), format_inr(m20.basic_tax) if m20 else "N/A"),
        (f"Surcharge ({m12.surcharge_rate_percent}%)", format_inr(m12.surcharge_amount), format_inr(m20.surcharge_amount) if m20 else "N/A"),
        (f"Health & Education Cess ({m12.cess_rate_percent}%)", format_inr(m12.cess_amount), format_inr(m20.cess_amount) if m20 else "N/A"),
        ("TOTAL TAX LIABILITY (u/s 288B)", format_inr(m12.total_tax_liability), format_inr(m20.total_tax_liability) if m20 else "N/A"),
    ]

    for item in rows_data:
        row_cells = table_comp.add_row().cells
        row_cells[0].text = item[0]
        row_cells[1].text = item[1]
        row_cells[2].text = item[2]
        if "TOTAL" in item[0] or "Net" in item[0] or "Taxable" in item[0]:
            for c in row_cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.bold = True

    doc.add_paragraph()

    # Section 3: Statutory Analysis
    doc.add_heading("3. Statutory Analysis & Grounds", level=2)
    p_law = doc.add_paragraph()
    p_law.add_run(comparison.detailed_statutory_note)

    if cg.pre_2001_notes:
        p_leg = doc.add_paragraph()
        run_leg_lbl = p_leg.add_run("Legacy Asset Grandfathering (Section 55(2)(b)): ")
        run_leg_lbl.bold = True
        p_leg.add_run(cg.pre_2001_notes)

    # Disclaimer
    doc.add_paragraph()
    p_disc = doc.add_paragraph()
    r_disc = p_disc.add_run(
        "“This report is a calculation aid and should be reviewed with reference to the applicable provisions "
        "of the Income-tax Act, Rules, notifications and circulars for the relevant Assessment Year.”"
    )
    r_disc.font.italic = True
    r_disc.font.size = Pt(9)
    r_disc.font.color.rgb = RGBColor(100, 116, 139)

    doc.save(output_path)
    return output_path
