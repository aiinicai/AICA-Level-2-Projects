"""Phase 4: iCalendar (RFC 5545, hand-written), Excel exports, reminders, KPIs,
client-letter content and rule-pack version snapshots/diffs."""
from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from typing import Any, Iterable

from sqlalchemy import select

from . import audit
from .models import (CLOSED, Entity, EntityPerson, Filing, Notification, Obligation, RuleVersion, User, db, utcnow)

# ================================================================ iCalendar
def ics_escape(text: str) -> str:
    """RFC 5545 §3.3.11 TEXT escaping."""
    return (text.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,")
            .replace("\r\n", "\\n").replace("\n", "\\n"))


def ics_fold(line: str) -> str:
    """RFC 5545 §3.1: lines longer than 75 octets are folded with CRLF + one space,
    never splitting a multi-byte UTF-8 character."""
    out, current, size = [], "", 0
    for ch in line:
        n = len(ch.encode("utf-8"))
        limit = 75 if not out else 74          # continuation lines start with a space
        if size + n > limit:
            out.append(current)
            current, size = "", 0
        current += ch
        size += n
    out.append(current)
    return "\r\n ".join(out)


def build_ics(obligations: Iterable[Obligation], cal_name: str, firm: str, today: date) -> str:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//MCA Compliance Mapper//EN", "CALSCALE:GREGORIAN",
             "METHOD:PUBLISH", f"X-WR-CALNAME:{ics_escape(cal_name)}", "X-WR-TIMEZONE:Asia/Kolkata"]
    for o in obligations:
        who = o.entity.name if o.entity else (f"{o.person.name} (DIN {o.person.din})" if o.person else "")
        status = "Filed/closed" if o.status in CLOSED else o.status.replace("_", " ").title()
        desc = (f"{o.title}\nPeriod: {o.period_key}\nDue: {o.due_date:%d-%m-%Y}\nStatus: {status}\nLaw: {o.law}\n"
                f"Rule: {o.rule_code} (pack {o.rule_version})" + ("\nProvisional date" if o.provisional else ""))
        lines += ["BEGIN:VEVENT",
                  f"UID:{o.key.replace(':', '-')}@mca-compliance-mapper",
                  f"DTSTAMP:{stamp}",
                  f"DTSTART;VALUE=DATE:{o.due_date:%Y%m%d}",
                  f"DTEND;VALUE=DATE:{o.due_date + timedelta(days=1):%Y%m%d}",
                  f"SUMMARY:{ics_escape(f'{o.form} — {who}')}",
                  f"DESCRIPTION:{ics_escape(desc)}",
                  f"CATEGORIES:{ics_escape('MCA compliance')}",
                  "TRANSP:TRANSPARENT",
                  f"STATUS:{'CANCELLED' if o.status in ('NOT_APPLICABLE', 'WAIVED') else 'CONFIRMED'}"]
        if o.status not in CLOSED and o.due_date >= today:
            lines += ["BEGIN:VALARM", "ACTION:DISPLAY", "TRIGGER:-P7D",
                      f"DESCRIPTION:{ics_escape(f'{o.form} due in 7 days')}", "END:VALARM"]
        lines.append("END:VEVENT")
    lines.append("END:VCALENDAR")
    return "\r\n".join(ics_fold(ln) for ln in lines) + "\r\n"


# ================================================================== Excel
NAVY, GOLD, PALE = "0B1D3A", "C8973F", "F8F6F1"


def _sheet(wb, title: str, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None,
           money_cols: Iterable[int] = (), date_cols: Iterable[int] = ()):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(bottom=Side(style="medium", color=GOLD))
    for r in rows:
        ws.append(r)
    for col in money_cols:
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = '[>=10000000]"₹"##\\,##\\,##\\,##0;[>=100000]"₹"##\\,##\\,##0;"₹"##,##0'
    for col in date_cols:
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = "DD-MM-YYYY"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def workbook_bytes(sheets: list[dict], meta: list[tuple[str, str]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    about = wb.active
    about.title = "About"
    about["A1"] = "MCA Compliance Mapper — export"
    about["A1"].font = Font(bold=True, size=14, color=NAVY)
    for i, (k, v) in enumerate(meta, start=3):
        about.cell(row=i, column=1, value=k).font = Font(bold=True)
        about.cell(row=i, column=2, value=v)
    about.column_dimensions["A"].width = 28
    about.column_dimensions["B"].width = 90
    for s in sheets:
        _sheet(wb, **s)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def register_rows(obls: list[Obligation], today: date, fee_for) -> list[list[Any]]:
    from .services import HEALTH_LABELS, health
    filings ={f.obligation_id: f for f in db.session.execute(select(Filing).where(
        Filing.obligation_id.in_([o.id for o in obls]))).scalars()} if obls else {}
    rows = []
    for o in obls:
        f = filings.get(o.id)
        fb = fee_for(o, today) if (o.fee_regime != "NONE" and o.status not in CLOSED and o.due_date < today) else None
        rows.append([o.entity.name if o.entity else (o.person.name if o.person else ""),
                     o.entity.cin if o.entity else (f"DIN {o.person.din}" if o.person else ""),
                     o.form, o.period_key, o.due_date, HEALTH_LABELS[health(o, today)], o.status.replace("_", " ").title(),
                     o.assignee.full_name if o.assignee else "", f.srn if f else "", f.filing_date if f else None,
                     (f.normal_fee_paid + f.additional_fee_paid) if f else None, fb.total if fb else None,
                     "Yes" if o.provisional else "", "Yes" if o.pre_engagement else "", o.law, o.rule_code,
                     o.rule_version])
    return rows


REGISTER_HEADERS = ["Entity / person", "CIN / DIN", "Form", "Period", "Due date", "Health", "Status", "Assignee",
                    "SRN", "Filed on", "Fees paid ₹", "Fee if filed today ₹", "Provisional", "Pre-engagement",
                    "Law", "Rule", "Rule-pack version"]
REGISTER_WIDTHS = [30, 24, 26, 18, 12, 14, 18, 20, 12, 12, 13, 15, 11, 13, 45, 18, 14]


# ============================================================== reminders
REMINDER_OFFSETS = {"T-30": 30, "T-7": 7, "T-1": 1, "OVERDUE": -1}


def reminder_kind(o: Obligation, today: date) -> str | None:
    days = (o.due_date - today).days
    return {30: "T-30", 7: "T-7", 1: "T-1", -1: "OVERDUE"}.get(days)


def run_reminders(today: date, actor=audit.SYSTEM, catch_up: bool = False) -> dict[str, int]:
    """Create in-app notifications at T-30, T-7, T-1 and on the first overdue day, for the
    assignee and their manager (the entity's relationship manager, else every active manager).
    Idempotent: one notification per (user, obligation, kind). With catch_up=True, missed
    windows (e.g. the server was off) are sent for the nearest threshold already passed."""
    managers = [u.id for u in db.session.execute(select(User).where(User.role == "MANAGER",
                                                                    User.is_active_flag.is_(True))).scalars()]
    q = select(Obligation).where(Obligation.superseded_at.is_(None), Obligation.status.not_in(list(CLOSED)),
                                 Obligation.due_date <= today + timedelta(days=30),
                                 Obligation.due_date >= today - timedelta(days=1 if not catch_up else 3650))
    made = 0
    existing = {(n.user_id, n.obligation_id, n.kind) for n in db.session.execute(select(Notification)).scalars()}
    for o in db.session.execute(q).scalars():
        if o.pre_engagement and o.due_date < today - timedelta(days=1) and not catch_up:
            continue
        kind = reminder_kind(o, today)
        if kind is None and catch_up:
            d = (o.due_date - today).days
            kind = "OVERDUE" if d < 0 else ("T-1" if d <= 1 else "T-7" if d <= 7 else "T-30")
        if kind is None:
            continue
        recipients = set()
        if o.assignee_id:
            recipients.add(o.assignee_id)
        ent = o.entity or (o.person.links[0].entity if o.person and o.person.links else None)
        if ent and ent.rm_id:
            recipients.add(ent.rm_id)
        elif managers:
            recipients.update(managers)
        who = o.entity.name if o.entity else (o.person.name if o.person else "")
        msg = (f"{o.form} for {who} ({o.period_key}) " + ("is overdue since " if kind == "OVERDUE" else "is due on ")
               + f"{o.due_date:%d-%m-%Y}.")
        for uid in recipients:
            if (uid, o.id, kind) in existing:
                continue
            db.session.add(Notification(user_id=uid, obligation_id=o.id, kind=kind, message=msg))
            existing.add((uid, o.id, kind))
            made += 1
    if made:
        audit.record(db.session, actor, "REMINDERS_CREATED", "notification", None, after={"count": made, "date": today})
    db.session.commit()
    return {"notifications": made}


def email_digest(user: User, notes: list[Notification]) -> str:
    lines = [f"Dear {user.full_name},", "", "MCA compliance reminders:", ""]
    lines += [f"- {n.message}" for n in notes]
    lines += ["", "Open the MCA Compliance Mapper for details.", "(Automated internal reminder — do not reply.)"]
    return "\n".join(lines)


def send_digests(notes_by_user: dict[int, list[Notification]]) -> int:
    """Optional SMTP digest (env SMTP_HOST etc.). Returns messages sent; 0 if SMTP is not configured."""
    import os
    import smtplib
    from email.message import EmailMessage
    host = os.environ.get("SMTP_HOST")
    if not host:
        return 0
    sent = 0
    with smtplib.SMTP(host, int(os.environ.get("SMTP_PORT", "587")), timeout=20) as s:
        s.starttls()
        if os.environ.get("SMTP_USER"):
            s.login(os.environ["SMTP_USER"], os.environ.get("SMTP_PASSWORD", ""))
        for uid, notes in notes_by_user.items():
            user = db.session.get(User, uid)
            if not user or not user.email or not notes:
                continue
            m = EmailMessage()
            m["Subject"] = f"MCA compliance: {len(notes)} reminder(s)"
            m["From"] = os.environ.get("SMTP_FROM", os.environ.get("SMTP_USER", "noreply@localhost"))
            m["To"] = user.email
            m.set_content(email_digest(user, notes))
            s.send_message(m)
            sent += 1
    return sent


def nightly(as_of: date, catch_up: bool = False) -> dict[str, int]:
    """Recompute every active entity (new FY / half-year / KYC cycle appear by themselves), then
    reminders, then the optional e-mail digest. Call inside an app context."""
    from .services import recompute_all
    stats = recompute_all(as_of)
    stats.update(run_reminders(as_of, catch_up=catch_up))
    fresh = db.session.execute(select(Notification).where(Notification.created_at >= utcnow() - timedelta(hours=1),
                                                          Notification.read_at.is_(None))).scalars().all()
    by_user: dict[int, list[Notification]] = defaultdict(list)
    for n in fresh:
        by_user[n.user_id].append(n)
    try:
        stats["emails"] = send_digests(by_user)
    except OSError as e:                     # SMTP down: in-app notifications still stand
        stats["emails"] = 0
        stats["email_error"] = str(e)[:200]
    return stats


def start_scheduler(app):
    """In-process nightly job at 01:00 IST (brief §3). Returns the scheduler (or None if disabled)."""
    from apscheduler.schedulers.background import BackgroundScheduler
    from apscheduler.triggers.cron import CronTrigger
    ist = timezone(timedelta(hours=5, minutes=30))

    def job():
        with app.app_context():
            from .auth import today
            stats = nightly(today())
            app.logger.info("nightly job: %s", stats)

    sched = BackgroundScheduler(timezone=ist, daemon=True)
    sched.add_job(job, CronTrigger(hour=1, minute=0, timezone=ist), id="nightly", replace_existing=True,
                  misfire_grace_time=6 * 3600, coalesce=True)
    sched.start()
    return sched


# ==================================================================== KPIs
def kpis(obls: list[Obligation], filings: dict[int, Filing], today: date, window_months: int = 24):
    """Entity compliance score = on-time filings ÷ due filings in the rolling window
    (excluding pre-engagement and non-filing items); staff on-time % by assignee."""
    start = today - timedelta(days=round(window_months * 30.44))
    per_entity: dict[int, list[int]] = defaultdict(lambda: [0, 0])
    per_staff: dict[int, list[int]] = defaultdict(lambda: [0, 0])
    for o in obls:
        if o.pre_engagement or o.superseded_at or o.fee_regime == "NONE" or not (start <= o.due_date < today):
            continue
        if o.status in ("NOT_APPLICABLE", "WAIVED"):
            continue
        f = filings.get(o.id)
        on_time = bool(f and f.filing_date <= o.due_date)
        if o.entity_id:
            per_entity[o.entity_id][1] += 1
            per_entity[o.entity_id][0] += on_time
        if o.assignee_id and f is not None:
            per_staff[o.assignee_id][1] += 1
            per_staff[o.assignee_id][0] += on_time
    return per_entity, per_staff, start


# ========================================================= client letter
DOCS_NEEDED: dict[str, list[str]] = {
    "AOC4": ["Signed audited financial statements and auditor's report", "Board's report with annexures",
             "Date of the AGM and the notice"],
    "AOC4_OPC": ["Signed audited financial statements and auditor's report", "Director's report"],
    "AOC4_CFS": ["Consolidated financial statements and auditor's report"],
    "MGT7": ["Shareholding pattern at 31 March and transfers during the year", "Dates of Board and general meetings",
             "Directors' and KMP remuneration details"],
    "MGT7_OPC": ["Shareholding at 31 March", "Dates of Board meetings", "Director's remuneration"],
    "MGT8": ["Registers and minutes for the practising company secretary's certification"],
    "ADT1": ["Auditor's consent and eligibility certificate", "Resolution appointing the auditor"],
    "ADT1_FIRST": ["Auditor's consent and eligibility certificate", "Board resolution appointing the first auditor"],
    "DPT3": ["Details of loans and deposits outstanding at 31 March", "Auditor's certificate (where deposits are held)"],
    "MSME1_H1": ["List of micro/small suppliers with dues outstanding over 45 days, with reasons for delay"],
    "MSME1_H2": ["List of micro/small suppliers with dues outstanding over 45 days, with reasons for delay"],
    "INC20A": ["Bank statement showing subscription money received", "Proof of registered office"],
    "DIR3KYC_TRIENNIAL": ["Director's current mobile number and e-mail (OTP)", "PAN, Aadhaar and address proof"],
    "DIR3KYC_CHANGE": ["Proof of the new address / mobile / e-mail", "Director available for OTP"],
    "DIR12": ["Consent (DIR-2) and DIR-8", "Resolution and appointment/resignation letter"],
    "PAS3": ["List of allottees", "Resolution and valuation report (where applicable)", "Bank statement"],
    "PAS3_PP": ["List of allottees", "PAS-4 offer letter and PAS-5 record", "Special resolution", "Bank statement"],
    "CHG1": ["Instrument creating the charge", "Sanction letter", "Particulars of property charged"],
    "CHG4": ["No-dues / satisfaction letter from the charge holder"],
    "MGT14": ["Certified copy of the resolution with explanatory statement"],
    "PAS6": ["Reconciliation of share capital audit report from the RTA / PCS"],
    "LLP_FORM11": ["Contribution of each partner at 31 March", "Changes in partners during the year"],
    "LLP_FORM8": ["Signed statement of account and solvency", "Audit report (if audit applies)"],
    "CSR2": ["CSR spending details, unspent amounts and impact assessment (if any)"],
    "BEN2": ["BEN-1 declaration received", "Particulars of the significant beneficial owner"],
}


def letter_items(obls: list[Obligation]) -> list[dict]:
    items = []
    for o in obls:
        items.append({"o": o, "docs": DOCS_NEEDED.get(o.rule_code, ["Information to be confirmed with our office"])})
    return items


# =================================================== rule-pack versions
def snapshot(rp) -> dict:
    return {
        "rules": {r.code: r.model_dump(mode="json") for r in rp.rules},
        "thresholds": [t.model_dump(mode="json") for t in rp.thresholds],
        "fees": [t.model_dump(mode="json") for t in rp._fee_rows],
        "schemes": [s.model_dump(mode="json") for s in rp.schemes],
        "hashes": {r.code: r.content_hash for r in rp.rules},
    }


IGNORED = {"verified", "verified_by", "verified_on"}


def diff_snapshots(old: dict, new: dict) -> dict:
    o_rules, n_rules = old.get("rules", {}), new.get("rules", {})
    added = sorted(set(n_rules) - set(o_rules))
    removed = sorted(set(o_rules) - set(n_rules))
    changed = []
    for code in sorted(set(o_rules) & set(n_rules)):
        a, b = o_rules[code], n_rules[code]
        fields = [(k, a.get(k), b.get(k)) for k in sorted(set(a) | set(b)) if k not in IGNORED and a.get(k) != b.get(k)]
        if fields:
            changed.append((code, fields))
    other = [k for k in ("thresholds", "fees", "schemes") if old.get(k) != new.get(k)]
    return {"added": added, "removed": removed, "changed": changed, "other": other}


def latest_version() -> RuleVersion | None:
    return db.session.execute(select(RuleVersion).order_by(RuleVersion.id.desc())).scalars().first()
